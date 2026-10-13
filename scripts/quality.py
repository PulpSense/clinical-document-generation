"""Fail-closed content, XML, renderer, and page-visual quality gates."""

from __future__ import annotations

import base64
import binascii
import copy
import hashlib
import hmac
import importlib
import importlib.util
import json
import math
import os
import platform
import re
import signal
import shutil
import struct
import subprocess
import sys
import tempfile
import time
import zipfile
import zlib
from pathlib import Path, PurePosixPath
from typing import Any, Iterable, Mapping, Sequence

from docx import Document
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph
from pypdf import PdfReader
from lxml import etree as ET

from contracts import APPROVED_PACKAGED_FONT_FALLBACKS, BOILERPLATE_VERSION, BUNDLED_FONT_FILES, ICF_RETAINED_SHELL_SECTIONS, RECOVERY_POLICIES, batch_plan, canonical_study_type, contracted_template_bundle, contextual_clinical_boilerplate, country_name_aliases, document_set, finding_requires_repair, get_path, icf_contract, icf_retained_sections, meaningful, narrative_review_finding, protocol_concept_ownership, protocol_contract, protocol_table_contracts, recovery_finding, section_applies, participant_followup_summary, semantic_evidence_contract, state_name_aliases, sterling_clause_contract, sterling_clause_text
from drafting import accepted_cross_section_duplicate_findings, evidence_coverage_finding, governing_resources, evidence_grounded, hypothesis_claim_issues, lexical_evidence_uncertainty, source_evidence_grounded, source_evidence_diagnostics, section_evidence_value
from prs_xml import screening_interval_requirement, validate_output as validate_prs_output
from rendering import audit_docx, refresh_toc_from_pdf, rendered_section_snapshot, template_paths
from contracts import benefit_context_contract, document_table_reference, source_text_values
from drafting import merged_drafts, sha256_value
from contracts import canonical_review_surface, reviewed_presentation_only, internal_presentation_note, review_concern


def _review_evidence_finding(artifact: str, section_id: str) -> dict[str, Any]:
    """Declare the construction-owned missing-evidence blocker in the audit seam."""
    return recovery_finding({
        "category": "document_structure", "code": "review_evidence_unavailable",
        "artifact": artifact, "field": section_id, "target_ids": [section_id],
        "publication_disposition": "blocking",
        "issue": "Review evidence has no current passage for the affected section; repair section localization before dispatch.",
    }, "deterministic_structure_defect")


class ReviewEvidenceError(ValueError):
    """Code could not bind a caution to the actual candidate section."""

    def __init__(self, artifact: str, section_id: str):
        self.finding = _review_evidence_finding(artifact, section_id)
        super().__init__(self.finding["issue"])


VERIFY_SCHEMA = "hermes-verification/v1"
RESPONSE_SCHEMA = "hermes-verification-response/v1"
INSPECTION_FAILURE_TYPES = {"native_inspection_incomplete", "inspection_unavailable", "permission_denied"}
FINAL_REVIEW_SCHEMA = "final-exact-artifact-review/v1"
FINAL_REVIEW_SCOPE = "complete_branch_document_set"
CONTENT_CHECKS = ("substantive", "source_supported", "no_invention", "no_internal_language", "cross_document_consistent")
CROSS_DOCUMENT_CHECKS = ("protocol_number", "study_title", "study_type", "population", "procedures", "risks_benefits")
VISUAL_CHECKS = (
    "clipping", "overlap", "overflow", "orphan_heading", "bad_table_split",
    "blank_page", "footer_collision", "unreadable_text", "duplicate_section",
    "inconsistent_style", "missing_header_footer", "toc_mismatch",
    "excessive_whitespace", "artificial_pagination",
)
UNCERTAINTY_CODES = {
    "narrative-review-uncertainty",
    "lexical-evidence-uncertainty", "icf-lexical-evidence-uncertainty",
    "retrospective-lexical-evidence-uncertainty", "icf-summary-grounding-review",
    "protocol-concept-repetition", "protocol-exact-repetition", "icf-concept-repetition",
    "visible-duplicate-word", "icf-exact-generated-duplication", "draft-repetition-uncertainty",
}
INTERNAL_DIAGNOSTIC_CODES = {"sterling-clause-word-goal", "nonmaterial-presentation"}
TRANSIENT_REVIEW_STATUSES = {"retryable_error", "transient_error", "unavailable", "temporarily_unavailable"}
_MAC_FONT_NAMES: set[str] | None = None
_WINDOWS_FONT_NAMES: set[str] | None = None
PDFIUM_WORKER_TIMEOUT_SECONDS = 120.0
PDFIUM_MAX_PAGES = 400
PDFIUM_MAX_DIMENSION_PIXELS = 20_000
PDFIUM_MAX_TOTAL_PIXELS = 1_000_000_000
PDFIUM_MAX_OUTPUT_BYTES = 1_000_000_000
PDFIUM_WORKER_MEMORY_BYTES = 2 * 1024 * 1024 * 1024
PDFIUM_WORKER_FILE_BYTES = PDFIUM_MAX_OUTPUT_BYTES + 64 * 1024 * 1024
RELEASE_CERTIFICATION_PUBLIC_KEY = "references/release-certification-public-key.json"
RELEASE_CERTIFICATION_SIGNATURE_ALGORITHM = "rsa-pkcs1-v1_5-sha256"
RELEASE_CERTIFICATION_TRUSTED_KEY_ID = "50aa8bde2e3f31c7c24984078dbe1d236f118e0744c43a84d0e4c77ffcc1107c"
_SHA256_DIGEST_INFO_PREFIX = bytes.fromhex("3031300d060960864801650304020105000420")
CERTIFICATION_EVIDENCE_MAX_FILES = 1024
CERTIFICATION_EVIDENCE_MAX_ITEM_BYTES = 32 * 1024 * 1024
CERTIFICATION_EVIDENCE_MAX_TOTAL_BYTES = 128 * 1024 * 1024
CERTIFICATION_CASE_ORDER = (
    "retrospective",
    "ambispective-sterling",
    "prospective-advarra",
    "prospective-sterling",
    "ambispective-advarra",
)
CERTIFICATION_DIAGNOSTIC_TARGET_SECONDS = 900.0
CERTIFICATION_EXTENDED_DIAGNOSTIC_SECONDS = 1080.0
CERTIFICATION_RUNTIME_CEILING_SECONDS = 2700.0
PROMOTED_ACTIVATION_COMMIT_POINT = "post_activation_smoke_passed_journal_commit"
DETERMINISTIC_BRANCH_ACCEPTANCE_CASES = (
    "prospective-advarra-sparse-complete", "prospective-advarra-rich-complete",
    "ambispective-sterling-sparse-complete", "ambispective-sterling-rich-complete",
    "prospective-sterling-sparse-complete", "prospective-sterling-rich-complete",
    "ambispective-advarra-sparse-complete", "ambispective-advarra-rich-complete",
    "retrospective-sparse-complete", "retrospective-rich-complete",
)
FORMAT_CONFORMANCE_BASELINE_CASES = {
    "retrospective-protocol": "retrospective-sparse-complete",
    "prospective-advarra": "prospective-advarra-sparse-complete",
    "prospective-sterling": "prospective-sterling-rich-complete",
    "ambispective-advarra": "ambispective-advarra-sparse-complete",
    "ambispective-sterling": "ambispective-sterling-rich-complete",
}


def certification_runtime_classification(elapsed_seconds: float) -> str:
    """Classify certification runtime without shortening the correctness ceiling."""
    if elapsed_seconds <= 0.0:
        return "invalid"
    if elapsed_seconds < CERTIFICATION_DIAGNOSTIC_TARGET_SECONDS:
        return "under_15_minutes"
    if elapsed_seconds <= CERTIFICATION_EXTENDED_DIAGNOSTIC_SECONDS:
        return "within_18_minute_diagnostic_window"
    if elapsed_seconds <= CERTIFICATION_RUNTIME_CEILING_SECONDS:
        return "above_18_minutes_within_correctness_ceiling"
    return "correctness_ceiling_exceeded"


def branch_acceptance_inventory(repo_root: Path) -> tuple[dict[str, Any], ...]:
    """Load and bind the ten repository-confined corpus descriptors and sources."""
    root = repo_root.resolve()
    catalog_path = root / "references/conformance-fixtures/branch-acceptance-corpus.json"
    catalog = json.loads(catalog_path.read_text(encoding="utf-8"))
    if not isinstance(catalog, list):
        raise ValueError("Branch Acceptance Corpus catalog must be a JSON list.")
    inventory = []
    for entry in catalog:
        if not isinstance(entry, Mapping):
            raise ValueError("Branch Acceptance Corpus entries must be JSON objects.")
        if entry.get("regression") is True:
            continue
        descriptor_path = (root / str(entry.get("path") or "")).resolve()
        try:
            descriptor_path.relative_to(root)
        except ValueError as exc:
            raise ValueError("Branch Acceptance Corpus descriptor escapes the repository.") from exc
        if not descriptor_path.is_file():
            raise ValueError(f"Branch Acceptance Corpus descriptor is missing: {descriptor_path}")
        descriptor = json.loads(descriptor_path.read_text(encoding="utf-8"))
        required = {"fixture_id", "study_type", "profile", "source_fixture"}
        if (
            not isinstance(descriptor, dict)
            or not required.issubset(descriptor)
            or not str(descriptor["profile"]).endswith("complete")
        ):
            raise ValueError(f"Invalid Branch Acceptance Corpus descriptor: {descriptor_path}")
        if descriptor["study_type"] != "Retrospective" and not descriptor.get("icf_template"):
            raise ValueError(f"ICF family is missing from corpus descriptor: {descriptor_path}")
        source_path = (root / str(descriptor["source_fixture"])).resolve()
        try:
            source_path.relative_to(root)
        except ValueError as exc:
            raise ValueError("Branch Acceptance Corpus source escapes the repository.") from exc
        if not source_path.is_file():
            raise ValueError(f"Branch Acceptance Corpus source is missing: {source_path}")
        inventory.append({
            "descriptor": descriptor,
            "descriptor_path": descriptor_path.relative_to(root).as_posix(),
            "descriptor_sha256": hashlib.sha256(descriptor_path.read_bytes()).hexdigest(),
            "source_path": source_path.relative_to(root).as_posix(),
            "source_sha256": hashlib.sha256(source_path.read_bytes()).hexdigest(),
        })
    case_ids = tuple(str(item["descriptor"]["fixture_id"]) for item in inventory)
    if case_ids != DETERMINISTIC_BRANCH_ACCEPTANCE_CASES:
        raise ValueError("Branch Acceptance Corpus does not match the canonical ten-case inventory.")
    return tuple(inventory)


def branch_acceptance_source(repo_root: Path, corpus_item: Mapping[str, Any]) -> dict[str, Any]:
    """Materialize the effective source shared by every corpus lifecycle runner."""
    case = corpus_item["descriptor"]
    study_type = str(case["study_type"])
    richness = str(case["profile"]).split("-", 1)[0]
    source_path = (repo_root / str(corpus_item["source_path"])).resolve()
    if not source_path.is_relative_to(repo_root.resolve()):
        raise ValueError("Branch Acceptance Corpus source escapes the repository.")
    reference = json.loads(source_path.read_text(encoding="utf-8"))
    if get_path(reference, "meta.study_type") != study_type:
        raise ValueError("Branch Acceptance Corpus descriptor and source study types disagree.")
    reference["meta"]["protocol_number"] = f"{study_type[:3].upper()}-{richness}-26"; reference["study"]["title"] = str(case.get("study_title") or f"{study_type} {richness} acceptance study"); reference["study"]["short_title"] = str(case.get("study_short_title") or f"{study_type} {richness} acceptance")
    if study_type != "Retrospective":
        reference["meta"]["icf_template"] = str(case.get("icf_template") or reference["meta"].get("icf_template") or ("Sterling" if richness == "rich" else "Advarra"))
    if richness == "rich":
        reference["study"]["background"] = f"{reference['study']['background'].rstrip('.')} with extended follow-up."
        reference["study"]["timeline"] = "6 months"
        reference["procedures"]["assessments"] = list(reference["procedures"].get("assessments", [])) + ["Extended follow-up assessment"]
        reference["procedures"]["visit_schedule_table"] = list(reference["procedures"].get("visit_schedule_table", [])) + [{"visitNumber": "4", "visitName": "Month 6", "visitWindow": "±14 days", "CRFnumber": "CRF-04"}]
        reference["population"]["sample_size"] = "80 participants"
        for item in reference["population"].get("sample_size_evidence", []):
            if isinstance(item, dict): item["value"] = "80 participants"
        for item in reference.get("statistics", {}).get("sample_size_evidence", []):
            if isinstance(item, dict): item["value"] = "80 participants"
        if study_type != "Retrospective":
            reference["endpoints"].setdefault("other", []).extend([{"label": "Extended follow-up outcome", "time_point": "Month 6"}, {"label": "Participant experience outcome", "time_point": "Month 6"}])
            reference["design"]["interventions"] = [{"name": "Sentinel Patch", "type": "Device", "description": "Primary monitoring configuration.", "arm_group_label": "Primary cohort"}, {"name": "Sentinel Patch Extended", "type": "Device", "description": "Extended monitoring configuration.", "arm_group_label": "Extended cohort"}]
            reference["design"]["arms"] = [{"name": "Primary cohort", "description": "Primary observational cohort."}, {"name": "Extended cohort", "description": "Extended observational cohort."}]
        if reference.get("sites"):
            second_site = copy.deepcopy(reference["sites"][0]); second_site["facility"]["name"] = "Site Two"; second_site["contact"]["name"] = "Taylor Coordinator"; second_site["contact"]["email"] = "taylor@example.org"; second_site["investigators"][0]["name"] = "Morgan Investigator"
            reference["sites"].append(second_site); reference["design"]["number_of_sites"] = 2
            reference["design"]["study_design"] = str(reference["design"]["study_design"]).replace("single-center", "multicenter")
    reference["approval"] = {"status": "draft"}
    return reference

GOVERNED_GATE_SEQUENCE = (
    "clinical_fidelity",
    "content_completeness_consistency",
    "docx_prs_structure",
    "exact_artifact_rendering",
    "every_page_visual_qa",
    "exact_byte_atomic_delivery",
)
FORMAT_CONFORMANCE_MATRIX = "references/format-conformance-matrix.json"
CERTIFICATION_VISUAL_CHECKS = {
    "artificial_pagination", "bad_table_split", "blank_page", "clipping",
    "duplicate_section", "excessive_whitespace", "footer_collision",
    "inconsistent_style", "missing_header_footer", "orphan_heading",
    "overflow", "overlap", "toc_mismatch", "unreadable_text",
}

SYMBOL_FONT_FALLBACKS = (
    "Apple Symbols",
    "Segoe UI Symbol",
    "Arial Unicode MS",
    "Arial Unicode",
    "Symbol",
    "DejaVu Sans",
    "Noto Sans",
    "Arial",
    "Helvetica",
    "Liberation Sans",
)
SANS_FONT_FALLBACKS = (
    "Arial",
    "Helvetica",
    "Aptos",
    "Calibri",
    "Liberation Sans",
    "DejaVu Sans",
    "Noto Sans",
    "Verdana",
)
SERIF_FONT_FALLBACKS = (
    "Times New Roman",
    "Times",
    "Liberation Serif",
    "DejaVu Serif",
    "Georgia",
    "Cambria",
    "Noto Serif",
)
MONOSPACE_FONT_FALLBACKS = (
    "Courier New",
    "Menlo",
    "Consolas",
    "Liberation Mono",
    "DejaVu Sans Mono",
    "Noto Sans Mono",
)


def verification_request_sha256(request: Mapping[str, Any]) -> str:
    unsigned = dict(request)
    unsigned.pop("request_sha256", None)
    return hashlib.sha256(json.dumps(unsigned, sort_keys=True, ensure_ascii=False).encode()).hexdigest()


def verification_request_hash_valid(request: Mapping[str, Any]) -> bool:
    supplied = str(request.get("request_sha256") or "")
    return bool(supplied) and supplied == verification_request_sha256(request)


def _revision_owned_nonsymlink_path(revision_dir: Path, path: Path) -> bool:
    """Require a lexical descendant with no symlink in its revision-owned chain."""
    try:
        relative = path.relative_to(revision_dir)
    except ValueError:
        return False
    absolute_revision = revision_dir.absolute()
    current = Path(absolute_revision.anchor)
    for part in absolute_revision.parts[1:]:
        current = current / part
        if current.is_symlink():
            return False
    for part in relative.parts:
        current = current / part
        if current.is_symlink():
            return False
    try:
        return path.resolve(strict=False).is_relative_to(revision_dir.resolve(strict=True))
    except OSError:
        return False


def verification_request_ledger_record(
    request_path: Path,
    request: Mapping[str, Any],
    *,
    request_bytes: bytes | None = None,
) -> dict[str, Any]:
    """Bind one workflow-owned request identity to its exact serialized bytes."""
    return {
        "schema_version": "verification-request-ledger/v1",
        "request_id": request.get("request_id"),
        "request_sha256": request.get("request_sha256"),
        "request_file_sha256": (
            hashlib.sha256(request_bytes).hexdigest()
            if request_bytes is not None else sha256_file(request_path)
        ),
        "task": request.get("task"),
        "revision_id": request.get("revision_id"),
        "review_set": request.get("review_set"),
    }


def verification_request_authentication_findings(
    revision_dir: Path,
    request_path: Path,
    request: Mapping[str, Any],
    *,
    request_id_counts: Mapping[str, int],
) -> list[dict[str, Any]]:
    """Authenticate a mutable verifier request against workflow-owned authority."""
    request_id = request.get("request_id")
    task = request.get("task")
    target = (
        "verification:visual"
        if task == "rendered_page_visual_verification"
        else "verification:content"
    )

    def finding(code: str, issue: str) -> dict[str, Any]:
        return recovery_finding({
            "category": "verification",
            "field": "canonical_request_ledger",
            "code": code,
            "target_ids": [target],
            "verification_request_id": request_id,
            "issue": issue,
        }, "document_structure_defect")

    if not isinstance(request_id, str) or not request_id or Path(request_id).name != request_id:
        return [finding(
            "verification_request_not_found",
            "Verification request identity is absent or is not a canonical ledger key.",
        )]
    canonical_request_path = (
        revision_dir / "hermes/verification-requests" / f"{request_id}.json"
    )
    if (
        request_path != canonical_request_path
        or not _revision_owned_nonsymlink_path(revision_dir, request_path)
    ):
        return [finding(
            "verification_request_noncanonical_path",
            "Verification request is not the canonical non-symlinked request path for its identity.",
        )]
    if request_id_counts.get(request_id, 0) != 1:
        return [finding(
            "verification_request_id_duplicated",
            "Verification request identity does not resolve exactly once in the active revision.",
        )]
    ledger_path = revision_dir / "request-ledger" / f"{request_id}.json"
    if (
        not ledger_path.is_file()
        or not _revision_owned_nonsymlink_path(revision_dir, ledger_path)
    ):
        return [finding(
            "verification_request_ledger_missing",
            "Workflow-owned canonical verification request ledger entry is missing.",
        )]
    try:
        ledger = _json(ledger_path)
    except (OSError, ValueError, json.JSONDecodeError):
        return [finding(
            "verification_request_ledger_invalid",
            "Workflow-owned canonical verification request ledger entry is invalid.",
        )]
    expected_ledger_keys = {
        "schema_version", "request_id", "request_sha256", "request_file_sha256",
        "task", "revision_id", "review_set",
    }
    if (
        set(ledger) != expected_ledger_keys
        or ledger.get("schema_version") != "verification-request-ledger/v1"
    ):
        return [finding(
            "verification_request_ledger_invalid",
            "Workflow-owned canonical verification request ledger has a noncanonical schema or shape.",
        )]
    if request.get("revision_id") != revision_dir.name or ledger.get("revision_id") != revision_dir.name:
        return [finding(
            "verification_request_wrong_revision",
            "Verification request is not bound to the active Run Revision.",
        )]
    if ledger.get("request_id") != request_id:
        return [finding(
            "verification_request_not_found",
            "Verification request identity does not match its canonical ledger entry.",
        )]
    if ledger.get("task") != task:
        return [finding(
            "verification_request_wrong_task",
            "Verification request task differs from its canonical ledger task.",
        )]
    if ledger.get("review_set") != request.get("review_set"):
        return [finding(
            "verification_request_stale_review_set",
            "Verification request review set differs from its canonical ledger binding.",
        )]
    canonical_hash = ledger.get("request_sha256")
    if canonical_hash != request.get("request_sha256"):
        return [finding(
            "verification_request_canonical_hash_mismatch",
            "Supplied verification request hash differs from the workflow-owned canonical hash.",
        )]
    if not verification_request_hash_valid(request):
        return [finding(
            "verification_request_integrity_failure",
            "Verification request body does not match its declared canonical payload identity.",
        )]
    if ledger.get("request_file_sha256") != sha256_file(request_path):
        return [finding(
            "verification_request_exact_bytes_mismatch",
            "Verification request bytes differ from the workflow-owned canonical request bytes.",
        )]
    for inspection in request.get("native_inspection", []):
        for path_key, hash_key in (("path", "sha256"), ("native_path", "native_sha256")):
            path = revision_dir / str(inspection.get(path_key) or "")
            if (not _revision_owned_nonsymlink_path(revision_dir, path)
                    or not path.is_file() or sha256_file(path) != inspection.get(hash_key)):
                return [finding("native_inspection_changed", "Native inspection evidence or its exact native file changed; rebuild the review from current bytes.")]
    for view in (request.get("review_plan") or {}).get("readable_native", []):
        path = revision_dir / str(view.get("path") or "")
        if (not _revision_owned_nonsymlink_path(revision_dir, path) or not path.is_file()
                or sha256_file(path) != view.get("sha256")):
            return [finding("readable_native_changed", "Readable native review evidence changed.")]
    transport = request.get("review_transport") or {}
    if transport:
        index = revision_dir / str(transport.get("index_path") or "")
        if (not _revision_owned_nonsymlink_path(revision_dir, index) or not index.is_file()
                or sha256_file(index) != transport.get("index_sha256")):
            return [finding("review_packet_changed", "Bounded review reading index changed.")]
    for view in transport.get("views", []):
        for packet in view.get("packets", []):
            path = revision_dir / str(packet.get("path") or "")
            if (not _revision_owned_nonsymlink_path(revision_dir, path)
                    or not path.is_file() or sha256_file(path) != packet.get("sha256")):
                return [finding("review_packet_changed", "Bounded review reading evidence changed.")]
    return []


PAGE_RENDERER_BACKENDS = ("pypdfium2",)


def _text(value: Any) -> str:
    return "" if value is None else str(value).strip()


def _json(path: Path) -> dict[str, Any]:
    value = json.loads(path.read_text(encoding="utf-8"))
    if not isinstance(value, dict): raise ValueError(f"Expected JSON object: {path}")
    return value


def _write(path: Path, value: Mapping[str, Any]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True); path.write_text(json.dumps(value, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""): digest.update(chunk)
    return digest.hexdigest()


def canonical_evidence_sha256(value: Any) -> str:
    """Hash a machine-readable evidence value using one canonical encoding."""
    return hashlib.sha256(json.dumps(
        value,
        sort_keys=True,
        separators=(",", ":"),
        ensure_ascii=False,
    ).encode("utf-8")).hexdigest()


def _normalized_ooxml_hash(element: ET._Element, *, redact_text: bool = True) -> str:
    clone = ET.fromstring(ET.tostring(element))
    font_attributes = {"ascii", "hAnsi", "eastAsia", "cs", "font", "name", "typeface"}
    for node in clone.iter():
        for attribute in list(node.attrib):
            local_name = ET.QName(attribute).localname
            if local_name.startswith("rsid"):
                del node.attrib[attribute]
            elif local_name in font_attributes:
                font = node.attrib[attribute].casefold()
                if font in APPROVED_PACKAGED_FONT_FALLBACKS:
                    node.attrib[attribute] = APPROVED_PACKAGED_FONT_FALLBACKS[font]
        if redact_text and ET.QName(node).localname in {"t", "delText"}:
            node.text = "#TEXT" if node.text else ""
    return hashlib.sha256(ET.tostring(clone, method="c14n")).hexdigest()


def normalized_docx_format_signature(path: Path) -> dict[str, Any]:
    """Return semantic, text-normalized OOXML format evidence for one DOCX."""
    with zipfile.ZipFile(path) as archive:
        document_xml = ET.fromstring(archive.read("word/document.xml"))
        styles_xml = ET.fromstring(archive.read("word/styles.xml"))
        numbering_xml = ET.fromstring(archive.read("word/numbering.xml"))
        settings_xml = ET.fromstring(archive.read("word/settings.xml"))
        headers_footers = []
        for name in sorted(
            item for item in archive.namelist()
            if re.fullmatch(r"word/(?:header|footer)\d+\.xml", item)
        ):
            root = ET.fromstring(archive.read(name))
            fields = [" ".join(str(node.text or "").split()) for node in root.xpath(".//*[local-name()='instrText']")]
            headers_footers.append({
                "part": name,
                "semantic_sha256": _normalized_ooxml_hash(root),
                "fields": fields,
                "table_count": len(root.xpath(".//*[local-name()='tbl']")),
            })

    section_geometry = []
    for section in document_xml.xpath(".//*[local-name()='sectPr']"):
        item: dict[str, Any] = {}
        for child_name in ("pgSz", "pgMar", "cols", "type"):
            matches = section.xpath(f"./*[local-name()='{child_name}']")
            child = matches[0] if matches else None
            item[child_name] = (
                {ET.QName(key).localname: value for key, value in sorted(child.attrib.items())}
                if child is not None else None
            )
        section_geometry.append(item)

    table_geometry = []
    for table in document_xml.xpath(".//*[local-name()='tbl']"):
        rows = table.xpath("./*[local-name()='tr']")
        table_geometry.append({
            "grid_widths": [node.get(qn("w:w"), "") for node in table.xpath("./*[local-name()='tblGrid']/*[local-name()='gridCol']")],
            "row_count": len(rows),
            "repeating_header_rows": sum(bool(row.xpath("./*[local-name()='trPr']/*[local-name()='tblHeader']")) for row in rows),
            "non_splitting_rows": sum(bool(row.xpath("./*[local-name()='trPr']/*[local-name()='cantSplit']")) for row in rows),
            "semantic_sha256": _normalized_ooxml_hash(table),
        })

    paragraphs = document_xml.xpath(".//*[local-name()='body']//*[local-name()='p']")
    paragraph_records = []
    visible_paragraphs = []
    for index, paragraph in enumerate(paragraphs):
        text = " ".join("".join(str(node.text or "") for node in paragraph.xpath(".//*[local-name()='t']")).split())
        style_nodes = paragraph.xpath("./*[local-name()='pPr']/*[local-name()='pStyle']")
        style = style_nodes[0].get(qn("w:val"), "") if style_nodes else ""
        ppr = paragraph.xpath("./*[local-name()='pPr']")
        paragraph_records.append({
            "style": style,
            "ppr_sha256": _normalized_ooxml_hash(ppr[0]) if ppr else None,
            "run_property_sha256": [_normalized_ooxml_hash(node) for node in paragraph.xpath("./*[local-name()='r']/*[local-name()='rPr']")],
        })
        if text:
            visible_paragraphs.append((index, text, style, paragraph))

    signature_pattern = re.compile(r"\b(?:signature|signed|participant name|investigator name)\b", re.I)
    legal_pattern = re.compile(r"\b(?:consent|authorization|privacy|confidential|injury|compensation|withdraw|voluntary)\b", re.I)
    signature_blocks = [
        {"paragraph": index, "text_sha256": hashlib.sha256(text.encode("utf-8")).hexdigest(), "style": style}
        for index, text, style, _ in visible_paragraphs if signature_pattern.search(text)
    ]
    legal_terms = ("consent", "authorization", "privacy", "confidential", "injury", "compensation", "withdraw", "voluntary")
    consent_legal_placement = [
        {
            "paragraph": index,
            "text_sha256": hashlib.sha256(text.encode("utf-8")).hexdigest(),
            "style": style,
            "categories": [term for term in legal_terms if re.search(rf"\b{term}\w*\b", text, re.I)],
        }
        for index, text, style, _ in visible_paragraphs if legal_pattern.search(text)
    ]
    headings = []
    protocol_document = path.name == "protocol.docx"
    style_page_breaks = {
        str(style_node.get(qn("w:styleId"), ""))
        for style_node in styles_xml.xpath(".//*[local-name()='style']")
        if style_node.xpath("./*[local-name()='pPr']/*[local-name()='pageBreakBefore']")
    }
    for index, text, style, paragraph in visible_paragraphs:
        if not style.casefold().startswith("heading"):
            continue
        ppr_matches = paragraph.xpath("./*[local-name()='pPr']")
        ppr = ppr_matches[0] if ppr_matches else None
        direct_page_break = bool(ppr is not None and ppr.xpath("./*[local-name()='pageBreakBefore']"))
        effective_page_boundary = direct_page_break or style in style_page_breaks
        if protocol_document and not effective_page_boundary:
            previous = paragraph.getprevious()
            while previous is not None and ET.QName(previous).localname == "p":
                if previous.xpath('.//*[local-name()="br" and @*[local-name()="type"]="page"]'):
                    effective_page_boundary = True
                    break
                section = previous.xpath("./*[local-name()='pPr']/*[local-name()='sectPr']")
                if section and not section[0].xpath(
                    "./*[local-name()='type' and @*[local-name()='val']='continuous']"
                ):
                    effective_page_boundary = True
                    break
                if previous.xpath('.//*[local-name()="t" and normalize-space(text())]'):
                    break
                previous = previous.getprevious()
        headings.append({
            "paragraph": index,
            "text": text,
            "style": style,
            "page_break_before": direct_page_break,
            **({"effective_page_boundary": effective_page_boundary} if protocol_document else {}),
            "keep_with_next": bool(ppr is not None and ppr.xpath("./*[local-name()='keepNext']")),
        })
    field_codes = [" ".join(str(node.text or "").split()) for node in document_xml.xpath(".//*[local-name()='instrText']")]
    page_fields = sum("PAGE" in value.upper() for value in field_codes) + sum(
        "PAGE" in value.upper() for part in headers_footers for value in part["fields"]
    )
    update_fields = bool(settings_xml.xpath(".//*[local-name()='updateFields' and (@*[local-name()='val']='true' or @*[local-name()='val']='1')]"))
    numbered_headings = [
        item for item in headings if re.match(r"^\d+(?:\.|\s)", item["text"])
    ]
    toc_index = next((
        index for index, item in enumerate(headings)
        if "table of contents" in re.sub(r"[^a-z0-9]+", " ", item["text"].casefold())
    ), None)
    body_headings = (
        [
            item for item in headings[toc_index + 1:]
            if re.match(r"^\d+(?:\.|\s)", item["text"])
        ]
        if toc_index is not None
        else numbered_headings
    )
    first_numbered = body_headings[0]["paragraph"] if body_headings else None
    later_body_headings = body_headings[1:] if toc_index is not None else body_headings
    return {
        "section_geometry": section_geometry,
        "styles": {"count": len(styles_xml.xpath(".//*[local-name()='style']")), "semantic_sha256": _normalized_ooxml_hash(styles_xml, redact_text=False)},
        "numbering": {
            "abstract_count": len(numbering_xml.xpath(".//*[local-name()='abstractNum']")),
            "number_count": len(numbering_xml.xpath(".//*[local-name()='num']")),
            "semantic_sha256": _normalized_ooxml_hash(numbering_xml, redact_text=False),
        },
        "headers_footers": headers_footers,
        "fields_toc": {"codes": field_codes, "update_fields": update_fields},
        "page_furniture": {
            "header_count": sum(item["part"].startswith("word/header") for item in headers_footers),
            "footer_count": sum(item["part"].startswith("word/footer") for item in headers_footers),
            "page_field_count": page_fields,
        },
        "table_geometry": table_geometry,
        "signature_blocks": signature_blocks,
        "consent_legal_placement": consent_legal_placement,
        "paragraph_rhythm": {"paragraph_count": len(paragraph_records), "semantic_sha256": canonical_evidence_sha256(paragraph_records)},
        "pagination_relations": {
            "headings": headings,
            "first_numbered_body_paragraph": first_numbered,
            "numbered_body_has_no_artificial_starts": all(
                not item["page_break_before"] for item in later_body_headings
            ),
            "all_headings_keep_with_next": all(item["keep_with_next"] for item in headings),
        },
    }


def _natural_section_three_flow(
    pages: Sequence[str], *, explicit_break_before: bool,
) -> bool:
    """Check the contracted front-matter boundary without demanding one page."""
    body_pages = [page for page in pages if "table of contents" not in page]
    section_two = next((index for index, page in enumerate(body_pages) if "2 investigator agreement" in page), None)
    section_three = next((index for index, page in enumerate(body_pages) if "3 general information" in page), None)
    if section_two is None or section_three is None:
        return False
    return (
        section_three == section_two + 1 if explicit_break_before
        else section_three in {section_two, section_two + 1}
    )


def _rendered_pagination_relations(docx_path: Path, pdf_path: Path) -> dict[str, Any]:
    document = Document(docx_path)
    pages = [
        re.sub(r"[^a-z0-9]+", " ", (page.extract_text() or "").casefold()).strip()
        for page in PdfReader(pdf_path).pages
    ]
    body = list(document.element.body)
    protocol = docx_path.name == "protocol.docx"
    headings = [
        paragraph for paragraph in document.paragraphs
        if paragraph.style.name.casefold().startswith("heading")
        and (not protocol or re.match(r"^\d+(?:\.\d+)*\.?\s+", paragraph.text.strip()))
        and (protocol or paragraph.style.name == "Heading ICF Section")
        and paragraph.text.strip()
        and "TITLE PAGE" not in paragraph.text.upper()
        and "TABLE OF CONTENTS" not in paragraph.text.upper()
    ]
    cohesion = []
    for heading in headings:
        heading_index = body.index(heading._p)
        first_content = ""
        for element in body[heading_index + 1:]:
            if element.tag == qn("w:tbl"):
                table = Table(element, document)
                first_content = " ".join(cell.text for cell in table.rows[0].cells) if table.rows else ""
                break
            if element.tag != qn("w:p"):
                continue
            paragraph = Paragraph(element, document)
            if paragraph.style.name.casefold().startswith("heading"):
                first_content = paragraph.text
                break
            if paragraph.text.strip():
                first_content = paragraph.text
                break
        heading_tokens = re.findall(r"[a-z0-9]+", heading.text.casefold())
        content_marker = " ".join(re.findall(r"[a-z0-9]+", first_content.casefold())[:4])
        together = False
        if first_content and heading_tokens and content_marker:
            for page in pages:
                if protocol:
                    heading_marker = " ".join(heading_tokens)
                    if heading_marker not in page:
                        continue
                    heading_end = page.index(heading_marker) + len(heading_marker)
                else:
                    first = page.find(heading_tokens[0])
                    if first < 0:
                        continue
                    heading_end = first + len(heading_tokens[0])
                    if len(heading_tokens) > 1:
                        second = page.find(heading_tokens[1], heading_end)
                        if second < 0:
                            continue
                        heading_end = second + len(heading_tokens[1])
                remainder = page[heading_end:]
                if protocol and content_marker in remainder:
                    together = True
                    break
                if not protocol:
                    content_tokens = [
                        token for token in re.findall(r"[a-z0-9]+", first_content.casefold())
                        if token not in {"a", "an", "the", "this"}
                    ][:3]
                    cursor = 0
                    for token in content_tokens:
                        position = remainder.find(token, cursor)
                        if position < 0:
                            break
                        cursor = position + len(token)
                    else:
                        if content_tokens:
                            together = True
                            break
        cohesion.append({
            "heading": heading.text.strip(),
            "first_content_sha256": hashlib.sha256(first_content.encode("utf-8")).hexdigest() if first_content else None,
            "same_page": together,
        })
    section_three = next((paragraph.text.strip() for paragraph in headings if "GENERAL INFORMATION" in paragraph.text.upper()), None)
    section_three_flow = True
    if section_three:
        signature = normalized_docx_format_signature(docx_path)
        explicit_break = any(
            item["page_break_before"]
            for item in signature["pagination_relations"]["headings"]
            if "GENERAL INFORMATION" in item["text"].upper()
        )
        section_three_flow = _natural_section_three_flow(
            pages, explicit_break_before=explicit_break,
        )
    return {
        "pdf_sha256": sha256_file(pdf_path),
        "heading_cohesion": cohesion,
        "all_headings_with_first_content": bool(cohesion) and all(item["same_page"] for item in cohesion),
        "natural_section_3_flow": section_three_flow,
    }


def audit_format_conformance_outputs(
    repo_root: Path,
    matrix: Mapping[str, Any],
    release_gate_report: Mapping[str, Any],
    evidence_root: Path,
) -> dict[str, Any]:
    """Compare generated deterministic DOCX semantics with approved baselines."""
    reported = {str(item.get("case")): item for item in release_gate_report.get("cases", []) if isinstance(item, Mapping)}
    results = []
    for case in matrix.get("cases", []):
        case_id = str(case.get("case_id") or "")
        lifecycle_case = FORMAT_CONFORMANCE_BASELINE_CASES.get(case_id, "")
        gate_case = reported.get(lifecycle_case, {})
        findings = []
        baseline_record = case.get("approved_output_baseline", {})
        baseline_path = repo_root / str(baseline_record.get("path") or "")
        if not baseline_path.is_file() or sha256_file(baseline_path) != baseline_record.get("sha256"):
            findings.append({"code": "FORMAT_BASELINE_IDENTITY_MISMATCH", "target": case_id})
            baseline = {"artifacts": {}}
        else:
            baseline = json.loads(baseline_path.read_text(encoding="utf-8"))
        revision_id = str((gate_case.get("result") or {}).get("revision_id") or "")
        revision_dir = evidence_root / lifecycle_case / "revisions" / revision_id
        candidate_dir = revision_dir / "candidate"
        actual_artifacts = {}
        for artifact, expected_signature in baseline.get("artifacts", {}).items():
            artifact_path = candidate_dir / artifact
            if not artifact_path.is_file():
                findings.append({"code": "FORMAT_CANDIDATE_MISSING", "target": f"{case_id}:{artifact}"})
                continue
            actual_signature = normalized_docx_format_signature(artifact_path)
            actual_artifacts[artifact] = {"sha256": sha256_file(artifact_path), "signature_sha256": canonical_evidence_sha256(actual_signature)}
            pdf_path = revision_dir / "rendered" / f"{artifact_path.stem}.pdf"
            if not pdf_path.is_file():
                findings.append({"code": "FORMAT_RENDERED_PDF_MISSING", "target": f"{case_id}:{artifact}"})
            else:
                relations = _rendered_pagination_relations(artifact_path, pdf_path)
                actual_artifacts[artifact]["rendered_pagination"] = relations
                if not relations["all_headings_with_first_content"]:
                    findings.append({"code": "FORMAT_ORPHAN_HEADING", "target": f"{case_id}:{artifact}"})
                if artifact == "protocol.docx" and not relations["natural_section_3_flow"]:
                    findings.append({"code": "FORMAT_SECTION_3_FLOW_BROKEN", "target": f"{case_id}:{artifact}#section=3"})
            if actual_signature != expected_signature:
                findings.append({
                    "code": "FORMAT_BASELINE_MISMATCH",
                    "target": f"{case_id}:{artifact}",
                    "expected_signature_sha256": canonical_evidence_sha256(expected_signature),
                    "actual_signature_sha256": canonical_evidence_sha256(actual_signature),
                })
        results.append({
            "case_id": case_id,
            "status": "passed" if not findings else "blocked",
            "baseline_sha256": baseline_record.get("sha256"),
            "artifacts": actual_artifacts,
            "findings": findings,
        })
    return {
        "status": "passed" if results and all(item["status"] == "passed" for item in results) else "blocked",
        "cases": results,
        "evidence_sha256": canonical_evidence_sha256(results),
    }


def load_format_conformance_matrix(repo_root: Path) -> dict[str, Any]:
    """Load and independently rehash the deterministic format matrix."""
    root = repo_root.resolve()
    matrix = _json(root / FORMAT_CONFORMANCE_MATRIX)
    if matrix.get("schema_version") != "clinical-format-conformance/v1":
        raise ValueError("Format conformance matrix schema is invalid.")
    unsigned = dict(matrix)
    declared_sha256 = str(unsigned.pop("matrix_sha256", ""))
    if not declared_sha256 or declared_sha256 != canonical_evidence_sha256(unsigned):
        raise ValueError("Format conformance matrix identity is invalid.")
    gates = matrix.get("gate_sequence")
    if not isinstance(gates, list) or tuple(
        item.get("gate_id") for item in gates if isinstance(item, Mapping)
    ) != GOVERNED_GATE_SEQUENCE:
        raise ValueError("Format conformance gate order is invalid.")
    cases = matrix.get("cases")
    expected_cases = (
        "retrospective-protocol", "prospective-advarra", "prospective-sterling",
        "ambispective-advarra", "ambispective-sterling",
    )
    if not isinstance(cases, list) or tuple(
        item.get("case_id") for item in cases if isinstance(item, Mapping)
    ) != expected_cases:
        raise ValueError("Format conformance case inventory is invalid.")
    for case in cases:
        resources = [
            case.get("fixture"),
            case.get("approved_output_baseline"),
            *list(case.get("baselines") or []),
        ]
        for resource in resources:
            if not isinstance(resource, Mapping):
                raise ValueError("Format conformance resource declaration is invalid.")
            path_text = str(resource.get("path") or "")
            relative = PurePosixPath(path_text)
            if relative.is_absolute() or ".." in relative.parts or relative.as_posix() != path_text:
                raise ValueError("Format conformance resource path is invalid.")
            target = root / relative
            if not target.is_file() or sha256_file(target) != resource.get("sha256"):
                raise ValueError(f"Format conformance resource identity is invalid: {relative}")
    return matrix


def validate_gate_ledger(repo_root: Path, ledger: Mapping[str, Any]) -> dict[str, Any]:
    """Reject reordered, stale, incomplete, or non-monotonic delivery gates."""
    matrix = load_format_conformance_matrix(repo_root)
    value = dict(ledger)
    if value.get("schema_version") != "clinical-gate-ledger/v1":
        raise ValueError("Gate ledger schema is invalid.")
    if value.get("matrix_sha256") != matrix.get("matrix_sha256"):
        raise ValueError("Gate ledger is not bound to the current conformance matrix.")
    attempt_id = str(value.get("attempt_id") or "")
    if not attempt_id:
        raise ValueError("Gate ledger attempt identity is missing.")
    unsigned = dict(value)
    declared_sha256 = str(unsigned.pop("ledger_sha256", ""))
    if not declared_sha256 or declared_sha256 != canonical_evidence_sha256(unsigned):
        raise ValueError("Gate ledger identity is invalid.")
    records = value.get("records")
    if not isinstance(records, list) or tuple(
        record.get("gate_id") for record in records if isinstance(record, Mapping)
    ) != GOVERNED_GATE_SEQUENCE:
        raise ValueError("Gate ledger sequence is invalid.")
    unresolved = False
    finding_fields = set(matrix["finding_contract"]["required"])
    predecessors = value.get("predecessors")
    if not isinstance(predecessors, list):
        raise ValueError("Gate ledger predecessor history is invalid.")
    predecessor_hashes = []
    gate_retry_owners = {
        str(item["gate_id"]): str(item["retry_owner"])
        for item in matrix["gate_sequence"]
    }
    for predecessor in predecessors:
        if not isinstance(predecessor, Mapping):
            raise ValueError("Gate ledger predecessor entry is invalid.")
        ledger_sha256 = str(predecessor.get("ledger_sha256") or "")
        blocked_findings = predecessor.get("blocked_findings")
        predecessor_code = f"{str(predecessor.get('terminal_gate')).upper()}_FAILED"
        if (
            predecessor.get("attempt_id") != attempt_id
            or not re.fullmatch(r"[0-9a-f]{64}", ledger_sha256)
            or predecessor.get("terminal_gate") not in GOVERNED_GATE_SEQUENCE
            or not isinstance(blocked_findings, list)
            or not blocked_findings
            or any(
                not isinstance(finding, Mapping)
                or not finding_fields <= set(finding)
                or not re.fullmatch(r"[A-Z][A-Z0-9_]+", str(finding.get("code") or ""))
                or finding.get("code") != predecessor_code
                or not str(finding.get("target") or "").strip()
                or not re.fullmatch(r"[0-9a-f]{64}", str(finding.get("evidence_sha256") or ""))
                or finding.get("retry_owner") != gate_retry_owners.get(str(predecessor.get("terminal_gate")))
                or finding.get("terminal_status") != "blocked"
                for finding in blocked_findings
            )
        ):
            raise ValueError("Gate ledger predecessor evidence is invalid.")
        predecessor_hashes.append(ledger_sha256)
    if len(predecessor_hashes) != len(set(predecessor_hashes)):
        raise ValueError("Gate ledger predecessor history contains duplicates.")
    for order, record in enumerate(records, start=1):
        if not isinstance(record, Mapping) or record.get("order") != order:
            raise ValueError("Gate ledger order is invalid.")
        status = str(record.get("terminal_status") or "")
        if status not in {"passed", "pending", "blocked"}:
            raise ValueError("Gate ledger terminal status is invalid.")
        if unresolved and status == "passed":
            raise ValueError(f"Gate {record.get('gate_id')} cannot pass after an unresolved earlier gate.")
        unresolved = unresolved or status != "passed"
        evidence_sha256 = str(record.get("evidence_sha256") or "")
        if not re.fullmatch(r"[0-9a-f]{64}", evidence_sha256):
            raise ValueError("Gate evidence identity is invalid.")
        expected_retry_owner = gate_retry_owners.get(str(record.get("gate_id")))
        if record.get("retry_owner") != expected_retry_owner:
            raise ValueError("Gate retry ownership does not match the governed matrix.")
        findings = record.get("findings")
        governed_finding_code = f"{str(record.get('gate_id')).upper()}_FAILED"
        if not isinstance(findings, list) or any(
            not isinstance(finding, Mapping)
            or not finding_fields <= set(finding)
            or not re.fullmatch(r"[A-Z][A-Z0-9_]+", str(finding.get("code") or ""))
            or finding.get("code") != governed_finding_code
            or not str(finding.get("target") or "").strip()
            or not re.fullmatch(r"[0-9a-f]{64}", str(finding.get("evidence_sha256") or ""))
            or not str(finding.get("retry_owner") or "").strip()
            or finding.get("retry_owner") != record.get("retry_owner")
            or finding.get("terminal_status") != "blocked"
            for finding in findings
        ):
            raise ValueError("Gate finding contract is invalid.")
        if (status == "blocked") != bool(findings):
            raise ValueError("Only blocked gates may carry findings, and blocked gates require them.")
    return value


def build_gate_ledger(
    repo_root: Path,
    evidence_by_gate: Mapping[str, Any],
    *,
    attempt_id: str,
    predecessors: Iterable[Mapping[str, Any]] = (),
    statuses: Mapping[str, str] | None = None,
    findings_by_gate: Mapping[str, Iterable[Mapping[str, Any]]] | None = None,
) -> dict[str, Any]:
    """Build a canonical monotonic ledger from exact evidence values."""
    matrix = load_format_conformance_matrix(repo_root)
    if not str(attempt_id).strip():
        raise ValueError("Gate ledger attempt identity is required.")
    gates = {str(item["gate_id"]): item for item in matrix["gate_sequence"]}
    if set(evidence_by_gate) != set(GOVERNED_GATE_SEQUENCE):
        raise ValueError("Evidence must cover the complete governed gate sequence.")
    records = []
    for order, gate_id in enumerate(GOVERNED_GATE_SEQUENCE, start=1):
        status = str((statuses or {}).get(gate_id, "passed"))
        records.append({
            "order": order,
            "gate_id": gate_id,
            "terminal_status": status,
            "evidence_sha256": canonical_evidence_sha256(evidence_by_gate[gate_id]),
            "retry_owner": str(gates[gate_id]["retry_owner"]),
            "findings": [dict(item) for item in (findings_by_gate or {}).get(gate_id, [])],
        })
    if all(record["terminal_status"] == "passed" for record in records):
        raise ValueError("A complete passed ledger must advance from its final pending delivery gate.")
    ledger = {
        "schema_version": "clinical-gate-ledger/v1",
        "matrix_sha256": matrix["matrix_sha256"],
        "attempt_id": str(attempt_id),
        "predecessors": [dict(item) for item in predecessors],
        "records": records,
    }
    ledger["ledger_sha256"] = canonical_evidence_sha256(ledger)
    return validate_gate_ledger(repo_root, ledger)


def retry_gate_ledger(
    repo_root: Path,
    blocked_ledger: Mapping[str, Any],
    evidence_by_gate: Mapping[str, Any],
    *,
    statuses: Mapping[str, str],
    findings_by_gate: Mapping[str, Iterable[Mapping[str, Any]]] | None = None,
) -> dict[str, Any]:
    """Create one linked retry while retaining the exact terminal blocked ledger."""
    prior = validate_gate_ledger(repo_root, blocked_ledger)
    blocked_records = [
        record for record in prior["records"]
        if record["terminal_status"] == "blocked"
    ]
    if len(blocked_records) != 1:
        raise ValueError("A retry requires exactly one terminal blocked predecessor gate.")
    blocked = blocked_records[0]
    predecessor = {
        "attempt_id": prior["attempt_id"],
        "ledger_sha256": prior["ledger_sha256"],
        "terminal_gate": blocked["gate_id"],
        "blocked_findings": [dict(item) for item in blocked["findings"]],
    }
    return build_gate_ledger(
        repo_root,
        evidence_by_gate,
        attempt_id=str(prior["attempt_id"]),
        predecessors=[*list(prior["predecessors"]), predecessor],
        statuses=statuses,
        findings_by_gate=findings_by_gate,
    )


def advance_gate_ledger(
    repo_root: Path,
    ledger: Mapping[str, Any],
    *,
    gate_id: str,
    terminal_status: str,
    evidence: Any,
    findings: Iterable[Mapping[str, Any]] = (),
) -> dict[str, Any]:
    """Advance only the first unresolved gate and bind its replacement evidence."""
    current = validate_gate_ledger(repo_root, ledger)
    records = [dict(record) for record in current["records"]]
    unresolved = next(
        (record for record in records if record["terminal_status"] != "passed"),
        None,
    )
    if unresolved is None or unresolved["gate_id"] != gate_id:
        raise ValueError(f"Only the first unresolved gate may advance; requested {gate_id}.")
    if unresolved["terminal_status"] == "blocked":
        raise ValueError(f"A terminal blocked gate cannot advance; start a new retained attempt for {gate_id}.")
    unresolved["terminal_status"] = terminal_status
    unresolved["evidence_sha256"] = canonical_evidence_sha256(evidence)
    unresolved["findings"] = [dict(item) for item in findings]
    advanced = {
        "schema_version": current["schema_version"],
        "matrix_sha256": current["matrix_sha256"],
        "attempt_id": current["attempt_id"],
        "predecessors": list(current["predecessors"]),
        "records": records,
    }
    advanced["ledger_sha256"] = canonical_evidence_sha256(advanced)
    return validate_gate_ledger(repo_root, advanced)


def release_certification_payload(report: Mapping[str, Any]) -> bytes:
    """Canonical bytes covered by the detached release-certification signature."""
    unsigned = dict(report)
    unsigned.pop("evidence_attestation", None)
    return json.dumps(
        unsigned,
        sort_keys=True,
        separators=(",", ":"),
        ensure_ascii=False,
    ).encode("utf-8")


def _canonical_sha256_value(value: Any) -> str:
    return hashlib.sha256(json.dumps(
        value,
        sort_keys=True,
        separators=(",", ":"),
        ensure_ascii=False,
    ).encode("utf-8")).hexdigest()


def release_certification_key_id(key: Mapping[str, Any]) -> str:
    """Return the canonical identity shared by certification signers and verifiers."""
    identity = {
        "algorithm": str(key.get("algorithm") or ""),
        "exponent": int(str(key.get("exponent") or "")),
        "modulus": format(int(str(key.get("modulus") or ""), 16), "x"),
    }
    return _canonical_sha256_value(identity)


def _validated_certification_evidence(bundle: Mapping[str, Any]) -> list[dict[str, Any]]:
    """Decode one bounded canonical evidence inventory and rehash every byte."""
    if bundle.get("schema_version") != "release-certification-evidence/v1":
        raise ValueError("Certification evidence bundle schema is invalid.")
    entries = bundle.get("entries")
    if not isinstance(entries, list) or not entries or len(entries) > CERTIFICATION_EVIDENCE_MAX_FILES:
        raise ValueError("Certification evidence bundle has an invalid file count.")
    identities: set[str] = set()
    paths: set[str] = set()
    total_bytes = 0
    decoded: list[dict[str, Any]] = []
    for raw_item in entries:
        if not isinstance(raw_item, Mapping):
            raise ValueError("Certification evidence inventory item is invalid.")
        item = dict(raw_item)
        identity = str(item.get("identity") or "")
        identity_key = identity.casefold()
        if not identity or identity_key in identities:
            raise ValueError("Certification evidence contains a duplicate evidence identity.")
        identities.add(identity_key)
        if not re.fullmatch(r"[a-z0-9][a-z0-9._-]{0,127}", identity):
            raise ValueError("Certification evidence identity is noncanonical.")
        path_text = str(item.get("path") or "")
        relative = PurePosixPath(path_text)
        if (
            not path_text
            or "\\" in path_text
            or any(part in {"", ".", ".."} for part in path_text.split("/"))
            or relative.is_absolute()
            or relative.as_posix() != path_text
            or relative.parts[0] not in {"global", "cases"}
        ):
            raise ValueError(f"Certification evidence has a noncanonical evidence path: {path_text}")
        path_key = path_text.casefold()
        if path_key in paths:
            raise ValueError("Certification evidence contains a duplicate evidence path.")
        paths.add(path_key)
        declared_bytes = item.get("bytes")
        if isinstance(declared_bytes, bool) or not isinstance(declared_bytes, int):
            raise ValueError("Certification evidence byte length is invalid.")
        if declared_bytes < 0 or declared_bytes > CERTIFICATION_EVIDENCE_MAX_ITEM_BYTES:
            raise ValueError("Certification evidence item exceeds the governed byte limit.")
        try:
            content = base64.b64decode(
                str(item.get("content_base64") or ""), validate=True
            )
        except (binascii.Error, ValueError) as exc:
            raise ValueError("Certification evidence content is not strict base64.") from exc
        if (
            len(content) != declared_bytes
            or hashlib.sha256(content).hexdigest() != item.get("sha256")
        ):
            raise ValueError("Certification evidence bytes do not match the inventory.")
        total_bytes += declared_bytes
        if total_bytes > CERTIFICATION_EVIDENCE_MAX_TOTAL_BYTES:
            raise ValueError("Certification evidence bundle exceeds the governed byte limit.")
        decoded.append({**item, "content": content})
    return decoded


def _certification_evidence_findings(
    report: Mapping[str, Any],
    manifest: Mapping[str, Any],
    manifest_bytes: bytes,
) -> list[str]:
    """Independently derive certification claims from retained evidence bytes."""
    bundle = report.get("evidence_bundle") or {}
    try:
        entries = _validated_certification_evidence(bundle)
    except ValueError as exc:
        return [str(exc)]
    metadata = [
        {key: value for key, value in item.items() if key not in {"content", "content_base64"}}
        for item in entries
    ]
    if (
        bundle.get("inventory_sha256") != _canonical_sha256_value(metadata)
        or bundle.get("total_bytes") != sum(item["bytes"] for item in entries)
    ):
        return ["Certification evidence inventory identity or total byte count is invalid."]
    allowed_kinds = {
        "release_manifest", "preflight", "preflight_log", "layout_preservation", "deterministic_corpus",
        "runtime_identity", "case_report", "desktop_operation_state", "delivery_manifest",
        "output", "delivery_confirmation", "drafting_request", "drafting_response",
        "verification_request", "verification_response", "delegated_page_review",
        "parent_page_review", "parent_process_marker",
        "fixture_manifest", "fixture_source", "approved_source", "approved_reference",
        "content_artifact", "pdf", "page_image",
    }
    findings: list[str] = []
    indexed: dict[tuple[str | None, str], list[dict[str, Any]]] = {}
    parsed: dict[str, dict[str, Any] | None] = {}
    consumed_identities: set[str] = set()
    for item in entries:
        identity_key = str(item["identity"])
        kind = str(item.get("kind") or "")
        case_id = item.get("case_id")
        path = str(item.get("path") or "")
        if kind not in allowed_kinds:
            findings.append(f"Evidence kind is not release-authorized: {kind!r}.")
            continue
        if case_id is None:
            if not path.startswith("global/"):
                findings.append(f"Global evidence has a case path: {path}.")
        elif case_id not in CERTIFICATION_CASE_ORDER or not path.startswith(f"cases/{case_id}/"):
            findings.append(f"Case evidence is mis-scoped: {path}.")
        indexed.setdefault((case_id, kind), []).append(item)
        if kind not in {
            "output", "content_artifact", "pdf", "page_image", "release_manifest"
        }:
            try:
                value = json.loads(item["content"])
                parsed[identity_key] = value if isinstance(value, dict) else None
            except (UnicodeDecodeError, json.JSONDecodeError):
                parsed[identity_key] = None

    def items(case_id: str | None, kind: str) -> list[dict[str, Any]]:
        return indexed.get((case_id, kind), [])

    def consume(evidence: Iterable[Mapping[str, Any]]) -> None:
        consumed_identities.update(str(item.get("identity") or "") for item in evidence)

    def one(case_id: str | None, kind: str) -> dict[str, Any] | None:
        matches = items(case_id, kind)
        if len(matches) != 1:
            findings.append(f"Exactly one {kind!r} item is required for {case_id or 'global'}.")
            return None
        consume(matches)
        return matches[0]

    def payload(item: Mapping[str, Any] | None) -> dict[str, Any]:
        return parsed.get(str((item or {}).get("identity") or "")) or {}

    manifest_item = one(None, "release_manifest")
    if manifest_item is not None and manifest_item["content"] != manifest_bytes:
        findings.append("Retained release-manifest bytes do not match the installed manifest.")
    identity = report.get("release_identity") or {}
    preflight_item = one(None, "preflight")
    preflight = payload(preflight_item)
    required_checks = {
        "layout_preservation_corpus", "deterministic_branch_acceptance_corpus",
        "repository_regression_suite", "static_release_checks",
    }
    check_values = preflight.get("checks") or {}
    if (
        preflight_item is None
        or preflight_item.get("sha256") != report.get("preflight_evidence_sha256")
        or preflight.get("status") != "passed"
        or any(
            (preflight.get("candidate") or {}).get(key) != identity.get(key)
            for key in ("package_fingerprint", "git_commit")
        )
        or set(check_values) != required_checks
        or any(
            not isinstance(value, Mapping)
            or value.get("status") != "passed"
            or value.get("returncode") != 0
            for value in check_values.values()
        )
    ):
        findings.append("Preflight evidence does not independently pass every required check.")
    layout_item = one(None, "layout_preservation")
    layout = payload(layout_item)
    layout_summary = report.get("layout_preservation_evidence") or {}
    layout_check = check_values.get("layout_preservation_corpus") or {}
    if (
        layout_item is None
        or layout_item.get("sha256") != layout_summary.get("sha256")
        or not (
            layout.get("status") == "passed" and layout.get("returncode") == 0
            or isinstance(layout_check, Mapping)
            and layout_check.get("status") == "passed"
            and layout_check.get("returncode") == 0
            and layout_check.get("sha256") == layout_item.get("sha256")
        )
        or set((layout or layout_check).get("coverage") or []) != set(layout_summary.get("coverage") or [])
    ):
        findings.append("Layout-preservation evidence is incomplete or stale.")
    deterministic_item = one(None, "deterministic_corpus")
    deterministic = payload(deterministic_item)
    deterministic_check = check_values.get("deterministic_branch_acceptance_corpus") or {}
    if (
        not (
            deterministic.get("status") == "passed"
            and set(deterministic.get("case_ids") or []) == set(DETERMINISTIC_BRANCH_ACCEPTANCE_CASES)
            or isinstance(deterministic_check, Mapping)
            and deterministic_check.get("status") == "passed"
            and deterministic_check.get("returncode") == 0
            and deterministic_item is not None
            and deterministic_check.get("sha256") == deterministic_item.get("sha256")
            and set(deterministic_check.get("case_ids") or []) == set(DETERMINISTIC_BRANCH_ACCEPTANCE_CASES)
        )
    ):
        findings.append("Deterministic branch-acceptance evidence is incomplete.")
    preflight_logs = {
        str(item.get("identity") or ""): item
        for item in items(None, "preflight_log")
    }
    consume(preflight_logs.values())
    required_preflight_logs = {
        "static_release_checks", "repository_regression_suite",
    }
    if (
        set(preflight_logs) != required_preflight_logs
        or any(
            not isinstance(check_values.get(name), Mapping)
            or check_values[name].get("sha256") != preflight_logs[name].get("sha256")
            for name in required_preflight_logs
        )
    ):
        findings.append("Static and repository-regression preflight logs are incomplete.")
    runtime = payload(one(None, "runtime_identity"))
    configurations = report.get("hermes_configurations") or {}
    expected_modules = {
        item["path"]: item["sha256"]
        for item in manifest.get("files") or []
        if str(item.get("path") or "").startswith("scripts/")
        and str(item.get("path") or "").endswith(".py")
    }
    if (
        runtime.get("release_identity") != identity
        or runtime.get("python") != preflight.get("python_runtime")
        or runtime.get("page_renderer") != (manifest.get("inventory") or {}).get("pdf_page_renderer")
        or runtime.get("production_modules") != expected_modules
        or not str((runtime.get("python") or {}).get("version") or "")
        or not re.fullmatch(r"[0-9a-f]{64}", str((runtime.get("python") or {}).get("executable_sha256") or ""))
        or runtime.get("hermes_configuration_sha256") != {
            fixture: _canonical_sha256_value(configuration) for fixture, configuration in configurations.items()
        }
    ):
        findings.append("Python, worker, skill, renderer, or Hermes configuration identity is invalid.")

    cases = {str(item.get("fixture_id") or ""): item for item in report.get("cases") or []}
    for fixture in CERTIFICATION_CASE_ORDER:
        case = cases.get(fixture) or {}
        models = set(case.get("model_identifiers") or [])
        case_item = one(fixture, "case_report")
        case_payload = payload(case_item)
        state = payload(one(fixture, "desktop_operation_state"))
        delivery_manifest = payload(one(fixture, "delivery_manifest"))
        delivery = payload(one(fixture, "delivery_confirmation"))
        raw_case_identity = case_payload.get("release_identity") or {}
        if (
            case_item is None or case_item.get("sha256") != case.get("report_sha256")
            or case_payload.get("status", case_payload.get("outcome")) != "passed"
            or {key: raw_case_identity.get(key) for key in identity} != identity
            or case_payload.get("hermes_configuration") != configurations.get(fixture)
            or set(case_payload.get("model_identifiers") or []) != models
        ):
            findings.append(f"Case report is incomplete or stale for {fixture}.")
        state_identity = state.get("release_identity") or {}
        if (
            state.get("status") != "passed"
            or (state.get("result") or {}).get("status") != "passed"
            or not ((state.get("result") or {}).get("delivery") or {}).get("confirmed")
            or not (state.get("cleanup") or {}).get("owned_processes_reaped")
            or {key: state_identity.get(key) for key in identity} != identity
            or state_identity.get("hermes_configuration") != configurations.get(fixture)
        ):
            findings.append(f"Desktop operation evidence does not pass for {fixture}.")
        expected_outputs = {str(item.get("path") or ""): item for item in case.get("output_evidence") or []}
        actual_outputs = {
            str(item["path"])[len(f"cases/{fixture}/"):]: item for item in items(fixture, "output")
        }
        consume(actual_outputs.values())
        if set(actual_outputs) != set(expected_outputs):
            findings.append(f"Output evidence does not cover the delivered set for {fixture}.")
        for path, expected in expected_outputs.items():
            actual = actual_outputs.get(path) or {}
            if actual.get("sha256") != expected.get("sha256") or actual.get("bytes") != expected.get("bytes") or expected.get("confirmed") is not True:
                findings.append(f"Output bytes do not match {fixture}:{path}.")
        manifest_outputs = {
            str(item.get("path") or ""): {
                key: item.get(key) for key in ("path", "sha256", "bytes")
            }
            for item in delivery_manifest.get("client_outputs") or []
        }
        expected_manifest_outputs = {
            path: {key: item.get(key) for key in ("path", "sha256", "bytes")}
            for path, item in expected_outputs.items()
        }
        opened_outputs = {str(item.get("path") or ""): item for item in delivery.get("opened") or []}
        if (
            delivery_manifest.get("status") != "passed"
            or (delivery_manifest.get("quality") or {}).get("status") != "passed"
            or manifest_outputs != expected_manifest_outputs
            or delivery.get("confirmed") is not True
            or opened_outputs != expected_outputs
        ):
            findings.append(f"Delivery evidence is invalid for {fixture}.")
        draft_request_items = items(fixture, "drafting_request")
        draft_response_items = items(fixture, "drafting_response")
        consume(draft_request_items)
        consume(draft_response_items)
        draft_requests = {payload(item).get("request_id"): payload(item) for item in draft_request_items}
        draft_responses = [payload(item) for item in draft_response_items]
        if (
            not draft_requests
            or len(draft_requests) != len(draft_responses)
            or any(
                not (request := draft_requests.get(response.get("request_id")))
                or response.get("request_sha256") != request.get("request_sha256")
                or response.get("status") not in {None, "passed"}
                or (response.get("producer") or {}).get("model_id") not in models
                for response in draft_responses
            )
        ):
            findings.append(f"Drafting worker evidence is incomplete for {fixture}.")
        fixture_manifest = payload(one(fixture, "fixture_manifest"))
        fixture_source = one(fixture, "fixture_source")
        approved_source = one(fixture, "approved_source")
        approved_reference = one(fixture, "approved_reference")
        provenance = fixture_manifest.get("provenance") or fixture_manifest
        if provenance.get("synthetic") is not True or provenance.get("contains_private_data") is not False:
            findings.append(f"Fixture evidence is not explicitly synthetic and non-private for {fixture}.")
        artifact_entries = {
            "source_input": fixture_source,
            "approved_source": approved_source,
            "approved_reference": approved_reference,
        }
        manifest_artifacts = fixture_manifest.get("artifacts") or {}
        if any(
            entry is None
            or (manifest_artifacts.get(name) or {}).get("sha256") != entry.get("sha256")
            for name, entry in artifact_entries.items()
        ):
            findings.append(f"Retained fixture source bytes do not match the governed fixture manifest for {fixture}.")
        normalization = fixture_manifest.get("approved_normalization") or {}
        if (
            normalization.get("source_input_sha256") != (fixture_source or {}).get("sha256")
            or normalization.get("approved_source_sha256") != (approved_source or {}).get("sha256")
        ):
            findings.append(f"Retained approved source normalization is invalid for {fixture}.")
        request_items = items(fixture, "verification_request")
        response_items = items(fixture, "verification_response")
        consume(request_items)
        consume(response_items)
        requests = [(item, payload(item)) for item in request_items]
        content_requests = [(item, value) for item, value in requests if value.get("task") == "clinical_content_verification"]
        content_responses = [(item, payload(item)) for item in response_items]
        content_artifacts = items(fixture, "content_artifact")
        consume(content_artifacts)
        if len(content_requests) != 1 or len(content_responses) != 1:
            findings.append(f"Content verification evidence is incomplete for {fixture}.")
        else:
            request = content_requests[0][1]
            response = content_responses[0][1]
            requested_artifacts = {
                str(item.get("path") or ""): item
                for item in request.get("artifacts") or []
                if isinstance(item, Mapping)
            }
            retained_artifacts = {
                str(item.get("path") or "")[len(f"cases/{fixture}/content/"):]: item
                for item in content_artifacts
                if str(item.get("path") or "").startswith(
                    f"cases/{fixture}/content/"
                )
            }
            if (
                response.get("request_id") != request.get("request_id")
                or response.get("request_sha256") != request.get("request_sha256")
                or response.get("task") != request.get("task")
                or response.get("status") != "passed"
                or (response.get("producer") or {}).get("model_id") not in models
                or not response.get("section_assessments")
                or any(item.get("status") != "passed" for item in response.get("section_assessments") or [])
                or not response.get("cross_document_assessments")
                or any(item.get("status") != "passed" for item in response.get("cross_document_assessments") or [])
                or not requested_artifacts
                or set(retained_artifacts) != set(requested_artifacts)
                or any(
                    retained_artifacts[path].get("sha256")
                    != requested_artifacts[path].get("sha256")
                    for path in requested_artifacts
                )
            ):
                findings.append(f"Content verification does not pass for {fixture}.")
        visual_requests = [(item, value) for item, value in requests if value.get("task") == "rendered_page_visual_verification"]
        parent_review_items = items(fixture, "parent_page_review")
        delegated_review_items = items(fixture, "delegated_page_review")
        consume(parent_review_items)
        consume(delegated_review_items)
        visual_reviews = [
            (item, payload(item))
            for item in [*parent_review_items, *delegated_review_items]
        ]
        parent_marker_items = items(fixture, "parent_process_marker")
        consume(parent_marker_items)
        parent_marker = payload(parent_marker_items[0]) if len(parent_marker_items) == 1 else {}
        visual_summary = case.get("visual_qa") or {}
        if len(visual_requests) != len(visual_summary) or len(visual_reviews) != len(visual_summary):
            findings.append(f"Every-page review set is incomplete for {fixture}.")
        if parent_review_items and (
            len(parent_marker_items) != 1
            or parent_marker.get("status") != "completed"
            or parent_marker.get("completion_requirement")
            != "Desktop parent must inspect every bound page image."
            or parent_marker.get("producer_model_policy") != "record_actual_nonempty_model_id"
        ):
            findings.append(f"Desktop-parent review completion evidence is invalid for {fixture}.")
        if not parent_review_items and parent_marker_items:
            findings.append(f"Unexpected Desktop-parent marker exists without fallback review evidence for {fixture}.")
        for artifact, visual in visual_summary.items():
            request_pair = next((pair for pair in visual_requests if pair[0].get("sha256") == visual.get("request_sha256")), None)
            review_pair = next((
                pair for pair in visual_reviews
                if any(part.get("artifact") == artifact for part in pair[1].get("page_assessments") or [])
            ), None)
            pdf = next((item for item in items(fixture, "pdf") if str(item["path"]).endswith(f"/{artifact}.pdf")), None)
            page_items = [item for item in items(fixture, "page_image") if f"/pages/{artifact}/" in str(item["path"])]
            consume([pdf] if pdf is not None else [])
            consume(page_items)
            expected_pages = list(visual.get("page_sha256") or [])
            request_item, request = request_pair or ({}, {})
            review_item, review = review_pair or ({}, {})
            assessments = review.get("page_assessments") or []
            docx = actual_outputs.get(f"output/{artifact}.docx") or {}
            retained_page_sha256 = [item.get("sha256") for item in page_items]
            request_artifacts = [
                item for item in request.get("artifacts") or []
                if item.get("artifact") == artifact
            ]
            request_artifact = request_artifacts[0] if len(request_artifacts) == 1 else {}
            if (
                request_item.get("sha256") != visual.get("request_sha256")
                or review_item.get("sha256") != visual.get("response_sha256")
                or review.get("request_id") != request.get("request_id")
                or review.get("request_sha256") != request.get("request_sha256")
                or request_artifact.get("docx_sha256") != docx.get("sha256")
                or request_artifact.get("pdf_sha256") != (pdf or {}).get("sha256")
                or [item.get("sha256") for item in request_artifact.get("pages") or []]
                != retained_page_sha256
                or review.get("status") != "passed"
                or (review.get("producer") or {}).get("model_id") not in models
                or (pdf or {}).get("sha256") != visual.get("pdf_sha256")
                or retained_page_sha256 != expected_pages
                or [item.get("sha256") for item in assessments] != expected_pages
                or any(item.get("status") != "passed" or set(item.get("checks") or []) != CERTIFICATION_VISUAL_CHECKS for item in assessments)
            ):
                findings.append(f"PDF or every-page review is invalid for {fixture}:{artifact}.")
    unconsumed = sorted(
        str(item.get("path") or item.get("identity") or "")
        for item in entries
        if str(item.get("identity") or "") not in consumed_identities
    )
    if unconsumed:
        findings.append(
            "Certification evidence contains unrelated or semantically unbound items: "
            + ", ".join(unconsumed)
            + "."
        )
    return findings


def release_certification_attestation_findings(
    report: Mapping[str, Any],
    manifest: Mapping[str, Any],
    skill_root: Path,
    *,
    trusted_key_id: str = RELEASE_CERTIFICATION_TRUSTED_KEY_ID,
) -> list[str]:
    """Verify report provenance against the manifest-bound release public key."""
    findings: list[str] = []
    key_path = skill_root / RELEASE_CERTIFICATION_PUBLIC_KEY
    try:
        key_bytes = key_path.read_bytes()
        key = json.loads(key_bytes)
    except (OSError, json.JSONDecodeError) as exc:
        return [f"Release certification public key is unavailable or invalid: {exc}"]
    manifest_entry = next(
        (
            item for item in manifest.get("files") or []
            if item.get("path") == RELEASE_CERTIFICATION_PUBLIC_KEY
        ),
        None,
    )
    if (
        manifest_entry is None
        or manifest_entry.get("bytes") != len(key_bytes)
        or manifest_entry.get("sha256") != hashlib.sha256(key_bytes).hexdigest()
    ):
        findings.append("Release certification public key is not bound to the immutable manifest.")
    try:
        modulus = int(str(key.get("modulus") or ""), 16)
        exponent = int(key.get("exponent"))
    except (TypeError, ValueError):
        return findings + ["Release certification public key parameters are invalid."]
    try:
        key_identity = release_certification_key_id(key)
    except (TypeError, ValueError):
        return findings + ["Release certification public key identity is invalid."]
    if (
        key.get("schema_version") != "release-certification-public-key/v1"
        or key.get("algorithm") != RELEASE_CERTIFICATION_SIGNATURE_ALGORITHM
        or key.get("key_id") != key_identity
        or key_identity != trusted_key_id
        or exponent < 3
        or exponent % 2 == 0
        or modulus.bit_length() < 2048
    ):
        findings.append("Release certification public key identity or strength is invalid.")
    attestation = report.get("evidence_attestation") or {}
    payload_digest = hashlib.sha256(release_certification_payload(report)).digest()
    if (
        attestation.get("schema_version") != "release-certification-attestation/v1"
        or attestation.get("algorithm") != RELEASE_CERTIFICATION_SIGNATURE_ALGORITHM
        or attestation.get("key_id") != key_identity
        or attestation.get("payload_sha256") != payload_digest.hex()
    ):
        findings.append("Release certification evidence attestation metadata is missing or invalid.")
        return findings
    try:
        signature = base64.b64decode(
            str(attestation.get("signature_base64") or ""),
            validate=True,
        )
    except (binascii.Error, ValueError):
        return findings + ["Release certification evidence signature is not strict base64."]
    encoded_bytes = (modulus.bit_length() + 7) // 8
    digest_info = _SHA256_DIGEST_INFO_PREFIX + payload_digest
    expected = b"\x00\x01" + b"\xff" * (encoded_bytes - len(digest_info) - 3) + b"\x00" + digest_info
    if len(signature) != encoded_bytes:
        findings.append("Release certification evidence signature has an invalid length.")
    else:
        recovered = pow(int.from_bytes(signature, "big"), exponent, modulus).to_bytes(encoded_bytes, "big")
        if not hmac.compare_digest(recovered, expected):
            findings.append("Release certification evidence signature does not verify.")
    return findings


def _manifest_package_fingerprint(manifest: Mapping[str, Any]) -> tuple[str, str]:
    """Return the recorded and canonical fingerprints for one release manifest."""
    payload = dict(manifest)
    recorded = str(payload.pop("package_fingerprint", ""))
    computed = hashlib.sha256(
        json.dumps(
            payload, sort_keys=True, separators=(",", ":"), ensure_ascii=False
        ).encode("utf-8")
    ).hexdigest()
    return recorded, computed


def _content_sha256(path: Path) -> str:
    """Hash visible DOCX text independently from its presentation properties."""
    if path.suffix.casefold() != ".docx":
        return sha256_file(path)
    try:
        stories: list[tuple[str, list[str]]] = []
        with zipfile.ZipFile(path) as package:
            for name in sorted(package.namelist()):
                if not name.startswith("word/") or not name.endswith(".xml"):
                    continue
                root = ET.fromstring(package.read(name))
                values = [
                    str(element.text or "")
                    for element in root.iter()
                    if ET.QName(element).localname in {"t", "instrText"}
                ]
                if values:
                    stories.append((name, values))
        payload = json.dumps(stories, ensure_ascii=False, separators=(",", ":")).encode("utf-8")
        return hashlib.sha256(payload).hexdigest()
    except (OSError, ET.XMLSyntaxError, zipfile.BadZipFile):
        return sha256_file(path)


def _executable_candidates(
    names: Iterable[str],
    *,
    environment: Mapping[str, str] | None = None,
    home: Path | None = None,
) -> Iterable[tuple[Path, str]]:
    """Yield PATH and common installation candidates without changing the host."""
    search_path = (environment or {}).get("PATH") if environment is not None else None
    yielded: set[str] = set()
    for name in names:
        if path := shutil.which(name, path=search_path):
            resolved = str(Path(path).resolve())
            yielded.add(resolved.casefold())
            yield Path(resolved), "PATH"

    user_home = (home or Path.home()).expanduser()
    directories = [
        user_home / ".hermes/bin",
        Path(sys.executable).resolve().parent,
        Path("/opt/homebrew/bin"),
        Path("/usr/local/bin"),
        Path("/usr/bin"),
        Path("/bin"),
    ]
    sources = {
        str(user_home / ".hermes/bin"): "Hermes bundled tools",
        str(Path(sys.executable).resolve().parent): "Python environment",
    }
    bundled_tools_root = user_home / ".cache/codex-runtimes"
    if bundled_tools_root.is_dir():
        for pattern in ("*/dependencies/bin/override", "*/dependencies/bin/fallback"):
            for directory in sorted(bundled_tools_root.glob(pattern)):
                directories.append(directory)
                sources[str(directory)] = "bundled workspace tools"
    for variable in ("ProgramFiles", "ProgramFiles(x86)", "ProgramData"):
        if value := (environment or os.environ).get(variable):
            base = Path(value)
            directories.extend((base / "LibreOffice/program", base / "chocolatey/bin"))

    for directory in directories:
        for name in names:
            candidates = (directory / name, directory / f"{name}.exe")
            for candidate in candidates:
                if not candidate.is_file() or not os.access(candidate, os.X_OK):
                    continue
                resolved = str(candidate.resolve())
                key = resolved.casefold()
                if key in yielded:
                    continue
                yielded.add(key)
                yield Path(resolved), sources.get(str(directory), "common installation path")


def _pdfium_integrity_failure(code: str, path: str, issue: str) -> dict[str, Any]:
    return {"status": "blocked", "finding": {
        "category": "renderer",
        "field": "page_renderer",
        "code": code,
        "path": path,
        "issue": issue,
    }}


def _pdfium_runtime_integrity(
    skill_root: Path,
    *,
    require_promoted_runtime: bool = True,
    allow_pending_activation: bool = False,
) -> dict[str, Any]:
    """Verify one installed runtime and explain the first exact mismatch."""
    resolved_root = Path(skill_root).resolve()
    active_lifecycle_evidence = any(
        path.exists() or path.is_symlink()
        for path in (
            resolved_root / "PROMOTION-RECORD.json",
            resolved_root / "INSTALLATION-ASSURANCE.json",
            resolved_root.parent / ".clinical-document-generation.activation.json",
        )
    )
    require_promoted_runtime = require_promoted_runtime or active_lifecycle_evidence
    runtime_root = resolved_root / "runtime"
    runtime_python = runtime_root / "python"
    identity_path = runtime_root / "PDF-RENDERER.json"
    if runtime_root.is_symlink():
        return _pdfium_integrity_failure(
            "renderer.pdfium_runtime_directory_symlink",
            "runtime",
            "The PDFium runtime directory must not be a symlink.",
        )
    if runtime_python.is_symlink():
        return _pdfium_integrity_failure(
            "renderer.pdfium_runtime_root_symlink",
            "runtime/python",
            "The PDFium runtime root must not be a symlink.",
        )
    if not runtime_python.is_dir():
        return _pdfium_integrity_failure(
            "renderer.pdfium_runtime_missing",
            "runtime/python",
            "The manifest-bound PDFium runtime directory is missing.",
        )
    if identity_path.is_symlink() or not identity_path.is_file():
        return _pdfium_integrity_failure(
            "renderer.pdfium_metadata_missing",
            "runtime/PDF-RENDERER.json",
            "The PDFium operational metadata is missing or symlinked.",
        )
    try:
        recorded = _json(identity_path)
        manifest = _json(resolved_root / "RELEASE-MANIFEST.json")
        recorded_fingerprint, computed_fingerprint = _manifest_package_fingerprint(manifest)
        if not recorded_fingerprint or computed_fingerprint != recorded_fingerprint:
            return _pdfium_integrity_failure(
                "renderer.pdfium_manifest_fingerprint_invalid",
                "RELEASE-MANIFEST.json",
                "The PDFium inventory is not bound to a valid release package fingerprint.",
            )
        promotion_path = resolved_root / "PROMOTION-RECORD.json"
        if require_promoted_runtime and not (
            promotion_path.exists() or promotion_path.is_symlink()
        ):
            return _pdfium_integrity_failure(
                "renderer.pdfium_promotion_record_invalid",
                "PROMOTION-RECORD.json",
                "The promoted release record is missing, malformed, or symlinked.",
            )
        if promotion_path.exists() or promotion_path.is_symlink():
            if promotion_path.is_symlink() or not promotion_path.is_file():
                return _pdfium_integrity_failure(
                    "renderer.pdfium_promotion_record_invalid",
                    "PROMOTION-RECORD.json",
                    "The promoted release record is missing, malformed, or symlinked.",
                )
            try:
                promotion = _json(promotion_path)
            except (OSError, ValueError, json.JSONDecodeError, TypeError):
                return _pdfium_integrity_failure(
                    "renderer.pdfium_promotion_record_invalid",
                    "PROMOTION-RECORD.json",
                    "The promoted release record is missing, malformed, or symlinked.",
                )
            if (
                promotion.get("schema_version") != "promoted-release/v1"
                or promotion.get("package_fingerprint") != recorded_fingerprint
            ):
                return _pdfium_integrity_failure(
                    "renderer.pdfium_promotion_fingerprint_mismatch",
                    "PROMOTION-RECORD.json",
                    "The runtime manifest does not match the promoted release fingerprint.",
                )
            promotion_status = promotion.get("status")
            if promotion_status not in {"active", "activation_pending"} or (
                promotion_status == "activation_pending"
                and not allow_pending_activation
            ):
                return _pdfium_integrity_failure(
                    "renderer.pdfium_activation_pending",
                    "PROMOTION-RECORD.json",
                    "The release has not passed the active-path smoke commit point.",
                )
            assurance_path = resolved_root / "INSTALLATION-ASSURANCE.json"
            expected_assurance_sha256 = promotion.get("runtime_assurance_sha256")
            if (
                not isinstance(expected_assurance_sha256, str)
                or re.fullmatch(r"[0-9a-f]{64}", expected_assurance_sha256) is None
                or assurance_path.is_symlink()
                or not assurance_path.is_file()
            ):
                return _pdfium_integrity_failure(
                    "renderer.pdfium_installation_assurance_mismatch",
                    "INSTALLATION-ASSURANCE.json",
                    "The promoted release does not match its bound installation assurance.",
                )
            try:
                actual_assurance_sha256 = sha256_file(assurance_path)
            except OSError:
                actual_assurance_sha256 = ""
            if actual_assurance_sha256 != expected_assurance_sha256:
                return _pdfium_integrity_failure(
                    "renderer.pdfium_installation_assurance_mismatch",
                    "INSTALLATION-ASSURANCE.json",
                    "The promoted release does not match its bound installation assurance.",
                )
            try:
                assurance = _json(assurance_path)
            except (OSError, ValueError, json.JSONDecodeError, TypeError):
                return _pdfium_integrity_failure(
                    "renderer.pdfium_installation_assurance_mismatch",
                    "INSTALLATION-ASSURANCE.json",
                    "The promoted release does not match its bound installation assurance.",
                )
            if promotion_status == "activation_pending":
                if (
                    assurance.get("schema_version") != "installation-assurance/v2"
                    or assurance.get("status") != "staging_passed"
                ):
                    return _pdfium_integrity_failure(
                        "renderer.pdfium_pending_assurance_invalid",
                        "INSTALLATION-ASSURANCE.json",
                        "The pending release does not match its staging assurance.",
                    )
            else:
                if promotion.get("activation_commit_point") != (
                    PROMOTED_ACTIVATION_COMMIT_POINT
                ):
                    return _pdfium_integrity_failure(
                        "renderer.pdfium_activation_commit_point_invalid",
                        "PROMOTION-RECORD.json",
                        "The promoted release lacks the active-path smoke commit point.",
                    )
                active_smoke = assurance.get("active_path_smoke")
                if (
                    assurance.get("schema_version") != "installation-assurance/v2"
                    or assurance.get("status") != "passed"
                    or not isinstance(active_smoke, Mapping)
                    or active_smoke.get("status") != "passed"
                ):
                    return _pdfium_integrity_failure(
                        "renderer.pdfium_active_smoke_assurance_invalid",
                        "INSTALLATION-ASSURANCE.json",
                        "The promoted release lacks a passing active-path smoke assurance.",
                    )
                active_identity = active_smoke.get("release_identity")
                expected_identity = {
                    "git_commit": manifest.get("git_commit"),
                    "package_fingerprint": recorded_fingerprint,
                }
                if not isinstance(active_identity, Mapping) or {
                    key: active_identity.get(key) for key in expected_identity
                } != expected_identity:
                    return _pdfium_integrity_failure(
                        "renderer.pdfium_active_smoke_identity_mismatch",
                        "INSTALLATION-ASSURANCE.json",
                        "The active-path smoke identity does not match the release manifest.",
                    )
                if promotion.get("active_path_smoke_sha256") != (
                    canonical_evidence_sha256(active_smoke)
                ):
                    return _pdfium_integrity_failure(
                        "renderer.pdfium_active_smoke_hash_mismatch",
                        "PROMOTION-RECORD.json",
                        "The active-path smoke evidence does not match its promotion binding.",
                    )
        manifest_identity = manifest["inventory"]["pdf_page_renderer"]
    except (OSError, ValueError, json.JSONDecodeError, KeyError, TypeError):
        return _pdfium_integrity_failure(
            "renderer.pdfium_manifest_invalid",
            "RELEASE-MANIFEST.json",
            "The release manifest does not contain a valid PDFium identity.",
        )
    identity_fields = ("kind", "version", "wheel", "wheel_sha256", "platform")
    if (
        recorded.get("kind") != "pypdfium2"
        or any(recorded.get(field) != manifest_identity.get(field) for field in identity_fields)
    ):
        return _pdfium_integrity_failure(
            "renderer.pdfium_metadata_mismatch",
            "runtime/PDF-RENDERER.json",
            "The PDFium operational identity does not match the release manifest.",
        )
    wheel_relative = Path(str(recorded.get("wheel") or ""))
    wheel_candidate = resolved_root / wheel_relative
    if wheel_candidate.is_symlink():
        return _pdfium_integrity_failure(
            "renderer.pdfium_wheel_symlink", wheel_relative.as_posix(),
            "The packaged PDFium wheel must not be a symlink.",
        )
    wheel = wheel_candidate.resolve()
    try:
        wheel.relative_to(resolved_root)
    except ValueError:
        return _pdfium_integrity_failure(
            "renderer.pdfium_wheel_escaping",
            wheel_relative.as_posix(),
            "The PDFium wheel path escapes the release root.",
        )
    if not wheel.is_file() or sha256_file(wheel) != recorded.get("wheel_sha256"):
        return _pdfium_integrity_failure(
            "renderer.pdfium_wheel_changed",
            wheel_relative.as_posix(),
            "The packaged PDFium wheel is missing, symlinked, or changed.",
        )
    expected_inventory = manifest_identity.get("runtime_inventory")
    if not isinstance(expected_inventory, list) or not expected_inventory:
        return _pdfium_integrity_failure(
            "renderer.pdfium_inventory_invalid",
            "RELEASE-MANIFEST.json",
            "The release manifest has no valid PDFium runtime inventory.",
        )
    expected_by_path: dict[str, dict[str, Any]] = {}
    for item in expected_inventory:
        if not isinstance(item, Mapping) or set(item) != {"path", "bytes", "sha256"}:
            return _pdfium_integrity_failure(
                "renderer.pdfium_inventory_invalid", "RELEASE-MANIFEST.json",
                "A PDFium inventory entry has an invalid shape.",
            )
        relative = Path(str(item.get("path") or ""))
        normalized = relative.as_posix()
        if (
            not normalized
            or relative.is_absolute()
            or ".." in relative.parts
            or normalized in expected_by_path
            or not isinstance(item.get("bytes"), int)
            or int(item["bytes"]) < 0
            or not re.fullmatch(r"[0-9a-f]{64}", str(item.get("sha256") or ""))
        ):
            return _pdfium_integrity_failure(
                "renderer.pdfium_inventory_escaping",
                normalized or "RELEASE-MANIFEST.json",
                "A PDFium inventory path is unsafe, duplicated, or malformed.",
            )
        expected_by_path[normalized] = {
            "bytes": int(item["bytes"]), "sha256": str(item["sha256"]),
        }
    if list(expected_by_path) != sorted(expected_by_path):
        return _pdfium_integrity_failure(
            "renderer.pdfium_inventory_invalid", "RELEASE-MANIFEST.json",
            "The PDFium runtime inventory is not deterministically ordered.",
        )
    actual_paths: dict[str, Path] = {}
    for path in runtime_python.rglob("*"):
        normalized = path.relative_to(runtime_python).as_posix()
        if path.is_symlink():
            return _pdfium_integrity_failure(
                "renderer.pdfium_runtime_symlink", normalized,
                f"The installed PDFium path is symlinked: {normalized}.",
            )
        if path.is_file():
            actual_paths[normalized] = path
    for normalized in expected_by_path:
        if normalized not in actual_paths:
            return _pdfium_integrity_failure(
                "renderer.pdfium_runtime_file_missing", normalized,
                f"The installed PDFium file is missing: {normalized}.",
            )
    for normalized in sorted(actual_paths):
        if normalized not in expected_by_path:
            return _pdfium_integrity_failure(
                "renderer.pdfium_runtime_file_extra", normalized,
                f"The installed PDFium runtime has an extra file: {normalized}.",
            )
        expected = expected_by_path[normalized]
        path = actual_paths[normalized]
        if path.stat().st_size != expected["bytes"] or sha256_file(path) != expected["sha256"]:
            return _pdfium_integrity_failure(
                "renderer.pdfium_runtime_file_changed", normalized,
                f"The installed PDFium file changed: {normalized}.",
            )
    return {"status": "passed", "identity": {
        "kind": "pypdfium2",
        "path": "python:pypdfium2",
        "module": "pypdfium2",
        "python_path": str(runtime_python),
        "version": str(recorded.get("version") or ""),
        "source": "release-owned runtime",
        "wheel": str(recorded.get("wheel") or ""),
        "wheel_sha256": str(recorded.get("wheel_sha256") or ""),
    }}


def page_renderers(
    *,
    environment: Mapping[str, str] | None = None,
    home: Path | None = None,
    skill_root: Path | None = None,
    require_promoted_runtime: bool = True,
    allow_pending_activation: bool = False,
) -> list[dict[str, Any]]:
    """Return the one release-owned PDFium page renderer, if verified."""
    del environment, home
    if skill_root is None:
        return []
    verification = _pdfium_runtime_integrity(
        Path(skill_root),
        require_promoted_runtime=require_promoted_runtime,
        allow_pending_activation=allow_pending_activation,
    )
    return [verification["identity"]] if verification["status"] == "passed" else []


def _write_pdfium_png(bitmap: Any, output: Path) -> None:
    """Write PDFium's reverse-byte-order RGB buffer with only the standard library."""
    width = int(bitmap.width)
    height = int(bitmap.height)
    stride = int(bitmap.stride)
    raw = bytes(bitmap.buffer)
    scanlines = b"".join(
        b"\x00" + raw[offset:offset + width * 3]
        for offset in range(0, stride * height, stride)
    )

    def chunk(kind: bytes, data: bytes) -> bytes:
        payload = kind + data
        return struct.pack(">I", len(data)) + payload + struct.pack(">I", zlib.crc32(payload) & 0xFFFFFFFF)

    header = struct.pack(">IIBBBBB", width, height, 8, 2, 0, 0, 0)
    output.write_bytes(
        b"\x89PNG\r\n\x1a\n"
        + chunk(b"IHDR", header)
        + chunk(b"IDAT", zlib.compress(scanlines))
        + chunk(b"IEND", b"")
    )


def page_renderer(
    *,
    environment: Mapping[str, str] | None = None,
    home: Path | None = None,
    skill_root: Path | None = None,
) -> dict[str, Any] | None:
    """Return the preferred installed PDF page renderer."""
    identities = page_renderers(environment=environment, home=home, skill_root=skill_root)
    return identities[0] if identities else None


def _is_release_owned_pdfium_identity(identity: Mapping[str, Any]) -> bool:
    """Return whether an identity names the installed release-owned runtime."""
    python_path = Path(str(identity.get("python_path") or "")).expanduser()
    return (
        identity.get("kind") == "pypdfium2"
        and identity.get("path") == "python:pypdfium2"
        and identity.get("module") == "pypdfium2"
        and identity.get("source") == "release-owned runtime"
        and python_path.is_absolute()
        and python_path.name == "python"
        and python_path.parent.name == "runtime"
    )


def _one_pdfium_renderer(
    identities: Iterable[Mapping[str, Any]],
    *,
    skill_root: Path | None = None,
    require_promoted_runtime: bool = True,
    allow_pending_activation: bool = False,
) -> list[dict[str, Any]]:
    """Keep one identity verified against this release's manifest and runtime."""
    candidates = [dict(identity) for identity in identities if _is_release_owned_pdfium_identity(identity)]
    if skill_root is None:
        return candidates[:1]
    verified = page_renderers(
        skill_root=skill_root,
        require_promoted_runtime=require_promoted_runtime,
        allow_pending_activation=allow_pending_activation,
    )
    identity_fields = (
        "kind", "path", "module", "python_path", "version", "source", "wheel", "wheel_sha256",
    )
    for identity in candidates:
        if any(all(identity.get(field) == known.get(field) for field in identity_fields) for known in verified):
            return [identity]
    return []


def _ordered_page_images(output_dir: Path) -> list[Path]:
    def key(path: Path) -> tuple[int, str]:
        match = re.search(r"(\d+)(?=\.png$)", path.name)
        return (int(match.group(1)) if match else -1, path.name)
    return sorted(output_dir.glob("page*.png"), key=key)


def _rasterize_pdfium_worker_render(
    pdf: Path,
    output_dir: Path,
    identity: Mapping[str, Any],
    *,
    first_page_only: bool = False,
    dpi: int = 130,
    require_promoted_runtime: bool = True,
    allow_pending_activation: bool = False,
) -> list[Path]:
    """Render inside the isolated worker after re-verifying its release."""
    output_dir.mkdir(parents=True, exist_ok=True)
    for stale in output_dir.glob("page*.png"):
        stale.unlink()
    kind = str(identity["kind"])
    if kind != "pypdfium2":
        raise RuntimeError(f"Unsupported page renderer: {kind}")

    release_root = Path(__file__).resolve().parents[1]
    governed = _one_pdfium_renderer(
        [identity],
        skill_root=release_root,
        require_promoted_runtime=require_promoted_runtime,
        allow_pending_activation=allow_pending_activation,
    )
    if not governed:
        raise RuntimeError(
            "PDF page rendering requires the manifest-verified release-owned runtime."
        )
    identity = governed[0]
    module_name = str(identity["module"])
    python_path = str(identity["python_path"])
    inserted = False
    previous_modules: dict[str, Any] = {}
    if python_path not in sys.path:
        sys.path.insert(0, python_path)
        inserted = True
    module_roots = ("pypdfium2", "pypdfium2_raw", "pypdfium2_cfg")
    for name in list(sys.modules):
        if any(name == root or name.startswith(root + ".") for root in module_roots):
            previous_modules[name] = sys.modules.pop(name)
    document = None
    previous_dont_write_bytecode = sys.dont_write_bytecode
    sys.dont_write_bytecode = True
    try:
        pdfium = importlib.import_module(module_name)
        module_file = Path(str(getattr(pdfium, "__file__", ""))).resolve()
        try:
            module_file.relative_to(Path(python_path).resolve())
        except ValueError as exc:
            raise RuntimeError(f"pypdfium2 resolved outside the release-owned runtime: {module_file}") from exc
        document = pdfium.PdfDocument(str(pdf))
        document_pages = len(document)
        if document_pages > PDFIUM_MAX_PAGES:
            raise RuntimeError(
                f"renderer.pdfium_page_limit_exceeded: PDF has {document_pages} pages; "
                f"the governed maximum is {PDFIUM_MAX_PAGES}."
            )
        count = min(document_pages, 1) if first_page_only else document_pages
        total_pixels = 0
        total_output_bytes = 0
        for index in range(count):
            page = document[index]
            width = max(1, int(math.ceil(float(page.get_width()) * dpi / 72.0)))
            height = max(1, int(math.ceil(float(page.get_height()) * dpi / 72.0)))
            if width > PDFIUM_MAX_DIMENSION_PIXELS or height > PDFIUM_MAX_DIMENSION_PIXELS:
                page.close()
                raise RuntimeError(
                    "renderer.pdfium_dimension_limit_exceeded: "
                    f"page {index + 1} would render at {width}x{height} pixels."
                )
            total_pixels += width * height
            if total_pixels > PDFIUM_MAX_TOTAL_PIXELS:
                page.close()
                raise RuntimeError(
                    "renderer.pdfium_pixel_limit_exceeded: "
                    f"rendering exceeded {PDFIUM_MAX_TOTAL_PIXELS} total pixels."
                )
            bitmap = page.render(scale=dpi / 72.0, rev_byteorder=True)
            try:
                if int(bitmap.format) != 2:
                    raise RuntimeError(f"PDFium returned unsupported bitmap format {bitmap.format}.")
                page_path = output_dir / f"page-{index + 1}.png"
                _write_pdfium_png(bitmap, page_path)
                total_output_bytes += page_path.stat().st_size
                if total_output_bytes > PDFIUM_MAX_OUTPUT_BYTES:
                    raise RuntimeError(
                        "renderer.pdfium_output_limit_exceeded: "
                        f"page images exceeded {PDFIUM_MAX_OUTPUT_BYTES} bytes."
                    )
            finally:
                bitmap.close()
                page.close()
    except ImportError as exc:
        raise RuntimeError(f"The release-owned pypdfium2 page renderer is unavailable: {exc}") from exc
    finally:
        sys.dont_write_bytecode = previous_dont_write_bytecode
        if document is not None:
            document.close()
        for name in list(sys.modules):
            if any(name == root or name.startswith(root + ".") for root in module_roots):
                sys.modules.pop(name, None)
        sys.modules.update(previous_modules)
        if inserted:
            sys.path.remove(python_path)

    pages = _ordered_page_images(output_dir)
    if not pages:
        raise RuntimeError(f"{kind} page-image export produced no PNG pages.")
    temporary = []
    for index, page in enumerate(pages, 1):
        target = output_dir / f".normalized-{index}.png"
        page.replace(target)
        temporary.append(target)
    normalized = []
    for index, page in enumerate(temporary, 1):
        target = output_dir / ("page.png" if first_page_only else f"page-{index}.png")
        page.replace(target)
        normalized.append(target)
    return normalized


def _required_pdfium_worker_resource_limits() -> list[str]:
    system = platform.system()
    if system == "Windows":
        return []
    names = ["RLIMIT_CPU", "RLIMIT_FSIZE", "RLIMIT_NOFILE"]
    if system == "Linux":
        names.insert(1, "RLIMIT_AS")
    return names


def _supported_pdfium_worker_resource_limits() -> list[str]:
    try:
        import resource
    except ImportError:
        return []
    return [
        name for name in _required_pdfium_worker_resource_limits()
        if getattr(resource, name, None) is not None
    ]


def _apply_pdfium_worker_resource_limits() -> list[str]:
    """Apply every supported hard process limit before loading PDFium."""
    required = _required_pdfium_worker_resource_limits()
    try:
        import resource
    except ImportError as exc:
        if not required:
            return []
        raise RuntimeError(
            "renderer.pdfium_resource_limits_unavailable: "
            "The worker host does not expose required process limits."
        ) from exc
    applied = []
    limits = (
        ("RLIMIT_CPU", max(1, int(PDFIUM_WORKER_TIMEOUT_SECONDS))),
        ("RLIMIT_AS", PDFIUM_WORKER_MEMORY_BYTES),
        ("RLIMIT_FSIZE", PDFIUM_WORKER_FILE_BYTES),
        ("RLIMIT_NOFILE", 128),
    )
    supported = set(_supported_pdfium_worker_resource_limits())
    missing = [name for name in required if name not in supported]
    if missing:
        raise RuntimeError(
            "renderer.pdfium_resource_limits_unavailable: "
            f"The worker host does not expose required process limits: {', '.join(missing)}."
        )
    for name, requested in limits:
        if name not in supported:
            continue
        kind = getattr(resource, name, None)
        if kind is None:
            continue
        try:
            _soft, hard = resource.getrlimit(kind)
            finite_hard = requested if hard == resource.RLIM_INFINITY else min(requested, hard)
            resource.setrlimit(kind, (finite_hard, finite_hard))
            applied.append(name)
        except (OSError, ValueError) as exc:
            raise RuntimeError(
                "renderer.pdfium_resource_limit_failed: "
                f"The worker could not apply supported process limit {name}."
            ) from exc
    return applied


def run_pdfium_worker(request_path: Path) -> dict[str, Any]:
    """Internal workflow action for one isolated, bounded PDFium render."""
    try:
        request = json.loads(request_path.read_text(encoding="utf-8"))
        pdf = Path(str(request["pdf"])).resolve()
        output_dir = Path(str(request["output_dir"])).resolve()
        require_promoted_runtime = request.get("require_promoted_runtime", True)
        allow_pending_activation = request.get("allow_pending_activation", False)
        if (
            type(require_promoted_runtime) is not bool
            or type(allow_pending_activation) is not bool
        ):
            raise RuntimeError(
                "renderer.pdfium_worker_protocol_invalid: Promotion mode must be boolean."
            )
        if not pdf.is_file() or sha256_file(pdf) != request.get("pdf_sha256"):
            raise RuntimeError(
                "renderer.pdfium_input_changed: The PDF bytes changed before worker rendering."
            )
        output_dir.mkdir(parents=True, exist_ok=True)
        if output_dir.is_symlink() or any(output_dir.iterdir()):
            raise RuntimeError(
                "renderer.pdfium_worker_output_invalid: Worker output staging was not empty."
            )
        applied_limits = _apply_pdfium_worker_resource_limits()
        pages = _rasterize_pdfium_worker_render(
            pdf,
            output_dir,
            request["identity"],
            first_page_only=bool(request.get("first_page_only")),
            dpi=int(request.get("dpi", 130)),
            require_promoted_runtime=require_promoted_runtime,
            allow_pending_activation=allow_pending_activation,
        )
        evidence = [
            {"name": page.name, "bytes": page.stat().st_size, "sha256": sha256_file(page)}
            for page in pages
        ]
        return {"status": "passed", "pages": evidence, "resource_limits": applied_limits}
    except Exception as exc:
        message = str(exc)
        code, separator, issue = message.partition(": ")
        if not separator or not code.startswith("renderer.pdfium_"):
            code = "renderer.pdfium_worker_crashed"
            issue = message or type(exc).__name__
        return {"status": "blocked", "findings": [{
            "category": "renderer",
            "field": "page_renderer",
            "code": code,
            "issue": issue,
        }]}


def _pdfium_worker_command(request_path: Path) -> list[str]:
    release_root = Path(__file__).resolve().parents[1]
    return [
        sys.executable,
        "-B",
        str(release_root / "scripts/workflow.py"),
        "--internal-pdfium-worker",
        str(request_path),
    ]


def _clear_page_output(output_dir: Path) -> None:
    output_dir.mkdir(parents=True, exist_ok=True)
    for path in output_dir.iterdir():
        if path.is_dir() and not path.is_symlink():
            shutil.rmtree(path)
        else:
            path.unlink(missing_ok=True)


def _terminate_pdfium_worker_group(process: subprocess.Popen[str]) -> None:
    """Terminate and reap the worker's complete process group after any exit."""
    if platform.system() != "Windows":
        try:
            os.killpg(process.pid, signal.SIGKILL)
        except ProcessLookupError:
            pass
    elif process.poll() is None:
        process.kill()
    process.communicate()


def rasterize_pdf(
    pdf: Path,
    output_dir: Path,
    identity: Mapping[str, Any],
    *,
    first_page_only: bool = False,
    dpi: int = 130,
    timeout_seconds: float = PDFIUM_WORKER_TIMEOUT_SECONDS,
    require_promoted_runtime: bool = True,
    allow_pending_activation: bool = False,
) -> list[Path]:
    """Render through one killable worker bound to the current immutable release."""
    output_dir = output_dir.resolve()
    _clear_page_output(output_dir)
    release_root = Path(__file__).resolve().parents[1]
    governed = _one_pdfium_renderer(
        [identity],
        skill_root=release_root,
        require_promoted_runtime=require_promoted_runtime,
        allow_pending_activation=allow_pending_activation,
    )
    if not governed:
        raise RuntimeError(
            "PDF page rendering requires the manifest-verified release-owned runtime."
        )
    timeout = min(PDFIUM_WORKER_TIMEOUT_SECONDS, float(timeout_seconds))
    if timeout <= 0:
        raise RuntimeError("renderer.pdfium_worker_timeout: No worker time remained.")
    with tempfile.TemporaryDirectory(prefix=".pdfium-worker-", dir=output_dir.parent) as temporary:
        temporary_root = Path(temporary)
        worker_output = temporary_root / "pages"
        worker_output.mkdir()
        request_path = temporary_root / "request.json"
        request_path.write_text(json.dumps({
            "pdf": str(pdf.resolve()),
            "pdf_sha256": sha256_file(pdf),
            "output_dir": str(worker_output),
            "identity": governed[0],
            "first_page_only": first_page_only,
            "dpi": dpi,
            "require_promoted_runtime": require_promoted_runtime,
            "allow_pending_activation": allow_pending_activation,
        }), encoding="utf-8")
        process = subprocess.Popen(
            _pdfium_worker_command(request_path),
            cwd=release_root,
            text=True,
            stdout=subprocess.PIPE,
            stderr=subprocess.PIPE,
            start_new_session=platform.system() != "Windows",
        )
        try:
            stdout, stderr = process.communicate(timeout=timeout)
        except subprocess.TimeoutExpired as exc:
            _terminate_pdfium_worker_group(process)
            _clear_page_output(output_dir)
            raise RuntimeError(
                f"renderer.pdfium_worker_timeout: PDFium exceeded {timeout:g} seconds."
            ) from exc
        _terminate_pdfium_worker_group(process)
        try:
            result = json.loads(stdout)
        except json.JSONDecodeError as exc:
            _clear_page_output(output_dir)
            raise RuntimeError(
                "renderer.pdfium_worker_crashed: "
                f"PDFium worker exited {process.returncode}: {stderr.strip() or stdout.strip()}"
            ) from exc
        if process.returncode or result.get("status") != "passed":
            _clear_page_output(output_dir)
            finding = next(iter(result.get("findings") or []), {})
            code = str(finding.get("code") or "renderer.pdfium_worker_crashed")
            issue = str(finding.get("issue") or stderr.strip() or "PDFium worker failed.")
            raise RuntimeError(f"{code}: {issue}")
        required_limits = set(_required_pdfium_worker_resource_limits())
        if not required_limits.issubset(set(result.get("resource_limits") or [])):
            _clear_page_output(output_dir)
            raise RuntimeError(
                "renderer.pdfium_resource_limits_unavailable: "
                "The worker could not apply required host process limits."
            )
        evidence = result.get("pages")
        actual = [
            {"name": page.name, "bytes": page.stat().st_size, "sha256": sha256_file(page)}
            for page in _ordered_page_images(worker_output)
            if page.is_file() and not page.is_symlink()
        ]
        if not isinstance(evidence, list) or actual != evidence or not actual:
            _clear_page_output(output_dir)
            raise RuntimeError(
                "renderer.pdfium_worker_output_invalid: Worker page evidence did not match output bytes."
            )
        pages = []
        for item in actual:
            source = worker_output / item["name"]
            target = output_dir / item["name"]
            os.replace(source, target)
            pages.append(target)
        return pages


def renderers(
    *,
    environment: Mapping[str, str] | None = None,
    skill_root: Path | None = None,
    deadline_monotonic: float | None = None,
    clock: Any = time.monotonic,
) -> list[dict[str, Any]]:
    """List supported host DOCX renderers in governed fidelity order."""
    search_path = (environment or {}).get("PATH") if environment is not None else None
    system = platform.system()
    identities: list[dict[str, Any]] = []
    yielded: set[str] = set()

    def append(identity: dict[str, Any]) -> None:
        key = str(identity.get("path", "")).casefold()
        if key and key not in yielded:
            yielded.add(key)
            identities.append(identity)

    def probe_timeout(cap: float = 20.0) -> float | None:
        if deadline_monotonic is None:
            return cap
        remaining = deadline_monotonic - clock()
        return min(cap, remaining) if remaining > 0 else None

    if system == "Windows":
        try:
            if importlib.util.find_spec("win32com.client") is not None:
                append({"kind": "Microsoft Word", "path": "COM:Word.Application", "version": "installed COM application", "platform": "Windows", "source": "host application"})
        except (ImportError, ValueError):
            pass
    if system == "Darwin" and Path("/Applications/Microsoft Word.app").exists():
        append({"kind": "Microsoft Word", "path": "/Applications/Microsoft Word.app", "version": "installed macOS application", "platform": "Darwin", "source": "host application"})
    office_candidates: list[tuple[Path, str]] = []
    if skill_root is not None:
        assurance_path = skill_root / "INSTALLATION-ASSURANCE.json"
        try:
            if assurance_path.is_file() and not assurance_path.is_symlink():
                installation = json.loads(assurance_path.read_text(encoding="utf-8"))
                active_smoke = installation.get("active_path_smoke") or {}
                smoke_assurance = active_smoke.get("assurance") or {}
                smoke_renderer = smoke_assurance.get("renderer") or {}
                renderer_path = Path(str(smoke_renderer.get("path") or ""))
                if (
                    installation.get("status") == "passed"
                    and active_smoke.get("status") == "passed"
                    and active_smoke.get("active_root") == str(skill_root.resolve())
                    and smoke_assurance.get("status") == "passed"
                    and smoke_renderer.get("kind") == "LibreOffice"
                    and renderer_path.is_absolute()
                    and not renderer_path.is_symlink()
                ):
                    office_candidates.append((renderer_path, "active installation smoke"))
        except (OSError, ValueError, AttributeError, TypeError, json.JSONDecodeError):
            pass
        if system == "Linux" and skill_root.parent.name == "skills":
            profile_root = skill_root.parent.parent.resolve()
            for name in ("libreoffice", "soffice"):
                candidate = profile_root / "clinical-office-runtime/bin" / name
                try:
                    resolved = candidate.resolve(strict=True)
                    if resolved.is_relative_to(profile_root) and resolved.is_file():
                        office_candidates.append((resolved, "Hermes profile office runtime"))
                except (OSError, ValueError):
                    pass
    office_candidates.extend(_executable_candidates(("libreoffice", "soffice"), environment=environment))
    if system == "Darwin":
        office_candidates.append((Path("/Applications/LibreOffice.app/Contents/MacOS/soffice"), "macOS application"))
    for candidate, source in office_candidates:
        if candidate.is_file() and os.access(candidate, os.X_OK):
            path = str(candidate)
            timeout = probe_timeout()
            if timeout is None:
                break
            try:
                result = subprocess.run([path, "--version"], text=True, capture_output=True, timeout=timeout, env=dict(environment) if environment else None)
            except (OSError, subprocess.TimeoutExpired):
                continue
            if result.returncode == 0:
                append({"kind": "LibreOffice", "path": path, "version": result.stdout.strip(), "platform": system, "source": source})
    return identities


def renderer(
    *,
    environment: Mapping[str, str] | None = None,
    skill_root: Path | None = None,
) -> dict[str, Any] | None:
    """Return the preferred supported host office renderer."""
    identities = renderers(environment=environment, skill_root=skill_root)
    return identities[0] if identities else None


def _template_fonts(path: Path) -> set[str]:
    """Return font families that can render visible text in the DOCX stories.

    Word stores fonts for dormant styles and alternate writing systems alongside
    the fonts used by visible runs.  Requiring every declaration makes renderer
    preflight depend on fonts that cannot affect an English document.  Resolve
    each visible run through its run, character-style, paragraph-style, and
    document-default hierarchy instead.
    """
    namespace = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
    ns = {"w": namespace}
    attr = lambda name: f"{{{namespace}}}{name}"
    fonts: set[str] = set()

    def font_values(element: ET._Element | None) -> dict[str, str]:
        if element is None:
            return {}
        node = element.find("w:rPr/w:rFonts", ns) if element.tag != attr("rPr") else element.find("w:rFonts", ns)
        if node is None:
            return {}
        return {
            key: value.strip()
            for key in ("ascii", "hAnsi", "eastAsia", "cs")
            if (value := node.get(attr(key))) and value.strip()
        }

    def scripts(text: str) -> set[str]:
        required: set[str] = set()
        for character in text:
            codepoint = ord(character)
            if character.isspace() or character.isascii():
                required.add("ascii")
            elif (
                0x3040 <= codepoint <= 0x30FF
                or 0x3400 <= codepoint <= 0x9FFF
                or 0xAC00 <= codepoint <= 0xD7AF
                or 0xF900 <= codepoint <= 0xFAFF
            ):
                required.add("eastAsia")
            elif (
                0x0590 <= codepoint <= 0x08FF
                or 0x0900 <= codepoint <= 0x0DFF
                or 0xFB1D <= codepoint <= 0xFEFC
            ):
                required.add("cs")
            else:
                required.add("hAnsi")
        return required or {"ascii"}

    with zipfile.ZipFile(path) as package:
        styles: dict[str, tuple[str | None, dict[str, str]]] = {}
        defaults: dict[str, str] = {}
        if "word/styles.xml" in package.namelist():
            styles_root = ET.fromstring(package.read("word/styles.xml"))
            defaults = font_values(styles_root.find("w:docDefaults/w:rPrDefault/w:rPr", ns))
            for style in styles_root.findall("w:style", ns):
                style_id = style.get(attr("styleId"))
                if not style_id:
                    continue
                based_on = style.find("w:basedOn", ns)
                styles[style_id] = (
                    based_on.get(attr("val")) if based_on is not None else None,
                    font_values(style),
                )

        def style_fonts(style_id: str | None) -> dict[str, str]:
            resolved: dict[str, str] = {}
            seen: set[str] = set()
            while style_id and style_id not in seen:
                seen.add(style_id)
                based_on, declared = styles.get(style_id, (None, {}))
                for key, value in declared.items():
                    resolved.setdefault(key, value)
                style_id = based_on
            return resolved

        story_names = [
            name for name in package.namelist()
            if name == "word/document.xml"
            or re.fullmatch(r"word/(?:header|footer)\d+\.xml", name)
            or name in {"word/footnotes.xml", "word/endnotes.xml", "word/comments.xml"}
        ]
        for name in story_names:
            root = ET.fromstring(package.read(name))
            for run in root.iter(attr("r")):
                text = "".join(node.text or "" for node in run.findall("w:t", ns))
                symbols = run.findall("w:sym", ns)
                if not text.strip() and not symbols:
                    continue
                for symbol in symbols:
                    if value := symbol.get(attr("font")):
                        fonts.add(value.strip())

                direct = font_values(run.find("w:rPr", ns))
                run_style = run.find("w:rPr/w:rStyle", ns)
                character = style_fonts(run_style.get(attr("val")) if run_style is not None else None)
                paragraph = run.getparent()
                while paragraph is not None and paragraph.tag != attr("p"):
                    paragraph = paragraph.getparent()
                paragraph_style = paragraph.find("w:pPr/w:pStyle", ns) if paragraph is not None else None
                inherited = style_fonts(paragraph_style.get(attr("val")) if paragraph_style is not None else None)
                sources = (direct, character, inherited, defaults)
                for script in scripts(text):
                    fallbacks = {
                        "ascii": ("ascii", "hAnsi"),
                        "hAnsi": ("hAnsi", "ascii"),
                        "eastAsia": ("eastAsia", "hAnsi", "ascii"),
                        "cs": ("cs", "hAnsi", "ascii"),
                    }[script]
                    selected = next(
                        (source[key] for source in sources for key in fallbacks if key in source),
                        None,
                    )
                    if selected:
                        fonts.add(selected)
    return fonts


def _font_probe(
    font: str,
    *,
    environment: Mapping[str, str] | None = None,
    timeout_seconds: float = 10.0,
) -> tuple[bool | None, str]:
    global _MAC_FONT_NAMES, _WINDOWS_FONT_NAMES
    executable = shutil.which("fc-match", path=(environment or {}).get("PATH"))
    if not executable and platform.system() == "Darwin":
        try:
            if _MAC_FONT_NAMES is None:
                result = subprocess.run(["/usr/sbin/system_profiler", "SPFontsDataType", "-json"], text=True, capture_output=True, timeout=max(0.001, min(10.0, timeout_seconds)))
                if result.returncode != 0:
                    raise RuntimeError(result.stderr.strip() or "system_profiler failed")
                entries = json.loads(result.stdout).get("SPFontsDataType", [])
                _MAC_FONT_NAMES = set()
                for item in entries:
                    if not isinstance(item, Mapping):
                        continue
                    _MAC_FONT_NAMES.add(str(item.get("_name", "")).casefold())
                    for face in item.get("typefaces", []):
                        if not isinstance(face, Mapping):
                            continue
                        for key in ("_name", "family", "fullname"):
                            if value := str(face.get(key, "")).strip():
                                _MAC_FONT_NAMES.add(value.casefold())
            normalized_font = font.casefold()
            available = any(
                name.removesuffix(".ttf").removesuffix(".otf").removesuffix(".ttc").removesuffix(" bold").removesuffix(" italic").strip() == normalized_font
                for name in _MAC_FONT_NAMES
            )
            return available, "macOS system font inventory"
        except (OSError, RuntimeError, subprocess.TimeoutExpired, ValueError, json.JSONDecodeError):
            pass
    if not executable and platform.system() == "Windows":
        try:
            if _WINDOWS_FONT_NAMES is None:
                import winreg
                _WINDOWS_FONT_NAMES = set()
                key_path = r"SOFTWARE\Microsoft\Windows NT\CurrentVersion\Fonts"
                for hive in (winreg.HKEY_LOCAL_MACHINE, winreg.HKEY_CURRENT_USER):
                    try:
                        with winreg.OpenKey(hive, key_path) as key:
                            for index in range(winreg.QueryInfoKey(key)[1]):
                                name, filename, _kind = winreg.EnumValue(key, index)
                                _WINDOWS_FONT_NAMES.add(re.sub(r"\s*\([^)]*\)\s*$", "", str(name)).casefold())
                                _WINDOWS_FONT_NAMES.add(Path(str(filename)).stem.casefold())
                    except OSError:
                        continue
            normalized_font = font.casefold()
            return normalized_font in _WINDOWS_FONT_NAMES, "Windows system font inventory"
        except (ImportError, OSError):
            pass
    if not executable:
        return None, "font inventory is unavailable; availability will be decided by render evidence."
    try:
        result = subprocess.run(
            [executable, "-f", "%{family}", font],
            text=True,
            capture_output=True,
            timeout=max(0.001, min(5.0, timeout_seconds)),
            env=dict(environment) if environment else None,
        )
    except (OSError, subprocess.TimeoutExpired) as exc:
        return None, f"font probe unavailable: {exc}"
    families = {part.strip().casefold() for part in result.stdout.split(",") if part.strip()}
    return result.returncode == 0 and font.casefold() in families, result.stdout.strip()


def _bundled_font_path(
    repo_root: Path,
    font: str,
    approved_font_plan: Mapping[str, Any] | None = None,
) -> Path | None:
    families = (
        approved_font_plan.get("packaged_families", {})
        if isinstance(approved_font_plan, Mapping)
        else BUNDLED_FONT_FILES
    )
    filename = families.get(font)
    if not filename:
        return None
    assets = (
        approved_font_plan.get("packaged_font_assets", {})
        if isinstance(approved_font_plan, Mapping)
        else {}
    )
    relative = next(
        (str(path) for path in assets if Path(str(path)).name == filename),
        f"assets/fallback-fonts/{filename}",
    )
    path = repo_root / relative
    return path if path.is_file() else None


def _runtime_environment(
    repo_root: Path,
    environment: Mapping[str, str] | None,
    contracted_bundle: Mapping[str, Any] | None = None,
) -> dict[str, str]:
    runtime = dict(environment or os.environ)
    approved_font_plan = contracted_bundle.get("approved_font_plan", {}) if isinstance(contracted_bundle, Mapping) else {}
    asset_paths = approved_font_plan.get("packaged_font_assets", {}) if isinstance(approved_font_plan, Mapping) else {}
    font_dirs = sorted({(repo_root / str(relative)).parent for relative in asset_paths})
    if not font_dirs:
        font_dirs = [repo_root / "assets/fallback-fonts"]
    available_dirs = [path for path in font_dirs if path.is_dir()]
    if available_dirs:
        existing = runtime.get("SAL_FONTPATH", "")
        runtime["SAL_FONTPATH"] = os.pathsep.join([
            *(str(path) for path in available_dirs),
            *([existing] if existing else []),
        ])
    return runtime


def _font_fallback_candidates(font: str) -> tuple[str, ...]:
    """Return visually compatible cross-platform fallbacks in preference order."""
    normalized = font.casefold()
    if "symbol" in normalized or "dingbat" in normalized or "wingding" in normalized:
        candidates = SYMBOL_FONT_FALLBACKS
    elif any(marker in normalized for marker in ("mono", "courier", "consolas", "menlo", "code")):
        candidates = MONOSPACE_FONT_FALLBACKS
    elif any(marker in normalized for marker in ("serif", "times", "georgia", "cambria", "garamond", "minion")):
        candidates = SERIF_FONT_FALLBACKS
    else:
        candidates = SANS_FONT_FALLBACKS
    seen = {normalized}
    ordered = []
    for candidate in candidates:
        key = candidate.casefold()
        if key in seen:
            continue
        seen.add(key)
        ordered.append(candidate)
    return tuple(ordered)


def _approved_packaged_font_fallback(
    font: str,
    approved_font_plan: Mapping[str, Any] | None = None,
) -> str:
    """Map a missing template font to one audited release-owned substitute."""
    normalized = font.casefold().strip()
    approved = (
        approved_font_plan.get("approved_fallbacks", {})
        if isinstance(approved_font_plan, Mapping)
        else APPROVED_PACKAGED_FONT_FALLBACKS
    )
    families = (
        approved_font_plan.get("packaged_families", {})
        if isinstance(approved_font_plan, Mapping)
        else BUNDLED_FONT_FILES
    )
    if normalized in approved:
        return str(approved[normalized])
    if any(marker in normalized for marker in ("mono", "courier", "consolas", "menlo", "code")):
        return "Liberation Mono" if "Liberation Mono" in families else next(iter(families))
    if any(marker in normalized for marker in ("serif", "times", "georgia", "cambria", "garamond", "minion")):
        return "Liberation Serif" if "Liberation Serif" in families else next(iter(families))
    return "Liberation Sans" if "Liberation Sans" in families else next(iter(families))


def preflight(
    repo_root: Path,
    reference: Mapping[str, Any],
    *,
    contracted_bundle: Mapping[str, Any] | None = None,
    deadline_seconds: float = 30.0,
    environment: Mapping[str, str] | None = None,
    renderer_identities: Iterable[Mapping[str, Any]] | None = None,
    page_renderer_identities: Iterable[Mapping[str, Any]] | None = None,
    clock: Any = time.monotonic,
) -> dict[str, Any]:
    """Resolve and smoke-test Render Assurance capabilities.

    Inventory uncertainty is evidence to test, not evidence of absence. The
    function only blocks after every local renderer/page-renderer combination
    has failed its disposable render.
    """
    started = clock()
    deadline = started + max(0.0, deadline_seconds)

    def remaining() -> float:
        return max(0.0, deadline - clock())

    bundle = contracted_bundle or contracted_template_bundle(repo_root, reference)
    approved_font_plan = bundle["approved_font_plan"]
    runtime_environment = _runtime_environment(repo_root, environment, bundle)
    renderer_candidates = (
        [dict(item) for item in renderer_identities]
        if renderer_identities is not None
        else renderers(
            environment=runtime_environment,
            skill_root=repo_root,
            deadline_monotonic=deadline,
            clock=clock,
        )
    )
    page_candidates = _one_pdfium_renderer(
        [dict(item) for item in page_renderer_identities]
        if page_renderer_identities is not None
        else page_renderers(environment=runtime_environment, skill_root=repo_root) if remaining() > 0 else [],
        skill_root=repo_root,
    )
    identity = None
    page_identity = None
    renderer_attempts: list[dict[str, Any]] = []
    page_attempts: list[dict[str, Any]] = []
    findings: list[dict[str, Any]] = []
    protocol_template, icf_template = template_paths(repo_root, reference, contracted_bundle=bundle)
    templates = [path for path in (protocol_template, icf_template) if path is not None]
    required_fonts = sorted({font for path in templates for font in _template_fonts(path)})
    font_results: dict[str, Any] = {}
    font_substitutions: dict[str, str] = {}
    bundled_substitution = False
    for font in required_fonts:
        probe_remaining = remaining()
        if probe_remaining <= 0:
            font_results[font] = {"available": None, "match": "Render Assurance deadline expired before font inventory."}
            continue
        available, detail = _font_probe(font, environment=runtime_environment, timeout_seconds=probe_remaining)
        font_results[font] = {"available": available, "match": detail}
        if available is not False:
            continue
        candidate = _approved_packaged_font_fallback(font, approved_font_plan)
        bundled = _bundled_font_path(repo_root, candidate, approved_font_plan)
        if bundled is not None:
            font_substitutions[font] = candidate
            font_results[font].update({
                "substitute": candidate,
                "substitute_match": f"bundled approved compatible font: {bundled.relative_to(repo_root)}",
            })
            bundled_substitution = True
        else:
            font_results[font]["attempted_fallbacks"] = [candidate]
    smoke: dict[str, Any] = {"status": "not_run"}
    if renderer_candidates and page_candidates and remaining() > 0:
        with tempfile.TemporaryDirectory(prefix="hermes-renderer-preflight-") as directory:
            root = Path(directory)
            source = root / "preflight.docx"
            document = Document()
            document.add_paragraph("Hermes renderer preflight")
            for font in required_fonts:
                selected = font_substitutions.get(font, font)
                run = document.add_paragraph().add_run(f"{selected}: Aa 123 •")
                run.font.name = selected
                fonts = run._element.get_or_add_rPr().get_or_add_rFonts()
                for attribute in ("ascii", "hAnsi", "eastAsia", "cs"):
                    fonts.set(qn(f"w:{attribute}"), selected)
            document.save(source)
            page_renderer_failed = False
            for candidate_renderer in renderer_candidates:
                attempt_remaining = remaining()
                if attempt_remaining <= 0:
                    break
                if bundled_substitution and candidate_renderer.get("kind") != "LibreOffice":
                    renderer_attempts.append({
                        "renderer": candidate_renderer,
                        "status": "skipped",
                        "issue": "Bundled fallback fonts are isolated to the verified LibreOffice runtime.",
                    })
                    continue
                try:
                    pdf = _render_pdf(
                        source,
                        root,
                        candidate_renderer,
                        environment=runtime_environment,
                        timeout_seconds=max(0.001, attempt_remaining),
                    )
                except (OSError, RuntimeError, subprocess.TimeoutExpired) as exc:
                    renderer_attempts.append({"renderer": candidate_renderer, "status": "failed", "issue": str(exc)})
                    continue
                page_dir = root / "pages"
                for candidate_page_renderer in page_candidates:
                    attempt_remaining = remaining()
                    if attempt_remaining <= 0:
                        break
                    try:
                        rasterize_pdf(
                            pdf,
                            page_dir,
                            candidate_page_renderer,
                            first_page_only=True,
                        )
                    except (OSError, RuntimeError, subprocess.TimeoutExpired) as exc:
                        page_attempts.append({"renderer": candidate_page_renderer, "status": "failed", "issue": str(exc)})
                        page_renderer_failed = True
                        break
                    identity = candidate_renderer
                    page_identity = candidate_page_renderer
                    renderer_attempts.append({"renderer": candidate_renderer, "status": "passed"})
                    page_attempts.append({"renderer": candidate_page_renderer, "status": "passed"})
                    smoke = {"status": "passed", "pdf": "disposable/preflight.pdf", "page_image": "disposable/preflight.png", "pages": len(PdfReader(pdf).pages)}
                    break
                if identity is not None:
                    break
                if page_renderer_failed:
                    renderer_attempts.append({
                        "renderer": candidate_renderer,
                        "status": "failed",
                        "issue": "The release-owned pypdfium2 page renderer failed; no office renderer retry is permitted.",
                    })
                    break
                renderer_attempts.append({
                    "renderer": candidate_renderer,
                    "status": "failed",
                    "issue": "Every available page renderer failed for this renderer.",
                })
    if identity is None:
        if not renderer_candidates:
            issue = "No supported Microsoft Word or LibreOffice renderer is available."
            field = "renderer"
        elif not page_candidates:
            issue = "The release-owned pypdfium2 page renderer is unavailable."
            field = "page_renderer"
        else:
            issue = (
                "The release-owned pypdfium2 page renderer failed; Render Assurance stopped exactly."
                if page_attempts else
                "Every supported office renderer failed its smoke render."
            )
            field = "preflight"
        findings.append({"category": "renderer", "field": field, "issue": issue})
        smoke = {"status": "blocked", "issue": issue}
    else:
        for font, result in font_results.items():
            if result.get("available") is None:
                result["resolution"] = "render_verified"
    elapsed = clock() - started
    if elapsed > deadline_seconds and identity is None:
        findings.append({"category": "renderer", "field": "preflight", "issue": f"Renderer preflight exceeded its {deadline_seconds:.1f}s deadline ({elapsed:.3f}s)."})
    ordered_renderers = ([identity] if identity else []) + [candidate for candidate in renderer_candidates if candidate != identity]
    ordered_page_renderers = ([page_identity] if page_identity else []) + [candidate for candidate in page_candidates if candidate != page_identity]
    return {
        "status": "passed" if not findings else "blocked",
        "contracted_template_bundle": dict(bundle),
        "renderer": identity,
        "renderer_candidates": ordered_renderers,
        "renderer_attempts": renderer_attempts,
        "page_renderer": page_identity,
        "page_renderer_candidates": ordered_page_renderers,
        "page_renderer_attempts": page_attempts,
        "required_fonts": required_fonts,
        "fonts": font_results,
        "font_substitutions": font_substitutions,
        "smoke": smoke,
        "deadline_seconds": deadline_seconds,
        "elapsed_seconds": round(elapsed, 3),
        "findings": findings,
    }


def _render_pdf(docx: Path, output_dir: Path, identity: Mapping[str, Any], *, environment: Mapping[str, str] | None = None, timeout_seconds: float = 180.0) -> Path:
    if identity["kind"] == "LibreOffice":
        result = subprocess.run([str(identity["path"]), "--headless", "--convert-to", "pdf", "--outdir", str(output_dir), str(docx)], text=True, capture_output=True, timeout=timeout_seconds, env=dict(environment) if environment else None)
        path = output_dir / f"{docx.stem}.pdf"
        if result.returncode or not path.is_file():
            raise RuntimeError(f"LibreOffice PDF export failed: {result.stderr or result.stdout}")
        return path
    if identity.get("platform") == "Windows":
        return _windows_word_pdf(docx, output_dir, timeout_seconds=timeout_seconds)
    return _mac_pdf(docx, output_dir, identity, timeout_seconds=timeout_seconds)


def _mac_pdf(docx: Path, output_dir: Path, identity: Mapping[str, Any], *, timeout_seconds: float = 180.0) -> Path:
    if identity.get("kind") != "Microsoft Word":
        raise RuntimeError("Only Microsoft Word or LibreOffice may convert DOCX to PDF.")
    app = "Microsoft Word"
    output = output_dir / f"{docx.stem}.pdf"
    script = 'on run argv\nset src to POSIX file (item 1 of argv)\nset dst to POSIX file (item 2 of argv)\ntell application "Microsoft Word"\nset d to open src\nsave as d file name dst file format format PDF\nclose d saving no\nend tell\nend run'
    result = subprocess.run(["osascript", "-e", script, str(docx), str(output)], text=True, capture_output=True, timeout=timeout_seconds)
    if result.returncode or not output.is_file(): raise RuntimeError(f"{app} PDF export failed: {result.stderr or result.stdout}")
    return output


def _windows_word_pdf(docx: Path, output_dir: Path, *, timeout_seconds: float = 180.0) -> Path:
    output = output_dir / f"{docx.stem}.pdf"
    script = (
        "import sys\nimport win32com.client\n"
        "word=win32com.client.DispatchEx('Word.Application'); word.Visible=False; document=None\n"
        "try:\n document=word.Documents.Open(sys.argv[1], ReadOnly=True); document.ExportAsFixedFormat(sys.argv[2], 17)\n"
        "finally:\n document.Close(False) if document is not None else None; word.Quit()\n"
    )
    result = subprocess.run(
        [sys.executable, "-c", script, str(docx.resolve()), str(output.resolve())],
        text=True,
        capture_output=True,
        timeout=timeout_seconds,
    )
    if result.returncode:
        raise RuntimeError(f"Microsoft Word PDF export failed: {result.stderr or result.stdout}")
    if not output.is_file(): raise RuntimeError("Microsoft Word PDF export did not produce a file.")
    return output


def _protocol_section15_opening_findings(docx_path: Path, pdf_path: Path) -> list[dict[str, Any]]:
    """Catch a split Section 15 opening in the exact rendered Protocol PDF."""
    heading = "15. STANDARD EVALUATION PROCEDURES"
    caption = "Table 15.1. Proposed Visits and Study Assessments"
    paragraphs = {paragraph.text.strip() for paragraph in Document(docx_path).paragraphs}
    if heading not in paragraphs or caption not in paragraphs:
        return []
    try:
        page_lines = [
            [re.sub(r"\s+", " ", line).strip() for line in (page.extract_text() or "").splitlines()]
            for page in PdfReader(pdf_path).pages
        ]
    except (OSError, ValueError, TypeError, AttributeError, KeyError):
        # The mandatory image reviewer still owns pages with unextractable text.
        return []
    heading_pages = [number for number, lines in enumerate(page_lines, 1) if heading in lines]
    caption_pages = [number for number, lines in enumerate(page_lines, 1) if caption in lines]
    if not heading_pages or not caption_pages:
        return []
    heading_page, caption_page = heading_pages[-1], caption_pages[-1]
    table_page = next((
        number for number, lines in enumerate(page_lines, 1)
        if number >= caption_page
        and any(re.match(r"^Activity(?:\s|$)", line) for line in lines)
    ), None)
    if heading_page == caption_page and table_page in {None, caption_page}:
        return []
    return [{
        "category": "visual",
        "field": "protocol",
        "artifact": "protocol",
        "page": heading_page,
        "check": "artificial_pagination",
        "element": heading,
        "target_ids": ["layout:protocol"],
        "issue": (
            "Section 15 heading, introduction, caption, and assessment-table opening "
            "must appear together on one rendered page."
        ),
    }]


def audit_final_toc_destinations(docx_path: Path, pdf_path: Path) -> dict[str, Any]:
    """Read final bytes independently of the renderer's TOC mapping algorithm.

    A bounded refresh loop is not convergence evidence. Only observed destination
    disagreements block here; unextractable/ambiguous text remains explicit
    unknown evidence for the mandatory exact-artifact visual review.
    """
    document = Document(docx_path)
    headings = [p.text.strip() for p in document.paragraphs
                if p.style.name in {"Heading 1", "Heading 2"} and p.text.strip()]
    if not any("TABLE OF CONTENTS" in heading.upper() for heading in headings):
        return {"status": "not_applicable", "destinations": [], "findings": []}

    def key(text: str) -> str:
        return " ".join(re.findall(r"\w+", text.casefold()))

    cache: dict[str, list[int]] = {}
    # Read cached field results, including hyperlink/SDT descendants. Do not
    # depend on Paragraph.text's treatment of hyperlink text in python-docx.
    for paragraph in document.element.body.iter(qn("w:p")):
        styles = paragraph.xpath("./w:pPr/w:pStyle")
        if not styles or not re.fullmatch(r"toc[12]", styles[0].get(qn("w:val"), "").replace(" ", ""), re.I):
            continue
        text = "".join("\t" if node.tag == qn("w:tab") else (node.text or "")
                       for node in paragraph.iter() if node.tag in {qn("w:t"), qn("w:tab")})
        match = re.fullmatch(r"(.+?)[\s.]+(\d+)\s*", text)
        if match:
            cache.setdefault(key(match[1]), []).append(int(match[2]))
    try:
        page_lines = [[key(line) for line in (page.extract_text() or "").splitlines() if key(line)]
                      for page in PdfReader(pdf_path).pages]
    except (OSError, ValueError, TypeError, AttributeError, KeyError) as exc:
        return {"status": "unknown", "destinations": [], "findings": [], "issue": str(exc)}
    destinations, findings = [], []
    for heading in headings:
        marker = key(heading)
        matches = []
        for number, lines in enumerate(page_lines, 1):
            # Whole lines distinguish a real heading from a TOC row or prose
            # citation. Permit wrapped headings, never arbitrary substrings.
            if any(" ".join(lines[start:start + size]) == marker
                   for start in range(len(lines)) for size in range(1, 5)):
                matches.append(number)
        cached = cache.get(marker, [])
        rendered = matches[0] if len(matches) == 1 else None
        toc_values = []
        in_toc = False
        body_markers = {key(value) for value in headings if "TABLE OF CONTENTS" not in value.upper()}
        for lines in page_lines:
            for line in lines:
                if line == key("TABLE OF CONTENTS"):
                    in_toc = True
                elif line in body_markers:
                    in_toc = False
                if in_toc:
                    entry = re.fullmatch(re.escape(marker) + r" (\d+)", line)
                    if entry:
                        toc_values.append(int(entry[1]))
        # The TOC title itself is a destination, not necessarily a displayed
        # TOC row. All other destinations require independent visible evidence.
        title_only = "TABLE OF CONTENTS" in heading.upper() and not toc_values
        observed = [*cached, *toc_values, *([rendered] if rendered is not None else [])]
        status = "mismatch" if len(set(observed)) > 1 else (
            "passed" if rendered is not None and cached == [rendered]
            and (toc_values == [rendered] or title_only) else "unknown")
        row = {"heading": heading, "cached_page": cached[0] if len(cached) == 1 else None,
               "rendered_page": rendered, "rendered_candidates": matches,
               "rendered_toc_page": toc_values[0] if len(toc_values) == 1 else None,
               "rendered_toc_candidates": toc_values, "status": status}
        destinations.append(row)
        if status == "mismatch":
            findings.append({"category": "visual", "field": docx_path.stem,
                             "artifact": docx_path.stem, "check": "toc_mismatch",
                             "element": heading, "page": rendered,
                             "target_ids": [f"layout:{docx_path.stem}"],
                             "issue": f"Final TOC destination mismatch for {heading}: cached {cached}, visible TOC {toc_values}, rendered page {rendered}."})
    status = "mismatch" if findings else ("passed" if destinations and all(
        row["status"] == "passed" for row in destinations) else "unknown")
    return {"status": status, "destinations": destinations, "findings": findings}


def _pdf_body_pages(path: Path, *, docx_path: Path | None = None) -> list[dict[str, Any]]:
    """Extract body text using source section margins, never guessed bands.

    Without source geometry all extracted text is conservatively retained. For
    render_pages' exact artifact layout, resolve the matching candidate DOCX.
    """
    if docx_path is None and path.parent.name == "rendered":
        candidate = path.parent.parent / "candidate" / f"{path.stem}.docx"
        if candidate.is_file():
            docx_path = candidate
    sections = list(Document(docx_path).sections) if docx_path is not None else []
    margins_known = bool(sections) and all(section.top_margin is not None and section.bottom_margin is not None for section in sections)
    # Union of section body areas: page-to-section assignment is not guessed.
    top_margin = min(section.top_margin.pt for section in sections) if margins_known else 0.0
    bottom_margin = min(section.bottom_margin.pt for section in sections) if margins_known else 0.0
    result = []
    for number, page in enumerate(PdfReader(path).pages, 1):
        fragments: list[tuple[str, float, float]] = []
        try:
            bottom, top = float(page.mediabox.bottom) + bottom_margin, float(page.mediabox.top) - top_margin

            def visitor(text: str, cm: Any, tm: Any, font: Any, size: float) -> None:
                y = float(cm[1] * tm[4] + cm[3] * tm[5] + cm[5])
                if text.strip() and bottom <= y <= top:
                    fragments.append((text.strip(), y, abs(float(size))))

            graphics_in_body = False

            def operand(operator: bytes, operands: Any, cm: Any, tm: Any) -> None:
                nonlocal graphics_in_body
                if operator not in {b"Do", b"INLINE IMAGE"}:
                    return
                if operator == b"Do":
                    objects = page.get("/Resources", {}).get("/XObject", {})
                    image = objects.get(operands[0])
                    if image is None or image.get_object().get("/Subtype") != "/Image":
                        # Nested forms need their own coordinate evidence. Unknown
                        # graphical content cannot prove a body-empty page.
                        graphics_in_body = True
                        return
                # pypdf exposes BI/ID/EI as INLINE IMAGE. Both image forms
                # paint a unit square transformed by the current matrix.
                ys = [float(cm[5]), float(cm[1] + cm[5]),
                      float(cm[3] + cm[5]), float(cm[1] + cm[3] + cm[5])]
                if max(ys) >= bottom and min(ys) <= top:
                    graphics_in_body = True

            page.extract_text(visitor_text=visitor, visitor_operand_before=operand)
            text = "\n".join(fragment[0] for fragment in fragments)
            span = (max(y + size for _, y, size in fragments) - min(y for _, y, _ in fragments)) if fragments else 0.0
            result.append({"page": number, "text": text, "geometry_known": True,
                           "body_height_fraction": min(1.0, span / max(1.0, top - bottom)),
                           "body_empty": not text.strip() and not graphics_in_body})
        except (TypeError, AttributeError, KeyError):
            # Compatibility with text-only adapters; no word-count heuristic.
            text = page.extract_text() or ""
            lines = [line for line in text.splitlines()
                     if not re.search(r"\bpage\s+\d+\s+of\s+\d+\b", line, re.I)
                     and not re.fullmatch(r"(?:v(?:ersion)?\s*)?\S*\s*\d{1,2}\s+[A-Za-z]{3,9}\s+\d{4}", line.strip(), re.I)]
            result.append({"page": number, "text": "\n".join(lines), "geometry_known": False,
                           "body_height_fraction": None, "body_empty": not text.strip()})
    return result


def _blank_pdf_pages(path: Path, *, docx_path: Path | None = None) -> list[int]:
    return [item["page"] for item in _pdf_body_pages(path, docx_path=docx_path) if item["body_empty"]]


def _sparse_pdf_pages(docx_path: Path, pdf_path: Path) -> list[dict[str, Any]]:
    """Nonblocking visual context: short body spans and source hard-break links."""
    endings = []
    previous = ""
    for paragraph in Document(docx_path).paragraphs:
        before = paragraph.paragraph_format.page_break_before
        breaks = paragraph._p.xpath('.//w:br[@w:type="page"]')
        sections = paragraph._p.xpath('./w:pPr/w:sectPr')
        section_break = any(not section.xpath('./w:type[@w:val="continuous"]') for section in sections)
        if before and previous:
            endings.append(previous)
        if (breaks or section_break) and (paragraph.text.strip() or previous):
            endings.append(paragraph.text.strip() or previous)
        if paragraph.text.strip():
            previous = paragraph.text.strip()
    normalized_endings = [" ".join(text.casefold().split()) for text in endings]
    rows = []
    for item in _pdf_body_pages(pdf_path, docx_path=docx_path):
        if not item["geometry_known"] or not item["text"] or item["body_height_fraction"] > 0.2:
            continue
        text = " ".join(item["text"].casefold().split())
        rows.append({"page": item["page"], "blocking": False,
                     "body_height_fraction": item["body_height_fraction"],
                     "text": item["text"],
                     "forced_break_after": any(text.endswith(ending) for ending in normalized_endings),
                     "issue": "Small occupied body span; inspect context and any forced break, not a universal sparse-page failure."})
    return rows



def _retained_render_artifacts(
    revision_dir: Path, report: Mapping[str, Any], docx_paths: list[Path],
    offices: list[dict[str, Any]], pages: list[dict[str, Any]],
    fonts: Mapping[str, Any], substitutions: Mapping[str, str],
) -> list[dict[str, Any]]:
    """Rehash complete render evidence under currently available renderer identities."""
    retained = []
    if report.get("status") != "passed":
        return retained
    for docx in docx_paths:
        rows = [row for row in report.get("artifacts", [])
                if isinstance(row, Mapping) and row.get("artifact") == docx.stem]
        if len(rows) != 1:
            continue
        row = rows[0]
        try:
            if (row.get("status") != "passed" or row.get("renderer") not in offices
                    or row.get("page_renderer") not in pages
                    or row.get("font_evidence") != dict(fonts)
                    or row.get("font_substitutions") != dict(substitutions)
                    or row.get("docx") != docx.relative_to(revision_dir).as_posix()
                    or not row.get("pages")
                    or row.get("page_count") != len(row["pages"])
                    or [page.get("page") for page in row["pages"]] != list(range(1, len(row["pages"]) + 1))):
                continue
            files = [(row[key], row[f"{key}_sha256"]) for key in ("docx", "pdf")]
            files.extend((page["path"], page["sha256"]) for page in row["pages"])
            if all(_revision_owned_nonsymlink_path(revision_dir, revision_dir / path)
                   and sha256_file(revision_dir / path) == digest for path, digest in files):
                retained.append(dict(row))
        except (OSError, ValueError, KeyError, TypeError):
            continue
    return retained

def render_pages(
    revision_dir: Path,
    *,
    artifact_names: Iterable[str] | None = None,
    reuse_render_report: Mapping[str, Any] | None = None,
    contracted_bundle: Mapping[str, Any] | None = None,
    renderer_identity: Mapping[str, Any] | None = None,
    page_renderer_identity: Mapping[str, Any] | None = None,
    renderer_identities: Iterable[Mapping[str, Any]] | None = None,
    page_renderer_identities: Iterable[Mapping[str, Any]] | None = None,
    deadline_monotonic: float | None = None,
    clock: Any = time.monotonic,
    font_evidence: Mapping[str, Any] | None = None,
    font_substitutions: Mapping[str, str] | None = None,
    office_exporter: Any = None,
    page_exporter: Any = None,
    blank_page_detector: Any = None,
    require_promoted_runtime: bool = True,
    allow_pending_activation: bool = False,
) -> dict[str, Any]:
    repo_root = Path(__file__).resolve().parents[1]
    runtime_environment = _runtime_environment(repo_root, None, contracted_bundle)
    export_docx = office_exporter or _render_pdf
    export_pages = page_exporter or rasterize_pdf
    find_blank_pages = blank_page_detector or _blank_pdf_pages
    render_root = revision_dir / "rendered"
    selected_artifacts = set(artifact_names or ())
    docx_paths = [
        path
        for path in sorted((revision_dir / "candidate").glob("*.docx"))
        if not selected_artifacts or path.stem in selected_artifacts
    ]
    renderer_candidates = [dict(item) for item in (renderer_identities or [])]
    if renderer_identity is not None:
        renderer_candidates = [dict(renderer_identity), *[item for item in renderer_candidates if dict(item) != dict(renderer_identity)]]
    if not renderer_candidates and renderer_identities is None and renderer_identity is None:
        renderer_candidates = renderers(environment=runtime_environment, skill_root=repo_root)
    page_candidates = [dict(item) for item in (page_renderer_identities or [])]
    if page_renderer_identity is not None:
        page_candidates = [dict(page_renderer_identity), *[item for item in page_candidates if dict(item) != dict(page_renderer_identity)]]
    if not page_candidates and page_renderer_identities is None and page_renderer_identity is None:
        page_candidates = page_renderers(
            environment=runtime_environment,
            skill_root=repo_root,
            require_promoted_runtime=require_promoted_runtime,
            allow_pending_activation=allow_pending_activation,
        )
    page_candidates = _one_pdfium_renderer(
        page_candidates,
        skill_root=repo_root,
        require_promoted_runtime=require_promoted_runtime,
        allow_pending_activation=allow_pending_activation,
    )
    if not renderer_candidates:
        return {"status": "blocked", "findings": [{"category": "renderer", "field": "renderer", "issue": "No supported Microsoft Word or LibreOffice renderer is available."}]}
    if not page_candidates:
        if page_renderer_identities is None and page_renderer_identity is None:
            integrity = _pdfium_runtime_integrity(
                repo_root,
                require_promoted_runtime=require_promoted_runtime,
                allow_pending_activation=allow_pending_activation,
            )
            if integrity["status"] == "blocked":
                return {
                    "status": "blocked",
                    "renderer": renderer_candidates[0],
                    "findings": [integrity["finding"]],
                }
        return {"status": "blocked", "renderer": renderer_candidates[0], "findings": [{"category": "renderer", "field": "page_renderer", "issue": "No supported PDF page renderer is available."}]}

    retained = _retained_render_artifacts(
        revision_dir, reuse_render_report or {}, docx_paths, renderer_candidates, page_candidates,
        font_evidence or {}, font_substitutions or {},
    )
    retained_names = {row["artifact"] for row in retained}
    docx_paths = [path for path in docx_paths if path.stem not in retained_names]
    if selected_artifacts or retained:
        for docx in docx_paths:
            (render_root / f"{docx.stem}.pdf").unlink(missing_ok=True)
            shutil.rmtree(render_root / docx.stem, ignore_errors=True)
    elif render_root.exists():
        shutil.rmtree(render_root)
    renderer_attempts: list[dict[str, Any]] = []
    page_attempts: list[dict[str, Any]] = []
    for renderer_index, identity in enumerate(renderer_candidates, 1):
        remaining = 180.0 if deadline_monotonic is None else deadline_monotonic - clock()
        if remaining <= 0:
            break
        attempt_root = revision_dir / ".render-attempts" / f"renderer-{renderer_index}"
        if attempt_root.exists():
            shutil.rmtree(attempt_root)
        attempt_root.mkdir(parents=True)
        pdfs: dict[Path, Path] = {}
        try:
            for docx in docx_paths:
                remaining = 180.0 if deadline_monotonic is None else deadline_monotonic - clock()
                if remaining <= 0:
                    raise subprocess.TimeoutExpired("DOCX rendering", 0)
                pdf = export_docx(docx, attempt_root, identity, environment=runtime_environment, timeout_seconds=min(180.0, remaining))
                for _ in range(3):
                    if not refresh_toc_from_pdf(docx, pdf):
                        break
                    pdf.unlink(missing_ok=True)
                    remaining = 180.0 if deadline_monotonic is None else deadline_monotonic - clock()
                    if remaining <= 0:
                        raise subprocess.TimeoutExpired("DOCX rendering", 0)
                    pdf = export_docx(docx, attempt_root, identity, environment=runtime_environment, timeout_seconds=min(180.0, remaining))
                pdfs[docx] = pdf
        except (OSError, RuntimeError, subprocess.TimeoutExpired) as exc:
            renderer_attempts.append({"renderer": identity, "status": "failed", "issue": str(exc)})
            continue

        selected_page_renderer = None
        selected_pages: dict[Path, list[Path]] = {}
        for candidate in page_candidates:
            current_pages: dict[Path, list[Path]] = {}
            try:
                for docx, pdf in pdfs.items():
                    remaining = 180.0 if deadline_monotonic is None else deadline_monotonic - clock()
                    if remaining <= 0:
                        raise subprocess.TimeoutExpired("PDF page rendering", 0)
                    page_dir = attempt_root / docx.stem / str(candidate["kind"])
                    page_kwargs = {
                        "dpi": 130,
                        "timeout_seconds": min(PDFIUM_WORKER_TIMEOUT_SECONDS, remaining),
                    }
                    if page_exporter is None:
                        page_kwargs["require_promoted_runtime"] = require_promoted_runtime
                        page_kwargs["allow_pending_activation"] = (
                            allow_pending_activation
                        )
                    pages = export_pages(pdf, page_dir, candidate, **page_kwargs)
                    expected = len(PdfReader(pdf).pages)
                    if len(pages) != expected or expected == 0:
                        raise RuntimeError(f"Page rendering failed for {docx.name}: expected {expected}, got {len(pages)}.")
                    current_pages[docx] = pages
            except (OSError, RuntimeError, subprocess.TimeoutExpired) as exc:
                page_attempts.append({"renderer": candidate, "status": "failed", "issue": str(exc)})
                continue
            selected_page_renderer = candidate
            selected_pages = current_pages
            page_attempts.append({"renderer": candidate, "status": "passed"})
            break
        if selected_page_renderer is None:
            terminal_issue = str(
                (page_attempts[-1] if page_attempts else {}).get("issue")
                or "The release-owned pypdfium2 page renderer failed."
            )
            terminal_code = terminal_issue.partition(": ")[0]
            if not terminal_code.startswith("renderer.pdfium_"):
                terminal_code = "renderer.pdfium_worker_failed"
            return {
                "status": "blocked",
                "renderer_attempts": renderer_attempts,
                "page_renderer_attempts": page_attempts,
                "findings": [{
                    "category": "renderer",
                    "field": "rendering",
                    "code": terminal_code,
                    "issue": terminal_issue,
                    "recovery_class": "adapter_fault",
                    "action": "stop",
                }],
            }

        findings: list[dict[str, Any]] = []
        artifacts = list(retained)
        if not selected_artifacts and not retained and render_root.exists():
            shutil.rmtree(render_root)
        render_root.mkdir(parents=True, exist_ok=True)
        for docx, source_pdf in pdfs.items():
            pdf = render_root / source_pdf.name
            pdf.unlink(missing_ok=True)
            shutil.copy2(source_pdf, pdf)
            page_dir = render_root / docx.stem
            shutil.rmtree(page_dir, ignore_errors=True)
            page_dir.mkdir()
            pages = []
            for index, source_page in enumerate(selected_pages[docx], 1):
                page = page_dir / f"page-{index}.png"
                shutil.copy2(source_page, page)
                pages.append(page)
            artifacts.append({
                "artifact": docx.stem,
                "status": "passed",
                "renderer": identity,
                "page_renderer": selected_page_renderer,
                "font_evidence": dict(font_evidence or {}),
                "font_substitutions": dict(font_substitutions or {}),
                "docx": docx.relative_to(revision_dir).as_posix(),
                "docx_sha256": sha256_file(docx),
                "pdf": pdf.relative_to(revision_dir).as_posix(),
                "pdf_sha256": sha256_file(pdf),
                "page_count": len(pages),
                "pages": [{"page": i, "path": page.relative_to(revision_dir).as_posix(), "sha256": sha256_file(page)} for i, page in enumerate(pages, 1)],
            })
        renderer_attempts.append({"renderer": identity, "status": "passed"})
        return {
            "status": "passed" if not findings else "blocked",
            "renderer": identity,
            "renderer_attempts": renderer_attempts,
            "page_renderer": selected_page_renderer,
            "page_renderer_attempts": page_attempts,
            "artifacts": artifacts,
            "reused_artifacts": sorted(retained_names),
            "findings": findings,
        }
    return {
        "status": "blocked",
        "renderer_attempts": renderer_attempts,
        "page_renderer_attempts": page_attempts,
        "findings": [{"category": "renderer", "field": "rendering", "issue": "Every local renderer/page-renderer combination failed."}],
    }


def render_assurance(
    repo_root: Path,
    revision_dir: Path,
    reference: Mapping[str, Any],
    *,
    contracted_bundle: Mapping[str, Any] | None = None,
    structural_validation: Mapping[str, Any],
    candidate_font_substitutions: Mapping[str, str] | None = None,
    rebuild_candidate: Any = None,
    artifact_names: Iterable[str] | None = None,
    reuse_render_report: Mapping[str, Any] | None = None,
    renderer_identities: Iterable[Mapping[str, Any]] | None = None,
    page_renderer_identities: Iterable[Mapping[str, Any]] | None = None,
    deadline_seconds: float = 180.0,
    deadline_monotonic: float | None = None,
    clock: Any = time.monotonic,
    font_probe: Any = None,
    office_exporter: Any = None,
    page_exporter: Any = None,
    blank_page_detector: Any = None,
    require_promoted_runtime: bool = True,
    allow_pending_activation: bool = False,
) -> dict[str, Any]:
    """Resolve fonts and render one complete candidate through one assurance seam."""
    started = clock()
    deadline = deadline_monotonic if deadline_monotonic is not None else started + max(0.0, deadline_seconds)
    bundle = contracted_bundle or contracted_template_bundle(repo_root, reference)
    approved_font_plan = bundle["approved_font_plan"]
    runtime_environment = _runtime_environment(repo_root, None, bundle)
    selected_artifacts = set(artifact_names or ())
    all_candidate_paths = [path for path in sorted((revision_dir / "candidate").glob("*")) if path.is_file()]
    structural_evidence, structural_finding = _validated_candidate_structure(
        revision_dir,
        structural_validation,
        all_candidate_paths,
    )
    candidate_paths = [
        path
        for path in all_candidate_paths
        if path.suffix.casefold() == ".docx"
        if not selected_artifacts or path.stem in selected_artifacts
    ]
    if structural_finding is not None or not candidate_paths:
        finding = structural_finding or recovery_finding({
                "category": "document-structure",
                "field": "candidate",
                "issue": "Render Assurance requires a complete structurally validated DOCX candidate.",
            }, "document_structure_defect")
        return {
            "schema_version": "render-assurance/v1",
            "status": "blocked",
            "fonts": {},
            "font_substitutions": {},
            "candidate": {"files": []},
            "structural_validation": structural_evidence,
            "render": {"status": "not_run", "findings": [finding]},
            "findings": [finding],
        }

    required_fonts = sorted({font for path in candidate_paths for font in _template_fonts(path)})
    probe = font_probe or _font_probe
    fonts: dict[str, dict[str, Any]] = {}
    substitutions: dict[str, str] = {}
    for font in required_fonts:
        remaining = max(0.0, deadline - clock())
        if remaining <= 0:
            available, detail = None, "Render Assurance deadline expired before font inventory."
        else:
            available, detail = probe(font, environment=runtime_environment, timeout_seconds=remaining)
        state = "available" if available is True else "missing-or-unusable" if available is False else "unknown"
        evidence = {"state": state, "match": detail}
        if state == "unknown":
            evidence.update({
                "recovery_class": "font_capability_uncertainty",
                "action": RECOVERY_POLICIES["font_capability_uncertainty"],
            })
        if state == "missing-or-unusable":
            substitute = _approved_packaged_font_fallback(font, approved_font_plan)
            bundled = _bundled_font_path(repo_root, substitute, approved_font_plan)
            if bundled is not None:
                if substitute.casefold() != font.casefold():
                    substitutions[font] = substitute
                evidence.update({
                    "substitute": substitute,
                    "substitute_match": f"bundled approved compatible font: {bundled.relative_to(repo_root)}",
                })
            else:
                evidence["attempted_fallbacks"] = [substitute]
        fonts[font] = evidence

    prior_substitutions = dict(candidate_font_substitutions or {})
    for source, target in prior_substitutions.items():
        if source not in required_fonts and target in required_fonts:
            substitutions[source] = target
            fonts[source] = {
                "state": "missing-or-unusable",
                "match": "Persisted candidate substitution.",
                "substitute": target,
                "resolution": "candidate_substitution_preserved",
            }
    if substitutions != prior_substitutions:
        if rebuild_candidate is None:
            finding = recovery_finding({
                "category": "document-structure",
                "field": "font_substitutions",
                "issue": "The candidate must be rebuilt with the selected approved font substitutions before rendering.",
            }, "document_structure_defect")
            return {
                "schema_version": "render-assurance/v1",
                "status": "blocked",
                "fonts": fonts,
                "font_substitutions": substitutions,
                "candidate": {"files": _artifact_hashes(revision_dir, all_candidate_paths)},
                "render": {"status": "not_run", "findings": [finding]},
                "findings": [finding],
            }
        rebuild_report = rebuild_candidate(substitutions)
        if rebuild_report.get("status") != "passed":
            findings = [recovery_finding(finding, "document_structure_defect") for finding in (rebuild_report.get("findings") or [{
                "category": "document-structure",
                "field": "font_substitutions",
                "issue": "The candidate could not be rebuilt with approved font substitutions.",
            }])]
            return {
                "schema_version": "render-assurance/v1",
                "status": "blocked",
                "fonts": fonts,
                "font_substitutions": substitutions,
                "candidate": {"files": _artifact_hashes(revision_dir, all_candidate_paths)},
                "render": {"status": "not_run", "findings": findings},
                "findings": findings,
            }
        candidate_paths = [
            path
            for path in sorted((revision_dir / "candidate").glob("*.docx"))
            if not selected_artifacts or path.stem in selected_artifacts
        ]
        all_candidate_paths = [path for path in sorted((revision_dir / "candidate").glob("*")) if path.is_file()]
        post_rebuild_validation = {
            **dict(structural_validation),
            "candidate_files": _artifact_hashes(revision_dir, all_candidate_paths),
        }
        structural_evidence, structural_finding = _validated_candidate_structure(
            revision_dir,
            post_rebuild_validation,
            all_candidate_paths,
        )
        structural_evidence["revalidated_after_substitution"] = structural_finding is None
        if structural_finding is not None or not candidate_paths:
            finding = structural_finding or recovery_finding({
                "category": "document-structure",
                "field": "candidate",
                "issue": "The rebuilt candidate does not contain the selected DOCX artifacts.",
            }, "document_structure_defect")
            return {
                "schema_version": "render-assurance/v1",
                "status": "blocked",
                "fonts": fonts,
                "font_substitutions": substitutions,
                "candidate": {"files": _artifact_hashes(revision_dir, all_candidate_paths)},
                "structural_validation": structural_evidence,
                "render": {"status": "not_run", "findings": [finding]},
                "findings": [finding],
            }

    office_candidates = (
        [dict(item) for item in renderer_identities]
        if renderer_identities is not None
        else renderers(environment=runtime_environment, skill_root=repo_root, deadline_monotonic=deadline, clock=clock)
    )
    page_candidates = _one_pdfium_renderer(
        [dict(item) for item in page_renderer_identities]
        if page_renderer_identities is not None
        else page_renderers(
            environment=runtime_environment,
            skill_root=repo_root,
            require_promoted_runtime=require_promoted_runtime,
            allow_pending_activation=allow_pending_activation,
        ),
        skill_root=repo_root,
        require_promoted_runtime=require_promoted_runtime,
        allow_pending_activation=allow_pending_activation,
    )
    render_report = render_pages(
        revision_dir,
        artifact_names=artifact_names,
        reuse_render_report=reuse_render_report,
        contracted_bundle=bundle,
        renderer_identities=office_candidates,
        page_renderer_identities=(
            page_candidates
            if page_renderer_identities is not None or page_candidates
            else None
        ),
        deadline_monotonic=deadline,
        clock=clock,
        font_evidence=fonts,
        font_substitutions=substitutions,
        office_exporter=office_exporter,
        page_exporter=page_exporter,
        blank_page_detector=blank_page_detector,
        require_promoted_runtime=require_promoted_runtime,
        allow_pending_activation=allow_pending_activation,
    )
    render_report["renderer_attempts"] = _governed_adapter_attempts(render_report.get("renderer_attempts", []))
    render_report["page_renderer_attempts"] = _governed_adapter_attempts(
        render_report.get("page_renderer_attempts", []),
        failure_action="stop",
    )
    governed_findings = []
    for raw in render_report.get("findings", []):
        finding = dict(raw)
        if finding.get("category") == "visual":
            finding.update({
                "recovery_class": "visual_defect",
                "action": RECOVERY_POLICIES["visual_defect"],
            })
        elif finding.get("category") == "renderer" and not finding.get("recovery_class"):
            finding = recovery_finding(finding, "adapter_fault")
        governed_findings.append(finding)
    render_report["findings"] = governed_findings
    if render_report.get("status") == "passed":
        for evidence in fonts.values():
            if evidence["state"] == "unknown":
                evidence["resolution"] = "render_verified"
        for artifact in render_report.get("artifacts", []):
            artifact["font_evidence"] = fonts
    report = {
        "schema_version": "render-assurance/v1",
        "status": render_report.get("status", "blocked"),
        "fonts": fonts,
        "font_substitutions": substitutions,
        "candidate": {
            "files": _artifact_hashes(revision_dir, all_candidate_paths),
            "font_evidence": fonts,
            "font_substitutions": substitutions,
        },
        "structural_validation": structural_evidence,
        "render": render_report,
        "findings": governed_findings,
        "elapsed_seconds": round(clock() - started, 3),
    }
    if report["status"] != "passed" and governed_findings and all(
        finding.get("recovery_class") == "adapter_fault" for finding in governed_findings
    ):
        report["diagnostic"] = {
            "recovery_class": "adapter_fault",
            "action": RECOVERY_POLICIES["adapter_fault"],
            "outcome": "adapters_exhausted",
            "candidate_disposition": "preserved",
            "publication": "blocked",
            "office_attempts": render_report["renderer_attempts"],
            "page_attempts": render_report["page_renderer_attempts"],
            "font_evidence": fonts,
            "font_substitutions": substitutions,
            "candidate_files": report["candidate"]["files"],
        }
    return report


def _validated_candidate_structure(
    revision_dir: Path,
    validation: Mapping[str, Any],
    candidate_paths: Iterable[Path],
) -> tuple[dict[str, Any], dict[str, Any] | None]:
    paths = list(candidate_paths)
    current = _artifact_hashes(revision_dir, paths)
    expected = {str(name) for name in validation.get("expected_files", []) if str(name)}
    actual = {path.name for path in paths}
    recorded = {
        str(item.get("path")): item
        for item in validation.get("candidate_files", [])
        if isinstance(item, Mapping) and item.get("path")
    }
    evidence = {
        "status": validation.get("status"),
        "expected_files": sorted(expected),
        "candidate_files": current,
    }
    complete = (
        validation.get("status") == "structurally_valid"
        and bool(expected)
        and actual == expected
        and all(
            (row := recorded.get(item["path"])) is not None
            and row.get("sha256") == item["sha256"]
            and int(row.get("bytes") or 0) == item["bytes"]
            for item in current
        )
    )
    if complete:
        return evidence, None
    return evidence, recovery_finding({
        "category": "document-structure",
        "field": "candidate",
        "issue": "Render Assurance requires the complete structurally validated Branch Document Set with matching hashes.",
    }, "document_structure_defect")


def _governed_adapter_attempts(
    attempts: Iterable[Mapping[str, Any]],
    *,
    failure_action: str | None = None,
) -> list[dict[str, Any]]:
    governed = []
    for raw in attempts:
        attempt = {"adapter": dict(raw.get("renderer") or raw.get("adapter") or {}), "status": raw.get("status")}
        if raw.get("status") in {"failed", "skipped"}:
            attempt.update({
                "recovery_class": "adapter_fault",
                "action": failure_action or RECOVERY_POLICIES["adapter_fault"],
                "issue": str(raw.get("issue") or "Adapter did not complete."),
            })
        governed.append(attempt)
    return governed


def _artifact_hashes(revision_dir: Path, paths: Iterable[Path]) -> list[dict[str, Any]]:
    return [
        {
            "path": path.relative_to(revision_dir).as_posix(),
            "sha256": sha256_file(path),
            "bytes": path.stat().st_size,
        }
        for path in paths
    ]


def _timeline_covered(approved_timeline: Any, visible_text: Any) -> bool:
    """Accept grammatical glue while keeping every milestone bound to its time value."""
    approved = re.sub(r"\s+", " ", _text(approved_timeline)).strip().casefold().rstrip(".")
    visible = re.sub(r"\s+", " ", _text(visible_text)).strip().casefold()
    if not approved or approved in visible:
        return True
    month_year = (
        r"(?:january|february|march|april|may|june|july|august|"
        r"september|october|november|december)\s+\d{4}"
    )
    dates = re.findall(month_year, approved)
    if (
        len(dates) == 2
        and re.search(r"\breview\b[^.;]{0,100}\brecords\b", approved)
        and "no new participant visits" in approved
    ):
        historical_window = (
            rf"\b(?:review|chart abstraction)\b[^.;]{{0,100}}\brecords\b[^.;]{{0,80}}"
            rf"\b{re.escape(dates[0])}\b[^.;]{{0,80}}\b{re.escape(dates[1])}\b"
            rf".{{0,160}}(?:\bno(?:\s+new participant visits|"
            rf"\s+participant contact(?:,| or)\s*new participant visits)\b|"
            rf"\breview does not involve new participant visits\b)"
        )
        if re.search(historical_window, visible):
            return True
    clauses = [clause.strip(" .") for clause in approved.split(";") if clause.strip(" .")]
    pairs: list[tuple[str, str]] = []
    for clause in clauses:
        if ":" not in clause:
            return False
        milestone, time_value = (part.strip(" ,.") for part in clause.split(":", 1))
        if not milestone or not time_value:
            return False
        pairs.append((milestone, time_value))
    return bool(pairs) and all(
        re.search(
            rf"\b{re.escape(milestone)}\b[^.;]{{0,80}}\b{re.escape(time_value)}\b",
            visible,
        )
        for milestone, time_value in pairs
    )


def _retrospective_lexical_coverage_finding(
    approved: Any, visible_text: str, *, field: str, issue: str,
) -> dict[str, Any] | None:
    """Use the shared evidence policy for retrospective narrative coverage."""
    expected = re.sub(r"\s+", " ", _text(approved)).strip().casefold().rstrip(".")
    visible = re.sub(r"\s+", " ", visible_text).strip().casefold()
    if not expected or expected in visible:
        return None
    diagnostic = source_evidence_diagnostics(visible, field, approved, all_items=True)
    if diagnostic["passed"]:
        # Different spelling with observable coverage is not itself a defect.
        return None
    return {**evidence_coverage_finding(field, [diagnostic], issue,
        "Preserve the demonstrated approved value.", category="content"),
        "source_excerpt": _text(approved)}


def audit_computed_protocol_fields(document: Path | Document, reference: Mapping[str, Any]) -> list[dict[str, Any]]:
    """Check code-owned synopsis destinations before an independent review call."""
    document = document if hasattr(document, "tables") else Document(document)
    expected = participant_followup_summary(reference)
    findings = []
    for table in document.tables:
        for row in table.rows:
            label = re.sub(r"[^a-z]", "", row.cells[0].text.casefold())
            if label not in {"durationfollowup", "durationfollwup"} or len(row.cells) < 2:
                continue
            actual = row.cells[1].text.strip()
            if re.sub(r"\s+", " ", actual) != re.sub(r"\s+", " ", expected):
                findings.append(recovery_finding({
                    "category": "content", "check": "computed_source_fidelity",
                    "field": "general-information", "artifact": "protocol",
                    "target_ids": ["general-information"], "source_path": "study.timeline",
                    "issue": "The code-owned Duration / Follow-up synopsis differs from the approved participant follow-up projection.",
                    "expected": expected, "actual": actual,
                }, "deterministic_structure_defect"))
    return findings


def audit_source_surfaces(path: Path, reference: Mapping[str, Any]) -> list[dict[str, Any]]:
    """Conservative source-leaf checks, not a second renderer/mapping contract.

    These catch literal front-matter loss/concatenation. Semantic coverage and
    correct placement still require the independent field-by-field review.
    """
    document = Document(path)
    findings: list[dict[str, Any]] = []
    if path.stem == "icf":
        sterling = str(get_path(reference, "meta.icf_template", "")).casefold() == "sterling"
        if sterling:
            # Sterling owns the final first-page mail merge. The approved
            # address belongs in the source and protocol, while these exact
            # tags must survive in the ICF for the IRB's merge process.
            first_page_paragraphs = document.paragraphs[:20]
            site_index = next((
                index for index, paragraph in enumerate(first_page_paragraphs)
                if paragraph.text.strip().startswith("STUDY SITE:")
            ), None)
            site_block = "\n".join(
                paragraph.text for paragraph in first_page_paragraphs[site_index:site_index + 3]
            ) if site_index is not None else ""
            for field, marker in (
                ("sites[0].facility.name", "«Company_Name»"),
                ("sites[0].facility.address", "«Address»"),
                ("sites[0].facility.address.location", "«City_State_ZIP»"),
            ):
                if site_block.count(marker) != 1:
                    findings.append({
                        "category": "template_contract",
                        "field": field,
                        "target_ids": ["layout:icf"],
                        "issue": f"Sterling first-page merge field {marker} must appear exactly once in the STUDY SITE block.",
                    })
        else:
            # Inspect the address destination, never an unrelated later mention.
            address_blocks = []
            if document.tables:
                for row in document.tables[0].rows:
                    if re.fullmatch(r"(?:study\s+site\s+)?address\s*:", row.cells[0].text.strip(), re.I):
                        address_blocks.append("\n".join(cell.text for cell in row.cells[1:]))
            if not address_blocks:
                collecting = False
                for paragraph in document.paragraphs:
                    text = paragraph.text
                    if re.match(r"^\s*(?:study\s+site|site\s+address)\s*:", text, re.I):
                        address_blocks.append(text)
                        collecting = True
                    elif collecting and text.startswith("\t"):
                        address_blocks.append(text)
                    elif collecting and text.strip():
                        break
            normalized = " ".join(re.findall(r"\w+", "\n".join(address_blocks).casefold()))
            inventory = _source_field_inventory(reference)
            source_country = next((
                str(item["value"]) for item in inventory
                if item["source_path"] in {
                    "sites[0].facility.country", "sites[0].facility.address.country"
                }
            ), "")
            for item in inventory:
                field = item["source_path"]
                if not re.fullmatch(
                    r"sites\[0\]\.facility\.(?:address(?:\.(?:line1|line2|address_line1|address_line2|street|city|state|country|zip|postal_code))?|city|state|country|zip|postal_code)", field
                ):
                    continue
                value = " ".join(re.findall(r"\w+", str(item["value"]).casefold()))
                aliases = (
                    country_name_aliases(str(item["value"]))
                    if field.endswith(".country")
                    else state_name_aliases(
                        str(item["value"]), source_country,
                    ) if field.endswith(".state")
                    else {value}
                )
                present = any(
                    alias and re.search(r"(?<!\w)" + re.escape(alias) + r"(?!\w)", normalized)
                    for alias in aliases
                )
                if value and not present:
                    findings.append({"category": "content", "field": field,
                                     "target_ids": ["layout:icf"],
                                     "issue": f"ICF omits supplied site-address component {field}: {item['value']}"})
    if path.stem == "protocol":
        references = get_path(reference, "study.references") or reference.get("references")
        references_path = "study.references" if get_path(reference, "study.references") else "references"
        entries = references if isinstance(references, list) else [references] if isinstance(references, str) else []
        visible = re.sub(r"\s+", " ", "\n".join(p.text for p in document.paragraphs))
        for entry in entries:
            if isinstance(entry, str) and entry.strip() and re.sub(r"\s+", " ", entry.strip()) not in visible:
                findings.append(recovery_finding({
                    "category": "content", "artifact": "protocol", "field": "protocol.shell",
                    "target_ids": ["protocol.shell"], "source_path": references_path,
                    "check": "computed_source_fidelity",
                    "issue": "The code-owned bibliography omits a supplied reference.",
                    "source_excerpt": entry,
                }, "deterministic_structure_defect"))
        sponsor_name = str(get_path(reference, "parties.sponsor.name", "") or "").strip()
        if sponsor_name:
            for table in document.tables:
                for row in table.rows:
                    if not row.cells or not re.search(r"\bsponsor\b", row.cells[0].text, re.I):
                        continue
                    value = "\n".join(cell.text for cell in row.cells[1:])
                    if re.search(r"\w" + re.escape(sponsor_name) + r"|" + re.escape(sponsor_name) + r"\w", value, re.I):
                        findings.append({"category": "content", "field": "parties.sponsor.name",
                                         "target_ids": ["title-page"],
                                         "issue": "Sponsor name is concatenated with adjacent front-matter text without a separator."})
    return findings


def _normalized_substantive_text(value: str) -> str:
    return re.sub(r"[^a-z0-9]+", " ", value.casefold()).strip()


def _sterling_clause_sections(document: Document) -> dict[str, list[str]]:
    """Read governed sections in body order, including table-contained text."""
    contract = sterling_clause_contract()
    headings = {
        _normalized_substantive_text(str(item["section"])): str(item["section"])
        for item in contract["clauses"]
    }
    sections: dict[str, list[str]] = {title: [] for title in headings.values()}
    current = ""
    front_matter_nonsections = {
        "participant informed consent form and", "protocol no", "study",
        "investigator", "telephone",
    }

    def consume(text: str, *, heading_style: bool = False) -> None:
        nonlocal current
        text = re.sub(r"\s+", " ", text).strip()
        if not text:
            return
        key = _normalized_substantive_text(text)
        if key in headings:
            current = headings[key]
            return
        if heading_style and not any(
            key == marker or key.startswith(f"{marker} ")
            for marker in front_matter_nonsections
        ):
            current = ""
            return
        if current:
            sections[current].append(text)

    for child in document.element.body.iterchildren():
        if child.tag == qn("w:p"):
            paragraph = Paragraph(child, document)
            consume(
                paragraph.text,
                heading_style=paragraph.style.name.casefold().startswith("heading"),
            )
        elif child.tag == qn("w:tbl"):
            table = Table(child, document)
            for row in table.rows:
                for cell in row.cells:
                    for paragraph in cell.paragraphs:
                        consume(
                            paragraph.text,
                            heading_style=paragraph.style.name.casefold().startswith("heading"),
                        )
    return sections


def _sterling_triggered(clause: Mapping[str, Any], reference: Mapping[str, Any]) -> bool:
    trigger = clause.get("trigger") if isinstance(clause.get("trigger"), Mapping) else {}
    rule = str(trigger.get("rule") or "")
    paths = [str(path) for path in trigger.get("paths", [])]
    if rule == "always":
        return True
    values = [get_path(reference, path) for path in paths]
    if rule == "truthy":
        return any(value is True for value in values)
    if rule == "meaningful":
        return any(meaningful(value) for value in values)
    return False


def validate_sterling_clause_contract(
    rendered_document: Path | Document,
    normalized_source: Mapping[str, Any],
) -> dict[str, Any]:
    """Validate substantive Sterling language, applicability and placement."""
    document = rendered_document if hasattr(rendered_document, "paragraphs") else Document(rendered_document)
    contract = sterling_clause_contract()
    sections = _sterling_clause_sections(document)
    normalized_sections = {
        title: _normalized_substantive_text(" ".join(values))
        for title, values in sections.items()
    }
    raw_sections = {title: " ".join(values) for title, values in sections.items()}
    findings: list[dict[str, Any]] = []
    branch = canonical_study_type(get_path(normalized_source, "meta.study_type")) or "Prospective"
    draftable_icf_ids = {
        item.section_id for item in icf_contract(branch, "Sterling")
    }

    def add(clause: Mapping[str, Any], code: str, **extra: Any) -> None:
        severity = str((clause.get("severity") or {}).get({
            "sterling-clause-missing": "absent",
            "sterling-clause-weakened": "altered",
            "sterling-clause-unsupported": "unsupported",
            "sterling-clause-misplaced": "misplaced",
        }[code], "blocking"))
        finding = {
            "code": code,
            "category": "content",
            "check": "sterling_clause_contract",
            "clause_id": clause["clause_id"],
            "classification": clause["classification"],
            "target_ids": [clause["section_id"]],
            "expected_section": clause["section"],
            "severity": severity,
            "publication_disposition": "blocking",
            "safety_critical": bool(clause.get("safety_critical")) or any(
                marker in str(clause["clause_id"])
                for marker in ("risk", "injury", "voluntary", "rights", "authorization", "privacy")
            ),
            "issue": {
                "sterling-clause-missing": "Required governed Sterling clause is absent.",
                "sterling-clause-weakened": "Governed Sterling clause is materially altered, incomplete, or overbroad.",
                "sterling-clause-unsupported": "Sterling conditional or source-dependent language appears without its approved trigger.",
                "sterling-clause-misplaced": "Governed Sterling clause is outside its required section.",
            }[code],
            **extra,
        }
        recovery = (
            "deterministic_structure_defect"
            if code == "sterling-clause-misplaced"
            else "drafting_defect"
            if str(clause["section_id"]) in draftable_icf_ids
            else "document_structure_defect"
        )
        findings.append(recovery_finding(finding, recovery))

    for clause in contract["clauses"]:
        section = str(clause["section"])
        text = normalized_sections.get(section, "")
        validation = clause.get("validation") or {}
        triggered = _sterling_triggered(clause, normalized_source)
        exact_texts = [str(value) for value in validation.get("required_exact_texts", [])]
        exact_boilerplate = validation.get("exact_boilerplate_id")
        if exact_boilerplate:
            exact_texts.append(sterling_clause_text(str(clause["clause_id"])))
        if validation.get("exact_template"):
            exact_texts.append(sterling_clause_text(str(clause["clause_id"])))
        normalized_exact = [_normalized_substantive_text(value) for value in exact_texts]
        placement_markers = (
            [_normalized_substantive_text(sterling_clause_text(str(clause["clause_id"])))]
            if validation.get("placement_boilerplate_id")
            else normalized_exact
        )
        term_groups = [
            [_normalized_substantive_text(str(term)) for term in group]
            for group in validation.get("required_term_groups", [])
        ]
        source_values = [
            _normalized_substantive_text(_text(get_path(normalized_source, str(path))))
            for path in validation.get("required_source_values", [])
        ]
        authority = clause.get("approved_source") if isinstance(clause.get("approved_source"), Mapping) else {}
        authority_paths = [str(path) for path in authority.get("paths", [])]
        authority_records = [
            (path, section_evidence_value(str(clause["section_id"]), path, get_path(normalized_source, path), normalized_source))
            for path in authority_paths
            if meaningful(get_path(normalized_source, path))
        ]
        unsupported_markers = [
            _normalized_substantive_text(str(value))
            for value in validation.get("unsupported_markers", [])
        ]

        if not triggered:
            unsupported = any(marker and marker in text for marker in unsupported_markers)
            prohibited_section = validation.get("prohibit_when_untriggered") and section in {
                "COMPENSATION TO YOU", "GENETIC INFORMATION NONDISCRIMINATION ACT"
            } and bool(text)
            if prohibited_section:
                add(clause, "sterling-clause-unsupported", actual_section=section)
            elif unsupported:
                findings.append({
                    "code": "icf-lexical-evidence-uncertainty", "category": "content",
                    "field": clause["section_id"], "target_ids": [clause["section_id"]],
                    "clause_id": clause["clause_id"], "publication_disposition": "warning",
                    "matched_markers": [marker for marker in unsupported_markers if marker and marker in text],
                    "issue": "A phrase may imply an unapproved conditional clause. Independently review its meaning and approved applicability; a phrase match alone does not establish that claim.",
                })
            continue

        has_body = bool(text)
        exact_ok = all(value in text for value in normalized_exact)
        terms_ok = all(any(term and term in text for term in group) for group in term_groups)
        sources_ok = all(value and value in text for value in source_values)
        grounding_results = [
            (
                re.search(r"\b(?:no|not|without|will not|none)\b", text) is not None
                if isinstance(value, str) and value.strip().casefold() == "none"
                else source_evidence_grounded(raw_sections.get(section, ""), path, value, reference=normalized_source)
            )
            for path, value in authority_records
        ]
        semantic_issues = [
            issue for path, value in authority_records if path == "study.hypothesis"
            for issue in hypothesis_claim_issues(raw_sections.get(section, ""), value)
        ]
        grounding_ok = (
            any(grounding_results)
            if validation.get("source_grounding") == "any" and grounding_results
            else all(grounding_results)
        )
        contradiction_patterns = {
            "sterling.risks.foreseeable": (r"\bno (?:foreseeable )?risks?\b",),
            "sterling.voluntary.core": (r"\bmust participate\b", r"\bcannot (?:leave|withdraw|stop)\b", r"\bwith penalty\b"),
            "sterling.injury.core": (r"(?<!not )\bwaive (?:all |your )?legal rights\b",),
            "sterling.privacy.authorization": (r"\bwill identify you\b", r"\bno confidentiality\b"),
            "sterling.contact.participant-rights": (r"\bno participant rights\b",),
        }
        contradiction = any(
            re.search(pattern, text)
            for pattern in contradiction_patterns.get(str(clause["clause_id"]), ())
        )
        word_count = len(text.split())
        length_ok = word_count >= int(validation.get("minimum_words", 0) or 0)
        block_count = len([value for value in sections.get(section, []) if value.strip()])
        block_ok = block_count >= int(validation.get("minimum_substantive_blocks", 0) or 0)
        maximum = int(validation.get("maximum_words", 0) or 0)
        maximum_ok = not maximum or word_count <= maximum
        substantive_ok = (
            has_body and exact_ok and terms_ok and sources_ok and grounding_ok
            and not semantic_issues and not contradiction and block_ok
        )
        if substantive_ok:
            if not length_ok or not maximum_ok:
                findings.append({
                    "code": "sterling-clause-word-goal",
                    "category": "editorial",
                    "check": "sterling_clause_contract",
                    "clause_id": clause["clause_id"],
                    "target_ids": [clause["section_id"]],
                    "publication_disposition": "warning",
                    "issue": "The clause is outside its reference word-count goal.",
                    "actual_words": word_count,
                    "minimum_words": int(validation.get("minimum_words", 0) or 0),
                    "maximum_words": maximum,
                })
            continue

        grounding_diagnostics = [source_evidence_diagnostics(raw_sections.get(section, ""), path, value, reference=normalized_source)
                                 for (path, value), passed in zip(authority_records, grounding_results)
                                 if not passed]
        if (has_body and exact_ok and sources_ok and block_ok
                and not semantic_issues and not contradiction
                and (grounding_ok or lexical_evidence_uncertainty(grounding_diagnostics))
                and (not terms_ok or lexical_evidence_uncertainty(grounding_diagnostics))):
            findings.append({
                "code": "icf-lexical-evidence-uncertainty", "category": "content",
                "field": clause["section_id"], "target_ids": [clause["section_id"]],
                "clause_id": clause["clause_id"], "publication_disposition": "warning",
                "evidence_diagnostics": grounding_diagnostics,
                "missing_term_groups": [group for group in term_groups if not any(term in text for term in group)],
                "issue": "Sterling meaning and safeguards require independent review; exact-text and quantitative checks remain enforced.",
            })
            continue

        # Exact or safeguard-bearing text in another governed section is a
        # deterministic placement defect; preserve its bytes and move it.
        marker_values = placement_markers
        actual = next((
            title for title, candidate in normalized_sections.items()
            if title != section and marker_values
            and sum(value in candidate for value in marker_values) >= max(1, len(marker_values) // 2)
        ), "")
        if actual:
            add(clause, "sterling-clause-misplaced", actual_section=actual)
        elif not has_body or not block_ok:
            add(clause, "sterling-clause-missing")
        else:
            add(
                clause,
                "sterling-clause-weakened",
                source_fidelity_details=semantic_issues,
                evidence_diagnostics=[
                    source_evidence_diagnostics(raw_sections.get(section, ""), path, value, reference=normalized_source)
                    for (path, value), passed in zip(authority_records, grounding_results)
                    if not passed
                ],
                missing_term_groups=[group for group in term_groups if not any(term in text for term in group)],
                missing_source_values=[value for value in source_values if value not in text],
                ungrounded_source_paths=[
                    path for path, value in authority_records
                    if not source_evidence_grounded(raw_sections.get(section, ""), path, value, reference=normalized_source)
                    and not (isinstance(value, str) and value.strip().casefold() == "none" and re.search(r"\b(?:no|not|without|will not|none)\b", text))
                ],
                contradiction=contradiction,
            )
    return {
        "schema_version": "sterling-clause-validation/v1",
        "status": "blocked" if _blocking_findings(findings) else "passed",
        "contract_version": contract["schema_version"],
        "findings": findings,
    }


def assess_icf_output(
    rendered_document: Path | Document,
    normalized_source: Mapping[str, Any],
    family_contract: Mapping[str, Any],
) -> dict[str, Any]:
    """Assess exact material duplication introduced in the completed ICF.

    The interface deliberately accepts only the rendered document plus its two
    governing inputs.  Detection is conservative: only KEY INFORMATION is
    compared with detailed generated sections, and repair always targets the
    summary rather than deleting a complete disclosure.
    """
    document: Any = rendered_document if hasattr(rendered_document, "paragraphs") else Document(rendered_document)
    heading_map = {
        "BACKGROUND": "icf.background",
        "PURPOSE": "icf.study-purpose",
        "DURATION": "icf.duration",
        "PROCEDURES": "icf.procedures",
        "POTENTIAL RISKS SIDE EFFECTS DISCOMFORTS INCONVENIENCES": "icf.risks",
        "POTENTIAL RISKS EFFECTS DISCOMFORTS INCONVENIENCES": "icf.risks",
        "POTENTIAL BENEFITS": "icf.benefits",
        "ALTERNATIVE TREATMENTS": "icf.alternatives",
        "COMPENSATION TO YOU": "icf.payment",
        "COSTS TO YOU": "icf.costs",
        "STUDY COMPLICATIONS AND COMPENSATION": "icf.injury",
        "STUDY COMPLICATIONS COMPENSATION": "icf.injury",
        "CONFIDENTIALITY AND AUTHORIZATION TO COLLECT USE AND DISCLOSE YOUR MEDICAL INFORMATION": "icf.privacy",
        "CONFIDENTIALITY AUTHORIZATION TO COLLECT USE DISCLOSE YOUR MEDICAL INFORMATION": "icf.privacy",
    }
    retained_heading_keys = {
        re.sub(r"[^A-Z0-9]+", " ", title.upper()).strip()
        for _section_id, title in ICF_RETAINED_SHELL_SECTIONS.get("sterling", ())
    }
    retained_heading_keys.update({
        "NEW INFORMATION", "PARTICIPANT STATEMENT AND AUTHORIZATION",
        "PARTICIPANT STATEMENT AUTHORIZATION",
    })
    current = ""
    summary_blocks: list[str] = []
    detail_blocks: dict[str, list[str]] = {}
    for paragraph in document.paragraphs:
        text = re.sub(r"\s+", " ", paragraph.text).strip()
        key = re.sub(r"[^A-Z0-9]+", " ", text.upper()).strip()
        if key == "KEY INFORMATION":
            current = "icf.key-information-summary"
            continue
        if key in heading_map:
            current = heading_map[key]
            continue
        if key in retained_heading_keys:
            current = ""
            continue
        if not text or text.casefold() == "things you should know":
            continue
        if current == "icf.key-information-summary":
            summary_blocks.append(text)
        elif current.startswith("icf."):
            detail_blocks.setdefault(current, []).append(text)

    repeatable_markers = (
        "signature of participant", "signature of principal investigator",
        "participant statement", "printed name", "date",
    )
    classified_repeatable_values = [
        value
        for key in ("repeatable_text", "required_static_text", "legal_clauses")
        for value in family_contract.get(key, [])
    ]
    repeatable_text = {
        _normalized_substantive_text(str(value))
        for value in classified_repeatable_values
        if _normalized_substantive_text(str(value))
    }

    def explicitly_repeatable(value: str) -> bool:
        return any(value == allowed or value in allowed or allowed in value for allowed in repeatable_text)

    def candidates(text: str) -> set[str]:
        normalized = _normalized_substantive_text(text)
        values = {normalized}
        values.update(
            _normalized_substantive_text(sentence)
            for sentence in re.split(r"(?<=[.!?])\s+", text)
        )
        words = normalized.split()
        span_words = 16
        values.update(
            " ".join(words[index:index + span_words])
            for index in range(max(0, len(words) - span_words + 1))
        )
        return {
            value for value in values
            if len(value) >= 80 and len(value.split()) >= 12
            and not explicitly_repeatable(value)
            and not any(marker in value for marker in repeatable_markers)
        }

    details: dict[str, set[str]] = {}
    for section_id, blocks in detail_blocks.items():
        for block in blocks:
            for value in candidates(block):
                details.setdefault(value, set()).add(section_id)
    findings = []
    seen_sections: set[str] = set()
    for block in summary_blocks:
        for value in sorted(candidates(block), key=lambda item: (-len(item.split()), item)):
            for section_id in sorted(details.get(value, set())):
                if section_id in seen_sections:
                    continue
                seen_sections.add(section_id)
                findings.append(recovery_finding({
                    "code": "icf-exact-generated-duplication",
                    "category": "editorial",
                    "check": "exact_rendered_duplication",
                    "publication_disposition": "warning",
                    "target_ids": ["icf.key-information-summary"],
                    "summary_section": "icf.key-information-summary",
                    "detail_section": section_id,
                    "normalized_text": value,
                    "repair_action": "redraft_summary",
                }, "drafting_defect"))

    # Section ownership also applies to paraphrases. BACKGROUND owns lens
    # context, comparative evidence, the evidence gap, and rationale; PURPOSE
    # owns only the concise purpose, hypothesis, and primary outcome.
    stop_words = {
        "a", "an", "and", "are", "as", "at", "be", "by", "for", "from",
        "in", "is", "it", "of", "on", "or", "that", "the", "this", "to",
        "was", "were", "will", "with",
    }

    def concept_tokens(value: str) -> set[str]:
        return {
            token for token in _normalized_substantive_text(value).split()
            if len(token) >= 4 and token not in stop_words
        }

    background = " ".join(detail_blocks.get("icf.background", ()))
    purpose = " ".join(detail_blocks.get("icf.study-purpose", ()))
    background_tokens = concept_tokens(background)
    purpose_tokens = concept_tokens(purpose)
    union = background_tokens | purpose_tokens
    overlap = len(background_tokens & purpose_tokens) / len(union) if union else 0.0
    normalized_background = _normalized_substantive_text(background)
    normalized_purpose = _normalized_substantive_text(purpose)
    rationale_markers = (
        "non diffractive", "dysphotops", "contrast sensitivity", "randomized",
        "monofocal", "evidence", "broader range", "fewer visual",
    )
    repeated_markers = [
        marker for marker in rationale_markers
        if marker in normalized_background and marker in normalized_purpose
    ]
    if len(purpose_tokens) >= 24 and (overlap >= 0.30 or len(repeated_markers) >= 3):
        findings.append(recovery_finding({
            "code": "icf-concept-repetition",
            "category": "editorial",
            "check": "concept_repetition",
            "concept_id": "clinical-background-and-rationale",
            "primary_section": "icf.background",
            "secondary_section": "icf.study-purpose",
            "primary_paragraphs": list(detail_blocks.get("icf.background", ())),
            "secondary_paragraphs": list(detail_blocks.get("icf.study-purpose", ())),
            "target_ids": ["icf.study-purpose"],
            "treatment": "excessive",
            "necessary": False,
            "concise": False,
            "material": False,
            "repair_action": "redraft_secondary_section",
            "severity": "warning",
            "publication_disposition": "warning",
        }, "drafting_defect"))
    return {
        "schema_version": "icf-output-assessment/v1",
        "status": "repairable" if _blocking_findings(findings) else "passed",
        "family": str(family_contract.get("family") or get_path(normalized_source, "meta.icf_template") or ""),
        "findings": findings,
    }


def assess_protocol_concept_repetition(
    section_paragraphs: Mapping[str, Iterable[str]],
    ownership: Mapping[str, Mapping[str, Iterable[str]]],
) -> list[dict[str, Any]]:
    """Return explainable manual-review findings from explicit concept rules."""
    concept_markers = {
        "clinical-rationale": (
            ("unmet need", "evidence gap", "remaining gap"),
            ("current care", "existing treatment", "current treatment"),
            ("variable", "inconsistent", "limited"),
            ("recovery", "outcome"),
            ("new approach", "different approach", "evaluates", "investigation"),
        ),
        "complete-visit-schedule": (
            ("screening",), ("baseline",), ("follow up", "follow-up"),
            ("visit", "appointment"), ("week", "month", "day"),
        ),
        "ae-sae-definitions": (
            ("adverse event",), ("serious adverse event",),
            ("hospitalization", "life threatening", "death", "disability"),
        ),
        "statistical-methods": (
            ("mean", "median"), ("standard deviation",),
            ("percentage", "frequency"), ("confidence interval",),
        ),
        "endpoint-inventory": (
            ("primary endpoint",), ("secondary endpoint", "secondary endpoints"),
            ("visual acuity", "symptom score", "questionnaire"),
            ("distance", "intermediate", "near", "defocus"), ("week", "month", "day"),
        ),
    }
    thresholds = {
        "clinical-rationale": (20, 3),
        "complete-visit-schedule": (90, 4),
        "ae-sae-definitions": (35, 3),
        "statistical-methods": (30, 3),
        "endpoint-inventory": (28, 4),
    }
    owners = {
        concept: section_id
        for section_id, policy in ownership.items()
        for concept in policy.get("owns", ())
    }
    findings: list[dict[str, Any]] = []
    for secondary, policy in ownership.items():
        prohibited = set(policy.get("do_not_restate", ()))
        for concept in sorted(prohibited):
            primary = owners.get(concept)
            if not primary or primary == secondary:
                continue
            markers = concept_markers.get(concept, ())
            for paragraph in section_paragraphs.get(secondary, ()):
                normalized = _normalized_substantive_text(str(paragraph))
                matched = sum(any(_normalized_substantive_text(term) in normalized for term in group) for group in markers)
                minimum_words, minimum_markers = thresholds.get(concept, (30, 3))
                if len(normalized.split()) < minimum_words or matched < minimum_markers:
                    continue
                primary_paragraphs = [str(item) for item in section_paragraphs.get(primary, ())]
                secondary_tokens = set(normalized.split())
                overlap = max((
                    len(secondary_tokens & set(_normalized_substantive_text(item).split()))
                    / max(1, len(secondary_tokens | set(_normalized_substantive_text(item).split())))
                    for item in primary_paragraphs
                ), default=0.0)
                role_specific_reference = (
                    concept == "endpoint-inventory"
                    and secondary == "study-procedure.measurements"
                    and any(term in normalized for term in ("measure", "measured", "measurement"))
                    and not ("defocus" in normalized and "questionnaire" in normalized)
                ) or (
                    concept == "endpoint-inventory"
                    and secondary == "analysis-plan.datasets"
                    and any(term in normalized for term in ("analysis data set", "observations enter", "analysis population"))
                    and matched <= minimum_markers
                ) or (
                    concept == "endpoint-inventory"
                    and secondary == "analysis-plan.methodology"
                    and any(term in normalized for term in ("summarized", "analysed", "analyzed"))
                    and overlap < 0.45
                )
                if role_specific_reference:
                    continue
                material = concept == "endpoint-inventory" and (
                    overlap >= 0.45 or (matched >= 5 and len(normalized.split()) >= 90)
                )
                findings.append({
                    "code": "protocol-concept-repetition",
                    "concept_id": concept,
                    "primary_section": primary,
                    "secondary_section": secondary,
                    "primary_paragraphs": primary_paragraphs,
                    "secondary_paragraphs": [str(paragraph)],
                    "treatment": "excessive",
                    "necessary": False,
                    "concise": False,
                    "material": material,
                    "contradiction": False,
                    "obscures_required_information": False,
                    "materially_unusable": False,
                    "target_ids": [secondary],
                    "repair_action": "redraft_secondary_section",
                    "disposition": "manual_review",
                    "severity": "warning",
                    "publication_disposition": "warning",
                    "action": "manual_review",
                })
                break
    return findings


def deterministic_content_check(revision_dir: Path, reference: Mapping[str, Any]) -> list[dict[str, Any]]:
    """Check only file usability. AI reviewers own generated meaning and layout."""
    findings = []
    for name in document_set(get_path(reference, "meta.study_type")):
        path = revision_dir / "candidate" / name
        if name.endswith(".docx"):
            findings.extend(audit_docx(path))
    return findings


def content_review_sections(reference: Mapping[str, Any]) -> list[dict[str, Any]]:
    branch = canonical_study_type(get_path(reference, "meta.study_type")) or ""
    sections = [
        {
            "artifact": "protocol",
            "section_id": section.section_id,
            "number": section.number,
            "title": section.title,
            "content_expectations": list(section.content_expectations),
        }
        for section in protocol_contract(branch)
        if section_applies(reference, section)
    ]
    if branch != "Retrospective":
        choice = str(get_path(reference, "meta.icf_template", "Advarra"))
        sections.extend({"artifact": "icf", "section_id": section.section_id, "number": section.number,
                         "title": section.title, "content_expectations": list(section.content_expectations)}
                        for section in icf_contract(branch, choice))
        sections.extend({"artifact": "icf", "section_id": section_id, "number": "", "title": title} for section_id, title in icf_retained_sections(branch, choice))
        sections.extend((
            {"artifact": "study.xml", "section_id": "prs.brief-summary", "number": "", "title": "Brief Summary"},
            {"artifact": "study.xml", "section_id": "prs.detailed-description", "number": "", "title": "Detailed Description"},
            {"artifact": "study.xml", "section_id": "prs.structured", "number": "", "title": "Structured PRS Fields"},
        ))
    sections.extend({"artifact": artifact, "section_id": f"{artifact}.shell", "number": "",
                     "title": "Document shell, running headers and footers"}
                    for artifact in ("protocol", "icf")
                    if artifact == "protocol" or branch != "Retrospective")
    return sections


def semantic_evidence_inventory(reference: Mapping[str, Any]) -> list[dict[str, Any]]:
    """Bind shared narrative obligations to the existing full content review."""
    branch = canonical_study_type(get_path(reference, "meta.study_type")) or ""
    sections = list(protocol_contract(branch))
    if branch != "Retrospective":
        sections.extend(icf_contract(branch, str(get_path(reference, "meta.icf_template", "Advarra"))))
    records: dict[str, dict[str, Any]] = {}
    for section in sections:
        if not section_applies(reference, section) or section.section_id == "introduction":
            continue
        for path in section.evidence:
            obligation = semantic_evidence_contract(path, get_path(reference, path))
            if not obligation:
                continue
            record = records.setdefault(path, {
                **obligation, "target_ids": [],
                "scope": "Follow section concept ownership: complete facts in their owner sections, concise references or summaries elsewhere; preserve qualifications wherever an expectation is asserted.",
            })
            record["target_ids"].append(section.section_id)
    return list(records.values())


def source_routing_obligations(reference: Mapping[str, Any]) -> list[dict[str, Any]]:
    """Audit raw evidence against destinations, independent of accepted projections."""
    branch = canonical_study_type(get_path(reference, "meta.study_type"))
    obligations = []
    procedures = {path: get_path(reference, path) for path in (
        "procedures.assessments", "procedures.evaluation", "procedures.visit_schedule",
        "procedures.visit_schedule_table") if meaningful(get_path(reference, path))}
    targets = ([('protocol', 'study-procedure.enrollment')] if branch == 'Retrospective'
               else [('protocol', 'study-procedure.visits'), ('protocol', 'evaluation-procedures'),
                     ('icf', 'icf.procedures')])
    if procedures:
        obligations.extend({"obligation_id": f"visit-relationships:{artifact}:{section}",
                            "kind": "visit_relationships", "artifact": artifact,
                            "section_id": section, "source_values": procedures}
                           for artifact, section in targets)
    benefit_targets = ([('protocol', 'ethics')] if branch == 'Retrospective'
                       else [('protocol', 'risks-benefits.benefits'), ('icf', 'icf.benefits')])
    for artifact, section in benefit_targets:
        if benefit_context_contract(reference, section):
            obligations.append({"obligation_id": f"benefit-context:{artifact}:{section}",
                                "kind": "benefit_context", "artifact": artifact,
                                "section_id": section, "source_values": {
                                    "study.background": get_path(reference, "study.background"),
                                    "risks_benefits.benefits": get_path(reference, "risks_benefits.benefits")}})
    return obligations


def _bound_source_routing_obligations(revision_dir: Path, reference: Mapping[str, Any]) -> list[dict[str, Any]]:
    obligations = []
    for item in source_routing_obligations(reference):
        path = revision_dir / f"candidate/{item['artifact']}.docx"
        if path.is_file():
            passage = rendered_section_snapshot(path, item["section_id"])
            obligations.append({**item, "artifact_passage": passage})
    return obligations


def source_routing_assessment_errors(request: Mapping[str, Any], response: Mapping[str, Any]) -> list[str]:
    """Require explicit, source- and destination-grounded independent decisions."""
    expected = {row["obligation_id"]: row for row in request.get("source_routing_obligations", [])}
    rows = response.get("source_routing_assessments", [])
    if (not isinstance(rows, list) or any(not isinstance(row, Mapping) for row in rows)
            or any(not isinstance(row.get("obligation_id"), str) for row in rows)
            or len(rows) != len(expected) or {row.get("obligation_id") for row in rows} != set(expected)):
        return ["Every source-routing obligation needs exactly one independent assessment."]
    errors = []
    for row in rows:
        obligation = expected[row["obligation_id"]]
        passage = _review_evidence_text(obligation["artifact_passage"])
        facts = row.get("facts")
        if (row.get("status") not in {"passed", "failed"}
                or not isinstance(row.get("rationale"), str) or not row["rationale"].strip()
                or not isinstance(facts, list) or not facts):
            errors.append(f"{row['obligation_id']} requires status, rationale and a fact-by-fact routing audit.")
            continue
        for fact in facts:
            if not isinstance(fact, Mapping):
                errors.append("Source-routing facts must be objects.")
                continue
            excerpt, observed = fact.get("source_excerpt"), fact.get("affected_passage")
            source_path = fact.get("source_path")
            if (not isinstance(source_path, str) or source_path not in obligation["source_values"]
                    or not isinstance(excerpt, str) or not excerpt.strip()
                    or not any(excerpt in text for text in source_text_values(obligation["source_values"].get(source_path)))
                    or not isinstance(observed, str) or not observed.strip() or observed not in passage
                    or fact.get("disposition") not in {"preserved", "not_applicable", "defect"}):
                errors.append(f"{row['obligation_id']} must quote exact raw source and this destination's actual passage.")
        reported = response.get("findings")
        sections = response.get("section_assessments")
        affected = [item for item in reported if isinstance(item, Mapping)
                    and isinstance(item.get("target_ids"), list)
                    and item.get("artifact") == obligation["artifact"]
                    and obligation["section_id"] in item["target_ids"]] if isinstance(reported, list) else []
        companion = any(not reviewed_presentation_only(item) for item in affected)
        presentation_only = bool(affected) and not companion
        section_passed = any(item.get("artifact") == obligation["artifact"]
                             and item.get("section_id") == obligation["section_id"]
                             and (item.get("status") == "passed"
                                  or (item.get("status") in ("blocked", "failed") and presentation_only))
                             for item in sections if isinstance(item, Mapping)) if isinstance(sections, list) else False
        defects = any(isinstance(fact, Mapping) and fact.get("disposition") == "defect" for fact in facts)
        if (row["status"] == "passed" and (not section_passed or defects or companion)
                or row["status"] == "failed" and (not companion or not defects)):
            errors.append(f"{row['obligation_id']} contradicts the section decision or lacks its repair finding.")
    return errors


def _source_field_inventory(reference: Mapping[str, Any]) -> list[dict[str, Any]]:
    """Inventory supplied leaves directly, independent of drafting/table mappings."""
    inventory: list[dict[str, Any]] = []

    def visit(value: Any, path: str) -> None:
        if isinstance(value, Mapping):
            for key, child in value.items():
                visit(child, f"{path}.{key}" if path else str(key))
        elif isinstance(value, list):
            for index, child in enumerate(value):
                visit(child, f"{path}[{index}]")
        elif meaningful(value):
            inventory.append({"source_path": path, "value": value})

    visit(reference, "")
    return inventory


def verification_recovery_request_findings(
    revision_dir: Path,
    request_id: str,
    *,
    expected_task: str,
    current_review_set: int,
    reference: Mapping[str, Any],
    render_report: Mapping[str, Any],
    bound_artifact: str = "",
    section_target_ids: Iterable[str] = (),
    artifact_path: str | None = None,
    allow_missing_response: bool = False,
    _validate_companions: bool = True,
    originating_finding: Mapping[str, Any] | None = None,
) -> list[dict[str, str]]:
    """Own canonical verification authority for one proposed recovery route."""
    request_root = revision_dir / "hermes/verification-requests"
    request_paths = sorted(request_root.iterdir()) if request_root.is_dir() else []
    records: list[tuple[Path, dict[str, Any]]] = []
    parsed_records: list[tuple[Path, dict[str, Any]]] = []
    all_requests: list[dict[str, Any]] = []
    counts: dict[str, int] = {}
    malformed_inventory = False
    for path in request_paths:
        if path.suffix != ".json" or path.is_symlink() or not path.is_file():
            malformed_inventory = True
            continue
        try:
            request = _json(path)
        except (OSError, ValueError, json.JSONDecodeError):
            malformed_inventory = True
            continue
        identity = str(request.get("request_id") or "")
        parsed_records.append((path, request))
        all_requests.append(request)
        counts[identity] = counts.get(identity, 0) + 1
        if identity == request_id:
            records.append((path, request))

    def issue(code: str, text: str) -> list[dict[str, str]]:
        return [{"code": code, "issue": text}]

    if malformed_inventory:
        return issue("verification_request_invalid_inventory", "The active verification request inventory contains an unreadable or non-object entry.")
    required_bindings = ("request_id", "task", "revision_id", "review_set", "response_path")
    for path, item in parsed_records:
        if str(item.get("request_id") or "") == request_id:
            continue
        if (
            any(key not in item or item.get(key) in (None, "") for key in required_bindings)
            or item.get("task") not in {
                "clinical_content_verification", "rendered_page_visual_verification",
            }
            or verification_request_authentication_findings(
                revision_dir, path, item, request_id_counts=counts,
            )
        ):
            return issue("verification_request_invalid_inventory", "The active verification request inventory contains a noncanonical or unauthenticated entry.")
        if _validate_companions and verification_recovery_request_findings(
            revision_dir,
            str(item.get("request_id") or ""),
            expected_task=str(item.get("task") or ""),
            current_review_set=current_review_set,
            reference=reference,
            render_report=render_report,
            allow_missing_response=True,
            _validate_companions=False,
        ):
            return issue("verification_request_invalid_inventory", "The active verification request inventory contains a task-incomplete companion request.")
    if not records:
        return issue("verification_request_not_found", "The declared verification request identity does not resolve to a current request.")
    if len(records) != 1 or counts.get(request_id, 0) != 1:
        return issue("duplicate_verification_request_id", "The declared verification request identity resolves to more than one current request.")
    request_path, request = records[0]
    authentication = verification_request_authentication_findings(
        revision_dir, request_path, request, request_id_counts=counts,
    )
    if authentication:
        code = str(authentication[0].get("code") or "verification_request_canonical_mismatch")
        return issue(code, str(authentication[0].get("issue") or "Verification request authentication failed."))
    if request.get("review_set") != current_review_set:
        return issue("verification_request_stale_review_set", "The declared verification request review set does not match the current generation review set.")
    if request.get("task") != expected_task:
        return issue("verification_request_wrong_task", "The declared verification request task does not match the normalized recovery route.")

    target_ids = {str(item) for item in section_target_ids}
    if expected_task == "clinical_content_verification":
        expected_id = f"{revision_dir.name}.review-{current_review_set}.verify.content"
        expected_sections = [
            dict(item) for item in content_review_sections(reference)
        ]
        expected_artifacts = [
            {"artifact": name if name.endswith(".xml") else Path(name).stem, "path": f"candidate/{name}"}
            for name in sorted(document_set(get_path(reference, "meta.study_type")))
        ]
        branch = canonical_study_type(get_path(reference, "meta.study_type")) or ""
        expected_cross_document_checks = [] if branch == "Retrospective" else list(CROSS_DOCUMENT_CHECKS)
        content_requests = [item for item in all_requests if item.get("task") == expected_task]
        if (
            len(content_requests) != 1
            or request.get("request_id") != expected_id
            or [dict(item) for item in request.get("sections", []) if isinstance(item, Mapping)]
            != expected_sections
            or [
                {"artifact": str(item.get("artifact") or ""), "path": str(item.get("path") or "")}
                for item in request.get("artifacts", []) if isinstance(item, Mapping)
            ] != expected_artifacts
            or list(request.get("checks") or []) != list(CONTENT_CHECKS)
            or list(request.get("cross_document_checks") or []) != expected_cross_document_checks
            or request.get("semantic_evidence") != semantic_evidence_inventory(reference)
            or request.get("approved_source") != reference
        ):
            return issue("verification_request_incomplete_scope", "The declared content verification request is not the unique complete package-wide request for the current branch and review set.")
        for artifact in request.get("artifacts", []):
            if not isinstance(artifact, Mapping):
                return issue("verification_request_stale_artifact", "The content verification artifact binding is malformed.")
            path = revision_dir / str(artifact.get("path") or "")
            exact_matches = (
                path.is_file()
                and _revision_owned_nonsymlink_path(revision_dir, path)
                and sha256_file(path) == artifact.get("sha256")
            )
            content_matches = (
                not artifact.get("content_sha256")
                or path.is_file()
                and _content_sha256(path) == artifact.get("content_sha256")
            )
            if not exact_matches or not content_matches:
                return issue("verification_request_stale_artifact", "The content verification request is not bound to the current exact candidate bytes.")
        if request.get('source_routing_obligations') != _bound_source_routing_obligations(revision_dir, reference):
            return issue('verification_request_incomplete_scope', 'The content review does not bind the current source-routing obligations.')
    else:
        expected_artifacts = {
            str(item.get("artifact") or ""): {
                key: value for key, value in item.items()
                if key not in {"renderer", "page_renderer"}
            }
            for item in render_report.get("artifacts", []) if isinstance(item, Mapping)
        }
        request_artifacts = [dict(item) for item in request.get("artifacts", []) if isinstance(item, Mapping)]
        artifact_ids = {str(item.get("artifact") or "") for item in request_artifacts if item.get("artifact")}
        artifact = next(iter(artifact_ids)) if len(artifact_ids) == 1 else ""
        slug = re.sub(r"[^a-z0-9]+", "-", artifact.casefold()).strip("-")
        expected_id = f"{revision_dir.name}.review-{current_review_set}.verify.visual.{slug}" if slug else ""
        visual_matches = []
        for path in request_paths:
            try:
                item = _json(path)
            except (OSError, ValueError, json.JSONDecodeError):
                continue
            if item.get("task") == expected_task and {
                str(value.get("artifact") or "")
                for value in item.get("artifacts", []) if isinstance(value, Mapping)
            } == {artifact}:
                visual_matches.append(item)
        if (
            not artifact
            or artifact not in expected_artifacts
            or len(visual_matches) != 1
            or request.get("request_id") != expected_id
            or list(request.get("checks") or []) != list(VISUAL_CHECKS)
            or request_artifacts != [expected_artifacts[artifact]]
        ):
            return issue("verification_request_incomplete_scope", "The declared visual verification request is not the unique complete document-scoped request for the current render and review set.")
        for visual_artifact in request_artifacts:
            for key in ("docx", "pdf"):
                path = revision_dir / str(visual_artifact.get(key) or "")
                if (
                    not path.is_file()
                    or not _revision_owned_nonsymlink_path(revision_dir, path)
                    or sha256_file(path) != visual_artifact.get(f"{key}_sha256")
                ):
                    return issue("verification_request_stale_artifact", "The visual verification request is not bound to the current exact DOCX/PDF bytes.")
            for page in visual_artifact.get("pages", []):
                if not isinstance(page, Mapping):
                    return issue("verification_request_stale_artifact", "The visual verification page binding is malformed.")
                path = revision_dir / str(page.get("path") or "")
                if (
                    not path.is_file()
                    or not _revision_owned_nonsymlink_path(revision_dir, path)
                    or sha256_file(path) != page.get("sha256")
                ):
                    return issue("verification_request_stale_artifact", "The visual verification request is not bound to the current exact page-image bytes.")

    canonical_target_artifacts = {
        "icf" if target.startswith("icf.") else "study.xml" if target.startswith("prs.") else "protocol"
        for target in target_ids
    }
    if bound_artifact and canonical_target_artifacts and canonical_target_artifacts != {bound_artifact}:
        return issue("recovery_target_wrong_artifact", "The finding's exact artifact does not own every normalized section recovery target.")
    if target_ids:
        def target_artifact(target: str) -> str:
            if target.startswith("icf."):
                return "icf"
            if target.startswith("prs."):
                return "study.xml"
            return "protocol"

        section_artifacts = {
            target: [
                str(item.get("artifact") or "")
                for item in request.get("sections", [])
                if isinstance(item, Mapping) and str(item.get("section_id") or "") == target
            ]
            for target in target_ids
        }
        if any(values != [target_artifact(target)] for target, values in section_artifacts.items()):
            return issue("recovery_target_wrong_artifact", "The authenticated content request does not bind every normalized section target to its canonical artifact.")
    request_artifact_ids = {
        str(item.get("artifact") or "")
        for item in request.get("artifacts", []) if isinstance(item, Mapping)
    }
    if bound_artifact and not target_ids and bound_artifact not in request_artifact_ids:
        return issue("verification_request_wrong_artifact", "The declared verification request is not bound to the finding's exact artifact.")
    if artifact_path is not None:
        path_artifacts = {bound_artifact} if bound_artifact else canonical_target_artifacts
        bound_paths = {
            str(value)
            for item in request.get("artifacts", []) if isinstance(item, Mapping)
            and (
                not path_artifacts
                or str(item.get("artifact") or "") in path_artifacts
            )
            for value in [item.get("path"), item.get("docx"), item.get("pdf")]
            if isinstance(value, str) and value
        }
        bound_paths.update(
            str(page.get("path"))
            for item in request.get("artifacts", []) if isinstance(item, Mapping)
            and (
                not path_artifacts
                or str(item.get("artifact") or "") in path_artifacts
            )
            for page in item.get("pages", []) if isinstance(page, Mapping) and page.get("path")
        )
        if not artifact_path or artifact_path not in bound_paths:
            return issue("verification_request_wrong_artifact_path", "The declared verification request does not contain the finding's exact artifact path.")
    if _validate_companions:
        expected_inventory = {
            f"{revision_dir.name}.review-{current_review_set}.verify.content",
            *{
                f"{revision_dir.name}.review-{current_review_set}.verify.visual."
                f"{re.sub(r'[^a-z0-9]+', '-', str(item.get('artifact') or '').casefold()).strip('-')}"
                for item in render_report.get("artifacts", [])
                if isinstance(item, Mapping) and str(item.get("artifact") or "")
            },
        }
        if set(counts) != expected_inventory:
            return issue("verification_request_invalid_inventory", "The active verification request inventory is not the complete expected content-and-visual request set.")
    response_path = revision_dir / str(request.get("response_path") or "")
    if not response_path.is_file():
        return [] if allow_missing_response else issue(
            "verification_response_missing",
            "Verification-derived recovery has no corresponding authenticated response.",
        )
    if not _revision_owned_nonsymlink_path(revision_dir, response_path):
        return issue("verification_response_invalid", "Verification response path is not revision-owned canonical evidence.")
    try:
        response = _json(response_path)
    except (OSError, ValueError, json.JSONDecodeError):
        return issue("verification_response_invalid", "Verification-derived recovery response is unreadable or malformed.")
    if any(
        response.get(key) != expected
        for key, expected in (
            ("schema_version", RESPONSE_SCHEMA),
            ("request_id", request.get("request_id")),
            ("request_sha256", request.get("request_sha256")),
            ("task", request.get("task")),
            ("revision_id", request.get("revision_id")),
        )
    ):
        return issue("verification_response_invalid", "Verification-derived recovery response does not bind the canonical request.")
    producer = response.get("producer") if isinstance(response.get("producer"), Mapping) else {}
    if not _text(producer.get("model_id")) or not _text(producer.get("reviewer_id")):
        return issue("verification_response_invalid", "Verification-derived recovery response lacks producer and reviewer identity.")
    if originating_finding is not None and originating_finding.get("recovery_class") in {
        "drafting_defect", "visual_defect", "deterministic_structure_defect",
    }:
        optional = internal_presentation_note(originating_finding)
        if str(response.get("status") or "").casefold() not in ({"passed", "blocked", "failed"} if optional else {"blocked", "failed"}):
            return issue("verification_response_invalid", "A substantive verification-derived recovery requires a blocked or failed response.")
        match_keys = (
            "issue", "target_ids", "artifact", "page", "check", "element", "location", "missing_content_anchor", "code",
            "category", "recommended_action",
        )
        expected = {
            key: originating_finding.get(key)
            for key in match_keys if key in originating_finding
        }
        response_findings = response.get("findings") if isinstance(response.get("findings"), list) else []
        source_finding = originating_finding.get("verification_source_finding")
        if isinstance(source_finding, Mapping):
            finding_matches = any(
                isinstance(item, Mapping) and dict(item) == dict(source_finding)
                for item in response_findings
            )
        else:
            finding_matches = any(
                isinstance(item, Mapping)
                and all(item.get(key) == value for key, value in expected.items())
                for item in response_findings
            )
        if not finding_matches:
            return issue("verification_response_finding_mismatch", "The proposed recovery finding is not present in the exact authenticated response.")
    return []


def _assert_safe_mutation_path(path: Path, *, directory: bool | None = None) -> None:
    """Reject symlinks and unexpected occupied types in every path component."""
    absolute = path.absolute()
    current = Path(absolute.anchor)
    for part in absolute.parts[1:]:
        current = current / part
        if not os.path.lexists(current):
            continue
        if current.is_symlink():
            raise ValueError(f"Verification request path is unsafe: {current}")
        if current == absolute and directory is not None:
            valid = current.is_dir() if directory else current.is_file()
            if not valid:
                raise ValueError(f"Verification request path is unsafe: {current}")
        elif current != absolute and not current.is_dir():
            raise ValueError(f"Verification request path is unsafe: {current}")


def _native_inspection_views(revision_dir: Path, artifacts: list[dict[str, Any]], review_set: int) -> list[dict[str, Any]]:
    """Expose complete native XML through the worker's ordinary read-file path."""
    views = []
    for artifact in artifacts:
        native = revision_dir / artifact["path"]
        if native.suffix.casefold() not in {".docx", ".xml"}:
            continue
        if native.suffix.casefold() == ".docx":
            with zipfile.ZipFile(native) as archive:
                if archive.testzip() is not None:
                    raise ValueError(f"Unreadable native package: {native.name}")
                parts = {name: archive.read(name).decode("utf-8-sig")
                         for name in archive.namelist() if name.endswith((".xml", ".rels"))}
        else:
            parts = {native.name: native.read_text(encoding="utf-8-sig")}
        readable_parts = {}
        for name, xml in parts.items():
            root = ET.fromstring(xml.encode("utf-8"), parser=ET.XMLParser(resolve_entities=False, no_network=True))
            # Word commonly stores a whole part on one enormous line. Expose
            # XML as short lines so read_file pagination cannot omit its tail.
            readable_parts[name] = ET.tostring(root, encoding="unicode", pretty_print=True).splitlines()
        path = revision_dir / f"hermes/native-inspection/review-{review_set}/{native.name}.json"
        payload = {"native_path": artifact["path"], "native_sha256": artifact["sha256"],
                   "xml_parts": readable_parts}
        _assert_safe_mutation_path(path, directory=False)
        if path.exists() and _json(path) != payload:
            raise ValueError("Native inspection identity collision; start a fresh review set.")
        _write(path, payload)
        views.append({"artifact": artifact["artifact"], **{key: payload[key] for key in ("native_path", "native_sha256")},
                      "path": path.relative_to(revision_dir).as_posix(), "sha256": sha256_file(path)})
    return views


def retain_verification_response(revision_dir: Path, request_path: Path) -> None:
    """Keep exact rejected bytes before retrying the immutable review request."""
    request = _json(request_path)
    response_path = revision_dir / str(request["response_path"])
    if not _revision_owned_nonsymlink_path(revision_dir, response_path):
        raise ValueError("Rejected verification response must be revision-owned.")
    _assert_safe_mutation_path(response_path, directory=False)
    if not response_path.is_file():
        return
    path = revision_dir / f"hermes/verification-rejections/{sha256_file(response_path)}.json"
    _assert_safe_mutation_path(path, directory=False)
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(response_path.read_bytes())



def _has_visual_concern(context: Mapping[str, Any], artifact: str) -> bool:
    for prior in context.get("prior_findings", []):
        if not isinstance(prior, Mapping):
            return True
        concern = review_concern(prior)
        if (concern.get("category") == "visual" or concern.get("check") in VISUAL_CHECKS):
            if not concern.get("artifact") or concern["artifact"] == artifact:
                return True
    return False


def _reuse_visual_request(revision_dir: Path, proposed: Mapping[str, Any]) -> dict[str, Any] | None:
    """Retain original independently passed evidence; never manufacture a new review."""
    if proposed.get("task") != "rendered_page_visual_verification":
        return None
    artifacts = proposed.get("artifacts") or []
    if len(artifacts) != 1 or not isinstance(artifacts[0], Mapping):
        return None
    artifact = str(artifacts[0].get("artifact") or "")
    if _has_visual_concern(proposed.get("review_context") or {}, artifact):
        return None
    bindings = ("task", "revision_id", "source_identity_sha256", "generation_authority_sha256",
                "artifacts", "renderer", "page_renderer", "checks", "contracted_template_bundle")
    for path in sorted((revision_dir / "hermes/verification-requests").glob("*.json")):
        if path.is_symlink():
            continue
        prior = _json(path)
        if (prior.get("task") != proposed["task"]
                or not isinstance(prior.get("review_set"), int)
                or not 1 <= prior["review_set"] < proposed["review_set"]
                or not prior.get("source_identity_sha256")
                or not prior.get("generation_authority_sha256")
                or any(prior.get(key) != proposed.get(key) for key in bindings)):
            continue
        response_path = revision_dir / str(prior.get("response_path") or "")
        if not _revision_owned_nonsymlink_path(revision_dir, response_path):
            continue
        response = _json(response_path)
        if (response.get("status") == "passed" and not response.get("findings")
                and not _prior_concern_review_errors(proposed, response)
                and verification_response_is_complete(revision_dir, path)):
            return prior
    return None

def _bounded_review_view(revision_dir: Path, directory: Path, name: str, value: Any) -> dict[str, Any]:
    """Lossless reading alternative: short JSON strings in small immutable files."""
    text = json.dumps(value, ensure_ascii=False, indent=2)
    # UTF-8 and JSON escaping can expand characters. These bounds also cover
    # non-ASCII sources and long XML text nodes without relying on line paging.
    chunks = [text[index:index + 64] for index in range(0, len(text), 64)]
    batches, batch, byte_count = [], [], 32
    for chunk in chunks:
        encoded_size = len(json.dumps(chunk, ensure_ascii=False).encode("utf-8")) + 6
        if batch and byte_count + encoded_size > 10000:
            batches.append(batch)
            batch, byte_count = [], 32
        batch.append(chunk)
        byte_count += encoded_size
    if batch:
        batches.append(batch)
    packets = []
    for batch in batches:
        payload = {"text_chunks": batch}
        data = (json.dumps(payload, ensure_ascii=False, indent=2) + "\n").encode("utf-8")
        digest = hashlib.sha256(data).hexdigest()
        path = directory / f"{digest}.json"
        _assert_safe_mutation_path(path, directory=False)
        if path.exists() and path.read_bytes() != data:
            raise ValueError("Review reading packet identity collision.")
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_bytes(data)
        packets.append({"path": path.relative_to(revision_dir).as_posix(), "sha256": digest})
    return {"name": name, "packets": packets}


def _review_transport(revision_dir: Path, request: Mapping[str, Any]) -> dict[str, Any]:
    directory = revision_dir / "hermes/review-packets" / str(request["request_id"])
    views = [_bounded_review_view(revision_dir, directory, "request", request)]
    for native in request.get("native_inspection", []):
        views.append(_bounded_review_view(revision_dir, directory, native["path"],
                                          _json(revision_dir / native["path"])))
    index_path = directory / "index.json"
    index = {"request_id": request["request_id"], "revision_id": request["revision_id"],
             "task": request["task"], "response_path": request["response_path"],
             "instructions_chunks": [request["instructions"][i:i + 64]
                                     for i in range(0, len(request["instructions"]), 64)],
             "views": views, "progress_path": f"hermes/review-progress/{request['request_id']}.json"}
    _assert_safe_mutation_path(index_path, directory=False)
    if index_path.exists() and _json(index_path) != index:
        raise ValueError("Review reading index identity collision.")
    _write(index_path, index)
    progress_parent = revision_dir / "hermes/review-progress"
    _assert_safe_mutation_path(progress_parent, directory=True)
    progress_parent.mkdir(parents=True, exist_ok=True)
    return {
        "schema_version": "bounded-review-reading/v1",
        "index_path": index_path.relative_to(revision_dir).as_posix(),
        "index_sha256": sha256_file(index_path),
        "views": views,
        "progress_path": f"hermes/review-progress/{request['request_id']}.json",
        "reading_rule": "Read packets in order per view. Concatenate text_chunks exactly to recover "
                        "the complete original JSON; no text is summarized or omitted. Use this route "
                        "when whole-file or line reads truncate. Save completed independent assessments "
                        "to progress_path after each small group using the bound response schema and "
                        "identity. Candidate and approved source remain read-only. Read saved progress "
                        "after context compaction; return remaining assessments and findings. A checkpoint "
                        "is progress, never clinical approval. Report substantive defects normally.",
    }


def _readable_native_view(revision_dir: Path, native: Mapping[str, Any]) -> dict[str, Any]:
    """Native clinical text with locations, excluding nonclinical formatting XML."""
    inspection = _json(revision_dir / native["path"])
    records = []
    for name, lines in inspection["xml_parts"].items():
        root = ET.fromstring("\n".join(lines).encode(), parser=ET.XMLParser(resolve_entities=False, no_network=True))
        tree = root.getroottree()
        if str(native["native_path"]).endswith(".xml"):
            records.extend({"location": f"{name}:{tree.getpath(node)}", "value": node.text or "",
                            "attributes": dict(node.attrib)} for node in root.iter() if len(node) == 0)
        elif name.startswith("word/") and any(part in name for part in (
                "document.xml", "header", "footer", "footnotes.xml", "endnotes.xml", "comments.xml")):
            for node in root.iter(qn("w:p")):
                text = "".join(node.xpath(".//w:t/text() | .//w:instrText/text() | .//w:delText/text()",
                                          namespaces={"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}))
                if text:
                    records.append({"location": f"{name}:{tree.getpath(node)}", "text": text})
            for table in root.iter(qn("w:tbl")):
                records.append({"location": f"{name}:{tree.getpath(table)}", "table": [
                    [{"location": f"{name}:{tree.getpath(cell)}",
                      "text": " ".join(cell.xpath(".//w:t/text()", namespaces={"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}))}
                     for cell in row.findall(qn("w:tc"))]
                    for row in table.findall(qn("w:tr"))]})
        elif name.endswith(".rels"):
            records.extend({"location": f"{name}:{tree.getpath(node)}", "relationship": dict(node.attrib)}
                           for node in root if node.get("TargetMode") == "External")
    return {"artifact": native["artifact"], "native_path": native["native_path"],
            "native_sha256": native["native_sha256"], "records": records}


def _content_review_plan(revision_dir: Path, request: Mapping[str, Any], render_report: Mapping[str, Any]) -> dict[str, Any]:
    views = []
    for native in request["native_inspection"]:
        value = _readable_native_view(revision_dir, native)
        path = revision_dir / f"hermes/readable-native/review-{request['review_set']}/{native['artifact']}.json"
        _assert_safe_mutation_path(path, directory=False)
        if path.exists() and _json(path) != value:
            raise ValueError("Readable native evidence identity collision.")
        _write(path, value)
        views.append({"artifact": native["artifact"], "path": path.relative_to(revision_dir).as_posix(),
                      "sha256": sha256_file(path)})
    return {"schema_version": "content-review-scheduler/v1", "max_sections": 3,
            "max_cross_checks": 2, "max_parallel_units": 3, "unit_budget_seconds": 180,
            "max_no_progress_attempts": 2, "readable_native": views,
            "page_count_by_artifact": {r["artifact"]: len(r.get("pages", [])) for r in render_report.get("artifacts", [])}}


def _prior_concern_in_review_scope(prior: Mapping[str, Any], parent: Mapping[str, Any],
                                   sections: Sequence[Mapping[str, Any]], checks: Sequence[str]) -> bool:
    """Give every prior concern an owner, including unknown native locations."""
    concern = review_concern(prior)
    artifact = concern.get("artifact")
    targets = set(concern.get("target_ids") or [])
    localized = {row["section_id"] for row in parent["sections"]
                 if row["section_id"] in targets and (not artifact or row["artifact"] == artifact)}
    if localized:
        return any(row["section_id"] in localized and (not artifact or row["artifact"] == artifact)
                   for row in sections)
    parent_checks = parent["cross_document_checks"]
    if concern.get("check") in parent_checks:
        return concern["check"] in checks
    owner = "prs.structured" if artifact == "study.xml" else f"{artifact}.shell"
    if any(row["section_id"] == owner for row in parent["sections"]):
        return any(row["section_id"] == owner for row in sections)
    if parent_checks:
        return parent_checks[0] in checks
    return any(row["section_id"] == "protocol.shell" for row in sections)


def create_content_review_unit(revision_dir: Path, parent_path: Path, sections: list[dict[str, Any]],
                               cross_checks: list[str], attempt: int) -> Path:
    """Create a small immutable independent task under a canonical parent review."""
    parent = _json(parent_path)
    if verification_request_authentication_findings(revision_dir, parent_path, parent,
            request_id_counts={parent["request_id"]: 1}):
        raise ValueError("Cannot schedule stale or unbound parent review.")
    plan = parent["review_plan"]
    if (not sections and not cross_checks or len(sections) > plan["max_sections"]
            or len(cross_checks) > plan["max_cross_checks"]
            or any(row not in parent["sections"] for row in sections)
            or not set(cross_checks) <= set(parent["cross_document_checks"])):
        raise ValueError("Review unit exceeds its parent scope or capacity.")
    scope = {"sections": sections, "cross_document_checks": cross_checks}
    unit_id = sha256_value(scope)[:16]
    request_id = f"{parent['request_id']}.unit-{unit_id}.a{attempt}"
    branch = parent["approved_source"]["meta"]["study_type"]
    specifications = {s.section_id: s for s in protocol_contract(branch)}
    if branch != "Retrospective":
        specifications.update({s.section_id: s for s in icf_contract(branch,
            parent["approved_source"]["meta"].get("icf_template", "Advarra"))})
    source = parent["approved_source"]
    full_source = bool(cross_checks) or any(not specifications.get(row["section_id"])
                    or not specifications[row["section_id"]].evidence for row in sections)
    scoped_source = copy.deepcopy(source) if full_source else {"meta": copy.deepcopy(source["meta"])}
    if not full_source:
        from contracts import set_path
        for row in sections:
            spec = specifications[row["section_id"]]
            for source_path in (*spec.evidence, *spec.fidelity_evidence, *spec.table_evidence):
                value = get_path(source, source_path)
                if value is not None:
                    set_path(scoped_source, source_path, copy.deepcopy(value))
    passages = []
    for row in sections:
        native = next((n for n in parent["artifacts"] if n["artifact"] == row["artifact"]), None)
        if not native:
            continue
        path = revision_dir / native["path"]
        if row["artifact"] == "study.xml" or row["section_id"].endswith(".shell"):
            view = next(v for v in plan["readable_native"] if v["artifact"] == row["artifact"])
            passage = {"staged_path": view["path"], "sha256": view["sha256"],
                       "reading_rule": "Read the complete native view in bounded groups; resolve all locations before judgment."}
        else:
            passage = rendered_section_snapshot(path, row["section_id"])
        passages.append({"artifact": row["artifact"], "section_id": row["section_id"],
                         "native_path": native["path"], "native_sha256": native["sha256"], "passage": passage})
    if cross_checks:
        passages = copy.deepcopy(plan["readable_native"])
    if full_source:
        scoped_source = {"meta": copy.deepcopy(source["meta"]),
            "study": {key: copy.deepcopy(source["study"][key]) for key in ('title', 'short_title') if key in source['study']}}
    context = copy.deepcopy(parent.get("review_context") or {})
    context["prior_findings"] = [finding for finding in context.get("prior_findings", [])
        if _prior_concern_in_review_scope(finding, parent, sections, cross_checks)]
    unit = {key: copy.deepcopy(parent[key]) for key in (
        "schema_version", "task", "revision_id", "review_set", "source_identity_sha256",
        "generation_authority_sha256", "artifacts", "checks", "template_preservation", "reviewer_policy")}
    unit.update(request_id=request_id, parent_request_id=parent["request_id"],
        parent_request_sha256=parent["request_sha256"], review_unit_id=unit_id, **scope,
        canonical_source_evidence={"parent_request_path": parent_path.relative_to(revision_dir).as_posix(),
            "parent_request_sha256": parent["request_sha256"],
            "bounded_index_path": parent["review_transport"]["index_path"],
            "bounded_index_sha256": parent["review_transport"]["index_sha256"]},
        approved_source=scoped_source, source_field_inventory=[{"source_path": row["source_path"]}
            for row in _source_field_inventory(scoped_source)],
        readable_native_evidence=passages, native_inspection=[],
        native_inspection_fallback=parent["native_inspection"],
        review_context=context, authorized_boilerplate=({"parent_request_path": parent_path.relative_to(revision_dir).as_posix(),
            "field": "authorized_boilerplate"} if full_source else parent["authorized_boilerplate"]),
        contracted_template_bundle=parent["contracted_template_bundle"],
        review_rules=parent["instructions"], sterling_clause_contract=({"parent_request_path": parent_path.relative_to(revision_dir).as_posix(),
            "field": "sterling_clause_contract"} if parent.get("sterling_clause_contract")
            and (cross_checks or any(row["artifact"] == "icf" for row in sections)) else None),
        response_path=f"hermes/content-review-responses/{request_id}.json",
        instructions="Complete ONLY the assigned sections and cross-document checks. Read readable_native_evidence "
          "first: inline passages or staged path/sha256 references are extracted from exact native files. "
          "Read staged references in bounded groups; do not infer a pass from the index. Full approved source "
          "is in canonical_source_evidence.parent_request_path and its lossless bounded index. Read any staged "
          "authorized_boilerplate and sterling_clause_contract parent fields before judging applicable language. Tables remain structured "
          "and shell/XML evidence includes native locations. Use full native inspection fallback or the canonical "
          "parent approved source if more context is necessary; missing expected content is reviewable. "
          "Do not review unrelated sections or reread all formatting XML. Return explicit independent judgments "
          "and findings for the assigned scope. Return partial completed judgments if interrupted; progress is "
          "durable and unfinished scope is rescheduled. Preserve all supplied clinical facts, qualifications, "
          "consent safeguards and required template/legal language. Ordinary stylistic preferences do not block. "
          "For each substantive finding provide finding_id, artifact, target_ids, check, issue, affected_passage "
          "or missing-content anchor, source_path/source_excerpt and recommended_action. Nonpassing judgments "
          "must identify a finding; never claim completion for uninspected evidence. Independently reassess "
          "prior concerns in review_context, explaining resolutions in assessment notes. " + parent["template_preservation"])
    unit["parent_request_path"] = parent_path.relative_to(revision_dir).as_posix()
    unit["workload"] = {"inline_serialized_bytes": len(json.dumps(unit, ensure_ascii=False).encode()),
        "staged_native_bytes": sum((revision_dir / path).stat().st_size for view in passages
            if (path := view.get("path") or (view.get("passage", {}).get("staged_path")
                if isinstance(view.get("passage"), Mapping) else None))),
        "canonical_parent_bytes": parent_path.stat().st_size,
        "assigned_sections": len(sections), "assigned_cross_checks": len(cross_checks),
        "available_page_count": sum(count for artifact, count in plan.get("page_count_by_artifact", {}).items()
                                    if cross_checks or any(row["artifact"] == artifact for row in sections)),
        "tokens": None, "read_calls": None,
        "measurement_note": "Bytes measure assigned evidence, not actual model consumption. Token/read counts require worker telemetry."}
    unit["review_transport"] = _review_transport(revision_dir, unit)
    unit["request_sha256"] = "0" * 64
    for _ in range(4):
        unit["workload"]["serialized_request_bytes"] = len((json.dumps(unit, ensure_ascii=False, indent=2) + "\n").encode())
    unit["request_sha256"] = verification_request_sha256(unit)
    path = revision_dir / f"hermes/content-review-requests/{request_id}.json"
    ledger = revision_dir / f"hermes/content-review-ledgers/{request_id}.json"
    for target in (path, ledger):
        _assert_safe_mutation_path(target, directory=False)
    if path.exists() and _json(path) != unit:
        raise ValueError("Scheduled review request identity collision.")
    _write(path, unit)
    _write(ledger, {"request_sha256": unit["request_sha256"], "request_file_sha256": sha256_file(path)})
    return path


def _content_assessment_consistency(sections: Sequence[Mapping[str, Any]],
                                    crosses: Sequence[Mapping[str, Any]],
                                    issues: Sequence[Mapping[str, Any]]) -> tuple[dict[Any, str], set[str]]:
    """One disposition rule for unit credit and final package validation."""
    errors, contradictory = {}, set()
    substantive = [f for f in issues if not reviewed_presentation_only(f)]
    for row in [*sections, *crosses]:
        section = "section_id" in row
        key = (row.get("artifact"), row["section_id"]) if section else row["check"]
        covering = [f for f in substantive if
            (section and (not f.get("artifact") or f["artifact"] == row.get("artifact"))
             and row["section_id"] in (f.get("target_ids") or []))
            or (not section and f.get("check") == row["check"])
            or (row["status"] != "passed" and _notes_cite_finding(row.get("notes"), f.get("finding_id", "")))]
        if row["status"] == "passed" and covering:
            errors[key] = f"Reported substantive findings contradict passing assessment {key}."
            contradictory.update(f.get("finding_id", "") for f in covering)
        elif row["status"] != "passed" and not covering:
            errors[key] = f"Every non-passing assessment must be covered by a substantive reported finding; uncovered assessment {key}."
    return errors, contradictory


def _record_review_unit_feedback(revision_dir: Path, unit: Mapping[str, Any], path: Path,
                                 errors: Sequence[str]) -> None:
    """Preserve rejected bytes and precise contract feedback without overwriting attempts."""
    raw = path.read_bytes()
    digest = hashlib.sha256(raw).hexdigest()
    directory = revision_dir / "hermes/review-feedback" / unit["request_id"]
    _assert_safe_mutation_path(directory, directory=True)
    directory.mkdir(parents=True, exist_ok=True)
    raw_path = directory / f"{digest}.bin"
    _assert_safe_mutation_path(raw_path, directory=False)
    raw_path.write_bytes(raw)
    _write(directory / f"{digest}.json", {"response_path": path.relative_to(revision_dir).as_posix(),
        "response_sha256": digest, "errors": list(errors)})


def _retain_content_finding(findings: dict[str, Any], finding: Mapping[str, Any]) -> None:
    """Keep substantive support for retained nonpassing judgments in this review set."""
    identity = finding["finding_id"]
    previous = findings.get(identity)
    if previous is None or reviewed_presentation_only(previous) or not reviewed_presentation_only(finding):
        findings[identity] = dict(finding)


def content_review_unit_result(revision_dir: Path, unit_path: Path, *,
                               _authenticated_parent: Mapping[str, Any] | None = None) -> dict[str, Any] | None:
    """Authenticate completed or partial worker judgments, never infer a pass."""
    try:
        unit = _json(unit_path)
        canonical = revision_dir / f"hermes/content-review-requests/{unit['request_id']}.json"
        ledger_path = revision_dir / f"hermes/content-review-ledgers/{unit['request_id']}.json"
        if unit_path != canonical or not _revision_owned_nonsymlink_path(revision_dir, unit_path):
            return None
        ledger = _json(ledger_path)
        if (not verification_request_hash_valid(unit) or ledger.get("request_sha256") != unit["request_sha256"]
                or ledger.get("request_file_sha256") != sha256_file(unit_path)):
            return None
        parent_path = revision_dir / unit["parent_request_path"]
        if (parent_path != revision_dir / f"hermes/verification-requests/{unit['parent_request_id']}.json"
                or not _revision_owned_nonsymlink_path(revision_dir, parent_path)):
            return None
        parent = _authenticated_parent or _json(parent_path)
        if (parent["request_id"] != unit["parent_request_id"]
                or parent["request_sha256"] != unit["parent_request_sha256"]
                or (_authenticated_parent is None and verification_request_authentication_findings(revision_dir, parent_path, parent,
                    request_id_counts={parent["request_id"]: 1}))):
            return None
        response_paths = [revision_dir / unit["review_transport"]["progress_path"],
                          revision_dir / unit["response_path"]]
        rows, cross, findings, proofs = {}, {}, {}, []
        accepted_producer = {}
        for path in response_paths:
            if not path.is_file() or not _revision_owned_nonsymlink_path(revision_dir, path):
                continue
            try:
                response = _json(path)
                if not isinstance(response, Mapping):
                    raise ValueError("Response must be a JSON object.")
            except (OSError, ValueError) as exc:
                _record_review_unit_feedback(revision_dir, unit, path, [f"Unreadable response: {exc}"])
                continue
            if (response.get("schema_version") != RESPONSE_SCHEMA
                    or any(response.get(key) != unit[key] for key in ("request_id", "request_sha256", "task", "revision_id"))
                    or response.get("status") not in {"passed", "blocked", "failed", "in_progress"}):
                continue
            producer = response.get("producer") or {}
            authors = (parent.get("review_context") or {}).get("repair_authors", [])
            if (not isinstance(producer, Mapping) or not _text(producer.get("model_id"))
                    or not _text(producer.get("reviewer_id"))
                    or any(a.get("reviewer_id") == producer["reviewer_id"] for a in authors)):
                continue
            assessments = response.get("section_assessments", [])
            crosses = response.get("cross_document_assessments", [])
            issues = response.get("findings", [])
            if not all(isinstance(v, list) for v in (assessments, crosses, issues)):
                continue
            expected = {(r["artifact"], r["section_id"]) for r in unit["sections"]}
            if any(not isinstance(r, Mapping) or not isinstance(r.get("artifact"), str) or not isinstance(r.get("section_id"), str)
                   or (r["artifact"], r["section_id"]) not in expected
                   or r.get("status") not in {"passed", "blocked", "failed"}
                   or not isinstance(r.get("checks"), list) or any(not isinstance(c, str) for c in r["checks"])
                   or set(r["checks"]) != set(CONTENT_CHECKS) for r in assessments):
                continue
            if any(not isinstance(r, Mapping) or not isinstance(r.get("check"), str) or r["check"] not in unit["cross_document_checks"]
                   or r.get("status") not in {"passed", "blocked", "failed"} for r in crosses):
                continue
            if any(not isinstance(r, Mapping) or not isinstance(r.get("finding_id"), str)
                   or not _valid_finding_id(r["finding_id"]) or not isinstance(r.get("issue"), str)
                   or not r["issue"].strip() or any(key in r and not isinstance(r[key], str)
                       for key in ("artifact", "check", "category", "surface"))
                   or ("target_ids" in r and (not isinstance(r["target_ids"], list)
                       or any(not isinstance(target, str) for target in r["target_ids"]))) for r in issues):
                continue
            if len({(r['artifact'], r['section_id']) for r in assessments}) != len(assessments) or len({r['check'] for r in crosses}) != len(crosses):
                continue
            if response["status"] in {"blocked", "failed"} and not response.get("error") and not issues and all(r["status"] == "passed" for r in [*assessments, *crosses]):
                continue
            invalid, contradictory = _content_assessment_consistency(assessments, crosses, issues)
            if invalid:
                _record_review_unit_feedback(revision_dir, unit, path, list(invalid.values()))
                assessments = [r for r in assessments if (r["artifact"], r["section_id"]) not in invalid]
                crosses = [r for r in crosses if r["check"] not in invalid]
                issues = [r for r in issues if r["finding_id"] not in contradictory]
            scoped_context = copy.deepcopy(unit.get("review_context") or {})
            scoped_context["prior_findings"] = [f for f in scoped_context.get("prior_findings", [])
                if _prior_concern_in_review_scope(f, parent, assessments, [r["check"] for r in crosses])]
            if _prior_concern_review_errors({**unit, "review_context": scoped_context}, response):
                continue
            if len({r["finding_id"] for r in issues}) != len(issues):
                continue
            proof = {"request": unit_path.relative_to(revision_dir).as_posix(), "request_sha256": sha256_file(unit_path),
                     "response": path.relative_to(revision_dir).as_posix(), "response_sha256": sha256_file(path)}
            proofs.append(proof)
            accepted_producer = producer
            prior_ids = {f.get("finding_id") for f in (parent.get("review_context") or {}).get("prior_findings", [])}
            # New identities include the attempt; explicit retained concerns keep
            # their identity through every unit and subsequent repair cycle.
            prefix = hashlib.sha256(unit["request_id"].encode()).hexdigest()[:16]
            identities = {}
            for issue in issues:
                identity = issue["finding_id"]
                mapped = identity if identity in prior_ids else prefix + "-" + identity
                while mapped in prior_ids and mapped != identity:
                    mapped = "new-" + mapped
                identities[identity] = mapped
            for r in assessments:
                key = (r["artifact"], r["section_id"])
                linked = [identities[issue["finding_id"]] for issue in issues
                          if _notes_cite_finding(r.get("notes"), issue["finding_id"])]
                if linked:
                    r = {**r, "notes": [_text(r.get("notes")), *linked]}
                if key not in rows or r["status"] != "passed": rows[key] = {**r, "producer": producer}
            for r in crosses:
                r = dict(r)
                linked = [identities[issue["finding_id"]] for issue in issues
                          if _notes_cite_finding(r.get("notes"), issue["finding_id"])]
                if linked:
                    r["notes"] = [_text(r.get("notes")), *linked]
                if r["check"] not in cross or r["status"] != "passed": cross[r["check"]] = {**r, "producer": producer}
            for r in issues:
                finding_id = identities[r["finding_id"]]
                _retain_content_finding(findings, {**r, "finding_id": finding_id})
        if not proofs or not (rows or cross or findings):
            return None
        return {"section_assessments": list(rows.values()), "cross_document_assessments": list(cross.values()),
                "findings": list(findings.values()), "proofs": proofs, "producer": accepted_producer,
                "complete": len(rows) == len(unit["sections"]) and len(cross) == len(unit["cross_document_checks"])}
    except (OSError, ValueError, KeyError, TypeError):
        return None


def content_review_schedule_result(revision_dir: Path, parent_path: Path,
                                   unit_paths: Sequence[Path], *, publish: bool = True) -> dict[str, Any]:
    """Assemble exact worker judgments and expose incomplete work separately."""
    parent = _json(parent_path)
    if verification_request_authentication_findings(revision_dir, parent_path, parent,
            request_id_counts={parent["request_id"]: 1}):
        raise ValueError("Review schedule parent evidence changed.")
    rows, cross, findings, proofs = {}, {}, {}, []
    producer = {"model_id": "workflow", "reviewer_id": "review-scheduler"}
    for path in unit_paths:
        result = content_review_unit_result(revision_dir, path, _authenticated_parent=parent)
        if not result:
            continue
        proofs.extend(p for p in result["proofs"] if p not in proofs)
        producer = result["producer"]
        for r in result["section_assessments"]:
            key = (r["artifact"], r["section_id"])
            if key not in rows or r["status"] != "passed": rows[key] = r
            producer = r["producer"]
        for r in result["cross_document_assessments"]:
            if r["check"] not in cross or r["status"] != "passed": cross[r["check"]] = r
            producer = r["producer"]
        for r in result["findings"]:
            _retain_content_finding(findings, r)
    missing = [r for r in parent["sections"] if (r["artifact"], r["section_id"]) not in rows]
    missing_cross = [r for r in parent["cross_document_checks"] if r not in cross]
    complete = not missing and not missing_cross
    response = {key: parent[key] for key in ("request_id", "request_sha256", "task", "revision_id")}
    response.update(schema_version=RESPONSE_SCHEMA, producer=producer,
        section_assessments=list(rows.values()), cross_document_assessments=list(cross.values()),
        findings=list(findings.values()), review_unit_evidence=proofs, scheduler_partial=not complete,
        status="passed" if complete and all(reviewed_presentation_only(f) for f in findings.values()) and all(r["status"] == "passed" for r in [*rows.values(), *cross.values()]) else "blocked")
    if not complete and all(reviewed_presentation_only(f) for f in findings.values()):
        response["error"] = {"type": "native_inspection_incomplete",
                             "message": f"Independent review incomplete: {len(missing)} sections and {len(missing_cross)} cross-document checks remain."}
    response["review_coverage"] = {"remaining_sections": missing, "remaining_cross_document_checks": missing_cross}
    if publish:
        _assert_safe_mutation_path(revision_dir / parent["response_path"], directory=False)
        _write(revision_dir / parent["response_path"], response)
    return {"response": response, "remaining_sections": missing, "remaining_cross_document_checks": missing_cross,
            "completed": len(rows) + len(cross), "complete": complete, "findings": list(findings.values())}


def create_verification_requests(
    revision_dir: Path,
    reference: Mapping[str, Any],
    render_report: Mapping[str, Any],
    *,
    contracted_bundle: Mapping[str, Any] | None = None,
    review_set: int = 1,
    review_context: Mapping[str, Any] | None = None,
) -> list[Path]:
    requests = revision_dir / "hermes/verification-requests"; responses = revision_dir / "hermes/verification-responses"
    content_files = []
    for path in sorted((revision_dir / "candidate").glob("*")):
        if not path.is_file():
            continue
        if not _revision_owned_nonsymlink_path(revision_dir, path):
            raise ValueError(f"Candidate artifact path is not revision-owned canonical evidence: {path}")
        artifact = {
            "artifact": path.name if path.suffix.casefold() == ".xml" else path.stem,
            "path": path.relative_to(revision_dir).as_posix(),
        }
        if path.suffix.casefold() == ".docx":
            artifact["content_sha256"] = _content_sha256(path)
        artifact["sha256"] = sha256_file(path)
        content_files.append(artifact)
    branch = canonical_study_type(get_path(reference, "meta.study_type")) or ""
    sections = content_review_sections(reference)
    repo_root = Path(__file__).resolve().parents[1]
    bundle = dict(contracted_bundle or contracted_template_bundle(repo_root, reference))
    boilerplate_path = repo_root / str(bundle["fixed_clinical_boilerplate"]["path"])
    authorized_boilerplate = _json(boilerplate_path)
    if authorized_boilerplate.get("version") != BOILERPLATE_VERSION:
        raise ValueError("Verification boilerplate does not match the content contract.")
    authorized_boilerplate["sections"] = contextual_clinical_boilerplate(
        authorized_boilerplate["sections"], reference)
    cross_document_checks = [] if branch == "Retrospective" else list(CROSS_DOCUMENT_CHECKS)
    content_request_id = f"{revision_dir.name}.review-{review_set}.verify.content"
    template_preservation = ""
    if branch != "Retrospective" and str(get_path(reference, "meta.icf_template", "")).casefold() == "sterling":
        template_preservation = (
            "Sterling owns its reserved title, protocol number, investigator, site, telephone and sponsor "
            "merge tags. Preserve them unfilled, including «Protocol_No» in repeated headers and the "
            "study-identification block. They are intentional template content, not unfinished fields or "
            "cross-document mismatches. Review study-specific content and usable signing space normally."
        )
    content_instructions = (
        f"This is complete review set {review_set}. Act as an independent read-only reviewer of the "
        "actual hash-bound native files in artifacts, not just accepted drafts or their construction model. "
        "Open every DOCX and inspect body paragraphs, tables including nested cells, headers, footers and "
        "text boxes; parse all XML values and relationships. Assess every listed section and cross-document "
        "check against approved_source, source_field_inventory and applicable authorized template language. "
        "Missing content is reviewable: use its expected section or native path as a missing-content anchor. "
        "Check clinical meaning, substantive omissions, unsupported claims, quantities, units, timing anchors, "
        "cohort totals and links, negations, qualifications, consent applicability, safeguards and consistency "
        "across Protocol, selected ICF and PRS values. XML structural validity does not establish value fidelity. "
        "Reconcile raw visit/assessment relationships and background benefits with actual tables and prose; "
        "Assess table usability independently of fact presence: explicit visit-to-activity associations "
        "must be readable as table cells, using an activity matrix or clearly labeled visit rows. "
        "Do not accept raw field/value dumps as a usable schedule merely because every fact is present. "
        "Unassigned narrative details may remain clearly labeled notes; do not invent visit assignments. "
        "Code-generated tables are subject to this review, not presumed correct. "
        "source_routing_obligations are evidence context, not a required separate adjudication payload. "
        "Review section purpose and content_expectations; identify redundant or self-referential text "
        "that obscures study facts. Preserve applicable consent-copy and signature language and "
        "source-supported boundaries on access to existing records. "
        "Do not infer defects from word overlap, short sections, repeated identifiers or a missing optional "
        "fact. Do not invent methods, policies or clinical requirements. Review sparse but complete input "
        "proportionately; concise faithful paraphrases and negations may pass. Distinguish source-supported "
        "hypotheses from established results; do not strengthen their claims. Review applicable consent and "
        "template language in context, including required legal text and its study-specific substitutions. "
        "For each substantive finding return unique finding_id, artifact, target_ids, check, issue, source_path "
        "and source_excerpt or applicable template evidence, affected_passage or missing-content anchor, "
        "document_locations when useful, and recommended_action. Describe the actual concern and correction "
        "rather than requiring a pre-registered edit surface. Batch findings from the whole package. "
        "Ordinary stylistic preferences are internal notes and require no repair or warning. If reporting a "
        "harmless presentation observation, explicitly set material, safety_critical, contradiction, "
        "obscures_required_information and materially_unusable to false and pass the affected assessments. "
        "Return every requested section_assessment and cross_document_assessment with the requested checks. "
        "Use top-level passed when substantive concerns are resolved; use blocked for substantive defects "
        "and identify their failed assessments. A generic pass cannot conceal an explicitly unresolved "
        "material finding. If a repair author disputes a concern, independently reassess its source evidence "
        "and the current native output; explain why it is resolved or preserve the concern. "
        "Never edit files or approve on the repair author's assertion alone."
        " The native_inspection files contain complete XML parts extracted directly from these exact "
        "native bytes, including headers, footers, text boxes and relationships. Read them using read_file "
        "with pagination as needed; this authorized read-only path needs no terminal, unzip or execute_code. "
        "These are native evidence, not drafting summaries. If this path is also unavailable, return blocked "
        "with error.type native_inspection_incomplete and error.message describing the denied operation; "
        "do not turn a tooling failure into a clinical finding or claim a completed review. "
        "Preserve explicit ethics approval status, qualified claims and supplied citation-to-bibliography links."
    )
    if branch != "Retrospective" and not meaningful(get_path(reference, "procedures.completion")):
        content_instructions += (
            " The approved source provides no participant-completion criterion for Protocol Section 18.1. "
            "Do not treat the visit schedule or exit form as a formal completion rule; flag any affirmative claim that does so. "
            "A concise description of the anticipated end of planned follow-up and a distinction from early "
            "discontinuation are appropriate when grounded in the supplied timeline."
        )
    if str(get_path(reference, "meta.icf_template", "Advarra")).casefold() == "sterling":
        content_instructions += (
            " Verify Sterling's site merge tags at their template positions; check the supplied site "
            "address in the Protocol and PRS location components. Retain its fixed introduction and "
            "INFORMATION clause. The exact registry statement appears only when its approved source trigger "
            "is present. Continue reviewing all study-specific ICF sections against approved evidence."
        )
    payloads = [{
        "schema_version": VERIFY_SCHEMA,
        "request_id": content_request_id,
        "task": "clinical_content_verification",
        "revision_id": revision_dir.name,
        "generation_authority_sha256": sha256_value(governing_resources(repo_root, reference)),
        "review_set": review_set,
        "artifacts": content_files,
        "native_inspection": _native_inspection_views(revision_dir, content_files, review_set),
        "approved_source": reference,
        "source_field_inventory": _source_field_inventory(reference),
        "source_routing_obligations": _bound_source_routing_obligations(revision_dir, reference),
        "semantic_evidence": semantic_evidence_inventory(reference),
        "protocol_concept_ownership": protocol_concept_ownership(branch),
        "authorized_boilerplate": authorized_boilerplate,
        "sterling_clause_contract": (
            sterling_clause_contract(repo_root)
            if str(get_path(reference, "meta.icf_template", "")).casefold() == "sterling"
            else None
        ),
        "sections": sections,
        "code_owned_tables": ({
            "visit-schedule": {"section_id": "study-procedure.visits", "caption": "Table 9.2-1: Visit Schedule"},
            **{name: {"section_id": table["section_id"], "caption": table["caption"]}
               for name, table in protocol_table_contracts(reference).items()},
        } if branch != "Retrospective" else {}),
        "checks": list(CONTENT_CHECKS),
        "cross_document_checks": cross_document_checks,
        "instructions": content_instructions,
        "reviewer_policy": {
            "native_file_inspection_required": True,
            "independent_from_repair_author": True,
            "producer_reviewer_id_required": True,
            "producer_model_id_required": True,
        },
        "response_path": f"hermes/verification-responses/{content_request_id}.json",
    }]
    visual_artifacts = list(render_report.get("artifacts", []))
    visual_batches = [[artifact] for artifact in visual_artifacts] or [[]]
    for index, artifacts in enumerate(visual_batches, start=1):
        artifact_name = str(artifacts[0].get("artifact", "documents")) if artifacts else "documents"
        artifact_id = re.sub(r"[^a-z0-9]+", "-", artifact_name.casefold()).strip("-") or f"document-{index}"
        request_id = f"{revision_dir.name}.review-{review_set}.verify.visual.{artifact_id}"
        request_artifacts = [
            {key: value for key, value in artifact.items() if key not in {"renderer", "page_renderer"}}
            for artifact in artifacts
        ]
        for request_artifact in request_artifacts:
            for key in ("docx", "pdf"):
                bound_path = revision_dir / str(request_artifact.get(key) or "")
                if not bound_path.is_file() or not _revision_owned_nonsymlink_path(revision_dir, bound_path):
                    raise ValueError(f"Visual artifact path is not revision-owned canonical evidence: {bound_path}")
            for page in request_artifact.get("pages", []):
                bound_path = revision_dir / str(page.get("path") or "")
                if not bound_path.is_file() or not _revision_owned_nonsymlink_path(revision_dir, bound_path):
                    raise ValueError(f"Visual page path is not revision-owned canonical evidence: {bound_path}")
        payloads.append({
            "schema_version": VERIFY_SCHEMA,
            "request_id": request_id,
            "task": "rendered_page_visual_verification",
            "revision_id": revision_dir.name,
            "review_set": review_set,
            "renderer": artifacts[0].get("renderer", render_report.get("renderer")) if artifacts else render_report.get("renderer"),
            "page_renderer": artifacts[0].get("page_renderer", render_report.get("page_renderer")) if artifacts else render_report.get("page_renderer"),
            "artifacts": request_artifacts,
            "checks": list(VISUAL_CHECKS),
            "instructions": f"This is complete review set {review_set}. Independently open and inspect every supplied page PNG for this document; file existence, extracted text and prior reviewer conclusions cannot establish visual approval. Assess every requested check on every page for readability, clipping, table usability, heading flow, headers/footers, signing space, template layout and actual TOC destinations. Check that visit-to-activity associations can be scanned in labeled cells; raw field/value dumps are not a usable substitute for a schedule table. Treat lost table usability as an actual defect, not an ordinary stylistic preference; do not assume a code-generated table is correct. Natural pagination and short title, signature or consent pages may be correct. Preserve applicable consent content and usable signing space. Report substantive visual defects with artifact, page, check, element (visible heading/caption or a descriptive missing-content anchor), issue and recommended_action. No pre-registered layout rule is required. Batch actual defects; omit ordinary stylistic preferences. Reassess repaired files from these fresh images, independently of the repair author.",
            "reviewer_policy": {
                "image_inspection_required": True,
                "delegated_failure_fallback": "parent_reviews_the_same_bound_page_images",
                "deterministic_checks_alone_can_pass": False,
                "producer_reviewer_id_required": True,
                "producer_model_id_required": True,
            },
            "response_path": f"hermes/verification-responses/{request_id}.json",
        })
    for payload in payloads:
        payload["source_identity_sha256"] = canonical_evidence_sha256(reference)
        payload["generation_authority_sha256"] = payloads[0]["generation_authority_sha256"]
        payload["template_preservation"] = template_preservation
        payload["instructions"] += " " + template_preservation
        if payload["task"] == "rendered_page_visual_verification":
            payload["instructions"] += " Judge layout from these bound page images and renderer; a different local renderer's pagination alone does not establish a defect."
        if review_context:
            payload["review_context"] = dict(review_context)
            payload["instructions"] += (
                " Independently reassess the prior concerns and repair-author evidence in review_context "
                "against these current native files/page images. Explain each substantive concern's "
                "resolution in the relevant assessment notes or preserve it as a finding; a generic "
                "pass or the repair author's claim alone does not resolve an explicit concern."
            )
        payload["contracted_template_bundle"] = bundle
        if payload["task"] == "clinical_content_verification":
            payload["review_plan"] = _content_review_plan(revision_dir, payload, render_report)
            payload["review_transport"] = _review_transport(revision_dir, payload)
        payload["request_sha256"] = verification_request_sha256(payload)
    payloads = [_reuse_visual_request(revision_dir, payload) or payload for payload in payloads]
    expected = {payload["request_id"]: payload for payload in payloads}
    mutation_roots = (
        revision_dir / "hermes",
        requests,
        responses,
        revision_dir / "request-ledger",
    )
    for root in mutation_roots:
        _assert_safe_mutation_path(root, directory=True)

    proposed_ledgers: dict[str, dict[str, Any]] = {}
    for payload in payloads:
        request_id = str(payload["request_id"])
        serialized = (json.dumps(payload, indent=2, ensure_ascii=False) + "\n").encode("utf-8")
        request_path = requests / f"{request_id}.json"
        ledger_path = revision_dir / "request-ledger" / f"{request_id}.json"
        response_path = responses / f"{request_id}.json"
        if payload.get("response_path") != response_path.relative_to(revision_dir).as_posix():
            raise ValueError(f"Verification request response path is unsafe: {request_id}")
        for path, label in (
            (request_path, "request"),
            (ledger_path, "ledger"),
            (response_path, "response"),
        ):
            try:
                _assert_safe_mutation_path(path, directory=False)
            except ValueError as exc:
                raise ValueError(
                    f"Verification request {label} path is unsafe: {request_id}"
                ) from exc
        ledger = verification_request_ledger_record(
            request_path,
            payload,
            request_bytes=serialized,
        )
        ledger_occupied = os.path.lexists(ledger_path)
        request_occupied = os.path.lexists(request_path)
        if request_occupied and not ledger_occupied:
            raise ValueError(f"Verification request canonical ledger is missing: {request_id}")
        if ledger_occupied and _json(ledger_path) != ledger:
            raise ValueError(f"Verification request identity collision: {request_id}")
        if request_occupied:
            try:
                occupied_request = _json(request_path)
            except (OSError, ValueError, json.JSONDecodeError) as exc:
                raise ValueError(f"Verification request identity collision: {request_id}") from exc
            authentication = verification_request_authentication_findings(
                revision_dir,
                request_path,
                occupied_request,
                request_id_counts={request_id: 1},
            )
            if authentication:
                raise ValueError(f"Verification request identity collision: {request_id}")
        proposed_ledgers[request_id] = ledger

    existing_records: list[tuple[Path, dict[str, Any]]] = []
    if requests.is_dir():
        entries = sorted(requests.iterdir())
        if any(path.is_symlink() or not path.is_file() or path.suffix != ".json" for path in entries):
            raise ValueError("Verification request inventory path is unsafe")
        for path in entries:
            try:
                request = _json(path)
            except (OSError, ValueError, json.JSONDecodeError) as exc:
                raise ValueError("Verification request identity collision: invalid inventory") from exc
            existing_records.append((path, request))
    request_id_counts: dict[str, int] = {}
    for _path, request in existing_records:
        request_id = str(request.get("request_id") or "")
        request_id_counts[request_id] = request_id_counts.get(request_id, 0) + 1

    cleanup: list[tuple[Path, Path]] = []
    for path, request in existing_records:
        request_id = str(request.get("request_id") or "")
        canonical_path = requests / f"{request_id}.json"
        if path != canonical_path:
            raise ValueError(f"Verification request identity collision: {request_id}")
        authentication = verification_request_authentication_findings(
            revision_dir,
            path,
            request,
            request_id_counts=request_id_counts,
        )
        if authentication:
            raise ValueError(f"Verification request identity collision: {request_id}")
        response_path = responses / f"{request_id}.json"
        if request.get("response_path") != response_path.relative_to(revision_dir).as_posix():
            raise ValueError(f"Verification request response path is unsafe: {request_id}")
        _assert_safe_mutation_path(response_path, directory=False)
        replacement = expected.get(request_id)
        if (
            replacement is not None
            and request.get("request_sha256") == replacement["request_sha256"]
        ):
            continue
        cleanup.append((path, response_path))

    requests.mkdir(parents=True, exist_ok=True)
    responses.mkdir(parents=True, exist_ok=True)
    for path, response_path in cleanup:
        response_path.unlink(missing_ok=True)
        path.unlink(missing_ok=True)
    result = []
    for payload in payloads:
        path = requests / f"{payload['request_id']}.json"
        _write(path, payload)
        ledger_path = revision_dir / "request-ledger" / f"{payload['request_id']}.json"
        _write(ledger_path, proposed_ledgers[str(payload["request_id"])])
        result.append(path)
    return result


def content_review_resume(revision_dir: Path, request_path: Path) -> dict[str, Any]:
    """Checkpoint independent judgments; assemble only complete, still-bound evidence.

    A tooling error never becomes a clinical assessment. Original responses are
    retained; only explicit section/check judgments can fill review coverage.
    """
    request = _json(request_path)
    counts: dict[str, int] = {}
    for active_path in (revision_dir / "hermes/verification-requests").glob("*.json"):
        active = _json(active_path)
        active_id = str(active.get("request_id") or "")
        counts[active_id] = counts.get(active_id, 0) + 1
    authentication = verification_request_authentication_findings(
        revision_dir, request_path, request, request_id_counts=counts)
    if authentication:
        raise ValueError(authentication[0]["issue"])
    transport = request.get("review_transport")
    if request.get("task") != "clinical_content_verification" or not isinstance(transport, Mapping):
        raise ValueError("Content review has no supported checkpoint transport.")
    progress_path = revision_dir / str(transport["progress_path"])
    directory = progress_path.with_suffix("")
    response_path = revision_dir / str(request["response_path"])
    for path in (progress_path, response_path):
        if not _revision_owned_nonsymlink_path(revision_dir, path):
            raise ValueError("Review progress must be revision-owned.")
        _assert_safe_mutation_path(path, directory=False)
    _assert_safe_mutation_path(directory, directory=True)
    directory.mkdir(parents=True, exist_ok=True)
    expected_sections = {(r["artifact"], r["section_id"]) for r in request["sections"]}
    expected_cross = set(request.get("cross_document_checks", []))
    identity = ("request_id", "request_sha256", "revision_id", "task")
    context = request.get("review_context") or {}
    prior_authors = context.get("repair_authors", [context.get("repair_author", {})])

    def valid(response: Any) -> bool:
        if not isinstance(response, Mapping) or response.get("schema_version") != RESPONSE_SCHEMA:
            return False
        producer = response.get("producer")
        if (any(response.get(key) != request[key] for key in identity)
                or not isinstance(producer, Mapping)
                or not _text(producer.get("model_id")) or not _text(producer.get("reviewer_id"))
                or any(isinstance(author, Mapping) and author.get("reviewer_id") == producer.get("reviewer_id")
                       for author in prior_authors)):
            return False
        if response.get("status") not in {"passed", "blocked", "failed", "in_progress"}:
            return False
        sections = response.get("section_assessments", [])
        cross = response.get("cross_document_assessments", [])
        findings = response.get("findings", [])
        if not all(isinstance(rows, list) for rows in (sections, cross, findings)):
            return False
        section_ids, cross_ids = [], []
        for row in sections:
            if (not isinstance(row, Mapping) or row.get("status") not in {"passed", "blocked", "failed"}
                    or not isinstance(row.get("checks"), list)
                    or any(not isinstance(check, str) for check in row["checks"])
                    or set(row["checks"]) != set(CONTENT_CHECKS)):
                return False
            if not isinstance(row.get("artifact"), str) or not isinstance(row.get("section_id"), str):
                return False
            section_ids.append((row.get("artifact"), row.get("section_id")))
        for row in cross:
            if not isinstance(row, Mapping) or row.get("status") not in {"passed", "blocked", "failed"}:
                return False
            if not isinstance(row.get("check"), str):
                return False
            cross_ids.append(row["check"])
        if (len(set(section_ids)) != len(section_ids) or not set(section_ids) <= expected_sections
                or len(set(cross_ids)) != len(cross_ids) or not set(cross_ids) <= expected_cross):
            return False
        if (response.get("status") in {"blocked", "failed"} and not response.get("error")
                and not findings and not any(row["status"] != "passed" for row in [*sections, *cross])):
            return False
        return all(isinstance(row, Mapping) and _text(row.get("issue"))
                   and _valid_finding_id(_text(row.get("finding_id"))) for row in findings)

    # Validate identities before retaining anything as reusable review progress.
    for path in (progress_path, response_path):
        if not path.is_file():
            continue
        try:
            raw = path.read_bytes()
            response = json.loads(raw)
        except (OSError, ValueError):
            continue
        if valid(response) and any(response.get(key) for key in (
                "section_assessments", "cross_document_assessments", "findings")):
            snapshot = directory / f"attempt-{hashlib.sha256(raw).hexdigest()}.json"
            _assert_safe_mutation_path(snapshot, directory=False)
            if not snapshot.exists():
                snapshot.write_bytes(raw)
    sections, cross, findings = {}, {}, {}
    evidence, producers = [], []
    for snapshot in sorted(directory.glob("attempt-*.json")):
        if not _revision_owned_nonsymlink_path(revision_dir, snapshot):
            raise ValueError("Review checkpoint evidence escapes its revision.")
        raw = snapshot.read_bytes()
        if snapshot.stem != "attempt-" + hashlib.sha256(raw).hexdigest():
            raise ValueError("Review checkpoint evidence changed.")
        response = json.loads(raw)
        if not valid(response):
            raise ValueError("Review checkpoint binding is invalid.")
        producer = dict(response["producer"])
        producers.append(producer)
        evidence.append({"path": snapshot.relative_to(revision_dir).as_posix(),
                         "sha256": sha256_file(snapshot), "producer": producer})
        for key, rows, destination in (
            (lambda r: (r["artifact"], r["section_id"]), response.get("section_assessments", []), sections),
            (lambda r: r["check"], response.get("cross_document_assessments", []), cross),
        ):
            for row in rows:
                row_key = key(row)
                # An unresolved blocked judgment cannot disappear through a
                # later partial pass. It goes to normal independent repair.
                if row_key not in destination or row["status"] != "passed":
                    destination[row_key] = {**row, "producer": producer}
        for row in response.get("findings", []):
            finding_id = row["finding_id"]
            if finding_id in findings and findings[finding_id] != row:
                raise ValueError("Independent checkpoints have conflicting finding identities.")
            findings[finding_id] = row
    remaining_sections = [r for r in request["sections"] if (r["artifact"], r["section_id"]) not in sections]
    remaining_cross = [r for r in request.get("cross_document_checks", []) if r not in cross]
    complete = not remaining_sections and not remaining_cross
    if complete and producers:
        merged = {key: request[key] for key in identity}
        merged.update(schema_version=RESPONSE_SCHEMA, producer=producers[-1],
                      section_assessments=list(sections.values()), cross_document_assessments=list(cross.values()),
                      findings=list(findings.values()), review_progress_evidence=evidence,
                      status="blocked" if findings or any(r["status"] != "passed" for r in [*sections.values(), *cross.values()]) else "passed")
        # Publish only explicit judgments; the normal complete response validator
        # still owns independence, prior-concern reassessment and publication.
        _write(response_path, merged)
    return {"reading_strategy": "bounded_packets", "progress_path": transport["progress_path"],
            "remaining_sections": remaining_sections, "remaining_cross_document_checks": remaining_cross,
            "completed_section_count": len(sections), "completed_cross_document_count": len(cross),
            "checkpoint_evidence": evidence, "coverage_complete": complete}


def pending_verifications(revision_dir: Path) -> list[Path]:
    pending = []
    for path in sorted((revision_dir / "hermes/verification-requests").glob("*.json")):
        request = _json(path)
        if not (revision_dir / request["response_path"]).is_file(): pending.append(path)
    return pending


def _is_publication_warning(finding: Mapping[str, Any]) -> bool:
    return not finding_requires_repair(finding)


def _blocking_findings(findings: Iterable[Mapping[str, Any]]) -> list[dict[str, Any]]:
    return [dict(finding) for finding in findings if not _is_publication_warning(finding)]


def _deterministic_review_partition(
    findings: Iterable[Mapping[str, Any]],
) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    """Production deterministic checks now establish objective usability only."""
    return _blocking_findings(findings), _warning_findings(findings)




def _warning_findings(findings: Iterable[Mapping[str, Any]]) -> list[dict[str, Any]]:
    return [
        {key: value for key, value in finding.items()
         if key != "verification_source_finding" or finding.get('code') == 'nonmaterial-presentation'}
        for finding in findings if _is_publication_warning(finding)
    ]


def _accepted_drafting_warnings(revision_dir: Path, reference: Mapping[str, Any] | None = None) -> list[dict[str, Any]]:
    """Legacy drafting heuristics do not create delivery warnings or extra review work."""
    return []






def _review_evidence_text(value: Any) -> str:
    """Quote actual text values rather than JSON-escaped serialization."""
    if isinstance(value, Mapping):
        return "\n".join(_review_evidence_text(child) for child in value.values())
    if isinstance(value, list):
        return "\n".join(_review_evidence_text(child) for child in value)
    return "" if value is None else str(value)




def _localize_recovery_findings(
    findings: list[dict[str, Any]], records: list[dict[str, Any]], request: Mapping[str, Any],
) -> None:
    """Reuse unique validated review evidence; ask the reviewer about ambiguity."""
    current = [row for row in findings if row.get('verification_request_id') == request['request_id']
               and row.get('recovery_class') in {'drafting_defect', 'deterministic_structure_defect'}
               and row.get('publication_disposition') != 'warning']
    # Routing an earlier finding must not make a later ambiguous one look unique.
    original_targets = [dict(row) for row in current]
    for finding in current:
        source = dict(finding.get('verification_source_finding') or {})
        targets = finding.get('target_ids', [])
        matches = [record for record in records
                   if record['disposition'] == 'confirmed_defect'
                   and record['artifact'] == finding.get('artifact')
                   and record['section_id'] in targets]
        companions = [row for row in original_targets if row.get('artifact') == finding.get('artifact')
                      and set(row.get('target_ids', [])) & set(targets)]
        passages = {record['affected_passage'] for record in matches}
        if not _text(source.get('affected_passage')) and len(passages) == len(companions) == 1:
            finding['affected_passage'] = matches[0]['affected_passage']
            finding['recovery_localization'] = {
                'caution_sha256s': sorted({record['caution_sha256'] for record in matches}),
                'request_sha256': matches[0]['request_sha256'],
            }
        native_target = any(target in {'protocol.shell', 'icf.shell', 'prs.structured'} for target in targets)
        missing = not _text(source.get('affected_passage') or finding.get('affected_passage'))
        if 'prs.structured' in targets:
            missing = missing or not _text(source.get('source_path'))
        if native_target and missing:
            finding.update(recovery_finding({
                **finding, 'target_ids': ['verification:content'],
                'code': 'recovery_location_required',
                'issue': 'Identify the exact affected_passage for this finding, and source_path for XML fields. '
                         'The current review does not uniquely localize the correction. ' + finding['issue'],
            }, 'verifier_transient'))






def _delivery_warning_partition(warnings: Iterable[Mapping[str, Any]]) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    """Only an explicitly assessed remaining external limitation accompanies delivery."""
    remaining, history = [], []
    for warning in warnings:
        (remaining if _remaining_client_limitation(warning) else history).append(dict(warning))
    return remaining, history


def _render_warnings(render_report: Mapping[str, Any]) -> list[dict[str, Any]]:
    warnings: list[dict[str, Any]] = []
    for artifact in render_report.get("artifacts") or []:
        if isinstance(artifact, Mapping) and isinstance(artifact.get("warnings"), list):
            warnings.extend(_warning_findings(
                value for value in artifact["warnings"] if isinstance(value, Mapping)
            ))
    return warnings


def _remaining_client_limitation(finding: Mapping[str, Any]) -> bool:
    """Allow only explicit, nonmaterial external limitations to accompany delivery."""
    return (
        finding.get("check") == "client_limitation"
        and finding.get("remaining_limitation") is True
        and finding.get("limitation_kind") in {"external_dependency", "client_environment"}
        and bool(_text(finding.get("why_not_repairable")))
        and bool(_text(finding.get("recommended_action")))
        and all(finding.get(flag) is False for flag in (
            "material", "safety_critical", "contradiction", "obscures_required_information", "materially_unusable",
        ))
    )


def _valid_finding_id(finding_id: str) -> bool:
    return bool(re.fullmatch(r"[A-Za-z0-9]+(?:[-_][A-Za-z0-9]+)*", finding_id))


def _notes_cite_finding(notes: Any, finding_id: str) -> bool:
    if not _valid_finding_id(finding_id):
        return False
    return bool(
        re.search(
            rf"(?<![A-Za-z0-9_-]){re.escape(finding_id)}(?![A-Za-z0-9_-])",
            _text(notes),
        )
    )


def verification_response_is_complete(revision_dir: Path, request_path: Path) -> bool:
    """Return whether one current, canonically authenticated response fully passes."""
    try:
        _json(request_path)
    except (OSError, ValueError, json.JSONDecodeError):
        return False
    findings, _ = validate_verifications(revision_dir, request_paths=[request_path])
    return not _blocking_findings(findings)


def verification_response_is_terminal(revision_dir: Path, request_path: Path) -> bool:
    """Return whether a verifier produced a bound response the workflow can consume."""
    try:
        request = _json(request_path)
        if not isinstance(request, Mapping):
            return False
    except (OSError, ValueError, json.JSONDecodeError):
        return False
    response_path = revision_dir / str(request.get("response_path") or "")
    if not _revision_owned_nonsymlink_path(revision_dir, response_path):
        return False
    try:
        response = _json(response_path)
    except (OSError, ValueError, json.JSONDecodeError):
        return False
    if not isinstance(response, Mapping):
        return False
    request_id_counts: dict[str, int] = {}
    for active_path in (revision_dir / "hermes/verification-requests").glob("*.json"):
        try:
            active_request = _json(active_path)
        except (OSError, ValueError, json.JSONDecodeError):
            continue
        active_request_id = str(active_request.get("request_id") or "")
        request_id_counts[active_request_id] = request_id_counts.get(active_request_id, 0) + 1
    if verification_request_authentication_findings(
        revision_dir,
        request_path,
        request,
        request_id_counts=request_id_counts,
    ):
        return False
    if any(
        response.get(key) != expected
        for key, expected in (
            ("schema_version", RESPONSE_SCHEMA),
            ("request_id", request.get("request_id")),
            ("request_sha256", request.get("request_sha256")),
            ("task", request.get("task")),
            ("revision_id", request.get("revision_id")),
        )
    ):
        return False
    producer = response.get("producer") if isinstance(response.get("producer"), Mapping) else {}
    if not _text(producer.get("model_id")) or not _text(producer.get("reviewer_id")):
        return False
    status = str(response.get("status") or "").casefold()
    error = response.get("error")
    if (status in {"blocked", "failed"} and isinstance(error, Mapping)
            and error.get("type") in INSPECTION_FAILURE_TYPES and _text(error.get("message"))):
        validation_findings, evidence = validate_verifications(revision_dir, request_paths=[request_path])
        return bool(evidence) and bool(validation_findings) and all(
            item.get("recovery_class") == "capability_gap" for item in validation_findings)
    if status == "passed":
        if verification_response_is_complete(revision_dir, request_path):
            return True
        validation_findings, evidence = validate_verifications(revision_dir, request_paths=[request_path])
        # Transport completion is not clinical approval. Retain an authenticated
        # review with only resolution metadata errors so recovery can explain
        # those errors to the reviewer; the completeness gate still rejects it.
        return bool(evidence and validation_findings) and all(
            item.get("code") == "uncertainty_resolution_invalid"
            and item.get("recovery_class") == "verifier_transient"
            for item in validation_findings)
    findings = response.get("findings") if isinstance(response.get("findings"), list) else []
    if status in {"blocked", "failed"}:
        if not findings or not all(
            isinstance(item, Mapping) and bool(_text(item.get("issue")))
            for item in findings
        ):
            return False
        validation_findings, evidence = validate_verifications(
            revision_dir, request_paths=[request_path]
        )
        if not evidence:
            return False
        return not any(
            item.get("recovery_class") in {
                "verifier_transient", "document_structure_defect",
            } and item.get("code") != "uncertainty_resolution_invalid"
            for item in validation_findings
        )
    return False


def _prior_concern_review_errors(request: Mapping[str, Any], response: Mapping[str, Any]) -> list[str]:
    """Require visible independent reassessment of actual prior concerns, not rule adjudications."""
    context = request.get("review_context")
    if not isinstance(context, Mapping):
        return []
    errors = []
    authors = context.get("repair_authors", [context.get("repair_author", {})])
    producer = response.get("producer", {})
    if any(isinstance(author, Mapping) and author.get("reviewer_id")
           and author.get("reviewer_id") == producer.get("reviewer_id") for author in authors):
        errors.append("A repair author cannot provide the independent final review.")
    visual = request.get("task") == "rendered_page_visual_verification"
    assessments = response.get("page_assessments" if visual else "section_assessments", [])
    reported = response.get("findings", [])
    for prior in context.get("prior_findings", []):
        if not isinstance(prior, Mapping):
            continue
        concern = review_concern(prior)
        if reviewed_presentation_only(concern):
            continue
        is_visual = concern.get("category") == "visual" or str(concern.get("check", "")) in VISUAL_CHECKS
        if visual != is_visual:
            continue
        artifact = concern.get("artifact")
        if visual and artifact and artifact not in {
            row.get("artifact") for row in request.get("artifacts", []) if isinstance(row, Mapping)
        }:
            continue  # each visual task reviews only its own native document/pages
        identity = str(concern.get("finding_id") or "")
        if not identity:
            continue
        if any(isinstance(row, Mapping) and row.get("finding_id") == identity for row in reported):
            continue  # the concern remains reported, rather than silently cleared
        targets = concern.get("target_ids") or []
        section_targets = {row.get("section_id") for row in request.get("sections", []) if isinstance(row, Mapping)}
        scoped_targets = set(targets) & section_targets
        check = concern.get("check")
        cross_check = not visual and not scoped_targets and check in request.get("cross_document_checks", [])
        matched = not cross_check and any(
            isinstance(row, Mapping)
            and (not artifact or row.get("artifact") == artifact)
            and (visual or not scoped_targets or row.get("section_id") in scoped_targets)
            and _notes_cite_finding(row.get("notes"), identity)
            for row in assessments
        )
        if not visual and not scoped_targets and not matched:
            matched = any(isinstance(row, Mapping)
                and (not cross_check or row.get("check") == check)
                and _notes_cite_finding(row.get("notes"), identity)
                for row in response.get("cross_document_assessments", []))
        if not matched:
            errors.append(f"Prior substantive concern {identity} requires independent reassessment in the relevant assessment notes or a retained finding.")
    return errors


def validate_verifications(
    revision_dir: Path,
    *,
    request_paths: Iterable[Path] | None = None,
) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    findings: list[dict[str, Any]] = []; evidence: dict[str, Any] = {}
    request_root = revision_dir / "hermes/verification-requests"
    active_entries = sorted(request_root.iterdir()) if request_root.is_dir() else []
    for path in active_entries:
        if path.suffix != ".json" or path.is_symlink() or not path.is_file():
            findings.append(recovery_finding({
                "category": "verification",
                "field": "verification_request_inventory",
                "code": "verification_request_invalid_inventory",
                "target_ids": ["verification:content"],
                "verification_request_id": None,
                "issue": f"Unexpected noncanonical verification request entry: {path.name}",
            }, "document_structure_defect"))
    selected_paths = sorted(
        request_paths
        if request_paths is not None
        else [
            path for path in active_entries
            if path.suffix == ".json" and not path.is_symlink() and path.is_file()
        ]
    )
    active_request_id_counts: dict[str, int] = {}
    active_requests: dict[Path, dict[str, Any]] = {}
    required_bindings = ("request_id", "task", "revision_id", "review_set", "response_path")
    for active_path in active_entries:
        if active_path.suffix != ".json" or active_path.is_symlink() or not active_path.is_file():
            continue
        try:
            active_request = _json(active_path)
        except (OSError, ValueError, json.JSONDecodeError) as exc:
            if active_path not in selected_paths:
                findings.append(recovery_finding({
                    "category": "verification", "field": "verification_request",
                    "code": "verification_request_invalid_json",
                    "target_ids": ["verification:content"], "verification_request_id": None,
                    "issue": f"Verification request is not a readable JSON object: {exc}",
                }, "document_structure_defect"))
            continue
        active_requests[active_path] = active_request
        active_request_id = str(active_request.get("request_id") or "")
        active_request_id_counts[active_request_id] = active_request_id_counts.get(active_request_id, 0) + 1
    for active_path, active_request in active_requests.items():
        if active_path in selected_paths:
            continue
        if any(key not in active_request or active_request.get(key) in (None, "") for key in required_bindings):
            findings.append(recovery_finding({
                "category": "verification", "field": "verification_request",
                "code": "verification_request_missing_binding",
                "target_ids": ["verification:content"],
                "verification_request_id": active_request.get("request_id"),
                "issue": "Companion verification request is missing a required binding.",
            }, "document_structure_defect"))
            continue
        findings.extend(verification_request_authentication_findings(
            revision_dir, active_path, active_request,
            request_id_counts=active_request_id_counts,
        ))
    request_records: list[tuple[Path, dict[str, Any]]] = []
    for request_path in selected_paths:
        try:
            request = _json(request_path)
        except (OSError, ValueError, json.JSONDecodeError) as exc:
            findings.append(recovery_finding({
                "category": "verification",
                "field": "verification_request",
                "code": "verification_request_invalid_json",
                "target_ids": ["verification:content"],
                "verification_request_id": None,
                "issue": f"Verification request is not a readable JSON object: {exc}",
            }, "document_structure_defect"))
            continue
        required_bindings = ("request_id", "task", "revision_id", "review_set", "response_path")
        if any(key not in request or request.get(key) in (None, "") for key in required_bindings):
            findings.append(recovery_finding({
                "category": "verification",
                "field": "verification_request",
                "code": "verification_request_missing_binding",
                "target_ids": [
                    "verification:visual"
                    if request.get("task") == "rendered_page_visual_verification"
                    else "verification:content"
                ],
                "verification_request_id": request.get("request_id"),
                "issue": "Verification request is missing a required identity, task, revision, review-set, or response-path binding.",
            }, "document_structure_defect"))
            continue
        request_records.append((request_path, request))
    task_counts: dict[str, int] = {}
    for _, request in request_records:
        task = str(request.get("task"))
        task_counts[task] = task_counts.get(task, 0) + 1
    for request_path, request in request_records:
        evidence_key = request["task"] if task_counts[str(request.get("task"))] == 1 else request["request_id"]
        verification_target = "verification:visual" if request["task"] == "rendered_page_visual_verification" else "verification:content"
        authentication_findings = verification_request_authentication_findings(
            revision_dir,
            request_path,
            request,
            request_id_counts=active_request_id_counts,
        )
        if authentication_findings:
            findings.extend(authentication_findings)
            continue
        response_path = revision_dir / request["response_path"]
        if not _revision_owned_nonsymlink_path(revision_dir, response_path):
            findings.append(recovery_finding({
                "category": "verification",
                "field": request["task"],
                "code": "verification_response_noncanonical_path",
                "target_ids": [verification_target],
                "verification_request_id": request["request_id"],
                "issue": "Independent Hermes verification response path is not revision-owned canonical evidence.",
            }, "document_structure_defect"))
            continue
        if not response_path.is_file():
            findings.append(recovery_finding({"category": "verification", "field": request["task"], "code": "verification_response_missing", "target_ids": [verification_target], "verification_request_id": request["request_id"], "issue": "Independent Hermes verification response is missing."}, "verifier_transient")); continue
        try: response = _json(response_path)
        except (OSError, ValueError, json.JSONDecodeError) as exc:
            findings.append(recovery_finding({"category": "verification", "field": request["task"], "target_ids": [verification_target], "verification_request_id": request["request_id"], "issue": f"Invalid verification response: {exc}"}, "verifier_transient")); continue
        malformed_response = False
        for key, expected in (("schema_version", RESPONSE_SCHEMA), ("request_id", request["request_id"]), ("request_sha256", request["request_sha256"]), ("task", request["task"]), ("revision_id", request.get("revision_id"))):
            if response.get(key) != expected:
                malformed_response = True
                findings.append(recovery_finding({
                    "category": "reviewer-transient",
                    "field": key,
                    "target_ids": [verification_target],
                    "verification_request_id": request["request_id"],
                    "issue": f"Verification response binding mismatch for {key}.",
                }, "verifier_transient"))
        producer = response.get("producer") if isinstance(response.get("producer"), Mapping) else {}
        if not _text(producer.get("model_id")):
            malformed_response = True
            findings.append(recovery_finding({"category": "verification", "field": "producer.model_id", "target_ids": [verification_target], "verification_request_id": request["request_id"], "issue": "Verifier model identity is missing."}, "verifier_transient"))
        if not _text(producer.get("reviewer_id")):
            malformed_response = True
            findings.append(recovery_finding({"category": "verification", "field": "producer.reviewer_id", "target_ids": [verification_target], "verification_request_id": request["request_id"], "issue": "Verifier reviewer identity is missing."}, "verifier_transient"))
        if malformed_response:
            evidence[evidence_key] = {"request": request_path.relative_to(revision_dir).as_posix(), "request_sha256": sha256_file(request_path), "response": response_path.relative_to(revision_dir).as_posix(), "response_sha256": sha256_file(response_path), "producer": producer, "status": "malformed", "contracted_template_bundle": request.get("contracted_template_bundle")}
            continue
        checkpoint_invalid = False
        checkpoint_rows = response.get("review_progress_evidence", [])
        if not isinstance(checkpoint_rows, list):
            checkpoint_invalid = True
            checkpoint_rows = []
        for checkpoint in checkpoint_rows:
            if not isinstance(checkpoint, Mapping):
                checkpoint_invalid = True
                continue
            checkpoint_path = revision_dir / str(checkpoint.get("path") or "")
            if (not _revision_owned_nonsymlink_path(revision_dir, checkpoint_path)
                    or not checkpoint_path.is_file() or sha256_file(checkpoint_path) != checkpoint.get("sha256")):
                checkpoint_invalid = True
        if checkpoint_invalid:
            findings.append(recovery_finding({"category": "verification", "field": request["task"],
                "code": "review_checkpoint_changed", "target_ids": [verification_target],
                "verification_request_id": request["request_id"],
                "issue": "Saved independent review evidence changed; complete fresh inspection is required."}, "verifier_transient"))
            continue
        if "review_unit_evidence" in response:
            try:
                proofs = response["review_unit_evidence"]
                if not isinstance(proofs, list) or any(not isinstance(proof, Mapping) for proof in proofs):
                    raise ValueError("Malformed scheduled review evidence.")
                for proof in proofs:
                    for file_key, hash_key in (("request", "request_sha256"), ("response", "response_sha256")):
                        path = revision_dir / str(proof.get(file_key) or "")
                        if (not _revision_owned_nonsymlink_path(revision_dir, path) or not path.is_file()
                                or sha256_file(path) != proof.get(hash_key)):
                            raise ValueError("Scheduled independent evidence changed.")
                unit_paths = [revision_dir / value for value in dict.fromkeys(proof["request"] for proof in proofs)]
                assembled = content_review_schedule_result(revision_dir, request_path, unit_paths, publish=False)["response"]
                if any(response.get(key) != assembled.get(key) for key in (
                        "section_assessments", "cross_document_assessments", "findings", "review_unit_evidence", "scheduler_partial", "producer")):
                    raise ValueError("Parent review differs from its independent unit judgments.")
            except (OSError, ValueError, KeyError, TypeError) as exc:
                findings.append(recovery_finding({"category": "verification", "field": request["task"],
                    "code": "scheduled_review_evidence_invalid", "target_ids": [verification_target],
                    "verification_request_id": request["request_id"], "issue": str(exc)}, "verifier_transient"))
                continue
        error = response.get("error") if isinstance(response.get("error"), Mapping) else {}
        error_type = _text(error.get("type")).casefold()
        if (response.get("status") in {"blocked", "failed"}
                and error_type in INSPECTION_FAILURE_TYPES and _text(error.get("message"))):
            findings.append(recovery_finding({
                "category": "environment", "field": request["task"], "code": error_type,
                "target_ids": [verification_target], "verification_request_id": request["request_id"],
                "issue": _text(error["message"]),
            }, "capability_gap"))
            evidence[evidence_key] = {"request": request_path.relative_to(revision_dir).as_posix(),
                "request_sha256": sha256_file(request_path), "response": response_path.relative_to(revision_dir).as_posix(),
                "response_sha256": sha256_file(response_path), "producer": producer, "status": "inspection_incomplete"}
            continue
        transient = str(response.get("status") or "").casefold() in TRANSIENT_REVIEW_STATUSES or error_type in {
            "api_unavailable", "connection_error", "rate_limit", "timeout", "service_unavailable"
        }
        if transient:
            findings.append(recovery_finding({
                "category": "reviewer-transient",
                "field": request["task"],
                "target_ids": [verification_target],
                "verification_request_id": request["request_id"],
                "issue": _text(error.get("message")) or "Independent Hermes verifier reported a transient API failure.",
            }, "verifier_transient"))
            evidence[evidence_key] = {"request": request_path.relative_to(revision_dir).as_posix(), "request_sha256": sha256_file(request_path), "response": response_path.relative_to(revision_dir).as_posix(), "response_sha256": sha256_file(response_path), "producer": producer, "status": "transient", "contracted_template_bundle": request.get("contracted_template_bundle")}
            continue
        if not isinstance(response.get('status'), str) or response['status'] not in {'passed', 'blocked', 'failed'}:
            findings.append(recovery_finding({
                'category': 'reviewer-transient', 'field': 'status',
                'target_ids': [verification_target], 'verification_request_id': request['request_id'],
                'issue': 'Independent review must use a recognized terminal status.',
            }, 'verifier_transient'))
            continue
        raw_issues = response.get("findings")
        issues: list[Any] = list(raw_issues) if isinstance(raw_issues, list) else []
        malformed_finding = not isinstance(raw_issues, list) or any(
            not isinstance(item, Mapping) or not isinstance(item.get('issue'), str)
            or not item['issue'].strip() or (
                any(key in item and not isinstance(item[key], str)
                    for key in ('category', 'artifact', 'check', 'surface'))
                or 'target_ids' in item and (
                    not isinstance(item['target_ids'], list)
                    or any(not isinstance(target, str) for target in item['target_ids'])
                )
            ) for item in issues
        )
        if malformed_finding:
            findings.append(recovery_finding({
                'category': 'reviewer-transient', 'field': 'findings',
                'target_ids': [verification_target], 'verification_request_id': request['request_id'],
                'issue': 'Review findings must be a list of objects with a nonempty issue string; labels must be strings and target_ids must be a list of strings.',
            }, 'verifier_transient'))
            continue
        if response.get("status") != "passed" or issues:
            for item in issues or [{"issue": "Verifier did not pass the artifact."}]:
                source = item if isinstance(item, Mapping) else {"issue": item}
                category = "visual" if request["task"] == "rendered_page_visual_verification" else "verification"
                finding = {
                    "partial_review": bool(response.get("scheduler_partial") and response.get("review_unit_evidence")),
                    "category": category,
                    "field": request["task"],
                    "verification_request_id": request["request_id"],
                    "verification_source_finding": dict(source),
                    "issue": _text(source.get("issue")),
                }
                for key in (
                    "target_ids", "artifact", "page", "check", "element", "code",
                    "concept_id", "primary_section", "secondary_section",
                    "primary_paragraphs", "secondary_paragraphs", "treatment",
                    "necessary", "concise", "material", "resolved", "repair_attempt",
                    "safety_critical",
                    "contradiction", "obscures_required_information", "materially_unusable",
                    "disposition",
                    "surface", "table_ids",
                ):
                    if key in source: finding[key] = source[key]
                if 'surface' in finding:
                    finding['surface'] = canonical_review_surface(finding['surface'])
                if category == "visual":
                    expected_artifacts = {
                        str(artifact.get("artifact"))
                        for artifact in request.get("artifacts", [])
                        if isinstance(artifact, Mapping)
                    }
                    artifact = _text(source.get("artifact"))
                    artifact_record = next((
                        item for item in request.get("artifacts", [])
                        if isinstance(item, Mapping)
                        and _text(item.get("artifact")) == artifact
                    ), {})
                    expected_pages = {
                        item.get("page")
                        for item in artifact_record.get("pages", [])
                        if isinstance(item, Mapping)
                        and isinstance(item.get("page"), int)
                        and not isinstance(item.get("page"), bool)
                    }
                    page = source.get("page")
                    page_is_bound = (
                        isinstance(page, int)
                        and not isinstance(page, bool)
                        and page in expected_pages
                    )
                    check = _text(source.get("check"))
                    element = _text(source.get("element"))
                    if (
                        artifact not in expected_artifacts
                        or not page_is_bound
                    ):
                        findings.append(recovery_finding({
                            **finding,
                            "target_ids": [verification_target],
                            "verification_request_id": request["request_id"],
                            "issue": f"Verifier finding does not identify an actual reviewed artifact/page. {finding['issue']}".strip(),
                        }, "verifier_transient"))
                    else:
                        finding["target_ids"] = [f"layout:{artifact}"]
                        findings.append(recovery_finding(finding, "visual_defect"))
                else:
                    expected_artifacts = {
                        str(row.get("artifact"))
                        for row in [*request.get("artifacts", []), *request.get("sections", [])]
                        if isinstance(row, Mapping) and row.get("artifact")
                    }
                    artifact = _text(source.get("artifact"))
                    if artifact and artifact not in expected_artifacts:
                        findings.append(recovery_finding({
                            **finding, "target_ids": [verification_target],
                            "issue": f"Verifier finding identifies an artifact outside this review. {finding['issue']}".strip(),
                        }, "verifier_transient"))
                    elif _remaining_client_limitation(source):
                        findings.append({
                            **finding, "category": "content", "check": "client_limitation",
                            "remaining_limitation": True,
                            "limitation_kind": source["limitation_kind"],
                            "why_not_repairable": source["why_not_repairable"],
                            "recommended_action": source["recommended_action"],
                            "publication_disposition": "warning", "severity": "warning",
                        })
                    else:
                        # Native locations, table names and checks are reviewer evidence,
                        # never an admission registry for general document repair.
                        findings.append(recovery_finding({
                            **finding,
                            "category": _text(source.get("category")) or "content",
                            "recommended_action": source.get("recommended_action"),
                        }, "drafting_defect"))
        # Apply disposition only after authenticated response and artifact/page
        # identity validation. The raw reviewer evidence remains unchanged and auditable.
        for finding in findings:
            source = finding.get('verification_source_finding')
            if (finding.get('verification_request_id') == request['request_id']
                    and finding.get('recovery_class') in {'drafting_defect', 'deterministic_structure_defect', 'visual_defect'}
                    and isinstance(source, Mapping) and reviewed_presentation_only(source)):
                finding.update(code='nonmaterial-presentation', publication_disposition='warning',
                               severity='internal', authority='independent_review')
        for artifact in request.get("artifacts", []):
            if request["task"] == "clinical_content_verification":
                path = revision_dir / str(artifact.get("path"))
                exact_matches = (
                    path.is_file()
                    and _revision_owned_nonsymlink_path(revision_dir, path)
                    and sha256_file(path) == artifact.get("sha256")
                )
                content_matches = (
                    not artifact.get("content_sha256")
                    or path.is_file()
                    and _content_sha256(path) == artifact.get("content_sha256")
                )
                if not exact_matches or not content_matches:
                    findings.append(recovery_finding({"category": "verification", "field": request["task"], "artifact": artifact.get("artifact"), "artifact_path": artifact.get("path"), "target_ids": [verification_target], "verification_request_id": request["request_id"], "issue": f"Verification request is stale for {artifact.get('path')}."}, "document_structure_defect"))
            else:
                for key in ("docx", "pdf"):
                    path = revision_dir / str(artifact.get(key))
                    if not path.is_file() or not _revision_owned_nonsymlink_path(revision_dir, path) or sha256_file(path) != artifact.get(f"{key}_sha256"):
                        findings.append(recovery_finding({"category": "verification", "field": artifact.get("artifact", key), "artifact": artifact.get("artifact"), "artifact_path": artifact.get(key), "target_ids": [verification_target], "verification_request_id": request["request_id"], "issue": f"Visual verification request is stale for {artifact.get(key)}."}, "document_structure_defect"))
                for page in artifact.get("pages", []):
                    path = revision_dir / str(page.get("path"))
                    if not path.is_file() or not _revision_owned_nonsymlink_path(revision_dir, path) or sha256_file(path) != page.get("sha256"):
                        findings.append(recovery_finding({"category": "verification", "field": artifact.get("artifact", "page"), "artifact": artifact.get("artifact"), "artifact_path": page.get("path"), "target_ids": [verification_target], "verification_request_id": request["request_id"], "issue": f"Visual verification request is stale for {page.get('path')}."}, "document_structure_defect"))
        if request["task"] == "clinical_content_verification":
            expected_sections = {(item["artifact"], item["section_id"]) for item in request.get("sections", [])}
            valid_section_rows = [
                item for item in response.get("section_assessments", [])
                if isinstance(item, Mapping)
                and item.get("status") in {"passed", "blocked", "failed"}
                and set(item.get("checks", [])) == set(CONTENT_CHECKS)
            ]
            assessed_sections = {(item.get("artifact"), item.get("section_id")) for item in valid_section_rows}
            if expected_sections != assessed_sections or len(valid_section_rows) != len(expected_sections):
                findings.append(recovery_finding({"category": "verification", "field": request["task"], "target_ids": [verification_target], "verification_request_id": request["request_id"], "code": "review_coverage_incomplete", "issue": f"Every contracted section and content check must be explicitly assessed; expected {len(expected_sections)}, accepted {len(assessed_sections)}."}, "verifier_transient"))
            expected_cross = set(request.get("cross_document_checks", []))
            valid_cross_rows = [
                item for item in response.get("cross_document_assessments", [])
                if isinstance(item, Mapping)
                and item.get("status") in {"passed", "blocked", "failed"}
            ]
            assessed_cross = {item.get("check") for item in valid_cross_rows}
            if expected_cross != assessed_cross or len(valid_cross_rows) != len(expected_cross):
                findings.append(recovery_finding({"category": "verification", "field": request["task"], "target_ids": [verification_target], "verification_request_id": request["request_id"], "code": "review_coverage_incomplete", "issue": f"Every cross-document check must be explicitly assessed; expected {len(expected_cross)}, accepted {len(assessed_cross)}."}, "verifier_transient"))
            reported_findings = [item for item in issues if isinstance(item, Mapping)]
            reported_ids = [_text(item.get("finding_id")) for item in reported_findings]
            if (
                len(reported_findings) != len(issues)
                or any(not _valid_finding_id(finding_id) for finding_id in reported_ids)
                or len(set(reported_ids)) != len(reported_ids)
            ):
                findings.append(recovery_finding({
                    "category": "verification",
                    "field": request["task"],
                    "target_ids": [verification_target],
                    "verification_request_id": request["request_id"],
                    "issue": "Every content finding must have a unique valid finding_id.",
                }, "verifier_transient"))
            consistency, _ = _content_assessment_consistency(valid_section_rows, valid_cross_rows, reported_findings)
            for issue in consistency.values():
                findings.append(recovery_finding({"category": "verification", "field": request["task"],
                    "target_ids": [verification_target], "verification_request_id": request["request_id"],
                    "issue": issue}, "verifier_transient"))
            if response.get("status") == "passed" and any(
                item.get("status") != "passed" for item in [*valid_section_rows, *valid_cross_rows]
            ):
                findings.append(recovery_finding({"category": "verification", "field": request["task"], "target_ids": [verification_target], "verification_request_id": request["request_id"], "issue": "A passing content response contains a non-passing section or cross-document assessment."}, "verifier_transient"))
        if request["task"] == "rendered_page_visual_verification":
            expected_pages = {(a["artifact"], p["page"], p["sha256"]) for a in request.get("artifacts", []) for p in a.get("pages", [])}
            valid_page_rows = [
                p for p in response.get("page_assessments", [])
                if isinstance(p, Mapping)
                and p.get("status") in {"passed", "blocked", "failed"}
                and set(p.get("checks", [])) == set(VISUAL_CHECKS)
            ]
            assessed = {(p.get("artifact"), p.get("page"), p.get("sha256")) for p in valid_page_rows}
            if expected_pages != assessed or len(valid_page_rows) != len(expected_pages): findings.append(recovery_finding({"category": "verification", "field": "page_assessments", "target_ids": [verification_target], "verification_request_id": request["request_id"], "issue": f"Every rendered page and every visual check must be explicitly assessed; expected {len(expected_pages)}, accepted {len(assessed)}."}, "verifier_transient"))
            if response.get("status") == "passed" and any(
                page.get("status") != "passed" for page in valid_page_rows
            ):
                findings.append(recovery_finding({"category": "verification", "field": "page_assessments", "target_ids": [verification_target], "verification_request_id": request["request_id"], "issue": "A passing visual response contains a non-passing page assessment."}, "verifier_transient"))
        findings.extend(recovery_finding({
            "category": "verification", "field": "review_context",
            "target_ids": [verification_target], "verification_request_id": request["request_id"],
            "issue": message,
        }, "verifier_transient") for message in _prior_concern_review_errors(request, response))
        resolution_records, resolution_errors = ([], [])
        # General native repair inspects actual files; no registered-surface localization is required.
        evidence[evidence_key] = {
            "uncertainty_resolutions": resolution_records if not resolution_errors else [],
            "source_routing_assessments": (response.get("source_routing_assessments", [])
                                           if not resolution_errors else []),
            "task": request.get("task"),
            "request_id": request.get("request_id"),
            "review_set": request.get("review_set"),
            "source_identity_sha256": request.get("source_identity_sha256"),
            "generation_authority_sha256": request.get("generation_authority_sha256"),
            "review_context": request.get("review_context", {}),
            "request": request_path.relative_to(revision_dir).as_posix(),
            "request_sha256": sha256_file(request_path),
            "request_body_sha256": request.get("request_sha256"),
            "response": response_path.relative_to(revision_dir).as_posix(),
            "response_sha256": sha256_file(response_path),
            "producer": producer,
            "producer_model_id": producer.get("model_id"),
            "producer_reviewer_id": producer.get("reviewer_id"),
            "renderer": request.get("renderer"),
            "page_renderer": request.get("page_renderer"),
            "artifacts": request.get("artifacts", []),
            "sections": request.get("sections", []),
            "approved_source_sha256": canonical_evidence_sha256(request["approved_source"]) if "approved_source" in request else None,
            "semantic_evidence": request.get("semantic_evidence", []),
            "lexical_review_cautions": request.get("lexical_review_cautions", []),
            "section_assessments": response.get("section_assessments", []) if request.get("task") == "clinical_content_verification" else [],
            "checks": request.get("checks", []),
            "cross_document_checks": request.get("cross_document_checks", []),
            "contracted_template_bundle": request.get("contracted_template_bundle"),
        }
    return findings, evidence


def _final_verification_scope_findings(
    revision_dir: Path,
    reference: Mapping[str, Any],
    render_report: Mapping[str, Any],
    evidence: Mapping[str, Any],
) -> list[dict[str, Any]]:
    """Require fresh package content review and exact independently reviewed visual evidence."""
    expected_names = set(document_set(get_path(reference, "meta.study_type")))
    expected_paths = {f"candidate/{name}" for name in expected_names}
    expected_visual = {Path(name).stem for name in expected_names if name.endswith(".docx")}
    records = [dict(item) for item in evidence.values() if isinstance(item, Mapping)]
    content = [item for item in records if item.get("task") == "clinical_content_verification"]
    visual = [item for item in records if item.get("task") == "rendered_page_visual_verification"]
    issues: list[str] = []
    if len(content) != 1:
        issues.append("Final review requires exactly one package-wide content review.")
    else:
        if (content[0].get("approved_source_sha256") is not None
                and content[0]["approved_source_sha256"] != canonical_evidence_sha256(reference)):
            issues.append("The final content review belongs to a different approved source.")
        content_paths = {
            str(artifact.get("path"))
            for artifact in content[0].get("artifacts", [])
            if isinstance(artifact, Mapping)
        }
        expected_sections = {
            (str(item["artifact"]), str(item["section_id"]))
            for item in content_review_sections(reference)
        }
        content_sections = {
            (str(item.get("artifact")), str(item.get("section_id")))
            for item in content[0].get("sections", [])
            if isinstance(item, Mapping)
        }
        expected_cross = set(CROSS_DOCUMENT_CHECKS) if len(expected_names) > 1 else set()
        if (
            content[0].get("semantic_evidence") != semantic_evidence_inventory(reference)
            or content_paths != expected_paths
            or content_sections != expected_sections
            or set(content[0].get("checks") or []) != set(CONTENT_CHECKS)
            or set(content[0].get("cross_document_checks") or []) != expected_cross
        ):
            issues.append("The final content review does not cover every section and cross-document check in the exact branch document set.")
    expected_render_artifacts = {
        str(artifact.get("artifact")): {
            key: value
            for key, value in artifact.items()
            if key not in {"renderer", "page_renderer"}
        }
        for artifact in render_report.get("artifacts", [])
        if isinstance(artifact, Mapping)
    }
    expected_renderers = {
        str(artifact.get("artifact")): artifact.get("renderer", render_report.get("renderer"))
        for artifact in render_report.get("artifacts", [])
        if isinstance(artifact, Mapping)
    }
    expected_page_renderers = {
        str(artifact.get("artifact")): artifact.get("page_renderer", render_report.get("page_renderer"))
        for artifact in render_report.get("artifacts", [])
        if isinstance(artifact, Mapping)
    }
    visual_artifacts = [
        dict(artifact)
        for item in visual
        for artifact in item.get("artifacts", [])
        if isinstance(artifact, Mapping)
    ]
    visual_names = [str(artifact.get("artifact")) for artifact in visual_artifacts]
    if (
        len(visual) != len(expected_visual)
        or any(len(item.get("artifacts", [])) != 1 for item in visual)
        or len(visual_artifacts) != len(expected_visual)
        or set(visual_names) != expected_visual
        or set(expected_render_artifacts) != expected_visual
        or any(
            artifact != expected_render_artifacts.get(str(artifact.get("artifact")))
            for artifact in visual_artifacts
        )
    ):
        issues.append("Final review requires one document-scoped visual review with the complete rendered page inventory for every final DOCX.")
    content_set = content[0].get("review_set") if len(content) == 1 else None
    if not isinstance(content_set, int) or content_set < 1:
        issues.append("Final content review must identify a valid current review set.")
    for item in visual:
        visual_set = item.get("review_set")
        if not isinstance(visual_set, int) or not isinstance(content_set, int) or not 1 <= visual_set <= content_set:
            issues.append("Visual review cannot belong to a future or invalid review set.")
        elif visual_set != content_set and (
            item.get("source_identity_sha256") != canonical_evidence_sha256(reference)
            or not item.get("generation_authority_sha256")
            or item.get("generation_authority_sha256") != content[0].get("generation_authority_sha256")
            or set(item.get("checks") or []) != set(VISUAL_CHECKS)
            or _has_visual_concern(content[0].get("review_context") or {},
                                   str((item.get("artifacts") or [{}])[0].get("artifact") or ""))
            or any(author.get("reviewer_id") == item.get("producer_reviewer_id")
                   for author in (content[0].get("review_context") or {}).get("repair_authors", [])
                   if isinstance(author, Mapping))
        ):
            issues.append("Reused visual review must preserve approved source, generation authority and every visual check.")
    if any(
        not str(item.get("request_id") or "").strip()
        or not str(item.get("producer_model_id") or "").strip()
        or not str(item.get("producer_reviewer_id") or "").strip()
        or not re.fullmatch(r"[0-9a-f]{64}", str(item.get("request_sha256") or ""))
        or not re.fullmatch(r"[0-9a-f]{64}", str(item.get("response_sha256") or ""))
        for item in records
    ):
        issues.append("Every final response must bind request, reviewer, model, and response identity.")
    for item in visual:
        artifacts = item.get("artifacts", [])
        artifact_name = str(artifacts[0].get("artifact")) if len(artifacts) == 1 and isinstance(artifacts[0], Mapping) else ""
        renderer_identity = item.get("renderer")
        page_renderer_identity = item.get("page_renderer")
        if (
            not artifact_name
            or not isinstance(renderer_identity, Mapping)
            or not str(renderer_identity.get("kind") or "").strip()
            or renderer_identity != expected_renderers.get(artifact_name)
            or not isinstance(page_renderer_identity, Mapping)
            or not str(page_renderer_identity.get("kind") or "").strip()
            or page_renderer_identity != expected_page_renderers.get(artifact_name)
        ):
            issues.append("Every final visual response must bind the exact renderer and page renderer identities that produced its artifact evidence.")
            break
    actual_names = {
        path.name for path in (revision_dir / "candidate").glob("*") if path.is_file()
    }
    if actual_names != expected_names:
        issues.append("The final candidate directory is not the exact branch document set.")
    return [
        recovery_finding({
            "category": "verification",
            "field": "final_exact_artifact_review",
            "target_ids": ["verification:content"],
            "issue": issue,
        }, "document_structure_defect")
        for issue in issues
    ]


def _exact_rendered_hashes(
    revision_dir: Path,
    render_report: Mapping[str, Any],
) -> dict[str, str]:
    paths = [
        str(item.get("path"))
        for artifact in render_report.get("artifacts", [])
        if isinstance(artifact, Mapping)
        for item in [
            {"path": artifact.get("docx")},
            {"path": artifact.get("pdf")},
            *[
                {"path": page.get("path")}
                for page in artifact.get("pages", [])
                if isinstance(page, Mapping)
            ],
        ]
        if item.get("path")
    ]
    return {
        path: sha256_file(revision_dir / path)
        for path in paths
        if (revision_dir / path).is_file()
    }


def _final_review_bindings(evidence: Mapping[str, Any]) -> list[dict[str, Any]]:
    return [
        {
            "task": item.get("task"),
            "request_id": item.get("request_id"),
            "request_sha256": item.get("request_sha256"),
            "request_body_sha256": item.get("request_body_sha256"),
            "response_sha256": item.get("response_sha256"),
            "review_set": item.get("review_set"),
            "reviewer_identity": item.get("producer"),
            "reviewer_id": item.get("producer_reviewer_id"),
            "model_identity": item.get("producer_model_id"),
            "renderer": item.get("renderer"),
            "page_renderer": item.get("page_renderer"),
        }
        for item in sorted(
            (dict(value) for value in evidence.values() if isinstance(value, Mapping)),
            key=lambda value: str(value.get("request_id") or ""),
        )
    ]


def final_exact_artifact_review_findings(
    revision_dir: Path,
    reference: Mapping[str, Any],
    render_report: Mapping[str, Any],
    quality: Mapping[str, Any],
) -> list[dict[str, Any]]:
    """Revalidate the non-bypassable final review against current exact bytes."""
    final_review = quality.get("final_exact_artifact_review")
    if not isinstance(final_review, Mapping):
        return [recovery_finding({
            "category": "verification",
            "field": "final_exact_artifact_review",
            "issue": "Final Exact-Artifact Review is required.",
        }, "document_structure_defect")]
    expected_values = {
        "schema_version": FINAL_REVIEW_SCHEMA,
        "scope": FINAL_REVIEW_SCOPE,
        "status": "passed",
        "branch_document_set": sorted(document_set(get_path(reference, "meta.study_type"))),
        "mandatory_visual_checks": list(VISUAL_CHECKS),
        "every_page": True,
        "section_three_and_orphan_heading_checks": True,
    }
    issues = []
    if quality.get("status") != "passed" or quality.get("findings"):
        issues.append("Final Exact-Artifact Review is not attached to passing quality evidence.")
    if any(final_review.get(key) != expected for key, expected in expected_values.items()):
        issues.append("Final Exact-Artifact Review is incomplete.")
    candidate_hashes = {
        path.relative_to(revision_dir).as_posix(): sha256_file(path)
        for path in sorted((revision_dir / "candidate").glob("*"))
        if path.is_file()
    }
    if candidate_hashes != dict(final_review.get("candidate_hashes") or {}):
        issues.append("Final candidate bytes changed after review.")
    if _exact_rendered_hashes(revision_dir, render_report) != dict(final_review.get("rendered_hashes") or {}):
        issues.append("Final rendered bytes changed after review.")
    deterministic_findings = deterministic_content_check(revision_dir, reference)
    deterministic_blockers, deterministic_warnings = _deterministic_review_partition(deterministic_findings)
    verification_findings, fresh_evidence = validate_verifications(revision_dir)
    fresh_warnings = (
        deterministic_warnings + _accepted_drafting_warnings(revision_dir, reference)
        + _render_warnings(render_report) + _warning_findings(render_report.get("findings", []))
        + _warning_findings(verification_findings)
    )
    verification_findings = [
        dict(item) for item in deterministic_blockers
    ] + _blocking_findings(verification_findings)
    verification_findings.extend(
        _final_verification_scope_findings(
            revision_dir, reference, render_report, fresh_evidence
        )
    )
    fresh_resolutions = []
    fresh_warnings, fresh_history = _delivery_warning_partition(fresh_warnings)
    if fresh_history != list(quality.get("internal_diagnostics") or []):
        issues.append("Final internal diagnostic evidence changed after review.")
    if (
        canonical_evidence_sha256(fresh_evidence)
        != final_review.get("verification_evidence_sha256")
        or fresh_evidence != quality.get("verification_evidence")
        or _final_review_bindings(fresh_evidence) != final_review.get("review_bindings")
    ):
        issues.append("Final request or reviewer evidence changed after review.")
    if fresh_warnings != list(quality.get("warnings") or []):
        issues.append("Final verification warnings changed after review.")
    if fresh_resolutions != list(quality.get("resolved_evidence_cautions") or []):
        issues.append("Final evidence-caution resolutions changed after review.")
    return [
        recovery_finding({
            "category": "verification",
            "field": "final_exact_artifact_review",
            "issue": issue,
        }, "document_structure_defect")
        for issue in issues
    ] + verification_findings


def quality_report(revision_dir: Path, reference: Mapping[str, Any], render_report: Mapping[str, Any], xml_report: Mapping[str, Any] | None) -> dict[str, Any]:
    deterministic_findings = deterministic_content_check(revision_dir, reference)
    deterministic_blockers, deterministic_warnings = _deterministic_review_partition(deterministic_findings)
    warnings = (
        deterministic_warnings + _accepted_drafting_warnings(revision_dir, reference)
        + _render_warnings(render_report) + _warning_findings(render_report.get("findings", []))
    )
    findings = deterministic_blockers
    findings.extend(_blocking_findings(render_report.get("findings", [])))
    if xml_report:
        findings.extend(
            recovery_finding(
                {**dict(item), "target_ids": ["prs.structured"]},
                "deterministic_structure_defect",
            )
            for item in xml_report.get("findings", [])
        )
    verification_findings, evidence = validate_verifications(revision_dir)
    warnings.extend(_warning_findings(verification_findings))
    findings.extend(_blocking_findings(verification_findings))
    findings.extend(_final_verification_scope_findings(revision_dir, reference, render_report, evidence))
    resolutions = []
    warnings, internal_history = _delivery_warning_partition(warnings)
    candidate_hashes = {
        path.relative_to(revision_dir).as_posix(): sha256_file(path)
        for path in sorted((revision_dir / "candidate").glob("*"))
        if path.is_file()
    }
    rendered_hashes = _exact_rendered_hashes(revision_dir, render_report)
    review_bindings = _final_review_bindings(evidence)
    final_review = {
        "schema_version": FINAL_REVIEW_SCHEMA,
        "scope": FINAL_REVIEW_SCOPE,
        "status": "passed" if not findings else "blocked",
        "branch_document_set": sorted(document_set(get_path(reference, "meta.study_type"))),
        "candidate_hashes": candidate_hashes,
        "rendered_hashes": rendered_hashes,
        "verification_evidence_sha256": canonical_evidence_sha256(evidence),
        "review_bindings": review_bindings,
        "mandatory_visual_checks": list(VISUAL_CHECKS),
        "every_page": not findings,
        "section_three_and_orphan_heading_checks": not findings,
    }
    return {"status": final_review["status"], "findings": findings, "warnings": warnings, "resolved_evidence_cautions": resolutions, "internal_diagnostics": internal_history, "renderer": render_report.get("renderer"), "verification_evidence": evidence, "final_exact_artifact_review": final_review}


__all__ = ["CONTENT_CHECKS", "CROSS_DOCUMENT_CHECKS", "FORMAT_CONFORMANCE_MATRIX", "GOVERNED_GATE_SEQUENCE", "ICF_RETAINED_SHELL_SECTIONS", "PAGE_RENDERER_BACKENDS", "RECOVERY_POLICIES", "RESPONSE_SCHEMA", "VISUAL_CHECKS", "advance_gate_ledger", "audit_format_conformance_outputs", "build_gate_ledger", "canonical_evidence_sha256", "content_review_sections", "create_verification_requests", "deterministic_content_check", "load_format_conformance_matrix", "normalized_docx_format_signature", "page_renderer", "page_renderers", "pending_verifications", "preflight", "quality_report", "rasterize_pdf", "recovery_finding", "render_assurance", "render_pages", "renderer", "renderers", "retry_gate_ledger", "sha256_file", "validate_gate_ledger", "verification_request_authentication_findings", "verification_request_hash_valid", "verification_request_ledger_record", "verification_recovery_request_findings", "verification_request_sha256", "verification_response_is_complete", "verification_response_is_terminal"]
