"""Source-bound AI content handoff. Native builders consume the plan literally.

This module checks transport and shapes, never clinical applicability or meaning.
It does not derive a plan, call a model, or fall back to source interpretation.
"""
from __future__ import annotations

import copy
import hashlib
import json
from pathlib import Path
from typing import Any, Mapping
from xml.etree import ElementTree as ET

from docx import Document

SCHEMA = "clinical-content-plan/v1"
KEY = "_ai_content_plan"
REPEATED = ("intervention", "arm_group", "primary_outcome", "secondary_outcome", "other_outcome", "location")


def digest(value: Any) -> str:
    return hashlib.sha256(json.dumps(value, sort_keys=True, separators=(",", ":"), ensure_ascii=False).encode()).hexdigest()


def plan_for(reference: Mapping[str, Any]) -> Mapping[str, Any] | None:
    plan = reference.get(KEY)
    return plan if isinstance(plan, Mapping) else None


def with_plan(reference: Mapping[str, Any], plan: Mapping[str, Any]) -> dict[str, Any]:
    result = copy.deepcopy(dict(reference))
    result[KEY] = copy.deepcopy(dict(plan))
    return result


def source_only(reference: Mapping[str, Any]) -> dict[str, Any]:
    return {key: copy.deepcopy(value) for key, value in reference.items() if key != KEY}


def _xml_leaf_paths(node: ET.Element, prefix: str = "") -> list[str]:
    result = []
    for child in node:
        path = f"{prefix}/{child.tag}" if prefix else child.tag
        result.extend(_xml_leaf_paths(child, path) if len(child) else [path])
    return result


def request_payload(repo: Path, reference: Mapping[str, Any], bundle: Mapping[str, Any], governing: Mapping[str, Any]) -> dict[str, Any]:
    from contracts import protocol_contract, icf_retained_sections
    from rendering import _all_paragraphs, TOKEN, _heading_level, _icf_heading_key, _STERLING_ICF_HEADING_KEYS, _ADVARRA_ICF_HEADING_KEYS

    branch = reference["meta"]["study_type"]
    templates, tokens, anchors = {}, {}, {}
    for artifact, record in bundle["contracted_templates"].items():
        document = Document(repo / record["path"])
        templates[artifact] = [{"index": index, "text": paragraph.text, "style": paragraph.style.name}
                               for index, paragraph in enumerate(document.paragraphs)]
        anchors[artifact] = [index for index, paragraph in enumerate(document.paragraphs)
                             if _heading_level(paragraph) is not None
                             or artifact == "icf" and _icf_heading_key(paragraph.text) in (_STERLING_ICF_HEADING_KEYS | _ADVARRA_ICF_HEADING_KEYS)
                             or "{visitsTable}" in paragraph.text]
        names = {token.strip("{}") for paragraph in _all_paragraphs(document) for token in TOKEN.findall(paragraph.text)}
        names.update(("date", "version", "investigatorName", "investigatorTitle", "facilityName", "facilityCity", "studySiteAddress") if artifact == "protocol" else ("studyCordinatorPhone", "studySiteAddress"))
        tokens[artifact] = sorted(names)
    xml_schema = None
    if bundle.get("prs_authority"):
        path = repo / bundle["prs_authority"]["generation_template"]["path"]
        root = ET.parse(path).getroot()
        study = root if root.tag == "clinical_study" else root.find(".//clinical_study")
        if study is None:
            raise ValueError("PRS template has no clinical_study element.")
        xml_schema = {"template": path.read_text(), "scalar_paths": _xml_leaf_paths(study),
                      "repeated": {tag: _xml_leaf_paths(study.find(tag)) for tag in REPEATED}}
    guidance = {name: (repo / "references" / name).read_text() for name in
                ("ai-content-plan.md", "quality-corrections.md", "sterling-clause-contract.json")}
    payload = {
        "schema_version": SCHEMA, "task": "document_content_planning", "batch_id": "document-content-plan",
        "approved_source": source_only(reference), "governing_resources": dict(governing),
        "templates": templates, "structural_paragraph_indices": anchors, "docx_field_names": tokens,
        "protocol_sections": [section.public() for section in protocol_contract(branch)],
        "consent_template_comments": json.loads((repo / "references/generation-only-n8n-prompts.json").read_text())["consent_template_comments"],
        "template_tables": {artifact: [[[cell.text for cell in row.cells] for row in table.rows] for table in Document(repo / record["path"]).tables] for artifact, record in bundle["contracted_templates"].items()},
        "retained_icf_sections": list(icf_retained_sections(branch, reference["meta"].get("icf_template", "Advarra"))),
        "fixed_boilerplate": json.loads((repo / "references/fixed-clinical-boilerplate.json").read_text()),
        "guidance": guidance, "xml_schema": xml_schema,
        "response_path": "content-plan/response.json",
        "instructions": "Read all guidance and the entire approved source. Return the complete content plan defined in ai-content-plan.md. AI owns every content decision; builders only render your explicit plan. The source remains unchanged. Section drafts will follow this shared plan and the original n8n prompts.",
    }
    payload["request_id"] = "content-plan-" + digest(payload)[:20]
    payload["request_sha256"] = digest(payload)
    return json.loads(json.dumps(payload))


def _text_map(value: Any, label: str) -> None:
    if not isinstance(value, dict) or any(not isinstance(k, str) or not isinstance(v, str) for k, v in value.items()):
        raise ValueError(f"{label} must be an object of string values, including explicit empty strings.")


def _blocks(value: Any, label: str) -> None:
    if not isinstance(value, list) or any(not isinstance(item, dict) or set(item) != {"text", "bullet"}
            or not isinstance(item["text"], str) or not isinstance(item["bullet"], bool) for item in value):
        raise ValueError(f"{label} must be an array of {{text: string, bullet: boolean}} blocks.")


def validate_plan(plan: Any, request: Mapping[str, Any]) -> dict[str, Any]:
    """Validate explicit structure. Empty/omitted clinical content is AI-owned."""
    keys = {"docx_fields", "visits", "tables", "source_sections", "investigator_agreement", "container_introductions", "shell_edits", "extra_icf_sections", "bibliography", "xml", "decisions"}
    if not isinstance(plan, dict) or set(plan) != keys:
        raise ValueError(f"Plan requires exactly these keys: {sorted(keys)}")
    if not isinstance(plan["docx_fields"], dict) or set(plan["docx_fields"]) != set(request["docx_field_names"]):
        raise ValueError("docx_fields must address every requested artifact.")
    for artifact, names in request["docx_field_names"].items():
        _text_map(plan["docx_fields"][artifact], f"docx_fields.{artifact}")
        if set(plan["docx_fields"][artifact]) != set(names):
            raise ValueError(f"docx_fields.{artifact} must contain every requested token exactly once; use explicit empty strings.")
    if not isinstance(plan["bibliography"], str) or not isinstance(plan["decisions"], list):
        raise ValueError("bibliography must be text and decisions an array of source-grounded rationales.")
    if not isinstance(plan["visits"], list):
        raise ValueError("visits must be an array.")
    for visit in plan["visits"]:
        _text_map(visit, "visit")
        if set(visit) != {"visit", "timing", "display_visit_number", "CRFnumber"}:
            raise ValueError("Each visit needs visit, timing, display_visit_number and CRFnumber strings.")
    if not isinstance(plan["tables"], dict) or set(plan["tables"]) != {"study-sites", "schedule-of-assessments", "sample-size-evidence"}:
        raise ValueError("Return all three table objects, with explicit empty rows for unused tables.")
    for name, table in plan["tables"].items():
        if not isinstance(table, dict) or set(table) != {"caption", "header_rows", "rows", "supplemental_notes"}:
            raise ValueError(f"Table {name} needs caption, header_rows, rows and supplemental_notes.")
        rows = table["rows"]
        if (not isinstance(table["caption"], str) or type(table["header_rows"]) is not int
                or table["header_rows"] not in (1, 2) or not isinstance(rows, list)
                or any(not isinstance(row, list) or not row or any(not isinstance(cell, str) for cell in row) for row in rows)
                or rows and (len({len(row) for row in rows}) != 1 or len(rows) < table["header_rows"])
                or not isinstance(table["supplemental_notes"], list)
                or any(not isinstance(note, str) for note in table["supplemental_notes"])):
            raise ValueError(f"Table {name} has malformed rectangular rows or notes.")
        if rows and name == "study-sites" and len(rows[0]) != 3:
            raise ValueError("The site-directory layout requires three columns.")
        if rows and table["header_rows"] == 2 and rows[1][0]:
            raise ValueError("The second header row starts with an empty merged-cell continuation.")
    assessment = plan["tables"]["schedule-of-assessments"]
    if (assessment["rows"] or assessment["supplemental_notes"]) and not any(
        "{visitsTable}" in paragraph["text"] for paragraph in request["templates"]["protocol"]
    ):
        raise ValueError("This template has no assessment-table anchor; use its requested methods prose section.")
    expected = {s["section_id"] for s in request["protocol_sections"] if s["role"] == "source"}
    if not isinstance(plan["source_sections"], dict) or set(plan["source_sections"]) != expected:
        raise ValueError("source_sections needs an explicit include decision for every listed source-role section.")
    for name, section in plan["source_sections"].items():
        if not isinstance(section, dict) or set(section) != {"include", "blocks"} or type(section["include"]) is not bool:
            raise ValueError(f"Section {name} needs include and blocks.")
        _blocks(section["blocks"], name)
    _blocks(plan["investigator_agreement"], "investigator_agreement")
    if not isinstance(plan["container_introductions"], dict):
        raise ValueError("container_introductions must be an object.")
    containers = {s["section_id"] for s in request["protocol_sections"] if s["role"] == "container"}
    if set(plan["container_introductions"]) - containers:
        raise ValueError("Unknown protocol container ID.")
    for name, blocks in plan["container_introductions"].items():
        _blocks(blocks, name)
    if not isinstance(plan["shell_edits"], dict) or set(plan["shell_edits"]) != set(request["templates"]):
        raise ValueError("shell_edits must address each selected template.")
    for artifact, edits in plan["shell_edits"].items():
        if not isinstance(edits, list):
            raise ValueError("shell_edits entries must be arrays.")
        indices = []
        for edit in edits:
            if (not isinstance(edit, dict) or set(edit) != {"index", "text"} or type(edit["index"]) is not int
                    or edit["index"] < 0 or edit["index"] >= len(request["templates"][artifact])
                    or edit["text"] is not None and not isinstance(edit["text"], str)):
                raise ValueError("Shell edits require a valid original paragraph index and string text (null removes it).")
            if edit["index"] in request["structural_paragraph_indices"][artifact] and edit["text"] != request["templates"][artifact][edit["index"]]["text"]:
                raise ValueError("Shell edits must preserve structural headings and table anchors; use source_sections for optional headings.")
            indices.append(edit["index"])
        if len(indices) != len(set(indices)):
            raise ValueError("Duplicate shell edit index.")
    if not isinstance(plan["extra_icf_sections"], list):
        raise ValueError("extra_icf_sections must be an array.")
    from rendering import _icf_heading_key
    icf_anchors = {
        _icf_heading_key(request["templates"]["icf"][index]["text"])
        for index in request["structural_paragraph_indices"].get("icf", [])
    }
    for section in plan["extra_icf_sections"]:
        if (not isinstance(section, dict) or set(section) != {"heading", "before", "blocks"}
                or not isinstance(section["heading"], str) or not isinstance(section["before"], str)):
            raise ValueError("Extra consent sections need heading, before and blocks.")
        if _icf_heading_key(section["before"]) not in icf_anchors:
            raise ValueError("Extra consent sections must name an existing native ICF heading as their insertion anchor.")
        _blocks(section["blocks"], section["heading"])
    xml = plan["xml"]
    if request["xml_schema"] is None:
        if xml is not None:
            raise ValueError("This branch has no XML output; xml must be null.")
    else:
        if not isinstance(xml, dict) or set(xml) != {"design_branch", "scalars", "repeated"}:
            raise ValueError("xml requires design_branch, scalars and repeated.")
        if xml["design_branch"] not in {"observational_design", "interventional_design"}:
            raise ValueError("Unknown XML design branch.")
        _text_map(xml["scalars"], "xml.scalars")
        omitted_branch = "interventional_design" if xml["design_branch"] == "observational_design" else "observational_design"
        allowed = {p for p in request["xml_schema"]["scalar_paths"] if p.split('/')[0] not in REPEATED and not p.startswith('study_design/' + omitted_branch + '/')}
        if set(xml["scalars"]) != allowed:
            raise ValueError("XML scalar paths must match every leaf of the selected template branch exactly.")
        if not isinstance(xml["repeated"], dict) or set(xml["repeated"]) != set(REPEATED):
            raise ValueError("Return each XML repeated-record array explicitly.")
        for tag, records in xml["repeated"].items():
            if not isinstance(records, list):
                raise ValueError(f"XML {tag} must be an array.")
            for record in records:
                if not isinstance(record, dict):
                    raise ValueError(f"XML {tag} record must be an object.")
                for path, value in record.items():
                    if tag == "intervention" and path == "arm_group_label" and isinstance(value, list) and all(isinstance(v, str) for v in value):
                        continue
                    if not isinstance(value, str):
                        raise ValueError(f"XML {tag}/{path} must be a string.")
                if set(record) != set(request["xml_schema"]["repeated"][tag]):
                    raise ValueError(f"XML {tag} record paths must match its template leaves.")
    return copy.deepcopy(plan)


def handoff(repo: Path, workspace: Path, reference: Mapping[str, Any], bundle: Mapping[str, Any], governing: Mapping[str, Any]) -> tuple[dict[str, Any], dict[str, Any] | None]:
    folder = workspace / "content-plan"
    folder.mkdir(parents=True, exist_ok=True)
    request = request_payload(repo, reference, bundle, governing)
    path = folder / "request.json"
    if path.exists() and json.loads(path.read_text()) != request:
        raise ValueError("Content-plan request changed; start a new generation attempt.")
    path.write_text(json.dumps(request, indent=2, ensure_ascii=False) + "\n")
    response_path = workspace / request["response_path"]
    if not response_path.is_file():
        return request, None
    response = json.loads(response_path.read_text())
    if not isinstance(response, dict):
        raise ValueError("Content-plan response must be a JSON object.")
    if (response.get("schema_version") != SCHEMA or response.get("request_id") != request["request_id"]
            or response.get("request_sha256") != request["request_sha256"]
            or not isinstance(response.get("producer"), dict)
            or not isinstance(response["producer"].get("model_id"), str) or not response["producer"]["model_id"].strip()):
        raise ValueError("Content-plan response identity or producer is missing or mismatched.")
    plan = validate_plan(response.get("plan"), request)
    receipt = {"request_sha256": request["request_sha256"], "response_sha256": digest(response), "plan_sha256": digest(plan)}
    accepted = folder / "accepted.json"
    if accepted.exists() and json.loads(accepted.read_text()) != receipt:
        raise ValueError("Accepted content plan changed; start a new generation attempt.")
    accepted.write_text(json.dumps(receipt, indent=2) + "\n")
    return request, plan


def read_plan(workspace: Path) -> dict[str, Any] | None:
    folder = workspace / "content-plan"
    if not (folder / "accepted.json").exists():
        return None
    request = json.loads((folder / "request.json").read_text())
    response = json.loads((folder / "response.json").read_text())
    receipt = json.loads((folder / "accepted.json").read_text())
    if (digest(response) != receipt["response_sha256"] or digest(response["plan"]) != receipt["plan_sha256"]
            or request["request_sha256"] != receipt["request_sha256"]):
        raise ValueError("Accepted content-plan integrity mismatch.")
    return validate_plan(response["plan"], request)
