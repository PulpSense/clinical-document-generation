#!/usr/bin/env python3
"""Refresh and align static DOCX table-of-contents entries from a rendered PDF."""

from __future__ import annotations

import argparse
import json
import re
import shutil
import tempfile
import zipfile
from pathlib import Path
from typing import Any

from audit_static_toc import audit, normalize


def require_document():
    try:
        from docx import Document
        from docx.enum.text import WD_TAB_ALIGNMENT, WD_TAB_LEADER
    except ImportError as exc:
        raise SystemExit("python-docx is required to refresh DOCX TOC entries.") from exc
    return Document, WD_TAB_ALIGNMENT, WD_TAB_LEADER


def set_word_field_refresh_flags(path: Path) -> None:
    def add_dirty_flag(match: re.Match[str]) -> str:
        field_start = re.sub(r"\s*/\s*$", "", match.group(1))
        return field_start + ' w:dirty="true"/>'

    with tempfile.NamedTemporaryFile(delete=False, suffix=".docx") as tmp:
        tmp_path = Path(tmp.name)
    try:
        with zipfile.ZipFile(path, "r") as zin, zipfile.ZipFile(tmp_path, "w", zipfile.ZIP_DEFLATED) as zout:
            for item in zin.infolist():
                data = zin.read(item.filename)
                if item.filename == "word/settings.xml":
                    xml = data.decode("utf-8")
                    if "<w:updateFields" not in xml:
                        xml = xml.replace("</w:settings>", '<w:updateFields w:val="true"/></w:settings>')
                    else:
                        xml = re.sub(r"<w:updateFields\b[^>]*/>", '<w:updateFields w:val="true"/>', xml)
                    data = xml.encode("utf-8")
                elif item.filename.startswith("word/header") and b"PAGE" in data:
                    xml = data.decode("utf-8")
                    xml = re.sub(
                        r'(<w:fldChar\b(?=[^>]*w:fldCharType="begin")(?:(?!w:dirty=)[^>])*)/?>',
                        add_dirty_flag,
                        xml,
                        count=1,
                    )
                    xml = re.sub(r'w:dirty="false"', 'w:dirty="true"', xml)
                    data = xml.encode("utf-8")
                zout.writestr(item, data)
        tmp_path.replace(path)
    finally:
        if tmp_path.exists():
            tmp_path.unlink()


def update_paragraph_text(paragraph: Any, new_text: str) -> None:
    if paragraph.runs:
        paragraph.runs[0].text = new_text
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(new_text)


def toc_tab_position(document: Any) -> Any:
    section = document.sections[0]
    return section.page_width - section.left_margin - section.right_margin


def has_right_dot_leader_tab(paragraph: Any, tab_position: Any, alignment: Any, leader: Any) -> bool:
    for tab in paragraph.paragraph_format.tab_stops:
        if tab.alignment != alignment.RIGHT or tab.leader != leader.DOTS:
            continue
        # Allow a tiny tolerance for renderer/library round-tripping.
        if abs(int(tab.position) - int(tab_position)) < 20:
            return True
    return False


def align_toc_entry(
    paragraph: Any,
    title: str,
    page: int,
    tab_position: Any,
    alignment: Any,
    leader: Any,
) -> bool:
    current = paragraph.text.replace("\u00a0", " ").strip()
    replacement = f"{normalize(title)}\t{page}"
    already_aligned = "\t" in current and has_right_dot_leader_tab(paragraph, tab_position, alignment, leader)

    update_paragraph_text(paragraph, replacement)
    tab_stops = paragraph.paragraph_format.tab_stops
    tab_stops.clear_all()
    tab_stops.add_tab_stop(tab_position, alignment.RIGHT, leader.DOTS)
    paragraph.paragraph_format.left_indent = None
    paragraph.paragraph_format.right_indent = None
    paragraph.paragraph_format.first_line_indent = None
    return current != replacement or not already_aligned


def refresh_docx(docx_path: Path, pdf_path: Path, output_path: Path) -> dict[str, Any]:
    result = audit(docx_path, pdf_path)
    if result["missing_count"]:
        missing = ", ".join(item["title"] for item in result["missing"])
        raise SystemExit(f"Cannot refresh TOC because rendered headings were not found: {missing}")
    if result["entry_count"] == 0:
        raise SystemExit("Cannot refresh TOC because no static TOC entries were found.")

    if output_path.resolve() != docx_path.resolve():
        output_path.parent.mkdir(parents=True, exist_ok=True)
        shutil.copyfile(docx_path, output_path)
    else:
        output_path = docx_path

    Document, WD_TAB_ALIGNMENT, WD_TAB_LEADER = require_document()
    document = Document(str(output_path))
    tab_position = toc_tab_position(document)
    updated = []
    aligned_count = 0
    for entry in result["entries"]:
        paragraph = document.paragraphs[entry["paragraph_index"]]
        actual_page = entry["actual_page"]
        if actual_page is None:
            continue
        if actual_page != entry["listed_page"]:
            updated.append(
                {
                    "title": entry["title"],
                    "old_page": entry["listed_page"],
                    "new_page": actual_page,
                }
            )
        if align_toc_entry(
            paragraph,
            entry["title"],
            actual_page,
            tab_position,
            WD_TAB_ALIGNMENT,
            WD_TAB_LEADER,
        ):
            aligned_count += 1

    document.save(output_path)
    set_word_field_refresh_flags(output_path)
    return {
        **result,
        "output_docx": str(output_path),
        "updated_count": len(updated),
        "aligned_count": aligned_count,
        "alignment": "right_tab_with_dot_leader",
        "updates": updated,
    }


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--docx", required=True, help="DOCX containing static TOC entries.")
    parser.add_argument("--pdf", required=True, help="Rendered PDF used as the page source of truth.")
    parser.add_argument("--output", help="Output DOCX path. Defaults to updating --docx in place.")
    parser.add_argument("--report", help="Optional JSON refresh report path.")
    args = parser.parse_args()

    docx_path = Path(args.docx).expanduser().resolve()
    pdf_path = Path(args.pdf).expanduser().resolve()
    output_path = Path(args.output).expanduser().resolve() if args.output else docx_path
    result = refresh_docx(docx_path, pdf_path, output_path)
    rendered = json.dumps(result, indent=2)
    if args.report:
        report_path = Path(args.report).expanduser().resolve()
        report_path.parent.mkdir(parents=True, exist_ok=True)
        report_path.write_text(rendered + "\n", encoding="utf-8")
    print(rendered)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
