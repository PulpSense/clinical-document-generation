#!/usr/bin/env python3
"""The single public lifecycle for the clinical-document skill."""

from __future__ import annotations

import argparse
import base64
import binascii
import copy
import hashlib
import json
import math
import os
import platform
import pwd
import re
import select
import shlex
import shutil
import signal
import socket
import socketserver
import subprocess
import sys
import tarfile
import tempfile
import threading
import time
import zipfile
from concurrent.futures import ThreadPoolExecutor
from dataclasses import dataclass
from datetime import date, datetime, timezone
from pathlib import Path, PurePosixPath
from typing import Any, Callable, Iterable, Mapping, Sequence

from docx import Document

SCRIPT_DIR = Path(__file__).resolve().parent
if str(SCRIPT_DIR) not in sys.path: sys.path.insert(0, str(SCRIPT_DIR))

from contracts import BUNDLED_FONT_FILES, RECOVERY_POLICIES, ContractedTemplateBundleError, LAYOUT_REPAIR_RULES, batch_plan, canonical_study_type, contracted_template_bundle, document_set, get_path, icf_contract, parse_source_truth, protocol_contract, recovery_finding, repair_report, set_path, source_contract, source_truth_markdown
from drafting import MAX_ATTEMPTS, accepted_cross_section_duplicate_findings, governing_resources, ingest_responses, invalidate_accepted_targets, merged_drafts, missing_drafts, pending_requests, recorded_acceptance_response, retry_attempts, schedule_requests, sha256_file, sha256_value
from prs_xml import generate as generate_xml
from quality import CERTIFICATION_CASE_ORDER, CERTIFICATION_EVIDENCE_MAX_FILES, CERTIFICATION_EVIDENCE_MAX_ITEM_BYTES, CERTIFICATION_EVIDENCE_MAX_TOTAL_BYTES, CERTIFICATION_VISUAL_CHECKS, CONTENT_CHECKS, DETERMINISTIC_BRANCH_ACCEPTANCE_CASES, GOVERNED_GATE_SEQUENCE, RELEASE_CERTIFICATION_PUBLIC_KEY, RELEASE_CERTIFICATION_SIGNATURE_ALGORITHM, RELEASE_CERTIFICATION_TRUSTED_KEY_ID, RESPONSE_SCHEMA, VISUAL_CHECKS, _approved_packaged_font_fallback, _certification_evidence_findings, _manifest_package_fingerprint, _pdfium_runtime_integrity, _template_fonts, _validated_certification_evidence, advance_gate_ledger, audit_format_conformance_outputs, build_gate_ledger, canonical_evidence_sha256, create_verification_requests, load_format_conformance_matrix, page_renderers, pending_verifications, quality_report, release_certification_attestation_findings, release_certification_key_id, release_certification_payload, render_assurance, renderer, renderers, run_pdfium_worker, sha256_file as quality_sha256, validate_gate_ledger, verification_response_is_complete
from rendering import render_documents


REFERENCE = Path("reference/study.reference.json")
MAX_VERIFICATION_ATTEMPTS = 3
DESKTOP_DELIVERY_RETRIES = 2
NORMAL_RUNTIME_TARGET_MIN_SECONDS = 600.0
NORMAL_RUNTIME_TARGET_MAX_SECONDS = 720.0
DESKTOP_OPERATION_BUDGET_SECONDS = 1800.0
FORMAT_CONFORMANCE_TIMEOUT_SECONDS = 10 * 60
DESKTOP_STAGE_SOFT_BUDGETS = {
    "drafting": 480.0,
    "candidate": 120.0,
    "render_assurance": 300.0,
    "independent_verification": 240.0,
    "delivery": 60.0,
    "desktop_delivery": 60.0,
}
RELEASE_MANIFEST = "RELEASE-MANIFEST.json"
RELEASE_CERTIFICATION = "RELEASE-CERTIFICATION.json"
CERTIFICATION_REPORT_MAX_BYTES = (
    (CERTIFICATION_EVIDENCE_MAX_TOTAL_BYTES * 4 + 2) // 3
    + 16 * 1024 * 1024
)
RELEASE_ARCHIVE_MAX_MEMBERS = 1024
RELEASE_ARCHIVE_MAX_MEMBER_BYTES = 64 * 1024 * 1024
RELEASE_ARCHIVE_MAX_TOTAL_BYTES = CERTIFICATION_REPORT_MAX_BYTES + 128 * 1024 * 1024
RELEASE_ARCHIVE_MAX_COMPRESSION_RATIO = 100
INSTALLATION_ASSURANCE = "INSTALLATION-ASSURANCE.json"
PROMOTION_RECORD = "PROMOTION-RECORD.json"
MINIMUM_PYTHON_VERSION = (3, 10)
PDF_PAGE_RENDERER = {
    "kind": "pypdfium2",
    "version": "5.13.0",
    "wheel": "assets/runtime-wheels/pypdfium2-5.13.0-py3-none-macosx_13_0_arm64.whl",
    "wheel_sha256": "da5c7b74eebf40b5c1fbe1de01aa1edc8827a79fb1efd999616bc20dcaf77ba4",
    "platform": "macosx_13_0_arm64",
}
PRODUCTION_MODULES = {
    "contracts.py", "drafting.py", "prs_xml.py", "quality.py", "rendering.py", "workflow.py",
}
CERTIFIED_HERMES_CONFIGURATION = {
    "source": "clinical-release-certification",
    "max_turns": 80,
    "skill": "clinical-document-generation",
    "safe_mode": True,
    "model_identifier": "gpt-5.6-sol",
    "reasoning_configuration": "Hermes Desktop governed default",
}
CERTIFICATION_GATES = {
    "source", "content", "document_structure", "prs_xml", "package",
    "cross_document_consistency", "render_assurance", "every_page_visual_qa",
    "delivery_confirmation",
}
class OperationDeadlineExpired(RuntimeError):
    """Raised before an atomic publication would cross the operation deadline."""


def _runtime_version(value: Mapping[str, Any]) -> tuple[int, int, int]:
    raw = value.get("version_info")
    if isinstance(raw, list) and len(raw) >= 2:
        parts = [int(item) for item in raw[:3]]
    else:
        parts = [int(item) for item in str(value.get("version") or "").split(".")[:3]]
    return tuple((parts + [0, 0, 0])[:3])


def _current_python_runtime() -> dict[str, Any]:
    runtime = {
        "executable": str(Path(sys.executable).absolute()),
        "implementation": platform.python_implementation(),
        "version": platform.python_version(),
        "version_info": list(sys.version_info[:3]),
    }
    if _runtime_version(runtime) < (*MINIMUM_PYTHON_VERSION, 0):
        raise RuntimeError("The Desktop operation requires an explicitly resolved Python 3.10+ runtime.")
    return runtime


def resolve_python_runtime(
    *,
    candidates: list[Path] | None = None,
    environment: Mapping[str, str] | None = None,
) -> dict[str, Any]:
    """Resolve and identify an explicit Python 3.10+ executable for Desktop use."""
    runtime_environment = os.environ if environment is None else environment
    if candidates is None:
        requested = runtime_environment.get("HERMES_PYTHON")
        discovered = [Path(requested)] if requested else []
        discovered.append(Path(sys.executable))
        search_path = runtime_environment.get("PATH")
        for name in ("python3.13", "python3.12", "python3.11", "python3.10", "python3"):
            if executable := shutil.which(name, path=search_path):
                discovered.append(Path(executable))
        candidates = discovered
    probe = (
        "import json,platform,sys;"
        "print(json.dumps({'executable':str(__import__('pathlib').Path(sys.executable).absolute()),"
        "'implementation':platform.python_implementation(),'version':platform.python_version(),"
        "'version_info':list(sys.version_info[:3])}))"
    )
    seen: set[str] = set()
    failures: list[str] = []
    for candidate in candidates:
        key = str(candidate)
        if key in seen:
            continue
        seen.add(key)
        try:
            completed = subprocess.run(
                [key, "-c", probe],
                capture_output=True,
                text=True,
                timeout=5.0,
                check=False,
                env=dict(runtime_environment),
            )
            runtime = json.loads(completed.stdout) if completed.returncode == 0 else {}
            if not isinstance(runtime, dict) or _runtime_version(runtime) < (*MINIMUM_PYTHON_VERSION, 0):
                failures.append(f"{key}: unsupported Python runtime")
                continue
            return {
                "executable": str(runtime["executable"]),
                "implementation": str(runtime["implementation"]),
                "version": str(runtime["version"]),
                "version_info": list(_runtime_version(runtime)),
            }
        except (OSError, ValueError, json.JSONDecodeError, subprocess.SubprocessError) as exc:
            failures.append(f"{key}: {exc}")
    detail = "; ".join(failures) or "no candidates were found"
    raise RuntimeError(f"No supported Python 3.10+ Desktop runtime is available: {detail}")


def _desktop_deadline_state(
    persisted: Mapping[str, Any],
    *,
    budget_seconds: float,
    epoch_now: float,
) -> tuple[float, float, float, list[dict[str, Any]], list[dict[str, Any]]]:
    """Validate portable deadline state, failing closed when it is unsafe."""
    if not persisted:
        return epoch_now, epoch_now + budget_seconds, budget_seconds, [], []
    fail_closed = (epoch_now - budget_seconds, epoch_now, budget_seconds, [], [])
    if "started_at_epoch" not in persisted or "deadline_at_epoch" not in persisted:
        return fail_closed
    try:
        started_at_epoch = float(persisted["started_at_epoch"])
        deadline_at_epoch = float(persisted["deadline_at_epoch"])
        persisted_budget = float(
            persisted.get("budget_seconds", deadline_at_epoch - started_at_epoch)
        )
        if not all(math.isfinite(value) for value in (
            started_at_epoch, deadline_at_epoch, persisted_budget,
        )):
            raise ValueError("Desktop deadline state contains a non-finite value.")
        if deadline_at_epoch < started_at_epoch or persisted_budget <= 0:
            raise ValueError("Desktop deadline state has an invalid interval.")
        raw_runtime_history = persisted.get("runtime_history", [])
        if not isinstance(raw_runtime_history, list) or any(
            not isinstance(item, Mapping) for item in raw_runtime_history
        ):
            raise ValueError("Desktop runtime history must be a list of objects.")
        runtime_history = [dict(item) for item in raw_runtime_history]
        raw_stage_history = persisted.get("stage_history", [])
        if not isinstance(raw_stage_history, list) or any(
            not isinstance(item, Mapping) for item in raw_stage_history
        ):
            raise ValueError("Desktop stage history must be a list of objects.")
        stage_history = [dict(item) for item in raw_stage_history]
        for field in ("attempt_counters", "stage_timings", "soft_budgets", "cleanup"):
            if field in persisted and not isinstance(persisted[field], Mapping):
                raise ValueError(f"Desktop operation {field} must be an object.")
        for field in ("soft_budget_events", "pending_handoffs"):
            if field in persisted and not isinstance(persisted[field], list):
                raise ValueError(f"Desktop operation {field} must be a list.")
    except (TypeError, ValueError, OverflowError):
        return fail_closed
    return (
        started_at_epoch,
        deadline_at_epoch,
        persisted_budget,
        runtime_history,
        stage_history,
    )


def performance_classification(elapsed_seconds: float) -> str:
    """Classify measured runtime without turning the target into a timeout."""
    if elapsed_seconds > DESKTOP_OPERATION_BUDGET_SECONDS:
        return "deadline_exceeded"
    if elapsed_seconds < NORMAL_RUNTIME_TARGET_MIN_SECONDS:
        return "below_target_window"
    if elapsed_seconds <= NORMAL_RUNTIME_TARGET_MAX_SECONDS:
        return "target_window"
    return "above_target_within_deadline"


def _active_release_identity(skill_root: Path) -> dict[str, Any]:
    """Identify the immutable release that owns a Desktop operation."""
    manifest_path = skill_root / RELEASE_MANIFEST
    if manifest_path.is_file():
        manifest = _read(manifest_path)
        fingerprint = str(manifest.get("package_fingerprint") or "")
        if not fingerprint:
            raise RuntimeError("The active release manifest has no package fingerprint.")
        return {
            "package_fingerprint": fingerprint,
            "git_commit": manifest.get("git_commit"),
            "source": "promoted_release",
        }
    implementation = [
        {
            "path": path.relative_to(skill_root).as_posix(),
            "sha256": sha256_file(path),
        }
        for path in sorted((skill_root / "scripts").glob("*.py"))
    ]
    return {
        "package_fingerprint": sha256_value(implementation),
        "git_commit": None,
        "source": "controlled_checkout",
    }


def _release_excluded(path: Path) -> bool:
    """Return whether a path belongs to development-only or sensitive data."""
    parts = set(path.parts)
    if parts & {".git", ".scratch", "runs", ".hermes", ".test-venv", ".venv", "venv", "runtime", "__pycache__", ".pytest_cache", ".mypy_cache", ".ruff_cache", "tests", "input", "inputs", "source-data", "patient-data", "evidence", "output"}:
        return True
    name = path.name.casefold()
    if name in {".env", ".env.local", "artifact.md", ".coverage", INSTALLATION_ASSURANCE.casefold()} or name.endswith((".pem", ".key", ".p12", ".pfx", ".sqlite", ".sqlite3")):
        return True
    return False


def _pdfium_wheel_inventory(wheel: Path) -> list[dict[str, Any]]:
    """Derive the deterministic extraction inventory trusted by the release."""
    entries: list[dict[str, Any]] = []
    seen: set[str] = set()
    with zipfile.ZipFile(wheel) as archive:
        for member in archive.infolist():
            raw = (
                member.filename[:-1]
                if member.is_dir() and member.filename.endswith("/")
                else member.filename
            )
            relative = PurePosixPath(raw)
            normalized = relative.as_posix()
            if (
                not normalized
                or raw != normalized
                or relative.is_absolute()
                or "\\" in raw
                or any(part in {"", ".", ".."} for part in raw.split("/"))
                or normalized.casefold() in seen
            ):
                raise ValueError(
                    "The pinned pypdfium2 wheel has an unsafe or duplicate extraction path."
                )
            seen.add(normalized.casefold())
            if member.is_dir():
                continue
            if (member.external_attr >> 16) & 0o170000 == 0o120000:
                raise ValueError("The pinned pypdfium2 wheel contains an unsupported symbolic link.")
            payload = archive.read(member)
            if len(payload) != member.file_size:
                raise ValueError("The pinned pypdfium2 wheel contains an inconsistent file size.")
            entries.append({
                "path": normalized,
                "bytes": len(payload),
                "sha256": hashlib.sha256(payload).hexdigest(),
            })
    if not entries:
        raise ValueError("The pinned pypdfium2 wheel has no extractable runtime files.")
    return sorted(entries, key=lambda item: item["path"])


def _package_release_tree(
    repo_root: Path,
    output_path: Path,
    *,
    git_commit: str,
    certification_public_key: Mapping[str, Any] | None = None,
) -> dict[str, Any]:
    """Create a deterministic, installable Hermes skill archive.

    The archive contains only the skill runtime and its governed resources. The
    manifest is written into the archive and hashes every other member, making
    the installed candidate auditable without including source inputs or runs.
    """
    repo_root = repo_root.resolve()
    output_path = output_path.expanduser().resolve()
    if output_path == repo_root or repo_root in output_path.parents:
        raise ValueError("Release archive must be outside the skill repository.")
    actual_modules = {path.name for path in (repo_root / "scripts").glob("*.py")}
    if actual_modules != PRODUCTION_MODULES:
        raise ValueError(
            "Release packaging requires exactly the six production Python modules: "
            + ", ".join(sorted(PRODUCTION_MODULES))
        )
    exact_files = {"SKILL.md", "README.md", "CONTEXT.md", "requirements.txt", "agents/openai.yaml"}
    files = [
        path for path in sorted(repo_root.rglob("*"))
        if path.is_file()
        and (
            path.relative_to(repo_root).as_posix() in exact_files
            or path.relative_to(repo_root).parts[0] in {"assets", "references"}
            or (
                path.relative_to(repo_root).parts[0] == "scripts"
                and path.name in PRODUCTION_MODULES
            )
        )
    ]
    if not files:
        raise ValueError("No release files were found.")
    file_bytes = {path: path.read_bytes() for path in files}
    if certification_public_key is not None:
        if "private_exponent" in certification_public_key:
            raise ValueError("Release packaging cannot receive certification private-key material.")
        key_path = repo_root / RELEASE_CERTIFICATION_PUBLIC_KEY
        if key_path not in file_bytes:
            raise ValueError("The release certification public-key resource is missing.")
        file_bytes[key_path] = (
            json.dumps(certification_public_key, indent=2, ensure_ascii=False) + "\n"
        ).encode("utf-8")
    entries = []
    for path in files:
        relative = path.relative_to(repo_root).as_posix()
        content = file_bytes[path]
        entries.append({"path": relative, "sha256": hashlib.sha256(content).hexdigest(), "bytes": len(content)})
    bundle_references = (
        {"meta": {"study_type": "Prospective", "icf_template": "Advarra"}},
        {"meta": {"study_type": "Prospective", "icf_template": "Sterling"}},
        {"meta": {"study_type": "Ambispective", "icf_template": "Advarra"}},
        {"meta": {"study_type": "Ambispective", "icf_template": "Sterling"}},
        {"meta": {"study_type": "Retrospective"}},
    )
    bundles = [contracted_template_bundle(repo_root, reference) for reference in bundle_references]
    governed_resources = sorted({
        relative
        for bundle in bundles
        for relative in bundle["resource_hashes"]
    })
    templates = [repo_root / relative for relative in governed_resources if Path(relative).suffix.casefold() == ".docx"]
    font_inventory = {}
    for path in templates:
        if path.suffix.casefold() == ".docx":
            font_inventory[path.relative_to(repo_root).as_posix()] = sorted(_template_fonts(path))
    required_font_names = sorted({font for fonts in font_inventory.values() for font in fonts})
    approved_font_plan = bundles[0]["approved_font_plan"]
    certification_configuration_sha256 = {}
    for fixture_id in CERTIFICATION_CASE_ORDER:
        fixture_path = repo_root / "tests/fixtures/release-certification" / fixture_id / "fixture.json"
        try:
            fixture = _read(fixture_path)
            certification_configuration_sha256[fixture_id] = sha256_value(fixture["hermes_configuration"])
        except (OSError, KeyError, TypeError, ValueError, json.JSONDecodeError) as exc:
            raise ValueError(f"Release packaging requires the governed certification fixture: {fixture_id}") from exc
    pdf_renderer: dict[str, Any] = dict(PDF_PAGE_RENDERER)
    pdf_renderer_wheel = repo_root / pdf_renderer["wheel"]
    if (
        not pdf_renderer_wheel.is_file()
        or sha256_file(pdf_renderer_wheel) != pdf_renderer["wheel_sha256"]
    ):
        raise ValueError(
            "The pinned pypdfium2 wheel is missing or does not match its governed hash."
        )
    pdf_renderer["runtime_inventory"] = _pdfium_wheel_inventory(pdf_renderer_wheel)
    implementation_files = sorted(f"scripts/{name}" for name in PRODUCTION_MODULES)
    manifest = {
        "schema_version": "hermes-release-manifest/v2",
        "git_commit": git_commit,
        "package_root": "clinical-document-generation",
        "package_purpose": "Installable runtime for the reviewed clinical document workflow.",
        "installation": {
            "entrypoint": "SKILL.md",
            "runtime": "Python 3.10+",
            "dependencies": "requirements.txt",
            "install_as_direct_child_of": "Hermes skills directory",
            "activation": "atomic after end-to-end Render Assurance smoke; previous verified release retained",
            "required_external_tools": ["Microsoft Word or LibreOffice"],
        },
        "inventory": {
            "implementation": implementation_files,
            "contracted_template_bundles": bundles,
            "governed_resources": governed_resources,
            "font_identities": font_inventory,
            "font_fallbacks": {font: [_approved_packaged_font_fallback(font, approved_font_plan)] for font in required_font_names},
            "renderer_at_packaging": renderer(environment=os.environ),
            "pdf_page_renderer": pdf_renderer,
            "harness": {"python": platform.python_version(), "platform": platform.platform()},
            "model": "Hermes Desktop runtime; model identity is recorded per generation evidence.",
            "certification_configuration_sha256": certification_configuration_sha256,
        },
        "excluded_classes": ["git metadata", "development virtual environments", "credentials", "patient/source data", "old run outputs", "development tests", "installed runtime and assurance evidence"],
        "files": entries,
    }
    manifest["package_fingerprint"] = _manifest_package_fingerprint(manifest)[1]
    manifest_text = json.dumps(manifest, indent=2, ensure_ascii=False) + "\n"
    output_path.parent.mkdir(parents=True, exist_ok=True)
    temporary = output_path.with_name(f".{output_path.name}.tmp")
    try:
        with zipfile.ZipFile(temporary, "w", compression=zipfile.ZIP_DEFLATED, compresslevel=9) as archive:
            for path in files:
                relative = path.relative_to(repo_root).as_posix()
                info = zipfile.ZipInfo(f"clinical-document-generation/{relative}", date_time=(1980, 1, 1, 0, 0, 0))
                info.compress_type = zipfile.ZIP_DEFLATED
                archive.writestr(info, file_bytes[path])
            info = zipfile.ZipInfo(f"clinical-document-generation/{RELEASE_MANIFEST}", date_time=(1980, 1, 1, 0, 0, 0))
            info.compress_type = zipfile.ZIP_DEFLATED
            archive.writestr(info, manifest_text.encode("utf-8"))
        os.replace(temporary, output_path)
    finally:
        temporary.unlink(missing_ok=True)
    return {
        "status": "passed",
        "package": output_path.as_posix(),
        "package_fingerprint": manifest["package_fingerprint"],
        "git_commit": git_commit,
        "file_count": len(entries),
        "archive_bytes": output_path.stat().st_size,
        "manifest": f"clinical-document-generation/{RELEASE_MANIFEST}",
    }


def package_release(
    repo_root: Path,
    output_path: Path,
    *,
    certification_public_key: Mapping[str, Any] | None = None,
) -> dict[str, Any]:
    """Package the exact committed tree, never mutable checkout bytes."""
    repo_root = repo_root.resolve()
    output_path = output_path.expanduser().resolve()
    if output_path == repo_root or repo_root in output_path.parents:
        raise ValueError("Release archive must be outside the skill repository.")
    try:
        git_root = Path(subprocess.run(
            ["git", "-C", str(repo_root), "rev-parse", "--show-toplevel"],
            check=True,
            capture_output=True,
            text=True,
        ).stdout.strip()).resolve()
        commit = subprocess.run(
            ["git", "-C", str(repo_root), "rev-parse", "--verify", "HEAD^{commit}"],
            check=True,
            capture_output=True,
            text=True,
        ).stdout.strip()
    except (OSError, subprocess.CalledProcessError) as exc:
        raise ValueError("A release must be built from a committed Git tree.") from exc
    if git_root != repo_root:
        raise ValueError("Release packaging must target the Git repository root.")
    try:
        status_output = subprocess.run(
            [
                "git",
                "-C",
                str(repo_root),
                "status",
                "--porcelain=v1",
                "-z",
                "--untracked-files=all",
            ],
            check=True,
            capture_output=True,
            text=True,
        ).stdout
    except (OSError, subprocess.CalledProcessError) as exc:
        raise ValueError("The release-owned worktree state could not be verified.") from exc
    dirty_paths: list[Path] = []
    records = status_output.split("\0")
    index = 0
    while index < len(records):
        record = records[index]
        index += 1
        if not record:
            continue
        status = record[:2]
        candidates = [record[3:]]
        if "R" in status or "C" in status:
            if index < len(records) and records[index]:
                candidates.append(records[index])
                index += 1
        for candidate in candidates:
            relative = Path(candidate)
            if not _release_excluded(relative):
                dirty_paths.append(relative)
    if dirty_paths:
        names = ", ".join(path.as_posix() for path in sorted(set(dirty_paths)))
        raise ValueError(f"Release-owned resources must be clean and committed: {names}")
    with tempfile.TemporaryDirectory(prefix="clinical-release-commit-") as temporary_dir:
        snapshot_root = Path(temporary_dir) / "snapshot"
        snapshot_root.mkdir()
        snapshot_root = snapshot_root.resolve()
        archive_path = Path(temporary_dir) / "commit.tar"
        try:
            with archive_path.open("wb") as archive_handle:
                subprocess.run(
                    ["git", "-C", str(repo_root), "archive", "--format=tar", commit],
                    check=True,
                    stdout=archive_handle,
                    stderr=subprocess.PIPE,
                )
            with tarfile.open(archive_path, "r") as archive:
                for member in archive.getmembers():
                    target = (snapshot_root / member.name).resolve()
                    try:
                        target.relative_to(snapshot_root)
                    except ValueError as exc:
                        raise ValueError("The committed release archive contains an escaping path.") from exc
                    if member.issym() or member.islnk():
                        raise ValueError("The committed release archive contains an unsupported link.")
                archive.extractall(snapshot_root)
        except (OSError, subprocess.CalledProcessError, tarfile.TarError, ValueError) as exc:
            raise ValueError("The committed release tree could not be materialized.") from exc
        return _package_release_tree(
            snapshot_root,
            output_path,
            git_commit=commit,
            certification_public_key=certification_public_key,
        )


def _manifest_integrity(skill_root: Path, *, allow_runtime_state: bool = True) -> list[dict[str, Any]]:
    if skill_root.is_symlink():
        return [{
            "category": "installation",
            "field": ".",
            "issue": "The installed skill root must not be a symbolic link.",
        }]

    def symlink_component(relative: Path) -> Path | None:
        current = skill_root
        for part in relative.parts:
            current = current / part
            if current.is_symlink():
                return current
        return None

    manifest_path = skill_root / RELEASE_MANIFEST
    if symlink_component(Path(RELEASE_MANIFEST)) is not None:
        return [{
            "category": "installation",
            "field": RELEASE_MANIFEST,
            "issue": "A manifest-owned path or ancestor is an unsupported symbolic link.",
        }]
    if not manifest_path.is_file():
        return [{"category": "installation", "field": RELEASE_MANIFEST, "issue": "Release manifest is missing."}]
    try:
        manifest = _read(manifest_path)
    except (OSError, ValueError, json.JSONDecodeError) as exc:
        return [{"category": "installation", "field": RELEASE_MANIFEST, "issue": str(exc)}]
    findings = []
    recorded_fingerprint, computed_fingerprint = _manifest_package_fingerprint(manifest)
    if not recorded_fingerprint or computed_fingerprint != recorded_fingerprint:
        findings.append({"category": "installation", "field": "package_fingerprint", "issue": "Package fingerprint does not match the release manifest."})
    declared: set[str] = set()
    for item in manifest.get("files", []):
        relative = str(item.get("path") or "")
        declared.add(relative)
        path = skill_root / relative
        if symlink_component(Path(relative)) is not None:
            findings.append({
                "category": "installation",
                "field": relative,
                "issue": "A manifest-owned path or ancestor is an unsupported symbolic link.",
            })
            continue
        try:
            path.resolve().relative_to(skill_root.resolve())
        except ValueError:
            findings.append({"category": "installation", "field": str(path), "issue": "Manifest path escapes the skill root."})
            continue
        if not path.is_file():
            findings.append({"category": "installation", "field": str(item.get("path")), "issue": "Packaged file is missing."})
        elif sha256_file(path) != item.get("sha256") or path.stat().st_size != int(item.get("bytes", -1)):
            findings.append({"category": "installation", "field": str(item.get("path")), "issue": "Packaged file hash does not match the release manifest."})
    permitted_state = {RELEASE_CERTIFICATION, RELEASE_MANIFEST}
    if allow_runtime_state:
        permitted_state.update({INSTALLATION_ASSURANCE, PROMOTION_RECORD})
    actual = {
        path.relative_to(skill_root).as_posix()
        for path in skill_root.rglob("*")
        if path.is_file() and not (
            allow_runtime_state
            and (
                "runtime" in path.relative_to(skill_root).parts
                or "__pycache__" in path.relative_to(skill_root).parts
                or path.suffix == ".pyc"
            )
        )
    }
    extras = sorted(actual - declared - permitted_state)
    if extras:
        findings.append({"category": "installation", "field": "inventory", "issue": f"Unlisted packaged files are present: {', '.join(extras)}"})
    return findings


def _is_sha256(value: Any) -> bool:
    text = str(value or "")
    return len(text) == 64 and all(character in "0123456789abcdef" for character in text.casefold())


def _utc_timestamp(value: Any) -> datetime | None:
    try:
        parsed = datetime.fromisoformat(str(value).replace("Z", "+00:00"))
    except ValueError:
        return None
    return parsed.astimezone(timezone.utc) if parsed.tzinfo is not None else None



def _certification_attestation(
    skill_root: Path,
    *,
    trusted_key_id: str = RELEASE_CERTIFICATION_TRUSTED_KEY_ID,
) -> tuple[dict[str, Any] | None, list[dict[str, Any]]]:
    """Verify that the embedded full-corpus report certifies this exact candidate."""
    report_path = skill_root / RELEASE_CERTIFICATION
    try:
        if report_path.stat().st_size > CERTIFICATION_REPORT_MAX_BYTES:
            raise ValueError("Release Certification report exceeds the governed encoded byte limit")
        report = _read(report_path)
        manifest_path = skill_root / RELEASE_MANIFEST
        manifest_bytes = manifest_path.read_bytes()
        manifest = json.loads(manifest_bytes)
    except (OSError, ValueError, json.JSONDecodeError) as exc:
        return None, [{"category": "installation", "field": RELEASE_CERTIFICATION, "issue": f"A valid embedded Release Certification report is required: {exc}"}]
    identity = report.get("release_identity") or {}
    evidence_bundle = report.get("evidence_bundle") or {}
    if evidence_bundle.get("schema_version") != "release-certification-evidence/v1":
        return None, [{
            "category": "installation",
            "field": RELEASE_CERTIFICATION,
            "issue": (
                "The embedded Release Certification evidence bundle is missing or invalid; "
                "digest-shaped summary fields cannot authorize this release."
            ),
        }]
    evidence_findings = release_certification_attestation_findings(
        report, manifest, skill_root, trusted_key_id=trusted_key_id,
    )
    evidence_findings.extend(_certification_evidence_findings(report, manifest, manifest_bytes))
    if evidence_findings:
        return None, [{
            "category": "installation",
            "field": RELEASE_CERTIFICATION,
            "issue": (
                "Certification evidence bundle failed independent verification: "
                + "; ".join(evidence_findings)
            ),
        }]
    cases = report.get("cases") or []
    case_order = report.get("case_order") or []
    configurations = report.get("hermes_configurations") or {}
    expected_configuration_hashes = (manifest.get("inventory") or {}).get("certification_configuration_sha256") or {}
    layout_evidence = report.get("layout_preservation_evidence") or {}
    layout_started = _utc_timestamp(layout_evidence.get("started_at"))
    layout_completed = _utc_timestamp(layout_evidence.get("completed_at"))
    report_completed = _utc_timestamp(report.get("completed_at"))
    valid = (
        report.get("schema_version") == "release-certification-corpus/v1"
        and report.get("status") == "passed"
        and report.get("certification_scope") == "complete_three_case_corpus"
        and not report.get("findings")
        and identity.get("package_fingerprint") == manifest.get("package_fingerprint")
        and identity.get("git_commit") == manifest.get("git_commit")
        and _is_sha256(report.get("preflight_evidence_sha256"))
        and layout_evidence.get("status") == "passed"
        and layout_evidence.get("returncode") == 0
        and _is_sha256(layout_evidence.get("sha256"))
        and set(layout_evidence.get("coverage") or []) == {
            "Prospective/Advarra", "Prospective/Sterling", "Ambispective/Advarra",
            "Ambispective/Sterling", "Retrospective/Protocol",
        }
        and layout_started is not None
        and layout_completed is not None
        and report_completed is not None
        and layout_started <= layout_completed <= report_completed
        and set(configurations) == set(CERTIFICATION_CASE_ORDER)
        and expected_configuration_hashes == {
            fixture_id: sha256_value(configurations.get(fixture_id) or {})
            for fixture_id in CERTIFICATION_CASE_ORDER
        }
        and all(
            all(configuration.get(key) == expected for key, expected in CERTIFIED_HERMES_CONFIGURATION.items())
            and bool(configuration.get("layout_preservation_notes"))
            for configuration in configurations.values()
        )
        and tuple(case_order) == CERTIFICATION_CASE_ORDER
        and tuple(case.get("fixture_id") for case in cases) == CERTIFICATION_CASE_ORDER
    )
    for case in cases:
        fixture_id = str(case.get("fixture_id") or "")
        outputs = case.get("output_evidence") or []
        expected_outputs = (
            {"output/protocol.docx"}
            if fixture_id == "retrospective"
            else {"output/protocol.docx", "output/icf.docx", "output/study.xml"}
        )
        output_by_path = {str(item.get("path") or ""): item for item in outputs}
        visual = case.get("visual_qa") or {}
        expected_visual = {Path(path).stem for path in expected_outputs if path.endswith(".docx")}
        try:
            elapsed = float(case.get("elapsed_seconds"))
            desktop_elapsed = float(case.get("desktop_operation_elapsed_seconds"))
        except (TypeError, ValueError):
            elapsed = desktop_elapsed = -1.0
        runtime_valid = (
            0.0 < elapsed < 900.0
            if fixture_id == "retrospective"
            else 0.0 < elapsed <= 1080.0
        )
        output_valid = (
            set(output_by_path) == expected_outputs
            and all(
                item.get("confirmed") is True
                and int(item.get("bytes") or 0) > 0
                and _is_sha256(item.get("sha256"))
                for item in output_by_path.values()
            )
        )
        visual_valid = (
            set(visual) == expected_visual
            and all(
                item.get("status") == "passed"
                and int(item.get("page_count") or 0) > 0
                and int(item.get("page_count") or 0) == len(item.get("page_sha256") or [])
                and all(_is_sha256(digest) for digest in item.get("page_sha256") or [])
                and set(item.get("checks") or []) == CERTIFICATION_VISUAL_CHECKS
                and _is_sha256(item.get("request_sha256"))
                and _is_sha256(item.get("response_sha256"))
                and _is_sha256(item.get("pdf_sha256"))
                and _is_sha256(item.get("docx_sha256"))
                and item.get("producer_model_id") == CERTIFIED_HERMES_CONFIGURATION["model_identifier"]
                and item.get("docx_sha256") == output_by_path.get(f"output/{artifact}.docx", {}).get("sha256")
                for artifact, item in visual.items()
            )
        )
        render_assurance_evidence = case.get("render_assurance") or {}
        page_renderer_evidence = render_assurance_evidence.get("active_page_renderer") or {}
        expected_page_renderer = (manifest.get("inventory") or {}).get("pdf_page_renderer") or {}
        expected_selection = {
            "retrospective": ("Retrospective", None),
            "ambispective-sterling": ("Ambispective", "Sterling"),
            "prospective-advarra": ("Prospective", "Advarra"),
        }.get(fixture_id)
        expected_bundle = next((
            bundle for bundle in (manifest.get("inventory") or {}).get("contracted_template_bundles", [])
            if (
                (bundle.get("selection") or {}).get("study_type"),
                (bundle.get("selection") or {}).get("icf_family"),
            ) == expected_selection
        ), {})
        expected_gate_statuses = {gate: "passed" for gate in CERTIFICATION_GATES}
        if fixture_id == "retrospective":
            expected_gate_statuses["prs_xml"] = "not_applicable"
        valid = valid and all((
            case.get("status") == "passed",
            not case.get("findings"),
            case.get("release_identity") == identity,
            case.get("hermes_configuration_sha256") == expected_configuration_hashes.get(fixture_id),
            set(case.get("model_identifiers") or []) == {CERTIFIED_HERMES_CONFIGURATION["model_identifier"]},
            _is_sha256(case.get("report_sha256")),
            runtime_valid,
            desktop_elapsed > 0.0,
            case.get("within_approved_runtime") is True,
            output_valid,
            case.get("gate_statuses") == expected_gate_statuses,
            case.get("layout_checks") == {"natural_section_3_flow": "passed", "no_orphan_headings": "passed"},
            visual_valid,
            (render_assurance_evidence.get("active_renderer") or {}).get("kind") in {"Microsoft Word", "LibreOffice"},
            all(page_renderer_evidence.get(key) == expected_page_renderer.get(key) for key in ("kind", "version", "wheel", "wheel_sha256")),
            page_renderer_evidence.get("source") == "release-owned runtime",
            case.get("contracted_template_bundle_identity") == expected_bundle.get("identity_sha256"),
            case.get("layout_preservation_baseline_identity") == (expected_bundle.get("layout_preservation_baseline") or {}).get("sha256"),
        ))
    if not valid:
        return None, [{"category": "installation", "field": RELEASE_CERTIFICATION, "issue": "The embedded Release Certification report does not pass and bind this exact commit, fingerprint, corpus, model, and configuration."}]
    return report, []


def _validated_archive_members(archive: zipfile.ZipFile, extraction_root: Path) -> list[zipfile.ZipInfo]:
    """Reject every ambiguous archive member before extraction."""
    root = extraction_root.resolve()
    targets: set[str] = set()
    members = archive.infolist()
    if len(members) > RELEASE_ARCHIVE_MAX_MEMBERS:
        raise ValueError("Release archive contains too many members.")
    total_uncompressed = 0
    for info in members:
        raw = info.filename[:-1] if info.is_dir() and info.filename.endswith("/") else info.filename
        relative_path = PurePosixPath(raw)
        if (
            not raw
            or "\\" in raw
            or any(part in {"", ".", ".."} for part in raw.split("/"))
            or relative_path.is_absolute()
            or relative_path.as_posix() != raw
        ):
            raise ValueError(f"Release archive contains a noncanonical member path: {info.filename}")
        target = (root / Path(*relative_path.parts)).resolve()
        try:
            relative = target.relative_to(root)
        except ValueError as exc:
            raise ValueError(f"Release archive path escapes the staging root: {info.filename}") from exc
        if not relative.parts or relative.parts[0] != "clinical-document-generation":
            raise ValueError(f"Release archive member is outside its one package root: {info.filename}")
        key = str(target).casefold()
        if key in targets:
            raise ValueError(f"Release archive contains duplicate normalized target: {info.filename}")
        targets.add(key)
        if (info.external_attr >> 16) & 0o170000 == 0o120000:
            raise ValueError(f"Release archive contains an unsupported symbolic link: {info.filename}")
        member_limit = (
            CERTIFICATION_REPORT_MAX_BYTES
            if relative.name == RELEASE_CERTIFICATION
            else RELEASE_ARCHIVE_MAX_MEMBER_BYTES
        )
        if info.file_size < 0 or info.file_size > member_limit:
            raise ValueError(f"Release archive member exceeds the governed byte limit: {info.filename}")
        total_uncompressed += info.file_size
        if total_uncompressed > RELEASE_ARCHIVE_MAX_TOTAL_BYTES:
            raise ValueError("Release archive total uncompressed size exceeds the governed byte limit.")
        if info.file_size and (
            info.compress_size <= 0
            or info.file_size / info.compress_size > RELEASE_ARCHIVE_MAX_COMPRESSION_RATIO
        ):
            raise ValueError(f"Release archive member exceeds the governed compression ratio: {info.filename}")
    return members


def _sign_release_certification(
    report: Mapping[str, Any],
    private_key_path: Path,
) -> dict[str, Any]:
    """Sign one canonical certification report with an external RSA private key."""
    try:
        private_key = json.loads(private_key_path.expanduser().read_text(encoding="utf-8"))
        modulus = int(str(private_key.get("modulus") or ""), 16)
        private_exponent = int(str(private_key.get("private_exponent") or ""), 16)
        exponent = int(private_key.get("exponent"))
    except (OSError, TypeError, ValueError, json.JSONDecodeError) as exc:
        raise ValueError(f"Release Certification signing key is unavailable or invalid: {exc}") from exc
    key_identity = release_certification_key_id(private_key)
    if (
        private_key.get("schema_version") != "release-certification-public-key/v1"
        or private_key.get("algorithm") != RELEASE_CERTIFICATION_SIGNATURE_ALGORITHM
        or private_key.get("key_id") != key_identity
        or modulus.bit_length() < 2048
    ):
        raise ValueError("Release Certification signing key identity or strength is invalid.")
    signed = copy.deepcopy(dict(report))
    signed.pop("evidence_attestation", None)
    payload_digest = hashlib.sha256(release_certification_payload(signed)).digest()
    encoded_bytes = (modulus.bit_length() + 7) // 8
    digest_info = bytes.fromhex("3031300d060960864801650304020105000420") + payload_digest
    encoded = b"\x00\x01" + b"\xff" * (encoded_bytes - len(digest_info) - 3) + b"\x00" + digest_info
    signature = pow(int.from_bytes(encoded, "big"), private_exponent, modulus).to_bytes(encoded_bytes, "big")
    signed["evidence_attestation"] = {
        "schema_version": "release-certification-attestation/v1",
        "algorithm": RELEASE_CERTIFICATION_SIGNATURE_ALGORITHM,
        "key_id": key_identity,
        "payload_sha256": payload_digest.hex(),
        "signature_base64": base64.b64encode(signature).decode("ascii"),
    }
    return signed


def bind_release_certification(
    archive_path: Path,
    report_path: Path,
    *,
    private_key_path: Path | None = None,
    trusted_certification_key_id: str = RELEASE_CERTIFICATION_TRUSTED_KEY_ID,
) -> dict[str, Any]:
    """Attach the passing report to its already-certified immutable package."""
    archive_path = archive_path.expanduser().resolve()
    report_path = report_path.expanduser().resolve()
    if report_path.stat().st_size > CERTIFICATION_REPORT_MAX_BYTES:
        raise ValueError("Release Certification report exceeds the governed encoded byte limit.")
    report_bytes = report_path.read_bytes()
    try:
        report = json.loads(report_bytes)
    except json.JSONDecodeError as exc:
        raise ValueError("Release Certification report is not valid JSON.") from exc
    configured_key = private_key_path
    if configured_key is None and os.environ.get("CLINICAL_DOCUMENT_CERTIFICATION_PRIVATE_KEY"):
        configured_key = Path(os.environ["CLINICAL_DOCUMENT_CERTIFICATION_PRIVATE_KEY"])
    if not report.get("evidence_attestation") and configured_key is not None:
        report = _sign_release_certification(report, configured_key)
        report_bytes = (json.dumps(report, indent=2, ensure_ascii=False) + "\n").encode("utf-8")
        if len(report_bytes) > CERTIFICATION_REPORT_MAX_BYTES:
            raise ValueError("Release Certification report exceeds the governed encoded byte limit.")
    with tempfile.TemporaryDirectory(prefix="clinical-certification-bind-") as directory:
        extracted = Path(directory) / "extracted"
        with zipfile.ZipFile(archive_path) as source:
            _validated_archive_members(source, extracted)
            source.extractall(extracted)
        candidate = extracted / "clinical-document-generation"
        (candidate / RELEASE_CERTIFICATION).write_bytes(report_bytes)
        integrity = _manifest_integrity(candidate, allow_runtime_state=False)
        _, certification_findings = _certification_attestation(
            candidate, trusted_key_id=trusted_certification_key_id,
        )
        if integrity or certification_findings:
            issues = integrity + certification_findings
            raise ValueError("Release Certification cannot be bound: " + "; ".join(str(item["issue"]) for item in issues))
        temporary = archive_path.with_name(f".{archive_path.name}.certified")
        try:
            with zipfile.ZipFile(archive_path) as source, zipfile.ZipFile(temporary, "w", compression=zipfile.ZIP_DEFLATED, compresslevel=9) as target:
                for info in source.infolist():
                    if info.filename.endswith("/" + RELEASE_CERTIFICATION):
                        continue
                    target.writestr(info, source.read(info.filename))
                info = zipfile.ZipInfo(f"clinical-document-generation/{RELEASE_CERTIFICATION}", date_time=(1980, 1, 1, 0, 0, 0))
                info.compress_type = zipfile.ZIP_DEFLATED
                target.writestr(info, report_bytes)
            os.replace(temporary, archive_path)
        finally:
            temporary.unlink(missing_ok=True)
    return {
        "status": "passed",
        "package": archive_path.as_posix(),
        "package_fingerprint": report["release_identity"]["package_fingerprint"],
        "certification_sha256": hashlib.sha256(report_bytes).hexdigest(),
    }


def _validate_hermes_discovery(config_path: Path, active: Path) -> list[dict[str, Any]]:
    """Require Hermes discovery to select this promoted path and no editable copy."""
    try:
        lines = config_path.read_text(encoding="utf-8").splitlines()
    except OSError as exc:
        return [{"category": "installation", "field": "hermes_configuration", "issue": f"Hermes configuration cannot be read: {exc}"}]
    entries: list[str] = []
    scalars: dict[tuple[str, ...], Any] = {}
    stack: list[tuple[int, str]] = []
    defined_paths: set[tuple[str, ...]] = set()
    duplicate_paths: set[tuple[str, ...]] = set()
    in_external = False
    base_indent = 0
    external_definitions = 0
    for line in lines:
        stripped = line.strip()
        indent = len(line) - len(line.lstrip())
        if stripped and not stripped.startswith(("#", "- ")) and ":" in stripped:
            key, raw_value = stripped.split(":", 1)
            while stack and stack[-1][0] >= indent:
                stack.pop()
            path = tuple(item[1] for item in stack) + (key.strip(),)
            if path in defined_paths:
                duplicate_paths.add(path)
            defined_paths.add(path)
            raw_scalar = raw_value.strip()
            if raw_scalar:
                if raw_scalar[:1] in {"'", '"'} and raw_scalar[-1:] == raw_scalar[:1]:
                    value: Any = raw_scalar[1:-1]
                elif raw_scalar.casefold() in {"true", "false"}:
                    value = raw_scalar.casefold() == "true"
                elif re.fullmatch(r"-?\d+", raw_scalar):
                    value = int(raw_scalar)
                else:
                    value = raw_scalar
                scalars[path] = value
            else:
                stack.append((indent, key.strip()))
        if stripped == "external_dirs:":
            exact_external = tuple(item[1] for item in stack) == ("skills", "external_dirs")
            in_external = exact_external
            if exact_external:
                external_definitions += 1
                base_indent = indent
            continue
        if in_external and stripped and indent <= base_indent:
            in_external = False
        if in_external and stripped.startswith("- "):
            entries.append(stripped[2:].strip().strip("'\""))
    governed = {
        key: scalars.get(("skills", "clinical_document_generation", key))
        for key in CERTIFIED_HERMES_CONFIGURATION
    }
    governed_matches = governed == CERTIFIED_HERMES_CONFIGURATION
    try:
        host_turns_sufficient = int(scalars.get(("agent", "max_turns"), "0")) >= int(CERTIFIED_HERMES_CONFIGURATION["max_turns"])
    except ValueError:
        host_turns_sufficient = False
    host_matches = all((
        scalars.get(("model", "default")) == CERTIFIED_HERMES_CONFIGURATION["model_identifier"],
        scalars.get(("agent", "reasoning_effort")) == "medium",
        host_turns_sufficient,
    ))
    normalized_entries = [str(Path(entry).expanduser().resolve()) for entry in entries]
    if duplicate_paths or external_definitions != 1 or normalized_entries != [str(active.resolve())] or not governed_matches or not host_matches:
        return [{"category": "installation", "field": "hermes_configuration", "issue": f"Hermes must select only {active} and match the certified model, medium reasoning, safe-mode, and 80-turn governed settings."}]
    return []


def verify_installation(skill_root: Path, *, deadline_seconds: float = 120.0) -> dict[str, Any]:
    """Prove that an extracted release owns a complete local assurance path."""
    skill_root = skill_root.resolve()
    findings = _manifest_integrity(skill_root)
    fallback_fonts = skill_root / "assets/fallback-fonts"
    if not any(fallback_fonts.glob("*.ttf")):
        findings.append({"category": "installation", "field": "fallback_fonts", "issue": "No packaged compatible fonts are present."})
    reference = {"meta": {"study_type": "Prospective", "icf_template": "Advarra"}}
    office_renderers = [
        item for item in renderers(skill_root=skill_root)
        if item.get("kind") in {"Microsoft Word", "LibreOffice"}
    ]
    page_renderer_identities = [
        item for item in page_renderers(
            skill_root=skill_root, require_promoted_runtime=False
        )
        if item.get("kind") == "pypdfium2" and item.get("source") == "release-owned runtime"
    ]
    with tempfile.TemporaryDirectory(prefix="clinical-installation-assurance-") as directory:
        revision_dir = Path(directory)
        candidate_dir = revision_dir / "candidate"
        candidate_dir.mkdir()
        document = Document()
        for font in BUNDLED_FONT_FILES:
            run = document.add_paragraph().add_run(f"{font}: Clinical document installation smoke")
            run.font.name = font
        document.save(candidate_dir / "installation-smoke.docx")
        smoke_path = candidate_dir / "installation-smoke.docx"
        assurance = render_assurance(
            skill_root,
            revision_dir,
            reference,
            structural_validation={
                "status": "structurally_valid",
                "expected_files": [smoke_path.name],
                "candidate_files": [{
                    "path": smoke_path.relative_to(revision_dir).as_posix(),
                    "sha256": sha256_file(smoke_path),
                    "bytes": smoke_path.stat().st_size,
                }],
            },
            deadline_seconds=deadline_seconds,
            renderer_identities=office_renderers,
            page_renderer_identities=page_renderer_identities,
            rebuild_candidate=lambda _substitutions: {"status": "passed"},
            require_promoted_runtime=False,
        )
    if assurance.get("status") != "passed":
        findings.extend(assurance.get("findings", []))
    render_evidence = assurance.get("render", {})
    candidates = [attempt.get("adapter") for attempt in render_evidence.get("renderer_attempts", []) if attempt.get("adapter")]
    if not office_renderers:
        findings.append({"category": "installation", "field": "office_renderer", "issue": "Microsoft Word or LibreOffice is required on the host."})
    if not page_renderer_identities:
        findings.append({"category": "installation", "field": "pdf_page_renderer", "issue": "The release-owned pypdfium2 page renderer was not discovered."})
    return {
        "status": "passed" if not findings else "blocked",
        "renderer": render_evidence.get("renderer"),
        "renderer_candidates": candidates,
        "page_renderer": render_evidence.get("page_renderer"),
        "fonts": assurance.get("fonts", {}),
        "smoke": {
            "status": assurance.get("status"),
            "pages": sum(len(item.get("pages", [])) for item in render_evidence.get("artifacts", [])),
        },
        "render_assurance": assurance,
        "findings": findings,
    }


def _provision_page_renderer(skill_root: Path) -> dict[str, Any]:
    """Install the manifest-bound pypdfium2 wheel without host discovery."""
    skill_root = skill_root.resolve()
    runtime_root = skill_root / "runtime"
    runtime_python = runtime_root / "python"
    try:
        manifest = _read(skill_root / RELEASE_MANIFEST)
        identity = dict(manifest["inventory"]["pdf_page_renderer"])
    except (OSError, KeyError, TypeError, ValueError, json.JSONDecodeError):
        return {"status": "blocked", "findings": [{
            "category": "installation",
            "field": "pdf_page_renderer",
            "issue": "The release manifest does not identify its one packaged pypdfium2 renderer.",
        }]}
    if identity.get("kind") != "pypdfium2":
        return {"status": "blocked", "findings": [{
            "category": "installation",
            "field": "pdf_page_renderer",
            "issue": "The release manifest must identify pypdfium2 as its only PDF page renderer.",
        }]}
    existing = page_renderers(
        skill_root=skill_root, require_promoted_runtime=False
    )
    if existing:
        return {"status": "passed", "page_renderer": existing[0], "provisioned": False}
    if (
        runtime_root.is_symlink()
        or runtime_python.exists()
        or runtime_python.is_symlink()
        or (runtime_root / "PDF-RENDERER.json").exists()
    ):
        finding = dict(_pdfium_runtime_integrity(
            skill_root, require_promoted_runtime=False
        )["finding"])
        finding["category"] = "installation"
        finding["field"] = "pdf_page_renderer"
        return {"status": "blocked", "findings": [finding]}
    wheel_relative = Path(str(identity.get("wheel") or ""))
    wheel = (skill_root / wheel_relative).resolve()
    try:
        wheel.relative_to(skill_root)
    except ValueError:
        wheel = Path()
    expected_hash = str(identity.get("wheel_sha256") or "")
    if not wheel.is_file() or not expected_hash or sha256_file(wheel) != expected_hash:
        return {"status": "blocked", "findings": [{
            "category": "installation",
            "field": "pdf_page_renderer",
            "issue": "The packaged pypdfium2 wheel is missing or does not match its release-manifest hash.",
        }]}
    try:
        expected_inventory = _pdfium_wheel_inventory(wheel)
    except (OSError, ValueError, zipfile.BadZipFile):
        expected_inventory = []
    if identity.get("runtime_inventory") != expected_inventory:
        return {"status": "blocked", "findings": [{
            "category": "installation",
            "field": "pdf_page_renderer",
            "code": "installation.pdfium_manifest_inventory_invalid",
            "issue": "The PDFium extraction inventory does not match the immutable packaged wheel.",
        }]}
    platform_tag = str(identity.get("platform") or "")
    host = f"{platform.system()} {platform.machine()}"
    compatible = (
        platform.system() == "Darwin"
        and "macosx" in platform_tag
        and platform.machine().casefold() in platform_tag.casefold()
    )
    if not compatible:
        return {"status": "blocked", "findings": [{
            "category": "installation",
            "field": "pdf_page_renderer",
            "issue": f"Packaged pypdfium2 targets {platform_tag}; this host is {host}.",
        }]}
    staged_python = runtime_root / ".pypdfium2-install"
    shutil.rmtree(staged_python, ignore_errors=True)
    staged_python.mkdir(parents=True)
    try:
        with zipfile.ZipFile(wheel) as archive:
            for member in archive.infolist():
                target = (staged_python / member.filename).resolve()
                try:
                    target.relative_to(staged_python.resolve())
                except ValueError as exc:
                    raise ValueError("The packaged pypdfium2 wheel contains an escaping path.") from exc
                if (member.external_attr >> 16) & 0o170000 == 0o120000:
                    raise ValueError("The packaged pypdfium2 wheel contains an unsupported symbolic link.")
            archive.extractall(staged_python)
        shutil.rmtree(runtime_python, ignore_errors=True)
        os.replace(staged_python, runtime_python)
        _write(runtime_root / "PDF-RENDERER.json", {
            "kind": "pypdfium2",
            "version": str(identity.get("version") or ""),
            "wheel": wheel_relative.as_posix(),
            "wheel_sha256": expected_hash,
            "platform": platform_tag,
            "status": "provisioned",
            "inventory_source": RELEASE_MANIFEST,
        })
    except (OSError, ValueError, zipfile.BadZipFile) as exc:
        shutil.rmtree(staged_python, ignore_errors=True)
        return {"status": "blocked", "findings": [{
            "category": "installation",
            "field": "pdf_page_renderer",
            "issue": f"The packaged pypdfium2 wheel could not be installed: {exc}",
        }]}
    installed = page_renderers(
        skill_root=skill_root, require_promoted_runtime=False
    )
    if len(installed) != 1:
        (skill_root / "runtime/PDF-RENDERER.json").unlink(missing_ok=True)
        shutil.rmtree(runtime_python, ignore_errors=True)
        return {"status": "blocked", "findings": [{
            "category": "installation",
            "field": "pdf_page_renderer",
            "issue": "The release-owned pypdfium2 runtime could not be discovered after installation.",
        }]}
    identity = installed[0]
    return {"status": "passed", "page_renderer": identity, "provisioned": True}


def provision_render_assurance(skill_root: Path) -> dict[str, Any]:
    """Provision PDFium offline and verify the required host office renderer."""
    skill_root = skill_root.resolve()
    manifest_findings = _manifest_integrity(skill_root, allow_runtime_state=True)
    if manifest_findings:
        return {
            "status": "blocked",
            "findings": manifest_findings,
            "renderer": None,
            "page_renderer": None,
        }
    office_renderers = [
        item for item in renderers(skill_root=skill_root)
        if item.get("kind") in {"Microsoft Word", "LibreOffice"}
    ]
    if not office_renderers:
        return {"status": "blocked", "findings": [{
            "category": "installation",
            "field": "office_renderer",
            "code": "installation.office_renderer_required",
            "issue": "Install or enable Microsoft Word or LibreOffice on the host, then rerun release installation.",
        }]}
    page_provision = _provision_page_renderer(skill_root)
    if page_provision.get("status") != "passed":
        return page_provision
    return {
        "status": "passed",
        "renderer": office_renderers[0],
        "page_renderer": page_provision["page_renderer"],
        "provisioned": {"renderer": False, "page_renderer": page_provision["provisioned"]},
    }


def _relocate_paths(value: Any, source_root: Path, destination_root: Path) -> Any:
    """Rewrite staged absolute paths to their post-activation location."""
    if isinstance(value, Mapping):
        return {str(key): _relocate_paths(item, source_root, destination_root) for key, item in value.items()}
    if isinstance(value, list):
        return [_relocate_paths(item, source_root, destination_root) for item in value]
    if isinstance(value, str):
        source = str(source_root)
        if value == source or value.startswith(source + os.sep):
            return str(destination_root) + value[len(source):]
    return value


def _installation_smoke_result(candidate: Path) -> dict[str, Any]:
    timeout_seconds = 180
    try:
        completed = subprocess.run(
            [sys.executable, str(candidate / "scripts/workflow.py"), "--verify-installation"],
            cwd=candidate,
            text=True,
            capture_output=True,
            timeout=timeout_seconds,
        )
    except subprocess.TimeoutExpired:
        return {
            "status": "blocked",
            "findings": [{
                "category": "installation",
                "field": "smoke",
                "code": "installation.smoke_timeout",
                "timeout_seconds": timeout_seconds,
                "issue": f"Installation smoke exceeded {timeout_seconds} seconds.",
            }],
        }
    try:
        assurance = json.loads(completed.stdout)
    except json.JSONDecodeError:
        assurance = {
            "status": "blocked",
            "findings": [{
                "category": "installation",
                "field": "smoke",
                "issue": completed.stderr or completed.stdout or "Installation smoke returned no JSON.",
            }],
        }
    if completed.returncode and assurance.get("status") == "passed":
        return {
            "status": "blocked",
            "findings": [{
                "category": "installation",
                "field": "smoke",
                "issue": completed.stderr or "Installation smoke process failed.",
            }],
        }
    return dict(assurance)


def _installation_pdfium_findings(candidate: Path) -> list[dict[str, Any]]:
    integrity = _pdfium_runtime_integrity(
        candidate, require_promoted_runtime=False
    )
    if integrity.get("status") == "passed":
        return []
    finding = dict(integrity.get("finding") or {})
    finding["category"] = "installation"
    finding["field"] = "pdf_page_renderer"
    return [finding]


def _write_atomic_installation_state(path: Path, value: Mapping[str, Any]) -> None:
    """Commit external installation state without exposing partial JSON."""
    path.parent.mkdir(parents=True, exist_ok=True)
    descriptor, temporary_name = tempfile.mkstemp(
        prefix=f".{path.name}.", suffix=".tmp", dir=path.parent
    )
    temporary = Path(temporary_name)
    try:
        with os.fdopen(descriptor, "w", encoding="utf-8") as stream:
            stream.write(json.dumps(value, indent=2, ensure_ascii=False) + "\n")
            stream.flush()
            os.fsync(stream.fileno())
        os.replace(temporary, path)
        _sync_directory(path.parent)
    finally:
        temporary.unlink(missing_ok=True)


def _sync_directory(path: Path) -> None:
    descriptor = os.open(path, os.O_RDONLY | getattr(os, "O_DIRECTORY", 0))
    try:
        os.fsync(descriptor)
    except OSError as exc:
        if exc.errno not in {22, 45}:
            raise
    finally:
        os.close(descriptor)


def _filesystem_identity(path: Path) -> tuple[int, int] | None:
    try:
        details = path.stat(follow_symlinks=False)
    except FileNotFoundError:
        return None
    return details.st_dev, details.st_ino


def _activation_journal_path(skills_dir: Path) -> Path:
    return skills_dir / ".clinical-document-generation.activation.json"


def _remove_activation_journal(skills_dir: Path) -> None:
    _activation_journal_path(skills_dir).unlink(missing_ok=True)
    _sync_directory(skills_dir)


def _recover_interrupted_activation(
    skills_dir: Path,
    *,
    trusted_certification_key_id: str,
) -> dict[str, Any] | None:
    journal_path = _activation_journal_path(skills_dir)
    if not journal_path.exists() and not journal_path.is_symlink():
        return None
    active = skills_dir / "clinical-document-generation"
    previous = skills_dir / ".clinical-document-generation.previous"
    try:
        if journal_path.is_symlink() or not journal_path.is_file():
            raise ValueError("activation journal is not a regular file")
        journal = _read(journal_path)
        if journal.get("schema_version") != "activation-transaction/v1":
            raise ValueError("unsupported activation journal schema")
        displaced_key = journal.get("displaced_key")
        if displaced_key is not None and (
            not isinstance(displaced_key, str)
            or _safe_release_identity_key(displaced_key) != displaced_key
        ):
            raise ValueError("invalid displaced release identity")
        displaced = (
            skills_dir / f".clinical-document-generation.displaced-{displaced_key}"
            if isinstance(displaced_key, str) and displaced_key
            else None
        )
        history_name = journal.get("history_name")
        expected_history_name = f"{displaced_key}.json" if displaced_key else None
        if history_name != expected_history_name:
            raise ValueError("activation journal history identity does not match")
        history = (
            skills_dir / "release-history" / history_name
            if isinstance(history_name, str)
            and history_name == Path(history_name).name
            and history_name.endswith(".json")
            else None
        )
        candidate_identity = journal.get("candidate_identity")
        if not (
            isinstance(candidate_identity, list)
            and len(candidate_identity) == 2
            and all(isinstance(part, int) for part in candidate_identity)
        ):
            raise ValueError("invalid candidate identity")
        candidate_committed = _filesystem_identity(active) == tuple(candidate_identity)
        if candidate_committed:
            _, active_findings = _installation_candidate_integrity(
                active,
                trusted_certification_key_id=trusted_certification_key_id,
            )
            if active_findings:
                return {
                    "status": "blocked",
                    "stage": "activation_recovery_integrity",
                    "findings": active_findings,
                    "active_release_retained": True,
                    "previous_release_retained": previous.is_dir(),
                }
            try:
                if displaced is not None and displaced.exists():
                    shutil.rmtree(displaced)
                    _sync_directory(skills_dir)
                _remove_activation_journal(skills_dir)
            except OSError:
                return {
                    "status": "passed",
                    "stage": "activated",
                    "activation_commit_point": "candidate_to_active_atomic_swap",
                    "active": str(active),
                    "previous": str(previous) if previous.exists() else None,
                    "cleanup": {
                        "status": "deferred",
                        "findings": [{
                            "category": "installation",
                            "field": "activation_recovery_cleanup",
                            "code": "installation.activation_recovery_cleanup_deferred",
                            "issue": "Activation was already committed; interrupted cleanup remains journaled for retry.",
                        }],
                    },
                }
            return None
        else:
            if bool(journal.get("active_had_release")) and not active.exists():
                if not previous.exists():
                    raise OSError("interrupted activation has no restorable active release")
                os.replace(previous, active)
                _sync_directory(skills_dir)
            if bool(journal.get("previous_had_release")) and not previous.exists():
                if displaced is None or not displaced.exists():
                    raise OSError("interrupted activation has no restorable previous release")
                os.replace(displaced, previous)
                _sync_directory(skills_dir)
            if bool(journal.get("history_created")) and history is not None:
                history.unlink(missing_ok=True)
                _sync_directory(history.parent)
        _remove_activation_journal(skills_dir)
        return None
    except (OSError, ValueError, json.JSONDecodeError):
        return {
            "status": "blocked",
            "stage": "activation_recovery",
            "findings": [{
                "category": "installation",
                "field": "activation_recovery",
                "code": "installation.activation_recovery_required",
                "issue": "An interrupted activation could not be reconciled; no new candidate was staged.",
            }],
            "active_release_retained": active.exists(),
            "previous_release_retained": previous.exists(),
        }


def _safe_release_identity_key(fingerprint: object) -> str:
    key = "".join(
        character for character in str(fingerprint)
        if character.isalnum() or character in "-_"
    )
    if not key:
        raise ValueError("Release package fingerprint has no safe filesystem identity.")
    return key


def _retain_lightweight_release_history(
    release: Path,
    skills_dir: Path,
) -> tuple[Path, bool]:
    identity = _active_release_identity(release)
    certification: dict[str, Any] = {}
    promotion_path = release / "PROMOTION-RECORD.json"
    if promotion_path.is_file():
        try:
            certification = dict(_read(promotion_path).get("certification") or {})
        except (OSError, TypeError, ValueError, json.JSONDecodeError):
            certification = {}
    fingerprint = str(identity["package_fingerprint"])
    history_key = _safe_release_identity_key(fingerprint)
    history_path = skills_dir / "release-history" / f"{history_key}.json"
    history_value = {
        "schema_version": "release-history/v1",
        "git_commit": identity.get("git_commit"),
        "package_fingerprint": fingerprint,
        "certification": certification,
    }
    created = not history_path.exists()
    if created:
        _write_atomic_installation_state(history_path, history_value)
    elif _read(history_path) != history_value:
        raise ValueError("Release history record collision has conflicting content.")
    return history_path, created


def _release_identity_key(release: Path) -> str:
    return _safe_release_identity_key(
        _active_release_identity(release)["package_fingerprint"]
    )


def _installation_candidate_integrity(
    candidate: Path,
    *,
    trusted_certification_key_id: str,
) -> tuple[dict[str, Any], list[dict[str, Any]]]:
    """Revalidate the complete candidate after one executable installer hook."""
    integrity = _manifest_integrity(candidate, allow_runtime_state=True)
    certification, certification_findings = _certification_attestation(
        candidate, trusted_key_id=trusted_certification_key_id,
    )
    runtime_findings = _installation_pdfium_findings(candidate)
    return certification or {}, integrity + certification_findings + runtime_findings


def _accepted_installation_python_runtime() -> dict[str, Any]:
    """Return the exact supported interpreter identity before installation staging."""
    runtime = _current_python_runtime()
    if _runtime_version(runtime) < (*MINIMUM_PYTHON_VERSION, 0):
        raise RuntimeError("unsupported Python runtime")
    implementation = str(runtime.get("implementation") or "")
    executable_value = str(runtime.get("executable") or "")
    if not implementation or not executable_value:
        raise RuntimeError("incomplete Python runtime identity")
    executable = Path(executable_value).expanduser().resolve(strict=True)
    if not executable.is_file():
        raise RuntimeError("Python executable identity is not a file")
    return {
        "implementation": implementation,
        "version": str(runtime.get("version") or ".".join(
            str(part) for part in _runtime_version(runtime)
        )),
        "version_info": list(_runtime_version(runtime)),
        "executable": {
            "path": str(executable),
            "sha256": sha256_file(executable),
        },
    }


def install_release(
    archive_path: Path,
    skills_dir: Path,
    *,
    hermes_config_path: Path,
    verifier: Callable[[Path], Mapping[str, Any]] | None = None,
    provisioner: Callable[[Path], Mapping[str, Any]] = provision_render_assurance,
    trusted_certification_key_id: str = RELEASE_CERTIFICATION_TRUSTED_KEY_ID,
) -> dict[str, Any]:
    """Smoke, then atomically activate an installable skill archive."""
    try:
        python_runtime = _accepted_installation_python_runtime()
    except (OSError, RuntimeError, TypeError, ValueError):
        return {
            "status": "blocked",
            "stage": "python_runtime",
            "findings": [{
                "category": "installation",
                "field": "python_runtime",
                "code": "installation.python_runtime_unsupported",
                "issue": "Installation requires Python 3.10 or newer before staging.",
            }],
        }
    archive_path = archive_path.expanduser().resolve()
    skills_dir = skills_dir.expanduser().resolve()
    skills_dir.mkdir(parents=True, exist_ok=True)
    recovery = _recover_interrupted_activation(
        skills_dir,
        trusted_certification_key_id=trusted_certification_key_id,
    )
    if recovery is not None:
        return recovery
    staging_root = Path(tempfile.mkdtemp(prefix=".clinical-document-generation.install-", dir=skills_dir))
    active = skills_dir / "clinical-document-generation"
    previous = skills_dir / ".clinical-document-generation.previous"
    candidate = staging_root / "clinical-document-generation"
    try:
        with zipfile.ZipFile(archive_path) as archive:
            _validated_archive_members(archive, staging_root)
            archive.extractall(staging_root)
        integrity = _manifest_integrity(candidate, allow_runtime_state=False)
        certification, certification_findings = _certification_attestation(
            candidate, trusted_key_id=trusted_certification_key_id,
        )
        discovery_findings = _validate_hermes_discovery(hermes_config_path.expanduser().resolve(), active)
        if integrity or certification_findings or discovery_findings:
            return {
                "status": "blocked",
                "stage": "promotion_eligibility",
                "findings": integrity + certification_findings + discovery_findings,
                "active_release_retained": active.is_dir(),
            }
        provision = dict(provisioner(candidate))
        if provision.get("status") != "passed":
            return {"status": "blocked", "stage": "provision", "findings": list(provision.get("findings", [])), "active_release_retained": active.is_dir()}
        certification, provision_findings = _installation_candidate_integrity(
            candidate,
            trusted_certification_key_id=trusted_certification_key_id,
        )
        if provision_findings:
            return {
                "status": "blocked",
                "stage": "provision_integrity",
                "findings": provision_findings,
                "active_release_retained": active.is_dir(),
            }
        assurance = (
            _installation_smoke_result(candidate)
            if verifier is None
            else dict(verifier(candidate))
        )
        if assurance.get("status") != "passed":
            return {"status": "blocked", "stage": "installation_smoke", "findings": list(assurance.get("findings", [])), "active_release_retained": active.is_dir()}
        certification, installation_findings = _installation_candidate_integrity(
            candidate,
            trusted_certification_key_id=trusted_certification_key_id,
        )
        if installation_findings:
            return {
                "status": "blocked",
                "stage": "installation_integrity",
                "findings": installation_findings,
                "active_release_retained": active.is_dir(),
            }
        recorded_provision = _relocate_paths(provision, candidate, active)
        recorded_assurance = _relocate_paths(assurance, candidate, active)
        _write(candidate / INSTALLATION_ASSURANCE, {
            "status": "passed",
            "verified_at": datetime.now(timezone.utc).isoformat(),
            "python_runtime": python_runtime,
            "provision": recorded_provision,
            "assurance": recorded_assurance,
        })
        assurance_sha256 = sha256_file(candidate / INSTALLATION_ASSURANCE)
        report_path = candidate / RELEASE_CERTIFICATION
        manifest = _read(candidate / RELEASE_MANIFEST)
        model_identifiers = sorted({
            str(model)
            for case in certification.get("cases", [])
            for model in case.get("model_identifiers", [])
        })
        configuration_hashes = sorted({
            str(case.get("hermes_configuration_sha256"))
            for case in certification.get("cases", [])
        })
        activated_at = datetime.now(timezone.utc).isoformat()
        _write(candidate / PROMOTION_RECORD, {
            "schema_version": "promoted-release/v1",
            "status": "active",
            "git_commit": manifest.get("git_commit"),
            "package_fingerprint": manifest.get("package_fingerprint"),
            "certification": {
                "status": "passed",
                "report_sha256": sha256_file(report_path),
                "preflight_evidence_sha256": certification.get("preflight_evidence_sha256"),
                "model_identifiers": model_identifiers,
                "hermes_configuration_sha256": configuration_hashes,
            },
            "runtime_assurance_sha256": assurance_sha256,
            "activated_at": activated_at,
            "hermes_discovery": str(active),
        })
        displaced_previous = None
        displaced_key = None
        retained_history = None
        retained_history_created = False
        try:
            if previous.exists():
                displaced_key = _release_identity_key(previous)
                displaced_previous = (
                    skills_dir
                    / f".clinical-document-generation.displaced-{displaced_key}"
                )
                if displaced_previous.exists():
                    return {
                        "status": "blocked",
                        "stage": "activation_precondition",
                        "findings": [{
                            "category": "installation",
                            "field": "displaced_release",
                            "code": "installation.displaced_release_exists",
                            "issue": "Deterministic displaced-release path already exists.",
                            "path": str(displaced_previous),
                        }],
                    }
            candidate_identity = _filesystem_identity(candidate)
            if candidate_identity is None:
                raise ValueError("candidate identity disappeared")
            history_name = f"{displaced_key}.json" if displaced_key else None
            journal = {
                "schema_version": "activation-transaction/v1",
                "phase": "prepared",
                "candidate_identity": list(candidate_identity),
                "active_had_release": active.exists(),
                "previous_had_release": previous.exists(),
                "displaced_key": displaced_key,
                "history_name": history_name,
                "history_created": bool(
                    history_name
                    and not (skills_dir / "release-history" / history_name).exists()
                ),
            }
            _write_atomic_installation_state(
                _activation_journal_path(skills_dir), journal
            )
        except (OSError, ValueError):
            return {
                "status": "blocked",
                "stage": "activation_journal",
                "findings": [{
                    "category": "installation",
                    "field": "activation_journal",
                    "code": "installation.activation_journal_failed",
                    "issue": "Activation journal could not be committed before namespace mutation.",
                }],
                "active_release_retained": active.is_dir(),
                "previous_release_retained": previous.is_dir(),
            }
        if previous.exists():
            try:
                assert displaced_previous is not None
                os.replace(previous, displaced_previous)
                _sync_directory(skills_dir)
                retained_history, retained_history_created = (
                    _retain_lightweight_release_history(displaced_previous, skills_dir)
                )

            except (OSError, ValueError):
                try:
                    if (
                        displaced_previous is not None
                        and displaced_previous.exists()
                        and not previous.exists()
                    ):
                        os.replace(displaced_previous, previous)
                        _sync_directory(skills_dir)
                    _remove_activation_journal(skills_dir)
                except OSError:
                    return {
                        "status": "blocked",
                        "stage": "activation_recovery",
                        "findings": [{
                            "category": "installation",
                            "field": "activation_recovery",
                            "code": "installation.activation_recovery_required",
                            "issue": "Prior release displacement failed and restoration must be retried.",
                        }],
                        "active_release_retained": active.is_dir(),
                        "previous_release_retained": previous.is_dir(),
                    }
                return {
                    "status": "blocked",
                    "stage": "activation_precondition",
                    "findings": [{
                        "category": "installation",
                        "field": "displaced_release",
                        "code": "installation.displacement_failed",
                        "issue": "Prior release displacement failed before commit; release state was restored.",
                    }],
                    "active_release_retained": active.is_dir(),
                    "previous_release_retained": previous.is_dir(),
                }
        try:
            if active.exists():
                os.replace(active, previous)
                _sync_directory(skills_dir)
            os.replace(candidate, active)
            _sync_directory(skills_dir)
        except OSError:
            if _filesystem_identity(active) == candidate_identity:
                return {
                    "status": "passed",
                    "stage": "activated",
                    "activation_commit_point": "candidate_to_active_atomic_swap",
                    "active": str(active),
                    "previous": str(previous) if previous.exists() else None,
                    "retained_history": (
                        str(retained_history) if retained_history is not None else None
                    ),
                    "cleanup": {
                        "status": "deferred",
                        "record": str(_activation_journal_path(skills_dir)),
                        "findings": [{
                            "category": "installation",
                            "field": "activation_commit_confirmation",
                            "code": "installation.activation_reconciliation_deferred",
                            "issue": "The candidate swap committed; transaction reconciliation is deferred.",
                        }],
                    },
                }
            try:
                if previous.exists() and not active.exists():
                    os.replace(previous, active)
                    _sync_directory(skills_dir)
                if (
                    displaced_previous is not None
                    and displaced_previous.exists()
                    and not previous.exists()
                ):
                    os.replace(displaced_previous, previous)
                    _sync_directory(skills_dir)
                if retained_history is not None and retained_history_created:
                    retained_history.unlink(missing_ok=True)
                    _sync_directory(retained_history.parent)
                _remove_activation_journal(skills_dir)
            except OSError:
                return {
                    "status": "blocked",
                    "stage": "activation_recovery",
                    "findings": [{
                        "category": "installation",
                        "field": "activation_recovery",
                        "code": "installation.activation_recovery_required",
                        "issue": "Activation failed before commit and restoration must be retried.",
                    }],
                    "active_release_retained": active.is_dir(),
                    "previous_release_retained": previous.is_dir(),
                }
            return {
                "status": "blocked",
                "stage": "activation_swap",
                "findings": [{
                    "category": "installation",
                    "field": "activation_swap",
                    "code": "installation.activation_swap_failed",
                    "issue": "Activation swap failed before commit; the prior release state was restored.",
                }],
                "active_release_retained": active.is_dir(),
                "previous_release_retained": previous.is_dir(),
            }
        cleanup: dict[str, Any] = {"status": "passed", "findings": []}
        retain_activation_journal = False
        if displaced_previous is not None:
            try:
                shutil.rmtree(displaced_previous)
                _sync_directory(skills_dir)
            except OSError:
                finding = {
                    "category": "installation",
                    "field": "post_commit_cleanup",
                    "code": "installation.cleanup_deferred",
                    "issue": "Activation succeeded; displaced release cleanup is deferred.",
                    "path": str(displaced_previous),
                }
                deferred_record = (
                    skills_dir / "release-history"
                    / f"deferred-cleanup-{displaced_key}.json"
                )
                cleanup_findings = [finding]
                recorded_path: str | None = str(deferred_record)
                try:
                    _write_atomic_installation_state(deferred_record, {
                        "schema_version": "deferred-installation-cleanup/v1",
                        "status": "deferred",
                        "finding": finding,
                    })
                except OSError:
                    recorded_path = None
                    retain_activation_journal = True
                    cleanup_findings.append({
                        "category": "installation",
                        "field": "post_commit_cleanup_record",
                        "code": "installation.cleanup_record_failed",
                        "issue": "Activation succeeded; deferred-cleanup evidence could not be persisted.",
                        "path": str(deferred_record),
                    })
                    cleanup_findings.append({
                        "category": "installation",
                        "field": "activation_journal",
                        "code": "installation.activation_journal_retained",
                        "issue": "The activation journal remains as durable deferred-cleanup evidence.",
                        "path": str(_activation_journal_path(skills_dir)),
                    })
                cleanup = {
                    "status": "deferred",
                    "record": recorded_path,
                    "findings": cleanup_findings,
                }
        try:
            if not retain_activation_journal:
                _remove_activation_journal(skills_dir)
        except OSError:
            journal_finding = {
                "category": "installation",
                "field": "activation_journal_cleanup",
                "code": "installation.activation_journal_cleanup_deferred",
                "issue": "Activation succeeded; transaction-journal cleanup is deferred.",
                "path": str(_activation_journal_path(skills_dir)),
            }
            cleanup = {
                "status": "deferred",
                "record": cleanup.get("record"),
                "findings": [*cleanup.get("findings", []), journal_finding],
            }
        return {
            "status": "passed",
            "stage": "activated",
            "activation_commit_point": "candidate_to_active_atomic_swap",
            "active": str(active),
            "previous": str(previous) if previous.exists() else None,
            "retained_history": str(retained_history) if retained_history else None,
            "assurance": recorded_assurance,
            "promotion_record": str(active / PROMOTION_RECORD),
            "cleanup": cleanup,
        }
    finally:
        shutil.rmtree(staging_root, ignore_errors=True)


def _rollback_journal_path(skills_dir: Path) -> Path:
    return skills_dir / ".clinical-document-generation.rollback.json"


def _remove_rollback_journal(skills_dir: Path) -> None:
    _rollback_journal_path(skills_dir).unlink(missing_ok=True)
    _sync_directory(skills_dir)


def _disposable_rollback_smoke(
    release: Path,
    *,
    verifier: Callable[[Path], Mapping[str, Any]] | None,
    trusted_certification_key_id: str,
) -> tuple[dict[str, Any], list[dict[str, Any]], list[dict[str, Any]]]:
    probe_root = Path(tempfile.mkdtemp(prefix="clinical-document-rollback-smoke-"))
    probe = probe_root / "release"
    try:
        shutil.copytree(release, probe, symlinks=True)
        _, before_findings = _installation_candidate_integrity(
            probe,
            trusted_certification_key_id=trusted_certification_key_id,
        )
        if before_findings:
            return {}, before_findings, []
        assurance = (
            _installation_smoke_result(probe)
            if verifier is None
            else dict(verifier(probe))
        )
        _, after_findings = _installation_candidate_integrity(
            probe,
            trusted_certification_key_id=trusted_certification_key_id,
        )
        _, retained_release_findings = _installation_candidate_integrity(
            release,
            trusted_certification_key_id=trusted_certification_key_id,
        )
        return assurance, [], [*after_findings, *retained_release_findings]
    except OSError:
        return {}, [{
            "category": "installation",
            "field": "rollback_smoke_copy",
            "code": "installation.rollback_smoke_copy_failed",
            "issue": "A disposable rollback smoke copy could not be created.",
        }], []
    finally:
        shutil.rmtree(probe_root, ignore_errors=True)


def _recover_interrupted_rollback(
    skills_dir: Path,
    *,
    verifier: Callable[[Path], Mapping[str, Any]] | None,
    trusted_certification_key_id: str,
) -> dict[str, Any] | None:
    journal_path = _rollback_journal_path(skills_dir)
    if not journal_path.exists() and not journal_path.is_symlink():
        return None
    active = skills_dir / "clinical-document-generation"
    previous = skills_dir / ".clinical-document-generation.previous"
    try:
        if journal_path.is_symlink() or not journal_path.is_file():
            raise ValueError("rollback journal is not a regular file")
        journal = _read(journal_path)
        if journal.get("schema_version") != "rollback-transaction/v1":
            raise ValueError("unsupported rollback journal schema")
        quarantine_key = journal.get("quarantine_key")
        if (
            not isinstance(quarantine_key, str)
            or not quarantine_key
            or _safe_release_identity_key(quarantine_key) != quarantine_key
        ):
            raise ValueError("invalid rollback quarantine identity")
        quarantine = (
            skills_dir / f".clinical-document-generation.quarantine-{quarantine_key}"
        )
        previous_identity = journal.get("previous_identity")
        if not (
            isinstance(previous_identity, list)
            and len(previous_identity) == 2
            and all(isinstance(part, int) for part in previous_identity)
        ):
            raise ValueError("invalid rollback previous identity")
        if _filesystem_identity(active) == tuple(previous_identity):
            _, recovery_findings = _installation_candidate_integrity(
                active,
                trusted_certification_key_id=trusted_certification_key_id,
            )
            (
                recovery_assurance,
                smoke_copy_findings,
                post_recovery_findings,
            ) = _disposable_rollback_smoke(
                active,
                verifier=verifier,
                trusted_certification_key_id=trusted_certification_key_id,
            )
            if (
                recovery_findings
                or smoke_copy_findings
                or recovery_assurance.get("status") != "passed"
                or post_recovery_findings
            ):
                return {
                    "status": "blocked",
                    "stage": "rollback_recovery_integrity",
                    "findings": [
                        *recovery_findings,
                        *smoke_copy_findings,
                        *list(recovery_assurance.get("findings", [])),
                        *post_recovery_findings,
                    ],
                    "active_release_retained": True,
                    "previous_release_retained": previous.is_dir(),
                }
            cleanup_findings = []
            try:
                _remove_rollback_journal(skills_dir)
            except OSError:
                cleanup_findings.append({
                    "category": "installation",
                    "field": "rollback_recovery_cleanup",
                    "code": "installation.rollback_recovery_cleanup_deferred",
                    "issue": "Rollback was already committed; interrupted journal cleanup remains deferred.",
                })
            result = {
                "status": "passed",
                "stage": "rolled_back",
                "rollback_commit_point": "previous_to_active_atomic_swap",
                "active": str(active),
                "quarantined": str(quarantine),
                "assurance": {
                    **recovery_assurance,
                    "recovered_after_interruption": True,
                },
                "historical_run_revisions_rewritten": False,
            }
            if cleanup_findings:
                result["cleanup"] = {
                    "status": "deferred",
                    "findings": cleanup_findings,
                }
            return result
        if not active.exists() and quarantine.exists():
            os.replace(quarantine, active)
            _sync_directory(skills_dir)
        if not active.is_dir() or not previous.is_dir():
            raise OSError("interrupted rollback state is not restorable")
        _remove_rollback_journal(skills_dir)
        return None
    except (OSError, ValueError, json.JSONDecodeError):
        return {
            "status": "blocked",
            "stage": "rollback_recovery",
            "findings": [{
                "category": "installation",
                "field": "rollback_recovery",
                "code": "installation.rollback_recovery_required",
                "issue": "An interrupted rollback could not be reconciled; retry recovery before another swap.",
            }],
            "active_release_retained": active.is_dir(),
            "previous_release_retained": previous.is_dir(),
        }


def rollback_release(
    skills_dir: Path,
    *,
    verifier: Callable[[Path], Mapping[str, Any]] | None = None,
    trusted_certification_key_id: str = RELEASE_CERTIFICATION_TRUSTED_KEY_ID,
) -> dict[str, Any]:
    """Verify and atomically restore the previous release, quarantining active."""
    skills_dir = skills_dir.expanduser().resolve()
    active = skills_dir / "clinical-document-generation"
    previous = skills_dir / ".clinical-document-generation.previous"
    recovery = _recover_interrupted_rollback(
        skills_dir,
        verifier=verifier,
        trusted_certification_key_id=trusted_certification_key_id,
    )
    if recovery is not None:
        return recovery
    if not active.is_dir() or not previous.is_dir():
        return {
            "status": "blocked",
            "stage": "rollback_precondition",
            "findings": [{
                "category": "installation",
                "field": "rollback",
                "issue": "Rollback requires complete active and immediately previous releases.",
            }],
            "active_release_retained": active.is_dir(),
            "previous_release_retained": previous.is_dir(),
        }
    _, integrity_findings = _installation_candidate_integrity(
        previous,
        trusted_certification_key_id=trusted_certification_key_id,
    )
    if integrity_findings:
        return {
            "status": "blocked",
            "stage": "rollback_integrity",
            "findings": integrity_findings,
            "active_release_retained": True,
            "previous_release_retained": True,
        }
    assurance, smoke_copy_findings, post_smoke_findings = _disposable_rollback_smoke(
        previous,
        verifier=verifier,
        trusted_certification_key_id=trusted_certification_key_id,
    )
    if smoke_copy_findings:
        return {
            "status": "blocked",
            "stage": "rollback_smoke_copy_integrity",
            "findings": smoke_copy_findings,
            "active_release_retained": True,
            "previous_release_retained": True,
        }
    if assurance.get("status") != "passed":
        return {
            "status": "blocked",
            "stage": "rollback_smoke",
            "findings": list(assurance.get("findings", [])),
            "active_release_retained": True,
            "previous_release_retained": True,
        }
    if post_smoke_findings:
        return {
            "status": "blocked",
            "stage": "rollback_integrity_after_smoke",
            "findings": post_smoke_findings,
            "active_release_retained": True,
            "previous_release_retained": True,
        }
    try:
        manifest = _read(active / RELEASE_MANIFEST)
        fingerprint = str(manifest.get("package_fingerprint") or "")
    except (OSError, ValueError, json.JSONDecodeError):
        fingerprint = ""
    if not fingerprint:
        return {
            "status": "blocked",
            "stage": "rollback_identity",
            "findings": [{
                "category": "installation",
                "field": "package_fingerprint",
                "issue": "The active release has no package fingerprint and cannot be quarantined safely.",
            }],
            "active_release_retained": True,
            "previous_release_retained": True,
        }
    try:
        quarantine_key = _safe_release_identity_key(fingerprint)
    except ValueError:
        return {
            "status": "blocked",
            "stage": "rollback_identity",
            "findings": [{
                "category": "installation",
                "field": "package_fingerprint",
                "issue": "The active release fingerprint has no safe quarantine identity.",
            }],
            "active_release_retained": True,
            "previous_release_retained": True,
        }
    quarantine = skills_dir / f".clinical-document-generation.quarantine-{quarantine_key}"
    if quarantine.exists():
        return {
            "status": "blocked",
            "stage": "rollback_quarantine",
            "findings": [{
                "category": "installation",
                "field": "quarantine",
                "issue": f"Rollback quarantine already exists: {quarantine}",
            }],
            "active_release_retained": True,
            "previous_release_retained": True,
        }
    try:
        active_identity = _filesystem_identity(active)
        previous_identity = _filesystem_identity(previous)
        if active_identity is None or previous_identity is None:
            raise OSError("rollback release identity disappeared")
        _write_atomic_installation_state(_rollback_journal_path(skills_dir), {
            "schema_version": "rollback-transaction/v1",
            "active_identity": list(active_identity),
            "previous_identity": list(previous_identity),
            "quarantine_key": quarantine_key,
        })
    except OSError:
        return {
            "status": "blocked",
            "stage": "rollback_journal",
            "findings": [{
                "category": "installation",
                "field": "rollback_journal",
                "code": "installation.rollback_journal_failed",
                "issue": "Rollback journal could not be committed before namespace mutation.",
            }],
            "active_release_retained": active.is_dir(),
            "previous_release_retained": previous.is_dir(),
        }
    try:
        os.replace(active, quarantine)
        _sync_directory(skills_dir)
        os.replace(previous, active)
        _sync_directory(skills_dir)
    except OSError:
        if _filesystem_identity(active) == previous_identity:
            cleanup_findings = []
            try:
                _remove_rollback_journal(skills_dir)
            except OSError:
                cleanup_findings.append({
                    "category": "installation",
                    "field": "rollback_journal",
                    "code": "installation.rollback_journal_cleanup_failed",
                    "issue": "Rollback committed; journal cleanup is deferred.",
                })
            result = {
                "status": "passed",
                "stage": "rolled_back",
                "rollback_commit_point": "previous_to_active_atomic_swap",
                "active": str(active),
                "quarantined": str(quarantine),
                "assurance": assurance,
                "historical_run_revisions_rewritten": False,
            }
            if cleanup_findings:
                result["cleanup"] = {
                    "status": "deferred",
                    "findings": cleanup_findings,
                }
            return result
        try:
            if quarantine.exists() and not active.exists():
                os.replace(quarantine, active)
                _sync_directory(skills_dir)
            _remove_rollback_journal(skills_dir)
        except OSError:
            return {
                "status": "blocked",
                "stage": "rollback_recovery",
                "findings": [{
                    "category": "installation",
                    "field": "rollback_recovery",
                    "code": "installation.rollback_recovery_required",
                    "issue": "Rollback swap failed and restoration must be retried.",
                }],
                "active_release_retained": active.is_dir(),
                "previous_release_retained": previous.is_dir(),
            }
        return {
            "status": "blocked",
            "stage": "rollback_swap",
            "findings": [{
                "category": "installation",
                "field": "rollback_swap",
                "code": "installation.rollback_swap_failed",
                "issue": "Rollback swap failed; the prior active release was restored.",
            }],
            "active_release_retained": active.is_dir(),
            "previous_release_retained": previous.is_dir(),
        }
    cleanup_findings = []
    try:
        _remove_rollback_journal(skills_dir)
    except OSError:
        cleanup_findings.append({
            "category": "installation",
            "field": "rollback_journal",
            "code": "installation.rollback_journal_cleanup_failed",
            "issue": "Rollback committed; journal cleanup is deferred.",
        })
    result = {
        "status": "passed",
        "stage": "rolled_back",
        "rollback_commit_point": "previous_to_active_atomic_swap",
        "active": str(active),
        "quarantined": str(quarantine),
        "assurance": assurance,
        "historical_run_revisions_rewritten": False,
    }
    if cleanup_findings:
        result["cleanup"] = {"status": "deferred", "findings": cleanup_findings}
    return result


@dataclass(frozen=True)
class RetryTarget:
    category: str
    value: str

    @classmethod
    def parse(cls, target_id: str) -> "RetryTarget":
        category, separator, value = target_id.partition(":")
        return cls(category, value) if separator else cls("section", target_id)

    @property
    def target_id(self) -> str:
        return self.value if self.category == "section" else f"{self.category}:{self.value}"


VERIFICATION_TASK_BY_TARGET = {
    "content": "clinical_content_verification",
    "visual": "rendered_page_visual_verification",
}

LAYOUT_RULE_BY_VISUAL_CHECK = {
    "orphan_heading": "heading_cohesion",
    "artificial_pagination": "body_pagination",
    "bad_table_split": "table_pagination",
}
def _read(path: Path) -> dict[str, Any]:
    value = json.loads(path.read_text(encoding="utf-8"))
    if not isinstance(value, dict): raise ValueError(f"Expected JSON object: {path}")
    return value


def _read_corpus(path: Path) -> list[dict[str, Any]]:
    value = json.loads(path.read_text(encoding="utf-8"))
    if not isinstance(value, list) or any(not isinstance(item, dict) for item in value):
        raise ValueError(f"Expected a JSON list of objects: {path}")
    return value


def _write(path: Path, value: Mapping[str, Any]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True); path.write_text(json.dumps(value, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")


def _reference(run_dir: Path) -> tuple[Path, dict[str, Any]]:
    path = run_dir / REFERENCE
    if not path.is_file(): raise FileNotFoundError(f"Create {path} from the study inputs before running prepare.")
    return path, _read(path)


def _slug(value: Any) -> str:
    text = "".join(character.lower() if character.isalnum() else "-" for character in str(value or "study"))
    return "-".join(part for part in text.split("-") if part)[:60] or "study"


def desktop_operation_state_path(run_dir: Path, operation_id: str = "default") -> Path:
    """Return the canonical persisted-state path for one Desktop operation."""
    operation_key = _slug(operation_id)
    name = "desktop-operation.json" if operation_key == "default" else f"desktop-operation-{operation_key}.json"
    return Path(run_dir).resolve() / "logs" / name


def _source_path(run_dir: Path, reference: Mapping[str, Any]) -> Path:
    recorded = reference.get("source", {}).get("source_of_truth_file") if isinstance(reference.get("source"), Mapping) else None
    if recorded:
        path = Path(str(recorded)); return path if path.is_absolute() else run_dir / path
    protocol = reference.get("meta", {}).get("protocol_number") if isinstance(reference.get("meta"), Mapping) else "study"
    title = reference.get("study", {}).get("title") if isinstance(reference.get("study"), Mapping) else "study"
    return run_dir / "reference" / f"source-of-truth--{_slug(protocol)}--{_slug(title)}.md"


def _approved_payload(reference: Mapping[str, Any]) -> dict[str, Any]:
    """Return only reviewer-controlled and approval-bound data."""
    payload = copy.deepcopy(dict(reference))
    payload.pop("generation", None)
    payload.pop("approval", None)
    return payload


def _approval_valid(
    run_dir: Path,
    reference: Mapping[str, Any],
    *,
    contracted_bundle: Mapping[str, Any] | None = None,
) -> tuple[bool, str]:
    approval = reference.get("approval") if isinstance(reference.get("approval"), Mapping) else {}
    source = _source_path(run_dir, reference)
    if str(approval.get("status", "")).casefold() != "approved": return False, "The Source-of-Truth Markdown has not been explicitly approved."
    if not source.is_file(): return False, "The approved Source-of-Truth Markdown is missing."
    if approval.get("source_sha256") != sha256_file(source): return False, "The Source-of-Truth Markdown changed after approval; approve the current file again."
    revision_id = str(approval.get("revision_id") or "")
    snapshot_path = run_dir / "revisions" / revision_id / "approved-reference.json"
    approved_source = run_dir / "revisions" / revision_id / "approved-source.md"
    if not revision_id or not snapshot_path.is_file() or not approved_source.is_file():
        return False, "The write-once approved revision snapshot is missing."
    if sha256_file(approved_source) != approval.get("source_sha256"):
        return False, "The approved revision copy of the Source-of-Truth is not hash-identical."
    if approval.get("approved_reference_sha256") != sha256_file(snapshot_path):
        return False, "The approved reference snapshot changed after approval."
    snapshot = _read(snapshot_path)
    if _approved_payload(reference) != _approved_payload(snapshot):
        return False, "Study inputs changed after approval; prepare and approve a new Source-of-Truth revision."
    expected_governing = sha256_value(governing_resources(
        SCRIPT_DIR.parent,
        snapshot,
        contracted_bundle=contracted_bundle,
    ))
    if approval.get("governing_sha256") != expected_governing:
        return False, "Generation contracts, templates, or implementation changed after approval; approve the unchanged Source-of-Truth again to create a new immutable revision."
    return True, ""


_OPTIONAL_REVIEW_FIELDS = (
    "study.condition",
    "procedures.methods",
    "procedures.unscheduled_visits",
    "procedures.discontinuation",
    "procedures.termination",
    "procedures.study_termination",
    "statistics.methodology",
    "statistics.analysis_populations",
    "statistics.missing_data_handling",
    "safety.general_information",
    "safety.monitoring",
    "safety.adverse_events",
    "safety.follow_up",
    "confidentiality.data_handling",
    "confidentiality.publication",
    "risks_benefits.risks",
    "risks_benefits.benefits",
    "risks_benefits.costs",
    "risks_benefits.alternatives",
    "risks_benefits.privacy",
    "risks_benefits.injury_handling",
)
_OPTIONAL_PRS_REVIEW_FIELDS = (
    "regulatory.prs.observational_study_design",
    "regulatory.prs.time_perspective",
    "regulatory.prs.verification_date",
    "regulatory.prs.start_date",
    "regulatory.prs.start_date_type",
    "regulatory.prs.primary_completion_date",
    "regulatory.prs.primary_completion_date_type",
    "regulatory.prs.study_completion_date",
    "regulatory.prs.study_completion_date_type",
    "regulatory.prs.patient_registry",
    "regulatory.prs.target_duration_quantity",
    "regulatory.prs.target_duration_units",
)


def _ensure_optional_review_fields(reference: dict[str, Any], branch: str | None) -> None:
    missing = object()
    paths = _OPTIONAL_REVIEW_FIELDS + (_OPTIONAL_PRS_REVIEW_FIELDS if branch != "Retrospective" else ())
    for path in paths:
        if get_path(reference, path, missing) is missing:
            set_path(reference, path, None)
    for outcome_kind in ("primary", "secondary", "other"):
        outcomes = get_path(reference, f"endpoints.{outcome_kind}")
        if not isinstance(outcomes, list):
            continue
        for outcome in outcomes:
            if isinstance(outcome, dict):
                outcome.setdefault("description", "")
                outcome.setdefault("analysis_method", "")


def prepare(run_dir: Path, *, today: date | None = None, **_: Any) -> dict[str, Any]:
    """Validate mandatory inputs and create the reviewer-editable source file."""
    run_dir = run_dir.resolve(); reference_path, reference = _reference(run_dir)
    contract = source_contract(reference, derive_prs_study_type=True)
    if contract["status"] != "passed":
        path = run_dir / "reference/missing-inputs.md"; path.parent.mkdir(parents=True, exist_ok=True)
        path.write_text(repair_report(contract["blocking_findings"]), encoding="utf-8")
        return {"status": "blocked", "stage": "input_collection", "study_type": contract.get("study_type"), "missing": contract["blocking_findings"], "missing_inputs": path.relative_to(run_dir).as_posix(), "client_outputs": []}
    reference = contract["normalized_reference"]
    _ensure_optional_review_fields(reference, contract.get("study_type"))
    meta = reference.setdefault("meta", {})
    if not str(meta.get("date") or "").strip():
        meta["date"] = (today or date.today()).strftime("%d %b %Y")
    if meta.get("version") is None:
        meta["version"] = ""
    source = _source_path(run_dir, reference); source.parent.mkdir(parents=True, exist_ok=True)
    source.write_text(source_truth_markdown(reference), encoding="utf-8")
    reference.setdefault("source", {})["source_of_truth_file"] = source.relative_to(run_dir).as_posix()
    reference["approval"] = {"status": "awaiting_approval", "review_file": source.relative_to(run_dir).as_posix()}
    _write(reference_path, reference)
    (run_dir / "reference/missing-inputs.md").unlink(missing_ok=True)
    relative_source = source.relative_to(run_dir).as_posix()
    review_delivery = {
        "mode": "file",
        "role": "editable_source_of_truth",
        "path": relative_source,
        "absolute_path": source.resolve().as_posix(),
        "filename": source.name,
        "mime_type": "text/markdown",
        "editable": True,
        "inline_chat": False,
        "must_attach": True,
    }
    return {"status": "awaiting_approval", "stage": "source_review", "study_type": contract.get("study_type"), "source_of_truth": relative_source, "review_delivery": review_delivery, "client_outputs": []}


def approve(run_dir: Path, *, approved_by: str = "client", source_md: Path | None = None, **_: Any) -> dict[str, Any]:
    """Parse the current reviewer file exactly and bind explicit approval."""
    run_dir = run_dir.resolve(); reference_path, prior = _reference(run_dir)
    source = source_md.resolve() if source_md else _source_path(run_dir, prior)
    if not source.is_file(): raise FileNotFoundError(source)
    reference_dir = (run_dir / "reference").resolve()
    try: source.relative_to(reference_dir)
    except ValueError:
        canonical = reference_dir / "approved-source-of-truth.md"; shutil.copy2(source, canonical); source = canonical
    parsed = parse_source_truth(source.read_text(encoding="utf-8"), prior)
    contract = source_contract(parsed)
    report_path = run_dir / "reference/review-parse-report.md"
    report_path.write_text(repair_report(contract["blocking_findings"]), encoding="utf-8")
    if contract["status"] != "passed":
        return {"status": "blocked", "stage": "approval", "findings": contract["blocking_findings"], "client_outputs": []}
    reference = contract["normalized_reference"]
    reference.pop("generation", None)
    reference.setdefault("source", {})["source_of_truth_file"] = source.relative_to(run_dir).as_posix()
    digest = sha256_file(source)
    approval = {"status": "approved", "approved_by": approved_by, "approved_at": datetime.now(timezone.utc).isoformat(), "review_file": source.relative_to(run_dir).as_posix(), "source_sha256": digest}
    reference["approval"] = approval
    bundle = contracted_template_bundle(SCRIPT_DIR.parent, reference)
    governing = governing_resources(SCRIPT_DIR.parent, reference, contracted_bundle=bundle)
    governing_sha256 = sha256_value(governing)
    revision_identity = sha256_value({
        "approved_source_sha256": digest,
        "approved_reference": _approved_payload(reference),
        "governing_resources": governing,
    })
    revision_id = f"r-{revision_identity[:12]}"
    approval["revision_id"] = revision_id
    approval["governing_sha256"] = governing_sha256
    revision_dir = run_dir / "revisions" / revision_id
    snapshot_path = revision_dir / "approved-reference.json"
    approved_source_path = revision_dir / "approved-source.md"
    if revision_dir.exists():
        if not snapshot_path.is_file() or not approved_source_path.is_file():
            return {"status": "blocked", "stage": "approval", "findings": [{"category": "revision", "field": revision_id, "issue": "An incomplete revision already occupies the approved source hash."}], "client_outputs": []}
        existing = _read(snapshot_path)
        if _approved_payload(existing) != _approved_payload(reference) or sha256_file(approved_source_path) != digest:
            return {"status": "blocked", "stage": "approval", "findings": [{"category": "revision", "field": revision_id, "issue": "A different write-once revision already occupies the approved source hash."}], "client_outputs": []}
        reference["approval"] = dict(existing.get("approval", approval))
    else:
        revision_dir.mkdir(parents=True)
        _write(snapshot_path, reference)
        approved_source_path.write_text(source.read_text(encoding="utf-8"), encoding="utf-8")
    reference["approval"]["approved_reference_sha256"] = sha256_file(snapshot_path)
    _write(reference_path, reference)
    return {"status": "passed", "stage": "approval", "revision_id": revision_id, "source_sha256": digest, "client_outputs": []}


def validate(run_dir: Path, **_: Any) -> dict[str, Any]:
    run_dir = run_dir.resolve(); _, reference = _reference(run_dir)
    contract = source_contract(reference); approved, approval_issue = _approval_valid(run_dir, reference)
    findings = list(contract["blocking_findings"])
    if not approved: findings.append({"category": "approval", "field": "approval", "issue": approval_issue})
    return {"status": "passed" if not findings else "blocked", "stage": "readiness", "study_type": contract.get("study_type"), "findings": findings, "client_outputs": []}


def _awaiting(revision_dir: Path, *, stage: str, paths: list[Path], findings: list[Mapping[str, Any]] | None = None) -> dict[str, Any]:
    handoffs = []
    for path in paths:
        request = _read(path)
        handoff = {
            "request_path": path.relative_to(revision_dir).as_posix(),
            "response_path": str(request["response_path"]),
            "task": str(request["task"]),
            "batch_id": request.get("batch_id"),
        }
        if request.get("task") == "rendered_page_visual_verification":
            handoff["request_id"] = str(request["request_id"])
            handoff["request_sha256"] = str(request["request_sha256"])
            handoff["fallback_owner"] = "parent"
            handoff["completion_requirement"] = "inspect_every_bound_page_image"
        handoffs.append(handoff)
    return {
        "status": "awaiting_hermes",
        "stage": stage,
        "revision_id": revision_dir.name,
        "requests": [item["request_path"] for item in handoffs],
        "handoffs": handoffs,
        "findings": list(findings or []),
        "client_outputs": [],
    }


def _repair_block(
    run_dir: Path,
    stage: str,
    findings: list[Mapping[str, Any]],
    *,
    candidate_outputs: list[dict[str, Any]] | None = None,
) -> dict[str, Any]:
    path = run_dir / "reference/repair-report.md"
    path.write_text(repair_report(findings), encoding="utf-8")
    result: dict[str, Any] = {
        "status": "blocked",
        "stage": stage,
        "findings": [dict(finding) for finding in findings],
        "repair_report": path.relative_to(run_dir).as_posix(),
        "client_outputs": [],
    }
    if candidate_outputs is not None:
        result["candidate_outputs"] = candidate_outputs
    return result


def _document_report_failure(
    run_dir: Path,
    revision_dir: Path,
    document_report: Mapping[str, Any],
) -> tuple[list[dict[str, Any]], dict[str, Any] | None]:
    """Interpret one blocked DOCX report identically on every render path."""
    findings = []
    for item in document_report.get("artifacts", []):
        for raw in item.get("findings", []):
            finding = dict(raw)
            governed = (
                finding.get("recovery_class") in RECOVERY_POLICIES
                and finding.get("action") == RECOVERY_POLICIES.get(str(finding.get("recovery_class")))
            )
            findings.append(
                finding if governed
                else recovery_finding({**finding, "target_ids": [f"layout:{item['artifact']}"]}, "document_structure_defect")
            )
    classification = [
        finding for finding in findings
        if finding.get("category") == "layout-repair-classification"
    ]
    if not classification:
        return findings, None
    return findings, _repair_block(
        run_dir,
        "layout_repair_classification",
        classification,
        candidate_outputs=_candidate_outputs(revision_dir),
    )


def _candidate_fingerprint(
    repo_root: Path,
    revision_dir: Path,
    reference: Mapping[str, Any],
    model: Mapping[str, Any],
    *,
    contracted_bundle: Mapping[str, Any],
    font_substitutions: Mapping[str, str] | None = None,
    layout_repairs: Mapping[str, Iterable[Mapping[str, str]]] | None = None,
) -> tuple[str, dict[str, Any]]:
    implementation_files = [repo_root / "scripts" / name for name in ("workflow.py", "contracts.py", "drafting.py", "rendering.py", "quality.py", "prs_xml.py")]
    accepted_files = sorted((revision_dir / "hermes/accepted").glob("*.json"))
    payload = {
        "approved_reference_sha256": sha256_file(revision_dir / "approved-reference.json"),
        "approved_source_sha256": get_path(reference, "approval.source_sha256"),
        "contracted_template_bundle": dict(contracted_bundle),
        "implementation": {path.name: sha256_file(path) for path in implementation_files},
        "accepted_drafts": {path.relative_to(revision_dir).as_posix(): sha256_file(path) for path in accepted_files},
        "merged_model_sha256": sha256_value(model),
        "font_substitutions": dict(sorted((font_substitutions or {}).items())),
        "layout_repairs": {
            str(artifact): _normalized_layout_repair_records(repairs)
            for artifact, repairs in sorted(dict(layout_repairs or {}).items())
        },
    }
    return sha256_value(payload), payload


def _cached_build(revision_dir: Path, fingerprint: str) -> dict[str, Any] | None:
    path = revision_dir / "candidate-build.json"
    if not path.is_file():
        return None
    try:
        build = _read(path)
    except (OSError, ValueError, json.JSONDecodeError):
        return None
    if build.get("fingerprint") != fingerprint:
        return None
    for artifact in build.get("candidate_files", []):
        target = revision_dir / str(artifact.get("path"))
        if not target.is_file() or sha256_file(target) != artifact.get("sha256"):
            return None
    render_report = build.get("render_report", {})
    for artifact in render_report.get("artifacts", []):
        for key in ("docx", "pdf"):
            target = revision_dir / str(artifact.get(key))
            if not target.is_file() or sha256_file(target) != artifact.get(f"{key}_sha256"):
                return None
        for page in artifact.get("pages", []):
            target = revision_dir / str(page.get("path"))
            if not target.is_file() or sha256_file(target) != page.get("sha256"):
                return None
    return build


def _record_build(revision_dir: Path, fingerprint: str, governing: Mapping[str, Any], contracted_bundle: Mapping[str, Any], document_report: Mapping[str, Any], xml_report: Mapping[str, Any] | None, render_report: Mapping[str, Any]) -> dict[str, Any]:
    candidate_files = _candidate_file_records(revision_dir)
    build = {"fingerprint": fingerprint, "governing_resources": dict(governing), "contracted_template_bundle": dict(contracted_bundle), "candidate_files": candidate_files, "document_report": dict(document_report), "xml_report": dict(xml_report) if xml_report else None, "render_report": dict(render_report)}
    _write(revision_dir / "candidate-build.json", build)
    return build


def _candidate_file_records(revision_dir: Path) -> list[dict[str, Any]]:
    return [
        {
            "path": path.relative_to(revision_dir).as_posix(),
            "sha256": sha256_file(path),
            "bytes": path.stat().st_size,
        }
        for path in sorted((revision_dir / "candidate").glob("*"))
        if path.is_file()
    ]


def _cached_candidate_structure(revision_dir: Path, fingerprint: str) -> dict[str, Any] | None:
    path = revision_dir / "candidate-structure.json"
    if not path.is_file():
        return None
    try:
        snapshot = _read(path)
    except (OSError, ValueError, json.JSONDecodeError):
        return None
    if snapshot.get("fingerprint") != fingerprint:
        return None
    for artifact in snapshot.get("candidate_files", []):
        target = revision_dir / str(artifact.get("path"))
        if not target.is_file() or sha256_file(target) != artifact.get("sha256"):
            return None
    return snapshot


def _record_candidate_structure(
    revision_dir: Path,
    fingerprint: str,
    governing: Mapping[str, Any],
    contracted_bundle: Mapping[str, Any],
    document_report: Mapping[str, Any],
    xml_report: Mapping[str, Any] | None,
    expected_files: Iterable[str],
) -> dict[str, Any]:
    snapshot = {
        "fingerprint": fingerprint,
        "governing_resources": dict(governing),
        "contracted_template_bundle": dict(contracted_bundle),
        "candidate_files": _candidate_file_records(revision_dir),
        "document_report": dict(document_report),
        "xml_report": dict(xml_report) if xml_report else None,
        "font_substitutions": dict(governing.get("font_substitutions") or {}),
        "expected_files": sorted(str(name) for name in expected_files),
        "status": "structurally_valid",
    }
    _write(revision_dir / "candidate-structure.json", snapshot)
    return snapshot


def _merge_artifact_reports(prior: Mapping[str, Any], current: Mapping[str, Any]) -> dict[str, Any]:
    """Replace only newly generated artifact rows in a complete prior report."""
    def rows(report: Mapping[str, Any]) -> dict[str, dict[str, Any]]:
        result: dict[str, dict[str, Any]] = {}
        for item in report.get("artifacts", []):
            if not isinstance(item, Mapping) or not item.get("artifact"):
                continue
            row = dict(item)
            row.setdefault("renderer", report.get("renderer"))
            row.setdefault("page_renderer", report.get("page_renderer"))
            result[str(item["artifact"])] = row
        return result

    artifacts = rows(prior)
    artifacts.update(rows(current))
    merged = {**dict(prior), **dict(current), "artifacts": [artifacts[key] for key in sorted(artifacts)]}
    merged["status"] = "passed" if artifacts and all(item.get("status", "passed") == "passed" for item in artifacts.values()) else "blocked"
    return merged


def _merge_partial_assurance(
    prior_build: Mapping[str, Any],
    assurance: Mapping[str, Any],
    repaired_artifacts: set[str],
) -> tuple[dict[str, Any], dict[str, Any]]:
    """Return the complete exact-artifact evidence persisted after a local repair."""
    complete = dict(assurance)
    render_report = dict(complete.get("render") or {})
    if repaired_artifacts and render_report.get("status") == "passed":
        render_report = _merge_artifact_reports(prior_build.get("render_report", {}), render_report)
        complete["render"] = render_report
    return complete, render_report


def _unaffected_build_is_valid(
    revision_dir: Path,
    build: Mapping[str, Any],
    repaired_artifacts: set[str],
) -> bool:
    """Prove every retained artifact still matches the build being merged."""
    for item in build.get("candidate_files", []):
        path = revision_dir / str(item.get("path"))
        if path.suffix.casefold() == ".docx" and path.stem in repaired_artifacts:
            continue
        if not path.is_file() or sha256_file(path) != item.get("sha256"):
            return False
    for artifact in build.get("render_report", {}).get("artifacts", []):
        if str(artifact.get("artifact")) in repaired_artifacts:
            continue
        for key in ("docx", "pdf"):
            path = revision_dir / str(artifact.get(key))
            if not path.is_file() or sha256_file(path) != artifact.get(f"{key}_sha256"):
                return False
        for page in artifact.get("pages", []):
            path = revision_dir / str(page.get("path"))
            if not path.is_file() or sha256_file(path) != page.get("sha256"):
                return False
    return True


def _candidate_outputs(revision_dir: Path) -> list[dict[str, Any]]:
    """Describe retained artifacts that have not passed the Delivery Gate."""
    return [
        {
            "path": path.relative_to(revision_dir).as_posix(),
            "sha256": sha256_file(path),
            "bytes": path.stat().st_size,
            "delivery_status": "internal_candidate",
        }
        for path in sorted((revision_dir / "candidate").glob("*"))
        if path.is_file()
    ]


def desktop_attachment_reply(manifest: Mapping[str, Any], *, run_dir: Path | None = None) -> dict[str, Any]:
    """Build the supported file-link payload for the final Desktop reply.

    The manifest is the only source for attachments.  In particular, this
    deliberately does not scan the run directory, which could expose drafts,
    QA evidence, or rendered previews.
    """
    outputs = manifest.get("client_outputs")
    if not isinstance(outputs, list) or not outputs:
        raise ValueError("A passing Generation Manifest must contain client outputs.")
    attachments = []
    for item in outputs:
        if not isinstance(item, Mapping):
            raise ValueError("Generation Manifest contains an invalid client output.")
        path = str(item.get("path") or "")
        filename = Path(path).name
        path_parts = Path(path).parts
        digest = str(item.get("sha256") or "")
        if (
            not path
            or Path(path).is_absolute()
            or ".." in path_parts
            or not filename
            or filename != path.rsplit("/", 1)[-1]
            or len(digest) != 64
            or any(character not in "0123456789abcdef" for character in digest.lower())
            or int(item.get("bytes") or 0) < 1
        ):
            raise ValueError("Generation Manifest contains an invalid client output path.")
        absolute_path = (run_dir / path).resolve().as_posix() if run_dir is not None else path
        attachments.append({
            "filename": filename,
            "path": path,
            "absolute_path": absolute_path,
            "sha256": digest,
            "bytes": int(item.get("bytes") or 0),
            "mime_type": "application/xml" if filename.lower().endswith(".xml") else "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            "link": f"[{filename}](<{absolute_path}>)",
        })
    return {
        "type": "desktop_file_attachments",
        "attachments": attachments,
        "client_outputs_only": True,
    }


def confirm_desktop_delivery(
    manifest: Mapping[str, Any],
    reply: Mapping[str, Any],
    opener: Callable[[str], bytes],
    *,
    deadline: float | None = None,
    retries: int = DESKTOP_DELIVERY_RETRIES,
    clock: Callable[[], float] | None = None,
) -> dict[str, Any]:
    """Retrieve and open every attachment, preserving the manifest bytes.

    ``opener`` is the actual Desktop media boundary: it must return the bytes
    obtained by opening the link, rather than merely echoing a path. Retries
    reuse the same validated attachment and never invoke generation again.
    """
    clock = clock or time.monotonic
    expected = desktop_attachment_reply(manifest)["attachments"]
    actual = reply.get("attachments") if isinstance(reply, Mapping) else None
    findings = []
    def identity(item: Mapping[str, Any]) -> tuple[Any, ...]:
        return tuple(item.get(key) for key in ("filename", "path", "sha256", "bytes"))

    if not isinstance(actual, list) or [identity(item) for item in actual if isinstance(item, Mapping)] != [identity(item) for item in expected]:
        findings.append(recovery_finding({"category": "delivery", "field": "attachments", "issue": "Desktop reply attachments do not match the immutable Generation Manifest."}, "transport_fault"))
    if findings:
        return {"status": "blocked", "confirmed": False, "findings": findings, "attempts": 0}
    opened = []
    total_attempts = 0
    for item, delivered in zip(expected, actual, strict=True):
        attempts = 0
        last_issue = ""
        opened_current = False
        while attempts <= retries:
            attempts += 1
            total_attempts += 1
            if deadline is not None and clock() >= deadline:
                last_issue = "Desktop delivery deadline expired."
                break
            try:
                payload = opener(str(delivered.get("absolute_path") or delivered["path"]))
                if not isinstance(payload, bytes):
                    raise TypeError("Desktop opener did not return bytes.")
                if deadline is not None and clock() >= deadline:
                    raise TimeoutError("Desktop delivery deadline expired before retrieval was confirmed.")
                digest = hashlib.sha256(payload).hexdigest()
                if len(payload) != int(item["bytes"]) or digest != item["sha256"]:
                    raise ValueError("Retrieved file bytes do not match the immutable Generation Manifest.")
                opened.append({"filename": item["filename"], "sha256": digest, "bytes": len(payload), "attempts": attempts})
                opened_current = True
                break
            except Exception as exc:  # transport failures are reported, never certified
                last_issue = str(exc)
        else:
            last_issue = last_issue or "Desktop file transfer failed."
        if not opened_current:
            findings.append(recovery_finding({"category": "delivery", "field": item["filename"], "issue": last_issue or "Desktop file transfer failed."}, "transport_fault"))
            break
    return {
        "status": "confirmed" if not findings else "blocked",
        "confirmed": not findings,
        "opened": opened,
        "attempts": total_attempts,
        "findings": findings,
    }


def _desktop_delivery_set_finding(
    run_dir: Path,
    manifest: Mapping[str, Any],
) -> dict[str, Any] | None:
    try:
        _, delivery_reference = _reference(run_dir)
    except FileNotFoundError:
        return None
    expected_outputs = set(document_set(get_path(delivery_reference, "meta.study_type")))
    actual_outputs = {
        Path(str(item.get("path") or "")).name
        for item in manifest.get("client_outputs", [])
        if isinstance(item, Mapping)
    }
    if manifest.get("status") == "passed" and actual_outputs == expected_outputs:
        return None
    return {
        "category": "delivery",
        "field": "client_outputs",
        "issue": f"Desktop delivery requires the exact Branch Document Set; expected={sorted(expected_outputs)}, actual={sorted(actual_outputs)}.",
    }


def _handoff_response_is_bound(
    run_dir: Path,
    revision_id: str,
    handoff: Mapping[str, Any],
) -> bool:
    revision_dir = run_dir / "revisions" / revision_id
    request_path = revision_dir / str(handoff.get("request_path") or "")
    try:
        request = json.loads(request_path.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError):
        return False
    return (
        isinstance(request, Mapping)
        and request.get("request_id") == handoff.get("request_id")
        and request.get("request_sha256") == handoff.get("request_sha256")
        and request.get("response_path") == handoff.get("response_path")
        and request.get("task") == handoff.get("task")
        and verification_response_is_complete(revision_dir, request_path)
    )


def run_desktop_operation(
    run_dir: Path,
    *,
    handoff_runner: Callable[[list[Mapping[str, Any]], float], Any],
    fallback_handoff_runner: Callable[[list[Mapping[str, Any]], float], Any] | None = None,
    opener: Callable[[str], bytes],
    operation_id: str = "default",
    budget_seconds: float = DESKTOP_OPERATION_BUDGET_SECONDS,
    clock: Callable[[], float] | None = None,
    wall_clock: Callable[[], float] | None = None,
    runtime_identity: Mapping[str, Any] | None = None,
    release_identity: Mapping[str, Any] | None = None,
    progress: Callable[[str, float], Any] | None = None,
    cleanup: Callable[[str, float], Mapping[str, Any] | None] | None = None,
    stage_soft_budgets: Mapping[str, float] | None = None,
    require_promoted_runtime: bool = True,
) -> dict[str, Any]:
    """Run the post-approval lifecycle and confirm its Desktop file delivery.

    Hermes owns the ``handoff_runner`` boundary and Desktop owns ``opener``;
    this function only routes path metadata, advances the public workflow, and
    validates the immutable delivery manifest. The persisted deadline uses a
    UTC epoch so another process or compatible runtime cannot reset it.
    Monotonic time is converted to that epoch only within this process.
    """
    run_dir = run_dir.resolve()
    clock = clock or time.monotonic
    wall_clock = wall_clock or time.time
    if budget_seconds <= 0:
        raise ValueError("Desktop operation budget must be positive.")
    current_runtime = dict(runtime_identity or _current_python_runtime())
    if _runtime_version(current_runtime) < (*MINIMUM_PYTHON_VERSION, 0):
        raise RuntimeError("The Desktop operation requires an explicitly resolved Python 3.10+ runtime.")
    current_release_identity = dict(release_identity or _active_release_identity(SCRIPT_DIR.parent))
    if not str(current_release_identity.get("package_fingerprint") or ""):
        raise ValueError("Desktop operation release identity requires a package fingerprint.")
    try:
        _, operation_reference = _reference(run_dir)
    except FileNotFoundError:
        operation_reference = {}
    operation_approval = operation_reference.get("approval") or {}
    current_approval_identity = (
        {
            key: operation_approval.get(key)
            for key in (
                "status", "approved_by", "approved_at", "revision_id", "source_sha256",
                "approved_reference_sha256", "governing_sha256",
            )
        }
        if operation_approval
        else None
    )
    soft_budgets = {
        str(stage): float(seconds)
        for stage, seconds in (DESKTOP_STAGE_SOFT_BUDGETS if stage_soft_budgets is None else stage_soft_budgets).items()
        if float(seconds) > 0
    }
    process_monotonic_anchor = clock()
    process_epoch_anchor = wall_clock()

    def epoch_now() -> float:
        return process_epoch_anchor + max(0.0, clock() - process_monotonic_anchor)

    state_path = desktop_operation_state_path(run_dir, operation_id)
    try:
        persisted = _read(state_path) if state_path.is_file() else {}
    except (OSError, ValueError, json.JSONDecodeError):
        persisted = {
            "operation_id": operation_id,
            "budget_seconds": budget_seconds,
            "status": "unreadable",
        }
    if persisted.get("operation_id") not in {None, operation_id}:
        raise ValueError("Desktop operation state belongs to a different operation.")
    recorded_release_identity = persisted.get("release_identity")
    if isinstance(recorded_release_identity, Mapping) and dict(recorded_release_identity) != current_release_identity:
        return {
            "status": "blocked",
            "stage": "release_identity",
            "findings": [{
                "category": "release",
                "field": "release_identity",
                "issue": "Desktop operation evidence belongs to a different Promoted Release or governed configuration identity.",
            }],
            "client_outputs": [],
        }
    recorded_approval_identity = persisted.get("approval_identity")
    if isinstance(recorded_approval_identity, Mapping) and dict(recorded_approval_identity) != current_approval_identity:
        return {
            "status": "blocked",
            "stage": "approval_identity",
            "findings": [{
                "category": "approval",
                "field": "approval_identity",
                "issue": "Desktop operation evidence belongs to a different approved immutable revision.",
            }],
            "client_outputs": [],
        }

    terminal = persisted.get("status") in {"passed", "blocked", "timeout"}
    if terminal and isinstance(persisted.get("result"), Mapping):
        prior_result = dict(persisted["result"])
        current_reference = operation_reference
        recorded_bundle = prior_result.get("contracted_template_bundle")
        requires_bundle_validation = prior_result.get("status") == "passed"
        if requires_bundle_validation and not (
            isinstance(recorded_bundle, Mapping)
            and recorded_bundle.get("identity_sha256")
        ):
            return {
                "status": "blocked",
                "stage": "contracted_template_bundle",
                "findings": [{"category": "contract", "field": "contracted_template_bundle", "issue": "Persisted passing Desktop evidence lacks a complete Contracted Template Bundle identity."}],
                "client_outputs": [],
            }
        if requires_bundle_validation:
            try:
                _, current_reference = _reference(run_dir)
                current_bundle = contracted_template_bundle(SCRIPT_DIR.parent, current_reference)
            except (ContractedTemplateBundleError, OSError, ValueError, json.JSONDecodeError) as exc:
                return {
                    "status": "blocked",
                    "stage": "contracted_template_bundle",
                    "findings": [{"category": "contract", "field": "contracted_template_bundle", "issue": f"Persisted Desktop evidence cannot be validated against the current governed bundle: {exc}"}],
                    "client_outputs": [],
                }
            if current_bundle.get("identity_sha256") != recorded_bundle.get("identity_sha256"):
                return {
                    "status": "blocked",
                    "stage": "contracted_template_bundle",
                    "findings": [{"category": "contract", "field": "contracted_template_bundle", "issue": "Persisted Desktop delivery evidence belongs to a stale Contracted Template Bundle."}],
                    "client_outputs": [],
                }
        if requires_bundle_validation and not (
            isinstance(recorded_release_identity, Mapping)
            and recorded_release_identity.get("package_fingerprint")
        ):
            return {
                "status": "blocked",
                "stage": "release_identity",
                "findings": [{
                    "category": "release",
                    "field": "package_fingerprint",
                    "issue": "Persisted passing Desktop evidence lacks a Promoted Release fingerprint.",
                }],
                "client_outputs": [],
            }
        if requires_bundle_validation:
            manifest_name = str(prior_result.get("manifest") or "")
            manifest_path = (run_dir / manifest_name).resolve()
            try:
                manifest_path.relative_to(run_dir)
                manifest = _read(manifest_path)
                prepared_ledger = validate_gate_ledger(
                    SCRIPT_DIR.parent,
                    dict(manifest.get("gate_ledger") or {}) if isinstance(manifest.get("gate_ledger"), Mapping) else {},
                )
                final_ledger = validate_gate_ledger(
                    SCRIPT_DIR.parent,
                    dict(prior_result.get("gate_ledger") or {}) if isinstance(prior_result.get("gate_ledger"), Mapping) else {},
                )
                expected_attempts = list(
                    dict(current_reference.get("generation") or {}).get("gate_attempts") or []
                )
                _validate_expected_gate_attempts(manifest_path.parent, expected_attempts)
                if (
                    prepared_ledger["attempt_id"] != final_ledger["attempt_id"]
                    or prepared_ledger["predecessors"] != final_ledger["predecessors"]
                    or prepared_ledger["records"][:-1] != final_ledger["records"][:-1]
                    or prepared_ledger["records"][-1]["terminal_status"] != "pending"
                    or final_ledger["records"][-1]["terminal_status"] != "passed"
                    or len(final_ledger["predecessors"]) != len(expected_attempts)
                ):
                    raise ValueError("Persisted Desktop gate-ledger lineage or terminal transition is invalid.")
                set_finding = _desktop_delivery_set_finding(run_dir, manifest)
                if set_finding is not None:
                    raise ValueError(set_finding["issue"])
                reply = prior_result.get("desktop_reply") or desktop_attachment_reply(manifest, run_dir=run_dir)
                delivery = confirm_desktop_delivery(manifest, reply, opener)
                if not delivery.get("confirmed"):
                    issue = str((delivery.get("findings") or [{}])[0].get("issue") or "Terminal delivery could not be reconfirmed.")
                    raise ValueError(issue)
            except (OSError, TypeError, ValueError, json.JSONDecodeError) as exc:
                return {
                    "status": "blocked",
                    "stage": "terminal_delivery_validation",
                    "findings": [{
                        "category": "delivery",
                        "field": "terminal_result",
                        "issue": f"Persisted Desktop success no longer has accessible hash-matching outputs: {exc}",
                    }],
                    "client_outputs": [],
                }
        return prior_result
    started_at_epoch, deadline_at_epoch, persisted_budget, runtime_history, stage_history = (
        _desktop_deadline_state(
            persisted,
            budget_seconds=budget_seconds,
            epoch_now=epoch_now(),
        )
    )
    if isinstance(persisted.get("soft_budgets"), Mapping):
        soft_budgets = {
            str(stage): float(seconds)
            for stage, seconds in persisted["soft_budgets"].items()
            if float(seconds) > 0
        }
    current_stage = str(persisted.get("stage") or "approved")
    raw_attempt_counters = persisted.get("attempt_counters")
    if not isinstance(raw_attempt_counters, Mapping):
        raw_attempt_counters = {}
    attempt_counters = {
        "drafting": dict(raw_attempt_counters.get("drafting") or {}),
        "verification": dict(raw_attempt_counters.get("verification") or {}),
        "generation": dict(raw_attempt_counters.get("generation") or {}),
        "handoff_dispatches": dict(raw_attempt_counters.get("handoff_dispatches") or {}),
        "delivery": int(raw_attempt_counters.get("delivery") or 0),
    }
    raw_stage_timings = persisted.get("stage_timings")
    if not isinstance(raw_stage_timings, Mapping):
        raw_stage_timings = {}
    stage_timings = {
        str(stage): dict(timing)
        for stage, timing in raw_stage_timings.items()
        if isinstance(timing, Mapping)
    }
    raw_soft_budget_events = persisted.get("soft_budget_events")
    if not isinstance(raw_soft_budget_events, list):
        raw_soft_budget_events = []
    soft_budget_events = [
        dict(item)
        for item in raw_soft_budget_events
        if isinstance(item, Mapping)
    ]
    raw_cleanup = persisted.get("cleanup")
    cleanup_evidence = dict(raw_cleanup) if isinstance(raw_cleanup, Mapping) else {}
    raw_pending_handoffs = persisted.get("pending_handoffs")
    if not isinstance(raw_pending_handoffs, list):
        raw_pending_handoffs = []
    pending_handoffs = [
        dict(item)
        for item in raw_pending_handoffs
        if isinstance(item, Mapping)
    ]
    last_result: dict[str, Any] = {}
    process_deadline_monotonic = process_monotonic_anchor + max(0.0, deadline_at_epoch - process_epoch_anchor)

    def runtime_key(item: Mapping[str, Any]) -> tuple[Any, ...]:
        return tuple(item.get(key) for key in ("executable", "implementation", "version"))

    if not runtime_history or runtime_key(runtime_history[-1]) != runtime_key(current_runtime):
        runtime_history.append(current_runtime)

    def remaining_seconds() -> float:
        return deadline_at_epoch - epoch_now()

    def save(status: str, result: Mapping[str, Any] | None = None) -> dict[str, Any]:
        payload = {
            "operation_id": operation_id,
            "started_at_epoch": started_at_epoch,
            "deadline_at_epoch": deadline_at_epoch,
            "started_at": datetime.fromtimestamp(started_at_epoch, timezone.utc).isoformat(),
            "deadline_at": datetime.fromtimestamp(deadline_at_epoch, timezone.utc).isoformat(),
            "budget_seconds": persisted_budget,
            "runtime": current_runtime,
            "runtime_history": runtime_history,
            "release_identity": current_release_identity,
            "approval_identity": current_approval_identity,
            "status": status,
            "stage": current_stage,
            "stage_history": stage_history,
            "stage_timings": stage_timings,
            "attempt_counters": attempt_counters,
            "pending_handoffs": pending_handoffs,
            "soft_budgets": soft_budgets,
            "soft_budget_events": soft_budget_events,
        }
        if cleanup_evidence:
            payload["cleanup"] = cleanup_evidence
        if result is not None:
            payload["result"] = dict(result)
        _write(state_path, payload)
        return dict(result or payload)

    def finish(result: Mapping[str, Any]) -> dict[str, Any]:
        nonlocal current_stage, cleanup_evidence
        current_stage = str(result.get("pending_stage") or result.get("stage") or current_stage)
        elapsed = max(0.0, epoch_now() - started_at_epoch)
        measured = {
            **result,
            "elapsed_seconds": round(elapsed, 3),
            "performance_classification": performance_classification(elapsed),
            "target_window_seconds": [
                NORMAL_RUNTIME_TARGET_MIN_SECONDS,
                NORMAL_RUNTIME_TARGET_MAX_SECONDS,
            ],
            "operation_deadline_seconds": DESKTOP_OPERATION_BUDGET_SECONDS,
            "started_at_epoch": started_at_epoch,
            "deadline_at_epoch": deadline_at_epoch,
            "runtime": current_runtime,
            "release_identity": current_release_identity,
        }
        if cleanup is not None and not cleanup_evidence:
            cleanup_evidence = dict(cleanup(str(measured.get("status", "blocked")), max(0.0, remaining_seconds())) or {})
        return save(str(measured.get("status", "blocked")), measured)

    def record_timing(stage: str, elapsed: float) -> None:
        timing = stage_timings.setdefault(stage, {"elapsed_seconds": 0.0, "invocations": 0})
        timing["elapsed_seconds"] = round(float(timing.get("elapsed_seconds") or 0.0) + max(0.0, elapsed), 3)
        timing["invocations"] = int(timing.get("invocations") or 0) + 1

    def record_soft_budget_event(stage: str) -> None:
        elapsed = float(stage_timings.get(stage, {}).get("elapsed_seconds") or 0.0)
        if (
            stage in soft_budgets
            and elapsed >= soft_budgets[stage]
            and not any(item.get("stage") == stage for item in soft_budget_events)
        ):
            soft_budget_events.append({
                "stage": stage,
                "budget_seconds": soft_budgets[stage],
                "elapsed_seconds": round(elapsed, 3),
                "action": "early_parent_fallback" if stage == "independent_verification" else "record_diagnostic",
            })

    def refresh_generation_attempts() -> None:
        try:
            _, operation_reference = _reference(run_dir)
        except (OSError, ValueError, json.JSONDecodeError):
            return
        generation = operation_reference.get("generation")
        if not isinstance(generation, Mapping):
            return
        attempt_counters["generation"] = {
            str(target): int(attempt)
            for target, attempt in dict(generation.get("attempts") or {}).items()
        }
        attempt_counters["verification"].update({
            str(target): int(attempt)
            for target, attempt in dict(generation.get("verification_attempts") or {}).items()
        })

    # Persist before any operation work can be interrupted; otherwise a restart
    # could create a fresh correctness budget.
    save("running")
    while True:
        remaining = remaining_seconds()
        if remaining <= 0:
            retained = [
                dict(item)
                for item in last_result.get("candidate_outputs", [])
                if isinstance(item, Mapping)
            ]
            try:
                _, current_reference = _reference(run_dir)
                current_revision = str(current_reference.get("approval", {}).get("revision_id") or "")
                if current_revision and not retained:
                    retained = _candidate_outputs(run_dir / "revisions" / current_revision)
            except (OSError, ValueError, json.JSONDecodeError):
                pass
            return finish({
                "status": "timeout",
                "stage": "desktop_operation",
                "pending_stage": current_stage if current_stage != "approved" else "generate",
                "deadline_at_epoch": deadline_at_epoch,
                "findings": [{"category": "timeout", "field": "operation", "issue": "The post-approval Desktop operation deadline expired."}],
                "candidate_outputs": retained,
                "pending_result": last_result,
                "client_outputs": [],
            })
        generate_started = clock()
        observed_stages: set[str] = set()

        def observe_generate_stage(stage: str, elapsed: float) -> None:
            observed_stages.add(stage)
            record_timing(stage, elapsed)
            record_soft_budget_event(stage)

        try:
            result = generate(
                run_dir,
                operation_deadline=process_deadline_monotonic,
                clock=clock,
                stage_observer=observe_generate_stage,
                require_promoted_runtime=require_promoted_runtime,
            )
        finally:
            generate_elapsed = clock() - generate_started
            record_timing("generate", generate_elapsed)
            refresh_generation_attempts()
            save("running")
        last_result = dict(result)
        stage = str(result.get("stage") or "generate")
        current_stage = stage
        if stage not in observed_stages:
            # An awaiting result names the next stage; no time has been spent
            # waiting on that handoff yet. Upstream generation is tracked by
            # its observer callbacks and must not consume the next soft budget.
            stage_elapsed = 0.0 if result.get("status") == "awaiting_hermes" else generate_elapsed
            record_timing(stage, stage_elapsed)
            record_soft_budget_event(stage)
        if remaining_seconds() <= 0:
            continue
        stage_history.append({"stage": stage, "status": result.get("status"), "remaining_seconds": round(remaining, 3)})
        save("running")
        if progress is not None:
            progress(stage, max(0.0, remaining_seconds()))
        if result.get("status") == "awaiting_hermes":
            handoffs = result.get("handoffs")
            if not isinstance(handoffs, list) or not handoffs:
                return finish({**result, "status": "blocked", "stage": "hermes_handoff", "client_outputs": []})
            current_handoffs = [dict(item) for item in handoffs if isinstance(item, Mapping)]
            prior_by_response = {
                str(item.get("response_path") or ""): item
                for item in pending_handoffs
            }
            def handoff_identity(item: Mapping[str, Any]) -> tuple[Any, ...]:
                return tuple(item.get(key) for key in (
                    "request_id", "request_sha256", "request_path", "response_path", "task", "attempts",
                ))

            for handoff in current_handoffs:
                prior = prior_by_response.get(str(handoff.get("response_path") or ""))
                if prior is not None and handoff_identity(prior) != handoff_identity(handoff):
                    return finish({
                        "status": "blocked",
                        "stage": "hermes_handoff_integrity",
                        "findings": [{
                            "category": "hermes",
                            "field": str(handoff.get("request_path") or "handoff"),
                            "issue": "A missing worker response was replaced by a changed request instead of redispatching the exact original request.",
                        }],
                        "client_outputs": [],
                    })
            handoffs = [
                prior_by_response.get(str(item.get("response_path") or ""), item)
                for item in current_handoffs
            ]
            pending_handoffs = [dict(item) for item in handoffs]
            save("running")
            remaining = remaining_seconds()
            if remaining <= 0:
                continue
            fallback_attempted: set[str] = set()
            try:
                for handoff in handoffs:
                    if not isinstance(handoff, Mapping):
                        continue
                    counter_name = "verification" if str(handoff.get("task") or "").endswith("verification") else "drafting"
                    for target, attempt in dict(handoff.get("attempts") or {}).items():
                        attempt_counters[counter_name][str(target)] = max(
                            int(attempt_counters[counter_name].get(str(target)) or 0),
                            int(attempt),
                        )
                    dispatch_key = str(handoff.get("request_sha256") or handoff.get("request_id") or handoff.get("request_path") or "")
                    if dispatch_key:
                        attempt_counters["handoff_dispatches"][dispatch_key] = int(
                            attempt_counters["handoff_dispatches"].get(dispatch_key) or 0
                        ) + 1
                stage_elapsed = float(stage_timings.get(stage, {}).get("elapsed_seconds") or 0.0)
                supports_parent_fallback = (
                    stage == "independent_verification"
                    and any(item.get("fallback_owner") == "parent" for item in handoffs)
                )
                revision_id = str(result.get("revision_id") or "")
                fallback_owned: list[Mapping[str, Any]] = [
                    item for item in handoffs
                    if item.get("fallback_owner") == "parent"
                ]
                non_fallback: list[Mapping[str, Any]] = [
                    item for item in handoffs
                    if item.get("fallback_owner") != "parent"
                ]
                dispatch_handoffs: list[Mapping[str, Any]] = [item for item in handoffs]
                fresh_fallback_owned: list[Mapping[str, Any]] = []
                replayed_fallback_owned: list[Mapping[str, Any]] = []
                for item in fallback_owned:
                    dispatch_key = str(
                        item.get("request_sha256") or item.get("request_id")
                        or item.get("request_path") or ""
                    )
                    target = (
                        fresh_fallback_owned
                        if int(attempt_counters["handoff_dispatches"].get(dispatch_key) or 0) == 1
                        else replayed_fallback_owned
                    )
                    target.append(item)
                soft_timeout = min(remaining, soft_budgets.get(stage, remaining))

                def fallback_or_raise(
                    dispatch_error: Exception,
                    dispatched: Sequence[Mapping[str, Any]],
                ) -> None:
                    incomplete = [
                        item for item in dispatched
                        if item.get("fallback_owner") == "parent"
                        and revision_id
                        and not _handoff_response_is_bound(run_dir, revision_id, item)
                    ]
                    if incomplete and fallback_handoff_runner is not None:
                        fallback_attempted.update(
                            str(item.get("response_path") or "") for item in incomplete
                        )
                        fallback_handoff_runner(incomplete, remaining_seconds())
                        return
                    raise dispatch_error

                handoff_started = clock()
                save("running")
                try:
                    if supports_parent_fallback and fallback_owned and non_fallback:
                        with ThreadPoolExecutor(max_workers=1) as executor:
                            hard_future = executor.submit(
                                handoff_runner, non_fallback, remaining,
                            )
                            for delegated, timeout in (
                                (replayed_fallback_owned, 0.0),
                                (fresh_fallback_owned, soft_timeout),
                            ):
                                if not delegated:
                                    continue
                                try:
                                    handoff_runner(delegated, timeout)
                                except Exception as exc:
                                    fallback_or_raise(exc, delegated)
                            hard_future.result()
                    elif supports_parent_fallback and fallback_owned:
                        for delegated, timeout in (
                            (replayed_fallback_owned, 0.0),
                            (fresh_fallback_owned, soft_timeout),
                        ):
                            if not delegated:
                                continue
                            try:
                                handoff_runner(delegated, timeout)
                            except Exception as exc:
                                fallback_or_raise(exc, delegated)
                    else:
                        try:
                            handoff_runner(
                                dispatch_handoffs,
                                soft_timeout if supports_parent_fallback else remaining,
                            )
                        except Exception as exc:
                            fallback_or_raise(exc, dispatch_handoffs)
                finally:
                    record_timing(stage, clock() - handoff_started)
                    record_soft_budget_event(stage)
                    save("running")
                fallback_handoffs = [
                    item for item in handoffs
                    if item.get("fallback_owner") == "parent"
                    and str(item.get("response_path") or "") not in fallback_attempted
                    and revision_id
                    and not _handoff_response_is_bound(run_dir, revision_id, item)
                ]
                if fallback_handoffs:
                    remaining = remaining_seconds()
                    if remaining > 0:
                        fallback_attempted.update(
                            str(item.get("response_path") or "")
                            for item in fallback_handoffs
                        )
                        if fallback_handoff_runner is not None:
                            fallback_handoff_runner(fallback_handoffs, remaining)
                        else:
                            parent_handoffs = [{**item, "reviewer_owner": "parent"} for item in fallback_handoffs]
                            handoff_runner(parent_handoffs, remaining)
            except Exception as exc:
                revision_id = str(result.get("revision_id") or "")
                fallback_handoffs = [
                    item for item in handoffs
                    if item.get("fallback_owner") == "parent"
                    and str(item.get("response_path") or "") not in fallback_attempted
                    and revision_id
                    and not _handoff_response_is_bound(run_dir, revision_id, item)
                ]
                if fallback_handoffs:
                    remaining = remaining_seconds()
                    if remaining > 0:
                        try:
                            fallback_attempted.update(
                                str(item.get("response_path") or "")
                                for item in fallback_handoffs
                            )
                            if fallback_handoff_runner is not None:
                                fallback_handoff_runner(fallback_handoffs, remaining)
                            else:
                                parent_handoffs = [{**item, "reviewer_owner": "parent"} for item in fallback_handoffs]
                                handoff_runner(parent_handoffs, remaining)
                        except Exception as fallback_exc:
                            exc = RuntimeError(f"Delegated reviewer failed ({exc}); parent reviewer fallback failed ({fallback_exc})")
                        else:
                            continue
                return finish({
                    "status": "blocked",
                    "stage": "hermes_handoff",
                    "findings": [{"category": "hermes", "field": "handoff", "issue": str(exc)}],
                    "client_outputs": [],
                })
            continue
        pending_handoffs = []
        if result.get("status") != "passed":
            return finish(result)
        if remaining_seconds() <= 0:
            continue
        manifest_name = str(result.get("manifest") or "")
        manifest_path = (run_dir / manifest_name).resolve()
        try:
            manifest_path.relative_to(run_dir)
            manifest = _read(manifest_path)
        except (OSError, ValueError, json.JSONDecodeError) as exc:
            return finish({
                "status": "blocked",
                "stage": "desktop_delivery",
                "findings": [{"category": "delivery", "field": "manifest", "issue": str(exc)}],
                "client_outputs": [],
            })
        if set_finding := _desktop_delivery_set_finding(run_dir, manifest):
            return finish({
                "status": "blocked",
                "stage": "desktop_delivery_set",
                "findings": [set_finding],
                "client_outputs": [],
            })
        reply = result.get("desktop_reply") or desktop_attachment_reply(manifest, run_dir=run_dir)
        delivery = confirm_desktop_delivery(
            manifest,
            reply,
            opener,
            deadline=process_deadline_monotonic,
            clock=clock,
        )
        attempt_counters["delivery"] = int(delivery.get("attempts") or 0)
        if remaining_seconds() <= 0:
            current_stage = "desktop_delivery"
            last_result = {
                **result,
                "stage": "desktop_delivery",
                "delivery": delivery,
                "candidate_outputs": _candidate_outputs(
                    run_dir / "revisions" / str(result.get("revision_id") or "")
                ) if result.get("revision_id") else [],
                "client_outputs": [],
            }
            continue
        final = {
            **result,
            "status": "passed" if delivery["confirmed"] else "blocked",
            "stage": "desktop_delivery",
            "delivery": delivery,
            "deadline_at_epoch": deadline_at_epoch,
        }
        prepared_ledger = manifest.get("gate_ledger")
        if isinstance(prepared_ledger, Mapping):
            gate_findings = [
                {
                    "code": "EXACT_BYTE_ATOMIC_DELIVERY_FAILED",
                    "target": str(finding.get("field") or "desktop_delivery"),
                    "evidence_sha256": canonical_evidence_sha256(dict(finding)),
                    "retry_owner": "desktop_transport",
                    "terminal_status": "blocked",
                }
                for finding in delivery.get("findings", [])
                if isinstance(finding, Mapping)
            ]
            final["gate_ledger"] = advance_gate_ledger(
                SCRIPT_DIR.parent,
                prepared_ledger,
                gate_id="exact_byte_atomic_delivery",
                terminal_status="passed" if delivery["confirmed"] else "blocked",
                evidence={
                    "opened": delivery.get("opened", []),
                    "attempts": delivery.get("attempts"),
                    "manifest_outputs": manifest.get("client_outputs", []),
                },
                findings=gate_findings,
            )
        if not delivery["confirmed"]:
            final["client_outputs"] = []
        return finish(final)


def _production_response_is_bound(
    revision_dir: Path,
    handoff: Mapping[str, Any],
    *,
    model_identifier: str,
) -> bool:
    """Authenticate one worker response before reaping its Hermes process."""
    request_path = revision_dir / str(handoff.get("request_path") or "")
    response_path = revision_dir / str(handoff.get("response_path") or "")
    try:
        request = _read(request_path)
        response = _read(response_path)
    except (OSError, ValueError, json.JSONDecodeError):
        return False
    if str((response.get("producer") or {}).get("model_id") or "") != model_identifier:
        return False
    if str(handoff.get("task") or "") in {
        "clinical_content_verification", "rendered_page_visual_verification",
    }:
        return verification_response_is_complete(revision_dir, request_path)
    return all((
        response.get("schema_version") == "hermes-response/v2",
        response.get("request_id") == request.get("request_id"),
        response.get("request_sha256") == request.get("request_sha256"),
        response.get("revision_id") == request.get("revision_id"),
        response.get("task") == request.get("task"),
        response.get("batch_id") == request.get("batch_id"),
    ))


def _production_agent_prompt(
    skill_root: Path,
    revision_dir: Path,
    handoff: Mapping[str, Any],
    configuration: Mapping[str, Any],
) -> str:
    request_path = revision_dir / str(handoff["request_path"])
    response_path = revision_dir / str(handoff["response_path"])
    task = str(handoff.get("task") or "")
    model_identifier = str(configuration["model_identifier"])
    if task == "rendered_page_visual_verification":
        task_rule = (
            "Act as an independent visual verifier. Load and inspect every supplied page PNG "
            "with the vision tool and assess every requested check for every page."
        )
    elif task == "clinical_content_verification":
        task_rule = (
            "Act as an independent clinical-content verifier. Assess every bound section and "
            "cross-document check directly from the request evidence."
        )
    else:
        task_rule = "Draft only the requested sections from the closed approved evidence package."
    verification_rule = (
        " Write the response exactly once. The Desktop parent validates it automatically; "
        "do not invoke terminal commands or wait for command approval."
        if task in {"clinical_content_verification", "rendered_page_visual_verification"}
        else " The next generate invocation is the authoritative response validator."
    )
    return (
        "Complete one isolated clinical-document Hermes handoff.\n"
        f"Certified skill: {skill_root}\nRun revision: {revision_dir}\n"
        f"Request: {request_path}\nResponse: {response_path}\nTask: {task}\n\n"
        f"Read {skill_root / 'SKILL.md'} and load the clinical-document-generation skill. "
        f"Read the request completely. {task_rule} Write exact JSON directly to the response "
        f"path and bind every schema, request ID, request hash, task, target, and evidence "
        f"reference exactly. producer.model_id must be exactly {model_identifier!r}."
        f"{verification_rule} The validator interpreter is dependency-complete; do not search "
        "the filesystem for Python or dependency paths. Do not modify production code or "
        "approved source material."
    )


def _production_subprocess_environment(skill_root: Path) -> dict[str, str]:
    """Bind Hermes process state to the profile containing the installed skill."""
    untrusted_root = skill_root.expanduser().absolute()
    profile_paths = (untrusted_root, untrusted_root.parent, untrusted_root.parent.parent)
    if any(path.is_symlink() for path in profile_paths) or untrusted_root.parent.name != "skills":
        raise ValueError(
            "Production Desktop execution requires a non-symlinked skill and profile under an isolated Hermes skills directory."
        )
    hermes_home = untrusted_root.parent.parent.resolve()
    environment = {
        name: os.environ[name]
        for name in (
            "LANG", "LC_ALL", "SSL_CERT_FILE", "SSL_CERT_DIR",
            "REQUESTS_CA_BUNDLE", "CURL_CA_BUNDLE",
        )
        if os.environ.get(name)
    }
    environment.update({
        "HERMES_HOME": str(hermes_home),
        "HOME": str(hermes_home),
        "TMPDIR": str(hermes_home / ".tmp"),
        "XDG_CACHE_HOME": str(hermes_home / ".cache"),
        "XDG_CONFIG_HOME": str(hermes_home / ".config"),
        "XDG_DATA_HOME": str(hermes_home / ".local/share"),
        "XDG_STATE_HOME": str(hermes_home / ".local/state"),
        "PATH": "/usr/bin:/bin:/usr/sbin:/sbin",
    })
    for name in (".tmp", ".cache", ".config", ".local/share", ".local/state"):
        (hermes_home / name).mkdir(parents=True, exist_ok=True)
    return environment


def _production_sandbox_executable() -> Path:
    sandbox = Path("/usr/bin/sandbox-exec")
    if sandbox.is_symlink() or not sandbox.is_file() or not os.access(sandbox, os.X_OK):
        raise RuntimeError("The governed OS sandbox executable is unavailable.")
    return sandbox


def _managed_hermes_pair() -> tuple[Path, Path]:
    account_home = Path(pwd.getpwuid(os.getuid()).pw_dir)
    launcher = account_home / ".hermes/hermes-agent/venv/bin/hermes"
    interpreter = launcher.parent / "python"
    if (
        launcher.is_symlink()
        or not launcher.is_file()
        or not os.access(launcher, os.X_OK)
        or not interpreter.is_symlink()
        or not interpreter.is_file()
        or not os.access(interpreter, os.X_OK)
    ):
        raise RuntimeError("The governed Hermes virtual-environment launcher is unavailable.")
    return launcher, interpreter


def _managed_hermes_identity() -> dict[str, str]:
    launcher, interpreter = _managed_hermes_pair()
    return {
        "launcher": str(launcher),
        "launcher_sha256": sha256_file(launcher),
        "interpreter": str(interpreter),
        "interpreter_target": str(interpreter.resolve(strict=True)),
        "interpreter_target_sha256": sha256_file(interpreter.resolve(strict=True)),
    }


def _production_read_boundaries() -> tuple[Path, ...]:
    """Return stable host trees denied before exact governed read exceptions."""
    return tuple(dict.fromkeys(
        path.resolve(strict=False)
        for path in (Path("/Users"), Path("/private/tmp"), Path(tempfile.gettempdir()), Path("/Volumes"))
    ))


PRODUCTION_HERMES_NETWORK_HOST = "chatgpt.com"
PRODUCTION_HERMES_NETWORK_PORT = 443


class _ProductionConnectProxyHandler(socketserver.BaseRequestHandler):
    def handle(self) -> None:
        client = self.request
        client.settimeout(10.0)
        request = b""
        while b"\r\n\r\n" not in request and len(request) <= 16 * 1024:
            chunk = client.recv(4096)
            if not chunk:
                return
            request += chunk
        if b"\r\n\r\n" not in request or len(request) > 16 * 1024:
            client.sendall(b"HTTP/1.1 431 Request Header Fields Too Large\r\n\r\n")
            return
        header, pending = request.split(b"\r\n\r\n", 1)
        try:
            method, target, version = header.split(b"\r\n", 1)[0].decode("ascii").split()
        except (UnicodeDecodeError, ValueError):
            client.sendall(b"HTTP/1.1 400 Bad Request\r\n\r\n")
            return
        server = self.server
        if not isinstance(server, _ProductionConnectProxyServer):
            return
        expected = f"{server.governed_host}:{server.governed_port}"
        if method != "CONNECT" or target.casefold() != expected.casefold() or version not in {"HTTP/1.0", "HTTP/1.1"}:
            client.sendall(b"HTTP/1.1 403 Forbidden\r\n\r\n")
            return
        try:
            upstream = socket.create_connection(
                (server.governed_host, server.governed_port), timeout=10.0,
            )
        except OSError:
            client.sendall(b"HTTP/1.1 502 Bad Gateway\r\n\r\n")
            return
        with upstream:
            client.sendall(b"HTTP/1.1 200 Connection Established\r\n\r\n")
            if pending:
                upstream.sendall(pending)
            client.settimeout(None)
            upstream.settimeout(None)
            peers = (client, upstream)
            while True:
                readable, _, _ = select.select(peers, (), (), 1.0)
                for source in readable:
                    payload = source.recv(64 * 1024)
                    if not payload:
                        return
                    (upstream if source is client else client).sendall(payload)


class _ProductionConnectProxyServer(socketserver.ThreadingTCPServer):
    allow_reuse_address = False
    daemon_threads = True

    def __init__(self, host: str, port: int):
        self.governed_host = host
        self.governed_port = port
        super().__init__(("127.0.0.1", 0), _ProductionConnectProxyHandler)


@dataclass
class _ProductionConnectProxy:
    server: _ProductionConnectProxyServer
    thread: threading.Thread

    @property
    def port(self) -> int:
        return int(self.server.server_address[1])

    def close(self) -> None:
        self.server.shutdown()
        self.server.server_close()
        self.thread.join(timeout=5.0)


def _start_production_connect_proxy(
    host: str = PRODUCTION_HERMES_NETWORK_HOST,
    port: int = PRODUCTION_HERMES_NETWORK_PORT,
) -> _ProductionConnectProxy:
    server = _ProductionConnectProxyServer(host, port)
    thread = threading.Thread(target=server.serve_forever, daemon=True)
    thread.start()
    return _ProductionConnectProxy(server=server, thread=thread)


def _release_production_worker_resources(
    stdout_handle: Any | None,
    stderr_handle: Any | None,
    profile: Path | None,
    proxy: _ProductionConnectProxy | None,
) -> list[BaseException]:
    errors: list[BaseException] = []
    for resource in (stdout_handle, stderr_handle):
        if resource is None:
            continue
        try:
            resource.close()
        except BaseException as exc:
            errors.append(exc)
    if profile is not None:
        try:
            profile.unlink(missing_ok=True)
        except BaseException as exc:
            errors.append(exc)
    if proxy is not None:
        try:
            proxy.close()
        except BaseException as exc:
            errors.append(exc)
    return errors


def _reap_production_worker(
    process: subprocess.Popen[str],
    stdout_handle: Any,
    stderr_handle: Any,
    profile: Path,
    proxy: _ProductionConnectProxy,
) -> None:
    try:
        if process.poll() is None:
            try:
                os.killpg(process.pid, signal.SIGTERM)
                process.wait(timeout=5)
            except (ProcessLookupError, subprocess.TimeoutExpired):
                if process.poll() is None:
                    try:
                        os.killpg(process.pid, signal.SIGKILL)
                    except ProcessLookupError:
                        pass
                    process.wait()
    finally:
        errors = _release_production_worker_resources(
            stdout_handle, stderr_handle, profile, proxy,
        )
    if errors:
        raise RuntimeError("Production worker resources could not be fully released.") from errors[0]


def _production_dispatch_handoffs(
    handoffs: Sequence[Mapping[str, Any]],
    remaining_seconds: float,
    revision_dir: Path,
    configuration: Mapping[str, Any],
    *,
    skill_root: Path,
    run_dir: Path,
    runtime_identity: Mapping[str, Any],
    expected_managed_hermes_identity: Mapping[str, str] | None = None,
) -> None:
    """Run one concurrent Hermes wave under a read-only candidate boundary."""
    processes: list[tuple[subprocess.Popen[str], Mapping[str, Any], Any, Any, Path, float, _ProductionConnectProxy]] = []
    deadline = time.monotonic() + max(0.0, remaining_seconds - 5.0)
    logs = run_dir / "logs/hermes-agents"
    logs.mkdir(parents=True, exist_ok=True)
    sandbox = _production_sandbox_executable()
    environment = _production_subprocess_environment(skill_root)
    hermes_home = Path(environment["HERMES_HOME"])
    hermes_launcher, managed_python = _managed_hermes_pair()
    if expected_managed_hermes_identity is not None and (
        _managed_hermes_identity() != dict(expected_managed_hermes_identity)
    ):
        raise RuntimeError("The managed Hermes launcher or interpreter changed after operation binding.")
    hermes_install_root = hermes_launcher.parent.parent.parent
    interpreter_link_target = Path(os.readlink(managed_python))
    if not interpreter_link_target.is_absolute():
        interpreter_link_target = managed_python.parent / interpreter_link_target
    interpreter_link_root = interpreter_link_target.absolute().parent.parent
    managed_interpreter_root = managed_python.resolve(strict=True).parent.parent
    runtime_executable = Path(str(runtime_identity["executable"])).absolute()
    runtime_root = runtime_executable.parent.parent

    for handoff in handoffs:
        started = time.monotonic()
        request_id = Path(str(handoff["request_path"])).stem
        profile = None
        proxy = None
        stdout_handle = None
        stderr_handle = None
        try:
            cache_dir = run_dir / ".hermes-cache"
            cache_dir.mkdir(parents=True, exist_ok=True)
            profile = tempfile.NamedTemporaryFile(
                "w", prefix="clinical-production-adapter-", suffix=".sb", delete=False,
            )
            proxy = _start_production_connect_proxy()
            profile.write("(version 1)\n(allow default)\n")
            profile.write("(deny network*)\n")
            profile.write(f"(allow network-outbound (remote tcp \"localhost:{proxy.port}\"))\n")
            readable_roots = (
                Path("/System"), Path("/usr"), Path("/bin"), Path("/sbin"),
                Path("/Library"), Path("/Applications/LibreOffice.app"),
                Path("/private/etc"), Path("/etc"), Path("/dev"), Path("/private/var/db"),
                hermes_install_root, hermes_home, skill_root, run_dir, runtime_root,
            )
            read_boundaries = _production_read_boundaries()
            for boundary in read_boundaries:
                profile.write(f"(deny file-read* (subpath {json.dumps(str(boundary))}))\n")
            for readable_root in dict.fromkeys(path.resolve(strict=False) for path in readable_roots):
                if any(readable_root == boundary or boundary in readable_root.parents for boundary in read_boundaries):
                    profile.write(f"(allow file-read* (subpath {json.dumps(str(readable_root))}))\n")
            profile.write(f"(allow file-read* (literal {json.dumps(str(interpreter_link_root))}))\n")
            profile.write(f"(allow file-read* (subpath {json.dumps(str(managed_interpreter_root))}))\n")
            profile.write(
                "(deny file-write* (require-not (require-any "
                f"(subpath {json.dumps(str(hermes_home.resolve()))}) "
                f"(subpath {json.dumps(str(run_dir.resolve()))}) "
                "(literal \"/dev/null\"))))\n"
            )
            profile.write(f"(deny file-write* (subpath {json.dumps(str(skill_root.resolve()))}))\n")
            for executable in ("pytest", "py.test", "pip", "pip3"):
                profile.write(f"(deny process-exec (literal {json.dumps(executable)}))\n")
                resolved = shutil.which(executable)
                if resolved:
                    profile.write(f"(deny process-exec (literal {json.dumps(resolved)}))\n")
            profile.close()
            command = [
                str(managed_python), str(hermes_launcher), "chat", "-q",
                _production_agent_prompt(skill_root, revision_dir, handoff, configuration),
                "--source", str(configuration["source"]),
                "--max-turns", str(configuration["max_turns"]),
                "--skills", str(configuration["skill"]),
            ]
            if configuration.get("safe_mode") is True:
                command.append("--safe-mode")
            stdout_handle = (logs / f"{request_id}.stdout.log").open("w", encoding="utf-8")
            stderr_handle = (logs / f"{request_id}.stderr.log").open("w", encoding="utf-8")
            worker_environment = dict(environment)
            worker_environment.update({
                "HTTPS_PROXY": f"http://127.0.0.1:{proxy.port}",
                "HTTP_PROXY": f"http://127.0.0.1:{proxy.port}",
                "ALL_PROXY": f"http://127.0.0.1:{proxy.port}",
                "NO_PROXY": "",
            })
            process = subprocess.Popen(
                [str(sandbox), "-f", profile.name, *command],
                cwd=skill_root,
                env=worker_environment,
                stdout=stdout_handle,
                stderr=stderr_handle,
                text=True,
                start_new_session=True,
            )
        except BaseException:
            if profile is not None:
                if not profile.closed:
                    try:
                        profile.close()
                    except BaseException:
                        pass
            _release_production_worker_resources(
                stdout_handle,
                stderr_handle,
                Path(profile.name) if profile is not None else None,
                proxy,
            )
            for prior_process, _, prior_stdout, prior_stderr, prior_profile, _, prior_proxy in processes:
                _reap_production_worker(
                    prior_process, prior_stdout, prior_stderr, prior_profile, prior_proxy,
                )
            raise
        assert profile is not None and proxy is not None
        assert stdout_handle is not None and stderr_handle is not None
        processes.append((process, handoff, stdout_handle, stderr_handle, Path(profile.name), started, proxy))
    try:
        pending = list(processes)
        while pending and time.monotonic() < deadline:
            for row in list(pending):
                process, handoff, _, _, _, _, _ = row
                if _production_response_is_bound(
                    revision_dir, handoff,
                    model_identifier=str(configuration["model_identifier"]),
                ):
                    if process.poll() is None:
                        try:
                            os.killpg(process.pid, signal.SIGTERM)
                        except ProcessLookupError:
                            pass
                    process.wait(timeout=5)
                    pending.remove(row)
                elif process.poll() is not None:
                    pending.remove(row)
            if pending:
                time.sleep(0.05)
        missing = [
            str(handoff.get("response_path") or "")
            for process, handoff, _, _, _, _, _ in processes
            if not _production_response_is_bound(
                revision_dir, handoff,
                model_identifier=str(configuration["model_identifier"]),
            )
        ]
        if missing:
            raise RuntimeError(
                "Hermes workers did not produce complete bound responses: " + ", ".join(missing)
            )
    finally:
        for process, handoff, stdout_handle, stderr_handle, profile, started, proxy in processes:
            _reap_production_worker(process, stdout_handle, stderr_handle, profile, proxy)
            ended = time.monotonic()
            event_path = run_dir / "logs/hermes-agent-events.jsonl"
            event_path.parent.mkdir(parents=True, exist_ok=True)
            response_path = revision_dir / str(handoff.get("response_path") or "")
            with event_path.open("a", encoding="utf-8") as handle:
                handle.write(json.dumps({
                    "request_path": str(handoff.get("request_path") or ""),
                    "response_path": str(handoff.get("response_path") or ""),
                    "task": handoff.get("task"),
                    "batch_id": handoff.get("batch_id"),
                    "started_monotonic": started,
                    "ended_monotonic": ended,
                    "elapsed_seconds": round(ended - started, 3),
                    "returncode": process.returncode,
                    "response_exists": response_path.is_file(),
                    "stdout_log": f"logs/hermes-agents/{Path(str(handoff['request_path'])).stem}.stdout.log",
                    "stderr_log": f"logs/hermes-agents/{Path(str(handoff['request_path'])).stem}.stderr.log",
                }, ensure_ascii=False) + "\n")


def _installed_release_identity(skill_root: Path) -> dict[str, Any]:
    untrusted_root = skill_root.expanduser().absolute()
    if any(path.is_symlink() for path in (
        untrusted_root, untrusted_root.parent, untrusted_root.parent.parent,
    )):
        raise ValueError("The installed release root and profile must not contain symlinks.")
    skill_root = untrusted_root.resolve()
    findings = _manifest_integrity(skill_root, allow_runtime_state=True)
    if findings:
        raise ValueError(
            "Production Desktop adapter requires an intact packaged release: "
            + "; ".join(str(item.get("issue") or "") for item in findings)
        )
    manifest = _read(skill_root / RELEASE_MANIFEST)
    return {
        "package_fingerprint": manifest["package_fingerprint"],
        "git_commit": manifest["git_commit"],
        "source": "shipped_production_adapter",
    }


def _certification_preflight_authorizes_candidate(
    skill_root: Path,
    preflight_path: Path,
    identity: Mapping[str, Any],
) -> bool:
    """Require lifecycle-owned preflight evidence outside the candidate tree."""
    supplied = preflight_path.expanduser().absolute()
    if supplied != supplied.resolve() or supplied.is_symlink() or not supplied.is_file():
        return False
    root = skill_root.resolve()
    try:
        supplied.relative_to(root)
        return False
    except ValueError:
        pass
    try:
        evidence = _read(supplied)
        trusted_key = _read(root / RELEASE_CERTIFICATION_PUBLIC_KEY)
    except (OSError, ValueError, json.JSONDecodeError, TypeError):
        return False
    candidate = evidence.get("candidate") or {}
    producer = evidence.get("producer") or {}
    checks = evidence.get("checks") or {}
    attestation = evidence.get("evidence_attestation") or {}
    try:
        modulus = int(str(trusted_key.get("modulus") or ""), 16)
        exponent = int(str(trusted_key.get("exponent") or ""))
        signature = base64.b64decode(
            str(attestation.get("signature_base64") or ""), validate=True,
        )
        unsigned = dict(evidence)
        unsigned.pop("evidence_attestation", None)
        payload_digest = hashlib.sha256(release_certification_payload(unsigned)).digest()
        digest_info = bytes.fromhex("3031300d060960864801650304020105000420") + payload_digest
        encoded_bytes = (modulus.bit_length() + 7) // 8
        expected = b"\x00\x01" + b"\xff" * (encoded_bytes - len(digest_info) - 3) + b"\x00" + digest_info
        verified_signature = pow(
            int.from_bytes(signature, "big"), exponent, modulus,
        ).to_bytes(encoded_bytes, "big")
    except (ValueError, TypeError, binascii.Error, OverflowError):
        return False
    return bool(
        evidence.get("schema_version") == "release-certification-preflight/v1"
        and evidence.get("status") == "passed"
        and evidence.get("repository_clean") is True
        and candidate.get("release_root") == str(root)
        and all(
            candidate.get(field) == identity.get(field)
            for field in ("package_fingerprint", "git_commit")
        )
        and producer.get("path") == "tests/hermes_e2e.py"
        and producer.get("git_commit") == identity.get("git_commit")
        and attestation.get("schema_version") == "release-certification-attestation/v1"
        and attestation.get("algorithm") == RELEASE_CERTIFICATION_SIGNATURE_ALGORITHM
        and attestation.get("key_id") == RELEASE_CERTIFICATION_TRUSTED_KEY_ID
        and attestation.get("key_id") == release_certification_key_id(trusted_key)
        and attestation.get("payload_sha256") == payload_digest.hex()
        and len(signature) == encoded_bytes
        and verified_signature == expected
        and set(checks) == {
            "static_release_checks", "layout_preservation_corpus",
            "deterministic_branch_acceptance_corpus", "repository_regression_suite",
        }
        and all(
            isinstance(item, Mapping)
            and item.get("status") == "passed"
            and item.get("returncode") == 0
            for item in checks.values()
        )
    )


def run_production_desktop_operation(
    run_dir: Path,
    *,
    dispatch_handoffs: Callable[[Sequence[Mapping[str, Any]], float, Path, Mapping[str, Any]], None] | None = None,
    opener: Callable[[str], Any] | None = None,
    parent_visual_reviewer: Callable[[Sequence[Mapping[str, Any]], float, Path, Mapping[str, Any]], None] | None = None,
    release_identity: Mapping[str, Any] | None = None,
    hermes_configuration: Mapping[str, Any] = CERTIFIED_HERMES_CONFIGURATION,
    operation_id: str = "default",
    skill_root: Path | None = None,
    certification_preflight: Path | None = None,
) -> dict[str, Any]:
    """Shipped host adapter for normal and certification Desktop execution."""
    run_dir = run_dir.expanduser().resolve()
    untrusted_root = (skill_root or SCRIPT_DIR.parent).expanduser().absolute()
    if any(path.is_symlink() for path in (
        untrusted_root, untrusted_root.parent, untrusted_root.parent.parent,
    )):
        raise ValueError("The installed release root and profile must not contain symlinks.")
    root = untrusted_root.resolve()
    installed_identity = _installed_release_identity(root)
    if release_identity is not None and any(
        release_identity.get(field) != installed_identity.get(field)
        for field in ("package_fingerprint", "git_commit")
    ):
        raise ValueError("Supplied release identity does not match the installed candidate manifest.")
    identity = installed_identity
    require_promoted_runtime = True
    if certification_preflight is not None:
        if not _certification_preflight_authorizes_candidate(
            root, certification_preflight, identity,
        ):
            raise ValueError(
                "Certification execution requires passing lifecycle-owned preflight evidence for the exact candidate root."
            )
        candidate_runtime = _pdfium_runtime_integrity(
            root, require_promoted_runtime=False,
        )
        if candidate_runtime.get("status") != "passed":
            raise ValueError(
                "Certification execution requires a verified provisioned certification candidate."
            )
        require_promoted_runtime = False
    configuration = dict(hermes_configuration)
    allowed_configuration_fields = set(CERTIFIED_HERMES_CONFIGURATION) | {
        "layout_preservation_notes",
    }
    if set(configuration) not in (
        set(CERTIFIED_HERMES_CONFIGURATION), allowed_configuration_fields,
    ):
        raise ValueError("Production Desktop execution requires the exact governed Hermes configuration.")
    if "layout_preservation_notes" in configuration and (
        not isinstance(configuration["layout_preservation_notes"], list)
        or any(
            not isinstance(note, str) or not note.strip()
            for note in configuration["layout_preservation_notes"]
        )
    ):
        raise ValueError("Production Desktop execution requires governed layout-preservation notes.")
    governed_configuration = {
        field: configuration.get(field)
        for field in CERTIFIED_HERMES_CONFIGURATION
    }
    if governed_configuration != CERTIFIED_HERMES_CONFIGURATION:
        raise ValueError("Production Desktop execution requires the exact governed Hermes configuration.")
    if opener is None:
        raise ValueError("Production Desktop execution requires the actual Desktop opener.")
    runtime_identity = resolve_python_runtime(environment=os.environ)
    managed_hermes_identity = _managed_hermes_identity()
    identity = {**identity, "managed_hermes_identity": managed_hermes_identity}

    def revision_dir() -> Path:
        reference = _read(run_dir / REFERENCE)
        revision_id = str((reference.get("approval") or {}).get("revision_id") or "")
        if not revision_id:
            raise RuntimeError("The approved Run Revision identity is missing.")
        return run_dir / "revisions" / revision_id

    def route(handoffs: list[Mapping[str, Any]], remaining_seconds: float) -> None:
        selected = dispatch_handoffs
        if selected is None:
            _production_dispatch_handoffs(
                handoffs, remaining_seconds, revision_dir(), configuration,
                skill_root=root, run_dir=run_dir, runtime_identity=runtime_identity,
                expected_managed_hermes_identity=managed_hermes_identity,
            )
        else:
            selected(handoffs, remaining_seconds, revision_dir(), configuration)

    def fallback(handoffs: list[Mapping[str, Any]], remaining_seconds: float) -> None:
        active_revision = revision_dir()
        if parent_visual_reviewer is None:
            raise RuntimeError(
                "Delegated Visual QA failed; a genuine Desktop-parent reviewer callback is required."
            )
        parent_visual_reviewer(
            handoffs, remaining_seconds, active_revision, configuration,
        )
        marker = run_dir / "logs/desktop-parent-visual-review.json"
        _write(marker, {
            "status": "completed",
            "revision_id": active_revision.name,
            "request_paths": [str(item.get("request_path") or "") for item in handoffs],
            "response_paths": [str(item.get("response_path") or "") for item in handoffs],
            "completion_requirement": "Desktop parent must inspect every bound page image.",
            "required_producer_model_id": str(configuration["model_identifier"]),
        })

    return run_desktop_operation(
        run_dir,
        handoff_runner=route,
        fallback_handoff_runner=fallback,
        opener=opener,
        operation_id=operation_id,
        runtime_identity=runtime_identity,
        release_identity={**identity, "hermes_configuration": configuration},
        cleanup=lambda _status, _remaining: {
            "owned_processes_reaped": True,
            "late_responses_ignored": True,
        },
        require_promoted_runtime=require_promoted_runtime,
    )


def command_desktop_opener(command_path: Path) -> Callable[[str], bytes]:
    """Build an external host opener that emits exact retrieved bytes."""
    expanded = command_path.expanduser()
    if not expanded.is_absolute():
        raise ValueError(
            "The Desktop opener command must be an absolute, non-symlinked executable file."
        )
    command = expanded.absolute()
    if command.is_symlink() or not command.is_file():
        raise ValueError(
            "The Desktop opener command must be an absolute, non-symlinked executable file."
        )
    if not os.access(command, os.X_OK):
        raise ValueError("The Desktop opener command is not executable.")

    def open_attachment(path: str) -> bytes:
        return subprocess.run(
            [str(command), path],
            check=True,
            stdout=subprocess.PIPE,
            stderr=subprocess.PIPE,
        ).stdout

    return open_attachment


def command_parent_visual_reviewer(
    command_path: Path,
) -> Callable[[Sequence[Mapping[str, Any]], float, Path, Mapping[str, Any]], None]:
    """Build the external Desktop-parent callback used by the public CLI."""
    expanded = command_path.expanduser()
    if not expanded.is_absolute():
        raise ValueError("The Desktop-parent reviewer command must be an absolute executable file.")
    command = expanded.absolute()
    if command.is_symlink() or not command.is_file() or not os.access(command, os.X_OK):
        raise ValueError("The Desktop-parent reviewer command must be an absolute executable file.")

    def review(
        handoffs: Sequence[Mapping[str, Any]],
        remaining_seconds: float,
        revision_dir: Path,
        configuration: Mapping[str, Any],
    ) -> None:
        request_path = revision_dir / "hermes/desktop-parent-visual-review-request.json"
        _write(request_path, {
            "schema_version": "desktop-parent-visual-review-request/v1",
            "revision_id": revision_dir.name,
            "handoffs": [dict(item) for item in handoffs],
            "required_producer_model_id": str(configuration["model_identifier"]),
        })
        subprocess.run(
            [str(command), str(request_path)],
            check=True,
            timeout=max(1.0, remaining_seconds),
            env={"PATH": "/usr/bin:/bin:/usr/sbin:/sbin"},
        )

    return review


def _drafting_evidence(revision_dir: Path) -> list[dict[str, Any]]:
    evidence = []
    for path in sorted((revision_dir / "hermes/accepted").glob("*.json")):
        item = _read(path)
        request_id = str(item.get("request_id") or "")
        accepted_request = revision_dir / "hermes/accepted-requests" / f"{request_id}.json"
        evidence.append({
            "path": path.relative_to(revision_dir).as_posix(),
            "sha256": sha256_file(path),
            "request_id": request_id,
            "request_sha256": item.get("request_sha256"),
            "accepted_request_path": accepted_request.relative_to(revision_dir).as_posix() if accepted_request.is_file() else None,
            "accepted_request_file_sha256": sha256_file(accepted_request) if accepted_request.is_file() else None,
            "producer": item.get("producer"),
        })
    return evidence


def _publication_evidence_findings(
    revision_dir: Path,
    build: Mapping[str, Any],
) -> list[dict[str, Any]]:
    """Rehash every candidate, PDF, and page bound by the accepted build."""
    expected: dict[str, str] = {}
    candidate_rows = list(build.get("candidate_files") or [])
    for item in candidate_rows:
        if isinstance(item, Mapping):
            expected[str(item.get("path") or "")] = str(item.get("sha256") or "")
    actual_candidate_paths = {
        path.relative_to(revision_dir).as_posix()
        for path in (revision_dir / "candidate").glob("*") if path.is_file()
    }
    declared_candidate_paths = {
        str(item.get("path") or "") for item in candidate_rows if isinstance(item, Mapping)
    }
    render_report = dict(build.get("render_report") or {})
    render_rows = list(render_report.get("artifacts") or [])
    expected_render_artifacts = {
        PurePosixPath(path).stem
        for path in actual_candidate_paths
        if PurePosixPath(path).suffix.casefold() == ".docx"
    }
    declared_render_artifacts = [
        str(artifact.get("artifact") or "")
        for artifact in render_rows if isinstance(artifact, Mapping)
    ]
    inventory_valid = (
        declared_candidate_paths == actual_candidate_paths
        and len(candidate_rows) == len(declared_candidate_paths)
        and all(isinstance(item, Mapping) for item in candidate_rows)
        and len(declared_render_artifacts) == len(set(declared_render_artifacts))
        and set(declared_render_artifacts) == expected_render_artifacts
        and render_report.get("status") == "passed"
    )
    all_declared_page_paths: list[str] = []
    for artifact in render_rows:
        if not isinstance(artifact, Mapping):
            inventory_valid = False
            continue
        pages = list(artifact.get("pages") or [])
        page_count = artifact.get("page_count")
        declared_page_paths = {
            str(page.get("path") or "") for page in pages if isinstance(page, Mapping)
        }
        declared_page_path_rows = [
            str(page.get("path") or "") for page in pages if isinstance(page, Mapping)
        ]
        all_declared_page_paths.extend(declared_page_path_rows)
        artifact_name = str(artifact.get("artifact") or "")
        actual_page_paths = {
            path.relative_to(revision_dir).as_posix()
            for path in (revision_dir / "rendered" / artifact_name).glob("*.png")
            if path.is_file()
        }
        if (
            artifact.get("status") != "passed"
            or not isinstance(page_count, int)
            or page_count <= 0
            or [page.get("page") for page in pages if isinstance(page, Mapping)] != list(range(1, page_count + 1))
            or len(pages) != page_count
            or len(declared_page_path_rows) != len(declared_page_paths)
            or declared_page_paths != actual_page_paths
        ):
            inventory_valid = False
        for kind in ("docx", "pdf"):
            expected[str(artifact.get(kind) or "")] = str(artifact.get(f"{kind}_sha256") or "")
        for page in artifact.get("pages", []):
            if isinstance(page, Mapping):
                expected[str(page.get("path") or "")] = str(page.get("sha256") or "")
    if len(all_declared_page_paths) != len(set(all_declared_page_paths)):
        inventory_valid = False
    findings = []
    if not inventory_valid:
        findings.append({
            "code": "INCOMPLETE_PUBLICATION_EVIDENCE_INVENTORY",
            "target": "candidate-build.json",
            "evidence_sha256": canonical_evidence_sha256({
                "candidate_files": candidate_rows,
                "render_report": render_report,
            }),
            "actual_candidate_paths": sorted(actual_candidate_paths),
            "declared_candidate_paths": sorted(declared_candidate_paths),
            "expected_render_artifacts": sorted(expected_render_artifacts),
            "declared_render_artifacts": declared_render_artifacts,
            "retry_owner": "render_assurance",
            "terminal_status": "blocked",
        })
    for relative_text, expected_sha256 in sorted(expected.items()):
        relative = PurePosixPath(relative_text)
        valid_path = (
            bool(relative_text)
            and not relative.is_absolute()
            and ".." not in relative.parts
            and relative.as_posix() == relative_text
        )
        target = revision_dir / relative if valid_path else revision_dir
        actual_sha256 = quality_sha256(target) if valid_path and target.is_file() else None
        if not re.fullmatch(r"[0-9a-f]{64}", expected_sha256) or actual_sha256 != expected_sha256:
            findings.append({
                "code": "STALE_PUBLICATION_EVIDENCE",
                "target": relative_text or "candidate-build.json",
                "evidence_sha256": expected_sha256,
                "actual_sha256": actual_sha256,
                "retry_owner": "render_assurance",
                "terminal_status": "blocked",
            })
    if not expected:
        findings.append({
            "code": "MISSING_PUBLICATION_EVIDENCE",
            "target": "candidate-build.json",
            "evidence_sha256": quality_sha256(revision_dir / "candidate-build.json") if (revision_dir / "candidate-build.json").is_file() else "",
            "actual_sha256": None,
            "retry_owner": "render_assurance",
            "terminal_status": "blocked",
        })
    return findings


def _prepared_gate_ledger(
    revision_dir: Path,
    build: Mapping[str, Any],
    quality: Mapping[str, Any],
    published: Iterable[Mapping[str, Any]],
    expected_attempts: Iterable[Mapping[str, Any]] = (),
) -> dict[str, Any]:
    _validate_expected_gate_attempts(revision_dir, expected_attempts)
    verification = dict(quality.get("verification_evidence") or {})
    content_evidence = {
        key: value for key, value in verification.items()
        if "content" in str(key) or "clinical" in str(key)
    }
    visual_evidence = {
        key: value for key, value in verification.items()
        if "visual" in str(key) or "page" in str(key)
    }
    candidate_files = list(build.get("candidate_files") or [])
    evidence = {
        "clinical_fidelity": {"gate": "clinical_fidelity", "verification": content_evidence or verification},
        "content_completeness_consistency": {
            "gate": "content_completeness_consistency",
            "verification": content_evidence or verification,
            "candidate_files": candidate_files,
        },
        "docx_prs_structure": {
            "gate": "docx_prs_structure",
            "document_report": build.get("document_report"),
            "xml_report": build.get("xml_report"),
            "candidate_files": candidate_files,
        },
        "exact_artifact_rendering": {
            "gate": "exact_artifact_rendering",
            "render_report": build.get("render_report"),
            "render_assurance": quality.get("render_assurance"),
        },
        "every_page_visual_qa": {
            "gate": "every_page_visual_qa",
            "verification": visual_evidence or verification,
            "render_report": build.get("render_report"),
        },
        "exact_byte_atomic_delivery": {
            "gate": "exact_byte_atomic_delivery",
            "prepared_outputs": [dict(item) for item in published],
        },
    }
    return build_gate_ledger(
        SCRIPT_DIR.parent,
        evidence,
        attempt_id=revision_dir.name,
        predecessors=_retained_gate_predecessors(revision_dir),
        statuses={"exact_byte_atomic_delivery": "pending"},
    )


def _publish(
    run_dir: Path,
    revision_dir: Path,
    reference: Mapping[str, Any],
    quality: Mapping[str, Any],
    *,
    expected_attempts: Iterable[Mapping[str, Any]] = (),
    operation_deadline: float | None = None,
    clock: Callable[[], float] = time.monotonic,
) -> dict[str, Any]:
    _validate_expected_gate_attempts(revision_dir, expected_attempts)
    _retained_gate_predecessors(revision_dir)
    build_path = revision_dir / "candidate-build.json"
    try:
        build = _read(build_path)
    except (OSError, ValueError, json.JSONDecodeError) as exc:
        raise RuntimeError(f"stale publication evidence: {exc}") from exc
    if evidence_findings := _publication_evidence_findings(revision_dir, build):
        raise RuntimeError(f"stale publication evidence: {evidence_findings}")
    sources = sorted((revision_dir / "candidate").glob("*.docx")) + sorted((revision_dir / "candidate").glob("*.xml"))
    expected = set(document_set(get_path(reference, "meta.study_type")))
    actual = {source.name for source in sources}
    if actual != expected:
        raise RuntimeError(f"Atomic publication requires the exact branch document set; expected={sorted(expected)}, actual={sorted(actual)}")
    output = run_dir / "output"
    staging = Path(tempfile.mkdtemp(prefix=".output-staging-", dir=run_dir))
    backup: Path | None = None
    try:
        if operation_deadline is not None and clock() >= operation_deadline:
            raise OperationDeadlineExpired("The Desktop operation expired before publication.")
        for source in sources:
            shutil.copy2(source, staging / source.name)
            expected = next(
                (item for item in build.get("candidate_files", []) if item.get("path") == source.relative_to(revision_dir).as_posix()),
                None,
            )
            if not isinstance(expected, Mapping) or quality_sha256(staging / source.name) != expected.get("sha256"):
                raise RuntimeError(f"stale publication evidence: staged bytes changed for {source.name}")
        if evidence_findings := _publication_evidence_findings(revision_dir, build):
            raise RuntimeError(f"stale publication evidence: {evidence_findings}")
        # The only client-visible mutation is the atomic swap below. Recheck
        # after staging so slow copies cannot publish after the hard ceiling.
        if operation_deadline is not None and clock() >= operation_deadline:
            raise OperationDeadlineExpired("The Desktop operation expired while staging publication.")
        if output.exists():
            backup = Path(tempfile.mkdtemp(prefix=".output-backup-", dir=run_dir))
            backup.rmdir()
            os.replace(output, backup)
        try:
            os.replace(staging, output)
        except Exception:
            if backup is not None and backup.exists() and not output.exists():
                os.replace(backup, output)
            raise
    except Exception:
        if staging.exists():
            shutil.rmtree(staging)
        raise
    if backup is not None and backup.exists():
        shutil.rmtree(backup)
    published = [
        {"path": (output / source.name).relative_to(run_dir).as_posix(), "sha256": quality_sha256(output / source.name), "bytes": (output / source.name).stat().st_size}
        for source in sources
    ]
    manifest = {"status": "passed", "revision_id": revision_dir.name, "study_type": canonical_study_type(reference.get("meta", {}).get("study_type")), "approved_source_sha256": reference.get("approval", {}).get("source_sha256"), "approved_reference_sha256": sha256_file(revision_dir / "approved-reference.json"), "candidate_build_sha256": sha256_file(build_path) if build_path.is_file() else None, "contracted_template_bundle": build.get("contracted_template_bundle", {}), "governing_resources": build.get("governing_resources", {}), "drafting_evidence": _drafting_evidence(revision_dir), "quality": quality, "client_outputs": published, "gate_ledger": _prepared_gate_ledger(revision_dir, build, quality, published, expected_attempts)}
    manifest["desktop_reply"] = desktop_attachment_reply(manifest, run_dir=run_dir)
    _write(revision_dir / "delivery-manifest.json", manifest); _write(run_dir / "logs/generation-report.json", manifest)
    return {"status": "passed", "stage": "delivery", "revision_id": revision_dir.name, "contracted_template_bundle": manifest["contracted_template_bundle"], "client_outputs": [item["path"] for item in published], "desktop_reply": manifest["desktop_reply"], "delivery_status": "prepared_unconfirmed", "manifest": (revision_dir / "delivery-manifest.json").relative_to(run_dir).as_posix()}


def _clear_verification_responses(revision_dir: Path, tasks: set[str] | None = None) -> None:
    for request_path in (revision_dir / "hermes/verification-requests").glob("*.json"):
        try:
            request = _read(request_path)
        except (OSError, ValueError, json.JSONDecodeError):
            continue
        if tasks is None or str(request.get("task")) in tasks:
            (revision_dir / str(request.get("response_path", ""))).unlink(missing_ok=True)


def _normalized_layout_repair_records(repairs: Iterable[Mapping[str, Any]]) -> list[dict[str, str]]:
    normalized: dict[tuple[str, str], dict[str, str]] = {}
    for repair in repairs:
        if not isinstance(repair, Mapping):
            continue
        rule = str(repair.get("rule") or "").strip()
        target = " ".join(str(repair.get("target") or "").split())
        if rule and target:
            normalized[(rule, target)] = {"rule": rule, "target": target}
    return [normalized[key] for key in sorted(normalized)]


def _layout_repair_plan(findings: Iterable[Mapping[str, Any]]) -> tuple[dict[str, list[dict[str, str]]], list[dict[str, Any]]]:
    plan: dict[str, list[dict[str, str]]] = {}
    unsupported: list[dict[str, Any]] = []
    for raw in findings:
        finding = dict(raw)
        parsed_targets = [RetryTarget.parse(str(target)) for target in finding.get("target_ids", [])]
        layout_targets = [target.value for target in parsed_targets if target.category == "layout"]
        if not layout_targets:
            continue
        artifact = str(finding.get("artifact") or layout_targets[0]).removesuffix(".docx")
        check = str(finding.get("check") or "")
        target = " ".join(str(finding.get("element") or "").split())
        rule = LAYOUT_RULE_BY_VISUAL_CHECK.get(check)
        if rule not in LAYOUT_REPAIR_RULES.get(artifact, ()) or not target:
            unsupported.append({
                **finding,
                "required": "Classify the visual defect with one supported artifact, Layout Contract check, and exact heading or table-caption element before deterministic repair.",
            })
            continue
        plan.setdefault(artifact, []).append({"rule": rule, "target": target})
    plan = {
        artifact: _normalized_layout_repair_records(repairs)
        for artifact, repairs in plan.items()
    }
    return plan, unsupported


def _invalidate_layout_artifact(revision_dir: Path, artifact: str) -> None:
    """Invalidate only one DOCX and its exact rendered-page review evidence."""
    (revision_dir / "candidate" / f"{artifact}.docx").unlink(missing_ok=True)
    (revision_dir / "rendered" / f"{artifact}.pdf").unlink(missing_ok=True)
    shutil.rmtree(revision_dir / "rendered" / artifact, ignore_errors=True)
    for request_path in (revision_dir / "hermes/verification-requests").glob("*.json"):
        try:
            request = _read(request_path)
        except (OSError, ValueError, json.JSONDecodeError):
            continue
        if request.get("task") != "rendered_page_visual_verification":
            continue
        artifacts = request.get("artifacts") if isinstance(request.get("artifacts"), list) else []
        if not any(str(item.get("artifact")) == artifact for item in artifacts if isinstance(item, Mapping)):
            continue
        (revision_dir / str(request.get("response_path", ""))).unlink(missing_ok=True)
        request_path.unlink(missing_ok=True)


def _retained_gate_predecessors(revision_dir: Path) -> list[dict[str, Any]]:
    attempt_root = revision_dir / "attempts"
    attempt_dirs = sorted(path for path in attempt_root.glob("*") if path.is_dir())
    if not attempt_dirs:
        return []
    validated = []
    for attempt_dir in attempt_dirs:
        manifest_path = attempt_dir / "attempt-manifest.json"
        ledger_path = attempt_dir / "gate-ledger.json"
        if not manifest_path.is_file() or not ledger_path.is_file():
            raise ValueError(f"Retained attempt ledger or manifest is missing: {attempt_dir.name}")
        manifest = _read(manifest_path)
        ledger_file_sha256 = sha256_file(ledger_path)
        declared_files = dict(manifest.get("files") or {})
        actual_files = {
            path.relative_to(attempt_dir).as_posix(): sha256_file(path)
            for path in sorted(attempt_dir.rglob("*"))
            if path.is_file() and path.name != "attempt-manifest.json"
        }
        if (
            manifest.get("gate_ledger_sha256") != _read(ledger_path).get("ledger_sha256")
            or dict(manifest.get("files") or {}).get("gate-ledger.json") != ledger_file_sha256
            or declared_files != actual_files
        ):
            raise ValueError(f"Retained attempt ledger identity is invalid: {attempt_dir.name}")
        validated.append(validate_gate_ledger(SCRIPT_DIR.parent, _read(ledger_path)))
    latest = max(validated, key=lambda ledger: len(ledger["predecessors"]))
    ledgers_by_hash = {ledger["ledger_sha256"]: ledger for ledger in validated}
    for ledger in validated:
        for predecessor in ledger["predecessors"]:
            parent = ledgers_by_hash.get(predecessor["ledger_sha256"])
            if parent is None:
                raise ValueError("Retained attempt ledger predecessor is missing.")
            blocked_parent = next(
                record for record in parent["records"]
                if record["terminal_status"] == "blocked"
            )
            expected_predecessor = {
                "attempt_id": parent["attempt_id"],
                "ledger_sha256": parent["ledger_sha256"],
                "terminal_gate": blocked_parent["gate_id"],
                "blocked_findings": [dict(item) for item in blocked_parent["findings"]],
            }
            if predecessor != expected_predecessor:
                raise ValueError("Retained attempt ledger predecessor semantics are invalid.")
    retained_hashes = {ledger["ledger_sha256"] for ledger in validated}
    chained_hashes = {
        *[item["ledger_sha256"] for item in latest["predecessors"]],
        latest["ledger_sha256"],
    }
    if len(retained_hashes) != len(validated) or retained_hashes != chained_hashes:
        raise ValueError("Retained attempt ledger chain is incomplete or forked.")
    blocked = next(
        record for record in latest["records"]
        if record["terminal_status"] == "blocked"
    )
    return [
        *[dict(item) for item in latest["predecessors"]],
        {
            "attempt_id": latest["attempt_id"],
            "ledger_sha256": latest["ledger_sha256"],
            "terminal_gate": blocked["gate_id"],
            "blocked_findings": [dict(item) for item in blocked["findings"]],
        },
    ]


def _validate_expected_gate_attempts(
    revision_dir: Path,
    expected_attempts: Iterable[Mapping[str, Any]],
) -> None:
    expected = [dict(item) for item in expected_attempts]
    journal_path = revision_dir / "gate-attempt-journal.json"
    if journal_path.is_file():
        journal = _read(journal_path)
        unsigned_journal = dict(journal)
        declared_journal_sha256 = str(unsigned_journal.pop("journal_sha256", ""))
        if (
            journal.get("schema_version") != "clinical-gate-attempt-journal/v1"
            or journal.get("attempt_id") != revision_dir.name
            or declared_journal_sha256 != canonical_evidence_sha256(unsigned_journal)
            or journal.get("entries") != expected
        ):
            raise ValueError("Retained gate attempt journal is invalid.")
    elif expected:
        raise ValueError("Retained gate attempt journal is missing.")
    actual_dirs = sorted(path for path in (revision_dir / "attempts").glob("*") if path.is_dir())
    if len(actual_dirs) != len(expected):
        raise ValueError("Retained gate attempt inventory count is invalid.")
    expected_paths = [str(entry.get("path") or "") for entry in expected]
    actual_paths = {path.relative_to(revision_dir).as_posix() for path in actual_dirs}
    if len(expected_paths) != len(set(expected_paths)) or set(expected_paths) != actual_paths:
        raise ValueError("Retained gate attempt inventory paths are duplicated or incomplete.")
    for entry in expected:
        relative_text = str(entry.get("path") or "")
        relative = PurePosixPath(relative_text)
        if relative.is_absolute() or ".." in relative.parts or relative.as_posix() != relative_text:
            raise ValueError("Retained gate attempt path is invalid.")
        attempt_dir = revision_dir / relative
        manifest_path = attempt_dir / "attempt-manifest.json"
        ledger_path = attempt_dir / "gate-ledger.json"
        if (
            not attempt_dir.is_dir()
            or not manifest_path.is_file()
            or not ledger_path.is_file()
            or sha256_file(manifest_path) != entry.get("attempt_manifest_sha256")
            or _read(ledger_path).get("ledger_sha256") != entry.get("gate_ledger_sha256")
        ):
            raise ValueError(f"Retained gate attempt evidence is missing or stale: {relative_text}")


def _failed_gate_for_stage(stage: str, findings: Iterable[Mapping[str, Any]]) -> str:
    if stage == "pre_render_content":
        return "content_completeness_consistency"
    if stage == "rendering":
        return "docx_prs_structure"
    if stage == "rendered_document_qa" or any(
        str(item.get("category") or "").casefold() == "visual" for item in findings
    ):
        return "every_page_visual_qa"
    return "clinical_fidelity"


def _ledger_findings(
    findings: Iterable[Mapping[str, Any]],
    *,
    gate_id: str,
    retry_owner: str,
) -> list[dict[str, Any]]:
    result = []
    for item in findings:
        result.append({
            "code": f"{gate_id.upper()}_FAILED",
            "target": str(item.get("artifact") or item.get("field") or item.get("check") or "attempt"),
            "evidence_sha256": canonical_evidence_sha256(item),
            "retry_owner": retry_owner,
            "terminal_status": "blocked",
        })
    return result or [{
        "code": f"{gate_id.upper()}_FAILED",
        "target": "attempt",
        "evidence_sha256": canonical_evidence_sha256([]),
        "retry_owner": retry_owner,
        "terminal_status": "blocked",
    }]


def _archive_failed_attempt(revision_dir: Path, stage: str, findings: list[Mapping[str, Any]]) -> Path:
    """Preserve the complete failed candidate and QA evidence before any retry mutation."""
    archive_root = revision_dir / "attempts"
    archive_root.mkdir(parents=True, exist_ok=True)
    safe_stage = _slug(stage)
    sequence = 1
    while (archive_root / f"{safe_stage}-a{sequence:02d}").exists():
        sequence += 1
    destination = archive_root / f"{safe_stage}-a{sequence:02d}"
    journal_path = revision_dir / "gate-attempt-journal.json"
    journal_entries = list(_read(journal_path).get("entries") or []) if journal_path.is_file() else []
    _validate_expected_gate_attempts(revision_dir, journal_entries)
    predecessors = _retained_gate_predecessors(revision_dir)
    destination.mkdir()
    retained = (
        Path("candidate"),
        Path("rendered"),
        Path("candidate-build.json"),
        Path("hermes/accepted"),
        Path("hermes/accepted-requests"),
        Path("hermes/verification-requests"),
        Path("hermes/verification-responses"),
    )
    for relative in retained:
        source = revision_dir / relative
        target = destination / relative
        if source.is_dir():
            shutil.copytree(source, target)
        elif source.is_file():
            target.parent.mkdir(parents=True, exist_ok=True)
            shutil.copy2(source, target)
    failed_gate = _failed_gate_for_stage(stage, findings)
    failed_index = GOVERNED_GATE_SEQUENCE.index(failed_gate)
    statuses = {
        gate_id: "passed" if index < failed_index else "blocked" if index == failed_index else "pending"
        for index, gate_id in enumerate(GOVERNED_GATE_SEQUENCE)
    }
    gate_contract = next(
        item for item in load_format_conformance_matrix(SCRIPT_DIR.parent)["gate_sequence"]
        if item["gate_id"] == failed_gate
    )
    ledger_findings = _ledger_findings(
        findings,
        gate_id=failed_gate,
        retry_owner=str(gate_contract["retry_owner"]),
    )
    attempt_evidence = {
        gate_id: {
            "revision_id": revision_dir.name,
            "stage": stage,
            "gate_id": gate_id,
            "findings_sha256": canonical_evidence_sha256(findings),
        }
        for gate_id in GOVERNED_GATE_SEQUENCE
    }
    failed_ledger = build_gate_ledger(
        SCRIPT_DIR.parent,
        attempt_evidence,
        attempt_id=revision_dir.name,
        predecessors=predecessors,
        statuses=statuses,
        findings_by_gate={failed_gate: ledger_findings},
    )
    _write(destination / "gate-ledger.json", failed_ledger)
    files = {
        path.relative_to(destination).as_posix(): sha256_file(path)
        for path in sorted(destination.rglob("*"))
        if path.is_file()
    }
    _write(destination / "attempt-manifest.json", {
        "revision_id": revision_dir.name,
        "stage": stage,
        "attempt": sequence,
        "archived_at": datetime.now(timezone.utc).isoformat(),
        "findings": [dict(item) for item in findings],
        "gate_ledger_sha256": failed_ledger["ledger_sha256"],
        "files": files,
    })
    journal_entries.append({
        "path": destination.relative_to(revision_dir).as_posix(),
        "attempt_manifest_sha256": sha256_file(destination / "attempt-manifest.json"),
        "gate_ledger_sha256": failed_ledger["ledger_sha256"],
    })
    gate_journal = {
        "schema_version": "clinical-gate-attempt-journal/v1",
        "attempt_id": revision_dir.name,
        "entries": journal_entries,
    }
    gate_journal["journal_sha256"] = canonical_evidence_sha256(gate_journal)
    _write(journal_path, gate_journal)
    return destination


def _quality_retry(
    run_dir: Path,
    reference_path: Path,
    working_reference: dict[str, Any],
    approved_reference: Mapping[str, Any],
    revision_dir: Path,
    prior_attempts: Mapping[str, int],
    findings: list[Mapping[str, Any]],
    stage: str,
    *,
    contracted_bundle: Mapping[str, Any] | None = None,
    operation_deadline: float | None = None,
    clock: Callable[[], float] = time.monotonic,
    stage_observer: Callable[[str, float], Any] | None = None,
    require_promoted_runtime: bool = True,
) -> dict[str, Any]:
    """Retry draftable targets; deterministic layout defects require an actual repair."""
    invalid_recovery = [
        {
            "category": "recovery-classification",
            "field": "recovery_class",
            "issue": "Every recoverable finding must carry one governed Recovery Class and its lawful action.",
            "finding": dict(item),
        }
        for item in findings
        if item.get("recovery_class") not in RECOVERY_POLICIES
        or item.get("action") != RECOVERY_POLICIES.get(str(item.get("recovery_class")))
    ]
    if invalid_recovery:
        return _repair_block(
            run_dir,
            "recovery_classification",
            invalid_recovery,
            candidate_outputs=_candidate_outputs(revision_dir),
        )
    structural = [item for item in findings if item.get("recovery_class") == "document_structure_defect"]
    if structural:
        return _repair_block(
            run_dir,
            "document_structure",
            structural,
            candidate_outputs=_candidate_outputs(revision_dir),
        )
    expected_target_category = {
        "visual_defect": "layout",
        "drafting_defect": "section",
        "verifier_transient": "verification",
    }
    explicit_route_errors = [
        {
            "category": "recovery-classification",
            "field": "target_ids",
            "issue": f"Recovery Class {item.get('recovery_class')} requires {expected} targets.",
            "finding": dict(item),
        }
        for item in findings
        if (expected := expected_target_category.get(str(item.get("recovery_class")))) is not None
        and isinstance(item.get("target_ids"), list)
        and item.get("target_ids")
        and any(RetryTarget.parse(str(target)).category != expected for target in item["target_ids"])
    ]
    if explicit_route_errors:
        return _repair_block(
            run_dir,
            "recovery_classification",
            explicit_route_errors,
            candidate_outputs=_candidate_outputs(revision_dir),
        )
    generation_state = working_reference.setdefault("generation", {})
    expected_gate_attempts = list(generation_state.get("gate_attempts") or [])
    _validate_expected_gate_attempts(revision_dir, expected_gate_attempts)
    _archive_failed_attempt(revision_dir, stage, findings)
    gate_journal = _read(revision_dir / "gate-attempt-journal.json")
    generation_state["gate_attempts"] = list(gate_journal["entries"])
    _write(reference_path, working_reference)
    transient = [item for item in findings if item.get("recovery_class") == "verifier_transient"]
    if transient:
        verification_attempts = working_reference.setdefault("generation", {}).setdefault("verification_attempts", {})
        exhausted = []
        tasks = set()
        for item in transient:
            target = str((item.get("target_ids") or [item.get("field")])[0])
            next_attempt = int(verification_attempts.get(target, 0)) + 1
            verification_attempts[target] = next_attempt
            retry_target = RetryTarget.parse(target)
            task = VERIFICATION_TASK_BY_TARGET.get(retry_target.value) if retry_target.category == "verification" else None
            if task:
                tasks.add(task)
            if next_attempt > MAX_VERIFICATION_ATTEMPTS:
                exhausted.append({
                    **dict(item),
                    "field": target,
                    "issue": f"Reviewer retry limit reached after {MAX_VERIFICATION_ATTEMPTS} attempts. {item.get('issue', '')}".strip(),
                })
        _write(reference_path, working_reference)
        if exhausted:
            path = run_dir / "reference/repair-report.md"
            path.write_text(repair_report(exhausted), encoding="utf-8")
            return {"status": "blocked", "stage": "reviewer_retry_limit", "findings": exhausted, "repair_report": path.relative_to(run_dir).as_posix(), "client_outputs": []}
        _clear_verification_responses(revision_dir, tasks)
        remaining = [item for item in findings if item.get("recovery_class") != "verifier_transient"]
        if not remaining:
            return _awaiting(
                revision_dir,
                stage="independent_verification_retry",
                paths=sorted((revision_dir / "hermes/verification-requests").glob("*.json")),
                findings=transient,
            )
        findings = remaining
    known_sections = {section.section_id for section in protocol_contract(str(get_path(approved_reference, "meta.study_type", "")))}
    if canonical_study_type(get_path(approved_reference, "meta.study_type")) != "Retrospective":
        known_sections.update(section.section_id for section in icf_contract(
            str(get_path(approved_reference, "meta.study_type", "")),
            str(get_path(approved_reference, "meta.icf_template", "Advarra")),
        ))
        known_sections.update(("prs.brief-summary", "prs.detailed-description"))
    draftable_sections = {
        section_id
        for batch in batch_plan(
            str(get_path(approved_reference, "meta.study_type", "")),
            str(get_path(approved_reference, "meta.icf_template", "Advarra")),
        )
        for section_id in batch.section_ids
    }
    normalized: list[dict[str, Any]] = []
    for raw in findings:
        finding = dict(raw)
        targets = [str(item) for item in finding.get("target_ids", []) if item] if isinstance(finding.get("target_ids"), list) else []
        if not targets:
            field = str(finding.get("field") or "")
            if field in known_sections:
                targets = [field]
            elif finding.get("recovery_class") == "visual_defect":
                targets = [f"layout:{finding.get('artifact') or 'documents'}"]
            elif finding.get("recovery_class") == "verifier_transient":
                targets = ["verification:content"]
            else:
                targets = ["layout:documents"]
        finding["target_ids"] = targets
        normalized.append(finding)
    targets = {RetryTarget.parse(target) for finding in normalized for target in finding["target_ids"]}
    route_errors = [
        {
            "category": "recovery-classification",
            "field": "target_ids",
            "issue": f"Recovery Class {finding.get('recovery_class')} requires {expected} targets.",
            "finding": finding,
        }
        for finding in normalized
        if (expected := expected_target_category.get(str(finding.get("recovery_class")))) is not None
        and any(RetryTarget.parse(str(target)).category != expected for target in finding["target_ids"])
    ]
    if route_errors:
        return _repair_block(
            run_dir,
            "recovery_classification",
            route_errors,
            candidate_outputs=_candidate_outputs(revision_dir),
        )
    section_targets = {
        target.target_id
        for target in targets
        if target.category == "section" and target.target_id in draftable_sections
    }
    deterministic_targets = {
        target.target_id
        for target in targets
        if target.category == "section"
        and target.target_id in known_sections
        and target.target_id not in draftable_sections
    }
    has_layout_target = any(target.category == "layout" for target in targets)
    if deterministic_targets and not section_targets:
        path = run_dir / "reference/repair-report.md"
        path.write_text(repair_report(normalized), encoding="utf-8")
        return {
            "status": "blocked",
            "stage": "deterministic_repair",
            "findings": normalized,
            "repair_report": path.relative_to(run_dir).as_posix(),
            "client_outputs": [],
        }
    attempts, exhausted = retry_attempts(normalized, prior_attempts)
    working_reference.setdefault("generation", {})["attempts"] = attempts
    _write(reference_path, working_reference)
    if exhausted:
        path = run_dir / "reference/repair-report.md"; path.write_text(repair_report(exhausted), encoding="utf-8")
        return {"status": "blocked", "stage": stage, "findings": exhausted, "repair_report": path.relative_to(run_dir).as_posix(), "client_outputs": []}
    if section_targets:
        invalidate_accepted_targets(revision_dir, section_targets)
        (revision_dir / "candidate-build.json").unlink(missing_ok=True)
        _clear_verification_responses(revision_dir)
        created = schedule_requests(repo_root=SCRIPT_DIR.parent, revision_dir=revision_dir, revision_id=revision_dir.name, reference=approved_reference, attempts=attempts, wave="quality-retry", findings=normalized, contracted_bundle=contracted_bundle)
        if created:
            return _awaiting(revision_dir, stage="drafting_retry", paths=created, findings=normalized)
    if has_layout_target:
        generation = working_reference.setdefault("generation", {})
        repair_plan, unsupported = _layout_repair_plan(normalized)
        if unsupported:
            return _repair_block(
                run_dir,
                "layout_repair_classification",
                unsupported,
                candidate_outputs=_candidate_outputs(revision_dir),
            )
        persisted_repairs = generation.setdefault("layout_repairs", {})
        for artifact, repairs in repair_plan.items():
            existing = persisted_repairs.get(artifact, ())
            persisted_repairs[artifact] = _normalized_layout_repair_records([*existing, *repairs])
            _invalidate_layout_artifact(revision_dir, artifact)
        generation["pending_layout_artifacts"] = sorted(repair_plan)
        _write(reference_path, working_reference)
    else:
        tasks = {
            task
            for target in targets
            if target.category == "verification"
            for task in [VERIFICATION_TASK_BY_TARGET.get(target.value)]
            if task is not None
        }
        _clear_verification_responses(revision_dir, tasks or None)
    retry_options: dict[str, Any] = {
        "operation_deadline": operation_deadline,
        "clock": clock,
    }
    if not require_promoted_runtime:
        retry_options["require_promoted_runtime"] = False
    if stage_observer is not None:
        retry_options["stage_observer"] = stage_observer
    return generate(run_dir, **retry_options)


def generate(
    run_dir: Path,
    *,
    operation_deadline: float | None = None,
    clock: Callable[[], float] = time.monotonic,
    stage_observer: Callable[[str, float], Any] | None = None,
    require_promoted_runtime: bool = True,
    **_: Any,
) -> dict[str, Any]:
    """Advance one approved revision until it needs Hermes work or passes."""
    observed_at = clock()

    def observe_stage(stage: str) -> None:
        nonlocal observed_at
        now = clock()
        if stage_observer is not None:
            stage_observer(stage, max(0.0, now - observed_at))
        observed_at = now

    run_dir = run_dir.resolve(); reference_path, working_reference = _reference(run_dir)
    try:
        bundle = contracted_template_bundle(SCRIPT_DIR.parent, working_reference)
    except ContractedTemplateBundleError as exc:
        return _repair_block(run_dir, "contracted_template_bundle", [exc.finding])
    approved, approval_issue = _approval_valid(run_dir, working_reference, contracted_bundle=bundle)
    revision_id = str(working_reference.get("approval", {}).get("revision_id") or "")
    revision_dir = run_dir / "revisions" / revision_id
    reference = _read(revision_dir / "approved-reference.json") if approved else working_reference
    contract = source_contract(reference)
    if contract["status"] != "passed" or not approved:
        findings = list(contract["blocking_findings"])
        if not approved: findings.append({"category": "approval", "field": "approval", "issue": approval_issue})
        if approved:
            return _repair_block(run_dir, "approval_gate", findings)
        return {"status": "blocked", "stage": "approval_gate", "findings": findings, "client_outputs": []}
    if not revision_id or not revision_dir.is_dir(): return {"status": "blocked", "stage": "revision", "findings": [{"category": "revision", "field": "revision_id", "issue": "Approved immutable revision is missing."}], "client_outputs": []}
    state = working_reference.setdefault("generation", {})
    expected_governing = governing_resources(SCRIPT_DIR.parent, reference, contracted_bundle=bundle)
    governing_sha256 = sha256_value(expected_governing)
    prior_governing_sha256 = state.get("governing_sha256")
    if prior_governing_sha256 is not None and prior_governing_sha256 != governing_sha256:
        return {
            "status": "blocked",
            "stage": "revision",
            "findings": [{
                "category": "revision",
                "field": revision_id,
                "issue": "Generation resources changed inside an immutable revision.",
                "required": "Approve the unchanged Source-of-Truth again to create a new revision.",
            }],
            "client_outputs": [],
        }
    elif prior_governing_sha256 is None:
        state["governing_sha256"] = governing_sha256
        _write(reference_path, working_reference)
    attempts = state.setdefault("attempts", {})
    persisted_exhaustion = [
        {"category": "retry", "field": str(target), "issue": f"Retry limit reached after {MAX_ATTEMPTS} attempts.", "required": "Reviewer intervention before a new approved revision."}
        for target, count in attempts.items() if int(count) > MAX_ATTEMPTS
    ]
    if persisted_exhaustion:
        path = run_dir / "reference/repair-report.md"; path.write_text(repair_report(persisted_exhaustion), encoding="utf-8")
        return {"status": "blocked", "stage": "retry_limit", "findings": persisted_exhaustion, "repair_report": path.relative_to(run_dir).as_posix(), "client_outputs": []}
    drafting_findings = [
        finding if finding.get("category") in {"source-evidence", "request-integrity"}
        else recovery_finding(finding, "drafting_defect")
        for finding in ingest_responses(revision_dir, expected_governing)
    ]
    if drafting_findings:
        attempts, exhausted = retry_attempts(drafting_findings, attempts); state["attempts"] = attempts; _write(reference_path, working_reference)
        source_gaps = [item for item in drafting_findings if item.get("category") in {"source-evidence", "request-integrity"}]
        if source_gaps or exhausted:
            path = run_dir / "reference/repair-report.md"; path.write_text(repair_report([*source_gaps, *exhausted]), encoding="utf-8")
            return {"status": "blocked", "stage": "drafting", "findings": [*source_gaps, *exhausted], "repair_report": path.relative_to(run_dir).as_posix(), "client_outputs": []}
        created = schedule_requests(repo_root=SCRIPT_DIR.parent, revision_dir=revision_dir, revision_id=revision_id, reference=reference, attempts=attempts, wave="retry", findings=drafting_findings, contracted_bundle=bundle)
        if created: return _awaiting(revision_dir, stage="drafting_retry", paths=created, findings=drafting_findings)
    created = schedule_requests(repo_root=SCRIPT_DIR.parent, revision_dir=revision_dir, revision_id=revision_id, reference=reference, attempts=attempts, wave="initial", contracted_bundle=bundle)
    pending = pending_requests(revision_dir, expected_governing)
    if created or pending: return _awaiting(revision_dir, stage="drafting", paths=pending or created)
    missing = missing_drafts(revision_dir, reference, SCRIPT_DIR.parent, contracted_bundle=bundle)
    if missing: return {"status": "blocked", "stage": "drafting", "findings": [{"category": "drafting", "field": item, "issue": "Required section has no accepted draft after all requests were processed."} for item in missing], "client_outputs": []}
    duplicate_findings = accepted_cross_section_duplicate_findings(
        revision_dir,
        reference,
        expected_governing,
    )
    if duplicate_findings:
        return _quality_retry(
            run_dir,
            reference_path,
            working_reference,
            reference,
            revision_dir,
            attempts,
            duplicate_findings,
            "pre_render_content",
            contracted_bundle=bundle,
            operation_deadline=operation_deadline,
            clock=clock,
            stage_observer=stage_observer,
            require_promoted_runtime=require_promoted_runtime,
        )

    observe_stage("drafting")
    model = merged_drafts(revision_dir, reference, expected_governing)
    font_substitutions = {
        str(source): str(target)
        for source, target in dict(state.get("font_substitutions") or {}).items()
    }
    layout_repairs = {
        str(artifact): tuple(_normalized_layout_repair_records(rules))
        for artifact, rules in dict(state.get("layout_repairs") or {}).items()
        if isinstance(rules, list)
    }
    fingerprint, governing = _candidate_fingerprint(
        SCRIPT_DIR.parent,
        revision_dir,
        reference,
        model,
        contracted_bundle=bundle,
        font_substitutions=font_substitutions,
        layout_repairs=layout_repairs,
    )
    try:
        prior_build = _read(revision_dir / "candidate-build.json")
    except (OSError, ValueError, json.JSONDecodeError):
        prior_build = None
    repair_artifacts = {
        str(artifact)
        for artifact in state.get("pending_layout_artifacts", [])
        if str(artifact) in layout_repairs
    }
    partial_repair = (
        repair_artifacts
        if prior_build and repair_artifacts and _unaffected_build_is_valid(revision_dir, prior_build, repair_artifacts)
        else set()
    )
    build = _cached_build(revision_dir, fingerprint)
    structure = _cached_candidate_structure(revision_dir, fingerprint)
    assurance_report: dict[str, Any] | None = None
    if build is None:
        if structure is not None:
            document_report = structure["document_report"]
            xml_report = structure.get("xml_report")
        else:
            document_report = render_documents(
                SCRIPT_DIR.parent,
                revision_dir,
                reference,
                model,
                contracted_bundle=bundle,
                font_substitutions=font_substitutions,
                artifact_names=partial_repair or None,
                layout_repairs=layout_repairs,
            )
            if partial_repair:
                document_report = _merge_artifact_reports(prior_build.get("document_report", {}), document_report)
            if document_report["status"] != "passed":
                findings, classification_block = _document_report_failure(run_dir, revision_dir, document_report)
                if classification_block is not None:
                    return classification_block
                return _quality_retry(run_dir, reference_path, working_reference, reference, revision_dir, attempts, findings, "rendering", contracted_bundle=bundle, operation_deadline=operation_deadline, clock=clock, stage_observer=stage_observer, require_promoted_runtime=require_promoted_runtime)
            xml_report = prior_build.get("xml_report") if partial_repair else None
            if canonical_study_type(reference.get("meta", {}).get("study_type")) != "Retrospective":
                if xml_report is None or not (revision_dir / "candidate/study.xml").is_file():
                    prs_authority = bundle["prs_authority"]
                    template = SCRIPT_DIR.parent / str(prs_authority["generation_template"]["path"])
                    xml_report = generate_xml(
                        template,
                        revision_dir / "candidate/study.xml",
                        reference,
                        model.get("prs", {}),
                        structural_template=SCRIPT_DIR.parent / str(prs_authority["structural_reference"]["path"]),
                    )
                if xml_report["status"] != "passed":
                    findings = [
                        {**finding, "category": "document-structure", "recovery_class": "document_structure_defect", "action": "preserve_and_stop"}
                        for finding in xml_report["findings"]
                    ]
                    return _repair_block(run_dir, "xml", findings, candidate_outputs=_candidate_outputs(revision_dir))
            structure = _record_candidate_structure(revision_dir, fingerprint, governing, bundle, document_report, xml_report, document_set(get_path(reference, "meta.study_type")))

        observe_stage("candidate")
        remaining = (
            DESKTOP_OPERATION_BUDGET_SECONDS
            if operation_deadline is None
            else operation_deadline - clock()
        )
        if remaining <= 0:
            return {
                "status": "timeout",
                "stage": "render_assurance",
                "findings": [{"category": "timeout", "field": "operation", "issue": "The persisted Desktop operation deadline expired before Render Assurance."}],
                "candidate_outputs": _candidate_outputs(revision_dir),
                "client_outputs": [],
            }

        def rebuild_with_substitutions(resolved: Mapping[str, str]) -> dict[str, Any]:
            nonlocal document_report
            document_report = render_documents(
                SCRIPT_DIR.parent,
                revision_dir,
                reference,
                model,
                contracted_bundle=bundle,
                font_substitutions=resolved,
                artifact_names=partial_repair or None,
                layout_repairs=layout_repairs,
            )
            if partial_repair:
                document_report = _merge_artifact_reports(prior_build.get("document_report", {}), document_report)
            return document_report

        assurance_report = render_assurance(
            SCRIPT_DIR.parent,
            revision_dir,
            reference,
            contracted_bundle=bundle,
            structural_validation=structure,
            candidate_font_substitutions=font_substitutions,
            rebuild_candidate=rebuild_with_substitutions,
            artifact_names=partial_repair or None,
            deadline_seconds=remaining,
            deadline_monotonic=operation_deadline,
            clock=clock,
            require_promoted_runtime=require_promoted_runtime,
        )
        resolved_substitutions = {
            str(source): str(target)
            for source, target in dict(assurance_report.get("font_substitutions") or {}).items()
        }
        if resolved_substitutions != font_substitutions:
            font_substitutions = resolved_substitutions
            state["font_substitutions"] = font_substitutions
            fingerprint, governing = _candidate_fingerprint(
                SCRIPT_DIR.parent,
                revision_dir,
                reference,
                model,
                contracted_bundle=bundle,
                font_substitutions=font_substitutions,
                layout_repairs=layout_repairs,
            )
            structure = _record_candidate_structure(revision_dir, fingerprint, governing, bundle, document_report, xml_report, document_set(get_path(reference, "meta.study_type")))
        assurance_report, render_report = _merge_partial_assurance(prior_build or {}, assurance_report, partial_repair)
        state["render_assurance"] = assurance_report
        state.pop("renderer_preflight", None)
        structure = _record_candidate_structure(revision_dir, fingerprint, governing, bundle, document_report, xml_report, document_set(get_path(reference, "meta.study_type")))
        _write(reference_path, working_reference)
        if assurance_report.get("status") != "passed":
            findings = [
                {**finding, "target_ids": finding.get("target_ids") or [f"layout:{finding.get('artifact') or 'documents'}"]}
                for finding in assurance_report.get("findings", [])
            ]
            if findings and all(finding.get("recovery_class") == "adapter_fault" for finding in findings):
                diagnostic_path = run_dir / "reference/render-assurance-diagnostic.md"
                diagnostic_path.write_text(repair_report(findings), encoding="utf-8")
                return {
                    "status": "blocked",
                    "stage": "render_assurance",
                    "findings": findings,
                    "render_assurance": assurance_report,
                    "repair_report": diagnostic_path.relative_to(run_dir).as_posix(),
                    "candidate_outputs": _candidate_outputs(revision_dir),
                    "client_outputs": [],
                }
            return _quality_retry(run_dir, reference_path, working_reference, reference, revision_dir, attempts, findings, "rendered_document_qa", contracted_bundle=bundle, operation_deadline=operation_deadline, clock=clock, stage_observer=stage_observer, require_promoted_runtime=require_promoted_runtime)
        build = _record_build(revision_dir, fingerprint, governing, bundle, document_report, xml_report, render_report)
    else:
        observe_stage("candidate")
        document_report = build["document_report"]
        xml_report = build.get("xml_report")
        render_report = build["render_report"]
        recorded_assurance = state.get("render_assurance")
        assurance_report = dict(recorded_assurance) if isinstance(recorded_assurance, Mapping) else {
            "schema_version": "render-assurance/v1",
            "status": "passed",
            "font_substitutions": font_substitutions,
            "render": render_report,
            "findings": [],
        }
    observe_stage("render_assurance")
    if state.pop("pending_layout_artifacts", None) is not None:
        _write(reference_path, working_reference)
    create_verification_requests(revision_dir, reference, render_report, contracted_bundle=bundle)
    pending_checks = pending_verifications(revision_dir)
    if pending_checks: return _awaiting(revision_dir, stage="independent_verification", paths=pending_checks)
    final_quality = quality_report(revision_dir, reference, render_report, xml_report)
    final_quality["render_assurance"] = assurance_report
    if final_quality["status"] != "passed":
        return _quality_retry(run_dir, reference_path, working_reference, reference, revision_dir, attempts, final_quality["findings"], "quality", contracted_bundle=bundle, operation_deadline=operation_deadline, clock=clock, stage_observer=stage_observer, require_promoted_runtime=require_promoted_runtime)
    observe_stage("independent_verification")
    try:
        published = _publish(
            run_dir,
            revision_dir,
            reference,
            final_quality,
            expected_attempts=list(state.get("gate_attempts") or []),
            operation_deadline=operation_deadline,
            clock=clock,
        )
    except OperationDeadlineExpired as exc:
        return {
            "status": "timeout",
            "stage": "delivery",
            "findings": [{"category": "timeout", "field": "operation", "issue": str(exc)}],
            "candidate_outputs": _candidate_outputs(revision_dir),
            "client_outputs": [],
        }
    observe_stage("delivery")
    return published


def _save_recorded_handoff(revision_dir: Path, request_path: Path) -> None:
    request = _read(request_path)
    response = recorded_acceptance_response(request)
    _write(revision_dir / request["response_path"], response)


def _save_verification_response(revision_dir: Path, request_path: Path, responder: Callable[[Mapping[str, Any]], Mapping[str, Any]]) -> None:
    request = _read(request_path)
    _write(revision_dir / request["response_path"], responder(request))


def run_release_gate(
    repo_root: Path,
    *,
    verification_responder: Callable[[Mapping[str, Any]], Mapping[str, Any]] | None = None,
    evidence_root: Path | None = None,
    require_promoted_runtime: bool = True,
) -> dict[str, Any]:
    """Exercise all six lifecycle cases; genuine verifier responses remain external."""
    timestamp = datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%SZ")
    if evidence_root is not None:
        root = evidence_root.resolve()
    else:
        base = (repo_root / ".scratch/release-gate" / timestamp).resolve()
        root = base
        suffix = 1
        while root.exists():
            root = base.with_name(f"{base.name}-{suffix}")
            suffix += 1
    root.mkdir(parents=True, exist_ok=True)
    results = []
    corpus = _read_corpus(repo_root / "references/conformance-fixtures/branch-acceptance-corpus.json")
    for case in [item for item in corpus if str(item.get("profile", "")).endswith("complete")]:
            branch = str(case["study_type"]); richness = str(case["profile"]).split("-", 1)[0]
            run_dir = root / str(case["fixture_id"])
            if not (run_dir / REFERENCE).is_file():
                reference = _read(repo_root / str(case["source_fixture"]))
                reference["meta"]["protocol_number"] = f"{branch[:3].upper()}-{richness}-26"; reference["study"]["title"] = f"{branch} {richness} acceptance study"; reference["study"]["short_title"] = f"{branch} {richness} acceptance"
                if richness == "rich" and branch != "Retrospective":
                    reference["meta"]["icf_template"] = "Sterling"
                if richness == "rich":
                    reference["study"]["background"] = f"{reference['study']['background'].rstrip('.')} with extended follow-up."
                    reference["study"]["timeline"] = "6 months"
                    reference["procedures"]["assessments"] = list(reference["procedures"].get("assessments", [])) + ["Extended follow-up assessment"]
                    reference["procedures"]["visit_schedule_table"] = list(reference["procedures"].get("visit_schedule_table", [])) + [{"visitNumber": "4", "visitName": "Month 6", "visitWindow": "±14 days", "CRFnumber": "CRF-04"}]
                    reference["population"]["sample_size"] = "80 participants"
                    for item in reference["population"].get("sample_size_evidence", []):
                        if isinstance(item, dict): item["value"] = "80 participants"
                    for item in reference.get("statistics", {}).get("sample_size_evidence", []):
                        if isinstance(item, dict): item["value"] = "80 participants"
                    if branch != "Retrospective":
                        reference["endpoints"].setdefault("other", []).extend([{"label": "Extended follow-up outcome", "time_point": "Month 6"}, {"label": "Participant experience outcome", "time_point": "Month 6"}])
                        reference["design"]["interventions"] = [{"name": "Sentinel Patch", "type": "Device", "description": "Primary monitoring configuration.", "arm_group_label": "Primary cohort"}, {"name": "Sentinel Patch Extended", "type": "Device", "description": "Extended monitoring configuration.", "arm_group_label": "Extended cohort"}]
                        reference["design"]["arms"] = [{"name": "Primary cohort", "description": "Primary observational cohort."}, {"name": "Extended cohort", "description": "Extended observational cohort."}]
                    if reference.get("sites"):
                        second_site = copy.deepcopy(reference["sites"][0]); second_site["facility"]["name"] = "Site Two"; second_site["contact"]["name"] = "Taylor Coordinator"; second_site["contact"]["email"] = "taylor@example.org"; second_site["investigators"][0]["name"] = "Morgan Investigator"
                        reference["sites"].append(second_site); reference["design"]["number_of_sites"] = 2
                        reference["design"]["study_design"] = str(reference["design"]["study_design"]).replace("single-center", "multicenter")
                reference["approval"] = {"status": "draft"}; _write(run_dir / REFERENCE, reference)
                first = prepare(run_dir)
                if first.get("status") != "awaiting_approval": results.append({"case": run_dir.name, "status": "failed", "result": first}); continue
                approval = approve(run_dir, approved_by="Recorded Acceptance")
                if approval.get("status") != "passed": results.append({"case": run_dir.name, "status": "failed", "result": approval}); continue
            reference = _read(run_dir / REFERENCE)
            case_bundle = contracted_template_bundle(repo_root, reference)
            result: dict[str, Any] = {}
            for _attempt in range(12):
                result = generate(run_dir, require_promoted_runtime=require_promoted_runtime)
                if result.get("status") != "awaiting_hermes": break
                revision_dir = run_dir / "revisions" / str(result["revision_id"])
                for relative in result.get("requests", []):
                    request_path = revision_dir / relative
                    if "verification-requests" in relative:
                        if verification_responder is not None:
                            _save_verification_response(revision_dir, request_path, verification_responder)
                    else:
                        _save_recorded_handoff(revision_dir, request_path)
                if result.get("stage") == "independent_verification" and verification_responder is None:
                    break
            results.append({"case": run_dir.name, "corpus": case, "icf_template": get_path(reference, "meta.icf_template"), "source_sha256": sha256_value(_approved_payload(reference)), "contracted_template_bundle": case_bundle, "status": "passed" if result.get("status") == "passed" else "failed", "result": result})
    passed = all(item["status"] == "passed" for item in results)
    awaiting = any(item.get("result", {}).get("status") == "awaiting_hermes" for item in results)
    synthetic_verification = bool(verification_responder is not None and getattr(verification_responder, "synthetic", False))
    recorded_drafting = True
    status = "structural_passed" if passed and recorded_drafting else "passed" if passed else "awaiting_hermes" if awaiting else "blocked"
    assurance = "synthetic-structural-only" if synthetic_verification else "recorded-drafting-structural-only"
    unique_bundles = {
        bundle["identity_sha256"]: bundle
        for item in results
        if isinstance((bundle := item.get("contracted_template_bundle")), Mapping)
    }
    report = {"status": status, "assurance": assurance, "recorded_drafting": recorded_drafting, "cases": results, "contracted_template_bundles": list(unique_bundles.values()), "distinct_source_count": len({item.get("source_sha256") for item in results}), "visual_evidence": "external image inspection required; no automated visual approval is fabricated", "evidence_root": root.as_posix()}
    _write(root / "release-gate-report.json", report); return report


def _synthetic_format_verification(request: Mapping[str, Any]) -> dict[str, Any]:
    """Produce explicit non-certifying structural responses for the matrix runner."""
    response: dict[str, Any] = {
        "schema_version": RESPONSE_SCHEMA,
        "request_id": request["request_id"],
        "request_sha256": request["request_sha256"],
        "task": request["task"],
        "producer": {"model_id": "DeterministicFormatConformance/v1"},
        "status": "passed",
        "findings": [],
    }
    if request["task"] == "rendered_page_visual_verification":
        response["page_assessments"] = [
            {
                "artifact": artifact["artifact"],
                "page": page["page"],
                "sha256": page["sha256"],
                "status": "passed",
                "checks": list(VISUAL_CHECKS),
            }
            for artifact in request.get("artifacts", [])
            for page in artifact.get("pages", [])
        ]
    else:
        response["section_assessments"] = [
            {
                "artifact": item["artifact"],
                "section_id": item["section_id"],
                "status": "passed",
                "checks": list(CONTENT_CHECKS),
            }
            for item in request.get("sections", [])
        ]
        response["cross_document_assessments"] = [
            {"check": check, "status": "passed"}
            for check in request.get("cross_document_checks", [])
        ]
    return response


setattr(_synthetic_format_verification, "synthetic", True)


def _run_format_conformance_in_disposable_candidate(
    repo_root: Path,
    evidence_root: Path,
) -> dict[str, Any]:
    """Run source conformance inside a governed disposable install candidate."""
    with tempfile.TemporaryDirectory(
        prefix=".clinical-document-generation.install-",
    ) as directory:
        staging_root = Path(directory).resolve()
        archive_path = staging_root / "candidate.zip"
        package_release(repo_root, archive_path)
        with zipfile.ZipFile(archive_path) as archive:
            _validated_archive_members(archive, staging_root)
            archive.extractall(staging_root)
        candidate = staging_root / "clinical-document-generation"
        integrity = _manifest_integrity(candidate, allow_runtime_state=False)
        if integrity:
            return {
                "status": "blocked",
                "assurance": "deterministic-structural-only",
                "live_certification_required": True,
                "findings": integrity,
            }
        provision = provision_render_assurance(candidate)
        if provision.get("status") != "passed":
            return {
                "status": "blocked",
                "assurance": "deterministic-structural-only",
                "live_certification_required": True,
                "findings": provision.get("findings", []),
            }
        command = [
            sys.executable,
            str(candidate / "scripts/workflow.py"),
            "--format-conformance",
            "--format-conformance-root",
            str(evidence_root),
        ]
        completed = subprocess.run(
            command,
            cwd=candidate,
            capture_output=True,
            text=True,
            timeout=FORMAT_CONFORMANCE_TIMEOUT_SECONDS,
            check=False,
        )
        try:
            result = json.loads(completed.stdout)
        except (TypeError, ValueError, json.JSONDecodeError):
            result = {
                "status": "blocked",
                "assurance": "deterministic-structural-only",
                "live_certification_required": True,
                "findings": [{
                    "category": "conformance",
                    "field": "disposable_candidate",
                    "issue": "The disposable format-conformance candidate did not return valid JSON.",
                    "exit_code": completed.returncode,
                }],
            }
        return dict(result)


def run_format_conformance(
    repo_root: Path,
    *,
    evidence_root: Path | None = None,
) -> dict[str, Any]:
    """Run all deterministic format cases without claiming live certification."""
    repo_root = repo_root.resolve()
    resolved_evidence_root = (
        evidence_root.resolve()
        if evidence_root is not None
        else (repo_root / ".scratch/format-conformance").resolve()
    )
    try:
        git_root = Path(subprocess.run(
            ["git", "-C", str(repo_root), "rev-parse", "--show-toplevel"],
            check=True,
            capture_output=True,
            text=True,
        ).stdout.strip()).resolve()
    except (OSError, subprocess.CalledProcessError):
        git_root = None
    if git_root == repo_root:
        return _run_format_conformance_in_disposable_candidate(
            repo_root,
            resolved_evidence_root,
        )
    available_page_renderers = page_renderers(
        skill_root=repo_root,
        require_promoted_runtime=False,
    )
    if not available_page_renderers:
        return _run_format_conformance_in_disposable_candidate(
            repo_root,
            resolved_evidence_root,
        )
    matrix = load_format_conformance_matrix(repo_root)
    report = run_release_gate(
        repo_root,
        verification_responder=_synthetic_format_verification,
        evidence_root=resolved_evidence_root,
        require_promoted_runtime=False,
    )
    output_audit = audit_format_conformance_outputs(
        repo_root,
        matrix,
        report,
        resolved_evidence_root,
    )
    covered = set()
    all_cases_passed = True
    for item in report.get("cases", []):
        if not isinstance(item, Mapping):
            all_cases_passed = False
            continue
        corpus: Mapping[str, Any] = item.get("corpus") if isinstance(item.get("corpus"), Mapping) else {}
        case_result: Mapping[str, Any] = item.get("result") if isinstance(item.get("result"), Mapping) else {}
        study_type = str(canonical_study_type(
            corpus.get("study_type") or case_result.get("study_type") or str(item.get("case") or "").split("-", 1)[0]
        ) or "")
        if study_type == "Retrospective":
            covered.add("retrospective-protocol")
        else:
            icf_template = str(item.get("icf_template") or "").strip().casefold()
            if icf_template:
                covered.add(f"{study_type.casefold()}-{icf_template}")
        all_cases_passed = all_cases_passed and item.get("status") == "passed"
    expected = {str(case["case_id"]) for case in matrix["cases"]}
    structural_passed = (
        report.get("status") == "structural_passed"
        and report.get("assurance") == "synthetic-structural-only"
        and all_cases_passed
        and covered == expected
        and output_audit.get("status") == "passed"
    )
    result = {
        "schema_version": "clinical-format-conformance-report/v1",
        "status": "structural_passed" if structural_passed else "blocked",
        "assurance": "deterministic-structural-only",
        "live_certification_required": True,
        "matrix_sha256": matrix["matrix_sha256"],
        "covered_cases": sorted(covered),
        "expected_cases": sorted(expected),
        "release_gate_status": report.get("status"),
        "release_gate_assurance": report.get("assurance"),
        "release_gate_report_sha256": canonical_evidence_sha256(report),
        "output_baseline_status": output_audit.get("status"),
        "output_baseline_evidence_sha256": output_audit.get("evidence_sha256"),
        "output_baseline_cases": output_audit.get("cases", []),
        "evidence_root": report.get("evidence_root"),
    }
    root = Path(str(report.get("evidence_root") or resolved_evidence_root))
    root.mkdir(parents=True, exist_ok=True)
    _write(root / "format-conformance-report.json", result)
    return result


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--run-dir")
    parser.add_argument("--stage", choices=("prepare", "approve", "validate", "generate"))
    parser.add_argument("--approved-by", default="client")
    parser.add_argument("--source-md")
    parser.add_argument("--desktop-operation", action="store_true", help="run the shipped real-Hermes Desktop adapter")
    parser.add_argument("--desktop-opener-command")
    parser.add_argument("--parent-visual-review-command")
    parser.add_argument("--operation-id", default="default")
    parser.add_argument("--release-gate", action="store_true")
    parser.add_argument("--release-gate-root")
    parser.add_argument("--format-conformance", action="store_true", help="run the non-certifying deterministic format matrix")
    parser.add_argument("--format-conformance-root")
    parser.add_argument("--package-release", metavar="ARCHIVE", help="create an immutable candidate archive")
    parser.add_argument("--provision-candidate", action="store_true", help="install the packaged PDFium runtime into an extracted certification candidate")
    parser.add_argument("--bind-certification", metavar="REPORT", help="embed a passing full-corpus report in --release-archive")
    parser.add_argument("--release-archive", help="candidate archive used with --bind-certification")
    parser.add_argument("--verify-installation", action="store_true", help="smoke-test this installed release")
    parser.add_argument("--install-release", metavar="ARCHIVE", help="atomically install and activate a certified release archive")
    parser.add_argument("--rollback-release", action="store_true", help="verify and atomically restore the immediately previous release")
    parser.add_argument("--skills-dir", help="Hermes skills directory for install or rollback")
    parser.add_argument("--hermes-config", help="Hermes config.yaml whose discovery path must select only the Promoted Release")
    parser.add_argument("--internal-pdfium-worker", metavar="REQUEST", help=argparse.SUPPRESS)
    args = parser.parse_args(argv)
    if args.internal_pdfium_worker: result = run_pdfium_worker(Path(args.internal_pdfium_worker))
    elif args.package_release: result = package_release(SCRIPT_DIR.parent, Path(args.package_release))
    elif args.provision_candidate:
        result = provision_render_assurance(SCRIPT_DIR.parent)
        shutil.rmtree(SCRIPT_DIR / "__pycache__", ignore_errors=True)
    elif args.bind_certification:
        if not args.release_archive: parser.error("--release-archive is required with --bind-certification")
        result = bind_release_certification(Path(args.release_archive), Path(args.bind_certification))
    elif args.verify_installation: result = verify_installation(SCRIPT_DIR.parent)
    elif args.install_release:
        if not args.skills_dir or not args.hermes_config: parser.error("--skills-dir and --hermes-config are required with --install-release")
        result = install_release(Path(args.install_release), Path(args.skills_dir), hermes_config_path=Path(args.hermes_config))
    elif args.rollback_release:
        if not args.skills_dir: parser.error("--skills-dir is required with --rollback-release")
        result = rollback_release(Path(args.skills_dir))
    elif args.desktop_operation:
        if not args.run_dir: parser.error("--run-dir is required with --desktop-operation")
        if not args.desktop_opener_command: parser.error("--desktop-opener-command is required with --desktop-operation")
        if not args.parent_visual_review_command: parser.error("--parent-visual-review-command is required with --desktop-operation")
        result = run_production_desktop_operation(
            Path(args.run_dir),
            operation_id=args.operation_id,
            opener=command_desktop_opener(Path(args.desktop_opener_command)),
            parent_visual_reviewer=command_parent_visual_reviewer(
                Path(args.parent_visual_review_command)
            ),
        )
    elif args.release_gate: result = run_release_gate(SCRIPT_DIR.parent, evidence_root=Path(args.release_gate_root) if args.release_gate_root else None)
    elif args.format_conformance: result = run_format_conformance(SCRIPT_DIR.parent, evidence_root=Path(args.format_conformance_root) if args.format_conformance_root else None)
    else:
        if not args.run_dir or not args.stage: parser.error("--run-dir and --stage are required unless a release or conformance operation is used")
        run_dir = Path(args.run_dir).expanduser().resolve()
        result = {"prepare": prepare, "approve": approve, "validate": validate, "generate": generate}[args.stage](run_dir, approved_by=args.approved_by, source_md=Path(args.source_md).expanduser() if args.source_md else None)
    print(json.dumps(result, indent=2, ensure_ascii=False)); return 0 if result.get("status") in {"passed", "structural_passed", "awaiting_approval", "awaiting_hermes"} else 1


__all__ = ["approve", "bind_release_certification", "command_desktop_opener", "command_parent_visual_reviewer", "confirm_desktop_delivery", "desktop_attachment_reply", "desktop_operation_state_path", "generate", "install_release", "package_release", "performance_classification", "prepare", "provision_render_assurance", "resolve_python_runtime", "rollback_release", "run_desktop_operation", "run_format_conformance", "run_production_desktop_operation", "run_release_gate", "validate", "verify_installation"]


if __name__ == "__main__": raise SystemExit(main())
