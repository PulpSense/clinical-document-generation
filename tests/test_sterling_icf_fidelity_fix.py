import copy
import json
import re
from pathlib import Path

import pytest
from docx import Document
from docx.oxml.ns import qn

import contracts
import prs_xml
import quality
from rendering import render_documents


ROOT = Path(__file__).resolve().parents[1]
STUDY_RESULTS_DEFAULT = (
    "You will not routinely receive the results of testing performed for this study. "
    "If a study test identifies a new medical condition or concern, the study doctor "
    "will inform you and recommend appropriate follow-up with your treating doctor. "
    "The finding may affect whether you can continue in the study."
)
AUTHORIZATION_DURATION_DEFAULT = (
    "This authorization does not expire unless you revoke it in writing."
)
POST_STUDY_HEADING = (
    "What happens to information about me after the study is over or if I cancel my "
    "permission to use my PHI?"
)


def source(*, family="Sterling"):
    value = json.loads(
        (ROOT / "tests/fixtures/prospective-acceptance-source.json").read_text(
            encoding="utf-8"
        )
    )
    value["meta"]["icf_template"] = family
    value["design"]["intervention_name"] = (
        "TECNIS PureSee intraocular lens and TECNIS Odyssey intraocular lens"
    )
    value["design"]["intervention_type"] = "Medical device — intraocular lenses."
    value["procedures"]["assessments"] = (
        "Demographics, medical history, eye examinations, visual acuity tests, "
        "the AIOLIS questionnaire, adverse events, and device deficiencies."
    )
    value["procedures"]["termination"] = None
    value.setdefault("confidentiality", {})["data_handling"] = None
    return value


def model():
    return {
        "protocol": [],
        "prs": {},
        "icf": {
            "icf.procedures": {
                "paragraphs": [
                    {
                        "text": (
                            "You will attend the approved visits for study lens "
                            "implantation and study eye testing."
                        )
                    }
                ],
                "lists": [],
            },
            "icf.privacy": {
                "paragraphs": [
                    {
                        "text": (
                            "Study records will use appropriate identifiers and "
                            "access controls."
                        )
                    }
                ],
                "lists": [],
            },
        },
    }


def render(tmp_path, reference=None, draft=None):
    reference = reference or source()
    report = render_documents(
        ROOT,
        tmp_path,
        reference,
        draft or model(),
        artifact_names={"icf"},
    )
    assert report["status"] == "passed", report
    path = tmp_path / "candidate/icf.docx"
    return path, Document(path)


def normalized(text):
    return re.sub(r"\s+", " ", text).strip()


def visible(document):
    return normalized(" ".join(p.text for p in document.paragraphs if p.text.strip()))


def heading_index(document, title):
    return next(
        index
        for index, paragraph in enumerate(document.paragraphs)
        if normalized(paragraph.text) == title
    )


def section_paragraphs(document, title):
    start = heading_index(document, title)
    values = []
    for paragraph in document.paragraphs[start + 1 :]:
        if paragraph.text.strip() and paragraph.style.name.casefold().startswith("heading"):
            break
        values.append(paragraph)
    return values


def has_real_list(paragraph):
    properties = paragraph._p.pPr
    return properties is not None and properties.find(qn("w:numPr")) is not None


FIDELITY_MODULES = {
    "sterling.results.policy",
    "sterling.voluntary.core",
    "sterling.withdrawal.investigator-termination",
    "sterling.privacy.authorization",
    "sterling.privacy.data-categories",
    "sterling.privacy.authorized-recipients",
    "sterling.privacy.authorization-duration",
    "sterling.privacy.authorization-withdrawal",
    "sterling.privacy.post-study",
    "sterling.privacy.research-record-access",
    "sterling.privacy.registry-disclosure",
}


def fidelity_findings(path, reference):
    return [
        item
        for item in quality.validate_sterling_clause_contract(path, reference)["findings"]
        if item.get("module_id") in FIDELITY_MODULES
        and item.get("severity") == "blocking"
    ]


def test_sterling_contract_is_versioned_template_derived_module_authority():
    governed = contracts.sterling_clause_contract(ROOT)
    assert governed["schema_version"] == "sterling-icf-modules/v2"
    assert governed["authority"]["template"].endswith(
        "reference/sterling-icf-reference.docx"
    )
    modules = governed["modules"]
    assert modules
    assert {item["classification"] for item in modules} <= {
        "core",
        "conditional",
        "source-bound",
    }
    required = {
        "module_id",
        "provenance",
        "placement",
        "applicability",
        "classification",
        "substitutions",
        "expected_structure",
        "validation",
        "severity",
    }
    assert all(required <= set(item) for item in modules)
    assert len({item["module_id"] for item in modules}) == len(modules)


def test_study_results_default_is_immediately_after_procedures(tmp_path):
    _path, document = render(tmp_path)
    assert heading_index(document, "STUDY RESULTS") > heading_index(document, "PROCEDURES")
    headings = [
        normalized(p.text)
        for p in document.paragraphs
        if p.text.strip() and p.style.name.casefold().startswith("heading")
    ]
    assert headings.index("STUDY RESULTS") == headings.index("PROCEDURES") + 1
    assert STUDY_RESULTS_DEFAULT in visible(document)


def test_source_results_policy_overrides_default_without_duplication(tmp_path):
    reference = source()
    approved = (
        "A plain-language summary of overall study results will be offered after "
        "the study is complete."
    )
    reference["procedures"]["return_of_results"] = approved
    _path, document = render(tmp_path, reference)
    text = visible(document)
    assert text.count(approved) == 1
    assert STUDY_RESULTS_DEFAULT not in text


def test_withdrawal_restores_concepts_and_real_default_reason_bullets(tmp_path):
    _path, document = render(tmp_path)
    paragraphs = section_paragraphs(document, "VOLUNTARY PARTICIPATION/WITHDRAWAL")
    text = normalized(" ".join(p.text for p in paragraphs)).casefold()
    for concept in (
        "participation is voluntary",
        "decline",
        "withdraw",
        "without penalty",
        "health or welfare",
        "no longer eligible",
        "required instructions or procedures",
        "investigator",
        "sponsor",
        "irb",
        "information collected before",
    ):
        assert concept in text
    bullets = [normalized(p.text).casefold() for p in paragraphs if has_real_list(p)]
    assert len(bullets) >= 4
    assert any("health or welfare" in item for item in bullets)
    assert any("no longer eligible" in item for item in bullets)
    assert any("required instructions or procedures" in item for item in bullets)
    assert any("stops the study" in item for item in bullets)
    assert not any("drug" in item or "disease worsens" in item for item in bullets)


def test_source_withdrawal_reasons_replace_defaults_and_remain_real_bullets(tmp_path):
    reference = source()
    reference["procedures"]["termination"] = [
        "The study doctor determines that continued participation is unsafe.",
        "The sponsor closes the study.",
    ]
    _path, document = render(tmp_path, reference)
    paragraphs = section_paragraphs(document, "VOLUNTARY PARTICIPATION/WITHDRAWAL")
    bullets = [normalized(p.text) for p in paragraphs if has_real_list(p)]
    assert bullets == reference["procedures"]["termination"]
    assert not any("no longer eligible" in item.casefold() for item in bullets)


def test_confidentiality_restores_fixed_core_and_real_factual_lists(tmp_path):
    _path, document = render(tmp_path)
    paragraphs = section_paragraphs(
        document,
        "CONFIDENTIALITY AUTHORIZATION TO COLLECT, USE DISCLOSE YOUR MEDICAL INFORMATION",
    )
    text = normalized(" ".join(p.text for p in paragraphs)).casefold()
    for concept in (
        "health information",
        "research study",
        "may disclose",
        "may receive",
        "voluntary",
        "ordinary medical care or benefits",
        "cannot begin or continue",
        "does not expire unless you revoke it in writing",
        "already collected",
        "redisclosed",
        "after the study is over",
        "cancel my permission",
        "clinicaltrials.gov",
    ):
        assert concept in text
    bullets = [normalized(p.text).casefold() for p in paragraphs if has_real_list(p)]
    assert any("demographic" in item for item in bullets)
    assert any("medical history" in item for item in bullets)
    assert any("eye" in item or "vision" in item for item in bullets)
    assert any("sponsor" in item for item in bullets)
    assert any("institutional review board" in item or "irb" in item for item in bullets)
    assert all(not item.startswith("•") for item in bullets)
    assert AUTHORIZATION_DURATION_DEFAULT in visible(document)
    assert POST_STUDY_HEADING in visible(document)


def test_missing_optional_privacy_facts_do_not_block_and_are_not_invented(tmp_path):
    reference = source()
    reference["confidentiality"].pop("information_categories", None)
    reference["confidentiality"].pop("authorized_recipients", None)
    path, document = render(tmp_path, reference)
    assert fidelity_findings(path, reference) == []
    text = visible(document).casefold()
    assert "social security number" not in text
    assert "genetic information" not in text


def test_research_record_access_is_source_bound(tmp_path):
    _path, without = render(tmp_path / "without")
    assert "cannot see your research records until" not in visible(without).casefold()

    reference = source()
    reference["confidentiality"]["research_record_access"] = (
        "You cannot see your research records until the study is complete."
    )
    _path, with_access = render(tmp_path / "with", reference)
    assert (
        "You cannot see your research records until the study is complete."
        in visible(with_access)
    )

    injected = Document(_path)
    privacy = section_paragraphs(
        injected,
        "CONFIDENTIALITY AUTHORIZATION TO COLLECT, USE DISCLOSE YOUR MEDICAL INFORMATION",
    )
    privacy[0]._p.addprevious(
        injected.add_paragraph(
            "You may review your research records only after the trial concludes."
        )._p
    )
    unsupported = tmp_path / "unsupported-access.docx"
    injected.save(unsupported)
    finding = next(
        item
        for item in quality.validate_sterling_clause_contract(
            unsupported, source()
        )["findings"]
        if item.get("module_id") == "sterling.privacy.research-record-access"
    )
    assert finding["code"] == "sterling-module-unsupported"
    assert finding["severity"] == "warning"


def test_non_eye_sterling_study_uses_intervention_neutral_phi_category(tmp_path):
    reference = source()
    reference["design"]["intervention_name"] = "Wearable cardiac monitor"
    reference["design"]["intervention_type"] = "Medical device"
    reference["procedures"]["assessments"] = (
        "Blood pressure assessment, ECG testing, and a symptom questionnaire."
    )
    _path, document = render(tmp_path, reference)
    privacy = normalized(" ".join(
        paragraph.text
        for paragraph in section_paragraphs(
            document,
            "CONFIDENTIALITY AUTHORIZATION TO COLLECT, USE DISCLOSE YOUR MEDICAL INFORMATION",
        )
    ))
    assert "Results of study assessments, tests, and procedures." in privacy
    assert "Responses to study questionnaires or surveys." in privacy
    assert "eye examinations" not in privacy
    assert "vision tests" not in privacy


def test_lens_terminology_is_specific_and_generic_drug_product_terms_are_absent(tmp_path):
    path, document = render(tmp_path)
    text = visible(document).casefold()
    assert "study lens" in text or "study lenses" in text
    for prohibited in (
        "drug/device",
        "drug/device/product",
        "study product",
        "study drug",
        "receiving a drug",
        "discontinuing a drug",
    ):
        assert prohibited not in text
    document.add_paragraph("You will receive the study drug.")
    broken = tmp_path / "bad-terminology.docx"
    document.save(broken)
    report = quality.validate_sterling_clause_contract(broken, source())
    assert any(
        item.get("code") == "sterling-intervention-terminology-invalid"
        and "study drug" in item.get("prohibited_terms", [])
        for item in report["findings"]
    )


@pytest.mark.parametrize("surface", ["table", "header", "footer"])
def test_lens_terminology_validation_covers_every_docx_surface(tmp_path, surface):
    path, document = render(tmp_path / surface)
    if surface == "table":
        document.add_table(rows=1, cols=1).cell(0, 0).text = "Study drug"
    else:
        part = (
            document.sections[0].header
            if surface == "header"
            else document.sections[0].footer
        )
        part.add_paragraph("Study drug")
    broken = tmp_path / f"{surface}-terminology.docx"
    document.save(broken)

    report = quality.validate_sterling_clause_contract(broken, source())
    assert any(
        item.get("code") == "sterling-intervention-terminology-invalid"
        and "study drug" in item.get("prohibited_terms", [])
        for item in report["findings"]
    )


def test_no_drafting_instructions_merge_fields_unused_alternatives_or_duplication(tmp_path):
    path, document = render(tmp_path)
    text = visible(document)
    lowered = text.casefold()
    assert "[include" not in lowered
    assert "[insert" not in lowered
    assert "[if applicable" not in lowered
    assert "«" not in text and "»" not in text
    assert not re.search(r"\{[#/]?[A-Za-z_]", text)
    substantive = [
        normalized(p.text).casefold()
        for p in document.paragraphs
        if len(normalized(p.text).split()) >= 8
    ]
    assert len(substantive) == len(set(substantive))
    assert fidelity_findings(path, source()) == []


def test_missing_core_results_body_fails_with_precise_module_diagnostic(tmp_path):
    path, document = render(tmp_path)
    for paragraph in section_paragraphs(document, "STUDY RESULTS"):
        paragraph._element.getparent().remove(paragraph._element)
    broken = tmp_path / "broken.docx"
    document.save(broken)
    report = quality.validate_sterling_clause_contract(broken, source())
    finding = next(
        item
        for item in report["findings"]
        if item.get("module_id") == "sterling.results.policy"
    )
    assert finding["code"] == "sterling-module-missing"
    assert finding["severity"] == "blocking"
    assert finding["expected_section"] == "STUDY RESULTS"


def test_each_privacy_list_module_validates_its_own_word_list_items(tmp_path):
    path, document = render(tmp_path)
    target = next(
        paragraph
        for paragraph in document.paragraphs
        if paragraph.text == "Demographic information collected for the study."
    )
    recipient_intro = next(
        paragraph
        for paragraph in document.paragraphs
        if paragraph.text == (
            "The following people or organizations may receive and use your study information:"
        )
    )
    recipient_intro._p.addnext(target._p)
    broken = tmp_path / "misplaced-phi-category.docx"
    document.save(broken)

    report = quality.validate_sterling_clause_contract(broken, source())
    category = next(
        item
        for item in report["findings"]
        if item.get("module_id") == "sterling.privacy.data-categories"
    )
    assert category["code"] == "sterling-module-weakened"
    assert category["malformed_list"] is True
    assert "demographic information collected for the study" in category[
        "missing_expected_list_items"
    ]
    assert not any(
        item.get("module_id") == "sterling.privacy.authorized-recipients"
        for item in report["findings"]
    )


def test_non_sterling_rendering_is_unchanged_by_sterling_modules(tmp_path):
    reference = source(family="Advarra")
    path, document = render(tmp_path, reference)
    assert path.is_file()
    assert not any(
        normalized(paragraph.text) == "STUDY RESULTS"
        for paragraph in document.paragraphs
    )
    assert "RELEASE OF MEDICAL RECORDS AND PRIVACY" in visible(document)


def test_generated_lens_study_xml_remains_structurally_valid(tmp_path):
    reference = source()
    output = tmp_path / "study.xml"
    report = prs_xml.generate(
        ROOT
        / "assets/client-templates/prs/clinicaltrials_prs_full_placeholder_template.xml",
        output,
        reference,
        {
            "brief_summary": {
                "text": "This prospective study evaluates visual outcomes after lens implantation."
            },
            "detailed_description": {
                "text": "The study evaluates visual acuity and participant-reported outcomes after the approved intraocular lens procedures."
            },
        },
        structural_template=(
            ROOT / "assets/client-templates/reference/prs-manual-reference.xml"
        ),
    )
    assert report["status"] == "passed", report
    assert prs_xml.validate_output(
        output,
        reference,
        ROOT / "assets/client-templates/reference/prs-manual-reference.xml",
        generation_template=(
            ROOT
            / "assets/client-templates/prs/clinicaltrials_prs_full_placeholder_template.xml"
        ),
    ) == []
