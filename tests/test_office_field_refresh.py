"""One pagination pass supplies page evidence without office-roundtripping DOCX."""
import subprocess
from pathlib import Path

import pytest
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from pypdf import PdfWriter

import rendering


def protocol(path):
    doc = Document()
    doc.styles.add_style("toc 1", WD_STYLE_TYPE.PARAGRAPH)
    doc.styles.add_style("toc 2", WD_STYLE_TYPE.PARAGRAPH)
    doc.add_heading("1. TITLE PAGE", 1)
    doc.add_paragraph("Clinical facts must remain unchanged.")
    doc.add_heading("4. TABLE OF CONTENTS", 1)
    doc.add_heading("5. INTRODUCTION", 1)
    table = doc.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "Source-owned table value"
    rendering._populate_protocol_toc_cache(doc, {})
    doc.save(path)


def test_single_office_call_uses_outline_and_preserves_native_content(tmp_path, monkeypatch):
    path = tmp_path / "protocol.docx"
    protocol(path)
    before = Document(path)
    calls = []
    monkeypatch.setattr(rendering.shutil, "which", lambda _: "/controlled/soffice")
    class Office:
        returncode = 0
        def __init__(self, args, **kwargs):
            calls.append(args)
            writer = PdfWriter()
            for _ in range(3):
                writer.add_blank_page(width=612, height=792)
            for title, page in [("1. TITLE PAGE", 0), ("4. TABLE OF CONTENTS", 1), ("5. INTRODUCTION", 2)]:
                writer.add_outline_item(title, page)
            output = Path(args[args.index("--outdir") + 1]) / "protocol.pdf"
            writer.write(output)
        def communicate(self, timeout=None):
            return "converted", ""
    monkeypatch.setattr(rendering.subprocess, "Popen", Office)
    result = rendering.refresh_protocol_fields(path)
    after = Document(path)
    assert result["status"] == "passed"
    assert result["page_count"] == 3
    assert len(calls) == 1
    assert [p.text for p in after.paragraphs if p.style.name.casefold() in {"toc 1", "toc 2"}] == [
        "1. TITLE PAGE\t1", "4. TABLE OF CONTENTS\t2", "5. INTRODUCTION\t3"]
    assert [p.text for p in before.paragraphs if not p.style.name.casefold().startswith("toc")] == [
        p.text for p in after.paragraphs if not p.style.name.casefold().startswith("toc")]
    assert before.tables[0]._tbl.xml == after.tables[0]._tbl.xml
    assert not list(tmp_path.glob(".office-fields-*"))
    assert not list(tmp_path.rglob("*.pdf"))


def test_missing_office_leaves_original_bytes_unchanged(tmp_path, monkeypatch):
    path = tmp_path / "protocol.docx"
    protocol(path)
    original = path.read_bytes()
    monkeypatch.setattr(rendering.shutil, "which", lambda _: None)
    with pytest.raises(RuntimeError, match="CLINICAL_SOFFICE"):
        rendering.refresh_protocol_fields(path)
    assert path.read_bytes() == original


def test_timeout_stops_office_and_preserves_original(tmp_path, monkeypatch):
    path = tmp_path / "protocol.docx"
    protocol(path)
    original = path.read_bytes()
    monkeypatch.setattr(rendering.shutil, "which", lambda _: "/controlled/soffice")
    killed = []
    class Office:
        pid = 12345
        def __init__(self, *args, **kwargs):
            pass
        def communicate(self, timeout=None):
            if timeout is not None:
                raise subprocess.TimeoutExpired("office", timeout)
            return "", ""
    monkeypatch.setattr(rendering.subprocess, "Popen", Office)
    monkeypatch.setattr(rendering.os, "killpg", lambda pid, signal: killed.append(pid))
    with pytest.raises(RuntimeError, match="timed out"):
        rendering.refresh_protocol_fields(path, timeout_seconds=0.1)
    assert killed == [12345]
    assert path.read_bytes() == original
    assert not list(tmp_path.glob(".office-fields-*"))


def test_missing_heading_outline_keeps_original_bytes(tmp_path, monkeypatch):
    path = tmp_path / "protocol.docx"
    protocol(path)
    original = path.read_bytes()
    monkeypatch.setattr(rendering.shutil, "which", lambda _: "/controlled/soffice")
    class Office:
        returncode = 0
        def __init__(self, args, **kwargs):
            writer = PdfWriter()
            writer.add_blank_page(width=612, height=792)
            writer.write(Path(args[args.index("--outdir") + 1]) / "protocol.pdf")
        def communicate(self, timeout=None):
            return "converted", ""
    monkeypatch.setattr(rendering.subprocess, "Popen", Office)
    with pytest.raises(ValueError, match="did not identify one page"):
        rendering.refresh_protocol_fields(path)
    assert path.read_bytes() == original


def test_pagination_receipt_records_actual_heading_pages(tmp_path, monkeypatch):
    path = tmp_path / 'protocol.docx'
    protocol(path)
    monkeypatch.setattr(rendering.shutil, 'which', lambda _: '/controlled/soffice')
    class Office:
        returncode = 0
        def __init__(self, args, **kwargs):
            writer = PdfWriter()
            for _ in range(5):
                writer.add_blank_page(width=612, height=792)
            for title, page in [('1. TITLE PAGE', 0), ('4. TABLE OF CONTENTS', 1), ('5. INTRODUCTION', 4)]:
                writer.add_outline_item(title, page)
            writer.write(Path(args[args.index('--outdir') + 1]) / 'protocol.pdf')
        def communicate(self, timeout=None):
            return 'converted', ''
    monkeypatch.setattr(rendering.subprocess, 'Popen', Office)
    receipt = rendering.refresh_protocol_fields(path)
    assert receipt['heading_pages'] == {
        '1. TITLE PAGE': 1, '4. TABLE OF CONTENTS': 2, '5. INTRODUCTION': 5}


@pytest.mark.parametrize('branch', ['prospective', 'ambispective', 'retrospective'])
def test_real_office_final_render_matches_cached_toc(tmp_path, branch):
    """Compare the final cached DOCX with its own render on the same office host."""
    import json
    import os
    from pypdf import PdfReader
    office = os.environ.get('CLINICAL_SOFFICE')
    if not office:
        pytest.skip('Set CLINICAL_SOFFICE to run the real-office regression.')
    root = Path(__file__).resolve().parents[1]
    source = json.loads((root / f'tests/fixtures/{branch}-acceptance-source.json').read_text())
    rendering.render_documents(root, tmp_path, source, {'protocol': [], 'icf': {}, 'prs': {}},
                               artifact_names={'protocol'}, audit_output=False)
    path = tmp_path / 'candidate/protocol.docx'
    receipt = rendering.refresh_protocol_fields(path)
    out = tmp_path / 'final-render'
    out.mkdir()
    subprocess.run([office, f'-env:UserInstallation={(tmp_path / "office-profile").as_uri()}',
                    '--headless', '--convert-to', 'pdf:writer_pdf_Export', '--outdir', str(out), str(path)],
                   capture_output=True, text=True, check=True, timeout=120)
    reader = PdfReader(out / 'protocol.pdf')
    pages = {}
    def collect(items):
        for item in items:
            if isinstance(item, list):
                collect(item)
            else:
                pages[item.title] = reader.get_destination_page_number(item) + 1
    collect(reader.outline)
    cached = {p.text.rsplit('\t', 1)[0]: int(p.text.rsplit('\t', 1)[1])
              for p in Document(path).paragraphs if p.style.name.casefold() in {'toc 1', 'toc 2'}}
    assert cached == receipt['heading_pages']
    assert cached == {heading: pages[heading] for heading in cached}
    assert receipt['page_count'] == len(reader.pages)
