"""Source interpretation uses authored facts, independent of production projections."""
import copy
import json
from pathlib import Path

import pytest

import contracts
import drafting
import quality
import rendering
from docx import Document
from test_clean_delivery import deliver
from test_handoff_quality import acceptance_verification

ROOT = Path(__file__).resolve().parents[1]
BRANCHES = [('Prospective', 'Advarra'), ('Prospective', 'Sterling'),
            ('Ambispective', 'Advarra'), ('Ambispective', 'Sterling'),
            ('Retrospective', 'Advarra')]


def source_for_routing(branch='Prospective', family='Advarra'):
    case = next(item for item in quality.branch_acceptance_inventory(ROOT)
                if item['descriptor']['study_type'] == branch
                and (branch == 'Retrospective' or item['descriptor']['icf_template'] == family))
    source = copy.deepcopy(quality.branch_acceptance_source(ROOT, case))
    source['study']['background'] = (
        'Previous studies found better sleep after exercise. '
        'The new diary approach might reduce night waking, but this is unproven. '
        'Participants are not guaranteed any direct benefit.')
    source['risks_benefits']['benefits'] = None
    source['procedures'].update({
        'visits': None, 'visit_schedule_table': None,
        'visit_schedule': ['First contact', 'Final contact'],
        'assessments': 'At the first contact collect a sleep diary. At the final contact measure fatigue.',
        'evaluation': None,
    })
    if branch == 'Retrospective':
        source['procedures']['assessments'] = 'Abstract previously recorded sleep diary and fatigue measurements; no new participant contact.'
        source['procedures']['visit_schedule'] = None
    return source


def schedule_mapping():
    return {
        'assignments': [
            {'visit': 'First contact', 'activity': 'sleep diary',
             'source_path': 'procedures.assessments',
             'source_excerpt': 'At the first contact collect a sleep diary.'},
            {'visit': 'Final contact', 'activity': 'fatigue',
             'source_path': 'procedures.assessments',
             'source_excerpt': 'At the final contact measure fatigue.'},
        ],
        'unallocated': [],
        'rationale': 'Both contacts are explicitly associated with their assessments.',
    }


@pytest.mark.parametrize('target', ['evaluation-procedures', 'risks-benefits.benefits'])
def test_response_envelope_exposes_required_interpretation_without_default_decision(tmp_path, target):
    source = source_for_routing()
    batch = next(b for b in contracts.batch_plan('Prospective') if target in b.section_ids)
    path = drafting.create_drafting_request(
        repo_root=ROOT, revision_dir=tmp_path, revision_id='r-envelope', reference=source,
        batch=batch, target_ids=[target], attempts={target: 1}, wave='initial')
    request = json.loads(path.read_text())
    response = drafting.response_template(request)
    key = 'schedule_interpretation' if target == 'evaluation-procedures' else 'benefit_interpretation'
    assert response['section_results'][0][key]['rationale'] == ''
    accepted, findings = drafting.validate_response(request, response)
    assert not (accepted or {}).get('drafts')
    assert any(row.get('code') == 'source-interpretation-invalid' for row in findings)


@pytest.mark.parametrize('branch,family', BRANCHES)
def test_background_benefit_context_reaches_correct_surfaces(tmp_path, branch, family):
    source = source_for_routing(branch, family)
    target = 'ethics' if branch == 'Retrospective' else 'risks-benefits.benefits'
    batch = next(b for b in contracts.batch_plan(branch, family) if target in b.section_ids)
    path = drafting.create_drafting_request(
        repo_root=ROOT, revision_dir=tmp_path, revision_id='r-routing', reference=source,
        batch=batch, target_ids=[target], attempts={target: 1}, wave='initial')
    request = json.loads(path.read_text())
    section = request['section_contracts'][0]
    assert section['benefit_context']['source_path'] == 'study.background'
    assert section['benefit_context']['value'] == source['study']['background']
    assert 'study.background' in section['optional_evidence']
    assert 'study.background' not in section['minimum_evidence']
    review = quality.source_routing_obligations(source)
    destinations = {(item['artifact'], item['section_id']) for item in review
                    if item['kind'] == 'benefit_context'}
    assert ('protocol', target) in destinations
    assert (('icf', 'icf.benefits') in destinations) == (branch != 'Retrospective')


@pytest.mark.parametrize('form', ['prose', 'list', 'mixed'])
def test_model_maps_explicit_relationships_without_changing_approved_source(form):
    source = source_for_routing()
    if form == 'list':
        source['procedures']['assessments'] = [
            'At the first contact collect a sleep diary.',
            'At the final contact measure fatigue.']
    if form == 'mixed':
        source['procedures']['visit_schedule'][0] = {
            'visit': 'First contact', 'procedures': ['Medical history']}
    original = copy.deepcopy(source)
    projected = contracts.interpreted_schedule_reference(source, schedule_mapping())
    table = contracts.protocol_table_contracts(projected)['schedule-of-assessments']
    assert table['rows'][0] == ['Activity', 'First contact', 'Final contact']
    rows = {row[0]: row[1:] for row in table['rows'][2:]}
    assert rows['Sleep diary'] == ['X', '']
    assert rows['Fatigue'] == ['', 'X']
    if form == 'mixed':
        assert rows['Medical history'] == ['X', '']
    assert source == original


@pytest.mark.parametrize('mutation', ['unknown_visit', 'invented_activity', 'fabricated_quote', 'missing_disposition'])
def test_invalid_interpretations_are_rejected(mutation):
    source = source_for_routing()
    mapping = schedule_mapping()
    if mutation == 'unknown_visit': mapping['assignments'][0]['visit'] = 'Week 99'
    if mutation == 'invented_activity': mapping['assignments'][0]['activity'] = 'Blood draw'
    if mutation == 'fabricated_quote': mapping['assignments'][0]['source_excerpt'] = 'Invented source text'
    if mutation == 'missing_disposition': mapping.pop('unallocated')
    with pytest.raises(ValueError):
        contracts.interpreted_schedule_reference(source, mapping)


def test_unallocated_meaning_is_preserved_instead_of_guessed():
    source = source_for_routing()
    source['procedures']['assessments'] += ' A telephone safety check may be needed.'
    mapping = schedule_mapping()
    mapping['unallocated'] = [{'source_path': 'procedures.assessments',
                               'source_excerpt': 'A telephone safety check may be needed.'}]
    table = contracts.protocol_table_contracts(
        contracts.interpreted_schedule_reference(source, mapping))['schedule-of-assessments']
    assert 'A telephone safety check may be needed.' in table['supplemental_notes']
    assert not any(row[0] == 'A telephone safety check may be needed.' for row in table['rows'])


def test_explicit_contacts_can_be_extracted_when_no_separate_visit_inventory_exists():
    source = source_for_routing()
    source['procedures']['visit_schedule'] = None
    mapping = schedule_mapping()
    mapping['visits'] = [
        {'visit': 'first contact', 'source_path': 'procedures.assessments',
         'source_excerpt': 'At the first contact collect a sleep diary.'},
        {'visit': 'final contact', 'source_path': 'procedures.assessments',
         'source_excerpt': 'At the final contact measure fatigue.'}]
    for row in mapping['assignments']: row['visit'] = row['visit'].lower()
    projected = contracts.interpreted_schedule_reference(source, mapping)
    rows = contracts.protocol_table_contracts(projected)['schedule-of-assessments']['rows']
    assert rows[0] == ['Activity', 'first contact', 'final contact']
    assert ['Sleep diary', 'X', ''] in rows
    assert ['Fatigue', '', 'X'] in rows


@pytest.mark.parametrize('field', ['visit_schedule', 'visit_schedule_table'])
@pytest.mark.parametrize('form', ['prose', 'list'])
def test_accepted_alternate_narrative_schedule_has_a_repairable_mapping_path(field, form):
    source = source_for_routing()
    narrative = source['procedures']['assessments']
    source['procedures'].update({'assessments': None, 'visit_schedule': None,
                                 'visit_schedule_table': None, field: narrative})
    if field == 'visit_schedule_table':
        # This is optional supplementary prose, not an obligatory-input alias.
        source['procedures']['assessments'] = narrative
    assert not contracts.input_findings(source)
    mapping = schedule_mapping()
    mapping['visits'] = [
        {'visit': 'first contact', 'source_path': f'procedures.{field}',
         'source_excerpt': 'At the first contact collect a sleep diary.'},
        {'visit': 'final contact', 'source_path': f'procedures.{field}',
         'source_excerpt': 'At the final contact measure fatigue.'}]
    for row in mapping['assignments']:
        row['source_path'] = f'procedures.{field}'
        row['visit'] = row['visit'].lower()
    source['procedures'][field] += ' A telephone safety check may be needed.'
    mapping['unallocated'] = [{'source_path': f'procedures.{field}',
                               'source_excerpt': 'A telephone safety check may be needed.'}]
    if form == 'list':
        items = [row['source_excerpt'] for row in mapping['assignments']] + [mapping['unallocated'][0]['source_excerpt']]
        source['procedures'][field] = items
        mapping['narrative_items'] = [{'source_path': f'procedures.{field}', 'source_index': index,
                                     'source_excerpt': item} for index, item in enumerate(items)]
    if field == 'visit_schedule_table':
        mapping['assignments'].extend(schedule_mapping()['assignments'])
        for row in mapping['assignments']: row['visit'] = row['visit'].lower()
    contract = contracts.schedule_interpretation_contract(source)
    assert contract and f'procedures.{field}' in contract['sources']
    projected = contracts.interpreted_schedule_reference(source, mapping)
    table = contracts.protocol_table_contracts(projected)['schedule-of-assessments']
    assert ['Sleep diary', 'X', ''] in table['rows']
    assert ['Fatigue', '', 'X'] in table['rows']
    assert 'A telephone safety check may be needed.' in table['supplemental_notes']
    assert len(table['rows'][0]) == 3  # Prose sentences are not extra contacts.


def test_narrative_list_selection_preserves_structured_rows_and_true_contact_identity():
    source = source_for_routing()
    source['procedures']['assessments'] = None
    source['procedures']['visit_schedule'] = [
        'At the first contact collect a sleep diary.',
        {'visit': 'First contact', 'visitNumber': '11', 'procedures': ['Medical history']},
        'Telephone check', 'At the final contact measure fatigue.']
    mapping = schedule_mapping()
    for row in mapping['assignments']: row['source_path'] = 'procedures.visit_schedule'
    mapping['assignments'][1]['visit'] = 'final contact'
    mapping['visits'] = [{'visit': 'final contact', 'source_path': 'procedures.visit_schedule',
                          'source_excerpt': 'At the final contact measure fatigue.'}]
    mapping['narrative_items'] = [{'source_path': 'procedures.visit_schedule', 'source_index': index,
                                   'source_excerpt': source['procedures']['visit_schedule'][index]}
                                  for index in (0, 3)]
    before = copy.deepcopy(source)
    projected = contracts.interpreted_schedule_reference(source, mapping)
    visits = contracts.normalized_visit_records(projected)
    assert [(r['visit'], r['visitNumber']) for r in visits[:2]] == [('First contact', '11'), ('Telephone check', '3')]
    assert visits[0]['procedures'] == ['Medical history', 'sleep diary']
    assert visits[2]['procedures'] == ['fatigue']
    assert source == before
    mapping['narrative_items'][0]['source_index'] = 1  # Cannot replace a typed row.
    with pytest.raises(ValueError): contracts.interpreted_schedule_reference(source, mapping)


def test_editable_approval_preserves_mixed_records_and_narrative_strings():
    source = source_for_routing()
    source['procedures']['visit_schedule'] = [
        {'visit': 'First contact', 'visitNumber': '11', 'procedures': 'Medical history'},
        'Telephone check', 'At the final contact measure fatigue.']
    before = copy.deepcopy(source)
    markdown = contracts.source_truth_markdown(source)
    parsed = contracts.parse_source_truth(markdown, source)
    assert parsed['procedures']['visit_schedule'] == source['procedures']['visit_schedule']
    assert "{'visit':" not in markdown
    assert source == before
    changed = markdown.replace('Medical history', 'Medical history and medication review')
    updated = contracts.parse_source_truth(changed, source)
    assert updated['procedures']['visit_schedule'][0]['procedures'] == 'Medical history and medication review'


def test_editable_approval_rejects_unknown_mixed_entry_type():
    source = source_for_routing()
    source['procedures']['visit_schedule'] = [{'visit': 'Baseline'}, 'Telephone check']
    markdown = contracts.source_truth_markdown(source)
    with pytest.raises(ValueError):
        contracts.parse_source_truth(markdown.replace('| record |', '| unknown |'), source)


def test_repeated_contact_labels_use_existing_visit_identity():
    source = source_for_routing('Ambispective')
    source['procedures']['visit_schedule'] = [
        {'visit': 'Assessment', 'timing': 'Historical baseline', 'visitNumber': '1', 'procedures': []},
        {'visit': 'Assessment', 'timing': 'Prospective follow-up', 'visitNumber': '2', 'procedures': []}]
    source['procedures']['assessments'] = 'Historical baseline: sleep diary. Prospective follow-up: fatigue.'
    mapping = schedule_mapping()
    mapping['assignments'] = [
        {'visit': 'Assessment', 'visit_number': '1', 'activity': 'sleep diary',
         'source_path': 'procedures.assessments', 'source_excerpt': 'Historical baseline: sleep diary.'},
        {'visit': 'Assessment', 'visit_number': '2', 'activity': 'fatigue',
         'source_path': 'procedures.assessments', 'source_excerpt': 'Prospective follow-up: fatigue.'}]
    projected = contracts.interpreted_schedule_reference(source, mapping)
    rows = contracts.protocol_table_contracts(projected)['schedule-of-assessments']['rows']
    assert ['Sleep diary', 'X', ''] in rows
    assert ['Fatigue', '', 'X'] in rows


def test_schedule_interpretation_cannot_drop_an_entire_source_list_item():
    source = source_for_routing()
    source['procedures']['assessments'] = [
        'At the first contact collect a sleep diary.', 'At the final contact measure fatigue.']
    mapping = schedule_mapping()
    mapping['assignments'].pop()
    with pytest.raises(ValueError, match='Every supplied narrative'):
        contracts.interpreted_schedule_reference(source, mapping)


def test_uninterpreted_structured_evidence_is_retained():
    source = source_for_routing()
    source['procedures']['assessments'] = [
        {'activity': 'Unscheduled safety call', 'First contact': '', 'Final contact': ''},
        source['procedures']['assessments']]
    projected = contracts.interpreted_schedule_reference(source, schedule_mapping())
    assert projected['procedures']['assessments'][0] == source['procedures']['assessments'][0]


def test_unlocalizable_source_routing_destination_requires_construction_repair(tmp_path):
    candidate = tmp_path / 'candidate'
    candidate.mkdir()
    Document().save(candidate / 'protocol.docx')
    with pytest.raises(quality.ReviewEvidenceError):
        quality.create_verification_requests(tmp_path, source_for_routing(), {'artifacts': []})


def request_for(tmp_path, source, target):
    branch, family = source['meta']['study_type'], source['meta']['icf_template']
    batch = next(b for b in contracts.batch_plan(branch, family) if target in b.section_ids)
    return json.loads(drafting.create_drafting_request(
        repo_root=ROOT, revision_dir=tmp_path, revision_id=tmp_path.name, reference=source,
        batch=batch, target_ids=[target], attempts={target: 1}, wave='initial').read_text())


@pytest.mark.parametrize('branch,family', BRANCHES[:-1])
def test_authenticated_mapping_reaches_both_protocol_tables_and_quality(tmp_path, branch, family):
    source = source_for_routing(branch, family)
    original = copy.deepcopy(source)
    request = request_for(tmp_path, source, 'evaluation-procedures')
    response = drafting.recorded_acceptance_response(request)
    response['section_results'][0]['schedule_interpretation'] = schedule_mapping()
    path = tmp_path / request['response_path']
    path.write_text(json.dumps(response))
    governing = drafting.governing_resources(ROOT, source)
    assert not drafting.ingest_responses(tmp_path, governing)
    assert drafting.accepted_draft(tmp_path, 'evaluation-procedures', governing)
    model = drafting.merged_drafts(tmp_path, source, governing)
    rendering.render_documents(ROOT, tmp_path, source, model)
    document = Document(tmp_path / 'candidate/protocol.docx')
    table = next(t for t in document.tables if t.cell(0, 0).text.strip() == 'Activity')
    actual = {row.cells[0].text: [c.text for c in row.cells[1:]] for row in table.rows[2:]}
    assert actual['Sleep diary'] == ['X', '']
    assert actual['Fatigue'] == ['', 'X']
    visit_table = next(t for t in document.tables if t.cell(0, 0).text == 'Visit Number')
    assert visit_table.rows[1].cells[1].text == 'First contact'
    assert visit_table.rows[2].cells[1].text == 'Final contact'
    assert not [f for f in quality.deterministic_content_check(tmp_path, source)
                if f.get('surface') == 'table']
    assert json.loads((tmp_path / 'approved-reference.json').read_text()) == original
    assert source == original


def audit_request():
    return {'source_routing_obligations': [{
        'obligation_id': 'benefit-context:icf:icf.benefits', 'kind': 'benefit_context',
        'artifact': 'icf', 'section_id': 'icf.benefits',
        'source_values': {'study.background': 'The diary approach might reduce night waking, but this is unproven.'},
        'artifact_passage': [{'paragraph': 'The diary approach might reduce night waking. This benefit is unproven and is not guaranteed.'}],
    }]}


def audit_response():
    return {'findings': [], 'section_assessments': [
        {'artifact': 'icf', 'section_id': 'icf.benefits', 'status': 'passed'}],
        'source_routing_assessments': [{
            'obligation_id': 'benefit-context:icf:icf.benefits', 'status': 'passed',
            'rationale': 'The possible benefit remains qualified and no benefit is guaranteed.',
            'facts': [{'source_path': 'study.background',
                       'source_excerpt': 'The diary approach might reduce night waking, but this is unproven.',
                       'affected_passage': 'The diary approach might reduce night waking. This benefit is unproven and is not guaranteed.',
                       'disposition': 'preserved'}]}]}


def test_independent_routing_audit_accepts_qualified_paraphrase():
    assert not quality.source_routing_assessment_errors(audit_request(), audit_response())


@pytest.mark.parametrize('mutation', ['bare_pass', 'wrong_section', 'fake_source', 'ignored_defect', 'duplicate', 'malformed_id'])
def test_independent_routing_audit_rejects_missing_or_unbound_evidence(mutation):
    request, response = audit_request(), audit_response()
    row = response['source_routing_assessments'][0]
    if mutation == 'bare_pass': row.pop('facts')
    if mutation == 'wrong_section': row['facts'][0]['affected_passage'] = 'This quote is only in Background.'
    if mutation == 'fake_source': row['facts'][0]['source_excerpt'] = 'The treatment cures insomnia.'
    if mutation == 'ignored_defect': row['facts'][0]['disposition'] = 'defect'
    if mutation == 'duplicate': response['source_routing_assessments'].append(copy.deepcopy(row))
    if mutation == 'malformed_id': row['obligation_id'] = []
    assert quality.source_routing_assessment_errors(request, response)


def routed_response(request):
    response = drafting.recorded_acceptance_response(request)
    for result, contract in zip(response.get('section_results', []), request['section_contracts']):
        if contract.get('schedule_interpretation'):
            authority = contract['schedule_interpretation']
            mapping = schedule_mapping()
            mapping['assignments'] = []
            mapping['visits'] = []
            mapping['narrative_items'] = [item for item in authority['schedule_items']
                                        if item['source_excerpt'] in {
                                            'At the first contact collect a sleep diary.',
                                            'At the final contact measure fatigue.'}]
            for row in schedule_mapping()['assignments']:
                label = row['visit'] if row['visit'] in authority['visits'] else row['visit'].lower()
                for path, value in authority['sources'].items():
                    if any(row['source_excerpt'] in text for text in contracts.source_text_values(value)):
                        assignment = {**row, 'visit': label, 'source_path': path}
                        mapping['assignments'].append(assignment)
                if label not in authority['visits']:
                    assignment = mapping['assignments'][-1]
                    mapping['visits'].append({key: assignment[key] for key in ('visit', 'source_path', 'source_excerpt')})
            for path, value in authority['sources'].items():
                if path in {'procedures.assessments', 'procedures.evaluation'}:
                    for text in contracts.source_text_values(value):
                        if not any(row['source_path'] == path and row['source_excerpt'] in text
                                   for row in mapping['assignments']):
                            mapping['unallocated'].append({'source_path': path, 'source_excerpt': text})
            result['schedule_interpretation'] = mapping
        if contract.get('benefit_context') and request['branch']['study_type'] != 'Retrospective':
            result['outcome'] = 'drafted'
            result['benefit_interpretation'] = {
                'potential_benefits': [{'source_path': 'study.background',
                                       'source_excerpt': 'The new diary approach might reduce night waking, but this is unproven.'}],
                'rationale': 'This is the study-specific possibility; previous exercise results are historical context.'}
            result['paragraphs'].append({
                'text': 'The diary approach might reduce night waking, but this benefit remains unproven. Participants may receive no direct benefit.',
                'evidence_refs': ['source:study.background'], 'boilerplate_refs': []})
    return response


@pytest.mark.parametrize('branch,family', BRANCHES)
def test_public_lifecycle_preserves_routed_facts_and_clean_delivery(tmp_path, monkeypatch, branch, family):
    source = source_for_routing(branch, family)
    result, revision = deliver(tmp_path, monkeypatch, source, draft=routed_response)
    assert result['status'] == 'passed', result
    assert not result.get('warnings')
    approved = json.loads((revision / 'approved-reference.json').read_text())
    assert approved['study']['background'] == source['study']['background']
    assert approved['procedures']['assessments'] == source['procedures']['assessments']
    assert approved['risks_benefits']['benefits'] is None
    for name in contracts.document_set(branch):
        assert (tmp_path / 'output' / name).read_bytes() == (revision / 'candidate' / name).read_bytes()
    if branch != 'Retrospective':
        for artifact, section in [('protocol', 'risks-benefits.benefits'), ('icf', 'icf.benefits')]:
            passage = quality._review_evidence_text(rendering.rendered_section_snapshot(
                revision / f'candidate/{artifact}.docx', section))
            assert 'might reduce night waking' in passage
            assert 'remains unproven' in passage
        doc = Document(revision / 'candidate/protocol.docx')
        table = next(t for t in doc.tables if t.cell(0, 0).text.strip() == 'Activity')
        rows = {row.cells[0].text: [c.text for c in row.cells[1:]] for row in table.rows[2:]}
        assert rows['Sleep diary'] == ['X', '']
        assert rows['Fatigue'] == ['', 'X']


@pytest.mark.parametrize('field', ['assessments', 'visit_schedule', 'mixed_list'])
def test_independent_source_mapping_defect_redrafts_and_rechecks_final_bytes(tmp_path, monkeypatch, field):
    drafts, reviews, approval_bytes = [], [], []
    source = source_for_routing()
    if field == 'visit_schedule':
        source['procedures']['visit_schedule'] = source['procedures'].pop('assessments')
    if field == 'mixed_list':
        source['procedures'].pop('assessments')
        source['procedures']['visit_schedule'] = [
            {'visit': 'First contact', 'visitNumber': '11', 'procedures': ['Medical history']},
            'Telephone check', 'At the first contact collect a sleep diary.',
            'At the final contact measure fatigue.']
    source_path = 'procedures.visit_schedule' if field == 'mixed_list' else f'procedures.{field}'

    def draft(request):
        current_approval = next((tmp_path / 'revisions').glob('*/approved-reference.json')).read_bytes()
        if not approval_bytes: approval_bytes.append(current_approval)
        assert current_approval == approval_bytes[0]
        response = routed_response(request)
        for result in response.get('section_results', []):
            if result['section_id'] == 'evaluation-procedures':
                drafts.append(request['wave'])
                if len(drafts) == 1:
                    # Structurally grounded, but clinically incomplete: the
                    # raw source's second relationship was not interpreted.
                    mapping = result['schedule_interpretation']
                    omitted = [row for row in mapping['assignments'] if row['activity'] == 'fatigue']
                    mapping['assignments'] = [row for row in mapping['assignments'] if row['activity'] != 'fatigue']
                    mapping['unallocated'].extend({key: row[key] for key in ('source_path', 'source_excerpt')} for row in omitted)
        return response

    def review(request):
        response = acceptance_verification(request)
        if request['task'] != 'clinical_content_verification':
            return response
        reviews.append(request['request_id'])
        if len(reviews) > 1:
            return response
        response['status'] = 'blocked'
        target = 'evaluation-procedures'
        next(row for row in response['section_assessments'] if row['section_id'] == target)['status'] = 'failed'
        audit = next(row for row in response['source_routing_assessments']
                     if row['obligation_id'] == 'visit-relationships:protocol:evaluation-procedures')
        audit['status'] = 'failed'
        audit['rationale'] = 'The final-contact fatigue assessment is absent from the table.'
        audit['facts'] = [{'source_path': source_path,
                           'source_excerpt': 'At the final contact measure fatigue.',
                           'affected_passage': 'Sleep diary', 'disposition': 'defect'}]
        response['findings'] = [{
            'finding_id': 'missing-final-fatigue', 'category': 'content', 'check': 'source_supported',
            'artifact': 'protocol', 'surface': 'source_interpretation', 'target_ids': [target],
            'issue': 'Final-contact fatigue assessment is missing from the table.',
            'affected_passage': 'Sleep diary', 'source_excerpt': 'At the final contact measure fatigue.',
            'material': True, 'evidence_diagnostics': [{'source_path': source_path,
                'unobserved_values': [{'source_excerpt': 'At the final contact measure fatigue.'}]}],
        }]
        return response

    result, revision = deliver(tmp_path, monkeypatch, source, draft=draft, review=review)
    assert result['status'] == 'passed', result
    assert (revision / 'approved-reference.json').read_bytes() == approval_bytes[0]
    for name in contracts.document_set(source['meta']['study_type']):
        assert (tmp_path / 'output' / name).read_bytes() == (revision / 'candidate' / name).read_bytes()
    assert len(drafts) == 2
    assert len(reviews) == 2
    assert reviews[0] != reviews[1]
    document = Document(revision / 'candidate/protocol.docx')
    table = next(t for t in document.tables if t.cell(0, 0).text.strip() == 'Activity')
    expected = ['Fatigue', '', '', 'X'] if field == 'mixed_list' else ['Fatigue', '', 'X']
    assert any([c.text for c in row.cells] == expected for row in table.rows)
    if field == 'mixed_list':
        assert len(table.rows[0].cells) == 4
        assert any([c.text for c in row.cells] == ['Medical history', 'X', '', ''] for row in table.rows)
    assert not result.get('warnings')
