"""Clinical tables retain facts while separating display text from approved inputs."""
import copy
import json

import pytest
from docx import Document

import contracts
import drafting
import rendering
from test_shared_output_projection import ROOT, COMBINATIONS, source_for


def assessment_source(branch, family):
    source = source_for(branch, family)
    source['procedures']['visit_schedule'] = [
        {'visit': 'Screening', 'procedures': ['Informed consent']},
        {'visit': 'Week 8', 'procedures': ['Questionnaire', 'Complete Exit Form1', 'SF-36', 'Form 2']},
    ]
    source['procedures']['assessments'] = 'The questionnaire is optional; record why an item is missing.'
    source['safety']['monitoring'] = 'Report adverse events between scheduled contacts as well.'
    return source


@pytest.mark.parametrize('branch,family', COMBINATIONS[1:])
def test_generated_matrix_keeps_unique_notes_as_adjacent_prose(tmp_path, branch, family):
    source = assessment_source(branch, family)
    original = copy.deepcopy(source)
    rendering.render_documents(ROOT, tmp_path, source, {'protocol': [], 'icf': {}, 'prs': {}}, artifact_names={'protocol'})
    document = Document(tmp_path / 'candidate/protocol.docx')
    table = next(t for t in document.tables if t.cell(0, 0).text.strip() == 'Activity')
    cells = '\n'.join(c.text for row in table.rows for c in row.cells)
    assert 'Approved assessment and safety details' not in cells
    for note in (source['procedures']['assessments'], source['safety']['monitoring']):
        paragraph = next(p for p in document.paragraphs if p.text == note)
        assert note not in cells
        assert paragraph._p.getprevious() is table._tbl or paragraph._p.getprevious().tag.endswith('}p')
    assert source == original


@pytest.mark.parametrize('branch,family', COMBINATIONS[1:])
def test_source_linked_display_edits_clean_labels_and_redundant_notes_without_number_stripping(tmp_path, branch, family):
    source = assessment_source(branch, family)
    label = 'Operative visit (one per eye), as listed in the synopsis.'
    source['procedures']['visit_schedule'][1]['procedures'].append(label)
    source['procedures']['assessments'] = None
    source['safety']['monitoring'] = 'Questionnaire at Week 8, as indicated in the source assessment table.'
    original = copy.deepcopy(source)
    presentation = {
        'activity_labels': [
            {'source_text': label, 'display_text': 'Operative visit (one per eye)', 'rationale': 'Remove the source cross-reference; retain each-eye qualification.'},
            {'source_text': 'Complete Exit Form1', 'display_text': 'Complete Exit Form', 'rationale': 'The numeral is an unexplained footnote marker, not a numbered form.'},
        ],
        'notes': [{'source_text': source['safety']['monitoring'], 'display_text': '', 'rationale': 'The same questionnaire and Week 8 relationship is already in the matrix.'}],
    }
    model = {'protocol': [{'section_id': 'evaluation-procedures', 'assessment_presentation': presentation}], 'icf': {}, 'prs': {}}
    rendering.render_documents(ROOT, tmp_path, source, model, artifact_names={'protocol'})
    document = Document(tmp_path / 'candidate/protocol.docx')
    table = next(t for t in document.tables if t.cell(0, 0).text.strip() == 'Activity')
    rows = [[c.text for c in r.cells] for r in table.rows]
    assert ['Operative visit (one per eye)', '', 'X'] in rows
    assert ['Complete Exit Form', '', 'X'] in rows
    assert ['SF-36', '', 'X'] in rows
    assert ['Form 2', '', 'X'] in rows
    assert not any(source['safety']['monitoring'] in p.text for p in document.paragraphs)
    assert source == original


@pytest.mark.parametrize('branch,family', COMBINATIONS)
def test_public_drafting_request_teaches_standalone_source_preserving_wording(tmp_path, branch, family):
    source = assessment_source(branch, family)
    target = 'methods' if branch == 'Retrospective' else 'evaluation-procedures'
    sections = contracts.protocol_contract(branch)
    if not any(s.section_id == target for s in sections):
        target = next(s.section_id for s in sections if s.role != 'container')
    batch = next(b for b in contracts.batch_plan(branch, family or 'Advarra') if target in b.section_ids)
    path = drafting.create_drafting_request(repo_root=ROOT, revision_dir=tmp_path,
        revision_id='presentation', reference=source, batch=batch,
        target_ids=[target], attempts={target: 1}, wave='initial')
    request = json.loads(path.read_text())
    assert any('standalone clinical' in rule for rule in request['constraints'])
    section = request['section_contracts'][0]
    if branch != 'Retrospective':
        contract = section['assessment_presentation']
        assert 'Complete Exit Form1' in contract['activity_labels']
        assert 'uncertainty means retain' in contract['instructions']
        assert 'Independent final review' in contract['instructions']


@pytest.mark.parametrize('branch,family', COMBINATIONS[1:])
def test_authenticated_optional_presentation_survives_ingestion_and_model_assembly(tmp_path, branch, family):
    source = assessment_source(branch, family)
    batch = next(b for b in contracts.batch_plan(branch, family) if 'evaluation-procedures' in b.section_ids)
    request = json.loads(drafting.create_drafting_request(repo_root=ROOT, revision_dir=tmp_path,
        revision_id=tmp_path.name, reference=source, batch=batch,
        target_ids=['evaluation-procedures'], attempts={'evaluation-procedures': 1}, wave='initial').read_text())
    response = drafting.recorded_acceptance_response(request)
    response['section_results'][0]['assessment_presentation'] = {
        'activity_labels': [{'source_text': 'Complete Exit Form1', 'display_text': 'Complete Exit Form',
                             'rationale': 'Source context identifies an orphan marker.'}], 'notes': []}
    path = tmp_path / request['response_path']
    path.write_text(json.dumps(response))
    governing = drafting.governing_resources(ROOT, source)
    assert not drafting.ingest_responses(tmp_path, governing)
    model = drafting.merged_drafts(tmp_path, source, governing)
    rendering.render_documents(ROOT, tmp_path, source, model, artifact_names={'protocol'})
    table = next(t for t in Document(tmp_path / 'candidate/protocol.docx').tables if t.cell(0, 0).text.strip() == 'Activity')
    labels = [r.cells[0].text for r in table.rows[2:]]
    assert 'Complete Exit Form' in labels
    assert 'Complete Exit Form1' not in labels
    assert json.loads((tmp_path / 'approved-reference.json').read_text()) == source


@pytest.mark.parametrize('branch,family', COMBINATIONS[1:])
def test_absent_optional_polish_preserves_explained_footnotes_and_numbered_forms(tmp_path, branch, family):
    source = assessment_source(branch, family)
    source['procedures']['assessments'] = '1 Exit forms are completed only after the final safety assessment.'
    rendering.render_documents(ROOT, tmp_path, source, {'protocol': [], 'icf': {}, 'prs': {}}, artifact_names={'protocol'})
    document = Document(tmp_path / 'candidate/protocol.docx')
    table = next(t for t in document.tables if t.cell(0, 0).text.strip() == 'Activity')
    labels = [r.cells[0].text for r in table.rows[2:]]
    assert 'Complete Exit Form1' in labels
    assert 'SF-36' in labels and 'Form 2' in labels
    assert source['procedures']['assessments'] in [p.text for p in document.paragraphs]


def test_display_edit_cannot_target_a_label_outside_the_source_inventory():
    source = assessment_source('Prospective', 'Advarra')
    with pytest.raises(ValueError, match='exact existing label'):
        contracts.protocol_table_contracts(source, {'activity_labels': [
            {'source_text': 'Invented procedure', 'display_text': 'Questionnaire', 'rationale': 'Invalid source binding.'}]})


def test_cleaned_unique_note_survives_section_body_replacement(tmp_path):
    source = assessment_source('Prospective', 'Advarra')
    original_note = 'The questionnaire is optional, as indicated in the source assessment table.'
    source['procedures']['assessments'] = original_note
    cleaned_note = 'The questionnaire is optional.'
    model = {'protocol': [{'section_id': 'evaluation-procedures',
        'paragraphs': [{'text': 'The scheduled activities appear in Table 15.1.'}],
        'assessment_presentation': {'notes': [{'source_text': original_note,
            'display_text': cleaned_note, 'rationale': 'Keep optionality while removing the internal source reference.'}]}}], 'icf': {}, 'prs': {}}
    rendering.render_documents(ROOT, tmp_path, source, model, artifact_names={'protocol'})
    text = [p.text for p in Document(tmp_path / 'candidate/protocol.docx').paragraphs]
    assert cleaned_note in text
    assert source['safety']['monitoring'] in text


def test_narrative_only_assessments_do_not_reference_a_missing_matrix(tmp_path):
    source = assessment_source('Prospective', 'Sterling')
    source['procedures']['visit_schedule'] = []
    source['procedures']['visit_schedule_table'] = None
    rendering.render_documents(ROOT, tmp_path, source, {'protocol': [], 'icf': {}, 'prs': {}}, artifact_names={'protocol'})
    snapshot = rendering.rendered_section_snapshot(tmp_path / 'candidate/protocol.docx', 'evaluation-procedures')
    assert 'Table 15.1' not in str(snapshot)
    assert source['procedures']['assessments'] in str(snapshot)
