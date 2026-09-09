import json
import re
import zipfile
from pathlib import Path

import pytest
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt
from docx.table import Table
from docx.text.paragraph import Paragraph
from pypdf import PdfReader, PdfWriter
from lxml import etree as LET

from contracts import batch_plan
from drafting import create_drafting_request, evidence_grounded, recorded_acceptance_response, validate_response
from quality import deterministic_content_check, render_pages, sha256_file
from rendering import _has_page_boundary_before, _normalize_protocol_section_pagination, refresh_toc_from_pdf, render_documents, render_fields
import rendering
import workflow


ROOT = Path(__file__).resolve().parents[1]


def _visible_text(document: Document) -> str:
    values = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            values.extend(cell.text for cell in row.cells)
    for section in document.sections:
        values.extend(paragraph.text for paragraph in section.header.paragraphs)
        values.extend(paragraph.text for paragraph in section.footer.paragraphs)
        for container in (section.header, section.footer):
            for table in container.tables:
                for row in table.rows:
                    values.extend(cell.text for cell in row.cells)
    return "\n".join(values)


def _minimal_source_bound_icf_procedures():
    return {
        "paragraphs": [{"text": "You will complete the approved study visits and procedures."}],
        "lists": [],
    }


def _points(value):
    return None if value is None else round(value.pt, 2)


def _paragraph_rhythm(paragraph):
    formatting = paragraph.paragraph_format
    line_spacing = formatting.line_spacing
    if line_spacing is not None and not isinstance(line_spacing, float):
        line_spacing = _points(line_spacing)
    return {
        "space_before": _points(formatting.space_before),
        "space_after": _points(formatting.space_after),
        "line_spacing": line_spacing,
        "left_indent": _points(formatting.left_indent),
        "first_line_indent": _points(formatting.first_line_indent),
    }


def _run_typography(paragraph):
    run = next((item for item in paragraph.runs if item.text.strip()), None)
    if run is None:
        return None
    return {
        "font": run.font.name,
        "size": _points(run.font.size),
        "bold": run.bold,
        "italic": run.italic,
    }


def _doc_default_font(docx_path: Path):
    with zipfile.ZipFile(docx_path) as package:
        styles = package.read("word/styles.xml")
    import xml.etree.ElementTree as ET

    root = ET.fromstring(styles)
    namespace = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
    fonts = root.find("./w:docDefaults/w:rPrDefault/w:rPr/w:rFonts", namespace)
    if fonts is None:
        return None, None
    return fonts.get(qn("w:ascii")), fonts.get(qn("w:hAnsi"))


def _explicit_page_break_count(document: Document) -> int:
    return len(document.element.body.xpath('.//w:pageBreakBefore | .//w:br[@w:type="page"]'))


def _add_page_break(paragraph) -> None:
    run = paragraph.add_run()
    page_break = run._r.makeelement(qn("w:br"), {qn("w:type"): "page"})
    run._r.append(page_break)


def _page_boundary_before(paragraph) -> bool:
    if paragraph.paragraph_format.page_break_before is True:
        return True
    previous = paragraph._p.getprevious()
    while previous is not None and previous.tag == qn("w:p"):
        prior = Paragraph(previous, paragraph._parent)
        if previous.xpath('.//w:br[@w:type="page"]'):
            return True
        if prior.text.strip():
            break
        previous = previous.getprevious()
    return False


def _docx_parts_without_pagination_controls(path: Path):
    with zipfile.ZipFile(path) as package:
        parts = {name: package.read(name) for name in package.namelist()}
    document = LET.fromstring(parts.pop("word/document.xml"))
    for tag in ("keepNext", "keepLines", "widowControl", "pageBreakBefore"):
        for element in document.xpath(f"//w:{tag}", namespaces={"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}):
            element.getparent().remove(element)
    return parts, LET.tostring(document, method="c14n")


def _numbering_level_signature(document: Document, paragraph):
    numbering_properties = paragraph._p.get_or_add_pPr().find(qn("w:numPr"))
    assert numbering_properties is not None
    number_id = int(numbering_properties.find(qn("w:numId")).get(qn("w:val")))
    level_id = int(numbering_properties.find(qn("w:ilvl")).get(qn("w:val")))
    numbering_root = document.part.numbering_part.element
    number = next(
        item for item in numbering_root.findall(qn("w:num"))
        if int(item.get(qn("w:numId"))) == number_id
    )
    abstract_id = int(number.find(qn("w:abstractNumId")).get(qn("w:val")))
    abstract = next(
        item for item in numbering_root.findall(qn("w:abstractNum"))
        if int(item.get(qn("w:abstractNumId"))) == abstract_id
    )
    level = next(
        item for item in abstract.findall(qn("w:lvl"))
        if int(item.get(qn("w:ilvl"))) == level_id
    )
    marker = level.find(qn("w:lvlText"))
    indentation = level.find(f"{qn('w:pPr')}/{qn('w:ind')}")
    fonts = level.find(f"{qn('w:rPr')}/{qn('w:rFonts')}")
    return {
        "marker": None if marker is None else marker.get(qn("w:val")),
        "left": None if indentation is None else indentation.get(qn("w:left")),
        "hanging": None if indentation is None else indentation.get(qn("w:hanging")),
        "font": None if fonts is None else fonts.get(qn("w:ascii")),
    }


def test_recorded_acceptance_coverage_sentences_are_section_specific(tmp_path):
    reference = json.loads((ROOT / "tests/fixtures/prospective-acceptance-source.json").read_text(encoding="utf-8"))
    batch = next(item for item in batch_plan("Prospective") if item.batch_id == "protocol-operations")
    request_path = create_drafting_request(
        repo_root=ROOT,
        revision_dir=tmp_path,
        revision_id="r-section-specific",
        reference=reference,
        batch=batch,
        attempts={section_id: 1 for section_id in batch.section_ids},
        wave="initial",
    )
    request = json.loads(request_path.read_text(encoding="utf-8"))
    response = recorded_acceptance_response(request)
    paragraphs = {
        result["section_id"]: [item["text"] for item in result["paragraphs"]]
        for result in response["section_results"]
    }

    duplicates = {
        text
        for section_paragraphs in paragraphs.values()
        for text in section_paragraphs
        if sum(text in other for other in paragraphs.values()) > 1
    }

    assert duplicates == set()


def test_bundle04_operational_omissions_are_rejected_before_candidate_construction(tmp_path):
    reference = json.loads((ROOT / "tests/fixtures/ambispective-acceptance-source.json").read_text(encoding="utf-8"))
    reference["design"]["intervention_description"] = (
        "An upper-arm wearable sensor used for research measurement that does not direct treatment."
    )
    reference["risks_benefits"]["risk_mitigation"] = (
        "Trained staff insert and remove sensors, skin is assessed at visits, participants receive "
        "contact instructions, research results do not direct treatment, and access is restricted by role."
    )
    reference["risks_benefits"]["costs"] = (
        "The sensor and study-only procedures are free; usual care remains the participant's or insurer's responsibility."
    )

    cases = {
        "protocol-analysis-and-oversight": {
            "financial-injury": (
                "Compensation and injury handling follow the approved terms.",
                ["source:risks_benefits.compensation_or_reimbursement", "source:risks_benefits.injury_handling"],
            ),
            "risks-benefits.risks": (
                "The study risks include the approved discomforts and confidentiality risk.",
                ["source:risks_benefits.risks"],
            ),
        },
        "icf-narrative": {
            "icf.procedures": (
                "You will complete the approved study visits and procedures after the screening interval.",
                [
                    "source:procedures.assessments",
                    "source:procedures.visit_schedule",
                    "source:procedures.minimum_days_before_screening_without_participation",
                ],
            ),
            "icf.risks": (
                "You may experience the approved study risks and discomforts.",
                ["source:risks_benefits.risks"],
            ),
        },
    }
    rejected = set()
    findings = []
    for batch_id, replacements in cases.items():
        batch = next(item for item in batch_plan("Ambispective", "Sterling") if item.batch_id == batch_id)
        request_path = create_drafting_request(
            repo_root=ROOT,
            revision_dir=tmp_path / batch_id,
            revision_id=f"r-{batch_id}",
            reference=reference,
            batch=batch,
            attempts={section_id: 1 for section_id in batch.section_ids},
            wave="initial",
        )
        request = json.loads(request_path.read_text(encoding="utf-8"))
        response = recorded_acceptance_response(request)
        for result in response["section_results"]:
            if result["section_id"] not in replacements:
                continue
            text, evidence_refs = replacements[result["section_id"]]
            result["paragraphs"] = [{
                "text": text,
                "evidence_refs": evidence_refs,
                "boilerplate_refs": [],
            }]
            result["lists"] = []
        accepted, batch_findings = validate_response(request, response)
        accepted_ids = {draft["section_id"] for draft in (accepted or {}).get("drafts", [])}
        rejected.update(set(replacements) - accepted_ids)
        findings.extend(batch_findings)

    expected = {"financial-injury", "risks-benefits.risks", "icf.procedures", "icf.risks"}
    assert rejected == expected
    assert expected <= {item["field"] for item in findings}


@pytest.mark.parametrize(
    "exclusion_items",
    [
        [
            "Type 1 diabetes or gestational diabetes.",
            "Known allergy to medical-grade adhesive.",
            "A skin condition at a proposed sensor site.",
            "Dialysis.",
            "Pregnancy.",
            "Participation in another interventional study within 30 days before screening.",
            "Missing historical source records.",
            "An investigator-determined safety concern.",
        ],
        [
            "Type 1 diabetes or gestational diabetes.",
            "Known allergy to medical-grade adhesive.",
            "A skin condition at a proposed sensor site.",
            "Receiving dialysis.",
            "Being pregnant.",
            "Participation in another interventional study within 30 days before screening.",
            "Missing historical source records.",
            "An investigator-determined safety concern.",
        ],
    ],
)
def test_exclusion_drafts_accept_source_complete_clinical_wording(tmp_path, exclusion_items):
    """Reproduce the exact accepted facts and wording returned by the Hermes run."""
    reference = json.loads((
        ROOT / "tests/fixtures/prospective-acceptance-source.json"
    ).read_text(encoding="utf-8"))
    reference["meta"]["study_type"] = "Ambispective"
    reference["population"]["exclusion_criteria"] = [
        "Type 1 or gestational diabetes",
        "Known allergy to medical-grade adhesive",
        "Skin condition at proposed sensor sites",
        "Dialysis",
        "Pregnancy",
        "Participation in another interventional study within 30 days before screening",
        "Missing historical source records",
        "Investigator-determined safety concern",
    ]
    batch = next(
        item for item in batch_plan("Ambispective")
        if item.batch_id == "protocol-foundations"
    )
    request_path = create_drafting_request(
        repo_root=ROOT,
        revision_dir=tmp_path,
        revision_id="r-hermes-exclusion-regression",
        reference=reference,
        batch=batch,
        attempts={section_id: 1 for section_id in batch.section_ids},
        wave="initial",
    )
    request = json.loads(request_path.read_text(encoding="utf-8"))
    response = recorded_acceptance_response(request)
    exclusion = next(
        item for item in response["section_results"]
        if item["section_id"] == "subjects.exclusion"
    )
    exclusion["paragraphs"] = [{
        "text": "An individual who meets any of the following criteria will be excluded from the study:",
        "evidence_refs": ["source:population.exclusion_criteria"],
        "boilerplate_refs": [],
    }]
    exclusion["lists"] = [{
        "items": exclusion_items,
        "evidence_refs": ["source:population.exclusion_criteria"],
        "boilerplate_refs": [],
    }]

    accepted, findings = validate_response(request, response)

    assert not [
        finding for finding in findings
        if finding.get("field") == "subjects.exclusion"
    ]
    assert any(
        draft["section_id"] == "subjects.exclusion"
        for draft in (accepted or {}).get("drafts", [])
    )


def test_clinical_wording_equivalence_does_not_weaken_fact_or_numeric_grounding():
    assert evidence_grounded("Being pregnant.", "Pregnancy")
    assert evidence_grounded("Pregnancy.", "Pregnancy")
    assert not evidence_grounded("Being present.", "Pregnancy")
    assert not evidence_grounded(
        "Participation in another interventional study before screening.",
        "Participation in another interventional study within 30 days before screening",
    )


def test_participant_facing_adverse_event_assessment_is_grounded():
    assert evidence_grounded(
        "The study team will assess any health problem.",
        "Adverse-event assessment",
        all_items=True,
    )
    assert evidence_grounded(
        "The study team will assess any health problems.",
        "Adverse-event assessment",
        all_items=True,
    )
    assert evidence_grounded(
        "The study team will assess any unwanted effect.",
        "Adverse-event assessment",
        all_items=True,
    )
    assert evidence_grounded(
        "The study team will assess any unwanted effects.",
        "Adverse-event assessment",
        all_items=True,
    )
    assert not evidence_grounded(
        "The study team will review your health information.",
        "Adverse-event assessment",
        all_items=True,
    )


def test_recorded_retrospective_schedule_uses_readable_visit_list(tmp_path):
    reference = json.loads((ROOT / "tests/fixtures/retrospective-acceptance-source.json").read_text(encoding="utf-8"))
    reference["procedures"]["visit_schedule_table"].append({
        "visitNumber": "4",
        "visitName": "Month 6",
        "visitWindow": "±14 days",
        "CRFnumber": "CRF-04",
    })
    batch = next(item for item in batch_plan("Retrospective") if item.batch_id == "protocol-operations")
    request_path = create_drafting_request(
        repo_root=ROOT,
        revision_dir=tmp_path,
        revision_id="r-readable-schedule",
        reference=reference,
        batch=batch,
        attempts={section_id: 1 for section_id in batch.section_ids},
        wave="initial",
    )
    request = json.loads(request_path.read_text(encoding="utf-8"))

    response = recorded_acceptance_response(request)

    enrollment = next(item for item in response["section_results"] if item["section_id"] == "study-procedure.enrollment")
    visible = "\n".join(
        [paragraph["text"] for paragraph in enrollment["paragraphs"]]
        + [item for group in enrollment["lists"] for item in group["items"]]
    )
    assert "approved visit schedule table is" not in visible.casefold()
    assert enrollment["lists"] == [{
        "items": [
            "Visit 1: Baseline (Day 0; CRF BL).",
            "Visit 2: Month 3 (Day 90 +/- 7; CRF M3).",
            "Visit 4: Month 6 (±14 days; CRF-04).",
        ],
        "evidence_refs": ["source:procedures.visit_schedule_table"],
        "boilerplate_refs": [],
    }]


def test_retrospective_eligibility_preserves_approved_criteria_verbatim(tmp_path):
    reference = json.loads((
        ROOT / "tests/fixtures/retrospective-acceptance-source.json"
    ).read_text(encoding="utf-8"))
    reference["population"]["inclusion_criteria"] = (
        "Adults with eligible historical vitrectomy records during the study period."
    )
    reference["population"]["exclusion_criteria"] = (
        "Incomplete records or missing baseline and follow-up visual acuity documentation."
    )
    model = {
        "protocol": [{
            "section_id": "subjects.eligibility",
            "paragraphs": [{
                "text": (
                    "Adults with eligible historical vitrectomy records are included. "
                    "Incomplete records and records missing baseline or follow-up visual acuity "
                    "documentation are excluded."
                ),
                "evidence_refs": [
                    "source:population.inclusion_criteria",
                    "source:population.exclusion_criteria",
                ],
                "boilerplate_refs": [],
            }],
            "lists": [],
        }],
        "icf": {},
        "prs": {},
    }

    render_documents(ROOT, tmp_path, reference, model)

    visible = _visible_text(Document(tmp_path / "candidate/protocol.docx"))
    assert reference["population"]["inclusion_criteria"] in visible
    assert reference["population"]["exclusion_criteria"] in visible
    assert not any(
        finding.get("field") == "subjects.eligibility"
        for finding in deterministic_content_check(tmp_path, reference)
    )


def test_sterling_icf_removes_review_metadata_and_uses_heading_styles(tmp_path):
    reference = json.loads((ROOT / "tests/fixtures/prospective-acceptance-source.json").read_text(encoding="utf-8"))
    reference["meta"]["icf_template"] = "Sterling"

    report = render_documents(ROOT, tmp_path, reference, {
        "protocol": [],
        "icf": {"icf.procedures": _minimal_source_bound_icf_procedures()},
        "prs": {},
    })
    assert report["status"] == "passed"

    output_path = tmp_path / "candidate/icf.docx"
    output = Document(output_path)
    mapped_headings = {"PURPOSE", "DURATION", "PROCEDURES", "POTENTIAL BENEFITS", "COSTS TO YOU"}
    heading_paragraphs = {paragraph.text.strip(): paragraph for paragraph in output.paragraphs if paragraph.text.strip() in mapped_headings}
    assert set(heading_paragraphs) == mapped_headings
    assert all(paragraph.style.name.casefold().startswith("heading") for paragraph in heading_paragraphs.values())
    for index, paragraph in enumerate(output.paragraphs):
        normalized = re.sub(r"\s+", " ", paragraph.text).strip()
        if not normalized.startswith(("COMPENSATION TO YOU", "STUDY COMPLICATIONS", "PARTICIPANT STATEMENT")):
            continue
        if index and output.paragraphs[index - 1].text.strip():
            assert paragraph.paragraph_format.space_before is not None
            assert paragraph.paragraph_format.space_before >= Pt(6)

    with zipfile.ZipFile(output_path) as package:
        names = set(package.namelist())
        assert not {name for name in names if "comments" in name.casefold() or name == "word/people.xml"}
        package_xml = "\n".join(
            package.read(name).decode("utf-8", errors="ignore")
            for name in names
            if name.endswith((".xml", ".rels"))
        )
    assert "trackRevisions" not in package_xml
    assert "commentRangeStart" not in package_xml
    assert "commentReference" not in package_xml


def test_rendering_uses_prs_provider_study_id_when_protocol_number_is_absent(tmp_path):
    reference = json.loads((ROOT / "tests/fixtures/prospective-acceptance-source.json").read_text(encoding="utf-8"))
    reference["meta"].pop("protocol_number", None)
    reference.setdefault("regulatory", {}).setdefault("prs", {})["provider_study_id"] = "AS-SP-001"
    reference["meta"]["icf_template"] = "Sterling"

    report = render_documents(ROOT, tmp_path, reference, {
        "protocol": [],
        "icf": {"icf.procedures": _minimal_source_bound_icf_procedures()},
        "prs": {},
    })

    assert report["status"] == "passed"
    assert "AS-SP-001" in _visible_text(Document(tmp_path / "candidate/protocol.docx"))
    assert "AS-SP-001" in _visible_text(Document(tmp_path / "candidate/icf.docx"))


def test_sterling_icf_replaces_example_shell_policies_with_authorized_content(tmp_path):
    reference = json.loads((ROOT / "tests/fixtures/prospective-acceptance-source.json").read_text(encoding="utf-8"))
    reference["meta"]["icf_template"] = "Sterling"
    risk_text = (
        "Taking part may involve inconvenience or discomfort from the study procedures described in this consent form. "
        "There is also a risk that private information could be disclosed, although safeguards will be used to protect it."
    )
    privacy_text = (
        "By signing this form, you authorize the study team to collect, use, and disclose the medical and study information "
        "described in this form for this research. Records will use appropriate identifiers and access controls."
    )
    model = {
        "protocol": [],
        "prs": {},
        "icf": {
            "icf.risks": {"paragraphs": [{"text": risk_text}], "lists": []},
            "icf.privacy": {"paragraphs": [{"text": privacy_text}], "lists": []},
        },
    }

    render_documents(ROOT, tmp_path, reference, model)
    visible = _visible_text(Document(tmp_path / "candidate/icf.docx"))

    assert visible.count(risk_text) == 2  # key information and the full risks section
    assert "risks or inconveniences that are currently unknown" not in visible
    assert "in a timely manner" not in visible
    assert "Sterling Institutional Review Board" not in visible
    assert "Example IRB" in visible
    assert "555-0100" in visible
    assert "irb@example.org" in visible
    assert "no new routine research procedures will be performed unless you separately agree" in visible
    assert "the FDA" not in visible


def test_protocol_removes_unsupported_template_governance_claims(tmp_path):
    reference = json.loads((ROOT / "tests/fixtures/prospective-acceptance-source.json").read_text(encoding="utf-8"))

    render_documents(ROOT, tmp_path, reference, {"protocol": [], "icf": {}, "prs": {}})
    visible = _visible_text(Document(tmp_path / "candidate/protocol.docx")).casefold()

    assert "will be registered with clinicaltrials.gov" not in visible
    assert "will be conducted in compliance with the protocol, gcp" not in visible
    assert "the study will be conducted under the approved protocol" in visible


def test_sterling_icf_does_not_duplicate_flat_site_address(tmp_path):
    reference = json.loads((ROOT / "tests/fixtures/prospective-acceptance-source.json").read_text(encoding="utf-8"))
    reference["meta"]["icf_template"] = "Sterling"
    reference["sites"][0]["facility"] = {
        "name": "North Surgical Research Center",
        "address": "220 Clinic Way",
        "city": "Boston",
        "state": "MA",
        "country": "United States",
    }

    model = {"protocol": [], "icf": {}, "prs": {}}
    fields = render_fields(reference, model)
    render_documents(ROOT, tmp_path, reference, model)
    visible = _visible_text(Document(tmp_path / "candidate/icf.docx"))

    assert fields["facilityAddress"] == "220 Clinic Way"
    assert fields["facilityLocation"] == "Boston, MA, United States"
    assert "Boston, MA, United States" in visible


def test_protocol_general_information_preserves_every_supplied_site_and_full_address(tmp_path):
    reference = json.loads((ROOT / "tests/fixtures/prospective-acceptance-source.json").read_text(encoding="utf-8"))
    reference["sites"] = [
        {
            "site_id": "01",
            "facility": {
                "name": "Triangle Orthopedic Research Center",
                "address": "1010 Recovery Drive",
                "city": "Raleigh",
                "state": "NC",
                "country": "United States",
            },
        },
        {
            "site_id": "02",
            "facility": {
                "name": "Piedmont Joint Health Institute",
                "address": "880 Clinical Avenue",
                "city": "Greensboro",
                "state": "NC",
                "country": "United States",
            },
        },
        {
            "site_id": "03",
            "facility": {
                "name": "Complete Address Clinic",
                "address": "101 Main St, Raleigh, NC 27601, United States",
                "city": "Raleigh",
                "state": "NC",
                "zip": "27601",
                "country": "United States",
            },
        },
        {
            "site_id": "04",
            "facility": {
                "name": "Orange Recovery Center",
                "address": "100 Orange Street",
                "city": "Orange",
                "state": "CA",
                "zip": "92868",
                "country": "United States",
            },
        },
        {
            "site_id": "05",
            "facility": {
                "name": "国际康复中心",
                "address": "100 Main Street",
                "city": "北京",
                "state": "北京市",
                "country": "中国",
            },
        },
        {
            "site_id": "06",
            "facility": {
                "name": "Capital Recovery Center",
                "address": {"line1": "200 Main Street", "line2": "Suite CA"},
                "city": "Sacramento",
                "state": "CA",
                "zip": "95814",
                "country": "United States",
            },
        },
    ]
    reference["design"]["number_of_sites"] = 6
    original = json.loads(json.dumps(reference))

    render_documents(
        ROOT,
        tmp_path,
        reference,
        {"protocol": [], "icf": {}, "prs": {}},
        artifact_names={"protocol"},
    )

    document = Document(tmp_path / "candidate/protocol.docx")
    summary = next(
        table for table in document.tables
        if table.rows and table.rows[0].cells[0].text.strip() == "Objective"
    )
    visible_summary = "\n".join(cell.text for row in summary.rows for cell in row.cells)
    assert "Triangle Orthopedic Research Center, 1010 Recovery Drive, Raleigh, NC, United States" in visible_summary
    assert "Piedmont Joint Health Institute, 880 Clinical Avenue, Greensboro, NC, United States" in visible_summary
    complete = "Complete Address Clinic, 101 Main St, Raleigh, NC 27601, United States"
    site_row = next(row for row in summary.rows if row.cells[0].text.strip() == "Study sites")
    assert site_row.cells[1].text.splitlines() == [
        "Triangle Orthopedic Research Center, 1010 Recovery Drive, Raleigh, NC, United States",
        "Piedmont Joint Health Institute, 880 Clinical Avenue, Greensboro, NC, United States",
        complete,
        "Orange Recovery Center, 100 Orange Street, Orange, CA, United States, 92868",
        "国际康复中心, 100 Main Street, 北京, 北京市, 中国",
        "Capital Recovery Center, 200 Main Street, Suite CA, Sacramento, CA, United States, 95814",
    ]
    assert reference == original


@pytest.mark.parametrize("icf_template", ["Advarra", "Sterling"])
def test_icf_contacts_preserve_supplied_study_doctor_phone_with_coordinator_and_irb(tmp_path, icf_template):
    reference = json.loads((ROOT / "tests/fixtures/prospective-acceptance-source.json").read_text(encoding="utf-8"))
    reference["meta"]["icf_template"] = icf_template
    reference["parties"]["principal_investigator"] = {
        "name": "Elena Marquez",
        "title": "MD",
        "phone": "919-555-0140",
    }
    reference["parties"]["study_coordinator"]["business_phone"] = "919-555-0141"
    reference["parties"]["irb"]["phone"] = "800-555-0188"
    reference["risks_benefits"]["injury_handling"] = (
        "For an immediate medical concern, participants should seek appropriate clinical care "
        "and contact the study doctor."
    )
    original = json.loads(json.dumps(reference))

    render_documents(
        ROOT,
        tmp_path,
        reference,
        {"protocol": [], "icf": {}, "prs": {}},
        artifact_names={"icf"},
    )

    visible = " ".join(_visible_text(Document(tmp_path / "candidate/icf.docx")).split())
    assert "Elena Marquez" in visible
    assert "919-555-0140" in visible
    assert "919-555-0141" in visible
    assert "800-555-0188" in visible
    assert "24-hour" not in visible.casefold()
    assert "around-the-clock" not in visible.casefold()
    assert reference == original


def test_sterling_icf_uses_real_page_and_page_count_fields_in_every_footer(tmp_path):
    reference = json.loads((ROOT / "tests/fixtures/prospective-acceptance-source.json").read_text(encoding="utf-8"))
    reference["meta"]["icf_template"] = "Sterling"

    render_documents(ROOT, tmp_path, reference, {"protocol": [], "icf": {}, "prs": {}})

    with zipfile.ZipFile(tmp_path / "candidate/icf.docx") as package:
        footers = [package.read(name) for name in package.namelist() if name.startswith("word/footer") and name.endswith(".xml")]

    assert len(footers) == 2
    assert all(b"<w:pgNum" not in footer for footer in footers)
    assert all(b"> PAGE \\* arabic \\* MERGEFORMAT </w:instrText>" in footer for footer in footers)
    assert all(b"> NUMPAGES \\* arabic \\* MERGEFORMAT </w:instrText>" in footer for footer in footers)


def test_sterling_icf_retains_every_client_shell_section_and_signature_block(tmp_path):
    reference = json.loads((ROOT / "tests/fixtures/prospective-acceptance-source.json").read_text(encoding="utf-8"))
    reference["meta"]["icf_template"] = "Sterling"
    model = {
        "protocol": [],
        "prs": {},
        "icf": {
            section_id: {"paragraphs": [{"text": f"Approved content for {section_id}."}], "lists": []}
            for section_id in ("icf.costs", "icf.alternatives", "icf.privacy", "icf.injury")
        },
    }

    render_documents(ROOT, tmp_path, reference, model)
    visible = " ".join(_visible_text(Document(tmp_path / "candidate/icf.docx")).split())

    for title in (
        "AUTHORIZATION TO USE AND DISCLOSE MEDICAL INFORMATION",
        "KEY INFORMATION",
        "BACKGROUND",
        "INFORMATION",
        "VOLUNTARY PARTICIPATION/WITHDRAWAL",
        "QUESTIONS",
        "PARTICIPANT STATEMENT AUTHORIZATION",
        "Signature of Participant",
    ):
        assert title in visible


def test_sterling_icf_keeps_withdrawal_heading_with_body_and_avoids_forced_signature_page(tmp_path):
    reference = json.loads((ROOT / "tests/fixtures/prospective-acceptance-source.json").read_text(encoding="utf-8"))
    reference["meta"]["icf_template"] = "Sterling"

    render_documents(ROOT, tmp_path, reference, {"protocol": [], "icf": {}, "prs": {}})
    document = Document(tmp_path / "candidate/icf.docx")

    withdrawal = next(paragraph for paragraph in document.paragraphs if paragraph.text.strip() == "VOLUNTARY PARTICIPATION/WITHDRAWAL")
    assert withdrawal.paragraph_format.keep_with_next is True
    assert withdrawal.paragraph_format.keep_together is True
    withdrawal_index = next(index for index, paragraph in enumerate(document.paragraphs) if paragraph._p is withdrawal._p)
    spacer = document.paragraphs[withdrawal_index + 1]
    assert not spacer.text.strip()
    assert spacer.paragraph_format.keep_with_next is True
    assert spacer.paragraph_format.keep_together is True
    questions = next(paragraph for paragraph in document.paragraphs if paragraph.text.strip() == "QUESTIONS")
    assert questions.paragraph_format.page_break_before is not True

    signature_index = next(
        index
        for index, paragraph in enumerate(document.paragraphs)
        if " ".join(paragraph.text.split()) == "PARTICIPANT STATEMENT AUTHORIZATION"
    )
    signature = document.paragraphs[signature_index]
    assert not signature._p.xpath('.//w:br[@w:type="page"]')
    signature_tail = document.paragraphs[signature_index:]
    assert all(paragraph.paragraph_format.keep_with_next is True for paragraph in signature_tail[:-1])
    assert all(paragraph.paragraph_format.keep_together is True for paragraph in signature_tail)


def test_prospective_advarra_restores_source_bound_injury_section_before_legal_rights(tmp_path):
    reference = json.loads((ROOT / "tests/fixtures/prospective-acceptance-source.json").read_text(encoding="utf-8"))
    model = {
        "protocol": [],
        "prs": {},
        "icf": {
            "icf.injury": {
                "paragraphs": [{"text": "This draft must not create a template section."}],
                "lists": [],
            }
        },
    }

    render_documents(ROOT, tmp_path, reference, model)
    visible = _visible_text(Document(tmp_path / "candidate/icf.docx"))

    assert visible.count("IN CASE OF AN INJURY RELATED TO THIS RESEARCH STUDY") == 1
    assert "This draft must not create a template section." not in visible
    assert "The above statement" not in visible
    assert "You do not lose any legal rights by signing and dating this consent document." in visible
    injury = reference["risks_benefits"]["injury_handling"]
    assert visible.count(injury) == 1
    assert visible.index("IN CASE OF AN INJURY RELATED TO THIS RESEARCH STUDY") < visible.index(injury)
    assert visible.index(injury) < visible.index("LEGAL RIGHTS")
    assert visible.index("LEGAL RIGHTS") < visible.index(
        "You do not lose any legal rights by signing and dating this consent document."
    )
    assert visible.index(
        "You do not lose any legal rights by signing and dating this consent document."
    ) < visible.index("WHOM TO CONTACT ABOUT THIS STUDY")
    output = Document(tmp_path / "candidate/icf.docx")
    authority = Document(ROOT / "assets/client-templates/reference/advarra-icf-reference.docx")
    output_heading = next(
        paragraph for paragraph in output.paragraphs
        if " ".join(paragraph.text.split()) == "IN CASE OF AN INJURY RELATED TO THIS RESEARCH STUDY"
    )
    authority_heading = next(
        paragraph for paragraph in authority.paragraphs
        if " ".join(paragraph.text.split()) == "IN CASE OF AN INJURY RELATED TO THIS RESEARCH STUDY"
    )
    assert _run_typography(output_heading) == _run_typography(authority_heading)


def test_ambispective_advarra_places_approved_injury_handling_before_legal_rights(tmp_path):
    reference = json.loads((ROOT / "tests/fixtures/ambispective-acceptance-source.json").read_text(encoding="utf-8"))

    render_documents(ROOT, tmp_path, reference, {"protocol": [], "prs": {}, "icf": {}})
    visible = _visible_text(Document(tmp_path / "candidate/icf.docx"))

    injury = reference["risks_benefits"]["injury_handling"]
    assert visible.count("IN CASE OF AN INJURY RELATED TO THIS RESEARCH STUDY") == 1
    assert visible.count(injury) == 1
    assert visible.index("IN CASE OF AN INJURY RELATED TO THIS RESEARCH STUDY") < visible.index(injury)
    assert visible.index(injury) < visible.index("LEGAL RIGHTS")


def test_shallow_section_draft_cannot_pass_content_depth_gate(tmp_path):
    reference = json.loads((ROOT / "tests/fixtures/prospective-acceptance-source.json").read_text(encoding="utf-8"))
    batch = next(item for item in batch_plan("Prospective") if item.batch_id == "protocol-foundations")
    request_path = create_drafting_request(
        repo_root=ROOT,
        revision_dir=tmp_path,
        revision_id="r-depth",
        reference=reference,
        batch=batch,
        attempts={section_id: 1 for section_id in batch.section_ids},
        wave="initial",
    )
    request = json.loads(request_path.read_text(encoding="utf-8"))
    response = recorded_acceptance_response(request)
    introduction = next(item for item in response["section_results"] if item["section_id"] == "introduction")
    introduction["paragraphs"] = [
        {
            **introduction["paragraphs"][0],
            "text": "This study evaluates the approved clinical objective using the stated design.",
        }
    ]
    introduction["lists"] = []

    accepted, findings = validate_response(request, response)
    accepted_ids = {draft["section_id"] for draft in (accepted or {}).get("drafts", [])}

    assert "introduction" not in accepted_ids
    assert any(item["field"] == "introduction" for item in findings)


def test_completion_drafts_omitting_approved_follow_up_visits_are_rejected(tmp_path):
    reference = json.loads((
        ROOT / "tests/fixtures/release-certification/prospective-advarra/approved-reference.json"
    ).read_text(encoding="utf-8"))
    batch = next(item for item in batch_plan("Prospective", "Advarra") if item.batch_id == "protocol-operations")
    request_path = create_drafting_request(
        repo_root=ROOT,
        revision_dir=tmp_path,
        revision_id="r-completion-coverage",
        reference=reference,
        batch=batch,
        attempts={section_id: 1 for section_id in batch.section_ids},
        wave="initial",
    )
    request = json.loads(request_path.read_text(encoding="utf-8"))
    response = recorded_acceptance_response(request)
    completion_ids = {"endpoint-criteria.completion", "endpoint-criteria.study-completion"}
    for result in response["section_results"]:
        if result["section_id"] not in completion_ids:
            continue
        subject = "participant" if result["section_id"] == "endpoint-criteria.completion" else "study"
        result["paragraphs"] = [{
            "text": f"The {subject} reaches completion after the Baseline visit on Day 0 within the 3-month timeline.",
            "evidence_refs": [
                "source:study.timeline",
                "source:procedures.visit_schedule",
                "source:procedures.assessments",
            ],
            "boilerplate_refs": [],
        }]
        result["lists"] = []

    accepted, findings = validate_response(request, response)
    accepted_ids = {draft["section_id"] for draft in (accepted or {}).get("drafts", [])}

    assert not completion_ids & accepted_ids
    assert completion_ids <= {item["field"] for item in findings}
    assert all(
        any("material facts are not observable" in item["issue"] for item in findings if item["field"] == section_id)
        for section_id in completion_ids
    )


def test_retrospective_objectives_omitting_secondary_objective_and_hypothesis_are_rejected(tmp_path):
    reference = json.loads((
        ROOT / "tests/fixtures/release-certification/retrospective/approved-reference.json"
    ).read_text(encoding="utf-8"))
    batch = next(item for item in batch_plan("Retrospective") if item.batch_id == "protocol-foundations")
    request_path = create_drafting_request(
        repo_root=ROOT,
        revision_dir=tmp_path,
        revision_id="r-retrospective-objective-coverage",
        reference=reference,
        batch=batch,
        attempts={section_id: 1 for section_id in batch.section_ids},
        wave="initial",
    )
    request = json.loads(request_path.read_text(encoding="utf-8"))
    response = recorded_acceptance_response(request)
    objectives = next(item for item in response["section_results"] if item["section_id"] == "objectives")
    objectives["paragraphs"] = [{
        "text": "The primary objective is to describe recovery outcomes, measured as the Primary outcome at Month 3.",
        "evidence_refs": [
            f"source:{path}"
            for path in next(
                item for item in request["section_contracts"] if item["section_id"] == "objectives"
            )["minimum_evidence"]
        ],
        "boilerplate_refs": [],
    }]
    objectives["lists"] = []

    accepted, findings = validate_response(request, response)
    accepted_ids = {draft["section_id"] for draft in (accepted or {}).get("drafts", [])}

    assert "objectives" not in accepted_ids
    assert any(
        item["field"] == "objectives" and "material facts are not observable" in item["issue"]
        for item in findings
    )


def _retrospective_safety_request(tmp_path, reference, revision_id):
    batch = next(
        item for item in batch_plan("Retrospective")
        if item.batch_id == "protocol-analysis-and-oversight"
    )
    request_path = create_drafting_request(
        repo_root=ROOT,
        revision_dir=tmp_path,
        revision_id=revision_id,
        reference=reference,
        batch=batch,
        attempts={section_id: 1 for section_id in batch.section_ids},
        wave="initial",
    )
    return json.loads(request_path.read_text(encoding="utf-8"))


def _replace_safety_result(response, text):
    safety_result = next(
        item for item in response["section_results"]
        if item["section_id"] == "quality-safety"
    )
    safety_result["paragraphs"] = [{
        "text": text,
        "evidence_refs": ["source:safety.roles", "source:risks_benefits.risks"],
        "boilerplate_refs": [],
    }]
    safety_result["lists"] = []


def test_retrospective_safety_contract_binds_structured_approved_parties(tmp_path):
    reference = json.loads((
        ROOT / "tests/fixtures/release-certification/retrospective/approved-reference.json"
    ).read_text(encoding="utf-8"))
    request = _retrospective_safety_request(tmp_path, reference, "r-retrospective-safety-role")
    safety = next(
        item for item in request["section_contracts"]
        if item["section_id"] == "quality-safety"
    )

    assert "safety.roles" in safety["minimum_evidence"]
    safety_role = next(
        item for item in request["approved_input"]
        if item["path"] == "safety.roles"
    )
    assert safety_role["value"] == [{
        "party": "investigator",
        "responsibilities": "assess_safety_events; report_safety_events",
    }]

    response = recorded_acceptance_response(request)
    _replace_safety_result(
        response,
        "Safety events are assessed and reported. Risks include privacy loss from chart review.",
    )

    accepted, findings = validate_response(request, response)
    accepted_ids = {draft["section_id"] for draft in (accepted or {}).get("drafts", [])}

    assert "quality-safety" not in accepted_ids
    assert any(
        item["field"] == "quality-safety" and "approved safety role" in item["issue"]
        for item in findings
    )


def test_retrospective_safety_contract_requires_every_structured_party_but_allows_paraphrase(tmp_path):
    reference = json.loads((
        ROOT / "tests/fixtures/release-certification/retrospective/approved-reference.json"
    ).read_text(encoding="utf-8"))
    reference["safety"]["roles"] = [
        {"party": "study physician", "responsibilities": "assess_safety_events"},
        {"party": "sponsor", "responsibilities": "report_safety_events"},
    ]
    request = _retrospective_safety_request(
        tmp_path,
        reference,
        "r-retrospective-structured-safety-roles",
    )
    response = recorded_acceptance_response(request)
    _replace_safety_result(
        response,
        "The study physician evaluates adverse events. Approved risks include device discomfort and privacy risks.",
    )

    accepted, findings = validate_response(request, response)
    accepted_ids = {draft["section_id"] for draft in (accepted or {}).get("drafts", [])}

    assert "quality-safety" not in accepted_ids
    assert any(
        item["field"] == "quality-safety" and "approved safety role" in item["issue"]
        for item in findings
    )

    swapped = recorded_acceptance_response(request)
    _replace_safety_result(
        swapped,
        (
            "The study physician reports adverse events while the sponsor evaluates them. "
            "Approved risks include device discomfort and privacy risks."
        ),
    )

    accepted_swapped, swapped_findings = validate_response(request, swapped)
    accepted_swapped_ids = {
        draft["section_id"] for draft in (accepted_swapped or {}).get("drafts", [])
    }

    assert "quality-safety" not in accepted_swapped_ids
    assert any(
        item["field"] == "quality-safety" and "approved safety role" in item["issue"]
        for item in swapped_findings
    )

    negated = recorded_acceptance_response(request)
    _replace_safety_result(
        negated,
        (
            "The study physician evaluates adverse events. "
            "The sponsor does not report those events. "
            "Approved risks include device discomfort and privacy risks."
        ),
    )

    accepted_negated, negated_findings = validate_response(request, negated)
    accepted_negated_ids = {
        draft["section_id"] for draft in (accepted_negated or {}).get("drafts", [])
    }

    assert "quality-safety" not in accepted_negated_ids
    assert any(
        item["field"] == "quality-safety" and "approved safety role" in item["issue"]
        for item in negated_findings
    )

    invalid_assignment_phrasings = [
        (
            "The study physician evaluates adverse events. "
            "Safety events are reported to the sponsor. "
            "Approved risks include device discomfort and privacy risks."
        ),
        (
            "The study physician evaluates adverse events. "
            "The sponsor receives reports of adverse events. "
            "Approved risks include device discomfort and privacy risks."
        ),
        (
            "The study physician evaluates adverse events. "
            "The sponsor is explicitly not considered to be responsible for reporting adverse events. "
            "Approved risks include device discomfort and privacy risks."
        ),
        (
            "The study physician evaluates adverse events. "
            "The sponsor reviews reports of adverse events. "
            "Approved risks include device discomfort and privacy risks."
        ),
        (
            "The study physician evaluates adverse events. "
            "The sponsor reports enrollment metrics. "
            "Approved risks include device discomfort and privacy risks."
        ),
        (
            "The study physician evaluates adverse events. "
            "The sponsor evaluates and reports adverse events. "
            "Approved risks include device discomfort and privacy risks."
        ),
        (
            "The study physician evaluates adverse events. "
            "The sponsor reports enrollment metrics and is aware of safety events. "
            "Approved risks include device discomfort and privacy risks."
        ),
        (
            "The study physician evaluates adverse events. "
            "After safety events were reviewed, enrollment metrics were reported by the sponsor. "
            "Approved risks include device discomfort and privacy risks."
        ),
        (
            "The study physician evaluates adverse events. "
            "The sponsor reports enrollment metrics to the safety committee. "
            "Approved risks include device discomfort and privacy risks."
        ),
        (
            "The study physician evaluates adverse events. "
            "The sponsor receives enrollment data and reports it. "
            "Approved risks include device discomfort and privacy risks."
        ),
        (
            "The study physician evaluates adverse events. "
            "Enrollment metrics rather than safety events were reported by the sponsor. "
            "Approved risks include device discomfort and privacy risks."
        ),
        (
            "The study physician evaluates adverse events. "
            "Enrollment metrics instead of safety events were reported by the sponsor. "
            "Approved risks include device discomfort and privacy risks."
        ),
        (
            "The study physician evaluates adverse events. "
            "The sponsor reports safety events. The IRB reports safety events. "
            "Approved risks include device discomfort and privacy risks."
        ),
        (
            "The study physician evaluates adverse events. "
            "The sponsor and IRB report safety events. "
            "Approved risks include device discomfort and privacy risks."
        ),
        (
            "The study physician evaluates adverse events. "
            "Safety events are reported by the sponsor and IRB. "
            "Approved risks include device discomfort and privacy risks."
        ),
        (
            "The study physician evaluates adverse events. "
            "The sponsor reports safety events. Safety events are reviewed by the IRB. "
            "Approved risks include device discomfort and privacy risks."
        ),
        (
            "The study physician evaluates adverse events. "
            "The sponsor or IRB reports safety events. "
            "Approved risks include device discomfort and privacy risks."
        ),
        (
            "The study physician evaluates adverse events. "
            "The sponsor as well as the IRB reports safety events. "
            "Approved risks include device discomfort and privacy risks."
        ),
        (
            "The study physician evaluates adverse events. "
            "The sponsor together with the IRB reports safety events. "
            "Approved risks include device discomfort and privacy risks."
        ),
        (
            "The study physician evaluates adverse events. "
            "The sponsor & IRB report safety events. "
            "Approved risks include device discomfort and privacy risks."
        ),
        (
            "The study physician evaluates adverse events. "
            "The sponsor and Emily report safety events. "
            "Approved risks include device discomfort and privacy risks."
        ),
        (
            "The study physician evaluates adverse events. "
            "The sponsor and family report safety events. "
            "Approved risks include device discomfort and privacy risks."
        ),
        (
            "The study physician evaluates adverse events. "
            "The sponsor reports safety events alongside IRB. "
            "Approved risks include device discomfort and privacy risks."
        ),
        (
            "The study physician evaluates adverse events. "
            "Safety events are reported by the sponsor and Emily. "
            "Approved risks include device discomfort and privacy risks."
        ),
        (
            "The study physician evaluates adverse events. "
            "Safety events are reported by the sponsor and family. "
            "Approved risks include device discomfort and privacy risks."
        ),
        (
            "The study physician evaluates adverse events. "
            "The sponsor and the clinical study report safety events. "
            "Approved risks include device discomfort and privacy risks."
        ),
        (
            "The study physician evaluates adverse events. "
            "The sponsor and Will report safety events. "
            "Approved risks include device discomfort and privacy risks."
        ),
        (
            "The study physician evaluates adverse events. "
            "The sponsor and Doe report safety events. "
            "Approved risks include device discomfort and privacy risks."
        ),
        (
            "The study physician evaluates adverse events. "
            "Alongside IRB the sponsor reports safety events. "
            "Approved risks include device discomfort and privacy risks."
        ),
        (
            "The study physician evaluates adverse events. "
            "IRB alongside the sponsor reports safety events. "
            "Approved risks include device discomfort and privacy risks."
        ),
        (
            "The study physician evaluates adverse events. "
            "The sponsor plus IRB reports safety events. "
            "Approved risks include device discomfort and privacy risks."
        ),
        (
            "The study physician evaluates adverse events. "
            "The sponsor reports safety events and IRB does too. "
            "Approved risks include device discomfort and privacy risks."
        ),
        (
            "The study physician evaluates adverse events. "
            "The sponsor assesses safety events and IRB reports them. "
            "Approved risks include device discomfort and privacy risks."
        ),
        (
            "The study physician evaluates adverse events. "
            "The sponsor reports safety events and Will does. "
            "Approved risks include device discomfort and privacy risks."
        ),
        (
            "The study physician evaluates adverse events. "
            "The sponsor reports safety events and Doe does. "
            "Approved risks include device discomfort and privacy risks."
        ),
        (
            "The study physician evaluates adverse events. "
            "The sponsor reports safety events and Will reports them. "
            "Approved risks include device discomfort and privacy risks."
        ),
        (
            "The study physician evaluates adverse events. "
            "The sponsor reports safety events and the study reports them. "
            "Approved risks include device discomfort and privacy risks."
        ),
        (
            "The study physician evaluates adverse events. "
            "The sponsor reporting safety events. "
            "Approved risks include device discomfort and privacy risks."
        ),
        (
            "The study physician evaluates adverse events. "
            "The sponsor communication safety events. "
            "Approved risks include device discomfort and privacy risks."
        ),
        (
            "The study physician evaluates adverse events. "
            "The sponsor reports safety events. "
            "The IRB is responsible for safety-event management. "
            "Approved risks include device discomfort and privacy risks."
        ),
        (
            "The study physician evaluates adverse events. "
            "The sponsor reports safety events. "
            "The IRB communication of safety events. "
            "Approved risks include device discomfort and privacy risks."
        ),
        (
            "The study physician evaluates adverse events. "
            "The sponsor reports safety events. "
            "The IRB handles safety. "
            "Approved risks include device discomfort and privacy risks."
        ),
        (
            "The study physician evaluates adverse events. "
            "The sponsor reports safety events. "
            "The IRB is the safety authority. "
            "Approved risks include device discomfort and privacy risks."
        ),
        (
            "The study physician evaluates adverse events. "
            "The sponsor reports safety events. "
            "The sponsor reports enrollment metrics. "
            "Approved risks include device discomfort and privacy risks."
        ),
        (
            "The study physician evaluates adverse events. "
            "The sponsor reports safety events. "
            "The IRB oversees AE reporting. "
            "Approved risks include device discomfort and privacy risks."
        ),
        (
            "The study physician evaluates adverse events. "
            "The sponsor reports safety events. "
            "The IRB oversees safety reporting. "
            "Approved risks include device discomfort and privacy risks."
        ),
        (
            "The study physician evaluates adverse events. "
            "The sponsor reports safety events. "
            "The sponsor monitors AEs. "
            "Approved risks include device discomfort and privacy risks."
        ),
        (
            "The study physician evaluates adverse events. "
            "The sponsor reports safety events. "
            "The IRB monitors safety. "
            "Approved risks include device discomfort and privacy risks."
        ),
        (
            "The study physician evaluates adverse events. "
            "The sponsor reports safety events. "
            "The IRB reviews safety. "
            "Approved risks include device discomfort and privacy risks."
        ),
        (
            "The study physician evaluates adverse events. "
            "The sponsor reports safety events. "
            "The IRB reports safety. "
            "Approved risks include device discomfort and privacy risks."
        ),
        (
            "The study physician evaluates adverse events. "
            "The sponsor reports safety events. "
            "Safety is reported by the IRB. "
            "Approved risks include device discomfort and privacy risks."
        ),
        (
            "The study physician evaluates adverse events. "
            "The sponsor reports safety events. "
            "The IRB performs safety reporting. "
            "Approved risks include device discomfort and privacy risks."
        ),
        (
            "The study physician evaluates adverse events. "
            "The sponsor reports safety events. "
            "The IRB communicates safety information. "
            "Approved risks include device discomfort and privacy risks."
        ),
        (
            "The study physician evaluates adverse events. "
            "The sponsor reports safety events. "
            "The IRB has safety oversight. "
            "Approved risks include device discomfort and privacy risks."
        ),
        (
            "The study physician evaluates adverse events. "
            "The sponsor reports safety events. "
            "The IRB is the safety lead. "
            "Approved risks include device discomfort and privacy risks."
        ),
        (
            "The study physician evaluates adverse events. "
            "The sponsor reports safety events. "
            "The IRB record-reviews safety. "
            "Approved risks include device discomfort and privacy risks."
        ),
        (
            "The study physician evaluates adverse events. "
            "The sponsor reports safety events. "
            "The IRB chart-reviews safety. "
            "Approved risks include device discomfort and privacy risks."
        ),
        (
            "The study physician evaluates adverse events. "
            "The sponsor reports safety events. "
            "The IRB data-reviews safety. "
            "Approved risks include device discomfort and privacy risks."
        ),
        *(
            (
                "The study physician evaluates adverse events. "
                "The sponsor reports safety events. "
                f"The IRB {verb} safety. "
                "Approved risks include device discomfort and privacy risks."
            )
            for verb in ("supervises", "coordinates", "administers", "owns", "directs")
        ),
        (
            "The study physician evaluates adverse events. "
            "The sponsor reports safety events. "
            "Safety data are provided by the chart review conducted by the IRB. "
            "Approved risks include device discomfort and privacy risks."
        ),
        (
            "The study physician evaluates adverse events. "
            "The sponsor reports safety events. "
            "The chart review board supervises safety. "
            "Approved risks include device discomfort and privacy risks."
        ),
        *(
            (
                "The study physician evaluates adverse events. "
                "The sponsor reports safety events. "
                f"{unsupported_assignment} "
                "Approved risks include device discomfort and privacy risks."
            )
            for unsupported_assignment in (
                "Safety oversight belongs to the IRB.",
                "Safety responsibility rests with the IRB.",
                "Safety is the IRB responsibility.",
                "Safety has the IRB as lead.",
                "Safety is overseen through the IRB.",
                "Safety officer reports events.",
                "The safety officer reports events.",
                "Clinical safety lead monitors incidents.",
                "Safety committee is responsible for AE reporting.",
                (
                    "For quality complaints and adverse events the approved risks include device discomfort "
                    "and privacy risk and the IRB reports safety events."
                ),
                (
                    "For quality complaints and adverse events approved risks include device discomfort "
                    "while the IRB supervises safety."
                ),
                (
                    "For quality complaints and adverse events the IRB reports safety events and approved "
                    "risks include device discomfort."
                ),
            )
        ),
    ]
    for invalid_text in invalid_assignment_phrasings:
        invalid = recorded_acceptance_response(request)
        _replace_safety_result(invalid, invalid_text)
        accepted_invalid, invalid_findings = validate_response(request, invalid)
        accepted_invalid_ids = {
            draft["section_id"] for draft in (accepted_invalid or {}).get("drafts", [])
        }
        assert "quality-safety" not in accepted_invalid_ids, invalid_text
        assert any(
            item["field"] == "quality-safety" and "approved safety role" in item["issue"]
            for item in invalid_findings
        )

    coordinated_reference = json.loads(json.dumps(reference))
    coordinated_reference["safety"]["roles"] = [
        {
            "party": "study physician",
            "responsibilities": "assess_safety_events; report_safety_events",
        },
        {
            "party": "sponsor",
            "responsibilities": "assess_safety_events; report_safety_events",
        },
    ]
    coordinated_request = _retrospective_safety_request(
        tmp_path / "coordinated",
        coordinated_reference,
        "r-retrospective-coordinated-safety-roles",
    )
    coordinated = recorded_acceptance_response(coordinated_request)
    _replace_safety_result(
        coordinated,
        (
            "The study physician and sponsor evaluate and report adverse events. "
            "Approved risks include device discomfort and privacy risks."
        ),
    )

    accepted_coordinated, coordinated_findings = validate_response(
        coordinated_request,
        coordinated,
    )
    accepted_coordinated_ids = {
        draft["section_id"] for draft in (accepted_coordinated or {}).get("drafts", [])
    }

    assert "quality-safety" not in accepted_coordinated_ids
    assert any(item["field"] == "quality-safety" for item in coordinated_findings)

    passive = recorded_acceptance_response(request)
    _replace_safety_result(
        passive,
        (
            "Adverse events are evaluated by the study physician and reported by the sponsor. "
            "Approved risks include device discomfort and privacy risks."
        ),
    )

    accepted_passive, passive_findings = validate_response(request, passive)
    accepted_passive_ids = {
        draft["section_id"] for draft in (accepted_passive or {}).get("drafts", [])
    }

    assert "quality-safety" not in accepted_passive_ids
    assert any(item["field"] == "quality-safety" for item in passive_findings)

    modified_passive = recorded_acceptance_response(request)
    _replace_safety_result(
        modified_passive,
        (
            "Safety events are evaluated independently by the study physician and reported by the sponsor. "
            "Approved risks include device discomfort and privacy risks."
        ),
    )

    accepted_modified_passive, modified_passive_findings = validate_response(
        request,
        modified_passive,
    )
    accepted_modified_passive_ids = {
        draft["section_id"]
        for draft in (accepted_modified_passive or {}).get("drafts", [])
    }

    assert "quality-safety" not in accepted_modified_passive_ids
    assert any(
        item["field"] == "quality-safety" for item in modified_passive_findings
    )

    shared_passive_reference = json.loads(json.dumps(reference))
    shared_passive_reference["safety"]["roles"] = [{
        "party": "study physician",
        "responsibilities": "assess_safety_events; report_safety_events",
    }]
    shared_passive_request = _retrospective_safety_request(
        tmp_path / "shared-passive",
        shared_passive_reference,
        "r-retrospective-shared-passive-safety-role",
    )
    for invalid_direct_text in (
        "The study physician assesses or reports safety events.",
        "The study physician assessment and reporting safety events.",
        "The study physician reports and AEs.",
        "The study physician assesses and safety events.",
        "The study physician assesses and reports and safety events.",
    ):
        invalid_direct = recorded_acceptance_response(shared_passive_request)
        _replace_safety_result(
            invalid_direct,
            invalid_direct_text + " Approved risks include device discomfort and privacy risks.",
        )
        accepted_invalid_direct, invalid_direct_findings = validate_response(
            shared_passive_request,
            invalid_direct,
        )
        accepted_invalid_direct_ids = {
            draft["section_id"]
            for draft in (accepted_invalid_direct or {}).get("drafts", [])
        }
        assert "quality-safety" not in accepted_invalid_direct_ids
        assert any(
            item["field"] == "quality-safety" for item in invalid_direct_findings
        )
    shared_passive = recorded_acceptance_response(shared_passive_request)
    _replace_safety_result(
        shared_passive,
        (
            "Safety events are assessed and reported by the study physician. "
            "Approved risks include device discomfort and privacy risks."
        ),
    )

    accepted_shared_passive, shared_passive_findings = validate_response(
        shared_passive_request,
        shared_passive,
    )
    accepted_shared_passive_ids = {
        draft["section_id"]
        for draft in (accepted_shared_passive or {}).get("drafts", [])
    }

    assert "quality-safety" not in accepted_shared_passive_ids
    assert any(
        item["field"] == "quality-safety" for item in shared_passive_findings
    )

    shared_agents_reference = json.loads(json.dumps(reference))
    shared_agents_reference["safety"]["roles"] = [
        {"party": "study physician", "responsibilities": "assess_safety_events"},
        {"party": "sponsor", "responsibilities": "assess_safety_events"},
    ]
    shared_agents_request = _retrospective_safety_request(
        tmp_path / "shared-agents",
        shared_agents_reference,
        "r-retrospective-shared-passive-agents",
    )
    for shared_agents_text in (
        "Safety events are assessed by the study physician and sponsor.",
        "Safety events are assessed jointly by the study physician and sponsor.",
    ):
        shared_agents = recorded_acceptance_response(shared_agents_request)
        _replace_safety_result(
            shared_agents,
            shared_agents_text + " Approved risks include device discomfort and privacy risks.",
        )
        accepted_shared_agents, shared_agents_findings = validate_response(
            shared_agents_request,
            shared_agents,
        )
        accepted_shared_agents_ids = {
            draft["section_id"]
            for draft in (accepted_shared_agents or {}).get("drafts", [])
        }
        assert "quality-safety" not in accepted_shared_agents_ids
        assert any(
            item["field"] == "quality-safety" for item in shared_agents_findings
        )

    contrast = recorded_acceptance_response(request)
    _replace_safety_result(
        contrast,
        (
            "The study physician evaluates adverse events. "
            "The sponsor does not assess events but reports them. "
            "Approved risks include device discomfort and privacy risks."
        ),
    )

    accepted_contrast, contrast_findings = validate_response(request, contrast)
    accepted_contrast_ids = {
        draft["section_id"] for draft in (accepted_contrast or {}).get("drafts", [])
    }

    assert "quality-safety" not in accepted_contrast_ids
    assert any(item["field"] == "quality-safety" for item in contrast_findings)

    overlapping_roles = [
        {"party": "physician", "responsibilities": "assess_safety_events"},
        {"party": "study physician", "responsibilities": "report_safety_events"},
    ]
    overlap_reference = json.loads(json.dumps(reference))
    overlap_reference["safety"]["roles"] = overlapping_roles
    overlap_request = _retrospective_safety_request(
        tmp_path / "overlap-omission",
        overlap_reference,
        "r-retrospective-overlap-omission",
    )
    overlap = recorded_acceptance_response(overlap_request)
    _replace_safety_result(
        overlap,
        (
            "The study physician assesses and reports safety events. "
            "Approved risks include device discomfort and privacy risks."
        ),
    )
    accepted_overlap, overlap_findings = validate_response(overlap_request, overlap)
    accepted_overlap_ids = {
        draft["section_id"] for draft in (accepted_overlap or {}).get("drafts", [])
    }
    assert "quality-safety" not in accepted_overlap_ids
    assert any(item["field"] == "quality-safety" for item in overlap_findings)

    for index, ordered_roles in enumerate((overlapping_roles, list(reversed(overlapping_roles)))):
        ordered_reference = json.loads(json.dumps(reference))
        ordered_reference["safety"]["roles"] = ordered_roles
        ordered_request = _retrospective_safety_request(
            tmp_path / f"overlap-order-{index}",
            ordered_reference,
            f"r-retrospective-overlap-order-{index}",
        )
        ordered = recorded_acceptance_response(ordered_request)
        _replace_safety_result(
            ordered,
            (
                "The physician assesses safety events. "
                "The study physician reports safety events. "
                "Approved risks include device discomfort and privacy risks."
            ),
        )
        accepted_ordered, ordered_findings = validate_response(ordered_request, ordered)
        accepted_ordered_ids = {
            draft["section_id"]
            for draft in (accepted_ordered or {}).get("drafts", [])
        }
        assert "quality-safety" in accepted_ordered_ids
        assert not any(item["field"] == "quality-safety" for item in ordered_findings)

    unicode_reference = json.loads(json.dumps(reference))
    unicode_reference["safety"]["roles"] = [{
        "party": "José Müller",
        "responsibilities": "assess_safety_events; report_safety_events",
    }]
    unicode_request = _retrospective_safety_request(
        tmp_path / "unicode-party",
        unicode_reference,
        "r-retrospective-unicode-safety-role",
    )
    unicode_response = recorded_acceptance_response(unicode_request)
    _replace_safety_result(
        unicode_response,
        (
            "José Müller assesses and reports safety events. "
            "Approved risks include device discomfort and privacy risks."
        ),
    )
    accepted_unicode, unicode_findings = validate_response(
        unicode_request,
        unicode_response,
    )
    accepted_unicode_ids = {
        draft["section_id"] for draft in (accepted_unicode or {}).get("drafts", [])
    }
    assert "quality-safety" in accepted_unicode_ids
    assert not any(item["field"] == "quality-safety" for item in unicode_findings)

    period_party_reference = json.loads(json.dumps(reference))
    period_party_reference["safety"]["roles"] = [{
        "party": "Dr. José Müller",
        "responsibilities": "assess_safety_events; report_safety_events",
    }]
    period_party_request = _retrospective_safety_request(
        tmp_path / "period-party",
        period_party_reference,
        "r-retrospective-period-safety-role",
    )
    period_party = recorded_acceptance_response(period_party_request)
    _replace_safety_result(
        period_party,
        (
            "The Dr. José Müller assesses and reports safety events promptly. "
            "Approved risks include device discomfort and privacy risks."
        ),
    )
    accepted_period_party, period_party_findings = validate_response(
        period_party_request,
        period_party,
    )
    accepted_period_party_ids = {
        draft["section_id"]
        for draft in (accepted_period_party or {}).get("drafts", [])
    }
    assert "quality-safety" in accepted_period_party_ids
    assert not any(
        item["field"] == "quality-safety" for item in period_party_findings
    )

    for index, exact_party in enumerate(("Acme Inc.", "Dr.", "Straße Safety GmbH", "İrem")):
        exact_reference = json.loads(json.dumps(reference))
        exact_reference["safety"]["roles"] = [{
            "party": exact_party,
            "responsibilities": "assess_safety_events; report_safety_events",
        }]
        exact_request = _retrospective_safety_request(
            tmp_path / f"exact-party-{index}",
            exact_reference,
            f"r-retrospective-exact-safety-role-{index}",
        )
        exact_response = recorded_acceptance_response(exact_request)
        accepted_exact, exact_findings = validate_response(exact_request, exact_response)
        accepted_exact_ids = {
            draft["section_id"]
            for draft in (accepted_exact or {}).get("drafts", [])
        }
        assert "quality-safety" in accepted_exact_ids, exact_party
        assert not any(
            item["field"] == "quality-safety" for item in exact_findings
        ), exact_party

    truncated_unicode = recorded_acceptance_response(unicode_request)
    _replace_safety_result(
        truncated_unicode,
        (
            "Jos Müller assesses and reports safety events. "
            "Approved risks include device discomfort and privacy risks."
        ),
    )
    accepted_truncated, truncated_findings = validate_response(
        unicode_request,
        truncated_unicode,
    )
    accepted_truncated_ids = {
        draft["section_id"]
        for draft in (accepted_truncated or {}).get("drafts", [])
    }
    assert "quality-safety" not in accepted_truncated_ids
    assert any(item["field"] == "quality-safety" for item in truncated_findings)

    article_reference = json.loads(json.dumps(reference))
    article_reference["safety"]["roles"] = [{
        "party": "The sponsor",
        "responsibilities": "report_safety_events",
    }]
    article_request = _retrospective_safety_request(
        tmp_path / "article-party",
        article_reference,
        "r-retrospective-article-safety-role",
    )
    article_response = recorded_acceptance_response(article_request)
    article_result = next(
        item for item in article_response["section_results"]
        if item["section_id"] == "quality-safety"
    )
    article_text = " ".join(
        paragraph["text"] for paragraph in article_result["paragraphs"]
    )
    assert "The The sponsor" not in article_text
    accepted_article, article_findings = validate_response(
        article_request,
        article_response,
    )
    accepted_article_ids = {
        draft["section_id"] for draft in (accepted_article or {}).get("drafts", [])
    }
    assert "quality-safety" in accepted_article_ids
    assert not any(item["field"] == "quality-safety" for item in article_findings)

    for coordinated_objects in (
        "The sponsor reports AEs and SAEs.",
        "The sponsor reports adverse events and quality complaints.",
    ):
        object_coordination = recorded_acceptance_response(request)
        _replace_safety_result(
            object_coordination,
            (
                "The study physician evaluates adverse events. "
                f"{coordinated_objects} "
                "Approved risks include device discomfort and privacy risks."
            ),
        )
        accepted_objects, object_findings = validate_response(
            request,
            object_coordination,
        )
        accepted_object_ids = {
            draft["section_id"]
            for draft in (accepted_objects or {}).get("drafts", [])
        }
        assert "quality-safety" in accepted_object_ids, coordinated_objects
        assert not any(
            item["field"] == "quality-safety" for item in object_findings
        ), coordinated_objects

    for descriptive_clause in (
        "The chart review summarized safety events.",
        "Safety-event chart review was completed.",
        "Safety data are provided by the chart review.",
        "Retrospective safety data were summarized from the chart review.",
        "Available safety events were summarized from the record review.",
        "Historical safety information was included in the data review.",
    ):
        descriptive = recorded_acceptance_response(request)
        _replace_safety_result(
            descriptive,
            (
                "The study physician evaluates adverse events. "
                "The sponsor reports safety events. "
                f"{descriptive_clause} "
                "Approved risks include device discomfort and privacy risks."
            ),
        )
        accepted_descriptive, descriptive_findings = validate_response(
            request,
            descriptive,
        )
        accepted_descriptive_ids = {
            draft["section_id"]
            for draft in (accepted_descriptive or {}).get("drafts", [])
        }
        assert "quality-safety" in accepted_descriptive_ids, descriptive_clause
        assert not any(
            item["field"] == "quality-safety" for item in descriptive_findings
        ), descriptive_clause

    paraphrase = recorded_acceptance_response(request)
    _replace_safety_result(
        paraphrase,
        (
            "The study physician evaluates adverse events. "
            "The sponsor reports safety events. "
            "Approved risks include device discomfort and privacy risks."
        ),
    )

    accepted_paraphrase, paraphrase_findings = validate_response(request, paraphrase)
    accepted_paraphrase_ids = {
        draft["section_id"] for draft in (accepted_paraphrase or {}).get("drafts", [])
    }

    assert "quality-safety" in accepted_paraphrase_ids
    assert not any(item["field"] == "quality-safety" for item in paraphrase_findings)


def test_client_protocol_template_renders_source_supported_schedule_of_assessments(tmp_path):
    reference = json.loads((ROOT / "tests/fixtures/prospective-acceptance-source.json").read_text(encoding="utf-8"))
    report = render_documents(ROOT, tmp_path, reference, {
        "protocol": [{
            "section_id": "quality-safety.reporting",
            "paragraphs": [{"text": "Approved safety events will be recorded and reported under the study procedures."}],
            "lists": [],
        }],
        "icf": {},
        "prs": {},
    })
    protocol_report = next(item for item in report["artifacts"] if item["artifact"] == "protocol")
    assert protocol_report["status"] == "passed"

    output = Document(tmp_path / "candidate/protocol.docx")
    assessment = next(table for table in output.tables if "Device initiation" in "\n".join(cell.text for row in table.rows for cell in row.cells))
    assessment_table = "\n".join(cell.text for row in assessment.rows for cell in row.cells)

    assert "Activity" in assessment_table
    assert "Baseline" in assessment_table
    assert "Consent" in assessment_table
    assert "X" in assessment_table
    grid_widths = [int(column.w) for column in assessment._tbl.tblGrid.gridCol_lst]
    assert grid_widths[0] > max(grid_widths[1:])
    assert assessment.cell(0, 0)._tc is assessment.cell(1, 0)._tc
    first_column = [row.tc_lst[0] for row in assessment._tbl.tr_lst]
    merge_values = [
        None if (merge := cell.tcPr.find(qn("w:vMerge"))) is None else merge.get(qn("w:val"))
        for cell in first_column
    ]
    first_widths = [int(cell.tcPr.find(qn("w:tcW")).get(qn("w:w"))) for cell in first_column]
    assert merge_values[:2] == ["restart", None]
    assert merge_values[2:] == [None] * (len(merge_values) - 2)
    assert len(set(first_widths)) == 1
    assert first_widths[0] == round(grid_widths[0] / 635)
    continuation_borders = first_column[1].tcPr.find(qn("w:tcBorders"))
    continuation_top = None if continuation_borders is None else continuation_borders.find(qn("w:top"))
    assert continuation_top is None or continuation_top.get(qn("w:val")) == "nil"
    body_runs = [
        run
        for row in assessment.rows[2:]
        for cell in row.cells
        for paragraph in cell.paragraphs
        for run in paragraph.runs
        if run.text.strip()
    ]
    assert body_runs
    assert all(run.font.name == "Arial" for run in body_runs)
    assert all(_points(run.font.size) == 10 for run in body_runs)
    for cell in assessment._tbl.tr_lst[-1].tc_lst:
        bottom = cell.tcPr.find(f"{qn('w:tcBorders')}/{qn('w:bottom')}")
        assert bottom is not None
        assert bottom.get(qn("w:val")) == "double"

    contact = next(table for table in output.tables if table.rows[0].cells[0].text.strip() == "Study Staff")
    assert contact.rows[0]._tr.get_or_add_trPr().find(qn("w:tblHeader")) is not None
    assert all(
        row._tr.get_or_add_trPr().find(qn("w:tblHeader")) is None
        for row in contact.rows[1:]
    )
    assert all(row._tr.get_or_add_trPr().find(qn("w:cantSplit")) is not None for row in contact.rows)
    assert all(
        paragraph.paragraph_format.keep_together is True
        for row in contact.rows
        for cell in row.cells
        for paragraph in cell.paragraphs
    )
    assert all(
        paragraph.paragraph_format.keep_with_next is True
        for cell in contact.rows[0].cells
        for paragraph in cell.paragraphs
    )
    contact_caption_tail = Paragraph(contact._tbl.getprevious(), output)
    contact_caption_head = Paragraph(contact_caption_tail._p.getprevious(), output)
    assert contact_caption_head.text.strip().startswith("Table 13.3.-1")
    assert contact_caption_tail.text.strip() == "Contact Information for Study"
    assert contact_caption_head.paragraph_format.keep_with_next is True
    assert contact_caption_tail.paragraph_format.keep_with_next is True
    assert contact_caption_head.paragraph_format.page_break_before is not True


def test_table_specific_break_requires_a_classified_table_repair(tmp_path):
    reference = json.loads((ROOT / "tests/fixtures/prospective-acceptance-source.json").read_text(encoding="utf-8"))

    report = render_documents(
        ROOT,
        tmp_path,
        reference,
        {"protocol": [], "icf": {}, "prs": {}},
        artifact_names={"protocol"},
        layout_repairs={"protocol": ({"rule": "table_pagination", "target": "Table 13.3.-1"},)},
    )

    assert report["layout_repairs"] == {
        "protocol": [{"rule": "table_pagination", "target": "Table 13.3.-1"}],
    }
    protocol = Document(tmp_path / "candidate/protocol.docx")
    contact = next(table for table in protocol.tables if table.rows[0].cells[0].text.strip() == "Study Staff")
    caption_tail = Paragraph(contact._tbl.getprevious(), protocol)
    caption_head = Paragraph(caption_tail._p.getprevious(), protocol)
    assert caption_head.text.strip().startswith("Table 13.3.-1")
    assert not re.match(r"^\d+(?:\.\d+)*\.?\s+", caption_head.text.strip())
    assert caption_head.paragraph_format.page_break_before is True
    visits_caption = next(
        paragraph for paragraph in protocol.paragraphs
        if paragraph.text.strip().startswith("Table 9.2-1")
    )
    assert visits_caption.paragraph_format.page_break_before is not True


def test_known_plain_icf_normalization_preserves_static_tabs_and_style_lists():
    document = Document()
    ordinary = document.add_paragraph()
    ordinary.add_run().add_tab()
    ordinary.add_run("   Ordinary generated body prose.")
    static_tabbed = document.add_paragraph()
    static_tabbed.add_run().add_tab()
    static_tabbed.add_run("Intentionally tab-aligned template content.")
    styled_list = document.add_paragraph("Style-inherited list item.", style="List Bullet")
    static_xml = static_tabbed._p.xml
    list_xml = styled_list._p.xml

    rendering._normalize_known_icf_plain_paragraphs(
        document,
        {ordinary._p, styled_list._p},
    )

    assert ordinary.text == "Ordinary generated body prose."
    assert _points(ordinary.paragraph_format.left_indent) == 0.0
    assert _points(ordinary.paragraph_format.right_indent) == 0.0
    assert _points(ordinary.paragraph_format.first_line_indent) == 0.0
    assert ordinary.paragraph_format.alignment == WD_ALIGN_PARAGRAPH.JUSTIFY
    assert static_tabbed._p.xml == static_xml
    assert styled_list._p.xml == list_xml


def test_retained_agreement_role_excludes_warning_signature_and_sterling_xml():
    advarra = Document()
    advarra.add_heading("AGREEMENT TO BE IN THE STUDY", level=1)
    ordinary = advarra.add_paragraph("  Ordinary retained agreement prose.")
    warning = advarra.add_paragraph(
        "IF YOU DO NOT AGREE WITH THE STATEMENT ABOVE, YOU SHOULD NOT SIGN THIS INFORMED CONSENT DOCUMENT."
    )
    warning.paragraph_format.right_indent = Pt(2.8)
    signature = advarra.add_paragraph()
    signature.add_run().add_tab()
    signature.add_run("Signature of Participant")
    warning_xml = warning._p.xml
    signature_xml = signature._p.xml

    rendering._normalize_retained_icf_agreement_prose(advarra, sterling=False)

    assert ordinary.text == "Ordinary retained agreement prose."
    assert ordinary.paragraph_format.alignment == WD_ALIGN_PARAGRAPH.JUSTIFY
    assert _points(ordinary.paragraph_format.left_indent) == 0.0
    assert _points(ordinary.paragraph_format.right_indent) == 0.0
    assert _points(ordinary.paragraph_format.first_line_indent) == 0.0
    assert warning._p.xml == warning_xml
    assert signature._p.xml == signature_xml

    sterling = Document()
    sterling.add_heading("PARTICIPANT STATEMENT AUTHORIZATION", level=1)
    sterling.add_paragraph("I have read and agree to participate.")
    sterling_xml = sterling.element.xml

    rendering._normalize_retained_icf_agreement_prose(sterling, sterling=True)

    assert sterling.element.xml == sterling_xml


def test_icf_multi_paragraph_generated_prose_uses_native_paragraphs_across_families(tmp_path):
    cases = (
        ("prospective-acceptance-source.json", "Advarra"),
        ("prospective-acceptance-source.json", "Sterling"),
        ("ambispective-acceptance-source.json", "Advarra"),
        ("ambispective-acceptance-source.json", "Sterling"),
    )
    for index, (fixture, family) in enumerate(cases):
        reference = json.loads((ROOT / "tests/fixtures" / fixture).read_text(encoding="utf-8"))
        reference["meta"]["icf_template"] = family
        first = f"First generated benefit paragraph {index}."
        second = f"Second generated benefit paragraph {index}."
        model = {
            "protocol": [],
            "prs": {},
            "icf": {
                "icf.benefits": {
                    "paragraphs": [
                        {"text": first, "evidence_refs": [], "boilerplate_refs": []},
                        {"text": second, "evidence_refs": [], "boilerplate_refs": []},
                    ],
                    "lists": [],
                },
            },
        }
        case_root = tmp_path / f"multi-{index}-{family.casefold()}"

        render_documents(ROOT, case_root, reference, model)

        document = Document(case_root / "candidate/icf.docx")
        first_paragraph = next(item for item in document.paragraphs if item.text == first)
        second_paragraph = next(item for item in document.paragraphs if item.text == second)
        assert first_paragraph._p is not second_paragraph._p
        assert not first_paragraph._p.xpath('.//w:br')
        assert not second_paragraph._p.xpath('.//w:br')
        for paragraph in (first_paragraph, second_paragraph):
            assert _points(paragraph.paragraph_format.left_indent) == 0.0
            assert _points(paragraph.paragraph_format.right_indent) == 0.0
            assert _points(paragraph.paragraph_format.first_line_indent) == 0.0
            assert paragraph.paragraph_format.alignment == WD_ALIGN_PARAGRAPH.JUSTIFY
        assert not any(
            first in paragraph.text
            and second in paragraph.text
            and not rendering._paragraph_has_numbering(paragraph)
            for paragraph in document.paragraphs
        )


def test_icf_ordinary_body_paragraphs_are_flush_and_justified_across_families(tmp_path):
    cases = (
        ("prospective-acceptance-source.json", "Advarra"),
        ("prospective-acceptance-source.json", "Sterling"),
        ("ambispective-acceptance-source.json", "Advarra"),
        ("ambispective-acceptance-source.json", "Sterling"),
    )
    for index, (fixture, family) in enumerate(cases):
        reference = json.loads((ROOT / "tests/fixtures" / fixture).read_text(encoding="utf-8"))
        reference["meta"]["icf_template"] = family
        markers = {
            "icf.risks": f"Ordinary risk paragraph {index}.",
            "icf.benefits": f"Ordinary benefit paragraph {index}.",
            "icf.payment": f"Ordinary payment paragraph {index}.",
            "icf.costs": f"Ordinary cost paragraph {index}.",
        }
        bullet_text = f"Intentional risk bullet {index}."
        model = {
            "protocol": [],
            "prs": {},
            "icf": {
                section_id: {
                    "paragraphs": [{"text": text, "evidence_refs": [], "boilerplate_refs": []}],
                    "lists": ([{"items": [bullet_text]}] if section_id == "icf.risks" else []),
                }
                for section_id, text in markers.items()
            },
        }
        case_root = tmp_path / f"{index}-{family.casefold()}"
        render_documents(ROOT, case_root, reference, model)
        output_path = case_root / "candidate/icf.docx"
        document = Document(output_path)
        authority_name = "sterling-icf-reference.docx" if family == "Sterling" else "advarra-icf-reference.docx"
        authority_path = ROOT / "assets/client-templates/reference" / authority_name

        assert _doc_default_font(output_path) == _doc_default_font(authority_path)
        for text in markers.values():
            paragraph = next(
                item for item in document.paragraphs
                if item.text == text
                and (item._p.pPr is None or item._p.pPr.find(qn("w:numPr")) is None)
            )
            assert _points(paragraph.paragraph_format.left_indent) == 0.0
            assert _points(paragraph.paragraph_format.right_indent) == 0.0
            assert _points(paragraph.paragraph_format.first_line_indent) == 0.0
            assert paragraph.paragraph_format.alignment == WD_ALIGN_PARAGRAPH.JUSTIFY
            assert not paragraph.text[:1].isspace()
            assert not paragraph._p.xpath('./w:r[1]/w:tab')

        bullet = next(item for item in document.paragraphs if item.text == bullet_text)
        assert bullet._p.get_or_add_pPr().find(qn("w:numPr")) is not None
        assert bullet.paragraph_format.alignment != WD_ALIGN_PARAGRAPH.JUSTIFY
        tab_aligned = [paragraph for paragraph in document.paragraphs if paragraph._p.xpath('.//w:tab')]
        assert tab_aligned
        assert any(paragraph.text.strip().casefold().startswith("signature of") for paragraph in tab_aligned)


@pytest.mark.parametrize(
    ("fixture", "family", "agreement_prefix"),
    [
        ("prospective-acceptance-source.json", "Advarra", "By signing and dating this consent document"),
        ("prospective-acceptance-source.json", "Sterling", "I have read or have had read to me"),
        ("ambispective-acceptance-source.json", "Advarra", "By signing and dating this consent document"),
        ("ambispective-acceptance-source.json", "Sterling", "I have read or have had read to me"),
    ],
)
def test_retained_icf_agreement_prose_is_flush_and_justified_without_changing_warnings_or_signatures(
    tmp_path,
    fixture,
    family,
    agreement_prefix,
):
    reference = json.loads((ROOT / "tests/fixtures" / fixture).read_text(encoding="utf-8"))
    reference["meta"]["icf_template"] = family
    case_root = tmp_path / f"{fixture.removesuffix('.json')}-{family.casefold()}"
    render_documents(ROOT, case_root, reference, {"protocol": [], "icf": {}, "prs": {}})
    document = Document(case_root / "candidate/icf.docx")
    agreement = next(
        paragraph for paragraph in document.paragraphs
        if " ".join(paragraph.text.split()).startswith(agreement_prefix)
    )

    assert _points(agreement.paragraph_format.left_indent) in (None, 0.0)
    assert _points(agreement.paragraph_format.right_indent) in (None, 0.0)
    assert _points(agreement.paragraph_format.first_line_indent) in (None, 0.0)
    assert agreement.paragraph_format.alignment == WD_ALIGN_PARAGRAPH.JUSTIFY
    if family == "Advarra":
        warning = next(
            paragraph for paragraph in document.paragraphs
            if paragraph.text.strip().startswith("IF YOU DO NOT AGREE WITH THE STATEMENT ABOVE")
        )
        assert _points(warning.paragraph_format.right_indent) not in (None, 0.0)
        assert warning.paragraph_format.alignment != WD_ALIGN_PARAGRAPH.JUSTIFY
        signature = next(
            paragraph for paragraph in document.paragraphs
            if paragraph.text.strip().startswith("Signature of Participant")
        )
        assert signature._p.xpath('.//w:tab')
        assert signature.paragraph_format.alignment != WD_ALIGN_PARAGRAPH.JUSTIFY


@pytest.mark.parametrize(
    ("fixture", "family"),
    [
        ("prospective-acceptance-source.json", "Advarra"),
        ("prospective-acceptance-source.json", "Sterling"),
        ("ambispective-acceptance-source.json", "Advarra"),
        ("ambispective-acceptance-source.json", "Sterling"),
    ],
)
def test_generated_icf_privacy_prose_has_direct_widow_control(tmp_path, fixture, family):
    reference = json.loads((ROOT / "tests/fixtures" / fixture).read_text(encoding="utf-8"))
    reference["meta"]["icf_template"] = family
    privacy_text = (
        "Your study information will be handled according to the approved privacy terms. "
        "Study results may be published, but you will not be identified."
    )
    privacy_list_item = "A privacy list item remains a native list without direct widow control."
    benefit_text = "Unrelated generated benefit prose keeps its existing pagination properties."
    model = {
        "protocol": [],
        "prs": {},
        "icf": {
            "icf.privacy": {
                "paragraphs": [{"text": privacy_text, "evidence_refs": [], "boilerplate_refs": []}],
                "lists": [{"items": [privacy_list_item]}],
            },
            "icf.benefits": {
                "paragraphs": [{"text": benefit_text, "evidence_refs": [], "boilerplate_refs": []}],
                "lists": [],
            },
        },
    }
    case_root = tmp_path / f"{fixture.removesuffix('.json')}-{family.casefold()}"
    render_documents(ROOT, case_root, reference, model)
    document = Document(case_root / "candidate/icf.docx")
    privacy = next(paragraph for paragraph in document.paragraphs if paragraph.text == privacy_text)
    widow_control = privacy._p.get_or_add_pPr().find(qn("w:widowControl"))

    assert widow_control is not None
    assert widow_control.get(qn("w:val")) == "true"
    assert privacy.paragraph_format.keep_together is not True
    assert privacy.paragraph_format.keep_with_next is not True
    privacy_list = next(paragraph for paragraph in document.paragraphs if paragraph.text == privacy_list_item)
    unrelated = next(paragraph for paragraph in document.paragraphs if paragraph.text == benefit_text)
    assert privacy_list._p.get_or_add_pPr().find(qn("w:widowControl")) is None
    assert unrelated._p.get_or_add_pPr().find(qn("w:widowControl")) is None


def test_generated_privacy_enables_an_explicitly_disabled_widow_control():
    document = Document()
    document.add_heading("RELEASE OF MEDICAL RECORDS AND PRIVACY", level=1)
    privacy = document.add_paragraph("Generated privacy prose.")
    disabled = privacy._p.get_or_add_pPr().makeelement(
        qn("w:widowControl"),
        {qn("w:val"): "0"},
    )
    privacy._p.get_or_add_pPr().append(disabled)
    document.add_heading("LEGAL RIGHTS", level=1)
    model = {
        "icf": {
            "icf.privacy": {
                "paragraphs": [{"text": privacy.text}],
                "lists": [],
            },
        },
    }

    rendering._normalize_generated_icf_privacy_prose(document, model, sterling=False)

    widow_control = privacy._p.get_or_add_pPr().find(qn("w:widowControl"))
    assert widow_control is not None
    assert widow_control.get(qn("w:val")) == "true"


def test_long_generated_privacy_remains_splittable_without_a_large_blank(
    tmp_path,
    governed_pdfium,
):
    reference = json.loads(
        (ROOT / "tests/fixtures/prospective-acceptance-source.json").read_text(encoding="utf-8")
    )
    sentence = (
        "Authorized reviewers may inspect coded study records under the approved privacy terms "
        "without publicly identifying the participant. "
    )
    privacy_text = sentence * 90
    model = {
        "protocol": [],
        "prs": {},
        "icf": {
            "icf.privacy": {
                "paragraphs": [{"text": privacy_text, "evidence_refs": [], "boilerplate_refs": []}],
                "lists": [],
            },
        },
    }
    render_documents(ROOT, tmp_path, reference, model, artifact_names={"icf"})
    report = render_pages(tmp_path, page_renderer_identities=[governed_pdfium])

    assert report["status"] == "passed"
    artifact = next(item for item in report["artifacts"] if item["artifact"] == "icf")
    pages = [
        " ".join((page.extract_text() or "").split())
        for page in PdfReader(tmp_path / artifact["pdf"]).pages
    ]
    privacy_pages = [
        index
        for index, page in enumerate(pages)
        if "Authorized reviewers may inspect coded study records" in page
    ]
    assert len(privacy_pages) >= 2
    assert privacy_pages == list(range(privacy_pages[0], privacy_pages[-1] + 1))
    assert all(len(pages[index].split()) >= 80 for index in privacy_pages[:-1])


@pytest.mark.parametrize(
    "fixture",
    [
        "prospective-acceptance-source.json",
        "ambispective-acceptance-source.json",
    ],
)
def test_sterling_injury_structure_remains_family_specific(tmp_path, fixture):
    reference = json.loads((ROOT / "tests/fixtures" / fixture).read_text(encoding="utf-8"))
    reference["meta"]["icf_template"] = "Sterling"
    injury = reference["risks_benefits"]["injury_handling"]
    model = {
        "protocol": [],
        "prs": {},
        "icf": {
            "icf.injury": {
                "paragraphs": [{"text": injury, "evidence_refs": [], "boilerplate_refs": []}],
                "lists": [],
            },
        },
    }
    case_root = tmp_path / fixture.removesuffix(".json")
    render_documents(ROOT, case_root, reference, model)
    visible = " ".join(_visible_text(Document(case_root / "candidate/icf.docx")).split())

    assert visible.count("STUDY COMPLICATIONS COMPENSATION") == 1
    assert visible.count(injury) == 1
    assert visible.index("STUDY COMPLICATIONS COMPENSATION") < visible.index(injury)
    assert "IN CASE OF AN INJURY RELATED TO THIS RESEARCH STUDY" not in visible


def test_client_templates_normalize_visual_edge_cases(tmp_path):
    reference = json.loads((ROOT / "tests/fixtures/prospective-acceptance-source.json").read_text(encoding="utf-8"))
    model = {
        "protocol": [{
            "section_id": "analysis-plan.datasets",
            "paragraphs": [{"text": "The analysis data sets include all approved study measures."}],
            "lists": [],
        }],
        "icf": {
            "icf.procedures": _minimal_source_bound_icf_procedures(),
            "icf.risks": {
                "paragraphs": [{"text": "Taking part may involve inconvenience or discomfort from the approved procedures."}],
                "lists": [],
            }
        },
        "prs": {},
    }

    report = render_documents(ROOT, tmp_path, reference, model)
    assert report["status"] == "passed"

    icf = Document(tmp_path / "candidate/icf.docx")
    icf_reference = Document(ROOT / "assets/client-templates/reference/advarra-icf-reference.docx")
    risk = next(paragraph for paragraph in icf.paragraphs if paragraph.text.startswith("Taking part may involve"))
    reference_risk_heading_index = next(
        index for index, paragraph in enumerate(icf_reference.paragraphs)
        if paragraph.text.strip() == "SIDE EFFECTS AND OTHER RISKS"
    )
    reference_risk = next(
        paragraph for paragraph in icf_reference.paragraphs[reference_risk_heading_index + 1:]
        if paragraph.text.strip()
    )
    expected_risk_rhythm = _paragraph_rhythm(reference_risk)
    expected_risk_rhythm.update(left_indent=0.0, first_line_indent=0.0)
    assert _paragraph_rhythm(risk) == expected_risk_rhythm
    assert risk.paragraph_format.alignment == WD_ALIGN_PARAGRAPH.JUSTIFY
    preferences = [paragraph for paragraph in icf.paragraphs if paragraph.text.strip().startswith(("☐ Yes,", "☐ No,"))]
    assert len(preferences) == 2
    assert all(paragraph._p.pPr is None or paragraph._p.pPr.find(qn("w:numPr")) is None for paragraph in preferences)
    preference_heading = next(
        paragraph for paragraph in icf.paragraphs
        if paragraph.text.strip().startswith("Check your preference below")
    )
    preference_heading_index = next(
        index for index, paragraph in enumerate(icf.paragraphs)
        if paragraph._p is preference_heading._p
    )
    final_preference_index = next(
        index for index, paragraph in enumerate(icf.paragraphs)
        if paragraph._p is preferences[-1]._p
    )
    preference_block = icf.paragraphs[preference_heading_index:final_preference_index + 1]
    assert all(paragraph.paragraph_format.keep_with_next is True for paragraph in preference_block[:-1])
    assert all(
        paragraph.paragraph_format.keep_together is True
        for paragraph in preference_block
    )
    agreement_index = next(index for index, paragraph in enumerate(icf.paragraphs) if paragraph.text.strip() == "AGREEMENT TO BE IN THE STUDY")
    assert not (agreement_index >= 2 and not icf.paragraphs[agreement_index - 1].text.strip() and not icf.paragraphs[agreement_index - 2].text.strip())
    agreement = icf.paragraphs[agreement_index]
    reference_agreement = next(
        paragraph for paragraph in icf_reference.paragraphs
        if paragraph.text.strip() == "AGREEMENT TO BE IN THE STUDY"
    )
    expected_rhythm = _paragraph_rhythm(reference_agreement)
    expected_rhythm.update(left_indent=0.0, first_line_indent=0.0)
    assert _paragraph_rhythm(agreement) == expected_rhythm
    no_sign_index = next(
        index
        for index, paragraph in enumerate(icf.paragraphs)
        if paragraph.text.strip().startswith("IF YOU DO NOT AGREE WITH THE STATEMENT ABOVE")
    )
    final_copy_index = next(
        index
        for index, paragraph in enumerate(icf.paragraphs)
        if paragraph.text.strip() == "You will be given a signed and dated copy of this informed consent document to keep."
    )
    assert all(
        paragraph.paragraph_format.keep_with_next is True
        for paragraph in icf.paragraphs[no_sign_index:final_copy_index]
    )

    protocol = Document(tmp_path / "candidate/protocol.docx")
    protocol_reference = Document(ROOT / "assets/client-templates/reference/protocol-reference.docx")
    analysis_heading = next(paragraph for paragraph in protocol.paragraphs if paragraph.text.strip().startswith("10.1. Analysis Data Sets"))
    reference_analysis_heading = next(
        paragraph for paragraph in protocol_reference.paragraphs
        if paragraph.text.strip().startswith("10.1. Analysis Data Sets")
    )
    assert _paragraph_rhythm(analysis_heading) == _paragraph_rhythm(reference_analysis_heading)
    summary_tables = [
        table
        for table in protocol.tables
        if any(row.cells[0].text.strip() in {"Objective", "Variables"} for row in table.rows)
    ]
    assert len(summary_tables) == 1
    assert summary_tables[0].rows[0].cells[0].text.strip() == "Objective"
    assert any(row.cells[0].text.strip() == "Variables" for row in summary_tables[0].rows)
    assert all(row.height is None for row in summary_tables[0].rows)
    visit_caption = next(paragraph for paragraph in protocol.paragraphs if paragraph.text.strip().startswith("Table 9.2-1"))
    assert visit_caption.paragraph_format.keep_with_next is True
    toc_rows = [paragraph for paragraph in protocol.paragraphs if paragraph.style.name.casefold() in {"toc 1", "toc 2"}]
    for paragraph in toc_rows:
        assert "\t" not in paragraph.text
        assert "......" in paragraph.text
        assert paragraph.text.rstrip()[-1:].isdigit()
    contact_table = next(
        table for table in protocol.tables
        if any("Business" in cell.text and "Phone" in cell.text for row in table.rows for cell in row.cells)
    )
    assert contact_table.rows[0].cells[3].width >= contact_table.rows[0].cells[1].width
    for row in contact_table.rows[1:]:
        for cell_index in (1, 3):
            assert cell_index < len(row.cells)
            assert row.cells[cell_index]._tc.get_or_add_tcPr().find(qn("w:noWrap")) is not None


def test_generated_protocol_and_icf_lists_use_word_numbering_not_typed_bullet_characters(tmp_path):
    reference = json.loads((ROOT / "tests/fixtures/prospective-acceptance-source.json").read_text(encoding="utf-8"))
    protocol_item = "Approved protocol eligibility item for Word numbering."
    icf_item = "Approved participant procedure item for Word numbering."
    model = {
        "protocol": [{"section_id": "subjects.inclusion", "paragraphs": [], "lists": [{"items": [protocol_item]}]}],
        "icf": {"icf.procedures": {"paragraphs": [], "lists": [{"items": [icf_item]}]}},
        "prs": {},
    }

    render_documents(ROOT, tmp_path, reference, model)
    protocol = Document(tmp_path / "candidate/protocol.docx")
    icf = Document(tmp_path / "candidate/icf.docx")

    for document, item in ((protocol, protocol_item), (icf, icf_item)):
        paragraph = next(paragraph for paragraph in document.paragraphs if paragraph.text == item)
        assert not paragraph.text.startswith("•")
        assert paragraph._p.get_or_add_pPr().find(qn("w:numPr")) is not None
        assert len(paragraph._p.findall(qn("w:pPr"))) == 1


def test_icf_overview_and_visit_detail_placeholders_do_not_duplicate_the_same_draft(tmp_path):
    reference = json.loads((ROOT / "tests/fixtures/prospective-acceptance-source.json").read_text(encoding="utf-8"))
    model = {
        "protocol": [],
        "prs": {},
        "icf": {
            "icf.procedures": {
                "paragraphs": [{
                    "text": "Historical records are reviewed. At Baseline, the Sentinel Patch is initiated.",
                    "evidence_refs": ["source:procedures.visit_schedule"],
                    "boilerplate_refs": [],
                }],
                "lists": [],
            }
        },
    }
    render_documents(ROOT, tmp_path, reference, model)
    visible = _visible_text(Document(tmp_path / "candidate/icf.docx"))

    assert visible.count("Historical records are reviewed.") == 1
    assert visible.count("At Baseline, the Sentinel Patch is initiated.") == 1


def test_advarra_icf_uses_every_accepted_study_section_and_removes_example_study_prose(tmp_path):
    reference = json.loads((ROOT / "tests/fixtures/ambispective-acceptance-source.json").read_text(encoding="utf-8"))
    section_ids = (
        "icf.study-purpose", "icf.procedures", "icf.duration", "icf.risks", "icf.benefits",
        "icf.payment", "icf.costs", "icf.alternatives", "icf.privacy", "icf.injury",
    )
    markers = {
        section_id: f"Approved participant-facing content for {section_id.replace('.', ' ')} appears here."
        for section_id in section_ids
    }
    model = {
        "protocol": [],
        "prs": {},
        "icf": {
            section_id: {
                "paragraphs": [{"text": marker, "evidence_refs": [], "boilerplate_refs": []}],
                "lists": [],
            }
            for section_id, marker in markers.items()
        },
    }

    render_documents(ROOT, tmp_path, reference, model)
    output = Document(tmp_path / "candidate/icf.docx")
    visible = _visible_text(output)

    assert all(
        marker in visible
        for section_id, marker in markers.items()
        if section_id != "icf.injury"
    )
    assert markers["icf.injury"] not in visible
    assert reference["risks_benefits"]["injury_handling"] in visible
    assert "eye tests and procedures" not in visible.casefold()
    assert "no additional side effects or risks expected" not in visible.casefold()
    assert 'w:type="column"' not in output.element.xml


def test_quality_gate_rejects_example_study_prose_not_supported_by_the_source(tmp_path):
    reference = json.loads((ROOT / "tests/fixtures/ambispective-acceptance-source.json").read_text(encoding="utf-8"))
    render_documents(ROOT, tmp_path, reference, {"protocol": [], "icf": {}, "prs": {}})
    icf_path = tmp_path / "candidate/icf.docx"
    icf = Document(icf_path)
    icf.add_paragraph("You will have eye tests and procedures performed for this study.")
    icf.save(icf_path)

    findings = deterministic_content_check(tmp_path, reference)

    assert any("example-study" in item["issue"].casefold() for item in findings)


def test_protocol_summary_table_keeps_client_label_width_and_source_masking(tmp_path):
    reference = json.loads((ROOT / "tests/fixtures/prospective-acceptance-source.json").read_text(encoding="utf-8"))
    model = {
        "protocol": [
            {"section_id": "objectives", "paragraphs": [], "lists": [{"items": ["Describe recovery outcomes."]}]},
            {"section_id": "study-design.bias", "paragraphs": [{"text": "Bias-control narrative belongs in Section 8.2."}], "lists": []},
        ],
        "icf": {},
        "prs": {},
    }
    render_documents(ROOT, tmp_path, reference, model)
    output = Document(tmp_path / "candidate/protocol.docx")
    summary = next(table for table in output.tables if table.rows[0].cells[0].text.strip() == "Objective")
    values = {row.cells[0].text.strip(): row.cells[1].text.strip() for row in summary.rows}

    assert all(row.cells[0].paragraphs[0].paragraph_format.right_indent is None for row in summary.rows)
    assert not values["Objective"].startswith("•")
    assert values["Masking"] == "None"
    assert "Bias-control narrative" not in values["Masking"]


def test_protocol_visit_overview_and_detail_placeholders_do_not_duplicate_the_same_draft(tmp_path):
    reference = json.loads((ROOT / "tests/fixtures/prospective-acceptance-source.json").read_text(encoding="utf-8"))
    model = {
        "protocol": [{
            "section_id": "study-procedure.visits",
            "paragraphs": [{"text": "Four study time points are planned. Baseline includes device initiation."}],
            "lists": [],
        }],
        "icf": {},
        "prs": {},
    }
    render_documents(ROOT, tmp_path, reference, model)
    document = Document(tmp_path / "candidate/protocol.docx")
    visible = _visible_text(document)

    assert visible.count("Four study time points are planned.") == 1
    assert visible.count("Baseline includes device initiation.") == 1
    generated = [paragraph for paragraph in document.paragraphs if "study time points" in paragraph.text or "device initiation" in paragraph.text]
    assert all(paragraph.paragraph_format.left_indent is None for paragraph in generated)
    assert all(paragraph.paragraph_format.right_indent is None for paragraph in generated)


def test_visit_schedule_omits_crf_column_when_source_has_no_crf_values(tmp_path):
    reference = json.loads((ROOT / "tests/fixtures/prospective-acceptance-source.json").read_text(encoding="utf-8"))
    for visit in reference["procedures"]["visit_schedule_table"]:
        visit.pop("CRFnumber", None)
    render_documents(ROOT, tmp_path, reference, {"protocol": [], "icf": {}, "prs": {}})
    output = Document(tmp_path / "candidate/protocol.docx")
    schedule = next(table for table in output.tables if len(table.rows[0].cells) > 1 and "Visit Name" in table.rows[0].cells[1].text)

    assert len(schedule.columns) == 3
    assert all("CRF" not in cell.text for row in schedule.rows for cell in row.cells)
    authority = Document(ROOT / "assets/client-templates/reference/protocol-reference.docx")
    authority_schedule = next(
        table for table in authority.tables
        if len(table.columns) == 3 and table.rows[0].cells[0].text.strip() == "Visit Number"
    )
    output_grid = [column.w for column in schedule._tbl.tblGrid.gridCol_lst]
    authority_grid = [column.w for column in authority_schedule._tbl.tblGrid.gridCol_lst]
    output_width = schedule._tbl.tblPr.find(qn("w:tblW"))
    authority_width = authority_schedule._tbl.tblPr.find(qn("w:tblW"))
    output_fills = [cell._tc.tcPr.find(qn("w:shd")).get(qn("w:fill")) for cell in schedule.rows[0].cells]
    authority_fills = [cell._tc.tcPr.find(qn("w:shd")).get(qn("w:fill")) for cell in authority_schedule.rows[0].cells]

    assert schedule.style.style_id == authority_schedule.style.style_id
    assert output_grid == authority_grid
    assert output_width.get(qn("w:w")) == authority_width.get(qn("w:w"))
    assert output_fills == authority_fills
    output_second_data_borders = [
        cell._tc.tcPr.find(qn("w:tcBorders"))
        for cell in schedule.rows[2].cells
    ]
    authority_second_data_borders = [
        cell._tc.tcPr.find(qn("w:tcBorders"))
        for cell in authority_schedule.rows[2].cells
    ]
    border_signature = lambda borders: None if borders is None else [
        (border.tag, tuple(sorted(border.attrib.items())))
        for border in borders
    ]
    assert [border_signature(borders) for borders in output_second_data_borders] == [
        border_signature(borders) for borders in authority_second_data_borders
    ]


def test_protocol_uses_source_bound_section_drafts_not_example_study_body_text(tmp_path):
    reference = json.loads((ROOT / "tests/fixtures/prospective-acceptance-source.json").read_text(encoding="utf-8"))
    model = {
        "protocol": [
            {"section_id": "study-procedure.discontinued", "paragraphs": [{"text": "Participants may withdraw at any time; information collected before withdrawal may remain in the study records."}], "lists": []},
            {"section_id": "endpoint-criteria.discontinuation", "paragraphs": [{"text": "A participant may discontinue participation at any time without penalty."}], "lists": []},
            {"section_id": "quality-safety.general", "paragraphs": [{"text": "Safety information and product quality complaints will be documented and reviewed under the approved study safety procedures."}], "lists": []},
            {"section_id": "quality-safety.analysis", "paragraphs": [{"text": "Safety findings will be summarized using the event categories and timing available in the study data."}], "lists": []},
            {"section_id": "ethics.confidentiality", "paragraphs": [{"text": "Study records will use appropriate identifiers and access controls for authorized personnel."}], "lists": []},
        ],
        "icf": {},
        "prs": {},
    }
    render_documents(ROOT, tmp_path, reference, model)
    output = Document(tmp_path / "candidate/protocol.docx")
    visible = _visible_text(output).casefold()

    assert len(output.sections) == 1
    assert "safety findings will be summarized" in visible
    assert "ocular adverse" not in visible
    assert "performed on eyes" not in visible
    assert "prior to surgery" not in visible
    assert "prior to visit 7" not in visible


def test_observational_outputs_remove_unapproved_template_study_claims(tmp_path):
    reference = json.loads((ROOT / "tests/fixtures/ambispective-acceptance-source.json").read_text(encoding="utf-8"))
    reference["meta"].pop("version", None)
    reference["parties"]["irb"]["approval_status"] = "Pending"
    reference["safety"] = {}
    reference["ethics"] = {}
    reference["parties"]["funding_source"] = {"name": "ACME Foundation"}
    reference["risks_benefits"] = {"compensation_or_reimbursement": "None"}
    model = {
        "protocol": [],
        "icf": {
            "icf.procedures": _minimal_source_bound_icf_procedures(),
            "icf.injury": {"paragraphs": [{"text": "Study personnel will evaluate and document research-related medical concerns."}], "lists": []},
            "icf.costs": {"paragraphs": [{"text": "No study-related cost terms are specified in the approved study information."}], "lists": []},
        },
        "prs": {},
    }

    report = render_documents(ROOT, tmp_path, reference, model)
    assert report["status"] == "passed"
    protocol = Document(tmp_path / "candidate/protocol.docx")
    icf = Document(tmp_path / "candidate/icf.docx")
    protocol_visible = _visible_text(protocol)
    icf_visible = _visible_text(icf)

    title_table = next(table for table in protocol.tables if table.cell(0, 0).text.strip() == "Protocol Number")
    assert title_table.cell(1, 1).text.strip() == ""
    assert "An ambispective observational device study" in protocol_visible
    assert "investigator-initiated clinical trial" not in protocol_visible
    assert "All subjects will be monitored for adverse events" not in protocol_visible
    assert "1996 version of the Declaration of Helsinki" not in protocol_visible
    assert "approval prior to initiating the study" not in protocol_visible
    assert "funding only" not in protocol_visible.casefold()

    assert "NOT TO BE USED FOR PARTICIPANT ENROLLMENT" not in icf_visible
    assert "description of this clinical trial" not in icf_visible
    assert "Advarra Institutional Review Board" not in icf_visible
    assert reference["parties"]["irb"]["name"] in icf_visible
    assert "all charges for medical care" not in icf_visible
    assert "insurance company" not in icf_visible


def test_protocol_agreement_and_icf_shell_are_source_bound(tmp_path):
    reference = json.loads((ROOT / "tests/fixtures/ambispective-acceptance-source.json").read_text(encoding="utf-8"))
    reference["meta"]["version"] = ""
    reference["meta"]["date"] = "24 Aug 2026"
    reference["sites"][0]["facility"]["address"] = {
        "street": "220 Clinic Way", "city": "Boston", "state": "MA", "country": "United States"
    }
    model = {
        "protocol": [],
        "icf": {
            "icf.procedures": {"paragraphs": [{"text": "You will complete the approved study visits and procedures."}], "lists": []},
            "icf.privacy": {"paragraphs": [{"text": "Your records will use coded identifiers and may be reviewed only by authorized study and oversight personnel."}], "lists": []},
            "icf.injury": {"paragraphs": [{"text": "Contact the study doctor promptly if you believe a study activity caused an injury."}], "lists": []},
            "icf.costs": {"paragraphs": [{"text": "Before you sign, the study team will explain whether any research costs are your responsibility."}], "lists": []},
        },
        "prs": {},
    }

    render_documents(ROOT, tmp_path, reference, model)
    protocol = Document(tmp_path / "candidate/protocol.docx")
    icf = Document(tmp_path / "candidate/icf.docx")
    protocol_visible = _visible_text(protocol).casefold()
    icf_visible = _visible_text(icf).casefold()

    assert "study products" not in protocol_visible
    assert "serious adverse events defined in section 13" not in protocol_visible
    assert "firm in another country" not in icf_visible
    assert "health canada" not in icf_visible
    assert "as required by u.s. law" not in icf_visible
    title = next(table for table in protocol.tables if table.rows[0].cells[0].text.strip() == "Protocol Number")
    title_values = {row.cells[0].text.strip(): row.cells[1].text.strip() for row in title.rows}
    assert title_values["Protocol Date"] == "24 Aug 2026"
    assert title_values["Amendment Number"] == ""
    header_control = "\n".join(
        cell.text
        for section in icf.sections
        for table in section.header.tables
        for row in table.rows
        for cell in row.cells
    ).casefold()
    assert "24 aug 2026" not in header_control

    front = icf.tables[0]
    labels = [row.cells[0].text.strip() for row in front.rows]
    values = [row.cells[1].text.strip() for row in front.rows]
    assert "Study Coordinator Telephone:" in labels
    assert "Study Site Address:" in labels
    assert reference["parties"]["study_coordinator"]["business_phone"] in values
    assert any(reference["sites"][0]["facility"]["address"]["street"] in value for value in values)

    summary = next(table for table in protocol.tables if table.rows[0].cells[0].text.strip() == "Objective")
    summary_values = {row.cells[0].text.strip(): row.cells[1].text.strip() for row in summary.rows}
    assert "Hypothesis" not in summary_values
    assert summary_values["Test Article(s)"]


@pytest.mark.parametrize(
    "fixture_name",
    (
        "prospective-acceptance-source.json",
        "ambispective-acceptance-source.json",
        "retrospective-acceptance-source.json",
    ),
)
def test_protocol_version_uses_its_own_title_control(tmp_path, fixture_name):
    reference = json.loads((ROOT / "tests/fixtures" / fixture_name).read_text(encoding="utf-8"))
    reference["meta"]["version"] = "1.0"

    render_documents(
        ROOT,
        tmp_path,
        reference,
        {"protocol": [], "icf": {}, "prs": {}},
        artifact_names={"protocol"},
    )

    protocol = Document(tmp_path / "candidate/protocol.docx")
    title = next(
        table for table in protocol.tables
        if table.rows[0].cells[0].text.strip() == "Protocol Number"
    )
    title_values = {
        row.cells[0].text.strip(): row.cells[1].text.strip()
        for row in title.rows
    }
    assert title_values["Amendment Number"] == ""
    assert title_values["Protocol Version"] == "1.0"


@pytest.mark.parametrize(
    "template_name",
    (
        "prospective-protocol.template.docx",
        "ambispective-protocol.template.docx",
        "retrospective-protocol.template.docx",
    ),
)
def test_protocol_amendment_number_is_blank_when_date_and_version_are_blank(
    template_name,
):
    protocol = Document(ROOT / "assets/client-templates/docx" / template_name)

    rendering._normalize_protocol_title_controls(
        protocol,
        {"meta": {"date": "", "version": ""}},
    )

    title = next(
        table for table in protocol.tables
        if table.rows[0].cells[0].text.strip() == "Protocol Number"
    )
    title_values = {
        row.cells[0].text.strip(): row.cells[1].text.strip()
        for row in title.rows
    }
    assert title_values["Amendment Number"] == ""


@pytest.mark.parametrize(
    ("fixture_name", "icf_template"),
    (
        ("prospective-acceptance-source.json", "Advarra"),
        ("prospective-acceptance-source.json", "Sterling"),
        ("ambispective-acceptance-source.json", "Advarra"),
        ("ambispective-acceptance-source.json", "Sterling"),
    ),
)
def test_icf_procedures_replace_open_ended_template_eligibility(
    tmp_path, fixture_name, icf_template,
):
    reference = json.loads((ROOT / "tests/fixtures" / fixture_name).read_text(encoding="utf-8"))
    reference["meta"]["icf_template"] = icf_template
    procedures = (
        "You may participate if you are 18 through 80 years old and have the target condition. "
        "Before screening, you must have gone at least 30 days without participating in another study. "
        "You cannot participate if you are unable to complete follow-up. At Baseline on Day 0, you will "
        "provide consent and begin use of the study device. Assessments occur at Baseline, Month 1, and Month 3."
    )
    model = {
        "protocol": [],
        "icf": {"icf.procedures": {"paragraphs": [{"text": procedures}], "lists": []}},
        "prs": {},
    }

    render_documents(ROOT, tmp_path, reference, model)

    visible = _visible_text(Document(tmp_path / "candidate/icf.docx"))
    assert procedures in visible
    assert "include but are not limited to" not in visible.casefold()
    assert "there may be other reasons why you cannot participate" not in visible.casefold()


@pytest.mark.parametrize(
    ("fixture_name", "icf_template"),
    (
        ("prospective-acceptance-source.json", "Advarra"),
        ("prospective-acceptance-source.json", "Sterling"),
        ("ambispective-acceptance-source.json", "Advarra"),
        ("ambispective-acceptance-source.json", "Sterling"),
    ),
)
@pytest.mark.parametrize("procedures_draft", (None, {"paragraphs": [], "lists": []}))
def test_icf_rendering_blocks_when_source_bound_procedures_are_missing_or_empty(
    tmp_path, fixture_name, icf_template, procedures_draft,
):
    reference = json.loads((ROOT / "tests/fixtures" / fixture_name).read_text(encoding="utf-8"))
    reference["meta"]["icf_template"] = icf_template
    icf = {} if procedures_draft is None else {"icf.procedures": procedures_draft}

    report = render_documents(
        ROOT,
        tmp_path,
        reference,
        {"protocol": [], "icf": icf, "prs": {}},
    )

    assert report["status"] == "blocked"
    icf_artifact = next(item for item in report["artifacts"] if item["artifact"] == "icf")
    assert icf_artifact["status"] == "blocked"
    assert icf_artifact["path"] == "candidate/icf.docx"
    assert any(
        finding["field"] == "icf.procedures"
        and "source-bound" in finding["issue"]
        for finding in icf_artifact["findings"]
    )
    visible = _visible_text(Document(tmp_path / icf_artifact["path"])).casefold()
    assert "include but are not limited to" not in visible
    assert "there may be other reasons why you cannot participate" not in visible


def test_retrospective_enrollment_body_is_inserted_after_template_heading_normalization(tmp_path):
    reference = json.loads((ROOT / "tests/fixtures/retrospective-acceptance-source.json").read_text(encoding="utf-8"))
    expected = "This retrospective study uses existing records and the approved record-review sequence."
    model = {
        "protocol": [{
            "section_id": "study-procedure.enrollment",
            "paragraphs": [{"text": expected}],
            "lists": [],
        }],
        "icf": {},
        "prs": {},
    }

    report = render_documents(ROOT, tmp_path, reference, model)

    assert report["status"] == "passed"
    protocol = Document(tmp_path / "candidate/protocol.docx")
    paragraphs = protocol.paragraphs
    heading_index = next(index for index, paragraph in enumerate(paragraphs) if paragraph.text.strip().startswith("8.1. Informed Consent"))
    next_heading = next(
        index for index in range(heading_index + 1, len(paragraphs))
        if paragraphs[index].style.name.casefold().startswith("heading")
    )
    assert expected in "\n".join(paragraph.text for paragraph in paragraphs[heading_index + 1:next_heading])


def test_retrospective_inline_template_body_is_not_duplicated(tmp_path):
    reference = json.loads((ROOT / "tests/fixtures/retrospective-acceptance-source.json").read_text(encoding="utf-8"))
    expected = "The analysis data sets will be organized around the approved outcome summaries."
    model = {
        "protocol": [{
            "section_id": "analysis-plan.datasets",
            "paragraphs": [{"text": expected}],
            "lists": [],
        }],
        "icf": {},
        "prs": {},
    }

    render_documents(ROOT, tmp_path, reference, model)

    protocol = Document(tmp_path / "candidate/protocol.docx")
    assert _visible_text(protocol).count(expected) == 1
    heading_index = next(
        index for index, paragraph in enumerate(protocol.paragraphs)
        if paragraph.text.strip() == "9.1. Analysis Data Sets"
    )
    assert protocol.paragraphs[heading_index].style.name == "Heading 2"
    assert protocol.paragraphs[heading_index + 1].text == expected
    assert protocol.paragraphs[heading_index + 1].style.name == "Normal"


def test_ambispective_body_sections_follow_template_pagination_and_spacing(tmp_path):
    reference = json.loads((ROOT / "tests/fixtures/ambispective-acceptance-source.json").read_text(encoding="utf-8"))

    render_documents(ROOT, tmp_path, reference, {"protocol": [], "icf": {}, "prs": {}})

    protocol = Document(tmp_path / "candidate/protocol.docx")
    authority = Document(ROOT / "assets/client-templates/reference/protocol-reference.docx")
    for title in ("3. GENERAL INFORMATION", "16. CONFIDENTIALITY"):
        heading = next(
            paragraph for paragraph in protocol.paragraphs
            if paragraph.text.strip() == title
        )
        authority_heading = next(
            paragraph for paragraph in authority.paragraphs
            if " ".join(paragraph.text.split()) == title
        )
        if title == "3. GENERAL INFORMATION":
            assert heading.paragraph_format.page_break_before is True
        else:
            assert heading.paragraph_format.page_break_before is not True
        assert _paragraph_rhythm(heading) == _paragraph_rhythm(authority_heading)
    toc_index = next(
        index for index, paragraph in enumerate(protocol.paragraphs)
        if paragraph.text.strip() == "4. TABLE OF CONTENTS"
    )
    body_headings = [
        paragraph for paragraph in protocol.paragraphs[toc_index + 1:]
        if paragraph.style.name.casefold().startswith("heading")
    ]
    first_body = next(paragraph for paragraph in body_headings if "INTRODUCTION" in paragraph.text)
    assert _page_boundary_before(first_body)
    assert all(
        paragraph.paragraph_format.page_break_before is not True
        for paragraph in body_headings
        if paragraph is not first_body
    )


@pytest.mark.parametrize(
    ("fixture_name", "expected_template"),
    (
        ("retrospective-acceptance-source.json", "retrospective-protocol.template.docx"),
        ("ambispective-acceptance-source.json", "ambispective-protocol.template.docx"),
        ("prospective-acceptance-source.json", "prospective-protocol.template.docx"),
    ),
)
def test_protocol_toc_boundaries_preserve_selected_client_template_without_package_growth(
    tmp_path, fixture_name, expected_template,
):
    reference = json.loads((ROOT / "tests/fixtures" / fixture_name).read_text(encoding="utf-8"))

    report = render_documents(ROOT, tmp_path, reference, {"protocol": [], "icf": {}, "prs": {}})

    protocol_report = next(item for item in report["artifacts"] if item["artifact"] == "protocol")
    output_path = tmp_path / protocol_report["path"]
    template_path = ROOT / protocol_report["template"]
    protocol = Document(output_path)
    headings = [paragraph for paragraph in protocol.paragraphs if paragraph.style.name.casefold().startswith("heading")]
    toc = next(paragraph for paragraph in headings if "TABLE OF CONTENTS" in paragraph.text)
    first_body = next(paragraph for paragraph in headings if "INTRODUCTION" in paragraph.text)

    assert protocol_report["template"].endswith(expected_template)
    assert protocol_report["template_sha256"] == sha256_file(template_path)
    assert _page_boundary_before(toc)
    assert toc.paragraph_format.page_break_before is True
    assert _page_boundary_before(first_body)
    assert first_body.paragraph_format.page_break_before is not True
    assert all(
        paragraph.paragraph_format.page_break_before is not True
        for paragraph in headings[headings.index(first_body) + 1:]
    )
    with zipfile.ZipFile(template_path) as template_package, zipfile.ZipFile(output_path) as output_package:
        template_media = {name for name in template_package.namelist() if name.startswith("word/media/")}
        output_media = {name for name in output_package.namelist() if name.startswith("word/media/")}
    assert output_media == template_media
    assert output_path.stat().st_size <= template_path.stat().st_size + 256 * 1024


@pytest.mark.parametrize(
    "fixture_name",
    ("prospective-acceptance-source.json", "ambispective-acceptance-source.json"),
)
def test_leaf_body_replacement_removes_obsolete_template_break_before_section_15(
    tmp_path, fixture_name,
):
    reference = json.loads(
        (ROOT / "tests/fixtures" / fixture_name).read_text(encoding="utf-8")
    )
    report = render_documents(
        ROOT,
        tmp_path,
        reference,
        {
            "protocol": [{
                "section_id": "ethics.confidentiality",
                "paragraphs": [{"text": "Approved confidentiality replacement."}],
                "lists": [],
            }],
            "icf": {},
            "prs": {},
        },
        artifact_names={"protocol"},
    )

    assert report["status"] == "passed"
    document = Document(tmp_path / "candidate/protocol.docx")
    section_15 = next(
        paragraph for paragraph in document.paragraphs
        if paragraph.text.strip() == "15. STANDARD EVALUATION PROCEDURES"
    )
    assert _has_page_boundary_before(section_15) is False
    previous = section_15._p.getprevious()
    assert previous is not None
    assert not previous.xpath('.//w:br[@w:type="page"]')


def test_protocol_toc_boundaries_are_added_once_when_template_has_none_and_toc_spans_pages():
    document = Document()
    document.add_paragraph("3. GENERAL INFORMATION", style="Heading 1")
    toc = document.add_paragraph("4. TABLE OF CONTENTS", style="Heading 1")
    for index in range(80):
        entry = document.add_paragraph(f"{index + 1}. Generated TOC entry")
        if index == 39:
            _add_page_break(entry)
    first_body = document.add_paragraph("5. INTRODUCTION", style="Heading 1")
    later_body = document.add_paragraph("6. OBJECTIVES", style="Heading 1")

    _normalize_protocol_section_pagination(document)
    first_count = _explicit_page_break_count(document)
    _normalize_protocol_section_pagination(document)

    assert _page_boundary_before(toc)
    assert _page_boundary_before(first_body)
    assert later_body.paragraph_format.page_break_before is not True
    assert first_count == 3  # TOC start, simulated continuation page, first body start.
    assert _explicit_page_break_count(document) == first_count


def test_protocol_toc_boundaries_preserve_existing_template_breaks_without_duplicates():
    document = Document()
    document.add_paragraph("3. GENERAL INFORMATION", style="Heading 1")
    cosmetic_spacer_one = document.add_paragraph()
    cosmetic_spacer_two = document.add_paragraph()
    before_toc = document.add_paragraph()
    _add_page_break(before_toc)
    toc = document.add_paragraph("4. TABLE OF CONTENTS", style="Heading 1")
    document.add_paragraph("Cached TOC entry")
    before_body = document.add_paragraph()
    _add_page_break(before_body)
    first_body = document.add_paragraph("5. INTRODUCTION", style="Heading 1")

    _normalize_protocol_section_pagination(document)
    _normalize_protocol_section_pagination(document)

    assert _page_boundary_before(toc)
    assert _page_boundary_before(first_body)
    assert toc.paragraph_format.page_break_before is True
    assert first_body.paragraph_format.page_break_before is not True
    assert _explicit_page_break_count(document) == 2
    assert before_toc._p.getparent() is None
    assert cosmetic_spacer_one._p.getparent() is None
    assert cosmetic_spacer_two._p.getparent() is None


@pytest.mark.parametrize(
    ("fixture_name", "icf_family", "artifact", "target"),
    (
        ("retrospective-acceptance-source.json", None, "protocol", "5. OBJECTIVE(S)"),
        ("prospective-acceptance-source.json", None, "protocol", "6. OBJECTIVE(S)"),
        ("ambispective-acceptance-source.json", None, "protocol", "6. OBJECTIVE(S)"),
        ("prospective-acceptance-source.json", "Advarra", "icf", "INTRODUCTION"),
        ("prospective-acceptance-source.json", "Sterling", "icf", "KEY INFORMATION"),
    ),
)
def test_every_layout_family_repair_is_local_idempotent_and_renderer_verified(
    tmp_path,
    governed_pdfium,
    fixture_name,
    icf_family,
    artifact,
    target,
):
    reference = json.loads((ROOT / "tests/fixtures" / fixture_name).read_text(encoding="utf-8"))
    if icf_family is not None:
        reference["meta"]["icf_template"] = icf_family
    output = tmp_path / f"{reference['meta']['study_type']}-{icf_family or 'protocol'}"
    document_report = render_documents(
        ROOT,
        output,
        reference,
        {
            "protocol": [],
            "icf": {"icf.procedures": _minimal_source_bound_icf_procedures()},
            "prs": {},
        },
        artifact_names={artifact},
    )
    candidate = output / "candidate" / f"{artifact}.docx"
    document = Document(candidate)
    heading = rendering._target_heading(document, target, protocol=artifact == "protocol")
    first_block = rendering._first_substantive_block(document, heading)
    assert isinstance(first_block, (Paragraph, Table))
    first_text = (
        first_block.text
        if isinstance(first_block, Paragraph)
        else " ".join(cell.text for cell in first_block.rows[0].cells)
    )

    # Simulate the exact orphan-heading class while retaining the client's
    # content and all non-pagination formatting as the metamorphic baseline.
    heading.paragraph_format.keep_with_next = False
    heading.paragraph_format.keep_together = False
    heading.paragraph_format.widow_control = False
    if isinstance(first_block, Paragraph):
        first_block.paragraph_format.widow_control = False
    document.save(candidate)
    before_text = _visible_text(Document(candidate))
    before_parts, before_document = _docx_parts_without_pagination_controls(candidate)

    repaired = Document(candidate)
    rendering._repair_heading_cohesion(
        repaired,
        target,
        protocol=artifact == "protocol",
    )
    repaired.save(candidate)

    after = Document(candidate)
    repaired_heading = rendering._target_heading(after, target, protocol=artifact == "protocol")
    after_parts, after_document = _docx_parts_without_pagination_controls(candidate)
    assert document_report["status"] == "passed"
    assert repaired_heading.paragraph_format.keep_with_next is True
    assert repaired_heading.paragraph_format.keep_together is True
    assert _visible_text(after) == before_text
    assert after_parts == before_parts
    assert after_document == before_document

    with zipfile.ZipFile(candidate) as package:
        first_repair_xml = package.read("word/document.xml")
    rendering._repair_heading_cohesion(
        after,
        target,
        protocol=artifact == "protocol",
    )
    after.save(candidate)
    with zipfile.ZipFile(candidate) as package:
        assert package.read("word/document.xml") == first_repair_xml

    render_report = render_pages(output, page_renderer_identities=[governed_pdfium])
    assert render_report["status"] == "passed", render_report
    rendered = next(item for item in render_report["artifacts"] if item["artifact"] == artifact)
    pages = [" ".join((page.extract_text() or "").split()) for page in PdfReader(output / rendered["pdf"]).pages]
    target_marker = " ".join(target.split())
    content_marker = " ".join(first_text.split()[:4])
    assert any(target_marker in page and content_marker in page for page in pages)
    assert all(page["sha256"] == sha256_file(output / page["path"]) for page in rendered["pages"])


def test_every_protocol_and_icf_family_uses_natural_body_pagination(tmp_path, governed_pdfium):
    cases = (
        ("prospective-acceptance-source.json", "Advarra"),
        ("prospective-acceptance-source.json", "Sterling"),
        ("ambispective-acceptance-source.json", "Advarra"),
        ("ambispective-acceptance-source.json", "Sterling"),
        ("retrospective-acceptance-source.json", None),
    )
    for fixture_name, icf_family in cases:
        reference = json.loads((ROOT / "tests/fixtures" / fixture_name).read_text(encoding="utf-8"))
        if icf_family is not None:
            reference["meta"]["icf_template"] = icf_family
        output = tmp_path / f"{reference['meta']['study_type']}-{icf_family or 'none'}"

        icf = (
            {"icf.procedures": _minimal_source_bound_icf_procedures()}
            if icf_family is not None else {}
        )
        document_report = render_documents(
            ROOT, output, reference, {"protocol": [], "icf": icf, "prs": {}}
        )
        render_report = render_pages(output, page_renderer_identities=[governed_pdfium])

        assert document_report["status"] == "passed"
        assert render_report["status"] == "passed"
        expected_artifacts = {"protocol", "icf"} if icf_family is not None else {"protocol"}
        assert {artifact["artifact"] for artifact in render_report["artifacts"]} == expected_artifacts
        assert all(
            page["sha256"] == sha256_file(output / page["path"])
            for artifact in render_report["artifacts"]
            for page in artifact["pages"]
        )

        protocol = Document(output / "candidate/protocol.docx")
        rendered_protocol = next(artifact for artifact in render_report["artifacts"] if artifact["artifact"] == "protocol")
        protocol_pages = [
            " ".join((page.extract_text() or "").split())
            for page in PdfReader(output / rendered_protocol["pdf"]).pages
        ]
        protocol_number = reference["meta"]["protocol_number"]
        assert all(protocol_number in page and "Page" in page for page in protocol_pages)
        assert "1. TITLE PAGE" in protocol_pages[0]
        toc_heading = next(
            paragraph.text.strip() for paragraph in protocol.paragraphs
            if "TABLE OF CONTENTS" in paragraph.text
            and paragraph.style.name.casefold().startswith("heading")
        )
        introduction_heading = next(
            paragraph.text.strip() for paragraph in protocol.paragraphs
            if "INTRODUCTION" in paragraph.text
            and paragraph.style.name.casefold().startswith("heading")
        )
        toc_page = next(index for index, page in enumerate(protocol_pages) if toc_heading in page)
        assert toc_page > 0
        assert any(
            index > toc_page and introduction_heading in page
            for index, page in enumerate(protocol_pages)
        )
        sparse_pages = [page for page in protocol_pages if len(page.split()) < 50]
        assert all("Duration / Follow- up" in page for page in sparse_pages)
        if any("Table 9.2-1. Visit Schedule" in paragraph.text for paragraph in protocol.paragraphs):
            assert any(
                "Table 9.2-1. Visit Schedule" in page
                and "Visit Number Visit Name Visit Window CRF Number" in page
                for page in protocol_pages
            )
        if any("Table 13.3.-1" in paragraph.text for paragraph in protocol.paragraphs):
            assert any(
                "Table 13.3.-1" in page
                and "Study Staff Business Phone e-mail Office Phone" in page
                for page in protocol_pages
            )
        general_information = next((
            paragraph.text.strip() for paragraph in protocol.paragraphs
            if "GENERAL INFORMATION" in paragraph.text
            and paragraph.style.name.casefold().startswith("heading")
        ), None)
        if general_information is not None:
            section_three_page = next(
                page for page in protocol_pages
                if general_information in page and toc_heading not in page
            )
            assert "2. INVESTIGATOR AGREEMENT" not in section_three_page
            assert "Objective" in section_three_page
            assert re.search(r"Duration / Follow- ?up", section_three_page)
        numbered_body_headings = [
            paragraph
            for paragraph in protocol.paragraphs
            if paragraph.style.name.casefold().startswith("heading")
            and re.match(r"^\d+(?:\.\d+)*\.?\s+", paragraph.text.strip())
            and "TITLE PAGE" not in paragraph.text
            and "TABLE OF CONTENTS" not in paragraph.text
        ]
        assert numbered_body_headings
        body = list(protocol.element.body)
        normalized_pages = [re.sub(r"[^a-z0-9]+", " ", page.casefold()).strip() for page in protocol_pages]
        for heading in numbered_body_headings:
            heading_index = body.index(heading._p)
            first_content = None
            for element in body[heading_index + 1:]:
                if element.tag == qn("w:tbl"):
                    table = Table(element, protocol)
                    first_content = " ".join(cell.text for cell in table.rows[0].cells) if table.rows else ""
                    break
                if element.tag != qn("w:p"):
                    continue
                paragraph = Paragraph(element, protocol)
                if paragraph.style.name.casefold().startswith("heading"):
                    break
                if paragraph.text.strip():
                    first_content = paragraph.text
                    break
            if not first_content:
                continue
            heading_marker = " ".join(re.findall(r"[a-z0-9]+", heading.text.casefold()))
            content_marker = " ".join(re.findall(r"[a-z0-9]+", first_content.casefold())[:4])
            assert any(
                heading_marker in page
                and content_marker in page[page.index(heading_marker) + len(heading_marker):]
                for page in normalized_pages
            ), f"Rendered heading is orphaned from first content: {heading.text}"
        first_body_heading = next(paragraph for paragraph in numbered_body_headings if "INTRODUCTION" in paragraph.text)
        assert _page_boundary_before(first_body_heading)
        assert all(
            paragraph.paragraph_format.page_break_before is not True
            for paragraph in numbered_body_headings
            if paragraph is not first_body_heading
            and paragraph.text.strip() != "3. GENERAL INFORMATION"
        )
        assert all(
            paragraph.paragraph_format.keep_with_next is True
            or paragraph.style.paragraph_format.keep_with_next is True
            for paragraph in numbered_body_headings
        )
        long_body = next(
            paragraph
            for paragraph in sorted(protocol.paragraphs, key=lambda item: len(item.text), reverse=True)
            if len(paragraph.text) > 200 and not paragraph.style.name.casefold().startswith("heading")
        )
        assert long_body.paragraph_format.keep_together is not True

        if icf_family is not None:
            icf = Document(output / "candidate/icf.docx")
            rendered_icf = next(artifact for artifact in render_report["artifacts"] if artifact["artifact"] == "icf")
            icf_pages = [
                " ".join((page.extract_text() or "").split())
                for page in PdfReader(output / rendered_icf["pdf"]).pages
            ]
            assert all(protocol_number in page and "Page" in page for page in icf_pages)
            assert all(len(page.split()) >= 50 for page in icf_pages)
            icf_headings = [
                paragraph
                for paragraph in icf.paragraphs
                if paragraph.style.name == "Heading ICF Section"
            ]
            assert icf_headings
            assert all(paragraph.paragraph_format.keep_with_next is True for paragraph in icf_headings)


def test_protocol_headings_keep_their_first_content_and_front_matter_boundaries(tmp_path):
    reference = json.loads((ROOT / "tests/fixtures/ambispective-acceptance-source.json").read_text(encoding="utf-8"))

    render_documents(ROOT, tmp_path, reference, {"protocol": [], "icf": {}, "prs": {}})

    protocol = Document(tmp_path / "candidate/protocol.docx")
    headings = [
        paragraph for paragraph in protocol.paragraphs
        if paragraph.style.name.casefold().startswith("heading")
    ]
    assert headings
    assert all(
        paragraph.paragraph_format.keep_with_next is True
        or paragraph.style.paragraph_format.keep_with_next is True
        for paragraph in headings
    )
    summary_heading = next(
        paragraph for paragraph in headings
        if paragraph.text.strip() == "3. GENERAL INFORMATION"
    )
    summary_index = list(protocol.element.body).index(summary_heading._p)
    assert list(protocol.element.body)[summary_index + 1].tag == qn("w:tbl")

    title_heading = next(paragraph for paragraph in headings if paragraph.text.strip() == "1. TITLE PAGE")
    toc_heading = next(paragraph for paragraph in headings if paragraph.text.strip() == "4. TABLE OF CONTENTS")
    assert title_heading._p.xpath('following::w:br[@w:type="page"]')
    assert toc_heading._p.xpath('preceding::w:br[@w:type="page"]')

    visits_heading = next(
        paragraph for paragraph in headings
        if paragraph.text.strip().startswith("9.2. Visits and Examinations")
    )
    body = list(protocol.element.body)
    visits_index = body.index(visits_heading._p)
    table_index = next(
        index for index in range(visits_index + 1, len(body))
        if body[index].tag == qn("w:tbl")
    )
    protected_chain = [Paragraph(body[index], protocol) for index in range(visits_index, table_index)]
    assert all(paragraph.paragraph_format.keep_with_next is True for paragraph in protected_chain)


@pytest.mark.parametrize(
    "fixture",
    [
        "prospective-acceptance-source.json",
        "ambispective-acceptance-source.json",
    ],
)
def test_rendered_section_three_starts_after_investigator_agreement(
    tmp_path,
    governed_pdfium,
    fixture,
):
    reference = json.loads((ROOT / "tests/fixtures" / fixture).read_text(encoding="utf-8"))
    case_root = tmp_path / fixture.removesuffix(".json")
    render_documents(ROOT, case_root, reference, {"protocol": [], "icf": {}, "prs": {}})

    report = render_pages(case_root, page_renderer_identities=[governed_pdfium])

    assert report["status"] == "passed"
    protocol = next(item for item in report["artifacts"] if item["artifact"] == "protocol")
    assert protocol["docx_sha256"] == sha256_file(case_root / protocol["docx"])
    assert protocol["pdf_sha256"] == sha256_file(case_root / protocol["pdf"])
    assert all(page["sha256"] == sha256_file(case_root / page["path"]) for page in protocol["pages"])
    pages = [" ".join((page.extract_text() or "").split()) for page in PdfReader(case_root / protocol["pdf"]).pages]
    agreement_page_index = next(index for index, text in enumerate(pages) if "2. INVESTIGATOR AGREEMENT" in text)
    section_three_page_index = next(index for index, text in enumerate(pages) if "3. GENERAL INFORMATION" in text)
    toc_page_index = next(index for index, text in enumerate(pages) if "4. TABLE OF CONTENTS" in text)
    section_three_page = pages[section_three_page_index]
    section_sixteen_page = next(text for text in pages if "16. CONFIDENTIALITY" in text)
    visits_heading_page = next(
        text for text in pages
        if "9.2. Visits and Examinations" in text and "4. TABLE OF CONTENTS" not in text
    )
    assert section_three_page_index == agreement_page_index + 1
    assert "3. GENERAL INFORMATION" not in pages[agreement_page_index]
    assert "2. INVESTIGATOR AGREEMENT" not in section_three_page
    assert toc_page_index > section_three_page_index
    assert "Objective" in section_three_page
    assert re.search(r"Duration / Follow- ?up", section_three_page)
    assert "15. STANDARD EVALUATION PROCEDURES" in section_sixteen_page
    assert "Table 9.2-1. Visit Schedule" in visits_heading_page


def test_section_three_table_repair_moves_the_complete_block_before_the_toc(
    tmp_path,
    governed_pdfium,
):
    reference = json.loads(
        (ROOT / "tests/fixtures/prospective-acceptance-source.json").read_text(
            encoding="utf-8"
        )
    )
    reference["study"].update({
        "title": (
            "Prospective Evaluation of the NovaStep Activity Sensor in Adults "
            "Recovering From Total Knee Arthroplasty"
        ),
        "short_title": "NovaStep Recovery Study",
        "timeline": (
            "Enrollment is expected to last 8 months. Each participant is followed "
            "from screening through Week 12, for "
            + "additional scheduled follow-up context " * 16
            + "extra words approximately 11 weeks after the baseline device fitting."
        ),
    })
    reference["objectives"]["primary"] = [
        "Describe the change in average daily step count from baseline at Week 2 "
        "to Week 12 after total knee arthroplasty."
    ]
    reference["population"].update({
        "sample_size": "72 participants",
        "study_population": (
            "Adults recovering from primary unilateral total knee arthroplasty who "
            "can complete study visits and use the NovaStep sensor."
        ),
    })
    reference["design"].update({
        "number_of_sites": 2,
        "study_design": (
            "Prospective, multi-site, single-arm observational device study. The "
            "device is used only for measurement and does not direct treatment."
        ),
    })

    shared_finding = {
        "category": "visual",
        "artifact": "protocol",
        "element": "3. GENERAL INFORMATION – Variables / Secondary endpoint(s)",
        "target_ids": ["layout:protocol"],
        "recovery_class": "visual_defect",
        "action": "targeted_layout_repair",
    }
    layout_plan, unsupported = workflow._layout_repair_plan(
        [
            {
                **shared_finding,
                "check": "bad_table_split",
                "issue": "The Section 3 row continues alone on the next page.",
            },
            {
                **shared_finding,
                "check": "artificial_pagination",
                "issue": "The same continuation creates a nearly empty page before the TOC.",
            },
        ],
        study_type="Prospective",
    )
    assert unsupported == []
    assert layout_plan == {
        "protocol": [{"rule": "table_pagination", "target": "3. GENERAL INFORMATION"}],
    }

    document_report = render_documents(
        ROOT,
        tmp_path,
        reference,
        {"protocol": [], "icf": {}, "prs": {}},
        artifact_names={"protocol"},
        layout_repairs=layout_plan,
    )
    render_report = render_pages(
        tmp_path,
        page_renderer_identities=[governed_pdfium],
    )

    assert document_report["status"] == "passed"
    assert render_report["status"] == "passed", render_report
    protocol = next(
        item for item in render_report["artifacts"]
        if item["artifact"] == "protocol"
    )
    pages = [
        " ".join((page.extract_text() or "").split())
        for page in PdfReader(tmp_path / protocol["pdf"]).pages
    ]
    toc_page = next(
        index for index, text in enumerate(pages)
        if "4. TABLE OF CONTENTS" in text
    )
    section_three_page = pages[toc_page - 1]

    assert "3. GENERAL INFORMATION" in section_three_page
    assert "Duration / Follow-up" in section_three_page
    assert "device fitting." in section_three_page
    assert len(section_three_page.split()) >= 150


def test_content_gate_rejects_same_section_duplicates_and_flattened_schedule_prose(tmp_path):
    reference = json.loads((ROOT / "tests/fixtures/retrospective-acceptance-source.json").read_text(encoding="utf-8"))
    duplicate = "The analysis data sets will be organized around the approved outcome summaries."
    raw_schedule = (
        "For informed consent / subject enrollment, the approved visit schedule table is "
        "BL; Baseline; 1; Day 0; M3; Month 3; 2; Day 90 +/- 7."
    )
    model = {
        "protocol": [
            {"section_id": "study-procedure.enrollment", "paragraphs": [{"text": raw_schedule}], "lists": []},
            {
                "section_id": "analysis-plan.datasets",
                "paragraphs": [{"text": duplicate}, {"text": duplicate}],
                "lists": [],
            },
        ],
        "icf": {},
        "prs": {},
    }
    render_documents(ROOT, tmp_path, reference, model)

    findings = deterministic_content_check(tmp_path, reference)

    issues = "\n".join(item["issue"] for item in findings)
    assert "duplicated in protocol section analysis-plan.datasets" in issues
    assert "flattened visit-schedule serialization" in issues


def test_generated_documents_retain_client_typography_and_section_rhythm(tmp_path):
    reference = json.loads((ROOT / "tests/fixtures/ambispective-acceptance-source.json").read_text(encoding="utf-8"))
    protocol_text = "The approved analysis population includes source-supported study records and observations."
    risk_text = "Taking part may involve only the source-supported study risks described here."
    model = {
        "protocol": [{
            "section_id": "analysis-plan.datasets",
            "paragraphs": [{"text": protocol_text, "evidence_refs": [], "boilerplate_refs": []}],
            "lists": [],
        }],
        "icf": {
            "icf.risks": {
                "paragraphs": [{"text": risk_text, "evidence_refs": [], "boilerplate_refs": []}],
                "lists": [],
            }
        },
        "prs": {},
    }

    render_documents(ROOT, tmp_path, reference, model)

    protocol_path = tmp_path / "candidate/protocol.docx"
    protocol = Document(protocol_path)
    protocol_reference_path = ROOT / "assets/client-templates/reference/protocol-reference.docx"
    protocol_reference = Document(protocol_reference_path)
    assert _doc_default_font(protocol_path) == _doc_default_font(protocol_reference_path)
    assert protocol.styles["Normal"].font.size == protocol_reference.styles["Normal"].font.size
    protocol_body = next(paragraph for paragraph in protocol.paragraphs if paragraph.text == protocol_text)
    reference_analysis_index = next(
        index for index, paragraph in enumerate(protocol_reference.paragraphs)
        if paragraph.text.strip().startswith("10.1. Analysis Data Sets")
        and paragraph.style.name.casefold().startswith("heading")
    )
    reference_protocol_body = next(
        paragraph for paragraph in protocol_reference.paragraphs[reference_analysis_index + 1:]
        if paragraph.text.strip()
    )
    assert protocol_body.style.style_id == reference_protocol_body.style.style_id
    assert _paragraph_rhythm(protocol_body) == _paragraph_rhythm(reference_protocol_body)
    assert _run_typography(protocol_body) == _run_typography(reference_protocol_body)

    icf_path = tmp_path / "candidate/icf.docx"
    icf_reference_path = ROOT / "assets/client-templates/reference/advarra-icf-reference.docx"
    icf = Document(icf_path)
    icf_reference = Document(icf_reference_path)
    assert _doc_default_font(icf_path) == _doc_default_font(icf_reference_path)
    heading_text = "SIDE EFFECTS AND OTHER RISKS"
    output_heading_index = next(index for index, paragraph in enumerate(icf.paragraphs) if paragraph.text.strip() == heading_text)
    reference_heading_index = next(
        index for index, paragraph in enumerate(icf_reference.paragraphs) if paragraph.text.strip() == heading_text
    )
    output_heading = icf.paragraphs[output_heading_index]
    reference_heading = icf_reference.paragraphs[reference_heading_index]
    output_spacer = icf.paragraphs[output_heading_index + 1]
    reference_spacer = icf_reference.paragraphs[reference_heading_index + 1]
    output_body = next(paragraph for paragraph in icf.paragraphs[output_heading_index + 1:] if paragraph.text.strip())
    reference_body = next(
        paragraph for paragraph in icf_reference.paragraphs[reference_heading_index + 1:] if paragraph.text.strip()
    )

    assert output_body.text == risk_text
    assert not output_spacer.text.strip()
    assert _paragraph_rhythm(output_spacer) == _paragraph_rhythm(reference_spacer)
    assert output_heading.style.name == "Heading ICF Section"
    assert output_heading.paragraph_format.alignment == WD_ALIGN_PARAGRAPH.LEFT
    assert _points(output_heading.paragraph_format.right_indent) in (None, 0)
    assert _points(output_heading.paragraph_format.left_indent) == 0.0
    assert _points(output_heading.paragraph_format.first_line_indent) == 0.0
    assert _points(output_heading.paragraph_format.space_after) == _points(reference_heading.paragraph_format.space_after)
    expected_body_rhythm = _paragraph_rhythm(reference_body)
    expected_body_rhythm.update(left_indent=0.0, first_line_indent=0.0)
    assert _paragraph_rhythm(output_body) == expected_body_rhythm
    assert output_body.paragraph_format.alignment == WD_ALIGN_PARAGRAPH.JUSTIFY
    assert _run_typography(output_heading) == _run_typography(reference_heading)
    assert _run_typography(output_body) == _run_typography(reference_body)


def test_protocol_investigator_agreement_uses_client_intro_and_bullet_designs(tmp_path):
    reference = json.loads((ROOT / "tests/fixtures/ambispective-acceptance-source.json").read_text(encoding="utf-8"))
    render_documents(ROOT, tmp_path, reference, {"protocol": [], "icf": {}, "prs": {}})

    output = Document(tmp_path / "candidate/protocol.docx")
    authority = Document(ROOT / "assets/client-templates/reference/protocol-reference.docx")
    output_heading_index = next(
        index for index, paragraph in enumerate(output.paragraphs)
        if "INVESTIGATOR AGREEMENT" in paragraph.text and paragraph.style.name.casefold().startswith("heading")
    )
    authority_heading_index = next(
        index for index, paragraph in enumerate(authority.paragraphs)
        if "INVESTIGATOR AGREEMENT" in paragraph.text and paragraph.style.name.casefold().startswith("heading")
    )
    output_body = [paragraph for paragraph in output.paragraphs[output_heading_index + 1:] if paragraph.text.strip()][:4]
    authority_body = [paragraph for paragraph in authority.paragraphs[authority_heading_index + 1:] if paragraph.text.strip()][:4]

    assert len(output_body) == len(authority_body) == 4
    assert _paragraph_rhythm(output_body[0]) == _paragraph_rhythm(authority_body[0])
    assert _run_typography(output_body[0]) == _run_typography(authority_body[0])
    assert all(paragraph._p.get_or_add_pPr().find(qn("w:numPr")) is not None for paragraph in output_body[1:])
    assert all(_run_typography(output_body[index]) == _run_typography(authority_body[index]) for index in range(1, 4))
    assert all(
        _numbering_level_signature(output, output_body[index])
        == _numbering_level_signature(authority, authority_body[index])
        for index in range(1, 4)
    )


@pytest.mark.parametrize(
    "fixture",
    [
        "prospective-acceptance-source.json",
        "ambispective-acceptance-source.json",
    ],
)
def test_protocol_section_three_has_the_only_new_body_heading_page_break(tmp_path, fixture):
    reference = json.loads((ROOT / "tests/fixtures" / fixture).read_text(encoding="utf-8"))
    case_root = tmp_path / fixture.removesuffix(".json")
    render_documents(ROOT, case_root, reference, {"protocol": [], "icf": {}, "prs": {}})
    document = Document(case_root / "candidate/protocol.docx")
    numbered_headings = [
        paragraph
        for paragraph in document.paragraphs
        if paragraph.style.name.casefold().startswith("heading")
        and re.match(r"^\d+(?:\.\d+)*\.?\s+", paragraph.text.strip())
    ]
    section_three = next(
        paragraph for paragraph in numbered_headings
        if paragraph.text.strip() == "3. GENERAL INFORMATION"
    )

    assert section_three.paragraph_format.page_break_before is True
    assert all(
        paragraph.paragraph_format.page_break_before is not True
        for paragraph in numbered_headings
        if paragraph is not section_three
        and paragraph.text.strip() not in {"1. TITLE PAGE", "4. TABLE OF CONTENTS"}
    )


def test_general_information_table_targeting_uses_section_structure_not_optional_labels():
    document = Document()
    document.add_heading("3. GENERAL INFORMATION", level=1)
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Overview"
    table.cell(0, 1).text = "Study summary"
    table.cell(1, 0).text = "Duration"
    table.cell(1, 1).text = "Twelve weeks"
    document.add_heading("4. TABLE OF CONTENTS", level=1)

    rendering._normalize_protocol_summary_table(document, {"sites": []})

    for row in table.rows:
        assert row._tr.get_or_add_trPr().find(qn("w:cantSplit")) is not None
        for cell in row.cells:
            margins = cell._tc.get_or_add_tcPr().find(qn("w:tcMar"))
            assert margins.find(qn("w:top")).get(qn("w:w")) == "40"
            assert margins.find(qn("w:bottom")).get(qn("w:w")) == "40"


@pytest.mark.parametrize(
    "fixture",
    [
        "prospective-acceptance-source.json",
        "ambispective-acceptance-source.json",
    ],
)
def test_protocol_general_information_rows_have_small_uniform_vertical_padding(
    tmp_path,
    fixture,
):
    reference = json.loads((ROOT / "tests/fixtures" / fixture).read_text(encoding="utf-8"))
    case_root = tmp_path / fixture.removesuffix(".json")
    render_documents(ROOT, case_root, reference, {"protocol": [], "icf": {}, "prs": {}})
    document = Document(case_root / "candidate/protocol.docx")
    table = next(
        item for item in document.tables
        if item.rows
        and item.rows[0].cells[0].text.strip() == "Objective"
        and any(row.cells[0].text.strip() == "Variables" for row in item.rows)
    )

    assert table.style is not None
    duration_label = next(
        row.cells[0].text.strip()
        for row in table.rows
        if "duration" in row.cells[0].text.casefold()
    )
    assert duration_label == "Duration / Follow‑up"
    for row in table.rows:
        assert row.height is None
        assert row._tr.get_or_add_trPr().find(qn("w:cantSplit")) is not None
        for cell in row.cells:
            margins = cell._tc.get_or_add_tcPr().find(qn("w:tcMar"))
            assert margins is not None
            for side in ("top", "bottom"):
                margin = margins.find(qn(f"w:{side}"))
                assert margin is not None
                assert margin.get(qn("w:w")) == "40"
                assert margin.get(qn("w:type")) == "dxa"
            for paragraph in cell.paragraphs:
                assert paragraph.paragraph_format.space_before in (None, Pt(0))
                assert paragraph.paragraph_format.space_after is None


@pytest.mark.parametrize(
    "fixture",
    [
        "prospective-acceptance-source.json",
        "ambispective-acceptance-source.json",
        "retrospective-acceptance-source.json",
    ],
)
def test_protocol_investigator_obligations_have_paragraph_spacing_before_signature_table(
    tmp_path,
    fixture,
):
    reference = json.loads((ROOT / "tests/fixtures" / fixture).read_text(encoding="utf-8"))
    case_root = tmp_path / fixture.removesuffix(".json")
    render_documents(ROOT, case_root, reference, {"protocol": [], "icf": {}, "prs": {}})
    document = Document(case_root / "candidate/protocol.docx")
    table = next(
        item for item in document.tables
        if any("Signature of Investigator" in cell.text for row in item.rows for cell in row.cells)
    )
    previous = table._tbl.getprevious()

    assert previous is not None and previous.tag == qn("w:p")
    final_obligation = Paragraph(previous, document)
    assert final_obligation.text.strip()
    assert final_obligation._p.get_or_add_pPr().find(qn("w:numPr")) is not None
    assert _points(final_obligation.paragraph_format.space_after) == 12.0
    assert not final_obligation._p.xpath('.//w:br | .//w:tab')


def test_protocol_signature_block_retains_client_completion_lines(tmp_path):
    reference = json.loads((ROOT / "tests/fixtures/ambispective-acceptance-source.json").read_text(encoding="utf-8"))
    render_documents(ROOT, tmp_path, reference, {"protocol": [], "icf": {}, "prs": {}})

    output = Document(tmp_path / "candidate/protocol.docx")
    authority = Document(ROOT / "assets/client-templates/reference/protocol-reference.docx")
    output_table = next(
        table for table in output.tables
        if any("Signature of Investigator" in cell.text for row in table.rows for cell in row.cells)
    )
    authority_table = next(
        table for table in authority.tables
        if any("Signature of Investigator" in cell.text for row in table.rows for cell in row.cells)
    )

    output_text = " ".join(cell.text for row in output_table.rows for cell in row.cells)
    assert "Alex Investigator" in output_text
    assert "MD" in output_text
    assert "Site One" in output_text
    assert output_table.style.style_id == authority_table.style.style_id
    assert [column.w for column in output_table._tbl.tblGrid.gridCol_lst] == [
        column.w for column in authority_table._tbl.tblGrid.gridCol_lst
    ]


def test_protocol_omits_empty_references_section_without_inventing_citations(tmp_path):
    reference = json.loads((ROOT / "tests/fixtures/ambispective-acceptance-source.json").read_text(encoding="utf-8"))
    reference.pop("references", None)
    render_documents(ROOT, tmp_path, reference, {"protocol": [], "icf": {}, "prs": {}})

    output = Document(tmp_path / "candidate/protocol.docx")
    visible = "\n".join(paragraph.text for paragraph in output.paragraphs)
    assert "REFERENCES" not in visible
    assert "Source of Truth" not in visible


def test_sterling_icf_removes_template_review_highlighting(tmp_path):
    reference = json.loads((ROOT / "tests/fixtures/prospective-acceptance-source.json").read_text(encoding="utf-8"))
    reference["meta"]["icf_template"] = "Sterling"
    model = {"protocol": [], "prs": {}, "icf": {
        "icf.risks": {"paragraphs": [{"text": "Authorized sparse risk text."}], "lists": []},
        "icf.benefits": {"paragraphs": [{"text": "Authorized benefit text."}], "lists": []},
        "icf.alternatives": {"paragraphs": [{"text": "Authorized alternatives text."}], "lists": []},
        "icf.costs": {"paragraphs": [{"text": "Authorized cost text."}], "lists": []},
    }}
    render_documents(ROOT, tmp_path, reference, model)

    output = Document(tmp_path / "candidate/icf.docx")
    runs = [run for paragraph in output.paragraphs for run in paragraph.runs]

    assert all(run.font.highlight_color is None for run in runs)


def test_protocol_running_header_uses_concise_title_and_nonwrapping_page_control(tmp_path):
    reference = json.loads((ROOT / "tests/fixtures/ambispective-acceptance-source.json").read_text(encoding="utf-8"))
    reference["study"].pop("short_title", None)
    reference["study"]["title"] = "Ambispective Evaluation of the Sentinel Patch in Adults With Postoperative Pain"
    render_documents(ROOT, tmp_path, reference, {"protocol": [], "icf": {}, "prs": {}})

    output = Document(tmp_path / "candidate/protocol.docx")
    header = output.sections[0].header.tables[0]
    assert header.cell(0, 0).text.strip() == "Sentinel Patch in Adults With Postoperative Pain"
    assert header.cell(0, 1)._tc.tcPr.find(qn("w:noWrap")) is not None
    page_control = header.cell(0, 1).paragraphs[0]
    assert "PAGE" in page_control._p.xml
    assert "NUMPAGES" in page_control._p.xml
    assert not page_control._p.xpath('.//w:br')


def test_protocol_cached_toc_uses_client_dot_leaders_and_subsection_indent(tmp_path):
    reference = json.loads((ROOT / "tests/fixtures/ambispective-acceptance-source.json").read_text(encoding="utf-8"))
    render_documents(ROOT, tmp_path, reference, {"protocol": [], "icf": {}, "prs": {}})
    protocol_path = tmp_path / "candidate/protocol.docx"
    pdf_path = tmp_path / "protocol.pdf"
    writer = PdfWriter()
    writer.add_blank_page(width=612, height=792)
    with pdf_path.open("wb") as handle:
        writer.write(handle)

    assert refresh_toc_from_pdf(protocol_path, pdf_path) is True
    output = Document(protocol_path)
    authority = Document(ROOT / "assets/client-templates/reference/protocol-reference.docx")
    toc_rows = [paragraph for paragraph in output.paragraphs if paragraph.style.name.casefold() in {"toc 1", "toc 2"}]
    heading_one = next(paragraph for paragraph in toc_rows if paragraph.text.strip().startswith("7. SUBJECTS"))
    heading_two = next(paragraph for paragraph in toc_rows if paragraph.text.strip().startswith("7.1."))
    assert _points(heading_one.paragraph_format.left_indent) == _points(authority.styles["toc 1"].paragraph_format.left_indent)
    assert _points(heading_two.paragraph_format.left_indent) == _points(authority.styles["toc 2"].paragraph_format.left_indent)
    assert heading_one.text.endswith("\t")  # no invented page for blank PDF evidence
    assert "......" not in heading_one.text
    assert 'TOC \\o "1-2"' in "\n".join(paragraph._p.xml for paragraph in toc_rows)
    assert not any(
        [cell.text.strip() for cell in table.rows[0].cells] == ["Section", "Page"]
        for table in output.tables
    )


def test_protocol_duplicate_titles_and_table_sections_use_client_body_design(tmp_path):
    reference = json.loads((ROOT / "tests/fixtures/ambispective-acceptance-source.json").read_text(encoding="utf-8"))
    safety_text = "Source-supported general safety information appears in this section."
    assessment_text = "Source-supported assessment information introduces the schedule table."
    model = {
        "protocol": [
            {"section_id": "quality-safety.general", "paragraphs": [{"text": safety_text}], "lists": []},
            {"section_id": "evaluation-procedures", "paragraphs": [{"text": assessment_text}], "lists": []},
        ],
        "icf": {},
        "prs": {},
    }
    render_documents(ROOT, tmp_path, reference, model)

    output = Document(tmp_path / "candidate/protocol.docx")
    authority = Document(ROOT / "assets/client-templates/reference/protocol-reference.docx")
    authority_heading_index = next(
        index for index, paragraph in enumerate(authority.paragraphs)
        if paragraph.text.strip().startswith("9.1. Informed Consent")
        and paragraph.style.name.casefold().startswith("heading")
    )
    authority_body = next(
        paragraph for paragraph in authority.paragraphs[authority_heading_index + 1:]
        if paragraph.text.strip()
    )

    for text in (safety_text, assessment_text):
        paragraph = next(item for item in output.paragraphs if item.text == text)
        assert paragraph.style.style_id == authority_body.style.style_id
        assert _paragraph_rhythm(paragraph) == _paragraph_rhythm(authority_body)
        assert _run_typography(paragraph) == _run_typography(authority_body)
