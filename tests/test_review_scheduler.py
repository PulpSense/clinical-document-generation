"""Scheduler capacity and recovery through public request/Desktop interfaces."""
import json
import pytest
import quality
from visual_review import IncompleteVisualReview
from test_content_review_recovery import operation, candidate, FAMILIES


@pytest.mark.parametrize('family', FAMILIES)
def test_limited_reviewers_finish_every_required_assessment(tmp_path, monkeypatch, family):
    scopes = []
    def worker(request, response, revision, *args):
        assert len(request['sections']) <= 3, 'Worker capacity exceeded by oversized assignment'
        assert len(request['cross_document_checks']) <= 2
        assert request.get('readable_native_evidence'), 'Native passages must be supplied'
        scopes.append([(r['artifact'], r['section_id']) for r in request['sections']])
        # This worker can complete only one section per invocation.
        response['section_assessments'] = response['section_assessments'][:1]
        return response
    result, revision, fallback = operation(tmp_path, monkeypatch, family, worker, worker)
    assert result['status'] == 'passed', result
    parent_path = next((revision / 'hermes/verification-requests').glob('*content.json'))
    parent = json.loads(parent_path.read_text())
    response = json.loads((revision / parent['response_path']).read_text())
    assert len(response['section_assessments']) == len(parent['sections'])
    assert len(response['cross_document_assessments']) == len(parent['cross_document_checks'])
    assert quality.verification_response_is_complete(revision, parent_path)
    assert len(scopes) > 2  # Progress is not stopped by a package-wide retry cap.


@pytest.mark.parametrize('family', FAMILIES)
def test_partial_native_defect_reaches_repair_and_fresh_complete_review(tmp_path, monkeypatch, family):
    from docx import Document
    from test_clean_delivery import controlled_render
    from test_general_document_repair import repair_response
    corrupted = []
    reviewed_sets = []
    repairs = []
    omitted_reassessment = []
    def render(root, revision, *args, **kwargs):
        if not corrupted:
            path = revision / 'candidate/protocol.docx'
            doc = Document(path)
            doc.paragraphs[0].text = 'UNSUPPORTED INVENTED COVER TITLE'
            doc.save(path)
            corrupted.append(True)
        return controlled_render(root, revision, *args, **kwargs)
    def worker(request, response, revision, *args):
        reviewed_sets.append(request['review_set'])
        doc = Document(revision / 'candidate/protocol.docx')
        if any(p.text == 'UNSUPPORTED INVENTED COVER TITLE' for p in doc.paragraphs):
            row = next((r for r in response['section_assessments'] if r['section_id'] == 'title-page'), None)
            if row:
                row['status'] = 'blocked'
                response['status'] = 'blocked'
                response['section_assessments'] = [row]
                response['findings'] = [{'finding_id': 'cover-title', 'artifact': 'protocol',
                    'target_ids': ['title-page'], 'check': 'source_supported',
                    'issue': 'The native cover title is not approved.', 'material': True,
                    'affected_passage': 'UNSUPPORTED INVENTED COVER TITLE', 'source_path': 'study.title',
                    'source_excerpt': request['approved_source']['study']['title']}]
        skip_notes = (family[0] == 'Retrospective' and request['review_set'] > 1
                      and any(r['section_id'] == 'title-page' for r in request['sections'])
                      and not omitted_reassessment)
        if skip_notes:
            omitted_reassessment.append(True)
        for row in response['section_assessments']:
            row['notes'] = [f"Independently reassessed {f['finding_id']} against current native bytes."
                            for f in request.get('review_context', {}).get('prior_findings', [])] if not skip_notes else []
        return response
    def repair(request, revision):
        repairs.append(request)
        path = revision / request['workspace_path'] / 'protocol.docx'
        doc = Document(path)
        doc.paragraphs[0].text = request['approved_source']['study']['title']
        doc.save(path)
        return repair_response(request)
    result, revision, _ = operation(tmp_path, monkeypatch, family, worker, worker, repair=repair, render=render)
    assert result['status'] == 'passed', result
    assert len(repairs) == 1
    assert set(reviewed_sets) == {1, 2}
    first = json.loads(next((revision / 'hermes/review-schedules').glob('*review-1*')).read_text())
    assert first['waves'][-1]['completed_assessments'] < len(quality.content_review_sections(repairs[0]['approved_source']))
    assert all(p.text != 'UNSUPPORTED INVENTED COVER TITLE' for p in Document(tmp_path / 'output/protocol.docx').paragraphs)


@pytest.mark.parametrize('failure', ['wrong-hash', 'missing-check', 'missing-response', 'bare-blocked'])
def test_bad_unit_is_retried_without_discarding_other_progress(tmp_path, monkeypatch, failure):
    damaged = []
    scopes = []
    def worker(request, response, revision, *args):
        scopes.append(request['request_id'])
        if not damaged:
            damaged.append(request['review_unit_id'])
            if failure == 'wrong-hash': response['request_sha256'] = '0' * 64
            elif failure == 'missing-check': response['section_assessments'][0]['checks'] = []
            elif failure == 'missing-response': return None
            else: response['status'] = 'blocked'
        return response
    result, revision, _ = operation(tmp_path, monkeypatch, FAMILIES[-1], worker, worker)
    assert result['status'] == 'passed', result
    assert sum('.unit-' + damaged[0] + '.' in name for name in scopes) == 2
    ledger = json.loads(next((revision / 'hermes/review-schedules').glob('*.json')).read_text())
    assert ledger['waves'][-1]['completed_assessments'] == len(quality.content_review_sections(json.loads((revision / 'approved-reference.json').read_text())))


def test_total_worker_failure_stops_with_complete_review_only_package(tmp_path, monkeypatch):
    def failed(*args): raise RuntimeError('Provider unavailable')
    result, revision, _ = operation(tmp_path, monkeypatch, FAMILIES[-1], failed, failed)
    assert result['status'] == 'review_required', result
    assert (tmp_path / 'review-output/REVIEW-ONLY-protocol.docx').is_file()
    ledger = json.loads(next((revision / 'hermes/review-schedules').glob('*.json')).read_text())
    assert ledger['no_progress'] == 2
    assert all(w['gained_assessments'] == 0 for w in ledger['waves'])


def test_saved_unit_progress_survives_worker_context_loss(tmp_path, monkeypatch):
    def worker(request, response, revision, *args):
        checkpoint = dict(response)
        checkpoint['status'] = 'in_progress'
        checkpoint['section_assessments'] = response['section_assessments'][:1]
        checkpoint['cross_document_assessments'] = response['cross_document_assessments'][:1]
        path = revision / request['review_transport']['progress_path']
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_text(json.dumps(checkpoint))
        return None  # Killed before any final response.
    result, revision, _ = operation(tmp_path, monkeypatch, FAMILIES[-1], worker, worker)
    assert result['status'] == 'passed', result
    assert list((revision / 'hermes/review-progress').glob('*.json'))


def test_slow_limited_workers_finish_inside_original_deadline(tmp_path, monkeypatch):
    now = [0.0]
    def worker(request, response, revision, *args):
        now[0] += 2.0
        response['section_assessments'] = response['section_assessments'][:1]
        return response
    result, revision, _ = operation(tmp_path, monkeypatch, FAMILIES[0], worker, worker, clock=lambda: now[0])
    assert result['status'] == 'passed', result
    state = json.loads((tmp_path / 'logs/desktop-operation.json').read_text())
    assert state['budget_seconds'] == 180
    assert state['deadline_at_epoch'] - state['started_at_epoch'] == 180
    assert 0 < now[0] < 150


@pytest.mark.parametrize('tamper', ['unit-response', 'parent-assessment', 'native-file', 'readable-view'])
def test_delivered_review_cannot_survive_changed_evidence(tmp_path, monkeypatch, tamper):
    result, revision, _ = operation(tmp_path, monkeypatch, FAMILIES[-1],
                                    lambda r, s, d: s, lambda r, s, d, h: s)
    assert result['status'] == 'passed', result
    parent_path = next((revision / 'hermes/verification-requests').glob('*content.json'))
    parent = json.loads(parent_path.read_text())
    response_path = revision / parent['response_path']
    response = json.loads(response_path.read_text())
    if tamper == 'unit-response':
        path = revision / response['review_unit_evidence'][0]['response']
        path.write_bytes(path.read_bytes() + b' ')
    elif tamper == 'parent-assessment':
        response['section_assessments'][0]['notes'] = ['Invented approval']
        response_path.write_text(json.dumps(response))
    elif tamper == 'native-file':
        path = revision / parent['artifacts'][0]['path']
        path.write_bytes(path.read_bytes() + b'changed')
    else:
        path = revision / parent['review_plan']['readable_native'][0]['path']
        path.write_bytes(path.read_bytes() + b'changed')
    assert not quality.verification_response_is_complete(revision, parent_path)


def test_expired_operation_does_not_dispatch_more_units(tmp_path, monkeypatch):
    now = [0.0]
    dispatches = []
    def worker(request, response, revision, *args):
        dispatches.append(request['request_id'])
        now[0] = 181.0
        return None
    result, revision, _ = operation(tmp_path, monkeypatch, FAMILIES[-1], worker, worker, clock=lambda: now[0])
    assert result['status'] in {'review_required', 'timeout'}, result
    assert len(dispatches) <= 3  # Only the already dispatched concurrent wave.


def test_scheduler_finishes_concurrently_with_slow_visual_fallback(tmp_path, monkeypatch):
    import threading
    now = [0.0]
    last_cross_started = threading.Event()
    def worker(request, response, revision, *args):
        if 'risks_benefits' in request['cross_document_checks']:
            last_cross_started.set()
        return response
    def visual(request, response, revision, fallback):
        if not fallback:
            raise IncompleteVisualReview('Delegated visual worker did not respond')
        assert last_cross_started.wait(30), 'Content scheduler was blocked behind parent visual fallback'
        now[0] = 179.0
        return response
    result, revision, _ = operation(tmp_path, monkeypatch, FAMILIES[0], worker, worker,
        clock=lambda: now[0], visual_review=visual)
    assert result['status'] == 'passed', result


def test_generic_visual_execution_failure_remains_fatal(tmp_path, monkeypatch):
    calls = []
    def visual(request, response, revision, fallback):
        calls.append(fallback)
        raise RuntimeError('Controlled generic visual execution failure')
    result, _, _ = operation(tmp_path, monkeypatch, FAMILIES[-1],
        lambda request, response, revision: response,
        lambda request, response, revision, handoff: response,
        visual_review=visual)
    assert calls == [False], 'Generic execution failure must not invoke recovery fallback'
    assert result['status'] == 'review_required', result
    assert result['stage'] == 'parent_visual_integrity', result
    assert not result.get('client_outputs')
    assert any(row['issue'] == 'Controlled generic visual execution failure'
               for row in result['findings'])


def test_interrupted_parent_resumes_units_without_resetting_approval_or_deadline(tmp_path, monkeypatch):
    from pathlib import Path
    import workflow
    from test_handoff_quality import acceptance_verification
    class InterruptedParent(BaseException): pass
    completed = []
    def interrupt(request, response, revision, *args):
        completed.extend((r['artifact'], r['section_id']) for r in request['sections'])
        (revision / request['response_path']).parent.mkdir(parents=True, exist_ok=True)
        (revision / request['response_path']).write_text(json.dumps(response))
        raise InterruptedParent()
    with pytest.raises(InterruptedParent):
        operation(tmp_path, monkeypatch, FAMILIES[-1], interrupt, interrupt)
    state_path = tmp_path / 'logs/desktop-operation.json'
    original = json.loads(state_path.read_text())
    revision = next((tmp_path / 'revisions').iterdir())
    approved_bytes = (revision / 'approved-reference.json').read_bytes()
    inspected_after_resume = []
    def resume(items, remaining):
        for handoff in items:
            request = json.loads((revision / handoff['request_path']).read_text())
            inspected_after_resume.extend((r['artifact'], r['section_id']) for r in request.get('sections', []))
            (revision / request['response_path']).parent.mkdir(parents=True, exist_ok=True)
            (revision / request['response_path']).write_text(json.dumps(acceptance_verification(request)))
    result = workflow.run_desktop_operation(tmp_path, handoff_runner=resume,
        fallback_handoff_runner=resume, opener=lambda path: Path(path).read_bytes(),
        budget_seconds=180, require_promoted_runtime=False)
    assert result['status'] == 'passed', result
    assert not set(completed) & set(inspected_after_resume)
    final = json.loads(state_path.read_text())
    assert final['started_at_epoch'] == original['started_at_epoch']
    assert final['deadline_at_epoch'] == original['deadline_at_epoch']
    assert final['approval_identity'] == original['approval_identity']
    assert (revision / 'approved-reference.json').read_bytes() == approved_bytes


def test_nonmaterial_finding_does_not_stop_required_review(tmp_path, monkeypatch):
    emitted = []
    def worker(request, response, revision, *args):
        if not emitted:
            emitted.append(True)
            response['findings'] = [{'finding_id': 'editorial', 'artifact': 'protocol',
                'target_ids': [request['sections'][0]['section_id']], 'issue': 'Harmless style preference',
                **{key: False for key in ('material', 'safety_critical', 'contradiction',
                    'obscures_required_information', 'materially_unusable')}}]
        return response
    result, revision, _ = operation(tmp_path, monkeypatch, FAMILIES[-1], worker, worker)
    assert result['status'] == 'passed', result
    assert not result.get('warnings')
    parent = next((revision / 'hermes/verification-requests').glob('*content.json'))
    assert quality.verification_response_is_complete(revision, parent)


@pytest.mark.parametrize('family,location', [(family, location) for family in FAMILIES
    for location in ('cross', 'unknown-section', 'xml')
    if location != 'xml' or family[0] != 'Retrospective'])
def test_prior_concerns_reach_responsible_unit_and_require_reassessment(tmp_path, family, location):
    from test_handoff_quality import acceptance_verification
    _, _, source = candidate(tmp_path, family)
    concern = {'finding_id': 'prior-fact', 'issue': 'Prior substantive concern', 'material': True}
    if location == 'cross':
        concern['check'] = 'cross_document_consistency'
    elif location == 'unknown-section':
        concern.update(artifact='protocol', target_ids=['unknown-native-location'])
    else:
        concern.update(artifact='study.xml', target_ids=['unknown-xml-location'])
    paths = quality.create_verification_requests(tmp_path, source, {'artifacts': []}, review_set=2,
        review_context={'prior_findings': [concern]})
    parent_path = paths[0]
    parent = json.loads(parent_path.read_text())
    units = [quality.create_content_review_unit(tmp_path, parent_path, [row], [], 1)
             for row in parent['sections']]
    units += [quality.create_content_review_unit(tmp_path, parent_path, [], [check], 1)
              for check in parent['cross_document_checks']]
    responsible = []
    for path in units:
        request = json.loads(path.read_text())
        if not request['review_context']['prior_findings']:
            continue
        responsible.append(path)
        response = acceptance_verification(request)
        response_path = tmp_path / request['response_path']
        response_path.parent.mkdir(parents=True, exist_ok=True)
        response_path.write_text(json.dumps(response))
        assert quality.content_review_unit_result(tmp_path, path) is None
        for row in [*response['section_assessments'], *response['cross_document_assessments']]:
            row['notes'] = ['Independently reassessed prior-fact against the exact current native bytes.']
        response_path.write_text(json.dumps(response))
        assert quality.content_review_unit_result(tmp_path, path)['complete']
    assert responsible, 'No reviewer received the prior concern'
