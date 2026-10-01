"""Real table reconstruction receives the finding before a derivative layout stop."""
import json
from pathlib import Path

from docx import Document
import pytest

import quality
import rendering
import workflow
from test_handoff_quality import acceptance_verification
from test_mixed_visit_representations import failed_source


ROOT = Path(__file__).resolve().parents[1]


def review_case(tmp_path, *, table_error=True):
    source = failed_source()
    run = tmp_path / "run"
    revision = run / "revisions/r-owner"
    reference_path = run / "reference/study.reference.json"
    reference_path.parent.mkdir(parents=True)
    reference_path.write_text(json.dumps({"generation": {}}))
    rendering.render_documents(ROOT, revision, source, {"protocol": [], "icf": {}, "prs": {}})
    (revision / "candidate/study.xml").write_text("<clinical_study/>")
    accepted = revision / "hermes/accepted"
    accepted.mkdir(parents=True)
    (accepted / "study-procedure.visits.json").write_text(json.dumps({"paragraphs": [
        {"text": "Unchanged visits prose."}]}))
    protocol = revision / "candidate/protocol.docx"
    doc = Document(protocol)
    table = next(t for t in doc.tables if t.rows[0].cells[0].text == "Visit Number")
    if table_error:
        table.add_row().cells[1].text = "Duplicate Month 3"
        doc.save(protocol)
    render_report = {"artifacts": []}
    for artifact in ("protocol", "icf"):
        pdf = revision / f"rendered/{artifact}.pdf"
        page = revision / f"rendered/{artifact}/page-1.png"
        page.parent.mkdir(parents=True)
        pdf.write_bytes(b"bound-pdf")
        page.write_bytes(b"bound-page")
        render_report["artifacts"].append({"artifact": artifact, "docx": f"candidate/{artifact}.docx",
            "pdf": f"rendered/{artifact}.pdf", "docx_sha256": quality.sha256_file(revision / f"candidate/{artifact}.docx"),
            "pdf_sha256": quality.sha256_file(pdf), "pages": [{"page": 1,
            "path": f"rendered/{artifact}/page-1.png", "sha256": quality.sha256_file(page)}]})
    (revision / "candidate-build.json").write_text(json.dumps({"render_report": render_report}))
    paths = quality.create_verification_requests(revision, source, render_report)
    for path in paths:
        request = json.loads(path.read_text())
        response = acceptance_verification(request)
        if request["task"] == "clinical_content_verification" and table_error:
            response["status"] = "blocked"
            response["findings"] = [{"finding_id": "extra-month3-visit", "category": "content",
                "artifact": "protocol", "check": "source_supported", "surface": "table",
                "table_ids": ["visit-schedule"], "target_ids": ["study-procedure.visits"],
                "issue": "Table 9.2-1 invents a duplicate Month 3 visit.", "contradiction": True}]
            for row in response["section_assessments"]:
                if row["section_id"] == "study-procedure.visits": row["status"] = "blocked"
        elif request["task"] == "rendered_page_visual_verification" and request["artifacts"][0]["artifact"] == "protocol":
            response["status"] = "blocked"
            response["findings"] = [{"finding_id": "table-header-wrap", "category": "visual",
                "artifact": "protocol", "page": 1, "check": "unreadable_text",
                "element": "Table 15.1. Proposed Visits and Study Assessments",
                "target_ids": ["layout:protocol"], "issue": "The crowded header breaks a word."}]
            response["page_assessments"][0]["status"] = "blocked"
        (revision / request["response_path"]).write_text(json.dumps(response))
    findings, _ = quality.validate_verifications(revision)
    return run, revision, reference_path, source, protocol, findings


def test_table_finding_routes_to_real_reconstruction_before_visual_stop(tmp_path, monkeypatch):
    run, revision, reference_path, source, protocol, findings = review_case(tmp_path)
    table_finding = next(f for f in findings if f.get("surface") == "table")
    assert table_finding["recovery_class"] == "deterministic_structure_defect"
    assert not any(f.get("recovery_class") == "drafting_defect" for f in findings)

    def reconstruct(_run, **kwargs):
        rendering.render_documents(ROOT, revision, source, {"protocol": [], "icf": {}, "prs": {}}, artifact_names={"protocol"})
        state = json.loads(reference_path.read_text())
        workflow._complete_pending_deterministic_reconstructions(reference_path, state, revision, outcome="rebuilt")
        assert workflow._complete_pending_recovery_attempts(revision, reference_path, state, require_candidate_change=True) == []
        return {"status": "awaiting_hermes", "stage": "independent_verification"}

    monkeypatch.setattr(workflow, "generate", reconstruct)
    result = workflow._quality_retry(run, reference_path, {"generation": {}}, source, revision, {}, findings, "quality", require_promoted_runtime=False)
    assert result["stage"] == "independent_verification"
    table = next(t for t in Document(protocol).tables if t.rows[0].cells[0].text == "Visit Number")
    assert [r.cells[1].text for r in table.rows[1:]] == ["Preoperative screening", "Operative visit for each eye", "Month 3 postoperative"]
    state = json.loads(reference_path.read_text())
    assert state["generation"]["deferred_visual_findings"][0]["status"] == "requires_fresh_exact_artifact_review"
    archived = list((revision / "attempts").rglob("*verify.visual.protocol.json"))
    assert any("table-header-wrap" in p.read_text() for p in archived)


def test_visual_only_failure_is_not_silently_cleared(tmp_path):
    run, revision, reference_path, source, protocol, findings = review_case(tmp_path, table_error=False)
    result = workflow._quality_retry(run, reference_path, {"generation": {}}, source, revision, {}, findings, "quality", require_promoted_runtime=False)
    assert result["status"] == "blocked" and result["stage"] == "layout_repair_classification"
    assert not (revision / "attempts").exists()


def test_unknown_table_owner_requires_reviewer_localization(tmp_path):
    run, revision, reference_path, source, protocol, findings = review_case(tmp_path)
    request_path = next(p for p in (revision / "hermes/verification-requests").glob("*.json")
                        if json.loads(p.read_text())["task"] == "clinical_content_verification")
    request = json.loads(request_path.read_text())
    response_path = revision / request["response_path"]
    response = json.loads(response_path.read_text())
    response["findings"][0]["table_ids"] = ["guessed-table"]
    response_path.write_text(json.dumps(response))
    findings, _ = quality.validate_verifications(revision)
    table_finding = next(f for f in findings if f.get("surface") == "table")
    assert table_finding["recovery_class"] == "verifier_transient"
    assert table_finding["target_ids"] == ["verification:content"]
