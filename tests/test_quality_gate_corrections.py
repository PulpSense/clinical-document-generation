"""Independent, read-only quality regressions; fixtures contain no clinical data."""
import json
from pathlib import Path

import pytest
from docx import Document

import quality
from docx.enum.style import WD_STYLE_TYPE
from pypdf import PdfWriter
from pypdf.generic import DictionaryObject, NameObject, DecodedStreamObject


def write_pdf(path, pages):
    """Real PDF text/geometry, without an office installation or fake reader."""
    writer = PdfWriter()
    font = DictionaryObject({NameObject("/Type"): NameObject("/Font"),
                             NameObject("/Subtype"): NameObject("/Type1"),
                             NameObject("/BaseFont"): NameObject("/Helvetica")})
    for lines in pages:
        page = writer.add_blank_page(width=612, height=792)
        page[NameObject("/Resources")] = DictionaryObject({NameObject("/Font"): DictionaryObject({NameObject("/F1"): writer._add_object(font)})})
        stream = DecodedStreamObject()
        commands = []
        for text, x, y in lines:
            escaped = text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
            commands.append(f"BT /F1 12 Tf 1 0 0 1 {x} {y} Tm ({escaped}) Tj ET")
        stream.set_data("\n".join(commands).encode("ascii"))
        page[NameObject("/Contents")] = writer._add_object(stream)
    writer.write(path)


def toc_docx(path, cached_page=2):
    document = Document()
    document.styles.add_style("toc 1", WD_STYLE_TYPE.PARAGRAPH)
    document.add_heading("TABLE OF CONTENTS", level=1)
    document.add_paragraph("TABLE OF CONTENTS\t1", style="toc 1")
    document.add_paragraph(f"1. PURPOSE\t{cached_page}", style="toc 1")
    document.add_heading("1. PURPOSE", level=1)
    document.add_paragraph("Study purpose body.")
    document.save(path)


def test_final_toc_audit_uses_final_pdf_not_renderer_mapping(tmp_path):
    docx, pdf = tmp_path / "protocol.docx", tmp_path / "protocol.pdf"
    toc_docx(docx)
    write_pdf(pdf, [[("TABLE OF CONTENTS", 72, 700), ("1. PURPOSE 2", 72, 670)],
                    [("1. PURPOSE", 72, 700), ("Study purpose body.", 72, 675)]])
    original = docx.read_bytes(), pdf.read_bytes()
    assert quality.audit_final_toc_destinations(docx, pdf)["status"] == "passed"
    assert (docx.read_bytes(), pdf.read_bytes()) == original
    write_pdf(pdf, [[("TABLE OF CONTENTS", 72, 700), ("1. PURPOSE 2", 72, 670)],
                    [("Preface", 72, 700)],
                    [("1. PURPOSE", 72, 700), ("Study purpose body.", 72, 675)]])
    report = quality.audit_final_toc_destinations(docx, pdf)
    assert report["status"] == "mismatch"
    row = next(row for row in report["destinations"] if row["heading"] == "1. PURPOSE")
    assert row["cached_page"] == 2 and row["rendered_page"] == 3


def test_final_toc_audit_unknown_is_not_a_fabricated_mismatch(tmp_path):
    docx, pdf = tmp_path / "protocol.docx", tmp_path / "protocol.pdf"
    toc_docx(docx)
    write_pdf(pdf, [[]])
    report = quality.audit_final_toc_destinations(docx, pdf)
    assert report["status"] == "unknown"
    assert not report["findings"]


def test_render_report_checks_after_last_refresh_export(tmp_path, monkeypatch):
    candidate = tmp_path / "candidate"
    candidate.mkdir()
    toc_docx(candidate / "protocol.docx")
    calls = []

    def export(docx, output_dir, identity, **kwargs):
        calls.append(identity)
        path = output_dir / "protocol.pdf"
        write_pdf(path, [[("TABLE OF CONTENTS", 72, 700), ("1. PURPOSE 2", 72, 675)],
                         [("Preface", 72, 700)], [("1. PURPOSE", 72, 700)]])
        return path

    def raster(pdf, output_dir, identity, **kwargs):
        output_dir.mkdir(parents=True)
        paths = [output_dir / f"page-{i}.png" for i in range(1, 4)]
        for path in paths:
            path.write_bytes(b"test-only page exporter")
        return paths

    # Simulate a non-converging updater; the independent auditor is real.
    monkeypatch.setattr(quality, "refresh_toc_from_pdf", lambda *_: True)
    monkeypatch.setattr(quality, "page_renderers", lambda **_: [{"kind": "pypdfium2",
        "path": "python:pypdfium2", "module": "pypdfium2", "python_path": "/release/runtime/python",
        "source": "release-owned runtime"}])
    report = quality.render_pages(tmp_path, office_exporter=export, page_exporter=raster,
        renderer_identities=[{"kind": "LibreOffice", "path": "/test/office"}],
        page_renderer_identities=[{"kind": "pypdfium2", "path": "python:pypdfium2",
            "module": "pypdfium2", "python_path": "/release/runtime/python", "source": "release-owned runtime"}])
    assert len(calls) == 4
    assert report["status"] == "blocked"
    assert report["artifacts"][0]["toc_destinations"]["status"] == "mismatch"
    assert report["artifacts"][0]["sparse_body_context"]
    assert all(not row["blocking"] for row in report["artifacts"][0]["sparse_body_context"])
    assert any(finding.get("check") == "toc_mismatch" for finding in report["findings"])


def test_blank_page_detection_ignores_furniture_not_short_body(tmp_path):
    pdf = tmp_path / "body.pdf"
    write_pdf(pdf, [[("A long running study title", 72, 760), ("Page 1 of 3", 72, 30)],
                    [("A long running study title", 72, 760), ("Sign here", 72, 600)],
                    [("A long running study title", 72, 760), ("Yes No", 72, 600)]])
    document = Document()
    docx = tmp_path / "body.docx"
    document.save(docx)
    assert quality._blank_pdf_pages(pdf, docx_path=docx) == [1]


def test_body_empty_requires_known_geometry_not_assumed_margins(tmp_path):
    from docx.shared import Inches
    pdf, docx = tmp_path / "small.pdf", tmp_path / "small.docx"
    write_pdf(pdf, [[("Sign here", 72, 760)]])
    document = Document()
    document.sections[0].top_margin = Inches(0.25)
    document.save(docx)
    assert quality._blank_pdf_pages(pdf, docx_path=docx) == []
    assert quality._blank_pdf_pages(pdf) == []  # unknown margins, not blank proof


def test_header_logo_does_not_hide_empty_body(tmp_path):
    from pypdf.generic import NumberObject
    pdf, docx = tmp_path / "logo.pdf", tmp_path / "logo.docx"
    Document().save(docx)
    writer = PdfWriter()
    image = DecodedStreamObject()
    image.set_data(b"\x00\x00\x00")
    image.update({NameObject("/Type"): NameObject("/XObject"), NameObject("/Subtype"): NameObject("/Image"),
        NameObject("/Width"): NumberObject(1), NameObject("/Height"): NumberObject(1),
        NameObject("/ColorSpace"): NameObject("/DeviceRGB"), NameObject("/BitsPerComponent"): NumberObject(8)})
    for y in (750, 400):
        page = writer.add_blank_page(width=612, height=792)
        page[NameObject("/Resources")] = DictionaryObject({NameObject("/XObject"): DictionaryObject({NameObject("/Logo"): writer._add_object(image)})})
        stream = DecodedStreamObject()
        stream.set_data(f"q 20 0 0 20 72 {y} cm /Logo Do Q".encode())
        page[NameObject("/Contents")] = writer._add_object(stream)
    writer.write(pdf)
    assert quality._blank_pdf_pages(pdf, docx_path=docx) == [1]


def test_sparse_geometry_flags_forced_break_prose_without_blocking_consent(tmp_path):
    docx, pdf = tmp_path / "protocol.docx", tmp_path / "protocol.pdf"
    sentence = "Confidentiality protections remain described in Section 16."
    document = Document()
    document.add_paragraph(sentence)
    document.add_page_break()
    document.add_heading("15. ASSESSMENTS", level=1)
    document.save(docx)
    write_pdf(pdf, [[(sentence, 72, 710)], [("15. ASSESSMENTS", 72, 710)]])
    before = docx.read_bytes(), pdf.read_bytes()
    rows = quality._sparse_pdf_pages(docx, pdf)
    assert rows[0]["page"] == 1 and rows[0]["forced_break_after"] is True
    assert all(row["blocking"] is False for row in rows)
    assert (docx.read_bytes(), pdf.read_bytes()) == before
    assert quality._blank_pdf_pages(pdf) == []


def test_visual_request_retains_unknown_toc_and_sparse_context(tmp_path):
    artifact = {"artifact": "protocol", "toc_destinations": {"status": "unknown", "destinations": [
        {"heading": "1. PURPOSE", "status": "unknown"}]},
        "sparse_body_context": [{"page": 2, "blocking": False, "forced_break_after": True}],
        "pages": [{"page": 1, "sha256": "test-bound-hash", "path": "rendered/page-1.png"}]}
    paths = quality.create_verification_requests(tmp_path, {"meta": {"study_type": "Retrospective"}}, {"artifacts": [artifact]})
    visual = json.loads(paths[1].read_text())
    assert visual["artifacts"][0] == artifact
    assert "unknown toc" in visual["instructions"].casefold()
    assert quality.verification_request_hash_valid(visual)


CASES = [("Retrospective", "Advarra"), ("Prospective", "Advarra"),
         ("Prospective", "Sterling"), ("Ambispective", "Advarra"),
         ("Ambispective", "Sterling")]


@pytest.mark.parametrize("branch,family", CASES)
def test_direct_source_surface_audit_catches_front_matter_omissions(tmp_path, branch, family):
    source = {"meta": {"study_type": branch, "icf_template": family},
              "parties": {"sponsor": {"name": "Example Foundation"}},
              "sites": [{"facility": {"address": "101 Example Road", "city": "Example City",
                                       "state": "EX", "country": "Example Country"}}]}
    document = Document()
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Sponsor"
    table.cell(0, 1).text = "Example Foundation\nExample CountryExample Foundation"
    protocol = tmp_path / "protocol.docx"
    document.save(protocol)
    before = protocol.read_bytes()
    findings = quality.audit_source_surfaces(protocol, source)
    assert any(item["field"] == "parties.sponsor.name" for item in findings)
    assert protocol.read_bytes() == before
    table.cell(0, 1).text = "Example Foundation\nExample Country"
    document.save(protocol)
    assert quality.audit_source_surfaces(protocol, source) == []
    if branch != "Retrospective":
        icf = tmp_path / "icf.docx"
        document = Document()
        document.add_paragraph("Study site: 101 Example Road")
        document.save(icf)
        findings = quality.audit_source_surfaces(icf, source)
        assert {item["field"] for item in findings} == {
            "sites[0].facility.city", "sites[0].facility.state", "sites[0].facility.country"}
        document.add_paragraph("Example City, EX, Example Country")
        document.save(icf)
        # An unrelated later mention must not hide the front-matter omission.
        assert quality.audit_source_surfaces(icf, source)
        document.paragraphs[0].text = "Study site: 101 Example Road, Example City, EX, Example Country"
        document.save(icf)
        assert quality.audit_source_surfaces(icf, source) == []


@pytest.mark.parametrize("branch,family", CASES)
def test_review_requires_source_precedence_and_independent_field_inventory(tmp_path, branch, family):
    source = {"meta": {"study_type": branch, "icf_template": family},
              "procedures": {"assessments": "Download and pain rating", "completion": "Visit and return"},
              "safety": {"monitoring": "AE review at each contact"},
              "sites": [{"facility": {"address": "101 Example Road", "city": "Example City"}}]}
    paths = quality.create_verification_requests(tmp_path, source, {"artifacts": []})
    content, visual = [json.loads(path.read_text()) for path in paths]
    instruction = content["instructions"].casefold()
    assert "source takes precedence" in instruction
    assert "authority" in instruction and "discontinuation" in instruction
    assert "treat exact authorized fixed clinical boilerplate as approved" not in instruction
    fields = {entry["source_path"]: entry["value"] for entry in content["source_field_inventory"]}
    assert fields["procedures.assessments"] == "Download and pain rating"
    assert fields["safety.monitoring"] == "AE review at each contact"
    assert fields["sites[0].facility.city"] == "Example City"
    assert "each supplied source field" in instruction
    assert "protocol_table_contracts" in instruction
    assert "whole-word" in visual["instructions"]
    assert "originating approval" in visual["instructions"]
    assert "sentence continuing" in visual["instructions"]
    assert quality.verification_request_hash_valid(content)
    assert quality.verification_request_hash_valid(visual)
