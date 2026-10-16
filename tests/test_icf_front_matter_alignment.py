import json
from pathlib import Path

import pytest
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH

import rendering

ROOT = Path(__file__).resolve().parents[1]


def test_front_matter_labels_and_values_are_explicitly_left_aligned():
    document = Document()
    document.styles["Normal"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    table = document.add_table(rows=2, cols=2)
    for row, values in zip(table.rows, [("Protocol Number:", "ABC-123"), ("Principal Investigator:", "Investigator Name")]):
        for cell, text in zip(row.cells, values):
            cell.text = text
            cell.add_paragraph()
    rendering._normalize_icf_front_matter(document, {})
    for row in table.rows:
        for cell in row.cells:
            assert len(cell.paragraphs) == 1
            assert cell.paragraphs[0].alignment == WD_ALIGN_PARAGRAPH.LEFT
    assert table.cell(0, 1).text == "ABC-123"


@pytest.mark.parametrize("branch", ["prospective", "ambispective"])
def test_advarra_output_uses_left_aligned_labels_and_values(branch, tmp_path):
    reference = json.loads((ROOT / f"tests/fixtures/{branch}-acceptance-source.json").read_text())
    reference["meta"]["icf_template"] = "Advarra"
    rendering.render_documents(ROOT, tmp_path, reference, {"protocol": [], "icf": {}, "prs": {}},
                               artifact_names={"icf"}, audit_output=False)
    table = Document(tmp_path / "candidate/icf.docx").tables[0]
    for row in table.rows:
        for cell in row.cells:
            assert all(paragraph.alignment == WD_ALIGN_PARAGRAPH.LEFT for paragraph in cell.paragraphs)
            assert len(cell.paragraphs) == 1 or cell.paragraphs[-1].text.strip()


def test_nonmetadata_table_alignment_and_empty_cells_are_preserved():
    document = Document()
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Visit"
    table.cell(0, 1).text = "Window"
    for row in table.rows:
        for cell in row.cells:
            cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    before = table._tbl.xml
    rendering._normalize_icf_front_matter(document, {})
    assert table._tbl.xml == before


@pytest.mark.parametrize("family", ["Advarra", "Sterling"])
def test_supplied_conditional_facts_reach_the_appropriate_drafting_sections(tmp_path, family):
    import drafting
    from contracts import batch_plan
    reference = json.loads((ROOT / "tests/fixtures/prospective-acceptance-source.json").read_text())
    reference["meta"]["icf_template"] = family
    reference["procedures"]["return_of_results"] = "Results are available on request after sponsor release."
    reference["procedures"]["pregnancy_testing"] = "A pregnancy test is performed at screening when applicable."
    reference["risks_benefits"]["pregnancy"] = "Use the source-specified reproductive precautions."
    reference["confidentiality"] = {"authorization_duration": "Authorization ends at study closeout."}
    batch = next(b for b in batch_plan("Prospective", family) if b.artifact == "icf")
    path = drafting.create_drafting_request(repo_root=ROOT, revision_dir=tmp_path,
        revision_id="conditional-source", reference=reference, batch=batch,
        attempts={target: 1 for target in batch.section_ids}, wave="initial")
    request = json.loads(path.read_text())
    targets = {section["section_id"]: section for section in request["section_contracts"]}
    assert "procedures.return_of_results" in targets["icf.procedures"]["minimum_evidence"]
    assert "procedures.pregnancy_testing" in targets["icf.procedures"]["minimum_evidence"]
    assert "risks_benefits.pregnancy" in targets["icf.risks"]["minimum_evidence"]
    assert "confidentiality.authorization_duration" in targets["icf.privacy"]["minimum_evidence"]
    assert {"procedures.return_of_results", "procedures.pregnancy_testing", "risks_benefits.pregnancy"} <= {
        item["path"] for item in request["approved_input"]}
