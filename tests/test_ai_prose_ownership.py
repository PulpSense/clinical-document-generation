"""Accepted AI prose survives public native assembly without semantic patching."""
import json
from pathlib import Path

import pytest
import rendering

ROOT = Path(__file__).resolve().parents[1]


def source(family='Advarra', branch='Prospective'):
    value = json.loads((ROOT / f'tests/fixtures/{branch.lower()}-acceptance-source.json').read_text())
    value['meta']['icf_template'] = family
    return value


def draft(text):
    return {'paragraphs': [{'text': text}], 'lists': []}


def body(path, section_id):
    return [item['paragraph'] for item in rendering.rendered_section_snapshot(path, section_id)
            if 'paragraph' in item]


@pytest.mark.parametrize('family', ['Advarra', 'Sterling'])
def test_enrollment_paraphrase_is_not_completed_with_an_extra_sentence(tmp_path, family):
    approved = source(family)
    approved['population']['sample_size'] = '40 participants'
    authored = 'The study will include 40 adults. Your participation lasts three months.'
    result = rendering.render_documents(ROOT, tmp_path, approved,
        {'protocol': [], 'icf': {'icf.duration': draft(authored)}, 'prs': {}},
        artifact_names={'icf'}, audit_output=False)
    assert result['status'] == 'passed'
    assert [text for text in body(tmp_path / 'candidate/icf.docx', 'icf.duration') if text.strip()] == [authored]


@pytest.mark.parametrize('family', ['Advarra', 'Sterling'])
def test_full_withdrawal_paragraph_is_inserted_as_authored(tmp_path, family):
    from docx import Document
    approved = source(family)
    fixed = json.loads((ROOT / 'references/fixed-clinical-boilerplate.json').read_text())['sections']['icf-withdrawal']
    authored = fixed + ' Tell your enrolling site coordinator if you leave; earned payments remain due.'
    rendering.render_documents(ROOT, tmp_path, approved,
        {'protocol': [], 'icf': {'icf.withdrawal': draft(authored)}, 'prs': {}},
        artifact_names={'icf'}, audit_output=False)
    texts = [p.text for p in Document(tmp_path / 'candidate/icf.docx').paragraphs]
    assert authored in texts
    assert '\n'.join(texts).count(fixed) == 1


def test_protocol_matrix_introduction_is_authored_while_table_stays_source_owned(tmp_path):
    from docx import Document
    approved = source()
    authored = 'The table lists scheduled activities; the follow-up questionnaire is completed before dilation.'
    rendering.render_documents(ROOT, tmp_path, approved,
        {'protocol': [dict(section_id='evaluation-procedures', **draft(authored))], 'icf': {}, 'prs': {}},
        artifact_names={'protocol'}, audit_output=False)
    path = tmp_path / 'candidate/protocol.docx'
    assert authored in [p.text for p in Document(path).paragraphs]
    assert any('table' in item for item in rendering.protocol_section_snapshot(path, 'evaluation-procedures'))


def test_retrospective_enrollment_paraphrase_gets_no_literal_source_append(tmp_path):
    from docx import Document
    approved = source(branch='Retrospective')
    approved['procedures']['discontinuation'] = 'Stop new abstraction when permission is withdrawn; retain already abstracted data.'
    authored = 'After permission withdrawal, collect no further chart data. Previously abstracted information is retained.'
    rendering.render_documents(ROOT, tmp_path, approved,
        {'protocol': [dict(section_id='study-procedure.enrollment', **draft(authored))], 'icf': {}, 'prs': {}},
        artifact_names={'protocol'}, audit_output=False)
    path = tmp_path / 'candidate/protocol.docx'
    paragraphs = [p.text for p in Document(path).paragraphs]
    assert authored in paragraphs
    assert approved['procedures']['discontinuation'] not in paragraphs


@pytest.mark.parametrize('family', ['Advarra', 'Sterling'])
def test_authored_icf_text_is_not_filtered_or_rewritten_by_template_cleanup(tmp_path, family):
    from docx import Document
    approved = source(family)
    approved['regulatory']['prs']['study_type'] = 'Observational'
    procedures = 'You must tell the study doctor or study staff about all side effects at the scheduled call.'
    privacy = 'Your clinical trial records will be kept for five years. Prior reports (references 1 and 2) describe the context.'
    approved['study']['references'] = ['First reference.', 'Second reference.']
    rendering.render_documents(ROOT, tmp_path, approved,
        {'protocol': [], 'icf': {'icf.procedures': draft(procedures), 'icf.privacy': draft(privacy)}, 'prs': {}},
        artifact_names={'icf'}, audit_output=False)
    texts = [p.text for p in Document(tmp_path / 'candidate/icf.docx').paragraphs]
    assert procedures in texts
    assert privacy in texts


def test_sterling_summary_keeps_all_authored_blocks(tmp_path):
    from docx import Document
    texts = ['Study purpose summary.', 'Participation summary.', 'Risk summary.',
             'Benefit summary.', 'Alternatives summary.', 'Additional approved clarification.']
    rendering.render_documents(ROOT, tmp_path, source('Sterling'),
        {'protocol': [], 'icf': {'icf.key-information-summary': {'paragraphs': [{'text': text} for text in texts], 'lists': []}}, 'prs': {}},
        artifact_names={'icf'}, audit_output=False)
    visible = [p.text for p in Document(tmp_path / 'candidate/icf.docx').paragraphs]
    assert all(visible.count(text) == 1 for text in texts)
    assert [text for text in visible if text in texts] == texts


def test_ai_target_fields_do_not_fall_back_to_source_prose():
    approved = source()
    approved['risks_benefits']['risks'] = 'Source risk statement.'
    approved['population']['sample_justification'] = 'Source justification.'
    fields = rendering.render_fields(approved, {'protocol': [], 'icf': {}, 'prs': {}})
    assert fields['AI_risks'] == ''
    assert fields['sampleSizeJustification'] == ''
    assert fields['AI_inclusionCriteria'] == ''
    assert fields['sampleSize']  # Direct source-field mapping still belongs to code.


@pytest.mark.parametrize('family', ['Advarra', 'Sterling'])
def test_drafting_handoff_requests_complete_sections_and_authorized_fixed_text(tmp_path, family):
    import copy
    import drafting
    from contracts import batch_plan
    approved = source(family)
    before = copy.deepcopy(approved)
    (tmp_path / 'generation-only.json').write_text('{}')
    batch = next(b for b in batch_plan('Prospective', family) if b.batch_id == 'icf-narrative')
    targets = ['icf.duration', 'icf.withdrawal']
    path = drafting.create_drafting_request(repo_root=ROOT, revision_dir=tmp_path,
        revision_id='ai-prose-trial', reference=approved, batch=batch,
        target_ids=targets, attempts={t: 1 for t in targets}, wave='generation-only')
    request = json.loads(path.read_text())
    constraints = '\n'.join(request['constraints'])
    assert 'Assembly inserts accepted text as authored' in constraints
    corrections = request['generation_only_prompts']['editorial_corrections']
    assert 'Draft the complete withdrawal section' in ' '.join(corrections['icf.withdrawal'])
    assert 'expected enrollment once' in ' '.join(corrections['icf.duration'])
    withdrawal = next(c for c in request['section_contracts'] if c['section_id'] == 'icf.withdrawal')
    duration = next(c for c in request['section_contracts'] if c['section_id'] == 'icf.duration')
    assert 'expected enrollment once' in ' '.join(duration['content_expectations'])
    assert withdrawal['fixed_boilerplate']
    assert approved == before


def test_authored_child_section_survives_parent_body_insertion(tmp_path):
    from docx import Document
    model = {'protocol': [
        dict(section_id='ethics', **draft('The named review board has not yet reviewed this study.')),
        dict(section_id='ethics.confidentiality', **draft('Confidentiality protections are described in Section 16.'))],
        'icf': {}, 'prs': {}}
    rendering.render_documents(ROOT, tmp_path, source(), model, artifact_names={'protocol'}, audit_output=False)
    texts = [p.text for p in Document(tmp_path / 'candidate/protocol.docx').paragraphs]
    assert 'The named review board has not yet reviewed this study.' in texts
    assert 'Confidentiality protections are described in Section 16.' in texts
