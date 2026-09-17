import copy
import hashlib
import json
import zipfile
from pathlib import Path
from types import SimpleNamespace

import pytest
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from contracts import LAYOUT_FAMILY_ARTIFACTS, VISUAL_CHECK_DISPOSITIONS, contracted_template_bundle
from quality import RESPONSE_SCHEMA, VISUAL_CHECKS, deterministic_content_check, validate_verifications, verification_request_sha256
import quality
from rendering import render_documents, render_fields
import rendering
from prs_xml import generate as generate_xml
import workflow


ROOT = Path(__file__).resolve().parents[1]


def _source():
    return json.loads(
        (
            ROOT
            / "tests/fixtures/release-certification/prospective-advarra/approved-reference.json"
        ).read_text(encoding="utf-8")
    )


def _minimal_source_bound_icf_procedures():
    return {
        "paragraphs": [{"text": "You will complete the approved study visits and procedures."}],
        "lists": [],
    }


def _visible(document):
    paragraphs = [paragraph.text for paragraph in document.paragraphs]
    cells = [cell.text for table in document.tables for row in table.rows for cell in row.cells]
    return "\n".join(paragraphs + cells)


def _page_identity():
    return {
        "kind": "pypdfium2",
        "path": "python:pypdfium2",
        "module": "pypdfium2",
        "python_path": "/release/runtime/python",
        "source": "release-owned runtime",
    }


def test_renderer_preflight_reports_missing_host_office_when_no_renderer_exists(tmp_path, monkeypatch):
    reference = _source()
    monkeypatch.setattr(quality, "renderers", lambda **_: [])
    report = quality.preflight(ROOT, reference, deadline_seconds=2)

    assert report["status"] == "blocked"
    assert report["smoke"]["status"] == "blocked"
    assert any("renderer" == finding["field"] for finding in report["findings"])


def test_renderer_preflight_records_renderer_font_and_smoke_evidence(tmp_path, monkeypatch):
    reference = _source()
    identity = {"kind": "LibreOffice", "path": "/usr/bin/soffice", "version": "test", "platform": "Linux"}
    page_identity = _page_identity()
    monkeypatch.setattr(quality, "renderers", lambda **_: [identity])
    monkeypatch.setattr(quality, "page_renderers", lambda **_: [page_identity])
    monkeypatch.setattr(quality, "_font_probe", lambda *_args, **_kwargs: (True, "test-font"))

    def fake_export(_docx, output_dir, _identity, **_kwargs):
        from pypdf import PdfWriter
        output = output_dir / "preflight.pdf"
        writer = PdfWriter()
        writer.add_blank_page(width=612, height=792)
        with output.open("wb") as handle:
            writer.write(handle)
        return output

    monkeypatch.setattr(quality, "_render_pdf", fake_export)

    def fake_rasterize(_pdf, output_dir, _identity, **_kwargs):
        output_dir.mkdir(parents=True, exist_ok=True)
        page = output_dir / "page.png"
        page.write_bytes(b"png")
        return [page]

    monkeypatch.setattr(quality, "rasterize_pdf", fake_rasterize)
    report = quality.preflight(ROOT, reference, deadline_seconds=10)

    assert report["status"] == "passed"
    assert report["renderer"] == identity
    assert report["elapsed_seconds"] < report["deadline_seconds"]
    assert report["smoke"] == {"status": "passed", "pdf": "disposable/preflight.pdf", "page_image": "disposable/preflight.png", "pages": 1}


def test_page_renderer_has_one_governed_backend_and_ignores_host_tools(tmp_path):
    assert quality.PAGE_RENDERER_BACKENDS == ("pypdfium2",)
    assert quality.page_renderer(
        environment={"PATH": "/usr/bin:/opt/homebrew/bin"},
        home=tmp_path,
        skill_root=tmp_path,
    ) is None


def test_pdfium_identity_filter_rejects_ambient_interpreter_identity(tmp_path):
    runtime_python = tmp_path / "runtime/python"
    governed = {
        "kind": "pypdfium2",
        "path": "python:pypdfium2",
        "module": "pypdfium2",
        "python_path": str(runtime_python),
        "source": "release-owned runtime",
    }

    assert quality._one_pdfium_renderer([
        {"kind": "pypdfium2", "path": "python:pypdfium2"},
        governed,
    ]) == [governed]


def test_rasterize_pdf_rejects_ambient_pdfium_before_import(tmp_path):
    with pytest.raises(RuntimeError, match="manifest-verified release-owned runtime"):
        quality.rasterize_pdf(
            tmp_path / "unused.pdf",
            tmp_path / "pages",
            {"kind": "pypdfium2", "path": "python:pypdfium2"},
        )


def test_rasterize_pdf_verifies_against_the_current_release_root(
    tmp_path,
    monkeypatch,
    governed_pdfium,
):
    observed_roots = []

    def observe_release_root(_identities, *, skill_root, **_kwargs):
        observed_roots.append(Path(skill_root).resolve())
        return []

    monkeypatch.setattr(quality, "_one_pdfium_renderer", observe_release_root)

    with pytest.raises(RuntimeError, match="manifest-verified release-owned runtime"):
        quality.rasterize_pdf(
            tmp_path / "unused.pdf",
            tmp_path / "pages",
            governed_pdfium,
        )

    assert observed_roots == [ROOT.resolve()]


def test_renderer_discovery_ignores_a_release_local_office_suite(tmp_path, monkeypatch):
    bundled = tmp_path / "runtime/LibreOffice/program/soffice"
    bundled.parent.mkdir(parents=True)
    bundled.write_text("#!/bin/sh\nprintf 'LibreOffice test\\n'\n", encoding="utf-8")
    bundled.chmod(0o755)
    monkeypatch.setattr(quality.platform, "system", lambda: "Linux")
    monkeypatch.setattr(quality, "_executable_candidates", lambda *_args, **_kwargs: iter(()))

    assert quality.renderers(skill_root=tmp_path) == []


def test_renderer_preflight_uses_the_selected_page_renderer(tmp_path, monkeypatch):
    reference = _source()
    renderer_identity = {"kind": "LibreOffice", "path": "/usr/bin/soffice", "version": "test", "platform": "Linux"}
    page_renderer_identity = _page_identity()
    monkeypatch.setattr(quality, "renderers", lambda **_: [renderer_identity])
    monkeypatch.setattr(quality, "page_renderers", lambda **_: [page_renderer_identity])
    monkeypatch.setattr(quality, "_font_probe", lambda *_args, **_kwargs: (True, "test-font"))

    def fake_export(_docx, output_dir, _identity, **_kwargs):
        from pypdf import PdfWriter
        output = output_dir / "preflight.pdf"
        writer = PdfWriter()
        writer.add_blank_page(width=612, height=792)
        with output.open("wb") as handle:
            writer.write(handle)
        return output

    def fake_rasterize(_pdf, output_dir, identity, **_kwargs):
        output_dir.mkdir(parents=True, exist_ok=True)
        page = output_dir / "page.png"
        page.write_bytes(b"png")
        assert identity == page_renderer_identity
        return [page]

    monkeypatch.setattr(quality, "_render_pdf", fake_export)
    monkeypatch.setattr(quality, "rasterize_pdf", fake_rasterize)

    report = quality.preflight(ROOT, reference, deadline_seconds=10)

    assert report["status"] == "passed"
    assert report["page_renderer"] == page_renderer_identity
    assert report["smoke"]["status"] == "passed"


def test_renderer_preflight_stops_when_the_release_owned_pdfium_fails(monkeypatch):
    renderer_identity = {"kind": "LibreOffice", "path": "/usr/bin/soffice", "version": "test", "platform": "Linux"}
    word_identity = {"kind": "Microsoft Word", "path": "/Applications/Microsoft Word.app", "platform": "Darwin"}
    broken = _page_identity()
    unapproved = {"kind": "pdftoppm", "path": "/usr/bin/pdftoppm", "source": "PATH"}
    attempts = []
    monkeypatch.setattr(quality, "renderers", lambda **_: [renderer_identity, word_identity])
    monkeypatch.setattr(quality, "page_renderers", lambda **_: [broken, unapproved])
    monkeypatch.setattr(quality, "_font_probe", lambda *_args, **_kwargs: (True, "test-font"))

    def fake_export(_docx, output_dir, _identity, **_kwargs):
        from pypdf import PdfWriter
        output = output_dir / "preflight.pdf"
        writer = PdfWriter()
        writer.add_blank_page(width=612, height=792)
        with output.open("wb") as handle:
            writer.write(handle)
        return output

    def fake_rasterize(_pdf, output_dir, identity, **_kwargs):
        attempts.append(identity["kind"])
        raise RuntimeError("controlled PDFium failure")

    monkeypatch.setattr(quality, "_render_pdf", fake_export)
    monkeypatch.setattr(quality, "rasterize_pdf", fake_rasterize)

    report = quality.preflight(ROOT, _source(), deadline_seconds=10)

    assert report["status"] == "blocked"
    assert attempts == ["pypdfium2"]
    assert report["page_renderer"] is None
    assert report["page_renderer_attempts"] == [
        {"renderer": broken, "status": "failed", "issue": "controlled PDFium failure"},
    ]
    assert len(report["renderer_attempts"]) == 1


def test_renderer_preflight_uses_a_packaged_font_when_host_fonts_are_missing(monkeypatch):
    identity = {"kind": "LibreOffice", "path": "/usr/bin/soffice", "version": "test", "platform": "Linux"}
    page_identity = _page_identity()
    monkeypatch.setattr(quality, "renderers", lambda **_: [identity])
    monkeypatch.setattr(quality, "page_renderers", lambda **_: [page_identity])
    monkeypatch.setattr(quality, "_template_fonts", lambda _path: {"Missing Client Font"})
    monkeypatch.setattr(quality, "_font_probe", lambda *_args, **_kwargs: (False, "not installed"))

    def fake_export(_docx, output_dir, _identity, **_kwargs):
        from pypdf import PdfWriter
        output = output_dir / "preflight.pdf"
        writer = PdfWriter()
        writer.add_blank_page(width=612, height=792)
        with output.open("wb") as handle:
            writer.write(handle)
        return output

    def fake_rasterize(_pdf, output_dir, _identity, **_kwargs):
        output_dir.mkdir(parents=True, exist_ok=True)
        page = output_dir / "page.png"
        page.write_bytes(b"png")
        return [page]

    monkeypatch.setattr(quality, "_render_pdf", fake_export)
    monkeypatch.setattr(quality, "rasterize_pdf", fake_rasterize)

    report = quality.preflight(ROOT, _source(), deadline_seconds=10)

    assert report["status"] == "passed"
    assert report["smoke"]["status"] == "passed"
    assert report["required_fonts"] == ["Missing Client Font"]
    assert report["font_substitutions"] == {"Missing Client Font": "Liberation Sans"}


def test_every_current_template_font_has_an_explicit_packaged_fallback():
    required = {
        font
        for template in (ROOT / "assets/client-templates/reference").glob("*.docx")
        for font in quality._template_fonts(template)
    }

    assert required
    assert not {
        font for font in required
        if font.casefold() not in quality.APPROVED_PACKAGED_FONT_FALLBACKS
    }


def test_renderer_preflight_uses_only_the_approved_packaged_font_for_a_missing_client_font(monkeypatch):
    identity = {"kind": "LibreOffice", "path": "/verified/soffice", "version": "test", "platform": "Darwin", "source": "host prerequisite"}
    page_identity = _page_identity()
    probes = {
        "Noto Sans Symbols": (False, "not installed"),
        "Apple Symbols": (True, "macOS system font inventory"),
    }
    monkeypatch.setattr(quality, "renderers", lambda **_: [identity])
    monkeypatch.setattr(quality, "page_renderers", lambda **_: [page_identity])
    monkeypatch.setattr(quality, "_template_fonts", lambda _path: {"Noto Sans Symbols"})
    monkeypatch.setattr(quality, "_font_probe", lambda font, **_kwargs: probes.get(font, (False, "not installed")))

    def fake_export(_docx, output_dir, _identity, **_kwargs):
        from pypdf import PdfWriter
        output = output_dir / "preflight.pdf"
        writer = PdfWriter()
        writer.add_blank_page(width=612, height=792)
        with output.open("wb") as handle:
            writer.write(handle)
        return output

    monkeypatch.setattr(quality, "_render_pdf", fake_export)

    def fake_rasterize(_pdf, output_dir, _identity, **_kwargs):
        output_dir.mkdir(parents=True, exist_ok=True)
        page = output_dir / "page.png"
        page.write_bytes(b"png")
        return [page]

    monkeypatch.setattr(quality, "rasterize_pdf", fake_rasterize)

    report = quality.preflight(ROOT, _source(), deadline_seconds=10)

    assert report["status"] == "passed"
    assert report["required_fonts"] == ["Noto Sans Symbols"]
    assert report["font_substitutions"] == {"Noto Sans Symbols": "Liberation Sans"}
    assert report["fonts"]["Noto Sans Symbols"] == {
        "available": False,
        "match": "not installed",
        "substitute": "Liberation Sans",
        "substitute_match": "bundled approved compatible font: assets/fallback-fonts/LiberationSans-Regular.ttf",
    }
    assert report["smoke"]["status"] == "passed"


def test_macos_renderer_honors_the_callers_remaining_deadline(tmp_path, monkeypatch):
    source = tmp_path / "source.docx"
    source.write_bytes(b"docx")
    observed = []

    def fake_run(_command, **kwargs):
        observed.append(kwargs["timeout"])
        (tmp_path / "source.pdf").write_bytes(b"pdf")
        return SimpleNamespace(returncode=0, stderr="", stdout="")

    monkeypatch.setattr(quality.subprocess, "run", fake_run)

    quality._render_pdf(
        source,
        tmp_path,
        {"kind": "Microsoft Word", "path": "/Applications/Microsoft Word.app", "platform": "Darwin"},
        timeout_seconds=7.5,
    )

    assert observed == [7.5]


def test_layout_repair_is_scoped_to_one_artifact_and_rule(tmp_path):
    reference = json.loads((ROOT / "tests/fixtures/prospective-acceptance-source.json").read_text(encoding="utf-8"))
    model = {
        "protocol": [],
        "icf": {"icf.procedures": _minimal_source_bound_icf_procedures()},
        "prs": {},
    }

    baseline = render_documents(ROOT, tmp_path, reference, model)
    icf_before = (tmp_path / "candidate/icf.docx").read_bytes()
    baseline_protocol = Document(tmp_path / "candidate/protocol.docx")
    unaffected_paragraphs_before = tuple(
        paragraph._p.xml
        for paragraph in baseline_protocol.paragraphs
        if paragraph.text.strip() != "5. INTRODUCTION"
    )
    unaffected_tables_before = tuple(table._tbl.xml for table in baseline_protocol.tables)
    geometry_before = _visible_formatting_fingerprint(tmp_path / "candidate/protocol.docx")[1:]
    with zipfile.ZipFile(tmp_path / "candidate/protocol.docx") as package:
        protocol_xml_before = package.read("word/document.xml")
    repaired = render_documents(
        ROOT,
        tmp_path,
        reference,
        model,
        artifact_names={"protocol"},
        layout_repairs={"protocol": ({"rule": "heading_cohesion", "target": "5. INTRODUCTION"},)},
    )

    assert baseline["status"] == repaired["status"] == "passed"
    assert [item["artifact"] for item in repaired["artifacts"]] == ["protocol"]
    assert repaired["layout_repairs"] == {
        "protocol": [{"rule": "heading_cohesion", "target": "5. INTRODUCTION"}],
    }
    assert "layout_repair" not in repaired
    assert (tmp_path / "candidate/icf.docx").read_bytes() == icf_before
    with zipfile.ZipFile(tmp_path / "candidate/protocol.docx") as package:
        assert package.read("word/document.xml") != protocol_xml_before
    repaired_protocol = Document(tmp_path / "candidate/protocol.docx")
    assert tuple(
        paragraph._p.xml
        for paragraph in repaired_protocol.paragraphs
        if paragraph.text.strip() != "5. INTRODUCTION"
    ) == unaffected_paragraphs_before
    assert tuple(table._tbl.xml for table in repaired_protocol.tables) == unaffected_tables_before
    assert _visible_formatting_fingerprint(tmp_path / "candidate/protocol.docx")[1:] == geometry_before
    assert repaired["contracted_template_bundle"]["layout_preservation_baseline"] == baseline["contracted_template_bundle"]["layout_preservation_baseline"]


def test_artificial_pagination_fails_closed_instead_of_removing_template_breaks():
    finding = {
        "category": "visual",
        "artifact": "protocol",
        "check": "artificial_pagination",
        "element": "15. REFERENCES",
        "target_ids": ["layout:protocol"],
        "issue": "A later template-owned heading starts on a new page.",
    }

    plan, unsupported = workflow._layout_repair_plan(
        [finding],
        study_type="Prospective",
    )

    assert plan == {}
    assert unsupported == [{
        **finding,
        "contracted_layout_family": "prospective-protocol",
        "disposition": "fail_closed",
        "evidence_retained": True,
        "required": "This recognized visual defect has no safe deterministic repair. Preserve its exact rendered evidence for governed corpus expansion.",
    }]


def test_visual_disposition_authority_covers_every_check_for_every_layout_family():
    assert set(VISUAL_CHECK_DISPOSITIONS) == set(LAYOUT_FAMILY_ARTIFACTS)
    assert set(LAYOUT_FAMILY_ARTIFACTS) == {
        "retrospective-protocol",
        "prospective-protocol",
        "ambispective-protocol",
        "advarra-icf",
        "sterling-icf",
    }
    assert all(
        set(dispositions) == set(VISUAL_CHECKS)
        for dispositions in VISUAL_CHECK_DISPOSITIONS.values()
    )


@pytest.mark.parametrize(
    ("study_type", "icf_template", "artifact", "family", "element"),
    (
        ("Retrospective", None, "protocol", "retrospective-protocol", "4. INTRODUCTION"),
        ("Prospective", None, "protocol", "prospective-protocol", "5. INTRODUCTION"),
        ("Ambispective", None, "protocol", "ambispective-protocol", "5. INTRODUCTION"),
        ("Prospective", "Advarra", "icf", "advarra-icf", "INTRODUCTION"),
        ("Prospective", "Sterling", "icf", "sterling-icf", "BACKGROUND"),
    ),
)
def test_layout_repair_selection_is_explicit_for_all_five_contracted_families(
    study_type,
    icf_template,
    artifact,
    family,
    element,
):
    finding = {
        "category": "visual",
        "artifact": artifact,
        "check": "orphan_heading",
        "element": element,
        "target_ids": [f"layout:{artifact}"],
        "issue": "A heading is separated from its first substantive block.",
    }

    plan, unsupported = workflow._layout_repair_plan(
        [finding],
        study_type=study_type,
        icf_template=icf_template,
    )

    assert VISUAL_CHECK_DISPOSITIONS[family]["orphan_heading"] == "repair:heading_cohesion"
    assert plan == {artifact: [{"rule": "heading_cohesion", "target": element}]}
    assert unsupported == []


@pytest.mark.parametrize(
    ("check", "expected_disposition"),
    (
        ("blank_page", "prevention:render_audit"),
        ("clipping", "fail_closed"),
        ("not_a_mandatory_check", "fail_closed:unknown_visual_check"),
    ),
)
def test_nonrepair_visual_dispositions_fail_closed_with_reproducible_evidence(
    check,
    expected_disposition,
):
    finding = {
        "category": "visual",
        "artifact": "protocol",
        "check": check,
        "element": "5. INTRODUCTION",
        "target_ids": ["layout:protocol"],
        "issue": "Rendered evidence exposes the defect.",
    }

    plan, unsupported = workflow._layout_repair_plan(
        [finding],
        study_type="Prospective",
    )

    assert plan == {}
    assert unsupported == [{
        **finding,
        "contracted_layout_family": "prospective-protocol",
        "disposition": expected_disposition,
        "evidence_retained": True,
        "required": (
            "The declared prevention invariant failed. Preserve its exact rendered evidence "
            "and stop before publication."
            if expected_disposition.startswith("prevention:")
            else "This recognized visual defect has no safe deterministic repair. Preserve its exact rendered evidence for governed corpus expansion."
            if expected_disposition == "fail_closed"
            else "This visual check is not part of the governed Layout Contract. Preserve its exact rendered evidence for classification before retry."
        ),
    }]


def test_section_three_table_repair_subsumes_its_artificial_pagination_symptom():
    shared = {
        "category": "visual",
        "artifact": "protocol",
        "element": "3. GENERAL INFORMATION",
        "target_ids": ["layout:protocol"],
        "recovery_class": "visual_defect",
        "action": "targeted_layout_repair",
    }
    findings = [
        {
            **shared,
            "page": 2,
            "check": "bad_table_split",
            "issue": "The Duration / Follow-up row continues alone on page 3.",
        },
        {
            **shared,
            "page": 3,
            "check": "excessive_whitespace",
            "issue": "The continuation leaves nearly the entire page unused.",
        },
        {
            **shared,
            "page": 3,
            "check": "artificial_pagination",
            "issue": "The isolated continuation creates a standalone page before the TOC.",
        },
    ]

    plan, unsupported = workflow._layout_repair_plan(
        findings,
        study_type="Prospective",
    )

    assert plan == {
        "protocol": [
            {"rule": "heading_cohesion", "target": "3. GENERAL INFORMATION"},
            {"rule": "table_pagination", "target": "3. GENERAL INFORMATION"},
        ],
    }
    assert unsupported == []


def test_escalated_table_boundary_subsumes_the_same_artificial_pagination_symptom():
    shared = {
        "category": "visual",
        "artifact": "protocol",
        "element": "3. GENERAL INFORMATION",
        "target_ids": ["layout:protocol"],
        "recovery_class": "visual_defect",
        "action": "targeted_layout_repair",
    }
    plan, unsupported = workflow._layout_repair_plan(
        [
            {**shared, "check": "bad_table_split", "issue": "split summary table"},
            {**shared, "check": "artificial_pagination", "issue": "isolated continuation"},
        ],
        study_type="Prospective",
        existing_repairs={
            "protocol": [
                {"rule": "table_pagination", "target": "3. GENERAL INFORMATION"},
            ],
        },
    )

    assert plan == {
        "protocol": [
            {"rule": "table_page_boundary", "target": "3. GENERAL INFORMATION"},
        ],
    }
    assert unsupported == []


def test_section_three_endpoint_split_routes_to_the_exact_summary_table_heading():
    finding = {
        "category": "visual",
        "artifact": "protocol",
        "check": "bad_table_split",
        "element": "3. GENERAL INFORMATION – Variables / Secondary endpoint(s)",
        "target_ids": ["layout:protocol"],
        "issue": "The Secondary endpoint(s) label is separated from its first bullet.",
    }

    plan, unsupported = workflow._layout_repair_plan(
        [
            finding,
            {
                **finding,
                "check": "artificial_pagination",
                "issue": "The same split leaves a nearly empty continuation page before the TOC.",
            },
        ],
        study_type="Prospective",
    )

    assert plan == {
        "protocol": [{"rule": "table_pagination", "target": "3. GENERAL INFORMATION"}],
    }
    assert unsupported == []


def test_excessive_whitespace_at_exact_heading_uses_scoped_cohesion_repair():
    finding = {
        "category": "visual",
        "artifact": "icf",
        "check": "excessive_whitespace",
        "element": "DURATION",
        "target_ids": ["layout:icf"],
        "issue": "A large vertical gap separates DURATION from its first substantive paragraph.",
    }

    plan, unsupported = workflow._layout_repair_plan(
        [finding],
        study_type="Prospective",
        icf_template="Sterling",
    )

    assert plan == {"icf": [{"rule": "heading_whitespace_cohesion", "target": "DURATION"}]}
    assert unsupported == []


def test_excessive_whitespace_at_other_sterling_heading_remains_non_destructive():
    finding = {
        "category": "visual",
        "artifact": "icf",
        "check": "excessive_whitespace",
        "element": "RISKS",
        "target_ids": ["layout:icf"],
        "issue": "A large vertical gap follows an unrelated ICF heading.",
    }

    plan, unsupported = workflow._layout_repair_plan(
        [finding],
        study_type="Prospective",
        icf_template="Sterling",
    )

    assert plan == {"icf": [{"rule": "heading_cohesion", "target": "RISKS"}]}
    assert unsupported == []


def test_excessive_whitespace_at_advarra_duration_remains_non_destructive():
    finding = {
        "category": "visual",
        "artifact": "icf",
        "check": "excessive_whitespace",
        "element": "DURATION",
        "target_ids": ["layout:icf"],
        "issue": "A large vertical gap follows an Advarra ICF heading.",
    }

    plan, unsupported = workflow._layout_repair_plan(
        [finding],
        study_type="Prospective",
        icf_template="Advarra",
    )

    assert plan == {"icf": [{"rule": "heading_cohesion", "target": "DURATION"}]}
    assert unsupported == []


def test_protocol_excessive_whitespace_retains_non_destructive_heading_cohesion():
    finding = {
        "category": "visual",
        "artifact": "protocol",
        "check": "excessive_whitespace",
        "element": "19.2 Study Completion",
        "target_ids": ["layout:protocol"],
        "issue": "A heading is isolated above excessive remaining whitespace.",
    }

    plan, unsupported = workflow._layout_repair_plan(
        [finding],
        study_type="Prospective",
    )

    assert plan == {"protocol": [{"rule": "heading_cohesion", "target": "19.2 Study Completion"}]}
    assert unsupported == []


def test_heading_cohesion_removes_empty_template_paragraphs_before_its_first_block():
    document = Document()
    document.add_heading("DURATION", level=1)
    document.add_paragraph("")
    document.add_paragraph("")
    document.add_paragraph("")
    document.add_paragraph("The study lasts approximately 14 weeks.")

    rendering._repair_heading_cohesion(
        document,
        "DURATION",
        protocol=False,
        remove_empty_intervening_paragraphs=True,
    )

    duration = next(index for index, paragraph in enumerate(document.paragraphs) if paragraph.text == "DURATION")
    assert document.paragraphs[duration + 1].text == "The study lasts approximately 14 weeks."


def test_layout_invalidation_removes_only_the_changed_artifacts_render_and_visual_evidence(tmp_path):
    revision = tmp_path / "revision"
    candidate = revision / "candidate"
    rendered = revision / "rendered"
    requests = revision / "hermes/verification-requests"
    responses = revision / "hermes/verification-responses"
    for directory in (candidate, rendered / "protocol", rendered / "icf", requests, responses):
        directory.mkdir(parents=True, exist_ok=True)
    for name in ("protocol.docx", "icf.docx"):
        (candidate / name).write_bytes(name.encode())
    for name in ("protocol.pdf", "icf.pdf"):
        (rendered / name).write_bytes(name.encode())
    (rendered / "protocol/page-1.png").write_bytes(b"protocol page")
    (rendered / "icf/page-1.png").write_bytes(b"icf page")
    for artifact in ("protocol", "icf"):
        response_path = responses / f"{artifact}.json"
        response_path.write_text("{}", encoding="utf-8")
        (requests / f"{artifact}.json").write_text(json.dumps({
            "task": "rendered_page_visual_verification",
            "artifacts": [{"artifact": artifact}],
            "response_path": response_path.relative_to(revision).as_posix(),
        }), encoding="utf-8")

    workflow._invalidate_layout_artifact(revision, "protocol")

    assert not (candidate / "protocol.docx").exists()
    assert not (rendered / "protocol.pdf").exists()
    assert not (rendered / "protocol").exists()
    assert not (requests / "protocol.json").exists()
    assert not (responses / "protocol.json").exists()
    assert (candidate / "icf.docx").is_file()
    assert (rendered / "icf.pdf").is_file()
    assert (rendered / "icf/page-1.png").is_file()
    assert (requests / "icf.json").is_file()
    assert (responses / "icf.json").is_file()


def test_heading_whitespace_cohesion_removes_contracted_keep_together_spacers():
    document = Document()
    document.add_heading("DURATION", level=1)
    for _ in range(3):
        spacer = document.add_paragraph("")
        properties = spacer._p.get_or_add_pPr()
        properties.append(OxmlElement("w:keepNext"))
        properties.append(OxmlElement("w:keepLines"))
    document.add_paragraph("The study lasts approximately 14 weeks.")

    rendering._repair_heading_cohesion(
        document,
        "DURATION",
        protocol=False,
        remove_empty_intervening_paragraphs=True,
    )

    duration = next(index for index, paragraph in enumerate(document.paragraphs) if paragraph.text == "DURATION")
    assert document.paragraphs[duration + 1].text == "The study lasts approximately 14 weeks."


def test_sterling_duration_whitespace_repair_removes_contracted_template_spacers(tmp_path):
    reference = json.loads(
        (ROOT / "tests/fixtures/ambispective-acceptance-source.json").read_text(encoding="utf-8")
    )
    reference["meta"]["icf_template"] = "Sterling"
    model = {
        "protocol": [],
        "icf": {"icf.procedures": _minimal_source_bound_icf_procedures()},
        "prs": {},
    }
    baseline_dir = tmp_path / "baseline"
    repaired_dir = tmp_path / "repaired"

    baseline = render_documents(ROOT, baseline_dir, reference, model, artifact_names={"icf"})
    repaired = render_documents(
        ROOT,
        repaired_dir,
        reference,
        model,
        artifact_names={"icf"},
        layout_repairs={
            "icf": ({"rule": "heading_whitespace_cohesion", "target": "DURATION"},),
        },
    )

    def empty_paragraphs_after_duration(path):
        paragraphs = Document(path).paragraphs
        duration = next(index for index, paragraph in enumerate(paragraphs) if paragraph.text == "DURATION")
        count = 0
        for paragraph in paragraphs[duration + 1 :]:
            if paragraph.text.strip():
                break
            count += 1
        return count

    assert baseline["status"] == repaired["status"] == "passed"
    assert empty_paragraphs_after_duration(baseline_dir / "candidate/icf.docx") == 3
    assert empty_paragraphs_after_duration(repaired_dir / "candidate/icf.docx") == 0


def test_heading_whitespace_cohesion_preserves_empty_section_boundary_paragraph():
    document = Document()
    document.add_heading("DURATION", level=1)
    section_boundary = document.add_paragraph("")
    section_boundary._p.get_or_add_pPr().append(OxmlElement("w:sectPr"))
    document.add_paragraph("The study lasts approximately 14 weeks.")

    rendering._repair_heading_cohesion(
        document,
        "DURATION",
        protocol=False,
        remove_empty_intervening_paragraphs=True,
    )

    duration = next(index for index, paragraph in enumerate(document.paragraphs) if paragraph.text == "DURATION")
    retained = document.paragraphs[duration + 1]
    assert retained.text == ""
    assert retained._p.find(qn("w:pPr") + "/" + qn("w:sectPr")) is not None


def test_heading_whitespace_cohesion_preserves_empty_bookmark_paragraph():
    document = Document()
    document.add_heading("DURATION", level=1)
    bookmark_paragraph = document.add_paragraph("")
    bookmark = OxmlElement("w:bookmarkStart")
    bookmark.set(qn("w:id"), "7")
    bookmark.set(qn("w:name"), "duration-boundary")
    bookmark_paragraph._p.append(bookmark)
    document.add_paragraph("The study lasts approximately 14 weeks.")

    rendering._repair_heading_cohesion(
        document,
        "DURATION",
        protocol=False,
        remove_empty_intervening_paragraphs=True,
    )

    duration = next(index for index, paragraph in enumerate(document.paragraphs) if paragraph.text == "DURATION")
    retained = document.paragraphs[duration + 1]
    assert retained._p.find(qn("w:bookmarkStart")) is not None


def test_heading_whitespace_cohesion_preserves_empty_numbered_paragraph():
    document = Document()
    document.add_heading("DURATION", level=1)
    numbered_paragraph = document.add_paragraph("")
    numbering = OxmlElement("w:numPr")
    level = OxmlElement("w:ilvl")
    level.set(qn("w:val"), "0")
    number_id = OxmlElement("w:numId")
    number_id.set(qn("w:val"), "1")
    numbering.extend((level, number_id))
    numbered_paragraph._p.get_or_add_pPr().append(numbering)
    document.add_paragraph("The study lasts approximately 14 weeks.")

    rendering._repair_heading_cohesion(
        document,
        "DURATION",
        protocol=False,
        remove_empty_intervening_paragraphs=True,
    )

    duration = next(index for index, paragraph in enumerate(document.paragraphs) if paragraph.text == "DURATION")
    retained = document.paragraphs[duration + 1]
    assert retained._p.find(qn("w:pPr") + "/" + qn("w:numPr")) is not None


def test_heading_whitespace_cohesion_preserves_inherited_page_boundary():
    document = Document()
    boundary_style = document.styles.add_style("Boundary Style", WD_STYLE_TYPE.PARAGRAPH)
    boundary_style.paragraph_format.page_break_before = True
    document.add_heading("DURATION", level=1)
    boundary_paragraph = document.add_paragraph("")
    boundary_paragraph.style = boundary_style
    document.add_paragraph("The study lasts approximately 14 weeks.")

    rendering._repair_heading_cohesion(
        document,
        "DURATION",
        protocol=False,
        remove_empty_intervening_paragraphs=True,
    )

    duration = next(index for index, paragraph in enumerate(document.paragraphs) if paragraph.text == "DURATION")
    retained = document.paragraphs[duration + 1]
    assert retained.style.name == "Boundary Style"
    assert retained.style.paragraph_format.page_break_before is True


def _visible_formatting_fingerprint(path):
    document = Document(path)
    paragraph_layout = tuple(
        (
            paragraph.text,
            paragraph.style.name,
            paragraph.alignment,
            paragraph.paragraph_format.keep_with_next,
            paragraph.paragraph_format.page_break_before,
        )
        for paragraph in document.paragraphs
    )
    page_geometry = tuple(
        (
            section.page_width,
            section.page_height,
            section.left_margin,
            section.right_margin,
            section.top_margin,
            section.bottom_margin,
            section.header_distance,
            section.footer_distance,
        )
        for section in document.sections
    )
    tables = tuple(
        tuple(tuple((cell.text, cell.width) for cell in row.cells) for row in table.rows)
        for table in document.tables
    )
    running_text = tuple(
        (
            tuple(paragraph.text for paragraph in section.header.paragraphs),
            tuple(paragraph.text for paragraph in section.footer.paragraphs),
        )
        for section in document.sections
    )
    return paragraph_layout, page_geometry, tables, running_text


@pytest.mark.parametrize(
    ("fixture_name", "icf_family", "candidate_names"),
    (
        ("prospective-acceptance-source.json", "Advarra", ("protocol.docx", "icf.docx")),
        ("prospective-acceptance-source.json", "Sterling", ("protocol.docx", "icf.docx")),
        ("ambispective-acceptance-source.json", "Advarra", ("protocol.docx", "icf.docx")),
        ("ambispective-acceptance-source.json", "Sterling", ("protocol.docx", "icf.docx")),
        ("retrospective-acceptance-source.json", None, ("protocol.docx",)),
    ),
)
def test_parallel_bundle_identity_preserves_candidate_bytes_and_visible_formatting(
    tmp_path,
    monkeypatch,
    fixture_name,
    icf_family,
    candidate_names,
):
    reference = json.loads((ROOT / "tests/fixtures" / fixture_name).read_text(encoding="utf-8"))
    if icf_family is not None:
        reference["meta"]["icf_template"] = icf_family
    model = {
        "protocol": [],
        "icf": (
            {"icf.procedures": _minimal_source_bound_icf_procedures()}
            if icf_family is not None
            else {}
        ),
        "prs": {},
    }
    baseline_dir = tmp_path / "baseline"
    parallel_identity_dir = tmp_path / "parallel-identity"
    monkeypatch.setattr(zipfile.time, "time", lambda: 1_800_000_000.0)

    bundle = contracted_template_bundle(ROOT, reference)
    legacy_bundle = copy.deepcopy(bundle)
    study_type = reference["meta"]["study_type"]
    protocol_name = {
        "Prospective": "prospective-protocol.template.docx",
        "Ambispective": "ambispective-protocol.template.docx",
        "Retrospective": "retrospective-protocol.template.docx",
    }[study_type]
    legacy_bundle["contracted_templates"]["protocol"]["path"] = f"assets/client-templates/docx/{protocol_name}"
    legacy_bundle["client_template_authorities"]["protocol"]["path"] = "assets/client-templates/reference/protocol-reference.docx"
    if icf_family is not None:
        icf_name = (
            "sterling-icf.template.docx"
            if icf_family == "Sterling"
            else "ambispective-icf.template.docx"
            if study_type == "Ambispective"
            else "prospective-icf.template.docx"
        )
        legacy_bundle["contracted_templates"]["icf"]["path"] = f"assets/client-templates/docx/{icf_name}"
        legacy_bundle["client_template_authorities"]["icf"]["path"] = (
            f"assets/client-templates/reference/{icf_family.casefold()}-icf-reference.docx"
        )
    baseline = render_documents(
        ROOT,
        baseline_dir,
        reference,
        model,
        contracted_bundle=legacy_bundle,
    )
    parallel = render_documents(
        ROOT,
        parallel_identity_dir,
        reference,
        model,
        contracted_bundle=bundle,
    )

    assert baseline["status"] == parallel["status"] == "passed"
    assert len(bundle["identity_sha256"]) == 64
    assert parallel["contracted_template_bundle"] == bundle
    for artifact in parallel["artifacts"]:
        assert artifact["template"] == bundle["contracted_templates"][artifact["artifact"]]["path"]
        assert artifact["client_template_authority"] == bundle["client_template_authorities"][artifact["artifact"]]["path"]
    for candidate_name in candidate_names:
        baseline_path = baseline_dir / "candidate" / candidate_name
        parallel_path = parallel_identity_dir / "candidate" / candidate_name
        assert baseline_path.read_bytes() == parallel_path.read_bytes()
        assert _visible_formatting_fingerprint(baseline_path) == _visible_formatting_fingerprint(parallel_path)


def test_layout_retry_persists_scoped_rule_and_the_original_operation_deadline(tmp_path, monkeypatch):
    run_dir = tmp_path / "run"
    revision_dir = run_dir / "revisions/r-test"
    revision_dir.mkdir(parents=True)
    reference_path = run_dir / "reference/study.reference.json"
    reference_path.parent.mkdir(parents=True)
    working_reference = {"generation": {"layout_repairs": {"icf": [{"rule": "heading_cohesion", "target": "QUESTIONS"}]}}}
    reference = json.loads((ROOT / "tests/fixtures/prospective-acceptance-source.json").read_text(encoding="utf-8"))
    observed = {}

    def fake_generate(path, **kwargs):
        observed.update({"path": path, **kwargs})
        return {"status": "awaiting_hermes"}

    monkeypatch.setattr(workflow, "generate", fake_generate)
    clock = lambda: 10.0

    workflow._quality_retry(
        run_dir,
        reference_path,
        working_reference,
        reference,
        revision_dir,
        {},
        [{"category": "visual", "field": "protocol", "artifact": "protocol", "check": "orphan_heading", "element": "5. INTRODUCTION", "target_ids": ["layout:protocol"], "recovery_class": "visual_defect", "action": "targeted_layout_repair", "issue": "orphan heading"}],
        "rendered_document_qa",
        operation_deadline=99.0,
        clock=clock,
    )

    persisted = json.loads(reference_path.read_text(encoding="utf-8"))
    assert persisted["generation"]["layout_repairs"] == {
        "icf": [{"rule": "heading_cohesion", "target": "QUESTIONS"}],
        "protocol": [{"rule": "heading_cohesion", "target": "5. INTRODUCTION"}],
    }
    assert persisted["generation"]["pending_layout_artifacts"] == ["protocol"]
    assert "layout_repair_level" not in persisted["generation"]
    assert observed == {"path": run_dir, "operation_deadline": 99.0, "clock": clock}
    attempt_manifest = json.loads(
        next((revision_dir / "attempts").glob("*/attempt-manifest.json")).read_text(
            encoding="utf-8"
        )
    )
    action = attempt_manifest["recovery_actions"][0]
    assert action["triggering_finding"]["repair_rule"] == "heading_cohesion"
    assert action["triggering_finding"]["repair_target"] == "5. INTRODUCTION"
    assert action["strategy_id"].endswith(":heading_cohesion")


def test_layout_retry_requires_an_exact_repair_element(tmp_path):
    finding = {
        "category": "visual",
        "artifact": "protocol",
        "check": "orphan_heading",
        "target_ids": ["layout:protocol"],
        "issue": "orphan heading without a bound element",
    }

    plan, unsupported = workflow._layout_repair_plan(
        [finding],
        study_type="Prospective",
    )

    assert plan == {}
    assert unsupported[0]["required"].endswith(
        "exact heading or table-caption element before deterministic repair."
    )


def test_layout_repair_rejects_a_nonexact_table_target_as_classification(tmp_path):
    reference = json.loads((ROOT / "tests/fixtures/prospective-acceptance-source.json").read_text(encoding="utf-8"))

    report = render_documents(
        ROOT,
        tmp_path,
        reference,
        {"protocol": [], "icf": {}, "prs": {}},
        artifact_names={"protocol"},
        layout_repairs={"protocol": ({"rule": "table_pagination", "target": "Table"},)},
    )

    assert report["status"] == "blocked"
    assert report["artifacts"][0]["findings"] == [{
        "category": "layout-repair-classification",
        "field": "protocol",
        "artifact": "protocol",
        "issue": "Layout repair target table must match exactly once; found 0: Table",
    }]

    run_dir = tmp_path / "run"
    revision_dir = run_dir / "revisions/r-test"
    revision_dir.mkdir(parents=True)
    (run_dir / "reference").mkdir()
    findings, block = workflow._document_report_failure(run_dir, revision_dir, report)
    assert findings[0]["target_ids"] == ["layout:protocol"]
    assert block is not None
    assert block["stage"] == "layout_repair_classification"


def test_document_report_failure_preserves_a_governed_drafting_route_through_quality_retry(tmp_path, monkeypatch):
    report = {
        "status": "blocked",
        "artifacts": [{
            "artifact": "icf",
            "findings": [{
                "category": "content",
                "field": "icf.study-purpose",
                "target_ids": ["icf.study-purpose"],
                "issue": "Study-purpose content is missing.",
                "recovery_class": "drafting_defect",
                "action": "retry_drafting_target",
            }],
        }],
    }
    run_dir = tmp_path / "run"
    revision_dir = run_dir / "revisions/r-test"
    revision_dir.mkdir(parents=True)
    reference_path = run_dir / "reference/study.reference.json"
    reference_path.parent.mkdir(parents=True)
    reference_path.write_text(json.dumps({"generation": {}}), encoding="utf-8")

    def fake_schedule_requests(**_kwargs):
        path = revision_dir / "hermes/drafting-requests/icf-study-purpose.json"
        path.parent.mkdir(parents=True)
        path.write_text(json.dumps({
            "task": "draft_sections",
            "response_path": "hermes/drafting-responses/icf-study-purpose.json",
        }), encoding="utf-8")
        return [path]

    monkeypatch.setattr(workflow, "schedule_requests", fake_schedule_requests)

    findings, block = workflow._document_report_failure(run_dir, revision_dir, report)
    result = workflow._quality_retry(
        run_dir,
        reference_path,
        {"generation": {}},
        json.loads((ROOT / "tests/fixtures/prospective-acceptance-source.json").read_text(encoding="utf-8")),
        revision_dir,
        {},
        findings,
        "quality",
    )

    assert block is None
    assert findings == report["artifacts"][0]["findings"]
    assert result["status"] == "awaiting_hermes", result
    assert result["stage"] == "drafting_retry"


def test_partial_render_merge_keeps_each_artifacts_bound_renderer(tmp_path):
    prior = {
        "renderer": {"kind": "Microsoft Word"},
        "page_renderer": {"kind": "pypdfium2"},
        "artifacts": [
            {"artifact": "icf", "pages": []},
            {"artifact": "protocol", "pages": []},
        ],
    }
    initial_paths = quality.create_verification_requests(tmp_path, _source(), prior)
    initial_icf = next(path for path in initial_paths if "visual.icf" in path.name)
    initial_request = json.loads(initial_icf.read_text(encoding="utf-8"))
    response_path = tmp_path / initial_request["response_path"]
    response_path.parent.mkdir(parents=True, exist_ok=True)
    response_path.write_text("{}", encoding="utf-8")
    current = {
        "status": "passed",
        "renderer": {"kind": "LibreOffice"},
        "page_renderer": {"kind": "pypdfium2"},
        "artifacts": [
            {"artifact": "protocol", "renderer": {"kind": "LibreOffice"}, "page_renderer": {"kind": "pypdfium2"}, "pages": []},
        ],
    }

    merged = workflow._merge_artifact_reports(prior, current)
    assurance, assurance_render = workflow._merge_partial_assurance(
        {"render_report": prior},
        {"status": "passed", "render": current},
        {"protocol"},
    )
    paths = quality.create_verification_requests(tmp_path, _source(), merged)
    icf_request = json.loads(next(path for path in paths if "visual.icf" in path.name).read_text(encoding="utf-8"))
    protocol_request = json.loads(next(path for path in paths if "visual.protocol" in path.name).read_text(encoding="utf-8"))

    assert icf_request["request_sha256"] == initial_request["request_sha256"]
    assert response_path.is_file()
    assert icf_request["renderer"] == {"kind": "Microsoft Word"}
    assert protocol_request["renderer"] == {"kind": "LibreOffice"}
    assert assurance["render"] == assurance_render == merged


def test_renderer_preflight_render_verifies_an_uninspectable_font_instead_of_failing(monkeypatch):
    identity = {"kind": "LibreOffice", "path": "/test/soffice", "version": "test", "platform": "Linux"}
    page_identity = _page_identity()
    monkeypatch.setattr(quality, "renderers", lambda **_: [identity])
    monkeypatch.setattr(quality, "page_renderers", lambda **_: [page_identity])
    monkeypatch.setattr(quality, "_template_fonts", lambda _path: {"Client Sans"})
    monkeypatch.setattr(quality, "_font_probe", lambda *_args, **_kwargs: (None, "font inventory unavailable"))

    def fake_export(_docx, output_dir, _identity, **_kwargs):
        from pypdf import PdfWriter
        output = output_dir / "preflight.pdf"
        writer = PdfWriter()
        writer.add_blank_page(width=612, height=792)
        with output.open("wb") as handle:
            writer.write(handle)
        return output

    def fake_rasterize(_pdf, output_dir, _identity, **_kwargs):
        output_dir.mkdir(parents=True, exist_ok=True)
        page = output_dir / "page.png"
        page.write_bytes(b"png")
        return [page]

    monkeypatch.setattr(quality, "_render_pdf", fake_export)
    monkeypatch.setattr(quality, "rasterize_pdf", fake_rasterize)

    report = quality.preflight(ROOT, _source(), deadline_seconds=10)

    assert report["status"] == "passed"
    assert report["font_substitutions"] == {}
    assert report["fonts"]["Client Sans"] == {
        "available": None,
        "match": "font inventory unavailable",
        "resolution": "render_verified",
    }


def test_renderer_preflight_falls_through_a_broken_preferred_renderer(monkeypatch):
    word = {"kind": "Microsoft Word", "path": "/Applications/Microsoft Word.app", "version": "test", "platform": "Darwin"}
    fallback = {"kind": "LibreOffice", "path": "/host/soffice", "version": "test", "platform": "Darwin", "source": "host prerequisite"}
    page_identity = _page_identity()
    monkeypatch.setattr(quality, "renderers", lambda **_: [word, fallback])
    monkeypatch.setattr(quality, "page_renderers", lambda **_: [page_identity])
    monkeypatch.setattr(quality, "_template_fonts", lambda _path: set())

    def fake_export(_docx, output_dir, identity, **_kwargs):
        if identity == word:
            raise RuntimeError("Word automation unavailable")
        from pypdf import PdfWriter
        output = output_dir / "preflight.pdf"
        writer = PdfWriter()
        writer.add_blank_page(width=612, height=792)
        with output.open("wb") as handle:
            writer.write(handle)
        return output

    def fake_rasterize(_pdf, output_dir, _identity, **_kwargs):
        output_dir.mkdir(parents=True, exist_ok=True)
        page = output_dir / "page.png"
        page.write_bytes(b"png")
        return [page]

    monkeypatch.setattr(quality, "_render_pdf", fake_export)
    monkeypatch.setattr(quality, "rasterize_pdf", fake_rasterize)

    report = quality.preflight(ROOT, _source(), deadline_seconds=10)

    assert report["status"] == "passed"
    assert report["renderer"] == fallback
    assert report["renderer_attempts"] == [
        {"renderer": word, "status": "failed", "issue": "Word automation unavailable"},
        {"renderer": fallback, "status": "passed"},
    ]


def test_document_rendering_falls_through_tool_failure_to_the_next_renderer(tmp_path, monkeypatch):
    candidate = tmp_path / "candidate"
    candidate.mkdir()
    Document().save(candidate / "protocol.docx")
    word = {"kind": "Microsoft Word", "path": "/test/word", "platform": "Darwin"}
    fallback = {"kind": "LibreOffice", "path": "/test/soffice", "platform": "Darwin", "source": "host prerequisite"}
    page_identity = _page_identity()

    def fake_export(docx, output_dir, identity, **_kwargs):
        if identity == word:
            raise RuntimeError("Word automation unavailable")
        from pypdf import PdfWriter
        output = output_dir / f"{docx.stem}.pdf"
        writer = PdfWriter()
        writer.add_blank_page(width=612, height=792)
        with output.open("wb") as handle:
            writer.write(handle)
        return output

    def fake_rasterize(_pdf, output_dir, _identity, **_kwargs):
        output_dir.mkdir(parents=True, exist_ok=True)
        page = output_dir / "page-1.png"
        page.write_bytes(b"png")
        return [page]

    monkeypatch.setattr(quality, "_render_pdf", fake_export)
    monkeypatch.setattr(quality, "rasterize_pdf", fake_rasterize)
    monkeypatch.setattr(quality, "page_renderers", lambda **_: [_page_identity()])
    monkeypatch.setattr(quality, "refresh_toc_from_pdf", lambda *_args: False)
    monkeypatch.setattr(quality, "_blank_pdf_pages", lambda _pdf: [])

    report = quality.render_pages(
        tmp_path,
        renderer_identities=[word, fallback],
        page_renderer_identities=[page_identity],
    )

    assert report["status"] == "passed"
    assert report["renderer"] == fallback
    assert report["renderer_attempts"] == [
        {"renderer": word, "status": "failed", "issue": "Word automation unavailable"},
        {"renderer": fallback, "status": "passed"},
    ]


def test_real_visual_defect_does_not_switch_away_from_the_client_renderer(tmp_path, monkeypatch):
    candidate = tmp_path / "candidate"
    candidate.mkdir()
    Document().save(candidate / "protocol.docx")
    word = {"kind": "Microsoft Word", "path": "/test/word", "platform": "Darwin"}
    fallback = {"kind": "LibreOffice", "path": "/test/soffice", "platform": "Darwin", "source": "host prerequisite"}
    page_identity = _page_identity()
    rendered_by = []

    def fake_export(docx, output_dir, identity, **_kwargs):
        rendered_by.append(identity["kind"])
        from pypdf import PdfWriter
        output = output_dir / f"{docx.stem}.pdf"
        writer = PdfWriter()
        writer.add_blank_page(width=612, height=792)
        with output.open("wb") as handle:
            writer.write(handle)
        return output

    def fake_rasterize(_pdf, output_dir, _identity, **_kwargs):
        output_dir.mkdir(parents=True, exist_ok=True)
        page = output_dir / "page-1.png"
        page.write_bytes(b"png")
        return [page]

    monkeypatch.setattr(quality, "_render_pdf", fake_export)
    monkeypatch.setattr(quality, "rasterize_pdf", fake_rasterize)
    monkeypatch.setattr(quality, "page_renderers", lambda **_: [_page_identity()])
    monkeypatch.setattr(quality, "refresh_toc_from_pdf", lambda *_args: False)
    monkeypatch.setattr(quality, "_blank_pdf_pages", lambda _pdf: [1])

    report = quality.render_pages(
        tmp_path,
        renderer_identities=[word, fallback],
        page_renderer_identities=[page_identity],
    )

    assert report["status"] == "blocked"
    assert report["renderer"] == word
    assert rendered_by == ["Microsoft Word"]
    assert report["findings"][0]["category"] == "visual"


def test_layout_rerender_is_pinned_to_the_exposing_renderer_and_cannot_fall_through(tmp_path, monkeypatch):
    candidate = tmp_path / "candidate"
    candidate.mkdir()
    Document().save(candidate / "protocol.docx")
    word = {"kind": "Microsoft Word", "path": "/test/word", "platform": "Darwin"}
    fallback = {"kind": "LibreOffice", "path": "/test/soffice", "platform": "Darwin"}
    page_identity = _page_identity()
    prior_build = {
        "render_report": {
            "renderer": word,
            "page_renderer": page_identity,
            "artifacts": [{
                "artifact": "protocol",
                "renderer": word,
                "page_renderer": page_identity,
            }],
        },
    }
    office_identities, page_identities = workflow._retained_layout_renderer_pair(
        prior_build,
        {"protocol"},
    )
    rendered_by = []

    def fake_export(docx, output_dir, identity, **_kwargs):
        rendered_by.append(identity["kind"])
        if identity == word:
            raise RuntimeError("The exposing renderer is now unavailable")
        from pypdf import PdfWriter
        output = output_dir / f"{docx.stem}.pdf"
        writer = PdfWriter()
        writer.add_blank_page(width=612, height=792)
        with output.open("wb") as handle:
            writer.write(handle)
        return output

    monkeypatch.setattr(quality, "renderers", lambda **_: [word, fallback])
    monkeypatch.setattr(
        quality,
        "_one_pdfium_renderer",
        lambda identities, **_kwargs: list(identities),
    )
    report = quality.render_pages(
        tmp_path,
        renderer_identities=office_identities,
        page_renderer_identities=page_identities,
        office_exporter=fake_export,
        require_promoted_runtime=False,
    )

    assert office_identities == [word]
    assert page_identities == [page_identity]
    assert report["status"] == "blocked"
    assert rendered_by == ["Microsoft Word"]
    assert report["renderer_attempts"] == [{
        "renderer": word,
        "status": "failed",
        "issue": "The exposing renderer is now unavailable",
    }]


def test_render_pages_regenerates_only_the_selected_artifact(tmp_path, monkeypatch):
    candidate = tmp_path / "candidate"
    candidate.mkdir()
    for artifact in ("protocol", "icf"):
        Document().save(candidate / f"{artifact}.docx")
    retained_pdf = tmp_path / "rendered/icf.pdf"
    retained_page = tmp_path / "rendered/icf/page-1.png"
    retained_page.parent.mkdir(parents=True)
    retained_pdf.write_bytes(b"retained-icf-pdf")
    retained_page.write_bytes(b"retained-icf-page")
    rendered = []

    def fake_export(docx, output_dir, _identity, **_kwargs):
        rendered.append(docx.stem)
        from pypdf import PdfWriter
        output = output_dir / f"{docx.stem}.pdf"
        writer = PdfWriter()
        writer.add_blank_page(width=612, height=792)
        with output.open("wb") as handle:
            writer.write(handle)
        return output

    def fake_rasterize(_pdf, output_dir, _identity, **_kwargs):
        output_dir.mkdir(parents=True, exist_ok=True)
        page = output_dir / "page-1.png"
        page.write_bytes(b"protocol-page")
        return [page]

    monkeypatch.setattr(quality, "_render_pdf", fake_export)
    monkeypatch.setattr(quality, "rasterize_pdf", fake_rasterize)
    monkeypatch.setattr(quality, "page_renderers", lambda **_: [_page_identity()])
    monkeypatch.setattr(quality, "refresh_toc_from_pdf", lambda *_args: False)
    monkeypatch.setattr(quality, "_blank_pdf_pages", lambda _pdf: [])

    report = quality.render_pages(
        tmp_path,
        artifact_names={"protocol"},
        renderer_identities=[{"kind": "test", "path": "/test/renderer"}],
        page_renderer_identities=[_page_identity()],
    )

    assert report["status"] == "passed"
    assert rendered == ["protocol"]
    assert [item["artifact"] for item in report["artifacts"]] == ["protocol"]
    assert retained_pdf.read_bytes() == b"retained-icf-pdf"
    assert retained_page.read_bytes() == b"retained-icf-page"


def test_generated_docx_replaces_a_missing_client_font_with_the_selected_fallback(tmp_path):
    render_documents(
        ROOT,
        tmp_path,
        _source(),
        {"protocol": [], "icf": {}, "prs": {}},
        font_substitutions={"Noto Sans Symbols": "Apple Symbols"},
    )

    fonts = quality._template_fonts(tmp_path / "candidate/protocol.docx")

    assert "Noto Sans Symbols" not in fonts
    assert "Apple Symbols" in fonts


def test_symbol_and_text_fonts_have_multiple_cross_platform_fallbacks():
    assert quality._font_fallback_candidates("Noto Sans Symbols")[:5] == (
        "Apple Symbols",
        "Segoe UI Symbol",
        "Arial Unicode MS",
        "Arial Unicode",
        "Symbol",
    )
    assert quality._font_fallback_candidates("Verdana")[:4] == (
        "Arial",
        "Helvetica",
        "Aptos",
        "Calibri",
    )


def test_template_font_inventory_only_requires_fonts_used_by_visible_latin_text(tmp_path):
    document = Document()
    dormant = document.styles.add_style("Dormant client style", WD_STYLE_TYPE.CHARACTER)
    dormant.font.name = "Unavailable Dormant Font"
    run = document.add_paragraph().add_run("Visible Latin text")
    fonts = run._element.get_or_add_rPr().get_or_add_rFonts()
    fonts.set(qn("w:ascii"), "Arial")
    fonts.set(qn("w:hAnsi"), "Arial")
    fonts.set(qn("w:eastAsia"), "Unavailable East Asia Font")
    fonts.set(qn("w:cs"), "Unavailable Complex Script Font")
    path = tmp_path / "font-contract.docx"
    document.save(path)

    assert quality._template_fonts(path) == {"Arial"}


def test_renderer_preflight_can_run_an_actual_local_renderer_smoke_export(monkeypatch):
    identity = quality.renderer()
    page_identity = quality.page_renderer(skill_root=ROOT)
    if identity is None or page_identity is None:
        import pytest
        pytest.skip("No installed release runtime is available in the source checkout")
    monkeypatch.setattr(quality, "_font_probe", lambda *_args, **_kwargs: (True, "test-font"))
    report = quality.preflight(ROOT, _source(), deadline_seconds=30)

    assert report["status"] == "passed"
    assert report["smoke"]["status"] == "passed"
    assert report["smoke"]["pages"] >= 1


def test_approved_investigator_and_facility_values_populate_protocol_agreement(tmp_path):
    reference = _source()
    render_documents(ROOT, tmp_path, reference, {"protocol": [], "icf": {}, "prs": {}})

    document = Document(tmp_path / "candidate/protocol.docx")
    agreement = next(
        table for table in document.tables
        if any("Signature of Investigator" in cell.text for row in table.rows for cell in row.cells)
    )
    values = [cell.text for row in agreement.rows for cell in row.cells]
    assert "Alex Investigator" in values
    assert "MD" in values
    assert "Site One" in values


@pytest.mark.parametrize(
    "screening_interval",
    (90, "90 days", "90 days before screening", "At least 90 days before screening"),
)
def test_render_fields_normalize_markdown_email_population_and_day_units(screening_interval):
    reference = _source()
    reference["parties"]["study_coordinator"]["email"] = (
        "[jamie.chen@example.org](mailto:jamie.chen@example.org)"
    )
    reference["population"].pop("study_population", None)
    reference["population"]["inclusion_criteria"] = [
        "Adults 18 to 80 years old with eligible historical records.",
    ]
    reference["procedures"]["minimum_days_before_screening_without_participation"] = screening_interval

    fields = render_fields(reference, {"protocol": [], "icf": {}, "prs": {}})

    assert fields["studyCordinatorEmail"] == "jamie.chen@example.org"
    assert fields["AI_populationShort"] == (
        "Adults 18 to 80 years old with eligible historical records."
    )
    assert fields["daysBeforeScreening"] == "90"


def test_render_fields_prefers_approved_protocol_summary_and_named_articles():
    reference = _source()
    reference["design"].update({
        "study_design_summary": "Prospective randomized parallel-group study",
        "study_design": "A much longer approved design narrative for the body section.",
        "intervention_name": "Acoltremon 0.003%",
        "control": "Preservative-free artificial tears",
    })

    fields = render_fields(reference, {"protocol": [], "icf": {}, "prs": {}})

    assert fields["studyDesignShort"] == "Prospective randomized parallel-group study"
    assert fields["testArticle(s)"] == "Acoltremon 0.003%"
    assert fields["controlArticle(s)"] == "Preservative-free artificial tears"


def test_protocol_summary_rows_keep_together(tmp_path):
    reference = _source()
    render_documents(ROOT, tmp_path, reference, {"protocol": [], "icf": {}, "prs": {}})

    document = Document(tmp_path / "candidate/protocol.docx")
    summary = next(
        table for table in document.tables
        if table.rows and table.rows[0].cells[0].text.strip() == "Objective"
    )
    variables = next(row for row in summary.rows if row.cells[0].text.strip() == "Variables")

    assert all(
        paragraph.paragraph_format.keep_together is True
        for cell in variables.cells
        for paragraph in cell.paragraphs
    )
    assert all(
        cell._tc.tcPr.find(qn("w:tcMar")).find(qn(side)).get(qn("w:w")) == "40"
        for row in summary.rows
        for cell in row.cells
        for side in ("w:top", "w:bottom")
    )


def test_protocol_contact_table_does_not_invent_round_the_clock_availability(tmp_path):
    reference = _source()
    render_documents(ROOT, tmp_path, reference, {"protocol": [], "icf": {}, "prs": {}})

    document = Document(tmp_path / "candidate/protocol.docx")
    contact = next(
        table for table in document.tables
        if table.rows and table.rows[0].cells[0].text.strip() == "Study Staff"
    )
    headers = [" ".join(cell.text.split()) for cell in contact.rows[0].cells]

    assert headers == ["Study Staff", "Business Phone", "e-mail", "Office Phone"]


def test_assessment_table_splits_plain_language_schedule_into_readable_rows(tmp_path):
    reference = _source()
    reference["procedures"].pop("visit_schedule", None)
    reference["procedures"].pop("visit_schedule_table", None)
    reference["procedures"]["assessments"] = (
        "Historical chart abstraction, baseline prospective visit, Month 1 phone "
        "follow-up, and Month 3 clinic follow-up, including pain score review, device "
        "usage download, and usability questionnaire."
    )

    render_documents(ROOT, tmp_path, reference, {"protocol": [], "icf": {}, "prs": {}})

    document = Document(tmp_path / "candidate/protocol.docx")
    table = next(
        item for item in document.tables
        if item.rows and item.rows[0].cells[0].text.strip() == "Approved visit or assessment"
    )
    rows = [(row.cells[0].text.strip(), row.cells[1].text.strip()) for row in table.rows[1:]]

    assert rows == [
        ("Historical chart abstraction", "Historical record review"),
        ("baseline prospective visit", "Baseline"),
        ("Month 1 phone follow-up", "Month 1"),
        ("Month 3 clinic follow-up", "Month 3"),
        ("pain score review", "Per approved schedule"),
        ("device usage download", "Per approved schedule"),
        ("usability questionnaire", "Per approved schedule"),
    ]


def test_protocol_visit_schedule_has_rows_when_approved_source_has_assessments_only(tmp_path):
    reference = _source()
    reference["procedures"].pop("visit_schedule", None)
    reference["procedures"].pop("visit_schedule_table", None)
    reference["procedures"]["assessments"] = ["Screening", "Month 1", "Month 3"]
    render_documents(ROOT, tmp_path, reference, {"protocol": [], "icf": {}, "prs": {}})

    document = Document(tmp_path / "candidate/protocol.docx")
    schedule = next(table for table in document.tables if table.cell(0, 0).text.strip() == "Visit Number")
    rows = [[cell.text.strip() for cell in row.cells] for row in schedule.rows[1:]]
    assert [row[1] for row in rows] == ["Screening", "Month 1", "Month 3"]
    assert all(not row[0] for row in rows)


def test_assessments_only_do_not_become_synthetic_numbered_visits(tmp_path):
    reference = _source()
    reference["procedures"].pop("visit_schedule", None)
    reference["procedures"].pop("visit_schedule_table", None)
    reference["procedures"]["assessments"] = [
        "Screening, consent, and baseline visit",
        "Sensor wear on Days 1 to 14, Weeks 6 to 8, and Weeks 10 to 12",
        "Telephone contact at Week 3",
        "Clinic visits at Weeks 6 and 12",
        "Record abstraction",
        "Sensor insertion and removal",
        "Sensor data download",
        "Medication review",
        "Adverse-event assessment",
        "Hemoglobin A1c at Week 12",
        "Usability questionnaire",
    ]

    render_documents(ROOT, tmp_path, reference, {"protocol": [], "icf": {}, "prs": {}})

    document = Document(tmp_path / "candidate/protocol.docx")
    schedule = next(
        table for table in document.tables
        if table.rows and table.rows[0].cells[0].text.strip() == "Visit Number"
    )
    rows = [[cell.text.strip() for cell in row.cells] for row in schedule.rows[1:]]

    assert [(row[0], row[1]) for row in rows] == [
        ("", "Screening, consent, and baseline visit"),
        ("", "Telephone contact at Week 3"),
        ("", "Clinic visits at Weeks 6 and 12"),
    ]
    assert all("Record abstraction" not in row for row in rows)


def test_advarra_icf_contact_and_withdrawal_are_source_bound(tmp_path):
    reference = _source()
    model = {
        "protocol": [],
        "prs": {},
        "icf": {
            "icf.contact-information": {
                "paragraphs": [{"text": "Call the study doctor at the toll-free IRB number and ask the Study Subject Adviser."}],
                "lists": [],
            },
            "icf.leaving-study": {
                "paragraphs": [{"text": "Call to schedule study exit procedures."}],
                "lists": [],
            },
        },
    }
    render_documents(ROOT, tmp_path, reference, model)
    visible = _visible(Document(tmp_path / "candidate/icf.docx")).casefold()

    assert "toll-free" not in visible
    assert "study subject adviser" not in visible
    assert "schedule study exit procedures" not in visible
    assert "555-0101" in visible
    assert "no new routine research procedures" in visible


def test_advarra_icf_replaces_unsupported_introduction_and_risk_shell_text(tmp_path):
    reference = _source()
    risk_text = (
        "Taking part may involve inconvenience or discomfort from the study procedures described in this consent form. "
        "There is also a risk that private information could be disclosed, although safeguards will be used to protect it."
    )
    model = {
        "protocol": [],
        "prs": {},
        "icf": {
            "icf.risks": {"paragraphs": [{"text": risk_text}], "lists": []},
        },
    }

    render_documents(ROOT, tmp_path, reference, model)
    visible = _visible(Document(tmp_path / "candidate/icf.docx")).casefold()

    assert "you are being invited to take part in a research study" in visible
    assert "past and present drug use" not in visible
    assert "very serious health consequences" not in visible
    assert "all devices can have side effects" not in visible
    assert "presently unforeseen and unknown risks" not in visible
    assert visible.count(risk_text.casefold()) == 1


def test_protocol_test_article_falls_back_to_approved_intervention_name():
    reference = _source()
    reference["design"].pop("arms", None)

    fields = rendering.render_fields(reference, {"protocol": [], "icf": {}, "prs": {}})

    assert fields["testArticle(s)"] == "Sentinel Patch"


def test_transient_verifier_failure_is_classified_for_retry_without_accepting_qa(tmp_path):
    revision = tmp_path / "revision"
    requests = revision / "hermes/verification-requests"
    responses = revision / "hermes/verification-responses"
    requests.mkdir(parents=True)
    responses.mkdir()
    request = {
        "schema_version": "hermes-verification/v1",
        "request_id": "r.verify.content",
        "task": "clinical_content_verification",
        "revision_id": revision.name,
        "review_set": 1,
        "response_path": "hermes/verification-responses/r.verify.content.json",
        "sections": [],
        "cross_document_checks": [],
        "artifacts": [],
    }
    request["request_sha256"] = verification_request_sha256(request)
    (requests / f"{request['request_id']}.json").write_text(json.dumps(request), encoding="utf-8")
    _write_verification_ledger(revision, request)
    response = {
        "schema_version": RESPONSE_SCHEMA,
        "request_id": request["request_id"],
        "request_sha256": request["request_sha256"],
        "task": request["task"],
        "revision_id": request["revision_id"],
        "producer": {"model_id": "Verifier", "reviewer_id": "content-reviewer"},
        "status": "retryable_error",
        "error": {"type": "api_unavailable", "message": "temporary connection failure"},
    }
    (responses / "r.verify.content.json").write_text(json.dumps(response), encoding="utf-8")

    findings, evidence = validate_verifications(revision)

    assert evidence[request["task"]]
    assert len(findings) == 1
    assert findings[0]["category"] == "reviewer-transient"
    assert findings[0]["target_ids"] == ["verification:content"]
    assert findings[0]["recovery_class"] == "verifier_transient"
    assert findings[0]["action"] == "retry_verifier"
    assert findings[0]["verification_request_id"] == request["request_id"]


def test_content_finding_without_a_section_target_reprompts_the_reviewer(tmp_path):
    revision = tmp_path / "revision"
    requests = revision / "hermes/verification-requests"
    responses = revision / "hermes/verification-responses"
    requests.mkdir(parents=True)
    responses.mkdir()
    request = {
        "schema_version": "hermes-verification/v1",
        "request_id": "r.verify.content",
        "task": "clinical_content_verification",
        "revision_id": revision.name,
        "review_set": 1,
        "response_path": "hermes/verification-responses/r.verify.content.json",
        "sections": [],
        "cross_document_checks": [],
        "artifacts": [],
    }
    request["request_sha256"] = verification_request_sha256(request)
    (requests / f"{request['request_id']}.json").write_text(json.dumps(request), encoding="utf-8")
    _write_verification_ledger(revision, request)
    response = {
        "schema_version": RESPONSE_SCHEMA,
        "request_id": request["request_id"],
        "request_sha256": request["request_sha256"],
        "task": request["task"],
        "revision_id": request["revision_id"],
        "producer": {"model_id": "user-selected-local-model", "reviewer_id": "content-reviewer"},
        "status": "blocked",
        "findings": [{"issue": "A substantive source fact is missing."}],
        "section_assessments": [],
        "cross_document_assessments": [],
    }
    (responses / "r.verify.content.json").write_text(json.dumps(response), encoding="utf-8")

    findings, _evidence = validate_verifications(revision)

    routing = next(item for item in findings if "source fact" in item["issue"])
    assert routing["target_ids"] == ["verification:content"]
    assert routing["recovery_class"] == "verifier_transient"
    assert routing["action"] == "retry_verifier"


def test_content_verification_request_documents_warning_identity_contract(tmp_path):
    (tmp_path / "candidate").mkdir()

    paths = quality.create_verification_requests(
        tmp_path,
        _source(),
        {"status": "passed", "artifacts": []},
    )
    request_path = next(path for path in paths if path.name.endswith("verify.content.json"))
    request = json.loads(request_path.read_text(encoding="utf-8"))
    instructions = request["instructions"]

    assert "Every finding must include a unique finding_id" in instructions
    assert "ASCII letters or digits separated only by single hyphens or underscores" in instructions
    assert "ordinary substantive content omission" in instructions
    assert "failed `procedures` cross-document assessment" in instructions
    assert "notes must cite that finding_id as an exact token" in instructions
    assert "Safety, invention, contradiction, source-integrity" in instructions
    assert "safety_critical" in instructions


def test_verification_request_collision_preserves_canonical_request_and_response(tmp_path):
    candidate = tmp_path / "candidate"
    candidate.mkdir()
    (candidate / "study.xml").write_text("<study>first</study>", encoding="utf-8")
    request_path = quality.create_verification_requests(
        tmp_path, _source(), {"status": "passed", "artifacts": []},
    )[0]
    ledger_path = tmp_path / "request-ledger" / f"{json.loads(request_path.read_text())['request_id']}.json"
    response_path = tmp_path / json.loads(request_path.read_text())["response_path"]
    response_path.parent.mkdir(parents=True, exist_ok=True)
    response_path.write_bytes(b"preserve-response")
    original_request = request_path.read_bytes()
    original_ledger = ledger_path.read_bytes()
    (candidate / "study.xml").write_text("<study>changed</study>", encoding="utf-8")

    with pytest.raises(ValueError, match="Verification request identity collision"):
        quality.create_verification_requests(
            tmp_path, _source(), {"status": "passed", "artifacts": []},
        )

    assert request_path.read_bytes() == original_request
    assert ledger_path.read_bytes() == original_ledger
    assert response_path.read_bytes() == b"preserve-response"


def test_content_verification_request_binds_canonical_artifact_identity(tmp_path):
    candidate = tmp_path / "candidate"
    candidate.mkdir()
    Document().save(candidate / "protocol.docx")
    Document().save(candidate / "icf.docx")
    (candidate / "study.xml").write_bytes(b"study")

    request_path = quality.create_verification_requests(
        tmp_path,
        _source(),
        {"status": "passed", "artifacts": []},
    )[0]
    request = json.loads(request_path.read_text(encoding="utf-8"))

    assert {
        item["path"]: item["artifact"]
        for item in request["artifacts"]
    } == {
        "candidate/icf.docx": "icf",
        "candidate/protocol.docx": "protocol",
        "candidate/study.xml": "study.xml",
    }


def test_content_omission_warning_requires_explicit_non_safety_classification():
    ordinary = {
        "category": "content",
        "check": "substantive",
        "safety_critical": False,
        "contradiction": False,
        "obscures_required_information": False,
        "materially_unusable": False,
    }

    assert quality._governed_content_omission(ordinary, ["icf.procedures"]) is True
    assert quality._governed_content_omission(
        {**ordinary, "safety_critical": True},
        ["icf.procedures"],
    ) is False
    assert quality._governed_content_omission(
        {"category": "content", "check": "substantive"},
        ["icf.procedures"],
    ) is False


def test_governed_content_omission_is_reported_as_a_nonblocking_manual_review_warning(
    tmp_path,
    monkeypatch,
):
    revision = tmp_path / "revision"
    requests = revision / "hermes/verification-requests"
    responses = revision / "hermes/verification-responses"
    requests.mkdir(parents=True)
    responses.mkdir()
    request = {
        "schema_version": "hermes-verification/v1",
        "request_id": "r.verify.content",
        "task": "clinical_content_verification",
        "revision_id": revision.name,
        "response_path": "hermes/verification-responses/content.json",
        "sections": [{"artifact": "icf", "section_id": "icf.leaving-study"}],
        "cross_document_checks": ["procedures"],
        "artifacts": [],
        "checks": list(quality.CONTENT_CHECKS),
        "review_set": 1,
    }
    request["request_sha256"] = verification_request_sha256(request)
    request_path = requests / f"{request['request_id']}.json"
    request_path.write_text(json.dumps(request), encoding="utf-8")
    _write_verification_ledger(revision, request)
    response = {
        "schema_version": RESPONSE_SCHEMA,
        "request_id": request["request_id"],
        "request_sha256": request["request_sha256"],
        "task": request["task"],
        "revision_id": request["revision_id"],
        "producer": {"model_id": "manual-review-model", "reviewer_id": "content-reviewer"},
        "status": "blocked",
        "findings": [{
            "finding_id": "content-001",
            "category": "content",
            "check": "substantive",
            "target_ids": ["icf.leaving-study"],
            "issue": "The ICF omits the optional leaving-study explanation.",
            "recommended_action": "Add a source-grounded leaving-study explanation at manual review.",
            "safety_critical": False,
            "contradiction": False,
            "obscures_required_information": False,
            "materially_unusable": False,
        }],
        "section_assessments": [{
            "artifact": "icf",
            "section_id": "icf.leaving-study",
            "status": "blocked",
            "checks": list(quality.CONTENT_CHECKS),
        }],
        "cross_document_assessments": [{
            "check": "procedures",
            "status": "failed",
            "notes": "The same ordinary omission affects procedures; see content-001.",
        }],
    }
    (responses / "content.json").write_text(json.dumps(response), encoding="utf-8")
    monkeypatch.setattr(quality, "deterministic_content_check", lambda *_args: [])
    monkeypatch.setattr(quality, "_final_verification_scope_findings", lambda *_args: [])

    report = quality.quality_report(
        revision,
        _source(),
        {"status": "passed", "findings": [], "artifacts": []},
        None,
    )

    assert report["status"] == "passed"
    assert report["findings"] == []
    assert report["warnings"] == [{
        "category": "content",
        "field": "clinical_content_verification",
        "check": "substantive",
        "target_ids": ["icf.leaving-study"],
        "verification_request_id": request["request_id"],
        "issue": "The ICF omits the optional leaving-study explanation.",
        "recommended_action": "Add a source-grounded leaving-study explanation at manual review.",
        "safety_critical": False,
        "contradiction": False,
        "obscures_required_information": False,
        "materially_unusable": False,
        "severity": "warning",
        "publication_disposition": "warning",
        "action": "manual_review",
    }]
    assert quality.verification_response_is_complete(revision, request_path) is True
    assert quality.final_exact_artifact_review_findings(
        revision,
        _source(),
        {"status": "passed", "findings": [], "artifacts": []},
        report,
    ) == []
    tampered = copy.deepcopy(report)
    tampered["warnings"][0]["issue"] = "tampered warning"
    assert any(
        finding["issue"] == "Final verification warnings changed after review."
        for finding in quality.final_exact_artifact_review_findings(
            revision,
            _source(),
            {"status": "passed", "findings": [], "artifacts": []},
            tampered,
        )
    )
    forged_response = copy.deepcopy(response)
    forged_response["findings"][0]["finding_id"] = "."
    forged_response["cross_document_assessments"][0]["notes"] = "."
    (responses / "content.json").write_text(json.dumps(forged_response), encoding="utf-8")
    forged_findings, _evidence = validate_verifications(revision)
    assert any(
        finding.get("recovery_class") == "verifier_transient"
        and "cross-document" in finding["issue"]
        for finding in forged_findings
    )

    duplicate_response = copy.deepcopy(response)
    duplicate_response["findings"].append(dict(duplicate_response["findings"][0]))
    (responses / "content.json").write_text(json.dumps(duplicate_response), encoding="utf-8")
    assert quality.verification_response_is_complete(revision, request_path) is False


def test_unreported_blocked_section_assessment_cannot_hide_behind_content_warning(tmp_path):
    revision = tmp_path / "revision"
    requests = revision / "hermes/verification-requests"
    responses = revision / "hermes/verification-responses"
    requests.mkdir(parents=True)
    responses.mkdir()
    request = {
        "schema_version": "hermes-verification/v1",
        "request_id": "r.verify.content",
        "task": "clinical_content_verification",
        "response_path": "hermes/verification-responses/content.json",
        "sections": [
            {"artifact": "icf", "section_id": "icf.leaving-study"},
            {"artifact": "icf", "section_id": "icf.risks"},
        ],
        "cross_document_checks": [],
        "artifacts": [],
    }
    request["request_sha256"] = verification_request_sha256(request)
    request_path = requests / f"{request['request_id']}.json"
    request_path.write_text(json.dumps(request), encoding="utf-8")
    (responses / "content.json").write_text(json.dumps({
        "schema_version": RESPONSE_SCHEMA,
        "request_id": request["request_id"],
        "request_sha256": request["request_sha256"],
        "task": request["task"],
        "producer": {"model_id": "manual-review-model", "reviewer_id": "content-reviewer"},
        "status": "blocked",
        "findings": [{
            "finding_id": "content-001",
            "category": "content",
            "check": "substantive",
            "target_ids": ["icf.leaving-study"],
            "issue": "The leaving-study explanation is incomplete.",
            "recommended_action": "Add the source-grounded explanation at manual review.",
        }],
        "section_assessments": [
            {
                "artifact": "icf",
                "section_id": "icf.leaving-study",
                "status": "blocked",
                "checks": list(quality.CONTENT_CHECKS),
                "notes": "See content-001.",
            },
            {
                "artifact": "icf",
                "section_id": "icf.risks",
                "status": "blocked",
                "checks": list(quality.CONTENT_CHECKS),
                "notes": "A separate safety problem was not reported as a finding.",
            },
        ],
        "cross_document_assessments": [],
    }), encoding="utf-8")

    findings, _evidence = validate_verifications(revision)

    assert any(
        finding.get("recovery_class") == "verifier_transient"
        and "non-passing assessment" in finding["issue"]
        for finding in findings
    )
    assert quality.verification_response_is_complete(revision, request_path) is False


def test_safety_cross_assessment_cannot_be_covered_by_ordinary_content_warning(tmp_path):
    revision = tmp_path / "revision"
    requests = revision / "hermes/verification-requests"
    responses = revision / "hermes/verification-responses"
    requests.mkdir(parents=True)
    responses.mkdir()
    request = {
        "schema_version": "hermes-verification/v1",
        "request_id": "r.verify.content",
        "task": "clinical_content_verification",
        "revision_id": revision.name,
        "review_set": 1,
        "response_path": "hermes/verification-responses/content.json",
        "sections": [{"artifact": "icf", "section_id": "icf.leaving-study"}],
        "cross_document_checks": ["risks_benefits"],
        "artifacts": [],
    }
    request["request_sha256"] = verification_request_sha256(request)
    request_path = requests / f"{request['request_id']}.json"
    request_path.write_text(json.dumps(request), encoding="utf-8")
    _write_verification_ledger(revision, request)
    (responses / "content.json").write_text(json.dumps({
        "schema_version": RESPONSE_SCHEMA,
        "request_id": request["request_id"],
        "request_sha256": request["request_sha256"],
        "task": request["task"],
        "revision_id": request["revision_id"],
        "producer": {"model_id": "manual-review-model", "reviewer_id": "content-reviewer"},
        "status": "blocked",
        "findings": [{
            "finding_id": "content-001",
            "category": "content",
            "check": "substantive",
            "target_ids": ["icf.leaving-study"],
            "issue": "The leaving-study explanation is incomplete.",
            "recommended_action": "Add the source-grounded explanation at manual review.",
        }],
        "section_assessments": [{
            "artifact": "icf",
            "section_id": "icf.leaving-study",
            "status": "blocked",
            "checks": list(quality.CONTENT_CHECKS),
            "notes": "See content-001.",
        }],
        "cross_document_assessments": [{
            "check": "risks_benefits",
            "status": "failed",
            "notes": "Safety mismatch; see content-001.",
        }],
    }), encoding="utf-8")

    findings, _evidence = validate_verifications(revision)

    assert any(
        finding.get("recovery_class") == "verifier_transient"
        and "cross-document" in finding["issue"]
        for finding in findings
    )
    assert quality.verification_response_is_complete(revision, request_path) is False


@pytest.mark.parametrize(
    ("source_category", "source_check", "target"),
    [
        ("content", "substantive", "icf.risks"),
        ("content", "source_supported", "icf.leaving-study"),
        ("content", "no_invention", "icf.leaving-study"),
        ("content", "no_internal_language", "icf.leaving-study"),
        ("content", "cross_document_consistent", "icf.leaving-study"),
        ("safety", "substantive", "icf.leaving-study"),
        ("integrity", "substantive", "icf.leaving-study"),
    ],
)
def test_safety_and_policy_correction_content_findings_remain_blocking(
    tmp_path,
    source_category,
    source_check,
    target,
):
    revision = tmp_path / "revision"
    requests = revision / "hermes/verification-requests"
    responses = revision / "hermes/verification-responses"
    requests.mkdir(parents=True)
    responses.mkdir()
    request = {
        "schema_version": "hermes-verification/v1",
        "request_id": "r.verify.content",
        "task": "clinical_content_verification",
        "revision_id": revision.name,
        "review_set": 1,
        "response_path": "hermes/verification-responses/content.json",
        "sections": [{"artifact": "icf", "section_id": target}],
        "cross_document_checks": [],
        "artifacts": [],
    }
    request["request_sha256"] = verification_request_sha256(request)
    request_path = requests / f"{request['request_id']}.json"
    request_path.write_text(json.dumps(request), encoding="utf-8")
    _write_verification_ledger(revision, request)
    (responses / "content.json").write_text(json.dumps({
        "schema_version": RESPONSE_SCHEMA,
        "request_id": request["request_id"],
        "request_sha256": request["request_sha256"],
        "task": request["task"],
        "revision_id": request["revision_id"],
        "producer": {"model_id": "manual-review-model", "reviewer_id": "content-reviewer"},
        "status": "blocked",
        "findings": [{
            "finding_id": f"{source_category}-{source_check}",
            "category": source_category,
            "check": source_check,
            "target_ids": [target],
            "issue": "The ICF omits a known study risk.",
            "recommended_action": "Correct the safety content before publication.",
        }],
        "section_assessments": [{
            "artifact": "icf",
            "section_id": target,
            "status": "blocked",
            "checks": list(quality.CONTENT_CHECKS),
        }],
        "cross_document_assessments": [],
    }), encoding="utf-8")

    findings, _evidence = validate_verifications(revision)

    assert len(findings) == 1
    assert findings[0]["category"] == source_category
    assert findings[0]["check"] == source_check
    assert findings[0]["recovery_class"] == "drafting_defect"
    assert findings[0]["action"] == "retry_drafting_target"
    assert quality.verification_response_is_complete(revision, request_path) is False


def test_rehashed_verification_request_is_rejected_by_terminal_and_validation_paths(tmp_path):
    (tmp_path / "candidate").mkdir()
    request_path = quality.create_verification_requests(
        tmp_path,
        _source(),
        {"status": "passed", "artifacts": []},
    )[0]
    request = json.loads(request_path.read_text(encoding="utf-8"))
    request["instructions"] = "Altered after canonical ledger creation."
    request["request_sha256"] = verification_request_sha256(request)
    request_path.write_text(json.dumps(request), encoding="utf-8")
    response_path = tmp_path / request["response_path"]
    response_path.parent.mkdir(parents=True, exist_ok=True)
    response_path.write_text(json.dumps({
        "schema_version": RESPONSE_SCHEMA,
        "request_id": request["request_id"],
        "request_sha256": request["request_sha256"],
        "task": request["task"],
        "revision_id": request["revision_id"],
        "producer": {"model_id": "manual-review-model", "reviewer_id": "content-reviewer"},
        "status": "blocked",
        "findings": [{"issue": "Altered request response."}],
    }), encoding="utf-8")

    findings, _evidence = validate_verifications(tmp_path, request_paths=[request_path])

    assert quality.verification_response_is_terminal(tmp_path, request_path) is False
    assert quality.verification_response_is_complete(tmp_path, request_path) is False
    assert [item["field"] for item in findings] == ["canonical_request_ledger"]
    assert findings[0]["verification_request_id"] == request["request_id"]


def test_verification_request_integrity_finding_retains_request_identity(tmp_path):
    revision = tmp_path / "revision"
    requests = revision / "hermes/verification-requests"
    requests.mkdir(parents=True)
    request = {
        "schema_version": "hermes-verification/v1",
        "request_id": "r-test.review-1.verify.content",
        "task": "clinical_content_verification",
        "revision_id": "revision",
        "review_set": 1,
        "response_path": "hermes/verification-responses/content.json",
        "sections": [],
        "cross_document_checks": [],
        "artifacts": [],
    }
    request["request_sha256"] = verification_request_sha256(request)
    request_path = requests / f"{request['request_id']}.json"
    request_path.write_text(json.dumps(request), encoding="utf-8")
    _write_verification_ledger(revision, request)
    request["response_path"] = "hermes/verification-responses/tampered.json"
    request_path.write_text(json.dumps(request), encoding="utf-8")

    findings, _evidence = validate_verifications(revision)

    finding = next(
        item for item in findings
        if item.get("code") == "verification_request_integrity_failure"
    )
    assert finding["verification_request_id"] == request["request_id"]
    assert finding["recovery_class"] == "document_structure_defect"
    assert finding["action"] == "preserve_and_stop"


def test_stale_visual_structural_finding_retains_exact_artifact_binding(tmp_path):
    revision = tmp_path / "revision"
    requests = revision / "hermes/verification-requests"
    responses = revision / "hermes/verification-responses"
    requests.mkdir(parents=True)
    responses.mkdir()
    request = {
        "schema_version": "hermes-verification/v1",
        "request_id": "revision.review-1.verify.visual.icf",
        "task": "rendered_page_visual_verification",
        "revision_id": revision.name,
        "review_set": 1,
        "response_path": "hermes/verification-responses/visual.icf.json",
        "artifacts": [{
            "artifact": "icf",
            "docx": "candidate/icf.docx",
            "docx_sha256": "0" * 64,
            "pdf": "rendered/icf.pdf",
            "pdf_sha256": "0" * 64,
            "pages": [],
        }],
    }
    request["request_sha256"] = verification_request_sha256(request)
    (requests / f"{request['request_id']}.json").write_text(json.dumps(request), encoding="utf-8")
    _write_verification_ledger(revision, request)
    (responses / "visual.icf.json").write_text(json.dumps({
        "schema_version": RESPONSE_SCHEMA,
        "request_id": request["request_id"],
        "request_sha256": request["request_sha256"],
        "task": request["task"],
        "revision_id": request["revision_id"],
        "producer": {"model_id": "manual-review-model", "reviewer_id": "visual-reviewer"},
        "status": "passed",
        "findings": [],
        "page_assessments": [],
    }), encoding="utf-8")

    findings, _evidence = validate_verifications(revision)

    stale = [item for item in findings if "is stale" in item.get("issue", "")]
    assert stale
    assert all(item["artifact"] == "icf" for item in stale)
    assert {item["artifact_path"] for item in stale} == {
        "candidate/icf.docx", "rendered/icf.pdf",
    }
    assert all(item["verification_request_id"] == request["request_id"] for item in stale)
    assert all(item["recovery_class"] == "document_structure_defect" for item in stale)


def test_stale_content_structural_finding_retains_exact_artifact_binding(tmp_path):
    revision = tmp_path / "revision"
    requests = revision / "hermes/verification-requests"
    responses = revision / "hermes/verification-responses"
    requests.mkdir(parents=True)
    responses.mkdir()
    request = {
        "schema_version": "hermes-verification/v1",
        "request_id": "revision.review-1.verify.content",
        "task": "clinical_content_verification",
        "revision_id": revision.name,
        "review_set": 1,
        "response_path": "hermes/verification-responses/content.json",
        "sections": [],
        "cross_document_checks": [],
        "artifacts": [{
            "artifact": "protocol",
            "path": "candidate/protocol.docx",
            "sha256": "0" * 64,
        }],
    }
    request["request_sha256"] = verification_request_sha256(request)
    (requests / f"{request['request_id']}.json").write_text(json.dumps(request), encoding="utf-8")
    _write_verification_ledger(revision, request)
    (responses / "content.json").write_text(json.dumps({
        "schema_version": RESPONSE_SCHEMA,
        "request_id": request["request_id"],
        "request_sha256": request["request_sha256"],
        "task": request["task"],
        "revision_id": request["revision_id"],
        "producer": {"model_id": "manual-review-model", "reviewer_id": "content-reviewer"},
        "status": "passed",
        "findings": [],
        "section_assessments": [],
        "cross_document_assessments": [],
    }), encoding="utf-8")

    findings, _evidence = validate_verifications(revision)

    stale = [item for item in findings if "is stale" in item.get("issue", "")]
    assert stale
    assert all(item["artifact"] == "protocol" for item in stale)
    assert {item["artifact_path"] for item in stale} == {"candidate/protocol.docx"}
    assert all(item["verification_request_id"] == request["request_id"] for item in stale)
    assert all(item["recovery_class"] == "document_structure_defect" for item in stale)


@pytest.mark.parametrize("failure", ["malformed", "stale"])
def test_verification_integrity_failure_remains_blocking(tmp_path, failure):
    revision = tmp_path / "revision"
    requests = revision / "hermes/verification-requests"
    responses = revision / "hermes/verification-responses"
    candidate = revision / "candidate/icf.docx"
    requests.mkdir(parents=True)
    responses.mkdir()
    candidate.parent.mkdir()
    candidate.write_bytes(b"current candidate")
    request = {
        "schema_version": "hermes-verification/v1",
        "request_id": "r.verify.content",
        "task": "clinical_content_verification",
        "response_path": "hermes/verification-responses/content.json",
        "sections": [],
        "cross_document_checks": [],
        "artifacts": ([{
            "path": "candidate/icf.docx",
            "sha256": "0" * 64,
        }] if failure == "stale" else []),
    }
    request["request_sha256"] = verification_request_sha256(request)
    request_path = requests / f"{request['request_id']}.json"
    request_path.write_text(json.dumps(request), encoding="utf-8")
    response = {
        "schema_version": RESPONSE_SCHEMA,
        "request_id": request["request_id"],
        "request_sha256": request["request_sha256"],
        "task": request["task"],
        "producer": {"model_id": "manual-review-model", "reviewer_id": "content-reviewer"},
        "status": "passed",
        "section_assessments": [],
        "cross_document_assessments": [],
    }
    if failure == "malformed":
        response["producer"] = {}
    (responses / "content.json").write_text(json.dumps(response), encoding="utf-8")

    findings, _evidence = validate_verifications(revision)

    assert findings
    assert all(item.get("publication_disposition") != "warning" for item in findings)
    assert quality.verification_response_is_complete(revision, request_path) is False


@pytest.mark.parametrize("category", ["rendering", "delivery"])
def test_rendering_and_delivery_findings_remain_blocking(tmp_path, monkeypatch, category):
    finding = {"category": category, "field": category, "issue": f"{category} failed"}
    monkeypatch.setattr(quality, "deterministic_content_check", lambda *_args: [finding])
    monkeypatch.setattr(quality, "validate_verifications", lambda *_args, **_kwargs: ([], {}))
    monkeypatch.setattr(quality, "_final_verification_scope_findings", lambda *_args: [])

    report = quality.quality_report(
        tmp_path,
        _source(),
        {"status": "passed", "findings": [], "artifacts": []},
        None,
    )

    assert report["status"] == "blocked"
    assert report["findings"] == [finding]
    assert report["warnings"] == []


def test_non_verifier_finding_cannot_self_label_as_publication_warning(tmp_path, monkeypatch):
    forged = {
        "category": "integrity",
        "field": "source",
        "issue": "Source integrity failed.",
        "publication_disposition": "warning",
    }
    monkeypatch.setattr(quality, "deterministic_content_check", lambda *_args: [forged])
    monkeypatch.setattr(quality, "validate_verifications", lambda *_args, **_kwargs: ([], {}))
    monkeypatch.setattr(quality, "_final_verification_scope_findings", lambda *_args: [])

    report = quality.quality_report(
        tmp_path,
        _source(),
        {"status": "passed", "findings": [], "artifacts": []},
        None,
    )

    assert report["status"] == "blocked"
    assert report["findings"] == [forged]
    assert report["warnings"] == []


def test_sterling_background_finding_schedules_background_redraft(tmp_path):
    run_dir = tmp_path / "run"
    revision = run_dir / "revisions/r-test"
    reference_path = run_dir / "reference/study.reference.json"
    reference_path.parent.mkdir(parents=True)
    revision.mkdir(parents=True)
    reference = _source()
    reference["meta"]["icf_template"] = "Sterling"
    reference_path.write_text(json.dumps({"generation": {}}), encoding="utf-8")
    finding = {
        "category": "content",
        "field": "icf.background",
        "artifact": "icf",
        "check": "participant_facing_language",
        "target_ids": ["icf.background"],
        "recovery_class": "drafting_defect",
        "action": "retry_drafting_target",
        "issue": "BACKGROUND requires participant-facing language.",
    }

    result = workflow._quality_retry(
        run_dir, reference_path, {"generation": {}}, reference, revision, {}, [finding], "quality",
        require_promoted_runtime=False,
    )

    request_path = next(
        revision.glob("hermes/requests/*icf-narrative*.json")
    )
    request = json.loads(request_path.read_text(encoding="utf-8"))
    assert result["status"] == "awaiting_hermes"
    assert "icf.background" in {
        section["section_id"] for section in request["section_contracts"]
    }


def test_quality_retry_fails_closed_for_nondraftable_section_target(tmp_path, monkeypatch):
    run_dir = tmp_path / "run"
    revision = run_dir / "revisions/r-test"
    reference_path = run_dir / "reference/study.reference.json"
    reference_path.parent.mkdir(parents=True)
    revision.mkdir(parents=True)
    reference_path.write_text(json.dumps({"generation": {}}), encoding="utf-8")
    monkeypatch.setattr(
        workflow,
        "generate",
        lambda *_args, **_kwargs: {"status": "awaiting_hermes", "stage": "unexpected"},
    )
    finding = {
        "category": "content",
        "field": "icf.nonexistent",
        "artifact": "icf",
        "target_ids": ["icf.nonexistent"],
        "recovery_class": "drafting_defect",
        "action": "retry_drafting_target",
        "issue": "A nonexistent section cannot be repaired.",
    }

    result = workflow._quality_retry(
        run_dir, reference_path, {"generation": {}}, _source(), revision, {}, [finding], "quality",
        require_promoted_runtime=False,
    )

    assert result["status"] == "blocked"
    assert result["stage"] == "internal_recovery_routing"
    assert result["findings"][0]["target_id"] == "icf.nonexistent"
    assert result["findings"][0]["artifact"] == "icf"
    assert result["findings"][0]["original_finding"] == finding
    assert not (revision / "attempts").exists()
    assert not (revision / "gate-attempt-journal.json").exists()


def _quality_retry_fixture(
    tmp_path,
    findings,
    *,
    reference=None,
    request_records=(),
    request_ledgers=None,
    candidate_build=None,
    working=None,
):
    run_dir = tmp_path / "run"
    revision = run_dir / "revisions/r-test"
    reference_path = run_dir / "reference/study.reference.json"
    reference_path.parent.mkdir(parents=True)
    revision.mkdir(parents=True)
    state = working or {"generation": {"review_set": 1}}
    reference_path.write_text(json.dumps(state), encoding="utf-8")
    for _filename, request in request_records:
        request_was_self_valid = (
            request.get("request_sha256") == verification_request_sha256(request)
        )
        for artifact in request.get("artifacts", []):
            if not isinstance(artifact, dict):
                continue
            if request.get("task") == "clinical_content_verification":
                path = revision / str(artifact.get("path") or "")
                path.parent.mkdir(parents=True, exist_ok=True)
                if path.suffix.casefold() == ".docx":
                    Document().save(path)
                    artifact["content_sha256"] = quality._content_sha256(path)
                else:
                    path.write_text("<study />", encoding="utf-8")
                artifact["sha256"] = quality.sha256_file(path)
            else:
                for key in ("docx", "pdf"):
                    path = revision / str(artifact.get(key) or "")
                    path.parent.mkdir(parents=True, exist_ok=True)
                    path.write_bytes(key.encode("utf-8"))
                    artifact[f"{key}_sha256"] = quality.sha256_file(path)
                for page in artifact.get("pages", []):
                    if isinstance(page, dict):
                        path = revision / str(page.get("path") or "")
                        path.parent.mkdir(parents=True, exist_ok=True)
                        path.write_bytes(b"png")
                        page["sha256"] = quality.sha256_file(path)
        if request_was_self_valid:
            request["request_sha256"] = verification_request_sha256(request)
    if candidate_build is None:
        visual_artifacts = [
            dict(item)
            for _filename, request in request_records
            if request.get("task") == "rendered_page_visual_verification"
            for item in request.get("artifacts", [])
            if isinstance(item, dict)
        ]
        if visual_artifacts:
            candidate_build = {"render_report": {"artifacts": visual_artifacts}}
    if candidate_build is not None:
        (revision / "candidate-build.json").write_text(
            json.dumps(candidate_build), encoding="utf-8",
        )
    request_root = revision / "hermes/verification-requests"
    seen_request_ids: set[str] = set()
    for filename, request in request_records:
        request_root.mkdir(parents=True, exist_ok=True)
        request_hash_was_valid = verification_request_sha256(request) == request.get("request_sha256")
        if request.get("task") == "clinical_content_verification":
            for artifact in request.get("artifacts", []):
                if not isinstance(artifact, dict) or artifact.get("sha256"):
                    continue
                candidate_path = revision / str(artifact.get("path") or "")
                candidate_path.parent.mkdir(parents=True, exist_ok=True)
                candidate_path.write_bytes(str(artifact.get("artifact") or "artifact").encode())
                artifact["sha256"] = hashlib.sha256(candidate_path.read_bytes()).hexdigest()
            if request_hash_was_valid:
                request["request_sha256"] = verification_request_sha256(request)
        request_id = str(request.get("request_id") or "")
        request_path = (
            request_root / f"{request_id}.json"
            if request_id and request_id not in seen_request_ids
            else request_root / filename
        )
        seen_request_ids.add(request_id)
        request_path.write_text(json.dumps(request), encoding="utf-8")
        ledger = (
            request_ledgers.get(str(request.get("request_id") or ""))
            if isinstance(request_ledgers, dict)
            else quality.verification_request_ledger_record(request_path, request)
        )
        if ledger is not None and request.get("request_id"):
            ledger_path = revision / "request-ledger" / f"{request['request_id']}.json"
            ledger_path.parent.mkdir(parents=True, exist_ok=True)
            ledger_path.write_text(json.dumps(ledger), encoding="utf-8")
        response_finding = next((
            dict(item) for item in findings
            if item.get("verification_request_id") == request.get("request_id")
        ), None)
        if response_finding is not None:
            response_finding.setdefault("finding_id", "fixture-finding")
            response_path = revision / str(request.get("response_path") or "")
            response_path.parent.mkdir(parents=True, exist_ok=True)
            response_path.write_text(json.dumps({
                "schema_version": RESPONSE_SCHEMA,
                "request_id": request.get("request_id"),
                "request_sha256": request.get("request_sha256"),
                "task": request.get("task"),
                "revision_id": request.get("revision_id"),
                "producer": {"model_id": "fixture-model", "reviewer_id": "fixture-reviewer"},
                "status": "blocked",
                "findings": [response_finding],
            }), encoding="utf-8")
    result = workflow._quality_retry(
        run_dir,
        reference_path,
        state,
        reference or _source(),
        revision,
        {},
        findings,
        "quality",
        require_promoted_runtime=False,
    )
    return result, run_dir, revision, reference_path, state


def _prearchive_routing_fixture(tmp_path, finding, *, reference=None, companions=()):
    result, run_dir, revision, _reference_path, working = _quality_retry_fixture(
        tmp_path,
        [*companions, finding],
        reference=reference,
    )

    assert result["status"] == "blocked"
    assert working["generation"]["review_set"] == 1
    assert not (revision / "attempts").exists()
    assert not (revision / "gate-attempt-journal.json").exists()
    assert not list(revision.rglob("attempt-manifest.json"))
    assert not (revision / "rendered").exists()
    assert not (revision / "hermes/requests").exists()
    assert not (run_dir / "reference/repair-report.md").exists()
    return result, revision


def _verification_request_payload(
    request_id,
    task,
    *,
    revision_id="r-test",
    review_set=1,
    artifact=None,
):
    request = {
        "schema_version": "hermes-verification-request/v2",
        "request_id": request_id,
        "task": task,
        "revision_id": revision_id,
        "review_set": review_set,
        "response_path": f"hermes/verification-responses/{request_id}.json",
    }
    if task == "clinical_content_verification":
        request["approved_source"] = _source()
        request["artifacts"] = [
            {
                "artifact": name if name.endswith(".xml") else Path(name).stem,
                "path": f"candidate/{name}",
            }
            for name in sorted(("protocol.docx", "icf.docx", "study.xml"))
        ]
        request["sections"] = quality.content_review_sections(_source())
        request["checks"] = list(quality.CONTENT_CHECKS)
        request["cross_document_checks"] = list(quality.CROSS_DOCUMENT_CHECKS)
    elif artifact is not None:
        request["artifacts"] = [{
            "artifact": artifact,
            "docx": f"candidate/{artifact}.docx",
            "docx_sha256": "a" * 64,
            "pdf": f"rendered/{artifact}.pdf",
            "pdf_sha256": "b" * 64,
            "pages": [],
        }]
        request["checks"] = list(VISUAL_CHECKS)
    request["request_sha256"] = verification_request_sha256(request)
    return request


def _write_verification_ledger(revision, request):
    request_paths = [
        path for path in (revision / "hermes/verification-requests").glob("*.json")
        if json.loads(path.read_text(encoding="utf-8")).get("request_id") == request["request_id"]
    ]
    assert len(request_paths) == 1
    ledger_path = revision / "request-ledger" / f"{request['request_id']}.json"
    ledger_path.parent.mkdir(parents=True, exist_ok=True)
    ledger_path.write_text(json.dumps(
        quality.verification_request_ledger_record(request_paths[0], request)
    ), encoding="utf-8")


def _write_verification_request(revision, request_id, task, *, revision_id="r-test", artifact=None, filename=None):
    request_root = revision / "hermes/verification-requests"
    request_root.mkdir(parents=True, exist_ok=True)
    request = _verification_request_payload(
        request_id,
        task,
        revision_id=revision_id,
        artifact=artifact,
    )
    path = request_root / (filename or f"{request_id}.json")
    path.write_text(json.dumps(request), encoding="utf-8")
    _write_verification_ledger(revision, request)
    return request


@pytest.mark.parametrize(
    "finding",
    (
        {
            "category": "visual", "field": "rendered_page_visual_verification",
            "artifact": "protocol", "check": "orphan_heading", "element": "5. INTRODUCTION",
            "target_ids": ["layout:protocol"], "recovery_class": "visual_defect",
            "action": "targeted_layout_repair", "issue": "Missing visual request identity.",
        },
        {
            "category": "visual", "field": "rendered_page_visual_verification",
            "artifact": "protocol", "check": "orphan_heading", "element": "5. INTRODUCTION",
            "target_ids": ["layout:protocol"], "verification_request_id": "   ",
            "recovery_class": "visual_defect", "action": "targeted_layout_repair",
            "issue": "Blank visual request identity.",
        },
        {
            "category": "content", "field": "clinical_content_verification",
            "artifact": "icf", "target_ids": ["icf.procedures"],
            "recovery_class": "drafting_defect", "action": "retry_drafting_target",
            "issue": "Missing content request identity.",
        },
        {
            "category": "reviewer-transient", "field": "clinical_content_verification",
            "target_ids": ["verification:content"], "recovery_class": "verifier_transient",
            "action": "retry_verifier", "issue": "Missing transient request identity.",
        },
        {
            "category": "reviewer-transient", "field": "clinical_content_verification",
            "target_ids": ["verification:content"], "verification_request_id": 7,
            "recovery_class": "verifier_transient", "action": "retry_verifier",
            "issue": "Numeric transient request identity.",
        },
        {
            "category": "reviewer-transient", "field": "clinical_content_verification",
            "target_ids": ["verification:content"], "verification_request_id": [],
            "recovery_class": "verifier_transient", "action": "retry_verifier",
            "issue": "List transient request identity.",
        },
    ),
    ids=(
        "visual-missing", "visual-blank", "content-missing", "transient-missing",
        "transient-numeric", "transient-list",
    ),
)
def test_verifier_originated_finding_requires_nonempty_request_id_before_archival(tmp_path, finding):
    result, _revision = _prearchive_routing_fixture(tmp_path, finding)

    assert result["stage"] == "internal_recovery_routing"
    assert result["findings"][0]["routing_failure_code"] == "missing_verification_request_id"
    assert result["findings"][0]["original_finding"] == finding


def test_verification_request_from_another_revision_fails_before_archival(tmp_path):
    finding = {
        "category": "visual", "field": "rendered_page_visual_verification",
        "artifact": "protocol", "check": "orphan_heading", "element": "5. INTRODUCTION",
        "target_ids": ["layout:protocol"], "verification_request_id": "shared-request",
        "recovery_class": "visual_defect", "action": "targeted_layout_repair",
        "issue": "The request belongs to another revision.",
    }
    request = _verification_request_payload(
        "shared-request",
        "rendered_page_visual_verification",
        revision_id="r-other",
        artifact="protocol",
    )

    result, _run_dir, revision, _reference_path, _working = _quality_retry_fixture(
        tmp_path,
        [finding],
        request_records=(("shared-request.json", request),),
    )

    assert result["status"] == "blocked"
    assert result["findings"][0]["routing_failure_code"] == "verification_request_wrong_revision"
    assert result["findings"][0]["original_finding"] == finding
    assert not (revision / "attempts").exists()
    assert not (revision / "gate-attempt-journal.json").exists()


def test_whitespace_padded_request_id_is_not_normalized_to_a_match(tmp_path):
    request = _verification_request_payload(
        "verification-id",
        "clinical_content_verification",
    )
    finding = {
        "category": "reviewer-transient", "field": "clinical_content_verification",
        "target_ids": ["verification:content"],
        "verification_request_id": " verification-id ",
        "recovery_class": "verifier_transient", "action": "retry_verifier",
        "issue": "Whitespace must not normalize an authenticated identity.",
    }

    result, _run_dir, _revision, _reference_path, _working = _quality_retry_fixture(
        tmp_path,
        [finding],
        request_records=(("content.json", request),),
    )

    assert result["status"] == "blocked"
    assert result["findings"][0]["routing_failure_code"] == "verification_request_invalid_inventory"
    assert result["findings"][0]["original_finding"] == finding


def test_stale_review_set_request_id_fails_before_archival(tmp_path):
    request = _verification_request_payload(
        "stale-content",
        "clinical_content_verification",
        review_set=1,
    )
    finding = {
        "category": "reviewer-transient", "field": "clinical_content_verification",
        "target_ids": ["verification:content"], "verification_request_id": "stale-content",
        "recovery_class": "verifier_transient", "action": "retry_verifier",
        "issue": "The request belongs to an earlier review set.",
    }
    working = {"generation": {"review_set": 2}}

    result, run_dir, revision, _reference_path, state = _quality_retry_fixture(
        tmp_path,
        [finding],
        request_records=(("content.json", request),),
        working=working,
    )

    assert result["status"] == "blocked"
    assert result["findings"][0]["routing_failure_code"] == "verification_request_stale_review_set"
    assert result["findings"][0]["original_finding"] == finding
    assert state["generation"]["review_set"] == 2
    assert not (revision / "attempts").exists()
    assert not (revision / "gate-attempt-journal.json").exists()
    assert not (run_dir / "reference/repair-report.md").exists()


def test_duplicate_verification_request_id_fails_before_archival(tmp_path):
    finding = {
        "category": "reviewer-transient", "field": "clinical_content_verification",
        "target_ids": ["verification:content"], "verification_request_id": "duplicate-request",
        "recovery_class": "verifier_transient", "action": "retry_verifier",
        "issue": "The request identity is duplicated.",
    }
    request = _verification_request_payload(
        "duplicate-request",
        "clinical_content_verification",
    )

    result, _run_dir, revision, _reference_path, _working = _quality_retry_fixture(
        tmp_path,
        [finding],
        request_records=(("one.json", request), ("two.json", request)),
    )

    assert result["status"] == "blocked"
    assert result["findings"][0]["routing_failure_code"] == "duplicate_verification_request_id"
    assert result["findings"][0]["original_finding"] == finding
    assert not (revision / "attempts").exists()
    assert not (revision / "gate-attempt-journal.json").exists()


def test_verification_binding_failure_precedes_pending_attempt_finalization(tmp_path, monkeypatch):
    run_dir = tmp_path / "run"
    revision = run_dir / "revisions/r-test"
    reference_path = run_dir / "reference/study.reference.json"
    reference_path.parent.mkdir(parents=True)
    revision.mkdir(parents=True)
    working = {
        "generation": {
            "review_set": 1,
            "pending_recovery_attempts": ["attempts/prior"],
            "gate_attempts": [{"path": "attempts/prior"}],
        }
    }
    reference_path.write_text(json.dumps(working), encoding="utf-8")
    monkeypatch.setattr(
        workflow,
        "_complete_pending_recovery_attempts",
        lambda *_args, **_kwargs: pytest.fail("binding failure finalized a pending attempt"),
    )
    finding = {
        "category": "visual", "field": "rendered_page_visual_verification",
        "artifact": "protocol", "check": "orphan_heading", "element": "5. INTRODUCTION",
        "target_ids": ["layout:protocol"], "recovery_class": "visual_defect",
        "action": "targeted_layout_repair", "issue": "Missing visual request identity.",
    }

    result = workflow._quality_retry(
        run_dir, reference_path, working, _source(), revision, {}, [finding], "quality",
        require_promoted_runtime=False,
    )

    assert result["status"] == "blocked"
    assert result["stage"] == "internal_recovery_routing"
    assert working["generation"]["review_set"] == 1
    assert working["generation"]["pending_recovery_attempts"] == ["attempts/prior"]
    assert not (revision / "attempts").exists()
    assert not (revision / "gate-attempt-journal.json").exists()


@pytest.mark.parametrize("include_artifact", (True, False), ids=("declared", "from-layout-target"))
def test_visual_request_must_bind_the_findings_exact_artifact(tmp_path, include_artifact):
    request = _verification_request_payload(
        "r-test.review-1.verify.visual.icf",
        "rendered_page_visual_verification",
        artifact="icf",
    )
    finding = {
        "category": "visual", "field": "rendered_page_visual_verification",
        "artifact": "protocol", "check": "orphan_heading", "element": "5. INTRODUCTION",
        "target_ids": ["layout:protocol"], "verification_request_id": request["request_id"],
        "recovery_class": "visual_defect", "action": "targeted_layout_repair",
        "issue": "A protocol finding cannot bind to the ICF visual request.",
    }
    if not include_artifact:
        finding.pop("artifact")

    result, run_dir, revision, _reference_path, _working = _quality_retry_fixture(
        tmp_path,
        [finding],
        request_records=(("visual.json", request),),
    )

    assert result["status"] == "blocked"
    assert result["findings"][0]["routing_failure_code"] == "verification_request_wrong_artifact"
    assert result["findings"][0]["original_finding"] == finding
    assert not (revision / "attempts").exists()
    assert not (revision / "gate-attempt-journal.json").exists()
    assert not (run_dir / "reference/repair-report.md").exists()


def test_visual_request_artifact_binding_does_not_normalize_suffixes(tmp_path):
    request = _verification_request_payload(
        "r-test.review-1.verify.visual.protocol",
        "rendered_page_visual_verification",
        artifact="protocol",
    )
    finding = {
        "category": "visual", "field": "rendered_page_visual_verification",
        "artifact": "protocol.docx", "check": "orphan_heading", "element": "5. INTRODUCTION",
        "target_ids": ["layout:protocol"], "verification_request_id": request["request_id"],
        "recovery_class": "visual_defect", "action": "targeted_layout_repair",
        "issue": "Raw artifact identities must compare exactly.",
    }

    result, _run_dir, _revision, _reference_path, _working = _quality_retry_fixture(
        tmp_path,
        [finding],
        request_records=(("visual.json", request),),
    )

    assert result["status"] == "blocked"
    assert result["findings"][0]["routing_failure_code"] == "verification_request_wrong_artifact"
    assert result["findings"][0]["original_finding"] == finding


def test_terminal_content_route_must_bind_the_findings_exact_artifact(tmp_path):
    request = _verification_request_payload(
        "r-test.review-1.verify.content",
        "clinical_content_verification",
        artifact="icf",
    )
    finding = {
        "category": "verification",
        "field": "clinical_content_verification",
        "artifact": "unknown",
        "target_ids": ["verification:content"],
        "verification_request_id": request["request_id"],
        "recovery_class": "document_structure_defect",
        "action": "preserve_and_stop",
        "issue": "A protocol structural finding cannot bind to an ICF content request.",
    }

    result, run_dir, revision, _reference_path, _working = _quality_retry_fixture(
        tmp_path,
        [finding],
        request_records=(("content.json", request),),
    )

    assert result["status"] == "blocked"
    assert result["stage"] == "internal_recovery_routing"
    assert result["findings"][0]["routing_failure_code"] == "verification_request_wrong_artifact"
    assert result["findings"][0]["original_finding"] == finding
    assert not (revision / "attempts").exists()
    assert not (revision / "gate-attempt-journal.json").exists()
    assert not (run_dir / "reference/repair-report.md").exists()


@pytest.mark.parametrize("declare_artifact", (True, False), ids=("declared", "omitted"))
def test_content_route_artifact_must_match_its_section_target(tmp_path, declare_artifact):
    request = _verification_request_payload(
        "r-test.review-1.verify.content",
        "clinical_content_verification",
        artifact="icf",
    )
    for section in request["sections"]:
        if section["section_id"] == "icf.procedures":
            section["artifact"] = "protocol"
    request["request_sha256"] = verification_request_sha256(request)
    finding = {
        "category": "content",
        "field": "icf.procedures",
        "artifact": "protocol",
        "target_ids": ["icf.procedures"],
        "verification_request_id": request["request_id"],
        "recovery_class": "drafting_defect",
        "action": "retry_drafting_target",
        "issue": "An ICF section target cannot be attributed to the Protocol artifact.",
    }
    if not declare_artifact:
        finding.pop("artifact")

    result, run_dir, revision, _reference_path, _working = _quality_retry_fixture(
        tmp_path,
        [finding],
        request_records=(("content.json", request),),
    )

    assert result["status"] == "blocked"
    assert result["stage"] == "internal_recovery_routing"
    assert result["findings"][0]["routing_failure_code"] == "verification_request_incomplete_scope"
    assert result["findings"][0]["original_finding"] == finding
    assert not (revision / "attempts").exists()
    assert not (revision / "gate-attempt-journal.json").exists()
    assert not (revision / "hermes/requests").exists()
    assert not (run_dir / "reference/repair-report.md").exists()


def test_rehashed_content_request_cannot_replace_canonical_source_authority(tmp_path):
    source = _source()
    request = _verification_request_payload(
        "r-test.review-1.verify.content",
        "clinical_content_verification",
    )
    request["approved_source"] = source
    request["request_sha256"] = verification_request_sha256(request)
    canonical_request_file_sha256 = hashlib.sha256(
        json.dumps(request).encode("utf-8")
    ).hexdigest()
    canonical_ledger = {
        "schema_version": "verification-request-ledger/v1",
        "request_id": request["request_id"],
        "request_sha256": request["request_sha256"],
        "request_file_sha256": canonical_request_file_sha256,
        "task": request["task"],
        "revision_id": request["revision_id"],
        "review_set": request["review_set"],
    }
    request["approved_source"] = {**source, "study": {**source["study"], "title": "Altered"}}
    request["request_sha256"] = verification_request_sha256(request)
    finding = {
        "category": "content",
        "field": "icf.procedures",
        "artifact": "icf",
        "target_ids": ["icf.procedures"],
        "verification_request_id": request["request_id"],
        "recovery_class": "drafting_defect",
        "action": "retry_drafting_target",
        "issue": "A rehashed request cannot replace canonical source authority.",
    }

    result, run_dir, revision, _reference_path, _working = _quality_retry_fixture(
        tmp_path,
        [finding],
        request_records=(("content.json", request),),
        request_ledgers={request["request_id"]: canonical_ledger},
    )

    assert result["status"] == "blocked"
    assert result["stage"] == "internal_recovery_routing"
    assert result["findings"][0]["routing_failure_code"] == "verification_request_canonical_hash_mismatch"
    assert result["findings"][0]["original_finding"] == finding
    assert not (revision / "attempts").exists()
    assert not (revision / "gate-attempt-journal.json").exists()
    assert not (run_dir / "reference/repair-report.md").exists()


def test_partial_content_request_cannot_authorize_recovery(tmp_path):
    request = _verification_request_payload(
        "r-test.review-1.verify.content",
        "clinical_content_verification",
        artifact="icf",
    )
    request["sections"] = [{"artifact": "icf", "section_id": "icf.procedures"}]
    request["request_sha256"] = verification_request_sha256(request)
    finding = {
        "category": "content",
        "field": "icf.procedures",
        "artifact": "icf",
        "target_ids": ["icf.procedures"],
        "verification_request_id": request["request_id"],
        "recovery_class": "drafting_defect",
        "action": "retry_drafting_target",
        "issue": "A partial package-wide request cannot authorize repair.",
    }

    result, run_dir, revision, _reference_path, _working = _quality_retry_fixture(
        tmp_path,
        [finding],
        request_records=(("content.json", request),),
    )

    assert result["status"] == "blocked"
    assert result["stage"] == "internal_recovery_routing"
    assert result["findings"][0]["routing_failure_code"] == "verification_request_incomplete_scope"
    assert result["findings"][0]["original_finding"] == finding
    assert not (revision / "attempts").exists()
    assert not (revision / "gate-attempt-journal.json").exists()
    assert not (run_dir / "reference/repair-report.md").exists()


def test_correctly_bound_content_finding_proceeds_normally(tmp_path):
    source = _source()
    request = _verification_request_payload(
        "r-test.review-1.verify.content",
        "clinical_content_verification",
    )
    request["artifacts"] = [
        {
            "artifact": name if name.endswith(".xml") else Path(name).stem,
            "path": f"candidate/{name}",
        }
        for name in sorted(("protocol.docx", "icf.docx", "study.xml"))
    ]
    request["sections"] = quality.content_review_sections(source)
    request["request_sha256"] = verification_request_sha256(request)
    finding = {
        "category": "content",
        "field": "icf.procedures",
        "artifact": "icf",
        "target_ids": ["icf.procedures"],
        "verification_request_id": request["request_id"],
        "recovery_class": "drafting_defect",
        "action": "retry_drafting_target",
        "issue": "Repair the exact ICF procedure section.",
    }

    result, _run_dir, revision, _reference_path, _working = _quality_retry_fixture(
        tmp_path,
        [finding],
        request_records=(("content.json", request),),
    )

    assert result["status"] == "awaiting_hermes"
    assert result["stage"] == "drafting_retry"
    assert (revision / "attempts").is_dir()


def test_terminal_route_must_bind_the_findings_exact_artifact_path(tmp_path):
    request = _verification_request_payload(
        "r-test.review-1.verify.content",
        "clinical_content_verification",
        artifact="icf",
    )
    request["artifacts"][0]["path"] = "candidate/icf.docx"
    request["request_sha256"] = verification_request_sha256(request)
    finding = {
        "category": "verification",
        "field": "clinical_content_verification",
        "artifact": "icf",
        "artifact_path": "candidate/protocol.docx",
        "target_ids": ["verification:content"],
        "verification_request_id": request["request_id"],
        "recovery_class": "document_structure_defect",
        "action": "preserve_and_stop",
        "issue": "A stale protocol path cannot bind to the ICF content request.",
    }

    result, run_dir, revision, _reference_path, _working = _quality_retry_fixture(
        tmp_path,
        [finding],
        request_records=(("content.json", request),),
    )

    assert result["status"] == "blocked"
    assert result["stage"] == "internal_recovery_routing"
    assert result["findings"][0]["routing_failure_code"] == "verification_request_wrong_artifact_path"
    assert result["findings"][0]["original_finding"] == finding
    assert not (revision / "attempts").exists()
    assert not (revision / "gate-attempt-journal.json").exists()
    assert not (run_dir / "reference/repair-report.md").exists()


def test_omitted_artifact_still_binds_path_to_section_target(tmp_path):
    request = _verification_request_payload(
        "r-test.review-1.verify.content",
        "clinical_content_verification",
        artifact="icf",
    )
    request["request_sha256"] = verification_request_sha256(request)
    finding = {
        "category": "content",
        "field": "icf.procedures",
        "artifact_path": "candidate/protocol.docx",
        "target_ids": ["icf.procedures"],
        "verification_request_id": request["request_id"],
        "recovery_class": "drafting_defect",
        "action": "retry_drafting_target",
        "issue": "An ICF target cannot bind to the Protocol path.",
    }

    result, run_dir, revision, _reference_path, _working = _quality_retry_fixture(
        tmp_path,
        [finding],
        request_records=(("content.json", request),),
    )

    assert result["status"] == "blocked"
    assert result["stage"] == "internal_recovery_routing"
    assert result["findings"][0]["routing_failure_code"] == "verification_request_wrong_artifact_path"
    assert result["findings"][0]["original_finding"] == finding
    assert not (revision / "attempts").exists()
    assert not (revision / "gate-attempt-journal.json").exists()
    assert not (run_dir / "reference/repair-report.md").exists()


def test_partial_visual_request_cannot_authorize_recovery(tmp_path):
    request = _verification_request_payload(
        "r-test.review-1.verify.visual.protocol",
        "rendered_page_visual_verification",
        artifact="protocol",
    )
    request["checks"] = []
    request["request_sha256"] = verification_request_sha256(request)
    expected_artifact = {
        "artifact": "protocol",
        "docx": "candidate/protocol.docx",
        "docx_sha256": "a" * 64,
        "pdf": "rendered/protocol.pdf",
        "pdf_sha256": "b" * 64,
        "pages": [{"page": 1, "path": "rendered/protocol/page-1.png", "sha256": "c" * 64}],
    }
    finding = {
        "category": "visual", "field": "rendered_page_visual_verification",
        "artifact": "protocol", "check": "orphan_heading", "element": "5. INTRODUCTION",
        "target_ids": ["layout:protocol"], "verification_request_id": request["request_id"],
        "recovery_class": "visual_defect", "action": "targeted_layout_repair",
        "issue": "A partial visual request cannot authorize repair.",
    }

    result, run_dir, revision, _reference_path, _working = _quality_retry_fixture(
        tmp_path,
        [finding],
        request_records=(("visual.json", request),),
        candidate_build={"render_report": {"artifacts": [expected_artifact]}},
    )

    assert result["status"] == "blocked"
    assert result["stage"] == "internal_recovery_routing"
    assert result["findings"][0]["routing_failure_code"] == "verification_request_incomplete_scope"
    assert result["findings"][0]["original_finding"] == finding
    assert not (revision / "attempts").exists()
    assert not (revision / "gate-attempt-journal.json").exists()
    assert not (run_dir / "reference/repair-report.md").exists()


def test_correctly_bound_visual_request_proceeds_normally(tmp_path, monkeypatch):
    run_dir = tmp_path / "run"
    revision = run_dir / "revisions/r-test"
    revision.mkdir(parents=True)
    candidate = revision / "candidate"
    candidate.mkdir()
    for name in ("protocol.docx", "icf.docx"):
        document = Document()
        document.add_heading("5. INTRODUCTION", level=1)
        document.add_paragraph("Body text.")
        document.save(candidate / name)
    (candidate / "study.xml").write_text("<study />", encoding="utf-8")
    expected_artifact = {
        "artifact": "protocol",
        "docx": "candidate/protocol.docx",
        "pdf": "rendered/protocol.pdf",
        "pages": [{"page": 1, "path": "rendered/protocol/page-1.png"}],
    }
    for relative, data in (
        (expected_artifact["pdf"], b"pdf"),
        (expected_artifact["pages"][0]["path"], b"png"),
    ):
        path = revision / relative
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_bytes(data)
    expected_artifact["docx_sha256"] = quality.sha256_file(revision / expected_artifact["docx"])
    expected_artifact["pdf_sha256"] = quality.sha256_file(revision / expected_artifact["pdf"])
    expected_artifact["pages"][0]["sha256"] = quality.sha256_file(
        revision / expected_artifact["pages"][0]["path"]
    )
    render_report = {"status": "passed", "artifacts": [expected_artifact]}
    request_paths = quality.create_verification_requests(
        revision,
        _source(),
        render_report,
        review_set=1,
    )
    request_path = next(
        path for path in request_paths
        if json.loads(path.read_text(encoding="utf-8"))["task"]
        == "rendered_page_visual_verification"
    )
    request = json.loads(request_path.read_text(encoding="utf-8"))
    response_path = revision / request["response_path"]
    response_path.parent.mkdir(parents=True, exist_ok=True)
    response_path.write_text(json.dumps({
        "schema_version": RESPONSE_SCHEMA,
        "request_id": request["request_id"],
        "request_sha256": request["request_sha256"],
        "task": request["task"],
        "revision_id": request["revision_id"],
        "producer": {"model_id": "fixture-model", "reviewer_id": "fixture-reviewer"},
        "status": "blocked",
        "findings": [{
            "finding_id": "fixture-visual-finding",
            "category": "visual",
            "artifact": "protocol",
            "check": "orphan_heading",
            "element": "5. INTRODUCTION",
            "target_ids": ["layout:protocol"],
            "issue": "Repair the bound protocol heading.",
        }],
    }), encoding="utf-8")
    (revision / "candidate-build.json").write_text(json.dumps({
        "render_report": {"artifacts": [expected_artifact]},
    }), encoding="utf-8")
    reference_path = run_dir / "reference/study.reference.json"
    reference_path.parent.mkdir(parents=True)
    working = {"generation": {"review_set": 1}}
    reference_path.write_text(json.dumps(working), encoding="utf-8")
    monkeypatch.setattr(
        workflow, "generate",
        lambda *_args, **_kwargs: {"status": "awaiting_hermes", "stage": "post-layout-repair"},
    )
    finding = {
        "category": "visual", "field": "rendered_page_visual_verification",
        "artifact": "protocol", "check": "orphan_heading", "element": "5. INTRODUCTION",
        "target_ids": ["layout:protocol"], "verification_request_id": request["request_id"],
        "recovery_class": "visual_defect", "action": "targeted_layout_repair",
        "issue": "Repair the bound protocol heading.",
    }

    result = workflow._quality_retry(
        run_dir, reference_path, working, _source(), revision, {}, [finding], "quality",
        require_promoted_runtime=False,
    )

    assert result == {"status": "awaiting_hermes", "stage": "post-layout-repair"}
    assert (revision / "attempts").is_dir()


@pytest.mark.parametrize("target_ids", [None, []], ids=["missing", "empty"])
def test_missing_or_empty_target_fails_before_archival(tmp_path, target_ids):
    finding = {
        "category": "content",
        "field": "unrecognized-section",
        "artifact": "protocol",
        "recovery_class": "drafting_defect",
        "action": "retry_drafting_target",
        "issue": "The reviewer did not provide a routable target.",
    }
    if target_ids is not None:
        finding["target_ids"] = target_ids

    result, _revision = _prearchive_routing_fixture(tmp_path, finding)

    assert result["stage"] == "internal_recovery_routing"
    assert result["findings"][0]["original_finding"] == finding


@pytest.mark.parametrize(
    ("finding", "reference_update", "expected_disposition"),
    (
        (
            {
                "category": "layout", "field": "visual", "artifact": "protocol",
                "check": "orphan_heading", "element": "5. INTRODUCTION",
                "target_ids": ["layout:icf"], "recovery_class": "visual_defect",
                "action": "targeted_layout_repair", "issue": "The layout target conflicts with the artifact.",
            },
            {},
            "fail_closed:unsupported_layout_target",
        ),
        (
            {
                "category": "layout", "field": "protocol", "artifact": "protocol",
                "check": "not_a_mandatory_check", "element": "5. INTRODUCTION",
                "target_ids": ["layout:protocol"], "recovery_class": "visual_defect",
                "action": "targeted_layout_repair", "issue": "Unsupported visual check.",
            },
            {},
            "fail_closed:unknown_visual_check",
        ),
        (
            {
                "category": "layout", "field": "icf", "artifact": "icf",
                "check": "orphan_heading", "element": "PURPOSE",
                "target_ids": ["layout:icf"], "recovery_class": "visual_defect",
                "action": "targeted_layout_repair", "issue": "Incompatible ICF family.",
            },
            {"icf_template": "Unsupported"},
            "fail_closed:unknown_template_family",
        ),
        (
            {
                "category": "layout", "field": "protocol", "artifact": "protocol",
                "check": "orphan_heading", "element": "",
                "target_ids": ["layout:protocol"], "recovery_class": "visual_defect",
                "action": "targeted_layout_repair", "issue": "Missing localized element.",
            },
            {},
            "repair:heading_cohesion",
        ),
    ),
    ids=["unsupported-target", "unsupported-check", "incompatible-family", "missing-localization"],
)
def test_unsupported_visual_route_fails_before_archival(
    tmp_path, finding, reference_update, expected_disposition,
):
    reference = _source()
    reference["meta"].update(reference_update)

    result, _revision = _prearchive_routing_fixture(tmp_path, finding, reference=reference)

    assert result["stage"] == "layout_repair_classification"
    assert result["findings"][0]["disposition"] == expected_disposition
    assert result["findings"][0]["issue"] == finding["issue"]


def test_preview_unsupported_is_evaluated_before_archive(tmp_path, monkeypatch):
    finding = {
        "category": "layout", "field": "protocol", "artifact": "protocol",
        "check": "not_a_mandatory_check", "element": "5. INTRODUCTION",
        "target_ids": ["layout:protocol"], "recovery_class": "visual_defect",
        "action": "targeted_layout_repair", "issue": "Unsupported visual check.",
    }
    archive_called = False

    def forbidden_archive(*_args, **_kwargs):
        nonlocal archive_called
        archive_called = True
        pytest.fail("unsupported recovery must stop before archival")

    monkeypatch.setattr(workflow, "_archive_failed_attempt", forbidden_archive)

    result, _revision = _prearchive_routing_fixture(tmp_path, finding)

    assert result["stage"] == "layout_repair_classification"
    assert archive_called is False


def test_malformed_finding_preserves_valid_companion_before_archival(tmp_path):
    companion = {
        "category": "content", "field": "icf.procedures", "artifact": "icf",
        "target_ids": ["icf.procedures"], "recovery_class": "drafting_defect",
        "action": "retry_drafting_target", "issue": "Repair the procedure explanation.",
    }
    malformed = {
        "category": "content", "field": "unrecognized-section", "artifact": "protocol",
        "recovery_class": "drafting_defect", "action": "retry_drafting_target",
        "issue": "No repair target was supplied.",
    }

    result, revision = _prearchive_routing_fixture(
        tmp_path, malformed, companions=[companion],
    )

    assert result["stage"] == "internal_recovery_routing"
    assert result["preserved_companion_findings"] == [companion]
    assert not (revision / "hermes/requests").exists()


def test_unsupported_visual_finding_preserves_valid_companion_before_archival(tmp_path):
    companion = {
        "category": "content", "field": "icf.procedures", "artifact": "icf",
        "target_ids": ["icf.procedures"], "recovery_class": "drafting_defect",
        "action": "retry_drafting_target", "issue": "Repair the procedure explanation.",
    }
    unsupported = {
        "category": "layout", "field": "protocol", "artifact": "protocol",
        "check": "not_a_mandatory_check", "element": "5. INTRODUCTION",
        "target_ids": ["layout:protocol"], "recovery_class": "visual_defect",
        "action": "targeted_layout_repair", "issue": "Unsupported visual check.",
    }

    result, revision = _prearchive_routing_fixture(
        tmp_path, unsupported, companions=[companion],
    )

    assert result["stage"] == "layout_repair_classification"
    assert result["preserved_companion_findings"] == [companion]
    assert not (revision / "hermes/requests").exists()


def test_supported_missing_target_normalizes_before_archival_and_proceeds(tmp_path):
    finding = {
        "category": "content", "field": "icf.procedures", "artifact": "icf",
        "recovery_class": "drafting_defect", "action": "retry_drafting_target",
        "issue": "Repair the procedure explanation.",
    }

    result, _run_dir, revision, _reference_path, _working = _quality_retry_fixture(
        tmp_path,
        [finding],
    )

    assert result["status"] == "awaiting_hermes"
    attempt = json.loads(next((revision / "attempts").glob("*/attempt-manifest.json")).read_text())
    assert attempt["findings"][0]["target_ids"] == ["icf.procedures"]


def test_verifier_request_id_must_resolve_to_the_target_task_before_archival(tmp_path):
    mismatched_request = _verification_request_payload(
        "r-test.verify.visual",
        "rendered_page_visual_verification",
    )
    finding = {
        "category": "reviewer-transient",
        "field": "content",
        "target_ids": ["verification:content"],
        "verification_request_id": "r-test.verify.visual",
        "recovery_class": "verifier_transient",
        "action": "retry_verifier",
        "issue": "The request identity conflicts with the declared verification target.",
    }

    result, _run_dir, revision, _reference_path, _working = _quality_retry_fixture(
        tmp_path,
        [finding],
        request_records=(("visual.json", mismatched_request),),
    )

    assert result["status"] == "blocked"
    assert result["stage"] == "internal_recovery_routing"
    assert result["findings"][0]["original_finding"] == finding
    assert not (revision / "attempts").exists()
    assert not (revision / "gate-attempt-journal.json").exists()


def test_nonexistent_verifier_request_id_fails_before_archival(tmp_path):
    finding = {
        "category": "reviewer-transient",
        "field": "content",
        "target_ids": ["verification:content"],
        "verification_request_id": "missing-request",
        "recovery_class": "verifier_transient",
        "action": "retry_verifier",
        "issue": "The declared verifier request does not exist.",
    }

    result, _revision = _prearchive_routing_fixture(tmp_path, finding)

    assert result["stage"] == "internal_recovery_routing"
    assert result["findings"][0]["original_finding"] == finding


@pytest.mark.parametrize(
    ("request_id", "create_content_request"),
    (("missing-visual-request", False), ("r-test.verify.content", True)),
    ids=("nonexistent", "wrong-task"),
)
def test_visual_finding_request_identity_is_validated_before_archival(
    tmp_path, request_id, create_content_request,
):
    request_records = ()
    if create_content_request:
        request = _verification_request_payload(
            request_id,
            "clinical_content_verification",
        )
        request_records = (("content.json", request),)
    finding = {
        "category": "visual",
        "field": "protocol",
        "artifact": "protocol",
        "check": "orphan_heading",
        "element": "5. INTRODUCTION",
        "target_ids": ["layout:protocol"],
        "verification_request_id": request_id,
        "recovery_class": "visual_defect",
        "action": "targeted_layout_repair",
        "issue": "The visual finding must remain bound to its visual request.",
    }

    result, _run_dir, revision, _reference_path, _working = _quality_retry_fixture(
        tmp_path,
        [finding],
        request_records=request_records,
    )

    assert result["status"] == "blocked"
    assert result["stage"] == "internal_recovery_routing"
    assert result["findings"][0]["original_finding"] == finding
    assert not (revision / "attempts").exists()
    assert not (revision / "gate-attempt-journal.json").exists()


def test_structural_finding_request_identity_is_bound_before_terminal_block(tmp_path):
    request_id = "r-test.review-1.verify.visual.protocol"
    request = _verification_request_payload(
        request_id,
        "rendered_page_visual_verification",
        revision_id="r-other",
        artifact="protocol",
    )
    finding = {
        "category": "verification",
        "field": "request_sha256",
        "target_ids": ["verification:visual"],
        "verification_request_id": request_id,
        "recovery_class": "document_structure_defect",
        "action": "preserve_and_stop",
        "issue": "A structural finding is bound to the wrong revision.",
    }

    result, _run_dir, revision, _reference_path, _working = _quality_retry_fixture(
        tmp_path,
        [finding],
        request_records=(("visual.json", request),),
    )

    assert result["status"] == "blocked"
    assert result["stage"] == "internal_recovery_routing"
    assert result["findings"][0]["original_finding"] == finding
    assert not (revision / "attempts").exists()
    assert not (revision / "gate-attempt-journal.json").exists()


def test_verification_request_hash_must_remain_bound_before_archival(tmp_path):
    request_id = "r-test.review-1.verify.content"
    request = _verification_request_payload(
        request_id,
        "clinical_content_verification",
    )
    request["response_path"] = "hermes/verification-responses/tampered.json"
    finding = {
        "category": "reviewer-transient",
        "field": "content",
        "target_ids": ["verification:content"],
        "verification_request_id": request_id,
        "recovery_class": "verifier_transient",
        "action": "retry_verifier",
        "issue": "The verification request bytes changed after binding.",
    }

    result, _run_dir, revision, _reference_path, _working = _quality_retry_fixture(
        tmp_path,
        [finding],
        request_records=(("content.json", request),),
    )

    assert result["status"] == "blocked"
    assert result["stage"] == "internal_recovery_routing"
    assert result["findings"][0]["original_finding"] == finding
    assert not (revision / "attempts").exists()
    assert not (revision / "gate-attempt-journal.json").exists()


def test_bound_verifier_request_id_still_retries_normally(tmp_path):
    request_id = "r-test.review-1.verify.content"
    request = _verification_request_payload(
        request_id,
        "clinical_content_verification",
    )
    finding = {
        "category": "reviewer-transient",
        "field": "content",
        "target_ids": ["verification:content"],
        "verification_request_id": request_id,
        "recovery_class": "verifier_transient",
        "action": "retry_verifier",
        "issue": "Retry the exact content verifier request.",
    }

    result, _run_dir, revision, _reference_path, _working = _quality_retry_fixture(
        tmp_path,
        [finding],
        request_records=(("content.json", request),),
    )

    assert result["status"] == "awaiting_hermes"
    assert result["stage"] == "independent_verification_retry"
    assert result["requests"] == [
        "hermes/verification-requests/r-test.review-1.verify.content.json"
    ]
    assert (revision / "hermes/verification-requests/r-test.review-1.verify.content.json").is_file()


def test_invalid_route_does_not_drop_valid_recovery_target(tmp_path):
    run_dir = tmp_path / "run"
    revision = run_dir / "revisions/r-test"
    reference_path = run_dir / "reference/study.reference.json"
    reference_path.parent.mkdir(parents=True)
    revision.mkdir(parents=True)
    reference = _source()
    reference["meta"]["icf_template"] = "Sterling"
    working = {"generation": {"review_set": 1}}
    reference_path.write_text(json.dumps(working), encoding="utf-8")
    valid = {
        "category": "content", "field": "icf.background", "artifact": "icf",
        "target_ids": ["icf.background"], "recovery_class": "drafting_defect",
        "action": "retry_drafting_target", "issue": "Repair the participant-facing background.",
    }
    invalid = {
        "category": "content", "field": "icf", "artifact": "icf",
        "target_ids": ["layout:icf"], "recovery_class": "drafting_defect",
        "action": "retry_drafting_target", "issue": "This drafting class has the wrong target category.",
    }

    result = workflow._quality_retry(
        run_dir, reference_path, working, reference, revision, {}, [valid, invalid], "quality",
        require_promoted_runtime=False,
    )

    assert result["status"] == "blocked"
    assert result["stage"] == "internal_recovery_routing"
    assert result["findings"][0]["original_finding"] == invalid
    assert result["preserved_companion_findings"] == [valid]
    assert working["generation"]["review_set"] == 1
    assert not (revision / "attempts").exists()
    assert not (revision / "gate-attempt-journal.json").exists()
    assert not (revision / "hermes/requests").exists()


def _mixed_icf_recovery_fixture(tmp_path):
    run_dir = tmp_path / "run"
    revision = run_dir / "revisions/r-test"
    reference_path = run_dir / "reference/study.reference.json"
    accepted = revision / "hermes/accepted"
    candidate = revision / "candidate"
    reference_path.parent.mkdir(parents=True)
    accepted.mkdir(parents=True)
    candidate.mkdir(parents=True)
    (accepted / "icf.background.json").write_text(
        json.dumps({"section_id": "icf.background", "paragraphs": [{"text": "unchanged background"}], "lists": []}),
        encoding="utf-8",
    )
    (accepted / "icf.procedures.json").write_text(
        json.dumps({"section_id": "icf.procedures", "paragraphs": [{"text": "old procedures"}], "lists": []}),
        encoding="utf-8",
    )
    (candidate / "icf.docx").write_bytes(b"old icf candidate")
    findings = [
        {
            "category": "content", "field": target, "artifact": "icf",
            "target_ids": [target], "recovery_class": "drafting_defect",
            "action": "retry_drafting_target", "issue": f"Repair {target}.",
        }
        for target in ("icf.background", "icf.procedures")
    ]
    attempt = workflow._archive_failed_attempt(revision, "quality", findings)
    journal = json.loads((revision / "gate-attempt-journal.json").read_text(encoding="utf-8"))
    working = {
        "generation": {
            "review_set": 1,
            "gate_attempts": journal["entries"],
            "pending_recovery_attempts": [attempt.relative_to(revision).as_posix()],
            "pending_review_set_advance": {"from_review_set": 1, "strategy_ids": [], "finding_sha256": "a" * 64},
        }
    }
    reference_path.write_text(json.dumps(working), encoding="utf-8")
    (accepted / "icf.procedures.json").write_text(
        json.dumps({"section_id": "icf.procedures", "paragraphs": [{"text": "repaired procedures"}], "lists": []}),
        encoding="utf-8",
    )
    (candidate / "icf.docx").write_bytes(b"candidate changed only for procedures")
    return revision, reference_path, working


def test_mixed_recovery_does_not_mask_unchanged_target(tmp_path):
    revision, reference_path, working = _mixed_icf_recovery_fixture(tmp_path)

    no_progress = workflow._complete_pending_recovery_attempts(
        revision, reference_path, working, require_candidate_change=True,
    )

    assert [item["target_ids"] for item in no_progress] == [["icf.background"]]
    assert "pending_review_set_advance" not in working["generation"]


def test_visual_companion_progress_is_measured_per_exact_layout_target(tmp_path):
    revision = tmp_path / "revisions/r-test"
    reference_path = tmp_path / "reference/study.reference.json"
    candidate = revision / "candidate"
    reference_path.parent.mkdir(parents=True)
    candidate.mkdir(parents=True)
    document = Document()
    document.add_heading("5. INTRODUCTION", level=1)
    document.add_paragraph("Introduction body.")
    document.add_heading("6. OBJECTIVES", level=1)
    document.add_paragraph("Objectives body.")
    document.save(candidate / "protocol.docx")
    (revision / "candidate-build.json").write_text(json.dumps({
        "governing_resources": {"layout_repairs": {"protocol": []}},
    }), encoding="utf-8")
    findings = [
        {
            "category": "visual", "field": "protocol", "artifact": "protocol",
            "check": "orphan_heading", "element": element,
            "repair_rule": "heading_cohesion", "target_ids": ["layout:protocol"],
            "recovery_class": "visual_defect", "action": "targeted_layout_repair",
            "issue": f"Repair {element}.",
        }
        for element in ("5. INTRODUCTION", "6. OBJECTIVES")
    ]
    attempt = workflow._archive_failed_attempt(revision, "quality", findings)
    journal = json.loads((revision / "gate-attempt-journal.json").read_text(encoding="utf-8"))
    working = {
        "generation": {
            "gate_attempts": journal["entries"],
            "pending_recovery_attempts": [attempt.relative_to(revision).as_posix()],
            "pending_review_set_advance": {"from_review_set": 1},
        }
    }
    reference_path.write_text(json.dumps(working), encoding="utf-8")
    document = Document(candidate / "protocol.docx")
    document.paragraphs[0].add_run(" ")
    document.save(candidate / "protocol.docx")
    (revision / "candidate-build.json").write_text(json.dumps({
        "governing_resources": {
            "layout_repairs": {
                "protocol": [{"rule": "heading_cohesion", "target": "5. INTRODUCTION"}],
            },
        },
    }), encoding="utf-8")

    no_progress = workflow._complete_pending_recovery_attempts(
        revision, reference_path, working, require_candidate_change=True,
    )

    assert [item["element"] for item in no_progress] == ["6. OBJECTIVES"]
    assert "pending_review_set_advance" not in working["generation"]


def test_mixed_no_progress_continues_all_unresolved_repairable_companions(tmp_path, monkeypatch):
    run_dir = tmp_path / "run"
    revision = run_dir / "revisions/r-test"
    reference_path = run_dir / "reference/study.reference.json"
    reference_path.parent.mkdir(parents=True)
    revision.mkdir(parents=True)
    drafting = {
        "category": "content", "field": "icf.procedures", "artifact": "icf",
        "target_ids": ["icf.procedures"], "recovery_class": "drafting_defect",
        "action": "retry_drafting_target", "issue": "The draft did not change.",
    }
    visual = {
        "category": "visual", "field": "protocol", "artifact": "protocol",
        "check": "orphan_heading", "element": "5. INTRODUCTION",
        "target_ids": ["layout:protocol"], "recovery_class": "visual_defect",
        "action": "targeted_layout_repair", "issue": "The visual defect remains unresolved.",
    }

    expected = {"status": "awaiting_hermes", "stage": "recovery_no_progress"}
    observed = {}

    def continue_retry(*args, **kwargs):
        observed["findings"] = args[6]
        observed["stage"] = args[7]
        return expected

    monkeypatch.setattr(workflow, "_quality_retry", continue_retry)
    result = workflow._continue_repairable_no_progress(
        run_dir,
        reference_path,
        {"generation": {}},
        _source(),
        revision,
        {},
        [drafting, visual],
        contracted_bundle={},
        operation_deadline=None,
        clock=lambda: 0.0,
        stage_observer=None,
        require_promoted_runtime=False,
    )

    assert result == expected
    assert observed["stage"] == "recovery_no_progress"
    assert observed["findings"] == [drafting, visual]


def test_shared_prs_record_progress_is_measured_per_target(tmp_path):
    run_dir = tmp_path / "run"
    revision = run_dir / "revisions/r-test"
    reference_path = run_dir / "reference/study.reference.json"
    accepted = revision / "hermes/accepted"
    candidate = revision / "candidate"
    reference_path.parent.mkdir(parents=True)
    accepted.mkdir(parents=True)
    candidate.mkdir(parents=True)
    accepted_record = {
        "narrative": {
            "brief_summary": {"text": "old brief"},
            "detailed_description": {"text": "unchanged detail"},
        }
    }
    (accepted / "prs-narrative.json").write_text(json.dumps(accepted_record), encoding="utf-8")
    (candidate / "study.xml").write_bytes(b"old study")
    finding = {
        "category": "content",
        "field": "prs-narrative",
        "artifact": "study.xml",
        "target_ids": ["prs.brief-summary", "prs.detailed-description"],
        "recovery_class": "drafting_defect",
        "action": "retry_drafting_target",
        "issue": "Both PRS narratives require repair.",
    }
    attempt = workflow._archive_failed_attempt(revision, "quality", [finding])
    journal = json.loads((revision / "gate-attempt-journal.json").read_text(encoding="utf-8"))
    working = {
        "generation": {
            "review_set": 1,
            "gate_attempts": journal["entries"],
            "pending_recovery_attempts": [attempt.relative_to(revision).as_posix()],
            "pending_review_set_advance": {"from_review_set": 1, "strategy_ids": [], "finding_sha256": "b" * 64},
        }
    }
    reference_path.write_text(json.dumps(working), encoding="utf-8")
    accepted_record["narrative"]["brief_summary"] = {"text": "repaired brief"}
    (accepted / "prs-narrative.json").write_text(json.dumps(accepted_record), encoding="utf-8")
    (candidate / "study.xml").write_bytes(b"study changed for brief only")

    no_progress = workflow._complete_pending_recovery_attempts(
        revision, reference_path, working, require_candidate_change=True,
    )

    assert [item["target_ids"] for item in no_progress] == [["prs.detailed-description"]]
    assert "pending_review_set_advance" not in working["generation"]


def test_same_finding_and_same_target_hash_cannot_advance_review_set(tmp_path):
    revision, reference_path, working = _mixed_icf_recovery_fixture(tmp_path)
    workflow._complete_pending_recovery_attempts(
        revision, reference_path, working, require_candidate_change=True,
    )

    workflow._advance_pending_review_set(reference_path, working)

    assert working["generation"]["review_set"] == 1
    assert working["generation"]["no_progress_history"][0]["target_ids"] == ["icf.background"]


def test_recovery_archive_preserves_canonical_verification_ledger(tmp_path):
    revision = tmp_path / "revisions/r-test"
    ledger = revision / "request-ledger/r-test.review-1.verify.content.json"
    ledger.parent.mkdir(parents=True)
    ledger.write_text('{"schema_version":"verification-request-ledger/v1"}', encoding="utf-8")
    finding = {
        "category": "content",
        "field": "icf.procedures",
        "artifact": "icf",
        "target_ids": ["icf.procedures"],
        "recovery_class": "drafting_defect",
        "action": "retry_drafting_target",
        "issue": "Preserve canonical request authority with rejected evidence.",
    }

    attempt = workflow._archive_failed_attempt(revision, "quality", [finding])

    archived = attempt / ledger.relative_to(revision)
    assert archived.read_bytes() == ledger.read_bytes()


def test_reviewer_defect_starts_a_fresh_complete_review_set(tmp_path, monkeypatch):
    run_dir = tmp_path / "run"
    revision = run_dir / "revisions/r-test"
    reference_path = run_dir / "reference/study.reference.json"
    request_dir = revision / "hermes/verification-requests"
    response_dir = revision / "hermes/verification-responses"
    reference_path.parent.mkdir(parents=True)
    request_dir.mkdir(parents=True)
    response_dir.mkdir()
    reference_path.write_text(json.dumps({"generation": {"review_set": 1}}), encoding="utf-8")
    for name, task in (
        ("content", "clinical_content_verification"),
        ("protocol", "rendered_page_visual_verification"),
        ("icf", "rendered_page_visual_verification"),
    ):
        response_path = f"hermes/verification-responses/{name}.json"
        (request_dir / f"{name}.json").write_text(json.dumps({
            "task": task,
            "response_path": response_path,
            "artifacts": [{"artifact": name}] if task.startswith("rendered") else [],
        }), encoding="utf-8")
        (revision / response_path).write_text("{}", encoding="utf-8")
    (revision / "candidate-build.json").write_text("{}", encoding="utf-8")
    rerun = {"status": "awaiting_hermes", "stage": "independent_verification"}

    def rebuild(_run_dir, **_kwargs):
        candidate = revision / "candidate/protocol.docx"
        candidate.parent.mkdir(parents=True, exist_ok=True)
        repaired = Document()
        repaired.add_heading("5. INTRODUCTION", level=1)
        repaired.add_paragraph("Repaired protocol body.")
        repaired.save(candidate)
        latest = json.loads(reference_path.read_text(encoding="utf-8"))
        (revision / "candidate-build.json").write_text(json.dumps({
            "governing_resources": {
                "layout_repairs": latest["generation"]["layout_repairs"],
            },
        }), encoding="utf-8")
        assert workflow._complete_pending_recovery_attempts(
            revision, reference_path, latest, require_candidate_change=True,
        ) == []
        workflow._advance_pending_review_set(reference_path, latest)
        return rerun

    monkeypatch.setattr(workflow, "generate", rebuild)

    result = workflow._quality_retry(
        run_dir,
        reference_path,
        {"generation": {"review_set": 1}},
        _source(),
        revision,
        {},
        [{"category": "visual", "field": "protocol.docx:10", "artifact": "protocol", "check": "orphan_heading", "element": "5. INTRODUCTION", "target_ids": ["layout:protocol.docx"], "recovery_class": "visual_defect", "action": "targeted_layout_repair", "issue": "orphan heading"}],
        "quality",
    )

    assert result == rerun
    state = json.loads(reference_path.read_text(encoding="utf-8"))
    assert state["generation"]["review_set"] == 2
    assert list(request_dir.glob("*.json")) == []
    assert list(response_dir.glob("*.json")) == []


def test_generate_stops_at_expired_operation_deadline_before_creating_requests(tmp_path):
    run_dir = tmp_path / "run"
    reference_path = run_dir / "reference/study.reference.json"
    reference_path.parent.mkdir(parents=True)
    reference_path.write_text(json.dumps(_source()), encoding="utf-8")
    assert workflow.prepare(run_dir)["status"] == "awaiting_approval"
    approval = workflow.approve(run_dir, approved_by="reviewer")
    revision = run_dir / "revisions" / approval["revision_id"]

    result = workflow.generate(
        run_dir,
        operation_deadline=10.0,
        clock=lambda: 10.0,
        require_promoted_runtime=False,
    )

    assert result["status"] == "timeout"
    assert result["stage"] == "generation"
    assert not (revision / "hermes/requests").exists()


def test_quality_retry_stops_at_expired_deadline_before_archiving(tmp_path):
    run_dir = tmp_path / "run"
    revision = run_dir / "revisions/r-test"
    reference_path = run_dir / "reference/study.reference.json"
    reference_path.parent.mkdir(parents=True)
    working = {"generation": {}}
    reference_path.write_text(json.dumps(working), encoding="utf-8")
    finding = {
        "category": "visual",
        "field": "protocol",
        "artifact": "protocol",
        "check": "orphan_heading",
        "element": "6.2. Inclusion/Exclusion Criteria",
        "target_ids": ["layout:protocol"],
        "recovery_class": "visual_defect",
        "action": "targeted_layout_repair",
        "issue": "orphan heading",
    }

    result = workflow._quality_retry(
        run_dir,
        reference_path,
        working,
        _source(),
        revision,
        {},
        [finding],
        "quality",
        operation_deadline=10.0,
        clock=lambda: 10.0,
        require_promoted_runtime=False,
    )

    assert result["status"] == "timeout"
    assert result["stage"] == "quality"
    assert not (revision / "attempts").exists()


def test_repairable_third_review_set_continues_within_the_operation_deadline(tmp_path, monkeypatch):
    run_dir = tmp_path / "run"
    revision = run_dir / "revisions/r-test"
    reference_path = run_dir / "reference/study.reference.json"
    reference_path.parent.mkdir(parents=True)
    finding = {"category": "visual", "field": "protocol.docx:10", "artifact": "protocol", "check": "orphan_heading", "element": "5. INTRODUCTION", "target_ids": ["layout:protocol.docx"], "recovery_class": "visual_defect", "action": "targeted_layout_repair", "issue": "orphan heading"}
    generation = {
        "review_set": 3,
        "recovery_history": [{"review_set": 2, "strategy_ids": [workflow._recovery_strategy_id(finding)]}],
    }
    reference_path.write_text(json.dumps({"generation": generation}), encoding="utf-8")
    rerun = {"status": "awaiting_hermes", "stage": "independent_verification"}

    def rebuild(*_args, **_kwargs):
        candidate = revision / "candidate/protocol.docx"
        candidate.parent.mkdir(parents=True, exist_ok=True)
        repaired = Document()
        repaired.add_heading("5. INTRODUCTION", level=1)
        repaired.add_paragraph("Repaired protocol body.")
        repaired.save(candidate)
        latest = json.loads(reference_path.read_text(encoding="utf-8"))
        (revision / "candidate-build.json").write_text(json.dumps({
            "governing_resources": {
                "layout_repairs": latest["generation"]["layout_repairs"],
            },
        }), encoding="utf-8")
        assert workflow._complete_pending_recovery_attempts(
            revision, reference_path, latest, require_candidate_change=True,
        ) == []
        workflow._advance_pending_review_set(reference_path, latest)
        return rerun

    monkeypatch.setattr(workflow, "generate", rebuild)

    result = workflow._quality_retry(
        run_dir,
        reference_path,
        {"generation": generation},
        _source(),
        revision,
        {},
        [finding],
        "quality",
    )

    assert result == rerun
    state = json.loads(reference_path.read_text(encoding="utf-8"))["generation"]
    assert state["review_set"] == 4


def test_recovery_attempt_accounting_does_not_exhaust_after_three_attempts():
    finding = {
        "category": "visual",
        "field": "protocol",
        "artifact": "protocol",
        "check": "orphan_heading",
        "element": "6.2. Inclusion/Exclusion Criteria",
        "target_ids": ["layout:protocol"],
        "recovery_class": "visual_defect",
        "action": "targeted_layout_repair",
        "issue": "orphan heading",
    }

    attempts, strategies = workflow._advance_recovery_attempts(
        [finding],
        {"layout:protocol": 12},
        {},
    )

    assert attempts["layout:protocol"] == 13
    assert strategies["layout:protocol"]


def test_repeated_orphan_heading_escalates_to_an_untried_page_boundary():
    finding = {
        "category": "visual",
        "field": "protocol",
        "artifact": "protocol",
        "check": "orphan_heading",
        "element": "6.2. Inclusion/Exclusion Criteria",
        "target_ids": ["layout:protocol"],
        "recovery_class": "visual_defect",
        "action": "targeted_layout_repair",
        "issue": "orphan heading",
    }

    plan, unsupported = workflow._layout_repair_plan(
        [finding],
        study_type="Retrospective",
        existing_repairs={
            "protocol": [
                {"rule": "heading_cohesion", "target": "6.2. Inclusion/Exclusion Criteria"},
            ],
        },
    )

    assert unsupported == []
    assert plan == {
        "protocol": [
            {"rule": "heading_page_boundary", "target": "6.2. Inclusion/Exclusion Criteria"},
        ],
    }


def test_repeated_split_table_escalates_to_an_untried_table_boundary():
    finding = {
        "category": "visual",
        "field": "protocol",
        "artifact": "protocol",
        "check": "bad_table_split",
        "element": "Table 13.3.-1",
        "target_ids": ["layout:protocol"],
        "recovery_class": "visual_defect",
        "action": "targeted_layout_repair",
        "issue": "split table",
    }

    plan, unsupported = workflow._layout_repair_plan(
        [finding],
        study_type="Prospective",
        existing_repairs={
            "protocol": [
                {"rule": "table_pagination", "target": "Table 13.3.-1"},
            ],
        },
    )

    assert unsupported == []
    assert plan == {
        "protocol": [
            {"rule": "table_page_boundary", "target": "Table 13.3.-1"},
        ],
    }


def test_exhausted_internal_layout_route_without_request_id_blocks_before_archive(tmp_path):
    finding = {
        "category": "visual",
        "field": "protocol",
        "artifact": "protocol",
        "check": "orphan_heading",
        "element": "6.2. Inclusion/Exclusion Criteria",
        "target_ids": ["layout:protocol"],
        "recovery_class": "visual_defect",
        "action": "targeted_layout_repair",
        "issue": "Internal deterministic finding has no verifier identity.",
    }
    working = {
        "generation": {
            "review_set": 7,
            "layout_repairs": {
                "protocol": [
                    {"rule": "heading_cohesion", "target": finding["element"]},
                    {"rule": "heading_page_boundary", "target": finding["element"]},
                ],
            },
        }
    }

    result, run_dir, revision, _reference_path, state = _quality_retry_fixture(
        tmp_path,
        [finding],
        working=working,
    )

    assert result["status"] == "blocked"
    assert result["stage"] == "layout_repair_classification"
    assert result["findings"][0]["disposition"] == (
        "fail_closed:missing_exact_visual_request_binding"
    )
    assert result["findings"][0]["issue"] == finding["issue"]
    assert state["generation"]["review_set"] == 7
    assert not (revision / "attempts").exists()
    assert not (revision / "gate-attempt-journal.json").exists()
    assert not (run_dir / "reference/repair-report.md").exists()


def test_exhausted_safe_layout_ladder_reprompts_visual_reviewer_until_deadline(
    tmp_path, monkeypatch,
):
    run_dir = tmp_path / "run"
    revision = run_dir / "revisions/r-test"
    reference_path = run_dir / "reference/study.reference.json"
    reference_path.parent.mkdir(parents=True)
    finding = {
        "category": "visual",
        "field": "protocol",
        "artifact": "protocol",
        "check": "orphan_heading",
        "element": "6.2. Inclusion/Exclusion Criteria",
        "target_ids": ["layout:protocol"],
        "verification_request_id": "r-test.review-7.verify.visual.protocol",
        "recovery_class": "visual_defect",
        "action": "targeted_layout_repair",
        "issue": "orphan heading",
    }
    generation = {
        "review_set": 7,
        "layout_repairs": {
            "protocol": [
                {"rule": "heading_cohesion", "target": finding["element"]},
                {"rule": "heading_page_boundary", "target": finding["element"]},
            ],
        },
    }
    reference_path.write_text(json.dumps({"generation": generation}), encoding="utf-8")
    request_dir = revision / "hermes/verification-requests"
    response_dir = revision / "hermes/verification-responses"
    request_dir.mkdir(parents=True)
    response_dir.mkdir(parents=True)
    request_path = request_dir / f"{finding['verification_request_id']}.json"
    response_path = response_dir / f"{finding['verification_request_id']}.json"
    request = {
        "request_id": finding["verification_request_id"],
        "task": "rendered_page_visual_verification",
        "revision_id": revision.name,
        "review_set": 7,
        "response_path": response_path.relative_to(revision).as_posix(),
        "artifacts": [{"artifact": "protocol", "pages": []}],
    }
    request["request_sha256"] = verification_request_sha256(request)
    request_path.write_text(json.dumps(request), encoding="utf-8")
    response_path.write_text("{}", encoding="utf-8")

    def archive_then_remove_current(*_args, **_kwargs):
        archive = revision / "attempts/quality-a01"
        archived_request = archive / request_path.relative_to(revision)
        archived_request.parent.mkdir(parents=True)
        archived_request.write_text(json.dumps(request), encoding="utf-8")
        request_path.unlink()
        response_path.unlink()
        (revision / "gate-attempt-journal.json").write_text(json.dumps({"entries": []}), encoding="utf-8")
        return archive

    monkeypatch.setattr(workflow, "_archive_failed_attempt", archive_then_remove_current)
    monkeypatch.setattr(workflow, "_complete_pending_recovery_attempts", lambda *_args, **_kwargs: [])

    result = workflow._quality_retry(
        run_dir,
        reference_path,
        {"generation": generation},
        _source(),
        revision,
        {},
        [finding],
        "quality",
    )

    assert result["status"] == "awaiting_hermes"
    assert result["stage"] == "independent_verification_retry"
    assert [item["request_id"] for item in result["handoffs"]] == [finding["verification_request_id"]]
    assert request_path.is_file()
    assert not response_path.exists()
    state = json.loads(reference_path.read_text(encoding="utf-8"))["generation"]
    assert state["verification_attempts"][finding["verification_request_id"]] == 1
    assert "recovery_exhaustion" not in state


def test_generation_authority_rebind_archives_complete_recovery_state(tmp_path):
    revision = tmp_path / "revisions/r-source"
    attempt = revision / "attempts/quality-a01"
    staging = revision / ".attempt-staging/interrupted"
    attempt.mkdir(parents=True)
    staging.mkdir(parents=True)
    (attempt / "attempt-manifest.json").write_text("{}", encoding="utf-8")
    (staging / "partial.json").write_text("{}", encoding="utf-8")
    (revision / "gate-attempt-journal.json").write_text('{"entries": []}', encoding="utf-8")
    (revision / "recovery-archive-transaction.json").write_text("{}", encoding="utf-8")
    (revision / "recovery-finalization-transaction.json").write_text("{}", encoding="utf-8")
    working = {"generation": {"governing_sha256": "a" * 64, "gate_attempts": []}}

    assert workflow._rebind_generation_authority(revision, working, "b" * 64) is True
    workflow._complete_generation_authority_rebind(revision)

    archive = revision / "generation-authority-attempts" / ("a" * 12)
    for relative in (
        "attempts/quality-a01/attempt-manifest.json",
        ".attempt-staging/interrupted/partial.json",
        "gate-attempt-journal.json",
        "recovery-archive-transaction.json",
        "recovery-finalization-transaction.json",
    ):
        assert (archive / relative).is_file()
        assert not (revision / relative).exists()
    workflow._validate_generation_authority_attempts(revision, working["generation"])


def test_generation_authority_inventory_includes_nested_manifest_named_files(tmp_path):
    nested = tmp_path / "nested"
    nested.mkdir(parents=True)
    (tmp_path / "authority-attempt-manifest.json").write_text("root", encoding="utf-8")
    (nested / "authority-attempt-manifest.json").write_text("nested", encoding="utf-8")

    inventory = workflow._generation_authority_inventory(tmp_path)

    assert [item["path"] for item in inventory] == ["nested/authority-attempt-manifest.json"]


def test_verifier_transient_remains_retryable_after_three_attempts(tmp_path, monkeypatch):
    run_dir = tmp_path / "run"
    revision = run_dir / "revisions/r-test"
    reference_path = run_dir / "reference/study.reference.json"
    reference_path.parent.mkdir(parents=True)
    candidate = revision / "candidate"
    candidate.mkdir(parents=True)
    Document().save(candidate / "protocol.docx")
    Document().save(candidate / "icf.docx")
    (candidate / "study.xml").write_text("<study />", encoding="utf-8")
    request_path = quality.create_verification_requests(
        revision,
        _source(),
        {"status": "passed", "artifacts": []},
        review_set=1,
    )[0]
    request = json.loads(request_path.read_text(encoding="utf-8"))
    request_id = request["request_id"]
    generation = {"review_set": 1, "verification_attempts": {request_id: 3}}
    reference_path.write_text(json.dumps({"generation": generation}), encoding="utf-8")
    finding = {
        "category": "reviewer-transient",
        "field": "verification:content",
        "target_ids": ["verification:content"],
        "verification_request_id": request_id,
        "recovery_class": "verifier_transient",
        "action": "retry_verifier",
        "issue": "malformed reviewer response",
    }
    expected = {"status": "awaiting_hermes", "stage": "independent_verification_retry"}
    monkeypatch.setattr(workflow, "_awaiting", lambda *_args, **_kwargs: expected)

    result = workflow._quality_retry(
        run_dir,
        reference_path,
        {"generation": generation},
        _source(),
        revision,
        {},
        [finding],
        "quality",
    )

    assert result == expected
    state = json.loads(reference_path.read_text(encoding="utf-8"))["generation"]
    assert state["verification_attempts"][request_id] == 4


def test_third_review_set_continues_with_a_distinct_governed_strategy(tmp_path, monkeypatch):
    run_dir = tmp_path / "run"
    revision = run_dir / "revisions/r-test"
    reference_path = run_dir / "reference/study.reference.json"
    reference_path.parent.mkdir(parents=True)
    prior = {"category": "visual", "field": "protocol", "artifact": "protocol", "check": "orphan_heading", "element": "5. INTRODUCTION", "target_ids": ["layout:protocol"], "recovery_class": "visual_defect", "action": "targeted_layout_repair", "issue": "orphan heading"}
    current = {**prior, "check": "bad_table_split", "element": "3. GENERAL INFORMATION", "issue": "split table"}
    generation = {
        "review_set": 3,
        "recovery_history": [{"review_set": 2, "strategy_ids": [workflow._recovery_strategy_id(prior)]}],
    }
    reference_path.write_text(json.dumps({"generation": generation}), encoding="utf-8")
    rerun = {"status": "awaiting_hermes", "stage": "independent_verification"}

    def rebuild(*_args, **_kwargs):
        candidate = revision / "candidate/protocol.docx"
        candidate.parent.mkdir(parents=True, exist_ok=True)
        repaired = Document()
        repaired.add_heading("5. INTRODUCTION", level=1)
        repaired.add_paragraph("Repaired protocol body.")
        repaired.save(candidate)
        latest = json.loads(reference_path.read_text(encoding="utf-8"))
        (revision / "candidate-build.json").write_text(json.dumps({
            "governing_resources": {
                "layout_repairs": latest["generation"]["layout_repairs"],
            },
        }), encoding="utf-8")
        assert workflow._complete_pending_recovery_attempts(
            revision, reference_path, latest, require_candidate_change=True,
        ) == []
        workflow._advance_pending_review_set(reference_path, latest)
        return rerun

    monkeypatch.setattr(workflow, "generate", rebuild)

    result = workflow._quality_retry(
        run_dir, reference_path, {"generation": generation}, _source(), revision, {}, [current], "quality",
    )

    assert result == rerun
    state = json.loads(reference_path.read_text(encoding="utf-8"))["generation"]
    assert state["review_set"] == 4
    selected = {**current, "repair_rule": "table_pagination"}
    assert workflow._recovery_strategy_id(selected) in state["recovery_history"][-1]["strategy_ids"]


def test_deterministic_reconstruction_records_exact_target_and_measured_bytes(tmp_path, monkeypatch):
    run_dir = tmp_path / "run"
    revision = run_dir / "revisions/r-test"
    reference_path = run_dir / "reference/study.reference.json"
    candidate = revision / "candidate/study.xml"
    candidate.parent.mkdir(parents=True)
    candidate.write_bytes(b"before")
    reference_path.parent.mkdir(parents=True)
    reference_path.write_text(json.dumps({"generation": {}}), encoding="utf-8")
    approved = json.loads(
        (ROOT / "tests/fixtures/prospective-acceptance-source.json").read_text(encoding="utf-8")
    )
    finding = {
        "category": "xml",
        "field": "eligibility/criteria/textblock",
        "target_ids": ["prs.structured"],
        "recovery_class": "deterministic_structure_defect",
        "action": "rebuild_deterministic_structure",
        "issue": "The structured qualifier differs from the approved source.",
    }

    def rebuild(_run_dir, **_kwargs):
        candidate.write_bytes(b"after")
        latest = json.loads(reference_path.read_text(encoding="utf-8"))
        workflow._complete_pending_deterministic_reconstructions(
            reference_path,
            latest,
            revision,
            outcome="rebuilt",
        )
        assert workflow._complete_pending_recovery_attempts(
            revision, reference_path, latest, require_candidate_change=True,
        ) == []
        return {"status": "awaiting_hermes", "stage": "independent_verification"}

    monkeypatch.setattr(workflow, "generate", rebuild)
    result = workflow._quality_retry(
        run_dir, reference_path, {"generation": {}}, approved, revision, {}, [finding], "xml",
    )

    assert result["status"] == "awaiting_hermes"
    records = json.loads(reference_path.read_text(encoding="utf-8"))["generation"]["deterministic_reconstructions"]
    assert len(records) == 1
    assert records[0]["target"] == "prs.structured"
    assert records[0]["candidate"] == "candidate/study.xml"
    assert records[0]["before_sha256"] != records[0]["after_sha256"]
    assert records[0]["candidate_bytes_changed"] is True
    attempt_manifest = json.loads(next((revision / "attempts").glob("*/attempt-manifest.json")).read_text(encoding="utf-8"))
    action = attempt_manifest["recovery_actions"][0]
    assert action["target"] == ["prs.structured"]
    assert action["candidate_bytes_changed"] is True
    assert action["deterministic_structure_changed"] is True
    assert isinstance(action["prompt_evidence_changed"], bool)


def test_pending_recovery_is_measured_before_resume_or_publication(tmp_path):
    run_dir = tmp_path / "run"
    revision = run_dir / "revisions/r-test"
    reference_path = run_dir / "reference/study.reference.json"
    reference_path.parent.mkdir(parents=True)
    working = {"generation": {}}
    reference_path.write_text(json.dumps(working), encoding="utf-8")
    finding = {
        "category": "visual",
        "field": "protocol",
        "artifact": "protocol",
        "check": "orphan_heading",
        "element": "5. INTRODUCTION",
        "target_ids": ["layout:protocol"],
        "recovery_class": "visual_defect",
        "action": "targeted_layout_repair",
        "issue": "orphan heading",
    }
    attempt = workflow._archive_failed_attempt(revision, "quality", [finding])
    journal = json.loads((revision / "gate-attempt-journal.json").read_text(encoding="utf-8"))
    # Simulate interruption after the journal commit but before the reference
    # learned about the attempt.

    workflow._reconcile_pending_recovery_attempts(revision, reference_path, working)

    manifest = json.loads((attempt / "attempt-manifest.json").read_text(encoding="utf-8"))
    assert manifest["recovery_actions"][0]["outcome_status"] == "interrupted_no_action"
    updated = json.loads(reference_path.read_text(encoding="utf-8"))["generation"]["gate_attempts"]
    workflow._validate_expected_gate_attempts(revision, updated)


@pytest.mark.parametrize(
    "fail_target",
    ["attempt-manifest.json", "gate-attempt-journal.json", "study.reference.json"],
)
def test_recovery_finalization_transaction_recovers_each_commit_boundary(tmp_path, monkeypatch, fail_target):
    revision = tmp_path / "revision"
    reference_path = tmp_path / "reference/study.reference.json"
    reference_path.parent.mkdir(parents=True)
    attempt = workflow._archive_failed_attempt(
        revision,
        "quality",
        [{
            "category": "visual",
            "artifact": "protocol",
            "check": "orphan_heading",
            "element": "5. INTRODUCTION",
            "target_ids": ["layout:protocol"],
            "recovery_class": "visual_defect",
            "action": "targeted_layout_repair",
            "issue": "orphan heading",
        }],
    )
    journal = json.loads((revision / "gate-attempt-journal.json").read_text(encoding="utf-8"))
    working_reference = {"generation": {"gate_attempts": journal["entries"]}}
    reference_path.write_text(json.dumps(working_reference), encoding="utf-8")
    original_write = workflow._write
    interrupted = False

    def fail_once(path, value):
        nonlocal interrupted
        if path.name == fail_target and not interrupted:
            interrupted = True
            raise OSError("simulated interruption")
        original_write(path, value)

    monkeypatch.setattr(workflow, "_write", fail_once)
    with pytest.raises(OSError, match="simulated interruption"):
        workflow._finalize_recovery_attempt(
            revision, attempt, reference_path, working_reference,
        )
    monkeypatch.setattr(workflow, "_write", original_write)

    resumed_reference = json.loads(reference_path.read_text(encoding="utf-8"))
    workflow._recover_recovery_finalization(
        revision, reference_path, resumed_reference,
    )

    assert not (revision / "recovery-finalization-transaction.json").exists()
    journal = json.loads((revision / "gate-attempt-journal.json").read_text(encoding="utf-8"))
    workflow._validate_expected_gate_attempts(revision, journal["entries"])


def test_recovery_archive_transaction_recovers_before_journal_commit(tmp_path, monkeypatch):
    revision = tmp_path / "revision"
    original_write = workflow._write
    interrupted = False

    def fail_journal_once(path, value):
        nonlocal interrupted
        if path.name == "gate-attempt-journal.json" and not interrupted:
            interrupted = True
            raise OSError("simulated archive interruption")
        original_write(path, value)

    monkeypatch.setattr(workflow, "_write", fail_journal_once)
    with pytest.raises(OSError, match="simulated archive interruption"):
        workflow._archive_failed_attempt(
            revision,
            "quality",
            [{
                "category": "visual",
                "artifact": "protocol",
                "check": "orphan_heading",
                "element": "5. INTRODUCTION",
                "target_ids": ["layout:protocol"],
                "recovery_class": "visual_defect",
                "action": "targeted_layout_repair",
                "issue": "orphan heading",
            }],
        )
    monkeypatch.setattr(workflow, "_write", original_write)

    workflow._recover_recovery_archive(revision)

    assert not (revision / "recovery-archive-transaction.json").exists()
    journal = json.loads((revision / "gate-attempt-journal.json").read_text(encoding="utf-8"))
    workflow._validate_expected_gate_attempts(
        revision, journal["entries"], require_measured=False,
    )


def test_recovery_attempts_are_recorded_without_a_stable_target_cap():
    prior = {
        "category": "visual",
        "field": "protocol",
        "artifact": "protocol",
        "check": "orphan_heading",
        "element": "5. INTRODUCTION",
        "target_ids": ["layout:protocol"],
        "recovery_class": "visual_defect",
        "action": "targeted_layout_repair",
        "issue": "orphan heading",
    }
    current = {**prior, "check": "bad_table_split", "element": "3. GENERAL INFORMATION"}
    prior_strategy = workflow._recovery_strategy_id(prior)

    attempts, strategies = workflow._advance_recovery_attempts(
        [current],
        {"layout:protocol": 3},
        {"layout:protocol": {prior_strategy: 2}},
    )

    assert attempts["layout:protocol"] == 4
    assert strategies["layout:protocol"][workflow._recovery_strategy_id(current)] == 2


def test_recovery_wave_increments_stable_target_once_across_multiple_strategies():
    base = {
        "category": "visual",
        "artifact": "protocol",
        "target_ids": ["layout:protocol"],
        "recovery_class": "visual_defect",
        "action": "targeted_layout_repair",
    }
    orphan_heading = {
        **base,
        "check": "orphan_heading",
        "element": "5. INTRODUCTION",
        "issue": "orphan heading",
    }
    split_table = {
        **base,
        "check": "bad_table_split",
        "element": "3. GENERAL INFORMATION",
        "issue": "split table",
    }

    attempts, strategies = workflow._advance_recovery_attempts(
        [orphan_heading, split_table],
        {"layout:protocol": 1},
        {},
    )

    assert attempts["layout:protocol"] == 2
    assert strategies["layout:protocol"] == {
        workflow._recovery_strategy_id(orphan_heading): 2,
        workflow._recovery_strategy_id(split_table): 2,
    }


def test_visual_strategy_identity_includes_the_exact_repair_element():
    base = {
        "category": "visual",
        "artifact": "protocol",
        "check": "bad_table_split",
        "target_ids": ["layout:protocol"],
        "recovery_class": "visual_defect",
        "action": "targeted_layout_repair",
        "issue": "split table",
    }

    section_three = workflow._recovery_strategy_id({**base, "element": "3. GENERAL INFORMATION"})
    section_fifteen = workflow._recovery_strategy_id({**base, "element": "15. STANDARD EVALUATION PROCEDURES"})

    assert section_three != section_fifteen


def test_workflow_binds_missing_identity_but_never_rewrites_mismatched_identity(tmp_path):
    revision = tmp_path / "revision"
    requests = revision / "hermes/verification-requests"
    responses = revision / "hermes/verification-responses"
    requests.mkdir(parents=True)
    responses.mkdir()
    request = {
        "schema_version": "hermes-verification/v1",
        "request_id": "r.verify.content",
        "task": "clinical_content_verification",
        "revision_id": revision.name,
        "review_set": 1,
        "response_path": "hermes/verification-responses/content.json",
        "sections": [],
        "cross_document_checks": [],
        "artifacts": [],
    }
    request["request_sha256"] = verification_request_sha256(request)
    (requests / f"{request['request_id']}.json").write_text(json.dumps(request), encoding="utf-8")
    _write_verification_ledger(revision, request)
    semantic = {
        "producer": {"model_id": "client-selected-model", "reviewer_id": "content-reviewer"},
        "status": "passed",
        "section_assessments": [],
        "cross_document_assessments": [],
    }
    response_path = responses / "content.json"
    response_path.write_text(json.dumps(semantic), encoding="utf-8")

    workflow._bind_available_verification_responses(revision)

    bound = json.loads(response_path.read_text(encoding="utf-8"))
    assert bound["request_id"] == request["request_id"]
    findings, _evidence = validate_verifications(revision)
    assert findings == []
    bound["request_id"] = "forged"
    response_path.write_text(json.dumps(bound), encoding="utf-8")
    workflow._bind_available_verification_responses(revision)
    assert json.loads(response_path.read_text(encoding="utf-8"))["request_id"] == "forged"
    findings, _evidence = validate_verifications(revision)
    mismatch = next(item for item in findings if item["field"] == "request_id")
    assert mismatch["recovery_class"] == "verifier_transient"
    bound["status"] = "failed"
    bound["findings"] = [{
        "artifact": "protocol",
        "check": "bad_table_split",
        "element": "3. GENERAL INFORMATION",
        "target_ids": ["layout:protocol"],
        "issue": "forged visual finding",
    }]
    response_path.write_text(json.dumps(bound), encoding="utf-8")
    malformed_findings, _evidence = validate_verifications(revision)
    assert {item["recovery_class"] for item in malformed_findings} == {"verifier_transient"}
    candidate = revision / "candidate/protocol.docx"
    candidate.parent.mkdir()
    candidate.write_bytes(b"unchanged candidate")
    reference_path = tmp_path / "reference/study.reference.json"
    reference_path.parent.mkdir()
    working = {"generation": {"review_set": 1}}
    reference_path.write_text(json.dumps(working), encoding="utf-8")

    result = workflow._quality_retry(
        tmp_path,
        reference_path,
        working,
        {},
        revision,
        {},
        [mismatch],
        "quality",
    )

    assert result["stage"] == "internal_recovery_routing"
    assert candidate.read_bytes() == b"unchanged candidate"
    state = json.loads(reference_path.read_text(encoding="utf-8"))["generation"]
    assert state["review_set"] == 1
    assert response_path.exists()
    assert not (revision / "attempts").exists()


def test_mixed_content_and_visual_findings_queue_both_repairs(tmp_path, monkeypatch):
    run_dir = tmp_path / "run"
    revision = run_dir / "revisions/r-test"
    reference_path = run_dir / "reference/study.reference.json"
    reference_path.parent.mkdir(parents=True)
    reference_path.write_text(json.dumps({"generation": {"review_set": 1}}), encoding="utf-8")
    (revision / "candidate-build.json").parent.mkdir(parents=True)
    (revision / "candidate-build.json").write_text("{}", encoding="utf-8")
    (revision / "candidate/protocol.docx").parent.mkdir()
    (revision / "candidate/protocol.docx").write_bytes(b"before layout repair")
    approved = json.loads(
        (ROOT / "tests/fixtures/retrospective-acceptance-source.json").read_text(encoding="utf-8")
    )

    def fake_schedule_requests(**_kwargs):
        path = revision / "hermes/requests/introduction.json"
        path.parent.mkdir(parents=True)
        path.write_text(json.dumps({
            "task": "draft_sections",
            "section_contracts": [{"section_id": "introduction"}],
            "response_path": "hermes/responses/introduction.json",
        }), encoding="utf-8")
        return [path]

    monkeypatch.setattr(workflow, "schedule_requests", fake_schedule_requests)
    result = workflow._quality_retry(
        run_dir,
        reference_path,
        {"generation": {"review_set": 1}},
        approved,
        revision,
        {},
        [
            {"category": "verification", "field": "introduction", "target_ids": ["introduction"], "recovery_class": "drafting_defect", "action": "retry_drafting_target", "issue": "missing context"},
            {"category": "visual", "field": "protocol", "artifact": "protocol", "check": "orphan_heading", "element": "5. INTRODUCTION", "target_ids": ["layout:protocol"], "recovery_class": "visual_defect", "action": "targeted_layout_repair", "issue": "orphan heading"},
        ],
        "quality",
    )

    assert result["stage"] == "drafting_retry"
    state = json.loads(reference_path.read_text(encoding="utf-8"))["generation"]
    assert state["review_set"] == 1
    assert state["pending_review_set_advance"]["from_review_set"] == 1
    assert state["pending_layout_artifacts"] == ["protocol"]
    assert state["layout_repairs"] == {
        "protocol": [{"rule": "heading_cohesion", "target": "5. INTRODUCTION"}],
    }
    manifest_path = next((revision / "attempts").glob("*/attempt-manifest.json"))
    actions = json.loads(manifest_path.read_text(encoding="utf-8"))["recovery_actions"]
    assert {item["outcome_status"] for item in actions} == {"pending"}

    (revision / "candidate/protocol.docx").write_bytes(b"rebuilt after both repairs")
    (revision / "candidate-build.json").write_text('{"fingerprint":"rebuilt"}', encoding="utf-8")
    latest = json.loads(reference_path.read_text(encoding="utf-8"))
    assert workflow._complete_pending_recovery_attempts(
        revision,
        reference_path,
        latest,
        require_candidate_change=True,
    ) == []
    workflow._advance_pending_review_set(reference_path, latest)
    actions = json.loads(manifest_path.read_text(encoding="utf-8"))["recovery_actions"]
    drafting_action = next(item for item in actions if item["triggering_finding"]["recovery_class"] == "drafting_defect")
    visual_action = next(item for item in actions if item["triggering_finding"]["recovery_class"] == "visual_defect")
    assert drafting_action["prompt_evidence_changed"] is True
    assert drafting_action["candidate_bytes_changed"] is True
    assert drafting_action["deterministic_structure_changed"] is False
    assert visual_action["candidate_bytes_changed"] is True
    assert visual_action["deterministic_structure_changed"] is True
    assert json.loads(reference_path.read_text(encoding="utf-8"))["generation"]["review_set"] == 2


def test_unchanged_rebuilt_candidate_cannot_advance_the_review_set(tmp_path):
    run_dir = tmp_path / "run"
    revision = run_dir / "revisions/r-test"
    reference_path = run_dir / "reference/study.reference.json"
    reference_path.parent.mkdir(parents=True)
    candidate = revision / "candidate/protocol.docx"
    candidate.parent.mkdir(parents=True)
    candidate.write_bytes(b"same candidate")
    working = {"generation": {"review_set": 1}}
    reference_path.write_text(json.dumps(working), encoding="utf-8")
    finding = {
        "category": "visual",
        "artifact": "protocol",
        "check": "orphan_heading",
        "element": "5. INTRODUCTION",
        "target_ids": ["layout:protocol"],
        "recovery_class": "visual_defect",
        "action": "targeted_layout_repair",
        "issue": "orphan heading",
    }
    attempt = workflow._archive_failed_attempt(revision, "quality", [finding])
    journal = json.loads((revision / "gate-attempt-journal.json").read_text(encoding="utf-8"))
    working["generation"].update({
        "gate_attempts": journal["entries"],
        "pending_recovery_attempts": [attempt.relative_to(revision).as_posix()],
        "pending_review_set_advance": {
            "from_review_set": 1,
            "strategy_ids": [workflow._recovery_strategy_id(finding)],
            "finding_sha256": workflow.canonical_evidence_sha256([finding]),
        },
    })
    reference_path.write_text(json.dumps(working), encoding="utf-8")

    no_progress = workflow._complete_pending_recovery_attempts(
        revision,
        reference_path,
        working,
        require_candidate_change=True,
    )

    assert len(no_progress) == 1
    state = json.loads(reference_path.read_text(encoding="utf-8"))["generation"]
    assert state["review_set"] == 1
    assert "pending_review_set_advance" not in state


def test_terminal_timeout_measures_pending_recovery_without_advancing_review(tmp_path):
    revision = tmp_path / "revision"
    reference_path = tmp_path / "reference/study.reference.json"
    reference_path.parent.mkdir(parents=True)
    candidate = revision / "candidate/protocol.docx"
    candidate.parent.mkdir(parents=True)
    candidate.write_bytes(b"rebuilt candidate before timeout")
    finding = {
        "category": "visual",
        "artifact": "protocol",
        "check": "orphan_heading",
        "element": "5. INTRODUCTION",
        "target_ids": ["layout:protocol"],
        "recovery_class": "visual_defect",
        "action": "targeted_layout_repair",
        "issue": "orphan heading",
    }
    attempt = workflow._archive_failed_attempt(revision, "quality", [finding])
    journal = json.loads((revision / "gate-attempt-journal.json").read_text(encoding="utf-8"))
    working = {"generation": {
        "review_set": 1,
        "gate_attempts": journal["entries"],
        "pending_recovery_attempts": [attempt.relative_to(revision).as_posix()],
        "pending_review_set_advance": {
            "from_review_set": 1,
            "strategy_ids": [workflow._recovery_strategy_id(finding)],
            "finding_sha256": workflow.canonical_evidence_sha256([finding]),
        },
    }}
    reference_path.write_text(json.dumps(working), encoding="utf-8")

    workflow._finalize_pending_recovery_for_terminal(
        revision, reference_path, working,
    )

    action = json.loads((attempt / "attempt-manifest.json").read_text(encoding="utf-8"))["recovery_actions"][0]
    state = json.loads(reference_path.read_text(encoding="utf-8"))["generation"]
    assert action["outcome_status"] == "terminal_measured"
    assert isinstance(action["candidate_bytes_changed"], bool)
    assert isinstance(action["deterministic_structure_changed"], bool)
    assert isinstance(action["prompt_evidence_changed"], bool)
    assert state["review_set"] == 1
    assert "pending_review_set_advance" not in state


def test_timeline_coverage_allows_grammar_but_preserves_milestone_pairing():
    approved = (
        "IRB review and data access: Month 1; extraction and abstraction: Months 2 to 4; "
        "quality control and analysis: Months 5 to 6; final report: Month 7."
    )
    grammatical = (
        "The approved study timeline is IRB review and data access in Month 1; extraction and abstraction "
        "in Months 2 to 4; quality control and analysis in Months 5 to 6; and the final report in Month 7."
    )
    swapped = grammatical.replace("Month 1", "Month 7", 1)

    assert quality._timeline_covered(approved, grammatical)
    assert not quality._timeline_covered(approved, swapped)


def test_transient_verifier_failure_is_retried_with_a_bounded_counter(tmp_path):
    run_dir = tmp_path / "run"
    revision = run_dir / "revisions/r-test"
    requests = revision / "hermes/verification-requests"
    requests.mkdir(parents=True)
    request_id = "r-test.review-1.verify.content"
    request = _verification_request_payload(
        request_id,
        "clinical_content_verification",
    )
    (requests / f"{request['request_id']}.json").write_text(json.dumps(request), encoding="utf-8")
    _write_verification_ledger(revision, request)
    reference_path = run_dir / "reference/study.reference.json"
    reference_path.parent.mkdir(parents=True)
    working = {"generation": {"review_set": 1}}
    reference_path.write_text(json.dumps(working), encoding="utf-8")
    finding = {"category": "reviewer-transient", "field": "clinical_content_verification", "target_ids": ["verification:content"], "verification_request_id": request_id, "recovery_class": "verifier_transient", "action": "retry_verifier", "issue": "temporary"}

    result = workflow._quality_retry(run_dir, reference_path, working, _source(), revision, {}, [finding], "quality")

    assert result["status"] == "awaiting_hermes"
    assert result["stage"] == "independent_verification_retry"
    state = json.loads(reference_path.read_text(encoding="utf-8"))
    assert state["generation"]["verification_attempts"][request_id] == 1
    assert not (revision / request["response_path"]).exists()


def test_multiple_transient_findings_from_one_reviewer_consume_one_retry(tmp_path):
    run_dir = tmp_path / "run"
    revision = run_dir / "revisions/r-test"
    reference_path = run_dir / "reference/study.reference.json"
    reference_path.parent.mkdir(parents=True)
    working = {"generation": {"review_set": 1}}
    reference_path.write_text(json.dumps(working), encoding="utf-8")
    request_id = "r-test.review-1.verify.content"
    request_path = revision / "hermes/verification-requests/content.json"
    request_path.parent.mkdir(parents=True)
    request = _verification_request_payload(
        request_id,
        "clinical_content_verification",
    )
    request_path.write_text(json.dumps(request), encoding="utf-8")
    _write_verification_ledger(revision, request)
    response_path = revision / request["response_path"]
    response_path.parent.mkdir(parents=True, exist_ok=True)
    response_path.write_text("{}", encoding="utf-8")
    finding = {
        "category": "reviewer-transient",
        "field": "clinical_content_verification",
        "target_ids": ["verification:content"],
        "verification_request_id": request_id,
        "recovery_class": "verifier_transient",
        "action": "retry_verifier",
        "issue": "incomplete response",
    }

    result = workflow._quality_retry(
        run_dir,
        reference_path,
        working,
        _source(),
        revision,
        {},
        [finding, {**finding, "field": "section_assessments"}],
        "quality",
    )

    assert result["status"] == "awaiting_hermes"
    state = json.loads(reference_path.read_text(encoding="utf-8"))
    assert state["generation"]["verification_attempts"][request_id] == 1
    assert not response_path.exists()


def test_quality_retry_blocks_for_a_finding_without_governed_recovery_routing(tmp_path):
    run_dir = tmp_path / "run"
    revision = run_dir / "revisions/r-test"
    reference_path = run_dir / "reference/study.reference.json"
    reference_path.parent.mkdir(parents=True)
    working = {"generation": {"review_set": 1}}
    reference_path.write_text(json.dumps(working), encoding="utf-8")
    finding = {"category": "visual", "field": "protocol", "issue": "free text must not select recovery"}

    result = workflow._quality_retry(
        run_dir,
        reference_path,
        working,
        {},
        revision,
        {},
        [finding],
        "quality",
    )

    assert result["status"] == "blocked"
    assert result["stage"] == "internal_recovery_routing"
    assert result["findings"][0]["routing_failure_code"] == "invalid_recovery_contract"
    assert result["findings"][0]["original_finding"] == finding
    assert result["preserved_companion_findings"] == []
    assert working["generation"]["review_set"] == 1
    assert not (revision / "attempts").exists()
    assert not (revision / "gate-attempt-journal.json").exists()
    assert not (revision / "hermes/verification-requests").exists()


def test_quality_retry_blocks_before_archive_when_recovery_class_has_wrong_target_type(tmp_path):
    run_dir = tmp_path / "run"
    revision = run_dir / "revisions/r-test"
    reference_path = run_dir / "reference/study.reference.json"
    reference_path.parent.mkdir(parents=True)
    working = {"generation": {"review_set": 1}}
    reference_path.write_text(json.dumps(working), encoding="utf-8")
    finding = {
        "category": "verification",
        "field": "content",
        "target_ids": ["verification:content"],
        "recovery_class": "visual_defect",
        "action": "targeted_layout_repair",
        "issue": "category prose must not override the Recovery Class",
    }

    result = workflow._quality_retry(
        run_dir,
        reference_path,
        working,
        {},
        revision,
        {},
        [finding],
        "quality",
    )

    assert result["status"] == "blocked"
    assert result["stage"] == "internal_recovery_routing"
    assert result["findings"][0]["original_finding"] == finding
    assert working["generation"]["review_set"] == 1
    assert not (revision / "attempts").exists()
    assert not (revision / "gate-attempt-journal.json").exists()


def test_layout_failure_rebuilds_and_reverifies_before_retry_limit(tmp_path, monkeypatch):
    run_dir = tmp_path / "run"
    revision = run_dir / "revisions/r-test"
    reference_path = run_dir / "reference/study.reference.json"
    verification_response = revision / "hermes/verification-responses/visual.json"
    reference_path.parent.mkdir(parents=True)
    verification_response.parent.mkdir(parents=True)
    verification_request = revision / "hermes/verification-requests/visual.json"
    verification_request.parent.mkdir(parents=True)
    reference_path.write_text(json.dumps({"generation": {}}), encoding="utf-8")
    verification_response.write_text("{}", encoding="utf-8")
    verification_request.write_text(json.dumps({
        "task": "rendered_page_visual_verification",
        "response_path": "hermes/verification-responses/visual.json",
        "artifacts": [{"artifact": "protocol"}],
    }), encoding="utf-8")
    (revision / "candidate-build.json").write_text("{}", encoding="utf-8")
    rerun = {"status": "awaiting_hermes", "stage": "independent_verification"}
    monkeypatch.setattr(workflow, "generate", lambda _run_dir, **_kwargs: rerun)

    result = workflow._quality_retry(
        run_dir,
        reference_path,
        {"generation": {}},
        _source(),
        revision,
        {},
        [{"category": "visual", "field": "protocol.docx:10", "artifact": "protocol", "check": "orphan_heading", "element": "5. INTRODUCTION", "target_ids": ["layout:protocol.docx"], "recovery_class": "visual_defect", "action": "targeted_layout_repair", "issue": "orphan heading"}],
        "quality",
    )

    assert result == rerun
    state = json.loads(reference_path.read_text(encoding="utf-8"))
    assert state["generation"]["attempts"]["layout:protocol.docx"] == 2
    assert (revision / "candidate-build.json").exists()
    assert not verification_request.exists()
    assert not verification_response.exists()


def test_layout_failure_does_not_exhaust_the_same_strategy_after_three_attempts():
    finding = {"category": "visual", "field": "protocol.docx:10", "artifact": "protocol", "check": "orphan_heading", "element": "5. INTRODUCTION", "target_ids": ["layout:protocol.docx"], "recovery_class": "visual_defect", "action": "targeted_layout_repair", "issue": "orphan heading"}
    strategy_id = workflow._recovery_strategy_id(finding)

    attempts, strategies = workflow._advance_recovery_attempts(
        [finding],
        {"layout:protocol.docx": 3},
        {"layout:protocol.docx": {strategy_id: 3}},
    )

    assert attempts["layout:protocol.docx"] == 4
    assert strategies["layout:protocol.docx"][strategy_id] == 4


def test_prs_generation_validates_against_the_retained_client_manual_authority(tmp_path):
    reference = json.loads((ROOT / "tests/fixtures/prospective-acceptance-source.json").read_text(encoding="utf-8"))
    report = generate_xml(
        ROOT / "assets/client-templates/prs/clinicaltrials_prs_full_placeholder_template.xml",
        tmp_path / "study.xml",
        reference,
        {"brief_summary": {"text": "Approved summary."}, "detailed_description": {"text": "Approved description."}},
        structural_template=ROOT / "assets/client-templates/reference/prs-manual-reference.xml",
    )
    assert report["status"] == "passed"


def test_refresh_toc_uses_layout_extraction_to_find_body_heading(tmp_path, monkeypatch):
    document = Document()
    document.styles.add_style("toc 1", WD_STYLE_TYPE.PARAGRAPH)
    document.add_paragraph("4. TABLE OF CONTENTS", style="Heading 1")
    document.add_paragraph("15. STANDARD EVALUATION PROCEDURES\t5", style="toc 1")
    document.add_paragraph("15. STANDARD EVALUATION PROCEDURES", style="Heading 1")
    docx = tmp_path / "protocol.docx"
    document.save(docx)
    pdf = tmp_path / "protocol.pdf"
    pdf.write_bytes(b"pdf")

    class Page:
        def __init__(self, plain, layout):
            self.plain = plain
            self.layout = layout

        def extract_text(self, extraction_mode=None):
            return self.layout if extraction_mode == "layout" else self.plain

    class Reader:
        def __init__(self, _path):
            self.pages = [
                Page("4. TABLE OF CONTENTS", "4. TABLE OF CONTENTS"),
                Page("15. STANDARD EVALUATION PROCEDURES", "15. STANDARD EVALUATION PROCEDURES"),
                *[Page("", "") for _ in range(7)],
                Page("", "15. STANDARD EVALUATION PROCEDURES"),
            ]

    monkeypatch.setattr(rendering, "PdfReader", Reader)
    assert rendering.refresh_toc_from_pdf(docx, pdf)
    refreshed = Document(docx)
    toc = next(
        paragraph.text
        for paragraph in refreshed.paragraphs
        if paragraph.style.name == "toc 1" and paragraph.text.startswith("15.")
    )
    assert toc.endswith("\t10")


def test_protocol_omits_references_heading_when_no_references_are_supplied():
    document = Document()
    document.add_paragraph("REFERENCES", style="Heading 1")
    document.add_paragraph("")
    authority = Document()
    authority.add_paragraph("REFERENCES", style="Heading 1")

    rendering._ensure_protocol_references(document, authority, {})

    assert "REFERENCES" not in [paragraph.text.strip() for paragraph in document.paragraphs]


def test_layout_failure_repairs_only_the_affected_artifact_but_resets_all_review_evidence(tmp_path, monkeypatch):
    run_dir = tmp_path / "run"
    revision = run_dir / "revisions/r-test"
    reference_path = run_dir / "reference/study.reference.json"
    reference_path.parent.mkdir(parents=True)
    working = {"generation": {}}
    reference_path.write_text(json.dumps(working), encoding="utf-8")
    retained = (
        "candidate/protocol.docx",
        "candidate/icf.docx",
        "rendered/protocol.pdf",
        "rendered/protocol/page-1.png",
        "rendered/icf.pdf",
        "rendered/icf/page-1.png",
    )
    for relative in retained:
        path = revision / relative
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_bytes(relative.encode())
    requests = revision / "hermes/verification-requests"
    responses = revision / "hermes/verification-responses"
    requests.mkdir(parents=True)
    responses.mkdir(parents=True)
    for artifact in ("protocol", "icf"):
        request = {
            "task": "rendered_page_visual_verification",
            "request_id": f"visual-{artifact}",
            "response_path": f"hermes/verification-responses/visual-{artifact}.json",
            "artifacts": [{"artifact": artifact}],
        }
        (requests / f"visual-{artifact}.json").write_text(json.dumps(request), encoding="utf-8")
        (responses / f"visual-{artifact}.json").write_text("{}", encoding="utf-8")
    (revision / "candidate-build.json").write_text("{}", encoding="utf-8")
    resumed = {"status": "awaiting_hermes", "stage": "independent_verification"}
    monkeypatch.setattr(workflow, "generate", lambda _, **_kwargs: resumed)

    result = workflow._quality_retry(
        run_dir,
        reference_path,
        working,
        {"meta": {"study_type": "Retrospective"}},
        revision,
        {},
        [{"category": "visual", "field": "protocol", "artifact": "protocol", "check": "orphan_heading", "element": "5. INTRODUCTION", "target_ids": ["layout:protocol"], "recovery_class": "visual_defect", "action": "targeted_layout_repair", "issue": "orphaned heading"}],
        "quality",
    )

    assert result == resumed
    state = json.loads(reference_path.read_text(encoding="utf-8"))
    assert state["generation"]["attempts"]["layout:protocol"] == 2
    assert state["generation"]["layout_repairs"] == {
        "protocol": [{"rule": "heading_cohesion", "target": "5. INTRODUCTION"}],
    }
    assert not (revision / "candidate/protocol.docx").exists()
    assert not (revision / "rendered/protocol.pdf").exists()
    assert not (revision / "rendered/protocol").exists()
    assert not (requests / "visual-protocol.json").exists()
    assert not (responses / "visual-protocol.json").exists()
    assert (revision / "candidate/icf.docx").is_file()
    assert (revision / "rendered/icf.pdf").is_file()
    assert (revision / "rendered/icf/page-1.png").is_file()
    assert not (requests / "visual-icf.json").exists()
    assert not (responses / "visual-icf.json").exists()
    assert (revision / "candidate-build.json").is_file()


def test_protocol_renders_assignment_endpoint_hierarchy_and_both_operational_tables(tmp_path):
    reference = _source()
    reference["design"]["assignment_method"] = (
        "Subjects will be assigned 1:1 to TRYPTYR or control using the approved randomization schedule."
    )
    reference["endpoints"]["other"] = [
        {"category": "Powered exploratory endpoints in hierarchical order", "label": "Schirmer change versus control", "time_point": "Month 12"},
        {"category": "Descriptive exploratory endpoints", "label": "Ocular discomfort score", "time_point": "Months 1, 3, 6, 9, and 12"},
    ]
    visit_names = ["Baseline", "Day 14", "Month 1", "Month 3", "Month 6", "Month 9", "Month 12"]
    activities = [f"Assessment {index}" for index in range(1, 20)]
    # This synthetic case replaces the inherited three-visit study with seven
    # visits. Remove its obsolete parallel table rather than supply two studies'
    # schedules and expect the renderer to discard approved source rows.
    reference["procedures"].pop("visit_schedule_table", None)
    reference["procedures"]["visit_schedule"] = [
        {"visit": visit, "timing": visit, "procedures": activities if index == 0 else activities[index::3]}
        for index, visit in enumerate(visit_names)
    ]
    reference["statistics"]["sample_size_evidence"] = [
        {"study": "COMET-2", "timepoint": "Day 28", "mean_change_ods_vas": "-25.20", "se": "1.96", "estimated_sd": "30.0"},
        {"study": "COMET-3", "timepoint": "Day 90", "mean_change_ods_vas": "-29.60", "se": "1.97", "estimated_sd": "30.0"},
    ]
    reference["population"]["sample_size_evidence"] = []

    render_documents(ROOT, tmp_path, reference, {"protocol": [], "icf": {}, "prs": {}})

    document = Document(tmp_path / "candidate/protocol.docx")
    visible = _visible(document)
    headings = [p.text for p in document.paragraphs if p.style.name.casefold().startswith("heading")]
    assessment = next(table for table in document.tables if table.rows[0].cells[0].text.strip() == "Activity")
    sample_evidence = next(table for table in document.tables if table.rows[0].cells[0].text.strip() == "Study")

    assert "8.3. Method of Assigning Subjects to Treatment Arms" in headings
    assert reference["design"]["assignment_method"] in visible
    assert "Powered exploratory endpoints in hierarchical order" in visible
    assert "Descriptive exploratory endpoints" in visible
    assert len(assessment.rows) == 21
    assert len(assessment.columns) == 8
    assert len(sample_evidence.rows) == 3
    assert len(sample_evidence.columns) == 5
    guarded_fields = {"study-design.assignment", "procedures.visit_schedule", "statistics.sample_size_evidence"}
    assert not guarded_fields & {
        item["field"] for item in deterministic_content_check(tmp_path, reference)
    }

    assessment.cell(2, 1).text, assessment.cell(2, 2).text = assessment.cell(2, 2).text, assessment.cell(2, 1).text
    sample_evidence.cell(1, 0).text = "Changed study"
    document.save(tmp_path / "candidate/protocol.docx")
    guarded_findings = {
        item["field"] for item in deterministic_content_check(tmp_path, reference)
    }
    assert "procedures.visit_schedule" in guarded_findings
    assert "statistics.sample_size_evidence" in guarded_findings


def test_retrospective_quality_does_not_require_prospective_table_contracts(tmp_path):
    reference = json.loads(
        (ROOT / "tests/fixtures/release-certification/retrospective/approved-reference.json").read_text(
            encoding="utf-8"
        )
    )

    render_documents(ROOT, tmp_path, reference, {"protocol": [], "icf": {}, "prs": {}})

    fields = {item["field"] for item in deterministic_content_check(tmp_path, reference)}

    assert "procedures.visit_schedule" not in fields
    assert "statistics.sample_size_evidence" not in fields


def test_protocol_does_not_duplicate_a_model_echoed_visit_table_caption_or_details(tmp_path):
    reference = _source()
    model = {
        "protocol": [{
            "section_id": "study-procedure.visits",
            "paragraphs": [
                {"text": "Participants complete the approved visits in sequence."},
                {"text": "Table 9.2-1. Visit Schedule"},
                {"text": "Unique source-bound visit procedure detail."},
            ],
            "lists": [],
        }],
        "icf": {},
        "prs": {},
    }

    render_documents(ROOT, tmp_path, reference, model)

    document = Document(tmp_path / "candidate/protocol.docx")
    paragraphs = [paragraph.text.strip() for paragraph in document.paragraphs]
    assert paragraphs.count("Table 9.2-1. Visit Schedule") == 1
    assert paragraphs.count("Unique source-bound visit procedure detail.") == 1
    assert not any(
        "Table 9.2-1. Visit Schedule" in paragraph
        and "Unique source-bound visit procedure detail." in paragraph
        for paragraph in paragraphs
    )


def test_protocol_fidelity_gate_rejects_loss_of_operational_source_detail(tmp_path):
    reference = _source()
    reference["procedures"].update({
        "assessment_details": "Each visit includes ODS-VAS and unanesthetized Schirmer testing.",
        "intervention_management": "TRYPTYR 0.003% is administered twice daily and reconciled at every visit.",
        "discontinuation": "Participants may withdraw at any time without penalty.",
        "replacement": "Participants discontinued during enrollment will be replaced.",
    })
    reference["statistics"].update({
        "analysis_populations": "The intent-to-treat and per-protocol populations will be analyzed.",
        "methodology": "A mixed model for repeated measures will compare change from baseline.",
        "software": "Analyses will use R version 4.4.2.",
    })
    reference.setdefault("confidentiality", {})["retention"] = "Study records will be retained for 15 years after study closure."
    reference["risks_benefits"].update({
        "injury_handling": "Research-related injuries will receive immediate evaluation by the investigator.",
        "risks": "Transient ocular burning and privacy loss are foreseeable risks.",
        "benefits": "Participants may experience improved tear production, but benefit is not guaranteed.",
        "compensation_or_reimbursement": "Participants will receive $50 for each completed visit.",
    })
    model = {
        "protocol": [
            {"section_id": "study-procedure.visits", "paragraphs": [{"text": f'{reference["procedures"]["assessment_details"]} {reference["procedures"]["intervention_management"]}'}], "lists": []},
            {"section_id": "analysis-plan.datasets", "paragraphs": [{"text": reference["statistics"]["analysis_populations"]}], "lists": []},
            {"section_id": "analysis-plan.methodology", "paragraphs": [{"text": reference["statistics"]["methodology"]}], "lists": []},
            {"section_id": "analysis-plan.considerations", "paragraphs": [{"text": reference["statistics"]["software"]}], "lists": []},
            {"section_id": "confidentiality-publication", "paragraphs": [{"text": reference["confidentiality"]["retention"]}], "lists": []},
            {"section_id": "financial-injury", "paragraphs": [{"text": reference["risks_benefits"]["injury_handling"]}], "lists": []},
            {"section_id": "endpoint-criteria.discontinuation", "paragraphs": [{"text": f'{reference["procedures"]["discontinuation"]} {reference["procedures"]["replacement"]}'}], "lists": []},
            {"section_id": "risks-benefits.risks", "paragraphs": [{"text": reference["risks_benefits"]["risks"]}], "lists": []},
            {"section_id": "risks-benefits.benefits", "paragraphs": [{"text": f'{reference["risks_benefits"]["benefits"]} {reference["risks_benefits"]["compensation_or_reimbursement"]}'}], "lists": []},
        ],
        "icf": {},
        "prs": {},
    }

    render_documents(ROOT, tmp_path, reference, model)

    guarded = {
        "study-procedure.visits",
        "analysis-plan.datasets",
        "analysis-plan.methodology",
        "analysis-plan.considerations",
        "confidentiality-publication",
        "financial-injury",
        "endpoint-criteria.discontinuation",
        "risks-benefits.risks",
        "risks-benefits.benefits",
    }
    assert not guarded & {item["field"] for item in deterministic_content_check(tmp_path, reference)}

    document = Document(tmp_path / "candidate/protocol.docx")
    heading_titles = {
        "9.2. Visits and Examinations",
        "10.1. Analysis Data Sets",
        "10.2. Statistical Methodology",
        "10.3. General Statistical Considerations",
        "12. CONFIDENTIALITY/PUBLICATION OF THE STUDY",
        "17. FINANCIAL AND INSURANCE INFORMATION/STUDY RELATED INJURIES",
        "18.2. Patient Discontinuation",
        "19.1. Summary of risks",
        "19.2. Summary of benefits",
    }
    for index, paragraph in enumerate(document.paragraphs):
        if paragraph.text.strip() not in heading_titles:
            continue
        for target in document.paragraphs[index + 1:]:
            if target.style.name.casefold().startswith("heading"):
                break
            if target.text.strip():
                target.text = "Generic text that omits every approved operational qualifier."
                break
    document.save(tmp_path / "candidate/protocol.docx")

    assert guarded <= {item["field"] for item in deterministic_content_check(tmp_path, reference)}
