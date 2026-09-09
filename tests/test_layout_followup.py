"""Preserve supplied roles and meaningful ICF choice-block content."""
from pathlib import Path

import pytest
from lxml import etree
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

import rendering

ROOT = Path(__file__).resolve().parents[1]


def test_heading_cohesion_keeps_a_short_lead_in_with_its_first_list_item():
    document = Document()
    heading = document.add_paragraph('6.2. Inclusion/Exclusion Criteria', style='Heading 2')
    label = document.add_paragraph('Inclusion criteria:')
    first_item = document.add_paragraph('Age 18 years or older.', style='List Bullet')

    rendering._repair_heading_cohesion(
        document,
        heading.text,
        protocol=True,
    )

    assert heading.paragraph_format.keep_with_next is True
    assert label.paragraph_format.keep_with_next is True
    assert label.paragraph_format.keep_together is True
    assert first_item.paragraph_format.keep_with_next is not True


@pytest.mark.parametrize('funding_address', [None, '10 Main St, Country'])
def test_same_entity_sponsor_retains_funding_role(funding_address):
    funder = {'name': 'Sponsor Foundation'}
    if funding_address is not None:
        funder['address'] = funding_address
    fields = rendering.render_fields({'parties': {
        'sponsor': {'name': 'Sponsor Foundation', 'address': '10 Main St, Country'},
        'funding_source': funder,
    }}, {})
    doc = Document(ROOT / 'assets/client-templates/docx/prospective-protocol.template.docx')
    row = next(row for table in doc.tables for row in table.rows
               if any('{fundingSourceName}' in cell.text for cell in row.cells))
    for cell in row.cells:
        for paragraph in cell.paragraphs:
            rendering._replace_paragraph(paragraph, fields)
    text = '\n'.join(cell.text for cell in row.cells)
    assert 'Funding source: Sponsor' in text
    assert text.count('Sponsor Foundation') == 1
    assert text.count('10 Main St, Country') == 1
    separate = doc.add_paragraph('{fundingSourceName}')
    rendering._replace_paragraph(separate, fields)
    assert separate.text == 'Sponsor Foundation'


@pytest.mark.parametrize('content', ['simple_field', 'signing_spaces', 'page_break_before'])
def test_icf_preferences_retain_meaningful_blank_looking_paragraphs(content, tmp_path):
    doc = Document()
    doc.add_paragraph('Check your preference below:')
    yes = doc.add_paragraph('Yes, inform my primary care physician')
    cosmetic = doc.add_paragraph('   ')
    retained = doc.add_paragraph()
    if content == 'simple_field':
        field = OxmlElement('w:fldSimple')
        field.set(qn('w:instr'), 'DOCPROPERTY PatientInitials')
        run = OxmlElement('w:r')
        text = OxmlElement('w:t')
        text.text = 'AB'
        run.append(text)
        field.append(run)
        retained._p.append(field)
    elif content == 'signing_spaces':
        retained.add_run('          ').underline = True
    else:
        retained.paragraph_format.page_break_before = True
    no = doc.add_paragraph('No, do not inform my primary care physician')
    runs_before = [etree.tostring(node) for node in retained._p if node.tag != qn('w:pPr')]

    rendering._normalize_icf_preferences(doc)

    assert retained._p.getparent() is doc.element.body
    assert cosmetic._p.getparent() is None
    assert [etree.tostring(node) for node in retained._p if node.tag != qn('w:pPr')] == runs_before
    assert yes.text.startswith('☐ ') and no.text.startswith('☐ ')
    assert retained.paragraph_format.keep_together is True
    assert retained.paragraph_format.keep_with_next is True
    xml = doc.element.xml
    rendering._normalize_icf_preferences(doc)
    assert doc.element.xml == xml
    path = tmp_path / 'preferences.docx'
    doc.save(path)
    restored = Document(path).paragraphs[2]
    if content == 'simple_field':
        assert restored._p.xpath('./w:fldSimple/@w:instr') == ['DOCPROPERTY PatientInitials']
        assert restored._p.xpath('./w:fldSimple//w:t/text()') == ['AB']
    elif content == 'signing_spaces':
        assert restored.runs[0].text == '          '
        assert restored.runs[0].underline is True
    else:
        assert restored.paragraph_format.page_break_before is True
