"""Explicit source decisions survive intake and Sterling consent assembly."""
import json
from pathlib import Path

from docx import Document

from contracts import source_contract
from rendering import render_documents

ROOT = Path(__file__).resolve().parents[1]
DISCLOSURE_START = "A description of this clinical trial will be available"


def source():
    reference = json.loads((ROOT / "tests/fixtures/prospective-acceptance-source.json").read_text())
    reference["meta"]["icf_template"] = "Sterling"
    return reference


def render_icf(tmp_path, reference):
    render_documents(ROOT, tmp_path, reference, {"protocol": [], "icf": {}, "prs": {}}, artifact_names={"icf"})
    return Document(tmp_path / "candidate/icf.docx")


def test_existing_approved_disclosure_alias_is_included_once(tmp_path):
    reference = source()
    reference.setdefault("regulatory", {})["clinicaltrials_consent_disclosure_applies"] = (
        "Yes; sponsor requires public trial registration before enrollment and standard disclosure, no assigned NCT number."
    )
    document = render_icf(tmp_path, reference)
    assert sum(p.text.startswith(DISCLOSURE_START) for p in document.paragraphs) == 1


def test_prepare_exposes_normalized_disclosure_for_owner_review(tmp_path):
    import workflow
    reference = source()
    reference.setdefault("regulatory", {})["clinicaltrials_gov_disclosure_required"] = "Yes"
    path = tmp_path / "reference/study.reference.json"
    path.parent.mkdir(parents=True)
    path.write_text(json.dumps(reference))
    result = workflow.prepare(tmp_path)
    assert result["status"] == "awaiting_approval"
    prepared = json.loads(path.read_text())
    assert prepared["regulatory"]["prs"]["participant_registry_disclosure"] is True
    assert "<!-- field: regulatory.prs.participant_registry_disclosure -->\nTrue" in (tmp_path / result["source_of_truth"]).read_text()


def test_conflicting_decisions_require_clarification_and_do_not_render(tmp_path):
    import pytest
    reference = source()
    reference.setdefault("regulatory", {}).setdefault("prs", {})["participant_registry_disclosure"] = False
    reference["regulatory"]["clinicaltrials_consent_disclosure_applies"] = "Yes; requested by sponsor"
    result = source_contract(reference)
    assert result["status"] == "blocked"
    assert any("Conflicting participant registry disclosure" in f["issue"] for f in result["source_gaps"])
    with pytest.raises(ValueError, match="Conflicting participant registry disclosure"):
        render_icf(tmp_path, reference)
    assert not (tmp_path / "candidate/icf.docx").exists()


def test_negative_alias_keeps_disclosure_out_even_for_interventional_xml(tmp_path):
    reference = source()
    reference.setdefault("regulatory", {})["clinicaltrials_gov_disclosure_required"] = "No; not required under this plan"
    reference["regulatory"].setdefault("prs", {})["study_type"] = "Interventional"
    assert source_contract(reference)["normalized_reference"]["regulatory"]["prs"]["participant_registry_disclosure"] is False
    document = render_icf(tmp_path, reference)
    assert not any(p.text.startswith(DISCLOSURE_START) for p in document.paragraphs)


def test_missing_decision_does_not_infer_disclosure_from_registration_or_xml(tmp_path):
    reference = source()
    reference.setdefault("regulatory", {}).setdefault("prs", {})["study_type"] = "Interventional"
    reference["regulatory"]["registration_required"] = True
    assert source_contract(reference)["normalized_reference"]["regulatory"]["prs"].get("participant_registry_disclosure") is None
    assert not any(p.text.startswith(DISCLOSURE_START) for p in render_icf(tmp_path, reference).paragraphs)


def test_ambiguous_structured_decision_requires_intake_clarification():
    reference = source()
    reference.setdefault("regulatory", {})["clinicaltrials_consent_disclosure_applies"] = "Only if the sponsor later decides to register"
    result = source_contract(reference)
    assert result["status"] == "blocked"
    assert any(f["field"] == "regulatory.prs.participant_registry_disclosure" for f in result["source_gaps"])


def test_boolean_strings_and_matching_aliases_do_not_change_approved_source(tmp_path):
    import copy
    reference = source()
    reference.setdefault("regulatory", {}).setdefault("prs", {})["participant_registry_disclosure"] = "true"
    reference["regulatory"]["clinicaltrials_gov_disclosure_required"] = True
    before = copy.deepcopy(reference)
    document = render_icf(tmp_path, reference)
    authority = Document(ROOT / "assets/client-templates/reference/sterling-icf-reference.docx")
    exact = next(p.text for p in authority.paragraphs if p.text.startswith(DISCLOSURE_START))
    assert [p.text for p in document.paragraphs if p.text.startswith(DISCLOSURE_START)] == [exact]
    assert reference == before
    paragraph = next(p for p in document.paragraphs if p.text == exact)
    assert not paragraph._p.xpath(".//w:instrText | .//w:fldChar")
