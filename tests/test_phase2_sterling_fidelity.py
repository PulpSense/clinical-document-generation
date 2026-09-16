import copy
import json
from pathlib import Path

from docx import Document

import contracts
import quality
import rendering


ROOT = Path(__file__).resolve().parents[1]


def source():
    value = json.loads(
        (ROOT / "tests/fixtures/prospective-acceptance-source.json").read_text(
            encoding="utf-8"
        )
    )
    value["meta"]["icf_template"] = "Sterling"
    return value


def document_with_sections(sections):
    document = Document()
    for heading, paragraphs in sections:
        document.add_heading(heading, level=1)
        for paragraph in paragraphs:
            document.add_paragraph(paragraph)
    return document


def codes(report):
    return {item["code"] for item in report["findings"]}


def clause(report, clause_id):
    return next(item for item in report["findings"] if item.get("clause_id") == clause_id)


def test_missing_mandatory_sterling_clause_is_blocking():
    governed = contracts.sterling_clause_contract(ROOT)
    assert len(governed["clauses"]) == 28
    assert {item["classification"] for item in governed["clauses"]} == {
        "mandatory", "conditional", "source_dependent", "prohibited_from_invention"
    }
    assert all(
        set(item["severity"]) == {"absent", "altered", "unsupported", "misplaced"}
        and set(item["severity"].values()) == {"blocking"}
        for item in governed["clauses"]
    )

    document = document_with_sections([
        ("VOLUNTARY PARTICIPATION/WITHDRAWAL", []),
    ])

    report = quality.validate_sterling_clause_contract(document, source())

    finding = clause(report, "sterling.voluntary.core")
    assert finding["code"] == "sterling-clause-missing"
    assert finding["severity"] == "blocking"
    assert finding["target_ids"] == ["icf.voluntary-participation"]


def test_conditional_sterling_clause_is_rejected_without_trigger():
    document = document_with_sections([
        ("GENETIC INFORMATION NONDISCRIMINATION ACT", [
            "A federal law called GINA generally makes it illegal to discriminate based on genetic information."
        ]),
    ])

    report = quality.validate_sterling_clause_contract(document, source())

    finding = clause(report, "sterling.privacy.gina")
    assert finding["code"] == "sterling-clause-unsupported"
    assert finding["severity"] == "blocking"


def test_conditional_sterling_clause_is_required_when_trigger_exists():
    reference = source()
    reference.setdefault("regulatory", {})["genetic_information_collected"] = True
    document = document_with_sections([
        ("GENETIC INFORMATION NONDISCRIMINATION ACT", []),
    ])

    report = quality.validate_sterling_clause_contract(document, reference)

    finding = clause(report, "sterling.privacy.gina")
    assert finding["code"] == "sterling-clause-missing"
    assert finding["severity"] == "blocking"


def test_triggered_sterling_conditional_clause_is_inserted_by_renderer():
    reference = source()
    reference.setdefault("regulatory", {})["genetic_information_collected"] = True
    document = document_with_sections([("QUESTIONS", ["Questions body."])])

    rendering._insert_sterling_conditional_clauses(document, reference)
    report = quality.validate_sterling_clause_contract(document, reference)

    assert any(
        paragraph.text == "GENETIC INFORMATION NONDISCRIMINATION ACT"
        for paragraph in document.paragraphs
    )
    assert not any(
        item.get("clause_id") == "sterling.privacy.gina"
        for item in report["findings"]
    )


def test_source_dependent_payment_cannot_be_invented():
    reference = source()
    reference["risks_benefits"]["compensation_or_reimbursement"] = None
    document = document_with_sections([
        ("COMPENSATION TO YOU", ["You will receive $100 for each completed visit."]),
    ])

    report = quality.validate_sterling_clause_contract(document, reference)

    finding = clause(report, "sterling.compensation.terms")
    assert finding["code"] == "sterling-clause-unsupported"
    assert finding["severity"] == "blocking"


def test_required_sterling_clause_in_wrong_section_is_repairable_and_exactly_localized():
    voluntary = contracts.sterling_clause_text("sterling.voluntary.core")
    document = document_with_sections([
        ("COSTS TO YOU", [voluntary]),
        ("VOLUNTARY PARTICIPATION/WITHDRAWAL", []),
    ])

    report = quality.validate_sterling_clause_contract(document, source())

    finding = clause(report, "sterling.voluntary.core")
    assert finding["code"] == "sterling-clause-misplaced"
    assert finding["severity"] == "blocking"
    assert finding["recovery_class"] == "deterministic_structure_defect"
    assert finding["expected_section"] == "VOLUNTARY PARTICIPATION/WITHDRAWAL"
    assert finding["actual_section"] == "COSTS TO YOU"


def test_materially_weakened_voluntariness_and_participant_rights_are_blocking():
    weakened = "Taking part is voluntary. You may leave the study at any time."
    document = document_with_sections([
        ("VOLUNTARY PARTICIPATION/WITHDRAWAL", [weakened]),
        ("QUESTIONS", ["Contact the study staff with questions."]),
    ])

    report = quality.validate_sterling_clause_contract(document, source())

    assert clause(report, "sterling.voluntary.core")["code"] == "sterling-clause-weakened"
    rights = clause(report, "sterling.contact.participant-rights")
    assert rights["severity"] == "blocking"
    assert rights["safety_critical"] is True


def test_contradictory_safety_language_cannot_satisfy_keyword_validation():
    document = document_with_sections([
        ("POTENTIAL RISKS, EFFECTS, DISCOMFORTS, INCONVENIENCES", [
            "There are no foreseeable risks or discomforts because this study is perfectly safe for everyone forever."
        ]),
    ])

    report = quality.validate_sterling_clause_contract(document, source())

    finding = clause(report, "sterling.risks.foreseeable")
    assert finding["code"] == "sterling-clause-weakened"
    assert finding["contradiction"] is True


def test_unrelated_heading_cannot_lend_body_text_to_governed_section():
    voluntary = contracts.sterling_clause_text("sterling.voluntary.core")
    document = document_with_sections([
        ("VOLUNTARY PARTICIPATION/WITHDRAWAL", []),
        ("UNRELATED SECTION", [voluntary]),
    ])

    report = quality.validate_sterling_clause_contract(document, source())

    finding = clause(report, "sterling.voluntary.core")
    assert finding["code"] == "sterling-clause-missing"


def test_governed_body_language_inside_a_table_is_visible_to_validation():
    document = Document()
    document.add_heading("INFORMATION", level=1)
    table = document.add_table(rows=1, cols=1)
    table.cell(0, 0).text = contracts.sterling_clause_text(
        "sterling.information.new-findings"
    )

    report = quality.validate_sterling_clause_contract(document, source())

    assert not any(
        item.get("clause_id") == "sterling.information.new-findings"
        for item in report["findings"]
    )


def test_sterling_validation_checks_required_body_not_only_headings_and_signatures():
    document = document_with_sections([
        ("INFORMATION", []),
        ("VOLUNTARY PARTICIPATION/WITHDRAWAL", []),
        ("QUESTIONS", []),
        ("PARTICIPANT STATEMENT AUTHORIZATION", ["Signature of Participant Date"]),
    ])

    report = quality.validate_sterling_clause_contract(document, source())

    assert "sterling-clause-missing" in codes(report)
    assert {item["clause_id"] for item in report["findings"]} >= {
        "sterling.information.new-findings",
        "sterling.voluntary.core",
        "sterling.contact.participant-rights",
        "sterling.authorization.participant-statement",
    }


def test_recoverable_sterling_placement_is_repaired_and_revalidated():
    voluntary = contracts.sterling_clause_text("sterling.voluntary.core")
    document = document_with_sections([
        ("COSTS TO YOU", [voluntary]),
        ("VOLUNTARY PARTICIPATION/WITHDRAWAL", []),
    ])

    repair = rendering.repair_sterling_clause_placement(document, source())
    report = quality.validate_sterling_clause_contract(document, source())

    assert repair["status"] == "repaired"
    assert repair["revalidated"] is True
    assert not any(
        item.get("clause_id") == "sterling.voluntary.core"
        for item in report["findings"]
    )


def test_icf_background_and_purpose_material_concept_overlap_targets_purpose():
    background = (
        "The PureSee non-diffractive lens may reduce dysphotopsias and improve contrast sensitivity. "
        "A randomized trial found better intermediate and near vision than a monofocal lens, and the "
        "mix-and-match rationale is a broader range of vision with fewer disturbances."
    )
    purpose = (
        "PureSee is a non-diffractive lens that may reduce dysphotopsias and improve contrast sensitivity. "
        "A randomized trial found better intermediate and near vision than a monofocal lens. The rationale "
        "for mix-and-match implantation is broader vision with fewer disturbances before testing the hypothesis."
    )
    document = document_with_sections([
        ("BACKGROUND", [background]),
        ("PURPOSE", [purpose]),
    ])

    report = quality.assess_icf_output(document, source(), {"family": "Sterling"})

    finding = next(item for item in report["findings"] if item["code"] == "icf-concept-repetition")
    assert finding["primary_section"] == "icf.background"
    assert finding["secondary_section"] == "icf.study-purpose"
    assert finding["target_ids"] == ["icf.study-purpose"]
    assert finding["publication_disposition"] == "blocking"


def test_short_key_information_summary_to_detail_overlap_is_allowed():
    document = document_with_sections([
        ("KEY INFORMATION", ["The study measures vision three months after surgery."]),
        ("PURPOSE", [
            "The purpose is to test the hypothesis that the lens combination produces strong visual outcomes; "
            "the primary endpoint is corrected intermediate visual acuity at 66 cm three months after surgery."
        ]),
        ("DURATION", ["Your participation lasts through the three-month postoperative visit."]),
    ])

    report = quality.assess_icf_output(document, source(), {"family": "Sterling"})

    assert not any(item["code"] == "icf-concept-repetition" for item in report["findings"])


def test_protocol_endpoint_inventory_repetition_targets_secondary_sections():
    inventory = (
        "The primary endpoint is intermediate visual acuity at month 3. Secondary endpoints are corrected "
        "distance visual acuity, uncorrected intermediate visual acuity, near visual acuity, the defocus curve, "
        "and each questionnaire response at month 3."
    )
    sections = {
        "objectives": [inventory],
        "study-procedure.measurements": [
            inventory + " Each measurement is collected using the scheduled study assessments."
        ],
        "analysis-plan.datasets": [
            inventory + " These observations enter the analysis data set."
        ],
        "analysis-plan.methodology": [
            inventory + " Continuous values are summarized with means and standard deviations."
        ],
    }

    findings = quality.assess_protocol_concept_repetition(
        sections, contracts.protocol_concept_ownership("Prospective")
    )

    assert {item["secondary_section"] for item in findings} >= {
        "study-procedure.measurements",
        "analysis-plan.datasets",
        "analysis-plan.methodology",
    }
    assert all(item["target_ids"] == [item["secondary_section"]] for item in findings)
    assert all(item["publication_disposition"] == "blocking" for item in findings)


def test_legitimate_endpoint_references_with_section_specific_roles_are_allowed():
    sections = {
        "objectives": ["The primary endpoint is intermediate visual acuity at month 3."],
        "study-procedure.measurements": [
            "Intermediate visual acuity is measured at 66 cm during the month 3 visit."
        ],
        "analysis-plan.datasets": [
            "Month 3 observations from eyes completing the study enter the analysis data set."
        ],
        "analysis-plan.methodology": [
            "The primary endpoint will be summarized as a continuous variable as specified in Section 6."
        ],
    }

    assert quality.assess_protocol_concept_repetition(
        sections, contracts.protocol_concept_ownership("Prospective")
    ) == []


def test_safety_and_participant_rights_omissions_never_receive_generic_warning():
    for target in (
        "icf.risks",
        "icf.injury",
        "icf.voluntary-participation",
        "icf.contact.participant-rights",
        "icf.authorization.participant-statement",
    ):
        finding = {
            "category": "content",
            "check": "substantive",
            "target_ids": [target],
            "safety_critical": False,
            "contradiction": False,
            "obscures_required_information": False,
            "materially_unusable": False,
        }
        assert quality._governed_content_omission(finding, [target]) is False


def test_exact_substantive_protocol_duplication_remains_blocking(tmp_path):
    duplicate = (
        "The primary endpoint is intermediate visual acuity at month 3 and all secondary endpoints are "
        "evaluated at that same postoperative visit using the specified study assessments."
    )
    document = Document()
    contract = {item.section_id: item for item in contracts.protocol_contract("Prospective")}
    for section_id in ("objectives", "study-procedure.measurements"):
        spec = contract[section_id]
        document.add_heading(f"{spec.number} {spec.title}", level=1)
        document.add_paragraph(duplicate)
    candidate = tmp_path / "candidate"
    candidate.mkdir()
    document.save(candidate / "protocol.docx")
    Document().save(candidate / "icf.docx")

    findings = quality.deterministic_content_check(
        tmp_path, {"meta": {"study_type": "Prospective", "icf_template": "Sterling"}}
    )

    finding = next(item for item in findings if item.get("code") == "protocol-exact-repetition")
    assert finding["publication_disposition"] == "blocking"
    assert finding["recovery_class"] == "drafting_defect"
    assert finding["target_ids"] == ["study-procedure.measurements"]
