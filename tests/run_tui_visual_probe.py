"""Owner-authorized one synthetic callback/fallback probe; never a clinical run."""
import argparse
from datetime import datetime, timezone
import hashlib
import json
import os
from pathlib import Path
import subprocess
import sys
import time
from unittest.mock import patch

from docx import Document
from pypdf import PdfWriter
from pypdf.generic import DictionaryObject, NameObject, DecodedStreamObject


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('--run-dir', type=Path, required=True)
    parser.add_argument('--release-root', type=Path, required=True)
    parser.add_argument('--candidate-identity', type=Path, required=True)
    parser.add_argument('--authorize-paid-probe', action='store_true')
    args = parser.parse_args()
    if not args.authorize_paid_probe:
        parser.error('Explicit owner authorization required')
    REPO = args.release_root.resolve()
    sys.path.insert(0, str(REPO / 'scripts'))
    import workflow
    import quality
    from visual_review import visual_review_plan, create_visual_review_unit, callback_release_binding
    from tui_visual_callback import observe_unit
    identity = json.loads(args.candidate_identity.read_text())
    binding = callback_release_binding(REPO)
    if any(identity.get(k) != binding['release_identity'].get(k) for k in ('package_fingerprint', 'git_commit')):
        parser.error('Frozen identity differs from intact installed candidate')
    run = args.run_dir.absolute()
    if run.resolve().is_relative_to(REPO):
        parser.error('Probe evidence must remain outside installed candidate')
    run.mkdir(parents=True)  # Retain every first outcome; never overwrite a probe.
    revision = run / 'revisions/synthetic-visual-r1'
    (revision / 'candidate').mkdir(parents=True)
    (revision / 'rendered/protocol').mkdir(parents=True)
    lines = ['Synthetic Visual Tooling Probe', 'This is a one-page tooling fixture, not a clinical study.',
        'There are no patient records, study facts, or client deliverables.',
        'The page tests genuine image inspection and authenticated transport.',
        'End of synthetic visual fixture.']
    document = Document()
    document.add_heading(lines[0], 0)
    for line in lines[1:]: document.add_paragraph(line)
    docx = revision / 'candidate/protocol.docx'
    document.save(docx)
    writer = PdfWriter()
    page = writer.add_blank_page(width=612, height=792)
    font = writer._add_object(DictionaryObject({NameObject('/Type'): NameObject('/Font'),
        NameObject('/Subtype'): NameObject('/Type1'), NameObject('/BaseFont'): NameObject('/Helvetica')}))
    page[NameObject('/Resources')] = DictionaryObject({NameObject('/Font'): DictionaryObject({NameObject('/F1'): font})})
    stream = DecodedStreamObject()
    commands = ['BT /F1 14 Tf 60 720 Td']
    for i, line in enumerate(lines):
        if i: commands.append('0 -34 Td')
        commands.append('(' + line + ') Tj')
    commands.append('ET')
    stream.set_data('\n'.join(commands).encode())
    page[NameObject('/Contents')] = writer._add_object(stream)
    pdf = revision / 'rendered/protocol.pdf'
    writer.write(pdf)
    image = revision / 'rendered/protocol/page-1.png'
    painter = "import json, sys; from PIL import Image, ImageDraw, ImageFont; p, font, lines = json.loads(sys.argv[1]); im=Image.new('RGB',(1224,1584),'white'); d=ImageDraw.Draw(im); f=ImageFont.truetype(font,28); [d.text((120,144+68*i),line,font=f,fill='black') for i,line in enumerate(lines)]; im.save(p)"
    subprocess.run(['/opt/hermes/.venv/bin/python', '-B', '-c', painter,
        json.dumps([str(image), str(REPO / 'assets/fallback-fonts/LiberationSans-Regular.ttf'), lines])], check=True, timeout=10)
    source = {'meta': {'study_type': 'Retrospective'}, 'study': {'title': lines[0]}}
    request_id = revision.name + '.review-1.verify.visual.protocol'
    request = {'schema_version': quality.VERIFY_SCHEMA, 'request_id': request_id,
        'revision_id': revision.name, 'review_set': 1, 'task': 'rendered_page_visual_verification',
        'callback_release_binding': binding,
        'source_identity_sha256': quality.canonical_evidence_sha256(source),
        'generation_authority_sha256': hashlib.sha256(b'owner-authorized synthetic tooling probe').hexdigest(),
        'checks': list(quality.VISUAL_CHECKS),
        'instructions': 'This is a deliberately short, one-page synthetic tooling fixture. The declared layout is simple readable black text on white, with no branding, header/footer or TOC requirement. Judge actual applicability; natural short-page space is allowed. No clinical approval is requested.',
        'artifacts': [{'artifact': 'protocol', 'docx': 'candidate/protocol.docx', 'docx_sha256': quality.sha256_file(docx),
            'pdf': 'rendered/protocol.pdf', 'pdf_sha256': quality.sha256_file(pdf),
            'pages': [{'page': 1, 'path': 'rendered/protocol/page-1.png', 'sha256': quality.sha256_file(image)}]}],
        'response_path': f'hermes/verification-responses/{request_id}.json'}
    request['visual_review_plan'] = visual_review_plan(request)
    request['request_sha256'] = quality.verification_request_sha256(request)
    path = revision / f'hermes/verification-requests/{request_id}.json'
    path.parent.mkdir(parents=True)
    path.write_text(json.dumps(request, indent=2) + '\n')
    ledger = revision / f'request-ledger/{request_id}.json'
    ledger.parent.mkdir(parents=True)
    ledger.write_text(json.dumps(quality.verification_request_ledger_record(path, request)))
    handoff = {'task': request['task'], 'request_id': request_id, 'request_sha256': request['request_sha256'],
        'request_path': path.relative_to(revision).as_posix(), 'response_path': request['response_path'], 'fallback_owner': 'parent'}
    envelope = {'profile_path': os.environ['HERMES_HOME'], 'configuration': workflow.CERTIFIED_HERMES_CONFIGURATION,
        'callback_release_binding': binding}
    callback = workflow.command_parent_visual_reviewer(REPO / 'scripts/tui_visual_callback', skill_root=REPO)
    fallback_calls = [0]
    delegate_failures = []
    def delegated(_items, _budget):
        unit_path = create_visual_review_unit(revision, path, request['visual_review_plan']['targets'][:1], 1,
            observer_evidence_required=True)
        unit = json.loads(unit_path.read_text())
        item = {'request_path': unit_path.relative_to(revision).as_posix(), 'response_path': unit['response_path'],
            'task': unit['task'], 'visual_unit': True}
        try:
            # Deliberate synthetic injection; deployed 240s ceiling is unchanged.
            observe_unit(revision, item, 0.3, envelope)
        except Exception as exc:
            delegate_failures.append({'error_type': type(exc).__name__, 'cause_type': type(exc.__cause__).__name__ if exc.__cause__ else None})
            raise
    def fallback(items, budget):
        fallback_calls[0] += 1
        if fallback_calls[0] != 1:
            raise RuntimeError('Probe admits one paid callback only; retain first failure without a paid retry')
        callback(items, budget, revision, workflow.CERTIFIED_HERMES_CONFIGURATION)
    def generate(*_args, **_kwargs):
        if quality.verification_response_is_terminal(revision, path):
            return {'status': 'blocked', 'stage': 'synthetic_probe_done', 'client_outputs': [],
                'findings': [{'category': 'synthetic', 'issue': 'Tooling proof only; no clinical documents are offered.'}]}
        return {'status': 'awaiting_hermes', 'stage': 'independent_verification', 'revision_id': revision.name,
            'handoffs': [handoff]}
    started = time.monotonic()
    failure = None
    try:
        with patch.object(workflow, 'generate', generate):
            operation = workflow.run_desktop_operation(run, handoff_runner=delegated, fallback_handoff_runner=fallback,
                opener=lambda _: b'unused', budget_seconds=600,
                release_identity={**identity, 'host_mode': 'hermes-tui',
                    'callback_release_binding': binding,
                    'parent_visual_callback': {'path': str(REPO / 'scripts/tui_visual_callback'),
                        'sha256': binding['source_sha256']['scripts/tui_visual_callback']}})
    except Exception as exc:
        failure = {'error_type': type(exc).__name__, 'message': str(exc)}
        operation = None
    response_path = revision / request['response_path']
    response = json.loads(response_path.read_text()) if response_path.is_file() else None
    record = {'schema_version': 'tui-visual-callback-probe/v1', 'scope': 'synthetic callback transport; no clinical approval',
        'candidate_identity': identity, 'parent_request_path': path.relative_to(revision).as_posix(),
        'delegated_timeout_observed': any(d.get('error_type') in {'TimeoutExpired', 'TimeoutError'} for d in delegate_failures),
        'fallback_callback_completed': failure is None and fallback_calls[0] == 1,
        'started_at_utc': datetime.now(timezone.utc).isoformat(), 'elapsed_seconds': time.monotonic() - started,
        'callback_path': str(REPO / 'scripts/tui_visual_callback'), 'callback_sha256': quality.sha256_file(REPO / 'scripts/tui_visual_callback'),
        'callback_module_sha256': quality.sha256_file(REPO / 'scripts/tui_visual_callback.py'),
        'transport_module_sha256': quality.sha256_file(REPO / 'scripts/visual_review.py'),
        'run_dir': str(run), 'revision_path': str(revision), 'request_path': str(path),
        'request_sha256': quality.sha256_file(path), 'response_path': str(response_path),
        'response_sha256': quality.sha256_file(response_path) if response else None,
        'response_accepted': quality.verification_response_is_terminal(revision, path),
        'response_complete_pass': quality.verification_response_is_complete(revision, path),
        'actual_observers': [p.get('observer_evidence') for p in (response or {}).get('unit_proofs', [])],
        'delegated_failures': delegate_failures, 'callback_invocations': min(fallback_calls[0], 1),
        'operation': operation, 'exception': failure}
    (run / 'probe-result.json').write_text(json.dumps(record, indent=2) + '\n')
    print(json.dumps({k: record[k] for k in ('response_accepted', 'response_complete_pass', 'elapsed_seconds', 'delegated_failures', 'callback_invocations', 'exception')}, indent=2))
    return 0 if record['response_accepted'] and record['response_complete_pass'] else 1


if __name__ == '__main__':
    raise SystemExit(main())
