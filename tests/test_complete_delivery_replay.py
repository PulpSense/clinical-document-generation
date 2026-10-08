"""Synthetic minimized blocked-wording replay, not captured archive or fresh review.

Real ingestion, construction, exact-artifact checks and Desktop publication run.
Controlled responders provide no independent clinical or visual certification.
"""
import copy
import json
from pathlib import Path
from xml.etree import ElementTree as ET

from docx import Document
import pytest

import drafting
import quality
import workflow
from test_concept_reference_grounding import ANALYSIS_PLAN
from test_handoff_quality import acceptance_verification

ROOT = Path(__file__).resolve().parents[1]
TEXTS = (
    "Results will be interpreted descriptively to characterize the observed outcomes.",
    "Results will be interpreted within the descriptive scope of the approved analysis plan.",
)


def _run(tmp_path, text, mutation=None):
    run = tmp_path / "controlled-replay"
    reference_path = run / "reference/study.reference.json"
    reference_path.parent.mkdir(parents=True)
    source = json.loads((ROOT / "tests/fixtures/prospective-acceptance-source.json").read_text())
    source["meta"]["icf_template"] = "Sterling"
    source["statistics"].update(analysis_plan=ANALYSIS_PLAN, software=None)
    if mutation == "quantity":
        source["statistics"]["analysis_plan"] = ANALYSIS_PLAN + " Interpret results cautiously because the sample size is 40."
        text = "Interpret results cautiously because the sample size is 40."
    source["approval"] = {"status": "draft"}
    reference_path.write_text(json.dumps(source))
    assert workflow.prepare(run)["status"] == "awaiting_approval"
    assert workflow.approve(run, approved_by="Synthetic controlled replay")["status"] == "passed"
    approved = json.loads(reference_path.read_text())
    revision = run / "revisions" / approved["approval"]["revision_id"]
    approval = copy.deepcopy(approved["approval"])
    observed = {"draft_attempts": 0, "review_cautions": [], "blocked_before_repair": False}

    def runner(handoffs, _remaining):
        for handoff in handoffs:
            request_path = revision / handoff["request_path"]
            request = json.loads(request_path.read_text())
            if "verification-requests" in handoff["request_path"]:
                if request["task"] == "clinical_content_verification":
                    observed["review_cautions"].extend(request["lexical_review_cautions"])
                workflow._save_verification_response(revision, request_path, acceptance_verification)
                continue
            response = drafting.recorded_acceptance_response(request)
            for item in list(response.get("section_results", [])):
                if item["section_id"] == "analysis-plan.methodology":
                    item["paragraphs"][0]["text"] = source["statistics"]["analysis_plan"]
                if item["section_id"] != "analysis-plan.considerations":
                    continue
                observed["draft_attempts"] += 1
                item["paragraphs"] = [{"text": text, "evidence_refs": ["source:statistics.analysis_plan"], "boilerplate_refs": []}]
                item["lists"] = []
                if mutation and observed["draft_attempts"] == 1:
                    if mutation == "quantity":
                        item["paragraphs"][0]["text"] = "Interpret results cautiously because the sample size is 99."
                    elif mutation == "missing-section":
                        response["section_results"].remove(item)
                    elif mutation == "binding":
                        response["request_sha256"] = "0" * 64
                    # Observe real validation before handing the defect to the lifecycle.
                    accepted, findings = drafting.validate_response(request, response)
                    assert "analysis-plan.considerations" not in [d["section_id"] for d in (accepted or {}).get("drafts", [])]
                    assert findings
                    assert not (run / "output").exists()
                    observed["blocked_before_repair"] = True
            path = revision / request["response_path"]
            path.parent.mkdir(parents=True, exist_ok=True)
            path.write_text(json.dumps(response))

    result = workflow.run_desktop_operation(
        run, handoff_runner=runner, opener=lambda path: Path(path).read_bytes(),
        operation_id="synthetic-blocked-wording", release_identity={"package_fingerprint": "controlled-source-test"},
    )
    assert result["status"] == "passed", result
    assert result["stage"] == "desktop_delivery"
    assert result["delivery"]["confirmed"] is True
    assert sorted(result["client_outputs"]) == ["output/icf.docx", "output/protocol.docx", "output/study.xml"]
    assert json.loads(reference_path.read_text())["approval"] == approval
    state = json.loads((run / "logs/desktop-operation-synthetic-blocked-wording.json").read_text())
    assert state["deadline_at_epoch"] - state["started_at_epoch"] == 2700
    manifest = json.loads((revision / "delivery-manifest.json").read_text())
    for item in manifest["client_outputs"]:
        delivered = run / item["path"]
        assert delivered.read_bytes() == (revision / "candidate" / delivered.name).read_bytes()
        assert quality.sha256_file(delivered) == item["sha256"]
    protocol = Document(run / "output/protocol.docx")
    body = "\n".join(p.text for p in protocol.paragraphs)
    assert text in body
    assert "lexical-evidence-uncertainty" not in body
    assert source["study"]["title"] in body
    table_text = "\n".join(cell.text for table in protocol.tables for row in table.rows for cell in row.cells)
    assert "40 participants" in table_text
    assert "Primary outcome" in body + table_text
    assert "Month 3" in body + table_text
    xml = ET.parse(run / "output/study.xml")
    assert next(xml.getroot().iter("enrollment")).text == "40"
    consent = Document(run / "output/icf.docx")
    consent_text = "\n".join(p.text for p in consent.paragraphs).casefold()
    assert "voluntary" in consent_text
    assert "lexical-evidence-uncertainty" not in consent_text
    assert quality.deterministic_content_check(revision, source) == [] or all(
        f.get("publication_disposition") == "warning"
        for f in quality.deterministic_content_check(revision, source)
    )
    return result, observed, manifest


@pytest.mark.parametrize("text", TEXTS)
def test_synthetic_blocked_wording_reaches_complete_desktop_delivery(tmp_path, governed_pdfium, text):
    result, observed, manifest = _run(tmp_path, text)
    assert observed["draft_attempts"] == 1
    assert observed["review_cautions"]
    assert manifest["quality"]["resolved_evidence_cautions"]
    assert any(c.get("code") == "lexical-evidence-uncertainty" for c in result["desktop_reply"]["cautions"])


@pytest.mark.parametrize("mutation", ["quantity", "missing-section", "binding"])
def test_essential_mutation_requires_real_repair_before_delivery(tmp_path, governed_pdfium, mutation):
    _, observed, _ = _run(tmp_path, TEXTS[0], mutation)
    assert observed["blocked_before_repair"]
    assert observed["draft_attempts"] >= 2
