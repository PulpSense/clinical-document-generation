"""Exercise assembled artifacts across branches and source representations."""
import copy
import json
from pathlib import Path
from xml.etree import ElementTree as ET

import pytest
from docx import Document
from docx.oxml.ns import qn

import contracts
import prs_xml
import rendering
import quality

ROOT = Path(__file__).resolve().parents[1]
COMBINATIONS = [('Retrospective', None), ('Prospective', 'Advarra'),
                ('Prospective', 'Sterling'), ('Ambispective', 'Advarra'),
                ('Ambispective', 'Sterling')]


def source_for(branch, family):
    source = json.loads((ROOT / 'tests/fixtures/prospective-acceptance-source.json').read_text())
    source['meta']['study_type'] = branch
    if family:
        source['meta']['icf_template'] = family
    source['study']['title'] = 'Patient outcomes following treatment across participating clinical research facilities'
    source['study'].pop('short_title', None)
    source['study']['references'] = ['Alpha AB. Outcomes. Example Journal. 2024;1:1-9.',
                                     'Beta CD. Follow-up. Example Journal. 2025;2:10-19.']
    source['design']['study_design'] = f'Observational, {branch}, None; bilateral; investigator initiated study.'
    source['design']['control_arm'] = 'None'
    source['design'].pop('arms', None)
    source['population']['sample_size'] = '48 subjects (treatment); control: 0'
    source['parties']['principal_investigator']['name'] = 'Dr. Morgan Sample'
    source['sites'][0]['investigator'] = {'name': 'Professor Avery Sample', 'title': 'MD'}
    source['sites'][0].pop('investigators', None)
    source['procedures'].pop('visit_schedule_table', None)
    source['procedures']['visit_schedule'] = ['Screening', 'Week 8 follow-up']
    source['procedures']['assessments'] = 'Screening: history, medications. Week 8: questionnaire, examination.'
    return source


@pytest.mark.parametrize('branch,family', COMBINATIONS)
def test_public_assembly_preserves_references_and_uses_readable_fields(tmp_path, branch, family):
    source = source_for(branch, family)
    original = copy.deepcopy(source)
    names = {'protocol'} if branch == 'Retrospective' else {'protocol', 'icf'}
    model = {'protocol': [], 'icf': {
        'icf.background': {'paragraphs': [{'text': 'Published outcomes inform this study (references 1 and 2).'}], 'lists': []},
        'icf.duration': {'paragraphs': [{'text': 'Participation lasts 8 weeks. The study is expected to include 48 subjects (treatment); control: 0.'}], 'lists': []},
    }, 'prs': {}}
    rendering.render_documents(ROOT, tmp_path, source, model, artifact_names=names)
    protocol = Document(tmp_path / 'candidate/protocol.docx')
    prose = '\n'.join(p.text for p in protocol.paragraphs)
    assert all(ref in prose for ref in source['study']['references'])
    tables = [[[c.text for c in row.cells] for row in t.rows] for t in protocol.tables]
    assert not any('None;' in c or 'control: 0' in c for t in tables for row in t for c in row)
    if branch != 'Retrospective':
        inventory = next(t for t in tables if t[0][0] == 'Approved visit or assessment')
        assert inventory == [['Approved visit or assessment'], ['Screening'], ['Week 8 follow-up']]
        assert source['procedures']['assessments'] in prose
        inventory_table = next(t for t in protocol.tables if t.cell(0, 0).text == 'Approved visit or assessment')
        for row in inventory_table.rows[1:]:
            assert row.cells[0]._tc.tcPr.find(qn('w:shd')) is None
            assert all(run.bold is not True for paragraph in row.cells[0].paragraphs for run in paragraph.runs)
        visit_table = next(t for t in tables if t[0][0].replace(' ', '') == 'VisitNumber')
        assert all('Visit Window' not in row for row in visit_table)
    if family:
        icf = Document(tmp_path / 'candidate/icf.docx')
        text = '\n'.join(p.text for p in icf.paragraphs)
        assert 'control: 0' not in text
        assert '48 participants' in text
        assert '48 participants. There is no control group.' in text
        assert 'references 1 and 2' not in text
        if family == 'Sterling':
            assert 'Published outcomes inform this study' in text
    assert source == original


@pytest.mark.parametrize('branch,family', COMBINATIONS)
def test_supplied_bibliography_loss_is_construction_owned(tmp_path, branch, family):
    source = source_for(branch, family)
    rendering.render_documents(ROOT, tmp_path, source, {'protocol': [], 'icf': {}, 'prs': {}}, artifact_names={'protocol'})
    path = tmp_path / 'candidate/protocol.docx'
    assert not any(f.get('source_path') == 'study.references' for f in quality.audit_source_surfaces(path, source))
    document = Document(path)
    for paragraph in document.paragraphs:
        if 'Alpha AB.' in paragraph.text:
            paragraph.text = ''
    document.save(path)
    findings = [f for f in quality.audit_source_surfaces(path, source) if f.get('source_path') == 'study.references']
    assert findings and all(f['owner'] == 'construction' for f in findings)


def test_multiline_bibliography_entries_keep_their_citation_numbers(tmp_path):
    source = source_for('Retrospective', None)
    source['study']['references'] = [
        'Alpha AB. Outcomes. Example Journal. 2024;1:1-9.\ndoi:10.0000/example-alpha',
        '',
        'Beta CD. Follow-up. Example Journal. 2025;2:10-19.',
    ]
    rendering.render_documents(ROOT, tmp_path, source, {'protocol': []}, artifact_names={'protocol'})
    document = Document(tmp_path / 'candidate/protocol.docx')
    bibliography = next(p.text for p in document.paragraphs if 'Alpha AB.' in p.text)
    assert bibliography.startswith('1. Alpha AB.')
    assert '\n2. Beta CD.' in bibliography
    assert '\n2. doi:' not in bibliography
    assert not any(f.get('source_path') == 'study.references' for f in quality.audit_source_surfaces(tmp_path / 'candidate/protocol.docx', source))


@pytest.mark.parametrize('branch,family', COMBINATIONS[1:])
def test_structured_schedule_keeps_real_timing_and_assignments(tmp_path, branch, family):
    source = source_for(branch, family)
    source['procedures']['visit_schedule'] = [
        {'visit': 'Baseline', 'timing': 'Day 0', 'procedures': ['Blood draw (5 mL)']},
        {'visit': 'Telephone contact', 'timing': 'Week 8 ±7 days', 'procedures': ['Questionnaire']},
    ]
    source['procedures']['assessments'] = ['Blood draw (5 mL)', 'Questionnaire', 'Additional assessment if clinically indicated.']
    rendering.render_documents(ROOT, tmp_path, source, {'protocol': [], 'icf': {}, 'prs': {}}, artifact_names={'protocol'})
    document = Document(tmp_path / 'candidate/protocol.docx')
    table = next(t for t in document.tables if t.cell(0, 0).text.strip() == 'Activity')
    rows = {r.cells[0].text: [c.text for c in r.cells[1:]] for r in table.rows[2:]}
    assert rows['Blood draw (5 mL)'] == ['X', '']
    assert rows['Questionnaire'] == ['', 'X']
    assert 'Additional assessment if clinically indicated.' in [p.text for p in document.paragraphs]
    visits = next(t for t in document.tables if t.cell(0, 0).text == 'Visit Number')
    assert visits.cell(2, 2).text == 'Week 8 ±7 days'


def test_table_assignment_corruption_is_construction_owned(tmp_path):
    source = source_for('Prospective', 'Advarra')
    source['procedures']['visit_schedule'] = [
        {'visit': 'Baseline', 'timing': 'Day 0', 'procedures': ['Blood draw (5 mL)']},
        {'visit': 'Exit', 'timing': 'Week 8', 'procedures': ['Questionnaire']},
    ]
    source['procedures']['assessments'] = ['Blood draw (5 mL)', 'Questionnaire']
    rendering.render_documents(ROOT, tmp_path, source, {'protocol': [], 'icf': {}, 'prs': {}})
    prs_xml.generate(ROOT / 'assets/client-templates/prs/clinicaltrials_prs_full_placeholder_template.xml', tmp_path / 'candidate/study.xml', source, {'brief_summary': 'Summary.', 'detailed_description': 'Description.'})
    assert not any('schedule-of-assessments' in f.get('table_ids', []) for f in quality.deterministic_content_check(tmp_path, source))
    path = tmp_path / 'candidate/protocol.docx'
    document = Document(path)
    table = next(t for t in document.tables if t.cell(0, 0).text.strip() == 'Activity')
    table.cell(2, 1).text = ''
    table.cell(2, 2).text = 'X'
    document.save(path)
    findings = [f for f in quality.deterministic_content_check(tmp_path, source) if 'schedule-of-assessments' in f.get('table_ids', [])]
    assert findings and all(f['owner'] == 'construction' for f in findings)


@pytest.mark.parametrize('branch', ['Prospective', 'Ambispective'])
def test_xml_projection_preserves_identity_without_fabricating_groups(tmp_path, branch):
    source = source_for(branch, 'Sterling')
    output = tmp_path / 'study.xml'
    template = ROOT / 'assets/client-templates/prs/clinicaltrials_prs_full_placeholder_template.xml'
    report = prs_xml.generate(template, output, source, {'brief_summary': 'Summary.', 'detailed_description': 'Description.'})
    assert report['status'] == 'passed'
    study = next(ET.parse(output).getroot().iter('clinical_study'))
    assert study.findtext('overall_official/first_name') == 'Morgan'
    assert study.findtext('overall_official/middle_name') in ('', None)
    assert study.findtext('location/investigator/first_name') == 'Avery'
    assert study.findtext('brief_title') == rendering._protocol_short_title(source)
    assert study.findtext('brief_title') != source['study']['title']
    assert study.findtext('study_design/observational_design/number_of_groups') in ('', None)
    assert not study.findall('arm_group')


@pytest.mark.parametrize('target,value', [
    ('overall_official/first_name', 'Dr.'),
    ('location/investigator/first_name', ''),
    ('brief_title', 'Unrelated study title'),
    ('study_design/observational_design/number_of_groups', '0'),
])
def test_xml_scalar_corruption_is_detected(tmp_path, target, value):
    source = source_for('Prospective', 'Advarra')
    template = ROOT / 'assets/client-templates/prs/clinicaltrials_prs_full_placeholder_template.xml'
    output = tmp_path / 'study.xml'
    prs_xml.generate(template, output, source, {'brief_summary': 'Summary.', 'detailed_description': 'Description.'})
    tree = ET.parse(output)
    next(tree.getroot().iter('clinical_study')).find(target).text = value
    tree.write(output, encoding='utf-8', xml_declaration=True)
    assert prs_xml.validate_output(output, source, template)


@pytest.mark.parametrize('name,expected', [
    ('Dr. Morgan Sample, MD', ('Morgan', '', 'Sample')),
    ('Prof Avery Lee Sample', ('Avery', 'Lee', 'Sample')),
    ('María José Sample', ('María', 'José', 'Sample')),
    ('Sample', ('Sample', '', '')),
])
def test_person_name_projection_removes_only_honorifics(name, expected):
    assert prs_xml._name(name) == expected


def test_explicit_person_components_override_full_display_name(tmp_path):
    source = source_for('Prospective', 'Sterling')
    source['parties']['principal_investigator'].update({
        'first_name': 'María José', 'last_name': 'de la Cruz', 'middle_name': '', 'degrees': 'PhD'})
    fields = prs_xml._fields(source, {})
    assert fields['overallOfficialFirstName'] == 'María José'
    assert fields['overallOfficialLastName'] == 'de la Cruz'
    assert fields['overallOfficialDegrees'] == 'PhD'


def test_unknown_and_explicit_group_counts_are_distinct():
    source = source_for('Prospective', 'Sterling')
    assert prs_xml._fields(source, {})['numberOfGroups'] == ''
    source['design']['arms'] = [{'name': 'Group A'}, {'name': 'Group B'}]
    assert prs_xml._fields(source, {})['numberOfGroups'] == '2'


def test_presentation_cleanup_keeps_meaningful_negative_and_unrecognized_counts():
    source = source_for('Prospective', 'Advarra')
    source['design']['study_design'] = 'Observational, None; no masking; no control group'
    assert rendering.study_design_display(source) == 'Observational; no masking; no control group'
    source['population']['sample_size'] = '72 subjects total; 36 per cohort.'
    assert rendering.sample_size_display(source['population']['sample_size']) == '72 subjects total; 36 per cohort'
    assert rendering.sample_size_display('48 subjects (treatment); control: 24') == '48 participants in the treatment group and 24 in the control group'


# Authored citation text preservation is exercised at the public rendering seam
# in test_ai_prose_ownership.py; assembly no longer edits narrative citations.


@pytest.mark.parametrize('branch', ['Retrospective', 'Prospective', 'Ambispective'])
def test_followup_abbreviations_are_readable_without_losing_qualified_anchors(branch):
    source = {'meta': {'study_type': branch}, 'study': {'timeline': 'Enrollment: 8W; Follow-Up: 3M; Data Analysis: 1M.'}}
    assert contracts.participant_followup_summary(source) == '3 months'
    source['study']['timeline'] = 'Follow-up: 3 months after second-eye surgery.'
    assert contracts.participant_followup_summary(source) == '3 months after second-eye surgery'


def test_partial_timing_and_unassigned_assessments_are_preserved():
    source = source_for('Retrospective', None)
    source['procedures']['visit_schedule'] = [{'visit': 'Abstraction'}, {'visit': 'Follow-up', 'timing': 'Week 8 ±7 days'}]
    table = contracts.protocol_table_contracts(source)['schedule-of-assessments']
    assert table['rows'] == [['Approved visit or assessment', 'Approved timing'], ['Abstraction', ''], ['Follow-up', 'Week 8 ±7 days']]
    assert source['procedures']['assessments'] in table['supplemental_notes']
