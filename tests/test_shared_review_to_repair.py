"""Public Desktop review-to-repair interaction; controlled adapters only."""
import json
import copy
from pathlib import Path

import pytest

import test_repair_with_incomplete_visual as existing
from visual_review import IncompleteVisualReview


@pytest.mark.parametrize('family', existing.FAMILIES)
def test_clean_partial_parent_preserves_content_concerns_for_repair(tmp_path, monkeypatch, family):
    original = existing.operation
    calls = []

    def intercepted(*args, **kwargs):
        visual = kwargs['visual_review']

        def partial_parent(request, response, revision, fallback):
            calls.append((request['review_set'], fallback, request['request_id']))
            if request['review_set'] == 1 and fallback:
                raise IncompleteVisualReview('Clean authenticated parent retained incomplete scope')
            return visual(request, response, revision, fallback)

        kwargs['visual_review'] = partial_parent
        return original(*args, **kwargs)

    monkeypatch.setattr(existing, 'operation', intercepted)
    existing.test_content_defect_reaches_repair_despite_missing_visual_response(
        tmp_path, monkeypatch, family, 'initial-missing')
    assert any(review_set == 1 and fallback for review_set, fallback, _ in calls)
    assert any(review_set == 2 for review_set, _, _ in calls)
    state = json.loads((tmp_path / 'logs/desktop-operation.json').read_text())
    assert state['budget_seconds'] == 180
    assert state['deadline_at_epoch'] - state['started_at_epoch'] == 180


@pytest.mark.parametrize('family', existing.FAMILIES)
@pytest.mark.parametrize('effective', [True, False], ids=['native-correction', 'ineffective'])
def test_bounded_document_concern_to_native_repair_and_exact_delivery(tmp_path, monkeypatch, family, effective):
    """Controlled reviewers/rendering, real approval, scheduler, adoption and delivery.

    Five synthetic pages are a transport fixture, NOT office pagination or a
    clinical visual judgment. Page 3 is intentionally outside document anchors.
    The defect and its correction are actual Word-native paragraph formatting.
    """
    import drafting
    import quality
    import workflow
    from docx import Document
    from docx.shared import Pt
    from test_clean_delivery import ROOT, CASES, controlled_render
    from test_general_document_repair import repair_response
    from test_handoff_quality import acceptance_verification
    from test_visual_units import response_for
    from visual_review import run_visual_review, visual_response_coverage_errors

    case = next(c for c in CASES if c['descriptor']['study_type'] == family[0]
                and (family[1] is None or c['descriptor'].get('icf_template') == family[1]))
    source = quality.branch_acceptance_source(ROOT, case)
    reference = tmp_path / 'reference/study.reference.json'
    reference.parent.mkdir(parents=True)
    reference.write_text(json.dumps(source))
    workflow.prepare(tmp_path)
    approval = workflow.approve(tmp_path, approved_by='controlled-fixture-owner')
    revision = tmp_path / 'revisions' / approval['revision_id']
    approved = (revision / 'approved-reference.json').read_bytes()
    source_md = (revision / 'approved-source.md').read_bytes()
    renders, parents, units, contents, repairs, states = [], [], [], [], [], []
    original_format = []
    paragraph_index = []
    concern_id = 'controlled-native-interior-template'

    def broken(path):
        p = Document(path).paragraphs[paragraph_index[0]]
        return p.paragraph_format.space_before == Pt(144)

    def render(root, current, *args, **kwargs):
        path = current / 'candidate/protocol.docx'
        if not paragraph_index:
            doc = Document(path)
            index = next(i for i, p in enumerate(doc.paragraphs) if i > 10 and len(p.text) > 80)
            paragraph_index.append(index)
            original_format.append(doc.paragraphs[index].paragraph_format.space_before)
            doc.paragraphs[index].paragraph_format.space_before = Pt(144)
            doc.save(path)
        report = controlled_render(root, current, *args, **kwargs)
        for artifact in report['render']['artifacts']:
            # Deliberately synthetic renderer: bind native hash and render ordinal.
            # Distinct rerender bytes deliberately require fresh visual review
            # even for unchanged ICF, rather than testing lawful visual reuse.
            artifact['pages'] = []
            for number in range(1, 6):
                page = current / f"rendered/{artifact['artifact']}/page-{number}.png"
                page.write_bytes(f"CONTROLLED RENDER ONLY {artifact['docx_sha256']} {len(renders) + 1} {number}".encode())
                artifact['pages'].append({'page': number, 'path': page.relative_to(current).as_posix(),
                                          'sha256': quality.sha256_file(page)})
            artifact['page_count'] = 5
            # Distinguish exact controlled PDF bytes across native changes too.
            pdf = current / artifact['pdf']
            pdf.write_bytes(pdf.read_bytes() + f"\n% controlled {artifact['docx_sha256']} render {len(renders) + 1}\n".encode())
            artifact['pdf_sha256'] = quality.sha256_file(pdf)
        renders.append(copy.deepcopy(report['render']['artifacts']))
        return report

    def visual_dispatch(handoffs, budget):
        assert 0 < budget <= 180
        for handoff in handoffs:
            assert handoff['visual_unit'], 'Legacy response is not bounded-unit evidence'
            unit = json.loads((revision / handoff['request_path']).read_text())
            assert len({p['path'] for t in unit['visual_scope'] for p in t['pages']}) <= 3
            assert quality.verification_request_hash_valid(unit)
            for artifact in unit['artifacts']:
                assert quality.sha256_file(revision / artifact['docx']) == artifact['docx_sha256']
                assert quality.sha256_file(revision / artifact['pdf']) == artifact['pdf_sha256']
                for page in artifact['pages']:
                    assert quality.sha256_file(revision / page['path']) == page['sha256']
            response = response_for(unit)
            response['producer']['reviewer_id'] = 'controlled-independent-visual'
            response['status'] = 'passed'
            for row in response['visual_assessments']:
                row['notes'] = 'Controlled fixture: current native paragraph and synthetic page bindings inspected.'
                if unit['review_set'] > 1:
                    prior = unit['review_context']['prior_findings']
                    assert concern_id in {f['finding_id'] for f in prior}
                    row['notes'] += f' Independently reassessed {concern_id} against current native bytes.'
                if row['kind'] == 'document' and row['artifact'] == 'protocol' and broken(revision / 'candidate/protocol.docx'):
                    assert 3 not in {p['page'] for p in row['pages']}
                    evidence = unit['document_evidence']
                    assert quality.sha256_file(revision / evidence['path']) == evidence['sha256']
                    snapshot = json.loads((revision / evidence['path']).read_text())
                    assert {r['kind'] for r in snapshot['assessments']} == {'page', 'boundary'}
                    assert len(snapshot['assessments']) == 9
                    row['status'] = response['status'] = 'blocked'
                    response['findings'] = [{'finding_id': concern_id, 'target_id': row['target_id'],
                        'artifact': 'protocol', 'page': 3, 'check': 'template_preservation',
                        'evidence_basis': 'native', 'element': f'Controlled paragraph {paragraph_index[0]}',
                        'issue': 'Controlled native paragraph has 144pt spacing, violating the template hierarchy.',
                        'recommended_action': 'Restore the original paragraph spacing without changing text.'}]
            units.append(copy.deepcopy(unit))
            output = revision / unit['response_path']
            output.parent.mkdir(parents=True, exist_ok=True)
            output.write_text(json.dumps(response))

    def dispatch(handoffs, budget):
        for handoff in handoffs:
            request_path = revision / handoff['request_path']
            request = json.loads(request_path.read_text())
            states.append(json.loads((tmp_path / 'logs/desktop-operation.json').read_text()))
            if request['task'] == 'rendered_page_visual_verification':
                assert request.get('visual_review_plan'), 'Legacy path lacks actual bounded visual units'
                parents.append(copy.deepcopy(request))
                reviewed = run_visual_review(revision, request_path, visual_dispatch, budget)
                assert reviewed['complete'], reviewed
                response = json.loads((revision / request['response_path']).read_text())
                assert visual_response_coverage_errors(revision, request, response) == []
                continue
            if request['task'] == 'document_repair':
                repairs.append(copy.deepcopy(request))
                finding = next(f for f in request['findings'] if f['finding_id'] == concern_id)
                native = finding['verification_source_finding']
                assert native['evidence_basis'] == 'native' and native['page'] == 3
                assert native['target_id'] == 'protocol:document'
                assert native['check'] == 'template_preservation'
                assert request['approved_source'] == json.loads(approved)
                path = revision / request['workspace_path'] / 'protocol.docx'
                assert broken(path)
                if effective:
                    doc = Document(path)
                    doc.paragraphs[paragraph_index[0]].paragraph_format.space_before = original_format[0]
                    doc.save(path)
                response = repair_response(request, **({'status': 'needs_review',
                    'summary': 'Controlled ineffective edit: no productive safe correction remains.'}
                    if not effective and len(repairs) > 1 else {}))
            elif request['task'] in {'section_drafting', 'prs_narrative_drafting'}:
                response = drafting.recorded_acceptance_response(request)
            else:
                response = acceptance_verification(request)
                if request['task'] == 'clinical_content_verification':
                    contents.append(copy.deepcopy(request))
                    if request['review_set'] > 1:
                        parent = json.loads((revision / request['parent_request_path']).read_text()) if request.get('parent_request_path') else request
                        assert concern_id in {f['finding_id'] for f in parent['review_context']['prior_findings']}
                        assert broken(revision / 'candidate/protocol.docx') is (not effective)
                        for row in response['section_assessments']:
                            row['notes'] = [f'Independently reassessed {concern_id} in current native package.']
            output = revision / request['response_path']
            output.parent.mkdir(parents=True, exist_ok=True)
            output.write_text(json.dumps(response))

    monkeypatch.setattr(workflow, 'render_assurance', render)
    result = workflow.run_desktop_operation(tmp_path, handoff_runner=dispatch,
        fallback_handoff_runner=dispatch, opener=lambda path: Path(path).read_bytes(),
        budget_seconds=180, require_promoted_runtime=False,
        release_identity={'host_mode': 'hermes-tui', 'package_fingerprint': 'controlled-combined-regression',
                          'git_commit': '8e7a9e533745dd38ca067a87ec2902dd9d908507'})
    assert units, 'Must create authenticated bounded visual units, not legacy pass responses'
    assert repairs and all(r['revision_id'] == approval['revision_id'] for r in repairs)
    state = json.loads((tmp_path / 'logs/desktop-operation.json').read_text())
    for snapshot in [*states, state]:
        assert snapshot['budget_seconds'] == 180
        assert snapshot['deadline_at_epoch'] - snapshot['started_at_epoch'] == 180
        assert snapshot['deadline_at_epoch'] == states[0]['deadline_at_epoch']
        assert snapshot['approval_identity'] == states[0]['approval_identity']
    assert (revision / 'approved-reference.json').read_bytes() == approved
    assert (revision / 'approved-source.md').read_bytes() == source_md
    if not effective:
        assert result['status'] != 'passed' and not result.get('client_outputs'), result
        assert broken(revision / 'candidate/protocol.docx')
        assert concern_id in json.dumps(result), 'Ineffective repair erased original concern'
        assert any(p['review_set'] > 1 for p in parents), 'Ineffective attempt needs independent reassessment'
        return
    assert result['status'] == 'passed' and result['delivery']['confirmed'], result
    assert len(repairs) == 1 and len(renders) == 2
    assert {p['review_set'] for p in parents} == {1, 2}
    assert {c['review_set'] for c in contents} == {1, 2}
    final_content_path = next((revision / 'hermes/verification-requests').glob('*review-2*content.json'))
    final_content = json.loads(final_content_path.read_text())
    final_response = json.loads((revision / final_content['response_path']).read_text())
    assert final_response['status'] == 'passed'
    assert quality.verification_response_is_complete(revision, final_content_path)
    initial_content = next(c for c in contents if c['review_set'] == 1)
    assert initial_content['request_id'] != final_content['request_id']
    assert initial_content['source_identity_sha256'] == final_content['source_identity_sha256']
    assert initial_content['generation_authority_sha256'] == final_content['generation_authority_sha256']
    assert {r['section_id'] for r in final_response['section_assessments']} == {
        r['section_id'] for r in final_content['sections']}
    assert {r['check'] for r in final_response['cross_document_assessments']} == set(final_content['cross_document_checks'])
    for artifact in final_content['artifacts']:
        assert quality.sha256_file(revision / artifact['path']) == artifact['sha256']
    for artifact in {'protocol'} | ({'icf'} if family[1] else set()):
        initial = next(p for p in parents if p['review_set'] == 1 and p['artifacts'][0]['artifact'] == artifact)
        final = next(p for p in parents if p['review_set'] == 2 and p['artifacts'][0]['artifact'] == artifact)
        assert initial['request_id'] != final['request_id']
        assert initial['source_identity_sha256'] == final['source_identity_sha256']
        assert initial['generation_authority_sha256'] == final['generation_authority_sha256']
        assert {t['target_id'] for u in units if u['parent_request_id'] == final['request_id']
                for t in u['visual_scope']} == {t['target_id'] for t in final['visual_review_plan']['targets']}
        if artifact == 'protocol':
            for key in ('docx_sha256', 'pdf_sha256'):
                assert initial['artifacts'][0][key] != final['artifacts'][0][key]
            assert initial['artifacts'][0]['pages'] != final['artifacts'][0]['pages']
    expected = workflow.document_set(source['meta']['study_type'])
    assert sorted(p.name for p in (tmp_path / 'output').iterdir()) == sorted(expected)
    for name in expected:
        assert (tmp_path / 'output' / name).read_bytes() == (revision / 'candidate' / name).read_bytes()
        opened = next(row for row in result['delivery']['opened'] if row['filename'] == name)
        assert opened['sha256'] == quality.sha256_file(tmp_path / 'output' / name)
    assert not broken(tmp_path / 'output/protocol.docx')
    manifest = json.loads((tmp_path / result['manifest']).read_text())
    assert quality.final_exact_artifact_review_findings(revision, json.loads(approved),
        json.loads((revision / 'candidate-build.json').read_text())['render_report'], manifest['quality']) == []
