import json
import hashlib
import shutil
from datetime import date
from pathlib import Path

import pytest
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from pypdf import PdfWriter

from contracts import batch_plan, contracted_template_bundle
from drafting import accepted_cross_section_duplicate_findings, accepted_draft, create_drafting_request, governing_resources, ingest_responses, pending_requests, recorded_acceptance_response, response_template, retry_attempts, schedule_requests, sha256_value, validate_response
from quality import CONTENT_CHECKS, RESPONSE_SCHEMA, VISUAL_CHECKS, create_verification_requests, deterministic_content_check, validate_verifications, verification_request_sha256
from contracts import icf_retained_sections
from rendering import audit_docx
from workflow import approve, generate, prepare, validate
import workflow
import quality


ROOT = Path(__file__).resolve().parents[1]


def test_single_artifact_retrospective_verification_omits_cross_document_checks(tmp_path):
    revision_dir = tmp_path / "r-retrospective"
    revision_dir.mkdir()

    request_paths = create_verification_requests(
        revision_dir,
        {"meta": {"study_type": "Retrospective"}},
        {"artifacts": []},
    )
    content_request = json.loads(request_paths[0].read_text(encoding="utf-8"))

    assert content_request["task"] == "clinical_content_verification"
    assert content_request["cross_document_checks"] == []
    assert "assess every cross-document check" not in content_request["instructions"].casefold()


def test_prospective_content_verification_requires_all_prs_semantic_targets(tmp_path):
    tmp_path.mkdir(exist_ok=True)
    request_paths = create_verification_requests(
        tmp_path,
        fixture(),
        {"artifacts": []},
    )
    content_request = json.loads(request_paths[0].read_text(encoding="utf-8"))

    prs_sections = {
        item["section_id"]
        for item in content_request["sections"]
        if item["artifact"] == "study.xml"
    }
    assert prs_sections == {
        "prs.brief-summary", "prs.detailed-description", "prs.structured",
    }
    assert content_request["reviewer_policy"]["producer_reviewer_id_required"] is True


def test_content_verification_response_path_is_unique_to_each_review_set(tmp_path):
    first = create_verification_requests(
        tmp_path,
        {"meta": {"study_type": "Retrospective"}},
        {"artifacts": []},
        review_set=1,
    )[0]
    first_request = json.loads(first.read_text(encoding="utf-8"))
    second = create_verification_requests(
        tmp_path,
        {"meta": {"study_type": "Retrospective"}},
        {"artifacts": []},
        review_set=2,
    )[0]
    second_request = json.loads(second.read_text(encoding="utf-8"))

    assert first_request["response_path"] != second_request["response_path"]
    assert "review-1.verify.content" in first_request["response_path"]
    assert "review-2.verify.content" in second_request["response_path"]
    assert "Set top-level status exactly `passed`" in second_request["instructions"]


def test_failed_content_assessment_is_complete_negative_evidence_not_a_transient_gap(tmp_path):
    request_path = create_verification_requests(
        tmp_path,
        {"meta": {"study_type": "Retrospective"}},
        {"artifacts": []},
        review_set=2,
    )[0]
    request = json.loads(request_path.read_text(encoding="utf-8"))
    response = acceptance_verification(request)
    failed_section = response["section_assessments"][0]["section_id"]
    response["status"] = "failed"
    response["section_assessments"][0]["status"] = "failed"
    response["findings"] = [{
        "target_ids": [failed_section],
        "issue": "The section contains a source-fidelity defect.",
    }]
    response_path = tmp_path / request["response_path"]
    response_path.write_text(json.dumps(response), encoding="utf-8")

    findings, _evidence = validate_verifications(
        tmp_path,
        request_paths=[request_path],
    )

    assert any(item["recovery_class"] == "drafting_defect" for item in findings)
    assert not any(
        item["recovery_class"] == "verifier_transient"
        and "explicitly assessed" in item["issue"]
        for item in findings
    )


def test_verification_response_requires_distinct_reviewer_and_model_identities(tmp_path):
    request_path = create_verification_requests(
        tmp_path,
        {"meta": {"study_type": "Retrospective"}},
        {"artifacts": []},
    )[0]
    request = json.loads(request_path.read_text(encoding="utf-8"))
    response = acceptance_verification(request)
    response["producer"].pop("reviewer_id")
    response_path = tmp_path / request["response_path"]
    response_path.write_text(json.dumps(response), encoding="utf-8")

    findings, _evidence = validate_verifications(tmp_path, request_paths=[request_path])

    assert any(item["field"] == "producer.reviewer_id" for item in findings)
    assert all(item["recovery_class"] == "verifier_transient" for item in findings)


def test_prs_structured_content_finding_routes_to_deterministic_rebuild(tmp_path):
    request_path = create_verification_requests(
        tmp_path,
        fixture(),
        {"artifacts": []},
    )[0]
    request = json.loads(request_path.read_text(encoding="utf-8"))
    response = acceptance_verification(request)
    structured = next(
        item for item in response["section_assessments"]
        if item["section_id"] == "prs.structured"
    )
    structured["status"] = "failed"
    response["status"] = "blocked"
    response["findings"] = [{
        "target_ids": ["prs.structured"],
        "issue": "The structured eligibility qualifier differs from the approved source.",
    }]
    response_path = tmp_path / request["response_path"]
    response_path.write_text(json.dumps(response), encoding="utf-8")

    findings, _evidence = validate_verifications(tmp_path, request_paths=[request_path])

    routed = next(item for item in findings if item.get("target_ids") == ["prs.structured"])
    assert routed["recovery_class"] == "deterministic_structure_defect"
    assert routed["action"] == "rebuild_deterministic_structure"


def _require_renderer():
    assert workflow.renderer() is not None


def acceptance_verification(request):
    response = {
        "schema_version": RESPONSE_SCHEMA,
        "request_id": request["request_id"],
        "request_sha256": request["request_sha256"],
        "task": request["task"],
        "producer": {
            "model_id": "TestAcceptanceVerifier/v1",
            "reviewer_id": request["task"],
        },
        "status": "passed",
        "findings": [],
    }
    if request["task"] == "rendered_page_visual_verification":
        response["page_assessments"] = [
            {
                "artifact": artifact["artifact"],
                "page": page["page"],
                "sha256": page["sha256"],
                "status": "passed",
                "checks": list(VISUAL_CHECKS),
            }
            for artifact in request.get("artifacts", [])
            for page in artifact.get("pages", [])
        ]
    else:
        response["section_assessments"] = [
            {
                "artifact": item["artifact"],
                "section_id": item["section_id"],
                "status": "passed",
                "checks": list(CONTENT_CHECKS),
            }
            for item in request.get("sections", [])
        ]
        response["cross_document_assessments"] = [
            {"check": check, "status": "passed"}
            for check in request.get("cross_document_checks", [])
        ]
    return response


def fixture():
    return json.loads((ROOT / "tests/fixtures/prospective-acceptance-source.json").read_text(encoding="utf-8"))


def test_drafting_request_is_scoped_and_hash_bound(tmp_path):
    reference = fixture(); batch = batch_plan("Prospective")[0]
    path = create_drafting_request(repo_root=ROOT, revision_dir=tmp_path, revision_id="r-test", reference=reference, batch=batch, attempts={item: 1 for item in batch.section_ids}, wave="initial")
    request = json.loads(path.read_text(encoding="utf-8"))
    assert (tmp_path / request["response_path"]).parent.is_dir()
    assert set(request["approved_source"]) <= set(batch.field_families)
    assert "template_fields" not in request["approved_source"]
    assert {item["section_id"] for item in request["evidence_checklist"]} == set(batch.section_ids)
    for item in request["evidence_checklist"]:
        assert set(item["approved_values"]) == set(item["required_source_paths"])
    response = recorded_acceptance_response(request)
    accepted, findings = validate_response(request, response)
    assert accepted and not findings


def test_retry_guidance_identifies_the_uncited_paragraph_and_allowed_references(tmp_path):
    """Reproduce the final protocol-operations response from Hermes run __03."""
    reference = fixture()
    reference["meta"]["study_type"] = "Ambispective"
    reference["procedures"]["assessments"] = [
        "Screening, consent, and baseline visit",
        "Sensor wear on Days 1 to 14, Weeks 6 to 8, and Weeks 10 to 12",
        "Telephone contact at Week 3",
        "Clinic visits at Weeks 6 and 12",
        "Record abstraction",
        "Sensor insertion and removal",
        "Sensor data download",
        "Medication review",
        "Adverse-event assessment",
        "Hemoglobin A1c at Week 12",
        "Usability questionnaire",
    ]
    reference["procedures"].pop("visit_schedule", None)
    batch = next(
        item for item in batch_plan("Ambispective")
        if item.batch_id == "protocol-operations"
    )
    request_path = create_drafting_request(
        repo_root=ROOT,
        revision_dir=tmp_path,
        revision_id="r-hermes-03-evaluation",
        reference=reference,
        batch=batch,
        target_ids=("evaluation-procedures",),
        attempts={"evaluation-procedures": 3},
        wave="retry",
    )
    request = json.loads(request_path.read_text(encoding="utf-8"))
    response = response_template(request, model_id="Hermes reproduction")
    result = response["section_results"][0]
    result["paragraphs"] = [{
        "text": "The Schedule of Assessments includes the following approved visits, monitoring periods, contacts, and procedures.",
        "evidence_refs": [],
        "boilerplate_refs": [],
    }]
    result["lists"] = [{
        "items": reference["procedures"]["assessments"],
        "evidence_refs": ["source:procedures.assessments"],
        "boilerplate_refs": [],
    }]

    accepted, findings = validate_response(request, response)
    finding = next(
        item for item in findings
        if item["field"] == "evaluation-procedures"
        and "no approved evidence" in item["issue"]
    )

    assert "evaluation-procedures" not in {
        item["section_id"] for item in (accepted or {}).get("drafts", [])
    }
    assert "Paragraph 1" in finding["next_action"]
    assert "source:procedures.assessments" in finding["next_action"]
    assert any(
        "Every paragraph object and every list object" in constraint
        for constraint in request["constraints"]
    )

    result["paragraphs"][0]["evidence_refs"] = ["source:procedures.assessments"]
    accepted, findings = validate_response(request, response)

    assert not findings
    assert [item["section_id"] for item in accepted["drafts"]] == [
        "evaluation-procedures"
    ]


def test_safety_analysis_uses_only_its_material_analysis_plan_clause(tmp_path):
    """Reproduce the final analysis response from Hermes run __03."""
    reference = fixture()
    reference["meta"]["study_type"] = "Ambispective"
    reference["statistics"]["analysis_plan"] = (
        "Summarize historical and prospective measures descriptively. "
        "For participants with both measurements, report the mean within-participant "
        "change in hemoglobin A1c with a two-sided 95% confidence interval. "
        "Summarize time in range, usable sensor time, usability, missing data, and "
        "adverse events using descriptive statistics."
    )
    reference["safety"].pop("adverse_events", None)
    batch = next(
        item for item in batch_plan("Ambispective")
        if item.batch_id == "protocol-analysis-and-oversight"
    )
    request_path = create_drafting_request(
        repo_root=ROOT,
        revision_dir=tmp_path,
        revision_id="r-hermes-03-safety-analysis",
        reference=reference,
        batch=batch,
        target_ids=("quality-safety.analysis",),
        attempts={"quality-safety.analysis": 3},
        wave="retry",
    )
    request = json.loads(request_path.read_text(encoding="utf-8"))
    contract = request["section_contracts"][0]
    response = response_template(request, model_id="Hermes reproduction")
    response["section_results"][0]["paragraphs"] = [
        {
            "text": contract["fixed_boilerplate"][0]["text"],
            "evidence_refs": [],
            "boilerplate_refs": ["safety-analysis"],
        },
        {
            "text": "In accordance with the approved analysis plan, adverse events will be summarized using descriptive statistics.",
            "evidence_refs": ["source:statistics.analysis_plan"],
            "boilerplate_refs": [],
        },
    ]

    accepted, findings = validate_response(request, response)

    scope = contract["evidence_scopes"][0]
    assert scope["path"] == "statistics.analysis_plan"
    assert scope["value"] == (
        "Summarize adverse events using descriptive statistics."
    )
    assert "hemoglobin A1c" not in scope["value"]
    assert "sensor" not in scope["value"]
    assert scope["sha256"] == sha256_value(scope["value"])
    assert not findings
    assert [item["section_id"] for item in accepted["drafts"]] == [
        "quality-safety.analysis"
    ]


def test_safety_analysis_does_not_offer_an_efficacy_only_plan_as_evidence(tmp_path):
    reference = fixture()
    reference["meta"]["study_type"] = "Ambispective"
    reference["statistics"]["analysis_plan"] = (
        "Report the mean within-participant hemoglobin A1c change with a two-sided "
        "95% confidence interval."
    )
    reference["safety"].pop("adverse_events", None)
    batch = next(
        item for item in batch_plan("Ambispective")
        if item.batch_id == "protocol-analysis-and-oversight"
    )
    request_path = create_drafting_request(
        repo_root=ROOT,
        revision_dir=tmp_path,
        revision_id="r-efficacy-only-safety-analysis",
        reference=reference,
        batch=batch,
        target_ids=("quality-safety.analysis",),
        attempts={"quality-safety.analysis": 1},
        wave="initial",
    )
    request = json.loads(request_path.read_text(encoding="utf-8"))
    contract = request["section_contracts"][0]

    assert contract["evidence_scopes"][0]["value"] == ""
    assert "statistics.analysis_plan" not in contract["minimum_evidence"]


def test_long_approved_sections_receive_a_soft_depth_signal_but_fail_only_on_missing_facts(tmp_path):
    reference = fixture()
    reference["study"]["background"] = " ".join(
        f"approved-background-detail-{index}" for index in range(120)
    )
    batch = next(item for item in batch_plan("Prospective") if item.batch_id == "protocol-foundations")
    path = create_drafting_request(
        repo_root=ROOT,
        revision_dir=tmp_path,
        revision_id="r-depth",
        reference=reference,
        batch=batch,
        attempts={item: 1 for item in batch.section_ids},
        wave="initial",
    )
    request = json.loads(path.read_text(encoding="utf-8"))
    introduction_contract = next(
        item for item in request["section_contracts"] if item["section_id"] == "introduction"
    )
    assert introduction_contract["approved_source_word_count"] >= 120
    assert introduction_contract["reference_detail_target_words"] >= 80

    response = recorded_acceptance_response(request)
    introduction = next(
        item for item in response["section_results"] if item["section_id"] == "introduction"
    )
    introduction["paragraphs"] = [{
        "text": "Approved background detail supports the stated study title hypothesis and primary endpoint.",
        "evidence_refs": [f"source:{path}" for path in introduction_contract["minimum_evidence"]],
        "boilerplate_refs": [],
    }]
    introduction["lists"] = []

    _accepted, findings = validate_response(request, response)

    assert not any("source-proportional detail" in item.get("issue", "") for item in findings)
    assert any(
        item.get("field") == "introduction"
        and "material facts are not observable" in item.get("issue", "")
        for item in findings
    )


def test_governed_drafting_response_rejects_duplicate_prose_across_contracts(tmp_path):
    reference = fixture()
    batch = next(item for item in batch_plan("Prospective") if item.batch_id == "protocol-operations")
    path = create_drafting_request(
        repo_root=ROOT,
        revision_dir=tmp_path,
        revision_id="r-duplicate-reproduction",
        reference=reference,
        batch=batch,
        attempts={item: 1 for item in batch.section_ids},
        wave="initial",
    )
    request = json.loads(path.read_text(encoding="utf-8"))
    assert any(
        "exact sentence or paragraph" in item and "exact list item" in item
        for item in request["constraints"]
    )
    response = recorded_acceptance_response(request)
    visits = next(item for item in response["section_results"] if item["section_id"] == "study-procedure.visits")
    evaluation = next(item for item in response["section_results"] if item["section_id"] == "evaluation-procedures")
    evaluation["paragraphs"] = json.loads(json.dumps(visits["paragraphs"]))
    evaluation["lists"] = json.loads(json.dumps(visits["lists"]))

    accepted, findings = validate_response(request, response)

    assert accepted is not None
    assert {
        item["section_id"] for item in accepted["drafts"]
    }.isdisjoint({"evaluation-procedures", "study-procedure.visits"})
    assert {
        item["section_id"] for item in accepted["drafts"]
    } == set(batch.section_ids) - {"evaluation-procedures", "study-procedure.visits"}
    duplicate_findings = [item for item in findings if "duplicated across separately contracted" in item["issue"]]
    assert duplicate_findings
    assert duplicate_findings[0]["target_ids"] == ["evaluation-procedures", "study-procedure.visits"]


def test_duplicate_ingestion_preserves_both_retry_targets_and_accepts_the_rest(tmp_path):
    reference = fixture()
    batch = next(item for item in batch_plan("Prospective") if item.batch_id == "protocol-operations")
    request_path = create_drafting_request(
        repo_root=ROOT,
        revision_dir=tmp_path,
        revision_id="r-duplicate-local-retry",
        reference=reference,
        batch=batch,
        attempts={item: 1 for item in batch.section_ids},
        wave="initial",
    )
    request = json.loads(request_path.read_text(encoding="utf-8"))
    response = recorded_acceptance_response(request)
    visits = next(item for item in response["section_results"] if item["section_id"] == "study-procedure.visits")
    evaluation = next(item for item in response["section_results"] if item["section_id"] == "evaluation-procedures")
    evaluation["paragraphs"] = json.loads(json.dumps(visits["paragraphs"]))
    evaluation["lists"] = json.loads(json.dumps(visits["lists"]))
    response_path = tmp_path / request["response_path"]
    response_path.parent.mkdir(parents=True, exist_ok=True)
    response_path.write_text(json.dumps(response), encoding="utf-8")

    governing = governing_resources(ROOT, reference)
    findings = ingest_responses(tmp_path, governing)

    assert findings[0]["target_ids"] == ["evaluation-procedures", "study-procedure.visits"]
    accepted_ids = {
        section_id for section_id in batch.section_ids
        if accepted_draft(tmp_path, section_id, governing) is not None
    }
    assert accepted_ids == set(batch.section_ids) - {"evaluation-procedures", "study-procedure.visits"}
    attempts, exhausted = retry_attempts(findings, {})
    assert exhausted == []
    created = schedule_requests(
        repo_root=ROOT,
        revision_dir=tmp_path,
        revision_id="r-duplicate-local-retry",
        reference=reference,
        attempts=attempts,
        wave="retry",
        findings=findings,
    )
    retry = next(
        json.loads(path.read_text(encoding="utf-8"))
        for path in created
        if json.loads(path.read_text(encoding="utf-8"))["batch_id"] == batch.batch_id
    )
    assert [item["section_id"] for item in retry["section_contracts"]] == [
        "study-procedure.visits",
        "evaluation-procedures",
    ]


def test_authenticated_cross_batch_duplicate_prose_is_found_before_rendering(tmp_path):
    reference = fixture()
    for batch_id, section_id in (
        ("protocol-foundations", "objectives"),
        ("protocol-operations", "study-procedure.measurements"),
    ):
        batch = next(item for item in batch_plan("Prospective") if item.batch_id == batch_id)
        path = create_drafting_request(
            repo_root=ROOT,
            revision_dir=tmp_path,
            revision_id="r-cross-batch-duplicate",
            reference=reference,
            batch=batch,
            attempts={item: 1 for item in batch.section_ids},
            wave="initial",
        )
        request = json.loads(path.read_text(encoding="utf-8"))
        response = recorded_acceptance_response(request)
        target = next(item for item in response["section_results"] if item["section_id"] == section_id)
        target["paragraphs"].append({
            "text": "The approved hypothesis states that prospective monitoring will describe recovery.",
            "evidence_refs": ["source:study.hypothesis"],
            "boilerplate_refs": [],
        })
        response_path = tmp_path / request["response_path"]
        response_path.parent.mkdir(parents=True, exist_ok=True)
        response_path.write_text(json.dumps(response), encoding="utf-8")

    governing = governing_resources(ROOT, reference)
    assert ingest_responses(tmp_path, governing) == []

    findings = accepted_cross_section_duplicate_findings(tmp_path, reference, governing)

    assert len(findings) == 1
    assert findings[0]["target_ids"] == ["objectives", "study-procedure.measurements"]
    assert findings[0]["recovery_class"] == "drafting_defect"
    assert findings[0]["action"] == "retry_drafting_target"


def test_cross_batch_duplicate_routes_localized_retry_before_candidate_render(tmp_path, monkeypatch):
    run_dir = tmp_path / "run"
    reference_path = run_dir / "reference/study.reference.json"
    reference_path.parent.mkdir(parents=True)
    reference_path.write_text(json.dumps(fixture()), encoding="utf-8")
    assert prepare(run_dir)["status"] == "awaiting_approval"
    approval = approve(run_dir, approved_by="reviewer")
    revision_dir = run_dir / "revisions" / approval["revision_id"]
    shared_text = "The approved hypothesis states that prospective monitoring will describe recovery."

    drafting = generate(run_dir)
    for relative in drafting["requests"]:
        request = json.loads((revision_dir / relative).read_text(encoding="utf-8"))
        response = recorded_acceptance_response(request)
        for result in response.get("section_results", []):
            if result["section_id"] in {"objectives", "study-procedure.measurements"}:
                result["paragraphs"].append({
                    "text": shared_text,
                    "evidence_refs": ["source:study.hypothesis"],
                    "boilerplate_refs": [],
                })
        response_path = revision_dir / request["response_path"]
        response_path.parent.mkdir(parents=True, exist_ok=True)
        response_path.write_text(json.dumps(response), encoding="utf-8")

    dependent = generate(run_dir)
    for relative in dependent["requests"]:
        request = json.loads((revision_dir / relative).read_text(encoding="utf-8"))
        response_path = revision_dir / request["response_path"]
        response_path.parent.mkdir(parents=True, exist_ok=True)
        response_path.write_text(json.dumps(recorded_acceptance_response(request)), encoding="utf-8")

    monkeypatch.setattr(
        workflow,
        "render_documents",
        lambda *_args, **_kwargs: (_ for _ in ()).throw(AssertionError("rendered too early")),
    )
    result = generate(run_dir)

    assert result["status"] == "awaiting_hermes"
    assert result["stage"] == "drafting_retry"
    assert {item["field"] for item in result["findings"]} == {"study-procedure.measurements"}
    assert not (revision_dir / "candidate").exists()


def test_fixed_boilerplate_outcome_cannot_bypass_duplicate_prose_gate(tmp_path):
    reference = fixture()
    batch = next(item for item in batch_plan("Prospective") if item.batch_id == "protocol-operations")
    path = create_drafting_request(
        repo_root=ROOT,
        revision_dir=tmp_path,
        revision_id="r-duplicate-boilerplate-bypass",
        reference=reference,
        batch=batch,
        attempts={item: 1 for item in batch.section_ids},
        wave="initial",
    )
    request = json.loads(path.read_text(encoding="utf-8"))
    response = recorded_acceptance_response(request)
    visits = next(item for item in response["section_results"] if item["section_id"] == "study-procedure.visits")
    evaluation = next(item for item in response["section_results"] if item["section_id"] == "evaluation-procedures")
    evaluation["paragraphs"] = json.loads(json.dumps(visits["paragraphs"]))
    evaluation["lists"] = json.loads(json.dumps(visits["lists"]))
    visits["outcome"] = "fixed_boilerplate"
    evaluation["outcome"] = "fixed_boilerplate"

    accepted, findings = validate_response(request, response)

    assert accepted is None
    assert {
        item["field"]
        for item in findings
        if "does not authorize the fixed_boilerplate outcome" in item["issue"]
    } == {"study-procedure.visits", "evaluation-procedures"}
    assert any("duplicated across separately contracted" in item["issue"] for item in findings)


def test_fixed_boilerplate_outcome_rejects_case_mutated_contract_text(tmp_path):
    reference = fixture()
    batch = next(item for item in batch_plan("Prospective") if item.batch_id == "protocol-operations")
    path = create_drafting_request(
        repo_root=ROOT,
        revision_dir=tmp_path,
        revision_id="r-mutated-boilerplate",
        reference=reference,
        batch=batch,
        attempts={item: 1 for item in batch.section_ids},
        wave="initial",
    )
    request = json.loads(path.read_text(encoding="utf-8"))
    response = recorded_acceptance_response(request)
    consent = next(item for item in response["section_results"] if item["section_id"] == "study-procedure.consent")
    assert consent["outcome"] == "fixed_boilerplate"
    consent["paragraphs"][0]["text"] = consent["paragraphs"][0]["text"].upper()

    accepted, findings = validate_response(request, response)

    assert accepted is None
    assert any(
        item["field"] == "study-procedure.consent"
        and "not its exact authorized boilerplate" in item["issue"]
        for item in findings
    )


def test_authorized_boilerplate_does_not_hide_an_unauthorized_cross_section_copy(tmp_path):
    reference = fixture()
    batch = next(item for item in batch_plan("Prospective") if item.batch_id == "protocol-foundations")
    path = create_drafting_request(
        repo_root=ROOT,
        revision_dir=tmp_path,
        revision_id="r-asymmetric-boilerplate-copy",
        reference=reference,
        batch=batch,
        attempts={item: 1 for item in batch.section_ids},
        wave="initial",
    )
    request = json.loads(path.read_text(encoding="utf-8"))
    response = recorded_acceptance_response(request)
    bias_contract = next(item for item in request["section_contracts"] if item["section_id"] == "study-design.bias")
    introduction = next(item for item in response["section_results"] if item["section_id"] == "introduction")
    introduction["paragraphs"].append({
        "text": bias_contract["fixed_boilerplate"][0]["text"],
        "evidence_refs": ["source:study.background"],
        "boilerplate_refs": [],
    })

    accepted, findings = validate_response(request, response)

    assert accepted is not None
    assert {
        item["section_id"] for item in accepted["drafts"]
    }.isdisjoint({"introduction", "study-design.bias"})
    duplicate = next(item for item in findings if "duplicated across separately contracted" in item["issue"])
    assert duplicate["target_ids"] == ["introduction", "study-design.bias"]


def test_mutated_drafting_request_with_stale_hash_is_blocked_before_response_ingestion(tmp_path):
    reference = fixture()
    batch = batch_plan("Prospective")[0]
    path = create_drafting_request(
        repo_root=ROOT,
        revision_dir=tmp_path,
        revision_id="r-integrity",
        reference=reference,
        batch=batch,
        attempts={item: 1 for item in batch.section_ids},
        wave="initial",
    )
    request = json.loads(path.read_text(encoding="utf-8"))
    request["approved_source"]["study"]["title"] = "Injected unapproved title"
    path.write_text(json.dumps(request), encoding="utf-8")
    response_path = tmp_path / request["response_path"]
    response_path.parent.mkdir(parents=True, exist_ok=True)
    response_path.write_text(json.dumps(recorded_acceptance_response(request)), encoding="utf-8")

    findings = ingest_responses(tmp_path, governing_resources(ROOT, reference))

    assert any(item["category"] == "request-integrity" for item in findings)
    assert not (tmp_path / "hermes/accepted-requests" / path.name).exists()


def test_mutated_drafting_request_cannot_be_resigned_with_unapproved_content(tmp_path):
    reference = fixture()
    batch = batch_plan("Prospective")[0]
    path = create_drafting_request(
        repo_root=ROOT,
        revision_dir=tmp_path,
        revision_id="r-resigned-integrity",
        reference=reference,
        batch=batch,
        attempts={item: 1 for item in batch.section_ids},
        wave="initial",
    )
    request = json.loads(path.read_text(encoding="utf-8"))
    request["approved_source"]["study"]["title"] = "Injected unapproved title"
    request.pop("request_sha256")
    request["request_sha256"] = sha256_value(request)
    path.write_text(json.dumps(request), encoding="utf-8")
    ledger = tmp_path / "request-ledger" / f"{request['request_id']}.json"
    ledger.write_text(json.dumps({
        "request_id": request["request_id"],
        "request_sha256": request["request_sha256"],
    }), encoding="utf-8")
    response_path = tmp_path / request["response_path"]
    response_path.parent.mkdir(parents=True, exist_ok=True)
    response_path.write_text(json.dumps(recorded_acceptance_response(request)), encoding="utf-8")

    findings = ingest_responses(tmp_path, governing_resources(ROOT, reference))

    assert any(item["category"] == "request-integrity" for item in findings)
    assert not (tmp_path / "hermes/accepted-requests" / path.name).exists()


def test_recorded_acceptance_spells_out_the_screening_interval_with_units(tmp_path):
    reference = fixture()
    batch = next(item for item in batch_plan("Prospective") if item.batch_id == "protocol-foundations")
    path = create_drafting_request(
        repo_root=ROOT,
        revision_dir=tmp_path,
        revision_id="r-screening",
        reference=reference,
        batch=batch,
        attempts={item: 1 for item in batch.section_ids},
        wave="initial",
    )
    response = recorded_acceptance_response(json.loads(path.read_text(encoding="utf-8")))
    inclusion = next(item for item in response["section_results"] if item["section_id"] == "subjects.inclusion")
    visible = "\n".join(
        [paragraph["text"] for paragraph in inclusion["paragraphs"]]
        + [item for group in inclusion["lists"] for item in group["items"]]
    )

    assert "30 days" in visible


def test_sparse_complete_approval_cannot_create_post_approval_source_questions(tmp_path, monkeypatch):
    _require_renderer()
    reference = fixture()
    reference["risks_benefits"] = {
        "compensation_or_reimbursement": reference["risks_benefits"]["compensation_or_reimbursement"]
    }
    run_dir = tmp_path / "run"
    reference_path = run_dir / "reference/study.reference.json"
    reference_path.parent.mkdir(parents=True)
    reference_path.write_text(json.dumps(reference), encoding="utf-8")

    assert prepare(run_dir)["status"] == "awaiting_approval"
    approval = approve(run_dir, approved_by="reviewer")
    assert approval["status"] == "passed"

    result = generate(run_dir)
    assert result["status"] == "awaiting_hermes"
    revision_dir = run_dir / "revisions" / approval["revision_id"]
    requests = [json.loads((revision_dir / relative).read_text(encoding="utf-8")) for relative in result["requests"]]
    serialized = json.dumps(requests).casefold()

    assert "source_gap" not in serialized
    assert "source question" not in serialized
    risk_benefit_contracts = {
        section["section_id"]: section
        for request in requests
        for section in request.get("section_contracts", [])
        if section["section_id"] in {"risks-benefits.risks", "risks-benefits.benefits", "icf.risks", "icf.benefits"}
    }
    assert set(risk_benefit_contracts) == {"risks-benefits.risks", "risks-benefits.benefits", "icf.risks", "icf.benefits"}
    assert all(contract["fixed_boilerplate"] for contract in risk_benefit_contracts.values())


def test_desktop_drafting_requests_bind_the_complete_contracted_template_bundle(tmp_path, monkeypatch):
    _require_renderer()
    reference = fixture()
    run_dir = tmp_path / "run"
    reference_path = run_dir / "reference/study.reference.json"
    reference_path.parent.mkdir(parents=True)
    reference_path.write_text(json.dumps(reference), encoding="utf-8")

    assert prepare(run_dir)["status"] == "awaiting_approval"
    assert approve(run_dir, approved_by="reviewer")["status"] == "passed"
    result = generate(run_dir)

    assert result["status"] == "awaiting_hermes"
    revision_dir = run_dir / "revisions" / result["revision_id"]
    requests = [
        json.loads((revision_dir / relative).read_text(encoding="utf-8"))
        for relative in result["requests"]
    ]
    expected_bundle = contracted_template_bundle(ROOT, reference)
    for request in requests:
        governing = request["governing_resources"]
        assert governing["contracted_template_bundle"] == expected_bundle
        assert "template_sha256" not in governing
        assert "boilerplate_sha256" not in governing


def test_incomplete_contracted_template_bundle_blocks_before_drafting(tmp_path, monkeypatch):
    run_dir = tmp_path / "run"
    reference_path = run_dir / "reference/study.reference.json"
    reference_path.parent.mkdir(parents=True)
    reference_path.write_text(json.dumps(fixture()), encoding="utf-8")
    assert prepare(run_dir)["status"] == "awaiting_approval"
    assert approve(run_dir, approved_by="reviewer")["status"] == "passed"

    broken_release = tmp_path / "broken-release"
    shutil.copytree(ROOT / "assets", broken_release / "assets")
    shutil.copytree(ROOT / "references", broken_release / "references")
    scripts = broken_release / "scripts"
    scripts.mkdir()
    (broken_release / "assets/client-templates/docx/prospective-protocol.template.docx").unlink()
    monkeypatch.setattr(workflow, "SCRIPT_DIR", scripts)

    result = generate(run_dir)

    assert result["status"] == "blocked"
    assert result["stage"] == "contracted_template_bundle"
    assert len(result["findings"]) == 1
    assert result["findings"][0]["field"] == "contracted_template_bundle"
    assert not list((run_dir / "revisions").glob("*/hermes/requests/*.json"))


def test_awaiting_hermes_exposes_path_only_response_bound_handoffs(tmp_path, monkeypatch):
    _require_renderer()
    run_dir = tmp_path / "run"
    reference_path = run_dir / "reference/study.reference.json"
    reference_path.parent.mkdir(parents=True)
    reference_path.write_text(json.dumps(fixture()), encoding="utf-8")
    assert prepare(run_dir)["status"] == "awaiting_approval"
    approval = approve(run_dir, approved_by="reviewer")

    result = generate(run_dir)

    revision_dir = run_dir / "revisions" / approval["revision_id"]
    expected = []
    for relative in result["requests"]:
        request = json.loads((revision_dir / relative).read_text(encoding="utf-8"))
        expected.append({
            "request_path": relative,
            "response_path": request["response_path"],
            "task": request["task"],
            "batch_id": request["batch_id"],
        })
    assert result["status"] == "awaiting_hermes"
    assert result["handoffs"] == expected


def test_environment_failure_retains_a_complete_candidate_built_before_render_assurance(tmp_path, monkeypatch):
    reference = json.loads((ROOT / "tests/fixtures/retrospective-acceptance-source.json").read_text(encoding="utf-8"))
    run_dir = tmp_path / "run"
    reference_path = run_dir / "reference/study.reference.json"
    reference_path.parent.mkdir(parents=True)
    reference_path.write_text(json.dumps(reference), encoding="utf-8")
    assert prepare(run_dir)["status"] == "awaiting_approval"
    approval = approve(run_dir, approved_by="reviewer")

    drafting = generate(run_dir)
    revision_dir = run_dir / "revisions" / approval["revision_id"]
    for relative in drafting["requests"]:
        request = json.loads((revision_dir / relative).read_text(encoding="utf-8"))
        response = recorded_acceptance_response(request)
        response_path = revision_dir / request["response_path"]
        response_path.parent.mkdir(parents=True, exist_ok=True)
        response_path.write_text(json.dumps(response), encoding="utf-8")

    monkeypatch.setattr(workflow, "render_assurance", lambda *_args, **_kwargs: {
        "schema_version": "render-assurance/v1",
        "status": "blocked",
        "fonts": {},
        "font_substitutions": {},
        "candidate": {"files": []},
        "render": {"status": "blocked", "renderer_attempts": [], "page_renderer_attempts": [], "findings": []},
        "findings": [{"category": "renderer", "field": "renderer", "recovery_class": "adapter_fault", "action": "advance_adapter", "issue": "fallback stack unavailable"}],
        "diagnostic": {"recovery_class": "adapter_fault", "action": "advance_adapter", "outcome": "adapters_exhausted", "candidate_disposition": "preserved", "office_attempts": [], "page_attempts": []},
    })

    result = generate(run_dir)

    assert result["status"] == "blocked"
    assert result["stage"] == "render_assurance"
    assert (revision_dir / "candidate/protocol.docx").is_file()
    assert [item["path"] for item in result["candidate_outputs"]] == ["candidate/protocol.docx"]
    assert result["client_outputs"] == []


def test_repeated_adapter_exhaustion_reuses_candidate_without_drafting_or_regeneration(tmp_path, monkeypatch):
    reference = json.loads((ROOT / "tests/fixtures/retrospective-acceptance-source.json").read_text(encoding="utf-8"))
    run_dir = tmp_path / "run"
    reference_path = run_dir / "reference/study.reference.json"
    reference_path.parent.mkdir(parents=True)
    reference_path.write_text(json.dumps(reference), encoding="utf-8")
    assert prepare(run_dir)["status"] == "awaiting_approval"
    approval = approve(run_dir, approved_by="reviewer")
    drafting = generate(run_dir)
    revision_dir = run_dir / "revisions" / approval["revision_id"]
    for relative in drafting["requests"]:
        request = json.loads((revision_dir / relative).read_text(encoding="utf-8"))
        response_path = revision_dir / request["response_path"]
        response_path.parent.mkdir(parents=True, exist_ok=True)
        response_path.write_text(json.dumps(recorded_acceptance_response(request)), encoding="utf-8")

    renders = 0
    real_render_documents = workflow.render_documents

    def counted_render(*args, **kwargs):
        nonlocal renders
        renders += 1
        return real_render_documents(*args, **kwargs)

    def exhausted(*_args, **_kwargs):
        candidate = revision_dir / "candidate/protocol.docx"
        return {
            "schema_version": "render-assurance/v1",
            "status": "blocked",
            "fonts": {},
            "font_substitutions": {},
            "candidate": {"files": [{"path": "candidate/protocol.docx", "sha256": hashlib.sha256(candidate.read_bytes()).hexdigest(), "bytes": candidate.stat().st_size}]},
            "render": {"status": "blocked", "renderer_attempts": [], "page_renderer_attempts": [], "findings": []},
            "findings": [{"category": "renderer", "field": "rendering", "recovery_class": "adapter_fault", "action": "advance_adapter", "issue": "all adapters exhausted"}],
            "diagnostic": {"recovery_class": "adapter_fault", "action": "advance_adapter", "outcome": "adapters_exhausted", "candidate_disposition": "preserved", "office_attempts": [], "page_attempts": []},
        }

    monkeypatch.setattr(workflow, "render_documents", counted_render)
    monkeypatch.setattr(workflow, "render_assurance", exhausted)

    first = generate(run_dir)
    candidate_hash = hashlib.sha256((revision_dir / "candidate/protocol.docx").read_bytes()).hexdigest()
    second = generate(run_dir)
    state = json.loads(reference_path.read_text(encoding="utf-8"))["generation"]

    assert first["stage"] == second["stage"] == "render_assurance"
    assert first["repair_report"] == second["repair_report"] == "reference/render-assurance-diagnostic.md"
    assert hashlib.sha256((revision_dir / "candidate/protocol.docx").read_bytes()).hexdigest() == candidate_hash
    assert renders == 1
    assert state["attempts"] == {}


def test_generation_does_not_reopen_source_intake_for_missing_optional_prs_study_type(tmp_path):
    run_dir = tmp_path / "run"
    reference_path = run_dir / "reference/study.reference.json"
    reference_path.parent.mkdir(parents=True)
    reference_path.write_text(json.dumps(fixture()), encoding="utf-8")
    assert prepare(run_dir)["status"] == "awaiting_approval"
    approval = approve(run_dir, approved_by="reviewer")
    revision_dir = run_dir / "revisions" / approval["revision_id"]

    working = json.loads(reference_path.read_text(encoding="utf-8"))
    approved = json.loads((revision_dir / "approved-reference.json").read_text(encoding="utf-8"))
    working["regulatory"]["prs"].pop("study_type")
    approved["regulatory"]["prs"].pop("study_type")
    (revision_dir / "approved-reference.json").write_text(json.dumps(approved), encoding="utf-8")
    working["approval"]["approved_reference_sha256"] = hashlib.sha256(
        (revision_dir / "approved-reference.json").read_bytes()
    ).hexdigest()
    reference_path.write_text(json.dumps(working), encoding="utf-8")

    result = generate(run_dir)

    assert result["status"] == "awaiting_hermes"
    assert result["stage"] != "approval_gate"
    assert not (revision_dir / "attempts").exists()


def test_source_gap_response_is_a_retryable_drafting_failure_not_a_reviewer_question(tmp_path):
    reference = fixture()
    batch = next(item for item in batch_plan("Prospective") if item.batch_id == "icf-narrative")
    path = create_drafting_request(
        repo_root=ROOT,
        revision_dir=tmp_path,
        revision_id="r-test",
        reference=reference,
        batch=batch,
        attempts={item: 1 for item in batch.section_ids},
        wave="initial",
    )
    request = json.loads(path.read_text(encoding="utf-8"))
    response = recorded_acceptance_response(request)
    risk = next(item for item in response["section_results"] if item["section_id"] == "icf.risks")
    risk.update({"outcome": "source_gap", "paragraphs": [], "question": "Supply more information."})

    _, findings = validate_response(request, response)

    risk_findings = [item for item in findings if item["field"] == "icf.risks"]
    assert risk_findings
    assert all(item["category"] == "drafting" for item in risk_findings)
    assert "ask" not in json.dumps(risk_findings).casefold()


def test_legacy_question_capable_handoff_is_ignored_after_contract_upgrade(tmp_path):
    reference = fixture()
    batch = batch_plan("Prospective")[0]
    current_path = create_drafting_request(
        repo_root=ROOT,
        revision_dir=tmp_path,
        revision_id="r-test",
        reference=reference,
        batch=batch,
        attempts={item: 1 for item in batch.section_ids},
        wave="initial",
    )
    legacy = json.loads(current_path.read_text(encoding="utf-8"))
    legacy.update({
        "schema_version": "hermes-request/v1",
        "prompt_version": "section-drafting-v1",
        "request_id": "legacy-question-capable-request",
        "response_path": "hermes/responses/legacy-question-capable-request.json",
    })
    legacy["governing_resources"]["prompt_version"] = "section-drafting-v1"
    legacy.pop("request_sha256")
    legacy["request_sha256"] = sha256_value(legacy)
    legacy_path = tmp_path / "hermes/requests/legacy-question-capable-request.json"
    legacy_path.write_text(json.dumps(legacy), encoding="utf-8")
    legacy_response = tmp_path / legacy["response_path"]
    legacy_response.parent.mkdir(parents=True, exist_ok=True)
    legacy_response.write_text(json.dumps({"outcome": "source_gap", "question": "Supply more information."}), encoding="utf-8")

    expected = governing_resources(ROOT, reference)
    assert pending_requests(tmp_path, expected) == [current_path]
    assert ingest_responses(tmp_path, expected) == []


def test_prepare_requires_editable_source_truth_file_delivery(tmp_path):
    run_dir = tmp_path / "run"
    reference_path = run_dir / "reference/study.reference.json"
    reference_path.parent.mkdir(parents=True)
    reference_path.write_text(json.dumps(fixture()), encoding="utf-8")

    result = prepare(run_dir)
    delivery = result["review_delivery"]
    source = Path(delivery["absolute_path"])

    assert result["status"] == "awaiting_approval"
    assert delivery["mode"] == "file"
    assert delivery["path"] == result["source_of_truth"]
    assert delivery["mime_type"] == "text/markdown"
    assert delivery["editable"] is True
    assert delivery["inline_chat"] is False
    assert source.is_file()
    assert source.read_text(encoding="utf-8").count("<!-- field:") >= 35


def test_prepare_defaults_editable_document_control_date_without_inventing_version(tmp_path):
    reference = fixture()
    reference["meta"].pop("date", None)
    reference["meta"].pop("version", None)
    run_dir = tmp_path / "run"
    reference_path = run_dir / "reference/study.reference.json"
    reference_path.parent.mkdir(parents=True)
    reference_path.write_text(json.dumps(reference), encoding="utf-8")

    result = prepare(run_dir, today=date(2026, 8, 24))
    prepared = json.loads(reference_path.read_text(encoding="utf-8"))
    source = (run_dir / result["source_of_truth"]).read_text(encoding="utf-8")

    assert prepared["meta"]["date"] == "24 Aug 2026"
    assert prepared["meta"]["version"] == ""
    assert "<!-- field: meta.date -->\n24 Aug 2026\n<!-- /field -->" in source
    assert "<!-- field: meta.version -->\n\n<!-- /field -->" in source
    assert "<!-- field: meta.icf_template -->" in source


def test_generation_resource_change_cannot_overwrite_an_approved_revision(tmp_path, monkeypatch):
    run_dir = tmp_path / "run"
    reference_path = run_dir / "reference/study.reference.json"
    reference_path.parent.mkdir(parents=True)
    reference_path.write_text(json.dumps(fixture()), encoding="utf-8")
    prepare(run_dir)
    approval = approve(run_dir, approved_by="reviewer")
    original = workflow.governing_resources

    def changed_resources(repo_root, reference, **kwargs):
        resources = original(repo_root, reference, **kwargs)
        resources["prompt_version"] = "changed-after-approval"
        return resources

    monkeypatch.setattr(workflow, "governing_resources", changed_resources)

    result = generate(run_dir)

    assert result["status"] == "blocked"
    assert result["stage"] == "approval_gate"
    assert approval["revision_id"] in json.dumps(json.loads(reference_path.read_text(encoding="utf-8")))


def test_prepare_exposes_optional_clinical_and_prs_fields_in_the_editable_source(tmp_path):
    reference = fixture()
    for path in ("condition",):
        reference["study"].pop(path, None)
    run_dir = tmp_path / "run"
    reference_path = run_dir / "reference/study.reference.json"
    reference_path.parent.mkdir(parents=True)
    reference_path.write_text(json.dumps(reference), encoding="utf-8")

    result = prepare(run_dir, today=date(2026, 8, 24))
    source = (run_dir / result["source_of_truth"]).read_text(encoding="utf-8")

    assert "<!-- field: study.condition -->" in source
    assert "<!-- field: risks_benefits.costs -->" in source
    assert "<!-- field: risks_benefits.injury_handling -->" in source
    assert "<!-- field: regulatory.prs.time_perspective -->" in source
    assert "<!-- field: regulatory.prs.start_date -->" in source
    assert "<!-- field: regulatory.prs.study_completion_date -->" in source
    assert "<!-- field: statistics.analysis_populations -->" in source
    assert "<!-- field: statistics.missing_data_handling -->" in source
    assert "analysis_method" in source
    assert "description" in source


def test_prepare_derives_prs_study_type_from_unambiguous_design_evidence(tmp_path):
    reference = fixture()
    reference["regulatory"]["prs"].pop("study_type")
    run_dir = tmp_path / "run"
    reference_path = run_dir / "reference/study.reference.json"
    reference_path.parent.mkdir(parents=True)
    reference_path.write_text(json.dumps(reference), encoding="utf-8")

    result = prepare(run_dir, today=date(2026, 8, 24))
    prepared = json.loads(reference_path.read_text(encoding="utf-8"))
    source = (run_dir / result["source_of_truth"]).read_text(encoding="utf-8")

    assert result["status"] == "awaiting_approval"
    assert prepared["regulatory"]["prs"]["study_type"] == "Observational"
    assert "<!-- field: regulatory.prs.study_type -->\nObservational\n<!-- /field -->" in source


def test_prepare_does_not_invent_optional_unknown_prs_study_type(tmp_path):
    reference = fixture()
    reference["regulatory"]["prs"].pop("study_type")
    reference["design"]["study_design"] = "Prospective, single-center device study."
    run_dir = tmp_path / "run"
    reference_path = run_dir / "reference/study.reference.json"
    reference_path.parent.mkdir(parents=True)
    reference_path.write_text(json.dumps(reference), encoding="utf-8")

    result = prepare(run_dir, today=date(2026, 8, 24))
    assert result["status"] == "awaiting_approval"
    assert result["stage"] == "source_review"
    prepared = json.loads(reference_path.read_text(encoding="utf-8"))
    assert "study_type" not in prepared["regulatory"]["prs"]
    assert not (run_dir / "revisions").exists()


def test_approval_does_not_restore_a_reviewer_cleared_prs_study_type(tmp_path):
    reference = fixture()
    reference["regulatory"]["prs"].pop("study_type")
    run_dir = tmp_path / "run"
    reference_path = run_dir / "reference/study.reference.json"
    reference_path.parent.mkdir(parents=True)
    reference_path.write_text(json.dumps(reference), encoding="utf-8")

    prepared = prepare(run_dir, today=date(2026, 8, 24))
    source_path = run_dir / prepared["source_of_truth"]
    source = source_path.read_text(encoding="utf-8")
    source_path.write_text(
        source.replace(
            "<!-- field: regulatory.prs.study_type -->\nObservational\n<!-- /field -->",
            "<!-- field: regulatory.prs.study_type -->\n\n<!-- /field -->",
        ),
        encoding="utf-8",
    )

    result = approve(run_dir, approved_by="reviewer")

    assert result["status"] == "passed"
    approved = json.loads((run_dir / "revisions" / result["revision_id"] / "approved-reference.json").read_text())
    assert approved["regulatory"]["prs"]["study_type"] is None


def test_protocol_leaf_heading_without_body_content_is_blocked(tmp_path):
    reference = json.loads((ROOT / "tests/fixtures/retrospective-acceptance-source.json").read_text(encoding="utf-8"))
    candidate = tmp_path / "candidate"
    candidate.mkdir()
    document = __import__("docx").Document()
    for section in __import__("contracts").protocol_contract("Retrospective"):
        style = "Heading 1" if section.number.count(".") == 1 else "Heading 2"
        document.add_paragraph(f"{section.number} {section.title}", style=style)
        if section.role == "leaf" and section.section_id != "introduction":
            document.add_paragraph(f"Complete source-grounded content for {section.title.lower()} is present in this section.")
    document.save(candidate / "protocol.docx")

    findings = deterministic_content_check(tmp_path, reference)

    assert any(
        item.get("field") == "introduction" and "no substantive content" in item.get("issue", "").casefold()
        for item in findings
    )


def test_unlocalized_protocol_duplicate_stops_as_a_structural_defect(tmp_path):
    reference = json.loads((ROOT / "tests/fixtures/retrospective-acceptance-source.json").read_text(encoding="utf-8"))
    run_dir = tmp_path / "run"
    revision = run_dir / "revisions/r-test"
    candidate = revision / "candidate"
    candidate.mkdir(parents=True)
    document = Document()
    document.add_paragraph(reference["study"]["title"])
    repeated = "This repeated paragraph appears before any contracted section heading is recognized."
    document.add_paragraph(repeated)
    document.add_paragraph(repeated)
    for section in __import__("contracts").protocol_contract("Retrospective"):
        style = "Heading 1" if section.number.count(".") == 1 else "Heading 2"
        document.add_paragraph(f"{section.number} {section.title}", style=style)
        if section.role == "leaf":
            document.add_paragraph(f"Complete source-grounded content for {section.title.lower()} is present in this section.")
    document.save(candidate / "protocol.docx")

    findings = deterministic_content_check(revision, reference)
    duplicate = next(item for item in findings if "duplicated in protocol section protocol" in item["issue"])
    reference_path = run_dir / "reference/study.reference.json"
    reference_path.parent.mkdir(parents=True)
    reference_path.write_text(json.dumps({"generation": {}}), encoding="utf-8")
    result = workflow._quality_retry(
        run_dir,
        reference_path,
        {"generation": {}},
        reference,
        revision,
        {},
        [duplicate],
        "quality",
    )

    assert (duplicate["recovery_class"], duplicate["action"]) == ("document_structure_defect", "preserve_and_stop")
    assert result["stage"] == "document_structure"


def test_stale_icf_cost_language_uses_the_supported_branch_drafting_target(tmp_path):
    from rendering import render_documents

    reference = fixture()
    render_documents(ROOT, tmp_path, reference, {"protocol": [], "icf": {}, "prs": {}})
    icf_path = tmp_path / "candidate/icf.docx"
    document = Document(icf_path)
    document.add_paragraph("All charges for medical care will be billed to your insurance company.")
    document.save(icf_path)

    findings = deterministic_content_check(tmp_path, reference)
    stale = next(item for item in findings if "all charges for medical care" in item["issue"])

    assert stale["target_ids"] == ["icf.costs"]
    assert (stale["recovery_class"], stale["action"]) == ("drafting_defect", "retry_drafting_target")


def test_complete_bundle_is_part_of_governing_resources():
    resources = governing_resources(ROOT, fixture())
    bundle = resources["contracted_template_bundle"]
    assert "assets/client-templates/prs/clinicaltrials_prs_full_placeholder_template.xml" in bundle["resource_hashes"]
    assert "assets/client-templates/reference/protocol-reference.docx" in bundle["resource_hashes"]
    assert "assets/client-templates/reference/advarra-icf-reference.docx" in bundle["resource_hashes"]
    assert set(resources["implementation_sha256"]) == {
        f"scripts/{name}"
        for name in ("contracts.py", "drafting.py", "prs_xml.py", "quality.py", "rendering.py", "workflow.py")
    }


def test_content_verifier_receives_authorized_boilerplate_and_blank_field_policy(tmp_path):
    candidate = tmp_path / "candidate"
    candidate.mkdir()
    for name in ("protocol.docx", "icf.docx", "study.xml"):
        (candidate / name).write_bytes(b"candidate")

    paths = create_verification_requests(tmp_path, fixture(), {"renderer": {}, "artifacts": []})
    content_path = next(path for path in paths if json.loads(path.read_text())["task"] == "clinical_content_verification")
    request = json.loads(content_path.read_text(encoding="utf-8"))

    assert (tmp_path / request["response_path"]).parent.is_dir()
    assert request["authorized_boilerplate"]["version"]
    assert request["authorized_boilerplate"]["sections"]["costs"]
    assert "Do not fail optional fields" in request["instructions"]
    assert "unknown version" in request["instructions"]
    assessed_ids = {item["section_id"] for item in request["sections"] if item["artifact"] == "icf"}
    expected_retained = {section_id for section_id, _title in icf_retained_sections("Prospective", "Advarra")}
    assert expected_retained <= assessed_ids
    assert "icf.introduction" in assessed_ids
    assert "icf.leaving-study" in assessed_ids


@pytest.mark.parametrize(
    ("assignment_method", "expected_in_scope"),
    [
        (None, False),
        ("Single observational cohort; no treatment assignment", True),
    ],
)
def test_content_verifier_scope_matches_optional_source_sections(
    tmp_path, assignment_method, expected_in_scope
):
    reference = fixture()
    if assignment_method is None:
        reference["design"].pop("assignment_method", None)
    else:
        reference["design"]["assignment_method"] = assignment_method
    candidate = tmp_path / "candidate"
    candidate.mkdir()
    for name in ("protocol.docx", "icf.docx", "study.xml"):
        (candidate / name).write_bytes(b"candidate")

    paths = create_verification_requests(
        tmp_path,
        reference,
        {"renderer": {}, "artifacts": []},
    )
    content_path = next(
        path for path in paths
        if json.loads(path.read_text())["task"] == "clinical_content_verification"
    )
    request = json.loads(content_path.read_text(encoding="utf-8"))
    protocol_section_ids = {
        item["section_id"]
        for item in request["sections"]
        if item["artifact"] == "protocol"
    }

    assert ("study-design.assignment" in protocol_section_ids) is expected_in_scope


def test_visual_verification_is_split_by_document_for_concurrent_review(tmp_path):
    candidate = tmp_path / "candidate"
    candidate.mkdir()
    for artifact in ("protocol", "icf"):
        document = Document()
        document.add_paragraph(f"Stable {artifact} content")
        document.save(candidate / f"{artifact}.docx")
    (candidate / "study.xml").write_bytes(b"candidate")
    rendered = tmp_path / "rendered"
    rendered.mkdir()
    artifacts = []
    for artifact, page_count in (("protocol", 11), ("icf", 6)):
        pdf = rendered / f"{artifact}.pdf"
        pdf.write_bytes(f"{artifact}-pdf".encode())
        pages = []
        for page_number in range(1, page_count + 1):
            page = rendered / artifact / f"page-{page_number}.png"
            page.parent.mkdir(parents=True, exist_ok=True)
            page.write_bytes(f"{artifact}-{page_number}".encode())
            pages.append({
                "page": page_number,
                "path": page.relative_to(tmp_path).as_posix(),
                "sha256": hashlib.sha256(page.read_bytes()).hexdigest(),
            })
        docx = candidate / f"{artifact}.docx"
        artifacts.append({
            "artifact": artifact,
            "docx": docx.relative_to(tmp_path).as_posix(),
            "docx_sha256": hashlib.sha256(docx.read_bytes()).hexdigest(),
            "pdf": pdf.relative_to(tmp_path).as_posix(),
            "pdf_sha256": hashlib.sha256(pdf.read_bytes()).hexdigest(),
            "pages": pages,
        })
    render_report = {
        "renderer": {"kind": "test"},
        "artifacts": artifacts,
    }

    paths = create_verification_requests(tmp_path, fixture(), render_report)
    requests = [json.loads(path.read_text(encoding="utf-8")) for path in paths]
    visual = [item for item in requests if item["task"] == "rendered_page_visual_verification"]

    assert len(visual) == 2
    assert {item["artifacts"][0]["artifact"] for item in visual} == {"protocol", "icf"}
    assert all(len(item["artifacts"]) == 1 for item in visual)
    assert len({item["response_path"] for item in visual}) == 2
    for request in requests:
        response_path = tmp_path / request["response_path"]
        response_path.parent.mkdir(parents=True, exist_ok=True)
        response_path.write_text(json.dumps(acceptance_verification(request)), encoding="utf-8")

    findings, evidence = validate_verifications(tmp_path)

    assert findings == []
    assert set(evidence) == {item["request_id"] for item in visual} | {"clinical_content_verification"}
    for request in visual:
        assert evidence[request["request_id"]]["artifacts"] == request["artifacts"]

    protocol_artifact = next(item for item in artifacts if item["artifact"] == "protocol")
    protocol_pdf = tmp_path / protocol_artifact["pdf"]
    protocol_page = tmp_path / protocol_artifact["pages"][0]["path"]
    protocol_pdf.write_bytes(b"repaired-protocol-pdf")
    protocol_page.write_bytes(b"repaired-protocol-page")
    protocol_docx = tmp_path / protocol_artifact["docx"]
    formatted = Document(protocol_docx)
    formatted.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    formatted.save(protocol_docx)
    protocol_artifact["docx_sha256"] = hashlib.sha256(protocol_docx.read_bytes()).hexdigest()
    protocol_artifact["pdf_sha256"] = hashlib.sha256(protocol_pdf.read_bytes()).hexdigest()
    protocol_artifact["pages"][0]["sha256"] = hashlib.sha256(protocol_page.read_bytes()).hexdigest()

    create_verification_requests(tmp_path, fixture(), render_report)

    response_by_artifact = {
        item["artifacts"][0]["artifact"]: tmp_path / item["response_path"]
        for item in visual
    }
    content_response = tmp_path / next(
        item["response_path"] for item in requests if item["task"] == "clinical_content_verification"
    )
    assert not response_by_artifact["protocol"].exists()
    assert response_by_artifact["icf"].is_file()
    assert not content_response.exists()

    content_changed = Document(protocol_docx)
    content_changed.paragraphs[0].text = "Changed protocol content"
    content_changed.save(protocol_docx)
    create_verification_requests(tmp_path, fixture(), render_report)
    assert not content_response.exists()


def test_response_with_wrong_request_hash_is_rejected(tmp_path):
    reference = fixture(); batch = batch_plan("Prospective")[0]
    path = create_drafting_request(repo_root=ROOT, revision_dir=tmp_path, revision_id="r-test", reference=reference, batch=batch, attempts={item: 1 for item in batch.section_ids}, wave="initial")
    request = json.loads(path.read_text(encoding="utf-8")); response = recorded_acceptance_response(request)
    response["request_sha256"] = "wrong"
    accepted, findings = validate_response(request, response)
    assert accepted is None
    assert any(item["field"] == "request_sha256" for item in findings)


def test_none_source_value_is_grounded_by_an_explicit_negative_statement(tmp_path):
    reference = fixture()
    reference["risks_benefits"]["compensation_or_reimbursement"] = "None"
    batch = next(item for item in batch_plan("Prospective") if item.batch_id == "icf-narrative")
    path = create_drafting_request(
        repo_root=ROOT,
        revision_dir=tmp_path,
        revision_id="r-none",
        reference=reference,
        batch=batch,
        attempts={item: 1 for item in batch.section_ids},
        wave="initial",
    )
    request = json.loads(path.read_text(encoding="utf-8"))
    response = recorded_acceptance_response(request)
    payment = next(item for item in response["section_results"] if item["section_id"] == "icf.payment")
    payment["paragraphs"] = [{
        "text": "You will not be paid for taking part in this study.",
        "evidence_refs": ["source:risks_benefits.compensation_or_reimbursement"],
        "boilerplate_refs": [],
    }]

    accepted, findings = validate_response(request, response)

    assert "icf.payment" in {draft["section_id"] for draft in (accepted or {}).get("drafts", [])}
    assert not [item for item in findings if item["field"] == "icf.payment"]


def test_invalid_section_preserves_other_sections_and_increments_once(tmp_path):
    reference = fixture(); batch = batch_plan("Prospective")[0]
    path = create_drafting_request(repo_root=ROOT, revision_dir=tmp_path, revision_id="r-test", reference=reference, batch=batch, attempts={item: 1 for item in batch.section_ids}, wave="initial")
    request = json.loads(path.read_text(encoding="utf-8")); response = recorded_acceptance_response(request)
    failed = response["section_results"][0]["section_id"]
    response["section_results"][0]["paragraphs"][0]["text"] = "Too short"
    accepted, findings = validate_response(request, response)
    assert accepted is not None
    assert failed not in {draft["section_id"] for draft in accepted["drafts"]}
    for finding in findings: finding["target_ids"] = [failed]
    attempts, exhausted = retry_attempts(findings, {})
    assert attempts[failed] == 2
    assert exhausted == []


def test_partial_batch_ingestion_retries_only_the_failed_section(tmp_path):
    reference = fixture()
    batch = batch_plan("Prospective")[0]
    path = create_drafting_request(
        repo_root=ROOT,
        revision_dir=tmp_path,
        revision_id="r-partial-ingest",
        reference=reference,
        batch=batch,
        attempts={item: 1 for item in batch.section_ids},
        wave="initial",
    )
    request = json.loads(path.read_text(encoding="utf-8"))
    response = recorded_acceptance_response(request)
    failed = response["section_results"][0]["section_id"]
    response["section_results"][0]["paragraphs"][0]["text"] = "Too short"
    response_path = tmp_path / request["response_path"]
    response_path.parent.mkdir(parents=True, exist_ok=True)
    response_path.write_text(json.dumps(response), encoding="utf-8")

    findings = ingest_responses(tmp_path, governing_resources(ROOT, reference))
    valid = set(batch.section_ids) - {failed}
    assert valid
    assert all(accepted_draft(tmp_path, section_id, governing_resources(ROOT, reference)) for section_id in valid)
    attempts, exhausted = retry_attempts(findings, {})
    assert exhausted == []
    created = schedule_requests(
        repo_root=ROOT,
        revision_dir=tmp_path,
        revision_id="r-partial-ingest",
        reference=reference,
        attempts=attempts,
        wave="retry",
        findings=findings,
    )
    retried_batch = [
        json.loads(item.read_text(encoding="utf-8"))
        for item in created
        if json.loads(item.read_text(encoding="utf-8"))["batch_id"] == batch.batch_id
    ]
    assert len(retried_batch) == 1
    assert [item["section_id"] for item in retried_batch[0]["section_contracts"]] == [failed]


def test_prs_retry_preserves_the_accepted_narrative_target(tmp_path):
    reference = fixture()
    batch = next(item for item in batch_plan("Prospective") if item.batch_id == "prs-narrative")
    request_path = create_drafting_request(
        repo_root=ROOT,
        revision_dir=tmp_path,
        revision_id="r-prs-partial",
        reference=reference,
        batch=batch,
        attempts={item: 1 for item in batch.section_ids},
        wave="initial",
    )
    request = json.loads(request_path.read_text(encoding="utf-8"))
    response = recorded_acceptance_response(request)
    response["narrative"]["detailed_description"]["text"] = "Too short"

    accepted, findings = validate_response(request, response)

    assert set((accepted or {})["narrative"]) == {"brief_summary"}
    assert {target for finding in findings for target in finding.get("target_ids", [])} == {"prs.detailed-description"}

    retry_path = create_drafting_request(
        repo_root=ROOT,
        revision_dir=tmp_path,
        revision_id="r-prs-partial",
        reference=reference,
        batch=batch,
        target_ids=("prs.detailed-description",),
        attempts={"prs.detailed-description": 2},
        wave="retry",
        findings=findings,
    )
    retry = json.loads(retry_path.read_text(encoding="utf-8"))
    assert [item["section_id"] for item in retry["section_contracts"]] == ["prs.detailed-description"]
    assert set(response_template(retry)["narrative"]) == {"detailed_description"}


def test_prs_narrative_rejects_evidence_outside_the_target_contract(tmp_path):
    reference = fixture()
    batch = next(item for item in batch_plan("Prospective") if item.batch_id == "prs-narrative")
    request_path = create_drafting_request(
        repo_root=ROOT,
        revision_dir=tmp_path,
        revision_id="r-prs-scope",
        reference=reference,
        batch=batch,
        attempts={item: 1 for item in batch.section_ids},
        wave="initial",
    )
    request = json.loads(request_path.read_text(encoding="utf-8"))
    response = recorded_acceptance_response(request)
    response["narrative"]["brief_summary"]["evidence_refs"].append("source:study.background")

    accepted, findings = validate_response(request, response)

    assert set((accepted or {})["narrative"]) == {"detailed_description"}
    assert any(
        item["field"] == "prs.brief-summary" and "not approved for this target" in item["issue"]
        for item in findings
    )


def test_visual_gate_rejects_unassessed_pages(tmp_path):
    request_dir = tmp_path / "hermes/verification-requests"; response_dir = tmp_path / "hermes/verification-responses"
    request_dir.mkdir(parents=True); response_dir.mkdir(parents=True)
    request = {"schema_version": "hermes-verification/v1", "request_id": "r.verify.visual", "task": "rendered_page_visual_verification", "response_path": "hermes/verification-responses/r.verify.visual.json", "artifacts": [{"artifact": "protocol", "pages": [{"page": 1, "sha256": "one"}, {"page": 2, "sha256": "two"}]}]}
    request["request_sha256"] = verification_request_sha256(request)
    response = {"schema_version": RESPONSE_SCHEMA, "request_id": request["request_id"], "request_sha256": request["request_sha256"], "task": request["task"], "producer": {"model_id": "test", "reviewer_id": "visual-reviewer"}, "status": "passed", "findings": [], "page_assessments": [{"artifact": "protocol", "page": 1, "sha256": "one", "status": "passed", "checks": list(VISUAL_CHECKS)}]}
    (request_dir / "r.verify.visual.json").write_text(json.dumps(request), encoding="utf-8")
    (response_dir / "r.verify.visual.json").write_text(json.dumps(response), encoding="utf-8")
    findings, _ = validate_verifications(tmp_path)
    incomplete = next(item for item in findings if item["field"] == "page_assessments")
    assert incomplete["recovery_class"] == "verifier_transient"
    assert incomplete["action"] == "retry_verifier"


def test_final_exact_artifact_review_rejects_missing_or_partial_review_requests(tmp_path, monkeypatch):
    revision = tmp_path / "revision"
    candidate = revision / "candidate"
    candidate.mkdir(parents=True)
    for name, payload in (
        ("protocol.docx", b"protocol"),
        ("icf.docx", b"icf"),
        ("study.xml", b"<clinical_study/>"),
    ):
        (candidate / name).write_bytes(payload)
    reference = fixture()
    monkeypatch.setattr(quality, "deterministic_content_check", lambda *_args: [])
    monkeypatch.setattr(quality, "validate_verifications", lambda *_args: ([], {}))

    report = quality.quality_report(
        revision,
        reference,
        {"status": "passed", "renderer": {"kind": "test"}, "artifacts": [], "findings": []},
        {"status": "passed", "findings": []},
    )

    assert report["status"] == "blocked"
    assert report["final_exact_artifact_review"]["status"] == "blocked"
    assert report["final_exact_artifact_review"]["every_page"] is False
    assert any("package-wide content review" in item["issue"] for item in report["findings"])
    assert any("document-scoped visual review" in item["issue"] for item in report["findings"])


def test_final_exact_artifact_review_rejects_grouped_or_empty_visual_requests(tmp_path, monkeypatch):
    revision = tmp_path / "revision"
    candidate = revision / "candidate"
    candidate.mkdir(parents=True)
    for name, payload in (
        ("protocol.docx", b"protocol"),
        ("icf.docx", b"icf"),
        ("study.xml", b"<clinical_study/>"),
    ):
        (candidate / name).write_bytes(payload)
    bound = {
        "review_set": 1,
        "request_sha256": "a" * 64,
        "response_sha256": "b" * 64,
        "producer_model_id": "client-selected-model",
        "producer_reviewer_id": "independent-reviewer",
    }
    evidence = {
        "content": {
            **bound,
            "task": "clinical_content_verification",
            "request_id": "review.content",
            "artifacts": [
                {"path": f"candidate/{name}"}
                for name in ("protocol.docx", "icf.docx", "study.xml")
            ],
        },
        "visual-grouped": {
            **bound,
            "task": "rendered_page_visual_verification",
            "request_id": "review.visual.grouped",
            "renderer": {"kind": "test-office"},
            "page_renderer": {"kind": "test-pages"},
            "artifacts": [{"artifact": "protocol"}, {"artifact": "icf"}],
        },
        "visual-empty": {
            **bound,
            "task": "rendered_page_visual_verification",
            "request_id": "review.visual.empty",
            "renderer": {"kind": "test-office"},
            "page_renderer": {"kind": "test-pages"},
            "artifacts": [],
        },
    }
    monkeypatch.setattr(quality, "deterministic_content_check", lambda *_args: [])
    monkeypatch.setattr(quality, "validate_verifications", lambda *_args: ([], evidence))

    report = quality.quality_report(
        revision,
        fixture(),
        {"status": "passed", "renderer": {"kind": "test"}, "artifacts": [], "findings": []},
        {"status": "passed", "findings": []},
    )

    assert report["status"] == "blocked"
    assert any("document-scoped visual review" in item["issue"] for item in report["findings"])


def test_final_exact_artifact_review_rejects_sampled_page_inventory(tmp_path):
    revision = tmp_path / "revision"
    candidate = revision / "candidate"
    rendered = revision / "rendered/protocol"
    candidate.mkdir(parents=True)
    rendered.mkdir(parents=True)
    (candidate / "protocol.docx").write_bytes(b"protocol")
    (revision / "rendered/protocol.pdf").write_bytes(b"pdf")
    for page_number in (1, 2):
        (rendered / f"page-{page_number}.png").write_bytes(f"page-{page_number}".encode())

    def digest(relative):
        return hashlib.sha256((revision / relative).read_bytes()).hexdigest()

    complete_artifact = {
        "artifact": "protocol",
        "status": "passed",
        "docx": "candidate/protocol.docx",
        "docx_sha256": digest("candidate/protocol.docx"),
        "pdf": "rendered/protocol.pdf",
        "pdf_sha256": digest("rendered/protocol.pdf"),
        "page_count": 2,
        "pages": [
            {
                "page": page_number,
                "path": f"rendered/protocol/page-{page_number}.png",
                "sha256": digest(f"rendered/protocol/page-{page_number}.png"),
            }
            for page_number in (1, 2)
        ],
    }
    bound = {
        "review_set": 1,
        "request_sha256": "a" * 64,
        "response_sha256": "b" * 64,
        "producer_model_id": "client-selected-model",
        "producer_reviewer_id": "protocol-visual-reviewer",
    }
    evidence = {
        "content": {
            **bound,
            "task": "clinical_content_verification",
            "request_id": "review.content",
            "artifacts": [{"path": "candidate/protocol.docx"}],
        },
        "visual": {
            **bound,
            "task": "rendered_page_visual_verification",
            "request_id": "review.visual.protocol",
            "renderer": {"kind": "test-office"},
            "page_renderer": {"kind": "test-pages"},
            "artifacts": [{**complete_artifact, "pages": complete_artifact["pages"][:1]}],
        },
    }

    render_report = {
        "status": "passed",
        "renderer": {"kind": "test-office"},
        "page_renderer": {"kind": "test-pages"},
        "artifacts": [complete_artifact],
    }
    findings = quality._final_verification_scope_findings(
        revision,
        {"meta": {"study_type": "Retrospective"}},
        render_report,
        evidence,
    )

    assert any("complete rendered page inventory" in item["issue"] for item in findings)

    evidence["visual"]["artifacts"] = [complete_artifact]
    evidence["visual"]["renderer"] = {"kind": "different-office"}
    findings = quality._final_verification_scope_findings(
        revision,
        {"meta": {"study_type": "Retrospective"}},
        render_report,
        evidence,
    )

    assert any("exact renderer and page renderer" in item["issue"] for item in findings)


def test_visual_gate_preserves_the_exact_failed_layout_element(tmp_path):
    request_dir = tmp_path / "hermes/verification-requests"; response_dir = tmp_path / "hermes/verification-responses"
    request_dir.mkdir(parents=True); response_dir.mkdir(parents=True)
    request = {"schema_version": "hermes-verification/v1", "request_id": "r.verify.visual", "task": "rendered_page_visual_verification", "response_path": "hermes/verification-responses/r.verify.visual.json", "artifacts": [{"artifact": "protocol", "pages": []}]}
    request["request_sha256"] = verification_request_sha256(request)
    response = {
        "schema_version": RESPONSE_SCHEMA,
        "request_id": request["request_id"],
        "request_sha256": request["request_sha256"],
        "task": request["task"],
        "producer": {"model_id": "test", "reviewer_id": "visual-reviewer"},
        "status": "failed",
        "findings": [{
            "artifact": "protocol",
            "page": 7,
            "check": "bad_table_split",
            "element": "Table 13.3.-1",
            "issue": "The contact table splits badly.",
        }],
        "page_assessments": [],
    }
    (request_dir / "r.verify.visual.json").write_text(json.dumps(request), encoding="utf-8")
    (response_dir / "r.verify.visual.json").write_text(json.dumps(response), encoding="utf-8")

    findings, _ = validate_verifications(tmp_path)

    failed_table = next(item for item in findings if item.get("check") == "bad_table_split")
    assert failed_table["element"] == "Table 13.3.-1"
    assert failed_table["target_ids"] == ["layout:protocol"]


def test_render_gate_detects_a_textless_pdf_page(tmp_path):
    from quality import _blank_pdf_pages

    path = tmp_path / "blank.pdf"
    writer = PdfWriter()
    writer.add_blank_page(width=612, height=792)
    with path.open("wb") as handle:
        writer.write(handle)

    assert _blank_pdf_pages(path) == [1]


def test_render_gate_does_not_call_a_sparse_body_page_textless(monkeypatch):
    import quality

    class Page:
        def extract_text(self):
            return (
                "Prospective sparse acceptance Page 3 of 11\n"
                "v 1.0 21 Aug 2026\n"
                "Variables • Primary outcome (Month 3)\n"
                "• Safety outcome (Month 3)\n"
                "Duration / Follow-up 3 months"
            )

    monkeypatch.setattr(quality, "PdfReader", lambda _: type("Reader", (), {"pages": [Page()]})())

    assert quality._blank_pdf_pages(Path("unused.pdf")) == []


def test_raw_document_word_count_alone_does_not_block_delivery(tmp_path):
    path = tmp_path / "short.docx"
    document = __import__("docx").Document()
    document.add_paragraph("Concise source-grounded content.")
    document.save(path)

    assert audit_docx(path) == []


def test_docx_audit_assigns_recovery_classes_at_the_finding_producer(tmp_path):
    from docx.oxml import OxmlElement

    path = tmp_path / "governed.docx"
    document = Document()
    document.add_paragraph("Unresolved {study_title}")
    document.settings.element.append(OxmlElement("w:trackRevisions"))
    document.save(path)

    findings = audit_docx(path)

    drafting = next(item for item in findings if "Unresolved template token" in item["issue"])
    structure = next(item for item in findings if "Tracked changes" in item["issue"])
    assert (drafting["recovery_class"], drafting["action"]) == ("document_structure_defect", "preserve_and_stop")
    assert (structure["recovery_class"], structure["action"]) == ("document_structure_defect", "preserve_and_stop")


def test_generic_content_pass_without_per_section_evidence_is_rejected(tmp_path):
    request_dir = tmp_path / "hermes/verification-requests"; response_dir = tmp_path / "hermes/verification-responses"
    request_dir.mkdir(parents=True); response_dir.mkdir(parents=True)
    request = {"schema_version": "hermes-verification/v1", "request_id": "r.verify.content", "task": "clinical_content_verification", "response_path": "hermes/verification-responses/r.verify.content.json", "artifacts": [], "sections": [{"artifact": "protocol", "section_id": "introduction"}], "checks": list(CONTENT_CHECKS), "cross_document_checks": ["study_title"]}
    request["request_sha256"] = verification_request_sha256(request)
    response = {"schema_version": RESPONSE_SCHEMA, "request_id": request["request_id"], "request_sha256": request["request_sha256"], "task": request["task"], "producer": {"model_id": "test", "reviewer_id": "content-reviewer"}, "status": "passed", "findings": []}
    (request_dir / "r.verify.content.json").write_text(json.dumps(request), encoding="utf-8")
    (response_dir / "r.verify.content.json").write_text(json.dumps(response), encoding="utf-8")
    findings, _ = validate_verifications(tmp_path)
    assert sum("explicitly assessed" in item["issue"] for item in findings) == 2


def test_visual_response_is_rejected_after_any_bound_artifact_changes(tmp_path):
    request_dir = tmp_path / "hermes/verification-requests"; request_dir.mkdir(parents=True)
    for relative, data in (("candidate/protocol.docx", b"docx"), ("rendered/protocol.pdf", b"pdf"), ("rendered/protocol/page-1.png", b"png")):
        path = tmp_path / relative; path.parent.mkdir(parents=True, exist_ok=True); path.write_bytes(data)
    digest = lambda relative: hashlib.sha256((tmp_path / relative).read_bytes()).hexdigest()
    request = {"schema_version": "hermes-verification/v1", "request_id": "r.verify.visual", "task": "rendered_page_visual_verification", "response_path": "hermes/verification-responses/r.verify.visual.json", "artifacts": [{"artifact": "protocol", "docx": "candidate/protocol.docx", "docx_sha256": digest("candidate/protocol.docx"), "pdf": "rendered/protocol.pdf", "pdf_sha256": digest("rendered/protocol.pdf"), "pages": [{"page": 1, "path": "rendered/protocol/page-1.png", "sha256": digest("rendered/protocol/page-1.png")}]}], "checks": list(VISUAL_CHECKS)}
    request["request_sha256"] = verification_request_sha256(request)
    request_path = request_dir / "r.verify.visual.json"; request_path.write_text(json.dumps(request), encoding="utf-8")
    response_path = tmp_path / request["response_path"]; response_path.parent.mkdir(parents=True); response_path.write_text(json.dumps(acceptance_verification(request)), encoding="utf-8")
    for relative in (
        "candidate/protocol.docx",
        "rendered/protocol.pdf",
        "rendered/protocol/page-1.png",
    ):
        path = tmp_path / relative
        original = path.read_bytes()
        path.write_bytes(b"changed")
        findings, _ = validate_verifications(tmp_path)
        stale = next(item for item in findings if "stale" in item["issue"] and relative in item["issue"])
        assert stale["recovery_class"] == "document_structure_defect"
        assert stale["action"] == "preserve_and_stop"
        path.write_bytes(original)


def test_visual_review_contract_rejects_artificial_pagination_defects():
    assert "excessive_whitespace" in VISUAL_CHECKS
    assert "artificial_pagination" in VISUAL_CHECKS


def test_generation_rejects_study_input_mutation_after_approval(tmp_path):
    run_dir = tmp_path / "run"; reference_path = run_dir / "reference/study.reference.json"; reference_path.parent.mkdir(parents=True)
    reference_path.write_text(json.dumps(fixture()), encoding="utf-8")
    assert prepare(run_dir)["status"] == "awaiting_approval"
    assert approve(run_dir, approved_by="reviewer")["status"] == "passed"
    changed = json.loads(reference_path.read_text(encoding="utf-8")); changed["study"]["title"] = "Changed after approval"
    reference_path.write_text(json.dumps(changed), encoding="utf-8")
    assert validate(run_dir)["status"] == "blocked"
    result = generate(run_dir)
    assert result["status"] == "blocked"
    assert result["stage"] == "approval_gate"


def test_generation_keeps_persisted_recovery_exhaustion_blocked_across_restart(tmp_path, monkeypatch):
    _require_renderer()
    run_dir = tmp_path / "run"; reference_path = run_dir / "reference/study.reference.json"; reference_path.parent.mkdir(parents=True)
    reference_path.write_text(json.dumps(fixture()), encoding="utf-8")
    assert prepare(run_dir)["status"] == "awaiting_approval"
    assert approve(run_dir, approved_by="reviewer")["status"] == "passed"
    reference = json.loads(reference_path.read_text(encoding="utf-8"))
    reference["generation"] = {
        "attempts": {"introduction": 4},
        "recovery_exhaustion": [{
            "category": "drafting",
            "field": "introduction",
            "strategy_id": "drafting_defect:retry_drafting_target:introduction:introduction",
            "issue": "The same recovery strategy made no progress after 3 attempts.",
        }],
    }
    reference_path.write_text(json.dumps(reference), encoding="utf-8")
    first = generate(run_dir); second = generate(run_dir)
    assert first["status"] == second["status"] == "blocked"
    assert first["stage"] == second["stage"] == "retry_limit"


def test_public_generation_recovers_a_visual_finding_without_changing_approval(tmp_path, monkeypatch):
    _require_renderer()
    reference = fixture()
    run_dir = tmp_path / "run"
    reference_path = run_dir / "reference/study.reference.json"
    reference_path.parent.mkdir(parents=True)
    reference_path.write_text(json.dumps(reference), encoding="utf-8")
    assert prepare(run_dir)["status"] == "awaiting_approval"
    approval = approve(run_dir, approved_by="reviewer")
    revision_id = approval["revision_id"]
    revision_dir = run_dir / "revisions" / revision_id
    approved_sha256 = json.loads(reference_path.read_text(encoding="utf-8"))["approval"]["approved_reference_sha256"]

    def deterministic_assurance(_repo_root, current_revision, *_args, **_kwargs):
        rendered = current_revision / "rendered"
        artifacts = []
        for docx in sorted((current_revision / "candidate").glob("*.docx")):
            name = docx.stem
            page_dir = rendered / name
            page_dir.mkdir(parents=True, exist_ok=True)
            pdf = rendered / f"{name}.pdf"
            writer = PdfWriter()
            writer.add_blank_page(width=612, height=792)
            with pdf.open("wb") as stream:
                writer.write(stream)
            page = page_dir / "page-1.png"
            page.write_bytes(f"deterministic {name} page image".encode())
            artifacts.append({
                "artifact": name,
                "docx": f"candidate/{name}.docx",
                "docx_sha256": hashlib.sha256(docx.read_bytes()).hexdigest(),
                "pdf": f"rendered/{name}.pdf",
                "pdf_sha256": hashlib.sha256(pdf.read_bytes()).hexdigest(),
                "page_count": 1,
                "pages": [{
                    "page": 1,
                    "path": f"rendered/{name}/page-1.png",
                    "sha256": hashlib.sha256(page.read_bytes()).hexdigest(),
                }],
                "renderer": {"kind": "test-office"},
                "page_renderer": {"kind": "test-page-renderer"},
                "status": "passed",
            })
        render = {
            "status": "passed",
            "renderer": {"kind": "test-office"},
            "page_renderer": {"kind": "test-page-renderer"},
            "artifacts": artifacts,
            "findings": [],
            "renderer_attempts": [],
            "page_renderer_attempts": [],
        }
        return {
            "schema_version": "render-assurance/v1",
            "status": "passed",
            "fonts": {},
            "font_substitutions": {},
            "candidate": {"files": []},
            "render": render,
            "findings": [],
        }

    monkeypatch.setattr(workflow, "render_assurance", deterministic_assurance)
    real_render_documents = workflow.render_documents
    forced_candidate = {"enabled": False, "bytes": b""}

    def controlled_render(*args, **kwargs):
        report = real_render_documents(*args, **kwargs)
        if forced_candidate["enabled"]:
            (revision_dir / "candidate/protocol.docx").write_bytes(forced_candidate["bytes"])
        return report

    monkeypatch.setattr(workflow, "render_documents", controlled_render)

    def accept_all(current_run, current_result):
        current_revision = current_run / "revisions" / current_result["revision_id"]
        for relative in current_result["requests"]:
            request = json.loads((current_revision / relative).read_text(encoding="utf-8"))
            response = (
                recorded_acceptance_response(request)
                if request["task"] in {"section_drafting", "prs_narrative_drafting"}
                else acceptance_verification(request)
            )
            response_path = current_revision / request["response_path"]
            response_path.parent.mkdir(parents=True, exist_ok=True)
            response_path.write_text(json.dumps(response), encoding="utf-8")

    clean_run = tmp_path / "clean-run"
    clean_reference = clean_run / "reference/study.reference.json"
    clean_reference.parent.mkdir(parents=True)
    clean_reference.write_text(json.dumps(fixture()), encoding="utf-8")
    assert prepare(clean_run)["status"] == "awaiting_approval"
    approve(clean_run, approved_by="reviewer")
    clean_result = generate(clean_run, require_promoted_runtime=False)
    while clean_result.get("stage") in {"drafting", "drafting_retry"}:
        accept_all(clean_run, clean_result)
        clean_result = generate(clean_run, require_promoted_runtime=False)
    assert clean_result["stage"] == "independent_verification"
    accept_all(clean_run, clean_result)
    clean_result = generate(clean_run, require_promoted_runtime=False)
    assert clean_result["status"] == "passed"
    assert set(clean_result["client_outputs"]) == {
        "output/icf.docx", "output/protocol.docx", "output/study.xml",
    }

    result = generate(run_dir, require_promoted_runtime=False)
    injected_drafting_failure = False
    saw_drafting_retry = False
    while result.get("stage") in {"drafting", "drafting_retry"}:
        saw_drafting_retry = saw_drafting_retry or result.get("stage") == "drafting_retry"
        for relative in result["requests"]:
            request = json.loads((revision_dir / relative).read_text(encoding="utf-8"))
            response_path = revision_dir / request["response_path"]
            response_path.parent.mkdir(parents=True, exist_ok=True)
            response = recorded_acceptance_response(request)
            if not injected_drafting_failure and response.get("section_results"):
                paragraph = next(
                    (item
                    for section in response["section_results"]
                    for item in section.get("paragraphs", [])),
                    None,
                )
                if paragraph is not None:
                    paragraph["evidence_refs"] = []
                    paragraph["boilerplate_refs"] = []
                    injected_drafting_failure = True
            response_path.write_text(
                json.dumps(response), encoding="utf-8"
            )
        result = generate(run_dir, require_promoted_runtime=False)

    assert result["stage"] == "independent_verification"
    assert saw_drafting_retry is True
    first_candidate = hashlib.sha256(
        (revision_dir / "candidate/protocol.docx").read_bytes()
    ).hexdigest()

    content_repair_run = tmp_path / "successful-content-repair"
    shutil.copytree(run_dir, content_repair_run)
    content_revision = content_repair_run / "revisions" / revision_id
    for relative in result["requests"]:
        request = json.loads((content_revision / relative).read_text(encoding="utf-8"))
        response = acceptance_verification(request)
        if request["task"] == "clinical_content_verification":
            failed_section = next(
                item for item in response["section_assessments"]
                if item["section_id"] == "introduction"
            )
            failed_section["status"] = "failed"
            response["status"] = "blocked"
            response["findings"] = [{
                "target_ids": [failed_section["section_id"]],
                "issue": "The section requires a source-grounded content repair.",
            }]
        response_path = content_revision / request["response_path"]
        response_path.parent.mkdir(parents=True, exist_ok=True)
        response_path.write_text(json.dumps(response), encoding="utf-8")
    content_result = generate(content_repair_run, require_promoted_runtime=False)
    assert content_result["stage"] == "drafting_retry"
    while content_result.get("stage") in {"drafting", "drafting_retry"}:
        for relative in content_result["requests"]:
            request = json.loads((content_revision / relative).read_text(encoding="utf-8"))
            response = recorded_acceptance_response(request)
            for section in response.get("section_results", []):
                if section.get("section_id") == "introduction" and section.get("paragraphs"):
                    section["paragraphs"][0]["text"] += (
                        " A prospective evaluation of a wearable monitoring device is the approved study background."
                    )
            response_path = content_revision / request["response_path"]
            response_path.parent.mkdir(parents=True, exist_ok=True)
            response_path.write_text(json.dumps(response), encoding="utf-8")
        content_result = generate(content_repair_run, require_promoted_runtime=False)
    assert content_result["stage"] == "independent_verification"
    assert hashlib.sha256(
        (content_revision / "candidate/protocol.docx").read_bytes()
    ).hexdigest() != first_candidate
    accept_all(content_repair_run, content_result)
    content_result = generate(content_repair_run, require_promoted_runtime=False)
    assert content_result["status"] == "passed"
    assert set(content_result["client_outputs"]) == {
        "output/icf.docx", "output/protocol.docx", "output/study.xml",
    }

    for relative in result["requests"]:
        request_path = revision_dir / relative
        request = json.loads(request_path.read_text(encoding="utf-8"))
        response = acceptance_verification(request)
        if request["task"] == "rendered_page_visual_verification":
            response["status"] = "failed"
            response["page_assessments"][0]["status"] = "failed"
            response["findings"] = [{
                "artifact": "protocol",
                "page": response["page_assessments"][0]["page"],
                "check": "bad_table_split",
                "element": "3. GENERAL INFORMATION",
                "target_ids": ["layout:protocol"],
                "issue": "The Section 3 summary table splits across a page boundary.",
            }]
        response_path = revision_dir / request["response_path"]
        response_path.parent.mkdir(parents=True, exist_ok=True)
        response_path.write_text(json.dumps(response), encoding="utf-8")

    recovered = generate(run_dir, require_promoted_runtime=False)

    assert recovered["status"] == "awaiting_hermes"
    assert recovered["stage"] == "independent_verification"
    state = json.loads(reference_path.read_text(encoding="utf-8"))
    assert state["approval"]["revision_id"] == revision_id
    assert state["approval"]["approved_reference_sha256"] == approved_sha256
    assert state["generation"]["review_set"] == 2
    assert state["generation"]["gate_attempts"]
    attempts = [
        json.loads(path.read_text(encoding="utf-8"))
        for path in sorted((revision_dir / "attempts").glob("*/attempt-manifest.json"))
    ]
    assert {item["stage"] for item in attempts} == {"drafting", "quality"}
    assert all(
        action["outcome_status"] == "measured"
        for attempt in attempts
        for action in attempt["recovery_actions"]
    )
    assert hashlib.sha256(
        (revision_dir / "candidate/protocol.docx").read_bytes()
    ).hexdigest() != first_candidate

    successful_repair_run = tmp_path / "successful-progressive-repair"
    shutil.copytree(run_dir, successful_repair_run)
    successful_revision = successful_repair_run / "revisions" / revision_id
    for relative in recovered["requests"]:
        request = json.loads((successful_revision / relative).read_text(encoding="utf-8"))
        response_path = successful_revision / request["response_path"]
        response_path.parent.mkdir(parents=True, exist_ok=True)
        response_path.write_text(json.dumps(acceptance_verification(request)), encoding="utf-8")
    successful = generate(successful_repair_run, require_promoted_runtime=False)
    assert successful["status"] == "passed"
    assert set(successful["client_outputs"]) == {
        "output/icf.docx", "output/protocol.docx", "output/study.xml",
    }
    successful_quality = json.loads(
        (successful_revision / "delivery-manifest.json").read_text(encoding="utf-8")
    )["quality"]
    assert successful_quality["final_exact_artifact_review"]["status"] == "passed"
    assert successful_quality["final_exact_artifact_review"]["every_page"] is True

    forced_candidate["bytes"] = (revision_dir / "candidate/protocol.docx").read_bytes()
    forced_candidate["enabled"] = True

    for relative in recovered["requests"]:
        request = json.loads((revision_dir / relative).read_text(encoding="utf-8"))
        response = acceptance_verification(request)
        if request["task"] == "rendered_page_visual_verification" and any(
            item.get("artifact") == "protocol" for item in request.get("artifacts", [])
        ):
            response["status"] = "failed"
            response["page_assessments"][0]["status"] = "failed"
            response["findings"] = [{
                "artifact": "protocol",
                "page": response["page_assessments"][0]["page"],
                "check": "bad_table_split",
                "element": "3. GENERAL INFORMATION",
                "target_ids": ["layout:protocol"],
                "issue": "The same Section 3 split remains after repair.",
            }]
        response_path = revision_dir / request["response_path"]
        response_path.parent.mkdir(parents=True, exist_ok=True)
        response_path.write_text(json.dumps(response), encoding="utf-8")

    no_progress = generate(run_dir, require_promoted_runtime=False)
    repeated = generate(run_dir, require_promoted_runtime=False)

    assert no_progress["status"] == "blocked"
    assert no_progress["stage"] == "recovery_no_progress"
    assert repeated["status"] == "blocked"
    assert repeated["stage"] == "retry_limit"


def test_public_generation_adopts_an_interrupted_attempt_and_preserves_approval(tmp_path):
    reference = json.loads(
        (ROOT / "tests/fixtures/retrospective-acceptance-source.json").read_text(encoding="utf-8")
    )
    run_dir = tmp_path / "run"
    reference_path = run_dir / "reference/study.reference.json"
    reference_path.parent.mkdir(parents=True)
    reference_path.write_text(json.dumps(reference), encoding="utf-8")
    assert prepare(run_dir)["status"] == "awaiting_approval"
    approval = approve(run_dir, approved_by="reviewer")
    revision_id = approval["revision_id"]
    revision_dir = run_dir / "revisions" / revision_id
    approved_sha256 = json.loads(reference_path.read_text(encoding="utf-8"))["approval"]["approved_reference_sha256"]
    workflow._archive_failed_attempt(
        revision_dir,
        "quality",
        [{
            "category": "visual",
            "artifact": "protocol",
            "check": "orphan_heading",
            "element": "5. INTRODUCTION",
            "target_ids": ["layout:protocol"],
            "recovery_class": "visual_defect",
            "action": "targeted_layout_repair",
            "issue": "interrupted before recovery state commit",
        }],
    )

    resumed = generate(run_dir, require_promoted_runtime=False)

    assert resumed["status"] == "awaiting_hermes"
    state = json.loads(reference_path.read_text(encoding="utf-8"))
    assert state["approval"]["revision_id"] == revision_id
    assert state["approval"]["approved_reference_sha256"] == approved_sha256
    assert state["generation"]["gate_attempts"]
    manifest = json.loads(next((revision_dir / "attempts").glob("*/attempt-manifest.json")).read_text(encoding="utf-8"))
    assert manifest["recovery_actions"][0]["outcome_status"] == "interrupted_no_action"


def test_public_generation_replays_a_committed_recovery_plan_after_interruption(tmp_path, monkeypatch):
    reference = json.loads(
        (ROOT / "tests/fixtures/retrospective-acceptance-source.json").read_text(encoding="utf-8")
    )
    run_dir = tmp_path / "run"
    reference_path = run_dir / "reference/study.reference.json"
    reference_path.parent.mkdir(parents=True)
    reference_path.write_text(json.dumps(reference), encoding="utf-8")
    assert prepare(run_dir)["status"] == "awaiting_approval"
    approval = approve(run_dir, approved_by="reviewer")
    revision_dir = run_dir / "revisions" / approval["revision_id"]
    approved = json.loads((revision_dir / "approved-reference.json").read_text(encoding="utf-8"))
    working = json.loads(reference_path.read_text(encoding="utf-8"))
    real_apply = workflow._apply_pending_recovery_plan
    monkeypatch.setattr(
        workflow,
        "_apply_pending_recovery_plan",
        lambda *_args, **_kwargs: (_ for _ in ()).throw(RuntimeError("interrupted after plan commit")),
    )
    with pytest.raises(RuntimeError, match="interrupted after plan commit"):
        workflow._quality_retry(
            run_dir,
            reference_path,
            working,
            approved,
            revision_dir,
            {},
            [{
                "category": "visual",
                "artifact": "protocol",
                "check": "orphan_heading",
                "element": "5. INTRODUCTION",
                "target_ids": ["layout:protocol"],
                "recovery_class": "visual_defect",
                "action": "targeted_layout_repair",
                "issue": "orphan heading",
            }],
            "quality",
        )
    monkeypatch.setattr(workflow, "_apply_pending_recovery_plan", real_apply)

    resumed = generate(run_dir, require_promoted_runtime=False)

    assert resumed["status"] == "awaiting_hermes"
    assert resumed["stage"] == "drafting_retry"
    state = json.loads(reference_path.read_text(encoding="utf-8"))
    assert state["approval"]["revision_id"] == approval["revision_id"]
    assert state["generation"]["pending_recovery_plan"]["applied"] is True
    assert state["generation"]["pending_recovery_attempts"]
