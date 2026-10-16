"""Visible native TOCs must not depend on a later office-rendering stage."""
import json
from pathlib import Path

import pytest
from docx import Document
from docx.oxml.ns import qn

import rendering

ROOT = Path(__file__).resolve().parents[1]


@pytest.mark.parametrize("branch", ["prospective", "ambispective", "retrospective"])
def test_first_native_assembly_has_visible_linked_toc_without_guessed_pages(tmp_path, branch):
    reference = json.loads((ROOT / f"tests/fixtures/{branch}-acceptance-source.json").read_text())
    rendering.render_documents(ROOT, tmp_path, reference, {"protocol": [], "icf": {}, "prs": {}},
                               artifact_names={"protocol"}, audit_output=False)
    document = Document(tmp_path / "candidate/protocol.docx")
    rows = [p for p in document.paragraphs if p.style.name.casefold() in {"toc 1", "toc 2"}]
    headings = [p.text.strip() for p in document.paragraphs if p.style.name in {"Heading 1", "Heading 2"} and p.text.strip()]
    assert rows, "First DOCX assembly leaves the table-of-contents page empty"
    assert [p.text.rstrip("\t") for p in rows] == headings
    bookmarks = set(document.element.body.xpath('.//w:bookmarkStart/@w:name'))
    assert all(link.get(qn("w:anchor")) in bookmarks for row in rows for link in row._p.xpath('./w:hyperlink'))
    assert all(row._p.xpath('./w:hyperlink') for row in rows)
    instructions = document.element.body.xpath('.//w:instrText/text()')
    assert sum(' TOC ' in text for text in instructions) == 1
    assert sum('PAGEREF ' in text for text in instructions) == len(rows)
    # Word can calculate real page references on update; native assembly has
    # no renderer evidence from which it could truthfully cache page numbers.
    assert all(not row._p.xpath('./w:r/w:t/text()') for row in rows)
