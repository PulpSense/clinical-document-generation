"""Accepted native concerns must not be hidden by missing independent visual review."""
import json
import pytest
import quality
from test_content_review_recovery import operation, FAMILIES

@pytest.mark.parametrize('family', FAMILIES)
@pytest.mark.parametrize('visual_failure', ['initial-missing', 'persistent-missing', 'authentication', 'cleanup', 'complete-concern',
    'aliased-foreign-page', 'aliased-foreign-artifact', 'aliased-source-clarification'])
def test_content_defect_reaches_repair_despite_missing_visual_response(tmp_path, monkeypatch, family, visual_failure):
    from docx import Document
    from test_clean_delivery import controlled_render
    from test_general_document_repair import repair_response
    corrupted = []
    reviewed_sets = []
    repairs = []
    omitted_reassessment = []
    def render(root, revision, *args, **kwargs):
        if not corrupted:
            path = revision / 'candidate/protocol.docx'
            doc = Document(path)
            doc.paragraphs[0].text = 'UNSUPPORTED INVENTED COVER TITLE'
            doc.save(path)
            corrupted.append(True)
        return controlled_render(root, revision, *args, **kwargs)
    def worker(request, response, revision, *args):
        reviewed_sets.append(request['review_set'])
        doc = Document(revision / 'candidate/protocol.docx')
        if any(p.text == 'UNSUPPORTED INVENTED COVER TITLE' for p in doc.paragraphs):
            row = next((r for r in response['section_assessments'] if r['section_id'] == 'title-page'), None)
            if row:
                row['status'] = 'blocked'
                response['status'] = 'blocked'
                response['section_assessments'] = [row]
                response['findings'] = [{'finding_id': 'cover-title', 'artifact': 'protocol',
                    'target_ids': ['title-page'], 'check': 'source_supported',
                    'issue': 'The native cover title is not approved.', 'material': True,
                    'affected_passage': 'UNSUPPORTED INVENTED COVER TITLE', 'source_path': 'study.title',
                    'source_excerpt': request['approved_source']['study']['title']}]
        skip_notes = (family[0] == 'Retrospective' and request['review_set'] > 1
                      and any(r['section_id'] == 'title-page' for r in request['sections'])
                      and not omitted_reassessment)
        if skip_notes:
            omitted_reassessment.append(True)
        for row in response['section_assessments']:
            row['notes'] = [f"Independently reassessed {f['finding_id']} against current native bytes."
                            for f in request.get('review_context', {}).get('prior_findings', [])] if not skip_notes else []
        return response
    def repair(request, revision):
        repairs.append(request)
        if visual_failure == 'complete-concern':
            assert {'source_supported', 'template_preservation'} <= {
                f['verification_source_finding']['check'] for f in request['findings']}
        path = revision / request['workspace_path'] / 'protocol.docx'
        doc = Document(path)
        doc.paragraphs[0].text = request['approved_source']['study']['title']
        doc.save(path)
        return repair_response(request)
    def visual(request, response, revision, fallback):
        if visual_failure == 'authentication':
            import workflow
            raise workflow.WorkerAuthenticationError('Controlled visual authentication rejection')
        if visual_failure == 'cleanup':
            from owned_observer_runtime import OwnedObserverLifetimeUnresolved
            raise OwnedObserverLifetimeUnresolved('Controlled unresolved descendant')
        if visual_failure.startswith('aliased-'):
            if not request['request_id'].endswith('.protocol'):
                return response
            response['status'] = response['page_assessments'][0]['status'] = 'blocked'
            finding = {'artifact': 'protocol', 'page': 1, 'check': 'template_preservation',
                'element': 'Controlled cover', 'issue': 'Controlled stop concern',
                'code': 'review_coverage_incomplete'}
            if visual_failure == 'aliased-foreign-page':
                finding['page'] = 999
            elif visual_failure == 'aliased-foreign-artifact':
                finding.update(artifact='icf', code='verification_response_missing')
            else:
                finding['source_clarification_required'] = True
            response['findings'] = [finding]
            return response
        if visual_failure == 'complete-concern' and request['review_set'] == 1:
            if not request['request_id'].endswith('.protocol'):
                return response
            doc = Document(revision / 'candidate/protocol.docx')
            if not any(p.text == 'UNSUPPORTED INVENTED COVER TITLE' for p in doc.paragraphs):
                return response
            response['status'] = 'blocked'
            response['page_assessments'][0]['status'] = 'blocked'
            response['findings'] = [{'artifact': 'protocol', 'page': 1,
                'check': 'template_preservation', 'element': 'Cover title',
                'issue': 'The rendered cover contains the unsupported replacement title.'}]
            return response
        if visual_failure == 'complete-concern' and request['review_set'] > 1:
            for row in response['page_assessments']:
                row['notes'] = [f"Independently reassessed {f['finding_id']} against the repaired rendered cover."
                                for f in request.get('review_context', {}).get('prior_findings', [])]
        if visual_failure == 'persistent-missing' or request['review_set'] == 1:
            return None
        return response
    if visual_failure == 'complete-concern':
        import workflow
        select = workflow._accepted_partial_repair_findings
        def select_with_transport_check(revision, **kwargs):
            for path in (revision / 'hermes/verification-requests').glob('*.verify.content.json'):
                request = json.loads(path.read_text())
                response_path = revision / request['response_path']
                if response_path.exists() and json.loads(response_path.read_text()).get('scheduler_partial'):
                    assert quality.verification_response_is_terminal(revision, path)
                    assert not quality.verification_response_is_complete(revision, path)
            return select(revision, **kwargs)
        monkeypatch.setattr(workflow, '_accepted_partial_repair_findings', select_with_transport_check)
    result, revision, _ = operation(tmp_path, monkeypatch, family, worker, worker, repair=repair, render=render, visual_review=visual)
    if visual_failure.startswith('aliased-'):
        assert result['status'] != 'passed', result
        assert not result.get('client_outputs'), result
        assert not repairs, 'Rejected binding or source stop reached native repair'
        assert result.get('findings'), 'Stop concern was erased'
        if visual_failure == 'aliased-source-clarification':
            assert result['stage'] == 'source_review', result
            assert any(f.get('source_clarification_required') or
                       f.get('verification_source_finding', {}).get('source_clarification_required')
                       for f in result['findings'])
        return
    if visual_failure in {'authentication', 'cleanup'}:
        assert result['status'] != 'passed', result
        assert len(repairs) == 0
        assert result['stage'] in {'parent_visual_integrity', 'worker_authentication'}, result
        return
    if visual_failure == 'persistent-missing':
        assert result['status'] == 'review_required', result
        assert result['stage'] == 'parent_visual_review', result
        assert not result['client_outputs']
        assert len(repairs) == 1
        assert set(reviewed_sets) == {1, 2}
        assert all(p.text != 'UNSUPPORTED INVENTED COVER TITLE'
                   for p in Document(revision / 'candidate/protocol.docx').paragraphs)
        return
    assert result['status'] == 'passed', result
    assert len(repairs) == 1
    assert set(reviewed_sets) == {1, 2}
    first = json.loads(next((revision / 'hermes/review-schedules').glob('*review-1*')).read_text())
    assert first['waves'][-1]['completed_assessments'] < len(quality.content_review_sections(repairs[0]['approved_source']))
    assert all(p.text != 'UNSUPPORTED INVENTED COVER TITLE' for p in Document(tmp_path / 'output/protocol.docx').paragraphs)
