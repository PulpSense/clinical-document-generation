"""Development replays of supplied COMFORT facts; no model-quality certification."""
import copy
import json
from pathlib import Path
from xml.etree import ElementTree as ET

import pytest
from docx import Document

import contracts
import prs_xml
import rendering

ROOT = Path(__file__).resolve().parents[1]
TRIALS = ROOT / 'docs/experiments/prompt-trials'
TEMPLATE = ROOT / 'assets/client-templates/prs/clinicaltrials_prs_full_placeholder_template.xml'


def comfort_source():
    return copy.deepcopy(json.loads((TRIALS / 'comfort.request.json').read_text())['source'])


@pytest.mark.parametrize('article', ['an', 'another', 'a'])
def test_xml_preserves_scoped_exclusion_without_adding_any_study_ban(article, tmp_path):
    source = comfort_source()
    rule = f'Participation in {article} interventional research study in the 60 days before screening, or planned interventional research participation during follow-up.'
    source['population']['inclusion_criteria'] = ['Able to consent personally.']
    source['population']['exclusion_criteria'] = [rule]
    source['procedures']['minimum_days_before_screening_without_participation'] = '60 days'
    original = copy.deepcopy(source)
    output = tmp_path / 'study.xml'
    prs_xml.generate(TEMPLATE, output, source, {'brief_summary': 'Summary.', 'detailed_description': 'Description.'}, check_output=False)
    text = next(ET.parse(output).getroot().iter('clinical_study')).findtext('eligibility/criteria/textblock')
    assert text.count('60 days') == 1
    assert rule in text
    assert 'without participation in another study' not in text
    assert source == original


def test_historical_record_column_does_not_shift_participant_visit_numbers(tmp_path):
    source = comfort_source()
    historical = {'visit': 'Historical record abstraction — staff task, not a participant visit',
                  'timing': 'Records dated 1 January 2025 through 30 September 2026; reviewed after consent',
                  'procedures': ['Abstract existing surgery dates']}
    source['procedures']['visit_schedule_table'].insert(0, historical)
    original = copy.deepcopy(source)
    rendering.render_documents(ROOT, tmp_path, source, {'protocol': [], 'icf': {}, 'prs': {}},
                               artifact_names={'protocol'}, audit_output=False)
    doc = Document(tmp_path / 'candidate/protocol.docx')
    schedule = next(t for t in doc.tables if t.rows[0].cells[0].text.strip() == 'Visit Number')
    assert [row.cells[0].text for row in schedule.rows[1:]] == ['', '1', '2', '3']
    matrix = next(t for t in doc.tables if t.rows[0].cells[0].text.strip() == 'Activity')
    assert [c.text.strip() for c in matrix.rows[1].cells] == ['Activity', 'Record review', 'Visit 1', 'Visit 2', 'Visit 3']
    assert [c.text.strip() for c in matrix.rows[2].cells] == ['Abstract existing surgery dates', 'X', '', '', '']
    assert source == original


def test_sterling_multiline_prose_uses_plain_body_type_and_removes_cosmetic_spacers(tmp_path):
    source = comfort_source()
    risk = 'Vision testing may cause fatigue.\n\nSaliva collection may cause discomfort.\n\nYou may stop testing.'
    model = {'protocol': [], 'prs': {}, 'icf': {
        'icf.background': {'paragraphs': [{'text': 'Existing records and future assessments describe comfort.'}], 'lists': []},
        'icf.risks': {'paragraphs': [{'text': risk}], 'lists': []}}}
    rendering.render_documents(ROOT, tmp_path, source, model, artifact_names={'icf'}, audit_output=False)
    doc = Document(tmp_path / 'candidate/icf.docx')
    target = [p for p in doc.paragraphs if p.text.strip() in risk.split('\n\n')]
    assert len(target) == 3
    for p in target:
        assert all(not run.bold and not run.underline and run.font.name == 'Arial' for run in p.runs)
    background = next(i for i,p in enumerate(doc.paragraphs) if p.text.strip() == 'BACKGROUND')
    purpose = next(i for i,p in enumerate(doc.paragraphs) if p.text.strip() == 'PURPOSE')
    assert [p.text for p in doc.paragraphs[background+1:purpose]] == ['Existing records and future assessments describe comfort.']


def actual_comfort():
    return json.loads((ROOT / 'tests/fixtures/comfort-hermes-v25.json').read_text())


def test_saved_hermes_reference_keeps_site_investigators_in_xml(tmp_path):
    source = actual_comfort()['source']
    original = copy.deepcopy(source)
    output = tmp_path / 'study.xml'
    prs_xml.generate(TEMPLATE, output, source, {'brief_summary': 'Summary.', 'detailed_description': 'Description.'}, check_output=False)
    xml = ET.parse(output).getroot()
    text = ' '.join(xml.itertext())
    for name in ['Ellis', 'Bennett', 'Foster']:
        assert name in text
    assert source == original


@pytest.mark.parametrize('flag,expected', [(True, True), ('True', True), ('Yes', True), (False, False), ('False', False), ('No', False)])
def test_explicit_genetic_flags_control_authorized_sterling_clause(flag, expected, tmp_path):
    run = actual_comfort()
    source = run['source']
    source['regulatory']['genetic_information_collected'] = flag
    rendering.render_documents(ROOT, tmp_path, source, run['drafts'], artifact_names={'icf'}, audit_output=False)
    doc = Document(tmp_path / 'candidate/icf.docx')
    heading = 'GENETIC INFORMATION NONDISCRIMINATION ACT'
    assert (heading in [p.text for p in doc.paragraphs]) is expected
    if expected:
        clause = next(c for c in contracts.sterling_clause_contract()['clauses'] if c['clause_id'] == 'sterling.privacy.gina')
        for text in clause['authorized_texts']:
            assert [p.text for p in doc.paragraphs].count(text) == 1


def test_qualifier_rows_become_bound_table_notes_without_losing_facts(tmp_path):
    run = actual_comfort()
    section = next(s for s in run['drafts']['protocol'] if s['section_id'] == 'evaluation-procedures')
    section['assessment_presentation'] = {'qualifiers': [
        {'source_text': 'Same wording at Day 0, Day 30 and Day 90',
         'display_text': 'Ocular comfort rating: use the same wording at Day 0, Day 30 and Day 90.',
         'rationale': 'A measurement instruction, not a separate participant assessment.'},
        {'source_text': 'Not presented as a validated instrument',
         'display_text': 'The study-specific ocular comfort rating is not presented as a validated instrument.',
         'rationale': 'Instrument qualification retained beneath the matrix.'}]}
    original = copy.deepcopy(run['source'])
    rendering.render_documents(ROOT, tmp_path, run['source'], run['drafts'], artifact_names={'protocol'}, audit_output=False)
    doc = Document(tmp_path / 'candidate/protocol.docx')
    matrix = next(t for t in doc.tables if t.rows[0].cells[0].text.strip() == 'Activity')
    labels = [row.cells[0].text for row in matrix.rows]
    for edit in section['assessment_presentation']['qualifiers']:
        assert edit['source_text'] not in labels
        assert edit['display_text'] in [p.text for p in doc.paragraphs]
    assert any('Ocular comfort rating' in label or 'ocular comfort rating' in label for label in labels)
    assert run['source'] == original


def test_historical_staff_task_is_excluded_from_participant_count():
    source = actual_comfort()['source']
    assert contracts.computed_study_fields(source)['visit_count'] == 3


@pytest.mark.parametrize('edit', [
    {'source_text': 'Not presented as a validated instrument', 'display_text': '', 'rationale': 'Qualification.'},
    {'source_text': 'A label absent from the approved schedule', 'display_text': 'Keep it.', 'rationale': 'Qualification.'},
])
def test_invalid_qualifier_edit_cannot_remove_an_unbound_or_unretained_fact(edit, tmp_path):
    run = actual_comfort()
    section = next(s for s in run['drafts']['protocol'] if s['section_id'] == 'evaluation-procedures')
    section['assessment_presentation'] = {'qualifiers': [edit]}
    with pytest.raises(ValueError, match='exact existing label/note'):
        rendering.render_documents(ROOT, tmp_path, run['source'], run['drafts'], artifact_names={'protocol'}, audit_output=False)


def test_actual_saved_icf_drafts_get_plain_typography_and_compact_background(tmp_path):
    run = actual_comfort()
    rendering.render_documents(ROOT, tmp_path, run['source'], run['drafts'], artifact_names={'icf'}, audit_output=False)
    doc = Document(tmp_path / 'candidate/icf.docx')
    risk_texts = [p['text'] for p in run['drafts']['icf']['icf.risks']['paragraphs']]
    risk_paragraphs = [p for p in doc.paragraphs if p.text in risk_texts]
    assert len(risk_paragraphs) == 4
    assert all(not r.bold and not r.underline and r.font.name == 'Arial' for p in risk_paragraphs for r in p.runs)
    background = next(i for i,p in enumerate(doc.paragraphs) if p.text == 'BACKGROUND')
    purpose = next(i for i,p in enumerate(doc.paragraphs) if p.text == 'PURPOSE')
    assert len(doc.paragraphs[background+1:purpose]) == 1


def test_blank_primary_placeholder_does_not_promote_subinvestigator_over_named_site_investigator(tmp_path):
    source = actual_comfort()['source']
    source['sites'][1]['investigators'][1] = {'name': 'Jordan Example', 'first_name': 'Jordan', 'last_name': 'Example', 'title': 'Sub-Investigator'}
    output = tmp_path / 'study.xml'
    prs_xml.generate(TEMPLATE, output, source, {'brief_summary': 'Summary.', 'detailed_description': 'Description.'}, check_output=False)
    locations = next(ET.parse(output).getroot().iter('clinical_study')).findall('location')
    assert 'Bennett' in ' '.join(locations[1].find('investigator').itertext())
