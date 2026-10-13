"""Concrete owner campaign seeds at the declared native/recipient boundaries."""
import json
from pathlib import Path
from zipfile import ZipFile

from docx import Document
from lxml import etree as ET
import pytest

import ai_quality_campaign as campaign
from readiness_injection import campaign_adapter, inject_native
import workflow


def plan_case(case):
    return next(c for c in campaign.readiness_plan({'commit': 'controlled-candidate'})['supplemental_operations'] if c['case_id'] == case)


@pytest.mark.parametrize('existing_title', ['NovaStep Recovery Study', 'Unrelated stock template title'])
def test_header_seed_changes_only_predeclared_source_owned_text(tmp_path, existing_title):
    planned = plan_case('Q3')
    seed = planned['native_seeds'][1]
    candidate = tmp_path / 'candidate'
    candidate.mkdir()
    path = candidate / 'protocol.docx'
    document = Document()
    document.sections[0].header.paragraphs[0].text = existing_title
    document.add_paragraph('Unrelated source content')
    document.save(path)
    before = path.read_bytes()
    if existing_title != seed['before']:
        with pytest.raises(ValueError, match='missing or ambiguous'):
            inject_native(candidate, seed, tmp_path / 'ledger')
        assert path.read_bytes() == before
    else:
        inject_native(candidate, seed, tmp_path / 'ledger')
        assert Document(path).sections[0].header.paragraphs[0].text == seed['after']
        assert Document(path).paragraphs[0].text == 'Unrelated source content'
        assert (tmp_path / 'ledger/Q3.2/before.bin').read_bytes() == before


def test_clipping_seed_selects_populated_source_schedule_row(tmp_path):
    seed = plan_case('Q4')['native_seeds'][0]
    candidate = tmp_path / 'candidate'
    candidate.mkdir()
    path = candidate / 'protocol.docx'
    document = Document()
    document.add_table(rows=1, cols=1)  # A blank layout row must be skipped.
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = 'Visit'
    table.cell(1, 0).text = 'Month 3'
    table.cell(1, 1).text = 'Day 90 +/- 7'
    document.save(path)
    inject_native(candidate, seed, tmp_path / 'ledger')
    with ZipFile(path) as archive:
        root = ET.fromstring(archive.read('word/document.xml'))
    rows = root.xpath('//w:tr', namespaces=seed['namespaces'])
    assert not rows[0].xpath('.//w:trHeight', namespaces=seed['namespaces'])
    clipped = root.xpath(seed['xpath'], namespaces=seed['namespaces'])[0]
    height = clipped.xpath('./w:trPr/w:trHeight', namespaces=seed['namespaces'])[0]
    assert height.get('{'+seed['namespaces']['w']+'}val') == '1'
    assert 'Day 90 +/- 7' in ''.join(clipped.xpath('.//w:t/text()', namespaces=seed['namespaces']))


def test_actual_recipient_wrapper_retries_exact_bytes_after_declared_transient_failure(tmp_path):
    from test_desktop_delivery import _manifest
    from readiness_injection import campaign_adapter
    calls = []
    def opener(path):
        calls.append(path)
        return b'1234' if path.endswith('Protocol final.docx') else b'567'
    opener.capability_executable = Path('/controlled/recipient-opener')
    reply = workflow.desktop_attachment_reply(_manifest(), run_dir=tmp_path)
    with campaign_adapter(workflow, plan_case('F2'), tmp_path, opener) as (retrieval, events):
        result = workflow.confirm_desktop_delivery(_manifest(), reply, retrieval, retries=1)
        assert retrieval.capability_executable == opener.capability_executable
    assert result['confirmed']
    assert result['opened'][0]['attempts'] == 2
    assert len(calls) == 2  # The declared transient fails before the real recipient call.
    assert [e['event'] for e in events] == ['transient_actual_recipient_transport_failure']
