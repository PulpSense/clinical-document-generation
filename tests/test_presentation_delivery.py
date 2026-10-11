"""Public lifecycle checks: harmless observations cannot veto verified delivery."""
import json

import pytest

from docx import Document

import quality
import drafting
import rendering
from test_clean_delivery import controlled_render, deliver, with_material_repair
from test_construction_reliability import source_for, FAMILIES, ROOT
from test_handoff_quality import acceptance_verification
from test_editorial_publication import repetition_review


@pytest.mark.parametrize('branch,family', FAMILIES)
def test_harmless_uneditable_observation_delivers_complete_clean_package(tmp_path, monkeypatch, branch, family):
    source = source_for(branch, family)
    title = source['study']['title']

    def review(request):
        response = acceptance_verification(request)
        if request['task'] == 'clinical_content_verification':
            response['status'] = 'blocked'
            response['findings'] = [{
                'finding_id': 'cover-style', 'category': 'content', 'artifact': 'protocol',
                'check': 'substantive', 'surface': 'shell', 'target_ids': ['protocol.shell'],
                'affected_passage': title, 'issue': 'A shorter display of the approved title would read better.',
                'material': False, 'safety_critical': False, 'contradiction': False,
                'obscures_required_information': False, 'materially_unusable': False,
            }]
            for row in response['section_assessments']:
                if row['section_id'] == 'protocol.shell': row['status'] = 'blocked'
        return response

    result, revision = deliver(tmp_path, monkeypatch, source, review=review)
    assert result['status'] == 'passed', result
    assert result['warnings'] == [] and not result['desktop_reply'].get('cautions')
    expected = ['protocol.docx'] if branch == 'Retrospective' else ['icf.docx', 'protocol.docx', 'study.xml']
    assert sorted(path.name for path in (tmp_path / 'output').iterdir()) == expected
    assert title in '\n'.join(p.text for p in Document(tmp_path / 'output/protocol.docx').paragraphs)
    manifest = json.loads((tmp_path / result['manifest']).read_text())
    assert manifest['quality']['internal_diagnostics']
    assert quality.final_exact_artifact_review_findings(
        revision, json.loads((revision / 'approved-reference.json').read_text()),
        json.loads((revision / 'candidate-build.json').read_text())['render_report'], manifest['quality']) == []
    # A harmless note never grants permission to reuse stale page review.
    page = next((revision / 'rendered').glob('*/page-1.png'))
    page.write_bytes(page.read_bytes() + b'changed after final review')
    assert quality.final_exact_artifact_review_findings(
        revision, json.loads((revision / 'approved-reference.json').read_text()),
        json.loads((revision / 'candidate-build.json').read_text())['render_report'], manifest['quality'])


@pytest.mark.parametrize('edit', ['unchanged', 'malformed', 'changed_quantity'])
def test_ineffective_optional_polish_finishes_fresh_reviews_and_delivers(tmp_path, monkeypatch, edit):
    source = source_for('Prospective', 'Sterling')
    source['population']['sample_size'] = '40 (planned)'
    passage = 'The study is expected to include 40 (planned).'
    edits, reviews = [], []

    def review(request):
        response = acceptance_verification(request)
        if request['task'] == 'clinical_content_verification':
            reviews.append(request['request_id'])
            response['status'] = 'blocked'
            response['findings'] = [{
                'finding_id': 'enrollment-wording', 'category': 'content', 'artifact': 'icf',
                'check': 'substantive', 'surface': 'shell', 'target_ids': ['icf.shell'],
                'affected_passage': passage, 'issue': 'The enrollment noun would improve readability.',
                'material': False, 'safety_critical': False, 'contradiction': False,
                'obscures_required_information': False, 'materially_unusable': False,
            }]
            for row in response['section_assessments']:
                if row['section_id'] == 'icf.shell': row['status'] = 'blocked'
        return response

    def draft(request):
        if not request.get('construction_repair'): return drafting.recorded_acceptance_response(request)
        edits.append(request)
        response = drafting.response_template(request, model_id='Controlled optional editor')
        response['section_results'][0]['paragraphs'] = [{
            'text': passage.replace('40', '41') if edit == 'changed_quantity' else passage,
            'evidence_refs': ['source:population.sample_size'],
        }]
        if edit == 'malformed': response['schema_version'] = 'wrong'
        return response

    review, draft = with_material_repair(review, draft)
    result, revision = deliver(tmp_path, monkeypatch, source, review=review, draft=draft)
    assert result['status'] == 'passed', result
    assert len(edits) == 1 and len(reviews) == 2
    assert result['warnings'] == []
    assert passage in '\n'.join(p.text for p in Document(tmp_path / 'output/icf.docx').paragraphs)
    assert list((revision / 'hermes/rejected').glob('*.json'))
    manifest = json.loads((tmp_path / result['manifest']).read_text())
    assert manifest['quality']['internal_diagnostics']
    assert quality.final_exact_artifact_review_findings(
        revision, json.loads((revision / 'approved-reference.json').read_text()),
        json.loads((revision / 'candidate-build.json').read_text())['render_report'], manifest['quality']) == []


@pytest.mark.parametrize('surface', ['shell', 'document_shell', 'Document Shell', 'document-shell', 'body'])
@pytest.mark.parametrize('status', ['passed', 'blocked'])
def test_presentation_review_labels_and_passing_assessments_are_accepted(tmp_path, surface, status):
    path, response_path, response = repetition_review(tmp_path, 'Prospective', 'Sterling')
    response['status'] = status
    response['findings'] = [{
        'finding_id': 'cover-note', 'category': 'content', 'artifact': 'protocol',
        'check': 'substantive', 'surface': surface, 'target_ids': ['protocol.shell'],
        'affected_passage': 'A prospective observational study of Other',
        'issue': 'A generic descriptor would read better.',
        'material': False, 'safety_critical': False, 'contradiction': False,
        'obscures_required_information': False, 'materially_unusable': False,
    }]
    for row in response['section_assessments']:
        row['status'] = status if row['section_id'] == 'protocol.shell' else 'passed'
    response_path.write_text(json.dumps(response))
    findings, _ = quality.validate_verifications(tmp_path, request_paths=[path])
    assert len(findings) == 1, findings
    assert findings[0]['code'] == 'nonmaterial-presentation'
    assert findings[0]['verification_source_finding']['surface'] == surface
    assert quality.verification_response_is_complete(tmp_path, path)


@pytest.mark.parametrize('mutation', ['material', 'safety_critical', 'contradiction',
    'obscures_required_information', 'materially_unusable', 'missing-flag',
    'source-fidelity', 'unknown', 'unknown-status', 'missing-section', 'stale', 'material-companion'])
def test_presentation_policy_cannot_hide_essential_or_incomplete_review(tmp_path, mutation):
    path, response_path, response = repetition_review(tmp_path, 'Prospective', 'Sterling')
    finding = response['findings'][0]
    if mutation in {'material', 'safety_critical', 'contradiction', 'obscures_required_information', 'materially_unusable'}:
        finding[mutation] = True
    elif mutation == 'missing-flag': del finding['material']
    elif mutation == 'source-fidelity': finding['check'] = 'source_fidelity'
    elif mutation == 'unknown': finding['check'] = 'unknown'
    elif mutation == 'unknown-status': response['status'] = 'something-else'
    elif mutation == 'missing-section': response['section_assessments'].pop()
    elif mutation == 'stale': response['request_sha256'] = 'stale'
    else:
        response['findings'].append({**finding, 'finding_id': 'material-companion', 'material': True})
    response_path.write_text(json.dumps(response))
    assert not quality.verification_response_is_complete(tmp_path, path)


@pytest.mark.parametrize('field,value', [('status', []), ('status', {}),
    ('category', []), ('category', {}), ('artifact', []), ('artifact', {}),
    ('target_ids', None), ('target_ids', 1), ('target_ids', [{}])])
def test_malformed_presentation_metadata_is_a_governed_review_error(tmp_path, field, value):
    path, response_path, response = repetition_review(tmp_path, 'Prospective', 'Sterling')
    if field == 'status': response[field] = value
    else: response['findings'][0][field] = value
    response_path.write_text(json.dumps(response))
    assert not quality.verification_response_is_complete(tmp_path, path)
    findings, _ = quality.validate_verifications(tmp_path, request_paths=[path])
    assert any(row.get('recovery_class') == 'verifier_transient' for row in findings)


@pytest.mark.parametrize('deadline', [None, 100.001])
@pytest.mark.parametrize('review_mode', ['complete', 'transient'])
def test_verified_registered_note_does_not_depend_on_optional_editor(tmp_path, monkeypatch, deadline, review_mode):
    source = source_for('Prospective', 'Sterling')
    source['population']['sample_size'] = '40 (planned)'
    reviews = []

    def review(request):
        response = acceptance_verification(request)
        if request['task'] != 'clinical_content_verification': return response
        reviews.append(request['request_id'])
        if review_mode == 'transient' and len(reviews) == 1:
            response['section_assessments'].pop()
        response['findings'] = [{
            'finding_id': 'optional-polish', 'category': 'content', 'artifact': 'icf',
            'check': 'substantive', 'surface': 'shell', 'target_ids': ['icf.shell'],
            'affected_passage': 'The study is expected to include 40 (planned).',
            'issue': 'An enrollment noun would improve readability.',
            'material': False, 'safety_critical': False, 'contradiction': False,
            'obscures_required_information': False, 'materially_unusable': False,
        }]
        return response

    def draft(request):
        if request.get('construction_repair'):
            raise TimeoutError('An optional editor is unavailable; delivery must not depend on it.')
        return drafting.recorded_acceptance_response(request)

    options = {'clock': lambda: 100.0, 'operation_deadline': deadline} if deadline else None
    result, _ = deliver(tmp_path, monkeypatch, source, review=review, draft=draft, generate_options=options)
    assert result['status'] == 'passed', result
    assert len(reviews) == (2 if review_mode == 'transient' else 1) and result['warnings'] == []


@pytest.mark.parametrize('branch,family', FAMILIES)
@pytest.mark.parametrize('category', ['Other', 'Behavioral', 'Diagnostic Test', 'Procedure'])
def test_admin_categories_are_not_used_as_observational_subjects(tmp_path, branch, family, category):
    source = source_for(branch, family)
    source.setdefault('regulatory', {}).setdefault('prs', {})['study_type'] = 'Observational'
    source['design']['intervention_type'] = category
    rendering.render_documents(ROOT, tmp_path, source, {'protocol': [], 'icf': {}, 'prs': {}}, artifact_names={'protocol'})
    text = '\n'.join(p.text for p in Document(tmp_path / 'candidate/protocol.docx').paragraphs)
    expected = f"{'An' if branch == 'Ambispective' else 'A'} {branch.lower()} observational study"
    assert expected in text
    assert f'study of {category}' not in text


def test_source_supported_device_descriptor_is_preserved(tmp_path):
    source = source_for('Prospective', 'Sterling')
    source['regulatory']['prs']['study_type'] = 'Observational'
    source['design']['intervention_type'] = 'Device'
    rendering.render_documents(ROOT, tmp_path, source, {'protocol': [], 'icf': {}, 'prs': {}}, artifact_names={'protocol'})
    text = '\n'.join(p.text for p in Document(tmp_path / 'candidate/protocol.docx').paragraphs)
    assert 'A prospective observational study involving a medical device' in text


@pytest.mark.parametrize('branch,family', FAMILIES)
def test_generated_cover_cautions_include_approved_display_evidence(tmp_path, branch, family):
    source = source_for(branch, family)
    source.setdefault('regulatory', {}).setdefault('prs', {})['study_type'] = 'Observational'
    source.setdefault('design', {})['intervention_type'] = 'Other'
    model = {'protocol': [], 'icf': {}, 'prs': {}}
    report = rendering.render_documents(ROOT, tmp_path, source, model, artifact_names={'protocol'})
    passage = next(p for row in report['artifacts'] for p in row['construction_passages'])
    rendering.render_documents(ROOT, tmp_path, source, {**model, 'construction_repairs': [{
        **passage, 'before': passage['text'], 'after': passage['text'] + ' study.',
    }]}, artifact_names={'protocol'})
    # Supply the heuristic through the public render report, without depending
    # on unrelated missing narrative bodies in this native cover fixture.
    render_report = controlled_render(ROOT, tmp_path)['render']
    render_report['artifacts'][0]['warnings'] = [{
        'category': 'editorial', 'field': 'protocol.docx', 'code': 'visible-duplicate-word',
        'affected_passage': passage['text'] + ' study.', 'publication_disposition': 'warning',
        'issue': 'A duplicated noun is a presentation observation.',
    }]
    path = quality.create_verification_requests(tmp_path, source, render_report)[0]
    request = json.loads(path.read_text())
    caution = next(row for row in request['lexical_review_cautions']
                   if row['caution'].get('code') == 'visible-duplicate-word')
    assert caution['source_evidence'], caution
    assert source['study']['title'] in quality._review_evidence_text(caution['source_evidence'])
