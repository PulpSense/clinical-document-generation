"""Native Word pagination rules override inherited template split settings."""
import json
from pathlib import Path

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

import rendering


@pytest.mark.parametrize("disabled", ["0", "false", "off"])
def test_contact_table_overrides_disabled_row_protection_and_keeps_rows_together(disabled):
    doc = Document()
    caption = doc.add_paragraph("Table 13.3.-1:")
    title = doc.add_paragraph("Contact Information for Study")
    table = doc.add_table(rows=3, cols=4)
    for cell, value in zip(table.rows[0].cells, ["Study Staff", "Business Phone", "e-mail", "Office Phone"]):
        cell.text = value
    for index, row in enumerate(table.rows):
        rule = OxmlElement("w:cantSplit")
        rule.set(qn("w:val"), disabled)
        row._tr.get_or_add_trPr().append(rule)
        for cell in row.cells:
            cell.paragraphs[0].paragraph_format.keep_with_next = True
    doc.add_paragraph("Following section")
    values = [[cell.text for cell in row.cells] for row in table.rows]
    rendering._normalize_protocol_table_pagination(doc)
    assert all(row._tr.trPr.find(qn("w:cantSplit")).get(qn("w:val")) == "true" for row in table.rows)
    for index, row in enumerate(table.rows):
        assert all(p.paragraph_format.keep_with_next is (index < len(table.rows) - 1)
                   for cell in row.cells for p in cell.paragraphs)
    assert caption.paragraph_format.keep_with_next
    assert title.paragraph_format.keep_with_next
    assert values == [[cell.text for cell in row.cells] for row in table.rows]


def test_native_protocol_contact_rows_cannot_split(tmp_path):
    root = Path(__file__).resolve().parents[1]
    source = json.loads((root / "tests/fixtures/prospective-acceptance-source.json").read_text())
    rendering.render_documents(root, tmp_path, source, {"protocol": [], "icf": {}, "prs": {}},
                               artifact_names={"protocol"}, audit_output=False)
    doc = Document(tmp_path / "candidate/protocol.docx")
    table = next(t for t in doc.tables if t.cell(0, 0).text.strip() == "Study Staff")
    assert all(row._tr.trPr.find(qn("w:cantSplit")).get(qn("w:val")) == "true" for row in table.rows)
