"""Canonical fact equivalence at the public Desktop delivery boundary.

Synthetic approved sources and controlled external review, real generation gates.
"""
import json
from pathlib import Path

import pytest
from test_prs_xml import fixture, TEMPLATE, MANUAL_REFERENCE
from prs_xml import generate, validate_output


def test_textual_screening_interval_preserves_prs_fact(tmp_path):
    source = fixture()
    source['procedures']['minimum_days_before_screening_without_participation'] = 'At least ninety days before screening'
    source['population']['exclusion_criteria'].append('Participation in another interventional study within 90 days before screening.')
    output = tmp_path / 'study.xml'
    generate(TEMPLATE, output, source, {'brief_summary': {'text': 'Summary.'}, 'detailed_description': {'text': 'Description.'}})
    assert 'At least 90 days without participation in another interventional study before screening' in output.read_text()
    assert not validate_output(output, source, MANUAL_REFERENCE, generation_template=TEMPLATE)

import copy
from docx import Document
import drafting
import quality
import workflow
from test_handoff_quality import acceptance_verification

ROOT = Path(__file__).resolve().parents[1]


def _desktop_interval(run, source_value, mutation=None):
    source = json.loads((ROOT / 'tests/fixtures/prospective-acceptance-source.json').read_text())
    source['meta']['icf_template'] = 'Sterling'
    source['procedures']['minimum_days_before_screening_without_participation'] = source_value
    source['population']['exclusion_criteria'].append('Participation in another interventional study within 90 days before screening.')
    reference_path = run / 'reference/study.reference.json'
    reference_path.parent.mkdir(parents=True)
    reference_path.write_text(json.dumps(source))
    assert workflow.prepare(run)['status'] == 'awaiting_approval'
    assert workflow.approve(run, approved_by='Synthetic structured-fact replay')['status'] == 'passed'
    approval = copy.deepcopy(json.loads(reference_path.read_text())['approval'])
    revision = run / 'revisions' / approval['revision_id']
    observed = {'mutated': False, 'blocked_fields': set(), 'review_sets': set()}

    def runner(handoffs, remaining):
        for handoff in handoffs:
            path = revision / handoff['request_path']
            request = json.loads(path.read_text())
            if 'verification-requests' in handoff['request_path']:
                observed['review_sets'].add(request['review_set'])
                if mutation and request['task'] == 'clinical_content_verification' and not observed['mutated']:
                    from xml.etree import ElementTree as ET
                    xml_path = revision / 'candidate/study.xml'
                    tree = ET.parse(xml_path)
                    study = next(tree.getroot().iter('clinical_study'))
                    criteria = study.find('eligibility/criteria/textblock')
                    criteria.text = criteria.text.replace('Adults with the target condition.', 'Unapproved eligibility replacement.')
                    study.find('enrollment').text = '99'
                    study.find('primary_outcome/outcome_time_frame').text = 'Month 99'
                    study.remove(study.find('overall_status'))
                    tree.write(xml_path, encoding='utf-8', xml_declaration=True)
                    defects = quality.deterministic_content_check(revision, source)
                    observed['blocked_fields'] = {finding.get('field') for finding in defects}
                    assert {'eligibility/criteria/textblock', 'enrollment', 'primary_outcome[1].outcome_time_frame', 'structure'} <= observed['blocked_fields']
                    assert not (run / 'output').exists()
                    assert all(finding.get('recovery_class') == 'deterministic_structure_defect' for finding in defects if finding.get('category') == 'xml')
                    observed['mutated'] = True
                workflow._save_verification_response(revision, path, acceptance_verification)
                continue
            response = drafting.recorded_acceptance_response(request)
            response_path = revision / request['response_path']
            response_path.parent.mkdir(parents=True, exist_ok=True)
            response_path.write_text(json.dumps(response))

    result = workflow.run_desktop_operation(run, handoff_runner=runner,
        opener=lambda path: Path(path).read_bytes(), operation_id='structured-interval',
        release_identity={'package_fingerprint': 'synthetic-structured-facts'})
    assert result['status'] == 'passed', result.get('stage')
    assert result['delivery']['confirmed']
    assert sorted(result['client_outputs']) == ['output/icf.docx', 'output/protocol.docx', 'output/study.xml']
    assert json.loads(reference_path.read_text())['approval'] == approval
    for filename in ('protocol.docx', 'icf.docx'):
        body = '\n'.join(p.text for p in Document(run / 'output' / filename).paragraphs)
        assert 'At least 90 days without participation in another interventional study before screening' in body
    assert 'At least 90 days without participation in another interventional study before screening' in (run / 'output/study.xml').read_text()
    manifest = json.loads((revision / 'delivery-manifest.json').read_text())
    for item in manifest['client_outputs']:
        delivered = run / item['path']
        assert delivered.read_bytes() == (revision / 'candidate' / delivered.name).read_bytes()
    return run, revision, source, observed


def test_qualified_text_interval_reaches_complete_desktop_delivery(tmp_path, governed_pdfium):
    _desktop_interval(tmp_path / 'run', 'At least ninety days before screening')

@pytest.mark.parametrize('value', ['90', '90 days', 'ninety days', 'At least ninety days before screening', 'At least 90 days before screening'])
def test_equivalent_approved_intervals_reach_complete_desktop_delivery(tmp_path, governed_pdfium, value):
    _, revision, source, _ = _desktop_interval(tmp_path / 'run', value)
    if value == '90':
        protocol_path = revision / 'candidate/protocol.docx'
        original = protocol_path.read_bytes()
        document = Document(protocol_path)
        document.add_paragraph('At least 99 days without participation in another interventional study before screening.')
        document.save(protocol_path)
        findings = quality.deterministic_content_check(revision, source)
        assert any(finding.get('field') == 'subjects.inclusion' for finding in findings)
        protocol_path.write_bytes(original)


def test_supported_xml_textual_interval_is_equivalent_without_waiving_clinical_mutations(tmp_path):
    source = fixture()
    output = tmp_path / 'study.xml'
    generate(TEMPLATE, output, source, {'brief_summary': {'text': 'Summary.'}, 'detailed_description': {'text': 'Description.'}})
    text = output.read_text().replace('At least 90 days without participation', 'At least ninety days without participation')
    output.write_text(text)
    assert not validate_output(output, source, MANUAL_REFERENCE, generation_template=TEMPLATE)


def test_source_backed_xml_mutations_require_real_reconstruction_before_delivery(tmp_path, governed_pdfium):
    _, _, _, observed = _desktop_interval(tmp_path / 'run', '90 days', mutation=True)
    assert observed['mutated']
    assert observed['review_sets'] == {1, 2}


def test_decimal_day_representation_preserves_interval_without_crashing(tmp_path):
    source = fixture()
    source['procedures']['minimum_days_before_screening_without_participation'] = '90.0 days'
    output = tmp_path / 'study.xml'
    generate(TEMPLATE, output, source, {'brief_summary': {'text': 'Summary.'}, 'detailed_description': {'text': 'Description.'}})
    assert 'At least 90 days without participation' in output.read_text()
    assert not validate_output(output, source, MANUAL_REFERENCE, generation_template=TEMPLATE)

@pytest.mark.parametrize('xml', ['<clinical_study>', '<study_collection />'])
def test_malformed_or_missing_prs_study_is_reported_as_a_blocker(tmp_path, xml):
    output = tmp_path / 'study.xml'
    output.write_text(xml)
    findings = validate_output(output, fixture(), MANUAL_REFERENCE, generation_template=TEMPLATE)
    assert findings and all(finding['category'] == 'xml' for finding in findings)
