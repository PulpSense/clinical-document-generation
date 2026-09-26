import copy
import json
from pathlib import Path

import pytest
import drafting
import rendering
from docx import Document

ROOT = Path(__file__).resolve().parents[1]
DATA = json.loads((ROOT / 'tests/fixtures/reliability-replays/run03-duration-references.json').read_text())


def duration_replay(text):
    request = copy.deepcopy(DATA['request'])
    response = copy.deepcopy(DATA['response'])
    response['section_results'][0]['paragraphs'][0]['text'] = text
    return drafting.validate_response(request, response)[1]


GOOD = 'Your follow-up period in this study will last 3 months. The research team plans to enroll participants over 6 months and spend 1 month analyzing the study results.'


def test_archived_correct_duration_passes():
    assert not duration_replay(GOOD)


@pytest.mark.parametrize('text', [GOOD.replace('3 months', '6 months').replace('over 6 months', 'over 3 months'), GOOD.replace('3 months', '3 weeks'), GOOD.replace('and spend 1 month analyzing the study results', ''), GOOD.replace('6 months', '9 months')])
def test_duration_preserves_phase_value_and_unit(text):
    assert duration_replay(text)


def test_number_words_duration_passes():
    assert not duration_replay(GOOD.replace('3 months', 'three months').replace('6 months', 'six months').replace('1 month', 'one month'))


def test_background_bibliography_is_rendered_verbatim():
    reference = copy.deepcopy(DATA['reference'])
    document = Document()
    document.add_paragraph('Study text')
    rendering._ensure_protocol_references(document, Document(), reference)
    text = '\n'.join(p.text for p in document.paragraphs)
    assert 'REFERENCES' in text
    assert '1. Hacopian AG, Brunson PB, Hall B.' in text
    assert 'doi:10.1007/s40123-025-01235-7' in text
    assert 'doi:10.1038/s41433-024-03039-8' in text
    assert 'Mix-and-match implantation' not in text


@pytest.mark.parametrize('text', [
    'Enrollment: 6 months; follow-up: 3 months; data analysis: 1 month.',
    'Enrollment is planned over 6 months (6M). Your follow-up period is 3 months (3M). The study timeline also includes 1 month (1M) for data analysis by the research team.',
])
def test_already_accepted_duration_styles_remain_valid(text):
    assert not duration_replay(text)


@pytest.mark.parametrize('timeline,text', [
    ('Enrollment: 8W; Follow-Up: 2W; Data Analysis: 1W', 'Enrollment lasts 8 weeks. Follow-up lasts 2 weeks. Data analysis lasts 1 week.'),
    ('Enrollment: 2 years; Follow-Up: 6 months; Data Analysis: 10 days', 'Participants enroll over 2 years. Follow-up takes 6 months. Analysis takes 10 days.'),
])
def test_duration_coverage_supports_other_units_and_values(timeline, text):
    assert drafting._timeline_grounded(text, timeline)


def test_unrecognized_timeline_retains_original_grounding_behavior():
    timeline = 'A final interview occurs after the last visit.'
    assert drafting._timeline_grounded('A final interview occurs after the last visit.', timeline)
    assert not drafting._timeline_grounded('Taking part is voluntary.', timeline)


def test_no_bibliography_does_not_create_references():
    document = Document()
    document.add_paragraph('REFERENCES')
    rendering._ensure_protocol_references(document, Document(), {'study': {'background': 'There are 3 visits and 40 participants.'}})
    assert not document.paragraphs


def test_explicit_references_are_preserved_without_duplicate_embedded_entries():
    reference = copy.deepcopy(DATA['reference'])
    entry = '1. Hacopian AG, Brunson PB, Hall B. Patient Reported Outcomes and Visual Acuity After Bilateral Implantation of a Next Generation Presbyopia Correcting Intraocular Lens. OPTH. 2026;20:1-8. doi:10.2147/OPTH.S572703'
    reference['references'] = [entry]
    result = rendering._supplied_protocol_references(reference)
    assert result.count(entry) == 1
    assert 'doi:10.1038/s41433-024-03039-8' in result


def test_heading_delimited_bibliography_without_doi_is_preserved():
    reference = {'study': {'background': 'Clinical background.\nREFERENCES\nSmith A. A supplied source. Journal. 2025;1:2.'}}
    assert rendering._supplied_protocol_references(reference) == 'Smith A. A supplied source. Journal. 2025;1:2.'


def test_references_reach_assembled_protocol(tmp_path):
    reference = json.loads((ROOT / 'tests/fixtures/prospective-acceptance-source.json').read_text())
    reference['study']['background'] = DATA['reference']['study']['background']
    rendering.render_documents(ROOT, tmp_path, reference, {'protocol': [], 'icf': {}, 'prs': {}})
    output = Document(tmp_path / 'candidate/protocol.docx')
    text = '\n'.join(p.text for p in output.paragraphs)
    assert text.count('REFERENCES') == 1
    assert text.count('doi:10.2147/OPTH.S572703') == 1
    assert 'doi:10.1038/s41433-024-03039-8' in text


def test_number_words_in_source_timeline_are_equivalent():
    assert drafting._timeline_grounded(GOOD, 'Enrollment: six months; Follow-Up: three months; Data Analysis: one month')


def test_wrapped_bibliography_entry_is_preserved():
    entry = '1. Smith A. Supplied study. Journal. 2025;1:2.\n doi:10.1234/example'
    result = rendering._supplied_protocol_references({'study': {'background': 'Clinical context.\n' + entry}})
    assert '1. Smith A. Supplied study. Journal. 2025;1:2.' in result
    assert 'doi:10.1234/example' in result


def test_inline_references_heading_is_preserved():
    assert rendering._supplied_protocol_references({'study': {'background': 'Clinical context.\nReferences: Smith A. Supplied study. Journal. 2025.'}}) == 'Smith A. Supplied study. Journal. 2025.'



def test_coordinated_duration_list_uses_its_own_phase_values():
    assert drafting._timeline_grounded('We plan six months of recruitment, three months of follow-up, and one month of analysis.', 'Enrollment: 6M; Follow-up: 3M; Analysis: 1M')
    assert not drafting._timeline_grounded('We plan three months of recruitment, six months of follow-up, and one month of analysis.', 'Enrollment: 6M; Follow-up: 3M; Analysis: 1M')


def test_numbered_journal_reference_without_identifier_is_preserved():
    entry = '1. Smith A. A supplied source. Journal. 2025;1:2.'
    assert rendering._supplied_protocol_references({'study': {'background': 'Clinical background.\n' + entry}}) == entry


def test_numbered_clinical_fact_is_not_a_bibliography():
    assert not rendering._supplied_protocol_references({'study': {'background': '1. Adults age 20 to 80 participate.\n2. Follow-up starts in 2025;1:2 visits are planned.'}})


def test_narrative_after_a_citation_is_not_copied_to_references():
    entry = '1. Smith A. Supplied source. Journal. 2025;1:2. doi:10.1234/test'
    background = entry + '\nParticipants will attend three visits.'
    assert rendering._supplied_protocol_references({'study': {'background': background}}) == entry


def test_conflicting_repeated_phase_duration_is_rejected():
    assert duration_replay(GOOD + ' Enrollment lasts 9 months.')
