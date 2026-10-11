"""Public intake/review handoff tests; controlled adapters do not prove live judgment."""
import json
from pathlib import Path

import workflow
from test_handoff_quality import fixture


def test_source_review_handoff_is_clickable_and_bound_to_editable_file_with_spaces(tmp_path):
    run = tmp_path / 'a study with spaces'
    reference = run / 'reference/study.reference.json'
    reference.parent.mkdir(parents=True)
    reference.write_text(json.dumps(fixture()))
    result = workflow.prepare(run)
    delivery = result['review_delivery']
    source = Path(delivery['absolute_path'])
    assert delivery['link'] == f'[{source.name}](<{source.as_posix()}>)'
    assert delivery['link'] in result['source_review_reply']['message']
    assert result['source_review_reply']['attachments'][0] == delivery
    assert 'approve' in result['source_review_reply']['message'].lower()
    assert result['client_outputs'] == []
    assert source.is_file()


def test_intake_handoff_exposes_obligatory_checklist_and_raw_conflict_evidence_without_parser_veto(tmp_path):
    source = fixture()
    source.setdefault('source', {})['field_candidates'] = {
        'population.sample_size': [{'value': '40 participants', 'source': 'brief'}, {'value': '80 participants', 'source': 'amendment'}],
    }
    source['study']['timeline'] = 'See the supplied follow-up schedule.'
    reference = tmp_path / 'reference/study.reference.json'
    reference.parent.mkdir()
    reference.write_text(json.dumps(source))
    result = workflow.prepare(tmp_path)
    assert result['status'] == 'awaiting_approval'
    guidance = result['intake_review']
    assert any(row['field'] == 'population.sample_size' for row in guidance['obligatory_inputs'])
    assert guidance['field_candidates'] == source['source']['field_candidates']
    assert 'one consolidated question' in guidance['instructions']
    assert 'optional' in guidance['instructions']


def test_content_review_receives_actual_native_file_even_when_a_section_is_missing(tmp_path):
    from docx import Document
    import quality
    candidate = tmp_path / 'candidate'
    candidate.mkdir()
    document = Document()
    document.add_paragraph('A concise candidate needing substantive review.')
    document.sections[0].header.paragraphs[0].text = 'Current native header'
    document.sections[0].footer.paragraphs[0].text = 'Current native footer'
    path = candidate / 'protocol.docx'
    document.save(path)
    requests = quality.create_verification_requests(tmp_path, fixture(), {'artifacts': []})
    request = json.loads(requests[0].read_text())
    artifact = request['artifacts'][0]
    assert artifact['sha256'] == quality.sha256_file(path)
    assert request['reviewer_policy']['native_file_inspection_required'] is True
    assert 'headers' in request['instructions'] and 'footers' in request['instructions']
    assert 'missing-content anchor' in request['instructions']
    assert request['approved_source']['study']['title'] == fixture()['study']['title']


def test_generic_content_pass_cannot_clear_an_explicit_substantive_finding(tmp_path):
    import quality
    from test_handoff_quality import acceptance_verification
    source = {'meta': {'study_type': 'Retrospective'}}
    request_path = quality.create_verification_requests(tmp_path, source, {'artifacts': []})[0]
    request = json.loads(request_path.read_text())
    response = acceptance_verification(request)
    response['findings'] = [{
        'finding_id': 'missing-study-rationale', 'artifact': 'protocol', 'target_ids': ['introduction'],
        'check': 'substantive', 'material': True, 'issue': 'The approved rationale is absent.',
        'missing_content_anchor': 'Introduction', 'recommended_action': 'Restore the approved rationale.',
    }]
    (tmp_path / request['response_path']).write_text(json.dumps(response))
    findings, _ = quality.validate_verifications(tmp_path, request_paths=[request_path])
    assert any(row.get('verification_source_finding', {}).get('finding_id') == 'missing-study-rationale' for row in findings)
    assert not quality.verification_response_is_complete(tmp_path, request_path)


def test_review_context_is_bound_with_source_and_artifact_request_identity(tmp_path):
    import quality
    from test_handoff_quality import acceptance_verification
    context = {'prior_findings': [{'finding_id': 'timing', 'artifact': 'protocol', 'target_ids': ['introduction'], 'issue': 'Review timing.', 'material': True}], 'repair_author_evidence': 'The supplied timeline is preserved.'}
    request_path = quality.create_verification_requests(
        tmp_path, {'meta': {'study_type': 'Retrospective'}}, {'artifacts': []}, review_context=context,
    )[0]
    request = json.loads(request_path.read_text())
    assert request['review_context'] == context
    response = acceptance_verification(request)
    response.pop('source_routing_assessments', None)
    (tmp_path / request['response_path']).write_text(json.dumps(response))
    assert not quality.verification_response_is_complete(tmp_path, request_path)
    for row in response['section_assessments']:
        if row['section_id'] == 'introduction':
            row['notes'] = 'timing: independently checked against the approved timeline; current text preserves its exact meaning.'
    (tmp_path / request['response_path']).write_text(json.dumps(response))
    assert quality.verification_response_is_complete(tmp_path, request_path)
    request['review_context']['repair_author_evidence'] = 'A different assertion.'
    request['request_sha256'] = quality.verification_request_sha256(request)
    request_path.write_text(json.dumps(request))
    assert not quality.verification_response_is_complete(tmp_path, request_path)
