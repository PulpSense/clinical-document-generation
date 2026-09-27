import copy
import json
from pathlib import Path
from docx import Document
import pytest
import contracts
import drafting
import quality

FIXTURE = Path(__file__).parent / 'fixtures/reliability-replays/sterling-rendered-scope'
SOURCE = json.loads((FIXTURE / 'source.json').read_text())
CLAUSES = {'sterling.background.context', 'sterling.procedures.activities'}


@pytest.mark.parametrize('path', sorted(FIXTURE.glob('candidate-*.docx')))
def test_rendered_gate_uses_same_clinical_evidence_as_drafting(path):
    findings = quality.validate_sterling_clause_contract(path, SOURCE)['findings']
    assert not [f for f in findings if f['clause_id'] in CLAUSES]


def test_missing_clinical_distance_remains_blocking():
    doc = Document(FIXTURE / 'candidate-3.docx')
    for paragraph in doc.paragraphs:
        paragraph.text = paragraph.text.replace('66 cm', 'far away')
    findings = quality.validate_sterling_clause_contract(doc, SOURCE)['findings']
    assert any(f['clause_id'] == 'sterling.procedures.activities' for f in findings)


def test_purpose_budget_reserves_fixed_source_enrollment_sentence():
    section = next(s for s in contracts.icf_contract('Prospective', 'Sterling') if s.section_id == 'icf.study-purpose')
    payload = drafting._section_payload(section, {}, SOURCE)
    assert payload['maximum_draft_words'] < 140
    assert any(str(payload['maximum_draft_words']) in s for s in payload['content_expectations'])


def test_full_candidate_gate_passes_after_concise_purpose_and_shared_scopes():
    doc = Document(FIXTURE / 'candidate-3.docx')
    replacement = [
        'The purpose is to describe and compare binocular vision, glasses use, satisfaction and visual disturbances after bilateral PureSee or Odyssey implantation through 6 months after second-eye surgery.',
        SOURCE['study']['hypothesis'],
        'The main outcome is mean binocular distance-corrected photopic intermediate visual acuity at 66 cm, measured 6 months after second-eye surgery and summarized separately for each cohort.',
    ]
    inside = False; index = 0
    for paragraph in doc.paragraphs:
        if paragraph.text == 'PURPOSE': inside = True; continue
        if inside and paragraph.text == 'DURATION': break
        if inside and paragraph.text.strip() and not paragraph.text.startswith('The study is expected to include'):
            paragraph.text = replacement[index]; index += 1
    assert index == 3
    assert quality.validate_sterling_clause_contract(doc, SOURCE)['status'] == 'passed'


def test_purpose_draft_cannot_pass_then_fail_the_same_rendered_word_limit():
    data = json.loads((Path(__file__).parent / 'fixtures/reliability-replays/bilateral-review-block.json').read_text())
    request = copy.deepcopy(data['request'])
    request['branch']['icf_template'] = 'Sterling'
    request['approved_source'] = copy.deepcopy(SOURCE)
    section = next(s for s in contracts.icf_contract('Prospective', 'Sterling') if s.section_id == 'icf.study-purpose')
    payload = drafting._section_payload(section, {}, SOURCE)
    request['section_contracts'] = [payload]
    response = drafting.response_template(request)
    result = response['section_results'][0]
    result['outcome'] = 'drafted'
    result['paragraphs'] = [{'text': 'Purpose ' + ('word ' * payload['maximum_draft_words']), 'evidence_refs': [], 'boilerplate_refs': []}]
    findings = drafting.validate_response(request, response)[1]
    assert any('normalized words' in f['issue'] for f in findings)
