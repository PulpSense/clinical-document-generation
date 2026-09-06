"""Regression coverage for source-faithful fields and narrow layout corrections."""
from pathlib import Path

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

import rendering

ROOT = Path(__file__).resolve().parents[1]


def test_assessment_supplemental_notes_are_full_width_body_paragraphs(monkeypatch, tmp_path):
    real_contracts = rendering.protocol_table_contracts
    note = 'Synthetic source note: review the measure at the approved final contact.'
    def contracts_with_notes(reference):
        contracts = real_contracts(reference)
        contracts['schedule-of-assessments']['supplemental_notes'] = [note]
        return contracts
    monkeypatch.setattr(rendering, 'protocol_table_contracts', contracts_with_notes)
    doc = Document()
    doc.add_paragraph('{visitsTable}')
    reference = {'procedures': {'visit_schedule': [{'visit': 'Final', 'procedures': ['Measure']}]}}
    rendering._assessment_matrix(doc, reference, ROOT / 'assets/client-templates/reference/protocol-reference.docx')
    following = doc.tables[0]._tbl.getnext()
    assert following.tag == qn('w:p')
    from docx.text.paragraph import Paragraph
    paragraph = Paragraph(following, doc)
    assert paragraph.text == note and paragraph.style.name == 'Normal'
    assert paragraph.paragraph_format.left_indent is None
    assert all(note not in cell.text for row in doc.tables[0].rows for cell in row.cells)
    import json
    reference = json.loads((ROOT / 'tests/fixtures/prospective-acceptance-source.json').read_text())
    rendering.render_documents(ROOT, tmp_path, reference, {'protocol': [{'section_id': 'evaluation-procedures', 'paragraphs': [{'text': 'Synthetic approved assessment narrative.'}], 'lists': []}], 'icf': {}, 'prs': {}}, artifact_names={'protocol'})
    generated = Document(tmp_path / 'candidate/protocol.docx')
    assert sum(p.text == note for p in generated.paragraphs) == 1



def test_real_office_refresh_smoke_with_wide_tables(tmp_path):
    import shutil
    import subprocess
    from pypdf import PdfReader
    office = shutil.which('libreoffice') or shutil.which('soffice')
    if office is None:
        pytest.skip('Real office renderer unavailable')
    doc = Document(ROOT / 'assets/client-templates/docx/prospective-protocol.template.docx')
    for child in list(doc.element.body):
        if child.tag != qn('w:sectPr'):
            doc.element.body.remove(child)
    doc.add_paragraph('4. TABLE OF CONTENTS', style='Heading 1')
    intro = doc.add_paragraph('5. INTRODUCTION', style='Heading 1')
    intro.paragraph_format.page_break_before = True
    doc.add_paragraph('Synthetic layout test body.')
    doc.add_paragraph('{visitsTable}')
    doc.add_paragraph('12. CONFIDENTIALITY/PUBLICATION OF THE STUDY', style='Heading 1')
    doc.add_paragraph('Synthetic publication test body.')
    reference = {'procedures': {'visit_schedule': [
        {'visit': label, 'timing': 'Day 1', 'procedures': ['Assessment']} for label in ['Screening', 'Baseline', 'Follow-up', 'Completion', 'Final contact']
    ]}, 'statistics': {'sample_size_evidence': [{'study': 'Example', 'timepoint': 'Day 1', 'mean_change_ods_vas': '10', 'se': '2', 'estimated_sd': '4', 'evidence': 'Synthetic'}]}}
    authority = ROOT / 'assets/client-templates/reference/protocol-reference.docx'
    rendering._assessment_matrix(doc, reference, authority)
    rendering._sample_size_evidence_table(doc, reference, authority)
    assert len(doc.tables) == 2
    for table in doc.tables:
        grid_width = sum(int(col.get(qn('w:w'))) for col in table._tbl.tblGrid)
        assert int(table._tbl.tblPr.find(qn('w:tblW')).get(qn('w:w'))) == grid_width
    assert all(run.font.size != Pt(7) for table in doc.tables for row in table.rows for cell in row.cells for p in cell.paragraphs for run in p.runs)
    path = tmp_path / 'smoke.docx'; doc.save(path)
    for cycle in range(3):
        out = tmp_path / str(cycle); out.mkdir()
        result = subprocess.run([office, f'-env:UserInstallation={(tmp_path / "office-profile").as_uri()}', '--headless', '--convert-to', 'pdf', '--outdir', str(out), str(path)], capture_output=True, text=True, timeout=90)
        assert result.returncode == 0, result.stderr
        pdf = out / 'smoke.pdf'
        assert pdf.is_file(), result.stdout + result.stderr
        text = '\n'.join(page.extract_text() or '' for page in PdfReader(pdf).pages)
        assert 'Synthetic' in text and 'Screening' in text and 'Baseline' in text
        assert rendering.refresh_toc_from_pdf(path, pdf)
        current = Document(path)
        rows = [p for p in current.paragraphs if p.style.name.casefold() in {'toc 1', 'toc 2'}]
        assert len(rows) == 3
        assert sum(' TOC ' in instruction for instruction in current.element.body.xpath('.//w:instrText/text()')) == 1
        assert len(current.element.body.xpath('.//w:bookmarkStart')) == 3



@pytest.mark.parametrize('shape', ['visit_schedule', 'visit_schedule_table'])
@pytest.mark.parametrize('identifiers', [[None, None], ['A', '9']])
def test_visit_rows_share_identifiers_with_assessment_matrix(shape, identifiers):
    from contracts import protocol_table_contracts
    doc = Document()
    table = doc.add_table(rows=2, cols=4)
    for cell, value in zip(table.rows[0].cells, ['Number', 'Visit', 'Timing', 'CRF']):
        cell.text = value
    for cell in table.rows[1].cells:
        cell.text = '{AI_visitName}'
    visits = [{'visit': label, 'visitNumber': number, 'timing': 'Day 1', 'procedures': ['Assessment']} for label, number in zip(['Screening', 'Final'], identifiers)]
    reference = {'procedures': {shape: visits}}
    rendering._visit_rows(doc, reference)
    expected = protocol_table_contracts(reference)['schedule-of-assessments']['rows'][1][1:]
    assert [f'Visit {row.cells[0].text}' for row in table.rows[1:]] == expected


def test_narrative_assessments_do_not_become_visit_rows():
    doc = Document()
    table = doc.add_table(rows=1, cols=4)
    table.cell(0, 0).text = '{AI_visitName}'
    rendering._visit_rows(doc, {'procedures': {'assessments': ['Pain rating', 'Device download']}})
    assert len(table.rows) <= 1
    assert not any('Pain rating' in cell.text or 'Device download' in cell.text for row in table.rows for cell in row.cells)



def test_street_containing_city_does_not_suppress_locality_or_postal_alias():
    facility = {'address': '10 Example City Road', 'city': 'Example City', 'postal_code': '12345'}
    assert rendering._facility_address(facility) == '10 Example City Road, Example City, 12345'
    fields = rendering.render_fields({'sites': [{'facility': facility}]}, {})
    assert fields['studySiteAddress'] == '10 Example City Road, Example City, 12345'



def test_small_assessment_table_retains_activity_column_proportion():
    doc = Document()
    doc.add_paragraph('{visitsTable}')
    reference = {'procedures': {'visit_schedule': [
        {'visit': label, 'timing': 'Day 1', 'procedures': ['Assessment']} for label in ['First', 'Second', 'Last']
    ]}}
    rendering._assessment_matrix(doc, reference, ROOT / 'assets/client-templates/reference/protocol-reference.docx')
    widths = [int(col.get(qn('w:w'))) for col in doc.tables[0]._tbl.tblGrid]
    assert widths[0] == round(sum(widths) * .40)



def test_sponsor_funding_composition_preserves_distinct_address_and_clarification():
    fields = rendering.render_fields({'parties': {
        'sponsor': {'name': 'Example Foundation', 'address': '1 Example Street'},
        'funding_source': {'name': 'Example Foundation', 'address': '2 Example Street', 'clarification': 'Grant support'},
    }}, {})
    doc = Document()
    combined = doc.add_paragraph('{sponsortAdress}{fundingSourceClarification}{fundingSourceName}{fundingSourceAdress}')
    separate = doc.add_paragraph('{fundingSourceName}')
    rendering._replace_paragraph(combined, fields)
    rendering._replace_paragraph(separate, fields)
    assert combined.text == '1 Example Street\nFunding source: Sponsor; Grant support; 2 Example Street'
    assert separate.text == 'Example Foundation'



def test_nested_facility_address_retains_second_line_and_supplied_postal_code():
    facility = {'address': {'address': '10 Main St', 'line2': 'Suite 9', 'city': 'Example Town', 'zip': '12345'}, 'state': 'EX', 'country': 'Example Country'}
    value = rendering._facility_address(facility)
    assert all(part in value for part in ['10 Main St', 'Suite 9', 'Example Town', '12345', 'EX', 'Example Country'])
    assert rendering._facility_address({'address': '10 Main St'}) == '10 Main St'



@pytest.mark.parametrize('branch', ['prospective', 'ambispective'])
@pytest.mark.parametrize('family', ['Advarra', 'Sterling'])
def test_rendered_icf_families_preserve_complete_site_address(tmp_path, branch, family):
    import json
    reference = json.loads((ROOT / f'tests/fixtures/{branch}-acceptance-source.json').read_text())
    reference['meta']['icf_template'] = family
    reference['sites'][0]['facility'] = {'name': 'Test facility', 'address': '123 Example Lane', 'city': 'Example City', 'state': 'EX', 'country': 'Example Country'}
    rendering.render_documents(ROOT, tmp_path, reference, {'protocol': [], 'icf': {}, 'prs': {}}, artifact_names={'icf'})
    doc = Document(tmp_path / 'candidate/icf.docx')
    text = '\n'.join(p.text for p in rendering._all_paragraphs(doc))
    assert '123 Example Lane, Example City, EX, Example Country' in text



def test_toc_refresh_retains_live_field_heading_links_and_updates_cache(tmp_path):
    from pypdf import PdfWriter
    from pypdf.generic import DictionaryObject, NameObject, DecodedStreamObject
    doc = Document(ROOT / 'assets/client-templates/docx/prospective-protocol.template.docx')
    # Use the client TOC styles on a small real document/PDF pair.
    for child in list(doc.element.body):
        if child.tag != qn('w:sectPr'):
            doc.element.body.remove(child)
    doc.add_paragraph('4. TABLE OF CONTENTS', style='Heading 1')
    doc.add_paragraph('5. INTRODUCTION', style='Heading 1')
    path = tmp_path / 'protocol.docx'; doc.save(path)
    pdf = tmp_path / 'protocol.pdf'
    for expected_page in (2, 3, 3):
        writer = PdfWriter()
        for index in range(expected_page):
            page = writer.add_blank_page(width=612, height=792)
            font = DictionaryObject({NameObject('/Type'): NameObject('/Font'), NameObject('/Subtype'): NameObject('/Type1'), NameObject('/BaseFont'): NameObject('/Helvetica')})
            page[NameObject('/Resources')] = DictionaryObject({NameObject('/Font'): DictionaryObject({NameObject('/F1'): writer._add_object(font)})})
            text = '4. TABLE OF CONTENTS' if index == 0 else '5. INTRODUCTION' if index == expected_page - 1 else 'Body material'
            stream = DecodedStreamObject(); stream.set_data(f'BT /F1 12 Tf 50 700 Td ({text}) Tj ET'.encode())
            page[NameObject('/Contents')] = writer._add_object(stream)
        writer.write(pdf)
        assert rendering.refresh_toc_from_pdf(path, pdf)
        updated = Document(path)
        instructions = updated.element.body.xpath('.//w:instrText/text()')
        assert sum(' TOC ' in value for value in instructions) == 1
        assert any('PAGEREF ' in value for value in instructions)
        rows = [p for p in updated.paragraphs if p.style.name.casefold() in {'toc 1', 'toc 2'}]
        assert [p.text for p in rows] == ['4. TABLE OF CONTENTS\t1', f'5. INTRODUCTION\t{expected_page}']
        bookmarks = set(updated.element.body.xpath('.//w:bookmarkStart/@w:name'))
        assert all(link.get(qn('w:anchor')) in bookmarks for p in rows for link in p._p.xpath('./w:hyperlink'))
        assert len(updated.element.body.xpath('.//w:bookmarkStart')) == 2



def test_wide_assessment_table_preserves_font_and_coherent_widths():
    doc = Document()
    doc.add_paragraph('{visitsTable}')
    reference = {'meta': {'study_type': 'Prospective'}, 'procedures': {'visit_schedule': [
        {'visit': label, 'timing': 'Day 1', 'procedures': ['Assessment']}
        for label in ['Screening', 'Baseline', 'Follow-up', 'Completion', 'Final contact']
    ]}}
    authority = ROOT / 'assets/client-templates/reference/protocol-reference.docx'
    rendering._assessment_matrix(doc, reference, authority)
    table = doc.tables[0]
    assert all(run.font.size != Pt(7) for row in table.rows for cell in row.cells for p in cell.paragraphs for run in p.runs)
    widths = [int(col.get(qn('w:w'))) for col in table._tbl.tblGrid]
    assert int(table._tbl.tblPr.find(qn('w:tblW')).get(qn('w:w'))) == sum(widths)
    for row in table._tbl.tr_lst:
        assert [int(cell.tcPr.find(qn('w:tcW')).get(qn('w:w'))) for cell in row.tc_lst] == widths
    assert min(widths[1:]) >= 1100  # usable whole-word contact labels, not a 40% activity allocation



def test_only_section_15_template_break_is_removed_and_short_tail_has_context():
    doc = Document()
    toc = doc.add_paragraph('4. TABLE OF CONTENTS', style='Heading 1')
    intro = doc.add_paragraph('5. INTRODUCTION', style='Heading 1')
    doc.add_paragraph('14.1 Confidentiality', style='Heading 2')
    context = doc.add_paragraph('A substantive confidentiality paragraph provides protections and describes access.')
    tail = doc.add_paragraph('Confidentiality protections are described in Section 16.')
    spacer = doc.add_paragraph()
    br = OxmlElement('w:br'); br.set(qn('w:type'), 'page'); spacer.add_run()._r.append(br)
    evaluation = doc.add_paragraph('15. STANDARD EVALUATION PROCEDURES', style='Heading 1')
    doc.add_paragraph('Body')
    retained = doc.add_paragraph()
    br = OxmlElement('w:br'); br.set(qn('w:type'), 'page'); retained.add_run()._r.append(br)
    doc.add_paragraph('16. DATA MANAGEMENT', style='Heading 1')
    rendering._normalize_protocol_section_pagination(doc)
    assert not rendering._has_page_boundary_before(evaluation)
    assert rendering._has_page_boundary_before(toc) and rendering._has_page_boundary_before(intro)
    assert retained._p.xpath('.//w:br[@w:type="page"]')
    assert context.paragraph_format.keep_with_next and tail.paragraph_format.keep_together
    assert tail.paragraph_format.keep_with_next is not True
    xml = doc.element.xml
    rendering._normalize_protocol_section_pagination(doc)
    assert doc.element.xml == xml



def test_preference_block_removes_only_internal_cosmetic_spacers():
    doc = Document()
    signature = doc.add_paragraph('Signature: __________________')
    signing_space = doc.add_paragraph()
    prompt = doc.add_paragraph('Check your preference below:')
    blank = doc.add_paragraph()
    yes = doc.add_paragraph('Yes, inform my primary care physician')
    yes.runs[0].font.size = Pt(12)
    yes.runs[0].bold = True
    doc.add_paragraph()
    no = doc.add_paragraph('No, do not inform my primary care physician')
    trailing = doc.add_paragraph()
    rendering._normalize_icf_preferences(doc)
    assert blank._p.getparent() is None
    assert signing_space._p.getparent() is not None and trailing._p.getparent() is not None
    assert signature.text == 'Signature: __________________'
    assert yes.text.startswith('☐ Yes,') and no.text.startswith('☐ No,')
    assert yes.runs[0].font.size == Pt(12) and yes.runs[0].bold
    assert prompt.paragraph_format.keep_with_next and yes.paragraph_format.keep_with_next
    assert no.paragraph_format.keep_with_next is not True
    xml = doc.element.xml
    rendering._normalize_icf_preferences(doc)
    assert doc.element.xml == xml



@pytest.mark.parametrize('sterling', [False, True])
def test_peer_icf_headings_share_indentation_without_changing_typography(sterling):
    doc = Document()
    headings = []
    for index, text in enumerate(['LENGTH OF STUDY', 'POSSIBLE BENEFITS', 'PAYMENT']):
        p = doc.add_paragraph()
        r = p.add_run(text)
        r.bold = r.underline = True
        r.font.name = 'Arial'
        r.font.size = Pt(12)
        p.paragraph_format.left_indent = Pt(index * 11)
        p.paragraph_format.first_line_indent = Pt(index * 3)
        headings.append(p)
        doc.add_paragraph('Body paragraph.')
    rendering._normalize_icf_heading_styles(doc, sterling=sterling)
    assert len({(p.paragraph_format.left_indent, p.paragraph_format.first_line_indent) for p in headings}) == 1
    for p in headings:
        assert p.runs[0].bold and p.runs[0].underline
        assert p.runs[0].font.name == 'Arial' and p.runs[0].font.size == Pt(12)
        assert p.paragraph_format.page_break_before is not True
    assert doc.paragraphs[1].paragraph_format.left_indent is None


@pytest.mark.parametrize('branch', ['prospective', 'ambispective', 'retrospective'])
@pytest.mark.parametrize('funding', ['Sponsor Foundation', 'Other Foundation', ''])
def test_sponsor_block_separates_funding_without_duplicate_entity(branch, funding):
    doc = Document(ROOT / f'assets/client-templates/docx/{branch}-protocol.template.docx')
    fields = rendering.render_fields({'parties': {
        'sponsor': {'name': 'Sponsor Foundation', 'address': '10 Main St, Country'},
        'funding_source': {'name': funding},
    }}, {})
    for paragraph in rendering._all_paragraphs(doc):
        rendering._replace_paragraph(paragraph, fields)
    cell = next(cell for row in doc.tables[0].rows for cell in row.cells if '10 Main St' in cell.text)
    assert cell.text.count('Sponsor Foundation') == 1
    assert '10 Main St, Country' in cell.text
    if funding == 'Other Foundation':
        assert '\nFunding source: Other Foundation' in cell.text
    assert 'CountryOther' not in cell.text


@pytest.mark.parametrize('street', ['123 Example Lane', {'street': '123 Example Lane'}, '123 Example Lane, Example City, EX, Example Country'])
def test_facility_front_matter_preserves_sibling_locality_without_duplication(street):
    doc = Document()
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = 'Address:'
    reference = {'sites': [{'facility': {'address': street, 'city': 'Example City', 'state': 'EX', 'country': 'Example Country'}}]}
    rendering._normalize_icf_front_matter(doc, reference)
    assert table.cell(0, 1).text == '123 Example Lane, Example City, EX, Example Country'
    assert not any(character.isdigit() for character in table.cell(0, 1).text.split('Lane')[1])
