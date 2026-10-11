"""Public generation/recovery replays; controlled reviews are not release evidence."""
import copy
import json
from pathlib import Path

import pytest
from docx import Document

import quality
import drafting
import rendering
import workflow
from test_clean_delivery import deliver, with_material_repair
from test_handoff_quality import acceptance_verification

ROOT = Path(__file__).resolve().parents[1]
FAMILIES = [('Prospective', 'Sterling'), ('Prospective', 'Advarra'),
            ('Ambispective', 'Sterling'), ('Ambispective', 'Advarra'), ('Retrospective', '')]


def source_for(branch, family='Advarra'):
    case = next(case for case in quality.branch_acceptance_inventory(ROOT)
                if case['descriptor']['study_type'] == branch
                and (branch == 'Retrospective' or case['descriptor']['icf_template'] == family))
    return copy.deepcopy(quality.branch_acceptance_source(ROOT, case))


@pytest.mark.parametrize('branch,family', FAMILIES)
@pytest.mark.parametrize('count', ['40', 40])
def test_bare_approved_count_generates_complete_enrollment_sentence(tmp_path, monkeypatch, branch, family, count):
    source = source_for(branch, family)
    source['population']['sample_size'] = count
    result, revision = deliver(tmp_path, monkeypatch, source)
    assert result['status'] == 'passed', result
    name = 'protocol.docx' if branch == 'Retrospective' else 'icf.docx'
    document = Document(tmp_path / 'output' / name)
    text = '\n'.join([p.text for p in document.paragraphs] +
                     [cell.text for table in document.tables for row in table.rows for cell in row.cells])
    assert ('sample size is 40.' if branch == 'Retrospective' else '40 participants') in text
    assert '40 participants participants' not in text
    if family == 'Sterling':
        assert 'The study is expected to include 40 participants.' in text
    # Intake presents scalar numbers as editable source text before approval.
    assert json.loads((revision / 'approved-reference.json').read_text())['population']['sample_size'] == str(count)
    assert sorted(p.name for p in (tmp_path / 'output').iterdir()) == sorted(workflow.document_set(branch))


@pytest.mark.parametrize('repair_mode', ['faithful', 'multiple_findings', 'eventual', 'unchanged', 'changed_quantity', 'changed_meaning', 'malformed'])
def test_generated_display_prose_is_repaired_persisted_and_fully_reviewed(tmp_path, monkeypatch, repair_mode):
    source = source_for('Prospective', 'Sterling')
    source['population']['sample_size'] = '40 (planned)'
    reviews = []
    repairs = []

    def review(request):
        response = acceptance_verification(request)
        if request['task'] == 'clinical_content_verification':
            text = Document(tmp_path / 'revisions' / request['revision_id'] / 'candidate/icf.docx')
            incomplete = 'The study is expected to include 40 (planned).'
            reviews.append(request['request_id'])
            passage = next((p.text for p in text.paragraphs if p.text in {
                incomplete, 'The study is expected to include 40 animals (planned).'}), None)
            if passage:
                response['status'] = 'blocked'
                response['findings'] = [{
                    'finding_id': 'enrollment-unit', 'category': 'content', 'artifact': 'icf',
                    'target_ids': ['icf.shell'], 'check': 'editorial_relevance', 'surface': 'document_shell',
                    'affected_passage': passage, 'issue': 'The enrollment sentence must describe participants.',
                    'recommended_action': 'Identify the counted entity while preserving the planned count.',
                    'material': 'animals' in passage, 'safety_critical': False, 'contradiction': 'animals' in passage,
                    'obscures_required_information': False, 'materially_unusable': False,
                }]
                if repair_mode == 'multiple_findings':
                    response['findings'] = [{**response['findings'][0], 'finding_id': f'enrollment-{i}',
                                             'issue': f'Editorial concern {i} on this enrollment sentence.'} for i in range(4)]
                for row in response['section_assessments']:
                    if row['section_id'] == 'icf.shell':
                        row['status'] = 'blocked'
        return response

    def draft(request):
        if request.get('construction_repair'):
            repairs.append(request)
            response = drafting.response_template(request, model_id='Controlled construction editor')
            mode = ('unchanged' if len(repairs) == 1 else 'faithful') if repair_mode == 'eventual' else repair_mode
            if mode == 'multiple_findings': mode = 'faithful'
            response['section_results'][0]['paragraphs'] = [{
                'text': {
                    'faithful': 'The study is expected to include 40 participants (planned).',
                    'unchanged': 'The study is expected to include 40 (planned).',
                    'changed_quantity': 'The study is expected to include 41 participants (planned).',
                    'changed_meaning': 'The study is expected to include 40 animals (planned).',
                    'malformed': 'The study is expected to include 40 participants (planned).',
                }[mode],
                'evidence_refs': ['source:population.sample_size'], 'boilerplate_refs': [],
            }]
            if mode == 'malformed':
                response['schema_version'] = 'wrong-schema'
            return response
        return drafting.recorded_acceptance_response(request)

    review, draft = with_material_repair(review, draft)
    result, revision = deliver(tmp_path, monkeypatch, source, review=review, draft=draft)
    if repair_mode == 'changed_meaning':
        assert result['status'] != 'passed', result
        assert not result.get('client_outputs') and not (tmp_path / 'output/icf.docx').exists()
        # An essential meaning failure receives the shared bounded construction route.
        assert len(repairs) == 3 and len(reviews) == 2
        return
    assert result['status'] == 'passed', result
    assert len(repairs) == 1
    if repair_mode == 'multiple_findings': assert len(repairs[0]['findings']) == 4
    assert len(reviews) == 2
    expected = ('The study is expected to include 40 participants (planned).'
                if repair_mode in {'faithful', 'multiple_findings'} else 'The study is expected to include 40 (planned).')
    assert expected in '\n'.join(
        p.text for p in Document(tmp_path / 'output/icf.docx').paragraphs)
    assert json.loads((revision / 'approved-reference.json').read_text())['population']['sample_size'] == '40 (planned)'
    manifest = json.loads((tmp_path / result['manifest']).read_text())
    assert quality.final_exact_artifact_review_findings(
        revision, json.loads((revision / 'approved-reference.json').read_text()),
        json.loads((revision / 'candidate-build.json').read_text())['render_report'], manifest['quality']) == []
    if repair_mode not in {'faithful', 'multiple_findings'}:
        assert manifest['quality']['internal_diagnostics']
        assert result['warnings'] == []
        assert list((revision / 'hermes/rejected').glob('*.json'))
        return
    # An on-disk correction is never trusted independently of its bound response.
    record_path = next((revision / 'hermes/accepted').glob('construction-*.json'))
    record = json.loads(record_path.read_text())
    record_path.write_text(json.dumps({**record, 'after': record['after'].replace('participants', 'animals')}))
    errors = drafting.ingest_responses(revision, record['governing_resources'])
    assert any(row['category'] == 'request-integrity' for row in errors)
    record_path.write_text(json.dumps(record))
    archived_manifest = revision / repairs[0]['construction_repair']['attempt_path'] / 'attempt-manifest.json'
    retained = json.loads(archived_manifest.read_text())
    for broken in ([], [None], [{'verification_source_finding': 'invalid'}]):
        archived_manifest.write_text(json.dumps({**retained, 'findings': broken}))
        errors = drafting.ingest_responses(revision, record['governing_resources'])
        assert any(row['category'] == 'request-integrity' for row in errors)
    archived_manifest.write_text(json.dumps(retained))


def test_unrelated_document_change_does_not_count_as_localized_shell_progress(tmp_path, monkeypatch):
    source = source_for('Prospective', 'Sterling')
    actual_renderer = workflow.render_documents
    builds, reviews, editors = [], [], []

    def constructor(root, revision, reference, model, **kwargs):
        report = actual_renderer(root, revision, reference, model, **kwargs)
        builds.append(report)
        if len(builds) > 1:
            path = revision / 'candidate/icf.docx'
            doc = Document(path)
            doc.add_paragraph('A separate changed paragraph in this controlled recovery replay.')
            doc.save(path)
        return report

    def review(request):
        response = acceptance_verification(request)
        if request['task'] == 'clinical_content_verification':
            reviews.append(request['request_id'])
            doc = Document(tmp_path / 'revisions' / request['revision_id'] / 'candidate/icf.docx')
            protected = next(p.text for p in doc.paragraphs if '«Protocol_Title»' in p.text)
            response['status'] = 'blocked'
            response['findings'] = [{
                'finding_id': 'protected-shell-replay', 'category': 'content', 'artifact': 'icf',
                'target_ids': ['icf.shell'], 'check': 'editorial_relevance', 'surface': 'document_shell',
                'affected_passage': protected, 'issue': 'Controlled finding on protected template text.',
                'material': True, 'safety_critical': False, 'contradiction': False,
                'obscures_required_information': False, 'materially_unusable': False,
            }]
            for row in response['section_assessments']:
                if row['section_id'] == 'icf.shell': row['status'] = 'blocked'
        return response

    def draft(request):
        if request.get('construction_repair'): editors.append(request)
        return drafting.recorded_acceptance_response(request)

    monkeypatch.setattr(workflow, 'render_documents', constructor)
    result, revision = deliver(tmp_path, monkeypatch, source, review=review, draft=draft)
    assert result['status'] != 'passed'
    assert not editors and len(reviews) == 1
    actions = json.loads(next((revision / 'attempts').glob('quality-*/attempt-manifest.json')).read_text())['recovery_actions']
    assert any(a['candidate_bytes_changed'] is True and a['target_bytes_changed'] is False for a in actions)


@pytest.mark.parametrize('value,expected', [
    (1, '1 participant'), ('1,000', '1,000 participants'),
    ('40 records', '40 records'), ('40 participants', '40 participants'),
    ('40 subjects total; 20 per cohort', '40 subjects total; 20 per cohort'),
    ('40 subjects (treatment); control: 0', '40 participants. There is no control group'),
])
def test_native_template_preserves_count_units_breakdowns_and_other_measurements(tmp_path, value, expected):
    source = source_for('Ambispective', 'Sterling')
    source['population']['sample_size'] = value
    background = 'A reading of 40 mmHg and a 40-minute session are separate measurements.'
    model = {'protocol': [], 'prs': {}, 'icf': {
        'icf.background': {'paragraphs': [{'text': background}], 'lists': []},
        'icf.procedures': {'paragraphs': [{'text': 'The planned assessment lasts 40 minutes.'}], 'lists': []},
    }}
    rendering.render_documents(ROOT, tmp_path, source, model, artifact_names={'icf'})
    text = '\n'.join(p.text for p in Document(tmp_path / 'candidate/icf.docx').paragraphs)
    assert f'The study is expected to include {expected}.' in text
    assert background in text


@pytest.mark.parametrize('branch,family', FAMILIES)
def test_localized_repair_covers_each_document_family(tmp_path, monkeypatch, branch, family):
    """Inject a stable construction typo at the public document-renderer boundary."""
    source = source_for(branch, family)
    actual_renderer = workflow.render_documents
    injected = []
    reviewed = []
    edits = []

    def constructor(root, revision, reference, model, **kwargs):
        if not injected:
            report = actual_renderer(root, revision, reference, model, **kwargs)
            passage = next(p for row in report['artifacts'] if row['artifact'] == 'protocol'
                           for p in row['construction_passages'])
            # Only presentation changes. This is a controlled fault adapter,
            # not a fixture accepted by the production editor trust boundary.
            injected.append({**passage, 'before': passage['text'], 'after': passage['text'] + ' study.'})
        return actual_renderer(root, revision, reference,
                               {**model, 'construction_repairs': [*injected, *model.get('construction_repairs', [])]}, **kwargs)

    def review(request):
        response = acceptance_verification(request)
        if request['task'] != 'clinical_content_verification':
            return response
        reviewed.append(request['request_id'])
        document = Document(tmp_path / 'revisions' / request['revision_id'] / 'candidate/protocol.docx')
        if any(p.text == injected[0]['after'] for p in document.paragraphs):
            response['status'] = 'blocked'
            response['findings'] = [{
                'finding_id': 'descriptor-duplicate-word', 'category': 'content', 'artifact': 'protocol',
                'target_ids': ['protocol.shell'], 'check': 'editorial_relevance', 'surface': 'document_shell',
                'affected_passage': injected[0]['after'], 'issue': 'The descriptor duplicates its terminal study noun.',
                'material': False, 'safety_critical': False, 'contradiction': False,
                'obscures_required_information': False, 'materially_unusable': False,
            }]
            for row in response['section_assessments']:
                if row['section_id'] == 'protocol.shell': row['status'] = 'blocked'
        return response

    def draft(request):
        if not request.get('construction_repair'):
            return drafting.recorded_acceptance_response(request)
        edits.append(request)
        response = drafting.response_template(request, model_id='Controlled editor')
        response['section_results'][0]['paragraphs'] = [{
            'text': injected[0]['before'],
            'evidence_refs': [f"source:{item['path']}" for item in request['approved_input'] if item['value']],
            'boilerplate_refs': [],
        }]
        return response

    monkeypatch.setattr(workflow, 'render_documents', constructor)
    review, draft = with_material_repair(review, draft)
    result, revision = deliver(tmp_path, monkeypatch, source, review=review, draft=draft)
    assert result['status'] == 'passed', result
    assert len(edits) == 1 and len(reviewed) == 2
    final_text = '\n'.join(p.text for p in Document(tmp_path / 'output/protocol.docx').paragraphs)
    assert injected[0]['before'] in final_text and injected[0]['after'] not in final_text
    assert sorted(p.name for p in (tmp_path / 'output').iterdir()) == sorted(workflow.document_set(branch))
    manifest = json.loads((tmp_path / result['manifest']).read_text())
    assert quality.final_exact_artifact_review_findings(
        revision, json.loads((revision / 'approved-reference.json').read_text()),
        json.loads((revision / 'candidate-build.json').read_text())['render_report'], manifest['quality']) == []
