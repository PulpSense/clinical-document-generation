import copy
import json
from pathlib import Path
import rendering

ROOT = Path(__file__).resolve().parents[1]
REPLAY = json.loads((ROOT / 'tests/fixtures/reliability-replays/synopsis-matrix-recovery.json').read_text())


def test_actual_matrix_source_produces_participant_followup_and_visit_count():
    source = copy.deepcopy(REPLAY['source'])
    fields = rendering.render_fields(source, {})
    assert fields['AI_duration'] == '6 months after second-eye surgery'
    assert fields['totalVisits'] == '5'
    assert source == REPLAY['source']


def review_result(tmp_path, finding):
    import quality
    request_path = quality.create_verification_requests(tmp_path, REPLAY['source'], {'artifacts': []})[0]
    request = json.loads(request_path.read_text())
    finding = {**finding, 'finding_id': 'synopsis-editorial'}
    response = {
        'schema_version': quality.RESPONSE_SCHEMA,
        **{key: request[key] for key in ('request_id','request_sha256','task','revision_id')},
        'producer': {'model_id':'fixture-reviewer','reviewer_id':'independent-fixture-reviewer'},
        'status':'blocked', 'findings':[finding],
        'section_assessments':[
            {'artifact':item['artifact'],'section_id':item['section_id'],
             'status':'failed' if item['section_id'] == 'general-information' else 'passed',
             'checks':list(quality.CONTENT_CHECKS),'notes':'synopsis-editorial'}
            for item in request['sections']],
        'cross_document_assessments':[{'check':check,'status':'passed'} for check in request['cross_document_checks']],
    }
    (tmp_path/request['response_path']).write_text(json.dumps(response))
    return quality.validate_verifications(tmp_path, request_paths=[request_path])


def test_explicitly_nonmaterial_editorial_review_is_a_retained_warning(tmp_path):
    finding = {**REPLAY['finding'],'material':False}
    findings,evidence = review_result(tmp_path,finding)
    warning = next(item for item in findings if item.get('check') == 'editorial_relevance')
    assert warning['publication_disposition'] == 'warning'
    assert warning['action'] == 'manual_review'
    assert evidence


import pytest


@pytest.mark.parametrize('flag', ['material','safety_critical','contradiction','obscures_required_information','materially_unusable'])
def test_material_or_uncertain_editorial_findings_cannot_become_warnings(tmp_path, flag):
    finding = {**REPLAY['finding'],'material':False,flag:True}
    findings,_ = review_result(tmp_path,finding)
    assert any(item.get('check') == 'editorial_relevance' and item.get('publication_disposition') != 'warning' for item in findings)


@pytest.mark.parametrize('check,target', [('source_supported','general-information'),('no_invention','general-information'),('editorial_relevance','icf.risks')])
def test_factual_or_safety_findings_remain_blocking(tmp_path, check, target):
    finding = {**REPLAY['finding'],'material':False,'check':check,'target_ids':[target]}
    if target.startswith('icf.'):
        finding['artifact']='icf'
        finding['safety_critical']=True
    findings,_ = review_result(tmp_path,finding)
    assert any(item.get('issue') == finding['issue'] and item.get('publication_disposition') != 'warning' for item in findings)


def test_computed_synopsis_is_checked_before_independent_review(tmp_path):
    from docx import Document
    import quality
    doc=Document();table=doc.add_table(rows=1,cols=2)
    table.cell(0,0).text='Duration / Follow‑up'
    table.cell(0,1).text=REPLAY['source']['study']['timeline']
    findings=quality.audit_computed_protocol_fields(doc, REPLAY['source'])
    assert findings[0]['target_ids']==['general-information']
    assert findings[0]['expected']=='6 months after second-eye surgery'
    table.cell(0,1).text='6 months after second-eye surgery'
    assert not quality.audit_computed_protocol_fields(doc, REPLAY['source'])


@pytest.mark.parametrize('omitted', ['material','affected_passage','recommended_action'])
def test_incomplete_editorial_severity_evidence_remains_blocking(tmp_path, omitted):
    finding={**REPLAY['finding'],'material':False}
    finding.pop(omitted)
    findings,_=review_result(tmp_path,finding)
    assert not any(item.get('publication_disposition') == 'warning' for item in findings)


@pytest.mark.parametrize('timeline,expected', [
    ('Enrollment: 8 months; follow-up: 6 months after second-eye surgery; data analysis: 2 months.', '6 months after second-eye surgery'),
    ('Follow-up: 1.5 months after surgery, data analysis: 2 months.', '1.5 months after surgery'),
])
def test_timeline_projection_preserves_anchor_and_excludes_other_phases(timeline, expected):
    source=copy.deepcopy(REPLAY['source']);source['study']['timeline']=timeline
    assert rendering.render_fields(source,{})['AI_duration']==expected


def test_reconstruction_changes_the_actual_synopsis_target_and_allows_fresh_review(tmp_path, monkeypatch):
    from docx import Document
    import workflow
    import quality
    source=copy.deepcopy(REPLAY['source'])
    run=tmp_path/'run';revision=run/'revisions/r-test';reference_path=run/'reference/study.reference.json'
    reference_path.parent.mkdir(parents=True);reference_path.write_text(json.dumps({'generation':{}}))
    rendering.render_documents(ROOT,revision,source,{'protocol':[],'icf':{},'prs':{}},artifact_names={'protocol'})
    path=revision/'candidate/protocol.docx'
    doc=Document(path)
    for table in doc.tables:
        for row in table.rows:
            if row.cells[0].text=='Duration / Follow‑up':row.cells[1].text=source['study']['timeline']
    doc.save(path)
    finding=quality.audit_computed_protocol_fields(path,source)[0]

    def reconstruct(_run, **kwargs):
        rendering.render_documents(ROOT,revision,source,{'protocol':[],'icf':{},'prs':{}},artifact_names={'protocol'})
        latest=json.loads(reference_path.read_text())
        workflow._complete_pending_deterministic_reconstructions(reference_path,latest,revision,outcome='rebuilt')
        assert not workflow._complete_pending_recovery_attempts(revision,reference_path,latest,require_candidate_change=True)
        assert not quality.audit_computed_protocol_fields(path,source)
        return {'status':'awaiting_hermes','stage':'independent_verification'}

    monkeypatch.setattr(workflow,'generate',reconstruct)
    result=workflow._quality_retry(run,reference_path,{'generation':{}},source,revision,{},[finding],'quality',require_promoted_runtime=False)
    assert result['stage']=='independent_verification'
    manifest=json.loads(next((revision/'attempts').glob('*/attempt-manifest.json')).read_text())
    assert manifest['recovery_actions'][0]['target_bytes_changed'] is True


def test_unrelated_document_change_does_not_count_as_synopsis_repair(tmp_path, monkeypatch):
    from docx import Document
    import workflow
    import quality
    source=copy.deepcopy(REPLAY['source']);run=tmp_path/'run';revision=run/'revisions/r-test'
    reference_path=run/'reference/study.reference.json';reference_path.parent.mkdir(parents=True)
    reference_path.write_text(json.dumps({'generation':{}}))
    rendering.render_documents(ROOT,revision,source,{'protocol':[],'icf':{},'prs':{}},artifact_names={'protocol'})
    path=revision/'candidate/protocol.docx';doc=Document(path)
    for table in doc.tables:
        for row in table.rows:
            if row.cells[0].text=='Duration / Follow‑up':row.cells[1].text=source['study']['timeline']
    doc.save(path);finding=quality.audit_computed_protocol_fields(path,source)[0]

    def ineffective(_run, **kwargs):
        doc.add_paragraph('An unrelated later paragraph changed.');doc.save(path)
        latest=json.loads(reference_path.read_text())
        workflow._complete_pending_deterministic_reconstructions(reference_path,latest,revision,outcome='rebuilt')
        no_progress=workflow._complete_pending_recovery_attempts(revision,reference_path,latest,require_candidate_change=True)
        assert no_progress
        assert quality.audit_computed_protocol_fields(path,source)
        return {'status':'blocked','stage':'measured_no_progress'}
    monkeypatch.setattr(workflow,'generate',ineffective)
    result=workflow._quality_retry(run,reference_path,{'generation':{}},source,revision,{},[finding],'quality',require_promoted_runtime=False)
    assert result['stage']=='measured_no_progress'


def test_retained_icf_section_is_observed_without_an_ai_draft(tmp_path):
    from docx import Document
    path=tmp_path/'icf.docx';doc=Document()
    doc.add_paragraph('VOLUNTARY PARTICIPATION/WITHDRAWAL')
    doc.add_paragraph('You may leave the study.')
    doc.add_paragraph('QUESTIONS');doc.add_paragraph('Call the study team.')
    doc.save(path)
    before=rendering.rendered_section_snapshot(path,'icf.voluntary-participation')
    doc.paragraphs[3].text='Call another study contact.';doc.save(path)
    assert rendering.rendered_section_snapshot(path,'icf.voluntary-participation')==before
    doc.paragraphs[1].text='You may leave the study at any time.';doc.save(path)
    assert rendering.rendered_section_snapshot(path,'icf.voluntary-participation')!=before


def test_recorded_editorial_warning_allows_complete_review_and_atomic_delivery(governed_pdfium):
    import workflow
    import quality
    from test_release_gate import acceptance_verification

    def reviewer(request):
        response=acceptance_verification(request)
        if request['task']=='clinical_content_verification':
            response['status']='blocked'
            response['findings']=[{
                'finding_id':'minor-synopsis-style','category':'content','artifact':'protocol',
                'check':'editorial_relevance','target_ids':['general-information'],
                'issue':'The synopsis has an optional wording improvement.',
                'affected_passage':'Duration / Follow-up', 'recommended_action':'Optional concise wording.',
                'material':False,'safety_critical':False,'contradiction':False,
                'obscures_required_information':False,'materially_unusable':False,
            }]
            for item in response['section_assessments']:
                if item['section_id']=='general-information':
                    item['status']='failed';item['notes']='minor-synopsis-style'
        return response

    report=workflow.run_release_gate(ROOT,verification_responder=reviewer,case_ids=('prospective-advarra-rich-complete',))
    assert report['status']=='structural_passed',report
    case=report['cases'][0]
    root=Path(report['evidence_root'])/case['fixture_id'] if 'fixture_id' in case else Path(report['evidence_root'])/'prospective-advarra-rich-complete'
    revision=root/'revisions'/case['result']['revision_id']
    manifest=json.loads((revision/'delivery-manifest.json').read_text())
    evidence=manifest['quality']
    assert evidence['status']=='passed'
    assert any(item['check']=='editorial_relevance' for item in evidence['warnings'])
    assert (root/'output/protocol.docx').is_file()
    for item in manifest['client_outputs']:
        assert (root/item['path']).read_bytes()==(revision/'candidate'/Path(item['path']).name).read_bytes()


def test_rendered_methodology_does_not_reacquire_endpoint_inventory_obligations(tmp_path):
    from docx import Document
    import quality
    source=copy.deepcopy(REPLAY['source'])
    source['statistics']['methodology']='Descriptive summaries use mean, standard deviation and median.'
    source['endpoints']['other']=[{'name':'Adherence','time_point':'Month 6','description':'Device use days.'}]
    doc=Document();doc.add_heading('10. ANALYSIS PLAN',level=1)
    doc.add_heading('10.2. Statistical Methodology',level=2)
    doc.add_paragraph('Descriptive summaries use mean, standard deviation and median.')
    (tmp_path/'candidate').mkdir();doc.save(tmp_path/'candidate/protocol.docx');Document().save(tmp_path/'candidate/icf.docx')
    findings=quality.deterministic_content_check(tmp_path,source)
    assert not [f for f in findings if f.get('field')=='analysis-plan.methodology' and 'endpoints.other' in f['issue']]


def test_source_validation_exposes_computed_fields_before_drafting():
    import contracts
    result=contracts.source_contract(REPLAY['source'])
    assert result['status']=='passed'
    assert result['computed_fields']['participant_followup']=='6 months after second-eye surgery'
    assert result['computed_fields']['visit_count']==5
    assert result['computed_fields']['assessment_table']['rows'][0][-1]=='Month 6'


def test_other_endpoint_inventory_is_checked_in_its_owner_not_methodology(tmp_path):
    from docx import Document
    import quality
    source=copy.deepcopy(REPLAY['source'])
    source['endpoints']['other']=[{'name':'Adherence','time_point':'Month 6','description':'Device use days.'}]
    doc=Document();doc.add_heading('8. STUDY DESIGN',level=1)
    doc.add_heading('8.1. Study Design',level=2)
    doc.add_paragraph('This descriptive study compares two lens cohorts.')
    (tmp_path/'candidate').mkdir();doc.save(tmp_path/'candidate/protocol.docx');Document().save(tmp_path/'candidate/icf.docx')
    findings=quality.deterministic_content_check(tmp_path,source)
    assert any(f.get('field')=='study-design.design' and 'endpoints.other' in f['issue'] for f in findings)


def test_followup_fallback_preserves_second_eye_anchor():
    source = copy.deepcopy(REPLAY['source'])
    source['study']['timeline'] = 'Study lasts 12 months.'
    source['procedures']['visit_schedule'] = []
    source['procedures']['visit_schedule_table'] = [
        {'visitNumber': '1', 'visitName': 'Month 6 postoperative assessment',
         'timing': '6 months after second-eye surgery', 'procedures': []},
    ]
    assert rendering.render_fields(source, {})['AI_duration'] == '6 months after second-eye surgery'


def test_followup_phase_preserves_intervening_qualifications():
    source = copy.deepcopy(REPLAY['source'])
    source['study']['timeline'] = (
        'Follow-up: 6 months after second-eye surgery; visits may occur within a 2-week window; analysis: 2 months.'
    )
    assert rendering.render_fields(source, {})['AI_duration'] == (
        '6 months after second-eye surgery; visits may occur within a 2-week window'
    )
