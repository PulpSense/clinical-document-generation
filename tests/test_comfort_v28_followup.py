"""Regression coverage for the user's saved v28 COMFORT run."""
import copy
import json
from pathlib import Path
from xml.etree import ElementTree as ET

import pytest
from docx import Document
from docx.oxml.ns import qn

import contracts
import drafting
import prs_xml
import rendering

ROOT = Path(__file__).resolve().parents[1]
TEMPLATE = ROOT / 'assets/client-templates/prs/clinicaltrials_prs_full_placeholder_template.xml'


def source():
    return json.loads((ROOT / 'tests/fixtures/comfort-hermes-v25.json').read_text())['source']


@pytest.mark.parametrize('explicit_role', [None, 'Study Chair'])
def test_xml_uses_supplied_title_when_role_is_absent_and_preserves_explicit_role(explicit_role, tmp_path):
    ref = source()
    if explicit_role:
        ref['sites'][1]['investigator']['role'] = explicit_role
    original = copy.deepcopy(ref)
    output = tmp_path / 'study.xml'
    prs_xml.generate(TEMPLATE, output, ref, {'brief_summary': 'Summary.', 'detailed_description': 'Details.'}, check_output=False)
    study = next(ET.parse(output).getroot().iter('clinical_study'))
    assert [n.findtext('investigator/role') for n in study.findall('location')] == [
        'Principal Investigator', explicit_role or 'Site Investigator', 'Site Investigator']
    assert ref == original


@pytest.mark.parametrize('family', ['Sterling', 'Advarra'])
def test_generated_risk_lists_keep_native_bullets_without_annotation_emphasis(family, tmp_path):
    ref = source()
    ref['meta']['icf_template'] = family
    risk = json.loads((ROOT / 'tests/fixtures/comfort-hermes-v28-icf-risks.json').read_text())['section']
    model = {'protocol': [], 'prs': {}, 'icf': {'icf.risks': risk}}
    rendering.render_documents(ROOT, tmp_path, ref, model, artifact_names={'icf'}, audit_output=False)
    doc = Document(tmp_path / 'candidate/icf.docx')
    items = risk['lists'][0]['items']
    paragraphs = [p for p in doc.paragraphs if p.text in items]
    assert len(paragraphs) == 5
    for p in paragraphs:
        assert p._p.pPr.find(qn('w:numPr')) is not None
        assert all(not r.bold and not r.underline and r.font.highlight_color is None for r in p.runs)
        assert not p._p.xpath('.//w:highlight[@w:val="yellow"]')


@pytest.mark.parametrize('branch,owner', [('Prospective', 'study-procedure.measurements'), ('Ambispective', 'study-procedure.measurements'), ('Retrospective', 'study-procedure.enrollment'), ('Ambispective', 'icf.procedures')])
def test_drafting_handoff_assigns_supplied_specimen_and_results_policies(branch, owner, tmp_path):
    ref = source()
    ref['meta']['study_type'] = branch
    original = copy.deepcopy(ref)
    revision = tmp_path / 'revision'
    revision.mkdir()
    (revision / 'generation-only.json').write_text('{}')
    batch = next(b for b in contracts.batch_plan(branch, ref['meta']['icf_template']) if owner in b.section_ids)
    path = drafting.create_drafting_request(repo_root=ROOT, revision_dir=revision, revision_id='test-revision', reference=ref, batch=batch,
        target_ids=[owner], attempts={owner: 1}, wave='generation-only')
    request = json.loads(path.read_text())
    section = request['section_contracts'][0]
    for field in ['procedures.specimens', 'procedures.return_of_results']:
        assert field in section['minimum_evidence']
        assert request['evidence_checklist'][0]['approved_values'][field] == contracts.get_path(ref, field)
        destinations = next(e for e in request['source_evidence_coverage_map']['entries'] if e['source_path'] == field)['destinations']
        assert any(d['artifact'] == ('icf.docx' if owner.startswith('icf.') else 'protocol.docx')
                   and d['section_id'] == owner for d in destinations)
    assert ref == original


def test_triggered_gina_clause_keeps_its_heading_and_authorized_paragraphs_together(tmp_path):
    ref = source()
    rendering.render_documents(ROOT, tmp_path, ref, {'protocol': [], 'icf': {}, 'prs': {}}, artifact_names={'icf'}, audit_output=False)
    doc = Document(tmp_path / 'candidate/icf.docx')
    index = next(i for i,p in enumerate(doc.paragraphs) if p.text == 'GENETIC INFORMATION NONDISCRIMINATION ACT')
    clause = next(c for c in contracts.sterling_clause_contract()['clauses'] if c['clause_id'] == 'sterling.privacy.gina')
    block = doc.paragraphs[index:index+1+len(clause['authorized_texts'])]
    assert [p.text for p in block[1:]] == clause['authorized_texts']
    assert all(p.paragraph_format.keep_with_next for p in block[:-1])
    assert not block[-1].paragraph_format.keep_with_next
    assert all(p.paragraph_format.keep_together for p in block)


def test_absent_optional_policies_stay_absent_in_the_drafting_handoff(tmp_path):
    ref = source()
    ref['procedures'].pop('specimens')
    ref['procedures'].pop('return_of_results')
    original = copy.deepcopy(ref)
    owner = 'study-procedure.measurements'
    revision = tmp_path / 'revision'
    revision.mkdir()
    (revision / 'generation-only.json').write_text('{}')
    batch = next(b for b in contracts.batch_plan('Ambispective', 'Sterling') if owner in b.section_ids)
    request_path = drafting.create_drafting_request(repo_root=ROOT, revision_dir=revision,
        revision_id='test-revision', reference=ref, batch=batch, target_ids=[owner],
        attempts={owner: 1}, wave='generation-only')
    request = json.loads(request_path.read_text())
    for field in ['procedures.specimens', 'procedures.return_of_results']:
        assert request['evidence_checklist'][0]['approved_values'][field] is None
        assert contracts.get_path(request['approved_source'], field) is None
    assert ref == original
