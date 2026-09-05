import json
import hashlib
from pathlib import Path

from docx import Document

from contracts import contracted_template_bundle
import quality
from quality import CONTENT_CHECKS, RESPONSE_SCHEMA, VISUAL_CHECKS
from rendering import template_paths
from workflow import run_release_gate
import workflow


ROOT = Path(__file__).resolve().parents[1]


def _require_renderer(governed_pdfium):
    assert workflow.renderer() is not None
    assert governed_pdfium["source"] == "release-owned runtime"


def acceptance_verification(request):
    response = {
        "schema_version": RESPONSE_SCHEMA,
        "request_id": request["request_id"],
        "request_sha256": request["request_sha256"],
        "task": request["task"],
        "producer": {
            "model_id": "TestAcceptanceVerifier/v1",
            "reviewer_id": request["task"],
        },
        "status": "passed",
        "findings": [],
    }
    if request["task"] == "rendered_page_visual_verification":
        response["page_assessments"] = [
            {"artifact": artifact["artifact"], "page": page["page"], "sha256": page["sha256"], "status": "passed", "checks": list(VISUAL_CHECKS)}
            for artifact in request.get("artifacts", []) for page in artifact.get("pages", [])
        ]
    else:
        response["section_assessments"] = [
            {"artifact": item["artifact"], "section_id": item["section_id"], "status": "passed", "checks": list(CONTENT_CHECKS)}
            for item in request.get("sections", [])
        ]
        response["cross_document_assessments"] = [{"check": check, "status": "passed"} for check in request.get("cross_document_checks", [])]
    return response


acceptance_verification.synthetic = True


def acceptance_repair_verification(request):
    response = acceptance_verification(request)
    if (
        request["task"] == "rendered_page_visual_verification"
        and ".review-1." in request["request_id"]
        and any(
            artifact.get("artifact") == "protocol"
            for artifact in request.get("artifacts", [])
        )
    ):
        study_type = request["contracted_template_bundle"]["selection"]["study_type"]
        check = "orphan_heading" if study_type == "Retrospective" else "bad_table_split"
        element = "4. INTRODUCTION" if study_type == "Retrospective" else "3. GENERAL INFORMATION"
        response["status"] = "blocked"
        response["findings"] = [{
            "artifact": "protocol",
            "check": check,
            "element": element,
            "issue": (
                "Synthetic corpus fault: the semantic block requires a "
                "governed pagination repair."
            ),
        }]
        protocol_page = next(
            page
            for page in response["page_assessments"]
            if page["artifact"] == "protocol"
        )
        protocol_page["status"] = "blocked"
    return response


acceptance_repair_verification.synthetic = True


def novastep_regression_verification(request):
    response = acceptance_verification(request)
    if (
        request["task"] == "rendered_page_visual_verification"
        and ".review-1." in request["request_id"]
        and any(
            artifact.get("artifact") == "protocol"
            for artifact in request.get("artifacts", [])
        )
    ):
        shared = {
            "artifact": "protocol",
            "element": "3. GENERAL INFORMATION",
        }
        response["status"] = "blocked"
        response["findings"] = [
            {
                **shared,
                "page": 2,
                "check": "bad_table_split",
                "element": "3. GENERAL INFORMATION – Variables / Secondary endpoint(s)",
                "issue": "The endpoint label is separated from its first bullet.",
            },
            {
                **shared,
                "page": 3,
                "check": "excessive_whitespace",
                "issue": "The table continuation leaves nearly the entire page unused.",
            },
            {
                **shared,
                "page": 3,
                "check": "artificial_pagination",
                "issue": "The continuation creates a standalone page before the TOC.",
            },
        ]
        for page in response["page_assessments"]:
            if page["artifact"] == "protocol" and page["page"] in {2, 3}:
                page["status"] = "blocked"
    return response


novastep_regression_verification.synthetic = True


def external_verification(request):
    return acceptance_verification(request)


def _section_geometry(document: Document) -> tuple[tuple[int | None, ...], ...]:
    return tuple((
        section.page_width, section.page_height, section.top_margin, section.right_margin,
        section.bottom_margin, section.left_margin, section.header_distance, section.footer_distance,
    ) for section in document.sections)


def test_branch_acceptance_corpus_public_lifecycles_pass_and_publish_exact_sets(monkeypatch, governed_pdfium):
    _require_renderer(governed_pdfium)
    report = run_release_gate(ROOT, verification_responder=acceptance_repair_verification)
    assert report["status"] == "structural_passed"
    assert report["assurance"] == "synthetic-structural-only"
    assert len(report["cases"]) == 10
    assert tuple(case["case"] for case in report["cases"]) == (
        quality.DETERMINISTIC_BRANCH_ACCEPTANCE_CASES
    )
    assert report["distinct_source_count"] == 10
    assert {case["icf_template"] for case in report["cases"] if case["icf_template"]} == {"Advarra", "Sterling"}
    for case in report["cases"]:
        assert case["status"] == "passed"
        descriptor_path = (
            ROOT
            / "references/conformance-fixtures/branch-acceptance-corpus"
            / f"{case['case']}.json"
        )
        assert case["descriptor_sha256"] == hashlib.sha256(
            descriptor_path.read_bytes()
        ).hexdigest()
        assert case["corpus"] == json.loads(descriptor_path.read_text(encoding="utf-8"))
        source_fixture = ROOT / case["corpus"]["source_fixture"]
        assert case["source_fixture_sha256"] == hashlib.sha256(
            source_fixture.read_bytes()
        ).hexdigest()
        outputs = case["result"]["client_outputs"]
        case_root = Path(report["evidence_root"]) / case["case"]
        reference = json.loads((case_root / "reference/study.reference.json").read_text(encoding="utf-8"))
        expected_bundle = contracted_template_bundle(ROOT, reference)
        revision_dir = case_root / "revisions" / case["result"]["revision_id"]
        manifest = json.loads((case_root / case["result"]["manifest"]).read_text(encoding="utf-8"))
        build = json.loads((revision_dir / "candidate-build.json").read_text(encoding="utf-8"))
        assert reference["generation"]["review_set"] == 2
        is_retrospective = reference["meta"]["study_type"] == "Retrospective"
        repair = (
            {"rule": "heading_cohesion", "target": "4. INTRODUCTION"}
            if is_retrospective
            else {"rule": "table_pagination", "target": "3. GENERAL INFORMATION"}
        )
        assert reference["generation"]["layout_repairs"]["protocol"] == [
            repair
        ]
        review_two_responses = revision_dir / "hermes/verification-responses"
        expected_review_two = {
            f"{case['result']['revision_id']}.review-2.verify.content.json",
            f"{case['result']['revision_id']}.review-2.verify.visual.protocol.json",
        }
        if reference["meta"]["study_type"] != "Retrospective":
            expected_review_two.add(
                f"{case['result']['revision_id']}.review-2.verify.visual.icf.json"
            )
        assert expected_review_two <= {
            path.name for path in review_two_responses.glob("*.json")
        }
        assert [record["terminal_status"] for record in manifest["gate_ledger"]["records"]] == [
            "passed", "passed", "passed", "passed", "passed", "pending",
        ]
        assert quality.validate_gate_ledger(ROOT, manifest["gate_ledger"]) == manifest["gate_ledger"]
        verification_requests = [
            json.loads(path.read_text(encoding="utf-8"))
            for path in (revision_dir / "hermes/verification-requests").glob("*.json")
        ]
        assert case["contracted_template_bundle"] == expected_bundle
        assert build["contracted_template_bundle"] == expected_bundle
        assert manifest["contracted_template_bundle"] == expected_bundle
        assert all(request["contracted_template_bundle"] == expected_bundle for request in verification_requests)
        protocol_template, icf_template = template_paths(ROOT, reference)
        assert _section_geometry(Document(case_root / "output/protocol.docx")) == _section_geometry(Document(protocol_template))
        if case["case"].startswith("retrospective"):
            assert outputs == ["output/protocol.docx"]
            visible = "\n".join(paragraph.text for paragraph in Document(Path(report["evidence_root"]) / case["case"] / "output/protocol.docx").paragraphs)
            assert "Unable to complete follow-up" in visible
            if "rich" in case["case"]:
                assert "Month 6" in visible
                assert "6 months" in visible
        else:
            assert outputs == ["output/icf.docx", "output/protocol.docx", "output/study.xml"]
            assert icf_template is not None
            icf_output = Document(case_root / "output/icf.docx")
            assert _section_geometry(icf_output) == _section_geometry(Document(icf_template))
            visible = "\n".join(paragraph.text for paragraph in icf_output.paragraphs).casefold()
            assert any(phrase in visible for phrase in (
                "should not sign",
                "if you would like to participate, you will be asked to sign",
                "if you agree to participate, you will be asked to sign",
            ))
    assert {
        bundle["identity_sha256"]
        for bundle in report["contracted_template_bundles"]
    } == {
        case["contracted_template_bundle"]["identity_sha256"]
        for case in report["cases"]
    }


def test_novastep_regression_completes_public_repair_review_and_atomic_publication(governed_pdfium):
    _require_renderer(governed_pdfium)
    case_id = "prospective-advarra-rich-complete"
    report = run_release_gate(
        ROOT,
        verification_responder=novastep_regression_verification,
        case_ids=(case_id,),
    )

    assert report["status"] == "structural_passed", report
    case = report["cases"][0]
    case_root = Path(report["evidence_root"]) / case_id
    reference = json.loads(
        (case_root / "reference/study.reference.json").read_text(encoding="utf-8")
    )
    assert reference["study"]["title"] == (
        "Prospective Evaluation of the NovaStep Activity Sensor in Adults "
        "Recovering From Total Knee Arthroplasty"
    )
    assert reference["generation"]["review_set"] == 2
    assert reference["generation"]["layout_repairs"]["protocol"] == [
        {"rule": "heading_cohesion", "target": "3. GENERAL INFORMATION"},
        {"rule": "table_pagination", "target": "3. GENERAL INFORMATION"},
    ]
    revision = case_root / "revisions" / case["result"]["revision_id"]
    failed_attempt = revision / "attempts/quality-a01"
    assert quality.sha256_file(failed_attempt / "candidate/protocol.docx") != (
        quality.sha256_file(revision / "candidate/protocol.docx")
    )
    assert quality.sha256_file(failed_attempt / "rendered/protocol.pdf") != (
        quality.sha256_file(revision / "rendered/protocol.pdf")
    )
    responses = revision / "hermes/verification-responses"
    response_paths = sorted(responses.glob("*.review-2.verify.*.json"))
    assert len(response_paths) == 3
    for path in response_paths:
        response = json.loads(path.read_text(encoding="utf-8"))
        assert response["status"] == "passed"
        assert response["request_id"].endswith(path.name.split(".review-2.", 1)[1].removesuffix(".json"))
        if response["task"] == "rendered_page_visual_verification":
            assert response["page_assessments"]
            assert all(item["status"] == "passed" for item in response["page_assessments"])
    manifest = json.loads(
        (revision / "delivery-manifest.json").read_text(encoding="utf-8")
    )
    assert [record["terminal_status"] for record in manifest["gate_ledger"]["records"]] == [
        "passed", "passed", "passed", "passed", "passed", "pending",
    ]
    assert case["result"]["client_outputs"] == [
        "output/icf.docx", "output/protocol.docx", "output/study.xml",
    ]
    for item in manifest["client_outputs"]:
        output_path = case_root / item["path"]
        candidate_path = revision / "candidate" / output_path.name
        assert output_path.read_bytes() == candidate_path.read_bytes()
        assert quality.sha256_file(output_path) == item["sha256"]


def test_recorded_drafting_keeps_release_gate_assurance_structural_with_external_verification(monkeypatch, governed_pdfium):
    _require_renderer(governed_pdfium)
    report = run_release_gate(ROOT, verification_responder=external_verification)

    assert report["status"] == "structural_passed"
    assert report["assurance"] == "recorded-drafting-structural-only"


def test_controlled_release_adapter_drives_the_complete_desktop_operation(tmp_path, monkeypatch, governed_pdfium):
    _require_renderer(governed_pdfium)
    run_dir = tmp_path / "controlled-retrospective"
    reference_path = run_dir / "reference/study.reference.json"
    reference_path.parent.mkdir(parents=True)
    reference = json.loads((ROOT / "tests/fixtures/retrospective-acceptance-source.json").read_text(encoding="utf-8"))
    reference["meta"]["protocol_number"] = "RET-CONTROLLED-41"
    reference["study"]["title"] = "Controlled Desktop Operation study"
    reference["approval"] = {"status": "draft"}
    reference_path.write_text(json.dumps(reference), encoding="utf-8")
    assert workflow.prepare(run_dir)["status"] == "awaiting_approval"
    assert workflow.approve(run_dir, approved_by="Controlled Release Adapter")["status"] == "passed"

    def handoff_runner(handoffs, _remaining_seconds):
        approved = json.loads(reference_path.read_text(encoding="utf-8"))
        revision_dir = run_dir / "revisions" / approved["approval"]["revision_id"]
        for handoff in handoffs:
            request_path = revision_dir / handoff["request_path"]
            if "verification-requests" in handoff["request_path"]:
                workflow._save_verification_response(revision_dir, request_path, acceptance_verification)
            else:
                workflow._save_recorded_handoff(revision_dir, request_path)

    result = workflow.run_desktop_operation(
        run_dir,
        handoff_runner=handoff_runner,
        opener=lambda path: Path(path).read_bytes(),
        operation_id="controlled-release-adapter",
        release_identity={"package_fingerprint": "controlled-candidate-41"},
    )

    assert result["status"] == "passed"
    assert result["stage"] == "desktop_delivery"
    assert result["client_outputs"] == ["output/protocol.docx"]
    assert result["delivery"]["confirmed"] is True
    state = json.loads((run_dir / "logs/desktop-operation-controlled-release-adapter.json").read_text())
    assert state["release_identity"]["package_fingerprint"] == "controlled-candidate-41"
    assert {"drafting", "candidate", "render_assurance", "independent_verification", "delivery"} <= set(state["stage_timings"])
    assert state["result"] == result
