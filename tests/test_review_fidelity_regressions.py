"""Review regressions at canonical draft ingestion and rendered assurance boundaries."""
import json
from pathlib import Path

import pytest
import drafting
from contracts import batch_plan
from finding_policy import blocks_construction

ROOT = Path(__file__).resolve().parents[1]


def section_response(tmp_path, section_id, text, *, source_path, source_value, family="Advarra", extra_source=None):
    source = json.loads((ROOT / "tests/fixtures/prospective-acceptance-source.json").read_text())
    source["meta"]["icf_template"] = family
    parent, key = source_path.split(".")
    source[parent][key] = source_value
    for path, value in (extra_source or {}).items():
        parent, key = path.split(".")
        source[parent][key] = value
    batch = next(b for b in batch_plan("Prospective", family) if section_id in b.section_ids)
    path = drafting.create_drafting_request(repo_root=ROOT, revision_dir=tmp_path, revision_id="review-regression",
        reference=source, batch=batch, target_ids=(section_id,), attempts={section_id: 1}, wave="initial")
    request = json.loads(path.read_text())
    response = drafting.recorded_acceptance_response(request)
    if batch.artifact == "prs":
        key = "brief_summary" if section_id == "prs.brief-summary" else "detailed_description"
        response["narrative"][key] = {"text": text, "evidence_refs": [f"source:{source_path}"]}
    else:
        response["section_results"][0].update(paragraphs=[{"text": text,
            "evidence_refs": [f"source:{source_path}"], "boilerplate_refs": []}], lists=[])
    return request, response


def test_faithful_four_word_consent_duration_is_retained_for_review(tmp_path):
    request, response = section_response(tmp_path, "icf.duration", "Follow-up lasts three months.",
        source_path="study.timeline", source_value="Follow-up: 3 months.")
    accepted, findings = drafting.validate_response(request, response)
    assert [s["section_id"] for s in accepted["drafts"]] == ["icf.duration"]
    assert findings and all(not blocks_construction(f) for f in findings)


def test_faithful_concise_prs_summary_survives_ingestion(tmp_path):
    text = "Study of cataract outcomes."
    paths = ["study.title", "study.hypothesis", "objectives.primary", "endpoints.primary", "design.study_design"]
    request, response = section_response(tmp_path, "prs.brief-summary", text,
        source_path="study.title", source_value=text, extra_source={p: text for p in paths})
    response["narrative"]["brief_summary"]["evidence_refs"] = [f"source:{p}" for p in paths]
    accepted, findings = drafting.validate_response(request, response)
    assert accepted and accepted["narrative"]["brief_summary"]["text"] == text
    assert findings and all(not blocks_construction(f) for f in findings)


@pytest.mark.parametrize("section,path,source,text", [
    ("endpoint-criteria.completion", "procedures.completion", "Participant completion requires attendance at the final assessment.",
     "Subjects finish participation once they attend the last evaluation."),
    ("endpoint-criteria.study-completion", "study.completion", "Study closeout requires verification of all case records.",
     "The project ends after every participant file has been checked."),
    ("icf.duration", "study.timeline", "Follow-up: 3 months.", "Your participation lasts three months."),
])
def test_specialized_source_paraphrases_reach_independent_review(tmp_path, section, path, source, text):
    request, response = section_response(tmp_path, section, text, source_path=path, source_value=source)
    # Other evidence remains supplied by the recorded response when required.
    row = response["section_results"][0]
    for p in request["section_contracts"][0]["minimum_evidence"]:
        if f"source:{p}" not in row["paragraphs"][0]["evidence_refs"]:
            row["paragraphs"][0]["evidence_refs"].append(f"source:{p}")
            value = drafting.get_path(request["approved_source"], p)
            if value:
                row["paragraphs"][0]["text"] += " " + str(value)
    accepted, findings = drafting.validate_response(request, response)
    assert section in [s["section_id"] for s in (accepted or {}).get("drafts", [])], findings
    assert findings and all(not blocks_construction(f) for f in findings)


def test_retrospective_eligibility_paraphrase_needs_meaning_review_not_reconstruction(tmp_path):
    import quality
    from rendering import render_documents
    from docx import Document
    source = json.loads((ROOT / "tests/fixtures/retrospective-acceptance-source.json").read_text())
    source["population"]["inclusion_criteria"] = ["Records describe adults who underwent cataract extraction."]
    render_documents(ROOT, tmp_path, source, {"protocol": [], "icf": {}, "prs": {}})
    path = tmp_path / "candidate/protocol.docx"
    doc = Document(path)
    for paragraph in doc.paragraphs:
        if "Records describe adults who underwent cataract extraction." in paragraph.text:
            paragraph.text = paragraph.text.replace("Records describe adults who underwent cataract extraction.",
                "Files concern grown patients treated by removing a clouded lens.")
    doc.save(path)
    findings = quality.deterministic_content_check(tmp_path, source)
    lexical = [f for f in findings if f.get("code") == "retrospective-lexical-evidence-uncertainty"]
    assert lexical, findings
    assert all(f["finding_policy"]["classification"] == "semantic_uncertainty" and
               f["finding_policy"]["repair_owner"] == "independent_review" for f in lexical)


def test_actual_textless_page_keeps_rendering_owner_and_can_be_reassessed(tmp_path, governed_pdfium):
    import quality
    from docx import Document
    from pypdf import PdfReader, PdfWriter
    from finding_policy import normalize_finding
    from test_render_assurance import _bundle, _structural_validation, _write_text_pdf
    candidate = tmp_path / "candidate"
    candidate.mkdir()
    doc = Document()
    doc.add_paragraph("Complete candidate with a controlled external renderer.")
    doc.save(candidate / "protocol.docx")
    original = (candidate / "protocol.docx").read_bytes()
    blank = True

    def export(docx, output_dir, _identity, **_kwargs):
        pdf = output_dir / f"{docx.stem}.pdf"
        _write_text_pdf(pdf)
        if blank:
            writer = PdfWriter()
            writer.add_page(PdfReader(pdf).pages[0])
            writer.add_blank_page(width=612, height=792)
            with pdf.open("wb") as handle:
                writer.write(handle)
        return pdf

    def assess():
        return quality.render_assurance(ROOT, tmp_path, {}, contracted_bundle=_bundle(),
            structural_validation=_structural_validation(tmp_path, "protocol.docx"),
            renderer_identities=[{"kind": "LibreOffice", "path": "/controlled/soffice"}],
            page_renderer_identities=[governed_pdfium], font_probe=lambda *_a, **_kw: (True, "controlled font inventory"),
            office_exporter=export)

    report = assess()
    finding = next(f for f in report["render"]["findings"] if "textless page" in f["issue"])
    governed = normalize_finding(finding, producer="render-assurance")
    assert governed["finding_policy"]["classification"] == "essential_defect"
    assert governed["finding_policy"]["repair_owner"] == "rendering"
    assert governed["recovery_class"] == "visual_defect"
    assert governed["artifact"] == "protocol" and governed["page"] == 2
    assert len(report["render"]["artifacts"][0]["pages"]) == 2
    assert (candidate / "protocol.docx").read_bytes() == original
    blank = False
    repaired = assess()
    assert repaired["status"] == "passed", repaired
    assert len(repaired["render"]["artifacts"][0]["pages"]) == 1


def test_concise_prs_caution_reaches_exact_independent_registry_review(tmp_path):
    import quality
    text = "Study of cataract outcomes."
    paths = ["study.title", "study.hypothesis", "objectives.primary", "endpoints.primary", "design.study_design"]
    request, response = section_response(tmp_path, "prs.brief-summary", text,
        source_path="study.title", source_value=text, extra_source={p: text for p in paths})
    response["narrative"]["brief_summary"]["evidence_refs"] = [f"source:{p}" for p in paths]
    path = tmp_path / request["response_path"]
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(response))
    assert drafting.ingest_responses(tmp_path, request["governing_resources"]) == []
    source = json.loads((tmp_path / "approved-reference.json").read_text())
    review_path = quality.create_verification_requests(tmp_path, source, {"artifacts": []})[0]
    review = json.loads(review_path.read_text())
    cautions = [c for c in review["lexical_review_cautions"] if c["section_id"] == "prs.brief-summary"]
    assert cautions and all(c["artifact"] == "study.xml" for c in cautions)


@pytest.mark.parametrize("section,path,source,text", [
    ("icf.duration", "study.timeline", "Follow-up: 3 months.", "Follow-up lasts nine months."),
    ("icf.duration", "study.timeline", "Enrollment: 6 months; Follow-up: 3 months.", "Enrollment lasts three months; follow-up lasts six months."),
    ("endpoint-criteria.completion", "procedures.completion", "Completion requires attendance at 3 assessments.", "Completion requires attendance at nine assessments."),
    ("endpoint-criteria.study-completion", "study.completion", "Closeout requires 3 verified records.", "Closeout requires nine verified records."),
    ("icf.duration", "study.timeline", "Follow-up: 3 months.", ""),
])
def test_missing_or_wrong_source_quantities_remain_essential(tmp_path, section, path, source, text):
    request, response = section_response(tmp_path, section, text, source_path=path, source_value=source)
    accepted, findings = drafting.validate_response(request, response)
    assert section not in [s["section_id"] for s in (accepted or {}).get("drafts", [])]
    assert any(blocks_construction(f) for f in findings)


@pytest.mark.parametrize("text", ["", "<summary>Study of cataract outcomes.</summary>"])
def test_empty_or_markup_prs_is_essential(tmp_path, text):
    request, response = section_response(tmp_path, "prs.brief-summary", text,
        source_path="study.title", source_value="Study of cataract outcomes.")
    accepted, findings = drafting.validate_response(request, response)
    assert not accepted
    assert any(blocks_construction(f) and "empty or contains XML markup" in f["issue"] for f in findings)


@pytest.mark.parametrize("refs", [[], ["source:study.title"]])
def test_concise_consent_still_requires_its_approved_evidence(tmp_path, refs):
    request, response = section_response(tmp_path, "icf.duration", "Follow-up lasts three months.",
        source_path="study.timeline", source_value="Follow-up: 3 months.")
    response["section_results"][0]["paragraphs"][0]["evidence_refs"] = refs
    accepted, findings = drafting.validate_response(request, response)
    assert not accepted["drafts"] and any(blocks_construction(f) for f in findings)


def test_retrospective_missing_eligibility_quantity_keeps_narrative_repair_owner(tmp_path):
    import quality
    from rendering import render_documents
    from docx import Document
    source = json.loads((ROOT / "tests/fixtures/retrospective-acceptance-source.json").read_text())
    source["population"]["inclusion_criteria"] = ["Records of adults aged at least 73 years."]
    render_documents(ROOT, tmp_path, source, {"protocol": [], "icf": {}, "prs": {}})
    path = tmp_path / "candidate/protocol.docx"
    doc = Document(path)
    for p in doc.paragraphs:
        p.text = p.text.replace("73 years", "99 years")
    doc.save(path)
    findings = quality.deterministic_content_check(tmp_path, source)
    quantity = next(f for f in findings if f.get("field") == "subjects.eligibility" and f.get("missing_numbers") == ["73"])
    assert quantity["finding_policy"]["classification"] == "essential_defect"
    assert quantity["finding_policy"]["repair_owner"] == "drafting"


def test_consent_cannot_reassign_followup_duration_to_analysis(tmp_path):
    request, response = section_response(tmp_path, "icf.duration", "Analysis lasts three months.",
        source_path="study.timeline", source_value="Follow-up: 3 months.")
    accepted, findings = drafting.validate_response(request, response)
    assert not accepted["drafts"]
    assert any(blocks_construction(f) and f.get("evidence_diagnostics") for f in findings)


def test_placeholder_prs_text_is_not_a_concise_narrative(tmp_path):
    paths = ["study.title", "study.hypothesis", "objectives.primary", "endpoints.primary", "design.study_design"]
    request, response = section_response(tmp_path, "prs.brief-summary", "TBD",
        source_path="study.title", source_value="Cataract study", extra_source={p: "Cataract study" for p in paths})
    response["narrative"]["brief_summary"]["evidence_refs"] = [f"source:{p}" for p in paths]
    accepted, findings = drafting.validate_response(request, response)
    assert not accepted
    assert any(blocks_construction(f) and "placeholder" in f["issue"] for f in findings)
