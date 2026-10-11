"""Public lifecycle evidence using actual native editing and controlled AI/render adapters."""
import json
from pathlib import Path
from docx import Document
import drafting
import quality
import workflow
import pytest
from test_clean_delivery import controlled_render, ROOT, CASES
from test_handoff_quality import acceptance_verification


def lifecycle(tmp_path, monkeypatch, repair, *, retrospective=True, review=None, render=None, case=None, draft=None, visual_review=None):
    monkeypatch.setattr(workflow, 'render_assurance', render or controlled_render)
    source = quality.branch_acceptance_source(ROOT, case or next(case for case in CASES
        if (case['descriptor']['study_type'] == 'Retrospective') == retrospective))
    reference = tmp_path / 'reference/study.reference.json'
    reference.parent.mkdir(parents=True)
    reference.write_text(json.dumps(source))
    workflow.prepare(tmp_path)
    approved = workflow.approve(tmp_path, approved_by='fixture-owner')
    revision = tmp_path / 'revisions' / approved['revision_id']
    approved_bytes = (revision / 'approved-reference.json').read_bytes()
    content_reviews = []
    repair_requests = []
    result = workflow.generate(tmp_path, require_promoted_runtime=False)
    for _ in range(14):
        if result['status'] != 'awaiting_hermes':
            break
        for relative in result['requests']:
            request = json.loads((revision / relative).read_text())
            if request['task'] == 'document_repair':
                repair_requests.append(request)
                response = repair(request, revision)
            elif request['task'] in {'section_drafting', 'prs_narrative_drafting'}:
                response = (draft or drafting.recorded_acceptance_response)(request)
            else:
                response = acceptance_verification(request)
                if request['task'] == 'clinical_content_verification':
                    content_reviews.append(request)
                    if review:
                        response = review(request, response, revision, len(content_reviews))
                    elif len(content_reviews) == 1:
                        response['status'] = 'blocked'
                        response['findings'] = [{'finding_id': 'native-content', 'category': 'content',
                            'artifact': 'protocol', 'target_ids': ['introduction'], 'check': 'source_supported',
                            'issue': 'Correct the native content using approved source.', 'material': True}]
                        next(row for row in response['section_assessments'] if row['section_id'] == 'introduction')['status'] = 'blocked'
                    for row in response['section_assessments']:
                        row['notes'] = ['Independently reassessed native-content against the approved source.']
                if request['task'] == 'rendered_page_visual_verification' and visual_review:
                    response = visual_review(request, response, revision)
            path = revision / request['response_path']
            path.parent.mkdir(parents=True, exist_ok=True)
            path.write_text(json.dumps(response))
        result = workflow.generate(tmp_path, require_promoted_runtime=False)
    assert (revision / 'approved-reference.json').read_bytes() == approved_bytes
    return result, revision, repair_requests, content_reviews


def repair_response(request, **changes):
    return {key: request[key] for key in ('request_id', 'request_sha256', 'revision_id', 'task', 'batch_id')} | {
        'schema_version': 'hermes-response/v2', 'status': 'completed',
        'producer': {'model_id': 'controlled-editor', 'reviewer_id': 'repair-author'},
        'summary': 'Corrected the supplied concerns with native document tools.', **changes}


def test_ai_can_edit_native_paragraphs_tables_and_headers_without_registered_surfaces(tmp_path, monkeypatch):
    def repair(request, revision):
        path = revision / request['workspace_path'] / 'protocol.docx'
        doc = Document(path)
        doc.add_paragraph('Native correction retained for independent review.')
        doc.tables[0].cell(0, 0).text = 'Corrected visit association'
        doc.sections[0].header.paragraphs[0].text = 'Corrected running header'
        doc.sections[0].footer.paragraphs[0].text = 'Corrected running footer'
        doc.save(path)
        return repair_response(request)
    result, revision, repairs, reviews = lifecycle(tmp_path, monkeypatch, repair)
    assert result['status'] == 'passed', result
    assert len(repairs) == 1 and len(reviews) == 2
    doc = Document(tmp_path / 'output/protocol.docx')
    assert doc.paragraphs[-1].text == 'Native correction retained for independent review.'
    assert doc.tables[0].cell(0, 0).text == 'Corrected visit association'
    assert doc.sections[0].header.paragraphs[0].text == 'Corrected running header'
    assert doc.sections[0].footer.paragraphs[0].text == 'Corrected running footer'
    assert reviews[-1]['review_context']['prior_findings'][0]['finding_id'] == 'native-content'
    assert repairs[0]['approved_source'] == json.loads((revision / 'approved-reference.json').read_text())


def test_general_xml_edit_is_reviewed_and_new_defect_gets_another_repair(tmp_path, monkeypatch):
    from xml.etree import ElementTree as ET
    def repair(request, revision):
        path = revision / request['workspace_path'] / 'study.xml'
        tree = ET.parse(path)
        title = tree.getroot().find('.//brief_title')
        title.text = 'Introduced wrong title' if request['request_id'].endswith('001') else request['approved_source']['study']['title']
        tree.write(path, encoding='utf-8', xml_declaration=True)
        docpath = revision / request['workspace_path'] / 'protocol.docx'
        doc = Document(docpath)
        visit_row = next(row for table in doc.tables for row in table.rows
                         if len(row.cells) > 2 and row.cells[1].text == 'Month 1')
        visit_row.cells[2].text = ('Wrong visit window' if request['request_id'].endswith('001')
                                   else 'Day 30 +/- 7')
        if not any(p.text == 'Preserved earlier native correction.' for p in doc.paragraphs):
            doc.add_paragraph('Preserved earlier native correction.')
        doc.save(docpath)
        return repair_response(request)
    def review(request, response, revision, number):
        if number <= 2:
            response['status'] = 'blocked'
            response['findings'] = [{'finding_id': 'native-content', 'category': 'content',
                'artifact': 'study.xml', 'target_ids': ['prs.structured'], 'check': 'source_supported',
                'issue': 'The XML title needs source-grounded correction.', 'material': True}]
            for row in response['section_assessments']:
                if row['section_id'] == 'prs.structured': row['status'] = 'blocked'
        else:
            assert ET.parse(revision / 'candidate/study.xml').find('.//brief_title').text == request['approved_source']['study']['title']
        return response
    result, revision, repairs, reviews = lifecycle(tmp_path, monkeypatch, repair, retrospective=False, review=review)
    assert result['status'] == 'passed', result
    assert len(repairs) == 2 and len(reviews) == 3
    assert repairs[-1]['feedback']  # repeated concern reaches AI reassessment, regardless of DOCX ZIP churn
    assert 'Preserved earlier native correction.' in [p.text for p in Document(tmp_path / 'output/protocol.docx').paragraphs]
    assert ET.parse(tmp_path / 'output/study.xml').find('.//brief_title').text == repairs[-1]['approved_source']['study']['title']
    visit_row = next(row for table in Document(tmp_path / 'output/protocol.docx').tables for row in table.rows
                     if len(row.cells) > 2 and row.cells[1].text == 'Month 1')
    assert visit_row.cells[2].text == 'Day 30 +/- 7'


def test_false_concern_is_disputed_then_independently_reassessed(tmp_path, monkeypatch):
    def repair(request, revision):
        return repair_response(request, status='disputed', disputes=[{
            'finding_id': 'native-content', 'explanation': 'The approved source supports the current native wording.'}])
    result, revision, repairs, reviews = lifecycle(tmp_path, monkeypatch, repair)
    assert result['status'] == 'passed', result
    assert len(repairs) == 1 and len(reviews) == 2
    assert reviews[-1]['review_context']['disputes'][0]['finding_id'] == 'native-content'


def test_corrected_native_files_survive_later_font_reconstruction(tmp_path, monkeypatch):
    render_count = []
    def render(root, revision, *args, **kwargs):
        render_count.append(1)
        if len(render_count) == 2:
            assert kwargs['rebuild_candidate']({'Arial': 'Liberation Sans'})['status'] == 'passed'
            assert Document(revision / 'candidate/protocol.docx').paragraphs[-1].text == 'Persistent native correction.'
        return controlled_render(root, revision)
    def repair(request, revision):
        path = revision / request['workspace_path'] / 'protocol.docx'
        doc = Document(path)
        doc.add_paragraph('Persistent native correction.')
        doc.save(path)
        return repair_response(request)
    result, *_ = lifecycle(tmp_path, monkeypatch, repair, render=render)
    assert result['status'] == 'passed', result
    assert Document(tmp_path / 'output/protocol.docx').paragraphs[-1].text == 'Persistent native correction.'


def test_original_candidate_mutation_during_repair_is_rejected(tmp_path, monkeypatch):
    def repair(request, revision):
        path = revision / 'candidate/protocol.docx'
        doc = Document(path)
        doc.add_paragraph('Unauthorized base mutation.')
        doc.save(path)
        return repair_response(request)
    result, revision, repairs, _ = lifecycle(tmp_path, monkeypatch, repair)
    assert result['status'] == 'blocked' and result['stage'] == 'integrity', result
    assert not (tmp_path / 'output').exists()


def test_malformed_xml_edit_returns_actionable_feedback_without_adoption(tmp_path, monkeypatch):
    from xml.etree import ElementTree as ET
    attempts = []
    def repair(request, revision):
        attempts.append(1)
        path = revision / request['workspace_path'] / 'study.xml'
        if len(attempts) == 1:
            path.write_text('<not-valid')
        else:
            path.write_bytes((revision / 'candidate/study.xml').read_bytes())
            tree = ET.parse(path)
            title = tree.find('.//brief_title')
            parent = next(node for node in tree.iter() if title in list(node))
            position = list(parent).index(title)
            parent.remove(title)
            replacement = ET.Element('brief_title')
            replacement.text = request['approved_source']['study']['title']
            parent.insert(position, replacement)
            tree.write(path, encoding='utf-8', xml_declaration=True)
        return repair_response(request)
    result, _, repairs, reviews = lifecycle(tmp_path, monkeypatch, repair, retrospective=False)
    assert result['status'] == 'passed', result
    assert len(attempts) == 2 and len(reviews) == 2


def test_repair_author_cannot_supply_final_independent_review(tmp_path, monkeypatch):
    def repair(request, revision):
        return repair_response(request, status='disputed', disputes=[{'finding_id': 'native-content', 'explanation': 'Supported by approved source.'}])
    def review(request, response, revision, number):
        if number == 1:
            response['status'] = 'blocked'
            response['findings'] = [{'finding_id': 'native-content', 'category': 'content',
                'artifact': 'protocol', 'target_ids': ['introduction'], 'check': 'source_supported',
                'issue': 'Reassess this content.', 'material': True}]
            next(row for row in response['section_assessments'] if row['section_id'] == 'introduction')['status'] = 'blocked'
        else:
            response['producer']['reviewer_id'] = 'repair-author'
        return response
    result, *_ = lifecycle(tmp_path, monkeypatch, repair, review=review)
    assert result['status'] != 'passed', result
    assert not (tmp_path / 'output').exists()


def test_stale_repair_response_never_certifies_or_overwrites_native_candidate(tmp_path, monkeypatch):
    def repair(request, revision):
        return repair_response(request, request_sha256='stale')
    result, revision, repairs, reviews = lifecycle(tmp_path, monkeypatch, repair)
    assert result['status'] == 'blocked' and result['stage'] == 'document_repair_no_progress', result
    assert not (tmp_path / 'output').exists()
    assert len(reviews) == 1
    assert (revision / 'candidate/protocol.docx').read_bytes() == (revision / 'repairs/pass-001/base/protocol.docx').read_bytes()


def test_source_mutation_during_repair_remains_unavailable(tmp_path, monkeypatch):
    def repair(request, revision):
        source = revision / 'approved-source.md'
        source.write_text(source.read_text() + '\nChanged approved source.\n')
        return repair_response(request)
    result, *_ = lifecycle(tmp_path, monkeypatch, repair)
    assert result['status'] == 'blocked', result
    assert not (tmp_path / 'output').exists()


def test_repair_can_report_unresolved_concerns_without_destroying_usable_base(tmp_path, monkeypatch):
    def repair(request, revision):
        return repair_response(request, status='needs_review', summary='No safe productive correction remains.')
    result, revision, repairs, reviews = lifecycle(tmp_path, monkeypatch, repair)
    assert result['status'] == 'blocked' and result['stage'] == 'document_repair_unresolved', result
    assert result['candidate_outputs'] and (revision / 'candidate/protocol.docx').is_file()
    assert not (tmp_path / 'output').exists()


@pytest.mark.parametrize('case', [case for case in CASES if case['descriptor']['profile'].startswith('rich')], ids=lambda case: case['descriptor']['fixture_id'])
def test_real_injected_prose_is_removed_and_freshly_reviewed_across_families(tmp_path, monkeypatch, case):
    defect = 'Participants receive an unsupported experimental drug.'
    observations = []
    def draft(request):
        response = drafting.recorded_acceptance_response(request)
        for row in response.get('section_results', []):
            if row['section_id'] == 'introduction':
                row['paragraphs'][0]['text'] += ' ' + defect
        return response
    def review(request, response, revision, count):
        present = any(defect in p.text for p in Document(revision / 'candidate/protocol.docx').paragraphs)
        observations.append(present)
        if present:
            response['status'] = 'blocked'
            response['findings'] = [{'finding_id': 'native-content', 'category': 'content',
                'artifact': 'protocol', 'target_ids': ['paragraph:17'], 'check': 'source_supported',
                'location': 'Protocol introduction paragraph', 'issue': defect, 'material': True}]
            next(row for row in response['section_assessments'] if row['section_id'] == 'introduction')['status'] = 'blocked'
        return response
    def repair(request, revision):
        path = revision / request['workspace_path'] / 'protocol.docx'
        doc = Document(path)
        for paragraph in doc.paragraphs:
            if defect in paragraph.text:
                paragraph.text = paragraph.text.replace(' ' + defect, '')
        doc.save(path)
        return repair_response(request)
    result, revision, repairs, reviews = lifecycle(tmp_path, monkeypatch, repair, case=case, draft=draft, review=review)
    assert result['status'] == 'passed', result
    assert observations == [True, False]
    assert len(repairs) == 1 and len(reviews) == 2
    assert next(row['sha256'] for row in reviews[0]['artifacts'] if row['artifact'] == 'protocol') != next(row['sha256'] for row in reviews[-1]['artifacts'] if row['artifact'] == 'protocol')
    assert not any(defect in p.text for p in Document(tmp_path / 'output/protocol.docx').paragraphs)
    assert json.loads((revision / 'approved-reference.json').read_text()) == repairs[0]['approved_source'] == reviews[-1]['approved_source']
    expected = {'protocol.docx'} if case['descriptor']['study_type'] == 'Retrospective' else {'protocol.docx', 'icf.docx', 'study.xml'}
    assert expected <= {path.name for path in (tmp_path / 'output').iterdir()}


def test_layout_repair_rechecks_both_protocol_and_icf_against_new_native_files(tmp_path, monkeypatch):
    from docx.shared import Pt
    case = next(case for case in CASES if case['descriptor']['study_type'] == 'Prospective' and case['descriptor']['profile'].startswith('rich'))
    seen = {'protocol': [], 'icf': []}
    def visual(request, response, revision):
        artifact = request['artifacts'][0]['artifact']
        seen[artifact].append(request['artifacts'][0]['docx_sha256'])
        if artifact == 'protocol' and len(seen[artifact]) == 1:
            response['status'] = 'blocked'
            response['findings'] = [{'finding_id': 'native-layout', 'category': 'visual',
                'artifact': 'protocol', 'target_ids': ['layout:protocol'], 'check': 'native spacing concern',
                'issue': 'Repair the visibly crowded paragraph spacing.', 'material': True,
                'page': 1, 'element': 'First body paragraph', 'location': 'Protocol page 1, first body paragraph'}]
            response['page_assessments'][0]['status'] = 'blocked'
        elif artifact == 'protocol':
            response['page_assessments'][0]['notes'] = ['native-layout: independently reinspected repaired page spacing.']
        return response
    def repair(request, revision):
        path = revision / request['workspace_path'] / 'protocol.docx'
        doc = Document(path)
        doc.paragraphs[0].paragraph_format.space_after = Pt(12)
        doc.save(path)
        return repair_response(request)
    result, revision, repairs, reviews = lifecycle(tmp_path, monkeypatch, repair, case=case,
        review=lambda request, response, revision, count: response, visual_review=visual)
    assert result['status'] == 'passed', result
    assert len(repairs) == 1
    assert len(seen['protocol']) == len(seen['icf']) == 2
    assert seen['protocol'][0] != seen['protocol'][1]
    assert Document(tmp_path / 'output/protocol.docx').paragraphs[0].paragraph_format.space_after.pt == 12
