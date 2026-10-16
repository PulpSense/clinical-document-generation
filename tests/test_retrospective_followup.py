"""Public handoff/rendering regressions for the first CHART review."""
import copy
import json
from pathlib import Path

import pytest
from docx import Document
from docx.oxml.ns import qn

import contracts
import drafting
import rendering

ROOT = Path(__file__).resolve().parents[1]


def source():
    ref = json.loads((ROOT / 'tests/fixtures/release-certification/retrospective/approved-reference.json').read_text())
    ref['sites'] = [
        {'facility': {'name': 'Meridian Eye Clinic', 'address': '520 Charles Street, Suite 110, Baltimore, MD 21201'},
         'investigator': {'name': 'Dr. Avery Brooks', 'title': 'Principal Investigator'},
         'contact': {'name': 'Casey Reed', 'title': 'Research Data Coordinator', 'business_phone': '+1 202-555-0181', 'office_phone': '+1 202-555-0182', 'email': 'casey.reed@example.org'}},
        {'facility': {'name': 'Severn Ophthalmology Center', 'address': '630 Severn Avenue, Suite 220, Annapolis, MD 21401'},
         'investigator': {'name': 'Dr. Riley Chen', 'title': 'Site Investigator'},
         'contact': {'name': 'Jordan Blake', 'title': 'Site Data Coordinator', 'business_phone': '+1 202-555-0183', 'office_phone': '+1 202-555-0184', 'email': 'jordan.blake@example.org'}}]
    ref['design']['number_of_sites'] = 2
    return ref


def test_accepted_eligibility_qualifications_survive_retrospective_assembly(tmp_path):
    ref = source()
    qualification = 'Exclusions apply to individual eyes. Retain an eligible eye when the other eye is excluded; exclude the person only when both eyes are ineligible.'
    model = {'protocol': [{'section_id': 'subjects.eligibility', 'paragraphs': [{'text': qualification}],
        'lists': [{'label': 'Inclusion criteria', 'items': ref['population']['inclusion_criteria']},
                  {'label': 'Exclusion criteria', 'items': ref['population']['exclusion_criteria']}]}], 'icf': {}, 'prs': {}}
    rendering.render_documents(ROOT, tmp_path, ref, model, audit_output=False)
    doc = Document(tmp_path / 'candidate/protocol.docx')
    assert [p.text for p in doc.paragraphs].count(qualification) == 1


@pytest.mark.parametrize('branch', ['Prospective', 'Ambispective', 'Retrospective'])
def test_protocol_opening_table_has_visible_cell_lines(branch, tmp_path):
    ref = source()
    ref['meta']['study_type'] = branch
    ref['meta']['icf_template'] = 'Sterling'
    rendering.render_documents(ROOT, tmp_path, ref, {'protocol': [], 'icf': {}, 'prs': {}}, artifact_names={'protocol'}, audit_output=False)
    doc = Document(tmp_path / 'candidate/protocol.docx')
    for row in doc.tables[0].rows:
        for cell in row.cells:
            borders = cell._tc.get_or_add_tcPr().find(qn('w:tcBorders'))
            assert borders is not None
            for side in ['top', 'left', 'bottom', 'right']:
                edge = borders.find(qn('w:'+side))
                assert edge is not None and edge.get(qn('w:val')) == 'single'
                assert edge.get(qn('w:color')) == '000000'


def test_retrospective_directory_keeps_all_supplied_sites_investigators_and_contacts(tmp_path):
    ref = source()
    # Empty plural placeholders must not hide the populated singular record.
    for site in ref['sites']:
        site['investigators'] = [{'name': '', 'title': ''}]
    original = copy.deepcopy(ref)
    rendering.render_documents(ROOT, tmp_path, ref, {'protocol': [], 'icf': {}, 'prs': {}}, audit_output=False)
    doc = Document(tmp_path / 'candidate/protocol.docx')
    directory = next(t for t in doc.tables if t.rows[0].cells[0].text == 'Study site')
    assert len(directory.rows) == 3
    text = '\n'.join(c.text for r in directory.rows for c in r.cells)
    for site in ref['sites']:
        for value in [*site['facility'].values(), *site['investigator'].values(), *site['contact'].values()]:
            assert value in text
    assert 'For participants enrolled' not in text
    assert ref == original


def test_retrospective_handoff_preserves_eligibility_context_and_site_table_ownership(tmp_path):
    ref = source()
    ref['population']['eligibility_notes'] = 'Exclusions apply per eye; retain the eligible fellow eye.'
    owner = 'subjects.eligibility'
    batch = next(b for b in contracts.batch_plan('Retrospective', '') if owner in b.section_ids)
    path = drafting.create_drafting_request(repo_root=ROOT, revision_dir=tmp_path,
        revision_id='test', reference=ref, batch=batch, target_ids=[owner],
        attempts={owner: 1}, wave='generation-only')
    request = json.loads(path.read_text())
    assert 'population.eligibility_notes' in request['section_contracts'][0]['minimum_evidence']
    assert request['evidence_checklist'][0]['approved_values']['population.eligibility_notes'] == ref['population']['eligibility_notes']
    directory = contracts.protocol_table_contracts(ref)['study-sites']
    assert directory['section_id'] == 'study-design.design'
    assert 'sites' in next(s for s in contracts.protocol_contract('Retrospective') if s.section_id == directory['section_id']).table_evidence


def test_empty_draft_fallback_keeps_supplied_criteria_and_eligibility_notes(tmp_path):
    ref = source()
    ref['population']['eligibility_notes'] = 'Exclude the person only if no eligible eye remains.'
    rendering.render_documents(ROOT, tmp_path, ref, {'protocol': [], 'icf': {}, 'prs': {}}, audit_output=False)
    visible = '\n'.join(p.text for p in Document(tmp_path / 'candidate/protocol.docx').paragraphs)
    for text in [*ref['population']['inclusion_criteria'], *ref['population']['exclusion_criteria'], ref['population']['eligibility_notes']]:
        assert text in visible
