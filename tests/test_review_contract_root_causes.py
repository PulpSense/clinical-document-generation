"""Shared failure patterns from review ingestion and narrative schedules."""
import copy
import json

import pytest
from docx import Document

import contracts
import quality
import rendering
import workflow
from test_review_resolutions import review_candidate
from test_handoff_quality import ROOT


@pytest.mark.parametrize('branch', ['Prospective', 'Ambispective'])
def test_generated_summary_is_available_to_section_review(tmp_path, branch):
    source = json.loads((ROOT / 'tests/fixtures/prospective-acceptance-source.json').read_text())
    source['meta']['icf_template'] = 'Sterling'
    source['meta']['study_type'] = branch
    text = 'This study compares the approved treatments.'
    rendering.render_documents(ROOT, tmp_path, source, {
        'protocol': [], 'icf': {'icf.key-information-summary': {
            'paragraphs': [{'text': text}], 'lists': []}}, 'prs': {}}, artifact_names={'icf'})
    snapshot = rendering.rendered_section_snapshot(tmp_path / 'candidate/icf.docx', 'icf.key-information-summary')
    assert text in quality._review_evidence_text(snapshot)


def test_substantive_blocked_review_survives_invalid_resolution_metadata(tmp_path):
    path, request, response = review_candidate(tmp_path)
    response['status'] = 'blocked'
    response['uncertainty_resolutions'] = []
    response['findings'] = [{
        'finding_id': 'missing-reminder', 'artifact': 'protocol',
        'check': 'source_supported', 'target_ids': ['study-design.bias'],
        'issue': 'The approved reminder procedure is missing.',
    }]
    for row in response['section_assessments']:
        if row['section_id'] == 'study-design.bias':
            row['status'] = 'failed'
    log = tmp_path / 'worker.log'
    log.write_text(json.dumps(response))
    handoff = {'task': request['task'], 'request_path': str(path.relative_to(tmp_path)),
               'response_path': request['response_path']}
    assert workflow._production_publish_quiet_response(tmp_path, handoff, log)
    findings, _ = quality.validate_verifications(tmp_path, request_paths=[path])
    assert any(row.get('verification_source_finding', {}).get('finding_id') == 'missing-reminder' for row in findings)
    assert any(row.get('code') == 'uncertainty_resolution_invalid' for row in findings)
    assert not quality.verification_response_is_complete(tmp_path, path)
    (tmp_path / request['response_path']).unlink()
    response['request_sha256'] = 'forged'
    log.write_text(json.dumps(response))
    assert not workflow._production_publish_quiet_response(tmp_path, handoff, log)


@pytest.mark.parametrize('visits', [
    ['Screening', 'Operative contact for each eye', 'Month 3 follow-up'],
    ['Baseline telephone contact', 'Week 2 clinic contact', 'Week 8 exit'],
])
def test_string_visit_inventory_drives_both_tables_without_splitting_prose(tmp_path, visits):
    source = json.loads((ROOT / 'tests/fixtures/prospective-acceptance-source.json').read_text())
    narrative = 'Screening: history, medications, examination. Follow-up: distance (4 m), near (40 cm), questionnaire.'
    source['procedures'] = {**source['procedures'], 'visit_schedule': visits, 'assessments': narrative}
    source['procedures'].pop('visit_schedule_table', None)
    original = copy.deepcopy(source)
    assert [row['visit'] for row in contracts.normalized_visit_records(source)] == visits
    rendering.render_documents(ROOT, tmp_path, source, {'protocol': [], 'icf': {}, 'prs': {}}, artifact_names={'protocol'})
    document = Document(tmp_path / 'candidate/protocol.docx')
    tables = [[[cell.text for cell in row.cells] for row in table.rows] for table in document.tables]
    schedule = next(rows for rows in tables if rows[0][0] == 'Approved visit or assessment')
    inventory = next(rows for rows in tables if rows[0][0].replace(' ', '') == 'VisitNumber')
    assert [row[0] for row in schedule[1:]] == visits
    assert [row[1] for row in inventory[1:]] == visits
    assert narrative in '\n'.join(p.text for p in document.paragraphs)
    assert source == original


def test_unstructured_assessment_prose_is_atomic():
    narrative = 'At screening: history, medications. At Week 8: distance (4 m), near (40 cm).'
    assert rendering._plain_language_assessments(narrative) == [narrative]


@pytest.mark.parametrize('raises', [True, False])
def test_identical_unconsumable_review_does_not_spend_entire_operation(tmp_path, monkeypatch, raises):
    now = [0.0]
    calls = []
    handoff = {'request_id': 'r1.verify.content', 'request_sha256': 'a' * 64,
               'request_path': 'hermes/verification-requests/content.json',
               'response_path': 'hermes/verification-responses/content.json',
               'task': 'clinical_content_verification'}
    monkeypatch.setattr(workflow, 'generate', lambda *_args, **_kwargs: {
        'status': 'awaiting_hermes', 'stage': 'independent_verification',
        'revision_id': 'r1', 'handoffs': [handoff]})
    monkeypatch.setattr(workflow.time, 'sleep', lambda *_: None)
    def runner(*_):
        calls.append(1)
        now[0] += 10
        if raises:
            raise ValueError('Completed worker output violates the response contract')
    result = workflow.run_desktop_operation(tmp_path, handoff_runner=runner,
        opener=lambda _: b'', budget_seconds=2700, clock=lambda: now[0], wall_clock=lambda: 1000 + now[0])
    assert len(calls) == 2
    assert result['status'] == 'blocked'
    assert result['findings'][0]['code'] == 'handoff_no_progress'
    assert now[0] < 2700


def test_empty_section_evidence_is_rejected_before_dispatch(tmp_path, monkeypatch):
    path, request, _ = review_candidate(tmp_path)
    monkeypatch.setattr(quality, 'rendered_section_snapshot', lambda *_: [])
    with pytest.raises(quality.ReviewEvidenceError) as error:
        quality.create_verification_requests(tmp_path, request['approved_source'], {'artifacts': []}, review_set=2)
    assert error.value.finding['code'] == 'review_evidence_unavailable'
    assert error.value.finding['owner'] == 'construction'
    assert not list((tmp_path / 'hermes/verification-requests').glob('*.review-2.*'))


def test_passing_response_with_incomplete_resolution_is_consumed_only_for_retry(tmp_path):
    path, request, response = review_candidate(tmp_path)
    response['uncertainty_resolutions'] = []
    log = tmp_path / 'worker.log'
    log.write_text(json.dumps(response))
    handoff = {'task': request['task'], 'request_path': str(path.relative_to(tmp_path)),
               'response_path': request['response_path']}
    assert workflow._production_publish_quiet_response(tmp_path, handoff, log)
    assert not quality.verification_response_is_complete(tmp_path, path)
    findings, _ = quality.validate_verifications(tmp_path, request_paths=[path])
    assert findings and all(row.get('code') == 'uncertainty_resolution_invalid' for row in findings)


def test_retry_prompt_includes_specific_contract_feedback_without_changing_request(tmp_path, monkeypatch):
    handoff = {'request_id': 'r1.verify.content', 'request_sha256': 'a' * 64,
               'request_path': 'content.json', 'response_path': 'response.json',
               'task': 'clinical_content_verification'}
    findings = [{'verification_request_id': handoff['request_id'], 'recovery_class': 'verifier_transient',
                 'issue': 'Every presented uncertainty needs an explicit source-grounded decision.'}]
    monkeypatch.setattr(workflow, 'generate', lambda *_args, **_kwargs: {
        'status': 'awaiting_hermes', 'stage': 'independent_verification_retry',
        'revision_id': 'r1', 'handoffs': [handoff], 'findings': findings})
    prompts = []
    monkeypatch.setattr(workflow.time, 'sleep', lambda *_: None)
    def runner(items, _remaining):
        prompts.append(workflow._production_agent_prompt(ROOT, tmp_path, items[0], {}, workspace_root=tmp_path))
    result = workflow.run_desktop_operation(tmp_path, handoff_runner=runner, opener=lambda _: b'', budget_seconds=30)
    assert result['status'] == 'blocked'
    assert len(prompts) == 2
    assert findings[0]['issue'] in prompts[0]
    assert handoff['request_sha256'] == 'a' * 64


def test_content_retry_exposes_bound_identity_and_feedback(tmp_path):
    path, request, _ = review_candidate(tmp_path)
    finding = {'verification_request_id': request['request_id'], 'recovery_class': 'verifier_transient',
               'issue': 'Supply the missing uncertainty resolution.'}
    result = workflow._awaiting(tmp_path, stage='independent_verification_retry', paths=[path], findings=[finding])
    assert result['handoffs'][0]['request_id'] == request['request_id']
    assert result['handoffs'][0]['request_sha256'] == request['request_sha256']
    assert result['findings'] == [finding]


def test_evidence_construction_failure_is_returned_to_its_owner(tmp_path, monkeypatch):
    _, request, _ = review_candidate(tmp_path)
    monkeypatch.setattr(quality, 'rendered_section_snapshot', lambda *_: [])
    monkeypatch.setattr(workflow, 'audit_computed_protocol_fields', lambda *_: [])
    bundle = contracts.contracted_template_bundle(ROOT, request['approved_source'])
    paths, findings = workflow._begin_independent_review(tmp_path, request['approved_source'],
                                                        {'artifacts': []}, bundle, 2)
    assert paths == []
    assert findings[0]['code'] == 'review_evidence_unavailable'
    assert findings[0]['owner'] == 'construction'


def test_new_candidate_review_has_its_own_transport_retry_budget(tmp_path, monkeypatch):
    handoff = {'request_id': 'r1.review-1', 'request_sha256': 'a' * 64,
               'request_path': 'review-1.json', 'response_path': 'response-1.json',
               'task': 'clinical_content_verification'}
    fresh = {**handoff, 'request_id': 'r1.review-2', 'request_sha256': 'b' * 64,
             'request_path': 'review-2.json', 'response_path': 'response-2.json'}
    def waiting(item):
        return {'status': 'awaiting_hermes', 'stage': 'independent_verification',
                'revision_id': 'r1', 'handoffs': [item]}
    results = iter([waiting(handoff), waiting(handoff), waiting(fresh),
                    {'status': 'blocked', 'stage': 'test-end', 'findings': []}])
    monkeypatch.setattr(workflow, 'generate', lambda *_args, **_kwargs: next(results))
    monkeypatch.setattr(workflow.time, 'sleep', lambda *_: None)
    calls = []
    result = workflow.run_desktop_operation(tmp_path,
        handoff_runner=lambda items, _: calls.append(items[0]['request_id']),
        opener=lambda _: b'', budget_seconds=30)
    assert result['stage'] == 'test-end'
    assert calls == ['r1.review-1', 'r1.review-1', 'r1.review-2']


@pytest.mark.parametrize('branch', ['Prospective', 'Ambispective'])
@pytest.mark.parametrize('family', ['Sterling', 'Advarra'])
def test_every_applicable_generated_icf_section_can_be_localized(tmp_path, branch, family):
    source = json.loads((ROOT / 'tests/fixtures/prospective-acceptance-source.json').read_text())
    source['meta'].update(study_type=branch, icf_template=family)
    sections = contracts.icf_contract(branch, family)
    model = {'protocol': [], 'prs': {}, 'icf': {
        spec.section_id: {'paragraphs': [{'text': 'Sentinel for ' + spec.section_id}], 'lists': []}
        for spec in sections}}
    rendering.render_documents(ROOT, tmp_path, source, model, artifact_names={'icf'})
    for spec in sections:
        snapshot = rendering.rendered_section_snapshot(tmp_path / 'candidate/icf.docx', spec.section_id)
        # Advarra's injury addendum is deliberately rebuilt verbatim from source.
        expected = (source['risks_benefits']['injury_handling']
                    if family == 'Advarra' and spec.section_id == 'icf.injury'
                    else 'Sentinel for ' + spec.section_id)
        assert expected in quality._review_evidence_text(snapshot), spec.section_id


def test_resolution_metadata_retry_reaches_clean_delivery(tmp_path, monkeypatch):
    from test_clean_delivery import deliver, reminder_source, reminder_draft
    from test_handoff_quality import acceptance_verification
    text = 'The team phones people before appointments so they remember to attend.'
    reviews = []
    def review(request):
        response = acceptance_verification(request)
        if request['task'] == 'clinical_content_verification':
            reviews.append(request['request_id'])
            for decision in response.get('uncertainty_resolutions', []):
                caution = next(row for row in request['lexical_review_cautions']
                               if row['caution_sha256'] == decision['caution_sha256'])
                if caution['section_id'] == 'study-design.bias':
                    decision.update(affected_passage=text,
                        source_excerpt=reminder_source()['statistics']['bias_minimization'],
                        rationale='Calling before appointments expresses the supplied telephone reminders.')
            if len(reviews) == 1:
                assert request['lexical_review_cautions']
                response['uncertainty_resolutions'] = []
        return response
    result, _ = deliver(tmp_path, monkeypatch, reminder_source(), review=review,
                        draft=lambda request: reminder_draft(request, text))
    assert result['status'] == 'passed', result
    assert result['warnings'] == []
    assert len(reviews) == 2 and reviews[0] == reviews[1]


def test_changing_metadata_errors_recover_through_public_desktop_delivery(tmp_path, monkeypatch):
    from pathlib import Path
    from test_clean_delivery import controlled_render, reminder_source, reminder_draft
    from test_handoff_quality import acceptance_verification
    monkeypatch.setattr(workflow, 'render_assurance', controlled_render)
    reference = tmp_path / 'reference/study.reference.json'
    reference.parent.mkdir(parents=True)
    reference.write_text(json.dumps(reminder_source()))
    assert workflow.prepare(tmp_path)['status'] == 'awaiting_approval'
    approved = workflow.approve(tmp_path, approved_by='controlled-fixture-owner')
    revision = tmp_path / 'revisions' / approved['revision_id']
    reviews, feedback = [], []
    text = 'The team phones people before appointments so they remember to attend.'

    def runner(items, _remaining):
        for item in items:
            request = json.loads((revision / item['request_path']).read_text())
            if request['task'] in {'section_drafting', 'prs_narrative_drafting'}:
                response = reminder_draft(request, text)
                path = revision / request['response_path']
                path.parent.mkdir(parents=True, exist_ok=True)
                path.write_text(json.dumps(response))
                continue
            response = acceptance_verification(request)
            if request['task'] == 'clinical_content_verification':
                reviews.append(request['request_id'])
                feedback.append(item.get('response_repair_feedback', []))
                if len(reviews) == 1:
                    response['uncertainty_resolutions'] = []
                elif len(reviews) == 2:
                    for decision in response['uncertainty_resolutions']:
                        decision['affected_passage'] = 'Not a quotation in this document.'
            log = tmp_path / f"{request['request_id']}.{len(reviews)}.worker.json"
            log.write_text(json.dumps(response))
            assert workflow._production_publish_quiet_response(revision, item, log)

    result = workflow.run_desktop_operation(tmp_path, handoff_runner=runner,
        opener=lambda path: Path(path).read_bytes(), budget_seconds=2700, require_promoted_runtime=False)
    assert result['status'] == 'passed', result
    assert result['warnings'] == []
    assert len(reviews) == 3 and len(set(reviews)) == 1
    assert feedback[1] and feedback[2] and feedback[1] != feedback[2]
    assert sorted(path.name for path in (tmp_path / 'output').iterdir()) == ['icf.docx', 'protocol.docx', 'study.xml']
    assert result['deadline_at_epoch'] - result['started_at_epoch'] == 2700


@pytest.mark.parametrize('errors,expected_calls', [
    ([['A'], ['B'], ['A'], ['B'], ['A']], 4),
    ([['A', 'B'], ['B', 'A'], ['A', 'B']], 2),
    ([['A', 'A'], ['A'], ['A'], ['A']], 3),
])
def test_feedback_progress_preserves_cycles_order_and_error_count(tmp_path, monkeypatch, errors, expected_calls):
    handoff = {'request_id': 'r1.verify', 'request_sha256': 'a' * 64,
               'request_path': 'request.json', 'response_path': 'response.json',
               'task': 'clinical_content_verification'}
    steps, calls = iter(errors), []
    def generate(*_args, **_kwargs):
        return {'status': 'awaiting_hermes', 'stage': 'independent_verification_retry',
                'revision_id': 'r1', 'handoffs': [handoff], 'findings': [
                    {'verification_request_id': handoff['request_id'], 'issue': issue,
                     'recovery_class': 'verifier_transient'} for issue in next(steps)]}
    monkeypatch.setattr(workflow, 'generate', generate)
    monkeypatch.setattr(workflow.time, 'sleep', lambda *_: None)
    result = workflow.run_desktop_operation(tmp_path,
        handoff_runner=lambda items, _: calls.append(items[0]['response_repair_feedback']),
        opener=lambda _: b'', budget_seconds=30)
    assert result['findings'][0]['code'] == 'handoff_no_progress'
    assert len(calls) == expected_calls


@pytest.mark.parametrize('resumed_error,expected_calls', [('A', 2), ('B', 3)])
def test_feedback_progress_and_deadline_survive_process_resume(tmp_path, monkeypatch, resumed_error, expected_calls):
    handoff = {'request_id': 'r1.verify', 'request_sha256': 'a' * 64,
               'request_path': 'request.json', 'response_path': 'response.json',
               'task': 'clinical_content_verification'}
    error, now, calls = ['A'], [0.0], []
    monkeypatch.setattr(workflow, 'generate', lambda *_args, **_kwargs: {
        'status': 'awaiting_hermes', 'stage': 'independent_verification_retry',
        'revision_id': 'r1', 'handoffs': [handoff], 'findings': [{
            'verification_request_id': handoff['request_id'], 'issue': error[0],
            'recovery_class': 'verifier_transient'}]})
    monkeypatch.setattr(workflow.time, 'sleep', lambda *_: None)
    def interrupted(*_args):
        calls.append(error[0])
        now[0] += 1
        raise KeyboardInterrupt
    with pytest.raises(KeyboardInterrupt):
        workflow.run_desktop_operation(tmp_path, handoff_runner=interrupted,
            opener=lambda _: b'', budget_seconds=30, clock=lambda: now[0], wall_clock=lambda: 1000)
    now[0], error[0] = 500, resumed_error
    def runner(*_args):
        calls.append(error[0])
        now[0] += 1
    result = workflow.run_desktop_operation(tmp_path, handoff_runner=runner,
        opener=lambda _: b'', budget_seconds=2700, clock=lambda: now[0], wall_clock=lambda: 1010)
    assert result['findings'][0]['code'] == 'handoff_no_progress'
    assert len(calls) == expected_calls
    assert result['started_at_epoch'] == 1000 and result['deadline_at_epoch'] == 1030
