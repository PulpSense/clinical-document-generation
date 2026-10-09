"""Public workflow regressions for shared reliability, using synthetic studies."""
import json
from pathlib import Path

import pytest

import contracts
import drafting
import quality
import workflow
from test_clean_delivery import deliver
from test_shared_output_projection import source_for

ROOT = Path(__file__).resolve().parents[1]


def draft_response(tmp_path, target, source, text, refs):
    batch = next(batch for batch in contracts.batch_plan(
        source['meta']['study_type'], source['meta'].get('icf_template')) if target in batch.section_ids)
    path = drafting.create_drafting_request(repo_root=ROOT, revision_dir=tmp_path,
        revision_id='general-reliability', reference=source, batch=batch,
        target_ids=[target], attempts={target: 1}, wave='initial')
    request = json.loads(path.read_text())
    response = drafting.recorded_acceptance_response(request)
    if target.startswith('prs.'):
        narrative = response['narrative'][target.split('.', 1)[1].replace('-', '_')]
        narrative['text'] = narrative['text'].replace(source['objectives']['primary'], text)
    else:
        response['section_results'][0]['paragraphs'] = [
            {'text': text, 'evidence_refs': refs, 'boilerplate_refs': []}]
        response['section_results'][0]['lists'] = []
    return request, response


def test_introduction_paraphrase_reaches_review_instead_of_retrying(tmp_path):
    source = source_for('Prospective', 'Sterling')
    source['study']['background'] = 'Prior studies found improved visual acuity after lens implantation.'
    request, response = draft_response(tmp_path, 'introduction', source,
        'Earlier research reported better eyesight following intraocular lens placement.', ['source:study.background'])
    accepted, findings = drafting.validate_response(request, response)
    assert [row['section_id'] for row in accepted['drafts']] == ['introduction']
    assert findings and all(row.get('publication_disposition') == 'warning' for row in findings)
    assert findings[0]['evidence_diagnostics'][0]['source_path'] == 'study.background'


def test_prs_paraphrase_is_retained_with_its_review_evidence(tmp_path):
    source = source_for('Prospective', 'Advarra')
    source['objectives']['primary'] = 'Measure the burden of persistent discomfort among working adults.'
    request, response = draft_response(tmp_path, 'prs.brief-summary', source,
        'Quantify ongoing soreness experienced by currently employed people in adulthood.', ['source:objectives.primary'])
    accepted, findings = drafting.validate_response(request, response)
    assert findings and all(row.get('publication_disposition') == 'warning' for row in findings)
    assert accepted['narrative']['brief_summary']['warnings'] == findings


def test_source_faithful_objective_mentions_do_not_stop_drafting(tmp_path):
    source = source_for('Prospective', 'Sterling')
    source['objectives']['primary'] = 'Assess the primary endpoint of visual acuity.'
    source['objectives']['secondary'] = []
    request, response = draft_response(tmp_path, 'objectives', source,
        'The study will assess the primary endpoint of visual acuity.', ['source:objectives.primary'])
    accepted, findings = drafting.validate_response(request, response)
    assert [row['section_id'] for row in accepted['drafts']] == ['objectives']
    assert findings and all(row.get('publication_disposition') == 'warning' for row in findings)


@pytest.mark.parametrize('source_text,draft_text', [
    ('Prior research documented improved visual outcomes at 13 months.',
     'Prior research documented improved visual outcomes at thirteen months.'),
    ('Prior research documented improved visual outcomes at 12 months.',
     'Prior research documented improved visual outcomes at one year.'),
])
def test_narrative_quantity_representation_reaches_independent_review(tmp_path, source_text, draft_text):
    source = source_for('Prospective', 'Sterling')
    source['study']['background'] = source_text
    request, response = draft_response(tmp_path, 'icf.background', source, draft_text, ['source:study.background'])
    accepted, findings = drafting.validate_response(request, response)
    assert [row['section_id'] for row in accepted['drafts']] == ['icf.background']
    assert findings and all(row.get('publication_disposition') == 'warning' for row in findings)


def test_materially_improving_repairs_finish_despite_unchanged_finding_text(tmp_path, monkeypatch):
    source = json.loads((ROOT / 'tests/fixtures/prospective-acceptance-source.json').read_text())
    source['meta']['icf_template'] = 'Sterling'
    source['study']['background'] = 'Prior research used 5 mL blood sampling, a -4 mm alignment offset, and 2 IU injections.'
    attempts = []
    passages = [
        'Prior research used 5 ML blood sampling, a 4 mm alignment offset, and 2 iu injections.',
        'Prior research used 5 mL blood sampling, a 4 mm alignment offset, and 2 iu injections.',
        'Prior research used 5 mL blood sampling, a -4 mm alignment offset, and 2 iu injections.',
        source['study']['background'],
    ]
    def draft(request):
        response = drafting.recorded_acceptance_response(request)
        for row in response.get('section_results', []):
            if row['section_id'] == 'icf.background':
                row['paragraphs'] = [{'text': passages[min(len(attempts), 3)],
                    'evidence_refs': ['source:study.background'], 'boilerplate_refs': []}]
                row['lists'] = []
                attempts.append(row['paragraphs'][0]['text'])
        return response
    result, _ = deliver(tmp_path, monkeypatch, source, draft=draft)
    assert result['status'] == 'passed', result
    assert attempts == passages
    assert result['warnings'] == []


@pytest.mark.parametrize('source_text,changed', [
    ('The acuity chart is at 66 cm.', 'The acuity chart is at 65 cm.'),
    ('The administered dose is 1 mg.', 'The administered dose is 1 g.'),
])
def test_unambiguous_quantity_changes_still_require_correction(source_text, changed):
    detail = drafting.source_evidence_diagnostics(changed, 'procedures.assessments', source_text)
    assert detail['quantity_issues']
    assert not drafting.lexical_evidence_uncertainty([detail])


INDEPENDENT_CASES = sorted((ROOT / 'tests/fixtures/independent-reliability').glob('*.json'))


def independent_draft(request, case, *, omit_objective=False):
    # Only unrelated sections use the standard orchestration adapter. The tested
    # passages are fixed authored records, never rewritten by evidence_grounded.
    response = drafting.recorded_acceptance_response(request)
    for row in response.get('section_results', []):
        record = case['drafts'].get(row['section_id'])
        if record:
            row['paragraphs'] = [{'text': ('Study records will be stored securely for authorized staff.'
                if omit_objective and row['section_id'] == 'objectives' else record['text']),
                'evidence_refs': record['evidence_refs'], 'boilerplate_refs': []}]
            row['lists'] = []
    for row in response.get('narrative', {}).values():
        row['text'] = row['text'].replace(case['source']['objectives']['primary'], case['prs_objective'])
    return response


@pytest.mark.parametrize('case_path', INDEPENDENT_CASES, ids=lambda path: path.stem)
def test_independent_clinical_inputs_deliver_all_reviewed_artifacts_cleanly(tmp_path, monkeypatch, case_path):
    from docx import Document
    from xml.etree import ElementTree as ET
    case = json.loads(case_path.read_text())
    result, _ = deliver(tmp_path, monkeypatch, case['source'], draft=lambda request: independent_draft(request, case))
    assert result['status'] == 'passed', result
    assert result['warnings'] == [] and not result['desktop_reply'].get('cautions')
    protocol = '\n'.join(row.text for row in Document(tmp_path / 'output/protocol.docx').paragraphs)
    assert case['drafts']['introduction']['text'] in protocol
    assert case['drafts']['objectives']['text'] in protocol
    manifest = json.loads((tmp_path / result['manifest']).read_text())
    assert workflow.confirm_desktop_delivery(manifest, result['desktop_reply'], lambda path: Path(path).read_bytes())['confirmed']
    if case['study_type'] != 'Retrospective':
        brief = ET.parse(tmp_path / 'output/study.xml').findtext('.//brief_summary/textblock')
        assert case['prs_objective'] in brief
        assessments = manifest['quality']['verification_evidence']['clinical_content_verification']['section_assessments']
        assert any(row['artifact'] == 'study.xml' and row['section_id'] == 'prs.brief-summary'
                   and row['status'] == 'passed' for row in assessments)


@pytest.mark.parametrize('case_path', INDEPENDENT_CASES, ids=lambda path: path.stem)
def test_real_omission_is_repaired_and_reverified_for_each_family(tmp_path, monkeypatch, case_path):
    from test_handoff_quality import acceptance_verification
    from docx import Document
    case = json.loads(case_path.read_text())
    objective_attempts = []
    bad = 'Study records will be stored securely for authorized staff.'
    confirmed = []
    def draft(request):
        if any(row['section_id'] == 'objectives' for row in request.get('section_contracts', [])):
            objective_attempts.append(request['request_id'])
        return independent_draft(request, case, omit_objective=len(objective_attempts) <= 1)
    def review(request):
        response = acceptance_verification(request)
        if request['task'] != 'clinical_content_verification':
            return response
        # Fixed fixture judgement, independent of the lexical checker: record
        # storage does not express any of these five approved clinical aims.
        cautions = [row for row in request.get('lexical_review_cautions', [])
                    if row['section_id'] == 'objectives' and bad in quality._review_evidence_text(row['artifact_passage'])]
        if not cautions:
            return response
        confirmed.append(request['request_id'])
        response['status'] = 'blocked'
        response['findings'] = [{'finding_id': 'clinical-aim-omitted', 'artifact': 'protocol',
            'check': 'source_supported', 'target_ids': ['objectives'],
            'issue': 'Storage arrangements omit the approved clinical aim.',
            'affected_passage': bad, 'source_excerpt': case['source']['objectives']['primary'],
            'recommended_action': 'State the approved clinical aim in accurate language.'}]
        for row in response['section_assessments']:
            if row['section_id'] == 'objectives':
                row['status'] = 'blocked'
        for decision in response['uncertainty_resolutions']:
            if decision['caution_sha256'] in {row['caution_sha256'] for row in cautions}:
                decision.update(disposition='confirmed_defect', affected_passage=bad,
                    source_excerpt=case['source']['objectives']['primary'],
                    rationale='Record storage does not describe the supplied clinical aim.')
        return response
    result, revision = deliver(tmp_path, monkeypatch, case['source'], draft=draft, review=review)
    assert confirmed and len(objective_attempts) == 2
    assert result['status'] == 'passed', result
    assert result['warnings'] == [] and not result['desktop_reply'].get('cautions')
    body = '\n'.join(row.text for row in Document(tmp_path / 'output/protocol.docx').paragraphs)
    assert bad not in body and case['drafts']['objectives']['text'] in body
    assert any('clinical-aim-omitted' in path.read_text() for path in (revision / 'attempts').rglob('attempt-manifest.json'))


@pytest.mark.parametrize('cycling', [False, True])
def test_ineffective_or_cycling_quantity_repairs_never_publish(tmp_path, monkeypatch, cycling):
    source = json.loads((ROOT / 'tests/fixtures/prospective-acceptance-source.json').read_text())
    source['meta']['icf_template'] = 'Sterling'
    source['study']['background'] = 'Prior research used 5 mL blood sampling, a -4 mm alignment offset, and 2 IU injections.'
    worse = 'Prior research used 5 ML blood sampling, a 4 mm alignment offset, and 2 iu injections.'
    better = 'Prior research used 5 mL blood sampling, a 4 mm alignment offset, and 2 iu injections.'
    attempts = []
    def draft(request):
        response = drafting.recorded_acceptance_response(request)
        for row in response.get('section_results', []):
            if row['section_id'] == 'icf.background':
                text = better if cycling and len(attempts) % 2 else worse
                attempts.append(text)
                row['paragraphs'] = [{'text': text, 'evidence_refs': ['source:study.background'], 'boilerplate_refs': []}]
                row['lists'] = []
        return response
    result, _ = deliver(tmp_path, monkeypatch, source, draft=draft)
    assert result['status'] != 'passed' and result.get('stage') == 'internal_recovery_stalled'
    assert len(attempts) == (4 if cycling else 3)
    assert not (tmp_path / 'output').exists()


@pytest.mark.parametrize('source_text,incorrect', [
    ('Prior research administered 2 mg in the morning and 5 mg in the evening.',
     'Prior research administered 5 mg in the morning and 2 mg in the evening.'),
    ('Prior research assessed diary completion for 13 months.',
     'Prior research assessed diary completion for 13 days.'),
    ('Prior research reported a mean difference of 2 points with confidence interval [1–3].',
     'Prior research reported a mean difference of 2 points.'),
    ('Prior research measured visual acuity at 66 cm.',
     'Earlier studies assessed eyesight at 65 cm.'),
])
def test_independent_review_repairs_changed_quantity_relationships(tmp_path, monkeypatch, source_text, incorrect):
    from test_handoff_quality import acceptance_verification
    from rendering import rendered_section_snapshot
    source = json.loads((ROOT / 'tests/fixtures/prospective-acceptance-source.json').read_text())
    source['meta']['icf_template'] = 'Sterling'
    source['study']['background'] = source_text
    attempts, rejected, reviewed = [], [], []
    def draft(request):
        response = drafting.recorded_acceptance_response(request)
        for row in response.get('section_results', []):
            if row['section_id'] == 'icf.background':
                text = source_text if attempts else incorrect
                attempts.append(text)
                row['paragraphs'] = [{'text': text, 'evidence_refs': ['source:study.background'], 'boilerplate_refs': []}]
                row['lists'] = []
        return response
    def review(request):
        response = acceptance_verification(request)
        if request['task'] != 'clinical_content_verification':
            return response
        path = tmp_path / 'revisions' / request['revision_id'] / 'candidate/icf.docx'
        passage = quality._review_evidence_text(rendered_section_snapshot(path, 'icf.background'))
        reviewed.append(passage)
        if incorrect not in passage:
            return response
        rejected.append(request['request_id'])
        response['status'] = 'blocked'
        response['findings'] = [{'finding_id': 'quantity-role-changed', 'artifact': 'icf',
            'check': 'source_supported', 'target_ids': ['icf.background'],
            'issue': 'The quantity relationship differs from the approved background.',
            'affected_passage': incorrect, 'source_excerpt': source_text,
            'recommended_action': 'Restore the approved dose roles or observation duration.'}]
        for row in response['section_assessments']:
            if row['section_id'] == 'icf.background':
                row['status'] = 'blocked'
        for decision in response.get('uncertainty_resolutions', []):
            caution = next(row for row in request['lexical_review_cautions'] if row['caution_sha256'] == decision['caution_sha256'])
            if caution['section_id'] == 'icf.background':
                decision.update(disposition='confirmed_defect', affected_passage=incorrect,
                    source_excerpt=source_text, rationale='The supplied quantity relationship or interval was changed or omitted.')
        return response
    result, _ = deliver(tmp_path, monkeypatch, source, draft=draft, review=review)
    assert rejected and len(attempts) == 2
    assert result['status'] == 'passed', result
    assert source_text in reviewed[-1] and incorrect not in reviewed[-1]


def test_prs_uncertainty_is_bound_to_xml_and_resolved_before_clean_delivery(tmp_path, monkeypatch):
    source = json.loads((ROOT / 'tests/fixtures/prospective-acceptance-source.json').read_text())
    source['objectives']['primary'] = 'Measure the burden of persistent discomfort among working adults.'
    text = 'Quantify ongoing soreness experienced by currently employed people in adulthood.'
    def draft(request):
        response = drafting.recorded_acceptance_response(request)
        for row in response.get('narrative', {}).values():
            row['text'] = row['text'].replace(source['objectives']['primary'], text)
        return response
    result, revision = deliver(tmp_path, monkeypatch, source, draft=draft)
    assert result['status'] == 'passed', result
    assert result['warnings'] == [] and not result['desktop_reply'].get('cautions')
    manifest = json.loads((tmp_path / result['manifest']).read_text())
    decisions = [row for row in manifest['quality']['resolved_evidence_cautions'] if row['artifact'] == 'study.xml']
    assert decisions and all(row['disposition'] == 'faithful' for row in decisions)
    assert all(row['artifact_sha256'] == quality.sha256_file(revision / 'candidate/study.xml') for row in decisions)
    assert all(source['objectives']['primary'] in quality._review_evidence_text(row['caution']['evidence_diagnostics'])
               for row in decisions)


def test_quantity_from_another_role_does_not_create_an_objective_false_alarm():
    detail = drafting.source_evidence_diagnostics(
        'Five milligrams is the dose. The preparation contains 5 g of carrier material.',
        'procedures.assessments', 'The dose is 5 mg.')
    assert not detail.get('quantity_issues')
    assert detail['release_acceptance'] is False


def test_negative_wording_proxy_requests_review_instead_of_an_objective_veto():
    detail = drafting.source_evidence_diagnostics(
        'Participant remuneration is absent.', 'risks_benefits.compensation_or_reimbursement', 'None')
    assert not detail['passed']
    assert drafting.lexical_evidence_uncertainty([detail])


def test_completion_followup_accepts_spelled_duration_for_review(tmp_path):
    source = source_for('Prospective', 'Advarra')
    source['study']['timeline'] = '13 months'
    request, response = draft_response(tmp_path, 'endpoint-criteria.completion', source,
        'The planned participant follow-up lasts thirteen months. Completion and discontinuation are separate dispositions.',
        ['source:study.timeline'])
    accepted, findings = drafting.validate_response(request, response)
    assert accepted['drafts']
    assert findings and all(row.get('publication_disposition') == 'warning' for row in findings)
