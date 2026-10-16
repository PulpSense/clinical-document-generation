import json
from pathlib import Path

import pytest
from docx import Document
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.shared import Inches

import rendering

ROOT = Path(__file__).resolve().parents[1]


@pytest.mark.parametrize("sterling", [False, True])
def test_body_spacers_removed_without_changing_signatures_or_fields(sterling):
    document = Document()
    document.sections[0].top_margin = Inches(1.77)
    document.add_paragraph("PURPOSE", style="Heading 1")
    document.add_paragraph()
    document.add_paragraph()
    inherited_heading_spacer = document.add_paragraph()
    inherited_properties = OxmlElement("w:rPr")
    inherited_properties.append(OxmlElement("w:u"))
    inherited_heading_spacer._p.get_or_add_pPr().append(inherited_properties)
    document.add_paragraph("Retain this consent wording.")
    field = document.add_paragraph()
    field._p.append(OxmlElement("w:fldSimple"))
    page_break = document.add_paragraph()
    page_break.add_run().add_break(WD_BREAK.PAGE)
    document.add_paragraph("PARTICIPANT STATEMENT AUTHORIZATION" if sterling else "AGREEMENT TO BE IN THE STUDY")
    document.add_paragraph("I agree to participate.")
    signing_space = document.add_paragraph()
    document.add_paragraph("________________________________" if sterling else "\t\t\t")
    document.add_paragraph("Signature of Participant")
    tail = [p._p.xml for p in document.paragraphs[-3:]]
    wording = [p.text for p in document.paragraphs if p.text.strip()]

    rendering._normalize_icf_body_spacing(document, sterling=sterling)

    assert [p.text for p in document.paragraphs if p.text.strip()] == wording
    assert [p._p.xml for p in document.paragraphs[-3:]] == tail
    assert field._p.getparent() is not None
    assert inherited_heading_spacer._p.getparent() is None
    assert page_break._p.getparent() is not None
    assert signing_space._p.getparent() is not None
    assert document.sections[0].top_margin.inches == pytest.approx(1.77 if sterling else 1, abs=0.001)
    before = document._element.xml
    rendering._normalize_icf_body_spacing(document, sterling=sterling)
    assert document._element.xml == before


@pytest.mark.parametrize("family", ["Advarra", "Sterling"])
def test_assembled_icf_cleans_body_after_authority_layout_is_applied(family, tmp_path):
    reference = json.loads((ROOT / "tests/fixtures/prospective-acceptance-source.json").read_text())
    reference["meta"]["icf_template"] = family
    rendering.render_documents(ROOT, tmp_path, reference, {"protocol": [], "icf": {}, "prs": {}},
                               artifact_names={"icf"}, audit_output=False)
    document = Document(tmp_path / "candidate/icf.docx")
    if family == "Advarra":
        assert all(section.top_margin == Inches(1) for section in document.sections)
    else:
        authority = Document(ROOT / "assets/client-templates/reference/sterling-icf-reference.docx")
        assert document.sections[0].top_margin == authority.sections[0].top_margin
    before = document._element.xml
    rendering._normalize_icf_body_spacing(document, sterling=family == "Sterling")
    assert document._element.xml == before
    assert any("Signature of Participant" in p.text for p in document.paragraphs)


def test_signature_compaction_retains_lines_labels_and_preference_wording():
    document = Document(ROOT / "assets/client-templates/reference/advarra-icf-reference.docx")
    rendering._normalize_icf_preferences(document)
    before = [p.text for p in document.paragraphs if p.text.strip()]
    lines = sum(bool(p._p.xpath('.//w:tab')) for p in document.paragraphs)
    rendering._compact_icf_signature_end(document)
    assert [p.text for p in document.paragraphs if p.text.strip()] == before
    assert sum(bool(p._p.xpath('.//w:tab')) for p in document.paragraphs) == lines
    final = next(p for p in document.paragraphs if p.text.strip().casefold() ==
                 "you will be given a signed and dated copy of this informed consent document to keep.")
    assert final.paragraph_format.keep_with_next
