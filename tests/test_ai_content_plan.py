"""Public generation/render boundaries: explicit AI content survives native assembly."""
import copy
import json
from pathlib import Path
from xml.etree import ElementTree as ET

import pytest
from docx import Document

import content_plan
import contracts
import drafting
import prs_xml
import rendering
import workflow
from test_generation_only import approved_run

ROOT = Path(__file__).resolve().parents[1]


def plan_response(request):
    """Deliberately unusual authored values reveal any assembler interpretation."""
    tables = {name: {"caption": name, "header_rows": 1, "rows": [], "supplemental_notes": []}
              for name in ("study-sites", "schedule-of-assessments", "sample-size-evidence")}
    plan = {"docx_fields": {artifact: {name: "" for name in names} for artifact, names in request["docx_field_names"].items()}, "visits": [], "tables": tables,
            "source_sections": {s["section_id"]: {"include": False, "blocks": []}
                                for s in request["protocol_sections"] if s["role"] == "source"},
            "investigator_agreement": [{"text": "AI-authored agreement.", "bullet": False}],
            "container_introductions": {}, "shell_edits": {a: [] for a in request["templates"]},
            "extra_icf_sections": [], "bibliography": "", "xml": None, "decisions": []}
    if request["xml_schema"]:
        plan["xml"] = {"design_branch": "observational_design", "scalars": {
            p: "" for p in request["xml_schema"]["scalar_paths"] if p.split('/')[0] not in content_plan.REPEATED
            and not p.startswith('study_design/interventional_design/')},
            "repeated": {tag: [] for tag in content_plan.REPEATED}}
    return {"schema_version": content_plan.SCHEMA, "request_id": request["request_id"],
            "request_sha256": request["request_sha256"], "producer": {"model_id": "controlled-fixture"}, "plan": plan}


def accept_plan(result, edit=None):
    workspace = Path(result["workspace"])
    request = json.loads((workspace / result["requests"][0]).read_text())
    response = plan_response(request)
    if edit:
        edit(response["plan"])
    (workspace / request["response_path"]).write_text(json.dumps(response))
    return request, response


def test_generation_requires_source_bound_ai_plan_before_drafting(tmp_path):
    approved_run(tmp_path)
    result = workflow.generate_only(tmp_path)
    assert result["stage"] == "content_planning"
    request, response = accept_plan(result)
    assert request["approved_source"] == json.loads((tmp_path / "revisions" / result["revision_id"] / "approved-reference.json").read_text())
    assert request["guidance"]["sterling-clause-contract.json"]
    assert request["xml_schema"]["repeated"]["location"]
    second = workflow.generate_only(tmp_path)
    assert second["stage"] == "drafting"
    for path in second["requests"]:
        draft_request = json.loads((Path(second["workspace"]) / path).read_text())
        assert draft_request["content_plan"] == response["plan"]
        assert draft_request["approved_source"] == request["approved_source"]
        assert drafting._trusted_request_valid(Path(second["workspace"]), draft_request)


def test_changed_accepted_plan_cannot_reuse_drafts(tmp_path):
    approved_run(tmp_path)
    result = workflow.generate_only(tmp_path)
    request, response = accept_plan(result)
    workflow.generate_only(tmp_path)
    response["plan"]["bibliography"] = "changed"
    (Path(result["workspace"]) / request["response_path"]).write_text(json.dumps(response))
    assert workflow.generate_only(tmp_path)["stage"] == "content_plan_response"


def test_invalid_plan_has_format_feedback_without_fallback(tmp_path):
    approved_run(tmp_path)
    result = workflow.generate_only(tmp_path)
    accept_plan(result, lambda plan: plan.pop("tables"))
    assert workflow.generate_only(tmp_path)["stage"] == "content_plan_response"
    assert not list(Path(result["workspace"]).glob("hermes/requests/*.json"))


def test_planned_tables_visits_titles_and_optional_sections_are_literal():
    source = json.loads((ROOT / "tests/fixtures/prospective-acceptance-source.json").read_text())
    bundle = contracts.contracted_template_bundle(ROOT, source)
    req = content_plan.request_payload(ROOT, source, bundle, {})
    plan = plan_response(req)["plan"]
    plan["docx_fields"]["protocol"]["AI_shortTitle"] = "An explicit title ending in of"
    plan["docx_fields"]["protocol"]["sampleSize"] = "90 records; not participants"
    plan["tables"]["schedule-of-assessments"].update(rows=[["Activity", "Historical", "Call"], ["Keep this safety note", "", "X"]], supplemental_notes=["Keep this safety note"])
    plan["visits"] = [{"visit": "Historical staff review", "timing": "After consent", "display_visit_number": "H", "CRFnumber": ""}]
    planned = content_plan.with_plan(source, plan)
    assert rendering.render_fields(planned, {}) == plan["docx_fields"]["protocol"]
    assert contracts.normalized_visit_records(planned) == plan["visits"]
    assert contracts.protocol_table_contracts(planned) == plan["tables"]
    section = next(s for s in contracts.protocol_contract("Prospective") if s.role == "source")
    assert not contracts.section_applies(planned, section)
    plan["source_sections"][section.section_id]["include"] = True
    assert contracts.section_applies(content_plan.with_plan(source, plan), section)


def test_xml_serializes_ai_values_without_cohort_or_eligibility_inference(tmp_path):
    source = json.loads((ROOT / "tests/fixtures/prospective-acceptance-source.json").read_text())
    bundle = contracts.contracted_template_bundle(ROOT, source)
    req = content_plan.request_payload(ROOT, source, bundle, {})
    plan = plan_response(req)["plan"]
    plan["xml"]["scalars"]["eligibility/criteria/textblock"] = "AI eligibility, exactly."
    plan["xml"]["scalars"]["brief_title"] = "AI title & literal <value>"
    for tag in ("arm_group", "intervention", "location", "primary_outcome"):
        plan["xml"]["repeated"][tag] = [{p: "" for p in req["xml_schema"]["repeated"][tag]}]
    plan["xml"]["repeated"]["arm_group"][0]["arm_group_label"] = "Explicit cohort"
    output = tmp_path / "study.xml"
    prs_xml.generate(ROOT / bundle["prs_authority"]["generation_template"]["path"], output,
                     content_plan.with_plan(source, plan), {}, check_output=False)
    root = ET.parse(output).getroot()
    assert root.find('.//criteria/textblock').text == "AI eligibility, exactly."
    assert root.find('.//brief_title').text == "AI title & literal <value>"
    assert root.find('.//arm_group/arm_group_label').text == "Explicit cohort"
    assert len(root.findall('.//arm_group')) == 1


@pytest.mark.parametrize('family', ['Advarra', 'Sterling'])
def test_native_documents_use_planned_consent_and_optional_sections(tmp_path, monkeypatch, family):
    source = json.loads((ROOT / "tests/fixtures/prospective-acceptance-source.json").read_text())
    source['meta']['icf_template'] = family
    # Facts under other fields still reach the AI; no canonical trigger is necessary.
    source['ethics'] = {'genetic_information': True, 'gina_section_required': True}
    source['procedures']['termination'] = None
    bundle = contracts.contracted_template_bundle(ROOT, source)
    request = content_plan.request_payload(ROOT, source, bundle, {})
    plan = plan_response(request)['plan']
    plan['source_sections']['endpoint-criteria.termination'] = {
        'include': True, 'blocks': [{'text': 'The AI-defined termination rule.', 'bullet': False}]}
    model = {'protocol': [{'section_id': 'endpoint-criteria.completion', 'paragraphs': [
        {'text': 'See Section 18.5. Value 18.5 is unchanged.'}], 'lists': []}], 'icf': {}, 'prs': {}}
    if family == 'Sterling':
        plan['extra_icf_sections'] = [{'heading': 'GENETIC INFORMATION NONDISCRIMINATION ACT',
            'before': 'QUESTIONS', 'blocks': [{'text': text, 'bullet': False} for text in
                next(c for c in contracts.sterling_clause_contract()['clauses'] if c['clause_id']=='sterling.privacy.gina')['authorized_texts']]}]
    def legacy(*args, **kwargs):
        pytest.fail('The planned workflow called a legacy clinical decision helper.')
    for name in ('_insert_sterling_conditional_clauses', '_preserve_sterling_registry_statement',
                 '_normalize_source_bound_shell', '_normalize_advarra_contact_sections',
                 '_normalize_sterling_retained_sections', 'repair_sterling_clause_placement'):
        monkeypatch.setattr(rendering, name, legacy)
    report = rendering.render_documents(ROOT, tmp_path, content_plan.with_plan(source, plan), model, audit_output=False)
    assert report['status'] == 'passed'
    doc = Document(tmp_path / 'candidate/protocol.docx')
    visible = '\n'.join(p.text for p in doc.paragraphs)
    assert 'The AI-defined termination rule.' in visible
    assert 'See Section 18.4. Value 18.5 is unchanged.' in visible
    assert any(p.text.startswith('18.3. Patient Termination') for p in doc.paragraphs)
    if family == 'Sterling':
        paragraphs = [p.text for p in Document(tmp_path / 'candidate/icf.docx').paragraphs]
        heading = paragraphs.index('GENETIC INFORMATION NONDISCRIMINATION ACT')
        assert heading < paragraphs.index('QUESTIONS')
        for block in plan['extra_icf_sections'][0]['blocks']:
            assert paragraphs.count(block['text']) == 1
