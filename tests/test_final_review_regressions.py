"""Final independent review reproductions, preserving approved source data."""
import copy
import json
from pathlib import Path

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import contracts
import rendering
import quality
from test_quality_gate_corrections import toc_docx, write_pdf
from pypdf import PdfWriter
from pypdf.generic import NameObject, DecodedStreamObject


def test_dual_visit_sources_preserve_table_identities_and_baseline_only_consent():
    root = Path(__file__).resolve().parents[1]
    source = json.loads((root / 'tests/fixtures/prospective-acceptance-source.json').read_text())
    before = copy.deepcopy(source)
    visits = contracts.normalized_visit_records(source)
    assert [v['visitNumber'] for v in visits] == ['1', '2', '3']
    assert [v['CRFnumber'] for v in visits] == ['BL', 'M1', 'M3']
    matrix = contracts.protocol_table_contracts(source)['schedule-of-assessments']['rows']
    assert matrix[1] == ['Activity', 'Visit 1', 'Visit 2', 'Visit 3']
    assert next(row for row in matrix if row[0] == 'Consent') == ['Consent', 'X', '', '']
    doc = Document()
    table = doc.add_table(rows=2, cols=4)
    for cell, label in zip(table.rows[0].cells, ['Number', 'Visit', 'Timing', 'CRF']):
        cell.text = label
    for cell in table.rows[1].cells:
        cell.text = '{AI_visitName}'
    rendering._visit_rows(doc, source)
    assert [row.cells[0].text for row in table.rows[1:]] == ['1', '2', '3']
    assert [row.cells[3].text for row in table.rows[1:]] == ['BL', 'M1', 'M3']
    assert source == before


@pytest.mark.parametrize('table_id, schedule_id, schedule_name', [
    ('01', '1', 'Baseline'), ('A ', 'A', 'Baseline'), ('1', '1', 'Final')])
def test_conflicting_cross_source_visit_ids_are_reported_verbatim(table_id, schedule_id, schedule_name):
    source = {'meta': {'study_type': 'Prospective', 'icf_template': 'Advarra'}, 'procedures': {
        'visit_schedule_table': [{'visitNumber': table_id, 'visitName': 'Baseline', 'visitWindow': 'Day 0'}],
        'visit_schedule': [{'visitNumber': schedule_id, 'visit': schedule_name, 'timing': 'Day 0', 'procedures': ['Consent']}]}}
    before = copy.deepcopy(source)
    visits = contracts.normalized_visit_records(source)
    assert [v['visitNumber'] for v in visits] == [table_id, schedule_id]
    findings = contracts.input_findings(source)
    assert any('conflict' in f['issue'].lower() and 'visit' in f['field'] for f in findings)
    assert source == before


def test_additional_schedule_visit_is_preserved_without_cross_visit_procedures():
    source = {'procedures': {
        'visit_schedule_table': [{'visitNumber': '01', 'visitName': 'Baseline', 'CRFnumber': 'BL'}],
        'visit_schedule': [{'visit': 'Baseline', 'procedures': ['Consent']},
                           {'visitNumber': 'F ', 'visit': 'Final', 'procedures': ['Pain']}]}}
    visits = contracts.normalized_visit_records(source)
    assert [v['visitNumber'] for v in visits] == ['01', 'F ']
    assert [v['procedures'] for v in visits] == [['Consent'], ['Pain']]
    doc = Document()
    table = doc.add_table(rows=1, cols=4)
    for cell in table.rows[0].cells:
        cell.text = '{AI_visitName}'
    rendering._visit_rows(doc, source)
    assert [row.cells[0].text for row in table.rows] == ['01', 'F ']


def test_inline_body_image_is_not_a_blank_page(tmp_path):
    docx, pdf = tmp_path / 'image.docx', tmp_path / 'image.pdf'
    Document().save(docx)
    writer = PdfWriter()
    page = writer.add_blank_page(width=612, height=792)
    stream = DecodedStreamObject()
    stream.set_data(b'q 200 0 0 200 100 300 cm BI /W 1 /H 1 /CS /RGB /BPC 8 ID \xff\x00\x00 EI Q')
    page[NameObject('/Contents')] = writer._add_object(stream)
    writer.write(pdf)
    assert quality._blank_pdf_pages(pdf, docx_path=docx) == []


@pytest.mark.parametrize('visible, expected', [('1. PURPOSE 9', 'mismatch'), ('Unreadable TOC entry', 'unknown')])
def test_final_toc_requires_pdf_visible_destination(tmp_path, visible, expected):
    docx, pdf = tmp_path / 'toc.docx', tmp_path / 'toc.pdf'
    toc_docx(docx)
    write_pdf(pdf, [[('TABLE OF CONTENTS', 72, 700), ('TABLE OF CONTENTS 1', 72, 680), (visible, 72, 660)],
                    [('1. PURPOSE', 72, 700)]])
    report = quality.audit_final_toc_destinations(docx, pdf)
    row = next(row for row in report['destinations'] if row['heading'] == '1. PURPOSE')
    assert row['status'] == expected
    assert report['status'] == expected


def test_same_entity_funder_keeps_explicit_role():
    doc = Document()
    p = doc.add_paragraph('{sponsortName}\n{sponsortAdress}{fundingSourceName}{fundingSourceAdress}')
    fields = rendering.render_fields({'parties': {
        'sponsor': {'name': 'Sponsor Foundation', 'address': '10 Main St, Country'},
        'funding_source': {'name': 'Sponsor Foundation'}}}, {})
    rendering._replace_paragraph(p, fields)
    assert 'Funding source: Sponsor' in p.text
    assert p.text.count('Sponsor Foundation') == 1
    assert p.text.count('10 Main St, Country') == 1


@pytest.mark.parametrize('kind', ['simple-field', 'underline', 'page-break', 'legacy-image'])
def test_preference_cleanup_preserves_meaningful_empty_paragraphs(kind):
    doc = Document()
    doc.add_paragraph('Check your preference below:')
    doc.add_paragraph('Yes, inform my primary care physician')
    p = doc.add_paragraph()
    if kind == 'simple-field':
        field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), 'DOCPROPERTY PatientInitials')
        run = OxmlElement('w:r'); text = OxmlElement('w:t'); text.text = 'AB'
        run.append(text); field.append(run); p._p.append(field)
    elif kind == 'underline':
        p.add_run('      ').underline = True
    elif kind == 'page-break':
        p.paragraph_format.page_break_before = True
    else:
        p.add_run()._r.append(OxmlElement('w:pict'))
    cosmetic = doc.add_paragraph()
    doc.add_paragraph('No, do not inform my primary care physician')
    rendering._normalize_icf_preferences(doc)
    assert p._p.getparent() is not None
    assert cosmetic._p.getparent() is None


@pytest.mark.parametrize('text', [
    'Adverse events are reviewed only at the final visit rather than at each contact.',
    'Device function is checked at each contact;adverse events are reviewed only at the final visit.',
    'Device function rather than adverse events is reviewed at each contact.',
])
def test_final_only_safety_does_not_authorize_each_contact(text):
    source = {'meta': {'study_type': 'Prospective'}, 'procedures': {'visit_schedule': [
        {'visit': 'Baseline', 'procedures': ['Consent']},
        {'visit': 'Final', 'procedures': ['Pain rating']}]}, 'safety': {'monitoring': text}}
    before = copy.deepcopy(source)
    table = contracts.protocol_table_contracts(source)['schedule-of-assessments']
    assert table['each_contact_safety_evidence'] == []
    assert text in table['supplemental_notes']
    assert source == before


def test_multiple_relative_durations_are_not_assigned_to_first_event():
    source = {'study': {'timeline': 'Interim assessment 4 weeks after baseline with final contact 12 weeks after baseline.'},
              'procedures': {'visit_schedule': [
                  {'visit': 'Baseline', 'timing': 'Day 0'},
                  {'visit': 'Interim', 'timing': 'Week 4 after baseline'},
                  {'visit': 'Final', 'timing': 'Week 12 after baseline'}]}}
    assert contracts.timeline_findings(source) == []
