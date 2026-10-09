"""Bibliographic identifiers must not become clinical quantity requirements."""
import copy
import json

import pytest
from docx import Document

import contracts
import drafting
import quality
import rendering
from test_shared_output_projection import ROOT, COMBINATIONS, source_for


def assert_background_requires_quantity_review(findings):
    rows = [row for row in findings if row.get('clause_id') == 'sterling.background.context']
    assert rows and all(row.get('publication_disposition') == 'warning' for row in rows)
    assert any(value.get('missing_numbers') for row in rows
               for diagnostic in row.get('evidence_diagnostics', [])
               for value in diagnostic.get('unobserved_values', []))



def reference(branch='Prospective', family='Sterling'):
    source = source_for(branch, family)
    source['study']['background'] = 'Earlier research described improved visual outcomes.3'
    source['study']['references'] = [
        'Alpha AB. Outcomes. Example Journal. 2024;1:1-9.',
        'Beta CD. Follow-up. Example Journal. 2025;2:10-19.',
        'Gamma EF. Vision. Example Journal. 2025;3:20-29.',
    ]
    return source


@pytest.mark.parametrize('branch', ['Prospective', 'Ambispective'])
def test_sterling_background_does_not_require_an_approved_citation_numeral(branch):
    source = reference(branch)
    original = copy.deepcopy(source)
    document = Document()
    document.add_heading('BACKGROUND', level=1)
    document.add_paragraph('Earlier research described improved visual outcomes.')
    findings = quality.validate_sterling_clause_contract(document, source)['findings']
    assert not [f for f in findings if f.get('clause_id') == 'sterling.background.context'
                and f.get('publication_disposition') != 'warning']
    assert source == original


@pytest.mark.parametrize('marker', ['.1,2', '.3', ' [reference 3]', ' (reference 3)', ' (references 1–3)'])
def test_known_citation_styles_are_metadata_in_final_background(marker):
    source = reference()
    source['study']['background'] = 'Earlier research described improved visual outcomes' + marker
    document = Document()
    document.add_heading('BACKGROUND', level=1)
    document.add_paragraph('Earlier research described improved visual outcomes.')
    findings = quality.validate_sterling_clause_contract(document, source)['findings']
    assert not [f for f in findings if f.get('clause_id') == 'sterling.background.context'
                and f.get('publication_disposition') != 'warning']


@pytest.mark.parametrize('quantity', ['3 months', '0.3 mL', '66 cm', '120 participants'])
def test_unobserved_clinical_quantity_requires_final_background_review(quantity):
    source = reference()
    source['study']['background'] = f'Earlier research described improved visual outcomes at {quantity}.3'
    document = Document()
    document.add_heading('BACKGROUND', level=1)
    document.add_paragraph('Earlier research described improved visual outcomes.')
    findings = quality.validate_sterling_clause_contract(document, source)['findings']
    # Missing token evidence must reach review; it cannot establish clinical
    # equivalence or be silently erased by bibliography normalization.
    assert_background_requires_quantity_review(findings)


@pytest.mark.parametrize('marker', ['.3', ' [reference 3]', ' (reference 3)'])
def test_citation_cannot_satisfy_a_missing_clinical_quantity(marker):
    source = reference()
    source['study']['background'] = 'Earlier research described improved visual outcomes at 3 months.'
    document = Document()
    document.add_heading('BACKGROUND', level=1)
    document.add_paragraph('Earlier research described improved visual outcomes' + marker)
    findings = quality.validate_sterling_clause_contract(document, source)['findings']
    # Missing token evidence must reach review; it cannot establish clinical
    # equivalence or be silently erased by bibliography normalization.
    assert_background_requires_quantity_review(findings)


def test_unresolved_citation_number_is_not_silently_discarded():
    source = reference()
    source['study']['background'] = 'Earlier research described improved visual outcomes.9'
    document = Document()
    document.add_heading('BACKGROUND', level=1)
    document.add_paragraph('Earlier research described improved visual outcomes.')
    findings = quality.validate_sterling_clause_contract(document, source)['findings']
    # Missing token evidence must reach review; it cannot establish clinical
    # equivalence or be silently erased by bibliography normalization.
    assert_background_requires_quantity_review(findings)


@pytest.mark.parametrize('bibliography', [
    {'study': {'references': ['3. Gamma EF. Vision. Example Journal. 2025;3:20-29.']}},
    {'study': {'references': '3. Gamma EF. Vision. Example Journal. 2025;3:20-29.'}},
    {'study': {'references': [{'text': '3. Gamma EF. Vision. Example Journal. 2025;3:20-29.'}]}},
    {'references': ['3. Gamma EF. Vision. Example Journal. 2025;3:20-29.']},
    {'references': '3. Gamma EF. Vision. Example Journal. 2025;3:20-29.'},
    {'study': {'references': ''}, 'references': '3. Gamma EF. Vision. Example Journal. 2025;3:20-29.'},
    {'study': {'background': 'Earlier research described improved visual outcomes.3\nReferences:\n3. Gamma EF. Vision. Example Journal. 2025;3:20-29.'}},
    {'study': {'background': 'Earlier research described improved visual outcomes.3\n3. Gamma EF. Vision. Example Journal. 2025.\ndoi:10.0000/example'}},
])
def test_explicit_bibliography_identity_is_shared_with_icf_assembly(tmp_path, bibliography):
    source = reference()
    source['study'].pop('references')
    source['study'].update(bibliography.get('study', {}))
    if 'references' in bibliography:
        source['references'] = bibliography['references']
    document = Document()
    document.add_heading('BACKGROUND', level=1)
    document.add_paragraph('Earlier research described improved visual outcomes.')
    findings = quality.validate_sterling_clause_contract(document, source)['findings']
    assert not [f for f in findings if f.get('clause_id') == 'sterling.background.context'
                and f.get('publication_disposition') != 'warning']
    rendering.render_documents(ROOT, tmp_path, source, {'protocol': [], 'icf': {
        'icf.background': {'paragraphs': [{'text': source['study']['background']}], 'lists': []}}, 'prs': {}},
        artifact_names={'icf'})
    prose = '\n'.join(p.text for p in Document(tmp_path / 'candidate/icf.docx').paragraphs)
    assert 'Earlier research described improved visual outcomes.' in prose
    assert 'Earlier research described improved visual outcomes.3' not in prose


def test_embedded_bibliography_citation_cannot_replace_a_clinical_quantity():
    source = reference()
    source['study'].pop('references')
    source['study']['background'] = (
        'Earlier research described improved visual outcomes at 3 months.\nReferences:\n'
        '3. Gamma EF. Vision. Example Journal. 2025;3:20-29. doi:10.0000/example')
    document = Document()
    document.add_heading('BACKGROUND', level=1)
    document.add_paragraph('Earlier research described improved visual outcomes.3')
    findings = quality.validate_sterling_clause_contract(document, source)['findings']
    # Missing token evidence must reach review; it cannot establish clinical
    # equivalence or be silently erased by bibliography normalization.
    assert_background_requires_quantity_review(findings)


def test_citation_range_requires_every_bibliographic_identity():
    source = reference()
    source['study']['references'] = ['1. Alpha AB. Outcomes.', '3. Gamma EF. Vision.']
    source['study']['background'] = 'Earlier research described improved visual outcomes (references 1–3).'
    document = Document()
    document.add_heading('BACKGROUND', level=1)
    document.add_paragraph('Earlier research described improved visual outcomes.')
    findings = quality.validate_sterling_clause_contract(document, source)['findings']
    # Missing token evidence must reach review; it cannot establish clinical
    # equivalence or be silently erased by bibliography normalization.
    assert_background_requires_quantity_review(findings)


@pytest.mark.parametrize('branch,family', COMBINATIONS[1:])
def test_icf_assembly_omits_known_citations_and_keeps_real_quantities(tmp_path, branch, family):
    source = reference(branch, family)
    original = copy.deepcopy(source)
    target = 'icf.background' if family == 'Sterling' else 'icf.study-purpose'
    text = 'Earlier research measured 3 months, 0.3 mL, and 66 cm.3 Evidence [reference 2] (reference 1). Unknown.9'
    rendering.render_documents(ROOT, tmp_path, source, {'protocol': [], 'icf': {
        target: {'paragraphs': [{'text': text}], 'lists': []}}, 'prs': {}},
        artifact_names={'protocol', 'icf'})
    document = Document(tmp_path / 'candidate/icf.docx')
    prose = '\n'.join(p.text for p in document.paragraphs)
    assert 'Earlier research measured 3 months, 0.3 mL, and 66 cm. Evidence. Unknown.9' in prose
    protocol = '\n'.join(p.text for p in Document(tmp_path / 'candidate/protocol.docx').paragraphs)
    assert all(citation in protocol for citation in source['study']['references'])
    assert source == original


def test_square_bracket_clinical_interval_is_preserved_and_its_loss_requires_review(tmp_path):
    source = reference()
    source['study']['background'] = 'Prior research reported a mean difference of 2 points with confidence interval [1–3].'
    target = {'paragraphs': [{'text': source['study']['background']}], 'lists': []}
    rendering.render_documents(ROOT, tmp_path, source, {'protocol': [], 'icf': {'icf.background': target}, 'prs': {}},
        artifact_names={'icf'})
    document = Document(tmp_path / 'candidate/icf.docx')
    assert source['study']['background'] in '\n'.join(p.text for p in document.paragraphs)
    for paragraph in document.paragraphs:
        if '[1–3]' in paragraph.text:
            paragraph.text = paragraph.text.replace('[1–3]', '')
    findings = quality.validate_sterling_clause_contract(document, source)['findings']
    # Missing token evidence must reach review; it cannot establish clinical
    # equivalence or be silently erased by bibliography normalization.
    assert_background_requires_quantity_review(findings)


@pytest.mark.parametrize('branch,family', COMBINATIONS)
def test_drafting_accepts_clinical_background_without_its_bibliographic_identifiers(tmp_path, branch, family):
    source = reference(branch, family)
    original = copy.deepcopy(source)
    targets = ['introduction']
    if family:
        targets.append('icf.background' if family == 'Sterling' else 'icf.study-purpose')
        if family == 'Sterling':
            targets.append('icf.key-information-summary')
    for target in targets:
        batch = next(b for b in contracts.batch_plan(branch, family) if target in b.section_ids)
        path = drafting.create_drafting_request(
            repo_root=ROOT, revision_dir=tmp_path, revision_id='citations', reference=source,
            batch=batch, target_ids=[target], attempts={target: 1}, wave='initial')
        request = json.loads(path.read_text())
        response = drafting.recorded_acceptance_response(request)
        if target == 'introduction':
            response['section_results'][0]['paragraphs'] = [{
                'text': 'Earlier research described improved visual outcomes.',
                'evidence_refs': ['source:study.background'], 'boilerplate_refs': []}]
        for paragraph in response['section_results'][0]['paragraphs']:
            paragraph['text'] = paragraph['text'].replace('outcomes.3', 'outcomes.')
        accepted, findings = drafting.validate_response(request, response)
        assert accepted and any(d['section_id'] == target for d in accepted['drafts']), findings
    assert source == original
