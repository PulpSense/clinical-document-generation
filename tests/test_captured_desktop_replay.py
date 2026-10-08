"""Opt-in private archive replay through the public Desktop operation.

Set CLINICAL_BLOCKED_REPLAY_ARCHIVES to a JSON list of archive paths. Archive
clinical payloads stay outside Git; only selected evidence members are read.
Recorded drafting plus controlled verification prove the lifecycle, not fresh
independent clinical judgment or Linux host compatibility.
"""
from __future__ import annotations

import copy
import hashlib
import json
import os
import time
from pathlib import Path, PurePosixPath
import zipfile
from xml.etree import ElementTree as ET

from docx import Document
import pytest

import drafting
import quality
import workflow
from test_handoff_quality import acceptance_verification


ARCHIVES = json.loads(os.environ.get("CLINICAL_BLOCKED_REPLAY_ARCHIVES", "[]"))


def _safe_member(archive, name):
    path = PurePosixPath(name)
    if path.is_absolute() or ".." in path.parts or "\\" in name:
        raise ValueError("Replay evidence member is unsafe.")
    return archive.read(name)


def _recorded_archive(run, archive_path):
    """Retain immutable source approval; rebind only generation request envelopes."""
    digest = hashlib.sha256(archive_path.read_bytes()).hexdigest()
    with zipfile.ZipFile(archive_path) as archive:
        reference = json.loads(_safe_member(archive, "reference/study.reference.json"))
        operation = json.loads(_safe_member(archive, "logs/desktop-operation.json"))
        reference.pop("generation", None)
        approval = copy.deepcopy(reference["approval"])
        revision = run / "revisions" / approval["revision_id"]
        prefix = revision.relative_to(run).as_posix() + "/"
        for name in (
            prefix + "approved-reference.json", prefix + "approved-source.md",
            reference["source"]["source_of_truth_file"], approval["review_file"],
        ):
            payload = _safe_member(archive, name)
            path = run / name
            path.parent.mkdir(parents=True, exist_ok=True)
            path.write_bytes(payload)
        reference_path = run / "reference/study.reference.json"
        reference_path.parent.mkdir(parents=True, exist_ok=True)
        reference_path.write_text(json.dumps(reference))
        sections = {}
        narratives = {}
        provenance = {}
        names = sorted(name for name in archive.namelist()
                       if name.startswith(prefix + "hermes/responses/") and name.endswith(".json"))
        for name in names:
            response = json.loads(_safe_member(archive, name))
            request_name = prefix + "hermes/requests/" + Path(name).name
            request = json.loads(_safe_member(archive, request_name))
            for result in response.get("section_results", []):
                section_id = result["section_id"]
                rank = (int(request.get("attempts", {}).get(section_id, 1)), name)
                if section_id not in provenance or rank > provenance[section_id]:
                    sections[section_id] = copy.deepcopy(result)
                    provenance[section_id] = rank
            for key, result in response.get("narrative", {}).items():
                narratives[key] = copy.deepcopy(result)
        assert "analysis-plan.considerations" in sections
    original_operation_id = operation["operation_id"]
    operation["operation_id"] = "controlled-captured-replay"
    state_path = run / "logs/desktop-operation-controlled-captured-replay.json"
    state_path.parent.mkdir(parents=True, exist_ok=True)
    state_path.write_text(json.dumps(operation))
    operation["original_operation_id"] = original_operation_id
    return reference, approval, revision, sections, narratives, digest, operation


def replay_captured_desktop(run, archive_path):
    reference, approval, revision, sections, narratives, digest, original_operation = _recorded_archive(run, archive_path)
    observed = {"draft_handoffs": 0, "review_cautions": [], "considerations_attempts": 0}

    def runner(handoffs, _remaining):
        for handoff in handoffs:
            request_path = revision / handoff["request_path"]
            request = json.loads(request_path.read_text())
            if "verification-requests" in handoff["request_path"]:
                if request["task"] == "clinical_content_verification":
                    observed["review_cautions"].extend(request["lexical_review_cautions"])
                workflow._save_verification_response(revision, request_path, acceptance_verification)
                continue
            observed["draft_handoffs"] += 1
            response = drafting.response_template(request, model_id="RecordedArchiveReplay/v1")
            if request["task"] == "prs_narrative_drafting":
                keys = {"prs.brief-summary": "brief_summary", "prs.detailed-description": "detailed_description"}
                response["narrative"] = {keys[item["section_id"]]: copy.deepcopy(narratives[keys[item["section_id"]]])
                                         for item in request["section_contracts"]}
            else:
                response["section_results"] = [copy.deepcopy(sections[item["section_id"]])
                                                for item in request["section_contracts"]]
                observed["considerations_attempts"] += sum(
                    item["section_id"] == "analysis-plan.considerations" for item in request["section_contracts"]
                )
            path = revision / request["response_path"]
            path.parent.mkdir(parents=True, exist_ok=True)
            path.write_text(json.dumps(response))

    replay_started = time.monotonic()
    recorded_elapsed = original_operation["result"]["elapsed_seconds"]

    def historical_wall_clock():
        return original_operation["started_at_epoch"] + recorded_elapsed + (time.monotonic() - replay_started)

    result = workflow.run_desktop_operation(
        run, handoff_runner=runner, opener=lambda path: Path(path).read_bytes(),
        operation_id="controlled-captured-replay",
        release_identity={"package_fingerprint": "controlled-checkout-replay"},
        wall_clock=historical_wall_clock, budget_seconds=original_operation["budget_seconds"],
    )
    # Keep assertion failures private and compact: never dump approved clinical payloads.
    assert result["status"] == "passed", (result.get("stage"), [f.get("field") for f in result.get("findings", [])])
    assert result["stage"] == "desktop_delivery"
    assert result["delivery"]["confirmed"] is True
    assert sorted(result["client_outputs"]) == ["output/icf.docx", "output/protocol.docx", "output/study.xml"]
    latest = json.loads((run / "reference/study.reference.json").read_text())
    assert latest["approval"] == approval
    assert {k: v for k, v in latest.items() if k not in {"approval", "generation"}} == {
        k: v for k, v in reference.items() if k != "approval"
    }
    state = json.loads((run / "logs/desktop-operation-controlled-captured-replay.json").read_text())
    assert state["deadline_at_epoch"] - state["started_at_epoch"] == 2700
    assert state["started_at_epoch"] == original_operation["started_at_epoch"]
    assert state["deadline_at_epoch"] == original_operation["deadline_at_epoch"]
    manifest = json.loads((revision / "delivery-manifest.json").read_text())
    assert manifest["quality"]["status"] == "passed"
    assert observed["considerations_attempts"] == 1
    assert observed["review_cautions"]
    assert manifest["quality"]["resolved_evidence_cautions"]
    assert any(c.get("code") == "lexical-evidence-uncertainty" for c in result["desktop_reply"]["cautions"])
    assert hashlib.sha256(archive_path.read_bytes()).hexdigest() == digest
    for item in manifest["client_outputs"]:
        path = run / item["path"]
        assert path.read_bytes() == (revision / "candidate" / path.name).read_bytes()
        assert quality.sha256_file(path) == item["sha256"]
    protocol = Document(run / "output/protocol.docx")
    visible = "\n".join(p.text for p in protocol.paragraphs) + "\n" + "\n".join(
        cell.text for table in protocol.tables for row in table.rows for cell in row.cells
    )
    assert reference["study"]["title"] in visible
    assert str(reference["population"]["sample_size"]) in visible
    for endpoint in reference["endpoints"]["primary"]:
        assert str(endpoint.get("label") or endpoint.get("name")) in visible
        assert str(endpoint["time_point"]) in visible
    assert all(p["text"] in visible for p in sections["analysis-plan.considerations"]["paragraphs"])
    assert "lexical-evidence-uncertainty" not in visible
    assert "lexical-evidence-uncertainty" not in "\n".join(p.text for p in Document(run / "output/icf.docx").paragraphs)
    assert next(ET.parse(run / "output/study.xml").getroot().iter("enrollment")).text
    return {"assurance": "recorded-archive-controlled-desktop-replay", "archive_sha256": digest,
            "revision_id": approval["revision_id"], "source_sha256": approval["source_sha256"],
            "approval_retained": True, "complete_delivery": True,
            "exact_reviewed_bytes_delivered": True, "deadline_seconds": 2700, "clock": "historical resume wall clock; measured monotonic elapsed",
            "original_deadline_retained": True,
            "verification": "controlled; not fresh independent clinical or visual approval",
            "elapsed_seconds": result["elapsed_seconds"],
            "recorded_prior_elapsed_seconds": recorded_elapsed,
            "replay_elapsed_seconds": round(time.monotonic() - replay_started, 3),
            "warnings": len(result["desktop_reply"]["cautions"])}


@pytest.mark.skipif(not ARCHIVES, reason="Private captured archives require explicit opt-in.")
@pytest.mark.parametrize("archive_path", ARCHIVES, ids=[f"captured-{i + 1}" for i in range(len(ARCHIVES))])
def test_captured_run_reaches_public_desktop_delivery(tmp_path, governed_pdfium, archive_path):
    evidence = replay_captured_desktop(tmp_path / "private-run", Path(archive_path))
    (tmp_path / "replay-evidence.json").write_text(json.dumps(evidence, indent=2))
