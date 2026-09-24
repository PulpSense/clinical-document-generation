"""Open Sterling template comments must govern retained text deterministically."""

import json
from pathlib import Path

from docx import Document

from contracts import icf_contract
from quality import create_verification_requests, validate_sterling_clause_contract
from rendering import render_documents


ROOT = Path(__file__).resolve().parents[1]
REGISTRY_START = "A description of this clinical trial will be available"


def _source():
    source = json.loads((ROOT / "tests/fixtures/prospective-acceptance-source.json").read_text())
    source["meta"]["icf_template"] = "Sterling"
    return source


def _render(tmp_path, source):
    render_documents(
        ROOT, tmp_path, source,
        {"protocol": [], "icf": {"icf.procedures": {"paragraphs": [{"text": "Approved study visits and assessments are described here."}], "lists": []}}, "prs": {}},
        artifact_names={"icf"},
    )
    return Document(tmp_path / "candidate/icf.docx")


def test_commented_fixed_sections_are_not_sent_for_ai_drafting():
    draft_ids = {section.section_id for section in icf_contract("Prospective", "Sterling")}
    assert {"icf.authorization-introduction", "icf.information", "icf.voluntary-participation"}.isdisjoint(draft_ids)


def test_fixed_introduction_and_new_information_keep_template_wording(tmp_path):
    authority = Document(ROOT / "assets/client-templates/reference/sterling-icf-reference.docx")
    document = _render(tmp_path, _source())
    visible = [paragraph.text for paragraph in document.paragraphs]
    assert authority.paragraphs[20].text in visible
    assert authority.paragraphs[22].text in visible
    assert authority.paragraphs[114].text in visible


def test_sterling_owned_front_matter_merge_fields_remain_unfilled(tmp_path):
    document = _render(tmp_path, _source())
    visible = [paragraph.text for paragraph in document.paragraphs[:20]]
    for field in (
        "«Protocol_Title»", "«Protocol_No»", "«First_Name»", "«Company_Name»",
        "«Address»", "«City_State_ZIP»", "«Telephone»", "«Sponsor»",
    ):
        assert any(field in paragraph for paragraph in visible)


def test_registry_statement_is_exact_template_text_when_approved(tmp_path):
    source = _source()
    source.setdefault("regulatory", {}).setdefault("prs", {})["participant_registry_disclosure"] = True
    authority = Document(ROOT / "assets/client-templates/reference/sterling-icf-reference.docx")
    exact = next(paragraph.text for paragraph in authority.paragraphs if paragraph.text.startswith(REGISTRY_START))

    document = _render(tmp_path, source)
    assert [paragraph.text for paragraph in document.paragraphs if paragraph.text.startswith(REGISTRY_START)] == [exact]
    findings = validate_sterling_clause_contract(document, source)["findings"]
    assert not any(item["clause_id"] == "sterling.privacy.registry-disclosure" for item in findings)


def test_registry_statement_is_absent_without_approved_trigger(tmp_path):
    document = _render(tmp_path, _source())
    text = "\n".join(paragraph.text for paragraph in document.paragraphs).casefold()
    assert REGISTRY_START.casefold() not in text
    assert "clinical trial may be registered on publicly accessible databases" not in text


def test_registry_statement_is_absent_when_source_says_no(tmp_path):
    source = _source()
    source.setdefault("regulatory", {}).setdefault("prs", {})["participant_registry_disclosure"] = False
    document = _render(tmp_path, source)
    assert not any(paragraph.text.startswith(REGISTRY_START) for paragraph in document.paragraphs)
    findings = validate_sterling_clause_contract(document, source)["findings"]
    assert not any(item["clause_id"] == "sterling.privacy.registry-disclosure" for item in findings)


def test_content_reviewer_knows_sterling_merge_fields_are_intentional(tmp_path):
    request_path = create_verification_requests(tmp_path, _source(), {"artifacts": []})[0]
    instructions = json.loads(request_path.read_text())["instructions"]
    assert "Sterling IRB owns the consent front-matter merge fields" in instructions
    assert "must remain unfilled" in instructions
