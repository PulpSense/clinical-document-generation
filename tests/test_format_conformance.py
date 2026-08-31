import json
from pathlib import Path

import pytest
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.shared import Inches, Pt

import quality
import workflow


ROOT = Path(__file__).resolve().parents[1]


def _gate_owner(gate_id):
    matrix = quality.load_format_conformance_matrix(ROOT)
    return next(item["retry_owner"] for item in matrix["gate_sequence"] if item["gate_id"] == gate_id)


def test_format_conformance_matrix_is_complete_hash_addressed_and_ordered():
    matrix = quality.load_format_conformance_matrix(ROOT)

    assert matrix["schema_version"] == "clinical-format-conformance/v1"
    assert [case["case_id"] for case in matrix["cases"]] == [
        "retrospective-protocol",
        "prospective-advarra",
        "prospective-sterling",
        "ambispective-advarra",
        "ambispective-sterling",
    ]
    assert [gate["gate_id"] for gate in matrix["gate_sequence"]] == list(quality.GOVERNED_GATE_SEQUENCE)
    assert matrix["finding_contract"]["required"] == [
        "code", "target", "evidence_sha256", "retry_owner", "terminal_status",
    ]
    required_checks = {
        "section_properties", "margins", "orientation", "styles", "numbering",
        "headers", "footers", "fields_toc", "page_furniture", "table_geometry",
        "signature_blocks", "consent_legal_language", "paragraph_rhythm",
        "natural_section_3_flow", "heading_cohesion", "front_matter_boundaries",
    }
    for case in matrix["cases"]:
        assert set(case["checks"]) == required_checks
        assert case["expected_outputs"] == (
            ["protocol.docx"] if case["study_type"] == "Retrospective"
            else ["icf.docx", "protocol.docx", "study.xml"]
        )
        for item in [case["fixture"], *case["baselines"]]:
            path = ROOT / item["path"]
            assert path.is_file()
            assert quality.sha256_file(path) == item["sha256"]
        approved = case["approved_output_baseline"]
        assert approved["path"].startswith("references/format-baselines/")
        approved_path = ROOT / approved["path"]
        assert quality.sha256_file(approved_path) == approved["sha256"]
        signature = json.loads(approved_path.read_text(encoding="utf-8"))
        required_signature_fields = {
            "section_geometry", "styles", "numbering", "headers_footers",
            "fields_toc", "page_furniture", "table_geometry", "signature_blocks",
            "consent_legal_placement", "paragraph_rhythm", "pagination_relations",
        }
        assert signature["artifacts"]
        assert all(set(value) == required_signature_fields for value in signature["artifacts"].values())
        for artifact in signature["artifacts"].values():
            assert artifact["section_geometry"]
            assert artifact["styles"]["count"] > 0
            assert artifact["headers_footers"]
            assert artifact["signature_blocks"]
            assert artifact["consent_legal_placement"]
            assert artifact["pagination_relations"]["numbered_body_has_no_artificial_starts"] is True
            assert artifact["pagination_relations"]["all_headings_keep_with_next"] is True

    canonical = dict(matrix)
    declared_sha256 = canonical.pop("matrix_sha256")
    assert declared_sha256 == quality.canonical_evidence_sha256(canonical)


def test_protocol_signature_treats_only_the_first_post_toc_boundary_as_governed(tmp_path):
    reference = json.loads(
        (ROOT / "references/conformance-fixtures/prospective-acceptance-source.json").read_text(
            encoding="utf-8"
        )
    )
    report = workflow.render_documents(
        ROOT, tmp_path, reference, {"protocol": [], "icf": {}, "prs": {}},
        artifact_names={"protocol"},
    )

    assert report["status"] == "passed"
    signature = quality.normalized_docx_format_signature(
        tmp_path / "candidate/protocol.docx"
    )
    pagination = signature["pagination_relations"]
    headings = pagination["headings"]
    toc_index = next(
        index for index, item in enumerate(headings)
        if "table of contents" in item["text"].casefold()
    )
    first_body = next(
        item for item in headings[toc_index + 1:]
        if item["text"][0].isdigit()
    )

    assert first_body["page_break_before"] is False
    assert first_body["effective_page_boundary"] is True
    assert pagination["first_numbered_body_paragraph"] == first_body["paragraph"]
    assert pagination["numbered_body_has_no_artificial_starts"] is True


def test_gate_ledger_is_hash_bound_and_failed_gates_are_monotonic():
    evidence = {
        gate_id: {"gate": gate_id, "artifact_sha256": str(index) * 64}
        for index, gate_id in enumerate(quality.GOVERNED_GATE_SEQUENCE, start=1)
    }
    pending = quality.build_gate_ledger(
        ROOT,
        evidence,
        attempt_id="test-root",
        statuses={"exact_byte_atomic_delivery": "pending"},
    )
    ledger = quality.advance_gate_ledger(
        ROOT,
        pending,
        gate_id="exact_byte_atomic_delivery",
        terminal_status="passed",
        evidence={"delivery": "confirmed"},
    )

    assert [record["gate_id"] for record in ledger["records"]] == list(quality.GOVERNED_GATE_SEQUENCE)
    assert all(record["terminal_status"] == "passed" for record in ledger["records"])
    assert all(len(record["evidence_sha256"]) == 64 for record in ledger["records"])
    unsigned = dict(ledger)
    declared_sha256 = unsigned.pop("ledger_sha256")
    assert declared_sha256 == quality.canonical_evidence_sha256(unsigned)
    assert quality.validate_gate_ledger(ROOT, ledger) == ledger

    bypass = json.loads(json.dumps(ledger))
    bypass["records"][2]["terminal_status"] = "blocked"
    bypass["records"][2]["findings"] = [{
        "code": "DOCX_PRS_STRUCTURE_FAILED",
        "target": "protocol.docx#section=3",
        "evidence_sha256": "f" * 64,
        "retry_owner": _gate_owner("docx_prs_structure"),
        "terminal_status": "blocked",
    }]
    bypass["records"][3]["terminal_status"] = "passed"
    unsigned = dict(bypass)
    unsigned.pop("ledger_sha256")
    bypass["ledger_sha256"] = quality.canonical_evidence_sha256(unsigned)
    with pytest.raises(ValueError, match="cannot pass after"):
        quality.validate_gate_ledger(ROOT, bypass)
    with pytest.raises(ValueError, match="must advance"):
        quality.build_gate_ledger(ROOT, evidence, attempt_id="forbidden-direct-pass")


@pytest.mark.parametrize("failed_gate", quality.GOVERNED_GATE_SEQUENCE)
def test_every_failed_gate_is_terminal_and_retained(failed_gate):
    evidence = {gate_id: {"gate": gate_id} for gate_id in quality.GOVERNED_GATE_SEQUENCE}
    failed_index = quality.GOVERNED_GATE_SEQUENCE.index(failed_gate)
    statuses = {
        gate_id: "passed" if index < failed_index else "blocked" if index == failed_index else "pending"
        for index, gate_id in enumerate(quality.GOVERNED_GATE_SEQUENCE)
    }
    finding = {
        "code": f"{failed_gate.upper()}_FAILED",
        "target": f"{failed_gate}:artifact#section=target",
        "evidence_sha256": "e" * 64,
        "retry_owner": _gate_owner(failed_gate),
        "terminal_status": "blocked",
    }
    ledger = quality.build_gate_ledger(
        ROOT,
        evidence,
        attempt_id=f"failed-{failed_gate}",
        statuses=statuses,
        findings_by_gate={failed_gate: [finding]},
    )
    with pytest.raises(ValueError, match="terminal blocked gate"):
        quality.advance_gate_ledger(
            ROOT,
            ledger,
            gate_id=failed_gate,
            terminal_status="passed",
            evidence={"replacement": True},
        )
    assert ledger["records"][failed_index]["findings"] == [finding]


def test_retry_ledger_retains_blocked_predecessor_and_findings():
    evidence = {gate_id: {"gate": gate_id} for gate_id in quality.GOVERNED_GATE_SEQUENCE}
    finding = {
        "code": "DOCX_PRS_STRUCTURE_FAILED",
        "target": "protocol.docx#section=3",
        "evidence_sha256": "f" * 64,
        "retry_owner": _gate_owner("docx_prs_structure"),
        "terminal_status": "blocked",
    }
    blocked = quality.build_gate_ledger(
        ROOT,
        evidence,
        attempt_id="revision-r1",
        statuses={
            "docx_prs_structure": "blocked",
            "exact_artifact_rendering": "pending",
            "every_page_visual_qa": "pending",
            "exact_byte_atomic_delivery": "pending",
        },
        findings_by_gate={"docx_prs_structure": [finding]},
    )

    retried = quality.retry_gate_ledger(
        ROOT,
        blocked,
        evidence,
        statuses={"exact_byte_atomic_delivery": "pending"},
    )

    assert retried["attempt_id"] == "revision-r1"
    assert retried["predecessors"][-1]["ledger_sha256"] == blocked["ledger_sha256"]
    assert retried["predecessors"][-1]["blocked_findings"] == [finding]
    assert quality.validate_gate_ledger(ROOT, retried) == retried


def test_only_next_pending_gate_advances_with_new_exact_evidence():
    evidence = {gate_id: {"gate": gate_id} for gate_id in quality.GOVERNED_GATE_SEQUENCE}
    pending = quality.build_gate_ledger(
        ROOT,
        evidence,
        attempt_id="delivery-attempt",
        statuses={"exact_byte_atomic_delivery": "pending"},
    )
    prior_hash = pending["records"][-1]["evidence_sha256"]

    passed = quality.advance_gate_ledger(
        ROOT,
        pending,
        gate_id="exact_byte_atomic_delivery",
        terminal_status="passed",
        evidence={"opened": [{"filename": "protocol.docx", "sha256": "a" * 64}]},
    )

    assert passed["records"][-1]["terminal_status"] == "passed"
    assert passed["records"][-1]["evidence_sha256"] != prior_hash
    assert quality.validate_gate_ledger(ROOT, passed) == passed
    with pytest.raises(ValueError, match="first unresolved"):
        quality.advance_gate_ledger(
            ROOT,
            pending,
            gate_id="every_page_visual_qa",
            terminal_status="passed",
            evidence={"invalid": True},
        )

    finding = {
        "code": "EVERY_PAGE_VISUAL_QA_FAILED",
        "target": "protocol.docx#page=4#section=3",
        "evidence_sha256": "c" * 64,
        "retry_owner": _gate_owner("every_page_visual_qa"),
        "terminal_status": "blocked",
    }
    blocked = quality.advance_gate_ledger(
        ROOT,
        quality.build_gate_ledger(
            ROOT,
            evidence,
            attempt_id="visual-attempt",
            statuses={"every_page_visual_qa": "pending", "exact_byte_atomic_delivery": "pending"},
        ),
        gate_id="every_page_visual_qa",
        terminal_status="blocked",
        evidence={"page_sha256": "c" * 64},
        findings=[finding],
    )
    with pytest.raises(ValueError, match="terminal blocked gate"):
        quality.advance_gate_ledger(
            ROOT,
            blocked,
            gate_id="every_page_visual_qa",
            terminal_status="passed",
            evidence={"page_sha256": "d" * 64},
        )
    assert blocked["records"][4]["findings"] == [finding]


def test_independent_runner_binds_exact_matrix_combinations(tmp_path, monkeypatch):
    cases = [
        {"status": "passed", "case": "retrospective", "result": {"study_type": "Retrospective"}, "icf_template": None},
        {"status": "passed", "case": "prospective-advarra", "result": {"study_type": "Prospective"}, "icf_template": "Advarra"},
        {"status": "passed", "case": "prospective-sterling", "result": {"study_type": "Prospective"}, "icf_template": "Sterling"},
        {"status": "passed", "case": "ambispective-advarra", "result": {"study_type": "Ambispective"}, "icf_template": "Advarra"},
        {"status": "passed", "case": "ambispective-sterling", "result": {"study_type": "Ambispective"}, "icf_template": "Sterling"},
    ]
    monkeypatch.setattr(
        workflow.subprocess,
        "run",
        lambda *_args, **_kwargs: (_ for _ in ()).throw(workflow.subprocess.CalledProcessError(1, "git")),
    )
    monkeypatch.setattr(workflow, "page_renderers", lambda **kwargs: [{"kind": "pypdfium2"}])
    monkeypatch.setattr(workflow, "run_release_gate", lambda *args, **kwargs: {
        "status": "structural_passed",
        "assurance": "synthetic-structural-only",
        "cases": cases,
        "evidence_root": tmp_path.as_posix(),
    })
    monkeypatch.setattr(
        workflow,
        "audit_format_conformance_outputs",
        lambda *args, **kwargs: {"status": "passed", "cases": []},
    )

    report = workflow.run_format_conformance(ROOT, evidence_root=tmp_path)

    assert report["status"] == "structural_passed"
    assert report["assurance"] == "deterministic-structural-only"
    assert report["live_certification_required"] is True
    assert report["output_baseline_status"] == "passed"
    assert report["matrix_sha256"] == quality.load_format_conformance_matrix(ROOT)["matrix_sha256"]
    assert report["covered_cases"] == [
        "ambispective-advarra", "ambispective-sterling", "prospective-advarra",
        "prospective-sterling", "retrospective-protocol",
    ]
    assert (tmp_path / "format-conformance-report.json").is_file()


def test_source_runner_bootstraps_disposable_release_candidate(monkeypatch, tmp_path):
    expected = {"status": "structural_passed", "assurance": "deterministic-structural-only"}
    calls = []
    evidence_root = (tmp_path / "evidence").resolve()
    monkeypatch.setattr(workflow, "page_renderers", lambda **kwargs: [])
    monkeypatch.setattr(
        workflow,
        "_run_format_conformance_in_disposable_candidate",
        lambda repo_root, target: calls.append((repo_root, target)) or expected,
        raising=False,
    )

    result = workflow.run_format_conformance(ROOT, evidence_root=evidence_root)

    assert result == expected
    assert calls == [(ROOT.resolve(), evidence_root)]


def test_disposable_bootstrap_uses_public_clean_tree_packager(tmp_path, monkeypatch):
    calls = []

    def fake_package(repo_root, output_path):
        calls.append((repo_root, output_path))
        raise RuntimeError("stop after public packager")

    monkeypatch.setattr(workflow, "package_release", fake_package)
    with pytest.raises(RuntimeError, match="stop after public packager"):
        workflow._run_format_conformance_in_disposable_candidate(
            ROOT,
            (tmp_path / "evidence").resolve(),
        )
    assert calls and calls[0][0] == ROOT


def test_normalized_signature_detects_semantic_format_mutation(tmp_path):
    source = ROOT / "assets/client-templates/docx/prospective-protocol.template.docx"
    candidate = tmp_path / "mutated.docx"
    document = Document(source)
    document.sections[0].left_margin = Inches(1.37)
    next(iter(document.styles)).font.size = Pt(13)
    document.sections[0].header.add_paragraph("Changed furniture")
    document.tables[0].columns[0].width = Inches(2.75)
    heading = next(paragraph for paragraph in document.paragraphs if paragraph.style.name.startswith("Heading"))
    heading.paragraph_format.page_break_before = True
    document.save(candidate)

    expected = quality.normalized_docx_format_signature(source)
    actual = quality.normalized_docx_format_signature(candidate)

    assert actual["section_geometry"] != expected["section_geometry"]
    assert actual["styles"] != expected["styles"]
    assert actual["headers_footers"] != expected["headers_footers"]
    assert actual["table_geometry"] != expected["table_geometry"]
    assert actual["paragraph_rhythm"] != expected["paragraph_rhythm"]
    assert actual["pagination_relations"] != expected["pagination_relations"]


def test_rendered_cohesion_rejects_trailing_protocol_heading(tmp_path, monkeypatch):
    docx_path = tmp_path / "protocol.docx"
    document = Document()
    document.add_heading("1. BODY", level=1)
    document.save(docx_path)
    pdf_path = tmp_path / "protocol.pdf"
    pdf_path.write_bytes(b"pdf")

    class Page:
        def extract_text(self):
            return "1. BODY"

    monkeypatch.setattr(quality, "PdfReader", lambda _path: type("Reader", (), {"pages": [Page()]})())
    result = quality._rendered_pagination_relations(docx_path, pdf_path)

    assert result["heading_cohesion"][0]["first_content_sha256"] is None
    assert result["all_headings_with_first_content"] is False


def test_icf_cohesion_requires_actual_first_content_not_unrelated_word_count(tmp_path, monkeypatch):
    docx_path = tmp_path / "icf.docx"
    document = Document()
    style = document.styles.add_style("Heading ICF Section", WD_STYLE_TYPE.PARAGRAPH)
    document.add_paragraph("AGREEMENT TO PARTICIPATE", style=style)
    document.add_paragraph("Actual first consent content must follow this heading.")
    document.save(docx_path)
    pdf_path = tmp_path / "icf.pdf"
    pdf_path.write_bytes(b"pdf")

    class Page:
        def extract_text(self):
            return "AGREEMENT TO PARTICIPATE " + "unrelated " * 100

    monkeypatch.setattr(quality, "PdfReader", lambda _path: type("Reader", (), {"pages": [Page()]})())
    result = quality._rendered_pagination_relations(docx_path, pdf_path)

    assert result["heading_cohesion"][0]["first_content_sha256"] is not None
    assert result["all_headings_with_first_content"] is False
