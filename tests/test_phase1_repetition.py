import copy
import json
from pathlib import Path

from docx import Document

import contracts
import quality
from contracts import batch_plan, icf_contract
from drafting import create_drafting_request, recorded_acceptance_response, validate_response
from quality import CONTENT_CHECKS, RESPONSE_SCHEMA, create_verification_requests, deterministic_content_check, validate_verifications
from rendering import render_documents


def icf_summary_obligations(reference):
    return contracts.icf_summary_obligations(reference)


def protocol_concept_ownership(study_type):
    return contracts.protocol_concept_ownership(study_type)


def assess_icf_output(*args, **kwargs):
    return quality.assess_icf_output(*args, **kwargs)


def assess_protocol_concept_repetition(*args, **kwargs):
    return quality.assess_protocol_concept_repetition(*args, **kwargs)


ROOT = Path(__file__).resolve().parents[1]


def fixture():
    return json.loads(
        (ROOT / "tests/fixtures/prospective-acceptance-source.json").read_text(
            encoding="utf-8"
        )
    )


def section(text):
    return {"paragraphs": [{"text": text}], "lists": []}


def visible_paragraphs(path):
    return [" ".join(item.text.split()) for item in Document(path).paragraphs if item.text.strip()]


def test_icf_contract_has_independent_key_information_summary_from_canonical_source(tmp_path):
    reference = fixture()
    reference["meta"]["icf_template"] = "Sterling"
    contracts = {item.section_id: item for item in icf_contract("Prospective", "Sterling")}

    assert "icf.key-information-summary" in contracts
    summary = contracts["icf.key-information-summary"]
    assert set(summary.summary_concepts) == {
        "study-purpose",
        "participation-duration",
        "principal-risks",
        "possible-benefit",
        "alternatives-voluntariness",
    }
    batch = next(item for item in batch_plan("Prospective", "Sterling") if item.batch_id == "icf-narrative")
    request_path = create_drafting_request(
        repo_root=ROOT,
        revision_dir=tmp_path,
        revision_id="r-summary",
        reference=reference,
        batch=batch,
        attempts={item: 1 for item in batch.section_ids},
        wave="initial",
    )
    request = json.loads(request_path.read_text(encoding="utf-8"))
    contract = next(item for item in request["section_contracts"] if item["section_id"] == "icf.key-information-summary")
    assert contract["summary_concepts"] == list(summary.summary_concepts)
    assert contract["minimum_evidence"]
    assert all(path.startswith(("study.", "objectives.", "endpoints.", "design.", "procedures.", "population.", "risks_benefits.")) for path in contract["minimum_evidence"])


def test_equivalent_narrative_and_structured_inputs_have_same_icf_summary_obligations():
    narrative = fixture()
    structured = copy.deepcopy(narrative)
    structured["procedures"]["assessments"] = [
        {"visit": "Screening", "procedures": ["Consent", "Medical history"]},
        {"visit": "Follow-up", "timing": "Week 4", "procedures": ["Assessment"]},
    ]
    structured["procedures"]["visit_schedule"] = [
        {"visit": "Screening", "timing": "Before treatment", "procedures": ["Consent", "Medical history"]},
        {"visit": "Follow-up", "timing": "Week 4", "procedures": ["Assessment"]},
    ]
    structured["design"]["interventions"] = [
        {"name": structured["design"]["intervention_name"], "type": structured["design"]["intervention_type"]}
    ]

    assert set(icf_summary_obligations(narrative)) == set(icf_summary_obligations(structured)) == {
        "study-purpose",
        "participation-duration",
        "principal-risks",
        "possible-benefit",
        "alternatives-voluntariness",
    }
    for obligations in (icf_summary_obligations(narrative), icf_summary_obligations(structured)):
        assert all(item["evidence_refs"] for item in obligations.values())


def test_sterling_key_information_does_not_reuse_detailed_purpose_risk_or_duration(tmp_path):
    reference = fixture()
    reference["meta"]["icf_template"] = "Sterling"
    purpose = "The detailed purpose explanation describes the complete scientific rationale and the primary research question in participant-facing language."
    risk = "The complete detailed risk explanation describes procedure discomfort, foreseeable device effects, and the safeguards used by the study team."
    duration = "Your complete participation includes screening, treatment visits, and follow-up for twelve months after the final procedure."
    model = {
        "protocol": [],
        "prs": {},
        "icf": {
            "icf.key-information-summary": {
                "paragraphs": [
                    {"text": "Researchers are studying the intervention.", "evidence_refs": ["source:objectives.primary"], "boilerplate_refs": []},
                    {"text": "You will attend study visits for about twelve months.", "evidence_refs": ["source:study.timeline"], "boilerplate_refs": []},
                    {"text": "The main risks include discomfort and possible loss of privacy.", "evidence_refs": ["source:risks_benefits.risks"], "boilerplate_refs": []},
                    {"text": "You may not benefit directly.", "evidence_refs": ["source:risks_benefits.benefits"], "boilerplate_refs": []},
                    {"text": "Taking part is voluntary; you may discuss other care options.", "evidence_refs": ["source:risks_benefits.alternatives"], "boilerplate_refs": []},
                ],
                "lists": [],
            },
            "icf.study-purpose": section(purpose),
            "icf.procedures": section("You will complete screening, treatment visits, and follow-up assessments."),
            "icf.duration": section(duration),
            "icf.risks": section(risk),
            "icf.benefits": section("The detailed benefits section states that direct benefit is not guaranteed."),
            "icf.alternatives": section("The detailed alternatives section explains available care choices."),
        },
    }

    report = render_documents(ROOT, tmp_path, reference, model, artifact_names={"icf"})
    assert report["status"] == "passed"
    paragraphs = visible_paragraphs(tmp_path / "candidate/icf.docx")
    assert sum(purpose in item for item in paragraphs) == 1
    assert sum(risk in item for item in paragraphs) == 1
    assert sum(duration in item for item in paragraphs) == 1
    assert any("Researchers are studying the intervention." in item for item in paragraphs)
    assert any("The main risks include discomfort" in item for item in paragraphs)


def test_rendered_icf_assessment_detects_renderer_only_generated_duplication(tmp_path):
    document = Document()
    document.add_heading("KEY INFORMATION", level=1)
    duplicate = "You will attend four research visits over twelve months and complete vision and symptom assessments at each scheduled visit."
    document.add_paragraph(duplicate)
    document.add_heading("PROCEDURES", level=1)
    document.add_paragraph(duplicate)
    path = tmp_path / "icf.docx"
    document.save(path)

    assessment = assess_icf_output(
        path,
        {"meta": {"study_type": "Prospective", "icf_template": "Sterling"}},
        {"family": "Sterling"},
    )

    assert assessment["status"] == "repairable"
    finding = assessment["findings"][0]
    assert finding["code"] == "icf-exact-generated-duplication"
    assert finding["target_ids"] == ["icf.key-information-summary"]
    assert finding["detail_section"] == "icf.procedures"
    assert finding["repair_action"] == "redraft_summary"


def test_rendered_icf_duplication_uses_governed_summary_drafting_recovery(tmp_path):
    document = Document()
    document.add_heading("KEY INFORMATION", level=1)
    duplicate = "You will attend four research visits over twelve months and complete vision and symptom assessments at each scheduled visit."
    document.add_paragraph(duplicate)
    document.add_heading("PROCEDURES", level=1)
    document.add_paragraph(duplicate)
    path = tmp_path / "icf.docx"
    document.save(path)

    finding = assess_icf_output(
        path,
        {"meta": {"study_type": "Prospective", "icf_template": "Sterling"}},
        {"family": "Sterling"},
    )["findings"][0]

    assert finding["target_ids"] == ["icf.key-information-summary"]
    assert finding["recovery_class"] == "drafting_defect"
    assert finding["action"] == contracts.RECOVERY_POLICIES["drafting_defect"]


def test_draft_time_icf_duplication_retries_summary_and_preserves_detail(tmp_path):
    reference = fixture()
    reference["meta"]["icf_template"] = "Sterling"
    batch = next(item for item in batch_plan("Prospective", "Sterling") if item.batch_id == "icf-narrative")
    request_path = create_drafting_request(
        repo_root=ROOT,
        revision_dir=tmp_path,
        revision_id="r-draft-duplicate-routing",
        reference=reference,
        batch=batch,
        attempts={item: 1 for item in batch.section_ids},
        wave="initial",
    )
    request = json.loads(request_path.read_text(encoding="utf-8"))
    response = recorded_acceptance_response(request)
    summary = next(
        item for item in response["section_results"]
        if item["section_id"] == "icf.key-information-summary"
    )
    risks = next(
        item for item in response["section_results"]
        if item["section_id"] == "icf.risks"
    )
    summary["paragraphs"][2]["text"] = risks["paragraphs"][0]["text"]

    accepted, findings = validate_response(request, response)

    accepted_ids = {item["section_id"] for item in (accepted or {}).get("drafts", [])}
    assert "icf.risks" in accepted_ids
    assert "icf.key-information-summary" not in accepted_ids
    duplicate = next(item for item in findings if "duplicated across separately contracted" in item["issue"])
    assert duplicate["target_ids"] == ["icf.key-information-summary"]


def test_rendered_icf_assessment_detects_long_generated_span_inside_distinct_paragraphs(tmp_path):
    shared = (
        "you will attend four research visits over twelve months and complete vision and symptom "
        "assessments at each scheduled visit"
    )
    document = Document()
    document.add_heading("KEY INFORMATION", level=1)
    document.add_paragraph(f"In brief, {shared}, as explained below.")
    document.add_heading("PROCEDURES", level=1)
    document.add_paragraph(f"During this study, {shared}, with additional details provided by the study team.")
    path = tmp_path / "icf.docx"
    document.save(path)

    assessment = assess_icf_output(path, {}, {"family": "Sterling"})

    assert assessment["status"] == "repairable"
    assert assessment["findings"][0]["detail_section"] == "icf.procedures"


def test_rendered_icf_assessment_exempts_family_classified_static_legal_text(tmp_path):
    legal = (
        "Taking part is voluntary and you may stop at any time without penalty or loss of benefits "
        "to which you are otherwise entitled."
    )
    document = Document()
    document.add_heading("KEY INFORMATION", level=1)
    document.add_paragraph(legal)
    document.add_heading("ALTERNATIVE TREATMENTS", level=1)
    document.add_paragraph(legal)
    path = tmp_path / "icf.docx"
    document.save(path)

    assessment = assess_icf_output(
        path,
        {},
        {"family": "Sterling", "required_static_text": [legal]},
    )

    assert assessment["status"] == "passed"


def test_deterministic_icf_check_uses_governed_static_legal_exemption(tmp_path):
    reference = fixture()
    reference["meta"]["icf_template"] = "Sterling"
    boilerplate = json.loads(
        (ROOT / "references/fixed-clinical-boilerplate.json").read_text(encoding="utf-8")
    )["sections"]["icf-voluntary"]
    model = {
        "protocol": [],
        "prs": {},
        "icf": {
            "icf.key-information-summary": {
                "paragraphs": [
                    {"text": "Researchers are studying the approved intervention."},
                    {"text": "You will attend the approved study visits."},
                    {"text": "The main risks are described below."},
                    {"text": "You may receive no direct benefit."},
                    {"text": boilerplate},
                ],
                "lists": [],
            },
            "icf.procedures": section("You will complete the approved study visits and procedures."),
            "icf.alternatives": section(boilerplate),
        },
    }
    render_documents(ROOT, tmp_path, reference, model)

    findings = deterministic_content_check(tmp_path, reference)

    assert not any(item.get("code") == "icf-exact-generated-duplication" for item in findings)


def test_deterministic_icf_check_does_not_exempt_generated_risk_boilerplate_duplication(tmp_path):
    reference = fixture()
    reference["meta"]["icf_template"] = "Sterling"
    boilerplate = json.loads(
        (ROOT / "references/fixed-clinical-boilerplate.json").read_text(encoding="utf-8")
    )["sections"]["icf-sparse-risks"]
    model = {
        "protocol": [],
        "prs": {},
        "icf": {
            "icf.key-information-summary": {
                "paragraphs": [
                    {"text": "Researchers are studying the approved intervention."},
                    {"text": "You will attend the approved study visits."},
                    {"text": boilerplate},
                    {"text": "You may receive no direct benefit."},
                    {"text": "Taking part is voluntary and other care options are available."},
                ],
                "lists": [],
            },
            "icf.procedures": section("You will complete the approved study visits and procedures."),
            "icf.risks": section(boilerplate),
        },
    }
    render_documents(ROOT, tmp_path, reference, model)

    findings = deterministic_content_check(tmp_path, reference)

    assert any(item.get("code") == "icf-exact-generated-duplication" for item in findings)


def test_rendered_icf_assessment_exempts_required_and_short_repeatable_text(tmp_path):
    document = Document()
    document.add_heading("KEY INFORMATION", level=1)
    document.add_paragraph("Week 4")
    document.add_paragraph("Study Device A")
    document.add_heading("PROCEDURES", level=1)
    document.add_paragraph("Week 4")
    document.add_paragraph("Study Device A")
    document.add_paragraph("Signature of Participant")
    document.add_paragraph("Signature of Participant")
    path = tmp_path / "icf.docx"
    document.save(path)

    assessment = assess_icf_output(path, {}, {"family": "Sterling"})
    assert assessment["status"] == "passed"
    assert assessment["findings"] == []


def test_protocol_concept_repetition_finds_paraphrased_secondary_explanation():
    sections = {
        "introduction": [
            "Existing treatment leaves an unmet need because recovery is slow and outcomes remain variable. The study evaluates a new approach intended to address that evidence gap."
        ],
        "objectives": [
            "Because current care has inconsistent outcomes and prolonged recovery, this investigation examines a different approach to fill the remaining evidence gap before stating its objective."
        ],
    }
    findings = assess_protocol_concept_repetition(sections, protocol_concept_ownership("Prospective"))

    assert len(findings) == 1
    finding = findings[0]
    assert finding["concept_id"] == "clinical-rationale"
    assert finding["primary_section"] == "introduction"
    assert finding["secondary_section"] == "objectives"
    assert finding["treatment"] == "excessive"
    assert finding["disposition"] == "manual_review"
    assert finding["target_ids"] == ["objectives"]


def test_protocol_concept_repetition_allows_endpoint_labels_visit_timing_and_brief_references():
    sections = {
        "objectives": ["The primary endpoint is change in symptom score at Week 4."],
        "study-procedure.visits": ["Participants return at Week 4 for the symptom-score assessment."],
        "analysis-plan.methodology": ["The Week 4 symptom-score endpoint will be summarized descriptively as defined in Section 6."],
    }
    findings = assess_protocol_concept_repetition(sections, protocol_concept_ownership("Prospective"))
    assert findings == []


def test_protocol_concept_repetition_allows_concise_completion_visit_reference():
    sections = {
        "study-procedure.visits": [
            "The complete visit schedule includes screening, baseline treatment, and follow-up assessments at Week 4 and Week 12."
        ],
        "endpoint-criteria.completion": [
            "A participant completes the study after screening, the treatment visit, and the final Week 12 follow-up visit."
        ],
    }
    findings = assess_protocol_concept_repetition(sections, protocol_concept_ownership("Prospective"))
    assert findings == []


def test_exact_protocol_repetition_across_sections_blocks_and_targets_secondary_only(tmp_path):
    duplicate = (
        "Existing treatment leaves an unmet need because recovery is slow and outcomes remain variable. "
        "The study evaluates a new approach intended to address that evidence gap."
    )
    document = Document()
    protocol = {item.section_id: item for item in contracts.protocol_contract("Retrospective")}
    for section_id in ("introduction", "objectives"):
        spec = protocol[section_id]
        document.add_heading(f"{spec.number} {spec.title}", level=1)
        document.add_paragraph(duplicate)
    candidate = tmp_path / "candidate"
    candidate.mkdir()
    document.save(candidate / "protocol.docx")

    findings = deterministic_content_check(
        tmp_path,
        {"meta": {"study_type": "Retrospective"}},
    )
    repeated = [item for item in findings if item.get("code") == "protocol-exact-repetition"]

    assert len(repeated) == 1
    assert repeated[0]["target_ids"] == ["objectives"]
    assert repeated[0]["recovery_class"] == "drafting_defect"
    assert repeated[0]["action"] == contracts.RECOVERY_POLICIES["drafting_defect"]


def test_protocol_concept_ownership_is_included_in_drafting_contract(tmp_path):
    reference = fixture()
    batch = next(item for item in batch_plan("Prospective") if item.batch_id == "protocol-foundations")
    request_path = create_drafting_request(
        repo_root=ROOT,
        revision_dir=tmp_path,
        revision_id="r-concepts",
        reference=reference,
        batch=batch,
        attempts={item: 1 for item in batch.section_ids},
        wave="initial",
    )
    request = json.loads(request_path.read_text(encoding="utf-8"))
    introduction = next(item for item in request["section_contracts"] if item["section_id"] == "introduction")
    objectives = next(item for item in request["section_contracts"] if item["section_id"] == "objectives")
    assert "clinical-rationale" in introduction["concept_ownership"]["owns"]
    assert "clinical-rationale" in objectives["concept_ownership"]["brief_reference_only"]
    assert "study-objectives" in objectives["concept_ownership"]["owns"]


def test_synthetic_icf_drafting_rendering_and_duplication_assessment_end_to_end(tmp_path):
    reference = fixture()
    reference["meta"]["icf_template"] = "Sterling"
    batch = next(item for item in batch_plan("Prospective", "Sterling") if item.batch_id == "icf-narrative")
    request_path = create_drafting_request(
        repo_root=ROOT,
        revision_dir=tmp_path,
        revision_id="r-e2e",
        reference=reference,
        batch=batch,
        attempts={item: 1 for item in batch.section_ids},
        wave="initial",
    )
    request = json.loads(request_path.read_text(encoding="utf-8"))
    response = recorded_acceptance_response(request)
    accepted, findings = validate_response(request, response)
    assert findings == []
    assert accepted is not None
    icf = {
        item["section_id"]: {"paragraphs": item["paragraphs"], "lists": item["lists"]}
        for item in accepted["drafts"]
    }
    report = render_documents(
        ROOT,
        tmp_path,
        reference,
        {"protocol": [], "icf": icf, "prs": {}},
        artifact_names={"icf"},
    )
    assert report["status"] == "passed"
    output = tmp_path / "candidate/icf.docx"
    assessment = assess_icf_output(output, reference, {"family": "Sterling"})
    assert assessment["status"] == "passed"
    visible = " ".join(visible_paragraphs(output)).casefold()
    assert "device discomfort" in visible
    assert "privacy risks" in visible
    assert "no direct benefit is guaranteed" in visible


def test_content_review_request_requires_structured_protocol_concept_findings(tmp_path):
    request_path = create_verification_requests(
        tmp_path,
        fixture(),
        {"artifacts": []},
    )[0]
    request = json.loads(request_path.read_text(encoding="utf-8"))
    ownership = request["protocol_concept_ownership"]
    assert ownership["introduction"]["owns"] == ["clinical-rationale"]
    assert ownership["objectives"]["do_not_restate"] == ["clinical-rationale"]
    instructions = request["instructions"]
    for field in (
        "concept_id",
        "primary_section",
        "secondary_section",
        "primary_paragraphs",
        "secondary_paragraphs",
    ):
        assert field in instructions
    assert "generic similarity percentage" in instructions
    assert "manual review" in instructions


def test_protocol_secondary_sections_do_not_own_complete_schedule_or_methods():
    protocol = {item.section_id: item for item in contracts.protocol_contract("Prospective")}
    assert "procedures.visit_schedule" not in protocol["study-procedure.measurements"].evidence
    assert "procedures.assessments" not in protocol["study-procedure.measurements"].evidence
    assert protocol["analysis-plan.considerations"].evidence == ("statistics.software",)
    assert protocol["analysis-plan.considerations"].boilerplate_key == "analysis-considerations-cross-reference"
    assert {"procedures.visit_schedule", "procedures.assessments"} <= set(protocol["endpoint-criteria.completion"].evidence)
    assert {"study.timeline", "procedures.visit_schedule", "procedures.assessments"} <= set(protocol["endpoint-criteria.study-completion"].evidence)
    assert protocol["endpoint-criteria.completion"].source_coverage == "concept_reference"
    assert protocol["endpoint-criteria.study-completion"].source_coverage == "concept_reference"
    assert "complete-visit-schedule" in protocol["endpoint-criteria.completion"].do_not_restate_concepts
    assert "complete-visit-schedule" in protocol["endpoint-criteria.study-completion"].do_not_restate_concepts
    assert "statistical-methods" in protocol["analysis-plan.considerations"].do_not_restate_concepts


def test_analysis_cross_reference_contains_no_source_gap_commentary():
    boilerplate = json.loads(
        (ROOT / "references/fixed-clinical-boilerplate.json").read_text(encoding="utf-8")
    )["sections"]["analysis-considerations-cross-reference"]

    assert "approved source" not in boilerplate.casefold()
    assert "when specified" not in boilerplate.casefold()


def test_key_information_rejects_generic_blocks_with_unearned_source_refs(tmp_path):
    reference = fixture()
    reference["meta"]["icf_template"] = "Sterling"
    batch = next(item for item in batch_plan("Prospective", "Sterling") if item.batch_id == "icf-narrative")
    request_path = create_drafting_request(
        repo_root=ROOT,
        revision_dir=tmp_path,
        revision_id="r-summary-grounding",
        reference=reference,
        batch=batch,
        target_ids=("icf.key-information-summary",),
        attempts={"icf.key-information-summary": 1},
        wave="initial",
    )
    request = json.loads(request_path.read_text(encoding="utf-8"))
    response = recorded_acceptance_response(request)
    summary = response["section_results"][0]
    for paragraph in summary["paragraphs"][:2]:
        paragraph["text"] = "This generic statement contains no approved study-specific fact."

    accepted, findings = validate_response(request, response)

    assert accepted is not None
    assert any(
        item["field"] == "icf.key-information-summary"
        and "not observable" in item["issue"]
        for item in findings
    )


def test_key_information_rejects_alternatives_without_voluntary_participation(tmp_path):
    reference = fixture()
    reference["meta"]["icf_template"] = "Sterling"
    reference.setdefault("risks_benefits", {})["alternatives"] = (
        "Continue routine clinical care instead of joining the study"
    )
    batch = next(item for item in batch_plan("Prospective", "Sterling") if item.batch_id == "icf-narrative")
    request_path = create_drafting_request(
        repo_root=ROOT,
        revision_dir=tmp_path,
        revision_id="r-summary-voluntary",
        reference=reference,
        batch=batch,
        target_ids=("icf.key-information-summary",),
        attempts={"icf.key-information-summary": 1},
        wave="initial",
    )
    request = json.loads(request_path.read_text(encoding="utf-8"))
    response = recorded_acceptance_response(request)
    alternatives = response["section_results"][0]["paragraphs"][4]
    alternatives["text"] = (
        "Other options include continuing routine clinical care instead of joining the study."
    )
    alternatives["evidence_refs"] = ["source:risks_benefits.alternatives"]
    alternatives["boilerplate_refs"] = ["icf-voluntary"]

    accepted, findings = validate_response(request, response)

    assert not (accepted or {}).get("drafts")
    assert any(
        item["field"] == "icf.key-information-summary"
        and "voluntary" in item["issue"].casefold()
        for item in findings
    )


def test_draft_time_icf_duplicate_targets_summary_and_preserves_detail(tmp_path):
    reference = fixture()
    reference["meta"]["icf_template"] = "Sterling"
    batch = next(item for item in batch_plan("Prospective", "Sterling") if item.batch_id == "icf-narrative")
    request_path = create_drafting_request(
        repo_root=ROOT,
        revision_dir=tmp_path,
        revision_id="r-icf-duplicate-routing",
        reference=reference,
        batch=batch,
        attempts={item: 1 for item in batch.section_ids},
        wave="initial",
    )
    request = json.loads(request_path.read_text(encoding="utf-8"))
    response = recorded_acceptance_response(request)
    summary = next(item for item in response["section_results"] if item["section_id"] == "icf.key-information-summary")
    detail = next(item for item in response["section_results"] if item["section_id"] == "icf.study-purpose")
    summary["paragraphs"][0]["text"] = detail["paragraphs"][0]["text"]

    accepted, findings = validate_response(request, response)

    duplicate = next(item for item in findings if "duplicated across separately contracted" in item["issue"])
    accepted_ids = {item["section_id"] for item in (accepted or {}).get("drafts", [])}
    assert duplicate["target_ids"] == ["icf.key-information-summary"]
    assert "icf.study-purpose" in accepted_ids
    assert "icf.key-information-summary" not in accepted_ids


def test_protocol_operational_design_contract_preserves_eye_specific_assignment():
    protocol = {item.section_id: item for item in contracts.protocol_contract("Prospective")}
    assert "design.intervention_description" in protocol["study-design.design"].evidence


def test_advarra_does_not_request_unrendered_key_information_summary():
    assert "icf.key-information-summary" not in {
        item.section_id for item in icf_contract("Prospective", "Advarra")
    }
    assert "icf.key-information-summary" not in {
        item.section_id for item in icf_contract("Ambispective", "Advarra")
    }


def test_retrospective_protocol_has_concept_ownership():
    ownership = protocol_concept_ownership("Retrospective")
    assert ownership["introduction"]["owns"] == ["clinical-rationale"]
    assert "clinical-rationale" in ownership["objectives"]["do_not_restate"]
    assert ownership["analysis-plan.methodology"]["owns"] == ["statistical-methods"]
    assert "statistical-methods" in ownership["analysis-plan.considerations"]["do_not_restate"]


def test_protocol_repetition_verifier_finding_remains_structured_manual_review(tmp_path):
    request_path = create_verification_requests(tmp_path, fixture(), {"artifacts": []})[0]
    request = json.loads(request_path.read_text(encoding="utf-8"))
    response = {
        "schema_version": RESPONSE_SCHEMA,
        "request_id": request["request_id"],
        "request_sha256": request["request_sha256"],
        "task": request["task"],
        "revision_id": request["revision_id"],
        "review_set": request["review_set"],
        "producer": {"model_id": "synthetic-reviewer", "reviewer_id": "content-reviewer"},
        "status": "blocked",
        "artifacts": request["artifacts"],
        "section_assessments": [
            {
                "artifact": item["artifact"],
                "section_id": item["section_id"],
                "status": (
                    "failed" if item["section_id"] == "analysis-plan.considerations"
                    else "passed"
                ),
                "checks": list(CONTENT_CHECKS),
            }
            for item in request["sections"]
        ],
        "cross_document_assessments": [
            {"check": check, "status": "passed"}
            for check in request["cross_document_checks"]
        ],
        "findings": [{
            "finding_id": "protocol-repeat-1",
            "category": "content",
            "check": "concept_repetition",
            "issue": "The secondary section repeats the complete statistical method.",
            "target_ids": ["analysis-plan.considerations"],
            "code": "protocol-concept-repetition",
            "concept_id": "statistical-methods",
            "primary_section": "analysis-plan.methodology",
            "secondary_section": "analysis-plan.considerations",
            "primary_paragraphs": ["Complete method."],
            "secondary_paragraphs": ["Repeated complete method."],
            "treatment": "excessive",
            "necessary": False,
            "concise": False,
            "material": False,
            "contradiction": False,
            "obscures_required_information": False,
            "materially_unusable": False,
            "disposition": "manual_review",
        }],
    }
    response_path = tmp_path / request["response_path"]
    response_path.parent.mkdir(parents=True, exist_ok=True)
    response_path.write_text(json.dumps(response), encoding="utf-8")

    findings, _evidence = validate_verifications(tmp_path, request_paths=[request_path])

    finding = next(item for item in findings if item.get("code") == "protocol-concept-repetition")
    assert finding["concept_id"] == "statistical-methods"
    assert finding["primary_section"] == "analysis-plan.methodology"
    assert finding["secondary_section"] == "analysis-plan.considerations"
    assert finding["publication_disposition"] == "warning"
    assert finding["action"] == "manual_review"


def test_repetition_finding_cannot_contradict_passing_section_assessment(tmp_path):
    request_path = create_verification_requests(tmp_path, fixture(), {"artifacts": []})[0]
    request = json.loads(request_path.read_text(encoding="utf-8"))
    response = {
        "schema_version": RESPONSE_SCHEMA,
        "request_id": request["request_id"],
        "request_sha256": request["request_sha256"],
        "task": request["task"],
        "revision_id": request["revision_id"],
        "review_set": request["review_set"],
        "producer": {"model_id": "synthetic-reviewer", "reviewer_id": "content-reviewer"},
        "status": "blocked",
        "artifacts": request["artifacts"],
        "section_assessments": [
            {
                "artifact": item["artifact"],
                "section_id": item["section_id"],
                "status": "passed",
                "checks": list(CONTENT_CHECKS),
            }
            for item in request["sections"]
        ],
        "cross_document_assessments": [
            {"check": check, "status": "passed"}
            for check in request["cross_document_checks"]
        ],
        "findings": [{
            "finding_id": "protocol-repeat-contradictory-envelope",
            "category": "content",
            "check": "concept_repetition",
            "issue": "The secondary section repeats the complete statistical method.",
            "target_ids": ["analysis-plan.considerations"],
            "code": "protocol-concept-repetition",
            "concept_id": "statistical-methods",
            "primary_section": "analysis-plan.methodology",
            "secondary_section": "analysis-plan.considerations",
            "primary_paragraphs": ["Complete method."],
            "secondary_paragraphs": ["Repeated complete method."],
            "treatment": "excessive",
            "necessary": False,
            "concise": False,
            "material": False,
            "contradiction": False,
            "obscures_required_information": False,
            "materially_unusable": False,
            "disposition": "manual_review",
        }],
    }
    response_path = tmp_path / request["response_path"]
    response_path.parent.mkdir(parents=True, exist_ok=True)
    response_path.write_text(json.dumps(response), encoding="utf-8")

    findings, _evidence = validate_verifications(tmp_path, request_paths=[request_path])

    assert any(
        item.get("recovery_class") == "verifier_transient"
        and "passing section assessment" in item.get("issue", "")
        for item in findings
    )


def test_protocol_repetition_verifier_cannot_downgrade_ungoverned_or_mistargeted_finding(tmp_path):
    request_path = create_verification_requests(tmp_path, fixture(), {"artifacts": []})[0]
    request = json.loads(request_path.read_text(encoding="utf-8"))
    response = {
        "schema_version": RESPONSE_SCHEMA,
        "request_id": request["request_id"],
        "request_sha256": request["request_sha256"],
        "task": request["task"],
        "revision_id": request["revision_id"],
        "review_set": request["review_set"],
        "producer": {"model_id": "synthetic-reviewer", "reviewer_id": "content-reviewer"},
        "status": "blocked",
        "artifacts": request["artifacts"],
        "section_assessments": [
            {
                "artifact": item["artifact"],
                "section_id": item["section_id"],
                "status": "passed",
                "checks": list(CONTENT_CHECKS),
            }
            for item in request["sections"]
        ],
        "cross_document_assessments": [
            {"check": check, "status": "passed"}
            for check in request["cross_document_checks"]
        ],
        "findings": [{
            "finding_id": "protocol-repeat-invalid",
            "category": "content",
            "check": "concept_repetition",
            "issue": "This finding tries to route an ungoverned ownership pair as a warning.",
            "target_ids": ["introduction"],
            "code": "protocol-concept-repetition",
            "concept_id": "statistical-methods",
            "primary_section": "objectives",
            "secondary_section": "introduction",
            "primary_paragraphs": ["Primary text."],
            "secondary_paragraphs": ["Secondary text."],
            "treatment": "excessive",
            "necessary": False,
            "concise": False,
            "disposition": "manual_review",
        }],
    }
    response_path = tmp_path / request["response_path"]
    response_path.parent.mkdir(parents=True, exist_ok=True)
    response_path.write_text(json.dumps(response), encoding="utf-8")

    findings, _evidence = validate_verifications(tmp_path, request_paths=[request_path])

    assert not any(item.get("publication_disposition") == "warning" for item in findings)
    assert any(item.get("recovery_class") == "verifier_transient" for item in findings)


def test_unresolved_material_repetition_cannot_be_downgraded_to_warning(tmp_path):
    request_path = create_verification_requests(tmp_path, fixture(), {"artifacts": []})[0]
    request = json.loads(request_path.read_text(encoding="utf-8"))
    response = {
        "schema_version": RESPONSE_SCHEMA,
        "request_id": request["request_id"],
        "request_sha256": request["request_sha256"],
        "task": request["task"],
        "revision_id": request["revision_id"],
        "review_set": request["review_set"],
        "producer": {"model_id": "synthetic-reviewer", "reviewer_id": "content-reviewer"},
        "status": "blocked",
        "artifacts": request["artifacts"],
        "section_assessments": [
            {
                "artifact": item["artifact"],
                "section_id": item["section_id"],
                "status": "passed",
                "checks": list(CONTENT_CHECKS),
            }
            for item in request["sections"]
        ],
        "cross_document_assessments": [
            {"check": check, "status": "passed"}
            for check in request["cross_document_checks"]
        ],
        "findings": [{
            "finding_id": "protocol-repeat-material",
            "category": "content",
            "check": "concept_repetition",
            "issue": "The unresolved repetition materially obscures the required statistical method.",
            "target_ids": ["analysis-plan.considerations"],
            "code": "protocol-concept-repetition",
            "concept_id": "statistical-methods",
            "primary_section": "analysis-plan.methodology",
            "secondary_section": "analysis-plan.considerations",
            "primary_paragraphs": ["Complete method."],
            "secondary_paragraphs": ["Repeated complete method."],
            "treatment": "excessive",
            "necessary": False,
            "concise": False,
            "material": True,
            "resolved": False,
            "repair_attempt": 1,
            "disposition": "manual_review",
        }],
    }
    response_path = tmp_path / request["response_path"]
    response_path.parent.mkdir(parents=True, exist_ok=True)
    response_path.write_text(json.dumps(response), encoding="utf-8")

    findings, _evidence = validate_verifications(tmp_path, request_paths=[request_path])

    finding = next(item for item in findings if item.get("code") == "protocol-concept-repetition")
    assert finding.get("publication_disposition") != "warning"
    assert finding["recovery_class"] == "drafting_defect"
    assert finding["target_ids"] == ["analysis-plan.considerations"]


def test_rendered_icf_assessment_uses_real_sterling_headings_and_explicit_exemption(tmp_path):
    duplicate = "This required legal disclosure is intentionally repeated to preserve the approved authorization language in both locations."
    document = Document()
    document.add_heading("KEY INFORMATION", level=1)
    document.add_paragraph(duplicate)
    document.add_heading("POTENTIAL RISKS, EFFECTS, DISCOMFORTS, INCONVENIENCES", level=1)
    document.add_paragraph(duplicate)
    path = tmp_path / "icf.docx"
    document.save(path)

    blocked = assess_icf_output(path, {}, {"family": "Sterling"})
    assert blocked["findings"][0]["detail_section"] == "icf.risks"

    exempt = assess_icf_output(
        path,
        {},
        {"family": "Sterling", "repeatable_text": [duplicate]},
    )
    assert exempt["status"] == "passed"


def test_final_exact_review_recomputes_deterministic_protocol_warnings(tmp_path, monkeypatch):
    (tmp_path / "candidate").mkdir()
    warning = {
        "code": "protocol-concept-repetition",
        "category": "content",
        "check": "concept_repetition",
        "target_ids": ["objectives"],
        "primary_section": "introduction",
        "secondary_section": "objectives",
        "publication_disposition": "warning",
        "action": "manual_review",
        "issue": "The secondary section repeats the complete rationale.",
    }
    monkeypatch.setattr(quality, "deterministic_content_check", lambda *_args: [warning])
    monkeypatch.setattr(quality, "validate_verifications", lambda *_args: ([], {}))
    monkeypatch.setattr(quality, "_final_verification_scope_findings", lambda *_args: [])
    reference = {"meta": {"study_type": "Retrospective"}}
    render_report = {"artifacts": []}

    report = quality.quality_report(tmp_path, reference, render_report, None)

    assert report["status"] == "passed"
    assert report["warnings"] == [warning]
    assert quality.final_exact_artifact_review_findings(
        tmp_path,
        reference,
        render_report,
        report,
    ) == []


def test_production_icf_assessment_receives_explicit_fixed_boilerplate_exemptions(tmp_path, monkeypatch):
    reference = fixture()
    reference["meta"]["icf_template"] = "Sterling"
    candidate = tmp_path / "candidate"
    candidate.mkdir()
    Document().save(candidate / "protocol.docx")
    Document().save(candidate / "icf.docx")
    captured = {}

    def capture_contract(_document, _source, family_contract):
        captured.update(family_contract)
        return {"status": "passed", "findings": []}

    monkeypatch.setattr(quality, "assess_icf_output", capture_contract)
    deterministic_content_check(tmp_path, reference)
    boilerplate = json.loads(
        (ROOT / "references/fixed-clinical-boilerplate.json").read_text(encoding="utf-8")
    )["sections"]

    assert boilerplate["icf-voluntary"] in captured["required_static_text"]
