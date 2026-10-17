"""Included subsection numbers, cross-references and linked TOC stay consistent."""
import json
from pathlib import Path

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

import drafting
import rendering

ROOT = Path(__file__).resolve().parents[1]


@pytest.mark.parametrize('branch,family', [
    ('Prospective', 'Advarra'), ('Prospective', 'Sterling'),
    ('Ambispective', 'Advarra'), ('Ambispective', 'Sterling'),
])
@pytest.mark.parametrize('patient,study', [(False, False), (True, False), (False, True), (True, True)])
def test_optional_sections_have_consecutive_headings_references_and_toc(tmp_path, branch, family, patient, study):
    source = json.loads((ROOT / f'tests/fixtures/{branch.lower()}-acceptance-source.json').read_text())
    source['meta']['icf_template'] = family
    source['procedures']['termination'] = 'Patient termination rule supplied.' if patient else None
    source['procedures']['study_termination'] = 'Study termination rule supplied.' if study else None
    reference = 'Study closeout is described in Section 18.5. Clinical values 18.5 and 18.4 are unchanged.'
    if study:
        reference += ' See Sections 18.4 and 18.5 for study termination and closeout, respectively.'
    model = {'protocol': [{'section_id': 'endpoint-criteria.completion',
                          'paragraphs': [{'text': reference}], 'lists': []}], 'icf': {}, 'prs': {}}
    rendering.render_documents(ROOT, tmp_path, source, model, artifact_names={'protocol'}, audit_output=False)
    path = tmp_path / 'candidate/protocol.docx'
    document = Document(path)
    headings = [p.text.strip() for p in document.paragraphs if p.style.name in {'Heading 1', 'Heading 2'}]
    expected = ['18. STUDY ENDPOINT CRITERIA', '18.1. Patient Completion of Study', '18.2. Patient Discontinuation']
    if patient:
        expected.append('18.3. Patient Termination')
    if study:
        expected.append(f'18.{3 + int(patient)}. Study Termination')
    closeout = f'18.{3 + int(patient) + int(study)}'
    expected.append(f'{closeout}. Study Completion')
    assert [h for h in headings if h.startswith('18.')] == expected
    visible = '\n'.join(p.text for p in document.paragraphs)
    assert f'Study closeout is described in Section {closeout}.' in visible
    assert 'Clinical values 18.5 and 18.4 are unchanged.' in visible
    if study:
        assert f'See Sections 18.{3 + int(patient)} and {closeout} for study termination and closeout' in visible
    assert ('Patient termination rule supplied.' in visible) == patient
    assert ('Study termination rule supplied.' in visible) == study
    rows = [p for p in document.paragraphs if p.style.name.casefold() in {'toc 1', 'toc 2'}]
    assert [p.text.rstrip('\t') for p in rows] == headings
    bookmarks = set(document.element.body.xpath('.//w:bookmarkStart/@w:name'))
    for row in rows:
        link = row._p.xpath('./w:hyperlink')[0]
        assert link.get(qn('w:anchor')) in bookmarks
        assert row._p.xpath('./w:r/w:instrText')
    snapshot = rendering.protocol_section_snapshot(path, 'endpoint-criteria.study-completion')
    assert snapshot[0]['paragraph'].strip() == expected[-1]
    plan = drafting._protocol_writing_plan(source)
    ids = {s['section_id'] for s in plan['sections']}
    assert ('endpoint-criteria.termination' in ids) == patient
    assert ('endpoint-criteria.study-termination' in ids) == study
    # Drafts refer to canonical identities; only the assembled display is renumbered.
    assert next(s['number'] for s in plan['sections'] if s['section_id'] == 'endpoint-criteria.study-completion') == '18.5.'


def test_number_resolution_preserves_split_runs_hyperlinks_bookmarks_and_clinical_values():
    doc = Document()
    doc.add_paragraph('18. STUDY ENDPOINT CRITERIA', 'Heading 1')
    doc.add_paragraph('18.1. Patient Completion', 'Heading 2')
    doc.add_paragraph('18.2. Patient Discontinuation', 'Heading 2')
    doc.add_paragraph('18.4. Study Termination', 'Heading 2')
    heading = doc.add_paragraph('18.5. Study Completion', 'Heading 2')
    bookmark = OxmlElement('w:bookmarkStart'); bookmark.set(qn('w:id'), '7'); bookmark.set(qn('w:name'), 'closeout')
    heading._p.insert(0, bookmark)
    p = doc.add_paragraph('See Sections 18.')
    bold = p.add_run('4'); bold.bold = True
    p.add_run(' and 18.5. Value 18.5 D stays unchanged. ')
    link = OxmlElement('w:hyperlink'); link.set(qn('w:anchor'), 'closeout')
    run = OxmlElement('w:r'); text = OxmlElement('w:t'); text.text = 'Section 18.5'; run.append(text); link.append(run); p._p.append(link)
    table = doc.add_table(1, 2)
    cell = table.cell(0, 0).merge(table.cell(0, 1)); cell.text = 'Section 18.5'
    rendering._renumber_protocol_subsections(doc)
    visible = ''.join(p._p.xpath('.//w:t/text()'))
    assert visible == 'See Sections 18.3 and 18.4. Value 18.5 D stays unchanged. Section 18.4'
    assert bold.bold and bold.text == '3'
    assert link.get(qn('w:anchor')) == 'closeout'
    assert heading._p.xpath('./w:bookmarkStart/@w:name') == ['closeout']
    assert cell.text == 'Section 18.4'
    once = doc.element.xml
    rendering._renumber_protocol_subsections(doc)
    assert doc.element.xml == once


def test_other_section_numbers_and_nested_subsections_use_the_same_rule():
    doc = Document()
    for text, level in [('12. Other Section', 1), ('12.1. First', 2),
                        ('12.3. Second', 2), ('12.3.1. First child', 3),
                        ('12.3.4. Second child', 3)]:
        doc.add_paragraph(text, f'Heading {level}')
    body = doc.add_paragraph('See Sections 12.3 and 12.3.4; measured value 12.3 remains unchanged.')
    rendering._renumber_protocol_subsections(doc)
    assert [p.text for p in doc.paragraphs[:-1]] == [
        '12. Other Section', '12.1. First', '12.2. Second',
        '12.2.1. First child', '12.2.2. Second child']
    assert body.text == 'See Sections 12.2 and 12.2.2; measured value 12.3 remains unchanged.'


def test_retrospective_numbering_and_section_snapshot_remain_stable(tmp_path):
    source = json.loads((ROOT / 'tests/fixtures/retrospective-acceptance-source.json').read_text())
    rendering.render_documents(ROOT, tmp_path, source, {'protocol': [], 'icf': {}, 'prs': {}},
                               artifact_names={'protocol'}, audit_output=False)
    doc = Document(tmp_path / 'candidate/protocol.docx')
    original = doc.element.xml
    rendering._renumber_protocol_subsections(doc)
    assert doc.element.xml == original
    snapshot = rendering.protocol_section_snapshot(tmp_path / 'candidate/protocol.docx', 'study-procedure.enrollment')
    assert snapshot[0]['paragraph'].startswith('8.1.')
