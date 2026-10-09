"""Controlled public delivery replays; renderer/review doubles are not clinical evidence."""
import json
from pathlib import Path

import pytest
from docx import Document
from pypdf import PdfWriter

import drafting
import quality
import workflow
from test_handoff_quality import acceptance_verification

ROOT = Path(__file__).resolve().parents[1]
CASES = quality.branch_acceptance_inventory(ROOT)


def controlled_render(_root, revision, *_args, **_kwargs):
    artifacts = []
    for docx in sorted((revision / 'candidate').glob('*.docx')):
        pdf = revision / f'rendered/{docx.stem}.pdf'
        page = revision / f'rendered/{docx.stem}/page-1.png'
        page.parent.mkdir(parents=True, exist_ok=True)
        writer = PdfWriter()
        writer.add_blank_page(width=612, height=792)
        with pdf.open('wb') as stream:
            writer.write(stream)
        page.write_bytes(b'controlled-render-double')
        artifacts.append({
            'artifact': docx.stem, 'docx': docx.relative_to(revision).as_posix(),
            'docx_sha256': quality.sha256_file(docx),
            'pdf': pdf.relative_to(revision).as_posix(), 'pdf_sha256': quality.sha256_file(pdf),
            'page_count': 1, 'pages': [{'page': 1, 'path': page.relative_to(revision).as_posix(),
                                       'sha256': quality.sha256_file(page)}],
            'renderer': {'kind': 'controlled-office-double'},
            'page_renderer': {'kind': 'controlled-page-double'}, 'status': 'passed',
        })
    render = {'status': 'passed', 'artifacts': artifacts, 'findings': [],
              'renderer': {'kind': 'controlled-office-double'},
              'page_renderer': {'kind': 'controlled-page-double'}}
    return {'schema_version': 'render-assurance/v1', 'status': 'passed', 'fonts': {},
            'font_substitutions': {}, 'candidate': {'files': []}, 'render': render, 'findings': []}


def deliver(tmp_path, monkeypatch, source, review=None, draft=None):
    monkeypatch.setattr(workflow, 'render_assurance', controlled_render)
    reference = tmp_path / 'reference/study.reference.json'
    reference.parent.mkdir(parents=True)
    reference.write_text(json.dumps(source))
    prepared = workflow.prepare(tmp_path)
    assert prepared['status'] == 'awaiting_approval', prepared
    approved = workflow.approve(tmp_path, approved_by='controlled-fixture-owner')
    revision = tmp_path / 'revisions' / approved['revision_id']
    result = workflow.generate(tmp_path, require_promoted_runtime=False)
    # A finite replay guard catches an accidental retry cycle without changing
    # the production deadline or repair policy.
    for _ in range(12):
        if result['status'] != 'awaiting_hermes':
            break
        for relative in result['requests']:
            request = json.loads((revision / relative).read_text())
            response = (draft or drafting.recorded_acceptance_response)(request) if request['task'] in {
                'section_drafting', 'prs_narrative_drafting'} else (review or acceptance_verification)(request)
            path = revision / request['response_path']
            path.parent.mkdir(parents=True, exist_ok=True)
            path.write_text(json.dumps(response))
        result = workflow.generate(tmp_path, require_promoted_runtime=False)
    return result, revision


@pytest.mark.parametrize('case', CASES, ids=lambda item: item['descriptor']['fixture_id'])
def test_varied_complete_sources_deliver_complete_exact_set_with_clean_message(tmp_path, monkeypatch, case):
    source = json.loads((ROOT / case['source_path']).read_text())
    result, revision = deliver(tmp_path, monkeypatch, source)
    assert result['status'] == 'passed', result
    assert result['warnings'] == []
    assert not result['desktop_reply'].get('cautions')
    manifest = json.loads((tmp_path / result['manifest']).read_text())
    confirmation = workflow.confirm_desktop_delivery(manifest, result['desktop_reply'],
                                                      lambda path: Path(path).read_bytes())
    assert confirmation['confirmed'] is True
    expected = workflow.document_set(source['meta']['study_type'])
    assert sorted(path.name for path in (tmp_path / 'output').iterdir()) == sorted(expected)
    for name in expected:
        assert (tmp_path / 'output' / name).read_bytes() == (revision / 'candidate' / name).read_bytes()
    protocol = Document(tmp_path / 'output/protocol.docx')
    assert source['study']['title'] in '\n'.join(p.text for p in protocol.paragraphs)


def reminder_source():
    source = json.loads((ROOT / CASES[0]['source_path']).read_text())
    source['statistics']['bias_minimization'] = (
        'Personnel will telephone enrolled participants to remind them about scheduled follow-up appointments.')
    return source


def reminder_draft(request, text):
    response = drafting.recorded_acceptance_response(request)
    for section in response.get('section_results', []):
        if section['section_id'] == 'study-design.bias':
            for paragraph in section['paragraphs']:
                if 'source:statistics.bias_minimization' in paragraph['evidence_refs']:
                    paragraph['text'] = text
    return response


def test_faithful_paraphrase_delivers_cleanly_and_retains_authenticated_resolution(tmp_path, monkeypatch):
    text = 'The team phones people before appointments so they remember to attend.'
    def review(request):
        response = acceptance_verification(request)
        for decision in response.get('uncertainty_resolutions', []):
            caution = next(row for row in request['lexical_review_cautions'] if row['caution_sha256'] == decision['caution_sha256'])
            if caution['section_id'] == 'study-design.bias':
                decision.update(affected_passage=text, source_excerpt=reminder_source()['statistics']['bias_minimization'],
                                rationale='Calling before appointments expresses the supplied telephone reminders.')
        return response
    result, revision = deliver(tmp_path, monkeypatch, reminder_source(), review=review,
                               draft=lambda request: reminder_draft(request, text))
    assert result['status'] == 'passed', result
    assert result['warnings'] == [] and not result['desktop_reply'].get('cautions')
    manifest = json.loads((tmp_path / result['manifest']).read_text())
    resolutions = manifest['quality']['resolved_evidence_cautions']
    assert any(row['affected_passage'] == text and row['disposition'] == 'faithful' for row in resolutions)
    assert text in '\n'.join(p.text for p in Document(tmp_path / 'output/protocol.docx').paragraphs)
    assert quality.final_exact_artifact_review_findings(
        revision, json.loads((revision / 'approved-reference.json').read_text()),
        json.loads((revision / 'candidate-build.json').read_text())['render_report'], manifest['quality']) == []


@pytest.mark.parametrize('repair_mode', ['faithful', 'ineffective', 'introduced_error'])
def test_confirmed_omission_is_repaired_then_verified_despite_different_words(tmp_path, monkeypatch, repair_mode):
    initial = 'Study documents will be kept securely.'
    corrected = {'faithful': 'The team phones people before appointments so they remember to attend.',
                 'ineffective': initial,
                 'introduced_error': 'The team phones people thirty times daily before appointments.'}[repair_mode]
    repair_requested = []

    def draft(request):
        if any(section.get('section_id') == 'study-design.bias' for section in request.get('section_contracts', [])):
            repair_requested.append(request['request_id'])
        return reminder_draft(request, initial if len(repair_requested) <= 1 else corrected)

    def review(request):
        response = acceptance_verification(request)
        if request['task'] != 'clinical_content_verification':
            return response
        for row in response['uncertainty_resolutions']:
            caution = next(item for item in request['lexical_review_cautions'] if item['caution_sha256'] == row['caution_sha256'])
            if caution['section_id'] == 'study-design.bias' and corrected in quality._review_evidence_text(caution['artifact_passage']):
                row.update(affected_passage=corrected, source_excerpt=reminder_source()['statistics']['bias_minimization'],
                           rationale='Calling ahead of appointments implements the approved reminder calls.')
        missing = [row for row in request['lexical_review_cautions']
                   if row['section_id'] == 'study-design.bias' and (initial in quality._review_evidence_text(row['artifact_passage'])
                       or (repair_mode == 'introduced_error' and corrected in quality._review_evidence_text(row['artifact_passage'])))]
        if not missing:
            return response
        passage = corrected if corrected in quality._review_evidence_text(missing[0]['artifact_passage']) else initial
        response['status'] = 'blocked'
        response['findings'] = [{
            'finding_id': 'missing-telephone-reminders', 'category': 'content', 'artifact': 'protocol',
            'check': 'source_supported', 'target_ids': ['study-design.bias'],
            'issue': 'Record storage does not describe the approved telephone reminder intervention.',
            'affected_passage': passage, 'source_excerpt': reminder_source()['statistics']['bias_minimization'],
            'recommended_action': 'Describe the supplied telephone reminders before appointments.',
        }]
        for row in response['section_assessments']:
            if row['section_id'] == 'study-design.bias':
                row['status'] = 'blocked'
        for row in response['uncertainty_resolutions']:
            if row['caution_sha256'] in {item['caution_sha256'] for item in missing}:
                row.update(disposition='confirmed_defect', affected_passage=passage,
                           source_excerpt=reminder_source()['statistics']['bias_minimization'],
                           rationale='Secure storage omits the planned reminder calls and must be corrected.')
        return response

    result, revision = deliver(tmp_path, monkeypatch, reminder_source(), draft=draft, review=review)
    assert len(repair_requested) >= 2
    if repair_mode != 'faithful':
        assert result['status'] != 'passed', result
        assert not (tmp_path / 'output').exists()
        return
    assert result['status'] == 'passed', result
    assert result['warnings'] == [] and not result['desktop_reply'].get('cautions')
    body = '\n'.join(p.text for p in Document(tmp_path / 'output/protocol.docx').paragraphs)
    assert corrected in body and initial not in body
    archives = list((revision / 'attempts').rglob('attempt-manifest.json'))
    assert archives and any('missing-telephone-reminders' in path.read_text() for path in archives)
    assert any(action['target_bytes_changed'] and action['candidate_bytes_changed']
               for path in archives for action in json.loads(path.read_text())['recovery_actions'])


@pytest.mark.parametrize('mutation', ['source', 'artifact', 'reviewer'])
def test_resolved_paraphrase_cannot_certify_changed_evidence(tmp_path, monkeypatch, mutation):
    text = 'The team phones people before appointments so they remember to attend.'
    result, revision = deliver(tmp_path, monkeypatch, reminder_source(),
                               draft=lambda request: reminder_draft(request, text))
    assert result['status'] == 'passed', result
    manifest = json.loads((tmp_path / result['manifest']).read_text())
    source = json.loads((revision / 'approved-reference.json').read_text())
    original = (tmp_path / 'output/protocol.docx').read_bytes()
    if mutation == 'source':
        source['statistics']['bias_minimization'] = 'No reminders are permitted.'
    elif mutation == 'artifact':
        path = revision / 'candidate/protocol.docx'
        document = Document(path)
        document.add_paragraph('A newly introduced unreviewed clinical claim.')
        document.save(path)
    else:
        response = next(path for path in (revision / 'hermes/verification-responses').glob('*content.json'))
        record = json.loads(response.read_text())
        record['producer']['reviewer_id'] = ''
        response.write_text(json.dumps(record))
    with pytest.raises(RuntimeError, match='stale publication evidence'):
        workflow._publish(tmp_path, revision, source, manifest['quality'])
    assert (tmp_path / 'output/protocol.docx').read_bytes() == original
