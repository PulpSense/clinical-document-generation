#!/usr/bin/env python3
"""Owner-authorized installed Linux probes; includes real Hermes/model calls.

Never invoked by ordinary generation or offline CI. Retains failed probes and
never installs, activates, repairs installed code or resets a study deadline.
"""
from __future__ import annotations

import argparse
import importlib
import json
import os
from pathlib import Path
import platform
import signal
import subprocess
import sys
import tempfile
import time
import uuid


def write(path, value):
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(value, indent=2) + '\n')


def boundary_probe(workflow, root, directory):
    """Actual ordinary prefix, installed read/write tools and harmless denial markers."""
    from worker_capabilities import digest
    directory.mkdir()
    checkpoint = directory / 'checkpoint.json'
    checkpoint.touch()
    native = directory / 'native-work'
    native.mkdir()
    tools_work = directory / 'tools-work'
    tools_work.mkdir()
    sentinels = [directory / name for name in ('source', 'installed-code', 'sibling-review', 'earlier-run')]
    for path in sentinels:
        path.write_text('sentinel')
    token = uuid.uuid4().hex
    evidence = directory / 'realistic-evidence.txt'
    evidence.write_text(''.join(f'{index:04d} {token} approved passage and structured visit association {index:04d}\n' for index in range(2048)))
    launcher, managed_python = workflow._managed_hermes_pair()
    hermes_root = launcher.parent.parent.parent
    environment = workflow._production_subprocess_environment(root)
    native_code = "from docx import Document; from pathlib import Path; import xml.etree.ElementTree as E; import sys; p=Path(sys.argv[1]); d=Document(); d.add_paragraph('native-probe'); d.save(p/'probe.docx'); E.ElementTree(E.Element('probe',value='native-probe')).write(p/'probe.xml')"
    script = '''import json,socket,subprocess,sys,time
from pathlib import Path
sys.path.insert(0,sys.argv[1])
from tools.file_tools import read_file_tool,write_file_tool
source,checkpoint,native,nonce,runtime,code,protected,role=sys.argv[2:]
read_bytes=0
for start in range(1,2049,128):
 value=read_file_tool(source,offset=start,limit=128)
 for index in range(start-1,start+127): assert f'{index:04d} {nonce}' in value
 read_bytes+=len(value.encode())
result={'read_calls':17 if role=='reviewer' else 16,'bounded_source_read_calls':16,'checkpoint_readback_calls':1 if role=='reviewer' else 0,'returned_bytes':read_bytes,'completed_lines':2048,'nonce':nonce}
if role=='reviewer':
 write_file_tool(checkpoint,json.dumps(result)); assert nonce in read_file_tool(checkpoint,offset=1,limit=20)
else:
 script=Path(native).parent/'tools-work/native-probe.py'; write_file_tool(str(script),code)
 subprocess.run([runtime,str(script),native],check=True)
for item in json.loads(protected):
 try: Path(item).write_text('marker')
 except OSError: pass
 else: raise AssertionError('Protected write succeeded')
try: socket.create_connection(('1.1.1.1',443),timeout=0.3)
except OSError: pass
else: raise AssertionError('Direct network escaped')
for target,expected in [('example.org:443',b'403'),('chatgpt.com:443',b'200')]:
 s=socket.create_connection(('127.0.0.1',3128),timeout=12);s.sendall(('CONNECT '+target+' HTTP/1.1\\r\\n\\r\\n').encode()); assert expected in s.recv(1024);s.close()
subprocess.Popen([runtime,'-c','import time;time.sleep(600)','child-'+role+'-'+nonce])
print(json.dumps({'status':'passed','nonce':nonce}),flush=True)
time.sleep(60)
'''
    observed_children = {}
    for role in ('reviewer', 'repair_author'):
        state = directory / (role + '-state')
        state.mkdir()
        for name in ('HOME', 'HERMES_HOME', 'TMPDIR', 'XDG_CACHE_HOME', 'XDG_CONFIG_HOME', 'XDG_DATA_HOME', 'XDG_STATE_HOME'):
            environment[name] = str(state)
        with tempfile.TemporaryDirectory(prefix='clinical-probe-') as temporary:
            policy = Path(temporary) / 'policy.json'
            proxy_path = Path(temporary) / 'provider.sock'
            proxy = workflow._start_production_connect_proxy(unix_path=proxy_path)
            write(policy, {'readable': [str(p.resolve()) for p in (root, directory, hermes_root, Path(sys.executable).parent.parent)],
                'writable': [str(p.resolve()) for p in ([state, checkpoint] if role == 'reviewer' else [state, native, tools_work])], 'protected': [],
                'cwd': str(directory), 'proxy_socket': str(proxy_path), 'proxy_port': proxy.port})
            environment.update(HTTPS_PROXY='http://127.0.0.1:3128', HTTP_PROXY='http://127.0.0.1:3128', ALL_PROXY='http://127.0.0.1:3128', NO_PROXY='')
            stdout = (directory / (role + '.stdout.log')).open('w')
            stderr = (directory / (role + '.stderr.log')).open('w')
            process = None
            child_pids = []
            try:
                command = [*workflow._production_sandbox_prefix(policy), str(managed_python), '-c', script,
                    str(hermes_root), str(evidence), str(checkpoint), str(native), token, sys.executable,
                    native_code, json.dumps([str(p) for p in (sentinels + ([native / 'denied-marker', tools_work / 'denied-marker'] if role == 'reviewer' else [checkpoint]))]), role]
                process = subprocess.Popen(command, cwd=directory, env=environment, stdout=stdout, stderr=stderr, start_new_session=True, text=True)
                deadline = time.monotonic() + 45
                while time.monotonic() < deadline:
                    if '"status": "passed"' in (directory / (role + '.stdout.log')).read_text():
                        break
                    if process.poll() is not None:
                        raise RuntimeError('Boundary/tool probe failed: inspect retained stderr.log')
                    time.sleep(.05)
                else:
                    raise RuntimeError('Installed boundary/tool probe did not complete within 45 seconds')
                marker = ('child-' + role + '-' + token).encode()
                for path in Path('/proc').glob('[0-9]*/cmdline'):
                    try:
                        if marker in path.read_bytes().split(b'\0'):
                            child_pids.append(int(path.parent.name))
                    except OSError:
                        pass
                if not child_pids:
                    raise RuntimeError('Child-cleanup probe failed to observe its real child')
            finally:
                if process is not None:
                    workflow._reap_production_worker(process, stdout, stderr, policy, proxy)
                else:
                    stdout.close(); stderr.close(); proxy.close()
            for pid in child_pids:
                for _ in range(40):
                    if not Path(f'/proc/{pid}').exists():
                        break
                    time.sleep(.05)
                else:
                    raise RuntimeError('Worker child was not reaped by ordinary cleanup')
        observed_children[role] = child_pids
    observed = json.loads(checkpoint.read_text())
    if observed.get('nonce') != token or observed.get('read_calls') != 17:
        raise ValueError('Installed checkpoint write/readback or bounded source-read accounting failed')
    if not all(p.read_text() == 'sentinel' for p in sentinels):
        raise ValueError('A protected sentinel changed during installed worker boundary probes')
    from docx import Document
    import xml.etree.ElementTree as ET
    if Document(native / 'probe.docx').paragraphs[0].text != 'native-probe' or ET.parse(native / 'probe.xml').getroot().get('value') != 'native-probe':
        raise ValueError('Installed native DOCX/XML execution did not preserve the expected probe values')
    result = {'status': 'passed', **observed, 'child_pids': observed_children, 'roles_exercised': ['reviewer', 'repair_author'],
              'provider_route': 'CONNECT chatgpt.com:443; arbitrary target and direct network denied',
              'evidence_sha256': digest(evidence)}
    write(directory / 'observations.json', result)
    return result


def collect(workflow, root, opener, parent_reviewer, approved_by):
    from worker_capabilities import REQUIRED_CAPABILITIES, digest
    integrity = workflow._manifest_integrity(root)
    if integrity:
        raise ValueError('Installed candidate inventory differs from its immutable manifest: ' + str(integrity))
    identity = workflow._worker_capability_identity(root)
    base = root.parent.parent / 'clinical-capabilities' / identity['package_manifest_sha256']
    attempt = base / ('attempt-' + uuid.uuid4().hex)
    attempt.mkdir(parents=True)
    capabilities = {}
    report = {'schema_version': 'worker-capabilities/v1', 'identity': identity, 'capabilities': capabilities,
              'evidence_kind': 'fresh_installed_linux', 'status': 'blocked', 'paid_model_calls': 'actual worker logs',
              'transports': {name: {'path': str(path), 'sha256': digest(path)} for name, path in
                  [('recipient_opener', opener[0]), ('parent_visual_reviewer', parent_reviewer[0])]}}
    def credit(names, files):
        proofs = [{'path': str(path.relative_to(base)), 'sha256': digest(path)} for path in files]
        for name in names:
            capabilities[name] = {'status': 'passed', 'evidence_files': proofs}
    try:
        boundary_probe(workflow, root, attempt / 'boundary')
        credit({'bounded_realistic_evidence_reads', 'checkpoint_write_readback', 'native_docx_xml_execution',
                'filesystem_denials', 'provider_route', 'worker_child_cleanup'},
               [p for p in (attempt / 'boundary').rglob('*') if p.is_file()])
        case = next(c for c in workflow.branch_acceptance_inventory(root)
                    if c['descriptor']['study_type'] == 'Retrospective' and c['descriptor']['profile'].startswith('rich'))
        source = workflow.branch_acceptance_source(root, case)
        for parent_fallback in (False, True):
            run = attempt / ('parent-visual' if parent_fallback else 'ordinary-visual')
            write(run / workflow.REFERENCE, source)
            workflow.prepare(run)
            approval = workflow.approve(run, approved_by=approved_by)
            revision = run / 'revisions' / approval['revision_id']
            def dispatch(handoffs, remaining):
                visual = [h for h in handoffs if h['task'] == 'rendered_page_visual_verification'] if parent_fallback else []
                workers = [h for h in handoffs if h not in visual]
                if workers:
                    workflow._production_dispatch_handoffs(workers, remaining, revision, workflow.CERTIFIED_HERMES_CONFIGURATION,
                        skill_root=root, run_dir=run, runtime_identity=workflow.resolve_python_runtime(environment=os.environ),
                        expected_managed_hermes_identity=identity['hermes'])
                if visual:
                    parent_reviewer[1](visual, remaining, revision, workflow.CERTIFIED_HERMES_CONFIGURATION)
            # Explicit installation probe bypasses clinical admission; its only
            # inputs are the immutable packaged synthetic fixture above.
            result = workflow.run_desktop_operation(run, handoff_runner=dispatch, fallback_handoff_runner=dispatch,
                opener=opener[1], budget_seconds=2700, require_promoted_runtime=False)
            write(run / 'capability-result.json', result)
            if result.get('status') != 'passed' or not result.get('delivery', {}).get('confirmed'):
                raise RuntimeError('Installed lifecycle capability probe did not produce confirmed complete client-ready synthetic files')
            credit({'parent_visual_fallback'} if parent_fallback else {'ordinary_worker_launch', 'page_image_inspection', 'recipient_attachment_retrieval'},
                   [p for p in run.rglob('*') if p.is_file() and not p.is_symlink() and '.hermes-cache' not in p.parts])
        if set(capabilities) != REQUIRED_CAPABILITIES:
            raise RuntimeError('Installed capability coverage is incomplete')
        if workflow._worker_capability_identity(root) != identity:
            raise RuntimeError('Installed dependencies changed during capability collection')
        report['status'] = 'passed'
        write(attempt / 'assurance.json', report)
        write(base / 'assurance.json', report)
    except Exception as exc:
        report['issue'] = str(exc)
        write(attempt / 'assurance.json', report)
    return report


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--release-root', type=Path, required=True)
    parser.add_argument('--desktop-opener-command', type=Path, required=True)
    parser.add_argument('--parent-visual-review-command', type=Path, required=True)
    parser.add_argument('--approved-by', required=True)
    parser.add_argument('--authorize-paid-probes', action='store_true')
    args = parser.parse_args()
    if not args.authorize_paid_probes:
        parser.error('Explicit owner authorization for installed/paid probes is required')
    if platform.system() != 'Linux':
        parser.error('These are actual installed Linux probes; use the declared Linux host')
    root = args.release_root.resolve()
    sys.path.insert(0, str(root / 'scripts'))
    workflow = importlib.import_module('workflow')
    if Path(workflow.__file__).resolve() != root / 'scripts/workflow.py':
        raise RuntimeError('Probe imported a different candidate workflow')
    report = collect(workflow, root,
        (args.desktop_opener_command.resolve(), workflow.command_desktop_opener(args.desktop_opener_command)),
        (args.parent_visual_review_command.resolve(), workflow.command_parent_visual_reviewer(args.parent_visual_review_command)), args.approved_by)
    print(json.dumps(report, indent=2))
    return 0 if report['status'] == 'passed' else 1


if __name__ == '__main__':
    raise SystemExit(main())
