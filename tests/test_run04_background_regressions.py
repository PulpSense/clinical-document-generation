"""Focused regressions from the September 25 fourth prospective run."""

import json
from pathlib import Path

import pytest
from docx import Document

from contracts import batch_plan
from drafting import _background_claim_findings, create_drafting_request, evidence_grounded, recorded_acceptance_response, validate_response
from quality import validate_sterling_clause_contract


ROOT = Path(__file__).resolve().parents[1]
CLINICAL_BACKGROUND = (
    "A large randomized controlled trial suggested better intermediate and near vision "
    "with PureSee than with a monofocal lens. There are minimal data on visual outcomes "
    "for the PureSee and Odyssey mix-and-match combination."
)
REFERENCES = " References: Doe A. OPTH. 2026;20:1-8. doi:10.2147/OPTH.S572703."


def reference():
    value = json.loads((ROOT / "tests/fixtures/prospective-acceptance-source.json").read_text())
    value["meta"]["icf_template"] = "Sterling"
    value["study"]["background"] = CLINICAL_BACKGROUND + REFERENCES
    return value


def request_for(tmp_path, artifact, target):
    value = reference()
    batch = next(item for item in batch_plan("Prospective", "Sterling") if item.artifact == artifact and target in item.section_ids)
    path = create_drafting_request(
        repo_root=ROOT, revision_dir=tmp_path, revision_id=f"r-run04-{artifact}",
        reference=value, batch=batch, target_ids=[target], attempts={target: 1}, wave="initial",
    )
    request = json.loads(path.read_text())
    return request, recorded_acceptance_response(request)


def test_sterling_background_accepts_clinical_facts_without_bibliography_numbers():
    document = Document()
    document.add_heading("BACKGROUND", level=1)
    document.add_paragraph(CLINICAL_BACKGROUND)

    findings = validate_sterling_clause_contract(document, reference())["findings"]

    assert not [item for item in findings if item["clause_id"] == "sterling.background.context"]


def test_run04_exact_reviewed_background_separates_qualification_from_bibliography():
    replay = json.loads((ROOT / "tests/fixtures/reliability-replays/run04-background.json").read_text())
    value = reference()
    value["study"]["background"] = replay["approved_background"]
    document = Document()
    document.add_heading("BACKGROUND", level=1)
    for paragraph in replay["candidate_paragraphs"]:
        document.add_paragraph(paragraph)

    clauses = validate_sterling_clause_contract(document, value)["findings"]
    claims = _background_claim_findings(
        "icf.background", "\n".join(replay["candidate_paragraphs"]), value,
    )

    assert not [item for item in clauses if item["clause_id"] == "sterling.background.context"]
    assert any("qualified trial evidence" in item["issue"] for item in claims)

    corrected = "\n".join(replay["candidate_paragraphs"]).replace(
        "PureSee provided better intermediate and near vision",
        "PureSee suggested better intermediate and near vision",
    ).replace(
        "Placing an EDOF lens in the dominant eye and a multifocal lens in the non-dominant eye may provide a wider range of vision with fewer visual disturbances than placing the same EDOF or trifocal lens in both eyes.",
        "The benefit of placing different lens types in the two eyes has not been established for this combination.",
    )
    assert _background_claim_findings("icf.background", corrected, value) == []


def test_sterling_background_keeps_decimal_clinical_values_intact():
    value = reference()
    value["study"]["background"] = (
        "In the prior study, mean visual acuity was 0.3 logMAR at the Month 3 visit."
    )
    document = Document()
    document.add_heading("BACKGROUND", level=1)
    document.add_paragraph(value["study"]["background"])

    findings = validate_sterling_clause_contract(document, value)["findings"]

    assert not [item for item in findings if item["clause_id"] == "sterling.background.context"]


def test_ignoring_bibliography_keeps_clinical_numbers_required():
    source = "Participants have three months of follow-up after cataract surgery in this study. " + REFERENCES

    assert evidence_grounded("Participants have 3 months of follow-up after cataract surgery in this study.", source)
    assert not evidence_grounded("Participants have 2 months of follow-up after cataract surgery in this study.", source)


@pytest.mark.parametrize("opening", [
    "A large randomized trial showed PureSee provided better intermediate and near vision than a monofocal lens.",
    "Current reports suggest good vision; a large randomized trial showed PureSee provided better intermediate and near vision than a monofocal lens.",
])
def test_icf_background_rejects_strengthened_trial_claim_before_rendering(tmp_path, opening):
    request, response = request_for(tmp_path, "icf", "icf.background")
    result = response["section_results"][0]
    result["paragraphs"] = [{
        "text": opening + " There are minimal data on the PureSee and Odyssey combination.",
        "evidence_refs": ["source:study.background"], "boilerplate_refs": [],
    }]
    result["lists"] = []

    _accepted, findings = validate_response(request, response)

    assert any(item.get("field") == "icf.background" and "qualified trial evidence" in item["issue"] for item in findings)


def test_icf_background_accepts_the_source_trial_qualification(tmp_path):
    request, response = request_for(tmp_path, "icf", "icf.background")
    result = response["section_results"][0]
    result["paragraphs"] = [{
        "text": CLINICAL_BACKGROUND,
        "evidence_refs": ["source:study.background"], "boilerplate_refs": [],
    }]
    result["lists"] = []

    _accepted, findings = validate_response(request, response)

    assert not [item for item in findings if item.get("field") == "icf.background"]


def test_icf_background_rejects_unproven_combination_superiority_early(tmp_path):
    request, response = request_for(tmp_path, "icf", "icf.background")
    result = response["section_results"][0]
    result["paragraphs"] = [{
        "text": CLINICAL_BACKGROUND + " Pairing the lens types may offer a wider range of vision "
                "with fewer disturbances than bilateral use of either lens type.",
        "evidence_refs": ["source:study.background"], "boilerplate_refs": [],
    }]
    result["lists"] = []

    _accepted, findings = validate_response(request, response)

    assert any(item.get("field") == "icf.background" and "untested combination" in item["issue"] for item in findings)


def test_prs_description_rejects_unproven_combination_superiority_early(tmp_path):
    value = reference()
    value["study"]["background"] = (
        CLINICAL_BACKGROUND + " Mix-and-match implantation may offer a wider range of vision "
        "with fewer visual disturbances than bilateral use of either lens type." + REFERENCES
    )
    batch = next(item for item in batch_plan("Prospective", "Sterling") if item.batch_id == "prs-narrative")
    path = create_drafting_request(
        repo_root=ROOT, revision_dir=tmp_path, revision_id="r-run04-prs", reference=value,
        batch=batch, target_ids=["prs.detailed-description"],
        attempts={"prs.detailed-description": 1}, wave="initial",
    )
    request = json.loads(path.read_text())
    response = recorded_acceptance_response(request)
    response["narrative"]["detailed_description"]["text"] = (
        "This study examines the PureSee and Odyssey combination. Pairing the lens types "
        "may offer a wider range of vision with fewer disturbances than bilateral use of either lens alone. "
        "There are minimal data for this combination."
    )

    _accepted, findings = validate_response(request, response)

    assert any(item.get("field") == "prs.detailed-description" and "untested combination" in item["issue"] for item in findings)
