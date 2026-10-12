"""Regressions for shared presentation of approved, heterogeneous inputs."""
import copy
import json
from xml.etree import ElementTree as ET

import pytest
from docx import Document

import prs_xml
import contracts
import drafting
import quality
import rendering
from test_shared_output_projection import ROOT, COMBINATIONS, source_for


@pytest.mark.parametrize('branch', ['Prospective', 'Ambispective'])
@pytest.mark.parametrize('title,explicit,expected', [
    ('MD, Principal Investigator', None, 'MD'),
    ('Clinical Research Coordinator', None, ''),
    ('M.D., Ph.D.; Site Investigator', None, 'M.D., Ph.D.'),
    ('PharmD, Study Coordinator', None, 'PharmD'),
    ('MD PhD, Principal Investigator', None, 'MD, PhD'),
    ('MD, Principal Investigator', 'DPhil', 'DPhil'),
    ('MD, Principal Investigator', '', ''),
])
def test_xml_credentials_are_separate_from_job_titles(tmp_path, branch, title, explicit, expected):
    source = source_for(branch, 'Sterling')
    source['sites'][0]['backup_contact'] = {'name': 'Jordan Sample'}
    people = [source['parties']['principal_investigator'],
              source['parties']['study_coordinator'], source['sites'][0]['contact'],
              source['sites'][0]['investigator'], source['sites'][0]['backup_contact']]
    for person in people:
        person.pop('degree', None)
        person.pop('degrees', None)
        person['title'] = title
        if explicit is not None:
            person['degrees'] = explicit
    original = copy.deepcopy(source)
    output = tmp_path / 'study.xml'
    template = ROOT / 'assets/client-templates/prs/clinicaltrials_prs_full_placeholder_template.xml'
    assert prs_xml.generate(template, output, source, {'brief_summary': 'Summary.', 'detailed_description': 'Description.'})['status'] == 'passed'
    study = next(ET.parse(output).getroot().iter('clinical_study'))
    for path in ['overall_official', 'overall_contact', 'location/contact', 'location/investigator',
                 'location/contact_backup']:
        assert (study.findtext(path + '/degrees') or '') == expected
    assert source == original


def test_xml_degree_corruption_remains_detectable(tmp_path):
    source = source_for('Prospective', 'Advarra')
    source['sites'][0]['investigator']['title'] = 'MD, Site Investigator'
    output = tmp_path / 'study.xml'
    template = ROOT / 'assets/client-templates/prs/clinicaltrials_prs_full_placeholder_template.xml'
    prs_xml.generate(template, output, source, {'brief_summary': 'Summary.', 'detailed_description': 'Description.'})
    tree = ET.parse(output)
    study = next(tree.getroot().iter('clinical_study'))
    assert study.findtext('location/investigator/degrees') == 'MD'
    study.find('location/investigator/degrees').text = 'Site Investigator'
    tree.write(output, encoding='utf-8', xml_declaration=True)
    assert prs_xml.validate_output(output, source, template)


@pytest.mark.parametrize('path', ['location/contact', 'location/contact_backup', 'location/investigator'])
def test_xml_does_not_accept_a_job_title_in_an_expected_blank_degree(tmp_path, path):
    source = source_for('Prospective', 'Advarra')
    for key in ['contact', 'contact_backup', 'investigator']:
        source['sites'][0][key] = {'name': 'Jordan Sample', 'title': 'Study Coordinator', 'degrees': ''}
    output = tmp_path / 'study.xml'
    template = ROOT / 'assets/client-templates/prs/clinicaltrials_prs_full_placeholder_template.xml'
    assert prs_xml.generate(template, output, source,
                            {'brief_summary': 'Summary.', 'detailed_description': 'Description.'})['status'] == 'passed'
    assert prs_xml.validate_output(output, source, template) == []
    tree = ET.parse(output)
    next(tree.getroot().iter('clinical_study')).find(path + '/degrees').text = 'Study Coordinator'
    tree.write(output, encoding='utf-8', xml_declaration=True)
    assert prs_xml.validate_output(output, source, template)


@pytest.mark.parametrize('branch,family', COMBINATIONS[1:])
@pytest.mark.parametrize('structured', [False, True, None])
def test_schedule_details_are_prose_without_inferred_assignments(tmp_path, branch, family, structured):
    source = source_for(branch, family)
    if structured:
        source['procedures']['visit_schedule'] = [
            {'visit': 'Screening', 'procedures': ['Consent']},
            {'visit': 'Week 8', 'procedures': ['Questionnaire']},
        ]
    elif structured is None:
        source['procedures']['visit_schedule'] = []
    original = copy.deepcopy(source)
    rendering.render_documents(ROOT, tmp_path, source, {'protocol': [], 'icf': {}, 'prs': {}})
    document = Document(tmp_path / 'candidate/protocol.docx')
    note = source['procedures']['assessments']
    assert note in [p.text for p in document.paragraphs]
    schedules = [t for t in document.tables if t.cell(0, 0).text.strip() in {'Activity', 'Approved visit or assessment'}]
    assert not any(note in c.text for t in schedules for row in t.rows for c in row.cells)
    if structured:
        schedule = next(t for t in document.tables if t.cell(0, 0).text.strip() == 'Activity')
        assert ['Consent', 'X', ''] in [[c.text for c in r.cells] for r in schedule.rows]
        assert ['Questionnaire', '', 'X'] in [[c.text for c in r.cells] for r in schedule.rows]
    assert source == original


@pytest.mark.parametrize('branch,family', COMBINATIONS)
def test_benefit_defaults_use_approved_context_and_remain_authorized(tmp_path, branch, family):
    source = source_for(branch, family)
    source['design']['intervention_name'] = 'Example monitoring device'
    source['risks_benefits'].pop('benefits', None)
    original = copy.deepcopy(source)
    targets = ['risks-benefits.benefits'] if branch != 'Retrospective' else []
    if family:
        targets.append('icf.benefits')
    model = {'protocol': [], 'icf': {}, 'prs': {}}
    # The same governed default must reach the writer and independent reviewer.
    for target in targets:
        batch = next(b for b in contracts.batch_plan(branch, family) if target in b.section_ids)
        request_path = drafting.create_drafting_request(
            repo_root=ROOT, revision_dir=tmp_path, revision_id='context', reference=source,
            batch=batch, target_ids=[target], attempts={target: 1}, wave='initial')
        request = json.loads(request_path.read_text())
        default = request['section_contracts'][0]['fixed_boilerplate'][0]
        assert 'outcomes associated with Example monitoring device' in default['text']
        assert 'may receive no direct benefit' in default['text']
        assert 'condition, intervention, or procedures' not in default['text']
        response = drafting.response_template(request, model_id='Synthetic test writer')
        result = response['section_results'][0]
        result.update(outcome='fixed_boilerplate', paragraphs=[{
            'text': default['text'], 'evidence_refs': [],
            'boilerplate_refs': [default['boilerplate_id']]}])
        accepted, findings = drafting.validate_response(request, response)
        assert accepted, findings
        if target.startswith('icf.'):
            model['icf'][target] = copy.deepcopy(result)
        else:
            model['protocol'].append(copy.deepcopy(result))
        paths = quality.create_verification_requests(tmp_path, source, {'artifacts': []})
        review = json.loads(paths[0].read_text())
        assert review['authorized_boilerplate']['sections'][default['boilerplate_id']] == default['text']
        result['paragraphs'][0]['text'] = 'You will benefit from Example monitoring device.'
        assert drafting.validate_response(request, response)[0] is None
    # Retrospective has no benefits section, but shares the controlled resource.
    defaults = drafting.load_boilerplate(ROOT, source)
    assert 'outcomes associated with Example monitoring device' in defaults['protocol-sparse-benefits']
    rendering.render_documents(ROOT, tmp_path, source, model)
    for artifact in ['protocol', 'icf'] if family else []:
        text = '\n'.join(p.text for p in Document(tmp_path / f'candidate/{artifact}.docx').paragraphs)
        assert 'outcomes associated with Example monitoring device' in text
        assert 'may receive no direct benefit' in text
    assert source == original


def test_benefit_defaults_do_not_invent_a_subject_when_no_intervention_is_supplied():
    source = source_for('Prospective', 'Advarra')
    source['design']['intervention_name'] = 'None'
    text = drafting.load_boilerplate(ROOT, source)['icf-sparse-benefits']
    assert 'the outcomes studied in this research' in text
    assert 'associated with None' not in text


@pytest.mark.parametrize('condition,expected', [
    ('chronic knee pain', 'outcomes related to chronic knee pain'),
    (['chronic knee pain', 'mobility limitation'],
     'outcomes related to chronic knee pain and mobility limitation'),
])
def test_benefit_defaults_use_condition_context_without_a_named_intervention(condition, expected):
    source = source_for('Retrospective', None)
    source['design']['intervention_name'] = 'None'
    source['study']['condition'] = condition
    defaults = drafting.load_boilerplate(ROOT, source)
    for key in ['protocol-sparse-benefits', 'icf-sparse-benefits']:
        assert expected in defaults[key]
        assert 'may receive no direct benefit' in defaults[key]


@pytest.mark.parametrize('branch,family', COMBINATIONS[1:])
def test_explicit_approved_benefits_are_preserved(tmp_path, branch, family):
    source = source_for(branch, family)
    supplied = 'Participants may receive no direct benefit. The study may inform future rehabilitation research.'
    source['risks_benefits']['benefits'] = supplied
    model = {'protocol': [{'section_id': 'risks-benefits.benefits',
                           'paragraphs': [{'text': supplied}], 'lists': []}],
             'icf': {'icf.benefits': {'paragraphs': [{'text': supplied}], 'lists': []}}}
    rendering.render_documents(ROOT, tmp_path, source, model)
    for artifact in ['protocol', 'icf']:
        assert supplied in '\n'.join(p.text for p in Document(tmp_path / f'candidate/{artifact}.docx').paragraphs)
