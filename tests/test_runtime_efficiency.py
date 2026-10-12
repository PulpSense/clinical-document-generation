"""Public runtime regressions; controlled adapters do not establish live performance."""
import copy
import json
import threading

import pytest

import drafting
import quality
import workflow
from contracts import batch_plan
from test_clean_delivery import ROOT, CASES
from test_desktop_delivery import _visual_handoff_fixture


@pytest.mark.parametrize('branch', ['Prospective', 'Ambispective', 'Retrospective'])
@pytest.mark.parametrize('citation', ['parties.irb', 'parties.irb.name'])
def test_declared_object_evidence_and_approved_children_are_citable(tmp_path, branch, citation):
    case = next(c for c in CASES if c['descriptor']['study_type'] == branch)
    source = quality.branch_acceptance_source(ROOT, case)
    source['ethics'] = {'irb_status': 'Not submitted', 'notice': 'Synthetic testing only.'}
    original = copy.deepcopy(source)
    batch = next(b for b in batch_plan(branch, source['meta'].get('icf_template', 'Advarra'))
                 if b.batch_id == 'protocol-analysis-and-oversight')
    path = drafting.create_drafting_request(repo_root=ROOT, revision_dir=tmp_path,
        revision_id='r-runtime', reference=source, batch=batch, target_ids=['ethics'],
        attempts={'ethics': 1}, wave='initial')
    request = json.loads(path.read_text())
    response = drafting.recorded_acceptance_response(request)
    response['section_results'][0]['paragraphs'] = [{
        'text': 'The designated review board will review the study; its current status is Not submitted.',
        'evidence_refs': ['source:' + citation, 'source:ethics'], 'boilerplate_refs': [],
    }]
    accepted, findings = drafting.validate_response(request, response)
    assert accepted and not findings, findings
    assert source == original


def test_visual_fallback_starts_while_content_reviewer_is_still_running(tmp_path, monkeypatch):
    visual, _, response = _visual_handoff_fixture(tmp_path, 'protocol')
    content = {'task': 'clinical_content_verification', 'request_id': 'r1.content',
               'request_sha256': 'a' * 64, 'request_path': 'content.json', 'response_path': 'content-response.json'}
    results = iter([
        {'status': 'awaiting_hermes', 'stage': 'independent_verification', 'revision_id': 'r1',
         'handoffs': [visual, content]},
        {'status': 'blocked', 'stage': 'quality', 'findings': [], 'client_outputs': []},
    ])
    monkeypatch.setattr(workflow, 'generate', lambda *_args, **_kwargs: next(results))
    fallback_started = threading.Event()
    content_observed_fallback = []

    def primary(items, remaining):
        if items[0]['task'] == 'clinical_content_verification':
            content_observed_fallback.append(fallback_started.wait(1))
        else:
            raise RuntimeError('Visual worker exited without a response.')

    def fallback(items, remaining):
        fallback_started.set()
        (tmp_path / 'revisions/r1' / visual['response_path']).write_text(json.dumps(response))

    workflow.run_desktop_operation(tmp_path, handoff_runner=primary,
        fallback_handoff_runner=fallback, opener=lambda _: b'', budget_seconds=30)
    assert content_observed_fallback == [True]


@pytest.mark.parametrize('family', [
    ('Prospective', 'Advarra'), ('Prospective', 'Sterling'),
    ('Ambispective', 'Advarra'), ('Ambispective', 'Sterling'), ('Retrospective', None),
])
def test_unchanged_visual_evidence_survives_a_native_repair(tmp_path, monkeypatch, family):
    from test_general_document_repair import lifecycle, repair_response
    from lxml import etree
    visual_reviews = []
    case = next(c for c in CASES if c['descriptor']['study_type'] == family[0]
                and (family[1] is None or c['descriptor'].get('icf_template') == family[1]))

    def review(request, response, revision, number):
        if number == 1:
            response['status'] = 'blocked'
            response['findings'] = [{'finding_id': 'native-content', 'category': 'content',
                'artifact': 'protocol' if family[0] == 'Retrospective' else 'study.xml',
                'target_ids': ['introduction' if family[0] == 'Retrospective' else 'prs.detailed-description'],
                'check': 'source_supported', 'issue': 'Restore the approved qualification.', 'material': True}]
            row = next(row for row in response['section_assessments']
                       if row['section_id'] == ('introduction' if family[0] == 'Retrospective' else 'prs.detailed-description'))
            row['status'] = 'blocked'
        return response

    def repair(request, revision):
        if family[0] == 'Retrospective':
            from docx import Document
            path = revision / request['workspace_path'] / 'protocol.docx'
            doc = Document(path)
            doc.add_paragraph('Approved source qualification restored.')
            doc.save(path)
        else:
            path = revision / request['workspace_path'] / 'study.xml'
            tree = etree.parse(str(path))
            node = tree.find('.//detailed_description')
            node.text = (node.text or '') + ' Approved source qualification restored.'
            tree.write(str(path), encoding='UTF-8', xml_declaration=True)
        return repair_response(request)

    def visual(request, response, revision):
        visual_reviews.append(request)
        return response

    result, revision, repairs, content = lifecycle(tmp_path, monkeypatch, repair,
        case=case, review=review, visual_review=visual)
    assert result['status'] == 'passed', result
    assert len(repairs) == 1
    assert len(content) == 2  # All content and cross-document checks are fresh.
    assert len(visual_reviews) == 2
    if family[0] != 'Retrospective':
        assert {r['review_set'] for r in visual_reviews} == {1}


@pytest.mark.parametrize('change', ['none', 'docx', 'pdf', 'page', 'renderer', 'substitutions', 'incomplete', 'mixed'])
def test_render_reuse_requires_exact_files_and_current_renderer(tmp_path, governed_pdfium, change):
    from docx import Document
    from test_render_assurance import _bundle, _structural_validation, _write_text_pdf
    candidate = tmp_path / 'candidate'
    candidate.mkdir()
    doc = Document()
    doc.add_paragraph('A complete candidate.')
    doc.save(candidate / 'protocol.docx')
    if change == 'mixed':
        doc.save(candidate / 'icf.docx')
    exports = []
    office = {'kind': 'LibreOffice', 'path': '/controlled/office'}

    def export(docx, output_dir, identity, **kwargs):
        exports.append(docx.name)
        pdf = output_dir / (docx.stem + '.pdf')
        _write_text_pdf(pdf)
        return pdf

    def rasterize(pdf, output_dir, identity, **kwargs):
        output_dir.mkdir(parents=True, exist_ok=True)
        page = output_dir / 'page-1.png'
        page.write_bytes(b'controlled page image')
        return [page]

    def run(reuse=None):
        return quality.render_assurance(ROOT, tmp_path, {}, contracted_bundle=_bundle(),
            structural_validation=_structural_validation(tmp_path, *sorted(p.name for p in candidate.iterdir())),
            renderer_identities=[office], page_renderer_identities=[governed_pdfium],
            font_probe=lambda *_args, **_kwargs: (True, 'available'),
            office_exporter=export, page_exporter=rasterize, blank_page_detector=lambda _: [],
            **({'reuse_render_report': reuse} if reuse else {}))

    first = run()
    assert first['status'] == 'passed'
    prior = copy.deepcopy(first['render'])
    row = next(row for row in prior['artifacts'] if row['artifact'] == 'protocol')
    if change in {'docx', 'mixed'}:
        doc.add_paragraph('Changed native content.')
        doc.save(candidate / 'protocol.docx')
    elif change in {'pdf', 'page'}:
        path = tmp_path / (row['pdf'] if change == 'pdf' else row['pages'][0]['path'])
        path.write_bytes(b'changed bytes')
    elif change == 'renderer':
        office['path'] = '/controlled/different-office'
    elif change == 'substitutions':
        row['font_substitutions'] = {'Previous font': 'Substitute'}
    elif change == 'incomplete':
        row['pages'] = []
    second = run(prior)
    assert second['status'] == 'passed', second
    assert len(exports) == (3 if change == 'mixed' else 1 if change == 'none' else 2)
    for artifact in second['render']['artifacts']:
        for key in ('docx', 'pdf'):
            assert quality.sha256_file(tmp_path / artifact[key]) == artifact[key + '_sha256']
        for page in artifact['pages']:
            assert quality.sha256_file(tmp_path / page['path']) == page['sha256']


def test_drafting_includes_declared_evidence_outside_legacy_batch_families(tmp_path):
    case = next(c for c in CASES if c['descriptor']['study_type'] == 'Prospective')
    source = quality.branch_acceptance_source(ROOT, case)
    source.setdefault('regulatory', {}).setdefault('prs', {})['oversight_info'] = {'status': 'Not submitted'}
    batch = next(b for b in batch_plan('Prospective') if b.batch_id == 'protocol-analysis-and-oversight')
    path = drafting.create_drafting_request(repo_root=ROOT, revision_dir=tmp_path,
        revision_id='r-runtime', reference=source, batch=batch, target_ids=['ethics'],
        attempts={'ethics': 1}, wave='initial')
    request = json.loads(path.read_text())
    response = drafting.recorded_acceptance_response(request)
    response['section_results'][0]['paragraphs'][0]['evidence_refs'] = ['source:regulatory.prs.oversight_info.status']
    response['section_results'][0]['paragraphs'][0]['boilerplate_refs'] = []
    accepted, findings = drafting.validate_response(request, response)
    assert accepted and not findings, findings


@pytest.mark.parametrize('change', ['none', 'source', 'renderer', 'visual-concern', 'repair-author',
                                    'response', 'missing-page', 'page-bytes'])
def test_visual_review_reuse_is_invalidated_by_new_authority_or_incomplete_evidence(tmp_path, change):
    from test_review_resolutions import review_candidate
    from test_clean_delivery import controlled_render
    from test_handoff_quality import acceptance_verification
    _, original, _ = review_candidate(tmp_path, branch='Retrospective')
    source = original['approved_source']
    render = controlled_render(ROOT, tmp_path)['render']
    paths = quality.create_verification_requests(tmp_path, source, render)
    visual_path = next(p for p in paths if '.verify.visual.protocol' in p.name)
    visual = json.loads(visual_path.read_text())
    response = acceptance_verification(visual)
    response_path = tmp_path / visual['response_path']
    response_path.write_text(json.dumps(response))
    assert quality.verification_response_is_complete(tmp_path, visual_path)
    context = None
    if change == 'source':
        source = copy.deepcopy(source)
        source['study']['title'] = 'Changed approved study title'
    elif change == 'renderer':
        render['artifacts'][0]['renderer']['version'] = 'new-version'
    elif change == 'visual-concern':
        context = {'prior_findings': [{'finding_id': 'table', 'category': 'visual',
            'artifact': 'protocol', 'check': 'bad_table_split', 'issue': 'Reassess table readability.'}]}
    elif change == 'repair-author':
        context = {'repair_authors': [response['producer']]}
    elif change == 'response':
        response['request_sha256'] = 'f' * 64
        response_path.write_text(json.dumps(response))
    elif change == 'missing-page':
        response['page_assessments'] = []
        response_path.write_text(json.dumps(response))
    elif change == 'page-bytes':
        page = tmp_path / render['artifacts'][0]['pages'][0]['path']
        page.write_bytes(b'new page bytes')
        render['artifacts'][0]['pages'][0]['sha256'] = quality.sha256_file(page)
    paths = quality.create_verification_requests(tmp_path, source, render, review_set=2, review_context=context)
    current_visual = next(p for p in paths if '.verify.visual.protocol' in p.name)
    assert json.loads(current_visual.read_text())['review_set'] == (1 if change == 'none' else 2)


@pytest.mark.parametrize('citation', ['parties', 'parties.sponsor.name', 'parties.irb.invented'])
def test_object_citations_do_not_authorize_unrelated_or_nonexistent_evidence(tmp_path, citation):
    case = next(c for c in CASES if c['descriptor']['study_type'] == 'Prospective')
    source = quality.branch_acceptance_source(ROOT, case)
    batch = next(b for b in batch_plan('Prospective') if b.batch_id == 'protocol-analysis-and-oversight')
    path = drafting.create_drafting_request(repo_root=ROOT, revision_dir=tmp_path,
        revision_id='r-runtime', reference=source, batch=batch, target_ids=['ethics'],
        attempts={'ethics': 1}, wave='initial')
    request = json.loads(path.read_text())
    response = drafting.recorded_acceptance_response(request)
    row = response['section_results'][0]['paragraphs'][0]
    row.update(evidence_refs=['source:' + citation], boilerplate_refs=[])
    _, findings = drafting.validate_response(request, response)
    assert findings


def test_content_only_review_wave_has_a_worker_budget_within_operation_deadline(tmp_path, monkeypatch):
    handoff = {'task': 'clinical_content_verification', 'request_id': 'r1.content',
               'request_sha256': 'a' * 64, 'request_path': 'content.json', 'response_path': 'content-response.json'}
    monkeypatch.setattr(workflow, 'generate', lambda *_args, **_kwargs: {
        'status': 'awaiting_hermes', 'stage': 'independent_verification',
        'revision_id': 'r1', 'handoffs': [handoff]})
    monkeypatch.setattr(workflow.time, 'sleep', lambda _: None)
    budgets = []
    def runner(items, remaining):
        budgets.append(remaining)
        raise RuntimeError('Controlled stalled reviewer.')
    result = workflow.run_desktop_operation(tmp_path, handoff_runner=runner,
        opener=lambda _: b'', budget_seconds=2700)
    assert result['status'] == 'blocked'
    assert budgets and all(0 < budget <= 720 for budget in budgets)
