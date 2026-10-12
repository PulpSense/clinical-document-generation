"""Approved relationships must reach Word regardless of the intake field used."""
import copy
import json
from pathlib import Path

import pytest
from docx import Document

from contracts import protocol_table_contracts
from rendering import render_documents

ROOT = Path(__file__).resolve().parents[1]


def records():
    return [
        {"visit": "Screening", "timing": "Preoperative", "procedures": ["Consent", "Demographics"]},
        {"visit": "Month 3", "timing": "3 months postoperative", "procedures": "Visual acuity at 66 cm; Questionnaire"},
    ]


@pytest.mark.parametrize("branch", ["Prospective", "Ambispective", "Retrospective"])
@pytest.mark.parametrize("field", ["assessments", "evaluation", "visit_schedule", "visit_schedule_table"])
def test_explicit_records_have_identical_visit_assignments(branch, field):
    source = {"meta": {"study_type": branch}, "procedures": {field: records()}}
    original = copy.deepcopy(source)
    table = protocol_table_contracts(source)["schedule-of-assessments"]
    assert table["rows"][0] == ["Activity", "Screening\n(Preoperative)", "Month 3\n(3 months postoperative)"]
    assert table["rows"][2:] == [
        ["Consent", "X", ""], ["Demographics", "X", ""],
        ["Visual acuity at 66 cm", "", "X"], ["Questionnaire", "", "X"],
    ]
    assert not table["supplemental_notes"]
    assert source == original


@pytest.mark.parametrize("branch,family", [("Prospective", "Advarra"), ("Prospective", "Sterling"),
                                          ("Ambispective", "Advarra"), ("Ambispective", "Sterling")])
def test_public_docx_assembly_constructs_real_matrix(tmp_path, branch, family):
    source = json.loads((ROOT / "tests/fixtures/prospective-acceptance-source.json").read_text())
    source["meta"].update(study_type=branch, icf_template=family)
    source["procedures"] = {"assessments": records()}
    render_documents(ROOT, tmp_path, source, {"protocol": [], "icf": {}, "prs": {}}, artifact_names={"protocol"})
    doc = Document(tmp_path / "candidate/protocol.docx")
    table = next(t for t in doc.tables if t.cell(0, 0).text.strip() == "Activity")
    assert len(table.columns) == 3
    assert [c.text for c in table.rows[2].cells] == ["Consent", "X", ""]
    assert not any("procedures:" in c.text or "timing:" in c.text for row in table.rows for c in row.cells)


def test_record_metadata_is_retained_without_duplicate_schedule_dump():
    source = {"procedures": {"assessments": [{**records()[0], "restriction": "Before dilation only"}]}}
    table = protocol_table_contracts(source)["schedule-of-assessments"]
    assert table["rows"][2] == ["Consent", "X"]
    assert "Before dilation only" in " ".join(table["supplemental_notes"])
    assert "Screening" in " ".join(table["supplemental_notes"])
    assert "procedures:" not in " ".join(table["supplemental_notes"])


def test_narrative_is_not_assigned_to_visits_by_guessing():
    source = {"procedures": {"assessments": "Collect 5 mL blood if clinically indicated."}}
    table = protocol_table_contracts(source)["schedule-of-assessments"]
    assert table["rows"] == []
    assert table["supplemental_notes"] == [source["procedures"]["assessments"]]


def test_duplicate_records_merge_but_different_timings_remain_distinct():
    source = {"procedures": {"visit_schedule": records(), "assessments": records()}}
    table = protocol_table_contracts(source)["schedule-of-assessments"]
    assert len(table["rows"][0]) == 3
    assert not table["supplemental_notes"]
    source["procedures"]["assessments"][1]["timing"] = "Month 6"
    table = protocol_table_contracts(source)["schedule-of-assessments"]
    assert len(table["rows"][0]) == 4


@pytest.mark.parametrize("field", ["assessments", "evaluation"])
def test_explicit_activity_matrix_needs_no_second_schedule_copy(field):
    source = {"procedures": {field: [
        {"activity": "Blood sample at 5 mL", "Screening": "X", "Week 6": ""},
        {"activity": "Questionnaire", "Screening": "", "Week 6": "X"},
    ]}}
    table = protocol_table_contracts(source)["schedule-of-assessments"]
    assert table["rows"][2:] == [["Blood sample at 5 mL", "X", ""], ["Questionnaire", "", "X"]]
    assert not table["supplemental_notes"]


@pytest.mark.parametrize("field", ["assessments", "evaluation"])
def test_single_explicit_record_uses_the_same_table_projection(field):
    table = protocol_table_contracts({"procedures": {field: records()[0]}})["schedule-of-assessments"]
    assert table["rows"][2:] == [["Consent", "X"], ["Demographics", "X"]]
    assert not table["supplemental_notes"]


@pytest.mark.parametrize("branch", ["Prospective", "Ambispective", "Retrospective"])
def test_independent_reviews_assess_table_usability_not_just_fact_presence(tmp_path, branch):
    from quality import create_verification_requests
    source = {"meta": {"study_type": branch, "icf_template": "Advarra"}, "procedures": {"assessments": records()}}
    content, visual = [json.loads(path.read_text()) for path in
                       create_verification_requests(tmp_path, source, {"artifacts": []})]
    assert "raw field/value dumps" in content["instructions"]
    assert "raw field/value dumps" in visual["instructions"]
    assert "visit-to-activity associations" in content["instructions"]
    assert "visit-to-activity associations" in visual["instructions"]
