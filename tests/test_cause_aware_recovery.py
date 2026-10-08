"""Recovery preserves independently usable evidence without acting on malformed items."""
from finding_policy import blocks_construction
import copy
import json
from pathlib import Path
from xml.etree import ElementTree as ET

from docx import Document
import pytest

import drafting
import quality
import workflow
from test_concept_reference_grounding import ANALYSIS_PLAN
from test_handoff_quality import acceptance_verification

ROOT = Path(__file__).resolve().parents[1]


@pytest.mark.parametrize("malformed_id", [None, "valid-companion", 123])
def test_malformed_finding_identity_does_not_redraft_valid_companions(tmp_path, malformed_id):
    request_path = quality.create_verification_requests(
        tmp_path, {"meta": {"study_type": "Retrospective"}}, {"artifacts": []},
    )[0]
    request = json.loads(request_path.read_text())
    response = acceptance_verification(request)
    targets = [row["section_id"] for row in response["section_assessments"]][:2]
    response["status"] = "blocked"
    response["findings"] = [
        {"finding_id": "valid-companion", "target_ids": [targets[0]], "issue": "Approved study purpose is materially omitted."},
        {"finding_id": malformed_id, "target_ids": [targets[1]], "issue": "Unlocalized malformed reviewer evidence."},
    ]
    for row in response["section_assessments"]:
        if row["section_id"] in targets:
            row["status"] = "blocked"
    (tmp_path / request["response_path"]).write_text(json.dumps(response))

    findings, _ = quality.validate_verifications(tmp_path)
    malformed = [item for item in findings if item.get("verification_source_finding", {}).get("issue") == "Unlocalized malformed reviewer evidence."]
    assert malformed and all(item["recovery_class"] == "verifier_transient" for item in malformed)
    assert all(blocks_construction(item) for item in malformed)
    assert all(item["finding_policy"]["classification"] == "technical_failure" for item in malformed)
    companions = [item for item in findings if item.get("verification_source_finding", {}).get("issue") == "Approved study purpose is materially omitted."]
    assert companions
    assert {item["recovery_class"] for item in companions} == ({"verifier_transient"} if malformed_id == "valid-companion" else {"drafting_defect"})



@pytest.mark.parametrize("interrupt", [False, True])
def test_public_mixed_recovery_requires_fresh_review_and_preserves_original_approval(tmp_path, governed_pdfium, interrupt):
    """Real quantity/XML mutations and controlled exact-artifact layout evidence.

    External reviewers are synthetic; this proves owner routing and measured
    repair, not independent clinical or visual certification.
    """
    run = tmp_path / "mixed"
    reference_path = run / "reference/study.reference.json"
    reference_path.parent.mkdir(parents=True)
    source = json.loads((ROOT / "tests/fixtures/prospective-acceptance-source.json").read_text())
    source["meta"]["icf_template"] = "Sterling"
    source["statistics"].update(analysis_plan=ANALYSIS_PLAN + " Interpret results cautiously because the sample size is 40.", software=None)
    reference_path.write_text(json.dumps(source))
    assert workflow.prepare(run)["status"] == "awaiting_approval"
    assert workflow.approve(run, approved_by="Controlled mixed recovery")["status"] == "passed"
    approval = copy.deepcopy(json.loads(reference_path.read_text())["approval"])
    revision = run / "revisions" / approval["revision_id"]
    observed = {"drafts": 0, "xml_mutated": False, "malformed": False, "visual": False, "review_sets": set(), "interrupted": False, "accepted": None}

    def runner(handoffs, remaining):
        assert remaining > 0
        for handoff in handoffs:
            path = revision / handoff["request_path"]
            request = json.loads(path.read_text())
            if "verification-requests" not in handoff["request_path"]:
                response = drafting.recorded_acceptance_response(request)
                for section in response.get("section_results", []):
                    if section["section_id"] == "analysis-plan.methodology":
                        section["paragraphs"][0]["text"] = source["statistics"]["analysis_plan"]
                    if section["section_id"] == "analysis-plan.considerations":
                        observed["drafts"] += 1
                        quantity = 99 if observed["drafts"] == 1 else 40
                        section["paragraphs"] = [{"text": f"Interpret results cautiously because the sample size is {quantity}.", "evidence_refs": ["source:statistics.analysis_plan"], "boilerplate_refs": []}]
                        section["paragraphs"].append(copy.deepcopy(section["paragraphs"][0]))
                        section["lists"] = []
                response_path = revision / request["response_path"]
                response_path.parent.mkdir(parents=True, exist_ok=True)
                response_path.write_text(json.dumps(response))
                continue
            observed["review_sets"].add(request["review_set"])
            if observed["accepted"] is None:
                observed["accepted"] = {p.name: p.read_bytes() for p in (revision / "hermes/accepted").glob("*.json")}
            response = acceptance_verification(request)
            if request["task"] == "clinical_content_verification" and not observed["xml_mutated"]:
                xml_path = revision / "candidate/study.xml"
                tree = ET.parse(xml_path)
                next(tree.getroot().iter("enrollment")).text = "99"
                tree.write(xml_path, encoding="utf-8", xml_declaration=True)
                assert any(f.get("field") == "enrollment" for f in quality.deterministic_content_check(revision, source))
                observed["xml_mutated"] = True
            elif request["task"] == "clinical_content_verification" and not observed["malformed"]:
                response["status"] = "blocked"
                response["findings"] = [{"finding_id": None, "target_ids": ["analysis-plan.considerations"], "issue": "Malformed reviewer finding must be localized again."}]
                observed["malformed"] = True
            if request["task"] == "rendered_page_visual_verification" and request["artifacts"][0]["artifact"] == "protocol" and request["review_set"] >= 2 and not observed["visual"]:
                response["status"] = "blocked"
                response["page_assessments"][0]["status"] = "blocked"
                response["findings"] = [{"finding_id": "localized-heading", "artifact": "protocol", "page": response["page_assessments"][0]["page"], "check": "orphan_heading", "element": "5. INTRODUCTION", "target_ids": ["layout:protocol"], "issue": "Keep the introduction heading with its first content block."}]
                observed["visual"] = True
            workflow._save_verification_response(revision, path, lambda _request, response=response: response)
            assert not (run / "output").exists()
        if interrupt and observed["malformed"] and not observed["interrupted"]:
            # Simulate a terminated caller after external evidence landed. Resume
            # must preserve its original deadline and recover malformed evidence.
            observed["interrupted"] = True
            content_request_path = next(p for p in (revision / "hermes/verification-requests").glob("*.json")
                if json.loads(p.read_text())["task"] == "clinical_content_verification")
            content_request = json.loads(content_request_path.read_text())
            (revision / content_request["response_path"]).write_text("{")
            raise KeyboardInterrupt("Interrupted external reviewer transport")

    options = dict(handoff_runner=runner, opener=lambda path: Path(path).read_bytes(),
        operation_id="mixed-cause-recovery", release_identity={"package_fingerprint": "controlled-mixed-recovery"})
    if interrupt:
        with pytest.raises(KeyboardInterrupt, match="Interrupted external reviewer transport"):
            workflow.run_desktop_operation(run, **options)
        original_deadline = json.loads((run / "logs/desktop-operation-mixed-cause-recovery.json").read_text())["deadline_at_epoch"]
    result = workflow.run_desktop_operation(run, **options)
    assert result["status"] == "passed", result
    assert result["delivery"]["confirmed"]
    assert sorted(result["client_outputs"]) == ["output/icf.docx", "output/protocol.docx", "output/study.xml"]
    state = json.loads(reference_path.read_text())
    assert state["approval"] == approval
    assert {p.name: p.read_bytes() for p in (revision / "hermes/accepted").glob("*.json")} == observed["accepted"]
    assert observed["drafts"] == 2
    assert observed["xml_mutated"] and observed["malformed"] and observed["visual"]
    assert len(observed["review_sets"]) >= 3
    attempts = [json.loads(p.read_text()) for p in (revision / "attempts").glob("*/attempt-manifest.json")]
    actions = [action for attempt in attempts for action in attempt["recovery_actions"]]
    assert {a["triggering_finding"]["recovery_class"] for a in actions} >= {"drafting_defect", "deterministic_structure_defect", "visual_defect", "verifier_transient"}
    assert all(a["outcome_status"] == "measured" for a in actions)
    assert state["generation"]["verification_attempts"]
    deadline = json.loads((run / "logs/desktop-operation-mixed-cause-recovery.json").read_text())
    assert deadline["deadline_at_epoch"] - deadline["started_at_epoch"] == 2700
    if interrupt:
        assert deadline["deadline_at_epoch"] == original_deadline
    manifest = json.loads((revision / "delivery-manifest.json").read_text())
    assert manifest["quality"]["final_exact_artifact_review"]["every_page"]
    for item in manifest["client_outputs"]:
        delivered = run / item["path"]
        assert delivered.read_bytes() == (revision / "candidate" / delivered.name).read_bytes()
        assert quality.sha256_file(delivered) == item["sha256"]
    assert next(ET.parse(run / "output/study.xml").getroot().iter("enrollment")).text == "40"
    body = "\n".join(p.text for p in Document(run / "output/protocol.docx").paragraphs)
    assert "sample size is 40" in body and "sample size is 99" not in body

@pytest.mark.parametrize("corruption", ["json", "binding", "producer"])
def test_malformed_response_authorizes_only_response_recovery_of_its_canonical_request(tmp_path, corruption):
    from test_owner_aware_recovery import review_case
    run, revision, reference_path, source, protocol, _ = review_case(tmp_path, table_error=False)
    request_path = next(p for p in (revision / "hermes/verification-requests").glob("*.json")
                        if json.loads(p.read_text())["task"] == "clinical_content_verification")
    request = json.loads(request_path.read_text())
    response_path = revision / request["response_path"]
    response = json.loads(response_path.read_text())
    if corruption == "json":
        response_path.write_text("{")
    else:
        if corruption == "binding":
            response["request_sha256"] = "0" * 64
        else:
            response["producer"] = {}
        response_path.write_text(json.dumps(response))
    findings, _ = quality.validate_verifications(revision)
    malformed = next(f for f in findings if f.get("verification_request_id") == request["request_id"])
    assert malformed["recovery_class"] == "verifier_transient"
    render_report = json.loads((revision / "candidate-build.json").read_text())["render_report"]
    assert quality.verification_recovery_request_findings(revision, request["request_id"],
        expected_task=request["task"], current_review_set=1, reference=source,
        render_report=render_report, originating_finding=malformed) == []
    # Corrupt evidence never authorizes a substantive clinical repair.
    assert quality.verification_recovery_request_findings(revision, request["request_id"],
        expected_task=request["task"], current_review_set=1, reference=source,
        render_report=render_report, originating_finding={**malformed, "recovery_class": "drafting_defect"})

@pytest.mark.parametrize("corruption", ["request", "artifact", "missing-document", "page"])
def test_response_recovery_rejects_forged_request_and_changed_artifact_before_archive(tmp_path, corruption):
    from test_owner_aware_recovery import review_case
    run, revision, reference_path, source, protocol, _ = review_case(tmp_path, table_error=False)
    request_path = next(p for p in (revision / "hermes/verification-requests").glob("*.json")
                        if json.loads(p.read_text())["task"] == "clinical_content_verification")
    request = json.loads(request_path.read_text())
    (revision / request["response_path"]).write_text("{")
    findings, _ = quality.validate_verifications(revision)
    malformed = next(f for f in findings if f.get("verification_request_id") == request["request_id"])
    if corruption == "request":
        request["approved_source"]["study"]["title"] = "Forged study authority"
        request["request_sha256"] = quality.verification_request_sha256(request)
        request_path.write_text(json.dumps(request))
    elif corruption == "artifact":
        document = Document(protocol)
        document.add_paragraph("Changed after the canonical request was bound.")
        document.save(protocol)
    elif corruption == "missing-document":
        protocol.unlink()
    else:
        (revision / "rendered/protocol/page-1.png").write_bytes(b"Changed page bytes")
    if corruption == "missing-document":
        retained, _ = quality.validate_verifications(revision)
        assert any("Visual verification request is stale for candidate/protocol.docx" in f["issue"] for f in retained)
    if corruption == "page":
        report = json.loads((revision / "candidate-build.json").read_text())["render_report"]
        assessed = quality.quality_report(revision, source, report, None)
        assert any("Visual verification request is stale for rendered/protocol/page-1.png" in f["issue"] for f in assessed["findings"])
        assert assessed["status"] == "blocked"
    before_reference = reference_path.read_bytes()
    result = workflow._quality_retry(run, reference_path, {"generation": {}}, source, revision,
                                    {}, [malformed], "quality", require_promoted_runtime=False)
    assert result["status"] == "blocked"
    assert result["stage"] == "internal_recovery_routing"
    assert result["client_outputs"] == []
    assert reference_path.read_bytes() == before_reference
    assert not (revision / "attempts").exists()
    assert not (revision / "gate-attempt-journal.json").exists()
