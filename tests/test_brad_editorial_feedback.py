"""Regression checks for section ownership exposed by reviewer feedback."""

import json
from pathlib import Path

from docx import Document

from contracts import batch_plan, input_findings, protocol_contract
from drafting import _coverage_findings, create_drafting_request
from quality import create_verification_requests
from rendering import _endpoint_synopsis, _protocol_followup_summary, render_documents


ROOT = Path(__file__).resolve().parents[1]


def _source():
    return json.loads((ROOT / "tests/fixtures/prospective-acceptance-source.json").read_text())


def test_objectives_request_asks_for_purpose_without_endpoint_inventory(tmp_path):
    source = _source()
    batch = next(item for item in batch_plan("Prospective") if item.batch_id == "protocol-foundations")
    path = create_drafting_request(
        repo_root=ROOT,
        revision_dir=tmp_path,
        revision_id="r-editorial-objectives",
        reference=source,
        batch=batch,
        attempts={section_id: 1 for section_id in batch.section_ids},
        wave="initial",
    )
    request = json.loads(path.read_text())
    sections = {item["section_id"]: item for item in request["section_contracts"]}
    objectives = sections["objectives"]
    design = sections["study-design.design"]

    assert "endpoint-inventory" not in objectives["concept_ownership"]["owns"]
    assert "endpoint-inventory" in design["concept_ownership"]["owns"]
    assert "endpoints.secondary" not in objectives["minimum_evidence"]
    assert "endpoints.other" not in objectives["minimum_evidence"]
    assert {"endpoints.primary", "endpoints.secondary", "endpoints.other"} <= set(design["minimum_evidence"])


def test_methods_and_schedule_requests_do_not_demand_irrelevant_restatement():
    sections = {item.section_id: item for item in protocol_contract("Prospective")}
    assert "study.hypothesis" not in sections["study-procedure.measurements"].evidence
    assert sections["evaluation-procedures"].source_coverage == "table_with_notes"
    assert "statistics.analysis_plan" in sections["analysis-plan.considerations"].evidence
    assert sections["analysis-plan.considerations"].boilerplate_key is None
    assert any("operational detail" in item for item in sections["confidentiality"].content_expectations)

    boilerplate = json.loads((ROOT / "references/fixed-clinical-boilerplate.json").read_text())
    assert "unscheduled visit" in boilerplate["sections"]["unscheduled"].casefold()
    assert "unscheduled contact" not in boilerplate["sections"]["unscheduled"].casefold()


def test_schedule_prose_accepts_brief_table_introduction():
    source = _source()
    request = {"approved_source": source}
    contract = {
        "source_coverage": "table_with_notes",
        "minimum_evidence": ["procedures.assessments", "procedures.visit_schedule_table"],
    }
    assert not _coverage_findings(
        request, contract, "evaluation-procedures",
        "The scheduled visits and study assessments are shown in Table 15.1.",
        ["source:procedures.visit_schedule_table"],
    )


def test_content_review_explicitly_checks_section_purpose_and_editorial_relevance(tmp_path):
    request = json.loads(create_verification_requests(tmp_path, _source(), {"artifacts": []})[0].read_text())
    instructions = request["instructions"].casefold()
    assert "section purpose" in instructions
    assert "redundant" in instructions
    assert "self-referential" in instructions
    assert "unsupported clinical product claims" in instructions
    assert "generic section 16 privacy prose" in instructions
    assert "section 18.5 must state the completion rule" in instructions


def test_full_title_cannot_double_as_approved_short_title():
    source = _source()
    source["study"]["short_title"] = source["study"]["title"]
    assert any(item["field"] == "study.short_title" for item in input_findings(source))
    source["study"].pop("short_title")
    source["study"]["title"] = "A very long full study title describing the intervention, population, setting, outcomes, and follow-up period"
    assert any(item["field"] == "study.short_title" for item in input_findings(source))


def test_general_information_uses_complete_endpoint_synopsis_and_final_visit():
    source = _source()
    synopsis = _endpoint_synopsis(source)
    assert "Primary endpoint:" in synopsis
    assert "Section 8.1" in synopsis
    assert "Section 6" not in synopsis
    source["study"]["timeline"] = "Enrollment: 6 months; follow-up: 3 months; data analysis: 1 month."
    source["procedures"]["visit_schedule_table"] = [
        {"visitNumber": "1", "visitName": "Preoperative screening", "timing": "Before surgery", "procedures": []},
        {"visitNumber": "2", "visitName": "3-month postoperative visit", "timing": "3 months postoperatively", "procedures": []},
    ]
    assert _protocol_followup_summary(source) == "3-month postoperative visit"


def test_completion_and_bias_contracts_do_not_require_repeated_design_or_visit_inventory():
    sections = {item.section_id: item for item in protocol_contract("Prospective")}
    assert "every approved visit" not in " ".join(sections["endpoint-criteria.study-completion"].content_expectations)
    assert "no masking" not in " ".join(sections["study-design.bias"].content_expectations)
    assert "study.timeline" in sections["endpoint-criteria.study-completion"].evidence


def test_generated_header_and_schedule_activity_follow_reviewed_case(tmp_path):
    source = _source()
    source["study"]["short_title"] = "Recovery Study"
    source["procedures"]["visit_schedule_table"] = [
        {"visitNumber": "1", "visitName": "Screening", "timing": "Preoperative", "procedures": ["demographics"]},
        {"visitNumber": "2", "visitName": "Postoperative visit", "timing": "Month 3", "procedures": ["Visual acuity"]},
    ]
    render_documents(ROOT, tmp_path, source, {"protocol": [], "icf": {}, "prs": {}}, artifact_names={"protocol"})
    document = Document(tmp_path / "candidate/protocol.docx")
    headers = [cell.text for section in document.sections for table in section.header.tables for row in table.rows for cell in row.cells]
    assert "Recovery Study" in headers
    assert source["study"]["title"] not in headers
    matrix = next(table for table in document.tables if any(cell.text == "Demographics" for row in table.rows for cell in row.cells))
    assert any(row.cells[0].text == "Demographics" for row in matrix.rows)
