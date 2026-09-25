"""Focused replay cases from the September 2026 blocked prospective run."""

import json
from pathlib import Path
from xml.etree import ElementTree as ET

import pytest
from docx import Document

import quality
import rendering
import workflow
from contracts import facility_projection, get_path, protocol_contract, protocol_table_contracts, source_contract
from drafting import _coverage_findings
from prs_xml import generate, validate_output


ROOT = Path(__file__).resolve().parents[1]
PRS_TEMPLATE = ROOT / "assets/client-templates/prs/clinicaltrials_prs_full_placeholder_template.xml"
PRS_REFERENCE = ROOT / "assets/client-templates/reference/prs-manual-reference.xml"
ADDRESS = "300 Test Clinic Road, Test City, New York 10003, USA"


def source():
    result = json.loads((ROOT / "tests/fixtures/prospective-acceptance-source.json").read_text())
    result["meta"]["icf_template"] = "Sterling"
    result["sites"][0]["facility"] = {
        "name": "Synthetic Ophthalmology Research Center",
        "address": ADDRESS,
    }
    return result


def test_sterling_site_merge_fields_are_the_required_address_destination(tmp_path):
    reference = source()
    rendering.render_documents(
        ROOT, tmp_path, reference, {"protocol": [], "icf": {}, "prs": {}},
        artifact_names={"icf"},
    )
    path = tmp_path / "candidate/icf.docx"
    assert not [
        item for item in quality.audit_source_surfaces(path, reference)
        if item["field"].startswith("sites[0].facility")
    ]

    document = Document(path)
    field = next(paragraph for paragraph in document.paragraphs if "«Address»" in paragraph.text)
    field.text = field.text.replace("«Address»", "")
    document.save(path)
    assert any(
        item["field"] == "sites[0].facility.address"
        for item in quality.audit_source_surfaces(path, reference)
    )


def test_complete_flat_address_projects_to_prs_components_and_is_validated(tmp_path):
    reference = source()
    projected = facility_projection(reference["sites"][0]["facility"])
    assert (projected["city"], projected["state"], projected["postal_code"], projected["country"]) == (
        "Test City", "New York", "10003", "USA",
    )
    output = tmp_path / "study.xml"
    report = generate(PRS_TEMPLATE, output, reference, {"brief_summary": {"text": "Summary."}, "detailed_description": {"text": "Description."}}, structural_template=PRS_REFERENCE)
    assert report["status"] == "passed"
    study = next(ET.parse(output).getroot().iter("clinical_study"))
    for tag, expected in (("city", "Test City"), ("state", "New York"), ("zip", "10003"), ("country", "United States")):
        node = study.find(f"location/facility/address/{tag}")
        assert node.text == expected
        broken = tmp_path / f"without-{tag}.xml"
        tree = ET.parse(output)
        next(tree.getroot().iter("clinical_study")).find(f"location/facility/address/{tag}").text = ""
        tree.write(broken, encoding="utf-8", xml_declaration=True)
        assert any(
            item["field"] == f"location[1].facility.address.{tag}"
            for item in validate_output(broken, reference, PRS_REFERENCE, generation_template=PRS_TEMPLATE)
        )


def test_flat_site_address_populates_protocol_investigator_agreement(tmp_path):
    reference = source()
    report = rendering.render_documents(
        ROOT, tmp_path, reference, {"protocol": [], "icf": {}, "prs": {}},
        artifact_names={"protocol"},
    )
    assert report["status"] == "passed"
    document = Document(tmp_path / "candidate/protocol.docx")
    populated_city = [
        cell.text
        for table in document.tables
        for row in table.rows
        for cell in row.cells
        if cell.text.startswith("Location of Facility (City)")
    ]
    assert populated_city == ["Location of Facility (City)\nTest City"]


def test_ambiguous_supplied_address_is_resolved_before_approval():
    reference = source()
    reference["sites"][0]["facility"]["address"] = "300 Test Clinic Road, Test City, New York, USA"
    report = source_contract(reference)
    assert any(
        item["field"] == "sites[0].facility.address"
        for item in report["source_gaps"]
    )
    reference["sites"][0]["facility"].update({"city": "Test City", "state": "New York", "country": "USA"})
    assert not [
        item for item in source_contract(reference)["source_gaps"]
        if item["field"] == "sites[0].facility.address"
    ]


def test_sterling_purpose_accepts_source_grounded_plain_language(tmp_path):
    reference = source()
    purpose = (
        "The purpose of this study is to describe recovery outcomes. "
        "Researchers expect prospective monitoring will describe recovery. "
        "The primary measure is the Primary outcome at Month 3."
    )
    rendering.render_documents(
        ROOT, tmp_path, reference,
        {"protocol": [], "icf": {"icf.study-purpose": {"paragraphs": [{"text": purpose}], "lists": []}}, "prs": {}},
        artifact_names={"icf"},
    )
    from quality import validate_sterling_clause_contract

    report = validate_sterling_clause_contract(Document(tmp_path / "candidate/icf.docx"), reference)
    assert not [
        finding for finding in report["findings"]
        if finding["clause_id"] == "sterling.purpose.study-purpose"
    ]


def test_statistical_methodology_covers_groups_without_endpoint_inventory():
    reference = source()
    reference["statistics"]["analysis_plan"] = (
        "Continuous visual-acuity and defocus outcomes will be summarized with mean and standard deviation. "
        "Questionnaire responses will be summarized by category with counts and percentages."
    )
    reference["endpoints"]["primary"] = [
        {"label": "Binocular intermediate visual acuity at 66 cm", "time_point": "Month 3"}
    ]
    reference["endpoints"]["secondary"] = [
        {"label": "Binocular distance visual acuity", "time_point": "Month 3"},
        {"label": "Symptom questionnaire response", "time_point": "Month 3"},
    ]
    contract = next(
        section.public() for section in protocol_contract("Prospective")
        if section.section_id == "analysis-plan.methodology"
    )
    contract["minimum_evidence"] = list(contract["evidence"])
    assert contract["source_coverage"] == "method_group_summary"
    content = (
        "Continuous visual-acuity outcomes will be summarized with mean and standard deviation. "
        "Questionnaire responses will be summarized by category with counts and percentages."
    )
    refs = ["source:statistics.analysis_plan", "source:endpoints.primary", "source:endpoints.secondary"]
    request = {
        "approved_source": reference,
        "approved_input": [
            {"path": path, "value": get_path(reference, path)}
            for path in contract["minimum_evidence"]
        ],
    }
    assert _coverage_findings(request, contract, "analysis-plan.methodology", content, refs) == []
    assert _coverage_findings(request, contract, "analysis-plan.methodology", content, refs[:-1])


def test_section_15_table_opening_has_a_governed_layout_repair():
    finding = {
        "category": "visual", "artifact": "protocol", "check": "artificial_pagination",
        "element": "15. STANDARD EVALUATION PROCEDURES",
        "target_ids": ["layout:protocol"],
    }
    plan, unsupported = workflow._layout_repair_plan([finding], study_type="Prospective")
    assert unsupported == []
    assert plan == {"protocol": [{"rule": "section15_table_opening", "target": finding["element"]}]}

    document = Document()
    heading = document.add_paragraph(finding["element"], style="Heading 1")
    introduction = document.add_paragraph("The schedule below identifies visits and procedures.")
    caption = document.add_paragraph("Table 15.1. Proposed Visits and Study Assessments")
    document.add_table(rows=2, cols=2)
    rendering._repair_section15_table_opening(document, finding["element"])
    assert heading.paragraph_format.keep_with_next
    assert introduction.paragraph_format.keep_with_next
    assert caption.paragraph_format.keep_with_next


def test_long_section_15_draft_is_reduced_to_table_connective(tmp_path):
    reference = source()
    long_intro = (
        "The Schedule of Assessments below identifies the preoperative screening, "
        "operative, and Month 3 postoperative visits and their associated procedures. "
        "The source-derived table and its accompanying notes provide the assessment and monitoring detail."
    )
    rendering.render_documents(
        ROOT, tmp_path, reference,
        {"protocol": [{"section_id": "evaluation-procedures", "paragraphs": [{"text": long_intro}], "lists": []}], "icf": {}, "prs": {}},
        artifact_names={"protocol"},
        layout_repairs={"protocol": [
            {"rule": "section15_table_opening", "target": "15. STANDARD EVALUATION PROCEDURES"},
            {"rule": "table_pagination", "target": "Table 15.1. Proposed Visits and Study Assessments"},
        ]},
    )
    document = Document(tmp_path / "candidate/protocol.docx")
    heading = next(paragraph for paragraph in document.paragraphs if paragraph.text == "15. STANDARD EVALUATION PROCEDURES")
    assert heading.paragraph_format.keep_with_next
    caption = next(paragraph for paragraph in document.paragraphs if paragraph.text == "Table 15.1. Proposed Visits and Study Assessments")
    assert caption.paragraph_format.page_break_before is False
    assert "The approved visits and assessments are summarized in Table 15.1." in [
        paragraph.text for paragraph in document.paragraphs
    ]
    assert long_intro not in [paragraph.text for paragraph in document.paragraphs]


def test_section_15_repair_distinguishes_table_lead_in_from_caption():
    document = Document()
    heading = document.add_paragraph("15. STANDARD EVALUATION PROCEDURES", style="Heading 1")
    lead_in = document.add_paragraph("Table 15.1 presents the approved visit schedule.")
    caption = document.add_paragraph("Table 15.1. Proposed Visits and Study Assessments", style="Heading 2")
    caption.paragraph_format.page_break_before = True
    document.add_table(rows=2, cols=2)

    rendering._repair_section15_table_opening(document, heading.text)

    assert heading.paragraph_format.page_break_before is True
    assert lead_in.paragraph_format.keep_with_next is True
    assert caption.paragraph_format.keep_with_next is True
    assert caption.paragraph_format.page_break_before is False


def test_section_15_matrix_owns_fully_allocated_visit_notes():
    reference = {
        "meta": {"study_type": "Prospective"},
        "procedures": {
            "visit_schedule": [
                {"visit": "Preoperative screening", "timing": "Before surgery", "procedures": ["Informed consent", "Demographics", "Medical history"]},
                {"visit": "Month 3 postoperative", "timing": "3 months postoperative", "procedures": ["Visual acuity", "AIOLIS questionnaire"]},
            ],
            "assessments": [
                "Screening: informed consent, demographics, and medical history.",
                "Month 3: visual acuity and AIOLIS questionnaire.",
                "Month 3: visual acuity and genetic sequencing.",
            ],
        },
    }
    table = protocol_table_contracts(reference)["schedule-of-assessments"]
    assert table["supplemental_notes"] == ["Month 3: visual acuity and genetic sequencing."]


def test_no_progress_reuses_authenticated_attempt_provenance_without_stale_request(tmp_path):
    revision = tmp_path / "revisions/r-test"
    candidate = revision / "candidate"
    accepted = revision / "hermes/accepted"
    candidate.mkdir(parents=True)
    accepted.mkdir(parents=True)
    (candidate / "protocol.docx").write_bytes(b"same candidate bytes")
    (accepted / "evaluation-procedures.json").write_text(json.dumps({
        "section_id": "evaluation-procedures", "paragraphs": [{"text": "Old introduction."}], "lists": [],
    }))
    finding = {
        "category": "content", "field": "clinical_content_verification", "artifact": "protocol",
        "verification_request_id": "r-test.review-1.verify.content",
        "target_ids": ["evaluation-procedures"], "check": "editorial_relevance",
        "recovery_class": "drafting_defect", "action": "retry_drafting_target",
        "issue": "The section repeats the table.",
    }
    attempt = workflow._archive_failed_attempt(revision, "quality", [finding])
    journal = json.loads((revision / "gate-attempt-journal.json").read_text())
    working = {"generation": {
        "gate_attempts": journal["entries"],
        "pending_recovery_attempts": [attempt.relative_to(revision).as_posix()],
    }}
    reference_path = tmp_path / "reference/study.reference.json"
    reference_path.parent.mkdir()
    reference_path.write_text(json.dumps(working))
    (candidate / "protocol.docx").write_bytes(b"same candidate bytes")
    (accepted / "evaluation-procedures.json").write_text(json.dumps({
        "section_id": "evaluation-procedures", "paragraphs": [{"text": "New introduction."}], "lists": [],
    }))

    no_progress = workflow._complete_pending_recovery_attempts(
        revision, reference_path, working, require_candidate_change=True,
    )
    internal = workflow._internalize_no_progress_findings(revision, working, no_progress)

    assert len(internal) == 1
    assert internal[0]["field"] == "recovery_no_progress"
    assert "verification_request_id" not in internal[0]
    assert internal[0]["prior_review_binding"]["verification_request_id"] == finding["verification_request_id"]
    no_progress[0]["triggering_finding_sha256"] = "0" * 64
    with pytest.raises(ValueError, match="archived reviewer finding"):
        workflow._internalize_no_progress_findings(revision, working, no_progress)


@pytest.mark.parametrize("outcome", ["construction_failed", "render_assurance_failed"])
def test_failed_build_records_cause_without_no_progress(tmp_path, outcome):
    revision = tmp_path / "revisions/r-test"
    (revision / "candidate").mkdir(parents=True)
    (revision / "candidate/protocol.docx").write_bytes(b"old candidate")
    finding = {
        "category": "visual", "field": "protocol", "artifact": "protocol",
        "check": "artificial_pagination", "target_ids": ["layout:protocol"],
        "recovery_class": "visual_defect", "action": "repair_layout",
        "issue": "Section 15 opening is split.",
    }
    attempt = workflow._archive_failed_attempt(revision, "quality", [finding])
    journal = json.loads((revision / "gate-attempt-journal.json").read_text())
    working = {"generation": {
        "gate_attempts": journal["entries"],
        "pending_recovery_attempts": [attempt.relative_to(revision).as_posix()],
    }}
    reference_path = tmp_path / "reference/study.reference.json"
    reference_path.parent.mkdir()
    reference_path.write_text(json.dumps(working))

    assert workflow._complete_pending_recovery_attempts(
        revision, reference_path, working,
        require_candidate_change=False, outcome_status=outcome,
    ) == []
    workflow._validate_expected_gate_attempts(revision, working["generation"]["gate_attempts"])
    manifest = json.loads((attempt / "attempt-manifest.json").read_text())
    assert manifest["recovery_actions"][0]["outcome_status"] == outcome


def test_protocol_benefits_does_not_claim_payment_as_a_benefit(tmp_path):
    reference = source()
    reference["risks_benefits"]["benefits"] = "Participants may receive no direct benefit."
    reference["risks_benefits"]["compensation_or_reimbursement"] = "No compensation or reimbursement will be provided."
    report = rendering.render_documents(
        ROOT, tmp_path, reference,
        {"protocol": [{
            "section_id": "risks-benefits.benefits",
            "paragraphs": [{"text": "Participants may receive no direct benefit. No compensation or reimbursement will be provided."}],
            "lists": [],
        }], "icf": {}, "prs": {}},
    )
    assert next(item for item in report["artifacts"] if item["artifact"] == "protocol")["status"] == "passed"
    findings = quality.deterministic_content_check(tmp_path, reference)
    assert any(
        item.get("field") == "risks-benefits.benefits" and "payment" in item.get("issue", "").casefold()
        for item in findings
    )


def test_study_completion_can_state_rule_without_visit_itinerary():
    reference = source()
    reference["study"]["timeline"] = "Study completion follows three months of participant follow-up and one month of data analysis."
    contract = next(
        section.public() for section in protocol_contract("Prospective")
        if section.section_id == "endpoint-criteria.study-completion"
    )
    contract["minimum_evidence"] = list(contract["evidence"])
    request = {
        "approved_source": reference,
        "approved_input": [
            {"path": path, "value": get_path(reference, path)}
            for path in contract["minimum_evidence"]
        ],
    }
    refs = [
        f"source:{path}" for path in contract["minimum_evidence"]
        if get_path(reference, path)
    ]
    content = reference["study"]["timeline"]
    assert _coverage_findings(request, contract, "endpoint-criteria.study-completion", content, refs) == []
