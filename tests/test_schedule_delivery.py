"""Synthetic schedule replay through the public Desktop delivery operation.

Controlled external responders are not independent clinical certification.
Internal construction, gates, cache invalidation and atomic publication are real.
"""
import copy
import json
import subprocess
from pathlib import Path

from docx import Document
import pytest

import contracts
import drafting
import quality
import workflow
from test_handoff_quality import acceptance_verification

ROOT = Path(__file__).resolve().parents[1]


def _restored_rows(contract, rows):
    restored = copy.deepcopy(rows)
    for plan in contracts.protocol_table_label_restorations(contract, rows):
        restored[plan['row']][0] = plan['source']
    return restored


@pytest.mark.parametrize("change", ["presentation", "facts"])
def test_schedule_variation_or_real_defect_reaches_verified_complete_delivery(tmp_path, governed_pdfium, monkeypatch, change):
    run = tmp_path / "run"
    reference_path = run / "reference/study.reference.json"
    reference_path.parent.mkdir(parents=True)
    source = json.loads((ROOT / "tests/fixtures/prospective-acceptance-source.json").read_text())
    source["procedures"]["visit_schedule_table"] = []
    source["procedures"]["visit_schedule"][0]["procedures"].extend([
        "demographics", "informed consent", "blood pressure in two positions", "As needed assessment", "Eye examination",
        "qPCR collection at 5 mL", "dose at 5 mcg", "heart rate at 60 bpm", "Hb assessment", "dose at 5 mg kg^-1 min^-1",
    ])
    reference_path.write_text(json.dumps(source))
    assert workflow.prepare(run)["status"] == "awaiting_approval"
    assert workflow.approve(run, approved_by="Synthetic schedule replay")["status"] == "passed"
    approval = copy.deepcopy(json.loads(reference_path.read_text())["approval"])
    revision = run / "revisions" / approval["revision_id"]
    observed = {"mutated": False, "drafts": None, "reviews": set()}

    def accepted_bytes():
        return {path.name: path.read_bytes() for path in (revision / "hermes/accepted").glob("*.json")}

    real_process = subprocess.run

    def office_process(command, *args, **kwargs):
        # Inject once at the external office boundary, before artifact review
        # hashes exist. Never mutate already-bound verifier inputs.
        if (isinstance(command, (list, tuple)) and "--convert-to" in command
                and "pdf" in command and Path(str(command[-1])) == revision / "candidate/protocol.docx"
                and not observed["mutated"]):
            protocol_path = revision / "candidate/protocol.docx"
            document = Document(protocol_path)
            table = next(t for t in document.tables if t.rows[0].cells[0].text.strip() == "Activity")
            replacements = {
                "presentation": [("demographics", "Demographics"), ("informed consent", "Informed Consent"),
                                 ("blood pressure in two positions", "Blood Pressure in Two Positions"),
                                 ("As needed assessment", "as needed assessment"), ("Eye examination", "eye examination")],
                "facts": [("qPCR collection at 5 mL", "qPCR collection at 99 mL"),
                          ("dose at 5 mcg", "dose at 5 Mcg"), ("heart rate at 60 bpm", "heart rate at 60 Bpm"),
                          ("Hb assessment", "hb assessment"), ("dose at 5 mg kg^-1 min^-1", "dose at 5 mg Kg^-1 min^-1")],
            }[change]
            for label, replacement in replacements:
                row = next(row for row in table.rows if row.cells[0].text.strip() == label)
                row.cells[0].text = replacement
            document.save(protocol_path)
            defects = [f for f in quality.deterministic_content_check(revision, source)
                       if f.get("field") == "procedures.visit_schedule"]
            # Raw altered units/identifiers are never declared equivalent. The
            # render boundary must restore exact source spelling before review.
            assert defects and all(f["owner"] == "construction"
                                  and f["action"] == "rebuild_deterministic_structure" for f in defects)
            assert not (run / "output").exists()
            observed["drafts"] = accepted_bytes()
            observed["mutated"] = True
        return real_process(command, *args, **kwargs)

    monkeypatch.setattr(subprocess, "run", office_process)

    def runner(handoffs, _remaining):
        for handoff in handoffs:
            request_path = revision / handoff["request_path"]
            request = json.loads(request_path.read_text())
            if "verification-requests" in handoff["request_path"]:
                observed["reviews"].add(request["review_set"])
                workflow._save_verification_response(revision, request_path, acceptance_verification)
            else:
                if observed["mutated"]:
                    pytest.fail("A generated table must not redispatch accepted prose.")
                response_path = revision / request["response_path"]
                response_path.parent.mkdir(parents=True, exist_ok=True)
                response_path.write_text(json.dumps(drafting.recorded_acceptance_response(request)))

    result = workflow.run_desktop_operation(run, handoff_runner=runner,
        opener=lambda path: Path(path).read_bytes(), operation_id="schedule-replay",
        release_identity={"package_fingerprint": "synthetic-schedule-test"})
    assert result["status"] == "passed", result
    assert result["delivery"]["confirmed"]
    assert sorted(result["client_outputs"]) == ["output/icf.docx", "output/protocol.docx", "output/study.xml"]
    assert result["desktop_reply"]["message"].startswith("Complete verified Branch Document Set delivered.")
    assert observed["mutated"] and observed["reviews"]
    assert accepted_bytes() == observed["drafts"]
    assert json.loads(reference_path.read_text())["approval"] == approval
    state = json.loads((run / "logs/desktop-operation-schedule-replay.json").read_text())
    assert state["deadline_at_epoch"] - state["started_at_epoch"] == 2700
    reconstructions = json.loads(reference_path.read_text()).get("generation", {}).get("deterministic_reconstructions", [])
    assurance = json.loads(reference_path.read_text())["generation"]["render_assurance"]
    if change == "facts":
        assert any(item["target"] == "evaluation-procedures" and item["outcome"] == "rebuilt"
                   and item["candidate_bytes_changed"] for item in reconstructions)
    else:
        assert reconstructions == []
        assert len(assurance["source_case_restorations"]) >= 5
    protocol = Document(run / "output/protocol.docx")
    table = next(t for t in protocol.tables if t.rows[0].cells[0].text.strip() == "Activity")
    assert contracts.protocol_table_rows_match(contracts.protocol_table_contracts(source)["schedule-of-assessments"],
                                             [[cell.text for cell in row.cells] for row in table.rows])
    assert "qPCR collection at 99 mL" not in [row.cells[0].text for row in table.rows]
    manifest = json.loads((revision / "delivery-manifest.json").read_text())
    for item in manifest["client_outputs"]:
        delivered = run / item["path"]
        assert delivered.read_bytes() == (revision / "candidate" / delivered.name).read_bytes()
        assert quality.sha256_file(delivered) == item["sha256"]


def test_narrative_only_assessments_reach_complete_public_delivery(tmp_path, governed_pdfium):
    run = tmp_path / 'run'
    reference_path = run / 'reference/study.reference.json'
    reference_path.parent.mkdir(parents=True)
    source = json.loads((ROOT / 'tests/fixtures/prospective-acceptance-source.json').read_text())
    source['procedures'].pop('visit_schedule', None)
    source['procedures'].pop('visit_schedule_table', None)
    reference_path.write_text(json.dumps(source))
    assert workflow.prepare(run)['status'] == 'awaiting_approval'
    assert workflow.approve(run, approved_by='Synthetic narrative assessments')['status'] == 'passed'
    approval = copy.deepcopy(json.loads(reference_path.read_text())['approval'])
    revision = run / 'revisions' / approval['revision_id']
    observed = {'accepted': None}

    def runner(handoffs, _remaining):
        for handoff in handoffs:
            request_path = revision / handoff['request_path']
            request = json.loads(request_path.read_text())
            if 'verification-requests' in handoff['request_path']:
                if observed['accepted'] is None:
                    observed['accepted'] = {p.name: p.read_bytes() for p in (revision / 'hermes/accepted').glob('*.json')}
                workflow._save_verification_response(revision, request_path, acceptance_verification)
            else:
                if observed['accepted'] is not None:
                    pytest.fail('Narrative assessment fallback must not redispatch accepted prose.')
                response_path = revision / request['response_path']
                response_path.parent.mkdir(parents=True, exist_ok=True)
                response_path.write_text(json.dumps(drafting.recorded_acceptance_response(request)))

    result = workflow.run_desktop_operation(run, handoff_runner=runner,
        opener=lambda path: Path(path).read_bytes(), operation_id='narrative-assessment-replay',
        release_identity={'package_fingerprint': 'synthetic-narrative-assessment-test'})
    assert result['status'] == 'passed', result
    assert result['delivery']['confirmed']
    assert sorted(result['client_outputs']) == ['output/icf.docx', 'output/protocol.docx', 'output/study.xml']
    assert result['desktop_reply']['message'].startswith('Complete verified Branch Document Set delivered.')
    assert json.loads(reference_path.read_text())['approval'] == approval
    assert observed['accepted'] == {p.name: p.read_bytes() for p in (revision / 'hermes/accepted').glob('*.json')}
    assert not json.loads(reference_path.read_text()).get('generation', {}).get('deterministic_reconstructions', [])
    document = Document(run / 'output/protocol.docx')
    assert any(t.rows[0].cells[0].text.strip() == 'Approved visit or assessment' for t in document.tables)
    for name in ('icf.docx', 'protocol.docx', 'study.xml'):
        assert (run / 'output' / name).read_bytes() == (revision / 'candidate' / name).read_bytes()


@pytest.mark.parametrize("label", ["demographics", "medical history", "distance-corrected defocus curve",
                                  "adverse event and device deficiency monitoring", "exit form", "blood pressure"])
def test_sentence_case_descriptor_is_not_a_clinical_change(label):
    contract = {"section_id": "evaluation-procedures", "header_rows": 1,
                "rows": [["Activity", "Visit 1"], [label, "X"]]}
    assert contracts.protocol_table_rows_match(contract, _restored_rows(contract, [["Activity", "Visit 1"], [label[0].upper() + label[1:], "X"]]))


def test_title_case_descriptive_words_preserve_technical_tokens():
    contract = {"section_id": "evaluation-procedures", "header_rows": 1,
                "rows": [["Activity", "Visit 1"], ["informed consent and qPCR collection at 5 mL", "X"]]}
    assert contracts.protocol_table_rows_match(contract, _restored_rows(contract,
        [["Activity", "Visit 1"], ["Informed Consent and qPCR Collection at 5 mL", "X"]]))


def test_function_words_are_not_mistaken_for_measurement_units():
    contract = {"section_id": "evaluation-procedures", "header_rows": 1,
                "rows": [["Activity", "Visit 1"], ["as needed assessment", "X"]]}
    assert contracts.protocol_table_rows_match(contract, _restored_rows(contract, [["Activity", "Visit 1"], ["As needed assessment", "X"]]))


@pytest.mark.parametrize("label,replacement", [("PCR collection", "Pcr collection"),
    ("sample at 5 mol", "Sample at 5 Mol"), ("sample at 5 mg", "Sample at 5 Mg"),
    ("brca1 assessment", "Brca1 Assessment"), ("eGFR assessment", "Egfr Assessment")])
def test_title_case_does_not_fold_identifiers_or_units(label, replacement):
    contract = {"section_id": "evaluation-procedures", "header_rows": 1,
                "rows": [["Activity", "Visit 1"], [label, "X"]]}
    assert not contracts.protocol_table_rows_match(contract, [["Activity", "Visit 1"], [replacement, "X"]])


@pytest.mark.parametrize("label,replacement", [("dose at 5 mcg", "Dose at 5 Mcg"),
    ("heart rate at 60 bpm", "Heart rate at 60 Bpm"), ("Hb assessment", "hb assessment"),
    ("dose at 5 mcg/kg/min", "Dose at 5 mcg/Kg/min"),
    ("dose at 5 mcg / kg / min", "Dose at 5 mcg / Kg / min")])
def test_quantity_units_and_short_technical_identifiers_are_exact(label, replacement):
    contract = {"section_id": "evaluation-procedures", "header_rows": 1,
                "rows": [["Activity", "Visit 1"], [label, "X"]]}
    assert not contracts.protocol_table_rows_match(contract, [["Activity", "Visit 1"], [replacement, "X"]])


@pytest.mark.parametrize("rows", [
    [["Activity", "Visit 1", "Visit 2"], ["qPCR collection at 99 mL", "X", ""], ["medical history", "", "X"]],
    [["Activity", "Visit 1", "Visit 2"], ["qPCR collection at 5 ML", "X", ""], ["medical history", "", "X"]],
    [["Activity", "Visit 1", "Visit 2"], ["QPCR collection at 5 mL", "X", ""], ["medical history", "", "X"]],
    [["Activity", "Visit 1", "Visit 2"], ["qPCR collection at 5 mL", "", "X"], ["medical history", "", "X"]],
    [["Activity", "Visit 2", "Visit 1"], ["qPCR collection at 5 mL", "X", ""], ["medical history", "", "X"]],
    [["Activity", "Visit 1", "Visit 2"], ["medical history", "", "X"], ["qPCR collection at 5 mL", "X", ""]],
    [["Activity", "Visit 1", "Visit 2"], ["qPCR collection at 5 mL", "X", ""]],
    [["Activity", "Visit 1", "Visit 2"], ["qPCR collection at 5 mL", "X"], ["medical history", "", "X"]],
    [["Activity", "Visit 1", "Visit 2"], ["qPCR collection at 5 mL", "X", ""], ["medical history", "", "X"], ["medical history", "", "X"]],
    [["Activity", "Visit 1", "Visit 2"], ["qPCR collection at 5 mL", "x", ""], ["medical history", "", "X"]],
])
def test_clinical_cells_and_source_relationships_remain_exact(rows):
    contract = {"section_id": "evaluation-procedures", "header_rows": 1,
        "rows": [["Activity", "Visit 1", "Visit 2"], ["qPCR collection at 5 mL", "X", ""], ["medical history", "", "X"]]}
    assert not contracts.protocol_table_rows_match(contract, rows)


def test_repeated_office_case_mutation_cannot_certify_stale_pages(tmp_path, governed_pdfium):
    import rendering
    from test_render_assurance import _bundle, _structural_validation, _write_text_pdf

    source = json.loads((ROOT / 'tests/fixtures/prospective-acceptance-source.json').read_text())
    source['procedures']['visit_schedule_table'] = []
    source['procedures']['visit_schedule'] = [{'visit': 'Screening', 'timing': 'Before enrollment',
                                             'procedures': ['medical history']}]
    rendering.render_documents(ROOT, tmp_path, source, {'protocol': [], 'icf': {}, 'prs': {}}, artifact_names={'protocol'})
    exports = []

    def export(docx, output_dir, _identity, **_kwargs):
        document = Document(docx)
        table = next(t for t in document.tables if t.rows[0].cells[0].text.strip() == 'Activity')
        from docx.oxml.ns import qn
        next(table.rows[2].cells[0]._tc.iter(qn('w:t'))).text = 'Medical history'
        document.save(docx)
        exports.append(docx.name)
        pdf = output_dir / 'protocol.pdf'
        _write_text_pdf(pdf)
        return pdf

    def pages(*_args, **_kwargs):
        pytest.fail('Persistently mutated DOCX/PDF pairs must not reach page verification.')

    report = quality.render_assurance(ROOT, tmp_path, source, contracted_bundle=_bundle(),
        structural_validation=_structural_validation(tmp_path, 'protocol.docx'),
        renderer_identities=[{'kind': 'LibreOffice', 'path': '/controlled/soffice'}],
        page_renderer_identities=[governed_pdfium], font_probe=lambda *_args, **_kwargs: (True, 'available'),
        office_exporter=export, page_exporter=pages, blank_page_detector=lambda _pdf: [])
    assert report['status'] == 'blocked'
    assert len(exports) == 4
    assert report['render'].get('artifacts', []) == []
    assert all(f['recovery_class'] == 'adapter_fault' for f in report['findings'])
