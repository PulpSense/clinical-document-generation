"""Source-grounded uncertainty decisions at the public verification boundary."""
import json
import pytest

import quality
from rendering import render_documents
from test_handoff_quality import acceptance_verification, ROOT


def review_candidate(tmp_path, branch='Prospective', template='Advarra'):
    source = json.loads((ROOT / 'tests/fixtures/prospective-acceptance-source.json').read_text())
    source['meta'].update(study_type=branch, icf_template=template)
    source['procedures']['retention'] = 'Personnel will telephone participants before scheduled appointments.'
    render_documents(ROOT, tmp_path, source, {'protocol': [{
        'section_id': 'study-design.bias',
        'paragraphs': [{'text': 'Staff call people ahead of visits so they remember to attend.'}],
        'lists': [],
    }], 'icf': {}, 'prs': {}}, artifact_names={'protocol'})
    accepted = tmp_path / 'hermes/accepted'
    accepted.mkdir(parents=True)
    (accepted / 'study-design.bias.json').write_text(json.dumps({'warnings': [{
        'code': 'lexical-evidence-uncertainty', 'field': 'study-design.bias',
        'publication_disposition': 'warning', 'issue': 'Reminder-call wording needs review.',
        'evidence_diagnostics': [{'source_path': 'procedures.retention',
                                  'source_excerpt': source['procedures']['retention']}],
    }]}))
    path = quality.create_verification_requests(tmp_path, source, {'artifacts': []})[0]
    request = json.loads(path.read_text())
    return path, request, acceptance_verification(request)


def test_bare_passing_review_cannot_dismiss_uncertainty(tmp_path):
    path, request, response = review_candidate(tmp_path)
    response.pop('uncertainty_resolutions', None)
    (tmp_path / request['response_path']).write_text(json.dumps(response))
    assert not quality.verification_response_is_complete(tmp_path, path)
    findings, _ = quality.validate_verifications(tmp_path, request_paths=[path])
    assert any(item.get('code') == 'uncertainty_resolution_invalid' for item in findings)


def faithful_decision(request):
    caution = request['lexical_review_cautions'][0]
    return {
        'caution_sha256': caution['caution_sha256'], 'disposition': 'faithful',
        'source_excerpt': 'Personnel will telephone participants before scheduled appointments.',
        'affected_passage': 'Staff call people ahead of visits so they remember to attend.',
        'rationale': 'Calling before visits conveys the supplied telephone reminders without adding an obligation.',
    }


def test_faithful_resolution_binds_source_section_rule_and_reviewer(tmp_path):
    path, request, response = review_candidate(tmp_path)
    response['uncertainty_resolutions'] = [faithful_decision(request)]
    (tmp_path / request['response_path']).write_text(json.dumps(response))
    findings, evidence = quality.validate_verifications(tmp_path, request_paths=[path])
    assert findings == []
    decision = evidence['clinical_content_verification']['uncertainty_resolutions'][0]
    assert decision['producer']['reviewer_id'] == 'clinical_content_verification'
    assert decision['approved_source_sha256'] == request['lexical_review_cautions'][0]['approved_source_sha256']
    assert decision['section_sha256'] and decision['artifact_sha256'] and decision['rule_sha256']
    # Repeating the same review preserves the resolution exactly.
    assert quality.validate_verifications(tmp_path, request_paths=[path])[1] == evidence


def test_resolution_cannot_quote_invented_passage_or_source(tmp_path):
    path, request, response = review_candidate(tmp_path)
    decision = faithful_decision(request)
    for field in ('source_excerpt', 'affected_passage', 'rationale'):
        response['uncertainty_resolutions'] = [{**decision, field: '' if field == 'rationale' else 'invented evidence'}]
        (tmp_path / request['response_path']).write_text(json.dumps(response))
        assert not quality.verification_response_is_complete(tmp_path, path)


@pytest.mark.parametrize('branch,template', [
    ('Prospective', 'Advarra'), ('Prospective', 'Sterling'),
    ('Ambispective', 'Advarra'), ('Ambispective', 'Sterling'), ('Retrospective', ''),
])
def test_validated_decision_localizes_missing_finding_passage_on_every_family(tmp_path, branch, template):
    path, request, response = review_candidate(tmp_path, branch, template)
    response['status'] = 'blocked'
    response['findings'] = [{'finding_id': 'repair-bias', 'artifact': 'protocol', 'category': 'content',
                             'check': 'substantive', 'target_ids': ['study-design.bias'],
                             'issue': 'Correct the supplied method.'}]
    for row in response['section_assessments']:
        if row['section_id'] == 'study-design.bias': row['status'] = 'blocked'
    caution = next(row for row in request['lexical_review_cautions'] if row['section_id'] == 'study-design.bias')
    for decision in response['uncertainty_resolutions']:
        if decision['caution_sha256'] == caution['caution_sha256']:
            decision.update(faithful_decision(request), caution_sha256=caution['caution_sha256'],
                            disposition='confirmed_defect')
    (tmp_path / request['response_path']).write_text(json.dumps(response))
    findings, _ = quality.validate_verifications(tmp_path, request_paths=[path])
    finding = next(row for row in findings if row.get('check') == 'substantive')
    assert finding['affected_passage'] == faithful_decision(request)['affected_passage']
    assert finding['recovery_localization']['request_sha256'] == request['request_sha256']
    # Preserve the original authenticated reviewer response.
    assert 'affected_passage' not in json.loads((tmp_path / request['response_path']).read_text())['findings'][0]


@pytest.mark.parametrize('target', ['protocol.shell', 'icf.shell', 'prs.structured'])
@pytest.mark.parametrize('mode', ['missing', 'ambiguous', 'faithful', 'other-artifact', 'other-section'])
def test_missing_native_location_requests_review_instead_of_guessing(target, mode):
    finding = {'verification_request_id': 'review-1', 'artifact': 'protocol', 'target_ids': [target],
               'recovery_class': 'deterministic_structure_defect', 'issue': 'Correct this field.',
               'verification_source_finding': {}}
    record = {'disposition': 'confirmed_defect', 'artifact': 'protocol', 'section_id': target,
              'affected_passage': 'Wrong value', 'caution_sha256': 'caution', 'request_sha256': 'bound-request'}
    records = [] if mode == 'missing' else [record]
    if mode == 'ambiguous': records.append(dict(record, affected_passage='Another value'))
    if mode == 'faithful': record['disposition'] = 'faithful'
    if mode == 'other-artifact': record['artifact'] = 'icf'
    if mode == 'other-section': record['section_id'] = 'subjects'
    quality._localize_recovery_findings([finding], records, {'request_id': 'review-1'})
    assert finding['recovery_class'] == 'verifier_transient'
    assert finding['target_ids'] == ['verification:content']
    assert finding['code'] == 'recovery_location_required'
    assert not finding['verification_source_finding'].get('affected_passage')


@pytest.mark.parametrize('target', ['protocol.shell', 'icf.shell', 'prs.structured'])
def test_repeated_cautions_share_one_location_without_changing_supplied_source_scope(target):
    finding = {'verification_request_id': 'review-1', 'artifact': 'protocol', 'target_ids': [target],
               'recovery_class': 'deterministic_structure_defect', 'issue': 'Correct this field.',
               'verification_source_finding': {'source_path': 'parties.principal_investigator.name'}}
    record = {'disposition': 'confirmed_defect', 'artifact': 'protocol', 'section_id': target,
              'affected_passage': 'Wrong value', 'caution_sha256': 'one', 'request_sha256': 'bound-request'}
    quality._localize_recovery_findings([finding], [record, dict(record, caution_sha256='two')],
                                       {'request_id': 'review-1'})
    assert finding['recovery_class'] == 'deterministic_structure_defect'
    assert finding['affected_passage'] == 'Wrong value'
    assert finding['verification_source_finding']['source_path'] == 'parties.principal_investigator.name'
    assert finding['recovery_localization']['caution_sha256s'] == ['one', 'two']


def test_one_decision_cannot_localize_two_findings_in_the_same_section():
    first = {'verification_request_id': 'review-1', 'artifact': 'protocol', 'target_ids': ['protocol.shell'],
             'recovery_class': 'deterministic_structure_defect', 'issue': 'Correct this field.',
             'verification_source_finding': {}}
    second = dict(first)
    record = {'disposition': 'confirmed_defect', 'artifact': 'protocol', 'section_id': 'protocol.shell',
              'affected_passage': 'Wrong value', 'caution_sha256': 'one', 'request_sha256': 'bound-request'}
    quality._localize_recovery_findings([first, second], [record], {'request_id': 'review-1'})
    assert first['recovery_class'] == second['recovery_class'] == 'verifier_transient'


def test_localization_preserves_explicit_finding_passage():
    finding = {'verification_request_id': 'review-1', 'artifact': 'icf', 'target_ids': ['icf.shell'],
               'recovery_class': 'deterministic_structure_defect', 'issue': 'Correct this field.',
               'verification_source_finding': {'affected_passage': 'Explicit quotation'}}
    quality._localize_recovery_findings([finding], [], {'request_id': 'review-1'})
    assert finding['verification_source_finding']['affected_passage'] == 'Explicit quotation'
    assert finding['recovery_class'] == 'deterministic_structure_defect'


def test_material_editorial_defect_requires_repair_instead_of_delivery_caution(tmp_path):
    from test_editorial_publication import repetition_review
    path, response_path, response = repetition_review(tmp_path, 'Prospective', 'Advarra')
    response['findings'][0]['material'] = True
    response_path.write_text(json.dumps(response))
    findings, _ = quality.validate_verifications(tmp_path, request_paths=[path])
    defect = next(item for item in findings if item.get('check') == 'concept_repetition')
    assert defect.get('publication_disposition') != 'warning'
    assert defect['recovery_class'] == 'drafting_defect'
    assert defect['target_ids'] == ['analysis-plan.considerations']


def test_unpresented_uncertainty_cannot_leak_as_client_warning(tmp_path, monkeypatch):
    # Supporting seam: isolate warning policy from already-covered package gates.
    monkeypatch.setattr(quality, 'deterministic_content_check', lambda *_: [{
        'code': 'lexical-evidence-uncertainty', 'field': 'study-design.bias',
        'publication_disposition': 'warning', 'issue': 'A new unreviewed uncertainty.',
    }])
    monkeypatch.setattr(quality, 'validate_verifications', lambda *_: ([], {}))
    monkeypatch.setattr(quality, '_final_verification_scope_findings', lambda *_: [])
    report = quality.quality_report(tmp_path, {'meta': {'study_type': 'Retrospective'}}, {}, None)
    assert report['status'] == 'blocked'
    assert any(item.get('code') == 'uncertainty_resolution_required' for item in report['findings'])


def test_source_complete_repetition_is_accepted_for_independent_review(tmp_path):
    import drafting
    from contracts import batch_plan
    from test_handoff_quality import fixture
    source = fixture()
    batch = next(item for item in batch_plan('Prospective') if item.batch_id == 'protocol-foundations')
    path = drafting.create_drafting_request(repo_root=ROOT, revision_dir=tmp_path,
        revision_id='r-repeated', reference=source, batch=batch,
        attempts={item: 1 for item in batch.section_ids}, wave='initial')
    request = json.loads(path.read_text())
    response = drafting.recorded_acceptance_response(request)
    # Both sections retain their complete original facts; only a harmless
    # cross-reference sentence repeats.
    for row in response['section_results'][:2]:
        row['paragraphs'].append({'text': 'Further approved study details are described in the following sections.',
                                 'evidence_refs': [], 'boilerplate_refs': []})
        # Use each section's existing citation without creating an intake fact.
        row['paragraphs'][-1]['evidence_refs'] = list(row['paragraphs'][0]['evidence_refs'])
    accepted, findings = drafting.validate_response(request, response)
    assert {row['section_id'] for row in accepted['drafts']} == set(batch.section_ids)
    assert all(row.get('publication_disposition') == 'warning' for row in findings)


def test_meaningful_remaining_limitation_is_disclosed_but_material_defect_is_not_waived(tmp_path):
    from test_editorial_publication import repetition_review
    path, response_path, response = repetition_review(tmp_path, 'Prospective', 'Advarra')
    response['findings'] = [{
        'finding_id': 'registry-upload-unavailable', 'category': 'content', 'artifact': 'protocol',
        'check': 'client_limitation', 'target_ids': ['analysis-plan.considerations'],
        'issue': 'The external registry upload service is unavailable; the complete reviewed files are available.',
        'remaining_limitation': True, 'limitation_kind': 'external_dependency',
        'why_not_repairable': 'Uploading to the external registry is unavailable during this service outage.',
        'recommended_action': 'Upload the reviewed XML once the registry is available.',
        'material': False, 'safety_critical': False, 'contradiction': False,
        'obscures_required_information': False, 'materially_unusable': False,
    }]
    response_path.write_text(json.dumps(response))
    findings, _ = quality.validate_verifications(tmp_path, request_paths=[path])
    assert len(findings) == 1
    assert findings[0].get('publication_disposition') == 'warning'
    assert quality.verification_response_is_complete(tmp_path, path)
    response['findings'][0]['material'] = True
    response_path.write_text(json.dumps(response))
    assert not quality.verification_response_is_complete(tmp_path, path)


def test_deterministic_assessment_cell_defect_routes_to_construction(tmp_path):
    from docx import Document
    from test_mixed_visit_representations import failed_source
    source = failed_source()
    render_documents(ROOT, tmp_path, source, {'protocol': [], 'icf': {}, 'prs': {}})
    path = tmp_path / 'candidate/protocol.docx'
    doc = Document(path)
    table = next(table for table in doc.tables if table.rows[0].cells[0].text.strip() == 'Activity')
    table.rows[1].cells[-1].text = 'Changed assignment'
    doc.save(path)
    findings = quality.deterministic_content_check(tmp_path, source)
    finding = next(row for row in findings if 'does not preserve the contracted headers' in row['issue'])
    assert finding['owner'] == 'construction'
    assert finding['recovery_class'] == 'deterministic_structure_defect'


def test_case_sensitive_quantity_change_is_an_objective_defect():
    import drafting
    valid = drafting.source_evidence_diagnostics('Blood draw at five mL.', 'procedures.assessments', 'Blood draw at 5 mL.')
    invalid = drafting.source_evidence_diagnostics('Blood draw at 5 ML.', 'procedures.assessments', 'Blood draw at 5 mL.')
    assert valid['passed'] is True
    assert invalid['passed'] is False
    assert not drafting.lexical_evidence_uncertainty([invalid])


def test_sterling_meaning_preserving_term_proxy_requests_review(tmp_path):
    from test_sterling_rendered_scope import SOURCE, FIXTURE
    from docx import Document
    doc = Document(FIXTURE / 'candidate-3.docx')
    active = False
    for paragraph in doc.paragraphs:
        if paragraph.text == 'PURPOSE':
            active = True
        elif paragraph.text == 'DURATION':
            active = False
        elif active:
            paragraph.text = paragraph.text.replace('primary endpoint', 'central result').replace('main outcome', 'central result')
    findings = quality.validate_sterling_clause_contract(doc, SOURCE)['findings']
    purpose = [row for row in findings if row['clause_id'] == 'sterling.purpose.study-purpose']
    assert purpose
    assert all(row.get('publication_disposition') == 'warning' for row in purpose)
    # Different clinical wording does not prove quantity-role identity; it must
    # be assessed against the approved facts, not cleared by token matching.
    for paragraph in doc.paragraphs:
        paragraph.text = paragraph.text.replace('66 cm', '65 cm')
    findings = quality.validate_sterling_clause_contract(doc, SOURCE)['findings']
    assert any(row.get('publication_disposition') == 'warning' and row.get('evidence_diagnostics') for row in findings)


def test_explicit_quantity_sign_change_is_not_a_paraphrase():
    import drafting
    source = 'The offset is -5 mm.'
    assert drafting.source_evidence_diagnostics('The offset is -5 mm.', 'procedures.assessments', source)['passed']
    invalid = drafting.source_evidence_diagnostics('The offset is 5 mm.', 'procedures.assessments', source)
    assert invalid['passed'] is False
    assert not drafting.lexical_evidence_uncertainty([invalid])


def test_quoted_duplicate_passage_is_located_for_review(tmp_path):
    from docx import Document
    path, request, _ = review_candidate(tmp_path)
    docx = tmp_path / 'candidate/protocol.docx'
    doc = Document(docx)
    passage = 'The "study study" descriptor is used here.'
    next(p for p in doc.paragraphs if p.text == 'Staff call people ahead of visits so they remember to attend.').text = passage
    doc.save(docx)
    cautions = quality._lexical_review_cautions(tmp_path, request['approved_source'], {'artifacts': [{
        'warnings': [{'code': 'visible-duplicate-word', 'field': 'protocol.docx',
                      'publication_disposition': 'warning', 'affected_passage': passage}]}]})
    caution = next(row for row in cautions if row['caution']['code'] == 'visible-duplicate-word')
    assert caution['section_id'] == 'study-design.bias'
    assert passage in quality._review_evidence_text(caution['artifact_passage'])


def test_header_duplicate_passage_has_explicit_construction_review_location(tmp_path):
    from docx import Document
    _, request, _ = review_candidate(tmp_path)
    docx = tmp_path / 'candidate/protocol.docx'
    doc = Document(docx)
    passage = 'Study study header'
    doc.sections[0].header.paragraphs[0].text = passage
    doc.save(docx)
    cautions = quality._lexical_review_cautions(tmp_path, request['approved_source'], {'artifacts': [{
        'warnings': [{'code': 'visible-duplicate-word', 'field': 'protocol.docx',
                      'publication_disposition': 'warning', 'affected_passage': passage}]}]})
    caution = next(row for row in cautions if row['caution']['code'] == 'visible-duplicate-word')
    assert caution['section_id'] == 'protocol.shell'
    assert passage in quality._review_evidence_text(caution['artifact_passage'])
    assert 'header' in quality._review_evidence_text(caution['artifact_passage'])
    assert any(row['section_id'] == 'protocol.shell' for row in quality.content_review_sections(request['approved_source']))

    from test_clean_delivery import controlled_render
    render = controlled_render(ROOT, tmp_path)['render']
    render['artifacts'][0]['warnings'] = [caution['caution']]
    path = quality.create_verification_requests(tmp_path, request['approved_source'], render, review_set=2)[0]
    request = json.loads(path.read_text())
    response = acceptance_verification(request)
    response['status'] = 'blocked'
    response['findings'] = [{'finding_id': 'duplicate-header', 'artifact': 'protocol',
        'category': 'content', 'check': 'substantive', 'target_ids': ['protocol.shell'],
        'issue': 'Remove the repeated word in the running header.'}]
    for row in response['section_assessments']:
        if row['section_id'] == 'protocol.shell':
            row['status'] = 'blocked'
    for decision in response['uncertainty_resolutions']:
        if decision['caution_sha256'] == caution['caution_sha256']:
            decision.update(disposition='confirmed_defect', affected_passage=passage,
                            rationale='The repeated header word is an editorial error.')
    (tmp_path / request['response_path']).write_text(json.dumps(response))
    findings, _ = quality.validate_verifications(tmp_path, request_paths=[path])
    assert len(findings) == 1, findings
    assert findings[0]['owner'] == 'construction'
    assert findings[0]['target_ids'] == ['protocol.shell']
