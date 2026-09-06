"""Integration source-loss regressions; no inferred visit assignments."""
import json
from pathlib import Path

from docx import Document

from rendering import render_documents
from contracts import normalized_visit_records, protocol_table_contracts, timeline_findings
import copy

ROOT = Path(__file__).resolve().parents[1]


def test_evaluation_only_notes_survive_without_structured_table_rows(tmp_path):
    reference = json.loads((ROOT / "tests/fixtures/prospective-acceptance-source.json").read_text())
    for key in ("visit_schedule", "visit_schedule_table", "assessments"):
        reference["procedures"].pop(key, None)
    note = "Download speed, distance, and cadence without changing treatment."
    reference["procedures"]["evaluation"] = note
    render_documents(ROOT, tmp_path, reference, {"protocol": [], "icf": {}, "prs": {}})
    document = Document(tmp_path / "candidate/protocol.docx")
    assert note in [paragraph.text for paragraph in document.paragraphs]


def test_both_schedule_tables_preserve_the_complete_approved_visit_inventory(tmp_path):
    reference = json.loads((ROOT / "tests/fixtures/prospective-acceptance-source.json").read_text())
    visits = reference["procedures"]["visit_schedule_table"]
    visits[0]["visitNumber"] = "BL-7"
    visits[1]["visitNumber"] = "FU-2"
    visits[2]["visitNumber"] = "FINAL"
    # The procedure schedule describes only Baseline, plus one extra contact;
    # neither source is permission to silently truncate the other.
    reference["procedures"]["visit_schedule"].append({
        "visitNumber": "CALL", "visit": "Telephone review", "timing": "Day 14",
        "procedures": ["Device review"],
    })
    original = copy.deepcopy(reference)
    rows = normalized_visit_records(reference)
    assert [row["visitNumber"] for row in rows] == ["BL-7", "FU-2", "FINAL", "CALL"]
    assert rows[0]["procedures"] == ["Consent", "Device initiation"]
    assert [row.get("CRFnumber", "") for row in rows] == ["BL", "M1", "M3", ""]
    matrix = protocol_table_contracts(reference)["schedule-of-assessments"]["rows"]
    assert matrix[1] == ["Activity", "Visit BL-7", "Visit FU-2", "Visit FINAL", "Visit CALL"]
    assert next(row for row in matrix if row[0] == "Consent")[1:] == ["X", "", "", ""]
    assert next(row for row in matrix if row[0] == "Device review")[1:] == ["", "", "", "X"]
    render_documents(ROOT, tmp_path, reference, {"protocol": [], "icf": {}, "prs": {}})
    document = Document(tmp_path / "candidate/protocol.docx")
    schedule = next(table for table in document.tables
                    if " ".join(table.cell(0, 0).text.split()) == "Visit Number")
    assert [[cell.text for cell in row.cells] for row in schedule.rows[1:]] == [
        ["BL-7", "Baseline", "Day 0", "BL"],
        ["FU-2", "Month 1", "Day 30 +/- 7", "M1"],
        ["FINAL", "Month 3", "Day 90 +/- 7", "M3"],
        ["CALL", "Telephone review", "Day 14", ""],
    ]
    assert reference == original


def test_timeline_trailing_labels_do_not_attach_to_the_next_duration():
    reference = {
        "study": {"timeline": "4 weeks after baseline (interim assessment) and 12 weeks after baseline (final contact)."},
        "procedures": {"visit_schedule": [
            {"visit": "Baseline", "timing": "Day 0"},
            {"visit": "Interim", "timing": "Week 4 after baseline"},
            {"visit": "Final", "timing": "Week 12 after baseline"},
        ]},
    }
    original = copy.deepcopy(reference)
    assert timeline_findings(reference) == []
    assert reference == original
