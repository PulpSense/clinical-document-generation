import hashlib
import json
import platform
import sys
import time
import types
from pathlib import Path

import pytest
from docx import Document
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

import quality
from quality import RECOVERY_POLICIES, page_renderers, rasterize_pdf, render_assurance


ROOT = Path(__file__).resolve().parents[1]


def test_release_owned_pdfium_is_the_only_page_renderer(tmp_path):
    runtime = tmp_path / "runtime"
    runtime_python = runtime / "python"
    (runtime_python / "pypdfium2").mkdir(parents=True)
    (runtime_python / "pypdfium2/__init__.py").write_text("", encoding="utf-8")
    (runtime_python / "pypdfium2_raw").mkdir()
    (runtime_python / "pypdfium2_raw/__init__.py").write_text("", encoding="utf-8")
    wheel = tmp_path / "assets/runtime-wheels/pdfium.whl"
    wheel.parent.mkdir(parents=True)
    wheel.write_bytes(b"governed wheel")
    files = {
        "pypdfium2/__init__.py": hashlib.sha256(b"").hexdigest(),
        "pypdfium2_raw/__init__.py": hashlib.sha256(b"").hexdigest(),
    }
    manifest_payload = {"inventory": {"pdf_page_renderer": {
            "kind": "pypdfium2", "version": "5.13.0",
            "wheel": "assets/runtime-wheels/pdfium.whl",
            "wheel_sha256": hashlib.sha256(b"governed wheel").hexdigest(),
            "platform": "macosx_13_0_arm64",
            "runtime_inventory": [
                {"path": path, "bytes": 0, "sha256": digest}
                for path, digest in sorted(files.items())
            ],
        }}}
    package_fingerprint = hashlib.sha256(
        json.dumps(
            manifest_payload, sort_keys=True, separators=(",", ":"), ensure_ascii=False
        ).encode("utf-8")
    ).hexdigest()
    manifest = {**manifest_payload, "package_fingerprint": package_fingerprint}
    (tmp_path / "RELEASE-MANIFEST.json").write_text(
        json.dumps(manifest), encoding="utf-8",
    )
    assurance_path = tmp_path / "INSTALLATION-ASSURANCE.json"
    assurance_path.write_text(json.dumps({"status": "passed"}), encoding="utf-8")
    (tmp_path / "PROMOTION-RECORD.json").write_text(json.dumps({
        "schema_version": "promoted-release/v1",
        "status": "active",
        "package_fingerprint": package_fingerprint,
        "runtime_assurance_sha256": hashlib.sha256(
            assurance_path.read_bytes()
        ).hexdigest(),
    }), encoding="utf-8")
    (runtime / "PDF-RENDERER.json").write_text(
        __import__("json").dumps({
            "kind": "pypdfium2", "version": "5.13.0",
            "wheel": "assets/runtime-wheels/pdfium.whl",
            "wheel_sha256": hashlib.sha256(b"governed wheel").hexdigest(),
            "platform": "macosx_13_0_arm64", "installed_files": files,
        }),
        encoding="utf-8",
    )

    assert page_renderers(
        environment={"PATH": "/usr/bin:/opt/homebrew/bin"},
        skill_root=tmp_path,
    ) == [
        {
            "kind": "pypdfium2",
            "path": "python:pypdfium2",
            "module": "pypdfium2",
            "python_path": str(runtime_python),
            "version": "5.13.0",
            "source": "release-owned runtime",
            "wheel": "assets/runtime-wheels/pdfium.whl",
            "wheel_sha256": hashlib.sha256(b"governed wheel").hexdigest(),
        }
    ]


def test_pdfium_rasterization_does_not_require_optional_pillow(tmp_path, monkeypatch, governed_pdfium):
    pdf = tmp_path / "one-page.pdf"
    writer = PdfWriter()
    writer.add_blank_page(width=72, height=72)
    with pdf.open("wb") as handle:
        writer.write(handle)
    monkeypatch.setitem(__import__("sys").modules, "PIL", None)
    pages = quality._rasterize_pdfium_worker_render(
        pdf,
        tmp_path / "pages",
        governed_pdfium,
    )

    assert pages[0].read_bytes().startswith(b"\x89PNG\r\n\x1a\n")


def test_pdfium_rasterization_kills_a_hung_worker_and_cleans_partial_output(
    tmp_path, monkeypatch, governed_pdfium
):
    pdf = tmp_path / "one-page.pdf"
    writer = PdfWriter()
    writer.add_blank_page(width=72, height=72)
    with pdf.open("wb") as handle:
        writer.write(handle)
    output_dir = tmp_path / "pages"
    output_dir.mkdir()
    (output_dir / "page-1.png").write_bytes(b"stale")
    monkeypatch.setattr(
        quality,
        "_pdfium_worker_command",
        lambda _request: [sys.executable, "-c", "import time; time.sleep(60)"],
        raising=False,
    )

    with pytest.raises(RuntimeError, match="renderer.pdfium_worker_timeout"):
        rasterize_pdf(
            pdf,
            output_dir,
            governed_pdfium,
            timeout_seconds=0.05,
        )

    assert list(output_dir.iterdir()) == []


def test_pdfium_rasterization_reports_worker_crash_and_cleans_output(
    tmp_path, monkeypatch, governed_pdfium
):
    pdf = tmp_path / "one-page.pdf"
    writer = PdfWriter()
    writer.add_blank_page(width=72, height=72)
    with pdf.open("wb") as handle:
        writer.write(handle)
    output_dir = tmp_path / "pages"
    crash_script = (
        "import json,sys; from pathlib import Path; "
        "request=json.loads(Path(sys.argv[1]).read_text()); "
        "Path(request['output_dir'],'page-1.png').write_bytes(b'partial'); "
        "raise SystemExit(17)"
    )
    monkeypatch.setattr(
        quality,
        "_pdfium_worker_command",
        lambda request: [sys.executable, "-c", crash_script, str(request)],
    )

    with pytest.raises(RuntimeError, match="renderer.pdfium_worker_crashed"):
        rasterize_pdf(pdf, output_dir, governed_pdfium)

    assert list(output_dir.iterdir()) == []


@pytest.mark.skipif(platform.system() == "Windows", reason="Windows has no POSIX process groups")
def test_pdfium_crash_kills_worker_descendants(tmp_path, monkeypatch, governed_pdfium):
    pdf = tmp_path / "one-page.pdf"
    writer = PdfWriter()
    writer.add_blank_page(width=72, height=72)
    with pdf.open("wb") as handle:
        writer.write(handle)
    orphan_marker = tmp_path / "crash-orphan-survived"
    child_script = (
        "import sys,time; from pathlib import Path; "
        "time.sleep(0.3); Path(sys.argv[1]).write_text('survived')"
    )
    parent_script = (
        "import subprocess,sys; "
        "subprocess.Popen([sys.executable,'-c',sys.argv[1],sys.argv[2]], "
        "stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL); "
        "raise SystemExit(17)"
    )
    monkeypatch.setattr(
        quality,
        "_pdfium_worker_command",
        lambda _request: [
            sys.executable, "-c", parent_script, child_script, str(orphan_marker)
        ],
    )

    with pytest.raises(RuntimeError, match="renderer.pdfium_worker_crashed"):
        rasterize_pdf(pdf, tmp_path / "pages", governed_pdfium)
    time.sleep(0.5)
    assert not orphan_marker.exists()


@pytest.mark.skipif(platform.system() == "Windows", reason="Windows has no POSIX process groups")
def test_pdfium_timeout_kills_worker_descendants(tmp_path, monkeypatch, governed_pdfium):
    pdf = tmp_path / "one-page.pdf"
    writer = PdfWriter()
    writer.add_blank_page(width=72, height=72)
    with pdf.open("wb") as handle:
        writer.write(handle)
    orphan_marker = tmp_path / "orphan-survived"
    child_script = (
        "import sys,time; from pathlib import Path; "
        "time.sleep(0.3); Path(sys.argv[1]).write_text('survived')"
    )
    parent_script = (
        "import subprocess,sys,time; "
        "subprocess.Popen([sys.executable,'-c',sys.argv[1],sys.argv[2]]); "
        "time.sleep(5)"
    )
    monkeypatch.setattr(
        quality,
        "_pdfium_worker_command",
        lambda _request: [
            sys.executable, "-c", parent_script, child_script, str(orphan_marker)
        ],
    )

    with pytest.raises(RuntimeError, match="renderer.pdfium_worker_timeout"):
        rasterize_pdf(
            pdf,
            tmp_path / "pages",
            governed_pdfium,
            timeout_seconds=0.05,
        )
    time.sleep(0.5)
    assert not orphan_marker.exists()


def test_pdfium_worker_rejects_pathological_page_dimensions(
    tmp_path, governed_pdfium
):
    pdf = tmp_path / "huge.pdf"
    writer = PdfWriter()
    writer.add_blank_page(width=2_000_000, height=2_000_000)
    with pdf.open("wb") as handle:
        writer.write(handle)

    with pytest.raises(RuntimeError, match="renderer.pdfium_dimension_limit_exceeded"):
        quality._rasterize_pdfium_worker_render(
            pdf,
            tmp_path / "pages",
            governed_pdfium,
        )


def test_pdfium_worker_rejects_excessive_page_count(tmp_path, governed_pdfium):
    pdf = tmp_path / "many-pages.pdf"
    writer = PdfWriter()
    for _ in range(quality.PDFIUM_MAX_PAGES + 1):
        writer.add_blank_page(width=72, height=72)
    with pdf.open("wb") as handle:
        writer.write(handle)

    with pytest.raises(RuntimeError, match="renderer.pdfium_page_limit_exceeded"):
        quality._rasterize_pdfium_worker_render(
            pdf,
            tmp_path / "pages",
            governed_pdfium,
        )


def test_pdfium_worker_rejects_excessive_output_bytes(
    tmp_path, monkeypatch, governed_pdfium
):
    pdf = tmp_path / "one-page.pdf"
    writer = PdfWriter()
    writer.add_blank_page(width=72, height=72)
    with pdf.open("wb") as handle:
        writer.write(handle)
    monkeypatch.setattr(quality, "PDFIUM_MAX_OUTPUT_BYTES", 1)

    with pytest.raises(RuntimeError, match="renderer.pdfium_output_limit_exceeded"):
        quality._rasterize_pdfium_worker_render(
            pdf,
            tmp_path / "pages",
            governed_pdfium,
        )


def test_pdfium_worker_rejects_excessive_total_pixels(
    tmp_path, monkeypatch, governed_pdfium
):
    pdf = tmp_path / "one-page.pdf"
    writer = PdfWriter()
    writer.add_blank_page(width=72, height=72)
    with pdf.open("wb") as handle:
        writer.write(handle)
    monkeypatch.setattr(quality, "PDFIUM_MAX_TOTAL_PIXELS", 1)

    with pytest.raises(RuntimeError, match="renderer.pdfium_pixel_limit_exceeded"):
        quality._rasterize_pdfium_worker_render(
            pdf,
            tmp_path / "pages",
            governed_pdfium,
        )


def test_pdfium_worker_applies_every_supported_resource_limit(monkeypatch):
    calls = []
    fake_resource = types.SimpleNamespace(
        RLIMIT_CPU=1,
        RLIMIT_AS=2,
        RLIMIT_FSIZE=3,
        RLIMIT_NOFILE=4,
        RLIM_INFINITY=-1,
        getrlimit=lambda _kind: (-1, -1),
        setrlimit=lambda kind, limits: calls.append((kind, limits)),
    )
    monkeypatch.setitem(sys.modules, "resource", fake_resource)
    monkeypatch.setattr(quality.platform, "system", lambda: "Linux")

    assert quality._supported_pdfium_worker_resource_limits() == [
        "RLIMIT_CPU", "RLIMIT_AS", "RLIMIT_FSIZE", "RLIMIT_NOFILE",
    ]
    assert quality._apply_pdfium_worker_resource_limits() == [
        "RLIMIT_CPU", "RLIMIT_AS", "RLIMIT_FSIZE", "RLIMIT_NOFILE",
    ]
    assert [kind for kind, _limits in calls] == [1, 2, 3, 4]


def test_pdfium_worker_fails_closed_when_supported_resource_limit_cannot_apply(monkeypatch):
    def reject_memory(kind, _limits):
        if kind == 2:
            raise OSError("controlled rejection")

    fake_resource = types.SimpleNamespace(
        RLIMIT_CPU=1,
        RLIMIT_AS=2,
        RLIMIT_FSIZE=3,
        RLIMIT_NOFILE=4,
        RLIM_INFINITY=-1,
        getrlimit=lambda _kind: (-1, -1),
        setrlimit=reject_memory,
    )
    monkeypatch.setitem(sys.modules, "resource", fake_resource)
    monkeypatch.setattr(quality.platform, "system", lambda: "Linux")

    with pytest.raises(RuntimeError, match="renderer.pdfium_resource_limit_failed"):
        quality._apply_pdfium_worker_resource_limits()


def test_pdfium_worker_requires_linux_address_space_limit(monkeypatch):
    fake_resource = types.SimpleNamespace(
        RLIMIT_CPU=1,
        RLIMIT_FSIZE=3,
        RLIMIT_NOFILE=4,
        RLIM_INFINITY=-1,
        getrlimit=lambda _kind: (-1, -1),
        setrlimit=lambda _kind, _limits: None,
    )
    monkeypatch.setitem(sys.modules, "resource", fake_resource)
    monkeypatch.setattr(quality.platform, "system", lambda: "Linux")

    with pytest.raises(
        RuntimeError,
        match="renderer.pdfium_resource_limits_unavailable.*RLIMIT_AS",
    ):
        quality._apply_pdfium_worker_resource_limits()


def test_pdfium_worker_excludes_nonoperational_darwin_address_space_limit(monkeypatch):
    calls = []
    fake_resource = types.SimpleNamespace(
        RLIMIT_CPU=1,
        RLIMIT_AS=2,
        RLIMIT_FSIZE=3,
        RLIMIT_NOFILE=4,
        RLIM_INFINITY=-1,
        getrlimit=lambda _kind: (-1, -1),
        setrlimit=lambda kind, limits: calls.append((kind, limits)),
    )
    monkeypatch.setitem(sys.modules, "resource", fake_resource)
    monkeypatch.setattr(quality.platform, "system", lambda: "Darwin")

    assert quality._required_pdfium_worker_resource_limits() == [
        "RLIMIT_CPU", "RLIMIT_FSIZE", "RLIMIT_NOFILE",
    ]
    assert quality._apply_pdfium_worker_resource_limits() == [
        "RLIMIT_CPU", "RLIMIT_FSIZE", "RLIMIT_NOFILE",
    ]
    assert [kind for kind, _limits in calls] == [1, 3, 4]


def test_pdfium_worker_accepts_exact_page_pixel_dimension_and_output_boundaries(
    tmp_path, monkeypatch, governed_pdfium
):
    pdf = tmp_path / "one-page.pdf"
    writer = PdfWriter()
    writer.add_blank_page(width=72, height=72)
    with pdf.open("wb") as handle:
        writer.write(handle)
    baseline = quality._rasterize_pdfium_worker_render(
        pdf,
        tmp_path / "baseline",
        governed_pdfium,
        dpi=130,
    )[0].read_bytes()
    monkeypatch.setattr(quality, "PDFIUM_MAX_PAGES", 1)
    monkeypatch.setattr(quality, "PDFIUM_MAX_DIMENSION_PIXELS", 130)
    monkeypatch.setattr(quality, "PDFIUM_MAX_TOTAL_PIXELS", 130 * 130)
    monkeypatch.setattr(quality, "PDFIUM_MAX_OUTPUT_BYTES", len(baseline))

    bounded = quality._rasterize_pdfium_worker_render(
        pdf,
        tmp_path / "bounded",
        governed_pdfium,
        dpi=130,
    )

    assert bounded[0].read_bytes() == baseline


def test_pdfium_worker_rejects_one_over_each_numeric_boundary(
    tmp_path, monkeypatch, governed_pdfium
):
    pdf = tmp_path / "one-page.pdf"
    writer = PdfWriter()
    writer.add_blank_page(width=72, height=72)
    with pdf.open("wb") as handle:
        writer.write(handle)
    baseline = quality._rasterize_pdfium_worker_render(
        pdf,
        tmp_path / "baseline",
        governed_pdfium,
        dpi=130,
    )[0].read_bytes()

    monkeypatch.setattr(quality, "PDFIUM_MAX_DIMENSION_PIXELS", 129)
    with pytest.raises(RuntimeError, match="renderer.pdfium_dimension_limit_exceeded"):
        quality._rasterize_pdfium_worker_render(
            pdf, tmp_path / "dimension", governed_pdfium, dpi=130
        )
    monkeypatch.setattr(quality, "PDFIUM_MAX_DIMENSION_PIXELS", 130)

    monkeypatch.setattr(quality, "PDFIUM_MAX_TOTAL_PIXELS", 130 * 130 - 1)
    with pytest.raises(RuntimeError, match="renderer.pdfium_pixel_limit_exceeded"):
        quality._rasterize_pdfium_worker_render(
            pdf, tmp_path / "pixels", governed_pdfium, dpi=130
        )
    monkeypatch.setattr(quality, "PDFIUM_MAX_TOTAL_PIXELS", 130 * 130)

    monkeypatch.setattr(quality, "PDFIUM_MAX_OUTPUT_BYTES", len(baseline) - 1)
    with pytest.raises(RuntimeError, match="renderer.pdfium_output_limit_exceeded"):
        quality._rasterize_pdfium_worker_render(
            pdf, tmp_path / "output", governed_pdfium, dpi=130
        )


@pytest.mark.parametrize(("constant", "replacement", "page_count", "page_size", "code"), [
    (
        "PDFIUM_MAX_PAGES = 400", "PDFIUM_MAX_PAGES = 1", 2, (72, 72),
        "renderer.pdfium_page_limit_exceeded",
    ),
    (
        "PDFIUM_MAX_DIMENSION_PIXELS = 20_000",
        "PDFIUM_MAX_DIMENSION_PIXELS = 100",
        1,
        (72, 72),
        "renderer.pdfium_dimension_limit_exceeded",
    ),
    (
        "PDFIUM_MAX_TOTAL_PIXELS = 1_000_000_000",
        "PDFIUM_MAX_TOTAL_PIXELS = 100",
        1,
        (72, 72),
        "renderer.pdfium_pixel_limit_exceeded",
    ),
    (
        "PDFIUM_MAX_OUTPUT_BYTES = 1_000_000_000",
        "PDFIUM_MAX_OUTPUT_BYTES = 1",
        1,
        (72, 72),
        "renderer.pdfium_output_limit_exceeded",
    ),
])
def test_real_worker_protocol_preserves_each_limit_finding(
    tmp_path,
    governed_pdfium,
    constant,
    replacement,
    page_count,
    page_size,
    code,
):
    release_root = Path(governed_pdfium["python_path"]).parents[1]
    worker_source = release_root / "scripts/quality.py"
    source = worker_source.read_text(encoding="utf-8")
    assert source.count(constant) == 1
    worker_source.write_text(source.replace(constant, replacement), encoding="utf-8")
    pdf = tmp_path / "bounded.pdf"
    writer = PdfWriter()
    for _ in range(page_count):
        writer.add_blank_page(width=page_size[0], height=page_size[1])
    with pdf.open("wb") as handle:
        writer.write(handle)

    with pytest.raises(RuntimeError, match=code):
        rasterize_pdf(pdf, tmp_path / "pages", governed_pdfium)


def test_recovery_classes_have_one_governed_action_each():
    assert RECOVERY_POLICIES == {
        "adapter_fault": "advance_adapter",
        "font_capability_uncertainty": "bounded_smoke_render",
        "deterministic_structure_defect": "rebuild_deterministic_structure",
        "document_structure_defect": "preserve_and_stop",
        "visual_defect": "targeted_layout_repair",
        "drafting_defect": "retry_drafting_target",
        "verifier_transient": "retry_verifier",
        "transport_fault": "retry_exact_bytes",
    }


def _write_text_pdf(path: Path) -> None:
    writer = PdfWriter()
    page = writer.add_blank_page(width=612, height=792)
    page[NameObject("/Resources")] = DictionaryObject({
        NameObject("/Font"): DictionaryObject({
            NameObject("/F1"): DictionaryObject({
                NameObject("/Type"): NameObject("/Font"),
                NameObject("/Subtype"): NameObject("/Type1"),
                NameObject("/BaseFont"): NameObject("/Helvetica"),
            })
        })
    })
    stream = DecodedStreamObject()
    stream.set_data(b"BT /F1 12 Tf 72 720 Td (Render Assurance) Tj ET")
    page[NameObject("/Contents")] = stream
    with path.open("wb") as handle:
        writer.write(handle)


def _bundle() -> dict:
    return {
        "approved_font_plan": {
            "approved_fallbacks": {"missing sans": "Liberation Sans"},
            "packaged_families": {"Liberation Sans": "LiberationSans-Regular.ttf"},
            "packaged_font_assets": {
                "Liberation Sans": "assets/fallback-fonts/LiberationSans-Regular.ttf"
            },
        }
    }


def _structural_validation(revision_dir: Path, *expected_files: str) -> dict:
    return {
        "status": "structurally_valid",
        "expected_files": list(expected_files),
        "candidate_files": [
            {
                "path": path.relative_to(revision_dir).as_posix(),
                "sha256": hashlib.sha256(path.read_bytes()).hexdigest(),
                "bytes": path.stat().st_size,
            }
            for path in sorted((revision_dir / "candidate").glob("*"))
            if path.is_file()
        ],
    }


def test_render_assurance_records_tri_state_fonts_and_binds_substitutions_to_exact_artifacts(tmp_path, governed_pdfium):
    candidate = tmp_path / "candidate"
    candidate.mkdir()
    document = Document()
    for font in ("Available Sans", "Missing Sans", "Unknown Sans"):
        run = document.add_paragraph().add_run(font)
        run.font.name = font
    document.save(candidate / "protocol.docx")

    states = {
        "Available Sans": (True, "host inventory"),
        "Missing Sans": (False, "host inventory"),
        "Unknown Sans": (None, "inventory unavailable"),
    }

    def rebuild(substitutions):
        rebuilt = Document(candidate / "protocol.docx")
        for paragraph in rebuilt.paragraphs:
            for run in paragraph.runs:
                if run.font.name in substitutions:
                    run.font.name = substitutions[run.font.name]
        rebuilt.save(candidate / "protocol.docx")
        return {"status": "passed"}

    def export(docx, output_dir, _identity, **_kwargs):
        pdf = output_dir / f"{docx.stem}.pdf"
        _write_text_pdf(pdf)
        return pdf

    def rasterize(_pdf, output_dir, _identity, **_kwargs):
        output_dir.mkdir(parents=True, exist_ok=True)
        page = output_dir / "page-1.png"
        page.write_bytes(b"page image")
        return [page]

    office = {"kind": "LibreOffice", "path": "/controlled/soffice"}
    pages = governed_pdfium
    report = render_assurance(
        ROOT,
        tmp_path,
        {},
        contracted_bundle=_bundle(),
        structural_validation=_structural_validation(tmp_path, "protocol.docx"),
        renderer_identities=[office],
        page_renderer_identities=[pages],
        font_probe=lambda font, **_kwargs: states[font],
        rebuild_candidate=rebuild,
        office_exporter=export,
        page_exporter=rasterize,
        blank_page_detector=lambda _pdf: [],
    )

    assert report["status"] == "passed"
    assert {font: evidence["state"] for font, evidence in report["fonts"].items()} == {
        "Available Sans": "available",
        "Missing Sans": "missing-or-unusable",
        "Unknown Sans": "unknown",
    }
    assert report["fonts"]["Unknown Sans"]["resolution"] == "render_verified"
    assert report["font_substitutions"] == {"Missing Sans": "Liberation Sans"}
    assert report["candidate"]["font_substitutions"] == report["font_substitutions"]
    assert report["candidate"]["font_evidence"] == report["fonts"]
    rebuilt_fonts = {run.font.name for paragraph in Document(candidate / "protocol.docx").paragraphs for run in paragraph.runs}
    assert rebuilt_fonts == {"Available Sans", "Liberation Sans", "Unknown Sans"}
    artifact = report["render"]["artifacts"][0]
    assert artifact["renderer"] == office
    assert artifact["page_renderer"] == pages
    assert artifact["font_substitutions"] == report["font_substitutions"]
    assert artifact["font_evidence"] == report["fonts"]
    for path_key, hash_key in (("docx", "docx_sha256"), ("pdf", "pdf_sha256")):
        assert hashlib.sha256((tmp_path / artifact[path_key]).read_bytes()).hexdigest() == artifact[hash_key]
    assert hashlib.sha256((tmp_path / artifact["pages"][0]["path"]).read_bytes()).hexdigest() == artifact["pages"][0]["sha256"]


def test_render_assurance_advances_ordered_adapters_with_governed_recovery_records(tmp_path, governed_pdfium):
    candidate = tmp_path / "candidate"
    candidate.mkdir()
    document = Document()
    document.add_paragraph("Complete candidate")
    document.save(candidate / "protocol.docx")
    original = (candidate / "protocol.docx").read_bytes()
    word = {"kind": "Microsoft Word", "path": "/controlled/word"}
    libreoffice = {"kind": "LibreOffice", "path": "/controlled/soffice"}
    pdfium = governed_pdfium

    def export(docx, output_dir, identity, **_kwargs):
        if identity == word:
            raise RuntimeError("controlled office failure")
        pdf = output_dir / f"{docx.stem}.pdf"
        _write_text_pdf(pdf)
        return pdf

    def rasterize(_pdf, output_dir, identity, **_kwargs):
        output_dir.mkdir(parents=True, exist_ok=True)
        page = output_dir / "page-1.png"
        page.write_bytes(b"page image")
        return [page]

    report = render_assurance(
        ROOT,
        tmp_path,
        {},
        contracted_bundle=_bundle(),
        structural_validation=_structural_validation(tmp_path, "protocol.docx"),
        renderer_identities=[word, libreoffice],
        page_renderer_identities=[pdfium],
        font_probe=lambda _font, **_kwargs: (True, "available"),
        office_exporter=export,
        page_exporter=rasterize,
        blank_page_detector=lambda _pdf: [],
    )

    assert report["status"] == "passed"
    assert (candidate / "protocol.docx").read_bytes() == original
    assert report["render"]["renderer_attempts"] == [
        {
            "adapter": word,
            "status": "failed",
            "recovery_class": "adapter_fault",
            "action": "advance_adapter",
            "issue": "controlled office failure",
        },
        {"adapter": libreoffice, "status": "passed"},
    ]
    assert report["render"]["page_renderer_attempts"] == [
        {"adapter": pdfium, "status": "passed"},
    ]


def test_render_assurance_stops_when_the_one_pdfium_renderer_fails(tmp_path, governed_pdfium):
    candidate = tmp_path / "candidate"
    candidate.mkdir()
    document = Document()
    document.add_paragraph("Complete candidate")
    document.save(candidate / "protocol.docx")
    pdfium = governed_pdfium
    unapproved_host_renderer = {
        "kind": "pdftoppm",
        "path": "/controlled/pdftoppm",
    }

    def export(docx, output_dir, _identity, **_kwargs):
        pdf = output_dir / f"{docx.stem}.pdf"
        _write_text_pdf(pdf)
        return pdf

    used = []

    def rasterize(_pdf, output_dir, identity, **_kwargs):
        used.append(identity)
        if identity == pdfium:
            raise RuntimeError("controlled PDFium failure")
        output_dir.mkdir(parents=True, exist_ok=True)
        page = output_dir / "page-1.png"
        page.write_bytes(b"unapproved page image")
        return [page]

    report = render_assurance(
        ROOT,
        tmp_path,
        {},
        contracted_bundle=_bundle(),
        structural_validation=_structural_validation(tmp_path, "protocol.docx"),
        renderer_identities=[{"kind": "LibreOffice", "path": "/controlled/soffice"}],
        page_renderer_identities=[pdfium, unapproved_host_renderer],
        font_probe=lambda _font, **_kwargs: (True, "available"),
        office_exporter=export,
        page_exporter=rasterize,
        blank_page_detector=lambda _pdf: [],
    )

    assert report["status"] == "blocked"
    assert used == [pdfium]
    assert report["render"]["page_renderer_attempts"] == [{
        "adapter": pdfium,
        "status": "failed",
        "recovery_class": "adapter_fault",
        "action": "stop",
        "issue": "controlled PDFium failure",
    }]
    assert report["findings"] == [{
        "category": "renderer",
        "field": "rendering",
        "code": "renderer.pdfium_worker_failed",
        "issue": "controlled PDFium failure",
        "recovery_class": "adapter_fault",
        "action": "stop",
    }]


def test_implicit_render_assurance_preserves_runtime_integrity_finding(
    tmp_path, monkeypatch
):
    candidate = tmp_path / "candidate"
    candidate.mkdir()
    document = Document()
    document.add_paragraph("Complete candidate")
    document.save(candidate / "protocol.docx")
    finding = {
        "category": "renderer",
        "field": "page_renderer",
        "code": "renderer.pdfium_runtime_file_changed",
        "path": "pypdfium2/__init__.py",
        "issue": "The installed PDFium file changed: pypdfium2/__init__.py.",
    }
    monkeypatch.setattr(quality, "page_renderers", lambda **_kwargs: [])
    monkeypatch.setattr(
        quality,
        "_pdfium_runtime_integrity",
        lambda _root, **_kwargs: {"status": "blocked", "finding": finding},
    )

    report = render_assurance(
        ROOT,
        tmp_path,
        {},
        contracted_bundle=_bundle(),
        structural_validation=_structural_validation(tmp_path, "protocol.docx"),
        renderer_identities=[{"kind": "LibreOffice", "path": "/controlled/soffice"}],
        font_probe=lambda _font, **_kwargs: (True, "available"),
    )

    assert report["status"] == "blocked"
    assert report["findings"][0]["code"] == finding["code"]
    assert report["findings"][0]["path"] == finding["path"]


def test_render_assurance_exhaustion_preserves_candidate_and_emits_one_diagnostic(tmp_path, governed_pdfium):
    candidate = tmp_path / "candidate"
    candidate.mkdir()
    document = Document()
    document.add_paragraph("Structurally valid candidate")
    document.save(candidate / "protocol.docx")
    stale_pdf = tmp_path / "rendered/protocol.pdf"
    stale_page = tmp_path / "rendered/protocol/page-1.png"
    stale_page.parent.mkdir(parents=True)
    stale_pdf.write_bytes(b"stale pdf")
    stale_page.write_bytes(b"stale page")
    original_hash = hashlib.sha256((candidate / "protocol.docx").read_bytes()).hexdigest()
    word = {"kind": "Microsoft Word", "path": "/controlled/word"}
    libreoffice = {"kind": "LibreOffice", "path": "/controlled/soffice"}

    report = render_assurance(
        ROOT,
        tmp_path,
        {},
        contracted_bundle=_bundle(),
        structural_validation=_structural_validation(tmp_path, "protocol.docx"),
        renderer_identities=[word, libreoffice],
        page_renderer_identities=[governed_pdfium],
        font_probe=lambda _font, **_kwargs: (True, "available"),
        office_exporter=lambda *_args, **_kwargs: (_ for _ in ()).throw(RuntimeError("controlled failure")),
        page_exporter=lambda *_args, **_kwargs: [],
    )

    assert report["status"] == "blocked"
    assert report["candidate"]["files"] == [{
        "path": "candidate/protocol.docx",
        "sha256": original_hash,
        "bytes": (candidate / "protocol.docx").stat().st_size,
    }]
    assert report["diagnostic"]["recovery_class"] == "adapter_fault"
    assert report["diagnostic"]["action"] == RECOVERY_POLICIES["adapter_fault"]
    assert report["diagnostic"]["outcome"] == "adapters_exhausted"
    assert report["diagnostic"]["candidate_disposition"] == "preserved"
    assert report["diagnostic"]["office_attempts"] == report["render"]["renderer_attempts"]
    assert report["diagnostic"]["page_attempts"] == []
    assert len(report["findings"]) == 1
    assert not stale_pdf.exists()
    assert not stale_page.exists()


def test_render_assurance_reuses_a_substituted_candidate_without_oscillating(tmp_path, governed_pdfium):
    candidate = tmp_path / "candidate"
    candidate.mkdir()
    document = Document()
    run = document.add_paragraph().add_run("Missing font")
    run.font.name = "Missing Sans"
    document.save(candidate / "protocol.docx")
    rebuilds = []

    def rebuild(substitutions):
        rebuilds.append(dict(substitutions))
        rebuilt = Document(candidate / "protocol.docx")
        for paragraph in rebuilt.paragraphs:
            for run in paragraph.runs:
                if run.font.name in substitutions:
                    run.font.name = substitutions[run.font.name]
        rebuilt.save(candidate / "protocol.docx")
        return {"status": "passed"}

    def probe(font, **_kwargs):
        return (False, "missing") if font == "Missing Sans" else (True, "available")

    def fail_export(*_args, **_kwargs):
        raise RuntimeError("controlled failure")

    first = render_assurance(
        ROOT,
        tmp_path,
        {},
        contracted_bundle=_bundle(),
        structural_validation=_structural_validation(tmp_path, "protocol.docx"),
        renderer_identities=[{"kind": "LibreOffice", "path": "/controlled/soffice"}],
        page_renderer_identities=[governed_pdfium],
        font_probe=probe,
        rebuild_candidate=rebuild,
        office_exporter=fail_export,
    )
    substituted_hash = hashlib.sha256((candidate / "protocol.docx").read_bytes()).hexdigest()
    second = render_assurance(
        ROOT,
        tmp_path,
        {},
        contracted_bundle=_bundle(),
        structural_validation=_structural_validation(tmp_path, "protocol.docx"),
        candidate_font_substitutions=first["font_substitutions"],
        renderer_identities=[{"kind": "LibreOffice", "path": "/controlled/soffice"}],
        page_renderer_identities=[governed_pdfium],
        font_probe=probe,
        rebuild_candidate=rebuild,
        office_exporter=fail_export,
    )

    assert first["font_substitutions"] == second["font_substitutions"] == {"Missing Sans": "Liberation Sans"}
    assert rebuilds == [{"Missing Sans": "Liberation Sans"}]
    assert hashlib.sha256((candidate / "protocol.docx").read_bytes()).hexdigest() == substituted_hash


def test_render_assurance_rejects_an_incomplete_branch_candidate(tmp_path):
    candidate = tmp_path / "candidate"
    candidate.mkdir()
    Document().save(candidate / "protocol.docx")

    report = render_assurance(
        ROOT,
        tmp_path,
        {},
        contracted_bundle=_bundle(),
        structural_validation=_structural_validation(
            tmp_path,
            "protocol.docx",
            "icf.docx",
            "study.xml",
        ),
        renderer_identities=[],
        page_renderer_identities=[],
    )

    assert report["status"] == "blocked"
    assert report["render"]["status"] == "not_run"
    assert report["findings"][0]["recovery_class"] == "document_structure_defect"
    assert "complete structurally validated" in report["findings"][0]["issue"]


def test_render_assurance_revalidates_the_exact_candidate_after_font_substitution(tmp_path, governed_pdfium):
    candidate = tmp_path / "candidate"
    candidate.mkdir()
    document = Document()
    run = document.add_paragraph().add_run("Missing font")
    run.font.name = "Missing Sans"
    document.save(candidate / "protocol.docx")
    (candidate / "study.xml").write_text("<study />", encoding="utf-8")

    def rebuild(_substitutions):
        (candidate / "study.xml").unlink()
        return {"status": "passed"}

    report = render_assurance(
        ROOT,
        tmp_path,
        {},
        contracted_bundle=_bundle(),
        structural_validation=_structural_validation(
            tmp_path,
            "protocol.docx",
            "study.xml",
        ),
        renderer_identities=[{"kind": "LibreOffice", "path": "/controlled/soffice"}],
        page_renderer_identities=[governed_pdfium],
        font_probe=lambda font, **_kwargs: (False, "missing") if font == "Missing Sans" else (True, "available"),
        rebuild_candidate=rebuild,
    )

    assert report["status"] == "blocked"
    assert report["render"]["status"] == "not_run"
    assert report["findings"][0]["recovery_class"] == "document_structure_defect"
    assert report["structural_validation"]["revalidated_after_substitution"] is False
