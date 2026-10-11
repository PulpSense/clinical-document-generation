"""Public delivery with source-bound construction recovery across source shapes."""
import copy
import json
from pathlib import Path
import pytest
from docx import Document
from xml.etree import ElementTree as ET

import drafting
import workflow
from test_clean_delivery import deliver
from test_construction_reliability import FAMILIES, source_for
from test_handoff_quality import acceptance_verification


@pytest.mark.parametrize('branch,family,fault', [
    ('Prospective', 'Advarra', 'email'), ('Prospective', 'Sterling', 'phone'),
    ('Ambispective', 'Advarra', 'quantity'), ('Ambispective', 'Sterling', 'title'),
    ('Prospective', 'Advarra', 'date'), ('Ambispective', 'Advarra', 'email'),
    ('Ambispective', 'Sterling', 'stale'),
    ('Prospective', 'Sterling', 'invalid-proposal'), ('Ambispective', 'Advarra', 'exhausted'),
])
def test_xml_field_defect_receives_source_bound_edit_and_clean_delivery(tmp_path, monkeypatch, branch, family, fault):
    from contracts import get_path
    source = source_for(branch, family)
    source['population']['sample_size'] = '40'
    source.setdefault('regulatory', {}).setdefault('prs', {})['verification_date'] = '2031-07-09'
    source['study']['short_title'] = 'Routine Function Follow-up'
    mode = fault
    fault = 'email' if mode in {'stale', 'invalid-proposal', 'exhausted'} else fault
    target, evidence, wrong = {
        'email': ('overall_contact/email', 'parties.study_coordinator.email', 'wrong@example.invalid'),
        'phone': ('overall_contact/phone', 'parties.study_coordinator.business_phone', '+1 202-555-0199'),
        'quantity': ('enrollment', 'population.sample_size', '49'),
        'title': ('brief_title', 'study.short_title', 'Incorrect study title'),
        'date': ('verification_date', 'regulatory.prs.verification_date', '2031-07-10'),
    }[fault]
    expected = get_path(source, evidence)
    original_generate = workflow.generate_xml
    edits, reviews = [], []

    def broken_generator(*args, **kwargs):
        report = original_generate(*args, **kwargs)
        tree = ET.parse(args[1])
        tree.getroot().find('.//clinical_study/' + target).text = 'changed-base@example.invalid' if mode == 'stale' and edits else wrong
        tree.write(args[1], encoding='utf-8', xml_declaration=True)
        return report

    monkeypatch.setattr(workflow, 'generate_xml', broken_generator)

    def review(request):
        response = acceptance_verification(request)
        if request['task'] != 'clinical_content_verification': return response
        reviews.append(request['request_id'])
        path = tmp_path / 'revisions' / request['revision_id'] / 'candidate/study.xml'
        node = ET.parse(path).getroot().find('.//clinical_study/' + target)
        if node.text != expected:
            response['status'] = 'blocked'
            response['findings'] = [{
                'finding_id': 'wrong-generated-field', 'category': 'content', 'check': 'source_supported',
                'artifact': 'study.xml', 'surface': 'structured', 'target_ids': ['prs.structured'],
                'material': True, 'contradiction': True,
                'source_path': evidence,
                'affected_passage': ET.tostring(node, encoding='unicode').strip(),
                'document_locations': ['study.xml — clinical_study/' + target],
                'issue': 'The generated scalar differs from its approved source value.',
            }]
            next(r for r in response['section_assessments'] if r['section_id'] == 'prs.structured')['status'] = 'failed'
        return response

    def draft(request):
        if not request.get('construction_repair'): return drafting.recorded_acceptance_response(request)
        edits.append(request)
        response = drafting.response_template(request, model_id='Controlled field editor')
        field = next(f for f in request['construction_repair']['fields'] if f['path'].endswith('/' + target.split('/')[-1] + '[1]'))
        after = ('invented@example.invalid' if mode == 'exhausted' or (mode == 'invalid-proposal' and len(edits) == 1) else
                 expected)
        response['field_edits'] = [{'path': field['path'], 'before': field['before'], 'after': after,
                                     'source_path': evidence}]
        return response

    result, revision = deliver(tmp_path, monkeypatch, source, draft=draft, review=review)
    if mode == 'exhausted':
        assert result['status'] == 'blocked' and result['stage'] == 'internal_recovery_stalled', result
        assert len(edits) == 3 and not result['client_outputs']
        assert not (tmp_path / 'output/study.xml').exists()
        return
    if mode == 'stale':
        assert result['status'] == 'blocked' and result['stage'] == 'source_integrity', result
        assert not result['client_outputs'] and not (tmp_path / 'output/study.xml').exists()
        return
    assert result['status'] == 'passed', result
    assert len(edits) == (2 if mode == 'invalid-proposal' else 1)
    assert len(reviews) == 2 and len(set(reviews)) == len(reviews)
    assert ET.parse(tmp_path / 'output/study.xml').getroot().findtext('.//clinical_study/' + target) == expected
    assert not result['warnings']
    assert get_path(json.loads((revision / 'approved-reference.json').read_text()), evidence) == expected


def test_explicit_compound_name_components_are_not_combined_with_guesses(tmp_path):
    from pathlib import Path
    from prs_xml import generate
    source = source_for('Ambispective', 'Advarra')
    source['sites'][0].pop('investigators', None)
    source['sites'][0]['investigator'] = {
        'first_name': 'Jordan', 'last_name': 'Example 1', 'name': 'Jordan Example 1', 'degrees': 'MD'}
    path = tmp_path / 'study.xml'
    generate(Path('assets/client-templates/prs/clinicaltrials_prs_full_placeholder_template.xml'), path,
             source, {'brief_summary': 'Summary.', 'detailed_description': 'Description.'})
    inv = ET.parse(path).getroot().find('.//location/investigator')
    assert inv.findtext('first_name') == 'Jordan'
    assert inv.findtext('last_name') == 'Example 1'
    assert inv.findtext('middle_name') == ''


@pytest.mark.parametrize('branch,family', FAMILIES)
def test_material_generated_prose_has_a_source_bound_repair_route(tmp_path, monkeypatch, branch, family):
    source = source_for(branch, family)
    source['design']['intervention_type'] = 'Other'
    source.setdefault('regulatory', {}).setdefault('prs', {})['study_type'] = 'Observational'
    expected = f"{'An' if branch == 'Ambispective' else 'A'} {branch.lower()} observational study"
    original_render = workflow.render_documents
    edits = []

    def broken_projection(root, revision, approved, model, **kwargs):
        faulty = copy.deepcopy(approved)
        faulty['design']['intervention_type'] = 'Device'
        return original_render(root, revision, faulty, model, **kwargs)
    monkeypatch.setattr(workflow, 'render_documents', broken_projection)

    def review(request):
        response = acceptance_verification(request)
        if request['task'] != 'clinical_content_verification': return response
        doc = Document(tmp_path / 'revisions' / request['revision_id'] / 'candidate/protocol.docx')
        passage = next((p.text for p in doc.paragraphs if 'observational study involving a medical device' in p.text), None)
        if passage:
            response['status'] = 'blocked'
            response['findings'] = [{
                'finding_id': 'invented-device', 'category': 'content', 'artifact': 'protocol',
                'check': 'source_supported', 'surface': 'document_shell', 'target_ids': ['protocol.shell'],
                'affected_passage': passage, 'issue': 'The cover invents a medical device.',
                'material': True, 'contradiction': True,
            }]
            next(r for r in response['section_assessments'] if r['section_id'] == 'protocol.shell')['status'] = 'failed'
        return response

    def draft(request):
        if not request.get('construction_repair'): return drafting.recorded_acceptance_response(request)
        edits.append(request)
        response = drafting.response_template(request, model_id='Controlled source-bound editor')
        response['section_results'][0]['paragraphs'] = [{
            'text': expected, 'evidence_refs': [f'source:{p}' for p in request['construction_repair']['source_paths']]}]
        return response

    result, revision = deliver(tmp_path, monkeypatch, source, review=review, draft=draft)
    assert result['status'] == 'passed', result
    assert len(edits) == 1 and not result['warnings']
    assert expected in [p.text for p in Document(tmp_path / 'output/protocol.docx').paragraphs]
    assert sorted(p.name for p in (tmp_path / 'output').iterdir()) == sorted(workflow.document_set(branch))


@pytest.mark.parametrize('branch,family', FAMILIES)
def test_generated_metadata_fields_recover_without_editing_approved_source(tmp_path, monkeypatch, branch, family):
    source = source_for(branch, family)
    source['meta']['date'] = '10 Oct 2026'
    original_render = workflow.render_documents
    edits = []

    def broken_projection(root, revision, approved, model, **kwargs):
        faulty = copy.deepcopy(approved)
        faulty['meta']['date'] = '01 Jan 1990'
        return original_render(root, revision, faulty, model, **kwargs)
    monkeypatch.setattr(workflow, 'render_documents', broken_projection)

    def review(request):
        response = acceptance_verification(request)
        if request['task'] != 'clinical_content_verification': return response
        doc = Document(tmp_path / 'revisions' / request['revision_id'] / 'candidate/protocol.docx')
        texts = [c.text for t in doc.tables for r in t.rows for c in r.cells]
        if '01 Jan 1990' in texts:
            response['status'] = 'blocked'
            response['findings'] = [{
                'finding_id': 'incorrect-control-date', 'category': 'content', 'artifact': 'protocol',
                'check': 'source_supported', 'surface': 'document_shell', 'target_ids': ['protocol.shell'],
                'source_path': 'meta.date', 'affected_passage': '01 Jan 1990',
                'issue': 'The document date differs from the approved date.', 'material': True, 'contradiction': True}]
            next(r for r in response['section_assessments'] if r['section_id'] == 'protocol.shell')['status'] = 'failed'
        return response

    def draft(request):
        if not request.get('construction_repair'): return drafting.recorded_acceptance_response(request)
        edits.append(request)
        response = drafting.response_template(request, model_id='Controlled metadata editor')
        response['field_edits'] = [{'path': f['path'], 'before': f['before'], 'after': '10 Oct 2026',
                                     'source_path': 'meta.date'} for f in request['construction_repair']['fields']]
        return response

    result, revision = deliver(tmp_path, monkeypatch, source, review=review, draft=draft)
    assert result['status'] == 'passed', result
    assert len(edits) == 1 and not result['warnings']
    doc = Document(tmp_path / 'output/protocol.docx')
    assert '10 Oct 2026' in [c.text for t in doc.tables for r in t.rows for c in r.cells]
    assert json.loads((revision / 'approved-reference.json').read_text())['meta']['date'] == '10 Oct 2026'


@pytest.mark.parametrize('branch,family', [case for case in FAMILIES if case[1] == 'Sterling'])
def test_material_display_quantity_is_corrected_from_approved_evidence(tmp_path, monkeypatch, branch, family):
    import rendering
    source = source_for(branch, family)
    source['population']['sample_size'] = '40 participants'
    original_fields = rendering.render_fields
    editors = []

    def faulty_fields(*args, **kwargs):
        fields = original_fields(*args, **kwargs)
        fields['sampleSize'] = '49 participants'
        return fields
    monkeypatch.setattr(rendering, 'render_fields', faulty_fields)

    def review(request):
        response = acceptance_verification(request)
        if request['task'] != 'clinical_content_verification': return response
        doc = Document(tmp_path / 'revisions' / request['revision_id'] / 'candidate/icf.docx')
        passage = next((p.text for p in doc.paragraphs if 'expected to include 49' in p.text), None)
        if passage:
            response['status'] = 'blocked'
            response['findings'] = [{'finding_id': 'wrong-display-count', 'category': 'content', 'artifact': 'icf',
                'target_ids': ['icf.shell'], 'check': 'source_supported', 'surface': 'document_shell',
                'affected_passage': passage, 'issue': 'The generated participant count is incorrect.',
                'material': True, 'contradiction': True}]
            next(r for r in response['section_assessments'] if r['section_id'] == 'icf.shell')['status'] = 'failed'
        return response

    def draft(request):
        if not request.get('construction_repair'): return drafting.recorded_acceptance_response(request)
        editors.append(request)
        response = drafting.response_template(request, model_id='Controlled count editor')
        response['section_results'][0]['paragraphs'] = [{'text': request['construction_repair']['before'].replace('49', '40'),
            'evidence_refs': ['source:population.sample_size']}]
        return response

    result, revision = deliver(tmp_path, monkeypatch, source, review=review, draft=draft)
    assert result['status'] == 'passed', result
    assert len(editors) == 1 and not result['warnings']
    assert any('expected to include 40 participants' in p.text for p in Document(tmp_path / 'output/icf.docx').paragraphs)


def test_metadata_repair_preserves_surrounding_native_formatting(tmp_path, monkeypatch):
    import rendering
    source = source_for('Prospective', 'Advarra')
    original_fields, original_document = rendering.render_fields, rendering.Document
    native_styles = []

    def paragraphs(container):
        yield from container.paragraphs
        for table in container.tables:
            for row in table.rows:
                for cell in row.cells:
                    yield from paragraphs(cell)

    def mixed_run_template(path):
        document = original_document(path)
        if isinstance(path, (str, Path)) and Path(path).name == 'prospective-protocol.template.docx':
            paragraph = next(p for p in paragraphs(document) if p.text.strip() == '{studyCordinatorName}')
            paragraph.clear()
            label = paragraph.add_run('Contact: ')
            label.bold, label.underline = True, True
            paragraph.add_run('{studyCordinatorName}')
        return document
    monkeypatch.setattr(rendering, 'Document', mixed_run_template)

    def faulty_fields(*args, **kwargs):
        fields = original_fields(*args, **kwargs)
        fields['studyCordinatorName'] = 'Incorrect coordinator'
        return fields
    monkeypatch.setattr(rendering, 'render_fields', faulty_fields)

    def review(request):
        response = acceptance_verification(request)
        if request['task'] != 'clinical_content_verification': return response
        doc = Document(tmp_path / 'revisions' / request['revision_id'] / 'candidate/protocol.docx')
        paragraph = next(p for p in paragraphs(doc) if p.text.startswith('Contact: '))
        native_styles.append([(r.text, r.bold, r.underline) for r in paragraph.runs if r.text == 'Contact: '])
        if 'Incorrect coordinator' in paragraph.text:
            response['status'] = 'blocked'
            response['findings'] = [{'finding_id': 'wrong-coordinator', 'category': 'content', 'artifact': 'protocol',
                'target_ids': ['protocol.shell'], 'check': 'source_supported', 'surface': 'document_shell',
                'source_path': 'parties.study_coordinator.name', 'affected_passage': 'Incorrect coordinator',
                'issue': 'The coordinator name is incorrect.', 'material': True, 'contradiction': True}]
            next(r for r in response['section_assessments'] if r['section_id'] == 'protocol.shell')['status'] = 'failed'
        return response

    def draft(request):
        if not request.get('construction_repair'): return drafting.recorded_acceptance_response(request)
        response = drafting.response_template(request, model_id='Controlled contact editor')
        response['field_edits'] = [{**f, 'after': source['parties']['study_coordinator']['name'],
            'source_path': 'parties.study_coordinator.name'} for f in request['construction_repair']['fields']]
        return response

    result, revision = deliver(tmp_path, monkeypatch, source, review=review, draft=draft)
    assert result['status'] == 'passed', result
    assert native_styles == [[('Contact: ', True, True)], [('Contact: ', True, True)]]
    from zipfile import ZipFile
    with ZipFile(tmp_path / 'output/protocol.docx') as archive:
        for name in archive.namelist():
            if name.startswith('word/') and name.endswith('.xml'):
                assert all('{urn:pulpsense:construction-recovery}fields' not in node.attrib
                           for node in ET.fromstring(archive.read(name)).iter())
