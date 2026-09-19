import copy
import json
import re
import zipfile
from pathlib import Path

import pytest
from docx import Document
from docx.oxml.ns import qn

import contracts
import quality
import rendering
from drafting import create_drafting_request, recorded_acceptance_response, validate_response


ROOT = Path(__file__).resolve().parents[1]
INCOMPLETE_LENS_SOURCE = ROOT / "tests/fixtures/sterling-lens-incomplete-source.json"
RESULTS_DEFAULT = (
    "You will not routinely receive the results of testing performed for this study. "
    "If a study test identifies a new medical condition or concern, the study doctor "
    "will inform you and recommend appropriate follow-up with your treating doctor. "
    "The finding may affect whether you can continue in the study."
)


def fixture(*, family="Sterling"):
    value = json.loads(
        (ROOT / "tests/fixtures/prospective-acceptance-source.json").read_text(
            encoding="utf-8"
        )
    )
    value["meta"]["icf_template"] = family
    value["study"]["title"] = "PureSee and Odyssey Visual Outcomes Study"
    value["study"]["background"] = (
        "Available reports describe visual outcomes associated with PureSee and Odyssey "
        "intraocular lenses, but evidence about their combined use is limited."
    )
    value["study"]["hypothesis"] = (
        "PureSee and Odyssey intraocular lenses are associated with useful visual outcomes, "
        "patient satisfaction, and limited visual disturbances at Month 3."
    )
    value["objectives"]["primary"] = [
        "Describe Month 3 visual outcomes, patient satisfaction, and visual disturbances associated with PureSee and Odyssey intraocular lenses."
    ]
    value["endpoints"]["primary"] = [
        {"label": "Month 3 binocular intermediate visual acuity", "time_point": "Month 3"}
    ]
    value["design"].update(
        {
            "study_design": "Prospective, single-center observational study.",
            "assignment_method": (
                "PureSee and Odyssey lens selections are made independently as part of "
                "routine clinical care before study participation; the research protocol "
                "does not assign a lens."
            ),
            "intervention_name": "PureSee and Odyssey intraocular lenses",
            "intervention_type": "Medical device — intraocular lenses",
        }
    )
    value["regulatory"]["prs"]["study_type"] = "Observational"
    value["population"]["sample_size"] = "40 participants"
    value["population"]["sample_justification"] = (
        "A feasibility review of the site's eligible cataract-surgery volume supports "
        "enrollment of 40 participants, with at least 35 evaluable participants expected "
        "after the source-supported allowance of five nonevaluable or withdrawn participants."
    )
    value["procedures"]["assessments"] = (
        "Preoperative screening with demographics, medical history, and an eye examination; "
        "one operative visit for each eye; Month 3 visual acuity, patient-satisfaction "
        "questionnaire, visual-disturbance questionnaire, and safety assessments."
    )
    value["procedures"]["visit_schedule"] = [
        {"visit": "Preoperative screening", "timing": "Preoperative", "procedures": ["Consent", "Eye examination"]},
        {"visit": "Operative visit for each eye", "timing": "One operative visit per eye", "procedures": ["Routine cataract surgery"]},
        {"visit": "Month 3 postoperative", "timing": "3 months postoperatively", "procedures": ["Visual outcomes", "Questionnaires"]},
    ]
    value["procedures"]["visit_schedule_table"] = [
        {"visitNumber": "1", "visitName": "Preoperative screening", "visitWindow": "Preoperative", "CRFnumber": "PRE"},
        {"visitNumber": "2", "visitName": "Operative visit for each eye", "visitWindow": "One operative visit per eye", "CRFnumber": "OP"},
        {"visitNumber": "3", "visitName": "Month 3 postoperative", "visitWindow": "3 months postoperatively", "CRFnumber": "M3"},
    ]
    value["risks_benefits"].update(
        {
            "risks": (
                "Foreseeable risks include the approved risks of cataract surgery; "
                "intraocular-lens and device-related risks; postoperative complications; "
                "visual symptoms or disturbances; discomfort from study eye examinations; "
                "and loss of confidentiality."
            ),
            "benefits": (
                "Direct benefit is not guaranteed. The study may improve understanding of "
                "visual outcomes, patient satisfaction, visual disturbances, and outcomes "
                "associated with the PureSee and Odyssey intraocular lenses."
            ),
            "costs": (
                "The sponsor pays for study-only visual testing and questionnaires. Routine "
                "cataract surgery, intraocular lenses selected for ordinary care, and ordinary-care "
                "visits may be billed to the participant or insurer according to usual coverage."
            ),
            "alternatives": (
                "Participation is optional. A person may decline the research and continue or "
                "discuss ordinary cataract care with the treating doctor."
            ),
            "injury_handling": (
                "The study doctor will arrange or provide appropriate care for a research-related "
                "injury. No additional compensation is available. The sponsor is financially "
                "responsible for reasonable costs of study-related injury care not covered by the participant's insurer."
            ),
        }
    )
    value.setdefault("confidentiality", {})["authorized_recipients"] = []
    return value


def model():
    return {
        "protocol": [
            {"section_id": "introduction", "paragraphs": [{"text": "Available reports describe visual outcomes associated with the studied lenses without identifying a specific evidence design or scale."}], "lists": []},
            {"section_id": "sample-size", "paragraphs": [{"text": "A site-feasibility review supports enrollment of 40 participants and the approved allowance for five nonevaluable or withdrawn participants supports the target of at least 35 evaluable participants."}], "lists": []},
            {"section_id": "risks-benefits.risks", "paragraphs": [{"text": "Foreseeable risks include cataract-surgery, intraocular-lens, postoperative, visual-disturbance, study-examination, and confidentiality risks described in the approved source."}], "lists": []},
            {"section_id": "risks-benefits.benefits", "paragraphs": [{"text": "Direct benefit is not guaranteed. The study may improve knowledge about visual outcomes, satisfaction, visual disturbances, and outcomes associated with PureSee and Odyssey intraocular lenses."}], "lists": []},
            {"section_id": "financial-injury", "paragraphs": [{"text": "The sponsor pays for study-only testing; routine cataract care may be billed to the participant or insurer. The study doctor arranges care for research-related injury, no additional compensation is available, and the sponsor has the source-specified financial responsibility."}], "lists": []},
            {"section_id": "endpoint-criteria.completion", "paragraphs": [{"text": "A participant completes preoperative screening, one operative visit for each eye, and the Month 3 postoperative visit."}], "lists": []},
        ],
        "prs": {},
        "icf": {
            "icf.key-information-summary": {"paragraphs": [
                {"text": "The study examines visual outcomes associated with PureSee and Odyssey intraocular lenses."},
                {"text": "About 40 participants will complete screening, operative care, and Month 3 research assessments."},
                {"text": "The foreseeable study-specific risks include the approved surgery, lens, postoperative, visual-symptom, examination, and privacy risks."},
                {"text": "Direct benefit is not guaranteed; the research may improve knowledge about visual outcomes, satisfaction, and visual disturbances."},
                {"text": "Taking part is voluntary; you may decline and continue or discuss ordinary cataract care with your treating doctor."},
            ], "lists": []},
            "icf.background": {"paragraphs": [{"text": "Available reports describe visual outcomes with PureSee and Odyssey lenses, but evidence about combined use is limited."}], "lists": []},
            "icf.study-purpose": {"paragraphs": [{"text": "The purpose is to describe binocular intermediate visual acuity at Month 3, participant satisfaction, and reported visual disturbances. The hypothesis is that the implanted lens combination will be associated with useful vision, satisfaction, and limited disturbances at Month 3. The primary endpoint is Month 3 binocular intermediate visual acuity."}], "lists": []},
            "icf.duration": {"paragraphs": [{"text": "Your participation continues through the Month 3 postoperative visit."}], "lists": []},
            "icf.procedures": {"paragraphs": [
                {"text": "You may be eligible if you are an adult with the target condition and can complete follow-up. You must have gone at least 90 days before screening without participating in another study."},
                {"text": "You will complete preoperative screening with consent, demographics, medical history, and an eye examination; ordinary cataract surgery with one operative visit for each eye; and Month 3 visual-acuity, questionnaire, visual-disturbance, and safety assessments. PureSee and Odyssey lens selections occur during routine clinical care before research participation and are not assigned by the research protocol."},
            ], "lists": []},
            "icf.risks": {"paragraphs": [{"text": "Foreseeable risks include the approved risks of cataract surgery, the intraocular lenses, postoperative complications, visual symptoms or disturbances, study eye examinations, and loss of confidentiality."}], "lists": []},
            "icf.benefits": {"paragraphs": [{"text": "You may not benefit directly. The study may improve understanding of visual outcomes, patient satisfaction, visual disturbances, and outcomes associated with PureSee and Odyssey intraocular lenses."}], "lists": []},
            "icf.payment": {"paragraphs": [{"text": "No compensation or reimbursement is planned."}], "lists": []},
            "icf.costs": {"paragraphs": [{"text": "The sponsor pays for study-only testing. Routine cataract surgery, lenses selected for ordinary care, and ordinary-care visits may be billed to you or your insurer."}], "lists": []},
            "icf.alternatives": {"paragraphs": [{"text": "Participation is optional. You may decline the research and continue or discuss ordinary cataract care with your treating doctor."}], "lists": []},
            "icf.privacy": {"paragraphs": [], "lists": []},
            "icf.injury": {"paragraphs": [{"text": "If you believe a study activity caused an injury or medical problem, contact the study doctor promptly and seek emergency care when needed. The study doctor will arrange appropriate care. No additional compensation is available. The sponsor is financially responsible for reasonable costs of study-related injury care not covered by your insurer. Signing this form does not waive your legal rights."}], "lists": []},
        },
    }


def normalized(text):
    return re.sub(r"\s+", " ", text).strip()


def visible(document):
    values = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            values.extend(cell.text for cell in row.cells)
    return normalized(" ".join(values))


def section(document, heading):
    start = next(i for i, p in enumerate(document.paragraphs) if normalized(p.text) == heading)
    values = []
    for paragraph in document.paragraphs[start + 1 :]:
        if paragraph.text.strip() and paragraph.style.name.casefold().startswith("heading"):
            break
        values.append(paragraph)
    return values


def real_bullet(paragraph):
    properties = paragraph._p.pPr
    return properties is not None and properties.find(qn("w:numPr")) is not None


def render_pair(tmp_path, reference=None, draft=None):
    reference = reference or fixture()
    report = rendering.render_documents(
        ROOT, tmp_path, reference, draft or model(), artifact_names={"protocol", "icf"}
    )
    assert report["status"] == "passed", report
    return (
        tmp_path / "candidate/protocol.docx",
        tmp_path / "candidate/icf.docx",
    )


def test_incomplete_lens_source_has_actionable_shared_release_blockers():
    reference = json.loads(INCOMPLETE_LENS_SOURCE.read_text(encoding="utf-8"))
    findings = contracts.release_source_findings(reference)
    by_field = {item["field"]: item for item in findings}
    assert {
        "study_design_classification",
        "study_specific_foreseeable_risks",
        "participant_cost_allocation",
        "research_injury_responsibility",
        "sample_size_justification",
    } <= set(by_field)
    assert all(item["category"] == "source-evidence" for item in by_field.values())
    assert "separately approved source correction" in by_field["study_design_classification"]["required"]
    assert set(by_field["study_specific_foreseeable_risks"]["affected_artifacts"]) == {"icf.docx", "protocol.docx"}
    assert set(by_field["participant_cost_allocation"]["affected_artifacts"]) == {"icf.docx", "protocol.docx"}
    assert set(by_field["research_injury_responsibility"]["affected_artifacts"]) == {"icf.docx", "protocol.docx"}


def test_complete_source_passes_release_contract_and_supported_classifications_are_consistent():
    observational = fixture()
    assert contracts.release_source_findings(observational) == []
    assert contracts.normalized_study_classification(observational) == "Observational"
    assert rendering.protocol_subtitle(observational) == "A prospective observational study of intraocular lenses"

    interventional = fixture()
    interventional["design"]["assignment_method"] = "PureSee and Odyssey lens assignments are determined by the research protocol."
    interventional["design"]["study_design"] = "Prospective, single-center interventional medical device study."
    interventional["regulatory"]["prs"]["study_type"] = "Interventional"
    assert contracts.release_source_findings(interventional) == []
    assert contracts.normalized_study_classification(interventional) == "Interventional"
    assert rendering.protocol_subtitle(interventional) == "A prospective interventional medical device study"

    ambiguous = fixture()
    ambiguous["design"].pop("assignment_method")
    assert rendering.protocol_subtitle(ambiguous) == "A prospective study"


def test_protocol_subtitle_preserves_unrelated_non_lens_behavior():
    reference = json.loads(
        (ROOT / "tests/fixtures/prospective-acceptance-source.json").read_text(encoding="utf-8")
    )
    reference["meta"]["icf_template"] = "Advarra"
    assert rendering.protocol_subtitle(reference) == "A prospective observational device study"


def test_classification_mismatch_blocks_without_silent_reclassification():
    reference = fixture()
    reference["regulatory"]["prs"]["study_type"] = "Interventional"
    findings = contracts.release_source_findings(reference)
    assert any(item["field"] == "study_design_classification" for item in findings)
    assert contracts.normalized_study_classification(reference) is None


@pytest.mark.parametrize(
    ("assignment", "declared", "expected"),
    [
        ("This study is not randomized.", "Observational", None),
        ("The investigator does not assign the intraocular lens.", "Observational", None),
        ("The lens was not selected as part of routine care.", "Observational", None),
        ("LENS SELECTION, DID NOT OCCUR INDEPENDENTLY OF THE RESEARCH!", "Observational", None),
        ("The lens was not selected before enrollment.", "Observational", None),
        ("This is not an ordinary-care lens selection.", "Observational", None),
        ("The investigator does not independently select the lens.", "Observational", None),
        ("THE STUDY PROTOCOL DOESN'T ASSIGN A LENS!", "Observational", None),
        ("The lens isn't selected as part of routine care before enrollment.", "Observational", None),
        ("The research protocol determines follow-up timing.", "Observational", None),
        ("Participants are randomized for questionnaire order.", "Observational", None),
        ("Lens selection occurs during routine care and the investigator independently reviews study data.", "Observational", None),
        ("The surgeon independently reviews eligibility while lens selection occurs during routine care.", "Observational", None),
        (
            "Lens selection occurs as part of routine clinical care before enrollment; the research protocol does not assign a lens.",
            "Observational",
            "Observational",
        ),
        (
            "Participants are assigned to receive PureSee or Odyssey by the study protocol.",
            "Interventional",
            "Interventional",
        ),
        (
            "Lens selection occurs during routine care, but participants are assigned to PureSee or Odyssey by the study protocol.",
            "Observational",
            None,
        ),
    ],
)
def test_lens_assignment_classification_is_polarity_aware(assignment, declared, expected):
    reference = fixture()
    reference["design"]["assignment_method"] = assignment
    reference["regulatory"]["prs"]["study_type"] = declared
    if declared == "Interventional":
        reference["design"]["study_design"] = (
            "Prospective, single-center interventional medical device study."
        )
    assert contracts.normalized_study_classification(reference) == expected
    findings = contracts.release_source_findings(reference)
    assert any(item["field"] == "study_design_classification" for item in findings) is (expected is None)


@pytest.mark.parametrize(
    ("assignment", "expected"),
    [
        (
            "The lens is selected independently during routine clinical care before study participation",
            "Observational",
        ),
        (
            "THE LENS IS SELECTED INDEPENDENTLY DURING ROUTINE CLINICAL CARE BEFORE STUDY PARTICIPATION!",
            "Observational",
        ),
        (
            "Before enrollment, the surgeon independently selects the lens as part of ordinary clinical care.",
            "Observational",
        ),
        (
            "Lens selection is independently completed in routine care before research participation.",
            "Observational",
        ),
        (
            "Lens selection is determined by the research protocol",
            "Interventional",
        ),
        (
            "LENS SELECTION IS DETERMINED BY THE RESEARCH PROTOCOL!",
            "Interventional",
        ),
        (
            "The study protocol determines lens selection before enrollment.",
            "Interventional",
        ),
        (
            "The intraocular lens is assigned by the study protocol.",
            "Interventional",
        ),
    ],
)
def test_affirmative_lens_assignment_forms_are_classified(assignment, expected):
    reference = fixture()
    reference["design"]["assignment_method"] = assignment
    reference["regulatory"]["prs"]["study_type"] = expected
    reference["design"]["study_design"] = (
        f"Prospective, single-center {expected.casefold()} medical device study."
    )
    assert contracts.normalized_study_classification(reference) == expected
    assert not any(
        item["field"] == "study_design_classification"
        for item in contracts.release_source_findings(reference)
    )


@pytest.mark.parametrize(
    "assignment",
    [
        "The lens is not selected independently during routine clinical care before study participation.",
        "THE LENS ISN'T SELECTED INDEPENDENTLY DURING ROUTINE CLINICAL CARE BEFORE STUDY PARTICIPATION!",
        "Before enrollment, the surgeon does not independently select the lens as part of ordinary clinical care.",
        "Lens selection is not independently completed in routine care before research participation.",
        "Lens selection is not determined by the research protocol.",
        "LENS SELECTION ISN'T DETERMINED BY THE RESEARCH PROTOCOL!",
        "The study protocol does not determine lens selection before enrollment.",
        "The intraocular lens is not assigned by the study protocol.",
        "The lens cannot be assigned by the study protocol.",
        "The lens hasn't been assigned by the study protocol.",
        "The lens mustn't be assigned by the study protocol.",
        "Nor is the lens assigned by the study protocol.",
        "It can't be true that lens selection is determined by the research protocol.",
        (
            "It is not the case, despite the preliminary planning discussion and review, "
            "that the intraocular lens is assigned by the study protocol."
        ),
        (
            "It is not the case, after the clinical team completed a detailed review of the "
            "screening schedule, operative workflow, follow-up plan, questionnaire sequence, "
            "and data-management procedures, that lens selection is determined by the research protocol."
        ),
        (
            "It is not true, after review of the study procedures, that the lens is determined "
            "by the research protocol."
        ),
        (
            "It is not the case, after the clinical team reviewed the screening schedule, "
            "operative workflow, and follow-up plan, that the lens is selected independently "
            "during routine clinical care before study participation."
        ),
    ],
)
def test_negated_affirmative_lens_assignment_forms_do_not_classify(assignment):
    reference = fixture()
    reference["design"]["assignment_method"] = assignment
    reference["design"]["study_design"] = "Prospective single-center lens study."
    reference["regulatory"]["prs"].pop("study_type")
    assert contracts.normalized_study_classification(reference) is None
    assert any(
        item["field"] == "study_design_classification"
        for item in contracts.release_source_findings(reference)
    )


@pytest.mark.parametrize(
    "assignment",
    [
        (
            "It is not the case—after the clinical team reviewed the screening schedule; "
            "the operative workflow; the follow-up plan; the questionnaire sequence; and "
            "the data-management procedures—that lens selection is determined by the "
            "research protocol."
        ),
        (
            "It is not true—after review of the screening schedule; operative workflow; "
            "and follow-up plan—that the intraocular lens is assigned by the study protocol."
        ),
        (
            "It isn't true—after review of the screening schedule; operative workflow; "
            "and follow-up plan—that the lens is selected independently during routine "
            "clinical care before study participation."
        ),
        (
            "It is not the case—after the team confirmed that screening was complete; "
            "the workflow was reviewed; and follow-up was approved—that lens selection "
            "is determined by the research protocol."
        ),
        (
            "It is not the case that, after review of the screening schedule; operative "
            "workflow; and follow-up plan, lens selection is determined by the research protocol."
        ),
        (
            "It wasn't true—after review of the screening schedule; operative workflow; "
            "and follow-up plan—that the lens is assigned by the study protocol."
        ),
        (
            "It is not the case, after the team reported, in a note, that screening was "
            "complete; and confirmed follow-up, that lens selection is determined by the "
            "research protocol."
        ),
        (
            "It is not true that questionnaire order is randomized; it wasn't true—after "
            "review of the screening schedule; operative workflow; and follow-up plan—that "
            "the lens is assigned by the study protocol."
        ),
        (
            "It is not true that the note says 'screening; lens selection is determined by "
            "the research protocol'."
        ),
        (
            "It was not the case that the note said \"screening; the lens is assigned by "
            "the study protocol.\""
        ),
        (
            "It is not true that questionnaire order is randomized, and it isn't true; "
            "lens selection is determined by the research protocol."
        ),
        (
            "It is not the case after the team said that screening was complete; follow-up "
            "was approved that lens selection is determined by the research protocol."
        ),
        (
            "It is not true that questionnaire order is randomized and it isn't true because "
            "the note says that follow-up was approved; lens selection is determined by the "
            "research protocol."
        ),
    ],
)
def test_governing_negation_preserves_internal_semicolons(assignment):
    reference = fixture()
    reference["design"]["assignment_method"] = assignment
    reference["design"]["study_design"] = "Prospective single-center lens study."
    reference["regulatory"]["prs"].pop("study_type")
    assert contracts.normalized_study_classification(reference) is None
    assert any(
        item["field"] == "study_design_classification"
        for item in contracts.release_source_findings(reference)
    )


def test_ambiguous_governing_negation_punctuation_blocks_classification():
    reference = fixture()
    reference["design"]["assignment_method"] = (
        "It is not the case; lens selection is determined by the research protocol."
    )
    reference["design"]["study_design"] = "Prospective single-center lens study."
    reference["regulatory"]["prs"].pop("study_type")
    assert contracts.normalized_study_classification(reference) is None
    assert any(
        item["field"] == "study_design_classification"
        for item in contracts.release_source_findings(reference)
    )


@pytest.mark.parametrize(
    ("assignment", "expected"),
    [
        (
            "Questionnaire order is not randomized, and lens selection is determined by the research protocol.",
            "Interventional",
        ),
        (
            "Questionnaire order isn't randomized, yet lens selection is determined by the research protocol.",
            "Interventional",
        ),
        (
            "Questionnaire order is not randomized; lens selection is determined by the research protocol.",
            "Interventional",
        ),
        (
            "Questionnaire order is not randomized and lens selection is determined by the research protocol.",
            "Interventional",
        ),
        (
            "Questionnaire order is not randomized and the study protocol determines lens selection.",
            "Interventional",
        ),
        (
            "Questionnaire order is not randomized and the surgeon independently selects the lens "
            "during routine clinical care before study participation.",
            "Observational",
        ),
        (
            "It is not the case that questionnaire order is randomized; "
            "lens selection is determined by the research protocol.",
            "Interventional",
        ),
        (
            "It is not true that questionnaire order is randomized, and lens selection "
            "is determined by the research protocol.",
            "Interventional",
        ),
        (
            "It is not true that questionnaire order is randomized and lens selection "
            "is determined by the research protocol.",
            "Interventional",
        ),
        (
            "It isn’t true that questionnaire order is randomized; lens selection is "
            "determined by the research protocol.",
            "Interventional",
        ),
        (
            "It wasn’t true that questionnaire order was randomized and lens selection is "
            "determined by the research protocol.",
            "Interventional",
        ),
        (
            "Follow-up timing is not assigned by the study, and the lens is selected independently "
            "during routine clinical care before study participation.",
            "Observational",
        ),
    ],
)
def test_unrelated_negative_clause_does_not_negate_separate_affirmative_assignment(
    assignment, expected
):
    reference = fixture()
    reference["design"]["assignment_method"] = assignment
    reference["regulatory"]["prs"]["study_type"] = expected
    reference["design"]["study_design"] = (
        f"Prospective, single-center {expected.casefold()} medical device study."
    )
    assert contracts.normalized_study_classification(reference) == expected
    assert not any(
        item["field"] == "study_design_classification"
        for item in contracts.release_source_findings(reference)
    )


def test_mixed_affirmative_lens_assignment_forms_block_classification():
    reference = fixture()
    reference["design"]["assignment_method"] = (
        "The lens is selected independently during routine clinical care before study participation; "
        "however, lens selection is determined by the research protocol."
    )
    assert contracts.normalized_study_classification(reference) is None
    assert any(
        item["field"] == "study_design_classification"
        for item in contracts.release_source_findings(reference)
    )


@pytest.mark.parametrize(
    ("costs", "valid"),
    [
        ("The sponsor pays study-only testing. Ordinary care may be billed to the participant or insurer.", True),
        ("The sponsor will not pay these costs; the participant or insurer is responsible.", True),
        ("The sponsor is not responsible; participant insurance terms are unresolved.", False),
        ("Costs are not specified.", False),
        ("The study team will explain costs later.", False),
        ("", False),
    ],
)
def test_cost_allocation_requires_complete_polarity_aware_policy(costs, valid):
    reference = fixture()
    reference["risks_benefits"]["costs"] = costs
    fields = {item["field"] for item in contracts.release_source_findings(reference)}
    assert ("participant_cost_allocation" not in fields) is valid


@pytest.mark.parametrize(
    ("injury", "valid"),
    [
        ("The study doctor will arrange injury care. No compensation is available. The participant or insurer is responsible for its cost.", True),
        ("The participant must obtain injury care. The sponsor will not pay and no compensation is available; the participant or insurer is responsible for costs.", True),
        ("The investigator will not provide care. No compensation is available. The sponsor is not financially responsible.", False),
        ("The investigator will not provide care. No compensation is available. The participant is responsible for costs.", False),
        ("Injury compensation is unknown.", False),
        ("The study team will explain injury terms later.", False),
        ("", False),
    ],
)
def test_injury_policy_requires_care_compensation_and_cost_responsibility(injury, valid):
    reference = fixture()
    reference["risks_benefits"]["injury_handling"] = injury
    fields = {item["field"] for item in contracts.release_source_findings(reference)}
    assert ("research_injury_responsibility" not in fields) is valid


@pytest.mark.parametrize(
    ("rationale", "valid"),
    [
        ("The sample size is sufficient.", False),
        ("Forty participants will be enrolled.", False),
        ("The sample size was selected for the study.", False),
        ("A site feasibility review of eligible cataract-surgery volume supports recruitment of 40 participants.", True),
        ("Forty participants provide 80% power for the approved effect-size and variance assumptions.", True),
        ("This exploratory pilot will characterize outcome variability for later planning.", True),
        ("Forty participants are planned to account for dropout from 35 evaluable participants.", False),
        ("Historical site retention data support a 5-participant attrition allowance from 40 enrolled to 35 evaluable.", True),
    ],
)
def test_sample_size_requires_substantive_approved_basis(rationale, valid):
    reference = fixture()
    reference["population"]["sample_justification"] = rationale
    fields = {item["field"] for item in contracts.release_source_findings(reference)}
    assert ("sample_size_justification" not in fields) is valid


def test_sterling_icf_semantics_lists_and_controlled_enrollment_placement(tmp_path):
    _protocol_path, icf_path = render_pair(tmp_path)
    document = Document(icf_path)
    text = visible(document)
    privacy = section(document, "CONFIDENTIALITY AUTHORIZATION TO COLLECT, USE DISCLOSE YOUR MEDICAL INFORMATION")
    privacy_text = normalized(" ".join(item.text for item in privacy))

    consequence = "cannot begin or continue in this research study"
    assert privacy_text.casefold().count(consequence) == 1
    assert privacy_text.count("protected health information (PHI)") == 1
    assert privacy_text.find("protected health information (PHI)") < privacy_text.find("use my PHI")
    recipient_bullets = [normalized(item.text) for item in privacy if real_bullet(item)]
    assert any("study doctor" in item.casefold() and "study team" in item.casefold() for item in recipient_bullets)
    assert any("sponsor" in item.casefold() and "authorized representatives" in item.casefold() for item in recipient_bullets)
    assert any("institutional review board" in item.casefold() or "ethics committee" in item.casefold() for item in recipient_bullets)
    assert any("regulatory or government authorities" in item.casefold() for item in recipient_bullets)
    assert "sponsor" in privacy_text.casefold() and "regulatory or government authorities" in privacy_text.casefold()

    for heading in ("VOLUNTARY PARTICIPATION/WITHDRAWAL",):
        assert len([item for item in section(document, heading) if real_bullet(item)]) >= 4
    assert len(recipient_bullets) >= 4
    phi_bullets = [item for item in privacy if real_bullet(item) and item not in []]
    assert len(phi_bullets) >= 8

    duration = normalized(" ".join(item.text for item in section(document, "DURATION")))
    assert "40" not in duration
    assert text.count("40 participants") == 2
    purpose = normalized(" ".join(item.text for item in section(document, "PURPOSE")))
    assert "40 participants" in purpose
    alternatives = normalized(" ".join(item.text for item in section(document, "ALTERNATIVE TREATMENTS")))
    assert "optional" in alternatives.casefold()
    assert "ordinary cataract care" in alternatives.casefold()
    benefits = normalized(" ".join(item.text for item in section(document, "POTENTIAL BENEFITS")))
    assert "directly" in benefits.casefold() or "direct benefit" in benefits.casefold()
    for phrase in ("visual outcomes", "patient satisfaction", "visual disturbances", "PureSee", "Odyssey"):
        assert phrase.casefold() in benefits.casefold()
    assert "condition, intervention, or procedures" not in benefits.casefold()

    headings = [normalized(p.text) for p in document.paragraphs if p.text.strip() and p.style.name.casefold().startswith("heading")]
    assert headings.index("STUDY RESULTS") == headings.index("PROCEDURES") + 1
    assert RESULTS_DEFAULT in text


def test_protocol_subtitle_visit_language_claims_benefits_and_toc(tmp_path):
    protocol_path, _icf_path = render_pair(tmp_path)
    document = Document(protocol_path)
    text = visible(document)
    assert "A prospective observational study of intraocular lenses" in text
    assert "medical device — intraocular lenses. study" not in text
    assert "Preoperative screening at the Preoperative time point" not in text
    assert "with One operative visit per eye" not in text
    assert "Preoperative screening" in text
    assert "One operative visit for each eye" in text
    assert "large randomized controlled trial" not in text.casefold()
    assert "35 evaluable participants are sufficient" not in text.casefold()
    visit_table = next(
        table for table in document.tables
        if table.rows and normalized(table.rows[0].cells[0].text) == "Visit Number"
    )
    assert [normalized(row.cells[1].text) for row in visit_table.rows[1:]] == [
        "Preoperative screening", "Operative visit for each eye", "Month 3 postoperative",
    ]
    benefit_text = normalized(" ".join(item.text for item in section(document, "19.2. Summary of benefits")))
    assert "compensation" not in benefit_text.casefold() and "reimbursement" not in benefit_text.casefold()
    assert "direct benefit is not guaranteed" in benefit_text.casefold()
    with zipfile.ZipFile(protocol_path) as package:
        document_xml = package.read("word/document.xml")
        header_xml = package.read("word/header1.xml")
    assert b"TOC" in document_xml
    assert b"PAGE" in header_xml and b"NUMPAGES" in header_xml


def test_sterling_headings_drop_hidden_client_smart_tag_fragments(tmp_path):
    _protocol_path, icf_path = render_pair(tmp_path)
    document = Document(icf_path)
    headings = [normalized(p.text) for p in document.paragraphs if p.style.name.casefold().startswith("heading")]
    assert "POTENTIAL RISKS, EFFECTS, DISCOMFORTS, INCONVENIENCES" in headings
    assert "INFORMATION" in headings
    with zipfile.ZipFile(icf_path) as package:
        document_xml = package.read("word/document.xml")
    assert b"INCONVENIENCESSIDE" not in document_xml
    assert b">NEW<" not in document_xml


def test_unsupported_specific_evidence_claim_is_a_release_finding(tmp_path):
    protocol_path, _icf_path = render_pair(tmp_path)
    document = Document(protocol_path)
    document.add_paragraph("A large randomized controlled trial proved superiority.")
    document.save(protocol_path)
    findings = quality.deterministic_content_check(tmp_path, fixture())
    assert any(item.get("field") == "unsupported_evidence_claim" for item in findings)


def test_unrelated_citation_does_not_authorize_evidence_design_claim(tmp_path):
    reference = fixture()
    reference["references"] = [{
        "citation": "Unrelated registry methods report.",
        "claim": "A registry report describes participant enrollment.",
    }]
    draft = model()
    draft["protocol"][0]["paragraphs"][0]["text"] = (
        "A large randomized controlled trial proved superiority."
    )
    protocol_path, _icf_path = render_pair(tmp_path, reference, draft)
    findings = quality.deterministic_content_check(tmp_path, reference)
    assert any(item.get("field") == "unsupported_evidence_claim" for item in findings)


def test_claim_linked_to_specific_citation_authorizes_design_assertion(tmp_path):
    reference = fixture()
    citation = "Smith et al. 2025. PureSee randomized controlled trial."
    reference["study"]["background_evidence"] = [{
        "claim": "A large randomized controlled trial proved superiority.",
        "citation": citation,
    }]
    draft = model()
    draft["protocol"][0]["paragraphs"][0]["text"] = (
        "A large randomized controlled trial proved superiority."
    )
    protocol_path, _icf_path = render_pair(tmp_path, reference, draft)
    visible_protocol = visible(Document(protocol_path))
    assert f"A large randomized controlled trial proved superiority ({citation})" in visible_protocol
    assert "REFERENCES" in visible_protocol
    assert citation in visible_protocol
    findings = quality.deterministic_content_check(tmp_path, reference)
    assert not any(item.get("field") == "unsupported_evidence_claim" for item in findings)


def test_sterling_privacy_deduplicates_by_module_identity_and_preserves_supplements(tmp_path):
    draft = model()
    unique = "The approved source requires records to be retained for seven years."
    duplicate_consequence = (
        "Your authorization is voluntary. Refusing or revoking authorization will not affect "
        "your ordinary medical care or benefits, but you cannot begin or continue in this "
        "research study without the authorization required for the research."
    )
    draft["icf"]["icf.privacy"] = {
        "paragraphs": [
            {"module_id": "source.retention", "text": unique},
            {"module_id": "source.retention", "text": unique},
            {"module_id": "sterling.privacy.authorization", "text": duplicate_consequence},
        ],
        "lists": [{
            "module_id": "source.additional-recipient",
            "items": ["A source-authorized data custodian."],
        }],
    }
    reference = fixture()
    report = rendering.render_documents(
        ROOT, tmp_path, reference, draft, artifact_names={"icf"}
    )
    assert report["status"] == "passed", report
    document = Document(tmp_path / "candidate/icf.docx")
    privacy = section(
        document,
        "CONFIDENTIALITY AUTHORIZATION TO COLLECT, USE DISCLOSE YOUR MEDICAL INFORMATION",
    )
    privacy_text = normalized(" ".join(item.text for item in privacy))
    assert privacy_text.count(unique) == 1
    assert privacy_text.casefold().count("cannot begin or continue in this research study") == 1
    assert privacy_text.count("protected health information (PHI)") == 1
    for required in (
        "does not expire unless you revoke it in writing",
        "What happens to information about me after the study is over",
        "ClinicalTrials.gov",
        "The study doctor and authorized study team",
    ):
        assert required in privacy_text
    custodian = next(item for item in privacy if item.text == "A source-authorized data custodian.")
    assert real_bullet(custodian)
    assert any(item.text.strip() for item in privacy)


def test_privacy_module_canonicalization_deduplicates_whitespace_only_variants():
    section = {
        "paragraphs": [
            {"module_id": "source.retention", "text": "Records are retained for seven years."},
            {"module_id": "source.retention", "text": "  Records   are retained for seven years.  "},
            {"module_id": "source.access", "text": "Authorized staff may review records."},
        ],
        "lists": [],
    }
    normalized_section, findings = contracts.normalize_privacy_modules(section)
    assert findings == []
    assert [item["module_id"] for item in normalized_section["paragraphs"]] == [
        "source.retention", "source.access",
    ]


def test_conflicting_privacy_module_is_order_independent_and_preserves_unrelated_modules():
    first = {"module_id": "source.retention", "text": "Records are retained for seven years."}
    second = {"module_id": "source.retention", "text": "Records are destroyed immediately."}
    unrelated = {"module_id": "source.access", "text": "Authorized staff may review records."}
    outcomes = []
    for paragraphs in ([first, second, unrelated], [second, unrelated, first]):
        normalized_section, findings = contracts.normalize_privacy_modules({
            "paragraphs": paragraphs,
            "lists": [],
        })
        outcomes.append((normalized_section, findings))
    for normalized_section, findings in outcomes:
        assert [item["module_id"] for item in normalized_section["paragraphs"]] == ["source.access"]
        assert [item["code"] for item in findings] == [
            "conflicting_privacy_module:source.retention"
        ]
        assert findings[0]["module_id"] == "source.retention"
    assert outcomes[0][1] == outcomes[1][1]


def test_conflicting_privacy_module_blocks_before_rendering(tmp_path):
    draft = model()
    draft["icf"]["icf.privacy"] = {
        "paragraphs": [
            {"module_id": "source.retention", "text": "Records are retained for seven years."},
            {"module_id": "source.retention", "text": "Records are destroyed immediately."},
            {"module_id": "source.access", "text": "Authorized staff may review records."},
        ],
        "lists": [],
    }
    report = rendering.render_documents(
        ROOT, tmp_path, fixture(), draft, artifact_names={"icf"}
    )
    assert report["status"] == "blocked"
    finding = report["artifacts"][0]["findings"][0]
    assert finding["code"] == "conflicting_privacy_module:source.retention"


def test_validated_privacy_draft_preserves_stable_module_identity(tmp_path):
    reference = fixture()
    batch = next(
        item for item in contracts.batch_plan("Prospective", "Sterling")
        if item.batch_id == "icf-narrative"
    )
    request_path = create_drafting_request(
        repo_root=ROOT,
        revision_dir=tmp_path,
        revision_id="r-privacy-module",
        reference=reference,
        batch=batch,
        attempts={section_id: 1 for section_id in batch.section_ids},
        wave="focused",
    )
    request = json.loads(request_path.read_text(encoding="utf-8"))
    response = recorded_acceptance_response(request)
    privacy = next(
        item for item in response["section_results"]
        if item["section_id"] == "icf.privacy"
    )
    privacy["paragraphs"][0]["module_id"] = "source.data-retention"
    accepted, findings = validate_response(request, response)
    assert not [item for item in findings if item.get("field") == "icf.privacy"]
    accepted_privacy = next(
        item for item in (accepted or {}).get("drafts", [])
        if item["section_id"] == "icf.privacy"
    )
    assert accepted_privacy["paragraphs"][0]["module_id"] == "source.data-retention"


def test_validated_privacy_draft_rejects_conflicting_same_id_payloads(tmp_path):
    reference = fixture()
    batch = next(
        item for item in contracts.batch_plan("Prospective", "Sterling")
        if item.batch_id == "icf-narrative"
    )
    request_path = create_drafting_request(
        repo_root=ROOT,
        revision_dir=tmp_path,
        revision_id="r-privacy-conflict",
        reference=reference,
        batch=batch,
        attempts={section_id: 1 for section_id in batch.section_ids},
        wave="focused",
    )
    request = json.loads(request_path.read_text(encoding="utf-8"))
    response = recorded_acceptance_response(request)
    privacy = next(
        item for item in response["section_results"]
        if item["section_id"] == "icf.privacy"
    )
    privacy["paragraphs"] = [
        {
            **privacy["paragraphs"][0],
            "module_id": "source.data-retention",
            "text": "Approved records are retained for seven years.",
        },
        {
            **privacy["paragraphs"][0],
            "module_id": "source.data-retention",
            "text": "Approved records are destroyed immediately.",
        },
    ]
    accepted, findings = validate_response(request, response)
    assert any(
        item.get("code") == "conflicting_privacy_module:source.data-retention"
        for item in findings
    )
    assert not any(
        item["section_id"] == "icf.privacy"
        for item in (accepted or {}).get("drafts", [])
    )


def test_non_sterling_family_isolated_from_sterling_module_changes(tmp_path):
    reference = fixture(family="Advarra")
    report = rendering.render_documents(ROOT, tmp_path, reference, model(), artifact_names={"icf"})
    assert report["status"] == "passed", report
    document = Document(tmp_path / "candidate/icf.docx")
    text = visible(document)
    assert "RELEASE OF MEDICAL RECORDS AND PRIVACY" in text
    assert "STUDY RESULTS" not in text
    assert "What happens to information about me after the study is over" not in text


def test_release_findings_do_not_change_prs_generation_inputs():
    reference = fixture()
    original = copy.deepcopy(reference)
    contracts.release_source_findings(reference)
    assert reference == original
